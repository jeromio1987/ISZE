"""Corrected age-floor scenarios: the floor must extend past age 19, otherwise vehicles
aged 20+ are still forced to zero and the floor does nothing where the ministry's
1999-2010 block actually sits (ages 16-27 at 30.06.2026).
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

BASE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Israel_Rebuild_20260922")
OUT = BASE / "rebuild_outputs"
OUT.mkdir(parents=True, exist_ok=True)

REG_SRC = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imie_mail_20260730\UIO ISUZU 2025.xlsx"
)
wb = openpyxl.load_workbook(REG_SRC, data_only=True)
ws = wb["גיליון1"]
BANDS = ["6.5 TON", "7.5 TON", "10TON", "11 TON", "12 TON"]
flow: dict[str, dict[int, int]] = {b: {} for b in BANDS}
for r in range(2, 34):
    y = int(ws.cell(row=r, column=1).value)
    for j, b in enumerate(BANDS, start=2):
        v = ws.cell(row=r, column=j).value
        flow[b][y] = int(v) if v else 0

SHEET_EMERGING = [1, .99, .98, .97, .96, .95, .9, .85, .8, .75, .7, .6, .5, .4, .3, .2, .15, .1, .05, 0]
BASE_CURVE = [1.0] + SHEET_EMERGING[:-1]   # index = ref - reg, 0 from index 20

MAX_AGE = 40


def floored(curve: list[float], floor: float, min_age: int, max_age: int = MAX_AGE) -> list[float]:
    """Extend the curve to max_age and hold at least `floor` from min_age upward."""
    c = list(curve) + [0.0] * max(0, max_age + 1 - len(curve))
    for age in range(min_age, max_age + 1):
        c[age] = max(c[age], floor)
    return c


def conv(ref: int, bands: list[str], curve: list[float]) -> float:
    tot = 0.0
    for b in bands:
        for rgy, units in flow[b].items():
            if not units:
                continue
            age = ref - rgy
            if 0 <= age < len(curve):
                tot += units * curve[age]
    return tot


OLD = {b: {y: flow[b][y] for y in range(1999, 2011)} for b in BANDS}
old_regs_total = sum(sum(OLD[b].values()) for b in BANDS)
MINISTRY_OLD = 1680
CUM_FLOW = sum(sum(flow[b].values()) for b in BANDS)
GOV_PARC = 14019


def conv_old(ref: int, curve: list[float]) -> float:
    tot = 0.0
    for b in BANDS:
        for rgy, units in OLD[b].items():
            age = ref - rgy
            if 0 <= age < len(curve):
                tot += units * curve[age]
    return tot


print(f"1999-2010 registrations (all bands): {old_regs_total:,}")
print(f"ministry: still on the road at 30.06.2026: {MINISTRY_OLD:,}  "
      f"({MINISTRY_OLD / GOV_PARC:.2%} of parc, implied survival {MINISTRY_OLD / old_regs_total:.1%})")
print(f"cumulative recorded flow 1994-2025, all bands: {CUM_FLOW:,}")
print(f"ministry parc at 30.06.2026: {GOV_PARC:,}  -> hard unreachable floor {GOV_PARC - CUM_FLOW:+,}")
print()

print("no-floor baseline:")
print(f"  fleet {conv(2026, BANDS, BASE_CURVE):,.0f}   old block {conv_old(2026, BASE_CURVE):,.0f}")
print()

print("=" * 104)
print("AGE-FLOOR SCENARIOS (floor extended to age 40, so 20+ year old trucks are not killed)")
print("=" * 104)
print(f"{'floor':>7}{'from age':>10}{'fleet 2026':>13}{'old block':>12}{'old survival':>14}"
      f"{'vs ministry':>14}{'still unreachable':>19}")
scen = []
for min_age, floor in [(15, 0.5), (15, 0.7), (15, 0.8), (15, 0.9), (15, 1.0),
                       (18, 0.9), (19, 0.9), (19, 1.0), (25, 0.8)]:
    c = floored(BASE_CURVE, floor, min_age)
    f26 = conv(2026, BANDS, c)
    o26 = conv_old(2026, c)
    scen.append({"floor": floor, "from_age": min_age, "fleet_2026": round(f26),
                 "old_block_2026": round(o26),
                 "old_survival": round(o26 / old_regs_total, 4),
                 "unreachable": round(GOV_PARC - f26)})
    print(f"{floor:>7.2f}{min_age:>10}{f26:>13,.0f}{o26:>12,.0f}"
          f"{o26 / old_regs_total:>14.1%}{o26 - MINISTRY_OLD:>+14,.0f}{GOV_PARC - f26:>19,.0f}")

print()
print("=" * 104)
print("BRITAIN V15 EQUIVALENT — 90 % owner-operator floor on cohorts older than 18 (age 19+)")
print("=" * 104)
uk = floored(BASE_CURVE, 0.9, 19)
uk_fleet = conv(2026, BANDS, uk)
uk_old = conv_old(2026, uk)
print(f"  fleet at 2026 with the UK-style floor : {uk_fleet:,.0f}")
print(f"  fleet at 2026 with no floor (today)   : {conv(2026, BANDS, BASE_CURVE):,.0f}")
print(f"  trucks the UK-style fix would add     : {uk_fleet - conv(2026, BANDS, BASE_CURVE):+,.0f}")
print(f"  as a share of the 3,259 gap           : "
      f"{(uk_fleet - conv(2026, BANDS, BASE_CURVE)) / (GOV_PARC - 10760):.1%}")

print()
print("=" * 104)
print("CALIBRATED CURVE — the floor strength that reproduces the ministry's own old-tail figure")
print("=" * 104)
for min_age in (15, 16, 17, 18, 19):
    lo, hi = 0.0, 1.0
    for _ in range(40):
        mid = (lo + hi) / 2
        if conv_old(2026, floored(BASE_CURVE, mid, min_age)) < MINISTRY_OLD:
            lo = mid
        else:
            hi = mid
    c = floored(BASE_CURVE, hi, min_age)
    print(f"  floor from age {min_age:>2}: {hi:.3f}  -> old block {conv_old(2026, c):,.0f} "
          f"(ministry 1,680), fleet {conv(2026, BANDS, c):,.0f}, "
          f"unreachable {GOV_PARC - conv(2026, BANDS, c):+,.0f}")

with (OUT / "floor_scenarios.csv").open("w", encoding="utf-8", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["floor", "from_age", "fleet_2026", "old_block_1999_2010_2026",
                "old_survival", "vs_ministry_1680", "unreachable_vs_14019"])
    for s in scen:
        w.writerow([s["floor"], s["from_age"], s["fleet_2026"], s["old_block_2026"],
                    s["old_survival"], s["old_block_2026"] - MINISTRY_OLD, s["unreachable"]])

(BASE / "floor_scenarios.json").write_text(json.dumps({
    "old_registrations_1999_2010": old_regs_total,
    "ministry_old_block": MINISTRY_OLD,
    "cumulative_flow_all_bands": CUM_FLOW,
    "gov_parc": GOV_PARC,
    "hard_unreachable_floor": GOV_PARC - CUM_FLOW,
    "no_floor": {"fleet_2026": round(conv(2026, BANDS, BASE_CURVE)),
                 "old_block_2026": round(conv_old(2026, BASE_CURVE))},
    "uk_style_floor_age19_90pct": {"fleet_2026": round(uk_fleet), "old_block_2026": round(uk_old),
                                   "trucks_added": round(uk_fleet - conv(2026, BANDS, BASE_CURVE))},
    "scenarios": scen,
}, indent=2), encoding="utf-8")
print("\nwrote:", OUT / "floor_scenarios.csv")
