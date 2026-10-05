"""Decompose the Israel gap at 30.06.2026 and test survival-curve / floor scenarios."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

BASE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Israel_Rebuild_20260922")
SRC = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imie_mail_20260730\UIO ISUZU 2025.xlsx"
)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["גיליון1"]
BANDS = ["6.5 TON", "7.5 TON", "10TON", "11 TON", "12 TON"]
reg: dict[str, dict[int, int]] = {b: {} for b in BANDS}
for r in range(2, 34):
    y = int(ws.cell(row=r, column=1).value)
    for j, b in enumerate(BANDS, start=2):
        v = ws.cell(row=r, column=j).value
        reg[b][y] = int(v) if v else 0
YEARS = sorted(reg["6.5 TON"])

# IML Emerging curve, workbook '3. UIO Calculation' Section A rows age 1..20
SHEET_EMERGING = [1, .99, .98, .97, .96, .95, .9, .85, .8, .75, .7, .6, .5, .4, .3, .2, .15, .1, .05, 0]
CURVE = [1.0] + SHEET_EMERGING[:-1]   # index = ref - reg ; index19 = .05, index>=20 = 0


def conv(ref: float, regs: dict[int, int], curve=CURVE, floor=None, floor_min_age=15) -> float:
    tot = 0.0
    for rgy, units in regs.items():
        if not units:
            continue
        age = ref - rgy
        if age < 0:
            continue
        if age >= len(curve):
            s = 0.0
        else:
            s = curve[int(age)]
        if floor is not None and age >= floor_min_age:
            s = max(s, floor)
        tot += units * s
    return tot


# ---- government snapshot, slide 14 (UTI July 2026 deck, "based on MOT data") ----
GOV = {
    "6.5-7.5 t": 8610,
    "7.6-10 t": 2485,
    "10.1-12 t": 2697,
    "12.1-16 t": 227,
}
GOV_TOTAL = sum(GOV.values())
print("government class sum:", GOV_TOTAL, "(spec says 14,019)")
print("government old tail 1999-2010 still on the road: 1,680  (12% of 14,019)")

LD_REG = {y: reg["6.5 TON"][y] + reg["7.5 TON"][y] for y in YEARS}
REG10 = dict(reg["10TON"])
REG11 = dict(reg["11 TON"])
REG12 = dict(reg["12 TON"])
MD_REG = {y: REG10[y] + REG11[y] + REG12[y] for y in YEARS}
ALL_REG = {y: LD_REG[y] + MD_REG[y] for y in YEARS}

print("\n=== cumulative registrations in the UTI/MOT flow table ===")
print("6.5 t :", sum(reg['6.5 TON'].values()))
print("7.5 t :", sum(reg['7.5 TON'].values()))
print("10 t  :", sum(REG10.values()))
print("11 t  :", sum(REG11.values()))
print("12 t  :", sum(REG12.values()))
print("6.5+7.5 (LD)  :", sum(LD_REG.values()))
print("10+11+12 (MD) :", sum(MD_REG.values()))
print("TOTAL 1994-2025:", sum(ALL_REG.values()))
print("GOVERNMENT PARC 30.06.2026:", GOV_TOTAL)
print("parc / cumulative flow =", round(GOV_TOTAL / sum(ALL_REG.values()), 4))

print("\n=== our convolution at ref years (IML Emerging, floor=none) ===")
hdr = f"{'year':>6}{'6.5+7.5':>10}{'10t':>8}{'11t':>8}{'12t':>8}{'MD 10+11+12':>13}{'total':>9}"
print(hdr)
conv_by_year = {}
for ref in (2019, 2020, 2021, 2022, 2023, 2024, 2025, 2026):
    row = {
        "LD_6.5+7.5": conv(ref, LD_REG),
        "MD_10t": conv(ref, REG10),
        "MD_11t": conv(ref, REG11),
        "MD_12t": conv(ref, REG12),
    }
    row["MD_10+11+12"] = row["MD_10t"] + row["MD_11t"] + row["MD_12t"]
    row["total"] = row["LD_6.5+7.5"] + row["MD_10+11+12"]
    conv_by_year[ref] = row
    print(f"{ref:>6}{row['LD_6.5+7.5']:>10.0f}{row['MD_10t']:>8.0f}{row['MD_11t']:>8.0f}"
          f"{row['MD_12t']:>8.0f}{row['MD_10+11+12']:>13.0f}{row['total']:>9.0f}")

c26 = conv_by_year[2026]
print("\n=== class-by-class gap at 30.06.2026 ===")
X = c26["MD_10t"]
Y = c26["MD_11t"] + c26["MD_12t"]
lines = [
    ("6.5-7.5 t  (our LD row)", GOV["6.5-7.5 t"], c26["LD_6.5+7.5"]),
    ("7.6-10 t   (our 10 t col)", GOV["7.6-10 t"], X),
    ("10.1-12 t  (our 11+12 t cols)", GOV["10.1-12 t"], Y),
    ("12.1-16 t  (no home in our model)", GOV["12.1-16 t"], 0.0),
]
gaps = {}
for label, gov, ours in lines:
    gap = gov - ours
    gaps[label] = gap
    print(f"  {label:<34} gov={gov:>7,}  ours={ours:>8.0f}  gap={gap:>+8.0f}  ({gap / gov * 100:>6.1f}% of gov class)")
print(f"  {'TOTAL':<34} gov={GOV_TOTAL:>7,}  ours={c26['total']:>8.0f}  gap={GOV_TOTAL - c26['total']:>+8.0f}")

print("\n=== old tail: cohorts first registered 1999-2010, retention at 30.06.2026 ===")
old_reg = {y: ALL_REG[y] for y in range(1999, 2011)}
print("  registrations 1999-2010 (all bands):", sum(old_reg.values()))
print("  our curve retains at 2026:", round(conv(2026, old_reg)))
print("  government says still on the road:", 1680)
print("  shortfall our curve creates:", 1680 - round(conv(2026, old_reg)))

print("\n=== hard ceiling test: could ANY curve on this flow reach the government parc? ===")
for label, gov_class, regs in [
    ("6.5-7.5 t", GOV["6.5-7.5 t"], LD_REG),
    ("7.6-10 t", GOV["7.6-10 t"], REG10),
    ("10.1-12 t", GOV["10.1-12 t"], {y: REG11[y] + REG12[y] for y in YEARS}),
    ("12.1-16 t", GOV["12.1-16 t"], {}),
]:
    cum = sum(regs.values())
    ratio = (gov_class / cum) if cum else None
    print(f"  {label:<12} cumulative flow={cum:>7,}  gov parc={gov_class:>7,}  "
          f"parc/flow={ratio if ratio is None else round(ratio, 4)}")

print("\n=== floor scenarios on the whole fleet (target 14,019 @ 30.06.2026) ===")
for floor_min_age in (15, 16, 17, 18, 19):
    for floor in (0.10, 0.20, 0.30, 0.50, 0.80, 1.00):
        t = conv(2026, ALL_REG, floor=floor, floor_min_age=floor_min_age)
        ld = conv(2026, LD_REG, floor=floor, floor_min_age=floor_min_age)
        print(f"  min_age={floor_min_age} floor={floor:>4}: total={t:>8.0f}  LD={ld:>8.0f}")

( BASE / "gap_decomposition.json").write_text(json.dumps({
    "gov_snapshot_2026_06_30": GOV,
    "our_conv_2026": c26,
    "our_conv_by_year": conv_by_year,
    "cumulative_flow": {b: sum(reg[b].values()) for b in BANDS},
    "cumulative_flow_LD": sum(LD_REG.values()),
    "cumulative_flow_MD": sum(MD_REG.values()),
    "old_tail_1999_2010_registrations": sum(old_reg.values()),
    "old_tail_our_curve_2026": round(conv(2026, old_reg)),
    "old_tail_government": 1680,
    "gaps": gaps,
}, indent=2), encoding="utf-8")
print("\nwrote:", BASE / "gap_decomposition.json")
