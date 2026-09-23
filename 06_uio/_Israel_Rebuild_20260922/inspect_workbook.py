"""READ-ONLY inspection of the live UIO workbook. Never saves."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

SRC = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"
)
OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Israel_Rebuild_20260922")

print("opening (read-only):", SRC, SRC.stat().st_size, "bytes")
wb = openpyxl.load_workbook(SRC, data_only=False, read_only=True)
print("sheets:")
for i, s in enumerate(wb.sheetnames):
    print(f"  {i:2d}. {s!r}")

# Sheet dimensions
for s in wb.sheetnames:
    ws = wb[s]
    print(f"dims {s!r}: {ws.max_row} x {ws.max_column}")

wb.close()

wb2 = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
s = "1.UIO Overview"
ws = wb2[s]
print("\n" + "=" * 110)
print(f"SHEET {s!r} — all rows col A..K")
print("=" * 110)
for r in range(1, min(ws.max_row, 160) + 1):
    vals = []
    for c in range(1, 12):
        v = ws.cell(row=r, column=c).value
        vals.append("" if v is None else str(v))
    line = " | ".join(vals).rstrip(" |")
    print(f"r{r:>3}: {line}")
wb2.close()
