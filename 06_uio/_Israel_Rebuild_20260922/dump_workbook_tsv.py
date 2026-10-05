"""READ-ONLY dump of the live UIO workbook sheets to TSV so the numbers can be read back cheaply."""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

SRC = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx")
OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Israel_Rebuild_20260922\_wb")
OUT.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
print("sheets:", wb.sheetnames)

SHEETS = ["1.UIO Overview", "2. Vehicle Volume", "3. UIO Calculation",
          "Assumptions & Mapping", "Data Sources", "Changelog", "Cover"]

for name in SHEETS:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    f = OUT / (name.replace(" ", "_").replace(".", "-") + ".tsv")
    with f.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        for i, r in enumerate(rows, 1):
            w.writerow(["" if v is None else v for v in r])
    print(f"{name!r}: {len(rows)} rows -> {f}")
wb.close()
