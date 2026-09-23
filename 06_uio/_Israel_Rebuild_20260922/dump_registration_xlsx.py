"""Dump the Israeli government registration band table (UTI extract) verbatim.

Source named in V20_SPEC.md Problem 9:
  01_data/90_INBOX/_imie_mail_20260730/UIO ISUZU 2025.xlsx
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

SRC = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imie_mail_20260730\UIO ISUZU 2025.xlsx"
)
OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Israel_Rebuild_20260922")

print("file:", SRC)
print("exists:", SRC.exists())
st = SRC.stat()
print("mtime:", datetime.fromtimestamp(st.st_mtime).isoformat(), "bytes:", st.st_size)

wb = openpyxl.load_workbook(SRC, data_only=True)
print("sheets:", wb.sheetnames)

payload = {
    "source_file": str(SRC),
    "source_mtime": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(),
    "read_utc": datetime.now(timezone.utc).isoformat(),
    "sheets": {},
}

for ws in wb.worksheets:
    rows = []
    for r in ws.iter_rows():
        row = []
        for c in r:
            v = c.value
            if isinstance(v, datetime):
                v = v.date().isoformat()
            row.append(v)
        rows.append(row)
    payload["sheets"][ws.title] = {
        "dims": ws.dimensions,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "rows": rows,
    }

    print("\n" + "=" * 100)
    print(f"SHEET {ws.title!r}  dims={ws.dimensions}  rows={ws.max_row} cols={ws.max_column}")
    print("=" * 100)
    for i, row in enumerate(rows, start=1):
        cells = " | ".join("" if v is None else str(v) for v in row)
        print(f"r{i:>3}: {cells}")

(OUT / "registration_band_table_raw.json").write_text(
    json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
)
print("\nwrote:", OUT / "registration_band_table_raw.json")
