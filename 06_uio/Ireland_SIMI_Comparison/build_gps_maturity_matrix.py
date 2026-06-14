"""
ISZE GPS Refinement Maturity Matrix — external format.

Five axes per distributor, each 0-3 (Axis B and E max at 2):
  A — UIO denominator quality
  B — Sales numerator depth
  C — PMS / maintenance data access
  D — Pricing / competitive aftermarket data
  E — Warranty / goodwill cleanliness

Axis A pre-filled from current ISZE methodology (IML / SMMT / SIMI etc.).
Other axes left blank for the GPS owner to fill — no fabricated values.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\UIO\Ireland_SIMI_Comparison\ISZE_GPS_DataMaturity_Matrix_2026-05-28.xlsx")

# Distributor list — match v17 + sales TSV
DISTRIBUTORS = [
    # (code, country, scope, axis_A_prefilled_level, axis_A_source_note, axis_B_prefilled)
    ("ITUK",     "United Kingdom",   "CV (N-series + F-Series)",  2, "SMMT/DVLA — full per-model registration cohorts 2008-2024",  1),  # user-confirmed sales-out
    ("IUK",      "United Kingdom",   "LCV (D-Max)",                1, "IMIT LCV study + IML records",                              None),
    ("UTI",      "Israel",            "CV trucks (LD + MD)",       1, "IML calculation — Developing Country survival curve",       None),
    ("UMI",      "Israel",            "LCV (D-Max)",                1, "IML calculation — LCV 12-yr survival",                       None),
    ("IRELAND",  "Ireland",           "CV + LCV",                  2, "SIMI (national reg authority) added 2026; IML pre-2010",     None),
    ("BELGIUM",  "Belgium / Benelux", "CV + LCV",                  1, "IML scrap-adjusted (Belgium = scrap-adj market) + IMIT LCV", None),
    ("GERMANY",  "Germany",           "CV + LCV",                  1, "IML scrap-adjusted + IMIT LCV (ISD Germany)",                None),
    ("PORTUGAL", "Portugal + Azores", "CV (LD + F-Series) + LCV",  1, "IML scrap-adjusted + IMIT LCV (Imotors)",                    None),
    ("SWEDEN",   "Sweden",            "CV + LCV",                  1, "IML + IMIT LCV (Isuzu Sverige)",                             None),
    ("NORWAY",   "Norway",            "CV + LCV",                  1, "IML + IMIT LCV (RSA)",                                       None),
    ("DENMARK",  "Denmark",           "CV + LCV",                  1, "IML scrap-adjusted + IMIT LCV (Isuzu Denmark)",              None),
    ("FINLAND",  "Finland",           "CV + LCV",                  1, "IML + IMIT LCV (Nordic Automotive)",                         None),
    ("ICELAND",  "Iceland",           "CV + LCV",                  1, "IML + IMIT LCV (BL)",                                        None),
    ("GREECE",   "Greece",            "CV + LCV",                  1, "IML scrap-adjusted + IMIT LCV (Petros Petropoulos)",         None),
    ("CYPRUS",   "Cyprus",            "CV + LCV",                  1, "IML scrap-adjusted + IMIT LCV (Loutsios & Sons)",            None),
    ("MALTA",    "Malta",             "CV + LCV",                  1, "IML scrap-adjusted + IMIT LCV (GasanZammit)",                None),
    ("ROMANIA",  "Romania",           "CV + LCV",                  1, "IML + IMIT LCV (Isuzu Rom Auto)",                            None),
    ("HUNGARY",  "Hungary",           "CV + LCV",                  1, "IML + IMIT LCV (AutoWallis)",                                None),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)

BLUE = PatternFill("solid", fgColor="1F4E78")
NAVY = PatternFill("solid", fgColor="2E75B6")
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
GREEN = PatternFill("solid", fgColor="E2EFDA")
RED0 = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREEN3 = PatternFill("solid", fgColor="C6EFCE")
HEADLINE = PatternFill("solid", fgColor="2E75B6")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
ITALIC_GREY = Font(italic=True, color="595959")
thin = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

LEVEL_FILL = {0: RED0, 1: AMBER, 2: GREEN, 3: GREEN3}

def style_header(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = BLUE; cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ============================================================
# Sheet 1: Framework
# ============================================================
ws = wb.create_sheet("Framework")
ws["A1"] = "GPS Refinement Maturity Matrix — Framework"
ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
ws.merge_cells("A1:F1")
ws["A2"] = "How accurate can the GPS calculation become for each distributor? Five axes scored 0-3 (B and E max 2)."
ws["A2"].font = Font(size=12, italic=True, color="595959")
ws.merge_cells("A2:F2")

intro = [
    ("", ""),
    ("Date prepared", "2026-05-28"),
    ("Purpose", "Map each ISZE distributor against five data-quality axes that determine how refined the GPS calculation can become. Identify quick wins (one-step upgrades) and structural limits."),
    ("", ""),
    ("How to read the matrix", ""),
    ("  Current level", "The realistic score today, based on data the GPS team can actually use."),
    ("  Target level", "The achievable ceiling given known data sources for that distributor."),
    ("  Gap", "Target − Current. Where >0, there is room to refine."),
    ("  Action", "What concretely needs to happen to close the gap."),
    ("  Owner", "Who controls the data needed for the next step."),
    ("", ""),
    ("Five axes (definitions on next page)", ""),
    ("  A — UIO denominator quality", "Where the units-in-operation count comes from. Higher = closer to national truth."),
    ("  B — Sales numerator depth", "Whose sales we see in the calculation. Higher = closer to actual parts consumption."),
    ("  C — PMS / maintenance data", "Service-interval and workshop visibility. Drives PMS capture rate refinement."),
    ("  D — Pricing / competitive data", "Aftermarket pricing visibility per market. Drives the EUR/unit wallet assumption."),
    ("  E — Warranty / goodwill data", "Cleanliness of the sales numerator (warranty parts excluded properly)."),
]
for i, (k, v) in enumerate(intro, 4):
    ws.cell(row=i, column=1, value=k).font = BOLD
    ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
autosize(ws, [32, 60, 18, 18, 18, 18])

# ============================================================
# Sheet 2: Axis definitions
# ============================================================
ws = wb.create_sheet("Axis_Definitions")
ws["A1"] = "Axis Definitions — Scoring Scale"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:E1")

r = 3
axes = [
    ("A — UIO Denominator Quality", "Where the units-in-operation number comes from.", [
        (0, "IML shipment records, lumped CV LD (no model split)", "Distributors with no scrap-adj logic and only IML's accumulated UIO model"),
        (1, "IML shipment records, per-model split (N3.5T / N7.5T / F / MD / LCV)", "Most ISZE distributors today after applying segment mapping"),
        (2, "National registration authority — SMMT, SIMI, KBA, Bilstatistik, etc. — true regs per model", "ITUK (SMMT), Ireland (SIMI from 2026)"),
        (3, "Per-VIN fleet register with age cohorts", "Workshop or distributor-system level; rare in ISZE network today"),
    ]),
    ("B — Sales Numerator Depth", "Whose parts sales we see in the calculation.", [
        (0, "ISZE → distributor only (wholesale-in, what is in our SAP export today)", "Default for all distributors today"),
        (1, "+ Distributor → dealer (distributor sell-out from their warehouse)", "ITUK confirmed (Niguri sales-out file). Others TBC."),
        (2, "+ Per-VIN parts attribution (which truck VIN consumed which part)", "Requires workshop integration — typically distributor-side"),
        (None, "(Dealer → end-customer / retail sell-through level is deliberately excluded — unreachable.)", ""),
    ]),
    ("C — PMS / Maintenance Data Access", "Visibility of service-interval and workshop behaviour.", [
        (0, "None — no PMS data shared", "Default"),
        (1, "One-off PMS data exchange (e.g., ITUK PMS study Apr 2026)", "ITUK has shared once"),
        (2, "Recurring PMS data feed", "Monthly/quarterly extracts from distributor"),
        (3, "Full workshop system integration (PMS captured per VIN automatically)", "Aspiration — DMS / workshop integration"),
    ]),
    ("D — Pricing / Competitive Data", "Aftermarket pricing visibility for that market.", [
        (0, "None — no aftermarket pricing visibility", "Default"),
        (1, "Anecdotal — known competitor parts cited by distributor", ""),
        (2, "TecDoc / Eucon coverage — structured competitive PNs and pricing", ""),
        (3, "Full competitive PR00 mapping with our genuine pricing benchmarked", "ITUK BVP work, ISZJ master data competitive parts"),
    ]),
    ("E — Warranty / Goodwill Cleanliness", "Whether warranty parts are properly excluded from sales numerator.", [
        (0, "Warranty / goodwill blended into total sales (numerator inflated)", ""),
        (1, "Warranty flagged in SAP but not consistently excluded from GPS calc", "Current state — gps_engine has exclude_warranty flag but inconsistent application"),
        (2, "Warranty excluded with full audit trail per distributor", "Target — clean numerator"),
    ]),
]

for axis_name, axis_desc, levels in axes:
    ws.cell(row=r, column=1, value=axis_name).font = Font(bold=True, size=12, color="1F4E78")
    ws.cell(row=r, column=2, value=axis_desc).font = ITALIC_GREY
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r += 1
    ws.cell(row=r, column=1, value="Level").font = BOLD
    ws.cell(row=r, column=2, value="What it means").font = BOLD
    ws.cell(row=r, column=3, value="Example / Note").font = BOLD
    for c in range(1, 6):
        ws.cell(row=r, column=c).fill = GREY
        ws.cell(row=r, column=c).border = BORDER
    r += 1
    for lvl, desc, ex in levels:
        if lvl is None:
            ws.cell(row=r, column=1, value="—").alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=desc).font = ITALIC_GREY
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        else:
            c1 = ws.cell(row=r, column=1, value=lvl)
            c1.font = BOLD
            c1.fill = LEVEL_FILL[lvl]
            c1.alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=desc).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=3, value=ex).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1

autosize(ws, [8, 42, 32, 20, 20])

# ============================================================
# Sheet 3: Matrix (the main grid)
# ============================================================
ws = wb.create_sheet("Maturity_Matrix")
ws["A1"] = "Maturity Matrix — All Distributors"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:O1")
ws["A2"] = "Score 0-3 per axis. Coloured by level: red=0 | amber=1 | green=2 | dark-green=3. Cells marked TBD need owner input."
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:O2")

# Column headers
hdr = [
    "Distributor", "Country", "Scope",
    "A — UIO\n(current)", "A target", "A gap",
    "B — Sales\n(current)", "B target", "B gap",
    "C — PMS\n(current)", "C target", "C gap",
    "D — Pricing\n(current)", "D target", "D gap",
    "E — Warranty\n(current)", "E target", "E gap",
    "Total current", "Total target", "Total gap",
]
for c, h in enumerate(hdr, 1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr))
ws.row_dimensions[4].height = 36

r = 5
for code, country, scope, a_cur, a_note, b_cur in DISTRIBUTORS:
    ws.cell(row=r, column=1, value=code).font = BOLD
    ws.cell(row=r, column=2, value=country)
    ws.cell(row=r, column=3, value=scope).font = ITALIC_GREY
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")

    # Axis A — pre-filled
    c = ws.cell(row=r, column=4, value=a_cur)
    c.fill = LEVEL_FILL[a_cur]; c.alignment = Alignment(horizontal="center"); c.font = BOLD
    ws.cell(row=r, column=5, value="").alignment = Alignment(horizontal="center")  # target TBD
    ws.cell(row=r, column=6, value="").alignment = Alignment(horizontal="center")  # gap TBD

    # Axis B — only ITUK confirmed
    if b_cur is not None:
        c = ws.cell(row=r, column=7, value=b_cur)
        c.fill = LEVEL_FILL[b_cur]; c.alignment = Alignment(horizontal="center"); c.font = BOLD
    else:
        c = ws.cell(row=r, column=7, value="TBD")
        c.fill = GREY; c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=8, value="").alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=9, value="").alignment = Alignment(horizontal="center")

    # Axes C, D, E — all TBD
    for col_offset in [10, 13, 16]:
        for off in range(3):
            cell = ws.cell(row=r, column=col_offset + off, value="TBD" if off == 0 else "")
            cell.alignment = Alignment(horizontal="center")
            if off == 0:
                cell.fill = GREY

    # Totals — leave blank until owner fills
    for col in [19, 20, 21]:
        ws.cell(row=r, column=col, value="").alignment = Alignment(horizontal="center")

    ws.row_dimensions[r].height = 24
    r += 1

# Notes row
r += 1
ws.cell(row=r, column=1, value="Notes on pre-filled values:").font = BOLD
r += 1
ws.cell(row=r, column=1, value="Axis A pre-filled from current ISZE methodology (v17 Data Sources / IML / SMMT / SIMI). All other cells = TBD until owner confirms.")
ws.cell(row=r, column=1).font = ITALIC_GREY
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=21)
r += 1
ws.cell(row=r, column=1, value="ITUK Axis B = 1 because ITUK shares sales-to-dealer data (Niguri file). All other distributors' B values must be filled in.")
ws.cell(row=r, column=1).font = ITALIC_GREY
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=21)

autosize(ws, [12, 18, 24,
              9, 9, 7,
              9, 9, 7,
              9, 9, 7,
              9, 9, 7,
              9, 9, 7,
              10, 10, 9])

# ============================================================
# Sheet 4: Refinement Checklist per Distributor
# ============================================================
ws = wb.create_sheet("Refinement_Checklist")
ws["A1"] = "Refinement Checklist — Per Distributor"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:F1")
ws["A2"] = "For each distributor: what concrete step would refine the GPS calculation, and who owns the data."
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:F2")

hdr_row = ["Distributor", "Axis", "Current → Target", "Action needed", "Data owner", "Status"]
for c, h in enumerate(hdr_row, 1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr_row))

r = 5
for code, country, scope, a_cur, a_note, b_cur in DISTRIBUTORS:
    # Group header per distributor
    ws.cell(row=r, column=1, value=f"{code} — {country}").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=1).fill = NAVY
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1

    # 5 rows per distributor (one per axis), pre-filled where confident
    axes_template = [
        ("A — UIO denominator", f"{a_cur} → ?", a_note, "ISZE Data team / IML / national reg authority", ""),
        ("B — Sales numerator", (f"{b_cur} → ?" if b_cur is not None else "TBD → ?"),
         ("ITUK shares Niguri sales-out file — confirm cadence and integrate into GPS calc."
          if code == "ITUK" else "Ask distributor whether they can share distributor-to-dealer sell-out."),
         "Distributor", ""),
        ("C — PMS data", "TBD → ?", "Ask distributor whether PMS / workshop visit data is shareable.",
         "Distributor", ""),
        ("D — Pricing data", "TBD → ?", "Identify whether TecDoc / Eucon coverage exists for this market; if so, integrate.",
         "ISZE Pricing / TecDoc / Eucon", ""),
        ("E — Warranty data", "TBD → ?", "Confirm warranty parts are flagged in SAP for this sold_to and exclude consistently.",
         "ISZE Finance / SAP team", ""),
    ]

    for axis, transition, action, owner, status in axes_template:
        ws.cell(row=r, column=1, value="").alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=axis).font = BOLD
        c3 = ws.cell(row=r, column=3, value=transition)
        c3.alignment = Alignment(horizontal="center")
        if "TBD" in transition:
            c3.fill = GREY
        ws.cell(row=r, column=4, value=action).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=5, value=owner).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=6, value=status).alignment = Alignment(horizontal="center")
        for cc in range(1, 7):
            ws.cell(row=r, column=cc).border = BORDER
        ws.row_dimensions[r].height = 32
        r += 1
    r += 1

autosize(ws, [16, 22, 14, 50, 28, 14])

# ============================================================
# Sheet 5: Scoreboard (auto-summary)
# ============================================================
ws = wb.create_sheet("Scoreboard")
ws["A1"] = "Scoreboard — Distributors by Current Level (Axis A only — pre-filled)"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:E1")
ws["A2"] = "Once you've filled in axes B-E, this sheet should be replaced with a live count across all five axes."
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:E2")

# Count by Axis A level
from collections import Counter
a_levels = Counter([row[3] for row in DISTRIBUTORS])

hdr = ["Axis A level", "Description", "Count", "Distributors"]
for c, h in enumerate(hdr, 1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr))

descriptions = {
    0: "IML lumped (no model split)",
    1: "IML per-model split (current standard)",
    2: "National registration authority (SMMT / SIMI etc.)",
    3: "Per-VIN cohort register",
}
r = 5
for lvl in [3, 2, 1, 0]:
    count = a_levels.get(lvl, 0)
    dists = ", ".join([d[0] for d in DISTRIBUTORS if d[3] == lvl]) or "—"
    c1 = ws.cell(row=r, column=1, value=lvl); c1.fill = LEVEL_FILL[lvl]
    c1.alignment = Alignment(horizontal="center"); c1.font = BOLD
    ws.cell(row=r, column=2, value=descriptions[lvl])
    ws.cell(row=r, column=3, value=count).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value=dists).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

r += 2
ws.cell(row=r, column=1, value="Quick-win observation").font = BOLD
r += 1
ws.cell(row=r, column=1, value="Two distributors already at Axis A = 2 (ITUK via SMMT, Ireland via SIMI). The biggest GPS refinement available across the network is lifting other markets from Axis A = 1 to 2 — needs each market's national registration authority. Distributors with scrap-adjusted markets (BE, CY, DE, DK, GR, MT, PT) are the closest candidates because the data is already part-way there.")
ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws.row_dimensions[r].height = 60

autosize(ws, [10, 42, 8, 60])

wb.save(OUT)
print(f"WROTE: {OUT}")
print(f"Sheets: {[s.title for s in wb.worksheets]}")
print(f"Distributors in matrix: {len(DISTRIBUTORS)}")
print(f"Pre-filled cells: Axis A = all {len(DISTRIBUTORS)}; Axis B = 1 (ITUK only)")
print(f"To fill: B for {len(DISTRIBUTORS)-1} distributors; C, D, E for all 18")
