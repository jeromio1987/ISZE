"""
In-place update of ISZE_GPS_Calculator.xlsx:
- Preserves all 6 existing sheets (Instructions, Assumptions, Shipment Data,
  UIO Calculation, GPS Model, Dashboard)
- Appends 6 new sheets at the end:
    GPS_Results_2026    — per-model methodology results 2022-2025
    GPS_Ranking_2026    — sorted by 2025 GPS%
    Maturity_Matrix     — 5-axis data model per distributor
    Axis_Definitions    — scoring scale 0-3
    Refinement_Checklist — per-distributor action list
    Maturity_Scoreboard — totals & gaps

Writes the updated file to all three identical copy locations.
"""
import openpyxl
import pandas as pd
import shutil
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter

# Source master + mirrors
MASTER = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Tools\isze-tools-github\isze-tools-github\tools\ISZE_GPS_Calculator.xlsx")
MIRRORS = [
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\Isuzu\Divers\ISZE_GPS_Calculator.xlsx"),
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\Master Dashboard\Divers\ISZE_GPS_Calculator.xlsx"),
]

V17_PATH = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Data\ISUZU_UIO_Overview_2026_v17.xlsx")
SALES_PATH = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Projects\Isuzu\prism\data\sales_detailed_for_claude_full_fixed.tsv")

T_N75 = 1120
T_FMD = 1344
T_LCV = 250
NON_PARTS = {"TRANSPORT", "FREIGHT", "DUTY", "PACKING"}
SALES_YEARS = [2022, 2023, 2024, 2025]

V17_DIST_TO_CODE = {
    "UTI": "UTI", "UMI": "UMI", "BL": "ICELAND", "RSA": "NORWAY",
    "Isuzu Sverige": "SWEDEN", "Isuzu Denmark": "DENMARK",
    "ITUK": "ITUK", "IUK": "IUK", "Isuzu Ireland": "IRELAND",
    "IBX": "BELGIUM", "ISD-Lang": "GERMANY", "Imotors": "PORTUGAL",
    "GasanZammit": "MALTA", "Nordic Auto.": "FINLAND", "AutoWallis": "HUNGARY",
    "Petros Petropoulos": "GREECE", "Isuzu Rom Auto": "ROMANIA", "Loutsios & Sons": "CYPRUS",
}
SOLD_TO_TO_CODE = {
    "ISUZU TRUCK (UK) LIMITED": "ITUK", "ISUZU TRUCK UK LIMITED": "ITUK",
    "I.M. PARTS & SERVICE Limited": "IUK",
    "UNIVERSAL MOTORS ISRAEL LTD": "UMI", "UNIVERSAL TRUCKS ISRAEL LTD": "UTI",
    "ISUZU IRELAND": "IRELAND", "ISUZU BENELUX NV": "BELGIUM",
    "ISUZU SALES DEUTSCHLAND GMBH": "GERMANY",
    "XXX ISUZU SALES DEUTSCHLAND GMBH": "GERMANY",
    "Isuzu Motors Germany GmbH": "EXCLUDE",
    "IMOTORS - Importacao e Comercio": "PORTUGAL",
    "ISUZU SVERIGE AB": "SWEDEN", "ISUZU MOTORS SWEDEN AB": "SWEDEN",
    "ISUZU DANMARK A/S": "DENMARK", "ISUZU ROM AUTO SRL": "ROMANIA",
    "RUTEBILEIERNES STANDARDISERING AS": "NORWAY",
    "PETROS PETROPOULOS S.A. (VEHICLE)": "GREECE",
    "PETROS PETROPOULOS S.A.": "GREECE",
    "LOUTSIOS A & SONS LTD": "CYPRUS",
    "Nordic Automotive trading oy.": "FINLAND",
    "GASANZAMMIT MOTORS LTD": "MALTA",
    "AW Distribution Kft.": "HUNGARY", "BL. EHF.": "ICELAND",
}
SEGMENT_ONLY = {"ITUK": "CV", "IUK": "LCV", "UTI": "CV", "UMI": "LCV"}
PROXY_YEAR = 2024

# ---------- Load v17 UIO ----------
wb17 = openpyxl.load_workbook(V17_PATH, read_only=True, data_only=True)
ws17 = wb17["1.UIO Overview"]
hdr17 = list(ws17.iter_rows(min_row=6, max_row=6, values_only=True))[0]
y2c = {}
for i, h in enumerate(hdr17):
    try:
        y = int(str(h).replace("E", "").replace("F", "").strip())
        y2c[y] = i
    except Exception:
        pass

uio = {}
for row in ws17.iter_rows(min_row=7, max_row=500, values_only=True):
    if not row[5]: continue
    v17_dist = str(row[5]).strip()
    if "MIDI" in v17_dist.upper(): continue
    code = V17_DIST_TO_CODE.get(v17_dist)
    if not code: continue
    cv = (str(row[2]) if row[2] else "").strip()
    mdl = (str(row[3]) if row[3] else "").strip()
    if cv == "CV" and mdl == "LD": seg = "LD"
    elif mdl == "F-Series": seg = "F"
    elif cv == "CV" and mdl == "MD": seg = "MD"
    elif cv == "LCV": seg = "LCV"
    else: continue
    for y in SALES_YEARS:
        if y not in y2c: continue
        v = row[y2c[y]] or 0
        if v:
            uio[(code, y, seg)] = uio.get((code, y, seg), 0) + v

def _raw(d, y, s):
    v = uio.get((d, y, s), 0)
    if v == 0:
        v = uio.get((d, PROXY_YEAR, s), 0)
    return v

def get_uio(d, y, s):
    if d in SEGMENT_ONLY:
        if SEGMENT_ONLY[d] == "CV" and s in ("LD", "F", "MD"):
            return _raw(d, y, s)
        if SEGMENT_ONLY[d] == "LCV" and s == "LCV":
            return _raw(d, y, s)
        return 0
    return _raw(d, y, s)

def wallet(d, y):
    return (get_uio(d, y, "LD") * T_N75 +
            get_uio(d, y, "F") * T_FMD +
            get_uio(d, y, "MD") * T_FMD +
            get_uio(d, y, "LCV") * T_LCV)

# ---------- Sales (parts only) ----------
sales = pd.read_csv(SALES_PATH, sep="\t", low_memory=False)
sales["net_sales_value"] = pd.to_numeric(sales["net_sales_value"], errors="coerce")
sales["year"] = pd.to_numeric(sales["year"], errors="coerce")
sales["part_upper"] = sales["part_number"].astype(str).str.upper().str.strip()
sales_parts = sales[~sales["part_upper"].isin(NON_PARTS)].copy()
sales_parts["dist_code"] = sales_parts["sold_to"].map(SOLD_TO_TO_CODE)
sales_by_dy = (sales_parts.dropna(subset=["dist_code"])
               .groupby(["dist_code", "year"])["net_sales_value"].sum().to_dict())

all_dists = sorted(set(d for (d,_,_) in uio.keys()) - {"EXCLUDE"})

rows = []
for d in all_dists:
    for y in SALES_YEARS:
        ld = get_uio(d, y, "LD"); f = get_uio(d, y, "F")
        md = get_uio(d, y, "MD"); lcv = get_uio(d, y, "LCV")
        total = ld + f + md + lcv
        wal = wallet(d, y)
        sv = sales_by_dy.get((d, y), 0.0)
        eu = sv/total if total else None
        gps = sv/wal if wal else None
        rows.append((d, y, ld, f, md, lcv, total, wal, sv, eu, gps))
dfm = pd.DataFrame(rows, columns=["dist","year","ld","f","md","lcv","total","wallet","sales","eu","gps"])

# ---------- Maturity matrix data ----------
MATURITY = [
    ("ITUK","United Kingdom","CV (N-series + F-Series)",       2,"CONFIRMED",1,"CONFIRMED",1,"CONFIRMED",2,"CONFIRMED",1,"DEFAULT", 2,2,2,3,2),
    ("IUK","United Kingdom","LCV (D-Max)",                     1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  1,"GUESS",    1,"DEFAULT", 2,1,1,2,2),
    ("UTI","Israel","CV trucks (LD + MD)",                     1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("UMI","Israel","LCV (D-Max)",                             1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("IRELAND","Ireland","CV + LCV",                           2,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("BELGIUM","Belgium (Benelux)","CV + LCV",                 1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("GERMANY","Germany","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("PORTUGAL","Portugal + Azores","CV (LD + F-Series) + LCV",1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("SWEDEN","Sweden","CV + LCV",                             1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("NORWAY","Norway","CV + LCV",                             1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("DENMARK","Denmark","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("FINLAND","Finland","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("ICELAND","Iceland","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("GREECE","Greece","CV + LCV",                             1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("CYPRUS","Cyprus","CV + LCV",                             1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("MALTA","Malta","CV + LCV",                               1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("ROMANIA","Romania","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
    ("HUNGARY","Hungary","CV + LCV",                           1,"CONFIRMED",0,"DEFAULT",  0,"DEFAULT",  0,"DEFAULT",  1,"DEFAULT", 2,1,1,2,2),
]

# ============================================================
# Open existing master, preserve sheets, append new ones
# ============================================================
wb = openpyxl.load_workbook(MASTER)
print(f"Existing sheets preserved: {wb.sheetnames}")

# Styles
BLUE = PatternFill("solid", fgColor="1F4E78")
NAVY = PatternFill("solid", fgColor="2E75B6")
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
RED0 = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREEN = PatternFill("solid", fgColor="E2EFDA")
GREEN3 = PatternFill("solid", fgColor="C6EFCE")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
ITALIC_GREY = Font(italic=True, color="595959")
thin = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LEVEL_FILL = {0: RED0, 1: AMBER, 2: GREEN, 3: GREEN3}
CONF_FONT = {"CONFIRMED": Font(bold=True, color="2E7D32"),
             "GUESS": Font(italic=True, color="E65100"),
             "DEFAULT": Font(italic=True, color="595959")}

def style_header(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = BLUE; cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def gps_fill(p):
    if p is None: return None
    if p >= 1.0: return GREEN3
    if p >= 0.33: return AMBER
    return RED0

# Remove old versions of our new sheets if they exist
for sn in ["GPS_Results_2026","GPS_Ranking_2026","Maturity_Matrix",
           "Axis_Definitions","Refinement_Checklist","Maturity_Scoreboard"]:
    if sn in wb.sheetnames:
        del wb[sn]

# ---------- GPS_Results_2026 ----------
ws = wb.create_sheet("GPS_Results_2026")
ws["A1"] = "GPS Results 2022-2025 — Per-Model Wallet Methodology"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:K1")
ws["A2"] = "Max wallet per vehicle: N 7.5T = EUR 1,120 | F-Series / CV MD = EUR 1,344 | LCV / D-Max / N 3.5T = EUR 250"
ws["A2"].font = ITALIC_GREY
ws["A3"] = "Sales: ISZE ERP net-sales export, parts-only (TRANSPORT / FREIGHT / DUTY / PACKING excluded). UIO source: current ISZE methodology (v17)."
ws["A3"].font = ITALIC_GREY
ws.merge_cells("A2:K2"); ws.merge_cells("A3:K3")

hdr = ["Distributor","Year","UIO CV LD","UIO F","UIO MD","UIO LCV","Total UIO","Wallet (€)","Sales parts (€)","€/unit","GPS %"]
for c,h in enumerate(hdr,1):
    ws.cell(row=5, column=c, value=h)
style_header(ws, 5, len(hdr))

r = 6
for _, row in dfm.iterrows():
    ws.cell(row=r, column=1, value=row["dist"]).font = BOLD
    ws.cell(row=r, column=2, value=int(row["year"])).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3, value=int(row["ld"]))
    ws.cell(row=r, column=4, value=int(row["f"]))
    ws.cell(row=r, column=5, value=int(row["md"]))
    ws.cell(row=r, column=6, value=int(row["lcv"]))
    c = ws.cell(row=r, column=7, value=int(row["total"])); c.font = BOLD
    ws.cell(row=r, column=8, value=row["wallet"]).number_format = "#,##0"
    ws.cell(row=r, column=9, value=row["sales"]).number_format = "#,##0"
    if row["eu"] is not None:
        c = ws.cell(row=r, column=10, value=row["eu"]); c.number_format = "€#,##0"
    if row["gps"] is not None:
        c = ws.cell(row=r, column=11, value=row["gps"])
        c.number_format = "0.0%"; c.fill = gps_fill(row["gps"]); c.font = BOLD
    for cc in range(3, 12):
        ws.cell(row=r, column=cc).alignment = Alignment(horizontal="right")
    r += 1
autosize(ws, [14,8,11,8,8,10,10,14,14,11,10])

# ---------- GPS_Ranking_2026 ----------
ws = wb.create_sheet("GPS_Ranking_2026")
ws["A1"] = "GPS Ranking — All Distributors (sorted by 2025 GPS%)"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:L1")
ws["A2"] = "Traffic lights: ≥100% green | 33-100% amber | <33% red"
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:L2")

hdr = ["Distributor","Scope","2022 €/unit","2022 GPS%","2023 €/unit","2023 GPS%","2024 €/unit","2024 GPS%","2025 €/unit","2025 GPS%","UIO 2024","2025 sales (€)"]
for c,h in enumerate(hdr,1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr))

rank = (dfm[dfm["year"] == 2025]
        .sort_values("gps", ascending=False, na_position="last").reset_index(drop=True))
r = 5
for _, rr in rank.iterrows():
    d = rr["dist"]
    ws.cell(row=r, column=1, value=d).font = BOLD
    ws.cell(row=r, column=2, value=SEGMENT_ONLY.get(d, "CV + LCV")).alignment = Alignment(horizontal="center")
    col = 3
    for y in SALES_YEARS:
        rec = dfm[(dfm["dist"]==d) & (dfm["year"]==y)].iloc[0]
        if rec["eu"] is not None:
            c1 = ws.cell(row=r, column=col, value=rec["eu"]); c1.number_format = "€#,##0"
            c1.alignment = Alignment(horizontal="right")
        col += 1
        if rec["gps"] is not None:
            c2 = ws.cell(row=r, column=col, value=rec["gps"])
            c2.number_format = "0.0%"; c2.fill = gps_fill(rec["gps"]); c2.font = BOLD
            c2.alignment = Alignment(horizontal="right")
        col += 1
    d24 = dfm[(dfm["dist"]==d) & (dfm["year"]==2024)].iloc[0]
    ws.cell(row=r, column=col, value=int(d24["total"])).alignment = Alignment(horizontal="right")
    col += 1
    ws.cell(row=r, column=col, value=rr["sales"]).number_format = "#,##0"
    ws.cell(row=r, column=col).alignment = Alignment(horizontal="right")
    r += 1
autosize(ws, [14,12,11,10,11,10,11,10,11,10,11,15])

# ---------- Maturity_Matrix ----------
ws = wb.create_sheet("Maturity_Matrix")
ws["A1"] = "Data Maturity Matrix — Per Distributor"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:T1")
ws["A2"] = "Five axes 0-3 (B and E max 2). Confidence: CONFIRMED (green) | GUESS (orange) | DEFAULT (grey)."
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:T2")

hdr = ["Distributor","Country","Scope",
       "A cur","A conf","A tgt", "B cur","B conf","B tgt",
       "C cur","C conf","C tgt", "D cur","D conf","D tgt",
       "E cur","E conf","E tgt", "Total cur","Total tgt"]
for c,h in enumerate(hdr,1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr))
ws.row_dimensions[4].height = 32

def put_level(r,c,lvl):
    cell = ws.cell(row=r, column=c, value=lvl)
    cell.fill = LEVEL_FILL[lvl]; cell.font = BOLD
    cell.alignment = Alignment(horizontal="center"); cell.border = BORDER

def put_conf(r,c,cf):
    cell = ws.cell(row=r, column=c, value=cf)
    cell.font = CONF_FONT[cf]; cell.alignment = Alignment(horizontal="center"); cell.border = BORDER

r = 5
for row in MATURITY:
    (code,country,scope,a_cur,a_conf,b_cur,b_conf,c_cur,c_conf,
     d_cur,d_conf,e_cur,e_conf,a_t,b_t,c_t,d_t,e_t) = row
    ws.cell(row=r, column=1, value=code).font = BOLD
    ws.cell(row=r, column=2, value=country)
    ws.cell(row=r, column=3, value=scope).font = ITALIC_GREY
    for cc in range(1,4): ws.cell(row=r, column=cc).border = BORDER
    put_level(r,4,a_cur); put_conf(r,5,a_conf); put_level(r,6,a_t)
    put_level(r,7,b_cur); put_conf(r,8,b_conf); put_level(r,9,b_t)
    put_level(r,10,c_cur); put_conf(r,11,c_conf); put_level(r,12,c_t)
    put_level(r,13,d_cur); put_conf(r,14,d_conf); put_level(r,15,d_t)
    put_level(r,16,e_cur); put_conf(r,17,e_conf); put_level(r,18,e_t)
    tot_cur = a_cur+b_cur+c_cur+d_cur+e_cur
    tot_t = a_t+b_t+c_t+d_t+e_t
    c1 = ws.cell(row=r, column=19, value=tot_cur); c1.fill = NAVY; c1.font = WHITE_BOLD
    c1.alignment = Alignment(horizontal="center"); c1.border = BORDER
    c2 = ws.cell(row=r, column=20, value=tot_t); c2.fill = NAVY; c2.font = WHITE_BOLD
    c2.alignment = Alignment(horizontal="center"); c2.border = BORDER
    r += 1
autosize(ws, [12,18,24,7,10,7,7,10,7,7,10,7,7,10,7,7,10,7,10,10])

# ---------- Axis_Definitions ----------
ws = wb.create_sheet("Axis_Definitions")
ws["A1"] = "Axis Definitions"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:E1")
r = 3
axes = [
    ("A — UIO Denominator","Where the units-in-operation number comes from.",[
        (0,"IML shipment records, lumped CV LD (no model split)","Distributors with no scrap-adj logic"),
        (1,"IML shipment records, per-model split","Current standard for most ISZE distributors"),
        (2,"National registration authority (SMMT / SIMI / KBA etc.)","ITUK, Ireland"),
        (3,"Per-VIN fleet register with age cohorts","Rare today"),
    ]),
    ("B — Sales Numerator","Whose parts sales we see.",[
        (0,"ISZE → distributor only (wholesale-in)","Default — what's in our SAP export"),
        (1,"+ Distributor → dealer (sell-out)","ITUK confirmed (Niguri)"),
        (2,"+ Per-VIN parts attribution","Requires workshop integration"),
    ]),
    ("C — PMS / Maintenance","Service-interval visibility.",[
        (0,"None","Default"),
        (1,"One-off PMS exchange","ITUK Apr 2026 study"),
        (2,"Recurring PMS feed","Quarterly extracts"),
        (3,"Full workshop integration","Aspiration"),
    ]),
    ("D — Pricing / Competitive","Aftermarket pricing visibility.",[
        (0,"None","Default"),
        (1,"Anecdotal","Competitor parts cited"),
        (2,"TecDoc / Eucon coverage","Structured competitive data"),
        (3,"Full competitive PR00 mapping","ITUK BVP work"),
    ]),
    ("E — Warranty / Goodwill","Warranty cleanly excluded.",[
        (0,"Blended into total sales","Numerator inflated"),
        (1,"Flagged in SAP, not consistently excluded","Current state"),
        (2,"Excluded with audit trail","Target"),
    ]),
]
for axis,desc,levels in axes:
    ws.cell(row=r, column=1, value=axis).font = Font(bold=True, size=12, color="1F4E78")
    ws.cell(row=r, column=2, value=desc).font = ITALIC_GREY
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r += 1
    for c,h in enumerate(["Level","Description","Example"], 1):
        ws.cell(row=r, column=c, value=h).font = BOLD
        ws.cell(row=r, column=c).fill = GREY
    r += 1
    for lvl,d,ex in levels:
        cell = ws.cell(row=r, column=1, value=lvl)
        cell.fill = LEVEL_FILL[lvl]; cell.font = BOLD
        cell.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=d)
        ws.cell(row=r, column=3, value=ex)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        r += 1
    r += 1
autosize(ws, [8,42,32,18,18])

# ---------- Refinement_Checklist ----------
ws = wb.create_sheet("Refinement_Checklist")
ws["A1"] = "Refinement Checklist — Action per Distributor"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:F1")
hdr = ["Distributor","Axis","Cur → Tgt","Action needed","Data owner","Status"]
for c,h in enumerate(hdr,1): ws.cell(row=4, column=c, value=h)
style_header(ws, 4, len(hdr))

axis_owners = {
    "A":"ISZE Data team / reg authority", "B":"Distributor",
    "C":"Distributor (workshop/DMS)", "D":"ISZE Pricing / TecDoc",
    "E":"ISZE Finance / SAP",
}
def action(code, axis, cur, tgt):
    if cur == tgt: return "Already at target."
    if axis == "A":
        if cur == 1 and tgt >= 2: return f"Identify and ingest national reg authority data for {code}'s market."
        if cur == 2 and tgt == 3: return "Move to per-VIN fleet register."
    if axis == "B":
        if cur == 0 and tgt >= 1: return f"Request distributor sell-out from {code}."
        if cur == 1 and tgt == 2: return "Workshop-level per-VIN parts attribution."
    if axis == "C":
        if cur == 0 and tgt >= 1: return f"Request one-off PMS study from {code}."
        if cur == 1 and tgt >= 2: return f"Convert one-off to recurring PMS feed for {code}."
    if axis == "D":
        if cur == 0 and tgt >= 1: return "Document competitor parts cited anecdotally."
        if cur == 1 and tgt >= 2: return f"Procure TecDoc / Eucon coverage for {code}."
        if cur == 2 and tgt == 3: return "Build competitive PR00 mapping."
    if axis == "E":
        if cur == 0 and tgt >= 1: return f"Confirm warranty flagged in SAP for {code}."
        if cur == 1 and tgt == 2: return f"Exclude warranty consistently from {code} GPS, document audit trail."
    return "—"

r = 5
for row in MATURITY:
    code = row[0]
    ws.cell(row=r, column=1, value=f"{code} — {row[1]}").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=1).fill = NAVY
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    axes_data = [
        ("A — UIO", row[3], row[13], axis_owners["A"]),
        ("B — Sales", row[5], row[14], axis_owners["B"]),
        ("C — PMS", row[7], row[15], axis_owners["C"]),
        ("D — Pricing", row[9], row[16], axis_owners["D"]),
        ("E — Warranty", row[11], row[17], axis_owners["E"]),
    ]
    for axis_label, cur, tgt, owner in axes_data:
        ws.cell(row=r, column=1, value="")
        ws.cell(row=r, column=2, value=axis_label).font = BOLD
        c3 = ws.cell(row=r, column=3, value=f"{cur} → {tgt}")
        c3.alignment = Alignment(horizontal="center"); c3.font = BOLD
        c3.fill = GREEN if cur == tgt else AMBER
        ws.cell(row=r, column=4, value=action(code, axis_label[0], cur, tgt)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=5, value=owner)
        ws.cell(row=r, column=6, value="")
        for cc in range(1, 7): ws.cell(row=r, column=cc).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
autosize(ws, [16,18,12,55,28,12])

# ---------- Maturity_Scoreboard ----------
ws = wb.create_sheet("Maturity_Scoreboard")
ws["A1"] = "Maturity Scoreboard"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:H1")

axis_cols = [("A — UIO",3,3,13),("B — Sales",5,2,14),("C — PMS",7,3,15),
             ("D — Pricing",9,3,16),("E — Warranty",11,2,17)]
hdr = ["Axis","Max","L0","L1","L2","L3","Cur total","Tgt total"]
for c,h in enumerate(hdr,1): ws.cell(row=3, column=c, value=h)
style_header(ws, 3, len(hdr))

r = 4
total_cur = 0; total_tgt = 0
for label, cur_idx, mx, tgt_idx in axis_cols:
    cur_vals = [m[cur_idx] for m in MATURITY]
    tgt_vals = [m[tgt_idx] for m in MATURITY]
    cnt = Counter(cur_vals)
    ws.cell(row=r, column=1, value=label).font = BOLD
    ws.cell(row=r, column=2, value=mx).alignment = Alignment(horizontal="center")
    for lvl in range(4):
        cell = ws.cell(row=r, column=3+lvl, value=cnt.get(lvl, 0) if lvl <= mx else "—")
        cell.alignment = Alignment(horizontal="center")
        if lvl <= mx: cell.fill = LEVEL_FILL[lvl]
        else: cell.fill = GREY
    cv = sum(cur_vals); tv = sum(tgt_vals)
    total_cur += cv; total_tgt += tv
    c1 = ws.cell(row=r, column=7, value=cv); c1.font = BOLD; c1.alignment = Alignment(horizontal="center")
    c2 = ws.cell(row=r, column=8, value=tv); c2.font = BOLD; c2.fill = GREEN; c2.alignment = Alignment(horizontal="center")
    r += 1

ws.cell(row=r, column=1, value="ALL AXES").font = BOLD
ws.cell(row=r, column=2, value=13).alignment = Alignment(horizontal="center")
for cc in range(3,7): ws.cell(row=r, column=cc, value="").fill = GREY
c1 = ws.cell(row=r, column=7, value=total_cur); c1.font = WHITE_BOLD; c1.fill = NAVY
c1.alignment = Alignment(horizontal="center")
c2 = ws.cell(row=r, column=8, value=total_tgt); c2.font = WHITE_BOLD; c2.fill = NAVY
c2.alignment = Alignment(horizontal="center")
r += 2

ws.cell(row=r, column=1, value="Distributors ranked by refinement gap").font = BOLD
r += 1
for c,h in enumerate(["Rank","Distributor","Current","Target","Gap"], 1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, r, 5)
r += 1
gaps = sorted([(m[0], m[3]+m[5]+m[7]+m[9]+m[11], m[13]+m[14]+m[15]+m[16]+m[17])
               for m in MATURITY], key=lambda x: -(x[2]-x[1]))
for i,(code,cur,tgt) in enumerate(gaps,1):
    gap = tgt - cur
    ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2, value=code).font = BOLD
    ws.cell(row=r, column=3, value=cur).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value=tgt).alignment = Alignment(horizontal="center")
    c = ws.cell(row=r, column=5, value=gap); c.alignment = Alignment(horizontal="center"); c.font = BOLD
    c.fill = RED0 if gap >= 6 else (AMBER if gap >= 4 else GREEN)
    r += 1
autosize(ws, [10,16,12,10,10,10,11,11])

# Save to master, then copy to mirrors
wb.save(MASTER)
print(f"\nSaved master: {MASTER}")
for m in MIRRORS:
    shutil.copy(MASTER, m)
    print(f"Mirrored to: {m}")

# Re-read to confirm
wb2 = openpyxl.load_workbook(MASTER)
print(f"\nFinal sheet list: {wb2.sheetnames}")
