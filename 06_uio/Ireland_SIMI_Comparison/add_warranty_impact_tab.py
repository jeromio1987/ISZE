"""
Add Warranty_Impact_2025 tab to ISZE_GPS_Calculator.xlsx + rename 'wallet'.

Warranty calc:
  - Raw CV warranty (Raw_CV_Warranty) and LCV warranty (Raw_LCV_Warranty)
    from Final_Complete_Warranty_Pricing_2026.xlsx, filtered to 2025
  - Each warranty line × PR00 price from Pricing_Data sheet (163k P/N price master)
  - Aggregated per DB
  - Shown alongside current GPS for context (NOT replacing GPS_Results)

Terminology change applied across file: 'wallet' / 'Wallet' → 'max spend per vehicle/year'.
"""
import openpyxl
import pandas as pd
import shutil
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MASTER = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Tools\isze-tools-github\isze-tools-github\tools\ISZE_GPS_Calculator.xlsx")
PUBLISHED = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Master Dashboard\Divers\ISZE_GPS_Calculator.xlsx")
WARR = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Data\Final_Complete_Warranty_Pricing_2026.xlsx")

# Mapping from warranty file distributor labels to canonical codes
CV_DST_MAP = {
    "ITUK": "ITUK",
    "UTI": "UTI",
    "IRELAND": "IRELAND",
    "ISD": "GERMANY",
    "PORTUGAL": "PORTUGAL",
    "CYPRUS": "CYPRUS",
    "GREECE": "GREECE",
    "MALTA": "MALTA",
    "MIDI": "EXCLUDE",
    "MIDI AIOS Products": "EXCLUDE",
}

# LCV distributor name -> code (full set from data discovery)
LCV_DIST_MAP = {
    # Core ISZE entities
    "BL ehf.": "ICELAND",
    "ISUZU BENELUX - BELGIUM": "BELGIUM",
    "ISUZU BENELUX - NETHERLANDS": "BELGIUM",          # rolled into BE
    "Isuzu Benelux NV": "BELGIUM",
    "Isuzu Benelux NV (Poland)": "BELGIUM",            # rolled into BE
    "ISUZU (UK)": "IUK",
    "I.M. Parts & Service Limited": "IUK",
    "Universal Motors": "UMI",
    "Universal Motors Israel Ltd.": "UMI",
    "Isuzu Ireland": "IRELAND",
    "Isuzu Ireland 2020 Unlimited Company": "IRELAND",
    "ISUZU SALES DEUTSCHLAND - ISD": "GERMANY",
    "ISUZU SALES DEUTSCHLAND - AUSTRIA": "GERMANY",    # ISD-Lang covers AT
    "ISUZU SALES DEUTSCHLAND - CZECH": "GERMANY",      # ISD-Lang covers CZ
    "ISUZU SALES DEUTSCHLAND - SLOVAKIA": "GERMANY",   # ISD-Lang covers SK
    "ISD-Lang": "GERMANY",
    "Imotors": "PORTUGAL",
    "IMOTORS LDA": "PORTUGAL",
    "Isuzu Sverige": "SWEDEN",
    "Isuzu Sverige AB": "SWEDEN",
    "RSA": "NORWAY",
    "Isuzu Denmark": "DENMARK",
    "ISUZU Danmark": "DENMARK",
    "Nordic Automotive": "FINLAND",
    "Isuzu Motor Finland OY": "FINLAND",
    "GasanZammit": "MALTA",
    "GASANZAMMIT MOTORS LIMITED": "MALTA",
    "Petros Petropoulos": "GREECE",
    "PETROS PETROPOULOS": "GREECE",
    "Loutsios & Sons": "CYPRUS",
    "Loutsios A & Sons Ltd.": "CYPRUS",
    "A. LOUTSIOS & SONS": "CYPRUS",
    "Isuzu Rom Auto": "ROMANIA",
    "ISUZU ROM AUTO SRL": "ROMANIA",
    "AutoWallis": "HUNGARY",
    "AW Distribution Kft.": "HUNGARY",
    # Exclusions (outside ISZE GPS scope)
    "MIDI Europe": "EXCLUDE",
    "MIDI Europe S.R.L.": "EXCLUDE",
    "MIDI EUROPE - FRANCE": "EXCLUDE",
    "MIDI EUROPE - ITALY": "EXCLUDE",
    "MIDI EUROPE - MACEDONIA": "EXCLUDE",
    "MIDI EUROPE - SPAIN": "EXCLUDE",
    "MIDI EUROPE - SWITZ": "EXCLUDE",
    "MIDI EUROPE S.R.L.(SLOVENIA)": "EXCLUDE",
    "Midi Europe S.P.L (Ser.mon.)": "EXCLUDE",
    "Midi Europe S.R.L (BOSNIA HERZEGOVINA)": "EXCLUDE",
    "WAE LTD.": "EXCLUDE",
    "ZAMZAM SPRING GENERAL TRADING CO., LTD": "EXCLUDE",
}

DB_INFO = {
    "ITUK":     {"db": "ITUK"},
    "IUK":      {"db": "IUK"},
    "UTI":      {"db": "UTI"},
    "UMI":      {"db": "UMI"},
    "IRELAND":  {"db": "Isuzu Ireland"},
    "BELGIUM":  {"db": "IBX"},
    "GERMANY":  {"db": "ISD-Lang"},
    "PORTUGAL": {"db": "Imotors"},
    "SWEDEN":   {"db": "Isuzu Sverige"},
    "NORWAY":   {"db": "RSA"},
    "DENMARK":  {"db": "Isuzu Denmark"},
    "FINLAND":  {"db": "Nordic Auto"},
    "ICELAND":  {"db": "BL"},
    "GREECE":   {"db": "Petros Petropoulos"},
    "CYPRUS":   {"db": "Loutsios & Sons"},
    "MALTA":    {"db": "GasanZammit"},
    "ROMANIA":  {"db": "Isuzu Rom Auto"},
    "HUNGARY":  {"db": "AutoWallis"},
}

def norm_pn(v):
    """Normalise part number to 10-char zero-padded string."""
    if v is None: return None
    s = str(v).strip()
    if s.endswith(".0"): s = s[:-2]
    if not s or s == "0": return None
    if s.isdigit():
        return s.zfill(10)
    return s

# ---------- Load PR00 price master ----------
print("Loading PR00 price master...")
wb_w = openpyxl.load_workbook(WARR, read_only=True, data_only=True)
ws_p = wb_w["Pricing_Data"]
prices = {}
header = list(ws_p.iter_rows(min_row=1, max_row=1, values_only=True))[0]
pn_idx = header.index("P/N")
price_idx = header.index("Price")
cur_idx = header.index("Currency")
for row in ws_p.iter_rows(min_row=2, max_row=ws_p.max_row, values_only=True):
    pn = norm_pn(row[pn_idx])
    if pn and row[price_idx] is not None:
        try:
            prices[pn] = float(row[price_idx])
        except (ValueError, TypeError):
            pass
print(f"  {len(prices):,} part numbers with PR00 prices loaded")

# ---------- CV warranty 2025 ----------
print("\nProcessing CV warranty (2025)...")
ws_cv = wb_w["Raw_CV_Warranty"]
hdr_cv = list(ws_cv.iter_rows(min_row=1, max_row=1, values_only=True))[0]
dst_idx = hdr_cv.index("DST_Name")
bill_idx = hdr_cv.index("BILLYMD")
pn1_idx = hdr_cv.index("PARTNBR")
q1_idx = hdr_cv.index("QTY1")
pn2_idx = hdr_cv.index("PARTNBR2")
q2_idx = hdr_cv.index("QTY2")

cv_agg = {}
cv_qty = {}
cv_lines = 0
cv_missing_price = 0
cv_missing_dist = set()
for row in ws_cv.iter_rows(min_row=2, max_row=ws_cv.max_row, values_only=True):
    bill = str(row[bill_idx]) if row[bill_idx] else ""
    if not bill.startswith("2025"): continue
    dst = str(row[dst_idx]).strip() if row[dst_idx] else ""
    code = CV_DST_MAP.get(dst)
    if code is None:
        cv_missing_dist.add(dst); continue
    if code == "EXCLUDE": continue
    for pn_i, q_i in [(pn1_idx, q1_idx), (pn2_idx, q2_idx)]:
        pn = norm_pn(row[pn_i])
        if not pn: continue
        try:
            qty = float(row[q_i] or 0)
        except (ValueError, TypeError):
            qty = 0
        if qty <= 0: continue
        price = prices.get(pn)
        if price is None:
            cv_missing_price += 1
            continue
        cv_agg[code] = cv_agg.get(code, 0) + qty * price
        cv_qty[code] = cv_qty.get(code, 0) + qty
        cv_lines += 1
print(f"  {cv_lines:,} CV warranty lines (2025) priced")
print(f"  Missing prices: {cv_missing_price}")
print(f"  Unmapped DST_Names: {sorted(cv_missing_dist)}")

# ---------- LCV warranty 2025 ----------
print("\nProcessing LCV warranty (2025)...")
ws_lcv = wb_w["Raw_LCV_Warranty"]
hdr_lcv = list(ws_lcv.iter_rows(min_row=1, max_row=1, values_only=True))[0]
am_idx = hdr_lcv.index("AssessMonth")
pn_idx = hdr_lcv.index("PartsNo")
pieces_idx = hdr_lcv.index("Pieces")
distname_idx = hdr_lcv.index("DistributorName")
distcode_idx = hdr_lcv.index("DistributorCode")

lcv_agg = {}
lcv_qty = {}
lcv_lines = 0
lcv_missing_price = 0
lcv_missing_dist = set()
for row in ws_lcv.iter_rows(min_row=2, max_row=ws_lcv.max_row, values_only=True):
    am = str(row[am_idx]) if row[am_idx] else ""
    if not am.startswith("2025"): continue
    dname = str(row[distname_idx]).strip() if row[distname_idx] else ""
    code = LCV_DIST_MAP.get(dname)
    if code is None:
        lcv_missing_dist.add(dname); continue
    if code == "EXCLUDE": continue
    pn = norm_pn(row[pn_idx])
    if not pn: continue
    try:
        qty = float(row[pieces_idx] or 0)
    except (ValueError, TypeError):
        qty = 0
    if qty <= 0: continue
    price = prices.get(pn)
    if price is None:
        lcv_missing_price += 1
        continue
    lcv_agg[code] = lcv_agg.get(code, 0) + qty * price
    lcv_qty[code] = lcv_qty.get(code, 0) + qty
    lcv_lines += 1
print(f"  {lcv_lines:,} LCV warranty lines (2025) priced")
print(f"  Missing prices: {lcv_missing_price}")
print(f"  Unmapped LCV DistributorNames: {sorted(lcv_missing_dist)}")

# ---------- Combined per DB ----------
all_codes = set(cv_agg.keys()) | set(lcv_agg.keys())
per_db = {}
for code in all_codes:
    cv_v = cv_agg.get(code, 0)
    cv_q = cv_qty.get(code, 0)
    lcv_v = lcv_agg.get(code, 0)
    lcv_q = lcv_qty.get(code, 0)
    per_db[code] = {
        "cv_qty": cv_q, "cv_val": cv_v,
        "lcv_qty": lcv_q, "lcv_val": lcv_v,
        "total_qty": cv_q + lcv_q,
        "total_val": cv_v + lcv_v,
    }

# Print summary
print(f"\n=== Warranty value 2025 at PR00 per DB ===")
print(f"{'DB code':<12}{'CV qty':>10}{'CV val €':>14}{'LCV qty':>10}{'LCV val €':>14}{'Total qty':>11}{'Total val €':>14}")
for code in sorted(per_db.keys()):
    d = per_db[code]
    print(f"{code:<12}{d['cv_qty']:>10.0f}{d['cv_val']:>14,.0f}{d['lcv_qty']:>10.0f}{d['lcv_val']:>14,.0f}{d['total_qty']:>11.0f}{d['total_val']:>14,.0f}")

# ---------- Load existing GPS results for the comparison ----------
SALES_PATH = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Projects\Isuzu\prism\data\sales_detailed_for_claude_full_fixed.tsv")
T_N75, T_FMD, T_LCV = 1120, 1344, 250
NON_PARTS = {"TRANSPORT","FREIGHT","DUTY","PACKING"}
V17_PATH = Path(r"C:\Users\jerom\Desktop\claude\ISZE\Data\ISUZU_UIO_Overview_2026_v17.xlsx")
SEGMENT_ONLY = {"ITUK":"CV","IUK":"LCV","UTI":"CV","UMI":"LCV"}
PROXY_YEAR = 2024

V17_DIST_TO_CODE = {"UTI":"UTI","UMI":"UMI","BL":"ICELAND","RSA":"NORWAY",
    "Isuzu Sverige":"SWEDEN","Isuzu Denmark":"DENMARK","ITUK":"ITUK","IUK":"IUK",
    "Isuzu Ireland":"IRELAND","IBX":"BELGIUM","ISD-Lang":"GERMANY","Imotors":"PORTUGAL",
    "GasanZammit":"MALTA","Nordic Auto.":"FINLAND","AutoWallis":"HUNGARY",
    "Petros Petropoulos":"GREECE","Isuzu Rom Auto":"ROMANIA","Loutsios & Sons":"CYPRUS"}
SOLD_TO_TO_CODE = {"ISUZU TRUCK (UK) LIMITED":"ITUK","ISUZU TRUCK UK LIMITED":"ITUK",
    "I.M. PARTS & SERVICE Limited":"IUK","UNIVERSAL MOTORS ISRAEL LTD":"UMI",
    "UNIVERSAL TRUCKS ISRAEL LTD":"UTI","ISUZU IRELAND":"IRELAND","ISUZU BENELUX NV":"BELGIUM",
    "ISUZU SALES DEUTSCHLAND GMBH":"GERMANY","XXX ISUZU SALES DEUTSCHLAND GMBH":"GERMANY",
    "Isuzu Motors Germany GmbH":"EXCLUDE",
    "IMOTORS - Importacao e Comercio":"PORTUGAL","ISUZU SVERIGE AB":"SWEDEN",
    "ISUZU MOTORS SWEDEN AB":"SWEDEN","ISUZU DANMARK A/S":"DENMARK",
    "ISUZU ROM AUTO SRL":"ROMANIA","RUTEBILEIERNES STANDARDISERING AS":"NORWAY",
    "PETROS PETROPOULOS S.A. (VEHICLE)":"GREECE","PETROS PETROPOULOS S.A.":"GREECE",
    "LOUTSIOS A & SONS LTD":"CYPRUS","Nordic Automotive trading oy.":"FINLAND",
    "GASANZAMMIT MOTORS LTD":"MALTA","AW Distribution Kft.":"HUNGARY","BL. EHF.":"ICELAND"}

print("\nLoading v17 UIO + sales for 2025 GPS comparison...")
wb17 = openpyxl.load_workbook(V17_PATH, read_only=True, data_only=True)
ws17 = wb17["1.UIO Overview"]
hdr17 = list(ws17.iter_rows(min_row=6, max_row=6, values_only=True))[0]
y2c = {}
for i,h in enumerate(hdr17):
    try: y2c[int(str(h).replace("E","").replace("F","").strip())] = i
    except: pass

uio = {}
for row in ws17.iter_rows(min_row=7, max_row=500, values_only=True):
    if not row[5]: continue
    v17_dist = str(row[5]).strip()
    if "MIDI" in v17_dist.upper(): continue
    code = V17_DIST_TO_CODE.get(v17_dist)
    if not code: continue
    cv = (str(row[2]) if row[2] else "").strip()
    mdl = (str(row[3]) if row[3] else "").strip()
    if cv == "CV" and mdl == "LD": seg = "LD"
    elif mdl == "F-Series": seg = "F"
    elif cv == "CV" and mdl == "MD": seg = "MD"
    elif cv == "LCV": seg = "LCV"
    else: continue
    for y in [2024, 2025]:
        if y not in y2c: continue
        v = row[y2c[y]] or 0
        if v: uio[(code,y,seg)] = uio.get((code,y,seg),0) + v

def _raw(d,y,s):
    v = uio.get((d,y,s),0)
    if v == 0: v = uio.get((d,PROXY_YEAR,s),0)
    return v

def get_uio(d,s):
    if d in SEGMENT_ONLY:
        if SEGMENT_ONLY[d]=="CV" and s in ("LD","F","MD"): return _raw(d,2025,s)
        if SEGMENT_ONLY[d]=="LCV" and s=="LCV": return _raw(d,2025,s)
        return 0
    return _raw(d,2025,s)

def max_spend(d):
    return (get_uio(d,"LD")*T_N75 + get_uio(d,"F")*T_FMD
          + get_uio(d,"MD")*T_FMD + get_uio(d,"LCV")*T_LCV)

# Sales 2025
sales = pd.read_csv(SALES_PATH, sep="\t", low_memory=False)
sales["net_sales_value"] = pd.to_numeric(sales["net_sales_value"], errors="coerce")
sales["year"] = pd.to_numeric(sales["year"], errors="coerce")
sales["part_upper"] = sales["part_number"].astype(str).str.upper().str.strip()
sales_parts = sales[~sales["part_upper"].isin(NON_PARTS)].copy()
sales_parts["dist_code"] = sales_parts["sold_to"].map(SOLD_TO_TO_CODE)
sales_2025 = (sales_parts[sales_parts["year"]==2025].dropna(subset=["dist_code"])
              .groupby("dist_code")["net_sales_value"].sum().to_dict())

# ============================================================
# Open existing GPS Calculator → add tab + rename "wallet"
# ============================================================
wb = openpyxl.load_workbook(MASTER)
print(f"\nLoaded calculator. Existing tabs: {wb.sheetnames}")

# ---- Rename "wallet" across all sheets ----
def replace_text(s):
    if s is None: return None
    if not isinstance(s, str): return s
    out = s
    out = out.replace("maximum wallet", "max spend per vehicle/year")
    out = out.replace("Maximum wallet", "Max spend per vehicle/year")
    out = out.replace("MAX WALLET", "MAX SPEND PER VEHICLE/YEAR")
    out = out.replace("max wallet", "max spend per vehicle/year")
    out = out.replace("Max wallet", "Max spend per vehicle/year")
    out = out.replace("MAXIMUM WALLET", "MAX SPEND PER VEHICLE/YEAR")
    out = out.replace("Wallet (€)", "Max spend (€)")
    out = out.replace("Wallet (EUR)", "Max spend (EUR)")
    out = out.replace("TOTAL wallet", "TOTAL max spend")
    out = out.replace("Total wallet", "Total max spend")
    out = out.replace(" wallet ", " max spend per vehicle ")
    out = out.replace(" wallet.", " max spend per vehicle/year.")
    out = out.replace("(wallet)", "(max spend per vehicle/year)")
    out = out.replace("wallet:", "max spend per vehicle/year:")
    out = out.replace("Wallet ", "Max spend ")
    out = out.replace(" wallet", " max spend per vehicle/year")
    return out

replacements = 0
for sn in wb.sheetnames:
    ws = wb[sn]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "allet" in cell.value:
                new = replace_text(cell.value)
                if new != cell.value:
                    cell.value = new
                    replacements += 1
print(f"Renamed {replacements} cell(s) containing 'wallet'")

# ---- Add Warranty_Impact_2025 tab ----
if "Warranty_Impact_2025" in wb.sheetnames:
    del wb["Warranty_Impact_2025"]
ws = wb.create_sheet("Warranty_Impact_2025")

BLUE = PatternFill("solid", fgColor="1F4E78")
NAVY = PatternFill("solid", fgColor="2E75B6")
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED0 = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFEB9C")
GREEN3 = PatternFill("solid", fgColor="C6EFCE")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
ITALIC_GREY = Font(italic=True, color="595959")
thin = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
def gps_fill(p):
    if p is None: return None
    if p >= 1.0: return GREEN3
    if p >= 0.33: return AMBER
    return RED0
def style_header(row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = BLUE; cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

ws["A1"] = "Warranty Impact 2025 — value at PR00 pricing per DB"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:J1")
ws["A2"] = "Raw warranty claims 2025 × PR00 price (Pricing_Data sheet, 163k P/N price master). Independent of sales TSV — warranty is not in the GPS numerator today."
ws["A2"].font = ITALIC_GREY
ws.merge_cells("A2:J2")
ws["A3"] = "Right-hand block shows what GPS 2025 would look like IF warranty value at PR00 were added — for comparison only, NOT the official GPS metric."
ws["A3"].font = Font(italic=True, color="C00000")
ws.merge_cells("A3:J3")

hdr = ["DB","CV qty","CV value PR00 (€)","LCV qty","LCV value PR00 (€)","TOTAL qty","TOTAL value PR00 (€)",
       "Sales 2025 (€)","GPS 2025 (current)","GPS 2025 IF warranty included"]
for c,h in enumerate(hdr,1):
    ws.cell(row=5, column=c, value=h)
style_header(5, len(hdr))
ws.row_dimensions[5].height = 32

r = 6
for code in sorted(per_db.keys()):
    if code not in DB_INFO:
        continue
    d = per_db[code]
    db_name = DB_INFO[code]["db"]
    ms = max_spend(code)
    sales_v = sales_2025.get(code, 0)
    gps_cur = sales_v / ms if ms else None
    gps_with = (sales_v + d["total_val"]) / ms if ms else None

    ws.cell(row=r, column=1, value=db_name).font = BOLD
    ws.cell(row=r, column=2, value=int(d["cv_qty"])).alignment = Alignment(horizontal="right")
    c1 = ws.cell(row=r, column=3, value=d["cv_val"]); c1.number_format = "#,##0"; c1.alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=4, value=int(d["lcv_qty"])).alignment = Alignment(horizontal="right")
    c2 = ws.cell(row=r, column=5, value=d["lcv_val"]); c2.number_format = "#,##0"; c2.alignment = Alignment(horizontal="right")
    c3 = ws.cell(row=r, column=6, value=int(d["total_qty"])); c3.font = BOLD; c3.alignment = Alignment(horizontal="right")
    c4 = ws.cell(row=r, column=7, value=d["total_val"]); c4.number_format = "#,##0"; c4.font = BOLD
    c4.fill = YELLOW; c4.alignment = Alignment(horizontal="right")
    c5 = ws.cell(row=r, column=8, value=sales_v); c5.number_format = "#,##0"; c5.alignment = Alignment(horizontal="right")
    if gps_cur is not None:
        c6 = ws.cell(row=r, column=9, value=gps_cur); c6.number_format = "0.0%"
        c6.fill = gps_fill(gps_cur); c6.font = BOLD; c6.alignment = Alignment(horizontal="right")
    if gps_with is not None:
        c7 = ws.cell(row=r, column=10, value=gps_with); c7.number_format = "0.0%"
        c7.fill = gps_fill(gps_with); c7.font = Font(bold=True, italic=True); c7.alignment = Alignment(horizontal="right")
    for cc in range(1,11): ws.cell(row=r, column=cc).border = BORDER
    r += 1

# Totals row
tot_cv_q = sum(d["cv_qty"] for d in per_db.values())
tot_cv_v = sum(d["cv_val"] for d in per_db.values())
tot_lcv_q = sum(d["lcv_qty"] for d in per_db.values())
tot_lcv_v = sum(d["lcv_val"] for d in per_db.values())
tot_q = tot_cv_q + tot_lcv_q
tot_v = tot_cv_v + tot_lcv_v

ws.cell(row=r, column=1, value="NETWORK TOTAL").font = WHITE_BOLD
ws.cell(row=r, column=1).fill = NAVY
ws.cell(row=r, column=2, value=int(tot_cv_q)).font = WHITE_BOLD
ws.cell(row=r, column=2).fill = NAVY
c = ws.cell(row=r, column=3, value=tot_cv_v); c.number_format = "#,##0"; c.font = WHITE_BOLD; c.fill = NAVY
ws.cell(row=r, column=4, value=int(tot_lcv_q)).font = WHITE_BOLD
ws.cell(row=r, column=4).fill = NAVY
c = ws.cell(row=r, column=5, value=tot_lcv_v); c.number_format = "#,##0"; c.font = WHITE_BOLD; c.fill = NAVY
c = ws.cell(row=r, column=6, value=int(tot_q)); c.font = WHITE_BOLD; c.fill = NAVY
c = ws.cell(row=r, column=7, value=tot_v); c.number_format = "#,##0"; c.font = WHITE_BOLD; c.fill = NAVY
for cc in range(1,11):
    ws.cell(row=r, column=cc).alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=cc).border = BORDER
r += 2

# Notes
ws.cell(row=r, column=1, value="Notes").font = BOLD
r += 1
notes = [
    "Source: Final_Complete_Warranty_Pricing_2026.xlsx — Raw_CV_Warranty + Raw_LCV_Warranty, filtered to 2025 (BILLYMD / AssessMonth).",
    "Pricing: PR00 price master (Pricing_Data sheet, " + f"{len(prices):,}" + " part numbers). The price in the warranty claim itself is NOT used — only quantity and part number.",
    "Critical context: warranty parts are NOT in the GPS sales TSV. The 'GPS 2025 (current)' column shows the official metric. The 'GPS 2025 IF warranty included' column is a what-if — methodologically weaker; only use for context, never as the headline GPS figure.",
    f"Coverage: {cv_lines + lcv_lines:,} warranty lines priced. {cv_missing_price + lcv_missing_price} lines could not be priced (PN not in PR00 master).",
]
for n in notes:
    ws.cell(row=r, column=1, value="• " + n).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    ws.row_dimensions[r].height = 32
    r += 1

for i,w in enumerate([22,10,16,10,16,11,16,16,14,18], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Save master + mirror
wb.save(MASTER)
print(f"\nSaved master: {MASTER}")
shutil.copy(MASTER, PUBLISHED)
print(f"Mirrored to published: {PUBLISHED}")

wb2 = openpyxl.load_workbook(MASTER)
print(f"\nFinal tabs ({len(wb2.sheetnames)}): {wb2.sheetnames}")
