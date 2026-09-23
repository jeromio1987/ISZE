"""_b_recon2.py -- country-code map, IML 1.UIO rows, and seam-reconstruction tests."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

IML = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\IML calculation (1).xlsx"
CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"

wb = openpyxl.load_workbook(IML, data_only=True)
ws = wb["4.Ctry Code"]
print("### IML 4.Ctry Code — rows whose text mentions the markets of interest ###")
for r in range(1, 250):
    row = [ws.cell(r, c).value for c in range(1, 7)]
    txt = " ".join(str(x) for x in row if x is not None).upper()
    if any(k in txt for k in ("GREECE", "GREEK", "CYPRUS", "PORTUG", "IRELAND", "BELGI", "ISRAEL", "NETHERL")):
        print("  r%-3d %s" % (r, row))
print("\n### IML 4.Ctry Code — first 25 rows ###")
for r in range(1, 26):
    print("  r%-3d %s" % (r, [ws.cell(r, c).value for c in range(1, 7)]))

ws = wb["1.UIO"]
print("\n### IML 1.UIO — rows whose Country Name (col B) matches ###")
found = 0
for r in range(6, ws.max_row + 1):
    nm = ws.cell(r, 2).value
    if isinstance(nm, str) and any(k in nm.upper() for k in ("GREECE", "CYPRUS", "PORTUG", "IRELAND", "BELGI", "ISRAEL")):
        vals = [ws.cell(r, c).value for c in range(1, 45)]
        print("  r%-4d %s" % (r, [None if x is None else (round(x, 1) if isinstance(x, float) else x) for x in vals]))
        found += 1
    if found > 25:
        break
print("  matches:", found)
wb.close()

# ---------------- reconstruction tests
CV_DEV = {1:1.00,2:0.99,3:0.98,4:0.97,5:0.96,6:0.95,7:0.90,8:0.80,9:0.70,10:0.55,
          11:0.40,12:0.30,13:0.20,14:0.10,15:0.05,16:0.00}
LCVC = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def conv(regs, c, t, y0=1990):
    return sum((regs.get(v) or 0.0)*S(c, t-v+1) for v in range(y0, t+1))

vol = json.load(open(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_b_iml_vol.json"))
CVV = {int(k): {int(y): v for y, v in d.items()} for k, d in vol['cv_vol'].items()}
LCVV = {int(k): {int(y): v for y, v in d.items()} for k, d in vol['lcv_vol'].items()}

wb = openpyxl.load_workbook(CAND, data_only=True)
ws = wb["1.UIO Overview"]
COLS = {y: 7 + (y-1990) for y in range(1990, 2031)}

SAP_PT_LCV = {2005:101,2006:462,2007:814,2008:476,2009:375,2010:472,2011:310,2012:163,2013:136,
              2014:210,2015:268,2016:372,2017:483,2018:609,2019:654,2020:327,2021:346,2022:394,
              2023:636,2024:975,2025:680}
SAP_PT_CV  = {2005:14,2006:74,2007:48,2008:77,2009:150,2010:168,2011:131,2012:52,2013:66,
              2014:54,2015:81,2016:173,2017:206,2018:267,2019:333,2020:208,2021:327,2022:314,
              2023:357,2024:265,2025:136}

for r, code, mt, src, alt in ((54, 217, "LCV", LCVV, SAP_PT_LCV),
                              (90, 233, "LCV", LCVV, None),
                              (79, 230, "CV", CVV, None)):
    c = LCVC if mt == "LCV" else CV_DEV
    ser = {y: float(ws.cell(r, COLS[y]).value or 0) for y in range(1990, 2031)}
    print("\n" + "=" * 100)
    print("r%d  ctry %d  model %s" % (r, code, mt))
    for label, regs in (("IML Vol only", src.get(code, {})),
                        ("IML Vol <=2004 + SAP alt 2005+", dict(list(src.get(code, {}).items()) + list((alt or {}).items())))):
        if not regs:
            continue
        print("  -- %s --" % label)
        for y in range(1990, 2026):
            g = conv(regs, c, y, 1990)
            print("     %d  conv %10.1f   V19 %10.0f   diff %+9.1f" % (y, g, ser[y], g-ser[y]))
wb.close()
