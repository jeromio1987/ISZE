"""_b_scope.py -- separate the two changes; measure each one's blast radius."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19W = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
CAND_A = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
YEARS = list(range(1990, 2031)); FY = [2026, 2027, 2028, 2029, 2030]
CV_DEV = {1:1.00,2:0.99,3:0.98,4:0.97,5:0.96,6:0.95,7:0.90,8:0.80,9:0.70,10:0.55,
          11:0.40,12:0.30,13:0.20,14:0.10,15:0.05,16:0.00}
LCV = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def curve(mt): return LCV if str(mt)=='LCV' else CV_DEV
OUT_OF_SCOPE={7,8}; LEAVE_AS_V19={22,24}; PLAN_ONLY_MD={27,52,80,89}
UK_LCV_REGS={2026:3020,2027:2400,2028:2400,2029:2400,2030:2400}

wb = openpyxl.load_workbook(V19W, data_only=True); ws = wb["1.UIO Overview"]
COLS={y:7+(y-1990) for y in YEARS}
ROWS=[{'row':r,'meta':[ws.cell(r,c).value for c in range(1,7)],
       'vals':{y: ws.cell(r,COLS[y]).value for y in YEARS}} for r in range(7,91)]
wb.close()
wba = openpyxl.load_workbook(CAND_A, data_only=True); wsa=wba['1.UIO Overview']
A_ROWS={r:[wsa.cell(r,COLS[y]).value for y in FY] for r in range(7,91)}
wba.close()

def bs(ser,c,start):
    R={}
    for t in YEARS:
        if t<start: R[t]=0.0; continue
        R[t]=ser[t]-sum(R[v]*S(c,t-v+1) for v in YEARS if start<=v<t)
    return R
def seam_of(ser,c,mode):
    if mode=='none': return 1990
    R=bs(ser,c,1990)
    for t in range(1991,2026):
        if ser[t]>0 and R[t]<0 and (R[t]<-1.0 if mode=='material' else abs(R[t])>0.25*ser[t]):
            return t
    return 1990
def state(ser,c,seam):
    R=bs(ser,c,seam)
    for v in YEARS:
        if R[v]<0: R[v]=0.0
    reb=sum(R[v]*S(c,2026-v) for v in YEARS if v<=2025)
    res=ser[2025]-reb
    if abs(res)>0.5 and ser[2025]>0:
        oldest=min([v for v in YEARS if S(c,2026-v)>0] or [2025])
        R[oldest]+=res/S(c,2026-oldest)
    return R
def smooth(x):
    n=len(x); sm=[]
    for i in range(n):
        if i==0: sm.append((x[0]+x[1])/2.0)
        elif i==n-1: sm.append((x[n-2]+x[n-1])/2.0)
        else: sm.append((x[i-1]+x[i]+x[i+1])/3.0)
    tx,ts=sum(x),sum(sm); f=(tx/ts) if ts else 1.0
    return [z*f for z in sm]

def run(precision, seam_scope=None):
    """precision: 'int' (=candidate a) or 'float'. seam_scope: None=all rows pct (candidate a),
       or a set of rows to which the material rule is applied."""
    res={}
    for row in ROWS:
        r=row['row']; mt=row['meta'][3]
        if r in OUT_OF_SCOPE or r in LEAVE_AS_V19: continue
        v={y: float(row['vals'].get(y) or 0) for y in YEARS}
        if all(v[t]==0 for t in FY) and v[2025]==0: continue
        c=curve(mt)
        if r==25:
            u={2025:v[2025]}; out={}
            for t in FY: u[t]=u[t-1]*(1-0.123)+UK_LCV_REGS[t]; out[t]=int(round(u[t]))
            res[r]=[out[t] for t in FY]; continue
        if r in PLAN_ONLY_MD:
            res[r]=[0]*5; continue
        raw={t:(v.get(t) or 0)-0.97*(v.get(t-1) or 0) for t in FY}
        if precision=='int': raw={t: round(raw[t]) for t in FY}
        prof=smooth([raw[t] for t in FY])
        prof=[round(z) for z in prof]
        if seam_scope and r in seam_scope: sm_ = 'material'
        else: sm_ = 'pct'
        R=state(v,c,seam_of(v,c,sm_))
        out={}
        for i,t in enumerate(FY):
            stock=sum(R[vg]*S(c,t-vg+1) for vg in range(1990,t))
            R[t]=prof[i]
            out[t]=int(round(stock+R[t]*S(c,1)))
        res[r]=[out[t] for t in FY]
    return res

BASE = run('int'            , None)          # candidate a
P    = run('float'          , None)          # precision only
PS   = run('float'          , {54, 90})      # precision + scoped seam
print("sanity: BASE reproduces candidate a on all rows?",
      all(BASE[r]==[round(float(x or 0)) for x in A_ROWS[r]] for r in BASE if r not in PLAN_ONLY_MD
          and r not in (22,24) and max([abs(float(x or 0)) for x in A_ROWS[r]] or [0])>0))

for tag, R in (("PRECISION ONLY (global)", P), ("PRECISION + seam on r54/r90", PS)):
    rows=[]; tot={y:0.0 for y in FY}
    for r in sorted(R):
        a=[round(float(x or 0)) for x in A_ROWS[r]]
        d=[R[r][i]-a[i] for i in range(5)]
        if any(abs(x)>=1 for x in d): rows.append((r, d))
        for i,y in enumerate(FY): tot[y]+=R[r][i]
    print("\n=== %s ===" % tag)
    print("  rows whose 2026-2030 cells move by >=1 unit: %d" % len(rows))
    for r,d in rows:
        print("     r%-3d delta=%s" % (r,d))
    print("  Grand Total: " + "  ".join("%d=%d" % (y, tot[y]) for y in FY))

# harness signatures under the final combination
def ratio(v):
    d=[v[i+1]-v[i] for i in range(4)]
    if len(set(1 if x>0 else (-1 if x<0 else 0) for x in d))>2: return None,d
    rs=[d[i+1]/d[i] for i in range(3) if abs(d[i])>1e-9]
    return (sum(rs)/len(rs) if rs else None), d
print("\n=== harness signature on the four adjudicated rows, final combination ===")
for r in (33,54,79,90):
    m,d = ratio(PS[r])
    f_rate = 0.03 if (m is not None and abs(m-0.97)<=0.01) else (
             0.123 if (m is not None and abs(m-0.877)<=0.01) else (
             0.0 if (m is not None and abs(m-1.0)<=0.01) else None))
    print("  r%-3d v20=%-38s deltas=%-22s mean=%s | C1 switch=%-5s C5 hit=%s" % (
        r, str(PS[r]), str(d), ("%.4f"%m) if m is not None else "n/a",
        (f_rate is not None and f_rate<0.08), (m is not None and abs(m-0.97)<=0.01)))
