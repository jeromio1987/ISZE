"""Referenced by V20_SPEC.md, section "Review of candidate c - decisions 24 September 2026".
r69 POLAND LCV: Jerome accepted this rebuild (24 Sep 2026) in place of candidate c's untraceable values.
r88 CYPRUS CV LD: reproduces candidate c exactly (lambda 0.7582) - kept.
Run:  python spec_20260924_r69_r88_trace.py
"""
"""Independent trace of candidate c rows r69 (Poland LCV) and r88 (Cyprus CV LD).

Method as stated in the candidate c log: back-solve yearly cohorts from V19's own
1990-2025 history, clip negative cohorts to 0, scale the surviving cohort state so
it reproduces the 2025 level (lambda), then roll 2026-2030 forward on the curve with
the registration input from sheet 'V20 REGS USED'. Written without reading the
builder's scripts.
"""
import openpyxl

WB = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22c.xlsx"
V19 = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"

CV = [1, .99, .98, .97, .96, .95, .9, .8, .7, .55, .4, .3, .2, .1, .05, 0]
LCV = [1, .98, .95, .92, .89, .84, .8, .65, .5, .4, .2, 0]


def S(curve, age):
    return curve[age - 1] if 1 <= age <= len(curve) else 0.0


def row_values(path, r):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["1.UIO Overview"]
    row = next(ws.iter_rows(min_row=r, max_row=r, values_only=True))
    # col G (idx 6) = 1990 ... col AP (idx 41) = 2025E, AQ:AU (42..46) = 2026-2030
    hist = {1990 + i: (row[6 + i] or 0) for i in range(36)}
    fut = [row[42 + i] for i in range(5)]
    return hist, fut


def regs_used(r):
    ws = openpyxl.load_workbook(WB, data_only=True, read_only=True)["V20 REGS USED"]
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == r:
            return [float(x) for x in row[5:10]]  # first block = regs actually used
    raise KeyError(r)


def trace(r, curve, label):
    hist, c_fut = row_values(V19, r)[0], row_values(WB, r)[1]
    R = {}
    for t in range(1990, 2026):
        R[t] = hist[t] - sum(R[v] * S(curve, t - v + 1) for v in R)
    neg = {v: round(x, 1) for v, x in R.items() if x < -0.5}
    Rc = {v: max(0.0, x) for v, x in R.items()}
    stock25 = sum(x * S(curve, 2025 - v + 1) for v, x in Rc.items())
    lam = hist[2025] / stock25 if stock25 else float("nan")
    Rs = {v: x * lam for v, x in Rc.items()}
    ru = regs_used(r)
    ru_clip = [max(0.0, x) for x in ru]
    out = []
    for i, t in enumerate(range(2026, 2031)):
        Rs[t] = ru_clip[i]
        out.append(round(sum(x * S(curve, t - v + 1) for v, x in Rs.items())))
    print(f"== r{r} {label}")
    print("  V19 2019-2025          :", [hist[y] for y in range(2019, 2026)])
    print("  negative cohorts (V19) :", neg)
    print("  regs used (sheet)      :", [round(x, 2) for x in ru], "-> clipped", [round(x, 2) for x in ru_clip])
    print(f"  lambda                 : {lam:.4f}")
    print("  traced 2026-2030       :", out)
    print("  candidate c 2026-2030  :", c_fut)
    print("  match                  :", out == c_fut)


trace(69, LCV, "POLAND LCV (IBX)")
trace(88, CV, "CYPRUS CV LD (Loutsios)")
