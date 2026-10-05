"""
fix_v20_row22_24_links.py -- restore the live sheet-4 cross-sheet formulas on rows 22 and 24
cols AQ:AU of the V20 candidate. The bulk block write replaced them with their cached values;
the numbers were right but the live links must survive. V19 is read-only.
"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19 = os.path.join(CUR, "ISUZU_UIO_Overview_2026.xlsx")
CAND = os.path.join(CUR, "ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx")

wbs = openpyxl.load_workbook(V19, data_only=False)
wss = wbs['1.UIO Overview']
FORM = {}
for r in (22, 24):
    for c in range(43, 48):                 # AQ..AU
        FORM[(r, c)] = wss.cell(row=r, column=c).value
wbs.close()
for k, v in FORM.items():
    print("source %s" % (k,), str(v)[:80])
    assert isinstance(v, str) and v.startswith('='), "expected formula at %s" % (k,)

import win32com.client as win32
excel = win32.DispatchEx("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
wb = None
try:
    wb = excel.Workbooks.Open(CAND, UpdateLinks=0, ReadOnly=False)
    sh = wb.Worksheets("1.UIO Overview")
    for (r, c), f in FORM.items():
        sh.Cells(r, c).Formula = f
        print("restored %s%d" % (get_column_letter(c), r))
    excel.CalculateFullRebuild()
    wb.Save()
    print("saved")
finally:
    if wb is not None:
        try: wb.Close(SaveChanges=False)
        except Exception: pass
    try: excel.Quit()
    except Exception: pass

# re-verify
wbv = openpyxl.load_workbook(CAND, data_only=False)['1.UIO Overview']
for r in (22, 24, 25, 26):
    print("R%-3d AQ = %s" % (r, str(wbv.cell(row=r, column=43).value)[:90]))
wbv2 = openpyxl.load_workbook(CAND, data_only=True)['1.UIO Overview']
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
print("\nGrand Total cached 2026-2030:", [wbv2.cell(row=2, column=COLS[y]).value for y in (2026,2030)])
