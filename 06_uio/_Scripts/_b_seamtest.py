"""_b_seamtest.py -- does detect_seam fire at all? what does a material rule change?"""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19W = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
YEARS = list(range(1990, 2031)); FY = [2026, 2027, 2028, 2029, 2030]
CV_DEV = {1:1.00,2:0.99,3:0.98,4:0.97,5:0.96,6:0.95,7:0.90,8:0.80,9:0.70,10:0.55,
          11:0.40,12:0.30,13:0.20,14:0.10,15:0.05,16:0.00}
LCV = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def curve(mt): return LCV if str(mt)=='LCV' else CV_DEV
OUT_OF_SCOPE={7,8}; LEAVE_AS_V19={22,24}; PLAN_ONLY_MD={27,52,80,89}
UK_LCV_REGS={2026:3020,2027:2400,2028:2400,2029:2400,2030:2400}

wb = openpyxl.load_workbook(V19W, data_only=True); ws = wb["1.UIO Overview"]
COLS = {y: 7+(y-1990) for y in YEARS}
ROWS=[{'row':r,'meta':[ws.cell(r,c).value for c in range(1,7)],
       'vals':{y: ws.cell(r,COLS[y]).value for y in YEARS}} for r in range(7,91)]
wb.close()

def bs(ser,c,start):
    R={}
    for t in YEARS:
        if t<start: R[t]=0.0; continue
        R[t]=ser[t]-sum(R[v]*S(c,t-v+1) for v in YEARS if start<=v<t)
    return R

def seam_pct(ser,c):
    R=bs(ser,c,1990)
    for t in range(1991,2026):
        if ser[t]>0 and R[t]<0 and abs(R[t])>0.25*ser[t]: return t,R
    return 1990,R
def seam_mat(ser,c,tol=1.0):
    R=bs(ser,c,1990)
    for t in range(1991,2026):
        if ser[t]>0 and R[t]<-tol: return t,R
    return 1990,R

print("%-5s %-9s %-4s %-6s | %-14s %-14s | %s" % ("row","country","mod","2025","seam_pct(engine)","seam_material","first negative implied cohort"))
for row in ROWS:
    r=row['row']; mt=row['meta'][3]
    if r in OUT_OF_SCOPE or r in LEAVE_AS_V19: continue
    v={y: float(row['vals'].get(y) or 0) for y in YEARS}
    if all(v[t]==0 for t in FY) and v[2025]==0: continue
    c=curve(mt)
    sp,Rp=seam_pct(v,c); sm,Rm=seam_mat(v,c)
    negs={t:round(Rm[t],1) for t in range(1991,2026) if Rm[t]<-1.0}
    if sp!=1990 or sm!=1990 or negs:
        print("r%-4d %-9s %-4s %6.0f | %-14s %-14s | %s" % (
            r,str(row['meta'][1])[:9],str(mt)[:4],v[2025],
            (sp if sp!=1990 else '-'), (sm if sm!=1990 else '-'), negs))
