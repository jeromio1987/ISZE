"""_b_iml2.py -- extract IML 6.CV Vol / 7.LCV Vol registration series by country code."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

IML = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\IML calculation (1).xlsx"
CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
V19 = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"

wb = openpyxl.load_workbook(IML, data_only=True)
print("SHEETS:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("   %-40s %s" % (sn, ws.calculate_dimension()))

# --- parse a "vol" sheet: find the year header row and the value columns
def parse_vol(ws, name):
    # locate row holding year numbers
    hdr_row = None
    hdr_cols = {}
    for r in range(1, 12):
        vals = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and 1990 <= v <= 2035:
                vals[c] = int(v)
        if len(vals) >= 5:
            hdr_row, hdr_cols = r, vals
            break
    print("\n%s: header row %s  n_year_cols=%d" % (name, hdr_row, len(hdr_cols)))
    if not hdr_row:
        return None, None
    cols = sorted(hdr_cols)
    print("   years %d..%d" % (hdr_cols[cols[0]], hdr_cols[cols[-1]]))
    # rows below: first col = country code
    series = {}
    for r in range(hdr_row + 1, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if not isinstance(code, (int, float)):
            continue
        code = int(code)
        d = {}
        for c in cols:
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                d[hdr_cols[c]] = float(v)
        if d:
            series.setdefault(code, {}).update(d)
    return hdr_row, series

ws = wb["6.CV Vol"]
_, cv_vol = parse_vol(ws, "6.CV Vol")
ws = wb["7.LCV Vol"]
_, lcv_vol = parse_vol(ws, "7.LCV Vol")
wb.close()

json.dump({'cv_vol': {str(k): v for k, v in cv_vol.items()},
           'lcv_vol': {str(k): v for k, v in lcv_vol.items()}},
          open(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_b_iml_vol.json", "w"), indent=1)
print("\nsaved _b_iml_vol.json ; cv codes %d, lcv codes %d" % (len(cv_vol), len(lcv_vol)))

for code, label in ((217, "PORTUGL 217"), (230, "GREECE 230"), (233, "CYPRUS 233"),
                    (206, "IRELAND 206"), (205, "U KING 205"), (208, "BELGIUM 208"),
                    (204, "DENMARK 204"), (201, "NETH? 201")):
    print("\n--- CV  %s ---" % label)
    print("   ", {y: cv_vol.get(code, {}).get(y) for y in range(1990, 2027)})
    print("--- LCV %s ---" % label)
    print("   ", {y: lcv_vol.get(code, {}).get(y) for y in range(1990, 2027)})

# cross-check: code list
print("\nCV codes:", sorted(cv_vol))
print("LCV codes:", sorted(lcv_vol))

# compare to candidate history
cwb = openpyxl.load_workbook(CAND, data_only=True)
cws = cwb["1.UIO Overview"]
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
for r, code, mt in ((79, 230, "CV"), (90, 233, "LCV"), (54, 217, "LCV"), (88, 233, "CV"), (51, 217, "CV")):
    hist = {y: cws.cell(r, COLS[y]).value for y in range(1990, 2026)}
    src = cv_vol if mt == "CV" else lcv_vol
    print("\n=== r%d code %d (%s) ===" % (r, code, mt))
    print("  candidate hist:", {y: hist[y] for y in hist if hist[y] is not None})
    print("  IML %s      :" % mt, {y: src.get(code, {}).get(y) for y in range(1990, 2026)})
cwb.close()
