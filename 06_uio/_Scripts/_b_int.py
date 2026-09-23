"""_b_int.py -- are AQ:AU row values integers or floats? which rows are fractional?"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
V19  = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
FY = [2026, 2027, 2028, 2029, 2030]

for path, tag in ((CAND, "CANDIDATE"), (V19, "V19")):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["1.UIO Overview"]
    frac, ints, blankish = [], 0, 0
    for r in range(7, 91):
        got = [ws.cell(r, COLS[y]).value for y in FY]
        nums = [v for v in got if isinstance(v, (int, float))]
        if not nums:
            blankish += 1; continue
        if any(abs(v - round(v)) > 1e-9 for v in nums):
            frac.append((r, ws.cell(r, 2).value, [round(v, 3) for v in got]))
        else:
            ints += 1
    print("### %s : rows all-integer %d ; fractional %d ; no numeric %d" % (tag, ints, len(frac), blankish))
    for r, ctry, v in frac[:30]:
        print("    r%-3d %-9s %s" % (r, ctry, v))
    # historical fractional too?
    histfrac = 0
    for r in range(7, 91):
        for y in range(2014, 2026):
            v = ws.cell(r, COLS[y]).value
            if isinstance(v, float) and abs(v - round(v)) > 1e-9:
                histfrac += 1
                break
    print("    rows with fractional historical values 2014-2025:", histfrac)
    print("    grand total 2026 = %r ; row2 formula? " % ws.cell(2, COLS[2026]).value)
    wb.close()

# formula view of row 2 and a couple of rows
wb = openpyxl.load_workbook(CAND, data_only=False)
ws = wb["1.UIO Overview"]
for cell in ("B2", "G2", "AQ2", "AQ26", "AQ79", "AQ33"):
    print("formula %s = %r" % (cell, ws[cell].value))
print("row 2 label:", ws["A2"].value, "| row3:", ws["A3"].value, "| row4:", ws["A4"].value)
wb.close()
