"""_b_recon.py -- read-only recon of V20 candidate: sheets, rows, AQ:AU, r79/r54/r90 history."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
V19 = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"

wb = openpyxl.load_workbook(CAND, data_only=True)
print("SHEETS:", wb.sheetnames)

ws = wb["1.UIO Overview"]
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
FY = [2026, 2027, 2028, 2029, 2030]

print("\n--- Grand Total row 2 (2026-2030) ---")
for y in FY:
    print(" %d %s" % (y, ws.cell(2, COLS[y]).value))

print("\n--- target rows: meta + AQ:AU + history tail ---")
for r in (33, 54, 79, 90, 88, 51, 26, 29, 25, 19):
    meta = [ws.cell(r, c).value for c in range(1, 7)]
    fu = [ws.cell(r, COLS[y]).value for y in FY]
    tail = [(y, ws.cell(r, COLS[y]).value) for y in range(1990, 2031)]
    print("\nr%d meta=%s" % (r, meta))
    print("   2025=%s  AQ:AU=%s" % (ws.cell(r, COLS[2025]).value, fu))
    print("   AQ:AU formulas? (need non-data_only) ")

wbf = openpyxl.load_workbook(CAND, data_only=False)
wsf = wbf["1.UIO Overview"]
for r in (33, 54, 79, 90, 26, 25):
    print("r%d formulas AQ:AU = %s" % (r, [wsf.cell(r, COLS[y]).value for y in FY]))

print("\n--- V20 flag col AV / AW ---")
for r in range(7, 91):
    a = ws.cell(r, 48).value
    if a:
        print("  r%d %s | AV=%s" % (r, ws.cell(r, 2).value, a))

# full history for target rows from candidate
out = {}
for r in (33, 54, 79, 90):
    out[r] = {str(y): ws.cell(r, COLS[y]).value for y in range(1990, 2031)}
json.dump(out, open(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_b_targets.json", "w"), indent=1)
print("\nsaved _b_targets.json")

# also check extra sheets content
for s in wb.sheetnames:
    if s.startswith("V20"):
        w = wb[s]
        print("\n=== sheet %s dims %s ===" % (s, w.dimensions))
        for i, row in enumerate(w.iter_rows(min_row=1, max_row=min(w.max_row, 12), values_only=True), 1):
            print("  ", [str(x)[:60] if x is not None else "" for x in row][:8])
