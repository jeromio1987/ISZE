"""V20 — dump supporting sheets of V19 (READ ONLY)."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

XL = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"
wbv = openpyxl.load_workbook(XL, data_only=True)
wbf = openpyxl.load_workbook(XL, data_only=False)

def dump(name, maxr=80, maxc=20, formulas=False):
    ws = (wbf if formulas else wbv)[name]
    print("\n" + "=" * 110)
    print("SHEET: %s   dims=%sx%s" % (name, ws.max_row, ws.max_column))
    print("=" * 110)
    for r in range(1, min(ws.max_row, maxr) + 1):
        cells = []
        for c in range(1, min(ws.max_column, maxc) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or str(v).strip() == '':
                continue
            s = str(v).replace('\n', ' ')
            cells.append("%s=%s" % (get_column_letter(c), s[:55]))
        if cells:
            print("R%-3d %s" % (r, " | ".join(cells)))

dump('3. UIO Calculation', maxr=60, maxc=16)
dump('4. ITUK UIO Projection', maxr=70, maxc=40)
