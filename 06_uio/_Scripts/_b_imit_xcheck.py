"""_b_imit_xcheck.py -- cross-check the fixed LCV rows against the IMIT LCV UIO study,
which the workbook itself cites as its LCV source."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
IMIT = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_SourceData\LCV UIO (November 2025)_EU_Study_Imit.xlsx"
CAND_A = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
CAND_B = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22b.xlsx"
V19    = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"

# --- IMIT: sheet 'ISZE', CY2019..CY2028 in cols B..K, distributor names col A
wb = openpyxl.load_workbook(IMIT, data_only=True)
ws = wb["ISZE"]
years = {}
for c in range(2, 12):
    v = ws.cell(2, c).value
    if isinstance(v, str) and v.startswith("CY"):
        years[int(v[2:6])] = c
print("IMIT 'ISZE' sheet years -> cols:", years)
imit = {}
for r in range(3, 60):
    name = ws.cell(r, 1).value
    if not isinstance(name, str):
        continue
    d = {}
    for y, c in years.items():
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)):
            d[y] = float(v)
    if d:
        imit[name.strip()] = d
wb.close()

print("IMIT distributors captured:", len(imit))
for k in sorted(imit):
    if any(t in k.lower() for t in ("petros", "loutsios", "imotor", "greece", "cyprus", "portug")):
        print("   %-40s %s" % (k, {y: round(v, 2) for y, v in sorted(imit[k].items())}))

def grab(p):
    w = openpyxl.load_workbook(p, data_only=True); s = w["1.UIO Overview"]
    out = {}
    for r in (54, 81, 90):
        out[r] = {y: s.cell(r, 7 + (y - 1990)).value for y in range(2019, 2031)}
    w.close(); return out

A, B, V = grab(CAND_A), grab(CAND_B), grab(V19)
PAIRS = [
    (90, "LOUTSIOS A & SONS LTD Cyprus", "CYPRUS LCV (Loutsios & Sons)"),
    (54, "IMOTORS Portugal",            "PORTUGL LCV (Imotors)"),
    (81, "PETROS Greece",               "GREECE LCV (Petros Petropoulos)"),
]
for r, ikey, label in PAIRS:
    print("\n" + "=" * 100)
    print("r%d  %s   vs IMIT '%s'" % (r, label, ikey))
    im = imit.get(ikey, {})
    print("  %-6s %10s %10s %10s %10s | %10s %10s" % ("year", "IMIT", "V19", "cand a", "cand b",
                                                      "b-IMIT %", "a-IMIT %"))
    for y in range(2025, 2031):
        i = im.get(y)
        v, a, b = V[r].get(y), A[r].get(y), B[r].get(y)
        if b is None:
            continue
        bs = ("%+.1f%%" % (100*(b/i - 1))) if i else "-"
        as_ = ("%+.1f%%" % (100*(a/i - 1))) if i else "-"
        print("  %-6d %10s %10s %10s %10s | %10s %10s" % (
            y, ("%.0f" % i) if i else "-", ("%.0f" % v) if v else "-",
            ("%.0f" % a) if a else "-", ("%.0f" % b) if b else "-", bs, as_))
