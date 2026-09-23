"""_b_sources.py -- read-only probe of registration sources for GR / PT / CY."""
import io, sys, json, traceback
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CANDS = [
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\07_UIO\isuzuregdata.xlsx",
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\07_UIO\Matriculas Isuzu SAP.xlsx",
    r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_SourceData\UIO.xlsx",
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\08_Distributor\Imotors_Portugal_UIO_GPS_PMS_Study.xlsx",
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imie_mail_20260730\UIO ISUZU 2025.xlsx",
]

for p in CANDS:
    print("=" * 110)
    print(p)
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        print("  sheets:", wb.sheetnames)
        for ws in wb.worksheets:
            print("  --- %s  dims=%s" % (ws.title, ws.calculate_dimension()))
            n = 0
            for row in ws.iter_rows(min_row=1, max_row=14, values_only=True):
                cells = [("" if c is None else str(c))[:26] for c in row[:12]]
                if any(cells):
                    print("      ", cells)
                n += 1
                if n >= 14:
                    break
        wb.close()
    except Exception as e:
        print("  ERROR", traceback.format_exception_only(type(e), e))
