"""V20 - post-build verification of the candidate (READ ONLY on candidate and V19)."""
import io, sys, os, json, zipfile, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19 = os.path.join(CUR, "ISUZU_UIO_Overview_2026.xlsx")
CAND = os.path.join(CUR, "ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx")
V19_ARCH = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Archive\ISUZU_UIO_Overview_2026_v18.xlsx"
FY = [2026, 2027, 2028, 2029, 2030]
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}

print("=" * 100)
print("1. FILE INTEGRITY")
print("=" * 100)
for label, p in [('V19 LIVE (must be untouched)', V19), ('V18 ARCHIVE (must be untouched)', V19_ARCH),
                 ('V20 CANDIDATE', CAND)]:
    st = os.stat(p)
    print("%-34s %10d bytes  mtime=%s" % (label, st.st_size, datetime.datetime.fromtimestamp(st.st_mtime)))

print("\nB4 on live V19 still =", openpyxl.load_workbook(V19, read_only=True)['1.UIO Overview']['B4'].value)

print("\n" + "=" * 100)
print("2. CANDIDATE - sheets and embedded assets")
print("=" * 100)
wbc = openpyxl.load_workbook(CAND, data_only=False)
print("sheets:", wbc.sheetnames)
z = zipfile.ZipFile(CAND)
for n in z.namelist():
    if n.startswith('xl/media') or n.startswith('xl/charts'):
        print("  asset:", n, z.getinfo(n).file_size)

print("\n" + "=" * 100)
print("3. CANDIDATE - cols AQ:AU cached values + scope flags")
print("=" * 100)
wbv = openpyxl.load_workbook(CAND, data_only=True)
ws = wbv['1.UIO Overview']
print("Row2 GRAND TOTAL 2026-2030 cached:", [ws.cell(row=2, column=COLS[y]).value for y in FY])
print("Row5 MIDI       2026-2030 cached:", [ws.cell(row=5, column=COLS[y]).value for y in FY])
print("Row3 growth     2026-2030 cached:", [ws.cell(row=3, column=COLS[y]).value for y in FY])
print("\n%-4s %-10s %-11s %9s %9s %9s %9s %9s  %s" % ("R","Country","Model","2026","2027","2028","2029","2030","flag"))
for r in range(7, 91):
    vals = [ws.cell(row=r, column=COLS[y]).value for y in FY]
    if any(v not in (None, 0) for v in vals):
        print("%-4d %-10s %-11s %9s %9s %9s %9s %9s  %s" % (
            r, str(ws.cell(row=r, column=2).value)[:10], str(ws.cell(row=r, column=4).value)[:11],
            *[('' if v is None else v) for v in vals], str(ws.cell(row=r, column=48).value)[:34]))

print("\n" + "=" * 100)
print("4. CHECK ROW TYPES (formulas vs values)")
print("=" * 100)
wbf = openpyxl.load_workbook(CAND, data_only=False)
wsf = wbf['1.UIO Overview']
for r in (22, 24, 25, 26, 27, 42, 7, 52):
    cells = [wsf.cell(row=r, column=COLS[y]).value for y in FY]
    kind = 'FORMULA' if isinstance(cells[0], str) and str(cells[0]).startswith('=') else 'value'
    print("R%-3d %-9s %-10s %-9s %s" % (r, str(wsf.cell(row=r,column=2).value)[:9],
          str(wsf.cell(row=r,column=4).value)[:10], kind, str(cells[0])[:60]))
print("AQ2 formula:", wsf['AQ2'].value)

print("\n" + "=" * 100)
print("5. NEW SHEETS CONTENT")
print("=" * 100)
for sn in ['V20 PLAN ONLY', 'V20 CHANGE LOG', 'V20 ASSUMPTIONS']:
    if sn in wbv.sheetnames:
        w = wbv[sn]
        print("\n--- %s (%dx%d) ---" % (sn, w.max_row, w.max_column))
        for r in range(1, min(w.max_row, 60) + 1):
            cells = []
            for c in range(1, min(w.max_column, 16) + 1):
                v = w.cell(row=r, column=c).value
                if v is None or str(v).strip() == '':
                    continue
                cells.append("%s=%s" % (get_column_letter(c), str(v)[:46]))
            if cells:
                print("  R%-3d %s" % (r, " | ".join(cells)))
    else:
        print("MISSING:", sn)

print("\n" + "=" * 100)
print("6. V19 HISTORICAL YEARS UNCHANGED IN CANDIDATE")
print("=" * 100)
v19v = openpyxl.load_workbook(V19, data_only=True)['1.UIO Overview']
diffs = 0
for r in range(7, 91):
    for y in range(1990, 2026):
        a = v19v.cell(row=r, column=COLS[y]).value
        b = ws.cell(row=r, column=COLS[y]).value
        if a != b:
            diffs += 1
            if diffs <= 10:
                print("  MISMATCH R%d %s: V19=%s CAND=%s" % (r, y, a, b))
print("historical cell mismatches 1990-2025: %d" % diffs)
