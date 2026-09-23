"""V20 — full series extraction + feasibility checks (READ ONLY)."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

XL = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"
wbv = openpyxl.load_workbook(XL, data_only=True)
ws = wbv['1.UIO Overview']
YEARS = list(range(1990, 2031))
COLS = {y: 7 + (y - 1990) for y in YEARS}

print("row  ctry_code country    cv    model       dist        firstYr firstVal  1990 1995 2000 2005 2010 2015 2020 2024 2025")
rows = []
for r in range(7, 91):
    vals = {y: ws.cell(row=r, column=COLS[y]).value for y in YEARS}
    nz = [y for y in YEARS if vals[y] not in (None, 0)]
    first = (nz[0], vals[nz[0]]) if nz else (None, None)
    meta = [ws.cell(row=r, column=c).value for c in range(1, 7)]
    rows.append({'row': r, 'meta': meta, 'vals': {str(y): vals[y] for y in YEARS}})
    print("%-4d %-9s %-10s %-5s %-11s %-11s %-9s %-9s %8s %8s %8s %8s %8s %8s %8s %8s %8s" % (
        r, str(meta[0])[:9], str(meta[1])[:10], str(meta[2])[:5], str(meta[3])[:11], str(meta[4] or '')[:11],
        str(first[0]), str(first[1]),
        *[str(vals[y] if vals[y] is not None else '')[:8] for y in (1990, 1995, 2000, 2005, 2010, 2015, 2020, 2024, 2025)]))

with open(r'C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_v20_v19_full.json', 'w', encoding='utf-8') as f:
    json.dump(rows, f, indent=1, default=str)
print("\nsaved _v20_v19_full.json")
