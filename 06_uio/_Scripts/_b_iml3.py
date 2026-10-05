"""_b_iml3.py -- robust parse of IML 6.CV Vol / 7.LCV Vol (code col may be text)."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

IML = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\IML calculation (1).xlsx"
wb = openpyxl.load_workbook(IML, data_only=True)

for sn in ("6.CV Vol", "7.LCV Vol"):
    ws = wb[sn]
    print("=" * 100)
    print(sn, ws.calculate_dimension())
    for r in range(1, 9):
        print("  row%-2d %s" % (r, [ws.cell(r, c).value for c in range(1, 12)]))
    print("  col A sample:", [ws.cell(r, 1).value for r in range(4, 12)])

def parse(ws):
    # years on row 3, value cols = those holding a year int in 1990..2035
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(3, c).value
        if isinstance(v, (int, float)) and 1990 <= v <= 2035:
            hdr[c] = int(v)
    out = {}
    for r in range(5, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        code = None
        if isinstance(raw, (int, float)):
            code = int(raw)
        elif isinstance(raw, str) and raw.strip().isdigit():
            code = int(raw.strip())
        if code is None:
            continue
        d = {}
        for c, y in hdr.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                d[y] = float(v)
        if d:
            out.setdefault(code, {}).update(d)
    return hdr, out

res = {}
for sn in ("6.CV Vol", "7.LCV Vol"):
    hdr, ser = parse(wb[sn])
    key = 'cv_vol' if sn.startswith('6') else 'lcv_vol'
    res[key] = {str(k): v for k, v in ser.items()}
    print("\n%s: n_cols=%d  n_codes=%d  codes=%s" % (sn, len(hdr), len(ser), sorted(ser)[:30]))
wb.close()
json.dump(res, open(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_b_iml_vol.json", "w"), indent=1)
print("saved")
