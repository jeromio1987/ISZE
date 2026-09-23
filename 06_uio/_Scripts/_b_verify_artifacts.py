"""_b_verify_artifacts.py -- verify the delivered workbook, log, changelog."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND_A = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
CAND_B = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22b.xlsx"
V19 = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
FY = [2026, 2027, 2028, 2029, 2030]
COLS = {y: 7 + (y - 1990) for y in FY}

print("### 1. workbook sheets + A/B/V19 untouched check")
for p, tag in ((V19, "V19"), (CAND_A, "A"), (CAND_B, "B")):
    wb = openpyxl.load_workbook(p, data_only=True)
    print("  %-4s %s" % (tag, wb.sheetnames))
    wb.close()

print("\n### 2. V20 REGS USED sheet in candidate b")
wb = openpyxl.load_workbook(CAND_B, data_only=True)
ws = wb["V20 REGS USED"]
print("  dims:", ws.dimensions)
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 14),
                                     max_col=12, values_only=True), 1):
    cells = [("" if c is None else str(c))[:22] for c in row]
    while cells and cells[-1] == "":
        cells.pop()
    if cells:
        print("   %3d| %s" % (i, " | ".join(cells)))
print("   ... total rows:", ws.max_row)

print("\n### 3. AV flag column, rows of interest")
ws = wb["1.UIO Overview"]
for r in (33, 42, 54, 79, 88, 90, 7, 22, 27):
    print("   r%-3d AV = %s" % (r, ws.cell(r, 48).value))

print("\n### 4. Grand Totals + A/B/V19 row values, five-year sums")
def grab(p):
    w = openpyxl.load_workbook(p, data_only=True); s = w["1.UIO Overview"]
    gt = {y: s.cell(2, COLS[y]).value for y in FY}
    rows = {r: [s.cell(r, COLS[y]).value for y in FY] for r in (33, 54, 79, 90, 25, 26, 29)}
    w.close(); return gt, rows
gA, rA = grab(CAND_A); gB, rB = grab(CAND_B); gV, rV = grab(V19)
print("   Grand Total 5y sum:  V19=%.0f  A=%.0f  B=%.0f   (B vs A %+.2f%% ; B vs V19 %+.2f%%)" % (
    sum(gV.values()), sum(gA.values()), sum(gB.values()),
    100*(sum(gB.values())/sum(gA.values())-1), 100*(sum(gB.values())/sum(gV.values())-1)))
for r in (25, 26, 29, 33, 54, 79, 90):
    print("   r%-3d  V19=%-32s A=%-32s B=%s" % (
        r, str([round(float(x),1) for x in rV[r]]), str([round(float(x),1) for x in rA[r]]),
        str([round(float(x),1) for x in rB[r]])))
print("\n   row 25 2026 locked value in B:", rB[25][0])
wb.close()

print("\n### 5. assumptions log headings")
txt = open(CUR + r"\UIO_V20_ASSUMPTIONS_2026-09-22b.md", encoding="utf-8").read()
for line in txt.splitlines():
    if line.startswith("#"):
        print("   " + line)

print("\n### 6. changelog b: rows and the three fixed rows")
import csv
rows = list(csv.DictReader(open(CUR + r"\UIO_V20_CHANGELOG_2026-09-22b.csv", encoding="utf-8-sig")))
print("   total changed cells:", len(rows))
for r in rows:
    if r['row'] in ('54', '79', '90'):
        print("   r%-3s %-3s %-9s %-14s %-8s -> %-8s  %s" % (
            r['row'], r['cell'], r['country'], r['model_type'], r['old_value_v19'],
            r['new_value_v20'], r['scope_status'][:46]))
