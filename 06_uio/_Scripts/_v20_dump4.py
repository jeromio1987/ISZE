"""V20 — dump Assumptions & Mapping + Cover (READ ONLY)."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

XL = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"
wbv = openpyxl.load_workbook(XL, data_only=True)
for name in ['Assumptions & Mapping', 'Cover']:
    ws = wbv[name]
    print("\n" + "=" * 110)
    print("SHEET: %s  (%sx%s)" % (name, ws.max_row, ws.max_column))
    print("=" * 110)
    for r in range(1, ws.max_row + 1):
        cells = []
        for c in range(1, min(ws.max_column, 12) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or str(v).strip() == '':
                continue
            cells.append("%s=%s" % (get_column_letter(c), str(v).replace('\n', ' ')[:100]))
        if cells:
            print("R%-3d %s" % (r, " | ".join(cells)))
