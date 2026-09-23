"""_b_sources2.py -- provenance sheets + LCV Imit + Greece/Cyprus/Portugal registration hunt."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19 = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
IMIT = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_SourceData\LCV UIO (November 2025)_EU_Study_Imit.xlsx"
IML = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\IML calculation (1).xlsx"

for path, sheets in ((V19, ["Assumptions & Mapping", "Data Sources"]),
                     (IMIT, None),
                     (IML, None)):
    print("#" * 112)
    print("WORKBOOK:", path)
    wb = openpyxl.load_workbook(path, data_only=True)
    print("SHEETS:", wb.sheetnames)
    targets = sheets if sheets else wb.sheetnames
    for sn in targets:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        print("\n---- %s  (%s) ----" % (sn, ws.calculate_dimension()))
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 70),
                                             max_col=min(ws.max_column, 22), values_only=True), 1):
            cells = [("" if c is None else (("%g" % c) if isinstance(c, float) else str(c)))[:34] for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                print("%3d| %s" % (i, " | ".join(cells)))
    wb.close()
