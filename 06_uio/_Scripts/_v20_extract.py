"""V20 — extract V19 cached values to JSON for analysis. READ ONLY on V19."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

XL = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"

# --- values (cached) ---
wbv = openpyxl.load_workbook(XL, data_only=True)
ws = wbv['1.UIO Overview']

# year columns: G=1990 ... AU=2030F
YEARS = list(range(1990, 2031))
COLS = {y: 7 + (y - 1990) for y in YEARS}
assert COLS[2026] == 43, COLS[2026]
assert COLS[2030] == 47

rows = []
for r in range(7, 91):
    rec = {
        'row': r,
        'ctry_code': ws.cell(row=r, column=1).value,
        'country': ws.cell(row=r, column=2).value,
        'cv_lcv': ws.cell(row=r, column=3).value,
        'model_type': ws.cell(row=r, column=4).value,
        'asc_db': ws.cell(row=r, column=5).value,
        'distributor': ws.cell(row=r, column=6).value,
    }
    rec['vals'] = {str(y): ws.cell(row=r, column=COLS[y]).value for y in YEARS}
    rows.append(rec)

grand = {str(y): ws.cell(row=2, column=COLS[y]).value for y in YEARS}
midi  = {str(y): ws.cell(row=5, column=COLS[y]).value for y in YEARS}
growth= {str(y): ws.cell(row=3, column=COLS[y]).value for y in YEARS}
b4 = ws['B4'].value

out = {'grand_total_row2': grand, 'midi_row5': midi, 'growth_row3': growth,
       'B4': b4, 'rows': rows}
with open(r'C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_v20_v19_dump.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, indent=1, default=str)

print("B4 =", b4)
print("\nV19 GRAND TOTAL (row 2) 2025-2030:")
for y in range(2025, 2031):
    print("  %d: %s" % (y, grand[str(y)]))
print("\nV19 MIDI contribution (row 5) 2025-2030:")
for y in range(2025, 2031):
    print("  %d: %s" % (y, midi[str(y)]))
print("\nRow-by-row 2025 / 2026 / 2030 (V19 cached):")
print("%-4s %-10s %-6s %-10s %-18s %9s %9s %9s" % ("R","Country","CV","Model","Distrib","2025E","2026F","2030F"))
for rec in rows:
    v = rec['vals']
    print("%-4d %-10s %-6s %-10s %-18s %9s %9s %9s" % (
        rec['row'], str(rec['country'])[:10], str(rec['cv_lcv'])[:6],
        str(rec['model_type'])[:10], str(rec['distributor'])[:18],
        v['2025'], v['2026'], v['2030']))
