"""_c_cache.py - one-pass cache of the UIO workbooks into JSON (read-only)."""
import io, sys, json, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r'C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT'
FILES = {
    'v19': CUR + r'\ISUZU_UIO_Overview_2026.xlsx',
    'a':   CUR + r'\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx',
    'b':   CUR + r'\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22b.xlsx',
}
OUT = r'C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_c_cache.json'
YEARS = list(range(1990, 2031))
COLS = {y: 7 + (y - 1990) for y in YEARS}

cache = {}
for tag, p in FILES.items():
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['1.UIO Overview']
    meta = {}
    vals = {}
    for r in range(1, 95):
        meta[r] = [ws.cell(r, c).value for c in range(1, 7)]
        vals[r] = {y: ws.cell(r, COLS[y]).value for y in YEARS}
    gt = {y: ws.cell(2, COLS[y]).value for y in YEARS}
    sheet_names = list(wb.sheetnames)
    extra = {}
    for sn in ('V20 REGS USED', 'V20 ASSUMPTIONS', 'V20 CHANGE LOG', 'V20 PLAN ONLY',
               'Assumptions & Mapping', '3. UIO Calculation'):
        if sn in sheet_names:
            s = wb[sn]
            grid = []
            for row in s.iter_rows(min_row=1, max_row=min(s.max_row, 120),
                                   max_col=min(s.max_column, 16), values_only=True):
                grid.append(list(row))
            extra[sn] = grid
    cache[tag] = {'meta': {str(k): v for k, v in meta.items()},
                  'vals': {str(k): {str(y): v for y, v in d.items()} for k, d in vals.items()},
                  'gt': {str(y): v for y, v in gt.items()},
                  'sheets': sheet_names, 'extra': extra, 'b4': ws['B4'].value}
    wb.close()
    print('cached', tag, len(sheet_names), 'sheets')

with open(OUT, 'w', encoding='utf-8') as fh:
    json.dump(cache, fh, default=str)
print('wrote', OUT, os.path.getsize(OUT), 'bytes')
