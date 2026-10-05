"""_b_recalc.py -- recalculate the cached Grand Total cells in candidate b via Excel, then close."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl, win32com.client as win32

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND_B = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22b.xlsx"
FY = [2026, 2027, 2028, 2029, 2030]
COLS = {y: 7 + (y - 1990) for y in FY}

wb = openpyxl.load_workbook(CAND_B, data_only=True)
ws = wb["1.UIO Overview"]
print("BEFORE (cached):", {y: ws.cell(2, COLS[y]).value for y in FY})
print("sheets:", wb.sheetnames)
wb.close()

excel = win32.DispatchEx("Excel.Application")
excel.Visible = False; excel.DisplayAlerts = False; excel.ScreenUpdating = False
wbk = None
try:
    wbk = excel.Workbooks.Open(CAND_B, UpdateLinks=0, ReadOnly=False)
    excel.CalculateFullRebuild()
    wbk.Save()
    sh = wbk.Worksheets("1.UIO Overview")
    print("AFTER (recalculated):", {y: round(float(sh.Cells(2, COLS[y]).Value), 2) for y in FY})
    for r in (54, 79, 90):
        print("  r%-3d cached cells = %s" % (r, [sh.Cells(r, COLS[y]).Value for y in FY]))
finally:
    if wbk is not None:
        try: wbk.Close(SaveChanges=False)
        except Exception: pass
    try: excel.Quit()
    except Exception: pass

wb = openpyxl.load_workbook(CAND_B, data_only=True)
ws = wb["1.UIO Overview"]
print("\nVERIFY reopened:", {y: ws.cell(2, COLS[y]).value for y in FY})
for r in (33, 54, 79, 90):
    print("  r%-3d %s" % (r, [ws.cell(r, COLS[y]).value for y in FY]))
print("  r25 2026 =", ws.cell(25, COLS[2026]).value)
wb.close()
