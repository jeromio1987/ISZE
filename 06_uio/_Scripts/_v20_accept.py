"""V20 - final acceptance verification against the spec's 'how we'll know it's right' list."""
import io, sys, os, json, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19 = os.path.join(CUR, "ISUZU_UIO_Overview_2026.xlsx")
CAND = os.path.join(CUR, "ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx")
FY = [2026, 2027, 2028, 2029, 2030]
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}

wf = openpyxl.load_workbook(CAND, data_only=False)['1.UIO Overview']
wv = openpyxl.load_workbook(CAND, data_only=True)['1.UIO Overview']
v19 = openpyxl.load_workbook(V19, data_only=True)['1.UIO Overview']

def gv(y): return wv.cell(row=2, column=COLS[y]).value
def vm(y): return wv.cell(row=5, column=COLS[y]).value

print("=" * 100)
print("V20 ACCEPTANCE CHECKLIST")
print("=" * 100)
ok = []

# 1 no jump at 2025/2026 boundary: 2025 values identical to V19
mism = sum(1 for r in range(7, 91) for y in range(1990, 2026)
           if v19.cell(row=r, column=COLS[y]).value != wv.cell(row=r, column=COLS[y]).value)
ok.append(("History 1990-2025 identical to V19", mism == 0, "%d mismatches" % mism))

# 2 IUK LCV locked
v = wv.cell(row=25, column=COLS[2026]).value
ok.append(("IUK LCV CY2026F == 44,926 (locked)", abs(v - 44926) < 0.5, "got %s" % v))
ok.append(("IUK LCV not V19's 49,280", abs(v - 49280) > 1, "V19 was 49280"))
v25 = [wv.cell(row=25, column=COLS[y]).value for y in FY]
ok.append(("IUK LCV declines 2026->2030", v25[-1] < v25[0], str(v25)))

# 3 Israel untouched
isr = [wv.cell(row=r, column=COLS[y]).value for r in (7, 8) for y in FY]
isr19 = [v19.cell(row=r, column=COLS[y]).value for r in (7, 8) for y in FY]
ok.append(("Israel rows 7/8 left as V19 (hook)", isr == isr19, "unchanged"))

# 4 UK CV untouched
uk = [wf.cell(row=r, column=COLS[y]).value for r in (22, 24) for y in FY]
allform = all(isinstance(x, str) and x.startswith('=') for x in uk)
ok.append(("UK CV rows 22/24 still live sheet-4 formulas", allform, "10 formulas"))

# 5 plan-only rows zero and out of GT
md = [wv.cell(row=r, column=COLS[y]).value for r in (27, 52, 80, 89) for y in FY]
ok.append(("Medium-duty rows 27/52/80/89 zeroed", all(x in (0, None) for x in md), "all zero"))

# 6 pickup fleet 2030 < 2025 (falling registrations cannot grow a fleet)
LCVROWS = [9, 12, 15, 18, 21, 25, 29, 35, 44, 54, 60, 63, 66, 69, 75, 81, 84, 90]
s25 = sum(v19.cell(row=r, column=COLS[2025]).value or 0 for r in LCVROWS)
s30 = sum(wv.cell(row=r, column=COLS[2030]).value or 0 for r in LCVROWS)
ok.append(("Pickup fleet 2030 < 2025", s30 < s25, "2025 %.0f -> 2030 %.0f (%.1f%%)" % (s25, s30, (s30/s25-1)*100)))

# 7 Ireland acceptance test vs spec's independent rebuild
ie_ld = wv.cell(row=26, column=COLS[2030]).value
ie_lcv = wv.cell(row=29, column=COLS[2030]).value
ok.append(("Ireland LD 2030 ~ spec's 1,382 (within 1%)", abs(ie_ld/1382-1) < 0.01, "got %s" % ie_ld))
ok.append(("Ireland LCV 2030 ~ spec's 2,527 (within 2%)", abs(ie_lcv/2527-1) < 0.02, "got %s" % ie_lcv))

# 8 overstatement vs the spec's 7.8% -> 37.7%
rows_all = [r for r in range(7, 91)]
def tot(src, y):
    return sum(src.cell(row=r, column=COLS[y]).value or 0 for r in rows_all)
print("\n   overstatement (incl MIDI, V19 vs V20, with plan-only removed from V20):")
for y in FY:
    a = tot(v19, y)
    b = tot(wv, y)
    print("     %d: V19 %10.0f  V20 %10.0f  = %5.1f%% too high in V19" % (y, a, b, (a/b-1)*100))

# 9 Grand Total + growth
print("\n   Grand Total (incl MIDI) 2026-2030 cached:", [round(gv(y)) for y in FY])
print("   Growth Y-o-Y 2025-2030 cached:", ["%.2f%%" % (wv.cell(row=3, column=COLS[y]).value*100) for y in FY])
print("   MIDI contribution 2026-2030 cached:", [round(vm(y)) for y in FY])

# 10 no row grows faster than regs / life
print("\n   rows still growing 2025->2030 (ex-2025-normal for stock rebuild):")
growing = []
for r in range(7, 91):
    a = wv.cell(row=r, column=COLS[2025]).value or 0
    b = wv.cell(row=r, column=COLS[2030]).value or 0
    if a > 0 and b > a * 1.05:
        growing.append((r, str(wv.cell(row=r, column=2).value), str(wv.cell(row=r, column=4).value), int(a), int(b)))
for g in growing:
    print("     R%-3d %-9s %-10s %6d -> %6d  (+%.0f%%)" % (g[0], g[1][:9], g[2][:10], g[3], g[4], (g[4]/g[3]-1)*100))

print("\n" + "=" * 100)
print("RESULT")
print("=" * 100)
for name, passed, note in ok:
    print("  [%s] %-52s %s" % ("PASS" if passed else "FAIL", name, note))

# integrity
print("\n" + "=" * 100)
print("FILE INTEGRITY")
print("=" * 100)
for label, p in [('V19 LIVE (untouched)', V19),
                 ('V18 ARCHIVE (untouched)', r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Archive\ISUZU_UIO_Overview_2026_v18.xlsx"),
                 ('V20 CANDIDATE', CAND)]:
    st = os.stat(p)
    print("  %-26s %9d bytes  %s" % (label, st.st_size, datetime.datetime.fromtimestamp(st.st_mtime)))
print("\n  PRISM share file mtime:", datetime.datetime.fromtimestamp(
    os.stat(r"C:\Users\jerom\Desktop\claude\ISZE\03_analytics\Isuzu\prism\data\uio_ld_n35_share.csv").st_mtime)
    if os.path.exists(r"C:\Users\jerom\Desktop\claude\ISZE\03_analytics\Isuzu\prism\data\uio_ld_n35_share.csv") else "n/a")
print("  V20 candidate sheets:", openpyxl.load_workbook(CAND, read_only=True).sheetnames)
