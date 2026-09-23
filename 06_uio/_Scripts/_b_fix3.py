"""_b_fix3.py -- decisive test: which LCV curve reproduces the IML/legacy history, and where the negatives come from."""
import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

IML = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\IML calculation (1).xlsx"
CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"

LCV_813 = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
LCV_823 = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.10,13:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def conv(regs, c, t, y0=1990):
    return sum((regs.get(v) or 0.0)*S(c, t-v+1) for v in range(y0, t+1))

vol = json.load(open(r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_Scripts\_b_iml_vol.json"))
LCVV = {int(k): {int(y): v for y, v in d.items()} for k, d in vol['lcv_vol'].items()}
CVV  = {int(k): {int(y): v for y, v in d.items()} for k, d in vol['cv_vol'].items()}

# IML 1.UIO level rows for the markets of interest
wb = openpyxl.load_workbook(IML, data_only=True)
ws = wb["1.UIO"]
Y0COL = 6          # 1990
LVL = {}
for r in range(6, ws.max_row + 1):
    c1, nm, cv, md = (ws.cell(r, c).value for c in (1, 2, 3, 4))
    if isinstance(c1, (int, float)) and int(c1) in (217, 233, 230, 206):
        d = {1990 + (c - Y0COL): ws.cell(r, c).value for c in range(Y0COL, Y0COL + 41)}
        LVL[(int(c1), str(cv), str(md))] = {y: v for y, v in d.items() if isinstance(v, (int, float))}
print("IML 1.UIO level rows captured:", list(LVL))

wb2 = openpyxl.load_workbook(CAND, data_only=True)
ws2 = wb2["1.UIO Overview"]
COLS = {y: 7 + (y-1990) for y in range(1990, 2031)}

def cand(r):
    return {y: ws2.cell(r, COLS[y]).value for y in range(1990, 2031)}

print("\n" + "#"*100)
for r, code, mt, lbl in ((54, 217, "LCV", "PORTUGL LCV"), (90, 233, "LCV", "CYPRUS LCV")):
    ser = {y: float(cand(r).get(y) or 0) for y in range(1990, 2031)}
    print("\n### r%d %s (code %d) ###" % (r, lbl, code))
    lvl = LVL.get((code, "LCV", "LCV"), {})
    print("   IML 1.UIO level equals candidate series?",
          "YES" if all(abs((lvl.get(y) or 0) - ser[y]) < 1.5 for y in range(1990, 2014)) else "NO")
    for y in (2005, 2006, 2012, 2014, 2019):
        print("     %d  cand %10.0f  IML level %10s  IML regs %8s" %
              (y, ser[y], (round(lvl[y],1) if lvl.get(y) is not None else "-"), LCVV.get(code, {}).get(y)))
    print("   IML 7.LCV Vol registrations 1990-2013:",
          {y: LCVV.get(code, {}).get(y) for y in range(1990, 2014)})
    for cname, c in (("8.13 (engine)", LCV_813), ("8.23 (workbook sheet 3)", LCV_823)):
        for yn, ysrc in (("IML regs", LCVV.get(code, {})),):
            bad = []
            for y in range(1991, 2026):
                g = conv(ysrc, c, y)
                if ser[y] > 0 and (g - ser[y]) > 0.25 * ser[y]:
                    bad.append((y, round(g - ser[y], 1)))
            print("   curve %-26s vs %s: divergence>25%% years: %s" % (cname, yn, bad[:8]))
    # back-solve implied births under both curves using the candidate's OWN series
    for cname, c in (("8.13 (engine)", LCV_813), ("8.23 (workbook sheet 3)", LCV_823)):
        R, neg = {}, {}
        for t in range(1990, 2026):
            R[t] = ser[t] - sum(R[v]*S(c, t-v+1) for v in range(1990, t))
            if R[t] < 0:
                neg[t] = round(R[t], 1)
        print("   implied-birth negatives under %-26s : %s" % (cname, neg))
wb.close()
wb2.close()

# what curve does the candidate sheet 3 actually declare?
wb = openpyxl.load_workbook(CAND, data_only=True)
ws = wb["3. UIO Calculation"]
print("\n### candidate '3. UIO Calculation' Section A ###")
for r in range(5, 26):
    print("   r%-3d %s" % (r, [ws.cell(r, c).value for c in range(2, 6)]))
wb.close()
