"""Portugal UIO (rows 51 LD, 53 F-Series, 54 pickups) rebuilt from registrations.

Referenced by V20_SPEC.md, section "Review of candidate c - decisions 24 September 2026".

Registrations: Matriculas Isuzu SAP.xlsx tab SUM, as mapped in
  01_data/08_Distributor/Imotors_Portugal_UIO_GPS_PMS_Study.xlsx sheet SAP_Raw_Registrations
  (LD = N35 3.5 t + N-Series 5 t / 7.5 t; F-Series = 11 t / 14 t; pickups = D-Max).
Survival curves: sheet '3. UIO Calculation' section A (CV developed 16-yr; LCV with age 12 = 0,
  the variant V20 uses because of the locked UK LCV arithmetic).
2026-2030 new registrations: the current forecast plan as V20 carries it (sheet 'V20 REGS USED'
  of candidate c, rows 51 / 53 / 54) - Jerome's decision 24 Sep 2026.

UIO(t) = sum over cohorts v <= t of R_v x S(t - v + 1)

Registrations start in 2005, so a year is fully covered only once every pre-2005 cohort has
reached zero survival: trucks from 2019 (S(16) = 0), pickups from 2015 (S(12) = 0). Earlier
years are floors, not full fleets.

Run:  python spec_20260924_pt_from_registrations.py
"""
import openpyxl

ISZE = r"C:\Users\jerom\Desktop\claude\ISZE"
STUDY = ISZE + r"\01_data\08_Distributor\Imotors_Portugal_UIO_GPS_PMS_Study.xlsx"
V19 = ISZE + r"\06_uio\_CURRENT\ISUZU_UIO_Overview_2026.xlsx"
V20C = ISZE + r"\06_uio\_CURRENT\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22c.xlsx"

CV = [1, .99, .98, .97, .96, .95, .9, .8, .7, .55, .4, .3, .2, .1, .05, 0]
LCV = [1, .98, .95, .92, .89, .84, .8, .65, .5, .4, .2, 0]


def S(curve, age):
    return curve[age - 1] if 1 <= age <= len(curve) else 0.0


def load_registrations():
    ws = openpyxl.load_workbook(STUDY, data_only=True, read_only=True)["SAP_Raw_Registrations"]
    reg = {}
    for row in ws.iter_rows(values_only=True):
        if isinstance(row[0], int) and 2005 <= row[0] <= 2025:
            reg[row[0]] = row[1:5]  # LCV_D_Max, N35_Grafter, N75_N50_N62, F_Series_MD
    return reg


def load_plan(r):
    ws = openpyxl.load_workbook(V20C, data_only=True, read_only=True)["V20 REGS USED"]
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == r:
            return [float(x) for x in row[5:10]]  # first block = registration input used
    raise KeyError(r)


def row_values(path, r, years):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["1.UIO Overview"]
    row = next(ws.iter_rows(min_row=r, max_row=r, values_only=True))
    return [row[6 + (y - 1990)] or 0 for y in years]


def uio(R, curve, years):
    return [round(sum(n * S(curve, t - v + 1) for v, n in R.items() if v <= t)) for t in years]


def main():
    reg = load_registrations()
    rows = [("LD", 51, CV, lambda v: v[1] + v[2], 2019),
            ("F-Series", 53, CV, lambda v: v[3], 2019),
            ("Pickups", 54, LCV, lambda v: v[0], 2015)]
    years = list(range(2012, 2031))
    print("year                  ", " ".join(f"{y:>6}" for y in years))
    for name, r, curve, pick, full_from in rows:
        R = {y: float(pick(v)) for y, v in reg.items()}
        plan = load_plan(r)
        for i, y in enumerate(range(2026, 2031)):
            R[y] = plan[i]
        print(f"r{r} {name:9s} regs+plan ", " ".join(f"{x:>6}" for x in uio(R, curve, years)),
              f"  (full coverage from {full_from}; plan 2026-30 {[round(x, 1) for x in plan]})")
        print(f"{'':14s} V19       ", " ".join(f"{x:>6}" for x in row_values(V19, r, years)))
        print(f"{'':14s} V20 c     ", " ".join(f"{x:>6}" for x in row_values(V20C, r, years)))


if __name__ == "__main__":
    main()
