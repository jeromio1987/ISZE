"""_b_scope2.py -- measure the exact fix: full-precision input + scoped material seam."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19W = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
CAND_A = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
YEARS=list(range(1990,2031)); FY=[2026,2027,2028,2029,2030]
CV_DEV={1:1.00,2:0.99,3:0.98,4:0.97,5:0.96,6:0.95,7:0.90,8:0.80,9:0.70,10:0.55,
        11:0.40,12:0.30,13:0.20,14:0.10,15:0.05,16:0.00}
LCV={1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def curve(mt): return LCV if str(mt)=='LCV' else CV_DEV
OUT={7,8}; LEAVE={22,24}; PLAN={27,52,80,89}
UKR={2026:3020,2027:2400,2028:2400,2029:2400,2030:2400}
wb=openpyxl.load_workbook(V19W,data_only=True); ws=wb["1.UIO Overview"]
COLS={y:7+(y-1990) for y in YEARS}
ROWS=[{'row':r,'meta':[ws.cell(r,c).value for c in range(1,7)],
       'vals':{y:ws.cell(r,COLS[y]).value for y in YEARS}} for r in range(7,91)]
wb.close()
wba=openpyxl.load_workbook(CAND_A,data_only=True); wsa=wba['1.UIO Overview']
A={r:[wsa.cell(r,COLS[y]).value for y in FY] for r in range(7,91)}
wba.close()

def bs(ser,c,start):
    R={}
    for t in YEARS:
        if t<start: R[t]=0.0; continue
        R[t]=ser[t]-sum(R[v]*S(c,t-v+1) for v in YEARS if start<=v<t)
    return R
def seam_of(ser,c,mode):
    if mode=='none': return 1990
    R=bs(ser,c,1990)
    for t in range(1991,2026):
        if ser[t]>0 and R[t]<0 and (R[t]<-1.0 if mode=='material' else abs(R[t])>0.25*ser[t]):
            return t
    return 1990
def state(ser,c,seam):
    R=bs(ser,c,seam)
    for v in YEARS:
        if R[v]<0: R[v]=0.0
    reb=sum(R[v]*S(c,2026-v) for v in YEARS if v<=2025)
    res=ser[2025]-reb
    if abs(res)>0.5 and ser[2025]>0:
        oldest=min([v for v in YEARS if S(c,2026-v)>0] or [2025])
        R[oldest]+=res/S(c,2026-oldest)
    return R
def smooth(x):
    n=len(x); sm=[]
    for i in range(n):
        if i==0: sm.append((x[0]+x[1])/2.0)
        elif i==n-1: sm.append((x[n-2]+x[n-1])/2.0)
        else: sm.append((x[i-1]+x[i]+x[i+1])/3.0)
    tx,ts=sum(x),sum(sm); f=(tx/ts) if ts else 1.0
    return [z*f for z in sm]

def run(full_precision, seam_rows):
    """full_precision: keep the recovered + smoothed input as float (no integer round).
       seam_rows: set of rows to which the material seam rule is applied."""
    res={}
    for row in ROWS:
        r=row['row']; mt=row['meta'][3]
        if r in OUT or r in LEAVE: continue
        v={y: float(row['vals'].get(y) or 0) for y in YEARS}
        if all(v[t]==0 for t in FY) and v[2025]==0: continue
        c=curve(mt)
        if r==25:
            u={2025:v[2025]}; out={}
            for t in FY: u[t]=u[t-1]*(1-0.123)+UKR[t]; out[t]=u[t]
            res[r]=[out[t] for t in FY]; continue
        if r in PLAN:
            res[r]=[0.0]*5; continue
        raw={t:(v.get(t) or 0)-0.97*(v.get(t-1) or 0) for t in FY}
        if not full_precision: raw={t: round(raw[t]) for t in FY}
        prof=smooth([raw[t] for t in FY])
        if not full_precision: prof=[float(round(z)) for z in prof]
        R=state(v,c,seam_of(v,c,'material' if r in seam_rows else 'pct'))
        out={}
        for i,t in enumerate(FY):
            stock=sum(R[vg]*S(c,t-vg+1) for vg in range(1990,t))
            R[t]=prof[i]
            out[t]=stock+R[t]*S(c,1)
        res[r]=[out[t] for t in FY]
    return res

def total(res,y):
    return sum(res[r][FY.index(y)] for r in res)

for tag,kw in (("1) candidate a            ", dict(full_precision=False, seam_rows=set())),
               ("2) + full-precision input ", dict(full_precision=True,  seam_rows=set())),
               ("3) + material seam r54/r90", dict(full_precision=True,  seam_rows={54,90}))):
    res=run(**kw)
    moved=[]
    for r in sorted(res):
        a=[float(x or 0) for x in A[r]]
        d=[res[r][i]-a[i] for i in range(5)]
        if any(abs(x)>=0.5 for x in d): moved.append((r,[round(x,1) for x in d]))
    print("\n=== %s ===" % tag)
    print("  rows moving >=0.5u : %s" % (moved if moved else "none"))
    print("  Grand Total: " + "  ".join("%d=%.1f" % (y,total(res,y)) for y in FY))

res=run(full_precision=True, seam_rows={54,90})
print("\n=== final values, four adjudicated rows ===")
for r in (33,54,79,90):
    print("  r%-3d %s" % (r, [round(x,1) for x in res[r]]))
print("\n=== candidate a values, same rows ===")
for r in (33,54,79,90):
    print("  r%-3d %s" % (r, [round(float(x or 0),1) for x in A[r]]))

def ratio(v):
    d=[v[i+1]-v[i] for i in range(4)]
    if len(set(1 if x>0 else (-1 if x<0 else 0) for x in d))>2: return None,d
    rs=[d[i+1]/d[i] for i in range(3) if abs(d[i])>1e-9]
    return (sum(rs)/len(rs) if rs else None), d
print("\n=== harness signature (rounded cells, as the harness reads them) ===")
for r in (33,54,79,90):
    vr=[float(round(x)) for x in res[r]]
    m,d=ratio(vr)
    fr=0.03 if (m is not None and abs(m-0.97)<=0.01) else (0.123 if (m is not None and abs(m-0.877)<=0.01) else (0.0 if (m is not None and abs(m-1.0)<=0.01) else None))
    print("  r%-3d cells=%-40s deltas=%-22s mean=%s C1=%-6s C5=%s" % (
        r,str([int(x) for x in vr]),str(d),("%.4f"%m) if m is not None else "n/a",
        (fr is not None and fr<0.08),(m is not None and abs(m-0.97)<=0.01)))
