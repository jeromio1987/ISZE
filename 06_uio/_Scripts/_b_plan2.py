"""_b_plan2.py -- the workbook's own plan sheets: 2. Vehicle Volume and 3. UIO Calculation."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = CUR + r"\ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx"
wb = openpyxl.load_workbook(CAND, data_only=True)

for sn in ("2. Vehicle Volume", "3. UIO Calculation"):
    ws = wb[sn]
    print("#" * 108)
    print("SHEET %s  %s" % (sn, ws.calculate_dimension()))
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 120),
                                         max_col=min(ws.max_column, 26), values_only=True), 1):
        cells = [("" if c is None else (("%g" % c) if isinstance(c, float) else str(c)))[:20] for c in row]
        while cells and cells[-1] == "":
            cells.pop()
        if cells:
            print("%3d| %s" % (i, " | ".join(cells)))
wb.close()
