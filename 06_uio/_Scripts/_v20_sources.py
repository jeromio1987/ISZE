"""V20 - source inventory dump (READ ONLY on every source)."""
import io, sys, json, glob, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = r"C:\Users\jerom\Desktop\claude"
PLAN = os.path.join(ROOT, r"ISZE\91_legacy\Isuzu\UIO\Script\ISZE CV_FY2025-FY2030_mid-term plan & 12RS.xlsx")
IMIT = os.path.join(ROOT, r"ISZE\06_uio\_SourceData\LCV UIO (November 2025)_EU_Study_Imit.xlsx")
VINSUM = os.path.join(ROOT, r"ISZE\01_data\07_UIO\vin\VIN_register_summary.json")
HELS = glob.glob(os.path.join(ROOT, r"ISZE\01_data\90_INBOX\_isd_std_mail_20260818\attachments\*Fahrzeugbestand*bersicht_D_.xlsx"))

def sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()

def show(path, sheetname, maxr, maxc):
    print("\n" + "=" * 110)
    print("FILE: %s" % os.path.basename(path))
    print("SHEET: %r" % sheetname)
    print("=" * 110)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheetname]
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc), start=1):
        cells = []
        for ci, cell in enumerate(row, start=1):
            v = cell.value
            if v is None or str(v).strip() == '':
                continue
            cells.append("%s=%s" % (get_column_letter(ci), str(v).replace('\n', ' ')[:44]))
        if cells:
            print("R%-3d %s" % (r, " | ".join(cells)))
    wb.close()

if os.path.exists(VINSUM):
    print("### VIN_register_summary.json"); print(open(VINSUM, encoding='utf-8').read())

for label, path in [("PLAN", PLAN), ("IMIT", IMIT)]:
    print("\n### %s sheets: %s" % (label, sheets(path)))

show(PLAN, 'mid-term plan to FY30 (Nov25)', 130, 12)
show(IMIT, 'ISZE', 60, 12)

print("\n### Helwig candidates: %s" % HELS)
for h in HELS:
    print("  sheets:", sheets(h))
    show(h, sheets(h)[0], 60, 12)
