"""V20 - is any row still RISING at 2030 in the candidate? Convergence check."""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = os.path.join(CUR, "ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx")
V19P = os.path.join(CUR, "ISUZU_UIO_Overview_2026.xlsx")
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
FY = [2025, 2026, 2027, 2028, 2029, 2030]
wv = openpyxl.load_workbook(CAND, data_only=True)['1.UIO Overview']
v1 = openpyxl.load_workbook(V19P, data_only=True)['1.UIO Overview']
OUT = {7, 8, 22, 24}
PLAN_ONLY = {27, 52, 80, 89}
print("rows still RISING 2029 -> 2030 in V20 (excluded Israel/UK-CV/plan-only):")
n = 0
for r in range(7, 91):
    if r in OUT or r in PLAN_ONLY: continue
    a = wv.cell(row=r, column=COLS[2029]).value or 0
    b = wv.cell(row=r, column=COLS[2030]).value or 0
    if b > a + 0.5:
        n += 1
        print("  R%-3d %-10s %-10s 2029 %8.0f -> 2030 %8.0f" % (
            r, str(wv.cell(row=r,column=2).value)[:10], str(wv.cell(row=r,column=4).value)[:10], a, b))
print("  count = %d" % n)

print("\nV19 equivalent count (rows still rising 2029->2030):")
n19 = 0
for r in range(7, 91):
    if r in OUT: continue
    a = v1.cell(row=r, column=COLS[2029]).value or 0
    b = v1.cell(row=r, column=COLS[2030]).value or 0
    if b > a + 0.5:
        n19 += 1
print("  count = %d (V19 had the fleet climbing in every LCV row)" % n19)

print("\nEvery in-scope row: 2025 -> 2030 direction")
rise = []
for r in range(7, 91):
    if r in OUT or r in PLAN_ONLY: continue
    a = wv.cell(row=r, column=COLS[2025]).value or 0
    b = wv.cell(row=r, column=COLS[2030]).value or 0
    if a and b > a:
        rise.append((r, str(wv.cell(row=r,column=2).value), str(wv.cell(row=r,column=4).value), a, b))
print("  rows with 2030 > 2025: %d" % len(rise))
for x in rise:
    print("    R%-3d %-10s %-10s %7.0f -> %7.0f" % x)
