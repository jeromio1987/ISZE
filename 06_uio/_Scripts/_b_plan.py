"""_b_plan.py -- ISZE CV FY2025-FY2030 mid-term plan: what registration input does it hold?"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

P = r"C:\Users\jerom\Desktop\claude\ISZE\91_legacy\Isuzu\UIO\Script\ISZE CV_FY2025-FY2030_mid-term plan & 12RS.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)
print("SHEETS:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("\n##### %s  (%s) #####" % (sn, ws.calculate_dimension()))
    n = 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 45),
                                         max_col=min(ws.max_column, 20), values_only=True), 1):
        cells = [("" if c is None else (("%g" % c) if isinstance(c, float) else str(c)))[:22] for c in row]
        while cells and cells[-1] == "":
            cells.pop()
        if cells:
            print("%3d| %s" % (i, " | ".join(cells)))
            n += 1
        if n > 42:
            print("   ... truncated")
            break
wb.close()
