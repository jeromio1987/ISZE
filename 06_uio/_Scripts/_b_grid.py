"""_b_grid.py -- brute force legitimate engine conventions; report harness signatures per row."""
import io, sys, itertools, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
V19W = CUR + r"\ISUZU_UIO_Overview_2026.xlsx"
YEARS = list(range(1990, 2031)); FY = [2026, 2027, 2028, 2029, 2030]
CV_DEV = {1:1.00,2:0.99,3:0.98,4:0.97,5:0.96,6:0.95,7:0.90,8:0.80,9:0.70,10:0.55,
          11:0.40,12:0.30,13:0.20,14:0.10,15:0.05,16:0.00}
LCV = {1:1.00,2:0.98,3:0.95,4:0.92,5:0.89,6:0.84,7:0.80,8:0.65,9:0.50,10:0.40,11:0.20,12:0.00}
def S(c,a): return c.get(a,0.0) if a>=1 else 0.0
def curve(mt): return LCV if str(mt)=='LCV' else CV_DEV
OUT_OF_SCOPE={7,8}; LEAVE_AS_V19={22,24}; PLAN_ONLY_MD={27,52,80,89}
UK_LCV_REGS={2026:3020,2027:2400,2028:2400,2029:2400,2030:2400}
TOL=0.01

wb = openpyxl.load_workbook(V19W, data_only=True); ws = wb["1.UIO Overview"]
COLS = {y: 7+(y-1990) for y in YEARS}
ROWS=[{'row':r,'meta':[ws.cell(r,c).value for c in range(1,7)],
       'vals':{y: ws.cell(r,COLS[y]).value for y in YEARS}} for r in range(7,91)]
wb.close()

def bs(ser,c,start):
    R={}
    for t in YEARS:
        if t<start: R[t]=0.0; continue
        R[t]=ser[t]-sum(R[v]*S(c,t-v+1) for v in YEARS if start<=v<t)
    return R
def seam_of(ser,c,mode):
    if mode=='none': return 1990
    R=bs(ser,c,1990)
    for t in range(1991,2026):
        if ser[t]>0 and R[t]<0 and (abs(R[t])>0.25*ser[t] if mode=='pct' else R[t]<-1.0):
            return t
    return 1990
def state(ser,c,seam):
    R=bs(ser,c,seam)
    for v in YEARS:
        if R[v]<0: R[v]=0.0
    reb=sum(R[v]*S(c,2026-v) for v in YEARS if v<=2025)
    res=ser[2025]-reb
    if abs(res)>0.5 and ser[2025]>0:
        oldest=min([v for v in YEARS if S(c,2026-v)>0] or [2025])
        R[oldest]+=res/S(c,2026-oldest)
    return R
def smooth(x, ends='halve'):
    n=len(x); sm=[]
    for i in range(n):
        if i==0:
            sm.append((x[0]+x[1])/2.0 if ends=='halve' else (2*x[0]+x[1])/3.0)
        elif i==n-1:
            sm.append((x[n-2]+x[n-1])/2.0 if ends=='halve' else (2*x[n-1]+x[n-2])/3.0)
        else:
            sm.append((x[i-1]+x[i]+x[i+1])/3.0)
    tx,ts=sum(x),sum(sm); f=(tx/ts) if ts else 1.0
    return [z*f for z in sm]

def build(reg_mode, out_mode, seam_mode, ends='halve'):
    res={}
    for row in ROWS:
        r=row['row']; mt=row['meta'][3]
        if r in OUT_OF_SCOPE or r in LEAVE_AS_V19: continue
        v={y: float(row['vals'].get(y) or 0) for y in YEARS}
        if all(v[t]==0 for t in FY) and v[2025]==0: continue
        c=curve(mt)
        if r==25:
            u={2025:v[2025]}; out={}
            for t in FY: u[t]=u[t-1]*(1-0.123)+UK_LCV_REGS[t]; out[t]=u[t]
            res[r]={'meta':row['meta'],'used':[UK_LCV_REGS[t] for t in FY],'v20':[out[t] for t in FY]}
            continue
        if r in PLAN_ONLY_MD:
            res[r]={'meta':row['meta'],'used':[0]*5,'v20':[0]*5}; continue
        raw={t: (v.get(t) or 0)-0.97*(v.get(t-1) or 0) for t in FY}
        if reg_mode.startswith('int'): raw={t: round(raw[t]) for t in FY}
        prof=[raw[t] for t in FY]
        if 'smooth' in reg_mode: prof=smooth(prof, ends)
        if reg_mode.endswith('int') or reg_mode=='nosmooth_int': prof=[round(z) for z in prof]
        R=state(v,c,seam_of(v,c,seam_mode))
        out={}
        for i,t in enumerate(FY):
            stock=sum(R[vg]*S(c,t-vg+1) for vg in range(1990,t))
            R[t]=prof[i]
            val=stock+R[t]*S(c,1)
            out[t]= int(round(val)) if out_mode=='int' else val
        res[r]={'meta':row['meta'],'used':prof,'v20':[out[t] for t in FY]}
    return res

def ratio(o):
    d=[o['v20'][i+1]-o['v20'][i] for i in range(4)]
    if len(set(1 if x>0 else (-1 if x<0 else 0) for x in d))>2: return None
    rs=[d[i+1]/d[i] for i in range(3) if abs(d[i])>1e-9]
    return (sum(rs)/len(rs), rs, d) if rs else None

MODES=['int_smooth_int','int_smooth','float_smooth','nosmooth_float','nosmooth_int']
print("%-16s %-8s %-9s | %s" % ("reg_mode","out","seam","r33 / r54 / r79 / r90  (mean ratio; C1=switches-flag, C5=hit)"))
for reg_mode in MODES:
    for out_mode in ('int','float'):
        for seam_mode in ('pct','material'):
            res=build(reg_mode,out_mode,seam_mode)
            cells=[]
            c1=0; c5=0
            for r in (33,54,79,90):
                o=res[r]; rr=ratio(o)
                if rr is None:
                    cells.append("r%d  none       " % r); continue
                m=rr[0]
                if abs(m-0.97)<=TOL: f_rate=0.03
                elif abs(m-0.877)<=TOL: f_rate=0.123
                elif abs(m-1.0)<=TOL: f_rate=0.0
                else: f_rate=None
                sw = (f_rate is not None and f_rate<0.08)
                hit = abs(m-0.97)<=TOL
                c1+= 1 if sw else 0; c5+= 1 if hit else 0
                cells.append("r%d %.4f %s%s" % (r,m,'C1' if sw else '  ','C5' if hit else '  '))
            print("%-16s %-8s %-9s | %s | switches=%d hits=%d" % (reg_mode,out_mode,seam_mode," | ".join(cells),c1,c5))
