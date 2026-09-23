"""V20 - growth ceiling check: no row may grow faster than its registration input / life."""
import io, sys, os, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

CUR = r"C:\Users\jerom\Desktop\claude\ISZE\06_uio\_CURRENT"
CAND = os.path.join(CUR, "ISUZU_UIO_Overview_2026_V20_CANDIDATE_2026-09-22.xlsx")
FY = [2026, 2027, 2028, 2029, 2030]
COLS = {y: 7 + (y - 1990) for y in range(1990, 2031)}
LIFE = {'LCV': 8.13, 'default': 9.85}
PLAN_ONLY = {27, 52, 80, 89}
OUT = {22, 24, 7, 8}

wv = openpyxl.load_workbook(CAND, data_only=True)['1.UIO Overview']
REG = json.load(open(os.path.join(CUR, "UIO_V20_VERIFY_2026-09-22.json"), encoding='utf-8'))['rows']

print("=" * 112)
print("GROWTH CEILING: steady state = avg annual registrations x average life")
print("=" * 112)
print("%-4s %-10s %-10s %8s %8s %9s %10s %9s  %s" % (
    "R", "Country", "Model", "2025", "2030", "avg regs", "ceiling", "2030/ceil", "verdict"))
bad = []
for r in range(7, 91):
    if r in PLAN_ONLY or r in OUT:
        continue
    p = REG.get(str(r))
    if not p or p['v20'] is None:
        continue
    mt = p['model']
    life = LIFE['LCV'] if mt == 'LCV' else LIFE['default']
    a = wv.cell(row=r, column=COLS[2025]).value or 0
    b = wv.cell(row=r, column=COLS[2030]).value or 0
    regs = p['regs_used']
    avg = sum(regs) / len(regs)
    ceil = avg * life
    ratio = (b / ceil) if ceil else 0
    verdict = 'ok'
    if ceil > 0 and ratio > 1.0:
        verdict = '*** EXCEEDS CEILING ***'
        bad.append(r)
    elif verdict == 'ok' and b > a * 1.05:
        verdict = 'ok (below ceiling)'
    if b > a * 1.05 or ratio > 1.0:
        print("%-4d %-10s %-10s %8.0f %8.0f %9.0f %10.0f %8.2f  %s" % (
            r, str(p['country'])[:10], str(mt)[:10], a, b, avg, ceil, ratio, verdict))
print("\nrows exceeding their own registration ceiling: %s" % (bad if bad else 'NONE'))
