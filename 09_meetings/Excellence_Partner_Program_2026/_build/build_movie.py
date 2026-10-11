"""Build the animated explainer (HTML, plays by double-click) and render it to MP4.

Reads the RECALCULATED workbook (cached values) so the movie always shows the model's numbers:
  Excellence_Partner_ROI_Model_2026-10-09.xlsx

Outputs (same folder as the workbook):
  Excellence_Partner_Movie_2026-10-09.html   self-contained animation with play / pause / scrubber
  Excellence_Partner_Movie_2026-10-09.mp4    90 s film, 1280x720, 20 fps

Run from the ISZE root:
  python 09_meetings/Excellence_Partner_Program_2026/_build/build_movie.py            # html + mp4
  python 09_meetings/Excellence_Partner_Program_2026/_build/build_movie.py --html     # html only
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PACK = HERE.parent
XLSX = PACK / "Excellence_Partner_ROI_Model_2026-10-09.xlsx"
HTML = PACK / "Excellence_Partner_Movie_2026-10-09.html"
MP4 = PACK / "Excellence_Partner_Movie_2026-10-09.mp4"
FPS, TOTAL = 20, 90

wb = load_workbook(XLSX, data_only=True)
inp, dist, var, uni, plan = wb["Inputs"], wb["Distributors"], wb["Design_variants"], wb["Universe"], wb["Plan_ROI"]


def v(ws, ref):
    x = ws[ref].value
    assert x is not None, f"{ws.title}!{ref} has no cached value - recalculate the workbook first"
    return x


short = {"Isuzu Truck UK": "ITUK", "Isuzu UK pickup": "IUK", "UTI (Israel trucks)": "UTI", "UMI (Israel light commercial)": "UMI"}
dlist = []
for r in range(5, 24):
    name = v(dist, f"B{r}")
    dlist.append({"name": short.get(name, name), "billed": v(dist, f"C{r}"),
                  "lv": [int(v(dist, f"I{r}")), int(v(dist, f"J{r}")), int(v(dist, f"K{r}"))]})
dlist.sort(key=lambda d: -d["billed"])

names = ["Deck", "Growth only", "Hybrid small", "Hybrid bigger", "Tiered"]
subs = ["flat on all stock orders", "only on stock orders above last year's", "flat + progressive bonus",
        "flat + progressive bonus", "rate on all by own growth"]
R0, C0 = 24, 10
designs = []
for k in range(5):
    designs.append({"name": names[k], "sub": subs[k], "y3": v(var, f"D{C0+k}"),
                    "net": [v(var, f"C{R0+k}"), v(var, f"D{R0+k}"), v(var, f"E{R0+k}")], "be": v(var, f"F{R0+k}")})


def hy(col):
    return {"f1": v(inp, f"{col}41"), "f2": v(inp, f"{col}42"), "l1": v(inp, f"{col}43"), "l2": v(inp, f"{col}44"),
            "r1": [v(inp, f"{col}{r}") for r in (45, 46, 47)], "r2": [v(inp, f"{col}{r}") for r in (48, 49, 50)]}


data = {
    "dist": dlist,
    "parts": [int(v(plan, f"{c}4")) for c in "BCD"],
    "deck": [v(var, f"{c}{C0}") for c in "BCD"],
    "rebate": v(inp, "B14"),
    "atstake": v(inp, "B15"),
    "designs": designs,
    "hy": [hy("B"), hy("C")],
    "deckRate": v(inp, "B4"),
    "tier": {"t": [v(inp, "B33"), v(inp, "B34"), v(inp, "B35")], "r": [v(inp, "C33"), v(inp, "C34"), v(inp, "C35"), v(inp, "C36")]},
    "mu": v(inp, "B12"), "sigma": v(inp, "B31"),
    "univ": {"camp": v(uni, "B4"), "enab": v(uni, "B5"), "eval": v(uni, "B6"), "cash": v(uni, "B8"),
             "be": int(v(uni, "B26")), "net": [v(uni, f"E{r}") for r in (22, 23, 24)],
             "act": [int(v(uni, f"B{r}")) for r in (12, 13, 14)]},
}
tpl = (HERE / "movie_template.html").read_text(encoding="utf-8")
HTML.write_text(tpl.replace("__DATA__", json.dumps(data)), encoding="utf-8")
print(HTML)

if "--html" in sys.argv:
    sys.exit(0)

from playwright.sync_api import sync_playwright

tmp = Path(tempfile.mkdtemp(prefix="movie_frames_"))
n = FPS * TOTAL
with sync_playwright() as p:
    b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
    pg = b.new_page(viewport={"width": 1280, "height": 720})
    pg.goto(f"file://{HTML}?rec=1&t=0")
    pg.wait_for_timeout(300)
    for i in range(n + 1):
        pg.evaluate(f"window.__render({i / FPS})")
        pg.screenshot(path=str(tmp / f"f{i:05d}.png"))
    b.close()
subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-framerate", str(FPS), "-i", str(tmp / "f%05d.png"),
                "-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", "20", str(MP4)], check=True)
shutil.rmtree(tmp, ignore_errors=True)
print(MP4, f"{MP4.stat().st_size/1e6:.1f} MB")
