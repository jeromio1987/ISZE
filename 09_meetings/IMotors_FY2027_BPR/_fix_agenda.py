from pptx import Presentation
from pathlib import Path

p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
slide = prs.slides[1]

repls = [
    ("Availability: fill rate and MOS", "Availability: fill rate, MOS and back orders"),
    ("slides 15–17", "slides 15–19"),
    ("slides 15-17", "slides 15–19"),
    ("slides 18–23", "slides 20–25"),
    ("slides 18-23", "slides 20–25"),
    ("slides 24–25", "slides 26–27"),
    ("slides 24-25", "slides 26–27"),
    ("slides 26–27", "slides 28–29"),
    ("slides 26-27", "slides 28–29"),
    ("slides 28–30", "slides 30–32"),
    ("slides 28-30", "slides 30–32"),
]

# Apply carefully in order that doesn't double-shift. Do title first, then ranges from bottom up.
ordered = [
    ("Availability: fill rate and MOS", "Availability: fill rate, MOS and back orders"),
    ("slides 28–30", "slides 30–32"),
    ("slides 28-30", "slides 30–32"),
    ("slides 26–27", "slides 28–29"),
    ("slides 26-27", "slides 28–29"),
    ("slides 24–25", "slides 26–27"),
    ("slides 24-25", "slides 26–27"),
    ("slides 18–23", "slides 20–25"),
    ("slides 18-23", "slides 20–25"),
    ("slides 15–17", "slides 15–19"),
    ("slides 15-17", "slides 15–19"),
]

for old, new in ordered:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    print("replaced", old, "->", new)

# also storyline observation if it mentions availability only
for sh in slide.shapes:
    if sh.has_text_frame and "then availability" in sh.text_frame.text:
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if "then availability" in run.text and "back order" not in run.text.lower():
                    run.text = run.text.replace(
                        "then availability, fleet and PMS",
                        "then availability and back orders, fleet and PMS",
                    )
                    print("storyline updated")

prs.save(str(p))
# print agenda again
for sh in prs.slides[1].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        t = sh.text_frame.text.strip().replace("\n", " | ")
        if "slide" in t.lower() or "Avail" in t or "Fleet" in t or "Priorit" in t or "Appendix" in t or "Programme" in t:
            print(t)
