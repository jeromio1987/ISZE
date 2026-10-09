# -*- coding: utf-8 -*-
"""Rebuild Sep-closed I-Motors BPR: lock Apr–Sep sales to sources, add replenishment ask + competitive one-pager."""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
PPTX = ROOT / "09_meetings/IMotors_FY2027_BPR/I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"
BAK = ROOT / "09_meetings/IMotors_FY2027_BPR/_bak_pre_rebuild_replenishment_20261006.pptx"
CMD = ROOT / "09_meetings/IMotors_FY2027_BPR/OPEN_I_Motors_BPR_2026-10-06_Sep_closed.cmd"
MARKS = ROOT / "09_meetings/IMotors_FY2027_BPR/_competitive_price_overlap_20261006/shortlist_bo_open_marks.csv"
SHORT = ROOT / "09_meetings/IMotors_FY2027_BPR/_competitive_price_overlap_20261006/shortlist_dropping_and_overpriced_vs_am.csv"

# Sold-to 519 Apr–Aug (sales_fact, CR/RE out) + Sep Billing MTD from Daily Sales Report 30 Sep 16:51
SALES_K = [140.1, 151.4, 153.2, 147.1, 109.4, 134.5]  # sum 835.7 → hero €836K
# Apr–Aug order receipt from prior closed pack chart; Sep = Daily Sales ORDER RECEIPT
ORDERS_K = [134.0, 159.0, 147.0, 153.0, 101.0, 154.7]  # sum 848.7 → hero €849K
# BP row IMOTORS from Daily Sales Report
BUDGET_K = [138.2, 149.5, 151.3, 145.3, 162.5, 177.7]  # sum 924.5 → €924K

YTD_SALES_K = round(sum(SALES_K))  # 836
YTD_ORDERS_K = round(sum(ORDERS_K))  # 849
YTD_BUDGET_K = round(sum(BUDGET_K))  # 924
SALES_GAP_K = YTD_BUDGET_K - YTD_SALES_K  # 88
ORDERS_GAP_K = YTD_BUDGET_K - YTD_ORDERS_K  # 75
SALES_GAP_PCT = (sum(SALES_K) / sum(BUDGET_K) - 1) * 100  # ~-9.6
ORDERS_GAP_PCT = (sum(ORDERS_K) / sum(BUDGET_K) - 1) * 100
FY_TARGET_K = 1858.0
REMAIN_K = FY_TARGET_K - sum(SALES_K)
PACE_FULL = sum(SALES_K) * 12 / 6
PACE_PCT = PACE_FULL / FY_TARGET_K * 100

BG, PANEL, BORDER = "262626", "303030", "6E6E6E"
WHITE, MUTED, GRID = "F2F2F2", "BDBDBD", "4D4D4D"
RED, AMBER = "E30613", "F5A000"
FONT = "Aptos"


def rgb(h):
    return RGBColor.from_string(h)


def txt(slide, x, y, w, h, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def hline(slide, x, y, w, color, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)


def panel(slide, points, x=9.86, y=1.46, w=3.25, h=5.58):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.28, y + 0.26, w - 0.5, 0.2, "OBSERVATIONS / QUESTIONS", 7.5, RED, True)
    yy = y + 0.72
    step = (h - 0.95) / max(len(points), 1)
    for i, pnt in enumerate(points):
        rect(slide, x + 0.28, yy + 0.07, 0.09, 0.09, RED if i == 0 else AMBER, shape=MSO_SHAPE.OVAL)
        txt(slide, x + 0.52, yy, w - 0.8, step - 0.08, pnt, 8.4, WHITE)
        yy += step


def replace_text(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    n += 1
        if old in sh.text_frame.text and n == 0:
            full = sh.text_frame.text.replace(old, new)
            first = True
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if first:
                        run.text = full
                        first = False
                        n += 1
                    else:
                        run.text = ""
    return n


def replace_anywhere(prs, old, new):
    return sum(replace_text(s, old, new) for s in prs.slides)


def set_chart(shape, categories, series_list):
    cd = CategoryChartData()
    cd.categories = list(categories)
    for name, vals in series_list:
        cd.add_series(name, list(vals))
    shape.chart.replace_data(cd)


def blank_layout(prs):
    for i in range(len(prs.slide_layouts)):
        layout = prs.slide_layouts[i]
        if "Blank" in (layout.name or "") or i == 6:
            return layout
    return prs.slide_layouts[0]


def base_slide(prs, section, title, sub, source, page_no):
    s = prs.slides.add_slide(blank_layout(prs))
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BG)
    bg.line.fill.background()
    spTree = s.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)

    txt(s, 0.25, 0.10, 3.8, 0.18, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 6.6, WHITE, True)
    txt(s, 0.25, 0.43, 7.5, 0.22, section.upper(), 7.5, RED, True)
    txt(s, 0.25, 0.68, 9.45, 0.66, title, 18 if len(title) > 70 else 20, WHITE, True)
    txt(s, 0.25, 1.42, 9.45, 0.24, sub, 8.5, MUTED)
    hline(s, 0.25, 7.18, 12.85, "575757")
    txt(s, 0.25, 7.24, 11.5, 0.16, source, 6.5, "8C8C8C")
    txt(s, 12.55, 7.24, 0.55, 0.16, str(page_no), 7, MUTED, align=PP_ALIGN.RIGHT)
    return s


def load_competitive_rows():
    by_pn = {}
    with SHORT.open(encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            by_pn[r["pn"]] = r
    marks = {}
    with MARKS.open(encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            marks[r["pn"]] = r
    open_now = []
    earlier = []
    for pn, m in marks.items():
        row = by_pn.get(pn, {})
        name = (row.get("description") or pn).replace(";", ",")
        open_eur = float(m.get("open_eur") or 0)
        if m.get("open_on_our_order") == "True":
            open_now.append((name, pn, open_eur))
        elif m.get("earlier_bo_2026") == "True":
            earlier.append((name, pn))
    open_now.sort(key=lambda x: -x[2])
    # top shortlist by impact from file order
    top = []
    with SHORT.open(encoding="utf-8", newline="") as f:
        for i, r in enumerate(csv.DictReader(f)):
            if i >= 6:
                break
            lost = float(r.get("std_list_lost") or 0)
            top.append(((r.get("description") or r["pn"]).replace(";", ","), r["pn"], lost))
    return top, open_now, earlier, len(marks), sum(1 for m in marks.values() if m.get("open_on_our_order") == "True" or m.get("earlier_bo_2026") == "True")


# ---------------------------------------------------------------------------
shutil.copy2(PPTX, BAK)
try:
    prs = Presentation(str(PPTX))
except PermissionError as e:
    alt = PPTX.with_name(PPTX.stem + "_REBUILD.pptx")
    print("LOCKED", e, "→ will write", alt)
    shutil.copy2(BAK, alt)
    PPTX = alt  # type: ignore
    prs = Presentation(str(PPTX))

print("YTD sales", sum(SALES_K), "->", YTD_SALES_K, "orders", sum(ORDERS_K), "->", YTD_ORDERS_K, "budget", sum(BUDGET_K))

# ===== SLIDE 4 — monthly Sales / Orders / Budget =====
s4 = prs.slides[3]
for sh in s4.shapes:
    if not sh.has_chart:
        continue
    from lxml import etree
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        sers.append(nm)
    out = []
    for nm in sers:
        low = nm.lower()
        if "budget" in low or "current" in low or "plan" in low:
            out.append((nm, BUDGET_K))
        elif "order" in low:
            out.append((nm, ORDERS_K))
        else:
            out.append((nm, SALES_K))
    set_chart(sh, ["Apr", "May", "Jun", "Jul", "Aug", "Sep"], out)
    print("S4", out)

aug_sep = SALES_K[4] + SALES_K[5]
bud_as = BUDGET_K[4] + BUDGET_K[5]
replace_text(
    s4,
    "YTD sales are €835K against €924K current budget: €89K below plan (-9.6%).",
    f"YTD sales are €{YTD_SALES_K}K against €{YTD_BUDGET_K}K current budget: €{SALES_GAP_K}K below plan ({SALES_GAP_PCT:.1f}%).",
)
replace_text(
    s4,
    "Orders are €846K YTD, €78K below the same budget phasing (-8.4%); the order pipeline is closer but still does not close the sales gap.",
    f"Orders are €{YTD_ORDERS_K}K YTD, €{ORDERS_GAP_K}K below the same budget phasing ({ORDERS_GAP_PCT:.1f}%); the order pipeline is closer but still does not close the sales gap.",
)
replace_text(
    s4,
    f"Aug–Sep sales: €{aug_sep:.0f}K vs €{bud_as:.0f}K budget ({(aug_sep/bud_as-1)*100:.1f}%)",
    f"Aug–Sep sales: €{aug_sep:.0f}K vs €{bud_as:.0f}K budget ({(aug_sep/bud_as-1)*100:.1f}%)",
)
# catch prior Aug-Sep wording variants
for old in [
    "Aug–Sep sales: €243K vs €340K budget (-28.5%)",
    "Aug–Sep sales: €244K vs €340K budget (-28.3%)",
]:
    replace_text(s4, old, f"Aug–Sep sales: €{aug_sep:.0f}K vs €{bud_as:.0f}K budget ({(aug_sep/bud_as-1)*100:.1f}%)")

replace_text(
    s4,
    "ISZE Daily Sales Report 30 Sep closed (Billing MTD / order receipt); Apr–Aug from ISZE billing series.",
    "Apr–Aug ISZE billing to I-Motors (sold-to 519); September Billing MTD / order receipt from ISZE Daily Sales Report 30 Sep closed.",
)

# ===== SLIDE 3 — hero tiles =====
s3 = prs.slides[2]
replace_anywhere(prs, "€835K", f"€{YTD_SALES_K}K")
replace_anywhere(prs, "€846K", f"€{YTD_ORDERS_K}K")
# gap lines on hero (may already be updated by above if embedded)
replace_text(
    s3,
    f"Apr–Sep sales vs €924K current budget | = €89K below plan (-9.6%)",
    f"Apr–Sep sales vs €{YTD_BUDGET_K}K current budget | = €{SALES_GAP_K}K below plan ({SALES_GAP_PCT:.1f}%)",
)
replace_text(
    s3,
    f"Apr–Sep orders vs €924K current budget | = €78K below plan (-8.4%)",
    f"Apr–Sep orders vs €{YTD_BUDGET_K}K current budget | = €{ORDERS_GAP_K}K below plan ({ORDERS_GAP_PCT:.1f}%)",
)
# also without pipe variants
replace_text(s3, "= €89K below plan (-9.6%)", f"= €{SALES_GAP_K}K below plan ({SALES_GAP_PCT:.1f}%)")
replace_text(s3, "= €78K below plan (-8.4%)", f"= €{ORDERS_GAP_K}K below plan ({ORDERS_GAP_PCT:.1f}%)")
replace_text(
    s3,
    "The full-year target is still €1.858M; after six months, €1.023M remains to be billed in Oct–Mar.",
    f"The full-year target is still €1.858M; after six months, €{REMAIN_K/1000:.3f}M remains to be billed in Oct–Mar.".replace("€0.", "€"),
)
# remain as €1.022M style
remain_m = REMAIN_K / 1000
replace_text(s3, "€1.023M remains", f"€{remain_m:.3f}M remains")
replace_text(s3, "€1.022M remains", f"€{remain_m:.3f}M remains")

# ===== SLIDE 5 — pace =====
s5 = prs.slides[4]
for sh in s5.shapes:
    if not sh.has_chart:
        continue
    from lxml import etree
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    cats = [n.text for n in root.findall(".//{%s}cat//{%s}pt/{%s}v" % (C, C, C))]
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    n = max(len(v) for _, v in sers)
    base_cats = cats[:n] if len(cats) >= n else ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"][:n]
    sales_eur = [v * 1000 for v in SALES_K]
    out = []
    for nm, vals in sers:
        if "2027" in nm:
            merged = list(vals) + [0.0] * n
            for i, v in enumerate(sales_eur):
                if i < n:
                    merged[i] = v
            for i in range(6, n):
                merged[i] = 0.0
            out.append((nm, merged[:n]))
        else:
            out.append((nm, (vals + [0] * n)[:n]))
    set_chart(sh, base_cats[:n], out)
    print("S5 FY27", sales_eur)

pace_m = PACE_FULL / 1000  # millions from K
short_k = FY_TARGET_K - PACE_FULL
replace_text(s5, "April–September billing totals €835K", f"April–September billing totals €{YTD_SALES_K}K")
replace_text(s5, "April–September billing totals €836K", f"April–September billing totals €{YTD_SALES_K}K")
for old_p in ["90.5%", "90.0%", "90.1%", "89.9%", "90.2%"]:
    replace_text(s5, old_p, f"{PACE_PCT:.1f}%")
replace_text(s5, "€1.67M", f"€{pace_m:.2f}M")
replace_text(s5, "€1.68M", f"€{pace_m:.2f}M")
replace_text(
    s5,
    "A straight line (× 12/6) gives €1.67M, €187K short of the €1.858M target.",
    f"A straight line (× 12/6) gives €{pace_m:.2f}M, €{abs(short_k):.0f}K short of the €1.858M target.",
)
replace_text(
    s5,
    "A straight line (× 12/6) gives €1.68M, €176K short of the €1.858M target.",
    f"A straight line (× 12/6) gives €{pace_m:.2f}M, €{abs(short_k):.0f}K short of the €1.858M target.",
)

# ===== SLIDE 6/7 remain language =====
for si in (5, 6):
    replace_text(prs.slides[si], "€1.023M", f"€{remain_m:.3f}M")
    replace_text(prs.slides[si], "€1.022M", f"€{remain_m:.3f}M")

# ===== SLIDE 9 — confirm rolling StD source language =====
s9 = prs.slides[8]
replace_text(
    s9,
    "Source: KPI Master File (Parts Sales DB to DLR €) Oct 2024–Aug 2026; Sep 2026 from KPI Check Sheet. Apr–Aug 2024 N/D in master — not used.",
    "Source: KPI Master File (Parts Sales DB to DLR €) Oct 2024–Aug 2026; Sep 2026 from KPI Check Sheet (€416.8k). Apr–Aug 2024 N/D in master — not used.",
)

# ===== NEW SLIDES: replenishment ask + competitive one-pager =====
top, open_now, earlier, n_marks, n_either = load_competitive_rows()

s_rep = base_slide(
    prs,
    "Availability  ·  discussion",
    "How does I-Motors replenish stock from ISZE?",
    "Meeting questions — we need your process in the room, not a lecture from our side",
    "Context: Sep dealer fill 81.0%, MOS 2.5, B/O 495 lines (KPI Check Sheet). BO part-number talk stays on the August book.",
    19,
)

questions = [
    ("Who triggers replenishment?", "Purchasing, warehouse, dealer demand, or a system rule?"),
    ("What signal do you use?", "MOS, fill rate, back orders, dealer orders, or a min/max list?"),
    ("How often do you reorder from ISZE?", "Daily, weekly, monthly — and does that change by part family?"),
    ("How do you clear back orders?", "Who prioritises aged lines, and when does a B/O become a lost sale at the dealer?"),
    ("Dealer fill vs own MOS", "When dealer fill falls (81% in Sep), do you buy more from ISZE even if MOS is still below 3.0?"),
]
for i, (q, hint) in enumerate(questions):
    y = 1.75 + i * 0.95
    rect(s_rep, 0.35, y, 9.25, 0.85, PANEL, BORDER)
    txt(s_rep, 0.55, y + 0.12, 8.8, 0.28, f"{i+1}.  {q}", 14, WHITE, True)
    txt(s_rep, 0.55, y + 0.45, 8.8, 0.30, hint, 10, MUTED)

panel(s_rep, [
    "We see dealer demand up and ISZE billing behind plan — that only works if stock is being used down. We need the replenishment rule that sits behind that.",
    "Please walk us through one recent week: what made you place (or not place) an order on ISZE.",
    "If MOS and B/O both move the wrong way, which KPI wins the buy decision?",
])

# Competitive / supply one-pager
s_comp = base_slide(
    prs,
    "Availability  ·  price and supply",
    "Competitive parts that are dropping at dealers and look expensive vs aftermarket brands",
    "26 part numbers on the shortlist; 9 also marked for open order or earlier 2026 back order — BO detail still August book",
    "Sources: Portugal sales to dealer through Aug 2026; aftermarket brand prices in the price ledger; open orders sold-to 519 (6 Oct export). Reference dealer price = list with channel %GM, not a Portugal shop quote.",
    20,
)

rect(s_comp, 0.35, 1.75, 3.00, 1.55, PANEL, BORDER)
txt(s_comp, 0.50, 1.90, 2.7, 0.22, "SHORTLIST", 7.5, RED, True)
txt(s_comp, 0.50, 2.20, 2.7, 0.45, "26 PNs", 28, WHITE, True)
txt(s_comp, 0.50, 2.75, 2.7, 0.35, "Dropping sales to dealer list € and above aftermarket brand median", 8.5, MUTED)

rect(s_comp, 3.50, 1.75, 3.00, 1.55, PANEL, BORDER)
txt(s_comp, 3.65, 1.90, 2.7, 0.22, "SUPPLY MARKS", 7.5, RED, True)
txt(s_comp, 3.65, 2.20, 2.7, 0.45, f"{n_either} of 26", 28, AMBER, True)
txt(s_comp, 3.65, 2.75, 2.7, 0.35, f"{len(open_now)} still open on our order · {len(earlier)} earlier BO 2026 only", 8.5, MUTED)

rect(s_comp, 6.65, 1.75, 3.00, 1.55, PANEL, BORDER)
txt(s_comp, 6.80, 1.90, 2.7, 0.22, "LIST € LOST (WORKING)", 7.5, RED, True)
txt(s_comp, 6.80, 2.20, 2.7, 0.45, "€13.8k", 28, WHITE, True)
txt(s_comp, 6.80, 2.75, 2.7, 0.35, "Combined list € lost on the 26 hits (FY / CY working cuts)", 8.5, MUTED)

txt(s_comp, 0.35, 3.50, 9.3, 0.25, "Examples (name · list € lost on the working cut)", 9, MUTED, True)
y = 3.80
for name, pn, lost in top[:5]:
    txt(s_comp, 0.45, y, 9.1, 0.28, f"•  {name}  ·  €{lost:,.0f} list € lost  ·  {pn}", 10, WHITE)
    y += 0.32

txt(s_comp, 0.35, 5.50, 9.3, 0.25, "Still open on our sales orders to I-Motors (sold-to 519)", 9, MUTED, True)
y = 5.80
for name, pn, open_eur in open_now[:4]:
    txt(s_comp, 0.45, y, 9.1, 0.28, f"•  {name}  ·  open €{open_eur:,.0f}  ·  {pn}", 10, AMBER if open_eur else WHITE)
    y += 0.30

panel(s_comp, [
    "This is a discussion shortlist, not a price change list. Aftermarket brand prices are mostly DE/NO/UK listings — not Portuguese shop prices.",
    "Ask: for the four open-order lines, is the delay ISZE supply, I-Motors order timing, or dealer wait?",
    "Ask: which of the top dropping body / filter / brake lines should we price-check with you first?",
    "IDS GPS for Portugal stays locked at 63.3% — separate from this competitive cut.",
])

# Move two new slides to after Lisbon (current index 17 = slide 18)
# New slides were appended at end — move both to position 18 (0-based insert index 18)
sldIdLst = prs.part._element.find(qn("p:sldIdLst"))
ids = list(sldIdLst)
# last two are competitive then replenishment? add_slide order: rep then comp → last two = rep, comp
# We want: ... Lisbon(17), rep, comp, Aug BO...
# Currently appended: [..., rep, comp] at end
new_els = ids[-2:]  # [rep, comp]
for el in new_els:
    sldIdLst.remove(el)
# insert at position 18 (after Lisbon which is index 17)
for offset, el in enumerate(new_els):
    sldIdLst.insert(18 + offset, el)

# Agenda renumber (+2) — temp tokens so later replaces do not cascade
s2 = prs.slides[1]
agenda_fixes = [
    ("slides 15–20", "@@A15"),
    ("slides 21–26", "@@A21"),
    ("slides 27–28", "@@A27"),
    ("slides 29–30", "@@A29"),
    ("slides 31–33", "@@A31"),
]
for old, new in agenda_fixes:
    replace_text(s2, old, new)
for old, new in [
    ("@@A15", "slides 15–22"),
    ("@@A21", "slides 23–28"),
    ("@@A27", "slides 29–30"),
    ("@@A29", "slides 31–32"),
    ("@@A31", "slides 33–35"),
]:
    replace_text(s2, old, new)

try:
    prs.save(str(PPTX))
except PermissionError:
    alt = PPTX.with_name(PPTX.stem + "_REBUILD.pptx")
    prs.save(str(alt))
    PPTX = alt
    print("SAVED ALT", PPTX)

# Renumber footers
prs2 = Presentation(str(PPTX))
for i, slide in enumerate(prs2.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and sh.left > Inches(12.0) and sh.top > Inches(7.0):
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.text = str(i)
                    break

# Point MOS/BO question toward replenishment slide
s17 = prs2.slides[16]
replace_text(
    s17,
    "Question: after the Lisbon strip, which aged August back-order euros should clear first while MOS rebuilds?",
    "Question: after the Lisbon strip and the replenishment discussion, which aged August back-order euros should clear first while MOS rebuilds?",
)

prs2.save(str(PPTX))
CMD.write_bytes(
    b'@echo off\r\nstart "" "%~dp0I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"\r\n'
)

print("SAVED", PPTX, "slides", len(prs2.slides), "size", PPTX.stat().st_size)
for i in range(2, 23):
    slide = prs2.slides[i]
    texts = [sh.text_frame.text.strip().replace("\n", " ")[:80]
             for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = texts[2] if len(texts) > 2 else (texts[1] if len(texts) > 1 else "?")
    print(f"{i+1:02d}", title)
