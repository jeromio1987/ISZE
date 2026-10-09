# -*- coding: utf-8 -*-
"""Build self-contained HTML for competitive vs AM brand price shortlist."""
from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from html import escape
from pathlib import Path

BASE = Path(__file__).resolve().parent
ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
OUT = BASE / "IMotors_Competitive_AM_Price_Shortlist_20261006.html"
OPEN_CMD = BASE / "OPEN_IMotors_Competitive_AM_Price_Shortlist.cmd"
MARKS_CSV = BASE / "shortlist_bo_open_marks.csv"

BO_XLSX = ROOT / "01_data/90_INBOX/_nuno_backordered_lines_20261006/20260929_1531_Nuno_Backordered lines 2026.xlsx"
ORDER_XLSX = ROOT / "01_data/90_INBOX/_imotors_bo_compare_20261006/EXPORT_ISZE_Order_20261006.xlsx"


def load(name: str) -> list[dict]:
    with (BASE / name).open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def norm_pn(pn) -> str:
    if pn is None:
        return ""
    s = str(pn).strip()
    s = re.sub(r"\.0$", "", s)
    s = s.lstrip("0") or "0"
    if s.lower() in ("nan", "none", ""):
        return ""
    return s


def fnum(x, d=0) -> str:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "—"
    if d == 0:
        return f"{v:,.0f}"
    return f"{v:,.{d}f}"


def feur(x, d=0) -> str:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "—"
    if d == 0:
        return f"€{v:,.0f}"
    return f"€{v:,.{d}f}"


def fpct(x) -> str:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "—"
    sign = "+" if v > 0 else ""
    return f"{sign}{v:.1f}%"


def idx(x) -> str:
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "—"
    return f"{v:.2f}×"


def brands_display(am_brands: str | None, am_sources: str | None = None) -> str:
    """Named brands when present; KS International when listing has no brand label."""
    s = (am_brands or "").strip()
    seen: set[str] = set()
    out: list[str] = []
    for b in [x.strip() for x in s.split(",") if x.strip()]:
        # skip empty / nan / bare loader codes that leaked into brand field
        if b.lower() in ("nan", "none", "ks", "ks_am", "ks_gen"):
            continue
        k = b.lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(b)
    if out:
        return ", ".join(out[:6]) + ("…" if len(out) > 6 else "")
    sources = (am_sources or "").lower()
    if "ks" in sources:
        return "KS International (no brand on listing)"
    return "—"


def load_supply_marks(short_pns: set[str]) -> dict[str, dict]:
    """Open orders sold-to 519 + any Nuno BO month in 2026 (same sources as BO AIM dive)."""
    import openpyxl

    open_eur: dict[str, float] = defaultdict(float)
    open_qty: dict[str, float] = defaultdict(float)
    open_lines: dict[str, int] = defaultdict(int)

    wb = openpyxl.load_workbook(ORDER_XLSX, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    hi = {str(h): i for i, h in enumerate(hdr)}
    for r in rows:
        if str(r[hi["Sold-to pt"]]) != "519":
            continue
        pn = norm_pn(r[hi["Material"]] or r[hi["Order Material"]])
        if not pn or pn not in short_pns:
            continue
        open_lines[pn] += 1
        try:
            open_eur[pn] += float(r[hi["EUR Amount"]] or 0)
        except (TypeError, ValueError):
            pass
        try:
            open_qty[pn] += float(r[hi["Order Quantity"]] or 0)
        except (TypeError, ValueError):
            pass
    wb.close()

    bo_months: dict[str, set[str]] = defaultdict(set)
    wb = openpyxl.load_workbook(BO_XLSX, data_only=True, read_only=True)
    for sn in wb.sheetnames:
        m = re.search(r"BO (\d{8})", sn)
        if not m:
            continue
        asof = m.group(1)
        ym = f"{asof[:4]}-{asof[4:6]}"
        ws = wb[sn]
        sheet_rows = list(ws.iter_rows(values_only=True))
        if not sheet_rows:
            continue
        hdr = [str(h).strip().upper() if h else "" for h in sheet_rows[0]]
        idx_map = {h: i for i, h in enumerate(hdr)}
        i_ref = idx_map.get("REFERENCIA") or idx_map.get("REF") or idx_map.get("PN")
        if i_ref is None:
            continue
        for r in sheet_rows[1:]:
            if not r:
                continue
            pn = norm_pn(r[i_ref])
            if pn and pn in short_pns:
                bo_months[pn].add(ym)
    wb.close()

    marks: dict[str, dict] = {}
    for pn in short_pns:
        is_open = pn in open_lines
        months = sorted(bo_months.get(pn, []))
        earlier_bo = bool(months) and not is_open
        # still show BO months when also open (for note), but label priority is open
        labels: list[str] = []
        if is_open:
            labels.append("Open on our order")
        if months and not is_open:
            labels.append("Earlier BO 2026")
        elif months and is_open:
            # open now; also had BO earlier — keep both facts readable
            labels.append("Earlier BO 2026")
        marks[pn] = {
            "pn": pn,
            "open_on_our_order": is_open,
            "earlier_bo_2026": bool(months) and not is_open,
            "had_bo_any_2026": bool(months),
            "bo_months": ";".join(months),
            "open_eur": round(open_eur.get(pn, 0.0), 2),
            "open_qty": open_qty.get(pn, 0.0),
            "open_lines": open_lines.get(pn, 0),
            "labels": labels,
        }
    return marks


def mark_tags_html(labels: list[str]) -> str:
    if not labels:
        return '<span class="sub">—</span>'
    bits = []
    for lab in labels:
        cls = "tag supply open" if lab.startswith("Open") else "tag supply bo"
        bits.append(f'<span class="{cls}">{escape(lab)}</span>')
    return '<div class="marks">' + "".join(bits) + "</div>"


def main() -> None:
    summary = json.loads((BASE / "summary.json").read_text(encoding="utf-8"))
    short = load("shortlist_dropping_and_overpriced_vs_am.csv")
    nm1 = load("near_miss_dropping_no_am_price.csv")[:15]
    nm2 = load("near_miss_overpriced_but_std_growing.csv")[:12]
    basket = load("basket_std_period_compare.csv")
    cov = summary["coverage"]
    cov_pct = cov["competitive_am_coverage_share_cy2025_list"] * 100

    short_pns = {norm_pn(r["pn"]) for r in short}
    marks = load_supply_marks(short_pns)

    # Persist marks CSV (no invented euros — open EUR from EXPORT only)
    with MARKS_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "pn",
                "open_on_our_order",
                "earlier_bo_2026",
                "had_bo_any_2026",
                "bo_months",
                "open_eur",
                "open_qty",
                "open_lines",
                "labels",
            ],
        )
        w.writeheader()
        for pn in sorted(short_pns):
            m = marks[pn]
            w.writerow(
                {
                    **{k: m[k] for k in w.fieldnames if k != "labels"},
                    "labels": " | ".join(m["labels"]),
                }
            )

    n_open = sum(1 for m in marks.values() if m["open_on_our_order"])
    n_bo_only = sum(1 for m in marks.values() if m["earlier_bo_2026"])
    n_either = sum(1 for m in marks.values() if m["labels"])
    n_neither = len(short) - n_either

    # Sanity vs agreed prior cut (fail loud if sources drift)
    expected_open = {"1476012190", "8982488951", "8983313830", "8981226920"}
    expected_bo_only = {
        "8974068144",
        "8980873901",
        "8983894341",
        "8983313840",
        "8982120690",
    }
    got_open = {pn for pn, m in marks.items() if m["open_on_our_order"]}
    got_bo_only = {pn for pn, m in marks.items() if m["earlier_bo_2026"]}
    if got_open != expected_open or got_bo_only != expected_bo_only:
        raise SystemExit(
            f"Supply-mark set drifted.\n"
            f" open got={sorted(got_open)} expected={sorted(expected_open)}\n"
            f" bo_only got={sorted(got_bo_only)} expected={sorted(expected_bo_only)}"
        )

    comp = next(b for b in basket if "Competitive" in b["basket"])
    fy26 = float(comp["std_list_eur_FY26_AprAug"])
    fy27 = float(comp["std_list_eur_FY27_AprAug"])
    cy25 = float(comp["std_list_eur_CY2025"])
    cy26 = float(comp["std_list_eur_CY2026_JanAug"])
    max_b = max(fy26, fy27, cy25, cy26 * 12 / 8)

    def bar_pct(v: float) -> float:
        return max(2.0, min(100.0, 100.0 * float(v) / max_b))

    impacts = [float(r["impact_score"] or 0) for r in short[:12]]
    max_imp = max(impacts) if impacts else 1.0

    short_rows = []
    for i, r in enumerate(short, 1):
        pn = norm_pn(r["pn"])
        m = marks[pn]
        short_rows.append(
            f"""<tr class="{'supply-hit' if m['labels'] else ''}">
<td class="num">{i}</td>
<td><strong>{escape(r['description'])}</strong><div class="sub mono">{escape(r['pn'])}</div>{mark_tags_html(m['labels'])}</td>
<td>{escape(r['aim_l1'])}<div class="sub">{escape(r.get('aim_l2') or '')}</div></td>
<td class="num">{feur(r['std_list_FY26_AprAug'])} → {feur(r['std_list_FY27_AprAug'])}<div class="sub">{fpct(r.get('list_chg_fy_pct'))} FY Apr–Aug</div></td>
<td class="num">{feur(r['std_list_lost'])}</td>
<td class="num">{feur(r['ref_dealer_eur'], 2)}</td>
<td class="num">{feur(r['am_price_median_eur'], 2)}</td>
<td class="num"><span class="tag bad">{idx(r['index_ref_vs_am_median'])}</span></td>
<td class="brands">{escape(brands_display(r.get('am_brands'), r.get('am_sources')))}</td>
</tr>"""
        )

    nm1_rows = []
    for r in nm1:
        nm1_rows.append(
            f"""<tr>
<td><strong>{escape(r['description'])}</strong><div class="sub mono">{escape(r['pn'])}</div></td>
<td>{escape(r['aim_l1'])}<div class="sub">{escape(r.get('aim_l2') or '')}</div></td>
<td class="num">{feur(r['std_list_lost'])}</td>
<td class="num">{fpct(r.get('list_chg_fy_pct'))}</td>
<td class="num">{fpct(r.get('list_chg_cy_runrate_pct'))}</td>
<td><span class="tag warn">No AM brand price</span></td>
</tr>"""
        )

    nm2_rows = []
    for r in nm2:
        nm2_rows.append(
            f"""<tr>
<td><strong>{escape(r['description'])}</strong><div class="sub mono">{escape(r['pn'])}</div></td>
<td>{escape(r['aim_l1'])}<div class="sub">{escape(r.get('aim_l2') or '')}</div></td>
<td class="num">{feur(r['ref_dealer_eur'], 2)}</td>
<td class="num">{feur(r['am_price_median_eur'], 2)}</td>
<td class="num"><span class="tag bad">{idx(r['index_ref_vs_am_median'])}</span></td>
<td class="num ok">{fpct(r.get('list_chg_fy_pct'))}</td>
<td class="brands">{escape(brands_display(r.get('am_brands'), r.get('am_sources')))}</td>
</tr>"""
        )

    impact_bars = []
    for r in short[:10]:
        sc = float(r["impact_score"] or 0)
        w = 100 * sc / max_imp
        pn = norm_pn(r["pn"])
        m = marks[pn]
        mark_note = (" · " + " / ".join(m["labels"])) if m["labels"] else ""
        impact_bars.append(
            f"""<div class="bar-row">
<div class="bar-lab">{escape(r['description'][:28])}<div class="sub mono">{escape(r['pn'])}{escape(mark_note)}</div></div>
<div class="bar-track"><div class="bar-fill" style="width:{w:.1f}%"></div></div>
<div class="bar-val">{feur(r['std_list_lost'])} lost · {idx(r['index_ref_vs_am_median'])}</div>
</div>"""
        )

    basket_bars = []
    for b in basket:
        a = float(b["std_list_eur_FY26_AprAug"])
        c = float(b["std_list_eur_FY27_AprAug"])
        label = (
            b["basket"]
            .replace(" (catalogue L1)", "")
            .replace(" (competitive OR contested)", "")
        )
        basket_bars.append(
            f"""<div class="bar-row"><div class="bar-lab">{escape(label)}<div class="sub">FY26</div></div><div class="bar-track"><div class="bar-fill navy" style="width:{bar_pct(a):.1f}%"></div></div><div class="bar-val">{feur(a)}</div></div>
<div class="bar-row"><div class="bar-lab"><div class="sub">FY27 · {fpct(b['list_FY27_vs_FY26_AprAug_pct'])}</div></div><div class="bar-track"><div class="bar-fill ok" style="width:{bar_pct(c):.1f}%"></div></div><div class="bar-val">{feur(c)}</div></div>"""
        )

    # Compact tables for embed (no fetch needed — also SSR above)
    short_embed = []
    for r in short:
        pn = norm_pn(r["pn"])
        m = marks[pn]
        row = dict(r)
        row["brands_display"] = brands_display(r.get("am_brands"), r.get("am_sources"))
        row["supply_labels"] = m["labels"]
        row["open_on_our_order"] = m["open_on_our_order"]
        row["earlier_bo_2026"] = m["earlier_bo_2026"]
        row["bo_months"] = m["bo_months"]
        row["open_eur"] = m["open_eur"]
        short_embed.append(row)

    embed = {
        "as_of": "2026-10-06",
        "distributor": "I-Motors Portugal",
        "coverage_share_cy2025": cov["competitive_am_coverage_share_cy2025_list"],
        "shortlist_hit_pns": cov["shortlist_hit_pns"],
        "shortlist_std_list_lost_sum": cov["shortlist_std_list_lost_sum"],
        "supply_marks": {
            "open_now_pns": n_open,
            "earlier_bo_only_pns": n_bo_only,
            "either_pns": n_either,
            "neither_pns": n_neither,
            "sources": {
                "open_orders": str(ORDER_XLSX),
                "nuno_bo": str(BO_XLSX),
                "sold_to": "519",
            },
        },
        "basket": basket,
        "shortlist": short_embed,
        "near_miss_dropping_no_am_top15": nm1,
        "near_miss_overpriced_growing_top12": nm2,
        "gaps": summary.get("gaps", []),
    }

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>I-Motors Portugal — competitive parts vs aftermarket brand prices</title>
<style>
:root {{
  --ink:#1a2332; --muted:#5c6b7a; --paper:#f6f3ee; --card:#fff;
  --navy:#15233c; --terra:#c0552d; --line:#d9d2c5; --ok:#2f6b4f; --bad:#9b2c2c;
  --serif:Georgia,"Iowan Old Style","Palatino Linotype",Palatino,serif;
  --sans:"Segoe UI",system-ui,sans-serif;
}}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--paper); color:var(--ink); font:15px/1.55 var(--sans); }}
.wrap {{ max-width:1240px; margin:0 auto; padding:28px 22px 80px; }}
header.hero {{
  background:linear-gradient(145deg,#15233c 0%,#243b5c 55%,#3a2a22 100%);
  color:#f7f4ee; border-radius:18px; padding:34px; margin-bottom:22px;
  box-shadow:0 18px 40px rgba(21,35,60,.22);
}}
.kicker {{ letter-spacing:.12em; text-transform:uppercase; font-size:11px; opacity:.75; }}
h1 {{ font-family:var(--serif); font-weight:600; font-size:clamp(1.45rem,3vw,2.05rem); margin:10px 0 12px; line-height:1.18; }}
.lede {{ font-size:1.05rem; max-width:72ch; opacity:.93; }}
.meta {{ display:flex; flex-wrap:wrap; gap:10px; margin-top:18px; }}
.chip {{ background:rgba(247,244,238,.1); border:1px solid rgba(247,244,238,.22); padding:6px 10px; border-radius:999px; font-size:12px; }}
.grid {{ display:grid; gap:14px; }}
.g4 {{ grid-template-columns:repeat(4,minmax(0,1fr)); }}
.g5 {{ grid-template-columns:repeat(5,minmax(0,1fr)); }}
@media (max-width:1100px) {{ .g4, .g5 {{ grid-template-columns:1fr 1fr; }} }}
@media (max-width:640px) {{ .g4, .g5 {{ grid-template-columns:1fr; }} }}
.kpi {{ background:var(--card); border:1px solid var(--line); border-radius:14px; padding:16px; }}
.kpi .lbl {{ font-size:11px; text-transform:uppercase; letter-spacing:.08em; color:var(--muted); }}
.kpi .val {{ font-family:var(--serif); font-size:1.65rem; margin-top:4px; color:var(--navy); }}
.kpi .note {{ font-size:12px; color:var(--muted); margin-top:6px; }}
.kpi.warn {{ border-color:#e2c089; background:#fff9ef; }}
.kpi.ok {{ border-color:#b7d3c3; background:#f3faf6; }}
.kpi.bad {{ border-color:#e2b0b0; background:#fff5f5; }}
section {{ background:var(--card); border:1px solid var(--line); border-radius:16px; padding:22px; margin:16px 0; }}
section h2 {{ font-family:var(--serif); font-size:1.35rem; margin:0 0 8px; color:var(--navy); }}
section h3 {{ font-size:1.02rem; margin:16px 0 8px; color:var(--navy); }}
.prose {{ max-width:78ch; }}
.prose p {{ margin:0 0 12px; }}
.callout {{ border-left:4px solid var(--terra); background:#faf6f2; padding:12px 14px; border-radius:0 10px 10px 0; margin:14px 0; }}
.callout.navy {{ border-left-color:var(--navy); background:#f3f6fa; }}
.callout.warn {{ border-left-color:#c9a227; background:#fff9ef; }}
.bar-row {{ display:grid; grid-template-columns:160px 1fr 170px; gap:10px; align-items:center; margin:7px 0; }}
@media (max-width:700px) {{ .bar-row {{ grid-template-columns:1fr; gap:2px; }} }}
.bar-lab {{ font-size:12.5px; }}
.bar-track {{ background:#ece7de; height:12px; border-radius:999px; overflow:hidden; }}
.bar-fill {{ height:100%; border-radius:999px; background:linear-gradient(90deg,var(--navy),var(--terra)); }}
.bar-fill.ok {{ background:var(--ok); }}
.bar-fill.navy {{ background:var(--navy); }}
.bar-val {{ font-variant-numeric:tabular-nums; font-size:12px; text-align:right; color:var(--muted); }}
table {{ width:100%; border-collapse:collapse; font-size:12.5px; }}
th, td {{ padding:8px 7px; border-bottom:1px solid var(--line); text-align:left; vertical-align:top; }}
th {{ font-size:10.5px; text-transform:uppercase; letter-spacing:.06em; color:#f7f4ee; background:var(--navy); position:sticky; top:0; z-index:1; }}
tbody tr:nth-child(even) {{ background:#faf8f4; }}
tbody tr:hover {{ background:#f3eee6; }}
tbody tr.supply-hit {{ background:#fff8f2; }}
tbody tr.supply-hit:nth-child(even) {{ background:#fff3ea; }}
.num {{ font-variant-numeric:tabular-nums; text-align:right; white-space:nowrap; }}
.mono {{ font-family:Consolas,"Cascadia Mono",monospace; font-size:11.5px; }}
.sub {{ font-size:11px; color:var(--muted); }}
.toc {{ display:flex; flex-wrap:wrap; gap:8px; margin:0 0 8px; }}
.toc a {{ color:var(--navy); text-decoration:none; background:#eef2f7; border:1px solid #d5deea; padding:6px 10px; border-radius:999px; font-size:12px; }}
.footer {{ color:var(--muted); font-size:12px; margin-top:28px; max-width:85ch; }}
.tag {{ display:inline-block; font-size:11px; padding:2px 7px; border-radius:6px; background:#eef2f7; color:var(--navy); }}
.tag.bad {{ background:#f8e0e0; color:#7a2222; }}
.tag.warn {{ background:#f7ecd0; color:#6a5210; }}
.tag.supply {{ margin:3px 4px 0 0; font-weight:600; }}
.tag.supply.open {{ background:#f3d9c8; color:#7a3418; }}
.tag.supply.bo {{ background:#dde8f5; color:#1e3a5f; }}
.marks {{ margin-top:4px; }}
.brands {{ font-size:11.5px; color:var(--muted); max-width:240px; }}
.scroll {{ overflow:auto; }}
.two {{ display:grid; grid-template-columns:1.05fr .95fr; gap:18px; }}
@media (max-width:960px) {{ .two {{ grid-template-columns:1fr; }} }}
.footnote {{ font-size:11.5px; color:var(--muted); margin-top:10px; }}
.ok {{ color:var(--ok); }}
ul {{ margin:8px 0 0 18px; padding:0; }}
ul li {{ margin:6px 0; }}
</style>
</head>
<body>
<div class="wrap">
<header class="hero">
  <div class="kicker">I-Motors · Portugal · competitive parts · aftermarket brand prices · 6 Oct 2026</div>
  <h1>Competitive sales to dealer is rising overall — the interesting hits are parts that are dropping and still expensive vs other aftermarket brands</h1>
  <p class="lede">Primary lens: Blue Print, Bosch, Delphi, Febi, Herth+Buss and similar brand prices in the price ledger. OEM peer / grey import is secondary and does not gate this shortlist. Aftermarket brand prices cover only about {cov_pct:.0f}% of Competitive-parts sales-to-dealer list € in CY2025 — treat hits as evidence where the join exists. Shortlist rows that are still open on our sales orders to I-Motors (sold-to 519), or that showed on Nuno’s BO extract earlier in 2026, are marked in the table.</p>
  <div class="meta">
    <span class="chip">AM brand cover: {cov_pct:.1f}% of CY2025 competitive StD list €</span>
    <span class="chip">Shortlist hits: {len(short)} PNs · {feur(cov['shortlist_std_list_lost_sum'])} list € lost</span>
    <span class="chip">Supply marks: {n_either} of {len(short)} (open {n_open} · earlier BO {n_bo_only})</span>
    <span class="chip">Competitive StD FY Apr–Aug: {fpct(comp['list_FY27_vs_FY26_AprAug_pct'])}</span>
    <span class="chip">Sales to dealer through 2026-08</span>
  </div>
</header>

<nav class="toc">
  <a href="#verdict">Verdict</a>
  <a href="#direction">Basket direction</a>
  <a href="#shortlist">Shortlist</a>
  <a href="#near-miss">Near-misses</a>
  <a href="#limits">Limits</a>
</nav>

<section id="verdict">
  <h2>1. Verdict</h2>
  <div class="prose">
    <p>At basket level, Competitive parts sales to dealer (list €) is <strong>up</strong>, not down: FY Apr–Aug {feur(fy26)} → {feur(fy27)} ({fpct(comp['list_FY27_vs_FY26_AprAug_pct'])}). The contested AIM families move the same way. So the commercial story is not “competitive is collapsing” — it is a short list of part numbers that are both losing sales-to-dealer list € and sitting well above the aftermarket brand median on our reference dealer price.</p>
    <p>Rule for a hit: sales to dealer list € dropping (FY Apr–Aug −15%+ or CY run-rate −20%+) <em>and</em> our reference dealer price more than 10% above the aftermarket brand median on that part (or its genuine/BVP pair). Ranked by impact = list € lost × overprice gap ratio. Combined list € lost on the {len(short)} hits: <strong>{feur(cov['shortlist_std_list_lost_sum'])}</strong>.</p>
    <p>Of those {len(short)} hits, <strong>{n_either}</strong> also sit on supply stress we can name from our books: <strong>{n_open}</strong> still open on our sales orders to I-Motors (sold-to 519, EXPORT 6 Oct), and <strong>{n_bo_only}</strong> only on Nuno’s BO extract earlier in 2026 (Jan–Aug tabs; Sep+ empty in the 29-Sep file). The other <strong>{n_neither}</strong> have neither mark — price pressure without a current open-order / earlier-BO flag.</p>
  </div>
  <div class="callout warn">
    <strong>Coverage caveat:</strong> only {cov['competitive_with_am_price_pns']} Competitive part numbers have an aftermarket brand price ({cov_pct:.1f}% of CY2025 competitive sales-to-dealer list € — {feur(cov['competitive_with_am_price_list_eur_cy2025'])} of {feur(cov['competitive_cy2025_list_eur'])}). Most observed brand prices are DE / NO / UK listings, not Portuguese workshop purchase prices.
  </div>
</section>

<div class="grid g5" style="margin:16px 0">
  <div class="kpi ok"><div class="lbl">Competitive StD FY Apr–Aug</div><div class="val">{fpct(comp['list_FY27_vs_FY26_AprAug_pct'])}</div><div class="note">{feur(fy26)} → {feur(fy27)} list €</div></div>
  <div class="kpi warn"><div class="lbl">AM brand coverage</div><div class="val">{cov_pct:.0f}%</div><div class="note">of CY2025 competitive StD list €</div></div>
  <div class="kpi bad"><div class="lbl">Shortlist hits</div><div class="val">{len(short)}</div><div class="note">dropping + expensive vs AM brands</div></div>
  <div class="kpi"><div class="lbl">List € lost on hits</div><div class="val">{feur(cov['shortlist_std_list_lost_sum'])}</div><div class="note">working cut · impact-ranked</div></div>
  <div class="kpi warn"><div class="lbl">Open / earlier BO</div><div class="val">{n_open} / {n_bo_only}</div><div class="note">{n_either} of {len(short)} shortlist PNs marked</div></div>
</div>

<section id="direction">
  <h2>2. Basket direction — sales to dealer list €</h2>
  <p class="prose">List € = qty × ISZE list (same basis as the AIM evolution pack). Pieces are not the KPI here.</p>
  <div class="two" style="margin-top:12px">
    <div>
      <h3>Competitive parts (catalogue L1)</h3>
      <div class="bar-row"><div class="bar-lab">CY2025</div><div class="bar-track"><div class="bar-fill navy" style="width:{bar_pct(cy25):.1f}%"></div></div><div class="bar-val">{feur(cy25)}</div></div>
      <div class="bar-row"><div class="bar-lab">2026 Jan–Aug</div><div class="bar-track"><div class="bar-fill navy" style="width:{bar_pct(cy26):.1f}%"></div></div><div class="bar-val">{feur(cy26)} <span class="sub">(8 mo)</span></div></div>
      <div class="bar-row"><div class="bar-lab">FY26 Apr–Aug</div><div class="bar-track"><div class="bar-fill" style="width:{bar_pct(fy26):.1f}%"></div></div><div class="bar-val">{feur(fy26)}</div></div>
      <div class="bar-row"><div class="bar-lab">FY27 Apr–Aug</div><div class="bar-track"><div class="bar-fill ok" style="width:{bar_pct(fy27):.1f}%"></div></div><div class="bar-val">{feur(fy27)} · {fpct(comp['list_FY27_vs_FY26_AprAug_pct'])}</div></div>
      <p class="footnote">Annualised 2026 Jan–Aug vs CY2025: {fpct(comp['list_runrate_2026_vs_CY2025_pct'])} on list €.</p>
    </div>
    <div>
      <h3>All three baskets (FY Apr–Aug)</h3>
      {''.join(basket_bars)}
    </div>
  </div>
</section>

<section id="shortlist">
  <h2>3. End-goal shortlist — dropping and expensive vs aftermarket brands</h2>
  <p class="prose">Reference dealer price = ISZE list × 1.40 × 1.275 (standard channel dials). Index = reference dealer price ÷ aftermarket brand median. Index above 1.10 means we look expensive on this working signal. Actual Portugal %GM may differ.</p>
  <div class="callout navy">
    <strong>Supply marks (Lisbon-readable):</strong>
    <span class="tag supply open">Open on our order</span> = still open on our sales orders to I-Motors (sold-to 519, EXPORT 6 Oct 2026).
    <span class="tag supply bo">Earlier BO 2026</span> = appeared on Nuno’s Backordered-lines extract in at least one populated 2026 month (Jan–Aug; Sep+ tabs empty in the 29-Sep file).
    Count for the room: <strong>{n_open}</strong> open now, <strong>{n_bo_only}</strong> earlier BO only, <strong>{n_neither}</strong> neither. Two open PNs also had BO earlier (wheel cylinder <span class="mono">1476012190</span>, front brake disc <span class="mono">8982488951</span>) — both tags show.
  </div>
  <h3>Top 10 by impact</h3>
  {''.join(impact_bars)}
  <h3 style="margin-top:22px">Full shortlist ({len(short)} part numbers)</h3>
  <div class="scroll">
  <table>
    <thead><tr>
      <th>#</th><th>Part · supply mark</th><th>Family</th><th>StD list € (FY Apr–Aug)</th><th>List € lost</th>
      <th>Our ref. dealer €</th><th>AM brand median €</th><th>Index</th><th>Brands (sample)</th>
    </tr></thead>
    <tbody>
{''.join(short_rows)}
    </tbody>
  </table>
  </div>
  <p class="footnote">AM brand prices are Eucon / KS observations (mostly DE/NO/UK). Where the Brands cell shows <em>KS International (no brand on listing)</em>, the price ledger has a KS shop observation with no competitor brand filled on the row. OEM peer median is available in the source CSV as a secondary column only. Open-order € in the embedded payload is from EXPORT only — Nuno BO euros are not reprinted here.</p>
</section>

<section id="near-miss">
  <h2>4. Near-misses</h2>
  <div class="two">
    <div>
      <h3>Dropping, but no aftermarket brand price</h3>
      <p class="prose">Priority to get a brand quote — these lose sales-to-dealer list € but we cannot price-check them yet. Showing top 15 of {cov['near_miss_dropping_no_am_pns']}.</p>
      <div class="scroll">
      <table>
        <thead><tr><th>Part</th><th>Family</th><th>List € lost</th><th>FY chg</th><th>CY run-rate</th><th>Status</th></tr></thead>
        <tbody>
{''.join(nm1_rows)}
        </tbody>
      </table>
      </div>
    </div>
    <div>
      <h3>Expensive vs AM brands, but StD growing</h3>
      <p class="prose">Price pressure signal, not volume loss yet. Showing top 12 of {cov['near_miss_overpriced_growing_pns']}.</p>
      <div class="scroll">
      <table>
        <thead><tr><th>Part</th><th>Family</th><th>Ref. dealer €</th><th>AM median €</th><th>Index</th><th>FY StD</th><th>Brands</th></tr></thead>
        <tbody>
{''.join(nm2_rows)}
        </tbody>
      </table>
      </div>
    </div>
  </div>
</section>

<section id="limits">
  <h2>5. What this does not prove</h2>
  <ul>
    <li>That Portuguese workshops pay the Eucon/KS listing — those are foreign aftermarket observations.</li>
    <li>That list € lost is all lost to aftermarket brands (stock, campaigns, and mix also move the line).</li>
    <li>That every Competitive part is overpriced — most have no brand price on file.</li>
    <li>OEM peer / Amayama grey as a substitute for aftermarket brand competition (kept in the full table as a secondary column only).</li>
    <li>That an “Earlier BO 2026” mark equals a current fill-rate miss — Nuno’s Sep/Oct+ BO tabs were empty in the 29-Sep file; open-order marks use our 6 Oct EXPORT only.</li>
  </ul>
  <div class="callout navy">
    <strong>Method note:</strong> Competitive = catalogue L1 “Competitive parts” (AIM remaps into Brakes / Suspension / …). Contested AIM L1 basket = Accident &amp; Body, Brakes, Clutch, Maintenance Service, Suspension. Sales to dealer from Portugal MONTH files through 2026-08 (no Sep PN grain). List from PR00 IE40. Supply marks join shortlist PNs to EXPORT sold-to 519 and Nuno Backordered lines 2026.
  </div>
</section>

<p class="footer">Source folder computed 2026-10-06 · shortlist_dropping_and_overpriced_vs_am.csv · shortlist_bo_open_marks.csv · near_miss_*.csv · basket_std_period_compare.csv · summary.json · Self-contained page (no fetch). Reference dealer price uses ×1.40 × 1.275 (= ×1.785) — spoken as reference dealer price / channel %GM, not as the working label.</p>
</div>
<script type="application/json" id="payload">{json.dumps(embed, ensure_ascii=False)}</script>
</body>
</html>
"""

    low = html.lower()
    for bad in ("sell-out", "sellout", "street", "mode a", "mode b", "wallet"):
        if bad in low:
            raise SystemExit(f"Banned token in HTML: {bad}")

    OUT.write_text(html, encoding="utf-8")
    # CRLF launcher
    OPEN_CMD.write_bytes(
        b"@echo off\r\n"
        b'start "" "%~dp0IMotors_Competitive_AM_Price_Shortlist_20261006.html"\r\n'
    )
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"wrote {OPEN_CMD}")
    print(f"wrote {MARKS_CSV}")
    print(
        f"shortlist={len(short)} open={n_open} earlier_bo={n_bo_only} "
        f"either={n_either} neither={n_neither} nm1={len(nm1)} nm2={len(nm2)}"
    )


if __name__ == "__main__":
    main()
