from pptx import Presentation
from pathlib import Path
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
print("slide_width", prs.slide_width, "height", prs.slide_height, "n", len(prs.slides))
for idx in [15, 16, 25, 26]:
    slide = prs.slides[idx]
    print(f"\n=== SLIDE {idx+1} shapes={len(slide.shapes)} ===")
    for sh in slide.shapes:
        name = sh.name
        try:
            l,t,w,h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            l=t=w=h=0
        kind = "chart" if getattr(sh,"has_chart",False) and sh.has_chart else ("table" if getattr(sh,"has_table",False) and sh.has_table else "text" if getattr(sh,"has_text_frame",False) and sh.has_text_frame else "other")
        preview = ""
        if getattr(sh,"has_text_frame",False) and sh.has_text_frame:
            preview = sh.text_frame.text.strip().replace("\n"," | ")[:110]
        print(f"  {kind:5} {name[:42]:42} ({l/914400:.2f},{t/914400:.2f},{w/914400:.2f}x{h/914400:.2f}) {preview}")

# sample colors from slide 17 first text
slide = prs.slides[16]
for sh in slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                try:
                    f = r.font
                    print("FONT", repr(f.name), f.size.pt if f.size else None, f.bold, f.color.rgb if f.color and f.color.rgb else None)
                except Exception as e:
                    print("font err", e)
                break
            break
        break
