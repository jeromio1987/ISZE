from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
print("slide_width", prs.slide_width, "height", prs.slide_height)
print("layouts", len(prs.slide_layouts), [l.name for l in prs.slide_layouts[:8]])
# slide 17 (index 16) and 16, 26, 27
for idx in [15, 16, 25, 26]:
    slide = prs.slides[idx]
    print(f"\n=== SLIDE {idx+1} shapes={len(slide.shapes)} ===")
    for sh in slide.shapes:
        name = sh.name
        try:
            l,t,w,h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            l=t=w=h=0
        kind = "chart" if sh.has_chart else ("table" if sh.has_table else "text" if sh.has_text_frame else "other")
        preview = ""
        if sh.has_text_frame:
            preview = sh.text_frame.text.strip().replace("\n"," | ")[:100]
        print(f"  {kind:5} {name[:40]:40} ({l/914400:.2f},{t/914400:.2f},{w/914400:.2f}x{h/914400:.2f}) {preview}")
