from pptx import Presentation
from pathlib import Path
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
# agenda slide 2
print("AGENDA texts:")
for sh in prs.slides[1].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        t = sh.text_frame.text.strip()
        if len(t) < 200:
            print(" -", t.replace("\n"," | "))
# fix source on 18/19
for idx in [17, 18]:
    for sh in prs.slides[idx].shapes:
        if sh.has_text_frame and "AIM catalogue" in sh.text_frame.text:
            print("SRC", idx+1, repr(sh.text_frame.text[:120]))
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.text = run.text.replace("catalogue out", "catalogue cut").replace("PN-out", "PN-cut")
prs.save(str(p))
print("saved")
