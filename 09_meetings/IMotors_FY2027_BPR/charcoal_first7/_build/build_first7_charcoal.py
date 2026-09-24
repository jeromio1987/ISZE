"""Slides 1-7 of the 685 KB base deck, rebuilt in the charcoal layout, plus the PMS service-demand cover slide.

Content source : ISZE/I_Motors_FY2027_Business_Performance_Review_DRAFT.pptx, slides 1-7 (text and figures kept as written)
PMS slide      : ISZE/I_Motors_FY2027_Business_Performance_Review_CHARCOAL_v8.pptx, slide 10 (table kept; observations
                 updated with the verified stock-build finding, see _imotors_bpr_calcs_2026-09-24)
Style          : charcoal draft (bg 262626, panel 303030/6E6E6E, red E30613, amber F5A000, Aptos)
Charts are native PowerPoint charts, so they can be edited.
"""
import pathlib, shutil
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE.parent / "I_Motors_FY2027_BPR_first7_plus_PMS_charcoal_2026-09-24.pptx"
LOGO = HERE / "isuzu_logo.png"

BG, PANEL, BORDER, ROW2 = "262626", "303030", "6E6E6E", "2B2B2B"
WHITE, MUTED, GRID = "F2F2F2", "BDBDBD", "4D4D4D"
RED, AMBER, GREY, GREEN = "E30613", "F5A000", "B7B7B7", "31A354"
FONT = "Aptos"
def rgb(h): return RGBColor.from_string(h)

prs = globals().get("PRS") or Presentation()   # PRS is injected when the merged deck build runs this file
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def txt(slide, x, y, w, h, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
    return tb

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def hline(slide, x, y, w, color, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)

PAGE = [1]
def base(section, title, sub, source):
    PAGE[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    txt(s, 0.25, 0.10, 3.8, 0.18, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 6.6, WHITE, True)
    txt(s, 0.25, 0.43, 6.0, 0.22, section.upper(), 7.5, RED, True)
    txt(s, 0.25, 0.68, 9.45, 0.66, title, 20 if len(title) <= 60 else 16.5, WHITE, True)
    txt(s, 0.25, 1.42, 9.45, 0.24, sub, 8.5, MUTED)
    hline(s, 0.25, 7.18, 12.85, "575757")
    txt(s, 0.25, 7.24, 11.5, 0.16, source, 6.5, "8C8C8C")
    txt(s, 12.55, 7.24, 0.55, 0.16, str(PAGE[0]), 7, MUTED, align=PP_ALIGN.RIGHT)
    return s

def panel(slide, points, x=9.86, y=1.46, w=3.25, h=5.58):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.28, y + 0.26, w - 0.5, 0.2, "OBSERVATIONS / QUESTIONS", 7.5, RED, True)
    yy = y + 0.72
    step = (h - 0.95) / max(len(points), 1)
    for i, pnt in enumerate(points):
        rect(slide, x + 0.28, yy + 0.07, 0.09, 0.09, RED if i == 0 else AMBER, shape=MSO_SHAPE.OVAL)
        txt(slide, x + 0.52, yy, w - 0.8, step - 0.08, pnt, 8.6, WHITE)
        yy += step

def tile(slide, x, y, w, h, label, big, sub, color=WHITE, big_size=24):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.18, y + 0.16, w - 0.3, 0.2, label, 7, RED, True)
    txt(slide, x + 0.18, y + 0.42, w - 0.3, 0.5, big, big_size, color, True)
    txt(slide, x + 0.18, y + 0.98, w - 0.3, h - 1.05, sub, 7.8, MUTED)

def style_chart(ch, num_fmt='#,##0', legend=True):
    ch.font.name = FONT; ch.font.size = Pt(8); ch.font.color.rgb = rgb(MUTED)
    ch.has_title = False
    if legend:
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.color.rgb = rgb(MUTED); ch.legend.font.size = Pt(8)
    else:
        ch.has_legend = False
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(GRID); va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.color.rgb = rgb(MUTED); va.tick_labels.font.size = Pt(7.5)
    va.tick_labels.number_format = num_fmt; va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(GRID); ca.tick_labels.font.color.rgb = rgb(MUTED); ca.tick_labels.font.size = Pt(7.5)

def labels(series, fmt, color=WHITE, pos=XL_LABEL_POSITION.OUTSIDE_END, size=7.5):
    dl = series.data_labels
    dl.show_value = True
    dl.number_format = fmt; dl.number_format_is_linked = False
    dl.font.size = Pt(size); dl.font.color.rgb = rgb(color); dl.font.bold = True
    dl.position = pos

def col_fill(series, c):
    series.format.fill.solid(); series.format.fill.fore_color.rgb = rgb(c); series.format.line.fill.background()

def line_fmt(series, c, width=2.25, marker=True, dash=False):
    series.smooth = False
    series.format.line.color.rgb = rgb(c); series.format.line.width = Pt(width)
    if dash: series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if marker:
        series.marker.style = XL_MARKER_STYLE.CIRCLE; series.marker.size = 6
        series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = rgb(c)
        series.marker.format.line.color.rgb = rgb(c)
    else:
        series.marker.style = XL_MARKER_STYLE.NONE

def chart(slide, kind, cd, x, y, w, h):
    return slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart

# ================================================================ 1 cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
rect(s, 0, 0, 0.10, 7.5, RED)
txt(s, 0.45, 0.32, 5, 0.2, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 7, WHITE, True)
txt(s, 0.45, 2.05, 10.5, 1.3, ["I Motors", "FY2027 Business Performance Review"], 30, WHITE, True)
txt(s, 0.45, 3.55, 10.5, 0.3, "FY27 plan  ·  Sales to dealer  ·  Stock  ·  Fill rate & MOS  ·  PMS & SPA", 11, MUTED)
txt(s, 0.45, 4.05, 10.5, 0.3, "Data through September 2026 where available  |  Review draft, 24 September 2026", 9, MUTED)
s.shapes.add_picture(str(LOGO), Inches(0.45), Inches(6.6), height=Inches(0.36))

# ================================================================ 2 executive view
s = base("Executive view", "Dealer demand is up, but I Motors is running its stock down: ISZE billing is behind plan and MOS is below target",
         "Periods differ by KPI: ISZE sales Apr–Sep 2026; sales-to-dealer/MOS through Aug 2026; commercial potential calendar 2025.",
         "Sources: I Motors sales plan; KPI Master File; ISZE derived commercial/fleet files.")
tiles = [
    ("ISZE SALES VS PLAN", "€774K", "Apr–Sep sales vs €924K current budget\n= €150K below plan (-16.2%)", RED),
    ("ORDERS", "€790K", "Apr–Sep orders vs €924K current budget\n= €134K below plan (-14.5%)", RED),
    ("SALES TO DEALER", "€1.864M", "Apr–Aug 2026 vs €1.491M Apr–Aug 2025\n= +25.0% year on year", GREEN),
    ("FILL RATE ISZE → I MOTORS", "87%", "Aug 2026, vs 93% target\nI Motors → dealer: 85.4% vs 85% target", RED),
    ("MOS", "2.2", "Aug 2026 (67 days of stock), vs 3.0 target\n3.7 in Sep 2025", RED),
    ("BACK-ORDER LINES", "282", "Aug 2026, down from 430 in Apr 2026", GREEN),
    ("TARGET-SPEND COVER", "56.0%", "CY2025 parts billed against modelled target spend\n€1.710M / €3.052M", AMBER),
    ("2025 FLEET BASE", "7,208", "Current fleet workbook: 2025 estimate\n2024 actual was 6,903", WHITE),
]
tw, th, gx, gy = 2.2, 1.6, 0.13, 0.15
for k, (lab, big, sub, c) in enumerate(tiles):
    x = 0.4 + (k % 4) * (tw + gx); y = 1.8 + (k // 4) * (th + gy)
    tile(s, x, y, tw, th, lab, big, sub.split("\n"), c)
rect(s, 0.4, 5.4, 4 * tw + 3 * gx, 1.1, PANEL, BORDER)
txt(s, 0.6, 5.58, 1.6, 0.25, "Management read", 9, RED, True)
txt(s, 2.25, 5.58, 7.0, 1.2, "Demand inside the dealer network is not the problem: sales to dealer are up. The plan gap is I Motors buying less than it sells while it uses up earlier stock. With MOS already below 3.0, replenishment has to restart, and that is what the second-half recovery depends on.", 9.5, WHITE)
panel(s, [
    "The full-year target is still €1.858M; after six months, €1.084M remains to be billed in Oct–Mar.",
    "Sales to dealer is up 25% on the comparable five months, while I Motors buys less than it sells to dealers in 6 of 7 PMS families: stock timing, not lost business.",
    "Stock is now 2.2 months against a 3.0 target. I Motors → dealer fill rate is back at its 85% target, so further stock cuts would put that at risk.",
    "The lasting gaps are brake shoes, cabin filters and AdBlue filters, low on sales to dealer as well. We propose an SPA there, passed on to dealer prices."])

# ================================================================ 3 FY27 plan monthly
s = base("FY2027 plan", "The first four months held plan; August and September created the YTD gap",
         "I Motors | Apr–Sep 2026 | EUR K", "Source: supplied I Motors FY2027 sales-plan report.")
cd = CategoryChartData(); cd.categories = ["Apr", "May", "Jun", "Jul", "Aug", "Sep"]
cd.add_series("Sales", [140, 151, 153, 147, 109, 73])
cd.add_series("Current budget", [138, 150, 151, 145, 162, 178])
cd.add_series("Orders", [134, 159, 147, 153, 101, 96])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 4.2)
style_chart(ch, '€#,##0"K"'); ch.plots[0].gap_width = 70; ch.plots[0].overlap = -5
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 200
for se, c in zip(ch.plots[0].series, [RED, GREY, AMBER]): col_fill(se, c); labels(se, '#,##0', MUTED, size=7)
rect(s, 0.4, 6.15, 9.2, 0.8, PANEL, BORDER)
txt(s, 0.65, 6.4, 4.3, 0.3, "Apr–Jul sales: €591K vs €584K budget (+1.2%)", 10, "55C271", True)
txt(s, 5.05, 6.4, 4.4, 0.3, "Aug–Sep sales: €182K vs €340K budget (-46.5%)", 10, RED, True)
panel(s, [
    "YTD sales are €774K against €924K current budget: €150K below plan (-16.2%).",
    "Orders are €790K YTD, €134K below the same budget phasing (-14.5%); the order pipeline does not yet close the sales gap.",
    "The issue is concentrated in the last two reported months rather than a weak start to FY27.",
    "The original BP and the current budget both total €1.858M for the year; the current budget is the relevant monthly phasing for this review."])

# ================================================================ 4 H2 recovery
s = base("FY2027 plan", "The second half now has to run above the original monthly pace to recover the full-year target",
         "Full-year target remains €1.858M", "Source: supplied FY2027 sales plan; arithmetic from reported Apr–Sep sales and full-year target.")
txt(s, 0.5, 1.95, 2.0, 0.25, "FY27 target", 10, WHITE, True)
txt(s, 7.6, 1.95, 2.0, 0.25, "€1.858M", 10, WHITE, True, PP_ALIGN.RIGHT)
bw = 9.1; done = 774 / 1858
rect(s, 0.5, 2.3, bw, 0.72, "4A4A4A")
rect(s, 0.5, 2.3, bw * done, 0.72, RED)
txt(s, 0.7, 2.3, 3.0, 0.72, "€774K billed", 13, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.5 + bw * done + 0.2, 2.3, 3.0, 0.72, "€1.084M remaining", 13, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
t3w = (9.1 - 2 * 0.2) / 3
for k, (lab, big, sub, c) in enumerate([
        ("CURRENT H2 PLAN", "€934K", ["Oct–Mar current budget", "Average €156K/month"], WHITE),
        ("REQUIRED H2 BILLING", "€1.084M", ["To still reach the FY27 target", "Average €181K/month"], RED),
        ("RECOVERY ABOVE PLAN", "+€150K", ["Equivalent to +16.1% versus the existing Oct–Mar budget"], RED)]):
    tile(s, 0.5 + k * (t3w + 0.2), 3.45, t3w, 1.75, lab, big, sub, c, 26)
txt(s, 0.5, 5.5, 9.1, 0.3, "This is a recovery requirement, not a forecast. It assumes the full-year target is held unchanged.", 9, MUTED)
panel(s, [
    "The remaining plan is materially harder than the first-half phasing: €181K/month is required versus a planned €156K/month.",
    "A recovery discussion therefore needs both demand actions and execution actions; simply waiting for normal seasonality is not enough evidence.",
    "Sales to dealer is already showing stronger network demand, so the first question is whether I Motors can translate that movement into replenishment from ISZE.",
    "The recovery plan should be built by category and dealer demand, not by applying a flat percentage uplift to every part family."])

# ================================================================ 5 sales to dealer
s = base("Sales to dealer", "Dealer-network sales are running ahead of last year, while ISZE billing is behind plan",
         "I Motors reported sales to dealer | Comparable Apr–Aug periods | EUR K",
         "Source: KPI Master File.xlsx, Parts Handling – Parts Sales DB to DLR in EUR; latest reported month Aug 2026.")
cd = CategoryChartData(); cd.categories = ["Apr", "May", "Jun", "Jul", "Aug"]
cd.add_series("2025", [300, 295, 299, 308, 289]); cd.add_series("2026", [297, 490, 380, 393, 304])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 3.75)
style_chart(ch, '€#,##0"K"'); ch.plots[0].gap_width = 80; ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 550
for se, c in zip(ch.plots[0].series, [GREY, RED]): col_fill(se, c); labels(se, '#,##0', MUTED, size=7.5)
t3w = (9.2 - 2 * 0.2) / 3
for k, (lab, big, sub, c) in enumerate([
        ("APR–AUG 2025", "€1.491M", ["Reported sales to dealer"], WHITE),
        ("APR–AUG 2026", "€1.864M", ["+25.0% vs comparable period"], "55C271"),
        ("LARGEST MONTH", "€490K", ["May 2026"], RED)]):
    tile(s, 0.4 + k * (t3w + 0.2), 5.62, t3w, 1.42, lab, big, sub, c, 22)
panel(s, [
    "Sales to dealer rose by €372K on the comparable five-month period, despite ISZE sales to I Motors being below FY27 plan.",
    "This is the most important disconnect in the review: distributor-to-dealer demand is stronger than the distributor’s replenishment from ISZE.",
    "Do not read the two euro levels as like-for-like margin: they are different sales stages and pricing bases.",
    "The next cut should identify which product families and dealers drove the May–July 2026 step-up, then test whether I Motors stock and replenishment followed."])

# ================================================================ 6 fill rate
s = base("Fill rate  ·  ISZE → I Motors", "Fill rate from ISZE to I Motors stays below the 93% target",
         "ISZE supplying I Motors  |  individual fill rate  |  Aug 2025 – Aug 2026  |  latest point 87%",
         "Source: supplied “INDIVIDUAL FILL RATE” chart, ISZE to I Motors. Target: 93%, the current ISZE internal fill-rate target (previously 95%).")
months = ["Aug 25", "Sep 25", "Oct 25", "Nov 25", "Dec 25", "Jan 26", "Feb 26", "Mar 26", "Apr 26", "May 26", "Jun 26", "Jul 26", "Aug 26"]
fr = [0.89, 0.90, 0.86, 0.85, 0.90, 0.91, 0.86, 0.84, 0.85, 0.86, 0.82, 0.89, 0.87]
cd = CategoryChartData(); cd.categories = months
cd.add_series("Fill rate ISZE → I Motors", fr); cd.add_series("Target 93%", [0.93] * 13)
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '0%'); ch.value_axis.minimum_scale = 0.80; ch.value_axis.maximum_scale = 1.0; ch.value_axis.major_unit = 0.05
a, t = ch.plots[0].series
line_fmt(a, RED, 2.5); line_fmt(t, AMBER, 1.75, marker=False, dash=True)
labels(a, '0%', WHITE, XL_LABEL_POSITION.ABOVE, 8)
rect(s, 0.4, 6.25, 9.2, 0.75, PANEL, BORDER)
txt(s, 0.65, 6.25, 8.8, 0.75, "ISZE → I Motors   |   Latest: 87%   |   High: 91%   |   Low: 82%   |   Target: 93%", 11, RED, True, anchor=MSO_ANCHOR.MIDDLE)
panel(s, [
    "This is our own supply to I Motors. It stays below the 93% target in every month shown; the best month is 91% (January 2026) and the latest 87%.",
    "The low point is 82% in June 2026, followed by 89% and 87%: better, but not yet a stable recovery.",
    "Every missed line here also limits what I Motors can deliver to its dealers (next slide).",
    "Question for ISZE and I Motors: which part numbers make up the missing 6 points, and are they supply constraints or order-timing issues?"])

# ================================================================ 7 stock & availability
s = base("Stock & availability  ·  I Motors", "MOS is 2.2, below the 3.0 target since April; back orders are at a 13-month low",
         "I Motors stock and back orders  |  KPI Master File  |  Aug 2025 – Aug 2026", "Source: KPI Master File.xlsx, Parts Handling – Days of Parts in Stock, Back Order Lines. Latest reported month Aug 2026. MOS target 3.0.")
m13 = ["Aug 25", "Sep 25", "Oct 25", "Nov 25", "Dec 25", "Jan 26", "Feb 26", "Mar 26", "Apr 26", "May 26", "Jun 26", "Jul 26", "Aug 26"]
days = [111, 111, 99, 78, 84, 75, 81, 96, 88, 85, 85, 79, 67]
mos = [round(d / 30, 1) for d in days]
txt(s, 0.5, 1.85, 4.3, 0.22, "MONTHS OF STOCK (MOS)  ·  TARGET 3.0", 8, RED, True)
cd = CategoryChartData(); cd.categories = m13; cd.add_series("MOS", mos); cd.add_series("Target 3.0", [3.0] * 13)
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 2.1, 4.55, 3.6)
style_chart(ch, '0.0', legend=True); ch.value_axis.minimum_scale = 2.0; ch.value_axis.maximum_scale = 4.0; ch.value_axis.major_unit = 0.5
ch.category_axis.tick_labels.font.size = Pt(6.5)
line_fmt(ch.plots[0].series[0], RED, 2.25); labels(ch.plots[0].series[0], '0.0', WHITE, XL_LABEL_POSITION.ABOVE, 7)
line_fmt(ch.plots[0].series[1], AMBER, 1.75, marker=False, dash=True)
txt(s, 0.5, 5.72, 4.3, 0.2, "MOS = days of parts in stock ÷ 30", 7.5, MUTED)
txt(s, 5.15, 1.85, 4.3, 0.22, "BACK-ORDER LINES", 8, RED, True)
bo = [451, 522, 421, 526, 402, 430, 550, 447, 430, 487, 385, 279, 282]
cd = CategoryChartData(); cd.categories = m13; cd.add_series("Back-order lines", bo)
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 5.05, 2.1, 4.55, 3.6)
style_chart(ch, '#,##0', legend=False); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 650; ch.plots[0].gap_width = 50
ch.category_axis.tick_labels.font.size = Pt(6.5)
col_fill(ch.plots[0].series[0], GREY); labels(ch.plots[0].series[0], '#,##0', WHITE, size=7)
rect(s, 0.4, 6.1, 9.2, 0.85, PANEL, BORDER)
txt(s, 0.65, 6.1, 8.8, 0.85, f"Aug 2026:  MOS {mos[-1]:.1f} vs 3.0 target ({3.0 - mos[-1]:.1f} months short)   |   back-order lines {bo[-1]}   |   fill rate to dealers 85.4%, at target", 10.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
panel(s, [
    "MOS was above the 3.0 target from August to October 2025 (3.3 to 3.7) and in March 2026 (3.2). Since April it has stayed below target and fallen further, to 2.2 in August.",
    "Back-order lines are the lowest of the 13 months (279 and 282 in July and August), and fill rate to dealers is at target. The stock cut has not hurt service yet.",
    "Below 3.0 there is little buffer left. Replenishment from ISZE should now return to the rate dealers buy at, which is what the second-half plan needs.",
    "Question: is 2.2 months a deliberate stock target, and which families will I Motors rebuild first?"])

# ================================================================ 8 PMS service-demand cover
import importlib.util
_spec = importlib.util.spec_from_file_location("pms_fy", str(pathlib.Path(__file__).resolve().parent / "pms_fy.py"))
pms_fy = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(pms_fy)
PT, PFL = pms_fy.build()
def cov(c, fy, m="cov_scrap"): return float(PT[(PT.cat == c) & (PT.fy == fy)][m].iloc[0])
def pctx(v): return f"{v*100:.0f}%"
fl_s = {fy: sum(PFL[(fy, "scrap")].values()) for fy in (2024, 2025, 2026)}
fl_t26 = sum(PFL[(2026, "ten")].values())
lower = 1 - cov("oil_filter", 2026, "cov_ten") / cov("oil_filter", 2026)
s = base("PMS service-demand cover", "Filters and brakes show opposite movements; the denominator is the modelled service demand",
         "Portugal / I Motors | fiscal years (April–March) | denominator: service demand of the registered fleet, scrap-rate model",
         "Sources: ISZE billing to I Motors (units, credit notes excluded); ISZE service-demand model; I Motors registration file (fleet per fiscal year).")
LAB = {"air_filter": "Air filter", "oil_filter": "Oil filter", "fuel_filter": "Fuel filter", "brake_pad": "Brake pad", "clutch_kit": "Clutch kit", "brake_shoe": "Brake shoe"}
OBS = {"air_filter": f"Rising every year ({pctx(cov('air_filter', 2024))} → {pctx(cov('air_filter', 2026))}); what practice can be copied?",
       "oil_filter": "FY2025 peak = bulk purchases in Jul 2024 and Jan 2025; FY2026 is that stock being used up",
       "fuel_filter": "Same pattern as oil filter: FY2025 peak from the Jul 2024 bulk purchase",
       "brake_pad": "FY2025 dip follows the earlier stock build; FY2026 includes the May–Sep 2025 campaign purchases",
       "clutch_kit": f"Flat at about {pctx(cov('clutch_kit', 2026))}; low but stable",
       "brake_shoe": "About 10%: persistent underdevelopment"}
def pc(v):
    x = round(v * 100); return GREEN if x >= 55 else (AMBER if x >= 35 else RED)
cols = [(0.4, 1.45), (1.85, 0.85), (2.7, 0.85), (3.55, 0.85), (4.4, 1.05), (5.45, 4.15)]
y = 1.85
for k, ((x, w), h) in enumerate(zip(cols, ["Category", "FY2024", "FY2025", "FY2026", "FY2026 on 10-yr car park", "Observation / question"])):
    rect(s, x, y, w, 0.42, RED if k != 4 else "8A0A10")
    txt(s, x + 0.08, y, w - 0.14, 0.42, h, 8 if k != 4 else 7, WHITE, True, PP_ALIGN.LEFT if k in (0, 5) else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
y += 0.46
for i, c in enumerate(pms_fy.CATS):
    fillc = ROW2 if i % 2 == 0 else PANEL
    vals = [LAB[c], cov(c, 2024), cov(c, 2025), cov(c, 2026), cov(c, 2026, "cov_ten"), OBS[c]]
    for j, ((x, w), v) in enumerate(zip(cols, vals)):
        rect(s, x, y, w, 0.52, fillc)
        if j in (1, 2, 3, 4):
            txt(s, x, y, w, 0.52, pctx(v), 11 if j < 4 else 9.5, pc(v) if j < 4 else MUTED, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        else:
            txt(s, x + 0.1, y, w - 0.18, 0.52, v, 9 if j == 0 else 8.2, WHITE, j == 0, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.55
txt(s, 0.4, y + 0.15, 9.2, 0.7,
    f"Denominator note: expected replacement units by service family from the ISZE service-demand model, applied to each fiscal year's fleet counted from the I Motors registration file (scrap-rate model: {fl_s[2024]:,.0f} → {fl_s[2025]:,.0f} → {fl_s[2026]:,.0f} vehicles; 10-year car park FY2026: {fl_t26:,.0f}). Numerator: units billed by ISZE to I Motors, Genuine and BVP, credit notes excluded. Not the same as full GPS target spend.", 7.8, MUTED)
panel(s, [
    "On fiscal years the pattern is the same: oil and fuel filters peak in FY2025 because of the July 2024 and January 2025 bulk purchases, then fall back in FY2026 while dealer sales of oil filters stay at about 300 a month.",
    "Brake pads mirror it: FY2025 dips after the earlier stock build, and FY2026 includes the May to September 2025 campaign purchases. Dealer sales stayed at about 80–90 pads a month.",
    f"The denominator now grows with the registered fleet each year. On the 10-year car park every share is about {lower*100:.0f}% lower; the pattern does not change.",
    "Question: how many months of filter and brake-pad stock does I Motors hold today, so the underlying share can be read without the stock effect?"])

if not globals().get("PRS"):
    prs.save(OUT)
print("saved", OUT)
