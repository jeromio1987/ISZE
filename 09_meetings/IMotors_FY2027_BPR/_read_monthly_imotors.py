from pathlib import Path
import openpyxl
from pptx import Presentation
from lxml import etree

# --- get Sep billing from daily report + prior months from deck chart 4 ---
sep_bill = 134461.35
sep_plan = 177732.91
kpi = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_nuno_kpi_portugal_20261001\09 - Set_KPI Check Sheet Monthly 26_Portugal.xlsx")
wb = openpyxl.load_workbook(kpi, data_only=True, read_only=True)
ws = wb["Checklist"]
# sales to dealer row
for row in ws.iter_rows(values_only=True):
    if row and row[1] and "Sales Distributor to all Dealers" in str(row[1]):
        # Jan..Sep at cols 2..
        months = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep"]
        vals = list(row[2:11])
        print("STD", dict(zip(months, vals)))
        # Apr-Sep
        print("STD Apr-Sep 2026", vals[3:9], "sum", sum(float(v) for v in vals[3:9]))
wb.close()

# Read chart series from slide 4 and 5 and 8
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}

def dump_chart(slide_idx, label):
    slide = prs.slides[slide_idx]
    for sh in slide.shapes:
        if not sh.has_chart: continue
        print(f"\n-- {label} chart {sh.name} --")
        # categories
        cats = [t.text for t in sh.chart._chartSpace.xpath(".//c:cat//c:pt/c:v", namespaces=NS)]
        print("cats", cats)
        for ser in sh.chart._chartSpace.xpath(".//c:ser", namespaces=NS):
            name = ser.xpath("./c:tx//c:v", namespaces=NS)
            name = name[0].text if name else "?"
            vals = [t.text for t in ser.xpath("./c:val//c:pt/c:v", namespaces=NS)]
            print("ser", name, vals)

for idx, lab in [(3,"s4 plan"), (4,"s5 pace"), (7,"s8 std"), (15,"s16 fill"), (16,"s17 mos/bo")]:
    dump_chart(idx, lab)

# executive tiles slide 3
print("\n=== EXEC tiles ===")
for sh in prs.slides[2].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        t = sh.text_frame.text.strip().replace("\n"," | ")
        if any(x in t for x in ["€", "%", "MOS", "plan", "Sales", "K", "M"]):
            print(t[:140])
