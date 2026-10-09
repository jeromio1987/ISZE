# -*- coding: utf-8 -*-
"""Rebuild ONLY slide 9 rolling sales-to-dealer line through Sep 2026.

Source months:
  Oct 24–Aug 26 from KPI Master File → Parts Handling → Parts Sales DB to DLR € → Imotors
  Sep 26 from KPI Check Sheet Sep 2026 (416806)
Apr–Aug 24 are N/D in the master — not used. No invented months.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt, Emu

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
PPTX = ROOT / "09_meetings/IMotors_FY2027_BPR/I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"
BAK = ROOT / "09_meetings/IMotors_FY2027_BPR/_bak_pre_slide9_rolling.pptx"

# Last 12 months Oct 2025 – Sep 2026 (EUR, then chart in EUR K)
LAST12_LABELS = ["Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
LAST12_EUR = [
    402361, 445067, 248410, 415020, 339278, 381438,
    297074, 489770, 379751, 393481, 303777, 416806,
]
LAST12_K = [round(v / 1000) for v in LAST12_EUR]
# Totals already on tiles: €4.51M / +21% vs €3.73M — verified from same series

RED = RGBColor.from_string("E30613")
MUTED = RGBColor.from_string("BDBDBD")
GRID = RGBColor.from_string("4D4D4D")
WHITE = RGBColor.from_string("F2F2F2")
FONT = "Aptos"


def replace_text(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    n += 1
    return n


shutil.copy2(PPTX, BAK)
prs = Presentation(str(PPTX))
slide = prs.slides[8]  # slide 9

# Remove frozen chart picture only
pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
if not pics:
    raise SystemExit("No picture on slide 9 — abort (unexpected layout)")
pic = pics[0]
left, top, width, height = pic.left, pic.top, pic.width, pic.height
sp = pic._element
sp.getparent().remove(sp)
print("removed picture", left, top, width, height)

# Live line chart in the same box
cd = CategoryChartData()
cd.categories = LAST12_LABELS
cd.add_series("Sales to dealer €K", LAST12_K)
chart = slide.shapes.add_chart(
    XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, cd
).chart

chart.has_title = False
chart.has_legend = False
chart.font.name = FONT
chart.font.size = Pt(8)
chart.font.color.rgb = MUTED

va = chart.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRID
va.major_gridlines.format.line.width = Pt(0.5)
va.format.line.fill.background()
va.tick_labels.font.color.rgb = MUTED
va.tick_labels.font.size = Pt(8)
va.tick_labels.number_format = "0"
va.tick_labels.number_format_is_linked = False

ca = chart.category_axis
ca.format.line.color.rgb = GRID
ca.tick_labels.font.color.rgb = MUTED
ca.tick_labels.font.size = Pt(9)

ser = chart.series[0]
ser.format.line.color.rgb = RED
ser.format.line.width = Pt(2.25)
try:
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 7
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = RED
    ser.marker.format.line.color.rgb = RED
except Exception as e:
    print("marker style warn", e)

# Label peaks + Sep end point
try:
    dl = ser.data_labels
    dl.show_value = False
    for idx in (1, 7, 11):  # Nov, May, Sep
        pt = ser.points[idx]
        pt.data_label.text_frame.clear()
        # python-pptx: show value via has_data_label
        pt.data_label.show_value = True
        pt.data_label.font.size = Pt(9)
        pt.data_label.font.bold = True
        pt.data_label.font.color.rgb = WHITE
        pt.data_label.number_format = "€#,##0\"K\""
        pt.data_label.number_format_is_linked = False
except Exception as e:
    print("labels warn", e)

# Observation: mention Sep recovery, not only weaker August
replace_text(
    slide,
    "Month-to-month movement is shown as one connected line so the pattern is visible: strong May, solid Jun-Jul, weaker August.",
    "Month-to-month movement is one line through September: strong May, solid Jun–Jul, weaker August, then September back at €417K.",
)

# Source note — Sep from Check Sheet, rest from Master
replace_text(
    slide,
    "Source: KPI Master File.xlsx, Parts Handling – Parts Sales DB to DLR in EUR; latest reported month Sep 2026 (KPI Check Sheet).",
    "Source: KPI Master File (Parts Sales DB to DLR €) Oct 2024–Aug 2026; Sep 2026 from KPI Check Sheet. Apr–Aug 2024 N/D in master — not used.",
)

prs.save(str(PPTX))
print("SAVED", PPTX)
print("series K", LAST12_K, "sum M", round(sum(LAST12_EUR) / 1e6, 2))
