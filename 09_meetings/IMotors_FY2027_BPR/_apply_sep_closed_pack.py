# -*- coding: utf-8 -*-
"""Apply Sep-closed P0/P1/P2 pack fixes to I-Motors review deck. One as-of: 30 Sep closed."""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
PPTX = ROOT / "09_meetings/IMotors_FY2027_BPR/I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"
BAK = ROOT / "09_meetings/IMotors_FY2027_BPR/_bak_pre_sep_closed_pack_apply.pptx"
CMD = ROOT / "09_meetings/IMotors_FY2027_BPR/OPEN_I_Motors_BPR_2026-10-06_Sep_closed.cmd"

# Daily sales report 30 Sep 16:51 — IMOTORS row
SEP_BILL = 134461.35
SEP_ORD = 154689.07
SEP_PLAN = 177732.91
SEP_BILL_K = round(SEP_BILL / 1000, 1)  # 134.5
SEP_ORD_K = round(SEP_ORD / 1000, 1)    # 154.7
SEP_PLAN_K = round(SEP_PLAN / 1000, 1)  # 177.7
PCT_PLAN = SEP_BILL / SEP_PLAN          # 75.7%

# KPI Check Sheet Sep 2026
STD_2026_K = [297.1, 489.8, 379.8, 393.5, 303.8, 416.8]
STD_2025_K = [300.0, 295.0, 299.0, 308.0, 289.0, 325.0]  # Sep = 1816 - 1491
FILL_DEALER_SEP = 0.8105
MOS_SEP = 2.49
BO_SEP = 495
FILL_ISZE_SEP = 0.87  # exec tile Sep closed

FY_TARGET_K = 1858.0

BG, PANEL, BORDER = "262626", "303030", "6E6E6E"
WHITE, MUTED, GRID = "F2F2F2", "BDBDBD", "4D4D4D"
RED, AMBER = "E30613", "F5A000"
FONT = "Aptos"


def rgb(h):
    return RGBColor.from_string(h)


def txt(slide, x, y, w, h, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def hline(slide, x, y, w, color, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)


def panel(slide, points, x=9.86, y=1.46, w=3.25, h=5.58):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.28, y + 0.26, w - 0.5, 0.2, "OBSERVATIONS / QUESTIONS", 7.5, RED, True)
    yy = y + 0.72
    step = (h - 0.95) / max(len(points), 1)
    for i, pnt in enumerate(points):
        rect(slide, x + 0.28, yy + 0.07, 0.09, 0.09, RED if i == 0 else AMBER, shape=MSO_SHAPE.OVAL)
        txt(slide, x + 0.52, yy, w - 0.8, step - 0.08, pnt, 8.4, WHITE)
        yy += step


def replace_text(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    n += 1
        # split across runs
        if old in sh.text_frame.text and n == 0:
            full = sh.text_frame.text.replace(old, new)
            first = True
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if first:
                        run.text = full
                        first = False
                        n += 1
                    else:
                        run.text = ""
    return n


def replace_anywhere(prs, old, new):
    total = 0
    for slide in prs.slides:
        total += replace_text(slide, old, new)
    return total


def charts_on(slide):
    return [sh for sh in slide.shapes if sh.has_chart]


def set_chart(shape, categories, series_list):
    """series_list = [(name, values), ...] in chart series order."""
    cd = CategoryChartData()
    cd.categories = list(categories)
    for name, vals in series_list:
        cd.add_series(name, list(vals))
    shape.chart.replace_data(cd)


def blank_layout(prs):
    for i in range(len(prs.slide_layouts)):
        layout = prs.slide_layouts[i]
        if "Blank" in (layout.name or "") or i == 6:
            return layout
    return prs.slide_layouts[0]


def base_slide(prs, section, title, sub, source, page_no):
    s = prs.slides.add_slide(blank_layout(prs))
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BG)
    bg.line.fill.background()
    spTree = s.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)

    txt(s, 0.25, 0.10, 3.8, 0.18, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 6.6, WHITE, True)
    txt(s, 0.25, 0.43, 7.5, 0.22, section.upper(), 7.5, RED, True)
    txt(s, 0.25, 0.68, 9.45, 0.66, title, 18 if len(title) > 70 else 20, WHITE, True)
    txt(s, 0.25, 1.42, 9.45, 0.24, sub, 8.5, MUTED)
    hline(s, 0.25, 7.18, 12.85, "575757")
    txt(s, 0.25, 7.24, 11.5, 0.16, source, 6.5, "8C8C8C")
    txt(s, 12.55, 7.24, 0.55, 0.16, str(page_no), 7, MUTED, align=PP_ALIGN.RIGHT)
    return s


def style_chart(ch, num_fmt="#,##0", legend=True):
    ch.font.name = FONT
    ch.font.size = Pt(8)
    ch.font.color.rgb = rgb(MUTED)
    ch.has_title = False
    if legend:
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.color.rgb = rgb(MUTED)
        ch.legend.font.size = Pt(8)
    else:
        ch.has_legend = False
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(GRID)
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.tick_labels.font.size = Pt(7.5)
    va.tick_labels.number_format = num_fmt
    va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(GRID)
    ca.tick_labels.font.color.rgb = rgb(MUTED)
    ca.tick_labels.font.size = Pt(7.5)


def col_fill(series, c):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(c)
    series.format.line.fill.background()


# ---------------------------------------------------------------------------
shutil.copy2(PPTX, BAK)
prs = Presentation(str(PPTX))

# ===== SLIDE 4 — plan bars =====
s4 = prs.slides[3]
for sh in charts_on(s4):
    # Keep Apr-Aug sales/orders from current chart; fix Sep
    from lxml import etree as ET
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = ET.fromstring(ET.tostring(sh.chart._chartSpace))
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    sales = None
    orders = None
    budget = None
    for nm, vals in sers:
        low = nm.lower()
        if "budget" in low or "current" in low or "plan" in low:
            budget = [138.0, 150.0, 151.0, 145.0, 162.0, SEP_PLAN_K]
        elif "order" in low:
            orders = list(vals[:5]) + [SEP_ORD_K]
        else:
            sales = list(vals[:5]) + [SEP_BILL_K]
    if sales is None:
        sales = [140.0, 151.0, 153.0, 147.0, 109.0, SEP_BILL_K]
    if orders is None:
        orders = [134.0, 159.0, 147.0, 153.0, 101.0, SEP_ORD_K]
    if budget is None:
        budget = [138.0, 150.0, 151.0, 145.0, 162.0, SEP_PLAN_K]
    out = []
    for nm, _ in sers:
        low = nm.lower()
        if "budget" in low or "current" in low or "plan" in low:
            out.append((nm, budget))
        elif "order" in low:
            out.append((nm, orders))
        else:
            out.append((nm, sales))
    set_chart(sh, ["Apr", "May", "Jun", "Jul", "Aug", "Sep"], out)
    print("S4", out)

ytd_k = sum(sales)
print("YTD sales K", ytd_k)

replace_text(s4, "Apr–Aug ISZE billing; September from ISZE Monthly Report VP (closed).",
             "ISZE Daily Sales Report 30 Sep closed (Billing MTD / order receipt); Apr–Aug from ISZE billing series.")
# Aug-Sep callout already ~243 vs 340 if Sep is 134.5 — verify
aug_sep = sales[4] + sales[5]
bud_as = budget[4] + budget[5]
replace_text(s4, "Aug–Sep sales: €243K vs €340K budget (-28.5%)",
             f"Aug–Sep sales: €{aug_sep:.0f}K vs €{bud_as:.0f}K budget ({(aug_sep/bud_as-1)*100:.1f}%)")

# ===== SLIDE 5 — pace =====
s5 = prs.slides[4]
apr_sep_eur = None
for sh in charts_on(s5):
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    cats = [n.text for n in root.findall(".//{%s}cat//{%s}pt/{%s}v" % (C, C, C))]
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    n = max(len(v) for _, v in sers)
    base_cats = cats[:n] if len(cats) >= n else ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"][:n]
    out = []
    for nm, vals in sers:
        if "2027" in nm:
            merged = list(vals)
            # first 5 already Apr-Aug euros; set/append Sep
            if len(merged) >= 5:
                if len(merged) == 5:
                    merged.append(SEP_BILL)
                else:
                    merged[5] = SEP_BILL
                # zero future months after Sep
                for i in range(6, len(merged)):
                    merged[i] = 0.0
            while len(merged) < n:
                merged.append(0.0)
            apr_sep_eur = merged[:6]
            out.append((nm, merged[:n]))
        else:
            out.append((nm, (vals + [0] * n)[:n]))
    set_chart(sh, base_cats[:n], out)
    print("S5 FY27", apr_sep_eur)

if apr_sep_eur is None:
    apr_sep_eur = [140054, 151368, 153167, 147111, 109377, SEP_BILL]
ytd_eur = sum(apr_sep_eur)
ytd_tile = round(ytd_eur / 1000)  # ~836
pace_full = ytd_eur * 12 / 6
pace_pct = pace_full / (FY_TARGET_K * 1000) * 100
print("YTD eur", ytd_eur, "pace%", pace_pct, "full", pace_full)

replace_text(s5, "April-August billing is flat on last year; a straight-line full year gives 90.5% of target",
             f"April–September billing totals €{ytd_tile}K; a straight-line full year gives {pace_pct:.1f}% of target")
replace_text(s5, "pace chart still April–August", "through September 2026 closed")
replace_text(s5, "April-August 2026", "April–September 2026")
replace_text(s5, "€701K", f"€{ytd_tile}K")
replace_text(s5, "90.5%", f"{pace_pct:.1f}%")
replace_text(s5, "€1.68M", f"€{pace_full/1e6:.2f}M")
replace_text(s5, "× 12/5", "× 12/6")
replace_text(s5, "(× 12/5)", "(× 12/6)")
replace_text(s5, "A straight line (× 12/5) gives €1.68M, €176K short of the €1.858M target.",
             f"A straight line (× 12/6) gives €{pace_full/1e6:.2f}M, €{(FY_TARGET_K*1000-pace_full)/1000:.0f}K short of the €1.858M target.")
replace_text(s5, "April-August was only 37.5% of the year",
             "April–September was half the year on the calendar; FY2026 was still back-loaded after September")
replace_text(s5, "to 30 September 2026 where stated; pace chart still April–August",
             "to 30 September 2026 closed")

# ===== SLIDE 8 — sales to dealer =====
s8 = prs.slides[7]
for sh in charts_on(s8):
    set_chart(sh, ["Apr", "May", "Jun", "Jul", "Aug", "Sep"], [
        ("2025", STD_2025_K),
        ("2026", STD_2026_K),
    ])
    print("S8 set")
replace_text(s8, "Comparable Apr–Sep periods", "Comparable Apr–Sep periods")
replace_text(s8, "latest reported month Sep 2026 (KPI Check Sheet).",
             "latest reported month Sep 2026 (KPI Check Sheet); chart now includes September.")

# ===== SLIDE 15 — ISZE fill: append Sep =====
s15 = prs.slides[14]
for sh in charts_on(s15):
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    n0 = len(sers[0][1])
    cats0 = [n.text for n in root.findall(".//{%s}cat//{%s}pt/{%s}v" % (C, C, C))][:n0]
    if cats0 and str(cats0[-1]).startswith("Aug"):
        cats = cats0 + ["Sep 26"]
        out = []
        for nm, vals in sers:
            if "target" in nm.lower():
                out.append((nm, list(vals) + [0.93]))
            else:
                out.append((nm, list(vals) + [FILL_ISZE_SEP]))
        set_chart(sh, cats, out)
        print("S15 appended Sep", FILL_ISZE_SEP)

# ===== SLIDE 16 — dealer fill =====
s16 = prs.slides[15]
for sh in charts_on(s16):
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    n0 = len(sers[0][1])
    cats0 = [n.text for n in root.findall(".//{%s}cat//{%s}pt/{%s}v" % (C, C, C))][:n0]
    if cats0 and str(cats0[-1]).startswith("Aug"):
        cats = cats0 + ["Sep 26"]
        out = []
        for nm, vals in sers:
            if "target" in nm.lower():
                out.append((nm, list(vals) + [0.85]))
            else:
                out.append((nm, list(vals) + [FILL_DEALER_SEP]))
        set_chart(sh, cats, out)
        print("S16 appended Sep", FILL_DEALER_SEP)

# ===== SLIDE 17 — MOS + BO =====
s17 = prs.slides[16]
for sh in charts_on(s17):
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    sers = []
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        sers.append((nm, vals))
    n0 = len(sers[0][1])
    cats0 = [n.text for n in root.findall(".//{%s}cat//{%s}pt/{%s}v" % (C, C, C))][:n0]
    if not (cats0 and str(cats0[-1]).startswith("Aug")):
        continue
    cats = cats0 + ["Sep 26"]
    out = []
    for nm, vals in sers:
        if "target" in nm.lower():
            out.append((nm, list(vals) + [3.0]))
        elif max(vals) > 50:
            out.append((nm, list(vals) + [float(BO_SEP)]))
        else:
            out.append((nm, list(vals) + [MOS_SEP]))
    set_chart(sh, cats, out)
    print("S17", sh.name, [(n, v[-1]) for n, v in out])

# ===== P2 — Your data (slide 10) =====
s10 = prs.slides[9]
replace_text(
    s10,
    "Sources: I Motors sales-to-dealer file (January 2025 to August 2026); I Motors registration file (to June 2026); ISZE billing to I Motors; ISZE PR00 price list of 10 August 2026.",
    "Sources: I Motors sales-to-dealer PN file (January 2025 to August 2026); KPI Check Sheet for Sep 2026 sales-to-dealer euros; I Motors registration file (to June 2026); ISZE billing; ISZE PR00 list 10 Aug 2026.",
)
replace_text(
    s10,
    "•  Sales to dealer: 26,966 lines, 7,262 part numbers, 91,153 units, January 2025 to August 2026 (month × part number × quantity).",
    "•  Sales to dealer PN file: 26,966 lines, 7,262 part numbers, 91,153 units, January 2025 to August 2026 (month × PN × qty). September euros come from the KPI Check Sheet (€416.8k), not yet from a Sep PN extract.",
)

# ===== P2 — Genuine share / GPS language (slide 21) =====
s21 = prs.slides[20]
replace_text(
    s21,
    "Genuine parts share rose to 58% in FY2026 on the scrap-rate model (51% on the 10-year count)",
    "Genuine parts share (scrap-rate model) rose to 58% in FY2026 — not the IDS locked GPS figure of 63.3%",
)
replace_text(
    s21,
    "* FY2027 pace = April-August 2026 billing × 12/5, measured against the FY2026 fleet (the registration file ends in June 2026).",
    "* FY2027 pace on this slide still uses the Apr–Aug cut. Room GPS for Portugal IDS stays locked at 63.3%; do not mix that with this scrap-rate share.",
)
replace_text(
    s21,
    "Source: ISZE billing to I Motors, invoice level, to 30 September 2026 where stated; pace chart still April–August; I Motors registration file; ISZE target-spend rates.",
    "Source: ISZE billing to I Motors; I Motors registration file; ISZE target-spend rates. This slide = scrap-rate genuine share. IDS / visit pack GPS = 63.3% (locked).",
)
replace_text(
    s21,
    "At the April-August pace, FY2027 would land at 52% (scrap-rate model), below FY2026.",
    "At the April–August pace on this slide, FY2027 would land at 52% (scrap-rate model), below FY2026. Use 63.3% only when speaking IDS GPS.",
)

# Exec genuine tile gloss (slide 3)
s3 = prs.slides[2]
replace_text(
    s3,
    "FY2026, warranty included, on the scrap-rate model\n51.4% on the 10-year car park",
    "FY2026 scrap-rate model (not IDS GPS 63.3%)\n51.4% on the 10-year car park",
)

# ===== P1 — Lisbon one-pager after slide 17 =====
s_lisbon = base_slide(
    prs,
    "Lisbon strip  ·  September closed",
    "One strip for Wednesday: Sep KPI, billing vs plan, and the August aged back-order pocket",
    "As-of 30 Sep 2026 closed  |  BO part-number deep-dive stays at 31 Aug",
    "Sources: KPI Check Sheet Sep 2026; ISZE Daily Sales Report 30 Sep 16:51; Aug BO book analysis (next slides).",
    18,
)

tiles = [
    ("DEALER FILL", "81.0%", "Sep · target 85%", RED),
    ("MOS", "2.5", "Sep · target 3.0", AMBER),
    ("B/O LINES", "495", "Sep · was 282 in Aug", RED),
    ("SALES TO DEALERS", "€417k", "Sep KPI euro", WHITE),
    ("ISZE BILLING", "€134.5k", f"vs plan €{SEP_PLAN_K:.1f}k · {PCT_PLAN*100:.1f}%", AMBER),
    ("AGED AUG BO", "€33.6k", "91+ days · Aug book only", RED),
]
for i, (lab, big, sub, col) in enumerate(tiles):
    x = 0.35 + (i % 3) * 3.15
    y = 1.75 + (i // 3) * 1.55
    rect(s_lisbon, x, y, 3.00, 1.40, PANEL, BORDER)
    txt(s_lisbon, x + 0.14, y + 0.12, 2.7, 0.20, lab, 7.5, RED, True)
    txt(s_lisbon, x + 0.14, y + 0.40, 2.7, 0.45, big, 26, col, True)
    txt(s_lisbon, x + 0.14, y + 0.95, 2.7, 0.35, sub, 8.5, MUTED)

panel(s_lisbon, [
    "Start here on Wednesday: demand is up (€417k to dealers) while fill fell to 81% and B/O jumped to 495.",
    "ISZE billing closed September at €134.5k against €177.7k plan (75.7%). Orders were €154.7k.",
    "Part-number back-order talk stays on the August book (next slides): half the Aug euro was older than 90 days (€33.6k).",
    "Ask: Sep PN tab for the 495 lines, and sales-to-dealer feed with dealer code.",
])

# Move Lisbon strip to after slide 17 (index 16) — currently last
sldIdLst = prs.part._element.find(qn("p:sldIdLst"))
ids = list(sldIdLst)
new_el = ids[-1]
sldIdLst.remove(new_el)
sldIdLst.insert(17, new_el)

# Agenda renumber for +1 slide
s2 = prs.slides[1]
agenda_fixes = [
    ("slides 15–19", "slides 15–20"),
    ("slides 20–25", "slides 21–26"),
    ("slides 26–27", "slides 27–28"),
    ("slides 28–29", "slides 29–30"),
    ("slides 30–32", "slides 31–33"),
]
for old, new in agenda_fixes:
    replace_text(s2, old, new)

prs.save(str(PPTX))

# Renumber page footers
prs2 = Presentation(str(PPTX))
for i, slide in enumerate(prs2.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and sh.left > Inches(12.0) and sh.top > Inches(7.0):
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.text = str(i)
                    break

# Point slide 17 question to Lisbon + Aug slides
s17b = prs2.slides[16]
replace_text(
    s17b,
    "Question: which aged back-order euros (see next slides, August book) should clear first while MOS rebuilds?",
    "Question: after the Lisbon strip, which aged August back-order euros should clear first while MOS rebuilds?",
)

prs2.save(str(PPTX))
CMD.write_bytes(
    b'@echo off\r\nstart "" "%~dp0I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"\r\n'
)

print("SAVED", PPTX, "slides", len(prs2.slides), "size", PPTX.stat().st_size)
for i in range(3, 21):
    slide = prs2.slides[i]
    texts = [sh.text_frame.text.strip().replace("\n", " ")[:70]
             for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = texts[2] if len(texts) > 2 else (texts[1] if len(texts) > 1 else "?")
    print(f"{i+1:02d}", title)
