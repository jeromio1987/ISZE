"""I Motors supporting slides: fleet (two counting methods), genuine parts share, FY2027 pace, stock movement.

Sources (all read-only):
  registrations : ISZE/01_data/07_UIO/Matriculas Isuzu SAP.xlsx (Astara, vehicle level, 2005 - 29 Jun 2026)
  survival      : UIO V19 curves, as copied in ISZE/01_data/08_Distributor/Imotors_Portugal_UIO_GPS_PMS_Study.xlsx
  billing       : ISZE/01_data/00_derived/sales_fact.parquet, sold-to 519 + 92, all order types (warranty parts included)
  sales to dealer: ISZE/01_data/00_derived/dealer_sellout_long_PORTUGAL_CY2025.csv + _CY2026YTD.csv (month rows only)
  families      : ISZE/01_data/00_derived/parts_catalog_dimension.csv (pms_category)
  campaign price: ISZE/01_data/00_derived/pms_campaign_suspect_lines_2024_2025.csv (>=20% under list)
  FY27 target   : ISZE/01_data/ISZE_GPS_v5.xlsx 'BP FY27 Reviewed' (EUR 1,858,269)
"""
import pathlib
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

R = pathlib.Path(r"C:\Users\jerom\Desktop\claude\ISZE")
DER = R / "01_data" / "00_derived"
OUT = R / "09_meetings" / "IMotors_FY2027_BPR" / "supporting_slides" / "I_Motors_Fleet_GPS_Stock_Supporting_Slides_2026-09-24.pptx"

BG, PANEL, BORDER = "262626", "303030", "6E6E6E"
WHITE, MUTED, GRID = "F2F2F2", "BDBDBD", "4D4D4D"
RED, AMBER, GREY = "E30613", "F5A000", "B7B7B7"
FONT = "Aptos"
def rgb(h): return RGBColor.from_string(h)

# ------------------------------------------------------------------ data: registrations and fleet
def load_regs():
    F = R / "01_data" / "07_UIO" / "Matriculas Isuzu SAP.xlsx"
    a = pd.read_excel(F, sheet_name="2005-2019", header=None, skiprows=2).iloc[:, 1:16]
    a.columns = ["vin", "model", "family", "Rodeo", "TF190", "RT56", "RT88", "RG", "N3.5", "N5", "N7.5", "F11", "F14", "plate", "reg_date"]
    b = pd.read_excel(F, sheet_name="2019-2026", header=None, skiprows=2).iloc[:, 1:16]
    b.columns = ["vin", "family", "Rodeo", "TF190", "RT56", "RT88", "RG", "N3.5", "N5", "N7.5", "F11", "F14", "model", "reg_date", "plate"]
    df = pd.concat([a, b], ignore_index=True)
    df = df[df.vin.notna()].drop_duplicates("vin").copy()
    df["reg_date"] = pd.to_datetime(df.reg_date, errors="coerce")
    seg_cls = {"Rodeo": "LCV", "TF190": "LCV", "RT56": "LCV", "RT88": "LCV", "RG": "LCV",
               "N3.5": "LD", "N5": "LD", "N7.5": "LD", "F11": "F", "F14": "F"}
    def cls(r):
        for c, k in seg_cls.items():
            if pd.notna(r[c]) and str(r[c]).strip() not in ("", "0"):
                return k
        m = str(r["model"]).upper()
        if m.startswith("D-MAX"): return "LCV"
        if m.startswith("N"): return "LD"      # unclassified N-series codes (NKR/NQR/NNR) are light trucks
        if m.startswith("F"): return "F"
        return None
    df["cls"] = df.apply(cls, axis=1)
    df = df[df.cls.notna()].copy()
    df["fy"] = df.reg_date.dt.year + (df.reg_date.dt.month >= 4).astype(int)
    return df

CV = [1, .99, .98, .97, .96, .95, .9, .8, .7, .55, .4, .3, .2, .1, .05]
LCV_C = [1, .98, .95, .92, .89, .84, .8, .65, .5, .4, .2]
def surv(c, age):
    cur = LCV_C if c == "LCV" else CV
    return cur[age - 1] if 1 <= age <= len(cur) else 0.0
RATE = {"LD": 739, "F": 887, "LCV": 250}

regs = load_regs()
coh = regs.groupby(["cls", "fy"]).size().rename("n").reset_index()

def fleet(FY):
    out = {}
    for c in ["LD", "F", "LCV"]:
        k = coh[(coh.cls == c) & (coh.fy <= FY)]
        out[c] = dict(scrap=sum(r.n * surv(c, FY - r.fy + 1) for r in k.itertuples()),
                      ten=float(k[k.fy > FY - 10].n.sum()),
                      scrap_in10=sum(r.n * surv(c, FY - r.fy + 1) for r in k.itertuples() if r.fy > FY - 10),
                      scrap_old=sum(r.n * surv(c, FY - r.fy + 1) for r in k.itertuples() if r.fy <= FY - 10))
    return out
FL = {fy: fleet(fy) for fy in (2024, 2025, 2026)}
def tot(fy, key): return sum(v[key] for v in FL[fy].values())
def wallet(fy, key): return sum(v[key] * RATE[c] for c, v in FL[fy].items())

# ------------------------------------------------------------------ data: billing
sf = pd.read_parquet(DER / "sales_fact.parquet", columns=["billing_date", "sold_to_sap", "part_number", "billed_qty", "net_sales_value"])
pt = sf[sf.sold_to_sap.isin([519, 92])].copy()
pt["fy"] = pt.billing_date.dt.year + (pt.billing_date.dt.month >= 4).astype(int)
pt["m"] = pt.billing_date.dt.to_period("M")
BILL = pt.groupby("fy").net_sales_value.sum()
YTD = pt[(pt.billing_date >= "2026-04-01") & (pt.billing_date <= "2026-08-31")].net_sales_value.sum()
LY = pt[(pt.billing_date >= "2025-04-01") & (pt.billing_date <= "2025-08-31")].net_sales_value.sum()
RUN = YTD * 12 / 5
TARGET = 1858269.405
MON = pt.groupby("m").net_sales_value.sum()
assert pt.billing_date.max() >= pd.Timestamp("2026-08-31")

# ------------------------------------------------------------------ data: families, sell-in vs sell-out
def norm(s): return s.astype(str).str.strip().str.replace(r"\.0$", "", regex=True).str.lstrip("0").replace("", "0")
cat = pd.read_csv(DER / "parts_catalog_dimension.csv", dtype=str, usecols=["part_number", "pms_category"])
cat["pn"] = norm(cat.part_number)
cat = cat[cat.pms_category.notna() & (cat.pms_category != "not_a_service_family")].drop_duplicates("pn")
def fam(c):
    c = str(c)
    if c.startswith("brake_pad"): return "brake_pad"
    if c.startswith("brake_rotor") or c.startswith("brake_disc"): return "brake_rotor"
    if c.startswith("brake_shoe") or c.startswith("brake_lining"): return "brake_shoe"
    if c in ("oil_filter", "fuel_filter", "air_filter", "clutch_kit"): return c
    return None
cat["fam"] = cat.pms_category.map(fam)
cat = cat[cat.fam.notna()][["pn", "fam"]]
FAMS = ["brake_pad", "brake_rotor", "brake_shoe", "oil_filter", "fuel_filter", "air_filter", "clutch_kit"]
FAM_LABEL = {"brake_pad": "Brake pads", "brake_rotor": "Brake rotors", "brake_shoe": "Brake shoes", "oil_filter": "Oil filters",
             "fuel_filter": "Fuel filters", "air_filter": "Air filters", "clutch_kit": "Clutch kits"}
si = pt.assign(pn=norm(pt.part_number)).merge(cat, on="pn")
SI = si.pivot_table(index="m", columns="fam", values="billed_qty", aggfunc="sum", fill_value=0).reindex(columns=FAMS, fill_value=0)
so = pd.concat([pd.read_csv(DER / f"dealer_sellout_long_PORTUGAL_{p}.csv", dtype={"part_number": str}) for p in ("CY2025", "CY2026YTD")])
so = so[so.period_grain == "month"].copy()
so["pn"] = norm(so.part_number); so["m"] = pd.PeriodIndex(so.period, freq="M")
SO = so.merge(cat, on="pn").pivot_table(index="m", columns="fam", values="qty", aggfunc="sum", fill_value=0).reindex(columns=FAMS, fill_value=0)
cl = pd.read_csv(DER / "pms_campaign_suspect_lines_2024_2025.csv", dtype={"part_number": str, "pn": str})
cl = cl[cl.distributor_code == "PORTUGAL"].copy()
cl["m"] = pd.to_datetime(cl.billing_date).dt.to_period("M"); cl["fam"] = cl.pms_category.map(fam)
CAMP = cl.pivot_table(index="m", columns="fam", values="billed_qty", aggfunc="sum", fill_value=0).reindex(columns=FAMS, fill_value=0)
MONTHS = pd.period_range("2024-01", "2026-08", freq="M")
SI = SI.reindex(MONTHS, fill_value=0); CAMP = CAMP.reindex(MONTHS, fill_value=0)
def fy_of(p): return p.year + (p.month >= 4)
def fysum(X, fy): return X[[fy_of(p) == fy for p in X.index]].sum()

# ------------------------------------------------------------------ pptx helpers
prs = globals().get("PRS") or Presentation()   # PRS is injected when the merged deck build runs this file
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def txt(slide, x, y, w, h, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
    return tb

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def hline(slide, x, y, w, color):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color); c.line.width = Pt(0.75)

def base(section, title, sub):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    txt(s, 0.25, 0.10, 3.8, 0.18, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 6.6, WHITE, True)
    txt(s, 0.25, 0.43, 6.0, 0.22, section.upper(), 7.5, RED, True)
    txt(s, 0.25, 0.68, 9.4, 0.66, title, 20 if len(title) <= 60 else 16.5, WHITE, True)
    txt(s, 0.25, 1.42, 9.4, 0.24, sub, 8.5, MUTED)
    hline(s, 0.25, 7.25, 12.85, "575757")
    return s

def panel(slide, points, x=9.86, y=1.46, w=3.25, h=5.58):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.28, y + 0.26, w - 0.5, 0.2, "OBSERVATIONS / QUESTIONS", 7.5, RED, True)
    yy = y + 0.72
    step = (h - 0.95) / max(len(points), 1)
    for i, pnt in enumerate(points):
        rect(slide, x + 0.28, yy + 0.07, 0.09, 0.09, RED if i == 0 else AMBER, shape=MSO_SHAPE.OVAL)
        txt(slide, x + 0.52, yy, w - 0.8, step - 0.08, pnt, 8.6, WHITE)
        yy += step

def kpi(slide, x, y, big, small, color=WHITE, w=2.2):
    txt(slide, x, y, w, 0.48, big, 24, color, True)
    txt(slide, x, y + 0.50, w, 0.30, small, 8, MUTED)

def style_chart(ch, legend=True, num_fmt='#,##0', labels=False, label_fmt=None):
    ch.font.name = FONT; ch.font.size = Pt(8); ch.font.color.rgb = rgb(MUTED)
    ch.has_title = False
    if legend:
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.color.rgb = rgb(MUTED); ch.legend.font.size = Pt(8)
    else:
        ch.has_legend = False
    try:
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(GRID); va.major_gridlines.format.line.width = Pt(0.5)
        va.format.line.fill.background()
        va.tick_labels.font.color.rgb = rgb(MUTED); va.tick_labels.font.size = Pt(7.5)
        va.tick_labels.number_format = num_fmt; va.tick_labels.number_format_is_linked = False
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb(GRID)
        ca.tick_labels.font.color.rgb = rgb(MUTED); ca.tick_labels.font.size = Pt(7.5)
        ca.has_major_gridlines = False
    except Exception:
        pass
    if labels:
        for pl in ch.plots:
            pl.has_data_labels = True
            dl = pl.data_labels
            dl.font.size = Pt(8); dl.font.color.rgb = rgb(WHITE)
            dl.number_format = label_fmt or num_fmt; dl.number_format_is_linked = False
            try: dl.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception: pass

def color_series(ch, colors, line=False, widths=None, dashes=None):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    for i, s in enumerate(ch.plots[0].series if len(ch.plots) == 1 else [x for p in ch.plots for x in p.series]):
        c = colors[i % len(colors)]
        if line:
            s.format.line.color.rgb = rgb(c); s.format.line.width = Pt((widths or [2.25] * 9)[i])
            s.smooth = False
            if dashes and dashes[i]: s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            try:
                s.marker.style = 8 if not (dashes and dashes[i]) else None
                s.marker.size = 4
                s.marker.format.fill.solid(); s.marker.format.fill.fore_color.rgb = rgb(c)
                s.marker.format.line.color.rgb = rgb(c)
            except Exception: pass
        else:
            s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(c)
            s.format.line.fill.background()

def add_chart(slide, kind, cd, x, y, w, h):
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    return gf.chart

def fmt_k(v): return f"€{v/1000:,.0f}K"
def fmt_m(v): return "€1.858M" if abs(v - TARGET) < 1 else f"€{v/1e6:,.2f}M"
def pct(v): return f"{v*100:.1f}%"

# ------------------------------------------------------------------ slide 1: cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
txt(s, 0.6, 0.5, 6, 0.25, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 8, WHITE, True)
rect(s, 0.6, 2.6, 0.12, 1.5, RED)
txt(s, 0.95, 2.5, 11, 0.8, "I Motors: fleet, genuine parts share and stock movement", 30, WHITE, True)
txt(s, 0.95, 3.35, 11, 0.4, "Supporting slides for the FY2027 business performance review", 14, MUTED)
txt(s, 0.95, 3.85, 11, 0.3, "Fiscal years April to March  ·  Data to 31 August 2026  ·  24 September 2026", 10, MUTED)

# ------------------------------------------------------------------ slide 2: fleet two methods
f24, f25, f26 = [(tot(y, "scrap"), tot(y, "ten")) for y in (2024, 2025, 2026)]
s = base("Fleet", f"Two ways to count the fleet: {tot(2026,'scrap'):,.0f} or {tot(2026,'ten'):,.0f} vehicles in FY2026",
         "I Motors registrations, vehicle level, 2005 to June 2026  |  trucks, F-Series and D-Max  |  fiscal years April to March")
cd = CategoryChartData(); cd.categories = ["FY2024", "FY2025", "FY2026"]
cd.add_series("Scrap-rate model (ISZE method)", [round(tot(y, "scrap")) for y in (2024, 2025, 2026)])
cd.add_series("10-year car park (I Motors method)", [round(tot(y, "ten")) for y in (2024, 2025, 2026)])
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.7, 9.2, 4.3)
style_chart(ch, labels=True); color_series(ch, [GREY, RED]); ch.plots[0].gap_width = 80
d26 = tot(2026, "ten") / tot(2026, "scrap") - 1
kpi(s, 0.6, 6.15, f"{tot(2026,'scrap'):,.0f}", "FY2026, scrap-rate model")
kpi(s, 3.1, 6.15, f"{tot(2026,'ten'):,.0f}", "FY2026, 10-year car park", RED)
kpi(s, 5.6, 6.15, f"+{d26*100:.0f}%", "10-year count above scrap-rate model", AMBER)
panel(s, [
    f"Both counts use the same I Motors registration file. The 10-year car park is {d26*100:.0f}% higher in FY2026 ({tot(2026,'ten')-tot(2026,'scrap'):,.0f} vehicles).",
    "10-year car park: every vehicle registered in the last ten years counts as one, and nothing older counts.",
    "Scrap-rate model: each registration year is reduced by a survival curve (trucks to age 15, D-Max to age 11), so older vehicles still count partly.",
    "Question: which base should both sides use for targets and for genuine parts share?"])

# ------------------------------------------------------------------ slide 3: why 10-year is higher (bridge)
ten, sin, sold = tot(2026, "ten"), tot(2026, "scrap_in10"), tot(2026, "scrap_old")
removed = ten - sin
s = base("Fleet", "Why the 10-year count is higher: scrapping inside the ten years outweighs the older vehicles still counted",
         "FY2026  |  bridge from the 10-year car park to the scrap-rate model  |  vehicles")
cats = ["10-year car park", "Scrapped within the\nlast ten years", "Older than ten years,\nstill on the road", "Scrap-rate model"]
invisible = [0, ten - removed, ten - removed, 0]
visible = [ten, removed, sold, ten - removed + sold]
cd = CategoryChartData(); cd.categories = cats
cd.add_series("base", [round(v) for v in invisible]); cd.add_series("Vehicles", [round(v) for v in visible])
ch = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, cd, 0.4, 1.7, 9.2, 4.9)
style_chart(ch, legend=False); ch.plots[0].gap_width = 60
b0, b1 = ch.plots[0].series
b0.format.fill.background(); b0.format.line.fill.background()
for i, c in enumerate([RED, AMBER, GREY, GREY]):
    pnt = b1.points[i]; pnt.format.fill.solid(); pnt.format.fill.fore_color.rgb = rgb(c)
b1.data_labels.show_value = True
b1.data_labels.font.size = Pt(9); b1.data_labels.font.color.rgb = rgb(WHITE); b1.data_labels.font.bold = True
b1.data_labels.number_format = '#,##0'; b1.data_labels.number_format_is_linked = False
b1.points[1].data_label.text_frame.text = f"−{removed:,.0f}"
b1.points[2].data_label.text_frame.text = " "
txt(s, 5.35, 2.12, 1.8, 0.25, f"+{sold:,.0f}", 9, WHITE, True, PP_ALIGN.CENTER)
for i in (1, 2):
    for r in b1.points[i].data_label.text_frame.paragraphs[0].runs:
        r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = rgb(WHITE)
old_regs = coh[(coh.fy >= 2011) & (coh.fy <= 2016)].n.sum()
recent_regs = coh[(coh.fy >= 2021) & (coh.fy <= 2026)].n.sum()
panel(s, [
    f"The scrap-rate model removes {removed:,.0f} vehicles from the last ten years (for example, a truck aged 10 counts as 0.55, a D-Max as 0.40), but only adds back {sold:,.0f} older vehicles.",
    f"Portugal registered few vehicles in the crisis years: {old_regs:,} in FY2011 to FY2016, against {recent_regs:,} in FY2021 to FY2026. There are few old vehicles left to add back.",
    "In a market with a large older fleet the result would flip, and the scrap-rate model would be the higher of the two.",
    "Neither method sees vehicles that were exported, imported used, or scrapped early. The file counts first registrations only."])

# ------------------------------------------------------------------ slide 4: cohort view
s = base("Fleet", "Each registration year: what the two methods count in FY2026",
         "Registrations by fiscal year and the vehicles each method counts in FY2026  |  FY2026 = April 2025 to March 2026")
yrs = list(range(2008, 2027))
regs_n = [int(coh[coh.fy == y].n.sum()) for y in yrs]
ten_n = [int(coh[coh.fy == y].n.sum()) if y > 2016 else 0 for y in yrs]
scr_n = [round(sum(r.n * surv(r.cls, 2026 - y + 1) for r in coh[coh.fy == y].itertuples())) for y in yrs]
cd = CategoryChartData(); cd.categories = [f"FY{str(y)[2:]}" for y in yrs]
cd.add_series("Registered", regs_n); cd.add_series("Counted: 10-year car park", ten_n); cd.add_series("Counted: scrap-rate model", scr_n)
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.7, 9.3, 5.3)
style_chart(ch); color_series(ch, ["5A5A5A", RED, GREY]); ch.plots[0].gap_width = 40; ch.plots[0].overlap = -10
panel(s, [
    "The 10-year car park counts every vehicle from FY2017 onwards in full and stops there. The scrap-rate model thins out each registration year as it gets older.",
    "The two methods agree on the newest vehicles. They differ most on vehicles aged 6 to 10, which the scrap-rate model already reduces by 5 to 45% (trucks) and 16 to 60% (D-Max).",
    "Vehicles older than ten years only count in the scrap-rate model, and in Portugal there are few of them.",
    "FY2027 registrations (April to June 2026) are in the file but not used here: the fleet is measured at the end of each full fiscal year."])

# ------------------------------------------------------------------ slide 5: GPS both bases
gps = {fy: (BILL[fy] / wallet(fy, "scrap"), BILL[fy] / wallet(fy, "ten")) for fy in (2024, 2025, 2026)}
g27 = (RUN / wallet(2026, "scrap"), RUN / wallet(2026, "ten"))
s = base("Genuine parts share", f"Genuine parts share rose to {gps[2026][0]*100:.0f}% in FY2026 on the scrap-rate model ({gps[2026][1]*100:.0f}% on the 10-year count)",
         "ISZE billing to I Motors, warranty parts included, divided by target spend (fleet × €739 light truck, €887 F-Series, €250 D-Max per year)")
cd = CategoryChartData(); cd.categories = ["FY2024", "FY2025", "FY2026", "FY2027 pace*"]
cd.add_series("On scrap-rate model", [round(gps[y][0], 3) for y in (2024, 2025, 2026)] + [round(g27[0], 3)])
cd.add_series("On 10-year car park", [round(gps[y][1], 3) for y in (2024, 2025, 2026)] + [round(g27[1], 3)])
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.7, 9.2, 4.2)
style_chart(ch, num_fmt='0%', labels=True, label_fmt='0.0%'); color_series(ch, [GREY, RED]); ch.plots[0].gap_width = 80
ch.value_axis.maximum_scale = 0.8; ch.value_axis.minimum_scale = 0
kpi(s, 0.6, 6.05, fmt_m(BILL[2026]), "ISZE billing FY2026")
kpi(s, 3.1, 6.05, fmt_m(wallet(2026, "scrap")), "target spend FY2026, scrap-rate model")
kpi(s, 5.6, 6.05, fmt_m(wallet(2026, "ten")), "target spend FY2026, 10-year car park")
txt(s, 0.4, 6.95, 9.2, 0.25, "* FY2027 pace = April-August 2026 billing × 12/5, measured against the FY2026 fleet (the registration file ends in June 2026).", 7, MUTED)
panel(s, [
    f"Billing grew from {fmt_m(BILL[2024])} (FY2024) to {fmt_m(BILL[2026])} (FY2026), faster than the fleet: +{(gps[2026][0]-gps[2024][0])*100:.0f} points on the scrap-rate model, +{(gps[2026][1]-gps[2024][1])*100:.0f} on the 10-year count.",
    f"The choice of fleet base moves the share by {(gps[2024][0]-gps[2024][1])*100:.0f} to {(gps[2026][0]-gps[2026][1])*100:.0f} points. The direction is the same on both.",
    "Part of FY2026 is brake-pad stock bought during the 2025 campaign (see the brake-pad slide), so FY2026 overstates the underlying level slightly.",
    f"At the April-August pace, FY2027 would land at {g27[0]*100:.0f}% (scrap-rate model), below FY2026."])

# ------------------------------------------------------------------ slide 6: FY2027 pace
mlabels = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
p26 = pd.period_range("2025-04", "2026-03", freq="M"); p27 = pd.period_range("2026-04", "2026-08", freq="M")
s = base("FY2027 pace", f"April-August billing is flat on last year; a straight-line full year gives {pct(RUN/TARGET)} of target",
         "ISZE billing to I Motors, € per month  |  FY2026 against FY2027 to 31 August 2026")
cd = CategoryChartData(); cd.categories = mlabels
cd.add_series("FY2026", [round(MON.get(p, 0)) for p in p26])
cd.add_series("FY2027", [round(MON.get(p, 0)) for p in p27] + [None] * 7)
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.7, 9.2, 4.1)
style_chart(ch, num_fmt='€#,##0,"K"'); color_series(ch, [GREY, RED]); ch.plots[0].gap_width = 60
kpi(s, 0.5, 5.95, fmt_k(YTD), "April-August 2026")
kpi(s, 2.45, 5.95, f"{(YTD/LY-1)*100:+.1f}%", f"against {fmt_k(LY)} a year earlier", AMBER)
kpi(s, 4.4, 5.95, fmt_m(RUN), "full year at the same pace (× 12/5)")
kpi(s, 6.8, 5.95, pct(RUN / TARGET), f"of the {fmt_m(TARGET)} FY2027 target", RED)
share = LY / BILL[2026]
panel(s, [
    f"A straight line (× 12/5) gives {fmt_m(RUN)}, {fmt_k(TARGET-RUN)} short of the {fmt_m(TARGET)} target.",
    f"FY2026 was back-loaded: April-August was only {share*100:.1f}% of the year. With the same seasonality FY2027 would reach {fmt_m(YTD/share)}.",
    f"August was the weak month: {fmt_k(MON[pd.Period('2026-08','M')])} against {fmt_k(MON[pd.Period('2025-08','M')])} a year earlier.",
    "Part of the gap is stock: in FY2027 I Motors buys less than it sells to dealers in most service families, most of all brake pads (see the stock slides)."])

# ------------------------------------------------------------------ slide 7: brake pads
lab = [p.strftime("%b %y") for p in MONTHS]
s = base("Stock movement", "Brake pads: campaign stock from 2025 is being sold through; dealer demand is stable",
         "Units per month  |  bought from ISZE (January 2024 to August 2026) against sold to dealers (January 2025 to August 2026)")
cd = CategoryChartData(); cd.categories = lab
cd.add_series("Bought from ISZE", [int(SI.loc[p, "brake_pad"]) for p in MONTHS])
cd.add_series("Sold to dealers", [int(SO.loc[p, "brake_pad"]) if p in SO.index else None for p in MONTHS])
cd.add_series("of which bought at campaign price", [int(CAMP.loc[p, "brake_pad"]) if CAMP.loc[p, "brake_pad"] > 0 else None for p in MONTHS])
ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.7, 9.2, 4.2)
style_chart(ch); color_series(ch, [GREY, RED, AMBER], line=True, widths=[2.0, 2.5, 1.5])
cs = ch.plots[0].series[2]
cs.format.line.fill.background(); cs.marker.style = 8; cs.marker.size = 9
b26, o26 = fysum(SI, 2026)["brake_pad"], fysum(SO, 2026)["brake_pad"]
b27, o27 = fysum(SI, 2027)["brake_pad"], fysum(SO, 2027)["brake_pad"]
kpi(s, 0.5, 6.05, f"{b26:,.0f} / {o26:,.0f}", "FY2026: bought / sold to dealers")
kpi(s, 3.3, 6.05, f"{b27:,.0f} / {o27:,.0f}", "FY2027 to August: bought / sold", RED)
kpi(s, 6.1, 6.05, f"{camp_pads:,.0f}" if (camp_pads := int(CAMP.loc["2025-05":"2025-09", "brake_pad"].sum())) else "0", "pads bought at campaign price, May-Sep 2025", AMBER)
ly5 = SO.loc["2025-04":"2025-08", "brake_pad"].sum()
panel(s, [
    f"FY2027 brake-pad purchases ({b27:,.0f} units) are well below dealer sales ({o27:,.0f}). I Motors is using up the stock it built in FY2026 ({b26-o26:+,.0f} units over dealer sales).",
    f"Dealer demand did not fall: {o27:,.0f} pads in April-August 2026 against {ly5:,.0f} a year earlier ({(o27/ly5-1)*100:+.0f}%).",
    "The drop in brake-part billing is therefore timing, not business lost to other suppliers. This matches what I Motors said in July.",
    "Question: how many months of brake-pad stock remain, and when do normal orders restart?"])

# ------------------------------------------------------------------ slide 8: oil & fuel filters
s = base("Stock movement", "Oil and fuel filters: large purchases in July 2024 and January 2025, steady dealer demand since",
         "Units per month  |  bought from ISZE against sold to dealers  |  sales to dealer available from January 2025")
for i, (f, x) in enumerate([("oil_filter", 0.4), ("fuel_filter", 5.05)]):
    cd = CategoryChartData(); cd.categories = lab
    cd.add_series("Bought from ISZE", [int(SI.loc[p, f]) for p in MONTHS])
    cd.add_series("Sold to dealers", [int(SO.loc[p, f]) if p in SO.index else None for p in MONTHS])
    ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, x, 1.95, 4.55, 4.6)
    style_chart(ch); color_series(ch, [GREY, RED], line=True, widths=[1.75, 2.25])
    txt(s, x + 0.1, 1.68, 4.3, 0.25, FAM_LABEL[f], 10, WHITE, True)
oil_norm = SO.loc["2025-01":"2026-08", "oil_filter"].mean()
panel(s, [
    f"July 2024: {int(SI.loc['2024-07','oil_filter']):,} oil and {int(SI.loc['2024-07','fuel_filter']):,} fuel filters in one month; January 2025: {int(SI.loc['2025-01','oil_filter']):,} oil filters. A normal month is roughly 200-300.",
    f"Dealers take about {oil_norm:,.0f} oil filters a month, with no break. Since FY2026 I Motors has bought fewer than it sold, so the earlier purchases are still being used up.",
    "The high 2024 filter share and the lower 2025 figure are therefore mostly timing: a stock build and its unwinding, not leakage.",
    "Limit: dealer sales before January 2025 are not in the file, so 2024 demand itself cannot be measured."])

# ------------------------------------------------------------------ slide 9: air filters & rotors (cumulative)
ov = pd.period_range("2025-01", "2026-08", freq="M"); lab2 = [p.strftime("%b %y") for p in ov]
s = base("Stock movement", "Air filters and brake rotors: bought well ahead of dealer sales in FY2026, now slowly unwinding",
         "Cumulative units since January 2025  |  bought from ISZE against sold to dealers")
gaps, peak = {}, {}
for f, x in [("air_filter", 0.4), ("brake_rotor", 5.05)]:
    cin = SI.loc[ov, f].cumsum(); cout = SO.reindex(ov, fill_value=0)[f].cumsum()
    gaps[f] = (int(cin.iloc[-1]), int(cout.iloc[-1]))
    g = cin - cout; peak[f] = (int(g.max()), g.idxmax().strftime("%B %Y"))
    cd = CategoryChartData(); cd.categories = lab2
    cd.add_series("Bought from ISZE (cumulative)", [int(v) for v in cin]); cd.add_series("Sold to dealers (cumulative)", [int(v) for v in cout])
    ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, x, 1.95, 4.55, 4.3)
    style_chart(ch); color_series(ch, [GREY, RED], line=True, widths=[2.25, 2.25])
    txt(s, x + 0.1, 1.68, 4.3, 0.25, f"{FAM_LABEL[f]}: {gaps[f][0]:,} bought, {gaps[f][1]:,} sold ({gaps[f][0]/gaps[f][1]:.2f}×)", 10, WHITE, True)
def r20(f): return SI.loc[ov, f].sum() / SO.reindex(ov, fill_value=0)[f].sum()
txt(s, 0.4, 6.35, 9.2, 0.6, [f"For comparison, over the same 20 months: oil filters {r20('oil_filter'):.2f}×, fuel filters {r20('fuel_filter'):.2f}×, brake pads {r20('brake_pad'):.2f}×. For those, purchases and dealer sales stay in step."], 8.5, MUTED)
a27, o27a = fysum(SI, 2027)["air_filter"], fysum(SO, 2027)["air_filter"]
r27, o27r = fysum(SI, 2027)["brake_rotor"], fysum(SO, 2027)["brake_rotor"]
panel(s, [
    (f"The gap between purchases and dealer sales peaked in {peak['air_filter'][1]} at {peak['air_filter'][0]:,} air filters and {peak['brake_rotor'][0]:,} rotors. Since April 2026 it has been narrowing."
     if peak['air_filter'][1] == peak['brake_rotor'][1] else
     f"The gap between purchases and dealer sales peaked at {peak['air_filter'][0]:,} air filters ({peak['air_filter'][1]}) and {peak['brake_rotor'][0]:,} rotors ({peak['brake_rotor'][1]}). It has been narrowing since."),
    f"FY2027 to August: {a27:,.0f} air filters bought against {o27a:,.0f} sold; {r27:,.0f} rotors bought against {o27r:,.0f} sold. Same pattern as brake pads, on a smaller scale.",
    f"Still open after 20 months: {gaps['air_filter'][0]-gaps['air_filter'][1]:,} air filters and {gaps['brake_rotor'][0]-gaps['brake_rotor'][1]:,} rotors. Either stock on hand, or sold outside the dealer file.",
    "Question for I Motors: do own workshops, fleet customers or other channels buy air filters and rotors directly from the central warehouse, and how many months of stock are on hand?"])

# ------------------------------------------------------------------ slide 10: family table
s = base("Stock movement", "Bought from ISZE against sold to dealers, by service family",
         "Units  |  FY2026 (April 2025 to March 2026) and FY2027 to 31 August 2026")
fams_t = ["brake_pad", "brake_rotor", "brake_shoe", "clutch_kit", "oil_filter", "fuel_filter", "air_filter"]
rows = [["Family", "FY2026 bought", "FY2026 sold", "Ratio", "FY2027 bought", "FY2027 sold", "Ratio", "Reading"]]
reading = {}
for f in fams_t:
    a26, s26, a27, s27 = fysum(SI, 2026)[f], fysum(SO, 2026)[f], fysum(SI, 2027)[f], fysum(SO, 2027)[f]
    r26, r27 = a26 / s26, a27 / s27
    if r26 > 1.15 and r27 < 0.95: rd = "Built in FY2026, running down"
    elif r27 < 0.95: rd = "Running down"
    elif r27 > 1.15: rd = "Buying ahead"
    else: rd = "In step"
    rows.append([FAM_LABEL[f], f"{a26:,.0f}", f"{s26:,.0f}", f"{r26:.2f}×", f"{a27:,.0f}", f"{s27:,.0f}", f"{r27:.2f}×", rd])
gft = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.4), Inches(1.85), Inches(9.2), Inches(0.42 * len(rows)))
gft._element.graphic.graphicData.tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId").text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
tbl = gft.table
widths = [1.5, 1.05, 1.0, 0.8, 1.05, 1.0, 0.8, 2.0]
for j, wd in enumerate(widths): tbl.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb("3A3A3A" if i == 0 else (PANEL if i % 2 else "2C2C2C"))
        tf = c.text_frame; tf.paragraphs[0].text = ""
        r = tf.paragraphs[0].add_run(); r.text = v
        r.font.name = FONT; r.font.size = Pt(9); r.font.bold = (i == 0 or j == 0)
        col = WHITE
        if i > 0 and j in (3, 6):
            x = float(v.replace("×", "")); col = RED if x < 0.95 else (AMBER if x > 1.15 else WHITE)
        r.font.color.rgb = rgb(MUTED if i == 0 else col)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j in (0, 7) else PP_ALIGN.RIGHT
txt(s, 0.4, 5.4, 9.2, 0.5, "Ratio = units bought from ISZE ÷ units sold to dealers. Red: below 0.95 (buying less than dealers take). Amber: above 1.15 (buying ahead of dealer sales).", 8, MUTED)
r27s = {f: fysum(SI, 2027)[f] / fysum(SO, 2027)[f] for f in fams_t}
under = [f for f in fams_t if r27s[f] < 0.95]
panel(s, [
    f"In FY2027 I Motors is buying less than it sells to dealers in {len(under)} of {len(fams_t)} families. This is broad destocking, not only brake pads.",
    f"Brake pads are the extreme case ({r27s['brake_pad']:.2f}×). The others sit between {min(r27s[f] for f in under if f!='brake_pad'):.2f}× and {max(r27s[f] for f in under if f!='brake_pad'):.2f}×.",
    "Brake pads, rotors, clutch kits and air filters were bought well ahead of dealer sales in FY2026. That stock is now being used up.",
    "Question for I Motors: is lower stock a deliberate target for FY2027, and at what level will orders return to the dealer-sales rate?"])

# ------------------------------------------------------------------ slide 11: method
s = base("Method", "How these figures were built",
         "Data used, what was assumed, and what was not")
items = [
    ("Fleet", "I Motors registration file, one line per vehicle, first registrations 2005 to June 2026. Light trucks, F-Series and D-Max are counted separately. Older N-series codes without a tonnage flag are counted as light trucks."),
    ("10-year car park", "All vehicles first registered in the last ten fiscal years, each counted once. Nothing older is counted."),
    ("Scrap-rate model", "Each registration year is reduced with the survival curve used in the ISZE fleet model: trucks 100% in year 1, 55% at age 10, 0% after age 15; D-Max 100% in year 1, 40% at age 10, 0% after age 11."),
    ("Genuine parts share", "All ISZE billing to I Motors in the fiscal year, warranty parts included, divided by target spend: fleet × €739 (light truck), €887 (F-Series), €250 (D-Max) per vehicle per year."),
    ("FY2027 pace", "April to August 2026 billing × 12/5 (rule of three), against the €1.858M FY2027 target and the FY2026 fleet."),
    ("Stock movement", "Units billed by ISZE to I Motors against units I Motors reported as sold to dealers, per part number and month (sales to dealer from January 2025). Parts are grouped into service families. No prices or margins of I Motors are used."),
    ("Campaign price", "Lines billed at least 20% under list price are marked as bought at campaign price."),
]
yy = 1.8
for h, t in items:
    txt(s, 0.4, yy, 2.2, 0.6, h, 10, RED, True)
    txt(s, 2.6, yy, 10.2, 0.7, t, 9.5, WHITE)
    yy += 0.72

if not globals().get("PRS"):
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
print("saved", OUT)
print("fleet FY2026 scrap/ten:", round(tot(2026, 'scrap')), round(tot(2026, 'ten')), "removed", round(removed), "old", round(sold))
print("gps:", {k: (round(v[0], 3), round(v[1], 3)) for k, v in gps.items()}, "fy27", tuple(round(x, 3) for x in g27))
print("ytd", round(YTD), "ly", round(LY), "run", round(RUN), "pct", round(RUN / TARGET, 3), "share", round(share, 3))
print("gaps", gaps)
