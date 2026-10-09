from pathlib import Path
import openpyxl

# Latest Sep closed + Oct reports
files = sorted(Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imotors_sales_check_20261006").glob("*.xlsx"))
for f in files[-3:]:
    print("FILE", f.name)
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    print(" sheets", wb.sheetnames[:15], "...", len(wb.sheetnames))
    # find IMOTORS on last day sheet
    for sn in wb.sheetnames:
        if sn.startswith("30.09") or sn.startswith("31.08") or sn.startswith("05.10") or sn.startswith("06.10") or "IMOTOR" in sn.upper():
            print("  try", sn)
    # scan first sheet headers + IMOTORS
    for sn in wb.sheetnames[:3] + [s for s in wb.sheetnames if "30.09" in s or "09" in s][:2]:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True, max_row=80))
        for i,r in enumerate(rows[:5]):
            print(" ", sn, i, r[:12] if r else None)
        for i,r in enumerate(rows):
            if r and any(str(v).upper().find("IMOTOR")>=0 for v in r if v is not None):
                print("  HIT", sn, i, r[:14])
    wb.close()
