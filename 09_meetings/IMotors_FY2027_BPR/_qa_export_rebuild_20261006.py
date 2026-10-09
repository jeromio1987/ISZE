# -*- coding: utf-8 -*-
"""Export key slides for Sep-closed rebuild QA."""
from __future__ import annotations

import json
import hashlib
from datetime import datetime, timezone
from pathlib import Path

import win32com.client
from pptx import Presentation
from lxml import etree

PPTX = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR"
    r"\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx"
)
QA = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\_qa")
QA.mkdir(exist_ok=True)
CMD = PPTX.parent / "OPEN_I_Motors_BPR_2026-10-06_Sep_closed.cmd"

# 1-based PowerPoint slide numbers
EXPORT = [3, 4, 9, 17, 18, 19, 20]
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"

pp = win32com.client.Dispatch("PowerPoint.Application")
pp.Visible = 1
pres = pp.Presentations.Open(str(PPTX), WithWindow=False)
print("slides", pres.Slides.Count)
taken = []
for n in EXPORT:
    out = QA / f"rebuild_slide_{n:02d}.png"
    pres.Slides(n).Export(str(out), "PNG", 1600, 900)
    print("exported", out.name, out.stat().st_size if out.exists() else 0)
    taken.append(str(out))
pres.Close()
pp.Quit()

# Cheap numeric checks
prs = Presentation(str(PPTX))
s4_sums = {}
for sh in prs.slides[3].shapes:
    if not sh.has_chart:
        continue
    root = etree.fromstring(etree.tostring(sh.chart._chartSpace))
    for ser in root.findall(".//{%s}ser" % C):
        nm = ser.find(".//{%s}tx//{%s}v" % (C, C))
        nm = nm.text if nm is not None else "?"
        vals = [float(n.text) for n in ser.findall("./{%s}val//{%s}pt/{%s}v" % (C, C, C))]
        s4_sums[nm] = round(sum(vals), 1)

hero_blob = " ".join(
    sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame
)
rep_blob = " ".join(
    sh.text_frame.text for sh in prs.slides[18].shapes if sh.has_text_frame
)
comp_blob = " ".join(
    sh.text_frame.text for sh in prs.slides[19].shapes if sh.has_text_frame
)

checks = [
    {"id": "exists", "ok": PPTX.exists(), "detail": str(PPTX)},
    {"id": "min_size", "ok": PPTX.stat().st_size > 200_000, "detail": f"{PPTX.stat().st_size} bytes"},
    {"id": "slide_count_35", "ok": len(prs.slides) == 35, "detail": str(len(prs.slides))},
    {"id": "s4_sales_sum_836", "ok": abs(s4_sums.get("Sales", 0) - 835.7) < 0.2, "detail": str(s4_sums)},
    {"id": "hero_836", "ok": "€836K" in hero_blob or "836K" in hero_blob, "detail": "hero sales tile"},
    {"id": "hero_849", "ok": "€849K" in hero_blob or "849K" in hero_blob, "detail": "hero orders tile"},
    {"id": "replenishment_slide", "ok": "replenish" in rep_blob.lower(), "detail": "slide 19"},
    {"id": "competitive_slide", "ok": "26" in comp_blob and "open" in comp_blob.lower(), "detail": "slide 20"},
    {"id": "gps_63_3", "ok": "63.3%" in comp_blob or "63.3%" in hero_blob, "detail": "locked GPS language present"},
    {"id": "launcher", "ok": CMD.exists() and b"\r\n" in CMD.read_bytes()[:80], "detail": str(CMD)},
    {"id": "pngs_exist", "ok": all(Path(p).exists() and Path(p).stat().st_size > 20_000 for p in taken), "detail": f"{len(taken)} files"},
]

ok = all(c["ok"] for c in checks)
receipt = {
    "schema": "ship_receipt/v2",
    "artifact": str(PPTX),
    "artifact_mtime_ms": int(PPTX.stat().st_mtime * 1000),
    "artifact_sha256": hashlib.sha256(PPTX.read_bytes()).hexdigest(),
    "written_at": datetime.now(timezone.utc).isoformat(),
    "ok": ok,
    "screenshots_taken": len(taken),
    "screenshots_read": 0,
    "screenshots": taken,
    "checks": checks,
    "notes": (
        "Rebuild 2026-10-06: locked Apr-Sep sales/orders to sold-to 519 + Daily Sales Sep; "
        "added replenishment ask (19) and competitive/supply one-pager (20). "
        "Parent agent must image-Read every PNG and bump screenshots_read."
    ),
}
rec = PPTX.with_name(PPTX.name.replace(".pptx", ".SHIP_RECEIPT.json"))
# keep same stem
rec = PPTX.parent / (PPTX.stem + ".SHIP_RECEIPT.json")
rec.write_text(json.dumps(receipt, indent=2), encoding="utf-8")
print("RECEIPT", rec, "ok", ok)
for c in checks:
    print(("PASS" if c["ok"] else "FAIL"), c["id"], c["detail"])
