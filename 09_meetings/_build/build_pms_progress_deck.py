# -*- coding: utf-8 -*-
"""
Build ISZE_PMS_Progress_2026-09-21.pptx

Sources merged (all read, nothing invented):
  - ISZE EU Conf 26 GPS/PMS definition and measure slides (Jerome)
  - ISZE P700 Parts support for scheduled PMS, 18 Aug 2026 (Jerome / FINAL pack)
  - ITUK PMS Project pack, 18 Aug 2026 (Duncan Young) - intervals, parts, KPI, roadmap, BVP list
  - ISZE/ITUK/ISZJ PMS meeting, 18 Aug 2026 - decisions and action items
  - ITUK dealer sell-out CY2025 and ISZE billing (coverage figures as published)
  - Data integrity open list 2026-09-19 (what is still blocked)

House style: navy #15233C, terracotta #C0552D, paper #F7F4EE.
Shared/external language only (no internal nicknames, no file paths).
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\ISZE_PMS_Progress_2026-09-21.pptx")

NAVY = RGBColor(0x15, 0x23, 0x3C)
NAVY2 = RGBColor(0x20, 0x37, 0x59)
TERRA = RGBColor(0xC0, 0x55, 0x2D)
PAPER = RGBColor(0xF7, 0xF4, 0xEE)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6B, 0x62, 0x56)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD3, 0xC8)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREEN = RGBColor(0x1D, 0x9E, 0x75)

HEAD_FONT = "Georgia"
BODY_FONT = "Segoe UI"

W = Inches(13.333)
H = Inches(7.5)
M = Inches(0.62)          # left/right margin
CW = W - 2 * M            # content width

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- primitives
def _rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    return sh


def _text(slide, x, y, w, h, runs=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=4, line_spacing=1.0):
    """runs: list of (text, size, bold, colour, font, indent_level) or str.

    Convenience: when the height is omitted and `runs` is passed in its place,
    the runs are detected and a default box height is applied.
    """
    if runs is None and isinstance(h, (list, str)):
        runs, h = h, None
    if h is None:
        # Estimate a height that fits the runs, so no box is sized blindly.
        est_lines = 0.0
        probe = runs if isinstance(runs, list) else [(runs, 12, False, INK, BODY_FONT, 0)]
        for item in probe:
            txt = item[0] if isinstance(item, (list, tuple)) else str(item)
            size = item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else 12
            chars_per_line = max(12.0, (w / 914400.0) * 96.0 / (size * 0.52))
            est_lines += max(1.0, -(-len(txt) // chars_per_line))
        h = Emu(int((est_lines * 1.3 * 12 + 6) * 12700))
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, 12, False, INK, BODY_FONT, 0)]
    first = True
    for item in runs:
        text, size, bold, colour, font, level = (list(item) + [0])[:6] if len(item) < 6 else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        p.level = level
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = font
    return tb


def _header(slide, kicker, title, subtitle=None):
    """Navy band with kicker + title; optional grey subtitle under it."""
    band_h = Inches(1.02)
    _rect(slide, 0, 0, W, band_h, NAVY)
    _rect(slide, 0, band_h, W, Inches(0.045), TERRA)
    runs = [(kicker.upper(), 9.5, True, RGBColor(0xB9, 0xC4, 0xD6), BODY_FONT, 0),
            (title, 23, False, WHITE, HEAD_FONT, 0)]
    _text(slide, M, Inches(0.14), CW, Inches(0.9), runs, space_after=2)
    if subtitle:
        _text(slide, M, band_h + Inches(0.14), CW, Inches(0.36),
              [(subtitle, 12, False, MUTED, BODY_FONT, 0)])
        return band_h + Inches(0.58)
    return band_h + Inches(0.22)


def _footer(slide, text):
    _rect(slide, M, H - Inches(0.52), CW, Emu(9525), LINE)
    _text(slide, M, H - Inches(0.45), CW, Inches(0.3),
          [(text, 8.5, False, MUTED, BODY_FONT, 0)])


def add_slide():
    return prs.slides.add_slide(BLANK)


# ---------------------------------------------------------------- slides
def slide_cover():
    s = add_slide()
    _rect(s, 0, 0, W, H, NAVY)
    _rect(s, 0, Inches(4.62), W, Inches(0.05), TERRA)
    _text(s, M, Inches(1.5), Inches(10.5),
          [("ISUZU MOTORS EUROPE \u00b7 PARTS", 11, True, RGBColor(0xB9, 0xC4, 0xD6), BODY_FONT, 0)])
    _text(s, M, Inches(1.95), Inches(11.2),
          [("Periodic Maintenance Service", 40, False, WHITE, HEAD_FONT, 0)])
    _text(s, M, Inches(3.15), Inches(10.4),
          [("Parts support and parts capture \u2014 one picture for ISZE, ITUK and IML",
            16, False, RGBColor(0xDD, 0xE4, 0xEF), BODY_FONT, 0)])
    _text(s, M, Inches(4.85), Inches(11.2), [
        ("Phase 1: N75 fleet \u00b7 NPR75 4HK1 and NPR88 4JZ1 \u00b7 40,000 km interval", 12, False, RGBColor(0xC9, 0xD4, 0xE3), BODY_FONT, 0),
        ("Prepared 21 September 2026 \u00b7 for discussion with IML and ITUK", 12, False, RGBColor(0xC9, 0xD4, 0xE3), BODY_FONT, 0),
    ], space_after=6)
    _text(s, M, Inches(6.35), Inches(11.2),
          [("Figures: ITUK dealer sell-out and ISZE billing, calendar 2025, on the IML service schedule.", 9, False, RGBColor(0x93, 0xA3, 0xB8), BODY_FONT, 0)])


def slide_why():
    s = add_slide()
    y = _header(s, "01  Why this matters", "The demand is predictable \u2014 seeing it is the problem",
                "Every Isuzu in service generates parts demand on a known schedule. Today neither ISZE nor the distributors can see how much of it we capture.")
    cards = [
        ("EUR 30.76M", "parts billed in 2025 across the 18-market network"),
        ("55.7%", "of the parts the fleet should need on the service schedule"),
        ("EUR 24.5M", "not captured \u2014 the size of the opportunity"),
    ]
    cw = (CW - Inches(0.3)) / 3
    for i, (val, sub) in enumerate(cards):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y, cw, Inches(1.15), PAPER)
        _rect(s, x, y, cw, Inches(0.05), TERRA if i == 2 else NAVY)
        _text(s, x + Inches(0.16), y + Inches(0.16), cw - Inches(0.32), Inches(0.5),
              [(val, 24, False, TERRA if i == 2 else NAVY, HEAD_FONT, 0)])
        _text(s, x + Inches(0.16), y + Inches(0.66), cw - Inches(0.32), Inches(0.42),
              [(sub, 10, False, MUTED, BODY_FONT, 0)], line_spacing=1.05)

    y2 = y + Inches(1.42)
    _text(s, M, y2, Inches(6.1),
          [("Where the euro sits", 14, True, NAVY, HEAD_FONT, 0)], space_after=8)
    _text(s, M, y2 + Inches(0.36), Inches(6.1), [
        ("Three families \u2014 clutch disc, brake pad and brake shoe \u2014 hold about 46% of the measured component gap.",
         11.5, False, INK, BODY_FONT, 0),
        ("Filters are the easiest to measure, not the largest in euro.", 11.5, False, INK, BODY_FONT, 0),
        ("PMS takes the interval-driven part of that gap and makes it actionable: filters, brake, clutch, urea.", 11.5, False, INK, BODY_FONT, 0),
    ], space_after=7, line_spacing=1.08)

    _text(s, M + Inches(6.5), y2, Inches(5.6),
          [("What PMS adds", 14, True, NAVY, HEAD_FONT, 0)], space_after=8)
    _text(s, M + Inches(6.5), y2 + Inches(0.36), Inches(5.6), [
        ("An interval per part family, from the service schedule \u2014 not an opinion.", 11.5, False, INK, BODY_FONT, 0),
        ("A denominator: units in service \u00d7 services per year.", 11.5, False, INK, BODY_FONT, 0),
        ("A cover % per category that every market can read the same way.", 11.5, False, INK, BODY_FONT, 0),
        ("A named gap in euro, so the commercial answer is a choice and not a guess.", 11.5, False, INK, BODY_FONT, 0),
    ], space_after=7, line_spacing=1.08)
    _footer(s, "ISZE model on the 18-market parts base, calendar 2025. Ceiling is modelled from the service schedule; the flat per-vehicle element is disclosed wherever the ceiling is quoted.")


def slide_two_lenses():
    s = add_slide()
    y = _header(s, "02  Method", "Two lenses, one objective: make the hidden parts potential visible",
                "GPS shows the whole genuine parts picture. PMS makes interval-driven leakage actionable, category by category.")
    colw = (CW - Inches(0.3)) / 2
    blocks = [
        ("GPS \u2014 Genuine Parts Share", NAVY,
         "Broad lens across all genuine part numbers.",
         "Actual genuine sales  \u00f7  estimated parts demand  =  GPS %",
         "Same structure as PMS: actual, estimate, coverage \u2014 applied to the full parts basket."),
        ("PMS \u2014 Periodic Maintenance Service", TERRA,
         "Focused lens on competitive, interval-driven parts.",
         "Actual PMS sales  \u00f7  estimated PMS demand  =  Coverage %",
         "Pinpoints leakage in filters, brake, clutch, urea and other scheduled items."),
    ]
    for i, (title, colour, sub, formula, note) in enumerate(blocks):
        x = M + i * (colw + Inches(0.3))
        _rect(s, x, y, colw, Inches(2.5), PAPER)
        _rect(s, x, y, Inches(0.06), Inches(2.5), colour)
        _text(s, x + Inches(0.24), y + Inches(0.2), colw - Inches(0.5), Inches(0.34),
              [(title, 15, True, colour, HEAD_FONT, 0)])
        _text(s, x + Inches(0.24), y + Inches(0.62), colw - Inches(0.5), Inches(0.3),
              [(sub, 11, False, MUTED, BODY_FONT, 0)])
        _rect(s, x + Inches(0.24), y + Inches(1.05), colw - Inches(0.5), Inches(0.62), WHITE)
        _text(s, x + Inches(0.38), y + Inches(1.22), colw - Inches(0.75), Inches(0.4),
              [(formula, 12, True, colour, BODY_FONT, 0)])
        _text(s, x + Inches(0.24), y + Inches(1.85), colw - Inches(0.5), Inches(0.5),
              [(note, 11, False, INK, BODY_FONT, 0)], line_spacing=1.08)

    y2 = y + Inches(2.78)
    steps = ["Units in service", "Service interval", "Expected demand", "Actual sales", "Cover % and gap"]
    sw = (CW - Inches(0.4)) / 5
    for i, st in enumerate(steps):
        x = M + i * (sw + Inches(0.1))
        _rect(s, x, y2, sw, Inches(0.62), NAVY if i == 4 else WHITE, None if i == 4 else LINE)
        _text(s, x + Inches(0.1), y2 + Inches(0.16), sw - Inches(0.2), Inches(0.34),
              [(st, 11, i == 4, WHITE if i == 4 else NAVY, BODY_FONT, 0)], align=PP_ALIGN.CENTER)
    _footer(s, "Both lenses are read on the same service schedule, so a distributor figure and a dealer figure mean the same thing.")


def slide_measure():
    s = add_slide()
    y = _header(s, "03  Method", "How the cover percentage is built \u2014 worked on ITUK N75",
                "Four inputs, no judgement calls. The denominator is the fleet, the schedule and the number of services a year.")
    # steps
    steps = [
        ("1.  Units in service", "9,359 N75", "N75 units in operation (DVLA 2024), both engine generations. Not the 22,389 full UK park."),
        ("2.  Services a year", "1.25", "50,000 km a year against a 40,000 km filter interval: 50,000 \u00f7 40,000 = 1.25."),
        ("3.  Expected units", "11,699", "9,359 \u00d7 1.25. The same expected quantity applies to oil, fuel and air. This is the denominator."),
    ]
    cw3 = (CW - Inches(0.3)) / 3
    for i, (head, val, body) in enumerate(steps):
        x = M + i * (cw3 + Inches(0.15))
        _rect(s, x, y, cw3, Inches(1.62), PAPER)
        _rect(s, x, y, cw3, Inches(0.05), NAVY)
        _text(s, x + Inches(0.18), y + Inches(0.16), cw3 - Inches(0.36), Inches(0.28),
              [(head, 12, True, NAVY, BODY_FONT, 0)])
        _text(s, x + Inches(0.18), y + Inches(0.5), cw3 - Inches(0.36), Inches(0.42),
              [(val, 26, False, NAVY, HEAD_FONT, 0)])
        _text(s, x + Inches(0.18), y + Inches(0.99), cw3 - Inches(0.36), Inches(0.58),
              [(body, 10, False, MUTED, BODY_FONT, 0)], line_spacing=1.05)

    y2 = y + Inches(1.9)
    _text(s, M, y2, CW, [("Cover %  =  units sold in calendar 2025  \u00f7  units the schedule expects", 14, True, TERRA, BODY_FONT, 0)])
    rows = [
        ("Oil filter", "5,335", "11,699", "45.6%", "4,691", "40.1%"),
        ("Fuel filter", "7,415", "11,699", "63.4%", "5,161", "44.1%"),
        ("Air filter", "2,284", "11,699", "19.5%", "2,204", "18.8%"),
    ]
    add_table(s, M, y2 + Inches(0.45), CW, ("Category", "ITUK to dealer", "Expected", "Cover", "ISZE to ITUK", "Cover"), rows,
              widths=[0.26, 0.16, 0.13, 0.11, 0.19, 0.15], highlight_rows={2})
    _footer(s, "Scope: the 37 N75 part numbers in the ITUK pack. Brake cover uses 0.5 events a year, not 1.25. Air on a 0.5 interval would read 39%. 50,000 km a year is the planning assumption, not a measured figure.")


def add_table(slide, x, y, w, headers, rows, widths=None, highlight_rows=(), font=10.5, head_font=9.5, row_h=Inches(0.3)):
    if widths:
        total = sum(widths)
        cols = [Emu(int(w * (wd / total))) for wd in widths]
    else:
        cols = [Emu(int(w / len(headers)))] * len(headers)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, row_h * (len(rows) + 1))
    tbl = shape.table
    tbl.first_row = True
    for i, cw in enumerate(cols):
        tbl.columns[i].width = cw
    for c, htxt in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = htxt
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if c else PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(head_font)
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.name = BODY_FONT
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri + 1, ci)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER if (ri % 2 == 0 or ri in highlight_rows) else WHITE
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if ci else PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size = Pt(font)
                r.font.bold = ri in highlight_rows
                r.font.color.rgb = TERRA if ci == 0 and ri in highlight_rows else INK
                r.font.name = BODY_FONT
    return shape


def slide_phase1():
    s = add_slide()
    y = _header(s, "04  Phase 1 scope", "N75 first \u2014 one interval, two engine families, one target fleet",
                "N75 is the only segment where both engine families share a 40,000 km interval and the fleet can be reached through existing customer records.")
    rows = [
        ("N75", "NPR75 / 4HK1", "40,000 km", "Phase 1 \u2014 target fleet"),
        ("N75", "NPR88 / 4JZ1", "40,000 km", "Phase 1 \u2014 target fleet"),
        ("N35.125", "NLR87 / RZ4E", "20,000 km", "Phase 2"),
        ("N35.150", "NLR88 / 4JZ1", "Flexible up to 40,000 km", "Phase 2"),
    ]
    add_table(s, M, y, CW, ("Segment", "Model / engine", "Interval basis", "Campaign"), rows,
              widths=[0.16, 0.24, 0.32, 0.28], highlight_rows={0, 1}, row_h=Inches(0.34))

    y2 = y + Inches(1.75)   # sit directly under the table (4 rows + header) and stay clear of it
    cx = M + Inches(0.3)
    colw = CW - Inches(0.6)
    _rect(s, cx, y2, colw, Inches(1.6), PAPER)
    _rect(s, cx, y2, Inches(0.06), Inches(1.6), GREEN)
    _text(s, cx + Inches(0.24), y2 + Inches(0.18), colw - Inches(0.5), Inches(0.3),
          [("Why N75 first", 13, True, GREEN, HEAD_FONT, 0)])
    _text(s, cx + Inches(0.24), y2 + Inches(0.58), colw - Inches(0.5), Inches(0.9), [
        ("Two principal engine families share a 40,000 km interval.", 11, False, INK, BODY_FONT, 0),
        ("The fleet is easier to target through existing customer records.", 11, False, INK, BODY_FONT, 0),
    ], space_after=6, line_spacing=1.06)

    _rect(s, cx, y2 + Inches(1.78), colw, Inches(1.6), PAPER)
    _rect(s, cx, y2 + Inches(1.78), Inches(0.06), Inches(1.6), TERRA)
    _text(s, cx + Inches(0.24), y2 + Inches(1.96), colw - Inches(0.5), Inches(0.3),
          [("Why N35 second", 13, True, TERRA, HEAD_FONT, 0)])
    _text(s, cx + Inches(0.24), y2 + Inches(2.36), colw - Inches(0.5), Inches(0.9), [
        ("Different interval logic, and weaker customer and mileage visibility.", 11, False, INK, BODY_FONT, 0),
        ("That raises the risk of targeting the wrong vehicle, so it follows.", 11, False, INK, BODY_FONT, 0),
    ], space_after=6, line_spacing=1.06)

    _footer(s, "Intervals as agreed with ITUK for the launch, checked against the workshop maintenance schedule for NPR75 / 4HK1.")


def slide_coverage_n75():
    s = add_slide()
    y = _header(s, "05  Coverage", "ITUK N75, calendar 2025 \u2014 distributor and dealer on one ruler",
                "Same method both layers. The dealer view sits above the distributor view, which stock and channel timing can explain in part.")
    bars = [
        ("Oil \u00b7 ISZE to ITUK", 40.1, "40.1%", NAVY),
        ("Oil \u00b7 ITUK to dealer", 45.6, "45.6%", BLUE),
        ("Fuel \u00b7 ISZE to ITUK", 44.1, "44.1%", NAVY),
        ("Fuel \u00b7 ITUK to dealer", 63.4, "63.4%", BLUE),
        ("Air \u00b7 ISZE to ITUK", 18.8, "18.8%", NAVY),
        ("Air \u00b7 ITUK to dealer", 19.5, "19.5%", BLUE),
    ]
    yb = y + Inches(0.1)
    label_w, val_w = Inches(2.55), Inches(0.85)
    track_x = M + label_w
    track_w = CW - label_w - val_w - Inches(0.1)
    for i, (label, pct, txt, colour) in enumerate(bars):
        yy = yb + i * Inches(0.44)
        _text(s, M, yy + Inches(0.04), label_w - Inches(0.12), Inches(0.3),
              [(label, 11, False, INK, BODY_FONT, 0)])
        _rect(s, track_x, yy, track_w, Inches(0.26), RGBColor(0xEE, 0xE9, 0xE0))
        _rect(s, track_x, yy, Emu(int(track_w * pct / 100)), Inches(0.26), colour)
        _text(s, track_x + track_w + Inches(0.1), yy + Inches(0.03), val_w, Inches(0.3),
              [(txt, 11, True, NAVY, BODY_FONT, 0)], align=PP_ALIGN.RIGHT)

    y2 = yb + Inches(2.78)
    _rect(s, M, y2, CW, Inches(1.45), PAPER)
    _rect(s, M, y2, Inches(0.06), Inches(1.45), TERRA)
    _text(s, M + Inches(0.24), y2 + Inches(0.16), CW - Inches(0.5), Inches(1.15), [
        ("Air filter is the largest single leak in the N75 pilot and the reason it is in phase 1.", 12, True, TERRA, BODY_FONT, 0),
        ("Oil and fuel are already moving; the air element is the one most often left off the scheduled service.", 11, False, INK, BODY_FONT, 0),
        ("Brake and clutch sit on a different denominator (events and manual units) and are shown on the next pages.", 11, False, INK, BODY_FONT, 0),
    ], space_after=6, line_spacing=1.06)
    _footer(s, "Expected 11,699 units per filter family: 9,359 N75 in operation at 1.25 services a year. Dealer view is the CY2025 sales-to-dealer extract; distributor view is ISZE billing. Do not add the two.")


def slide_full_uk():
    s = add_slide()
    y = _header(s, "06  Context", "Full UK park \u2014 the network ruler, kept separate from the pilot",
                "Use this when comparing markets or quoting a total gap. The pilot execution stays on N75. The two gaps must never be added.")
    rows = [
        ("Oil filter", "35.1%", "9,818", "27,986", "EUR 141k"),
        ("Fuel filter", "65.1%", "18,228", "27,986", "EUR 127k"),
        ("Air filter", "13.2%", "3,692", "27,986", "EUR 736k"),
        ("Brake", "35.3%", "3,953", "11,194", "EUR 642k"),
        ("Clutch  (provisional)", "7.8%", "411", "5,274", "EUR 393k"),
    ]
    add_table(s, M, y, CW, ("Category", "Cover", "Actual", "Expected", "Gap in euro"), rows,
              widths=[0.3, 0.13, 0.15, 0.16, 0.26], highlight_rows={2})
    y2 = y + Inches(2.55)
    cards = [("22,389", "UK park on this ruler (DVLA)"),
             ("~EUR 2.04M", "total CY2025 gap, five categories"),
             ("EUR 0.65M", "N75 pilot subset \u2014 a separate number")]
    cw = (CW - Inches(0.3)) / 3
    for i, (val, sub) in enumerate(cards):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y2, cw, Inches(1.05), PAPER)
        _rect(s, x, y2, cw, Inches(0.05), TERRA if i == 1 else NAVY)
        _text(s, x + Inches(0.18), y2 + Inches(0.14), cw - Inches(0.36), Inches(0.4),
              [(val, 20, False, TERRA if i == 1 else NAVY, HEAD_FONT, 0)])
        _text(s, x + Inches(0.18), y2 + Inches(0.6), cw - Inches(0.36), Inches(0.42),
              [(sub, 10, False, MUTED, BODY_FONT, 0)], line_spacing=1.05)
    _footer(s, "Clutch cover is provisional: the denominator assumes a manual share of the park that is not yet measured. Air filter is the largest single euro leak on the full park as well.")


def slide_bvp():
    s = add_slide()
    y = _header(s, "07  Genuine and BVP", "ITUK filter demand \u2014 genuine is not always the volume part",
                "ITUK year-to-date demand across both N75 engine slides. The picture matters for the price route, and it differs by family.")
    rows = [
        ("Oil filter", "1,338", "1,693", "3,031", "55.9%"),
        ("Fuel filter", "3,909", "499", "4,408", "11.3%"),
        ("Air filter", "1,287", "0", "1,287", "0.0%"),
    ]
    add_table(s, M, y, CW, ("Family", "Genuine", "BVP", "Total", "BVP share"), rows,
              widths=[0.24, 0.19, 0.19, 0.19, 0.19], highlight_rows={2})
    y2 = y + Inches(1.8)
    _rect(s, M, y2, CW, Inches(1.85), PAPER)
    _rect(s, M, y2, Inches(0.06), Inches(1.85), NAVY)
    _text(s, M + Inches(0.28), y2 + Inches(0.2), CW - Inches(0.6), Inches(1.4), [
        ("Oil: the BVP already outsells genuine, so there is nothing to win back on price.", 12, True, NAVY, BODY_FONT, 0),
        ("Fuel: genuine holds the volume; the BVP is a single part number with almost no demand.", 12, True, NAVY, BODY_FONT, 0),
        ("Air: the BVP has never sold. Air is also where the coverage gap is largest \u2014 so the answer there has to be made with the dealer, not assumed.", 12, True, TERRA, BODY_FONT, 0),
    ], space_after=8, line_spacing=1.06)
    _footer(s, "ITUK's own figures, year-to-date demand across NPR75 4HK1 and NPR88 4JZ1. Read together with the coverage table, not instead of it.")


def slide_price_route():
    s = add_slide()
    y = _header(s, "08  Price route", "Lower of genuine or BVP \u2014 and an SPA only where price is the barrier",
                "Genuine is proposed first. A BVP applies where genuine demand is quiet or the BVP is cheaper. Where neither works, a Special Price Agreement is the route.")
    cards = [
        ("Fuel", "Genuine demand is established.", "Retain genuine. A BVP equivalent is available if required."),
        ("Air", "Coverage is about 20%.", "BVP 5876102500, or an SPA on genuine where that is lower."),
        ("Oil", "The BVP is already supplied.", "Remains in the set. Not a separate campaign."),
        ("Pad", "No BVP is available.", "SPA on genuine only, and only where the repair requires the pad."),
    ]
    cw = (CW - Inches(0.45)) / 4
    for i, (head, sub, body) in enumerate(cards):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y, cw, Inches(1.95), PAPER)
        _rect(s, x, y, cw, Inches(0.05), TERRA if i in (1, 3) else NAVY)
        _text(s, x + Inches(0.18), y + Inches(0.16), cw - Inches(0.36), Inches(0.3),
              [(head, 15, True, NAVY, HEAD_FONT, 0)])
        _text(s, x + Inches(0.18), y + Inches(0.56), cw - Inches(0.36), Inches(0.36),
              [(sub, 10.5, True, MUTED, BODY_FONT, 0)], line_spacing=1.05)
        _text(s, x + Inches(0.18), y + Inches(1.0), cw - Inches(0.36), Inches(0.85),
              [(body, 10.5, False, INK, BODY_FONT, 0)], line_spacing=1.06)

    y2 = y + Inches(2.2)
    _text(s, M, y2, CW, [("One scheduled service, three filters \u2014 ISZE to ITUK, EUR", 13, True, NAVY, HEAD_FONT, 0)])
    rows = [
        ("Genuine set", "8983282070", "8975425400", "8981772710", "58.37"),
        ("BVP alternative", "1876102890", "1876102860", "5876102500", "35.50"),
    ]
    add_table(s, M, y2 + Inches(0.42), CW, ("", "Oil", "Fuel", "Air", "Set"), rows,
              widths=[0.26, 0.19, 0.19, 0.19, 0.17], highlight_rows={1})
    _footer(s, "The air BVP is 44% below the genuine element. Prices are ISZE to ITUK at the current list. A BVP applies only where genuine is not taken up, or is not the lower price.")


def slide_brakes():
    s = add_slide()
    y = _header(s, "09  Focus category", "Brakes \u2014 condition-based, and the main commercial blockage",
                "Brake price competitiveness is the dealer complaint that matters most. No schedule sets a pad replacement interval: brakes are replaced when the inspection says so.")
    left_w = CW * 0.52
    _text(s, M, y, left_w, Inches(0.32), [("What is agreed", 13, True, NAVY, HEAD_FONT, 0)])
    items = [
        ("Inspection every 20,000 km or 12 months \u2014 no promotional trigger, no volume pressure.", 0),
        ("Both N75 pad numbers are in scope: 8982447940 and 8982169220.", 0),
        ("ITUK states its margin on N75 pads is about 18% after logistics, with no further room on its side.", 0),
        ("So the support has to come from ISZE / IML, initially as an SPA on genuine.", 0),
        ("An acceptable gap against the independent aftermarket is roughly 15\u201320%; today the gap is wider.", 0),
    ]
    runs = []
    for txt, lvl in items:
        runs.append(("\u2022   " + txt, 11, False, INK, BODY_FONT, lvl))
    _text(s, M, y + Inches(0.4), left_w - Inches(0.2), Inches(3.2), runs, space_after=8, line_spacing=1.07)

    x2 = M + left_w + Inches(0.25)
    right_w = CW - left_w - Inches(0.25)
    _text(s, x2, y, right_w, Inches(0.32), [("Timing", 13, True, NAVY, HEAD_FONT, 0)])
    timeline = [
        ("Now", "SPA on genuine brake parts, until a BVP exists.", TERRA),
        ("December 2026", "BVP brake production in Japan.", NAVY),
        ("Around Q1 2027", "BVP realistically available to the market.", NAVY),
    ]
    for i, (when, what, colour) in enumerate(timeline):
        yy = y + Inches(0.45) + i * Inches(0.85)
        _rect(s, x2, yy, right_w, Inches(0.72), PAPER)
        _rect(s, x2, yy, Inches(0.06), Inches(0.72), colour)
        _text(s, x2 + Inches(0.22), yy + Inches(0.1), right_w - Inches(0.45), Inches(0.26),
              [(when, 11.5, True, colour, BODY_FONT, 0)])
        _text(s, x2 + Inches(0.22), yy + Inches(0.36), right_w - Inches(0.45), Inches(0.3),
              [(what, 10.5, False, INK, BODY_FONT, 0)])

    y3 = y + Inches(3.05)
    _rect(s, M, y3, CW, Inches(1.15), PAPER)
    _rect(s, M, y3, Inches(0.06), Inches(1.15), NAVY)
    _text(s, M + Inches(0.26), y3 + Inches(0.14), CW - Inches(0.55), Inches(0.9), [
        ("Proposal to IML: add the N75 brake rotors (8981919560 rear, 8982488951 front) to the R90 brake pack.", 11.5, True, NAVY, BODY_FONT, 0),
        ("They fit the N75 NPR, carry over to P700, sell steadily, and come only through ISZE \u2014 but they are not on the ITUK PMS list today.", 10.5, False, INK, BODY_FONT, 0),
    ], space_after=5, line_spacing=1.05)
    _footer(s, "Pad quantities ITUK to dealer 2,900 and ISZE to ITUK 2,730 across both numbers. Inspection interval from the N-series maintenance schedule for NPR75 / 4HK1.")


def slide_clutch():
    s = add_slide()
    y = _header(s, "10  Focus category", "Clutch \u2014 supplied for manual transmissions only, and the manual share is not measured",
                "No broad campaign. The kit goes to identified manual units. The blocker is that nobody has a transmission split we can stand behind.")
    colw = (CW - Inches(0.3)) / 2
    _rect(s, M, y, colw, Inches(2.55), PAPER)
    _rect(s, M, y, Inches(0.06), Inches(2.55), NAVY)
    _text(s, M + Inches(0.26), y + Inches(0.18), colw - Inches(0.55), Inches(0.32),
          [("What we sell now", 13, True, NAVY, HEAD_FONT, 0)])
    _text(s, M + Inches(0.26), y + Inches(0.58), colw - Inches(0.55), Inches(1.85), [
        ("Kit 5876102320 \u2014 one part number, ISZE to ITUK EUR 194.69.", 11, False, INK, BODY_FONT, 0),
        ("Genuine three pieces if the workshop needs them: disc 8981649171, pressure plate 8970388314, release bearing 1313100120 \u2014 EUR 282.08.", 11, False, INK, BODY_FONT, 0),
        ("The kit is about 31% lower than the three pieces separately.", 11, False, INK, BODY_FONT, 0),
        ("Disc 8981649181 is not to be included.", 11, False, INK, BODY_FONT, 0),
    ], space_after=7, line_spacing=1.06)

    x2 = M + colw + Inches(0.3)
    _rect(s, x2, y, colw, Inches(2.55), PAPER)
    _rect(s, x2, y, Inches(0.06), Inches(2.55), TERRA)
    _text(s, x2 + Inches(0.26), y + Inches(0.18), colw - Inches(0.55), Inches(0.32),
          [("Why the % cannot be trusted yet", 13, True, TERRA, HEAD_FONT, 0)])
    _text(s, x2 + Inches(0.26), y + Inches(0.58), colw - Inches(0.55), Inches(1.85), [
        ("Two different manual-share assumptions are in circulation: about 47% of the UK park, and an indication nearer 20% from the UK side.", 11, False, INK, BODY_FONT, 0),
        ("The chassis number does not carry transmission, and the UK vehicle list has no manual / automatic column either.", 11, False, INK, BODY_FONT, 0),
        ("So a clutch cover can read over 100% in one market and under 5% in another, purely because automatic vehicles sit in the denominator.", 11, False, INK, BODY_FONT, 0),
    ], space_after=7, line_spacing=1.06)

    y3 = y + Inches(2.85)
    _rect(s, M, y3, CW, Inches(1.15), PAPER)
    _rect(s, M, y3, Inches(0.06), Inches(1.15), TERRA)
    _text(s, M + Inches(0.26), y3 + Inches(0.15), CW - Inches(0.55), Inches(0.9), [
        ("What would close it", 12.5, True, TERRA, BODY_FONT, 0),
        ("A manual / automatic split by model and market from sales or shipment history \u2014 or a standard manual share by model family if a vehicle-level extract cannot be shared. Either one makes the clutch cover a measured number instead of a directional one.", 11, False, INK, BODY_FONT, 0),
    ], space_after=6, line_spacing=1.06)
    _footer(s, "Clutch is kept as a light-touch annual item rather than a phase-1 campaign. The same transmission question blocks the clutch denominator in every market, not only the UK.")


def slide_urea():
    s = add_slide()
    y = _header(s, "11  Parts in scope", "Urea filter is retained \u2014 but kept out of the campaign baseline",
                "The item stays in the pack at ITUK's request. It is excluded from the baseline metrics, because recent rework distorts the demand.")
    rows = [
        ("Genuine", "8983507161", "541", "506", "EUR 70.53"),
        ("BVP", "1876102341", "34", "46", "EUR 30.72"),
    ]
    add_table(s, M, y, CW, ("", "Part number", "ITUK to dealer", "ISZE to ITUK", "ISZE to ITUK EUR"), rows,
              widths=[0.15, 0.22, 0.2, 0.2, 0.23])
    y2 = y + Inches(1.65)
    cards = [
        ("Retained", "The DEF pump filter kit stays in the parts support set. It is the filter kit, not bulk fluid.", NAVY),
        ("Not a baseline metric", "About 6,000 vehicles had recent AdBlue rework, which inflates demand and would flatter any coverage figure.", TERRA),
        ("Interval", "The AdBlue filter runs on a much longer interval than the oil and fuel filters \u2014 it is not an every-service item.", NAVY),
    ]
    cw = (CW - Inches(0.3)) / 3
    for i, (head, body, colour) in enumerate(cards):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y2, cw, Inches(1.7), PAPER)
        _rect(s, x, y2, cw, Inches(0.05), colour)
        _text(s, x + Inches(0.18), y2 + Inches(0.18), cw - Inches(0.36), Inches(0.3),
              [(head, 13, True, colour, HEAD_FONT, 0)])
        _text(s, x + Inches(0.18), y2 + Inches(0.58), cw - Inches(0.36), Inches(1.0),
              [(body, 10.5, False, INK, BODY_FONT, 0)], line_spacing=1.06)
    _footer(s, "Quantities and prices from the ISZE parts support pack, ITUK to dealer and ISZE to ITUK, calendar 2025 basis.")


def slide_kpi():
    s = add_slide()
    y = _header(s, "12  Measurement", "What we will report \u2014 four blocks, agreed with ITUK",
                "One set of numbers for the campaign, so progress is visible without a second interpretation of it.")
    cols = [
        ("Commercial", ["Selected PMS quantity and value", "Dealer baseline uplift", "Air-filter inclusion rate", "Complete filter-pack rate", "Customer retention / next reminder"], TERRA),
        ("Contract maintenance", ["Penetration of registrations", "Active vehicles on contract", "Contract parts turnover", "Contract parts margin"], NAVY),
        ("Availability and quality", ["Parts availability / fill rate", "Back-order status", "Correct part by vehicle"], NAVY),
        ("Connected (future)", ["Vehicles with consented data", "Actionable odometer match", "Due reminders issued", "Parts pre-orders", "Missed-event reduction"], NAVY),
    ]
    cw = (CW - Inches(0.45)) / 4
    for i, (head, items, colour) in enumerate(cols):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y, cw, Inches(3.0), PAPER)
        _rect(s, x, y, cw, Inches(0.05), colour)
        _text(s, x + Inches(0.18), y + Inches(0.18), cw - Inches(0.36), Inches(0.5),
              [(head, 12.5, True, colour, HEAD_FONT, 0)], line_spacing=1.02)
        runs = []
        for it in items:
            runs.append(("\u2022  " + it, 10.5, False, INK, BODY_FONT, 0))
        _text(s, x + Inches(0.18), y + Inches(0.75), cw - Inches(0.36), Inches(2.1), runs,
              space_after=7, line_spacing=1.05)
    _footer(s, "Baseline per dealership is 2025 parts sales against units in service for the selected part numbers, so the uplift is measured against the dealer's own starting point.")


def slide_roadmap():
    s = add_slide()
    y = _header(s, "13  Plan", "From agreed scope to launch, then a scale decision on evidence",
                "The October go-live is on N75. N35 follows only where the data is reliable enough to target the right vehicle.")
    steps = [
        ("18 Aug", "Japan alignment", "Scope, target, support, KPI, BVP questions and roles agreed."),
        ("Sep", "Launch build", "Dealer baselines and targets, part packs, stock rules, one-page guide."),
        ("Oct 2026", "Go-live", "N75 fleet uptime PMS campaign launched."),
        ("Nov\u2013Dec", "Pilot review", "Conversion, availability, dealer performance, customer feedback."),
        ("FY27", "Scale decision", "N35 expansion where the data is reliable."),
    ]
    cw = (CW - Inches(0.6)) / 5
    for i, (when, head, body) in enumerate(steps):
        x = M + i * (cw + Inches(0.15))
        _rect(s, x, y, cw, Inches(1.85), PAPER if i != 2 else NAVY)
        _rect(s, x, y, cw, Inches(0.05), TERRA)
        _text(s, x + Inches(0.16), y + Inches(0.16), cw - Inches(0.32), Inches(0.28),
              [(when, 11, True, TERRA, BODY_FONT, 0)])
        _text(s, x + Inches(0.16), y + Inches(0.46), cw - Inches(0.32), Inches(0.28),
              [(head, 12.5, True, WHITE if i == 2 else NAVY, HEAD_FONT, 0)])
        _text(s, x + Inches(0.16), y + Inches(0.82), cw - Inches(0.32), Inches(0.95),
              [(body, 10, False, RGBColor(0xD8, 0xDF, 0xEA) if i == 2 else INK, BODY_FONT, 0)],
              line_spacing=1.05)

    y2 = y + Inches(2.15)
    roles = [
        ("ITUK", "Dealer execution, baselines, stock rules, service workflow, dashboard."),
        ("ISZE", "Commercial support, incentive mechanism, BVP supply support, category analysis."),
        ("IML / ISZJ", "Service schedule alignment, brake BVP timing, connected data requirements."),
    ]
    for i, (who, what) in enumerate(roles):
        yy = y2 + i * Inches(0.55)
        _rect(s, M, yy, CW, Inches(0.48), WHITE, LINE)
        _rect(s, M, yy, Inches(1.45), Inches(0.48), NAVY)
        _text(s, M + Inches(0.14), yy + Inches(0.11), Inches(1.2), Inches(0.28),
              [(who, 11, True, WHITE, BODY_FONT, 0)])
        _text(s, M + Inches(1.62), yy + Inches(0.12), CW - Inches(1.8), Inches(0.28),
              [(what, 11, False, INK, BODY_FONT, 0)])
    _footer(s, "Dates as agreed at the joint meeting on 18 August 2026.")


def slide_asks():
    s = add_slide()
    y = _header(s, "14  What we need", "Three data items would move this from directional to measured",
                "Everything else is in place. These are the gaps that decide whether a coverage percentage can be quoted as a fact.")
    rows = [
        ("Transmission split (manual / automatic)", "by model and market, from sales or shipment history \u2014 or a standard share by model family",
         "IML / ITUK", "Unlocks the clutch denominator in every market"),
        ("Measured annual distance per model family", "actual km a year, not a planning assumption",
         "ITUK / IML", "Governs the services-a-year figure; a wide swing today"),
        ("Sales-to-dealer extract for 2026 to date", "same columns as the calendar 2025 file: dealer, part number, quantity, value, period",
         "ITUK, then other distributors", "Keeps the dealer layer current instead of frozen at 2025"),
    ]
    add_table(s, M, y, CW, ("What", "In what shape", "Owner", "Why it matters"), rows,
              widths=[0.26, 0.34, 0.16, 0.24], row_h=Inches(0.58), font=10, head_font=9.5)

    y2 = y + Inches(2.55)
    _rect(s, M, y2, CW, Inches(1.25), PAPER)
    _rect(s, M, y2, Inches(0.06), Inches(1.25), NAVY)
    _text(s, M + Inches(0.26), y2 + Inches(0.16), CW - Inches(0.55), Inches(1.0), [
        ("Also open on the parts model, and worth one line each from IML:", 11.5, True, NAVY, BODY_FONT, 0),
        ("the air-to-oil replacement ratio, and the brake, rotor and clutch replacement rates \u2014 today these are internal assumptions rather than confirmed schedule values.", 11, False, INK, BODY_FONT, 0),
    ], space_after=6, line_spacing=1.06)
    _footer(s, "Until the transmission split exists, any clutch percentage in this pack is directional only. It is labelled as such wherever it appears.")


def slide_close():
    s = add_slide()
    _rect(s, 0, 0, W, H, NAVY)
    _rect(s, 0, Inches(1.35), W, Inches(0.05), TERRA)
    _text(s, M, Inches(0.75), CW, [("15  Decisions", 11, True, RGBColor(0xB9, 0xC4, 0xD6), BODY_FONT, 0)])
    _text(s, M, Inches(1.7), Inches(11.5), [("What we need agreed in the room", 30, False, WHITE, HEAD_FONT, 0)])
    items = [
        ("1", "Phase 1 stays N75, brake and air first, clutch light-touch."),
        ("2", "Brake support route: SPA on genuine now, with the brake BVP timing confirmed."),
        ("3", "The transmission split is raised with vehicle sales as a formal request, not a side note."),
        ("4", "Sales-to-dealer extract for 2026 to date, same shape as 2025."),
        ("5", "Air filter: promote inclusion when the scheduled menu calls for it, not on a forced interval."),
    ]
    y = Inches(2.75)
    for num, txt in items:
        _rect(s, M, y, Inches(0.42), Inches(0.42), TERRA)
        _text(s, M, y + Inches(0.09), Inches(0.42), Inches(0.3),
              [(num, 13, True, WHITE, BODY_FONT, 0)], align=PP_ALIGN.CENTER)
        _text(s, M + Inches(0.62), y + Inches(0.07), Inches(11.0), Inches(0.34),
              [(txt, 13, False, RGBColor(0xE6, 0xEB, 0xF3), BODY_FONT, 0)])
        y += Inches(0.62)
    _text(s, M, Inches(6.25), Inches(11.5),
          [("This is not a one-off analysis. It is the start of one shared parts opportunity model, refined with every market that joins.", 11, False, RGBColor(0xA9, 0xB7, 0xCB), BODY_FONT, 0)])


slide_cover()
slide_why()
slide_two_lenses()
slide_measure()
slide_phase1()
slide_coverage_n75()
slide_full_uk()
slide_bvp()
slide_price_route()
slide_brakes()
slide_clutch()
slide_urea()
slide_kpi()
slide_roadmap()
slide_asks()
slide_close()

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print("saved:", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
