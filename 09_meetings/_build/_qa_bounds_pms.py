from pptx import Presentation
from pptx.util import Emu
p = r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\ISZE_PMS_Progress_2026-09-21.pptx"
prs = Presentation(p)
W, H = prs.slide_width, prs.slide_height
print("slide size in:", round(W/914400,2), round(H/914400,2))
bad = 0
for i, s in enumerate(prs.slides, 1):
    n = len(s.shapes)
    over = []
    for sh in s.shapes:
        if sh.left is None: continue
        r = sh.left + (sh.width or 0)
        b = sh.top + (sh.height or 0)
        if sh.left < -9144 or sh.top < -9144 or r > W + 9144 or b > H + 9144:
            over.append((sh.shape_type, sh.name, round(sh.left/914400,2), round(sh.top/914400,2),
                         round(r/914400,2), round(b/914400,2)))
    if over:
        bad += 1
        print("slide", i, "shapes", n, "OUT OF BOUNDS:")
        for o in over: print("   ", o)
    else:
        print("slide", i, "shapes", n, "ok")
print("slides with overflow:", bad)
