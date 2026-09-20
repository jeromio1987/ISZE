"""Build ISZE_PMS_1pct_all_parts_ITUK_DLR.xlsx (shared, formulas in cells).

Study A: 1.00 percentage point of calendar 2025 ITUK-to-dealer sales
on the Duncan PMS part-number set (all scheduled-service families).
Study B: brake facts (pads, discs, R90 8982447940, BVP, cover vs 9359 x 0.5).

Does not edit any PowerPoint. Does not send mail.
"""
from __future__ import annotations

import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
sys.path.insert(0, str(ISZE / "03_analytics" / "Isuzu" / "prism"))

from config import DERIVED_DIR, SALES_FACT_PARQUET  # noqa: E402

DUNCAN = (
    ISZE
    / "01_data"
    / "90_INBOX"
    / "IML_ITUK_Meeting_20260813"
    / "N75_N35_Part_Numbers_-_PMS_Project.xlsx"
)
DEALER_CSV = DERIVED_DIR / "dealer_sellout_ituk_cy2025_by_pn.csv"
XREF_CSV = DERIVED_DIR / "genuine_bvp_crossref.csv"
OUT_MEET = ISZE / "09_meetings" / "ISZE_PMS_1pct_all_parts_ITUK_DLR.xlsx"
OUT_ROOT = ISZE / "ISZE_PMS_1pct_all_parts_ITUK_DLR.xlsx"

NAVY = "15233C"
PAPER = "F7F4EE"
TERR = "C0552D"
YELLOW = "FFF2CC"
GREEN = "E2EFDA"
GREY = "F2F2F2"
INK = "222222"
WHITE = "FFFFFF"
THIN = Border(
    left=Side(style="thin", color="C5C9D1"),
    right=Side(style="thin", color="C5C9D1"),
    top=Side(style="thin", color="C5C9D1"),
    bottom=Side(style="thin", color="C5C9D1"),
)
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_PAPER = PatternFill("solid", fgColor=PAPER)
FILL_YEL = PatternFill("solid", fgColor=YELLOW)
FILL_GRN = PatternFill("solid", fgColor=GREEN)
FILL_GREY = PatternFill("solid", fgColor=GREY)
FILL_TERR = PatternFill("solid", fgColor=TERR)
FONT_H = Font(name="Calibri", size=14, bold=True, color=WHITE)
FONT_SH = Font(name="Calibri", size=11, bold=True, color=NAVY)
FONT_TH = Font(name="Calibri", size=10, bold=True, color=WHITE)
FONT_B = Font(name="Calibri", size=10, color=INK)
FONT_SM = Font(name="Calibri", size=9, color="555555")
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")

MODEL_SHEETS = {
    "NLR87   N35.125   RZ4E": ("N35", "NLR87 N35.125 RZ4E"),
    "NLR88  N35.150  4JZ1": ("N35", "NLR88 N35.150 4JZ1"),
    "NPR75  N75  4HK1": ("N75", "NPR75 N75 4HK1"),
    "NPR88  N75  4JZ1": ("N75", "NPR88 N75 4JZ1"),
}

# IML / coverage N75 locked integers (April 2026 study / 15 Aug 2026 coverage file)
IML_N75_DLR = {"oil": 5404, "fuel": 11139, "air": 2289, "clutch": 340, "brake": 2923}
IML_N75_ISZE = {"oil": 4673, "fuel": 11021, "air": 2226, "clutch": 243, "brake": 2902}
IML_AVG_EUR = {"oil": 7.74, "fuel": 13.06, "air": 30.30, "clutch": 80.72, "brake": 88.63}
IML_FULL_ISZE = {"oil": 9818, "fuel": 18228, "air": 3692, "clutch": 411, "brake": 3953}


def norm_pn(val) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    digits = re.sub(r"\D", "", str(val).strip())
    if len(digits) < 8:
        return None
    return digits.zfill(10)


def family_from_duncan(cat: str) -> str:
    c = (cat or "").lower()
    if "oil" in c:
        return "oil"
    if "fuel" in c or "sedimenter" in c or "cartridge" in c:
        return "fuel"
    if "air" in c:
        return "air"
    if "brake" in c or "pad" in c or "shoe" in c:
        return "brake"
    if "clutch" in c:
        return "clutch"
    return "other_pms"


def family_label(fam: str) -> str:
    return {
        "oil": "Oil filter",
        "fuel": "Fuel filter",
        "air": "Air filter",
        "brake": "Brake",
        "clutch": "Clutch",
        "other_pms": "Other PMS",
    }.get(fam, fam)


def load_duncan() -> pd.DataFrame:
    wb = openpyxl.load_workbook(DUNCAN, read_only=True, data_only=True)
    rows = []
    for sheet, (series, model_label) in MODEL_SHEETS.items():
        ws = wb[sheet]
        header_idx = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            vals = [str(x).strip() if x is not None else "" for x in row]
            if any("PMS category" in v for v in vals):
                header_idx = i
                continue
            if header_idx is None:
                continue
            if row is None or all(x is None or str(x).strip() == "" for x in row[1:4]):
                continue
            cat = row[1]
            pn = norm_pn(row[2])
            if not pn:
                continue
            rows.append(
                {
                    "part_number": pn,
                    "duncan_category": str(cat).strip() if cat else "",
                    "family": family_from_duncan(str(cat) if cat else ""),
                    "description_duncan": str(row[3]).strip() if row[3] else "",
                    "campaign_priority": str(row[4]).strip() if row[4] else "",
                    "duncan_monthly_avg": row[5],
                    "duncan_ytd_demand": row[6],
                    "duncan_comment": str(row[7]).strip() if len(row) > 7 and row[7] else "",
                    "series": series,
                    "model_sheet": model_label,
                    "is_bvp_duncan": "bvp" in str(cat).lower(),
                }
            )
    wb.close()
    long = pd.DataFrame(rows)
    # one row per PN: collapse model flags
    g = (
        long.groupby("part_number", as_index=False)
        .agg(
            duncan_category=("duncan_category", "first"),
            family=("family", "first"),
            description_duncan=("description_duncan", "first"),
            campaign_priority=("campaign_priority", lambda s: " | ".join(sorted(set(x for x in s if x)))),
            duncan_comment=("duncan_comment", lambda s: " | ".join(sorted(set(x for x in s if x)))),
            is_bvp_duncan=("is_bvp_duncan", "any"),
            on_n35=("series", lambda s: int((s == "N35").any())),
            on_n75=("series", lambda s: int((s == "N75").any())),
            model_sheets=("model_sheet", lambda s: " | ".join(sorted(set(s)))),
            duncan_row_count=("part_number", "size"),
        )
    )
    return long, g


def load_dealer() -> pd.DataFrame:
    d = pd.read_csv(DEALER_CSV, dtype={"part_number": str})
    d["part_number"] = d["part_number"].map(norm_pn)
    d["qty"] = pd.to_numeric(d["qty"], errors="coerce").fillna(0)
    return d


def load_isze_ituk() -> pd.DataFrame:
    sales = pd.read_parquet(SALES_FACT_PARQUET)
    s = sales[sales["year"] == 2025].copy()
    sap = pd.to_numeric(s.get("sold_to_sap"), errors="coerce")
    name = s["sold_to"].astype(str).str.upper()
    is_ituk = (sap == 2575) | name.str.contains("ISUZU TRUCK", na=False)
    # keep IUK (2574 / I.M. PARTS) out
    is_iuk = (sap == 2574) | name.str.contains("I.M. PARTS", na=False)
    s = s[is_ituk & ~is_iuk]
    s["part_number"] = s["part_number"].map(norm_pn)
    s = s.dropna(subset=["part_number"])
    g = s.groupby("part_number", as_index=False).agg(
        isze_qty=("billed_qty", "sum"),
        isze_eur=("net_sales_value", "sum"),
    )
    g["isze_asp"] = g.apply(
        lambda r: (r["isze_eur"] / r["isze_qty"]) if r["isze_qty"] else None, axis=1
    )
    return g


def load_xref() -> pd.DataFrame:
    if not XREF_CSV.exists():
        return pd.DataFrame(columns=["genuine_pn", "bvp_pn", "pms_category"])
    x = pd.read_csv(XREF_CSV, dtype=str)
    x["genuine_pn"] = x["genuine_pn"].map(norm_pn)
    x["bvp_pn"] = x["bvp_pn"].map(norm_pn)
    return x


def brake_wear_flag(desc: str, pms_cat: str) -> str | None:
    d = (desc or "").upper()
    c = (pms_cat or "").lower()
    if "CLUTCH" in d:
        return None
    if "WHEEL; DISC" in d or d.startswith("WHEEL;"):
        return None
    if "EXH BRAKE" in d or "BRAKE UNIT; EXH" in d or "VALVE; EXH" in d:
        return None
    if "PARKING" in d:
        return None
    if "ROTOR; SENSOR" in d or "SENSOR" in d:
        return None
    if "CUP;" in d or "HOLDING" in d:
        return None
    if "BOLT" in d or "NUT" in d or "SPRING" in d or "HOSE" in d or "CLIP" in d:
        return None
    if "GASKET" in d or "BRACKET" in d or "SUPPORT" in d or "COVER; DUST" in d:
        return None
    if "REPAIR KIT" in d:
        return "repair_kit"
    if "CALIPER" in d and "PAD" not in d:
        return "caliper"
    if "PAD KIT" in d or c == "brake_pad":
        return "pad"
    if "ROTOR; DISC BRK" in d or "ROTOR; DISC BRAKE" in d or c == "brake_rotor":
        return "disc"
    if "SHOE ASM" in d or "SHOE KIT" in d or "SHOE; BRK" in d or c == "brake_shoe":
        return "shoe"
    return None


def style_header_row(ws, row, start, end):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = FILL_NAVY
        cell.font = FONT_TH
        cell.alignment = WRAP_C
        cell.border = THIN


def apply_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def fill_row(ws, row, start, end, fill, font=FONT_B, num_fmt=None):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = LEFT
        cell.border = THIN
        if num_fmt and cell.value is not None:
            pass


def build():
    duncan_long, duncan = load_duncan()
    dealer = load_dealer()
    isze = load_isze_ituk()
    xref = load_xref()

    pns = set(duncan["part_number"])
    dlr = dealer[dealer["part_number"].isin(pns)][
        ["part_number", "qty", "line_type", "pms_category", "description", "in_bvp_catalog",
         "catalog_paired_bvp", "catalog_paired_genuine"]
    ].copy()
    dlr = dlr.rename(columns={"qty": "dlr_qty", "description": "description_dealer"})

    m = duncan.merge(dlr, on="part_number", how="left")
    m = m.merge(isze, on="part_number", how="left")
    for c in ("dlr_qty", "isze_qty", "isze_eur", "isze_asp"):
        m[c] = pd.to_numeric(m[c], errors="coerce")
    m["dlr_qty"] = m["dlr_qty"].fillna(0)
    m["isze_qty"] = m["isze_qty"].fillna(0)
    m["isze_eur"] = m["isze_eur"].fillna(0)
    m["valued_eur"] = m.apply(
        lambda r: (r["dlr_qty"] * r["isze_asp"]) if pd.notna(r["isze_asp"]) else 0.0, axis=1
    )
    m["asp_missing"] = m["isze_asp"].isna() & (m["dlr_qty"] > 0)
    m["family_label"] = m["family"].map(family_label)
    m["line_type"] = m["line_type"].fillna("")
    m["on_n35"] = m["on_n35"].astype(int)
    m["on_n75"] = m["on_n75"].astype(int)

    # BVP twins from catalog
    twin_map = defaultdict(set)
    if not xref.empty:
        for _, r in xref.iterrows():
            if r.get("genuine_pn"):
                twin_map[r["genuine_pn"]].add(r.get("bvp_pn") or "")
            if r.get("bvp_pn"):
                twin_map[r["bvp_pn"]].add(r.get("genuine_pn") or "")
    m["catalog_bvp_twin"] = m["part_number"].map(
        lambda p: ", ".join(sorted(x for x in twin_map.get(p, set()) if x))
    )

    fam_order = ["oil", "fuel", "air", "brake", "clutch", "other_pms"]
    fam = (
        m.groupby("family", as_index=False)
        .agg(
            pn_count=("part_number", "nunique"),
            dlr_qty=("dlr_qty", "sum"),
            valued_eur=("valued_eur", "sum"),
            isze_qty=("isze_qty", "sum"),
            isze_eur=("isze_eur", "sum"),
            asp_missing_qty=("asp_missing", lambda s: m.loc[s.index, "dlr_qty"][s].sum() if s.any() else 0),
        )
    )
    fam["family"] = pd.Categorical(fam["family"], fam_order, ordered=True)
    fam = fam.sort_values("family").reset_index(drop=True)
    # ensure all families present
    for f in fam_order:
        if f not in set(fam["family"].astype(str)):
            fam = pd.concat(
                [fam, pd.DataFrame([{
                    "family": f, "pn_count": 0, "dlr_qty": 0, "valued_eur": 0,
                    "isze_qty": 0, "isze_eur": 0, "asp_missing_qty": 0,
                }])],
                ignore_index=True,
            )
    fam["family"] = pd.Categorical(fam["family"], fam_order, ordered=True)
    fam = fam.sort_values("family").reset_index(drop=True)
    fam["family_label"] = fam["family"].map(family_label)

    # Broader brake wear on dealer file (Study B)
    dealer = dealer.copy()
    dealer["wear_kind"] = dealer.apply(
        lambda r: brake_wear_flag(str(r.get("description") or ""), str(r.get("pms_category") or "")),
        axis=1,
    )
    dealer = dealer.merge(isze, on="part_number", how="left")
    dealer["isze_asp"] = pd.to_numeric(dealer["isze_asp"], errors="coerce")
    dealer["dlr_valued"] = dealer["qty"] * dealer["isze_asp"].fillna(0)
    wear = dealer[dealer["wear_kind"].isin(["pad", "disc", "shoe"])].copy()
    wear["in_duncan"] = wear["part_number"].isin(pns)
    wear = wear.merge(
        duncan[["part_number", "on_n35", "on_n75", "campaign_priority", "duncan_category"]],
        on="part_number",
        how="left",
    )

    # Hardware note (caliper / repair kit) - not in 1% basket
    hw = dealer[dealer["wear_kind"].isin(["caliper", "repair_kit"])].copy()

    totals = {
        "pn_count": int(m["part_number"].nunique()),
        "pn_with_dlr": int((m["dlr_qty"] > 0).sum()),
        "pn_with_isze": int((m["isze_qty"] != 0).sum()),
        "dlr_qty": float(m["dlr_qty"].sum()),
        "valued_eur": float(m["valued_eur"].sum()),
        "isze_qty": float(m["isze_qty"].sum()),
        "isze_eur": float(m["isze_eur"].sum()),
        "asp_missing_qty": float(m.loc[m["asp_missing"], "dlr_qty"].sum()),
        "n35_only_pns": int(((m["on_n35"] == 1) & (m["on_n75"] == 0)).sum()),
        "n75_only_pns": int(((m["on_n75"] == 1) & (m["on_n35"] == 0)).sum()),
        "both_pns": int(((m["on_n35"] == 1) & (m["on_n75"] == 1)).sum()),
    }

    r90 = m[m["part_number"] == "8982447940"]
    r90_dlr = float(r90["dlr_qty"].sum()) if not r90.empty else 0.0
    r90_isze_qty = float(r90["isze_qty"].sum()) if not r90.empty else 0.0
    r90_isze_eur = float(r90["isze_eur"].sum()) if not r90.empty else 0.0
    r90_val = float(r90["valued_eur"].sum()) if not r90.empty else 0.0
    r90_asp = float(r90["isze_asp"].iloc[0]) if not r90.empty and pd.notna(r90["isze_asp"].iloc[0]) else None

    pad_dlr = float(wear.loc[wear["wear_kind"] == "pad", "qty"].sum())
    disc_dlr = float(wear.loc[wear["wear_kind"] == "disc", "qty"].sum())
    shoe_dlr = float(wear.loc[wear["wear_kind"] == "shoe", "qty"].sum())
    pad_isze = float(wear.loc[wear["wear_kind"] == "pad", "isze_qty"].fillna(0).sum())
    disc_isze = float(wear.loc[wear["wear_kind"] == "disc", "isze_qty"].fillna(0).sum())
    shoe_isze = float(wear.loc[wear["wear_kind"] == "shoe", "isze_qty"].fillna(0).sum())
    pad_val = float(wear.loc[wear["wear_kind"] == "pad", "dlr_valued"].sum())
    disc_val = float(wear.loc[wear["wear_kind"] == "disc", "dlr_valued"].sum())
    shoe_val = float(wear.loc[wear["wear_kind"] == "shoe", "dlr_valued"].sum())
    pad_isze_eur = float(wear.loc[wear["wear_kind"] == "pad", "isze_eur"].fillna(0).sum())
    disc_isze_eur = float(wear.loc[wear["wear_kind"] == "disc", "isze_eur"].fillna(0).sum())
    shoe_isze_eur = float(wear.loc[wear["wear_kind"] == "shoe", "isze_eur"].fillna(0).sum())

    duncan_brake = m[m["family"] == "brake"]
    n35_pad_pns = {"8982169210", "8982169220"}
    n35_pad_dlr = float(m.loc[m["part_number"].isin(n35_pad_pns), "dlr_qty"].sum())

    bvp_brake = wear[(wear["line_type"] == "bvp") | wear["part_number"].astype(str).str.startswith(("1876", "5876"))]
    bvp_brake_qty = float(bvp_brake["qty"].sum()) if not bvp_brake.empty else 0.0

    top_wear = wear.sort_values(["qty", "dlr_valued"], ascending=False).head(20)

    # Print internal numbers for the agent
    print("=== DEFINITION ===")
    print("unique_pns", totals["pn_count"])
    print("n35_only", totals["n35_only_pns"], "n75_only", totals["n75_only_pns"], "both", totals["both_pns"])
    print(duncan[["part_number", "family", "on_n35", "on_n75", "duncan_category"]].to_string(index=False))
    print("=== STUDY A ===")
    print("dlr_qty", totals["dlr_qty"], "1pct_qty", totals["dlr_qty"] * 0.01)
    print("valued_eur", totals["valued_eur"], "1pct_eur", totals["valued_eur"] * 0.01)
    print("isze_qty", totals["isze_qty"], "1pct_isze_qty", totals["isze_qty"] * 0.01)
    print("isze_eur", totals["isze_eur"], "1pct_isze_eur", totals["isze_eur"] * 0.01)
    print("asp_missing_qty", totals["asp_missing_qty"])
    print(fam.to_string(index=False))
    print("=== IML N75 5-cat DLR (lock) ===", sum(IML_N75_DLR.values()), IML_N75_DLR)
    print("=== BRAKE ===")
    print("r90 dlr/isze/val", r90_dlr, r90_isze_qty, r90_val, "asp", r90_asp)
    print("duncan brake qty", float(duncan_brake["dlr_qty"].sum()), "eur", float(duncan_brake["valued_eur"].sum()))
    print("n35 pads 9210+9220 dlr", n35_pad_dlr)
    print("wear pad/disc/shoe dlr", pad_dlr, disc_dlr, shoe_dlr)
    print("wear pad/disc/shoe isze qty", pad_isze, disc_isze, shoe_isze)
    print("bvp brake wear qty", bvp_brake_qty)
    print("cover IML 2923/4679.5", 2923 / 4679.5)
    print("cover r90/4679.5", r90_dlr / 4679.5)
    print("cover duncan brake/4679.5", float(duncan_brake["dlr_qty"].sum()) / 4679.5)
    print("cover pads+shoes/4679.5", (pad_dlr + shoe_dlr) / 4679.5)
    print("top wear")
    print(top_wear[["part_number", "description", "wear_kind", "qty", "dlr_valued", "isze_qty", "isze_eur", "in_duncan"]].to_string(index=False))

    write_xlsx(m, fam, totals, wear, hw, top_wear, {
        "r90_dlr": r90_dlr, "r90_isze_qty": r90_isze_qty, "r90_isze_eur": r90_isze_eur,
        "r90_val": r90_val, "r90_asp": r90_asp,
        "pad_dlr": pad_dlr, "disc_dlr": disc_dlr, "shoe_dlr": shoe_dlr,
        "pad_isze": pad_isze, "disc_isze": disc_isze, "shoe_isze": shoe_isze,
        "pad_val": pad_val, "disc_val": disc_val, "shoe_val": shoe_val,
        "pad_isze_eur": pad_isze_eur, "disc_isze_eur": disc_isze_eur, "shoe_isze_eur": shoe_isze_eur,
        "n35_pad_dlr": n35_pad_dlr, "bvp_brake_qty": bvp_brake_qty,
        "duncan_brake_qty": float(duncan_brake["dlr_qty"].sum()),
        "duncan_brake_eur": float(duncan_brake["valued_eur"].sum()),
        "duncan_brake_isze_qty": float(duncan_brake["isze_qty"].sum()),
        "duncan_brake_isze_eur": float(duncan_brake["isze_eur"].sum()),
    })
    return totals


def write_xlsx(m: pd.DataFrame, fam: pd.DataFrame, totals: dict, wear: pd.DataFrame, hw: pd.DataFrame, top_wear: pd.DataFrame, br: dict):
    wb = openpyxl.Workbook()

    # ----- 00_Read_me -----
    ws = wb.active
    ws.title = "00_Read_me"
    ws.sheet_properties.tabColor = NAVY
    ws.merge_cells("A1:B1")
    ws["A1"] = "ITUK PMS  |  1 percentage point on all PMS parts  |  ITUK to dealer"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A1"].alignment = LEFT
    ws.row_dimensions[1].height = 28
    ws["A2"] = (
        "Audience: IML / ITUK / ISZE. File date: 15 August 2026. "
        "Calendar 2025 (1 January to 31 December). Every 1% and every cover % is an Excel formula."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 36

    read_rows = [
        ("What this file is",
         "One percentage point of calendar 2025 ITUK sales to dealers on the full PMS part-number set. "
         "The set is Duncan's N75 and N35 PMS project list (filters + brake + clutch, including BVP twins on that list). "
         "This is not 1% of expected service events, and it is not the fuel-and-air filter set only."),
        ("What 1.00 percentage point means",
         "1.00 percentage point = 1% of the calendar 2025 ITUK-to-dealer quantity (and of the labelled euro) on that part-number set. "
         "Formula: one_pct = total x 0.01. It is not 9,359 x 1.25 x 0.01 complete services."),
        ("Part-number definition",
         "Union of unique part numbers on Duncan's four model sheets: NLR87 N35.125 RZ4E, NLR88 N35.150 4JZ1, NPR75 N75 4HK1, NPR88 N75 4JZ1. "
         "Families follow Duncan's PMS category, mapped to the April / July IML study families (oil, fuel, air, brake, clutch). "
         "Air cleaner assemblies stay in air. Fuel sedimenter and fuel cartridge stay in fuel. Rear brake shoes stay in brake."),
        ("Sources",
         "1) Duncan N75/N35 PMS part-number workbook (13 August 2026 meeting inbox). "
         "2) April / July IML study (five families, N75 and full registered parc layers). "
         "3) N75 coverage calculations workbook (15 August 2026). "
         "4) dealer_sellout_ituk_cy2025_by_pn.csv (ITUK trucks). The IUK D-Max LCV file was not used."),
        ("Primary layer",
         "ITUK to dealer, calendar 2025. Quantity from the ITUK dealer invoice pivot (Customer Name x Item Number x month). "
         "That source is quantity only. It does not carry a dealer invoice euro."),
        ("Euro basis (labelled)",
         "Dealer invoice euro does not exist on the ITUK sell-out file. "
         "Valued euro = dealer quantity x ISZE calendar 2025 average selling price to ITUK (billing account 2575) for the same part number. "
         "This is an ISZE-to-ITUK price applied to dealer quantity. It is not a dealer invoice and not a workshop list price. "
         "ISZE-to-ITUK euro on sheet 05 is the actual billed net to ITUK on the same part numbers."),
        ("Secondary layer",
         "ISZE to ITUK billed quantity and billed euro, calendar 2025, same Duncan part-number set. Do not add the two layers."),
        ("Period",
         "Calendar 2025, 1 January to 31 December. Not 2026 year-to-date. Not warehouse issues."),
        ("UIO",
         "UIO is not used in the 1% calculation. "
         "N75 UIO 9,359 and wear factor 0.5 appear only on the brake cover sheet (Study B). "
         "22,389 is the complete UK registered parc (all models). Do not put 22,389 in the N75 cover denominator."),
        ("What is included",
         "N75 and N35 part numbers on Duncan's list. Genuine and BVP twins listed by Duncan. "
         "Supersession-related references on those four sheets."),
        ("What is excluded",
         "Lubricants (not on Duncan's list). Body panels. Urea filters. Wiper blades. "
         "Non-PMS hardware. IUK D-Max LCV. F-Series as a separate study. "
         "Catalog PMS codes that Duncan did not list (extra clutch kits, rotors) stay out of Study A; rotors are reported in Study B as a fact."),
        ("N35",
         "N35 is inside Study A because Duncan listed N35 scheduled-service part numbers. "
         "Study B leads with N75 (8982447940) and notes N35 pad kits where they appear."),
        ("How to read the sheets",
         "01_Definition = the part-number set. 02_Parameters = yellow inputs. "
         "03_One_pct = the 1% headline (formulas). 04_Family = 1% by family (formulas). "
         "05_PN_detail = every Duncan part number. 06_ISZE_layer = secondary ISZE to ITUK 1%. "
         "07_Brakes = Study B. 08_Brake_PNs = top wear part numbers."),
    ]
    ws["A4"] = "Topic"
    ws["B4"] = "Note"
    style_header_row(ws, 4, 1, 2)
    for i, (k, v) in enumerate(read_rows, start=5):
        ws.cell(i, 1, k).font = FONT_SH
        ws.cell(i, 1).fill = FILL_PAPER
        ws.cell(i, 1).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2, v).font = FONT_B
        ws.cell(i, 2).alignment = WRAP
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 52
    apply_widths(ws, {"A": 28, "B": 118})
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:4"
    ws.oddHeader.left.text = "ITUK PMS  |  1 percentage point  |  all PMS parts"

    # ----- 01_Definition -----
    ws = wb.create_sheet("01_Definition")
    ws.sheet_properties.tabColor = "C0552D"
    ws.merge_cells("A1:G1")
    ws["A1"] = "Definition  |  all PMS parts  |  Duncan N75 + N35 list"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Unique part numbers on Duncan's four model sheets. "
        "April / July IML study supplies the five families. Coverage file supplies the N75 cover method (used only in Study B)."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 28

    ws["A4"] = "Item"
    ws["B4"] = "Count / rule"
    style_header_row(ws, 4, 1, 2)
    def_items = [
        ("Unique PMS part numbers (Duncan union)", totals["pn_count"]),
        ("Of which sold by ITUK to dealers in 2025", totals["pn_with_dlr"]),
        ("Of which billed ISZE to ITUK in 2025", totals["pn_with_isze"]),
        ("N35 only (on N35 sheets, not on N75 sheets)", totals["n35_only_pns"]),
        ("N75 only (on N75 sheets, not on N35 sheets)", totals["n75_only_pns"]),
        ("On both N35 and N75 sheets", totals["both_pns"]),
        ("Other PMS family on Duncan's list", "None. Every Duncan row maps to oil, fuel, air, brake or clutch."),
        ("Lubricants", "Excluded. Not on Duncan's list."),
        ("Body", "Excluded. Not on Duncan's list."),
        ("Urea filter", "Excluded. Not on Duncan's list."),
        ("IUK D-Max LCV file", "Rejected. Wrong network."),
    ]
    for i, (k, v) in enumerate(def_items, start=5):
        ws.cell(i, 1, k).font = FONT_B
        ws.cell(i, 1).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2, v).font = FONT_B
        ws.cell(i, 2).alignment = WRAP
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 22
        if i <= 10:
            ws.cell(i, 2).fill = FILL_GRN
            if isinstance(v, (int, float)):
                ws.cell(i, 2).number_format = "#,##0"

    ws["A17"] = "Family map from Duncan PMS category"
    ws["A17"].font = FONT_SH
    ws.merge_cells("A17:B17")
    ws["A18"] = "Family"
    ws["B18"] = "Duncan category text (examples)"
    style_header_row(ws, 18, 1, 2)
    maps = [
        ("Oil filter", "Oil Filter, Oil Filter family, Oil Filter BVP"),
        ("Fuel filter", "Fuel Filter, Fuel Filter Element, Fuel Filter ASM, Fuel Cartridge, Fuel Sedimenter, and BVP twins"),
        ("Air filter", "Air Filter, Air Filter BVP, Air Cleaner ASM"),
        ("Brake", "Brake Pad Kit, Front/Rear Brake Pad, Brake Pad BVP, Rear Brake Shoe ASM / Kit"),
        ("Clutch", "Clutch Disc, Clutch Pressure Plate, and BVP twins. Manual only on Duncan's comments."),
        ("Other PMS", "Empty on this list. Kept so a later Duncan add would show."),
    ]
    for i, (k, v) in enumerate(maps, start=19):
        ws.cell(i, 1, k).border = THIN
        ws.cell(i, 2, v).border = THIN
        ws.cell(i, 1).alignment = WRAP
        ws.cell(i, 2).alignment = WRAP
        ws.row_dimensions[i].height = 28
    apply_widths(ws, {"A": 52, "B": 100, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18})
    ws.freeze_panes = "A5"

    # ----- 02_Parameters -----
    ws = wb.create_sheet("02_Parameters")
    ws.sheet_properties.tabColor = "F4C542"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Parameters  |  yellow = inputs  |  green = formulas"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Edit yellow only. The 1% sheets read PCT_POINT. "
        "UIO and wear factor are for Study B cover only. They do not change the 1% headline."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:F2")
    headers = ["ID", "Parameter", "Value", "Unit", "Excel name", "Notes"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 1, 6)

    params = [
        ("P01", "Percentage point", 0.01, "share", "PCT_POINT",
         "1.00 percentage point = 1% of the Duncan-set total. Yellow input."),
        ("P02", "Quantity period", "calendar 2025", "1 Jan 2025 to 31 Dec 2025", "QTY_PERIOD",
         "Both layers use this window."),
        ("P03", "N75 UIO (pilot parc)", 9359, "vehicles", "UIO_N75",
         "Study B only. Cleaned 2024A N75, both engine generations. Not 10,684. Not 22,389."),
        ("P04", "Wear events per vehicle per year", 0.5, "events / year", "WEAR_EVENTS_PER_YEAR",
         "Study B only. April 2026 analysis: 50,000 km / 100,000 km. Same as UIO x 0.5."),
        ("P05", "Complete UK registered parc (do not use as N75 UIO)", 22389, "vehicles", "UIO_COMPLETE_PARC",
         "Reference only. N35 + N75 + residual."),
        ("P06", "IML N75 dealer brake actual (locked integer)", 2923, "units", "IML_N75_DLR_BRAKE",
         "April 2026 N75 study / coverage file. N75 cut, not the full Duncan brake list."),
        ("P07", "IML N75 ISZE-to-ITUK brake actual (locked integer)", 2902, "units", "IML_N75_ISZE_BRAKE",
         "April 2026 N75 study integer lock."),
    ]
    for i, (pid, name, val, unit, xlname, note) in enumerate(params, start=5):
        ws.cell(i, 1, pid).border = THIN
        ws.cell(i, 2, name).border = THIN
        ws.cell(i, 2).alignment = WRAP
        c = ws.cell(i, 3, val)
        c.border = THIN
        c.fill = FILL_YEL
        c.font = Font(name="Calibri", size=10, bold=True, color=INK)
        if xlname == "PCT_POINT":
            c.number_format = "0.00%"
        elif isinstance(val, float):
            c.number_format = "0.00"
        elif isinstance(val, int):
            c.number_format = "#,##0"
        ws.cell(i, 4, unit).border = THIN
        ws.cell(i, 5, xlname).border = THIN
        ws.cell(i, 6, note).border = THIN
        ws.cell(i, 6).alignment = WRAP
        ws.row_dimensions[i].height = 36

    ws["A13"] = "D01"
    ws["B13"] = "Expected N75 brake quantity"
    ws["C13"] = "=C7*C8"
    ws["C13"].fill = FILL_GRN
    ws["C13"].number_format = "#,##0.0"
    ws["D13"] = "units / year"
    ws["E13"] = "EXPECTED_BRAKE_QTY"
    ws["F13"] = "9,359 x 0.5 = 4,679.5. Coverage file prints 4,680 as the integer. Cover formulas use 4,679.5."
    for col in range(1, 7):
        ws.cell(13, col).border = THIN
        ws.cell(13, col).alignment = WRAP
    ws.row_dimensions[13].height = 36

    # named ranges
    wb.defined_names.add(DefinedName(name="PCT_POINT", attr_text="'02_Parameters'!$C$5"))
    wb.defined_names.add(DefinedName(name="UIO_N75", attr_text="'02_Parameters'!$C$7"))
    wb.defined_names.add(DefinedName(name="WEAR_EVENTS_PER_YEAR", attr_text="'02_Parameters'!$C$8"))
    wb.defined_names.add(DefinedName(name="EXPECTED_BRAKE_QTY", attr_text="'02_Parameters'!$C$13"))
    wb.defined_names.add(DefinedName(name="IML_N75_DLR_BRAKE", attr_text="'02_Parameters'!$C$10"))
    wb.defined_names.add(DefinedName(name="IML_N75_ISZE_BRAKE", attr_text="'02_Parameters'!$C$11"))

    apply_widths(ws, {"A": 8, "B": 52, "C": 22, "D": 28, "E": 26, "F": 78})
    ws.freeze_panes = "A5"
    ws.row_dimensions[1].height = 24

    # ----- 05_PN_detail first (family/one_pct formulas point here) -----
    ws = wb.create_sheet("05_PN_detail")
    ws.sheet_properties.tabColor = "5B7C99"
    ws.merge_cells("A1:R1")
    ws["A1"] = "Part-number detail  |  Duncan set  |  calendar 2025  |  valued euro = dealer qty x ISZE ASP to ITUK"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Green cells are formulas. Quantity columns are source values. "
        "Valued euro is not a dealer invoice. Blank ASP means ISZE did not bill that part number to ITUK in 2025; valued euro is then 0."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:R2")
    ws.row_dimensions[2].height = 28

    pn_headers = [
        "Part number", "Family", "Family code", "Duncan category", "Description",
        "N35", "N75", "BVP on Duncan", "Catalog BVP twin", "Campaign priority (Duncan)",
        "ITUK to dealer qty", "ISZE ASP to ITUK EUR", "Valued EUR (dealer qty x ASP)",
        "1% dealer qty", "1% valued EUR",
        "ISZE to ITUK qty", "ISZE to ITUK EUR", "1% ISZE EUR",
    ]
    for i, h in enumerate(pn_headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 1, 18)

    m = m.sort_values(["family", "dlr_qty", "part_number"], ascending=[True, False, True]).reset_index(drop=True)
    first_data = 5
    for i, r in m.iterrows():
        row = first_data + i
        ws.cell(row, 1, r["part_number"]).number_format = "@"
        ws.cell(row, 2, r["family_label"])
        ws.cell(row, 3, r["family"])
        ws.cell(row, 4, r["duncan_category"])
        ws.cell(row, 5, r["description_duncan"] or r.get("description_dealer") or "")
        ws.cell(row, 6, int(r["on_n35"]))
        ws.cell(row, 7, int(r["on_n75"]))
        ws.cell(row, 8, "Yes" if r["is_bvp_duncan"] else "No")
        ws.cell(row, 9, r["catalog_bvp_twin"] or "")
        ws.cell(row, 10, r["campaign_priority"] or "")
        ws.cell(row, 11, float(r["dlr_qty"])).number_format = "#,##0.0"
        asp = r["isze_asp"]
        if pd.notna(asp):
            ws.cell(row, 12, float(asp)).number_format = "#,##0.00"
        else:
            ws.cell(row, 12, None)
        ws.cell(row, 13, f"=IF(L{row}=\"\",\"\",K{row}*L{row})")
        ws.cell(row, 13).number_format = "#,##0.00"
        ws.cell(row, 13).fill = FILL_GRN
        ws.cell(row, 14, f"=K{row}*PCT_POINT")
        ws.cell(row, 14).number_format = "#,##0.00"
        ws.cell(row, 14).fill = FILL_GRN
        ws.cell(row, 15, f"=IF(M{row}=\"\",\"\",M{row}*PCT_POINT)")
        ws.cell(row, 15).number_format = "#,##0.00"
        ws.cell(row, 15).fill = FILL_GRN
        ws.cell(row, 16, float(r["isze_qty"])).number_format = "#,##0.0"
        ws.cell(row, 17, float(r["isze_eur"])).number_format = "#,##0.00"
        ws.cell(row, 18, f"=Q{row}*PCT_POINT")
        ws.cell(row, 18).number_format = "#,##0.00"
        ws.cell(row, 18).fill = FILL_GRN
        for col in range(1, 19):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).font = FONT_B
            ws.cell(row, col).alignment = LEFT
        if r["part_number"] == "8982447940":
            ws.cell(row, 1).fill = FILL_YEL
            ws.cell(row, 2).fill = FILL_YEL
    last_pn = first_data + len(m) - 1
    tot_row = last_pn + 1
    ws.cell(tot_row, 1, "TOTAL Duncan set")
    ws.cell(tot_row, 11, f"=SUM(K{first_data}:K{last_pn})").number_format = "#,##0.0"
    ws.cell(tot_row, 13, f"=SUM(M{first_data}:M{last_pn})").number_format = "#,##0.00"
    ws.cell(tot_row, 14, f"=SUM(N{first_data}:N{last_pn})").number_format = "#,##0.00"
    ws.cell(tot_row, 15, f"=SUM(O{first_data}:O{last_pn})").number_format = "#,##0.00"
    ws.cell(tot_row, 16, f"=SUM(P{first_data}:P{last_pn})").number_format = "#,##0.0"
    ws.cell(tot_row, 17, f"=SUM(Q{first_data}:Q{last_pn})").number_format = "#,##0.00"
    ws.cell(tot_row, 18, f"=SUM(R{first_data}:R{last_pn})").number_format = "#,##0.00"
    for col in range(1, 19):
        ws.cell(tot_row, col).fill = FILL_NAVY
        ws.cell(tot_row, col).font = FONT_TH
        ws.cell(tot_row, col).border = THIN
    apply_widths(ws, {
        "A": 14, "B": 14, "C": 12, "D": 26, "E": 42, "F": 8, "G": 8, "H": 14,
        "I": 20, "J": 36, "K": 16, "L": 16, "M": 22, "N": 14, "O": 16, "P": 16, "Q": 16, "R": 14,
    })
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:R{last_pn}"
    ws.row_dimensions[4].height = 32

    # named totals
    wb.defined_names.add(DefinedName(name="DLR_QTY_TOTAL", attr_text=f"'05_PN_detail'!$K${tot_row}"))
    wb.defined_names.add(DefinedName(name="DLR_EUR_TOTAL", attr_text=f"'05_PN_detail'!$M${tot_row}"))
    wb.defined_names.add(DefinedName(name="ISZE_QTY_TOTAL", attr_text=f"'05_PN_detail'!$P${tot_row}"))
    wb.defined_names.add(DefinedName(name="ISZE_EUR_TOTAL", attr_text=f"'05_PN_detail'!$Q${tot_row}"))

    # ----- 03_One_pct -----
    ws = wb.create_sheet("03_One_pct", 2)
    ws.sheet_properties.tabColor = TERR
    ws.merge_cells("A1:D1")
    ws["A1"] = "Study A  |  1.00 percentage point  |  all PMS parts  |  ITUK to dealer"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Primary layer = ITUK to dealer. Secondary layer = ISZE to ITUK on the same part numbers. "
        "Euro on the primary layer is valued (dealer qty x ISZE ASP to ITUK), not a dealer invoice."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 32

    ws["A4"] = "Line"
    ws["B4"] = "ITUK to dealer (primary)"
    ws["C4"] = "ISZE to ITUK (secondary)"
    ws["D4"] = "Note"
    style_header_row(ws, 4, 1, 4)

    lines = [
        ("Duncan unique part numbers", f"={totals['pn_count']}", f"={totals['pn_count']}",
         "Same set on both layers."),
        ("Calendar 2025 total quantity", "=DLR_QTY_TOTAL", "=ISZE_QTY_TOTAL",
         "ITUK trucks. Not IUK D-Max."),
        ("Calendar 2025 euro", "=DLR_EUR_TOTAL", "=ISZE_EUR_TOTAL",
         "Primary = valued euro (dealer qty x ISZE ASP). Secondary = billed net to ITUK."),
        ("1.00 percentage point quantity", "=B6*PCT_POINT", "=C6*PCT_POINT",
         "Formula: total qty x 0.01."),
        ("1.00 percentage point euro", "=B7*PCT_POINT", "=C7*PCT_POINT",
         "Formula: total euro x 0.01."),
    ]
    fmts = {
        5: "0",
        6: "#,##0.0",
        7: "€#,##0.00",
        8: "#,##0.00",
        9: "€#,##0.00",
    }
    for i, (lab, b, c, note) in enumerate(lines, start=5):
        ws.cell(i, 1, lab).font = FONT_SH
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 4, note)
        for col in range(1, 5):
            ws.cell(i, col).border = THIN
            ws.cell(i, col).alignment = WRAP
        ws.cell(i, 2).fill = FILL_GRN
        ws.cell(i, 3).fill = FILL_GRN
        ws.cell(i, 2).number_format = fmts[i]
        ws.cell(i, 3).number_format = fmts[i]
        ws.row_dimensions[i].height = 28
    # highlight 1% rows
    for i in (8, 9):
        ws.cell(i, 1).fill = FILL_TERR
        ws.cell(i, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        ws.cell(i, 2).fill = FILL_YEL
        ws.cell(i, 3).fill = FILL_YEL

    ws["A11"] = "What this is not"
    ws["A11"].font = FONT_SH
    ws.merge_cells("A11:D11")
    nots = [
        "Not 117 additional complete services (that was 9,359 x 1.25 x 0.01 on the N75 filter expected).",
        "Not 1% of the fuel-and-air filter set only.",
        "Not a coverage-point move (actual / expected). Coverage stays on sheet 07 for brakes.",
        "Not IUK D-Max. Not 2026 warehouse year-to-date.",
    ]
    for i, t in enumerate(nots, start=12):
        ws.merge_cells(f"A{i}:D{i}")
        ws.cell(i, 1, t).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.row_dimensions[i].height = 20

    ws["A17"] = "IML N75 five-family lock (context only  |  not the 1% basis)"
    ws["A17"].font = FONT_SH
    ws.merge_cells("A17:D17")
    ws["A18"] = (
        "The April / July IML study N75 dealer five-family total is 22,095 units "
        f"(oil {IML_N75_DLR['oil']:,} + fuel {IML_N75_DLR['fuel']:,} + air {IML_N75_DLR['air']:,} + "
        f"clutch {IML_N75_DLR['clutch']:,} + brake {IML_N75_DLR['brake']:,}). "
        "That cut is N75 only and uses the study category classification. "
        "Study A uses Duncan's part numbers on all ITUK dealer lines (N35 + N75). The two totals are not the same number."
    )
    ws["A18"].alignment = WRAP
    ws.merge_cells("A18:D18")
    ws.row_dimensions[18].height = 56

    apply_widths(ws, {"A": 42, "B": 28, "C": 28, "D": 62})
    ws.freeze_panes = "A5"
    ws.row_dimensions[1].height = 24

    # ----- 04_Family -----
    ws = wb.create_sheet("04_Family", 3)
    ws.sheet_properties.tabColor = "2E6B4E"
    ws.merge_cells("A1:L1")
    ws["A1"] = "1% split by family  |  ITUK to dealer primary  |  ISZE to ITUK secondary"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = "Every total and every 1% cell is a formula against 05_PN_detail. Other PMS is on the list so a zero is visible."
    ws.merge_cells("A2:L2")
    ws["A2"].alignment = WRAP

    fam_headers = [
        "Family", "Family code", "PN count",
        "ITUK to dealer qty", "Valued EUR", "1% dealer qty", "1% valued EUR",
        "ISZE to ITUK qty", "ISZE to ITUK EUR", "1% ISZE qty", "1% ISZE EUR",
        "Share of dealer qty",
    ]
    for i, h in enumerate(fam_headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 1, 12)
    ws.row_dimensions[4].height = 32

    # family rows 5-10, SUMIF on family code col C of PN detail
    codes = ["oil", "fuel", "air", "brake", "clutch", "other_pms"]
    labels = [family_label(c) for c in codes]
    for i, (lab, code) in enumerate(zip(labels, codes), start=5):
        ws.cell(i, 1, lab)
        ws.cell(i, 2, code)
        ws.cell(i, 3, f"=COUNTIF('05_PN_detail'!$C$5:$C${last_pn},B{i})")
        ws.cell(i, 4, f"=SUMIF('05_PN_detail'!$C$5:$C${last_pn},B{i},'05_PN_detail'!$K$5:$K${last_pn})")
        ws.cell(i, 5, f"=SUMIF('05_PN_detail'!$C$5:$C${last_pn},B{i},'05_PN_detail'!$M$5:$M${last_pn})")
        ws.cell(i, 6, f"=D{i}*PCT_POINT")
        ws.cell(i, 7, f"=E{i}*PCT_POINT")
        ws.cell(i, 8, f"=SUMIF('05_PN_detail'!$C$5:$C${last_pn},B{i},'05_PN_detail'!$P$5:$P${last_pn})")
        ws.cell(i, 9, f"=SUMIF('05_PN_detail'!$C$5:$C${last_pn},B{i},'05_PN_detail'!$Q$5:$Q${last_pn})")
        ws.cell(i, 10, f"=H{i}*PCT_POINT")
        ws.cell(i, 11, f"=I{i}*PCT_POINT")
        ws.cell(i, 12, f"=IF(D$11=0,0,D{i}/D$11)")
        for col in range(1, 13):
            ws.cell(i, col).border = THIN
            ws.cell(i, col).font = FONT_B
            ws.cell(i, col).fill = FILL_GRN if col >= 3 else FILL_PAPER
        ws.cell(i, 3).number_format = "0"
        ws.cell(i, 4).number_format = "#,##0.0"
        ws.cell(i, 5).number_format = "€#,##0.00"
        ws.cell(i, 6).number_format = "#,##0.00"
        ws.cell(i, 7).number_format = "€#,##0.00"
        ws.cell(i, 8).number_format = "#,##0.0"
        ws.cell(i, 9).number_format = "€#,##0.00"
        ws.cell(i, 10).number_format = "#,##0.00"
        ws.cell(i, 11).number_format = "€#,##0.00"
        ws.cell(i, 12).number_format = "0.0%"
        ws.row_dimensions[i].height = 20

    # total row 11
    ws.cell(11, 1, "TOTAL")
    ws.cell(11, 2, "all")
    for col, letter in [(3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I"), (10, "J"), (11, "K")]:
        ws.cell(11, col, f"=SUM({letter}5:{letter}10)")
    ws.cell(11, 12, "=IF(D11=0,0,D11/D11)")
    for col in range(1, 13):
        ws.cell(11, col).fill = FILL_NAVY
        ws.cell(11, col).font = FONT_TH
        ws.cell(11, col).border = THIN
    ws.cell(11, 3).number_format = "0"
    for col in (4, 6, 8, 10):
        ws.cell(11, col).number_format = "#,##0.00"
    for col in (5, 7, 9, 11):
        ws.cell(11, col).number_format = "€#,##0.00"
    ws.cell(11, 12).number_format = "0.0%"

    ws["A13"] = "IML study category average euro (context  |  not used in the 1% formulas)"
    ws["A13"].font = FONT_SH
    ws.merge_cells("A13:L13")
    ws["A14"] = "Family"
    ws["B14"] = "IML study avg EUR"
    ws["C14"] = "N75 dealer qty (lock)"
    ws["D14"] = "N75 dealer x avg EUR"
    style_header_row(ws, 14, 1, 4)
    iml_rows = [("oil", "Oil filter"), ("fuel", "Fuel filter"), ("air", "Air filter"),
                ("brake", "Brake"), ("clutch", "Clutch")]
    for i, (code, lab) in enumerate(iml_rows, start=15):
        ws.cell(i, 1, lab).border = THIN
        ws.cell(i, 2, IML_AVG_EUR[code]).number_format = "€#,##0.00"
        ws.cell(i, 2).fill = FILL_YEL
        ws.cell(i, 2).border = THIN
        ws.cell(i, 3, IML_N75_DLR[code]).number_format = "#,##0"
        ws.cell(i, 3).fill = FILL_YEL
        ws.cell(i, 3).border = THIN
        ws.cell(i, 4, f"=B{i}*C{i}").number_format = "€#,##0.00"
        ws.cell(i, 4).fill = FILL_GRN
        ws.cell(i, 4).border = THIN
    ws.cell(20, 1, "N75 five-family lock")
    ws.cell(20, 3, "=SUM(C15:C19)").number_format = "#,##0"
    ws.cell(20, 4, "=SUM(D15:D19)").number_format = "€#,##0.00"
    for col in range(1, 5):
        ws.cell(20, col).fill = FILL_NAVY
        ws.cell(20, col).font = FONT_TH
        ws.cell(20, col).border = THIN
    ws["A22"] = (
        "The IML average euro is the study transfer-price convention (oil 7.74, fuel 13.06, air 30.30, clutch 80.72, brake 88.63). "
        "Study A 1% euro uses part-number ISZE ASP, not these averages. The N75 lock is a different cut from the Duncan all-models set."
    )
    ws["A22"].alignment = WRAP
    ws.merge_cells("A22:L22")
    ws.row_dimensions[22].height = 36

    apply_widths(ws, {
        "A": 18, "B": 14, "C": 12, "D": 18, "E": 16, "F": 14, "G": 16,
        "H": 16, "I": 16, "J": 12, "K": 14, "L": 16,
    })
    ws.freeze_panes = "A5"

    # ----- 06_ISZE_layer -----
    ws = wb.create_sheet("06_ISZE_layer")
    ws.sheet_properties.tabColor = "7A8B99"
    ws.merge_cells("A1:D1")
    ws["A1"] = "Secondary layer  |  ISZE to ITUK  |  same Duncan part-number set  |  calendar 2025"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Billing account 2575 (ISUZU TRUCK (UK) LIMITED). "
        "Quantity and euro are billed net. Do not add this layer to the dealer layer."
    )
    ws.merge_cells("A2:D2")
    ws["A2"].alignment = WRAP
    ws["A4"] = "Line"
    ws["B4"] = "Value"
    ws["C4"] = "Formula / source"
    style_header_row(ws, 4, 1, 3)
    isze_lines = [
        ("ISZE to ITUK quantity (Duncan set)", "=ISZE_QTY_TOTAL", "SUM of 05_PN_detail ISZE qty"),
        ("ISZE to ITUK billed euro (Duncan set)", "=ISZE_EUR_TOTAL", "SUM of 05_PN_detail ISZE euro"),
        ("1.00 percentage point quantity", "=B5*PCT_POINT", "ISZE qty x 0.01"),
        ("1.00 percentage point euro", "=B6*PCT_POINT", "ISZE euro x 0.01"),
    ]
    for i, (lab, val, src) in enumerate(isze_lines, start=5):
        ws.cell(i, 1, lab).border = THIN
        ws.cell(i, 2, val).border = THIN
        ws.cell(i, 2).fill = FILL_GRN
        ws.cell(i, 3, src).border = THIN
        ws.row_dimensions[i].height = 22
    ws.cell(5, 2).number_format = "#,##0.0"
    ws.cell(6, 2).number_format = "€#,##0.00"
    ws.cell(7, 2).number_format = "#,##0.00"
    ws.cell(8, 2).number_format = "€#,##0.00"
    ws.cell(7, 1).fill = FILL_TERR
    ws.cell(8, 1).fill = FILL_TERR
    ws.cell(7, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(8, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    apply_widths(ws, {"A": 48, "B": 22, "C": 40, "D": 20})
    ws.freeze_panes = "A5"

    # ----- 07_Brakes -----
    ws = wb.create_sheet("07_Brakes")
    ws.sheet_properties.tabColor = TERR
    ws.merge_cells("A1:F1")
    ws["A1"] = "Study B  |  brake parts  |  N75 first  |  calendar 2025"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Duncan Rev2 took brakes off the campaign headline. "
        "8982447940 remains the highest N75 focus item on Duncan's NPR75 / NPR88 sheets. Facts only."
    )
    ws.merge_cells("A2:F2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 28

    ws["A4"] = "Cover vs expected (N75 wear assumption)"
    ws["A4"].font = FONT_SH
    ws.merge_cells("A4:F4")
    bh = ["Lens", "Actual qty", "Expected qty", "Cover %", "Layer", "What it includes"]
    for i, h in enumerate(bh, 1):
        ws.cell(5, i, h)
    style_header_row(ws, 5, 1, 6)

    # rows 6-12 cover table
    # expected always = EXPECTED_BRAKE_QTY
    cover_rows = [
        ("N75 IML study lock (April 2026)", "=IML_N75_DLR_BRAKE",
         "ITUK to dealer",
         "N75 category cut from the April study / coverage file (2,923). Not VIN-tagged."),
        ("N75 IML study lock (April 2026)", "=IML_N75_ISZE_BRAKE",
         "ISZE to ITUK",
         "N75 share-scaled billed units (2,902)."),
        ("R90 pad 8982447940 only", br["r90_dlr"],
         "ITUK to dealer",
         "Highest N75 focus item on Duncan NPR75 / NPR88. Genuine. No BVP twin in the catalog pair file."),
        ("R90 pad 8982447940 only", br["r90_isze_qty"],
         "ISZE to ITUK",
         "Same part number, billed to ITUK."),
        ("Duncan brake family (all models on the list)", br["duncan_brake_qty"],
         "ITUK to dealer",
         "Includes N35 pad kits 8982169210 and 8982169220. Do not read this as N75-only cover."),
        ("All pad kits on the ITUK dealer file", br["pad_dlr"],
         "ITUK to dealer",
         "Catalog / description pad kits, including N35. Wider than Duncan."),
        ("Pad kits + shoes (no discs)", f"={br['pad_dlr']}+{br['shoe_dlr']}",
         "ITUK to dealer",
         "Wear items closest to the 0.5 kit assumption. Discs are reported separately."),
    ]
    for i, (lens, actual, layer, note) in enumerate(cover_rows, start=6):
        ws.cell(i, 1, lens).alignment = WRAP
        ws.cell(i, 2, actual)
        ws.cell(i, 2).number_format = "#,##0.0"
        if isinstance(actual, str) and actual.startswith("="):
            ws.cell(i, 2).fill = FILL_GRN
        else:
            ws.cell(i, 2).fill = FILL_YEL
        ws.cell(i, 3, "=EXPECTED_BRAKE_QTY")
        ws.cell(i, 3).fill = FILL_GRN
        ws.cell(i, 3).number_format = "#,##0.0"
        ws.cell(i, 4, f"=IF(C{i}=0,0,B{i}/C{i})")
        ws.cell(i, 4).fill = FILL_GRN
        ws.cell(i, 4).number_format = "0.0%"
        ws.cell(i, 5, layer)
        ws.cell(i, 6, note).alignment = WRAP
        for col in range(1, 7):
            ws.cell(i, col).border = THIN
            ws.cell(i, col).font = FONT_B
        ws.row_dimensions[i].height = 36

    ws["A14"] = "R90 8982447940 (N75 focus pad)"
    ws["A14"].font = FONT_SH
    ws.merge_cells("A14:F14")
    r90_lines = [
        ("ITUK to dealer qty", br["r90_dlr"], "#,##0.0"),
        ("ISZE to ITUK qty", br["r90_isze_qty"], "#,##0.0"),
        ("ISZE ASP to ITUK EUR", br["r90_asp"] if br["r90_asp"] is not None else "No 2025 bill",
         "€#,##0.00" if br["r90_asp"] is not None else "@"),
        ("Valued EUR (dealer qty x ASP)", br["r90_val"], "€#,##0.00"),
        ("ISZE to ITUK billed EUR", br["r90_isze_eur"], "€#,##0.00"),
        ("BVP twin in catalog pair file", "No", "@"),
        ("Duncan campaign priority", "Highest N75 focus item", "@"),
    ]
    ws["A15"] = "Fact"
    ws["B15"] = "Value"
    style_header_row(ws, 15, 1, 2)
    for i, (lab, val, fmt) in enumerate(r90_lines, start=16):
        ws.cell(i, 1, lab).border = THIN
        c = ws.cell(i, 2, val)
        c.border = THIN
        c.number_format = fmt
        c.fill = FILL_PAPER

    ws["A24"] = "N35 on the same list"
    ws["A24"].font = FONT_SH
    ws["A25"] = (
        f"Pad kits 8982169210 and 8982169220 are N35-primary on Duncan (NLR87 / NLR88) and also appear on NPR88 as VIN-specific. "
        f"ITUK to dealer quantity on those two part numbers in 2025: {br['n35_pad_dlr']:,.0f}. "
        "BVP twins 5876102380 and 5876170050 exist for 8982169220. Duncan year-to-date demand on 5876102380 was 0."
    )
    ws["A25"].alignment = WRAP
    ws.merge_cells("A25:F25")
    ws.row_dimensions[25].height = 48

    ws["A27"] = "Wear items beyond filters (what else exists)"
    ws["A27"].font = FONT_SH
    ws.merge_cells("A27:F27")
    ws["A28"] = "Kind"
    ws["B28"] = "ITUK to dealer qty"
    ws["C28"] = "Valued EUR"
    ws["D28"] = "ISZE to ITUK qty"
    ws["E28"] = "ISZE to ITUK EUR"
    ws["F28"] = "On Duncan list?"
    style_header_row(ws, 28, 1, 6)
    extra = [
        ("Pads (PAD KIT / brake_pad)", br["pad_dlr"], br["pad_val"], br["pad_isze"], br["pad_isze_eur"],
         "Main Duncan brake rows plus any extra pad kits on the dealer file"),
        ("Discs / rotors (ROTOR; DISC BRK)", br["disc_dlr"], br["disc_val"], br["disc_isze"], br["disc_isze_eur"],
         "Not on Duncan's PMS list. Reported as a fact."),
        ("Shoes (SHOE ASM / SHOE KIT)", br["shoe_dlr"], br["shoe_val"], br["shoe_isze"], br["shoe_isze_eur"],
         "N35 rear shoe is on Duncan (8973319320, 5878320440)."),
        ("BVP among pad/disc/shoe", br["bvp_brake_qty"], None, None, None,
         "BVP twins exist for 8982169220, not for 8982447940."),
    ]
    for i, row in enumerate(extra, start=29):
        ws.cell(i, 1, row[0]).border = THIN
        ws.cell(i, 2, row[1]).number_format = "#,##0.0"
        ws.cell(i, 2).border = THIN
        if row[2] is not None:
            ws.cell(i, 3, row[2]).number_format = "€#,##0.00"
        ws.cell(i, 3).border = THIN
        if row[3] is not None:
            ws.cell(i, 4, row[3]).number_format = "#,##0.0"
        ws.cell(i, 4).border = THIN
        if row[4] is not None:
            ws.cell(i, 5, row[4]).number_format = "€#,##0.00"
        ws.cell(i, 5).border = THIN
        ws.cell(i, 6, row[5]).border = THIN
        ws.cell(i, 6).alignment = WRAP
        ws.row_dimensions[i].height = 28

    ws["A34"] = (
        "Calipers and repair kits exist on the dealer file. They are hardware, not the 0.5 kit assumption, "
        "and they are not in the Study A 1% basket. See 08_Brake_PNs for the ranked wear list."
    )
    ws["A34"].alignment = WRAP
    ws.merge_cells("A34:F34")
    ws.row_dimensions[34].height = 28
    apply_widths(ws, {"A": 52, "B": 20, "C": 16, "D": 18, "E": 18, "F": 62})
    ws.freeze_panes = "A6"
    ws.row_dimensions[1].height = 24

    # ----- 08_Brake_PNs -----
    ws = wb.create_sheet("08_Brake_PNs")
    ws.sheet_properties.tabColor = "8B4513"
    ws.merge_cells("A1:L1")
    ws["A1"] = "Brake wear part numbers  |  pads, discs, shoes  |  ranked by ITUK to dealer qty"
    ws["A1"].fill = FILL_NAVY
    ws["A1"].font = FONT_H
    ws["A2"] = (
        "Wear items only (pad kit, rotor / disc, shoe). "
        "Valued euro = dealer qty x ISZE ASP to ITUK. Calipers and repair kits are listed below the wear block as hardware."
    )
    ws.merge_cells("A2:L2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 28

    bpn_headers = [
        "Part number", "Kind", "Description", "ITUK to dealer qty", "Valued EUR",
        "1% dealer qty", "1% valued EUR", "ISZE to ITUK qty", "ISZE to ITUK EUR",
        "On Duncan list", "N35", "N75",
    ]
    for i, h in enumerate(bpn_headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 1, 12)
    ws.row_dimensions[4].height = 32

    wear_sorted = wear.sort_values(["qty", "dlr_valued"], ascending=False).reset_index(drop=True)
    w_first = 5
    for i, r in wear_sorted.iterrows():
        row = w_first + i
        ws.cell(row, 1, r["part_number"]).number_format = "@"
        ws.cell(row, 2, r["wear_kind"])
        ws.cell(row, 3, r.get("description") or "")
        ws.cell(row, 4, float(r["qty"])).number_format = "#,##0.0"
        ws.cell(row, 5, float(r["dlr_valued"])).number_format = "€#,##0.00"
        ws.cell(row, 6, f"=D{row}*PCT_POINT").number_format = "#,##0.00"
        ws.cell(row, 6).fill = FILL_GRN
        ws.cell(row, 7, f"=E{row}*PCT_POINT").number_format = "€#,##0.00"
        ws.cell(row, 7).fill = FILL_GRN
        ws.cell(row, 8, float(r["isze_qty"] if pd.notna(r.get("isze_qty")) else 0)).number_format = "#,##0.0"
        ws.cell(row, 9, float(r["isze_eur"] if pd.notna(r.get("isze_eur")) else 0)).number_format = "€#,##0.00"
        ws.cell(row, 10, "Yes" if bool(r.get("in_duncan")) else "No")
        ws.cell(row, 11, int(r["on_n35"]) if pd.notna(r.get("on_n35")) else "")
        ws.cell(row, 12, int(r["on_n75"]) if pd.notna(r.get("on_n75")) else "")
        for col in range(1, 13):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).font = FONT_B
            ws.cell(row, col).alignment = LEFT
        if r["part_number"] == "8982447940":
            for col in range(1, 13):
                ws.cell(row, col).fill = FILL_YEL
                if col in (6, 7):
                    ws.cell(row, col).fill = FILL_GRN
    w_last = w_first + len(wear_sorted) - 1 if len(wear_sorted) else w_first
    if len(wear_sorted):
        wt = w_last + 1
        ws.cell(wt, 1, "TOTAL wear (pads + discs + shoes)")
        ws.cell(wt, 4, f"=SUM(D{w_first}:D{w_last})").number_format = "#,##0.0"
        ws.cell(wt, 5, f"=SUM(E{w_first}:E{w_last})").number_format = "€#,##0.00"
        ws.cell(wt, 6, f"=SUM(F{w_first}:F{w_last})").number_format = "#,##0.00"
        ws.cell(wt, 7, f"=SUM(G{w_first}:G{w_last})").number_format = "€#,##0.00"
        ws.cell(wt, 8, f"=SUM(H{w_first}:H{w_last})").number_format = "#,##0.0"
        ws.cell(wt, 9, f"=SUM(I{w_first}:I{w_last})").number_format = "€#,##0.00"
        for col in range(1, 13):
            ws.cell(wt, col).fill = FILL_NAVY
            ws.cell(wt, col).font = FONT_TH
            ws.cell(wt, col).border = THIN
        next_block = wt + 2
    else:
        next_block = 6

    ws.cell(next_block, 1, "Hardware on the dealer file (not in Study A, not in the 0.5 cover)")
    ws.cell(next_block, 1).font = FONT_SH
    ws.merge_cells(f"A{next_block}:L{next_block}")
    hh = next_block + 1
    for i, h in enumerate(["Part number", "Kind", "Description", "ITUK to dealer qty", "Valued EUR", "ISZE to ITUK qty"], 1):
        ws.cell(hh, i, h)
    style_header_row(ws, hh, 1, 6)
    hw_sorted = hw.sort_values("qty", ascending=False).reset_index(drop=True)
    for i, r in hw_sorted.iterrows():
        row = hh + 1 + i
        ws.cell(row, 1, r["part_number"]).number_format = "@"
        ws.cell(row, 2, r["wear_kind"])
        ws.cell(row, 3, r.get("description") or "")
        ws.cell(row, 4, float(r["qty"])).number_format = "#,##0.0"
        ws.cell(row, 5, float(r["dlr_valued"])).number_format = "€#,##0.00"
        ws.cell(row, 6, float(r["isze_qty"] if pd.notna(r.get("isze_qty")) else 0)).number_format = "#,##0.0"
        for col in range(1, 7):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).font = FONT_B

    apply_widths(ws, {
        "A": 14, "B": 10, "C": 42, "D": 18, "E": 16, "F": 14, "G": 16,
        "H": 16, "I": 16, "J": 14, "K": 8, "L": 8,
    })
    ws.freeze_panes = "A5"
    if len(wear_sorted):
        ws.auto_filter.ref = f"A4:L{w_last}"

    # print + freeze on all
    for name in wb.sheetnames:
        sh = wb[name]
        # Pre-delivery: col B is treated as a description column (min width 28, wrap)
        cur_b = sh.column_dimensions["B"].width or 10
        sh.column_dimensions["B"].width = max(cur_b, 28)
        for row in sh.iter_rows(min_col=2, max_col=4, max_row=sh.max_row):
            for cell in row:
                if cell.column in (2, 4):
                    cell.alignment = Alignment(
                        wrap_text=True,
                        vertical="center",
                        horizontal=cell.alignment.horizontal or "left",
                    )
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_view.showGridLines = False
        sh.oddFooter.left.text = "ITUK PMS  |  1 percentage point  |  all PMS parts"
        sh.oddFooter.right.text = "Page &P of &N"

    order = [
        "00_Read_me", "01_Definition", "02_Parameters", "03_One_pct",
        "04_Family", "05_PN_detail", "06_ISZE_layer", "07_Brakes", "08_Brake_PNs",
    ]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    OUT_MEET.parent.mkdir(parents=True, exist_ok=True)
    if OUT_MEET.exists():
        OUT_MEET.unlink()
    wb.save(OUT_MEET)
    shutil.copy2(OUT_MEET, OUT_ROOT)
    print("WROTE", OUT_MEET)
    print("COPIED", OUT_ROOT)


if __name__ == "__main__":
    build()
