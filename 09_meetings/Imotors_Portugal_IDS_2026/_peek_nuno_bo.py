from pathlib import Path
import openpyxl
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_nuno_backordered_lines_20261006\20260929_1531_Nuno_Backordered lines 2026.xlsx")
print("size", p.stat().st_size)
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("sheets", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("===", sn, "dims", ws.dimensions)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 8:
            break
        print(i, row[:20] if row else row)
    # count rows
    n = 0
    for _ in ws.iter_rows(values_only=True):
        n += 1
    print("nrows", n)
wb.close()
# also ISD file briefly
p2 = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_nuno_backordered_lines_20261006\20260729_1549_Nuno_ISD backorders.XLSX")
if p2.exists():
    wb = openpyxl.load_workbook(p2, read_only=True, data_only=True)
    print("ISD sheets", wb.sheetnames, "size", p2.stat().st_size)
    ws = wb[wb.sheetnames[0]]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 5: break
        print("ISD", i, row[:15] if row else row)
    wb.close()
