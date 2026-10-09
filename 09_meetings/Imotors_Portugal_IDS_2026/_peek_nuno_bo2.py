from pathlib import Path
import openpyxl
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_nuno_backordered_lines_20261006\20260929_1531_Nuno_Backordered lines 2026.xlsx")
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("sheets", wb.sheetnames)
for sn in ["BO 20260831", "BO 20260930"]:
    ws = wb[sn]
    print("===", sn)
    n = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        n += 1
        if i <= 6:
            vals = [v for v in (row or [])]
            print(i, vals[:18])
    print("nrows_incl_header", n)
wb.close()
