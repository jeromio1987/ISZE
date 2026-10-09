# -*- coding: utf-8 -*-
"""Historical I-Motors BO (Nuno) vs ISZE open orders (519) + AIM."""
from __future__ import annotations
from pathlib import Path
import json, csv, re
from collections import defaultdict
from datetime import datetime, date
import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BO_XLSX = ROOT / "01_data/90_INBOX/_nuno_backordered_lines_20261006/20260929_1531_Nuno_Backordered lines 2026.xlsx"
ORDER_XLSX = ROOT / "01_data/90_INBOX/_imotors_bo_compare_20261006/EXPORT_ISZE_Order_20261006.xlsx"
AIM = ROOT / "01_data/00_derived/_releases/aim_pn_classification_release_2026-10-04.csv"
KPI = ROOT / "01_data/90_INBOX/_nuno_kpi_portugal_20261001/09 - Set_KPI Check Sheet Monthly 26_Portugal.xlsx"
OUT = ROOT / "01_data/90_INBOX/_imotors_bo_compare_20261006/IMOTORS_BO_HISTORY_ANALYSIS_20261006"
OUT.mkdir(parents=True, exist_ok=True)

def norm_pn(x):
    if x is None:
        return ""
    s = str(x).strip().upper()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s

def parse_fecha(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    # YYYYMMDD
    if re.fullmatch(r"\d{8}", s):
        return datetime.strptime(s, "%Y%m%d").date()
    if re.fullmatch(r"\d{6}", s):
        return datetime.strptime(s, "%Y%m").date()
    try:
        return datetime.fromisoformat(s[:10]).date()
    except Exception:
        return None

# --- AIM ---
aim = {}
with AIM.open(encoding="utf-8", newline="") as f:
    r = csv.DictReader(f)
    for row in r:
        pn = norm_pn(row.get("part_number"))
        if not pn:
            continue
        aim[pn] = {
            "l1": row.get("aim_l1_fixed") or row.get("catalog_l1") or "NO_AIM",
            "l2": row.get("aim_l2_fixed") or row.get("catalog_l2") or "",
            "desc": row.get("description_norm") or "",
        }

# --- KPI monthly B/O counts ---
kpi_bo = {}
wbk = openpyxl.load_workbook(KPI, data_only=True, read_only=True)
ws = wbk["Checklist"]
months = ["2026-01","2026-02","2026-03","2026-04","2026-05","2026-06","2026-07","2026-08","2026-09"]
for row in ws.iter_rows(values_only=True):
    if row and row[1] and "B/O" in str(row[1]):
        for i, m in enumerate(months):
            kpi_bo[m] = row[2+i]
        break
wbk.close()

# --- Nuno BO sheets ---
wb = openpyxl.load_workbook(BO_XLSX, data_only=True, read_only=True)
monthly = []
all_lines = []
for sn in wb.sheetnames:
    m = re.search(r"BO (\d{8})", sn)
    if not m:
        continue
    asof = datetime.strptime(m.group(1), "%Y%m%d").date()
    ym = asof.strftime("%Y-%m")
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        monthly.append({"sheet": sn, "asof": asof.isoformat(), "ym": ym, "lines": 0, "lines_with_pn": 0, "eur_open": 0.0, "empty": True})
        continue
    hdr = [str(h).strip().upper() if h else "" for h in rows[0]]
    idx = {h:i for i,h in enumerate(hdr)}
    def col(*names):
        for n in names:
            if n in idx: return idx[n]
        return None
    i_ref = col("REFERENCIA", "REF", "PN")
    i_desc = col("DESCRIPCION", "DESCRIPTION")
    i_qty = col("QTY PDTE", "QTY_PDTE", "QTY PENDING")
    i_ord = col("QTY PEDIDA")
    i_pu = col("P.UNIT", "P UNIT", "PUNIT")
    i_fecha = col("FECHA")
    i_ped = col("PEDIDO")
    i_tipo = col("TIPO")
    i_prov = col("PROVEEDOR")
    i_refp = col("REF PROV")
    lines = 0
    with_pn = 0
    eur = 0.0
    ages = []
    by_l1 = defaultdict(lambda: {"lines":0,"eur":0.0,"qty":0.0})
    for r in rows[1:]:
        if not r or all(v is None or str(v).strip()=="" for v in r):
            continue
        pn = norm_pn(r[i_ref] if i_ref is not None else None)
        if not pn:
            # skip blank
            continue
        lines += 1
        with_pn += 1
        try:
            qty = float(r[i_qty] or 0)
        except Exception:
            qty = 0.0
        try:
            pu = float(r[i_pu] or 0)
        except Exception:
            pu = 0.0
        line_eur = qty * pu
        eur += line_eur
        fd = parse_fecha(r[i_fecha] if i_fecha is not None else None)
        age = (asof - fd).days if fd else None
        if age is not None:
            ages.append(age)
        a = aim.get(pn, {"l1":"NO_AIM","l2":"","desc":""})
        l1 = a["l1"] or "NO_AIM"
        by_l1[l1]["lines"] += 1
        by_l1[l1]["eur"] += line_eur
        by_l1[l1]["qty"] += qty
        all_lines.append({
            "asof": asof.isoformat(),
            "ym": ym,
            "pn": pn,
            "desc": (r[i_desc] if i_desc is not None else "") or a.get("desc") or "",
            "qty_pending": qty,
            "qty_ordered": float(r[i_ord] or 0) if i_ord is not None and r[i_ord] is not None else None,
            "unit_price": pu,
            "eur_open": line_eur,
            "order_date": fd.isoformat() if fd else "",
            "age_days": age,
            "pedido": r[i_ped] if i_ped is not None else "",
            "tipo": r[i_tipo] if i_tipo is not None else "",
            "proveedor": r[i_prov] if i_prov is not None else "",
            "ref_prov": norm_pn(r[i_refp]) if i_refp is not None else "",
            "aim_l1": l1,
            "aim_l2": a.get("l2") or "",
        })
    ages_sorted = sorted(ages)
    def pct(p):
        if not ages_sorted: return None
        k = int(round((p/100)*(len(ages_sorted)-1)))
        return ages_sorted[k]
    monthly.append({
        "sheet": sn,
        "asof": asof.isoformat(),
        "ym": ym,
        "lines": lines,
        "lines_with_pn": with_pn,
        "eur_open": round(eur, 2),
        "empty": lines == 0,
        "age_min": ages_sorted[0] if ages_sorted else None,
        "age_p50": pct(50),
        "age_p90": pct(90),
        "age_max": ages_sorted[-1] if ages_sorted else None,
        "kpi_bo_lines": kpi_bo.get(ym),
        "by_aim_l1": {k: {"lines": v["lines"], "eur": round(v["eur"],2), "qty": v["qty"]} for k,v in sorted(by_l1.items(), key=lambda kv: -kv[1]["eur"])},
    })
wb.close()

# --- ISZE open orders 519 ---
wb = openpyxl.load_workbook(ORDER_XLSX, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
hdr = list(next(rows))
hi = {str(h):i for i,h in enumerate(hdr)}
our = []
for r in rows:
    if str(r[hi["Sold-to pt"]]) != "519":
        continue
    pn = norm_pn(r[hi["Material"]] or r[hi["Order Material"]])
    a = aim.get(pn, {"l1":"NO_AIM","l2":"","desc":""})
    created = r[hi["Created on"]]
    if isinstance(created, datetime):
        cd = created.date()
    elif isinstance(created, date):
        cd = created
    else:
        cd = None
    try:
        eur = float(r[hi["EUR Amount"]] or 0)
    except Exception:
        eur = 0.0
    try:
        qty = float(r[hi["Order Quantity"]] or 0)
    except Exception:
        qty = 0.0
    our.append({
        "pn": pn,
        "desc": r[hi["Material Number"]] or a.get("desc") or "",
        "eur": eur,
        "qty": qty,
        "saty": r[hi["SaTy"]],
        "created": cd.isoformat() if cd else "",
        "aim_l1": a.get("l1") or "NO_AIM",
        "aim_l2": a.get("l2") or "",
        "sales_doc": r[hi["Sales Doc."]],
    })
wb.close()

our_by_pn = defaultdict(lambda: {"eur":0.0,"qty":0.0,"lines":0,"saty":set(),"aim_l1":"","aim_l2":"","desc":""})
for o in our:
    b = our_by_pn[o["pn"]]
    b["eur"] += o["eur"]
    b["qty"] += o["qty"]
    b["lines"] += 1
    b["saty"].add(str(o["saty"] or ""))
    b["aim_l1"] = o["aim_l1"]
    b["aim_l2"] = o["aim_l2"]
    b["desc"] = o["desc"]

# Latest populated Nuno month for match
populated = [m for m in monthly if not m["empty"]]
latest = populated[-1] if populated else None
latest_lines = [L for L in all_lines if latest and L["ym"] == latest["ym"]]
their_pn = defaultdict(lambda: {"eur":0.0,"qty":0.0,"lines":0,"age_max":0,"aim_l1":"","aim_l2":"","desc":"","pedido_dates":[]})
for L in latest_lines:
    b = their_pn[L["pn"]]
    b["eur"] += L["eur_open"]
    b["qty"] += L["qty_pending"]
    b["lines"] += 1
    b["age_max"] = max(b["age_max"], L["age_days"] or 0)
    b["aim_l1"] = L["aim_l1"]
    b["aim_l2"] = L["aim_l2"]
    b["desc"] = L["desc"]
    if L["order_date"]:
        b["pedido_dates"].append(L["order_date"])

both = sorted(set(their_pn) & set(our_by_pn))
only_them = sorted(set(their_pn) - set(our_by_pn))
only_us = sorted(set(our_by_pn) - set(their_pn))

def sum_eur(pns, side):
    if side == "them":
        return sum(their_pn[p]["eur"] for p in pns)
    return sum(our_by_pn[p]["eur"] for p in pns)

overlap = {
    "latest_their_ym": latest["ym"] if latest else None,
    "latest_their_asof": latest["asof"] if latest else None,
    "their_lines": latest["lines"] if latest else 0,
    "their_eur": latest["eur_open"] if latest else 0,
    "our_asof": "2026-10-06",
    "our_lines": len(our),
    "our_eur": round(sum(o["eur"] for o in our), 2),
    "pn_both": len(both),
    "pn_only_them": len(only_them),
    "pn_only_us": len(only_us),
    "eur_both_their": round(sum_eur(both, "them"), 2),
    "eur_both_our": round(sum_eur(both, "us"), 2),
    "eur_only_them": round(sum_eur(only_them, "them"), 2),
    "eur_only_us": round(sum_eur(only_us, "us"), 2),
}

# top both / only them / only us
def top_list(pns, side, n=25):
    rows = []
    for p in pns:
        if side == "them":
            t = their_pn[p]; o = our_by_pn.get(p)
            rows.append({
                "pn": p, "desc": t["desc"], "aim_l1": t["aim_l1"], "aim_l2": t["aim_l2"],
                "their_eur": round(t["eur"],2), "their_qty": t["qty"], "their_age_max": t["age_max"],
                "our_eur": round(o["eur"],2) if o else 0, "our_qty": o["qty"] if o else 0,
                "our_saty": ",".join(sorted(o["saty"])) if o else "",
            })
        else:
            o = our_by_pn[p]; t = their_pn.get(p)
            rows.append({
                "pn": p, "desc": o["desc"], "aim_l1": o["aim_l1"], "aim_l2": o["aim_l2"],
                "our_eur": round(o["eur"],2), "our_qty": o["qty"],
                "our_saty": ",".join(sorted(o["saty"])),
                "their_eur": round(t["eur"],2) if t else 0, "their_qty": t["qty"] if t else 0,
            })
    key = "their_eur" if side == "them" else "our_eur"
    rows.sort(key=lambda r: -r[key])
    return rows[:n]

# AIM history series for competitive share
history_aim = []
for m in monthly:
    if m["empty"]:
        history_aim.append({"ym": m["ym"], "lines": 0, "eur": 0, "competitive_eur": 0, "competitive_share": None})
        continue
    bl = m["by_aim_l1"]
    comp = bl.get("Competitive parts", {}).get("eur", 0)
    history_aim.append({
        "ym": m["ym"],
        "lines": m["lines"],
        "eur": m["eur_open"],
        "competitive_eur": comp,
        "competitive_share": round(comp / m["eur_open"], 4) if m["eur_open"] else None,
        "kpi_bo_lines": m.get("kpi_bo_lines"),
        "age_p50": m.get("age_p50"),
        "age_max": m.get("age_max"),
    })

summary = {
    "bo_file": str(BO_XLSX),
    "order_file": str(ORDER_XLSX),
    "note": "Nuno Backordered lines 2026.xlsx = Astara monthly BO extract (PN-level). Sep/Oct+ tabs empty in 29-Sep mail. Match vs ISZE open orders is as-of mismatch: their latest filled month vs our 6-Oct snapshot.",
    "monthly": monthly,
    "history_aim": history_aim,
    "overlap": overlap,
    "kpi_bo": kpi_bo,
}
(OUT / "summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")

# CSVs
with (OUT / "monthly_bo.csv").open("w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["ym","asof","lines","eur_open","kpi_bo_lines","age_p50","age_max","empty"])
    w.writeheader()
    for m in monthly:
        w.writerow({k: m.get(k) for k in w.fieldnames})

with (OUT / "overlap_both_top.csv").open("w", encoding="utf-8", newline="") as f:
    rows = top_list(both, "them", 40)
    if rows:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

with (OUT / "only_them_top.csv").open("w", encoding="utf-8", newline="") as f:
    rows = top_list(only_them, "them", 40)
    if rows:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

with (OUT / "only_us_top.csv").open("w", encoding="utf-8", newline="") as f:
    rows = top_list(only_us, "us", 40)
    if rows:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

# latest month line detail
with (OUT / "latest_their_lines.csv").open("w", encoding="utf-8", newline="") as f:
    if latest_lines:
        w = csv.DictWriter(f, fieldnames=list(latest_lines[0].keys())); w.writeheader(); w.writerows(latest_lines)

print("OUT", OUT)
print("populated months", [m["ym"] for m in populated])
for m in monthly:
    print(m["ym"], "lines", m["lines"], "eur", m["eur_open"], "kpi", m.get("kpi_bo_lines"), "empty", m["empty"], "age_p50", m.get("age_p50"))
print("OVERLAP", json.dumps(overlap, indent=2))
print("both sample", top_list(both, "them", 5))
print("only them sample", top_list(only_them, "them", 5))
print("only us sample", top_list(only_us, "us", 5))
