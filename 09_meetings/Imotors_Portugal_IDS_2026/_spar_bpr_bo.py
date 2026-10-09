from pathlib import Path
import json

bpr = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR")
print("BPR files:")
for p in sorted(bpr.glob("*")):
    if p.is_file():
        print(f"  {p.name}\t{p.stat().st_size}")

# try read pptx slide titles via python-pptx
pptx = list(bpr.glob("*.pptx"))
if pptx:
    from pptx import Presentation
    prs = Presentation(str(pptx[0]))
    print("PPTX", pptx[0].name, "slides", len(prs.slides))
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip().replace("\n", " | ")
                if t:
                    texts.append(t[:120])
        title = texts[0] if texts else "(no text)"
        print(f"  {i:02d}: {title}")
        for t in texts[1:3]:
            print(f"      · {t}")

# BO summary key points
s = json.loads(Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_imotors_bo_compare_20261006\IMOTORS_BO_HISTORY_ANALYSIS_20261006\summary.json").read_text(encoding="utf-8"))
ov = s["overlap"]
print("OVERLAP", ov)
print("months filled", [m["ym"] for m in s["monthly"] if not m.get("empty")])
