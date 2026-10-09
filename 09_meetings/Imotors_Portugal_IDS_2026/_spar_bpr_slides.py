from pptx import Presentation
from pathlib import Path
p = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\IMotors_FY2027_BPR\I_Motors_FY2027_Business_Performance_Review_2026-10-06_Sep_closed.pptx")
prs = Presentation(str(p))
print("slides", len(prs.slides), p.name)
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip().replace("\n", " | ")
            if t:
                texts.append(t[:160])
    blob = " ".join(texts).lower()
    if i in (1, 3, 15, 16, 17, 26, 27) or any(x in blob for x in ["fill", "back order", "back-order", "mos", "avail"]):
        print(f"{i:02d}: {texts[0] if texts else '-'}")
        for t in texts[1:5]:
            print("   ", t)
