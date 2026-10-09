# -*- coding: utf-8 -*-
from pathlib import Path
import json, csv, re
from collections import defaultdict
from datetime import datetime, date
import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BO_XLSX = ROOT / "01_data/90_INBOX/_nuno_backordered_lines_20261006/20260929_1531_Nuno_Backordered lines 2026.xlsx"
ORDER_XLSX = ROOT / "01_data/90_INBOX/_imotors_bo_compare_20261006/EXPORT_ISZE_Order_20261006.xlsx"
AIM = ROOT / "01_data/00_derived/_releases/aim_pn_classification_release_2026-10-04.csv"
OUT = ROOT / "01_data/90_INBOX/_imotors_bo_compare_20261006/IMOTORS_BO_HISTORY_ANALYSIS_20261006"
OUT.mkdir(parents=True, exist_ok=True)

def norm_pn(x):
    if x is None: return ""
    s = str(x).strip().upper()
    if s.endswith(".0") and s[:-2].isdigit(): s = s[:-2]
    return s

def parse_fecha(v):
    if v is None or v == "": return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    s = str(v).strip()
    if re.fullmatch(r"\d{8}", s): return datetime.strptime(s, "%Y%m%d").date()
    return None

aim = {}
with AIM.open(encoding="utf-8", newline="") as f:
    for row in csv.DictReader(f):
        pn = norm_pn(row.get("part_number"))
        if not pn: continue
        # commercial catalogue cut (same as prior Portugal AIM work)
        l1 = (row.get("catalog_l1") or row.get("aim_l1_fixed") or "NO_AIM").strip() or "NO_AIM"
        l2 = (row.get("catalog_l2") or row.get("aim_l2_fixed") or "").strip()
        aim[pn] = {"l1": l1, "l2": l2, "desc": row.get("description_norm") or ""}

KPI_BO = {"2026-01":430,"2026-02":550,"2026-03":447,"2026-04":430,"2026-05":487,"2026-06":385,"2026-07":279,"2026-08":282,"2026-09":495}
KPI_FILL = {"2026-01":0.7852,"2026-02":0.761,"2026-03":0.8406,"2026-04":0.8455,"2026-05":0.8395,"2026-06":0.8541,"2026-07":0.8362,"2026-08":0.8539,"2026-09":0.8105}
KPI_MOS = {"2026-01":2.5,"2026-02":2.7,"2026-03":3.2,"2026-04":2.87,"2026-05":2.83,"2026-06":2.77,"2026-07":2.62,"2026-08":2.20,"2026-09":2.49}
KPI_STD = {"2026-01":415020,"2026-02":339278,"2026-03":381438,"2026-04":297074,"2026-05":489770,"2026-06":379751,"2026-07":393481,"2026-08":303777,"2026-09":416806}

wb = openpyxl.load_workbook(BO_XLSX, data_only=True, read_only=True)
monthly = []
all_lines = []
for sn in wb.sheetnames:
    m = re.search(r"BO (\d{8})", sn)
    if not m: continue
    asof = datetime.strptime(m.group(1), "%Y%m%d").date()
    ym = asof.strftime("%Y-%m")
    rows = list(wb[sn].iter_rows(values_only=True))
    if not rows:
        monthly.append({"sheet":sn,"asof":asof.isoformat(),"ym":ym,"lines":0,"eur_open":0.0,"empty":True})
        continue
    hdr = [str(h).strip().upper() if h else "" for h in rows[0]]
    idx = {h:i for i,h in enumerate(hdr)}
    def c(*names):
        for n in names:
            if n in idx: return idx[n]
        return None
    i_ref,i_desc,i_qty,i_pu,i_fecha,i_ped,i_tipo = c("REFERENCIA"),c("DESCRIPCION"),c("QTY PDTE"),c("P.UNIT"),c("FECHA"),c("PEDIDO"),c("TIPO")
    lines=0; eur=0.0; ages=[]; by_l1=defaultdict(lambda:{"lines":0,"eur":0.0})
    by_l2=defaultdict(lambda:{"lines":0,"eur":0.0,"l1":""})
    for r in rows[1:]:
        if not r or all(v is None or str(v).strip()=="" for v in r): continue
        pn = norm_pn(r[i_ref] if i_ref is not None else None)
        if not pn: continue
        lines += 1
        qty = float(r[i_qty] or 0); pu = float(r[i_pu] or 0); line_eur = qty*pu; eur += line_eur
        fd = parse_fecha(r[i_fecha] if i_fecha is not None else None)
        age = (asof - fd).days if fd else None
        if age is not None: ages.append(age)
        a = aim.get(pn, {"l1":"NO_AIM","l2":"","desc":""})
        l1 = a["l1"] or "NO_AIM"; l2 = a["l2"] or ""
        by_l1[l1]["lines"]+=1; by_l1[l1]["eur"]+=line_eur
        by_l2[(l1,l2)]["lines"]+=1; by_l2[(l1,l2)]["eur"]+=line_eur; by_l2[(l1,l2)]["l1"]=l1
        all_lines.append({"asof":asof.isoformat(),"ym":ym,"pn":pn,"desc":(r[i_desc] if i_desc is not None else "") or a.get("desc") or "",
            "qty_pending":qty,"unit_price":pu,"eur_open":line_eur,"order_date":fd.isoformat() if fd else "",
            "age_days":age,"pedido":r[i_ped] if i_ped is not None else "","tipo":r[i_tipo] if i_tipo is not None else "",
            "aim_l1":l1,"aim_l2":l2})
    ages_sorted=sorted(ages)
    def pct(p):
        if not ages_sorted: return None
        return ages_sorted[int(round((p/100)*(len(ages_sorted)-1)))]
    monthly.append({
        "sheet":sn,"asof":asof.isoformat(),"ym":ym,"lines":lines,"eur_open":round(eur,2),"empty":lines==0,
        "age_min":ages_sorted[0] if ages_sorted else None,"age_p50":pct(50),"age_p90":pct(90),"age_max":ages_sorted[-1] if ages_sorted else None,
        "kpi_bo_lines":KPI_BO.get(ym),"kpi_fill":KPI_FILL.get(ym),"kpi_mos":KPI_MOS.get(ym),"kpi_std_eur":KPI_STD.get(ym),
        "by_aim_l1":{k:{"lines":v["lines"],"eur":round(v["eur"],2)} for k,v in sorted(by_l1.items(), key=lambda kv:-kv[1]["eur"])},
        "by_aim_l2":[{"l1":k[0],"l2":k[1],"lines":v["lines"],"eur":round(v["eur"],2)} for k,v in sorted(by_l2.items(), key=lambda kv:-kv[1]["eur"])[:12]],
    })
wb.close()

# ISZE open
wb=openpyxl.load_workbook(ORDER_XLSX, data_only=True, read_only=True)
ws=wb[wb.sheetnames[0]]; rows=ws.iter_rows(values_only=True); hdr=list(next(rows)); hi={str(h):i for i,h in enumerate(hdr)}
our=[]
for r in rows:
    if str(r[hi["Sold-to pt"]])!="519": continue
    pn=norm_pn(r[hi["Material"]] or r[hi["Order Material"]])
    a=aim.get(pn,{"l1":"NO_AIM","l2":"","desc":""})
    created=r[hi["Created on"]]; cd=created.date() if isinstance(created,datetime) else (created if isinstance(created,date) else None)
    eur=float(r[hi["EUR Amount"]] or 0); qty=float(r[hi["Order Quantity"]] or 0)
    our.append({"pn":pn,"desc":r[hi["Material Number"]] or a.get("desc") or "","eur":eur,"qty":qty,"saty":r[hi["SaTy"]],
        "created":cd.isoformat() if cd else "","aim_l1":a.get("l1") or "NO_AIM","aim_l2":a.get("l2") or ""})
wb.close()
our_by_pn=defaultdict(lambda:{"eur":0.0,"qty":0.0,"lines":0,"saty":set(),"aim_l1":"","aim_l2":"","desc":""})
for o in our:
    b=our_by_pn[o["pn"]]; b["eur"]+=o["eur"]; b["qty"]+=o["qty"]; b["lines"]+=1; b["saty"].add(str(o["saty"] or "")); b["aim_l1"]=o["aim_l1"]; b["aim_l2"]=o["aim_l2"]; b["desc"]=o["desc"]

populated=[m for m in monthly if not m["empty"]]; latest=populated[-1]
latest_lines=[L for L in all_lines if L["ym"]==latest["ym"]]
their_pn=defaultdict(lambda:{"eur":0.0,"qty":0.0,"lines":0,"age_max":0,"aim_l1":"","aim_l2":"","desc":""})
for L in latest_lines:
    b=their_pn[L["pn"]]; b["eur"]+=L["eur_open"]; b["qty"]+=L["qty_pending"]; b["lines"]+=1
    b["age_max"]=max(b["age_max"], L["age_days"] or 0); b["aim_l1"]=L["aim_l1"]; b["aim_l2"]=L["aim_l2"]; b["desc"]=L["desc"]

both=sorted(set(their_pn)&set(our_by_pn)); only_them=sorted(set(their_pn)-set(our_by_pn)); only_us=sorted(set(our_by_pn)-set(their_pn))

def top(pns, side, n=30):
    rows=[]
    for p in pns:
        if side=="them":
            t=their_pn[p]; o=our_by_pn.get(p)
            rows.append({"pn":p,"desc":t["desc"],"aim_l1":t["aim_l1"],"aim_l2":t["aim_l2"],"their_eur":round(t["eur"],2),"their_qty":t["qty"],"their_age_max":t["age_max"],
                "our_eur":round(o["eur"],2) if o else 0,"our_qty":o["qty"] if o else 0,"our_saty":",".join(sorted(o["saty"])) if o else ""})
        else:
            o=our_by_pn[p]; t=their_pn.get(p)
            rows.append({"pn":p,"desc":o["desc"],"aim_l1":o["aim_l1"],"aim_l2":o["aim_l2"],"our_eur":round(o["eur"],2),"our_qty":o["qty"],
                "our_saty":",".join(sorted(o["saty"])),"their_eur":round(t["eur"],2) if t else 0,"their_qty":t["qty"] if t else 0})
    rows.sort(key=lambda r: -r["their_eur" if side=="them" else "our_eur"]); return rows[:n]

# persistence: PNs open in Aug that were also open in prior months
aug_pns=set(their_pn)
persist={}
for ym in ["2026-07","2026-06","2026-05","2026-04","2026-03","2026-02","2026-01"]:
    prev=set(L["pn"] for L in all_lines if L["ym"]==ym)
    persist[ym]={"pn_overlap_with_aug":len(aug_pns & prev),"share_of_aug_pns":round(len(aug_pns & prev)/len(aug_pns),3) if aug_pns else None}

# age bands latest
bands=[("0-7",0,7),("8-14",8,14),("15-30",15,30),("31-60",31,60),("61-90",61,90),("91+",91,99999)]
age_band={"band":[], "lines":[], "eur":[]}
for name,a,b in bands:
    ls=[L for L in latest_lines if L["age_days"] is not None and a<=L["age_days"]<=b]
    age_band["band"].append(name); age_band["lines"].append(len(ls)); age_band["eur"].append(round(sum(L["eur_open"] for L in ls),2))

overlap={
    "latest_their_ym":latest["ym"],"latest_their_asof":latest["asof"],
    "their_lines":latest["lines"],"their_eur":latest["eur_open"],
    "our_asof":"2026-10-06","our_lines":len(our),"our_eur":round(sum(o["eur"] for o in our),2),
    "pn_both":len(both),"pn_only_them":len(only_them),"pn_only_us":len(only_us),
    "eur_both_their":round(sum(their_pn[p]["eur"] for p in both),2),
    "eur_both_our":round(sum(our_by_pn[p]["eur"] for p in both),2),
    "eur_only_them":round(sum(their_pn[p]["eur"] for p in only_them),2),
    "eur_only_us":round(sum(our_by_pn[p]["eur"] for p in only_us),2),
    "asof_gap_note":"Their filled tabs stop at 31 Aug 2026 (mail 29 Sep). Our open-order extract is 6 Oct. Sep KPI says 495 B/O lines but Sep tab in this workbook is empty — PN match for the Sep spike is not in this file.",
}

# competitive share history
for m in monthly:
    if m.get("empty"): 
        m["competitive_eur"]=0; m["competitive_share"]=None
    else:
        ce=m["by_aim_l1"].get("Competitive parts",{}).get("eur",0)
        m["competitive_eur"]=ce
        m["competitive_share"]=round(ce/m["eur_open"],4) if m["eur_open"] else None

summary={
    "bo_file":str(BO_XLSX),"order_file":str(ORDER_XLSX),
    "kpi_match":"Nuno BO line counts Jan–Aug equal KPI B/O lines exactly — KPI count is this extract.",
    "monthly":monthly,"overlap":overlap,"persist_aug_vs_prior":persist,"age_band_aug":age_band,
    "top_both":top(both,"them",30),"top_only_them":top(only_them,"them",30),"top_only_us":top(only_us,"us",30),
    "aug_by_l1":latest["by_aim_l1"],"aug_by_l2":latest["by_aim_l2"],
}
(OUT/"summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")

with (OUT/"monthly_bo.csv").open("w",encoding="utf-8",newline="") as f:
    w=csv.DictWriter(f, fieldnames=["ym","asof","lines","eur_open","kpi_bo_lines","kpi_fill","kpi_mos","kpi_std_eur","competitive_eur","competitive_share","age_p50","age_max","empty"])
    w.writeheader()
    for m in monthly:
        w.writerow({k:m.get(k) for k in w.fieldnames})

for name, rows in [("overlap_both_top.csv", top(both,"them",40)), ("only_them_top.csv", top(only_them,"them",40)), ("only_us_top.csv", top(only_us,"us",40))]:
    with (OUT/name).open("w",encoding="utf-8",newline="") as f:
        if rows:
            w=csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

with (OUT/"latest_their_lines.csv").open("w",encoding="utf-8",newline="") as f:
    if latest_lines:
        w=csv.DictWriter(f, fieldnames=list(latest_lines[0].keys())); w.writeheader(); w.writerows(latest_lines)

print(json.dumps({"overlap":overlap,"persist":persist,"age_band":age_band,"aug_l1":latest["by_aim_l1"],
    "monthly_compact":[{k:m.get(k) for k in ["ym","lines","eur_open","competitive_share","age_p50","age_max","kpi_fill"]} for m in monthly if not m.get("empty")]}, indent=2))
print("TOP BOTH", top(both,"them",8))
print("TOP ONLY THEM", top(only_them,"them",8))
