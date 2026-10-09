"""Dump slides 6, 7 and 10 of the v2 deck in detail (shape id, name, geometry, text runs)
so the pass-2 edit can target exact shapes."""
from __future__ import annotations
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
prs = Presentation(str(HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"))
EMU = 914400
for idx in (6, 7, 10):
    s = list(prs.slides)[idx - 1]
    print("=" * 100)
    print("SLIDE", idx, "shapes", len(s.shapes))
    for sh in s.shapes:
        try:
            geom = str(sh.auto_shape_type)
        except Exception:
            geom = "(n/a)"
        kind = type(sh).__name__
        print("  id=%-4s %-14s %-22s %-20s L%.3f T%.3f W%.3f H%.3f" % (
            sh.shape_id, kind, sh.name, geom,
            (sh.left or 0) / EMU, (sh.top or 0) / EMU, (sh.width or 0) / EMU, (sh.height or 0) / EMU))
        if sh.has_text_frame:
            for pi, p in enumerate(sh.text_frame.paragraphs):
                txt = "".join(r.text for r in p.runs)
                if not txt.strip():
                    continue
                r0 = p.runs[0]
                sz = r0.font.size.pt if r0.font.size else None
                col = None
                try:
                    col = str(r0.font.color.rgb)
                except Exception:
                    pass
                print("        P%d sz=%s b=%s col=%s align=%s :: %r" % (
                    pi, sz, r0.font.bold, col, p.alignment, txt))
        if getattr(sh, "has_chart", False):
            ch = sh.chart
            print("        CHART type=%s has_title=%s" % (ch.chart_type,
                  ch.has_title and ch.chart_title.text_frame.text))
