# -*- coding: utf-8 -*-
"""Pass 3d: update the speaker notes on the three slides whose content changed."""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
DECK = HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"

prs = Presentation(str(DECK))
sl = list(prs.slides)

EDITS = {
    5: (
        "PRESENTER\n"
        "The request is to retain the existing programme. This table is the financial reference: eleven "
        "customers, with the original plan, increase, rate and incentive values unchanged.\n\n"
        "EDITORIAL / CONFIRMATION NOTES (revision 2026-10-07, third pass)\n"
        "User feedback: the slide put too much emphasis on 'retain'. The heading now simply states the "
        "commitment; the closing sentence states the position once and stops. No amount, percentage, decimal "
        "precision, customer row or total has been recalculated or corrected; source-cent differences are "
        "retained. This detailed source labels EUR 393,000 as the FY2027 commitment. It does not, by itself, "
        "verify the FY2026 amount actually earned or paid.",
        350,
    ),
    8: (
        "PRESENTER\n"
        "Use the actual result as evidence of plan achievement, not as a return-on-incentive calculation. "
        "Explain the population: ITUK is an intercompany case and is not in this nine-distributor comparison.\n"
        "Then present the operating KPIs, which are a different question from the billing result: fill rate "
        "against the 95% standard and 92% floor, days of stock against each distributor's own supply-class band, "
        "and back orders against the 15-lines-per-EUR-100K ceiling. Say which distributors report and which do "
        "not, rather than leaving a blank.\n\n"
        "EDITORIAL / CONFIRMATION NOTES (revision 2026-10-07, third pass)\n"
        "User feedback: other KPIs give a better picture. A three-tile KPI band was added below the billing "
        "table. Figures come from KPI Master File august 2026.xlsx, sheet 'Parts Handling' - 'First Fill Rate', "
        "'Days of Parts in Stock', 'Back Order Lines' and 'Parts Sales DB to DLR in EUR'. The latest reported "
        "month per distributor was taken from the last fourteen reported months in that file; older cells in the "
        "same row belong to a period in which the distributor had gone dark and would be stale if quoted. Bands "
        "and ceilings come from ISZE_DB_KPI_Targets_FY2627_v1.xlsx ('Targets FY26-27' columns H-I and the "
        "Assumptions sheet). Distribution: fill rate - five of the six reporting distributors at or above the "
        "92% floor, Portugal (Imotors) below at 83.6% in July 2026, no value on file for UMI, Germany (ISD) or "
        "Benelux (IBX). Days of stock - four of the seven reporting sit above their own band, Benelux (IBX) 63 "
        "days and Portugal (Imotors) 79 days inside, Sweden 48 days below; no value for UMI or Germany (ISD). "
        "Back orders - four of the six with a usable sales base are above the 15-lines-per-EUR-100K ceiling, IUK "
        "134 lines and Ireland 20 lines inside; no value or no sales base for UMI, Germany (ISD) or Benelux "
        "(IBX). No KPI has been invented and no blank has been left without a stated reason. The former "
        "standalone footnote strip was folded into the source line to make room.",
        700,
    ),
    20: (
        "PRESENTER\n"
        "Close on the specific decision: retain the incentive at the current rate and total, with its target "
        "and reporting conditions. The cycle on the slide is the review mechanism, not the decision: confirm "
        "the monthly plan, check at month 6 and month 9 against the 10% tolerance, close the five reporting "
        "gaps, then reset the plan from the agreed monthly targets. The separate sales-to-dealer performance "
        "element is an approval item in its own right and needs agreed coverage before adoption.\n\n"
        "EDITORIAL / CONFIRMATION NOTES (revision 2026-10-07, third pass)\n"
        "User feedback: the slide could be much better, and suggested a circular arrangement for continuous "
        "improvement. The table was rebuilt as a four-step review cycle, with the plan reset feeding back into "
        "the next checkpoint. The fifth item - reviewing a separate sales-to-dealer performance element - was "
        "not put in the loop because it needs its own approval and has no owner in the source; it is shown "
        "beside the cycle instead, so nothing from the original five actions is lost. No change to the "
        "incentive rate or total is proposed by this slide.",
        600,
    ),
}

for n, (text, keep) in EDITS.items():
    sl[n - 1].notes_slide.notes_text_frame.text = text
    print("notes updated on slide", n)

prs.save(str(DECK))
print("saved")
