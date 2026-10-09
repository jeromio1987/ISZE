from pptx import Presentation
from pptx.util import Emu

path = r'C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx'
p = Presentation(path)
lines = []
lines.append('slide size EMU %d x %d = %.3f x %.3f in' % (p.slide_width, p.slide_height, p.slide_width / 914400, p.slide_height / 914400))

def hexof(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    try:
        bg = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        if bg is not None:
            return bg.get('val')
    except Exception:
        pass
    return None


def emu(v):
    return round(v / 914400, 2) if v is not None else None


for i, s in enumerate(p.slides, 1):
    lines.append('--- slide %d  layout=%s' % (i, s.slide_layout.name))
    lines.append('    bg xml: ' + (s._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}bg') is not None and 'yes' or 'no'))
    for sh in s.shapes:
        t = ''
        if sh.has_text_frame:
            t = ' | '.join(run.text for para in sh.text_frame.paragraphs for run in para.runs)[:120]
        fill = hexof(sh)
        lines.append('    [%s] name=%s pos=(%s,%s) size=(%s,%s) fill=%s txt=%s' % (
            sh.shape_type, sh.name, emu(sh.left), emu(sh.top), emu(sh.width), emu(sh.height), fill, t.replace('\n', ' ')))

open(r'C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\_build\_ref_geom.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('ok')
