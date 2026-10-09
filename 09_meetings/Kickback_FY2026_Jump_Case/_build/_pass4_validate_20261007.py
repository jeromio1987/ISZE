import hashlib
import math
import posixpath
import re
import sys
import zipfile

import lxml.etree as E
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

P = r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"
b = open(P, "rb").read()
print("bytes", len(b), "sha256", hashlib.sha256(b).hexdigest().upper())

prs = Presentation(P)
sl = list(prs.slides)
print("slides", len(sl))
SW, SH = prs.slide_width / 914400, prs.slide_height / 914400

off = []
for i, s in enumerate(sl, 1):
    for sh in s.shapes:
        if sh.left is None:
            continue
        l, t = (sh.left or 0) / 914400, (sh.top or 0) / 914400
        w, h = (sh.width or 0) / 914400, (sh.height or 0) / 914400
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            off.append(i)
print("off-slide shapes:", len(off), off)

est = []
for i, s in enumerate(sl, 1):
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        sz = None
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    sz = r.font.size.pt
                    break
            if sz:
                break
        if not sz:
            continue
        w = (sh.width or 0) / 914400
        if w <= 0.2:
            continue
        cpl = max(1, int(w * 96 / (sz * 0.523)))
        need = 0
        for p in sh.text_frame.paragraphs:
            n = len(p.text)
            need += max(1, math.ceil(n / cpl)) if n else 1
        if need * sz * 1.22 / 72.0 > (sh.height or 0) / 914400 + 0.06:
            est.append((i, sh.text_frame.text.strip()[:40]))
print("estimated overflow boxes:", len(est), est)

banned = ["sell-out", "sellout", "sell out", "mode a", "mode b", "0.82",
          "1,084", "reduced 20%", "top six", "minimum month", "sales-to-dealer evidence"]
hits = []
for i, s in enumerate(sl, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for bb in banned:
                if bb in sh.text_frame.text.lower():
                    hits.append((i, bb))
print("banned hits:", len(hits), hits)

follow = [i for i, s in enumerate(sl, 1)
          if any(sh.has_text_frame and "Confirm the remaining brake-pad stock" in sh.text_frame.text
                 for sh in s.shapes)]
print("follow-up present on slides:", follow)
for i, s in enumerate(sl, 1):
    for sh in s.shapes:
        if sh.has_text_frame and "Confirm the remaining brake-pad stock" in sh.text_frame.text:
            print("   paragraph count:", len(sh.text_frame.paragraphs))
            for j, p in enumerate(sh.text_frame.paragraphs):
                print("   p%d runs=%d %r" % (j, len(p.runs), p.text))

for i, s in enumerate(sl, 1):
    for sh in s.shapes:
        if sh.has_text_frame and "76%" in sh.text_frame.text:
            print("76%% on slide", i, "->", repr(sh.text_frame.text.strip()))

z = zipfile.ZipFile(P)
print("testzip", z.testzip())
names = set(z.namelist())
slides = [n for n in names if re.match(r"ppt/slides/slide[0-9]+[.]xml", n)]
print("slide parts", len(slides))
ct = z.read("[Content_Types].xml").decode("utf-8")
print("missing content types:", [n for n in slides if "/" + n not in ct])
print("missing rels:", [n for n in slides
                        if n.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels" not in names])
unres = 0
for n in names:
    if not n.endswith(".rels"):
        continue
    base = posixpath.dirname(posixpath.dirname(n))
    for rel in E.fromstring(z.read(n)):
        if rel.get("TargetMode") == "External":
            continue
        t = rel.get("Target")
        if t and posixpath.normpath(posixpath.join(base, t)) not in names:
            unres += 1
            print("   UNRESOLVED", n, t)
print("unresolved rels:", unres)
bad = 0
for n in slides + ["ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"]:
    try:
        E.fromstring(z.read(n))
    except Exception as e:
        bad += 1
        print("   MALFORMED", n, e)
print("malformed:", bad)
pres = E.fromstring(z.read("ppt/presentation.xml"))
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
ids = [e.get(R) for e in pres.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}sldId")]
tgt = {rel.get("Id"): rel.get("Target") for rel in E.fromstring(z.read("ppt/_rels/presentation.xml.rels"))}
print("sldIds", len(ids), "all resolve:", all(i in tgt for i in ids))
