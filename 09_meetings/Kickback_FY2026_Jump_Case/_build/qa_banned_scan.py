"""Banned-string scan helper (reusable) — text of every slide in a pptx."""
from __future__ import annotations
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BANNED = ["sell-out", "sell out", "sellout", "Mode A", "Mode B",
          "reduced 20%", "Maintenance family reduced", "minimum months of stock"]


def slide_texts(pptx):
    prs = Presentation(str(pptx))
    out = []
    for i, s in enumerate(prs.slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for c in row.cells:
                        parts.append(c.text)
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            parts.append("\nNOTES:\n" + s.notes_slide.notes_text_frame.text)
        out.append((i, "\n".join(parts)))
    return out


if __name__ == "__main__":
    p = Path(sys.argv[1])
    texts = slide_texts(p)
    print("slides:", len(texts))
    found = False
    for i, t in texts:
        for b in BANNED:
            if b.lower() in t.lower():
                for line in t.splitlines():
                    if b.lower() in line.lower():
                        print("  HIT S%02d [%s] %s" % (i, b, line.strip()[:160]))
                found = True
    print("BANNED_HITS:", "YES" if found else "NONE")
