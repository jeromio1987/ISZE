import io
from pptx import Presentation

path = r'C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx'
out = io.StringIO()
p = Presentation(path)


def a(x):
    return str(x).encode('ascii', 'replace').decode('ascii')


print('slide size', p.slide_width, p.slide_height, p.slide_width / 914400, p.slide_height / 914400, file=out)
for i, s in enumerate(p.slides, 1):
    print('=== slide', i, file=out)
    for sh in s.shapes:
        print('  ', sh.shape_type, a(sh.name), sh.left, sh.top, sh.width, sh.height, file=out)
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                t = ''.join(r.text for r in para.runs)
                if t.strip():
                    r0 = para.runs[0]
                    try:
                        col = r0.font.color.rgb
                    except Exception:
                        col = None
                    print('       |', a(repr(t[:150])), r0.font.size, r0.font.name, r0.font.bold, col, file=out)
        try:
            if sh.fill.type is not None and sh.fill.type == 1:
                print('        fill=', sh.fill.fore_color.rgb, file=out)
        except Exception:
            pass
    if s.has_notes_slide:
        nt = s.notes_slide.notes_text_frame.text
        print('   NOTES:', a(repr(nt[:200])), file=out)

open(r'C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\_build\_ref_deck_dump.txt', 'w', encoding='ascii', errors='replace').write(out.getvalue())
print('OK', len(p.slides))
