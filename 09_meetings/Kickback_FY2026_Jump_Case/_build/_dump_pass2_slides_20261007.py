"""Dump text + geometry of the new slides, and run off-slide / overflow checks on the whole deck."""
from __future__ import annotations
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
prs = Presentation(str(HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"))
EMU = 914400
SW, SH = prs.slide_width, prs.slide_height

for idx in (10, 11, 21, 22):
    s = list(prs.slides)[idx - 1]
    print("=" * 108)
    print("SLIDE", idx)
    for sh in s.shapes:
        try:
            geom = str(sh.auto_shape_type)
        except Exception:
            geom = "(n/a)"
        L, T = (sh.left or 0) / EMU, (sh.top or 0) / EMU
        W, H = (sh.width or 0) / EMU, (sh.height or 0) / EMU
        print("  id=%-4s %-14s %-22s L%.3f T%.3f W%.3f H%.3f B%.3f R%.3f" % (
            sh.shape_id, type(sh).__name__, sh.name, L, T, W, H, T + H, L + W))
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt.strip():
                    continue
                r0 = p.runs[0]
                sz = r0.font.size.pt if r0.font.size else None
                col = None
                try:
                    col = str(r0.font.color.rgb)
                except Exception:
                    pass
                print("        sz=%s b=%s col=%s :: %r" % (sz, r0.font.bold, col, txt[:150]))

print("\n" + "=" * 108)
print("OFF-SLIDE / OVERFLOW across the deck (8.5pt worst case)")
issues = []
for si, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.left is None:
            continue
        l, t, w, h = sh.left, sh.top, sh.width, sh.height
        if l < -10000 or t < -10000 or l + w > SW + 10000 or t + h > SH + 10000:
            issues.append("S%02d OUTSIDE %s l=%.3f t=%.3f r=%.3f b=%.3f" % (
                si, sh.name, l / EMU, t / EMU, (l + w) / EMU, (t + h) / EMU))
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            box_w, box_h = (w or 0) / EMU, (h or 0) / EMU
            total = 0.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    total += 0.10
                    continue
                size = None
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                size = size or 8.5
                line_h = size * 1.22 / 72
                cpl = max(1, int(box_w * 72 / (size * 0.52)))
                lines = max(1, -(-len(txt) // cpl))
                total += lines * line_h * 1.10
            if box_h > 0.05 and total > box_h * 1.02:
                issues.append("S%02d OVERFLOW %s est %.2f > box %.2f  %r" % (
                    si, sh.name, total, box_h, tf.text[:60].replace("\n", " / ")))
print("issues:", len(issues))
for i in issues:
    print("  ", i)
