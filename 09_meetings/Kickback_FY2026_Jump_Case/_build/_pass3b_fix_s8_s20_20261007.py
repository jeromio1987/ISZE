# -*- coding: utf-8 -*-
"""Pass 3b: enlarge the slide 8 KPI tiles, rebuild the slide 20 review cycle."""
from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
DECK = HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"

BG = RGBColor(0x26, 0x26, 0x26)
PANEL = RGBColor(0x30, 0x30, 0x30)
RED = RGBColor(0xE3, 0x06, 0x13)
INK = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xBD, 0xBD, 0xBD)
GREY = RGBColor(0x4B, 0x4B, 0x4B)
FONT = "Aptos"


def fill(sh, rgb):
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb
    sh.line.fill.background()
    sh.shadow.inherit = False


def say(shape, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.clear()
    for i, (txt, size, col, bold, sp) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if sp:
            p.space_before = Pt(sp)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = col
        r.font.bold = bold
        r.font.name = FONT


def canvas(slide, x, y, w, h, rgb=None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if rgb is None:
        sh.fill.background()
        sh.line.fill.background()
    else:
        fill(sh, rgb)
    sh.shadow.inherit = False
    return sh


def tbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb


prs = Presentation(str(DECK))
slides = list(prs.slides)

# ------------------------------------------------------------------ slide 8
s8 = slides[7]
# remove the previous, too-small tile band and its three accent bars
for sh in list(s8.shapes):
    if sh.top is None:
        continue
    top = sh.top / 914400.0
    if 5.99 <= top <= 6.90 and sh.left is not None and sh.left / 914400.0 > 0.5:
        sh._element.getparent().remove(sh._element)

canvas(s8, 0.42, 5.96, 12.49, 0.98, PANEL)
tiles = [
    ("FILL RATE  ·  STANDARD 95%, FLOOR 92%", "5 of 6",
     "reporting distributors at or above the 92% floor. Below it: Portugal (Imotors), 83.6%, Jul 2026. "
     "No value on file: UMI, Germany (ISD), Benelux (IBX)."),
    ("DAYS OF STOCK  ·  AGAINST SUPPLY-CLASS BAND", "4 of 7",
     "above their own band. Inside: Benelux (IBX) 63d, Portugal (Imotors) 79d. Below: Sweden 48d. "
     "No value on file: UMI, Germany (ISD)."),
    ("BACK ORDERS  ·  CEILING 15 LINES PER EUR 100K", "4 of 6",
     "above the ceiling. Inside: IUK 134 lines, Ireland 20 lines. No value or no sales base on file: "
     "UMI, Germany (ISD), Benelux (IBX)."),
]
for n, (head, val, sub) in enumerate(tiles):
    x = 0.62 + n * 4.21
    canvas(s8, x, 6.00, 0.045, 0.90, RED)
    tb = tbox(s8, x + 0.16, 6.02, 3.93, 0.90)
    say(tb, [(head, 7.8, MUTED, False, 0),
             (val, 16.0, INK, True, 1.5),
             (sub, 7.0, MUTED, False, 1.5)])
print("slide 8 tiles enlarged")

# ----------------------------------------------------------------- slide 20
s20 = slides[19]
for sh in list(s20.shapes):
    if sh.top is None:
        continue
    top = sh.top / 914400.0
    if top < 1.60 or top > 7.0:
        continue
    t = sh.text_frame.text.strip().lower() if sh.has_text_frame else ""
    if t.startswith("maintain the target"):
        continue
    sh._element.getparent().remove(sh._element)

cx, cy, R = 6.665, 4.72, 1.52

# the ring: an unfilled oval with a thick outline, so the centre stays clear
ring = canvas(s20, cx - R, cy - R, 2 * R, 2 * R, None, MSO_SHAPE.OVAL)
ring.fill.background()
ring.line.color.rgb = GREY
ring.line.width = Pt(7)

mid = tbox(s20, cx - 1.02, cy - 0.62, 2.04, 1.24)
say(mid, [("A cycle, not a\nsingle approval", 10.5, INK, True, 0),
          ("The plan is reset from the agreed monthly targets after every checkpoint.",
           7.0, MUTED, False, 3)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

nodes = [
    (-90, "1  ·  ISZE / Japan", "Confirm the monthly plan for each distributor"),
    (0, "2  ·  ISZE / Japan", "Check at month 6 against the 10% tolerance"),
    (90, "3  ·  ISZE / Japan", "Check at month 9 against the 10% tolerance"),
    (180, "4  ·  ISZE", "Close the five reporting gaps and load ISD Feb-Jul 2026"),
]
bw, bh = 2.34, 0.86
for a, tag, txt in nodes:
    rad = math.radians(a)
    px = cx + R * math.cos(rad) - bw / 2
    py = cy + R * math.sin(rad) - bh / 2
    canvas(s20, px, py, bw, bh, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE)
    canvas(s20, px, py, bw, 0.04, RED)
    tb = tbox(s20, px + 0.13, py + 0.14, bw - 0.26, bh - 0.28)
    say(tb, [(tag, 6.8, RED, True, 0), (txt, 7.8, INK, False, 1.6)])

for a in (-45, 45, 135, 225):
    rad = math.radians(a)
    ax, ay = cx + R * math.cos(rad), cy + R * math.sin(rad)
    arrow = canvas(s20, ax - 0.18, ay - 0.13, 0.36, 0.26, RED, MSO_SHAPE.RIGHT_ARROW)
    arrow.rotation = a + 90

left = tbox(s20, 0.42, 2.40, 2.95, 4.30)
say(left, [("A CONTINUOUS REVIEW CYCLE", 8.8, RED, True, 0),
           ("The four steps run as a loop, not a one-off approval. Each checkpoint uses the agreed "
            "monthly plan, and the reporting gaps are closed before the next cycle begins.", 8.4, MUTED, False, 4),
           ("No change to the incentive rate or total.", 8.8, INK, True, 7)])

right = tbox(s20, 9.90, 2.40, 3.00, 4.30)
say(right, [("SEPARATE APPROVAL ITEM", 8.8, RED, True, 0),
            ("Review a separate sales-to-dealer performance element. Coverage and measurement must be "
             "agreed before adoption.", 8.4, MUTED, False, 4),
            ("WHAT THE CYCLE PROTECTS", 8.8, RED, True, 8),
            ("Retain the incentive to support sales targets and secure the reporting.", 8.4, MUTED, False, 4),
            ("Use that reporting to manage stock risk and target commercial support.", 8.4, MUTED, False, 4)])

prs.save(str(DECK))
print("slide 20 rebuilt; saved")
