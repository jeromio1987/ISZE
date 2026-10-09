"""Confirm the only change to preserved slides is the page-number marker, plus the intended edits."""
from __future__ import annotations
import sys
import re
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
OLD = HERE / "_build" / "_backup" / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pass1.pptx"
NEW = HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"


def norm(path):
    """slide text with the small right-aligned page marker replaced by <PAGE>."""
    prs = Presentation(str(path))
    out = []
    for s in prs.slides:
        parts = []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text
            if sh.left is not None and abs(sh.left / 914400 - 12.20) < 0.02 \
                    and t.strip().isdigit() and len(t.strip()) == 2:
                t = "<PAGE>"
            parts.append(t)
        out.append("\n".join(parts))
    return prs, out


_, a = norm(OLD)
prs_new, b = norm(NEW)
print("pass1 slides:", len(a), " pass2 slides:", len(b))

# match each pass2 slide to a pass1 slide
unmatched_new = []
for i, t in enumerate(b, 1):
    if t in a:
        continue
    # try to find a pass1 slide whose text is a close match (the corrected wider review)
    best = None
    for j, u in enumerate(a, 1):
        if u[:400] == t[:400]:
            best = j
            break
    unmatched_new.append((i, best, t[:90].replace("\n", " / ")))

print("\nslides whose normalised text is not present in pass 1:")
for i, best, t in unmatched_new:
    print("   pass2 slide %02d  (closest pass1 %s)  %s" % (i, best, t))

# confirm 1..9 and the closing slides are exact
print("\nOK check: pass2 slides 1-9 == pass1 slides 1-9 ?", a[:9] == b[:9])
print("OK check: pass2 slide 23 == pass1 slide 21 (the tension chain) ?", a[20] == b[22])
print("OK check: pass2 slide 22 == pass1 slide 20 (decision) ?", a[19] == b[21])
print("OK check: pass2 slides 12-20 == pass1 slides 10-18 ?", a[9:18] == b[11:20])
