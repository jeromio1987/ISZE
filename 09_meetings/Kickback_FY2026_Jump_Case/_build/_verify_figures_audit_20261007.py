"""Standalone verification of the retired-figure list (2026-10-07).

Reads source only. Saves nothing. Prints raw numbers so the audit deck can be
built from confirmed figures rather than from the list on trust.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
DER = ROOT / "01_data" / "00_derived"
PLAN = ROOT / "01_data" / "11_planning"

CODES = {
    "ITUK": [2575], "IUK": [2574], "UMI": [1797], "UTI": [1799],
    "Benelux": [90], "Germany": [93, 2516, 601], "Portugal": [519, 92],
    "Ireland": [76], "Norway": [86], "Sweden": [2181],
}
DIST = ["IUK", "UTI", "UMI", "Germany", "Benelux", "Portugal", "Sweden", "Norway", "Ireland"]
FY26 = ["2025-04", "2025-05", "2025-06", "2025-07", "2025-08", "2025-09",
        "2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03"]

print("=" * 78)
print("A. PLAN COUNTS  (items 1 + 2)")
print("=" * 78)
p = PLAN / "budget_per_DB_reviewed_24042026.xlsx"
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
ws = wb["Incentive FY27"]
print("header row 1/2:")
for r in (1, 2):
    print("  ", r, [ws.cell(r, c).value for c in range(1, 10)])
name_map = {"ITUK": "ITUK", "IUK": "IUK", "UMI": "UMI", "UTI": "UTI",
            "IBX": "Benelux", "ISD": "Germany", "Portugal": "Portugal",
            "Ireland": "Ireland", "Norway": "Norway", "Sweden": "Sweden", "KTK": "KTK"}
inc = {}
for r in range(3, 22):
    label = ws.cell(r, 1).value
    if not label:
        continue
    key = str(label).strip()
    print(f"  row {r}: {key!r}  bp26={ws.cell(r,2).value}  bp27={ws.cell(r,4).value}  inc%={ws.cell(r,5).value}  incEUR={ws.cell(r,6).value}  incentive={ws.cell(r,8).value}")
    if key in name_map:
        inc[name_map[key]] = {
            "bp26": ws.cell(r, 2).value, "bp27": ws.cell(r, 4).value,
            "inc_pct": ws.cell(r, 5).value, "inc_eur": ws.cell(r, 6).value,
            "incentive": ws.cell(r, 8).value,
        }
wb.close()

sf = pd.read_parquet(DER / "sales_fact.parquet",
                     columns=["billing_date", "sold_to_sap", "net_sales_value"])
sf["billing_date"] = pd.to_datetime(sf["billing_date"], errors="coerce")
sf = sf.dropna(subset=["billing_date"])
sf["ym"] = sf["billing_date"].dt.to_period("M").astype(str)
sf["eur"] = pd.to_numeric(sf["net_sales_value"], errors="coerce").fillna(0.0)
sf["code"] = pd.to_numeric(sf["sold_to_sap"], errors="coerce").astype("Int64")

bill = {}
for name, codes in CODES.items():
    bill[name] = sf[sf["code"].isin(codes)].groupby("ym")["eur"].sum().reindex(FY26, fill_value=0.0)

print()
print(f"{'customer':10} {'bp26':>14} {'billed26':>14} {'pct':>8} {'bp27':>14} {'plan inc%':>10}")
above = 0
gt12 = []
for c in DIST:
    b = bill[c].sum()
    bp26 = inc[c]["bp26"]
    pct = (b / bp26 - 1) * 100
    if pct >= 0:
        above += 1
    if pct >= 12:
        gt12.append((c, pct))
    print(f"{c:10} {bp26:14,.0f} {b:14,.0f} {pct:+7.2f}% {inc[c]['bp27']:14,.0f} {inc[c]['inc_pct']*100:9.2f}%")
print()
print(f"distributors at/above plan : {above} of {len(DIST)}")
print(f"at/above +12%             : {len(gt12)} -> {[(c, round(p,2)) for c, p in gt12]}")

print()
print("=" * 78)
print("B. ITUK FEBRUARY  (item 3)")
print("=" * 78)
m = sf[sf["code"] == 2575].groupby("ym")["eur"].sum()
run = m.reindex(FY26[:10], fill_value=0.0).mean()
feb = m.get("2026-02", 0.0)
mar = m.get("2026-03", 0.0)
print("monthly ITUK billing:")
for k in FY26:
    print(f"  {k}  {m.get(k, 0.0):14,.2f}")
print(f"run rate Apr25-Jan26 (10 months) : {run:,.2f}")
print(f"February 2026 billed              : {feb:,.2f}")
print(f"February excess over run rate     : {feb - run:,.2f}")
print(f"March 2026 billed                 : {mar:,.2f}")
print(f"March excess over run rate        : {mar - run:,.2f}")
print()
# the wrong file's row
split = pd.read_csv(ROOT / "01_data" / "90_INBOX" / "_order_insights_20261007" /
                    "22_price_recon_attribution_split.csv")
ituk = split[split["cust"].str.contains("ITUK")].iloc[0]
parts = ituk["inc_Feb"] + ituk["inc_Mar1_12"] + ituk["inc_Mar13_31"]
print("wrong file, ITUK row:")
print(f"  inc_Feb        {ituk['inc_Feb']:,.2f}")
print(f"  inc_Mar1_12    {ituk['inc_Mar1_12']:,.2f}")
print(f"  inc_Mar13_31   {ituk['inc_Mar13_31']:,.2f}")
print(f"  sum of three   {parts:,.2f}")
print(f"  stated inc_total_FebMar {ituk['inc_total_FebMar']:,.2f}")
print(f"  does it foot?  {abs(parts - ituk['inc_total_FebMar']) < 1}")

print()
print("=" * 78)
print("C. PRICE 13 MARCH  (item 10)")
print("=" * 78)
meta = json.loads((DER / "price_bp_weighted_compare_meta.json").read_text(encoding="utf-8"))
print("weighted increase (billing qty):", meta["sensitivity"]["weight_base"]["weighted_increase_pct_billing_qty"])
print("value step eur (revised basis) :", meta["sensitivity"]["prior_vintage"]["value_step_eur_revised_basis"])
print("revision effective             :", meta["sensitivity"]["prior_vintage"]["revision_effective"])
print("window                         :", meta["window"])

print()
print("=" * 78)
print("D. KICKBACK TOTALS  (item 10 'sound')")
print("=" * 78)
tot26 = sum(v["bp26"] for v in inc.values())
tot27 = sum(v["bp27"] for v in inc.values())
tot_inc = sum(v["incentive"] or 0 for v in inc.values())
print(f"customers with plan rows: {len(inc)}")
print(f"plan FY26 {tot26:,.2f} -> FY27 {tot27:,.2f}  (+{(tot27/tot26-1)*100:.4f}%)")
print(f"incentive FY2027 total: {tot_inc:,.2f}")
for c, v in inc.items():
    if v["incentive"] and v["bp27"]:
        print(f"   {c:10} incentive {v['incentive']:12,.2f}  = {(v['incentive']/v['bp27'])*100:5.2f}% of plan")
