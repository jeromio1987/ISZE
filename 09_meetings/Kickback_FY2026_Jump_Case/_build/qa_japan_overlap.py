# -*- coding: utf-8 -*-
"""Computational overlap check between text boxes on each slide."""
import sys
from pptx import Presentation

DECK = r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case\ISZE_Kickback_Japan_Pack_20261007.pptx"
p = Presentation(DECK)
fails = []
for i, s in enumerate(p.slides, 1):
    tb = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.width and sh.height:
            tb.append((sh.name, sh.left / 914400, sh.top / 914400,
                       (sh.left + sh.width) / 914400, (sh.top + sh.height) / 914400,
                       sh.text_frame.text[:42].replace("\n", " ")))
    for a in range(len(tb)):
        for b in range(a + 1, len(tb)):
            n1, l1, t1, r1, b1, x1 = tb[a]
            n2, l2, t2, r2, b2, x2 = tb[b]
            ox = min(r1, r2) - max(l1, l2)
            oy = min(b1, b2) - max(t1, t2)
            if ox > 0.02 and oy > 0.02:
                # allow a small tolerance (boxes are deliberately generous)
                if ox > 0.05 and oy > 0.05:
                    fails.append("slide %d: %s/%s overlap %.2f x %.2f in  %r | %r" % (
                        i, n1, n2, ox, oy, x1, x2))
print("overlap fails:", len(fails))
for f in fails:
    print("  -", f)
sys.exit(1 if fails else 0)
