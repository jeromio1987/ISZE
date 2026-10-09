# -*- coding: utf-8 -*-
"""Pass 3c: remove the overlapping arrowheads from the slide 20 cycle."""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
DECK = HERE / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pptx"

INK = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xBD, 0xBD, 0xBD)
FONT = "Aptos"

prs = Presentation(str(DECK))
s20 = list(prs.slides)[19]

# the arrows are the only small autoshapes left on the slide
for sh in list(s20.shapes):
    if sh.shape_type is not None and str(sh.shape_type).startswith("AUTO_SHAPE"):
        l = (sh.left or 0) / 914400.0
        t = (sh.top or 0) / 914400.0
        w = (sh.width or 0) / 914400.0
        h = (sh.height or 0) / 914400.0
        if w < 0.45 and h < 0.35:
            sh._element.getparent().remove(sh._element)

mid = None
for sh in s20.shapes:
    if sh.has_text_frame and "A cycle, not a" in sh.text_frame.text:
        mid = sh
        break

if mid is not None:
    tf = mid.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.clear()
    for i, (txt, size, col, bold, sp) in enumerate([
        ("A cycle, not a\nsingle approval", 10.5, INK, True, 0),
        ("Four steps, clockwise from the top. The plan is reset from the agreed "
         "monthly targets after every checkpoint.", 7.0, MUTED, False, 3),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        if sp:
            p.space_before = Pt(sp)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = col
        r.font.bold = bold
        r.font.name = FONT
    print("centre caption updated")

prs.save(str(DECK))
print("saved")
