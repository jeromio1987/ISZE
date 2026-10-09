# -*- coding: utf-8 -*-
"""Probe 2: ARC angle control, rotation, arrowheads, DONUT band width."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(OUT_DIR, "_ring_probe2.pptx")

BG = RGBColor(0x26, 0x26, 0x26)
RED = RGBColor(0xE3, 0x06, 0x13)
WHITE = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xBD, 0xBD, 0xBD)
GREEN = RGBColor(0x31, 0xA3, 0x54)
AMBER = RGBColor(0xF5, 0xA0, 0x00)
PANEL = RGBColor(0x30, 0x30, 0x30)
FONT = "Aptos"

prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
s = prs.slides.add_slide(blank)
f = s.background.fill
f.solid()
f.fore_color.rgb = BG


def label(x, y, w, h, t, size=8.5, color=MUTED):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = False
    r = tf.paragraphs[0].add_run()
    r.text = t
    r.font.size = Pt(size)
    r.font.name = FONT
    r.font.color.rgb = color
    return tb


def arc(x, y, d, a1, a2, color=RED, lw=10, rot=0, tag=""):
    sh = s.shapes.add_shape(MSO_SHAPE.ARC, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.shadow.inherit = False
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(lw)
    try:
        sh.adjustments[0] = a1
        sh.adjustments[1] = a2
    except Exception as e:
        print("adj fail", tag, e)
    if rot:
        sh.rotation = rot
    vals = []
    for i in range(2):
        try:
            vals.append(sh.adjustments[i])
        except Exception:
            vals.append(None)
    print("arc %-14s req=(%s,%s) got=%s" % (tag, a1, a2, vals))
    return sh


def tri(x, y, w, h, color=AMBER, rot=0):
    sh = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    if rot:
        sh.rotation = rot
    return sh


def donut(x, y, d, hole, color=RED):
    sh = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.adjustments[0] = hole
    print("donut hole req", hole, "got", sh.adjustments[0])
    return sh


# A: does an ARC respond to angles?  270->360, 0->90 etc, rotated by hand
arc(0.6, 0.7, 2.4, 0, 90, RED, 10, 0, "0-90")
arc(3.4, 0.7, 2.4, 90, 180, RED, 10, 0, "90-180")
arc(6.2, 0.7, 2.4, 180, 270, RED, 10, 0, "180-270")
arc(9.0, 0.7, 2.4, 270, 360, RED, 10, 0, "270-360")
label(0.6, 3.2, 4, 0.3, "A: ARC with explicit angle pairs (no rotation)")

# B: default ARC, rotated 90/180/270 - does rotation move it?
arc(0.6, 3.6, 2.4, 270, 360, RED, 10, 90, "def+90")
arc(3.4, 3.6, 2.4, 270, 360, RED, 10, 180, "def+180")
arc(6.2, 3.6, 2.4, 270, 360, RED, 10, 270, "def+270")
label(0.6, 6.1, 6, 0.3, "B: same arc rotated 90 / 180 / 270 (does rotation apply?)")

# C: arrowheads by rotation
rots = [0, 45, 90, 135, 180, 225, 270, 315]
for i, r in enumerate(rots):
    tri(9.4 + (i % 4) * 0.85, 3.6 + (i // 4) * 0.95, 0.50, 0.40, AMBER, r)
label(9.4, 5.5, 4, 0.3, "C: ISOSCELES_TRIANGLE rotated 0..315 (a triangle points UP at 0)")

# D: donut hole widths
donut(10.2, 0.7, 2.4, 0.28, RED)
label(10.2, 3.2, 3, 0.3, "D: DONUT hole=0.28")

prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT))
