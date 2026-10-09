"""Static QA on the built deck: shapes outside the slide, and estimated text overflow."""
from __future__ import annotations
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
PPTX = HERE / "ISZE_Kickback_FY2026_Year_End_Jump_Case.pptx"

prs = Presentation(PPTX)
SW, SH = prs.slide_width, prs.slide_height
EMU_IN = 914400
issues = []
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if l is None:
            continue
        # bounds
        if l < -10000 or t < -10000 or l + w > SW + 10000 or t + h > SH + 10000:
            issues.append(f"S{si:02d} OUTSIDE  {sh.shape_type} l={l/EMU_IN:.2f} t={t/EMU_IN:.2f} "
                          f"r={(l+w)/EMU_IN:.2f} b={(t+h)/EMU_IN:.2f}")
        # text overflow estimate
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            box_w_in = (w or 0) / EMU_IN
            box_h_in = (h or 0) / EMU_IN
            total = 0.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    total += 0.12
                    continue
                size = None
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                size = size or 10
                # chars per line at this size, ~0.52*size pt average char width
                line_h = size * 1.22 / 72
                cpl = max(1, int(box_w_in * 72 / (size * 0.52)))
                lines = max(1, -(-len(txt) // cpl))
                total += lines * line_h
            if total > box_h_in * 1.12 and box_h_in > 0.05:
                issues.append(f"S{si:02d} OVERFLOW est {total:.2f}in > box {box_h_in:.2f}in  "
                              f"'{tf.text[:58]}...'")
print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}  issues: {len(issues)}")
for i in issues:
    print(" ", i)
