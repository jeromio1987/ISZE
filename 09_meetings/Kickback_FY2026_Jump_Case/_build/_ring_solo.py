# -*- coding: utf-8 -*-
"""Standalone render of the ring: step boxes OUTSIDE the ring radius, so the ring
and its accent arcs are never covered.  Used to verify before it goes in the deck."""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _ring import add_ring, step_radius, free_windows  # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_ring_solo.pptx")

BG = RGBColor(0x26, 0x26, 0x26)
PANEL = RGBColor(0x30, 0x30, 0x30)
RED = RGBColor(0xE3, 0x06, 0x13)
WHITE = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xBD, 0xBD, 0xBD)
AMBER = RGBColor(0xF5, 0xA0, 0x00)
GRID = RGBColor(0x4B, 0x4B, 0x4B)
FONT = "Aptos"

prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
f = s.background.fill
f.solid()
f.fore_color.rgb = BG

CX, CY, R = 6.6667, 4.15, 1.40
BW, BH = 2.25, 0.80
SR = step_radius(CX, CY, BW, BH)
print("ring radius %.2f in, step-box centre radius %.2f in" % (R, SR))
print("free ring windows (deg):", free_windows(R, BW, BH, SR))

add_ring(s, CX, CY, radius=R)


def box(cx, cy, title, body, accent=RED):
    x, y = cx - BW / 2.0, cy - BH / 2.0
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(BW), Inches(BH))
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.10
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = GRID
    sh.line.width = Pt(0.75)
    sh.text_frame.text = ""
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.13), Inches(y + 0.085),
                           Inches(0.05), Inches(0.028))
    a.shadow.inherit = False
    a.fill.solid()
    a.fill.fore_color.rgb = accent
    a.line.fill.background()

    tb = s.shapes.add_textbox(Inches(x + 0.11), Inches(y + 0.10), Inches(BW - 0.22),
                              Inches(BH - 0.20))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(7.8)
    r.font.name = FONT
    r.font.bold = True
    r.font.color.rgb = accent
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = body
    r2.font.size = Pt(9.6)
    r2.font.name = FONT
    r2.font.color.rgb = WHITE


box(CX, CY - SR, "1  ·  ISZE / JAPAN", "Confirm the monthly plan for each distributor")
box(CX + SR, CY, "2  ·  ISZE / JAPAN", "Check at month 6 against the 10% tolerance")
box(CX, CY + SR, "3  ·  ISZE / JAPAN", "Check at month 9 against the 10% tolerance")
box(CX - SR, CY, "4  ·  ISZE", "Close the five reporting gaps and load ISD Feb-Jul 2026",
    accent=AMBER)

# centre text, added after the ring so it sits above it
tb = s.shapes.add_textbox(Inches(CX - 0.72), Inches(CY - 0.42), Inches(1.44), Inches(0.84))
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
p = tf.paragraphs[0]
p.alignment = 2
r = p.add_run()
r.text = "A cycle, not a\nsingle approval"
r.font.size = Pt(11.5)
r.font.name = FONT
r.font.bold = True
r.font.color.rgb = WHITE
p2 = tf.add_paragraph()
p2.alignment = 2
r2 = p2.add_run()
r2.text = "clockwise from the top"
r2.font.size = Pt(7.8)
r2.font.name = FONT
r2.font.color.rgb = MUTED

prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT))
