"""Static QA on the FIGURES AUDIT deck: shapes outside the slide, estimated text overflow,
and a banned-word scan. Points at the new audit file only."""
from __future__ import annotations
import io
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
PPTX = HERE / "ISZE_Kickback_Figures_Audit_20261007.pptx"

prs = Presentation(PPTX)
SW, SH = prs.slide_width, prs.slide_height
EMU_IN = 914400
issues = []
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if l is None:
            continue
        if l < -10000 or t < -10000 or l + w > SW + 10000 or t + h > SH + 10000:
            issues.append(f"S{si:02d} OUTSIDE  {sh.shape_type} l={l/EMU_IN:.2f} t={t/EMU_IN:.2f} "
                          f"r={(l+w)/EMU_IN:.2f} b={(t+h)/EMU_IN:.2f}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            box_w_in = (w or 0) / EMU_IN
            box_h_in = (h or 0) / EMU_IN
            total = 0.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    total += 0.12
                    continue
                size = None
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                size = size or 10
                line_h = size * 1.22 / 72
                cpl = max(1, int(box_w_in * 72 / (size * 0.52)))
                lines = max(1, -(-len(txt) // cpl))
                total += lines * line_h
            if total > box_h_in * 1.12 and box_h_in > 0.05:
                issues.append(f"S{si:02d} OVERFLOW est {total:.2f}in > box {box_h_in:.2f}in  "
                              f"'{tf.text[:60]}...'")

BANNED = ["sell-out", "sellout", "sell out", "mode a", "mode b", "street index", "wallet"]
ban_hits = 0
for si, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            texts.append(("text", sh.text_frame.text))
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    texts.append(("cell", c.text))
    for kind, t in texts:
        low = t.lower()
        for b in BANNED:
            if b in low:
                ban_hits += 1
                print(f"  BAN slide {si} [{kind}] {b!r} -> {t[:90]}")

print(f"slides: {len(prs.slides._sldIdLst)}  layout issues: {len(issues)}  banned-word hits: {ban_hits}")
for i in issues:
    print(" ", i)
print("STATIC:", "CLEAN" if (not issues and ban_hits == 0) else "REVIEW")
