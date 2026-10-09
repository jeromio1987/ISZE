# -*- coding: utf-8 -*-
"""Build the ring on slide 20 in numbered steps so each can be open-tested."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Kickback_FY2026_Jump_Case")
SRC = HERE / "_build" / "_backup" / "ISZE_Kickback_Management_Revised_v2_2026-10-07.pre-dynamics-20261007-1512.pptx"
STEP = int(sys.argv[1])
OUT = HERE / "_build" / ("_bisect_%d.pptx" % STEP)

RED = RGBColor(0xE3, 0x06, 0x13)
INK = RGBColor(0xF2, 0xF2, 0xF2)
PANEL = RGBColor(0x30, 0x30, 0x30)

prs = Presentation(str(SRC))
slide = prs.slides[19]

if STEP >= 1:                      # clear the old flow
    for sh in list(slide.shapes):
        if sh.name in ("Oval 8", "TextBox 9"):
            sh._element.getparent().remove(sh._element)
        elif sh.name.startswith("Rounded Rectangle") or sh.name.startswith("Rectangle "):
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and sh.text_frame.text.strip().startswith(("1  ·", "2  ·", "3  ·", "4  ·")):
            sh._element.getparent().remove(sh._element)

if STEP >= 2:                      # plain oval track, no adjustments
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.38), Inches(2.44), Inches(4.56), Inches(4.56))
    o.fill.background(); o.line.color.rgb = RED; o.line.width = Pt(8)

if STEP >= 3:                      # rounded rect step box, no adjustments
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.30), Inches(2.60), Inches(2.70), Inches(1.02))
    r.fill.solid(); r.fill.fore_color.rgb = PANEL; r.line.color.rgb = RED; r.line.width = Pt(1.5)

if STEP >= 4:                      # step box WITH adjustment 0
    r2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.30), Inches(2.60), Inches(2.70), Inches(1.02))
    r2.adjustments[0] = 0.10
    r2.fill.solid(); r2.fill.fore_color.rgb = PANEL; r2.line.color.rgb = RED; r2.line.width = Pt(1.5)

if STEP >= 5:                      # chevron with rotation
    c = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(9.0), Inches(5.4), Inches(0.42), Inches(0.34))
    c.fill.solid(); c.fill.fore_color.rgb = RED; c.line.fill.background()
    c.rotation = 135

if STEP >= 6:                      # text written into a step box
    tb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.30), Inches(5.60), Inches(2.70), Inches(1.02))
    tb.adjustments[0] = 0.10
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = "HELLO"; run.font.size = Pt(8.6); run.font.name = "Aptos"
    run.font.color.rgb = INK

if STEP >= 7:                      # shadow inherit
    s7 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0))
    s7.shadow.inherit = False

if STEP >= 8:                      # textbox with auto_size None
    t8 = slide.shapes.add_textbox(Inches(5.30), Inches(6.80), Inches(2.70), Inches(0.40))
    t8.text_frame.auto_size = None
    p = t8.text_frame.paragraphs[0]
    run = p.add_run(); run.text = "CAPTION"

prs.save(str(OUT))
print("wrote", OUT.name, "step", STEP)
