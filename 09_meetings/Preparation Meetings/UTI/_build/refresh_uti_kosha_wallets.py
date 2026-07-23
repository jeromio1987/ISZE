# -*- coding: utf-8 -*-
"""Refresh max-spend wallet cells on GPS · UTI from UIO v19 (uio_2026.json)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
UIO_JSON = ROOT / "01_data" / "91_build" / "uio_2026.json"
PLAN_XLSX = (
    ROOT
    / "09_meetings"
    / "UTI_Presentation"
    / "All_Material"
    / "GPS_PMS_Rate_Models"
    / "GPS_PMS_UTI_Kosha_Model.xlsx"
)
SHEET = "GPS · UTI"

RATE_LD = 739.0
RATE_MD = 887.0

# Row 6 — max spend € columns per fiscal block
MAX_SPEND_CELLS = {2024: "F6", 2025: "J6", 2026: "N6", 2027: "R6"}


def uti_max_spend_by_year() -> dict[int, int]:
    data = json.loads(UIO_JSON.read_text(encoding="utf-8"))
    buckets: dict[int, dict[str, float]] = {}
    for rec in data:
        if rec.get("dist") != "UTI":
            continue
        mt = str(rec.get("mtype", "")).upper()
        for y in (2024, 2025, 2026, 2027):
            raw = rec.get(str(y))
            if raw is None:
                continue
            b = buckets.setdefault(y, {"LD": 0.0, "MD": 0.0})
            if mt == "LD":
                b["LD"] += float(raw)
            elif mt in ("MD", "F-SERIES", "F SERIES", "F"):
                b["MD"] += float(raw)
    out: dict[int, int] = {}
    for y, b in sorted(buckets.items()):
        out[y] = int(round(b["LD"] * RATE_LD + b["MD"] * RATE_MD))
    return out


def main() -> None:
    wallets = uti_max_spend_by_year()
    wb = openpyxl.load_workbook(PLAN_XLSX)
    ws = wb[SHEET]
    for fy, addr in MAX_SPEND_CELLS.items():
        if fy not in wallets:
            continue
        old = ws[addr].value
        ws[addr].value = wallets[fy]
        print(f"{addr} FY{fy}: {old:,} -> {wallets[fy]:,}")
    wb.save(PLAN_XLSX)
    print("saved", PLAN_XLSX)


if __name__ == "__main__":
    main()
