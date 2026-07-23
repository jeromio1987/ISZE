# -*- coding: utf-8 -*-
"""Single source for UTI GPS evolution + stretch slides (UIO v19)."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BILLING_JSON = ROOT / "01_data" / "00_derived" / "gps_billing_clean.json"
WARRANTY_CSV = ROOT / "01_data" / "00_derived" / "distributor_warranty_expense_fiscal.csv"
UIO_JSON = ROOT / "01_data" / "91_build" / "uio_2026.json"
PLAN_XLSX = (
    ROOT
    / "09_meetings"
    / "UTI_Presentation"
    / "All_Material"
    / "GPS_PMS_Rate_Models"
    / "GPS_PMS_UTI_Kosha_Model.xlsx"
)

RATE_LD = 739.0
RATE_MD = 887.0
STRETCH_PP = 1.8  # percentage points on plan GPS (UTI kickback policy)


def _uti_uio_by_year() -> dict[int, int]:
    data = json.loads(UIO_JSON.read_text(encoding="utf-8"))
    years: dict[int, float] = {}
    for rec in data:
        if rec.get("dist") != "UTI":
            continue
        for y in (2024, 2025, 2026, 2027):
            if str(y) in rec:
                years[y] = years.get(y, 0) + float(rec[str(y)])
    return {y: int(round(v)) for y, v in sorted(years.items())}


def _max_spend_fy27() -> float:
  buckets: dict[str, float] = {"LD": 0.0, "MD": 0.0}
  for rec in json.loads(UIO_JSON.read_text(encoding="utf-8")):
      if rec.get("dist") != "UTI":
          continue
      mt = str(rec.get("mtype", "")).upper()
      v = float(rec.get("2027", 0) or 0)
      if mt == "LD":
          buckets["LD"] += v
      elif mt in ("MD", "F-SERIES", "F SERIES", "F"):
          buckets["MD"] += v
  return buckets["LD"] * RATE_LD + buckets["MD"] * RATE_MD


def load_evolution_rows() -> list[dict]:
    bill = {int(k): float(v) for k, v in json.loads(BILLING_JSON.read_text(encoding="utf-8"))["UTI"].items()}
    warr: dict[int, float] = {}
    for line in WARRANTY_CSV.read_text(encoding="utf-8").splitlines()[1:]:
        parts = line.split(",")
        if parts[0] == "UTI":
            for fy, col in ((2024, 1), (2025, 2), (2026, 3)):
                warr[fy] = float(parts[col])

    wb = openpyxl.load_workbook(PLAN_XLSX, data_only=True)
    ws = wb["GPS · UTI"]
    max_spend = {
        2024: float(ws.cell(6, 6).value),
        2025: float(ws.cell(6, 10).value),
        2026: float(ws.cell(6, 14).value),
    }
    plan_bill = float(ws.cell(6, 17).value)
    plan_max = float(ws.cell(6, 18).value)
    wb.close()

    uio = _uti_uio_by_year()
    labels = ["FY24", "FY25", "FY26", "FY27 plan"]
    rows = []
    for i, (fy, is_plan) in enumerate(((2024, False), (2025, False), (2026, False), (2027, True))):
        if is_plan:
            rows.append(
                {
                    "fy": f"FY{fy}",
                    "xlabel": labels[i],
                    "is_plan": True,
                    "uio": uio.get(fy, 0),
                    "billing_eur": plan_bill,
                    "warranty_eur": 0.0,
                    "max_spend_eur": plan_max,
                    "gps_net_pct": plan_bill / plan_max * 100 if plan_max else 0.0,
                }
            )
            continue
        b, w, ms = bill[fy], warr.get(fy, 0.0), max_spend.get(fy, 0.0)
        rows.append(
            {
                "fy": f"FY{fy}",
                "xlabel": labels[i],
                "is_plan": False,
                "uio": uio.get(fy, 0),
                "billing_eur": b,
                "warranty_eur": w,
                "max_spend_eur": ms,
                "gps_net_pct": (b - w) / ms * 100 if ms else 0.0,
            }
        )
    return rows


def load_stretch_bundle() -> dict:
    rows = load_evolution_rows()
    r26 = rows[2]
    kosha_plan_gps = rows[3]["gps_net_pct"]
    max27 = _max_spend_fy27()
    bp_base = rows[3]["billing_eur"]
    gps_plan = 36.5  # commercial FY27 plan target (negotiated — differs from Kosha BP ÷ v19 wallet)
    gps_stretch = round(gps_plan + STRETCH_PP, 1)
    gps_stretch_eur = round(STRETCH_PP / 100 * max27)
    stretch_bp = round(bp_base + gps_stretch_eur)
    return {
        "gps_fy26": round(r26["gps_net_pct"], 1),
        "kosha_plan_gps": round(kosha_plan_gps, 1),
        "gps_plan": gps_plan,
        "gps_stretch": gps_stretch,
        "pp": STRETCH_PP,
        "bp_base": int(bp_base),
        "gps_stretch_eur": gps_stretch_eur,
        "stretch_bp": stretch_bp,
        "delta_pct": round((stretch_bp - bp_base) / bp_base * 100, 1) if bp_base else 0.0,
        "max_spend": int(round(max27)),
        "incentive_eur": 25_201,
        "regular_kickback_eur": 36_000,
    }
