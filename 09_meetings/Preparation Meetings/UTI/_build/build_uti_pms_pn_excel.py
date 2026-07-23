"""UTI PMS score workbook — category/subcategory evolution + PN-level detail, by FY.

Reuses uti_iml_fy_engine (IML v7 interval methodology, N75_only scope, 100% addressable)
for category-level capture. Adds a PN-level layer (qty/revenue only — expected demand
is only modelled at subcategory level, so no per-PN capture rate exists).
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from uti_iml_fy_engine import (  # noqa: E402
    ALL_CATS,
    BRAKE_CATS,
    CLUTCH_CATS,
    FILTER_CATS,
    compute_uti_fy,
    load_sales,
)
from modules.pms_engine import CATEGORY_LABELS, filter_pms_by_scope, pms_scope_for_distributor  # noqa: E402

OUT = Path(__file__).resolve().parent.parent / "UTI_PMS_Score_by_PN.xlsx"

# prism's native `fiscal_year` column already buckets Apr-Mar correctly (e.g. native FY2025
# = Apr 2025-Mar 2026), but is labeled one year behind Jerome's convention, where that same
# Apr 2025-Mar 2026 window is "FY26". NATIVE_FY_YEARS drives all data pulls; FY_LABEL converts
# to Jerome's convention for every displayed column.
NATIVE_FY_YEARS = [2023, 2024, 2025]
FY_LABEL = {fy: fy + 1 for fy in NATIVE_FY_YEARS}

FAMILY_OF = {}
for c in FILTER_CATS:
    FAMILY_OF[c] = "Filters"
for c in BRAKE_CATS:
    FAMILY_OF[c] = "Brake"
for c in CLUTCH_CATS:
    FAMILY_OF[c] = "Clutch"

NAVY = "1F3864"
TERRACOTTA = "C0653A"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAMILY_FILL = PatternFill("solid", fgColor="D9E2F3")


def style_header(ws, ncols, row=1):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def autosize(ws, df, extra=2):
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), df[col].astype(str).map(len).max() if len(df) else 0) + extra
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)


def build_category_frames(cat_by_fy: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (total-only, genuine/bvp split) capture-rate frames, subcategory rows x FY cols."""
    rows_total, rows_split = [], []
    for cat in ALL_CATS:
        base = {"Family": FAMILY_OF[cat], "Subcategory": CATEGORY_LABELS.get(cat, cat)}
        row_t = dict(base)
        for fy in NATIVE_FY_YEARS:
            lbl = FY_LABEL[fy]
            c = cat_by_fy[fy]["categories"].get(cat)
            row_t[f"FY{lbl} capture %"] = round(c["total"]["capture"] * 100, 1) if c else None
            row_t[f"FY{lbl} qty"] = round(c["total"]["qty"]) if c else None
            row_t[f"FY{lbl} revenue EUR"] = round(c["total"]["rev"]) if c else None
            row_t[f"FY{lbl} gap EUR"] = round(c["total"]["gap"]) if c else None
        rows_total.append(row_t)

        for lt, label in (("genuine", "Genuine"), ("bvp", "BVP")):
            row_s = dict(base, Type=label)
            for fy in NATIVE_FY_YEARS:
                lbl = FY_LABEL[fy]
                c = cat_by_fy[fy]["categories"].get(cat)
                row_s[f"FY{lbl} capture %"] = round(c[lt]["capture"] * 100, 1) if c else None
                row_s[f"FY{lbl} qty"] = round(c[lt]["qty"]) if c else None
                row_s[f"FY{lbl} revenue EUR"] = round(c[lt]["rev"]) if c else None
            rows_split.append(row_s)

    return pd.DataFrame(rows_total), pd.DataFrame(rows_split)


def build_pn_frame(sales: pd.DataFrame) -> pd.DataFrame:
    scope = pms_scope_for_distributor("UTI")
    parts = []
    for fy in NATIVE_FY_YEARS:
        fy_sales = sales[(sales["distributor_code"] == "UTI") & (sales["fiscal_year"] == fy)]
        sub = filter_pms_by_scope(fy_sales, scope)
        g = (
            sub.groupby(["pms_category", "part_number", "description", "pms_type"], dropna=False)
            .agg(qty=("billed_qty", "sum"), revenue_eur=("net_sales_value", "sum"))
            .reset_index()
        )
        g["fiscal_year"] = FY_LABEL[fy]
        parts.append(g)
    pn = pd.concat(parts, ignore_index=True)
    pn["Family"] = pn["pms_category"].map(FAMILY_OF)
    pn["Subcategory"] = pn["pms_category"].map(lambda c: CATEGORY_LABELS.get(c, c))
    pn["Type"] = pn["pms_type"].map({"genuine": "Genuine", "bvp": "BVP"}).fillna(pn["pms_type"])

    qty_wide = pn.pivot_table(
        index=["Family", "Subcategory", "part_number", "description", "Type"],
        columns="fiscal_year", values="qty", aggfunc="sum", fill_value=0,
    )
    qty_wide.columns = [f"FY{c} qty" for c in qty_wide.columns]

    rev_wide = pn.pivot_table(
        index=["Family", "Subcategory", "part_number", "description", "Type"],
        columns="fiscal_year", values="revenue_eur", aggfunc="sum", fill_value=0,
    )
    rev_wide.columns = [f"FY{c} revenue EUR" for c in rev_wide.columns]

    wide = qty_wide.join(rev_wide).reset_index()
    wide = wide.rename(columns={"part_number": "Part Number", "description": "Description"})

    qty_cols = [c for c in wide.columns if c.endswith(" qty")]
    wide["Total qty (all FY)"] = wide[qty_cols].sum(axis=1)
    wide = wide.sort_values(
        ["Family", "Subcategory", "Total qty (all FY)"], ascending=[True, True, False]
    ).drop(columns="Total qty (all FY)")

    for c in qty_wide.columns.tolist():
        wide[c] = wide[c].round().astype(int)
    for c in rev_wide.columns.tolist():
        wide[c] = wide[c].round().astype(int)

    return wide


def write_sheet(writer, name, df):
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    style_header(ws, len(df.columns))
    autosize(ws, df)


def main():
    print("Loading tagged UTI sales...")
    sales = load_sales()
    cat_by_fy = {fy: compute_uti_fy(fy, sales) for fy in NATIVE_FY_YEARS}

    cat_total, cat_split = build_category_frames(cat_by_fy)
    pn_detail = build_pn_frame(sales)

    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        readme = pd.DataFrame({
            "UTI PMS Score — Category, Subcategory & PN-level Evolution": [
                "Distributor: UTI (Israel, MENA) | Scope: N75_only, 100% addressable",
                "Demand model: IML v7 service intervals (filters 1.25x/yr, brake+clutch wear 0.5x/yr, 50,000 km/yr) — NOT the N75 0.5 svc/yr model used elsewhere in ISZE PMS briefs.",
                "Time base: ISZE fiscal year (Apr-Mar). FY2024=Apr 2023-Mar 2024, FY2025=Apr 2024-Mar 2025, FY2026=Apr 2025-Mar 2026 (ended 30 Mar 2026, latest available billing data).",
                "'Category - Total' = blended genuine+BVP capture rate per subcategory per FY.",
                "'Category - Genuine vs BVP' = same split by channel. UTI is ~98% BVP by volume — genuine-side capture is a thin sample and directionally indicative only, not precision-grade.",
                "'PN Detail' = part-number-level qty & revenue by FY. No per-PN capture rate: expected demand is only modelled at subcategory level (weighted UIO x IML frequency), not per PN, so a PN can't carry its own denominator.",
                "Known caveat: PMS classification (genuine/BVP tagging via part-catalog fig_key) has an open leakage issue under review across ISZE (assemblies vs elements, brake_rotor mapping) — category totals may shift once the part-master fix lands. Treat this file as current-best, not final.",
                "Source: sales_fact.parquet (SAP billing) tagged via prism pms_tagging/pms_engine, same pipeline as UTI_PMS_Evolution.html.",
            ]
        })
        readme.to_excel(writer, sheet_name="README", index=False, header=False)
        ws = writer.sheets["README"]
        ws.column_dimensions["A"].width = 130
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.font = Font(bold=True, size=13, color=NAVY)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[cell.row].height = 30

        write_sheet(writer, "Category - Total", cat_total)
        write_sheet(writer, "Category - Genuine vs BVP", cat_split)
        write_sheet(writer, "PN Detail", pn_detail)

    print(f"Wrote {OUT}")
    print(f"PN Detail rows: {len(pn_detail)}")


if __name__ == "__main__":
    main()
