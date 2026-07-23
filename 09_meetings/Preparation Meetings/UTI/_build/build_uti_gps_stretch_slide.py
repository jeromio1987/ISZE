"""
UTI GPS stretch proposal — meeting slide (Isuzu template v2, UTI-facing).

No internal ISZE references. Numbers from stretch workbook (2026-07-03).
"""
from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.dirname(HERE)
TEMPLATE_DIR = r"C:\Users\jerom\Desktop\claude\ISZE\04_tools\ppt_template_v2"
sys.path.insert(0, HERE)
sys.path.insert(0, TEMPLATE_DIR)

from uti_gps_figures import load_stretch_bundle  # noqa: E402

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

import isze_chart_theme as ct
import isze_deck_builder as db
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ct.FONT_FAMILY = "Calibri"
plt.rcParams["font.family"] = ct.FONT_FAMILY

CHART_DIR = os.path.join(HERE, "charts_out")
os.makedirs(CHART_DIR, exist_ok=True)

# ── figures (UIO v19 wallet — refresh via uti_gps_figures + refresh_uti_kosha_wallets) ──
UTI = load_stretch_bundle()

# Layout safe zone — logo ~6.85", footer ~7.40"
CONTENT_BOTTOM = Inches(6.42)

RED_PPTX = RGBColor(0xCC, 0x0E, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
LGRAY = RGBColor(0xF0, 0xF0, 0xF0)
MGRAY = RGBColor(0xD9, 0xD9, 0xD9)
NAVY = RGBColor(0x1A, 0x2B, 0x4A)


def cp(name: str) -> str:
    return os.path.join(CHART_DIR, name)


def _eur(v: float, *, k: bool = False) -> str:
    if k:
        return f"€{v / 1000:,.0f}K"
    if abs(v) >= 1_000_000:
        return f"€{v / 1_000_000:.2f}M"
    return f"€{v:,.0f}"


def _gps_ladder_chart(path: str) -> None:
    """Horizontal bars — FY26 charcoal, plan grey, stretch red."""
    labels = ["FY27 stretch", "FY27 plan", "FY26 actual"]
    values = [UTI["gps_stretch"], UTI["gps_plan"], UTI["gps_fy26"]]
    colors = [ct.RED, ct.GRAY_MED, ct.CHARCOAL]

    fig, ax = plt.subplots(figsize=(5.6, 2.6), dpi=200)
    fig.patch.set_facecolor(ct.WHITE)
    ax.set_facecolor(ct.WHITE)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color(ct.GRAY_LIGHT)
    ax.tick_params(axis="both", colors=ct.GRAY, labelsize=10, length=0)

    y = np.arange(len(labels))
    bars = ax.barh(y, values, height=0.55, color=colors, zorder=3)
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=11, color=ct.CHARCOAL)
    ax.set_xlim(0, 48)
    ax.set_xlabel("GPS % (parts billing ÷ max spend wallet)", fontsize=9.5, color=ct.GRAY)
    ax.xaxis.grid(True, color=ct.GRAY_BG, linewidth=1.2, zorder=0)
    ax.set_axisbelow(True)

    for bar, val in zip(bars, values):
        ax.annotate(
            f"{val:.1f}%",
            (bar.get_width(), bar.get_y() + bar.get_height() / 2),
            xytext=(6, 0),
            textcoords="offset points",
            ha="left",
            va="center",
            fontsize=11,
            fontweight="bold",
            color=ct.CHARCOAL,
        )

    ax.annotate(
        f"+{UTI['pp']:.1f}pp",
        (UTI["gps_stretch"], 0),
        xytext=(UTI["gps_stretch"] + 1.2, 0.35),
        fontsize=9,
        color=ct.RED,
        fontweight="bold",
        arrowprops=dict(arrowstyle="->", color=ct.RED, lw=1.2),
    )

    ax.set_title(
        "Your GPS target path, FY26 → FY27",
        fontsize=12,
        color=ct.CHARCOAL,
        loc="left",
        fontweight="bold",
        pad=8,
    )
    fig.tight_layout(pad=0.5)
    fig.savefig(path, dpi=200, facecolor=ct.WHITE)
    plt.close(fig)


def _bp_bridge_chart(path: str) -> None:
    """Clean waterfall — base, increment, stretch total."""
    labels = ["BP base\nFY27", f"+GPS stretch\n(+{UTI['pp']:.1f}pp)", "Stretch BP\nFY27"]
    values = [UTI["bp_base"], UTI["gps_stretch_eur"], UTI["stretch_bp"]]
    total_idx = {0, 2}

    fig, ax = plt.subplots(figsize=(5.6, 2.6), dpi=200)
    fig.patch.set_facecolor(ct.WHITE)
    ax.set_facecolor(ct.WHITE)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color(ct.GRAY_LIGHT)
    ax.tick_params(axis="y", colors=ct.GRAY, labelsize=9, length=0)
    ax.tick_params(axis="x", colors=ct.CHARCOAL, labelsize=10, length=0)
    ax.yaxis.grid(True, color=ct.GRAY_BG, linewidth=1.2, zorder=0)
    ax.set_axisbelow(True)

    running = 0
    for i, v in enumerate(values):
        if i in total_idx:
            bottom, height, color = 0, v, ct.GRAY_MED if i == 0 else ct.RED
            running = v
        else:
            bottom = running
            height = abs(v)
            color = ct.CHARCOAL
            running += v
        ax.bar(i, height, bottom=bottom, color=color, width=0.55, zorder=3)
        label_y = bottom + height
        lbl = _eur(v) if i in total_idx else f"+{_eur(v)}"
        ax.annotate(
            lbl,
            (i, label_y),
            xytext=(0, 5),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=10,
            fontweight="bold",
            color=ct.CHARCOAL,
        )
        if i < len(values) - 1:
            ax.plot([i + 0.28, i + 0.72], [running, running], color=ct.GRAY_LIGHT, lw=1)

    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels)
    ax.set_ylim(0, UTI["stretch_bp"] * 1.18)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"€{x/1e6:.1f}M"))
    ax.set_title(
        "Business plan bridge — how stretch is funded",
        fontsize=12,
        color=ct.CHARCOAL,
        loc="left",
        fontweight="bold",
        pad=8,
    )
    fig.tight_layout(pad=0.5)
    fig.savefig(path, dpi=200, facecolor=ct.WHITE)
    plt.close(fig)


# ── table helpers ─────────────────────────────────────────────────────────
def _style_cell(cell, text, *, size=10, bold=False, color=BLACK, fill=WHITE, align=PP_ALIGN.LEFT):
    cell.margin_top = Pt(2)
    cell.margin_bottom = Pt(2)
    cell.margin_left = Pt(6)
    cell.margin_right = Pt(6)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"


def _add_table(slide, rows, left, top, width, col_widths, *, row_h=0.34, hdr_size=9, body_size=9):
    nr, nc = len(rows), len(rows[0])
    tbl = slide.shapes.add_table(nr, nc, left, top, width, Inches(row_h * nr)).table
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = cw
    tbl.first_row = False
    tbl.horz_banding = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            align = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if i == 0:
                _style_cell(cell, val, size=hdr_size, bold=True, color=WHITE, fill=RED_PPTX, align=align)
            else:
                bold = str(val).startswith(("Stretch", "Total", "+"))
                fill = MGRAY if "Stretch" in str(row[0]) and j == 0 else (LGRAY if i % 2 == 0 else WHITE)
                if "Stretch" in str(row[0]):
                    fill = RGBColor(0xFD, 0xF2, 0xF2) if j > 0 else fill
                _style_cell(cell, val, size=body_size, bold=bold, fill=fill, align=align)
    return tbl


def _section_label(slide, text, left, top):
    db.txt(slide, text, left, top, Inches(5.5), Inches(0.22), size=10, bold=True, color=RED_PPTX)


# ── build charts ──────────────────────────────────────────────────────────
_gps_ladder_chart(cp("uti_gps_ladder_h.png"))
_bp_bridge_chart(cp("uti_bp_bridge_h.png"))

# ── assemble slide ────────────────────────────────────────────────────────
prs = db.new_presentation()
s = db.build_content(
    prs,
    title="FY27 stretch proposal — GPS target and business plan",
    eyebrow="UTI · Genuine Parts Share · FY2027",
    page_no=1,
)

M = Inches(0.32)
W_CHART = Inches(5.55)
W_TBL = Inches(6.95)
GAP = Inches(0.18)
TBL_LEFT = M + W_CHART + GAP

# Row 1 — GPS
r1_top = Inches(1.20)
_row_h = Inches(2.15)
_section_label(s, "1 · GPS TARGET", M, r1_top)
db.add_picture_fit(s, cp("uti_gps_ladder_h.png"), M, r1_top + Inches(0.24), W_CHART, _row_h)
_add_table(
    s,
    [
        ["Step", "GPS %", "vs prior"],
        ["FY26 actual", f"{UTI['gps_fy26']:.1f}%", "Starting point"],
        ["FY27 plan (base)", f"{UTI['gps_plan']:.1f}%", f"+{UTI['gps_plan'] - UTI['gps_fy26']:.1f}pp"],
        ["FY27 stretch", f"{UTI['gps_stretch']:.1f}%", f"+{UTI['pp']:.1f}pp on plan"],
    ],
    TBL_LEFT,
    r1_top + Inches(0.24),
    W_TBL,
    [Inches(2.35), Inches(1.55), Inches(2.95)],
    row_h=0.36,
)

# Row 2 — BP + incentive
r2_top = Inches(3.52)
_row_h2 = Inches(2.15)
_section_label(s, "2 · BUSINESS PLAN & INCENTIVE", M, r2_top)
db.add_picture_fit(s, cp("uti_bp_bridge_h.png"), M, r2_top + Inches(0.24), W_CHART, _row_h2)
_add_table(
    s,
    [
        ["Item", "Amount", "Notes"],
        ["BP FY27 (base)", _eur(UTI["bp_base"]), "Reviewed plan"],
        ["+ GPS stretch", f"+{_eur(UTI['gps_stretch_eur'])}", f"+{UTI['pp']:.1f}pp × {_eur(UTI['max_spend'])} wallet"],
        ["Stretch BP FY27", _eur(UTI["stretch_bp"]), f"+{UTI['delta_pct']:.1f}% vs base"],
        ["Regular kickback", _eur(UTI["regular_kickback_eur"]), "Unchanged"],
        ["GPS stretch bonus", _eur(UTI["incentive_eur"], k=True), "FY close — see below"],
    ],
    TBL_LEFT,
    r2_top + Inches(0.24),
    W_TBL,
    [Inches(2.35), Inches(1.75), Inches(2.75)],
    row_h=0.33,
)

# How it works — single line, grey box, above logo
how_top = Inches(5.88)
how_h = Inches(0.48)
db.rect(s, M, how_top, Inches(12.7), how_h, db.GRAY_BG)
db.txt(
    s,
    "How it works",
    M + Inches(0.14),
    how_top + Inches(0.05),
    Inches(1.35),
    Inches(0.18),
    size=9,
    bold=True,
    color=RED_PPTX,
)
db.txt(
    s,
    f"Stretch adds +{UTI['pp']:.1f} percentage points on your FY27 GPS plan "
    f"({UTI['gps_plan']:.1f}% → {UTI['gps_stretch']:.1f}%). "
    f"The purchase uplift is {UTI['pp']:.1f}% × your max-spend wallet ({_eur(UTI['max_spend'])}) = +{_eur(UTI['gps_stretch_eur'])} added to base BP. "
    f"€25K stretch bonus is paid at FY close if you reach stretch GPS ({UTI['gps_stretch']:.1f}%) or stretch BP ({_eur(UTI['stretch_bp'])}).",
    M + Inches(0.14),
    how_top + Inches(0.22),
    Inches(12.4),
    Inches(0.24),
    size=9,
    color=GRAY,
)

out_path = os.path.join(OUT_DIR, "UTI_GPS_Stretch_Proposal.pptx")
out_alt = os.path.join(OUT_DIR, "UTI_GPS_Stretch_Proposal_rebuild.pptx")
try:
    prs.save(out_path)
    print("Saved:", out_path)
except PermissionError:
    prs.save(out_alt)
    print("Saved (original open — close PPTX and rename):", out_alt)
