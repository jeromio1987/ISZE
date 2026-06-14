# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

RED       = RGBColor(0xCC, 0x0E, 0x1E)
LIGHT_RED = RGBColor(0xDC, 0x60, 0x6E)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
BLACK     = RGBColor(0x1A, 0x1A, 0x1A)
GRAY      = RGBColor(0x5A, 0x5A, 0x5A)

SW = Inches(13.33)
SH = Inches(7.5)
TEMPLATE = r"C:\Users\jerom\Desktop\claude\ISZE\Isuzu\isuzu-ppt-template\assets\template.pptx"

prs = Presentation(TEMPLATE)
while prs.slides:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

# ---------- helpers ----------
def set_white_bg(s):
    f = s.background.fill; f.solid(); f.fore_color.rgb = WHITE

def txt(s, text, left, top, width, height, size=14, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, font="Calibri", wrap=True):
    box = s.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = wrap
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
    return box

def rect(s, left, top, width, height, fill_color):
    sh = s.shapes.add_shape(1, left, top, width, height)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def oval(s, left, top, width, height, fill_color):
    sh = s.shapes.add_shape(9, left, top, width, height)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def add_logo(s, left, top, size=22):
    return txt(s, "ISUZU", left, top, Inches(2.0), Inches(0.55),
               size=size, bold=True, color=RED, font="Arial Black")

def add_footer_bar(s):
    h = Inches(0.13); rect(s, 0, SH - h, SW, h, RED)

def add_wave(s):
    oval(s, Inches(6.2), Inches(3.5), Inches(9.5), Inches(6.0), LIGHT_RED)
    oval(s, Inches(4.2), Inches(4.7), Inches(11.5), Inches(5.0), RED)
    add_footer_bar(s)

def cover(title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_white_bg(s)
    add_wave(s); add_logo(s, Inches(0.38), Inches(0.22), size=30)
    txt(s, title, Inches(0.45), Inches(2.45), Inches(7.6), Inches(2.0),
        size=40, bold=True, color=BLACK)
    txt(s, subtitle, Inches(0.45), Inches(5.30), Inches(7.0), Inches(0.5),
        size=15, color=GRAY)
    return s

def content(title):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_white_bg(s)
    add_logo(s, Inches(0.28), Inches(6.80), size=18)
    add_footer_bar(s)
    txt(s, title, Inches(0.28), Inches(0.30), Inches(12.77), Inches(0.72),
        size=30, bold=True, color=BLACK)
    rect(s, Inches(0.28), Inches(1.02), Inches(12.77), Emu(38100), RED)
    return s

def closing():
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_white_bg(s)
    txt(s, "Moving the World – for You", Inches(2.5), Inches(2.80),
        Inches(8.5), Inches(0.65), size=18, color=GRAY, align=PP_ALIGN.CENTER)
    txt(s, "ISUZU", Inches(1.5), Inches(3.20), Inches(10.5), Inches(2.1),
        size=100, bold=True, color=RED, align=PP_ALIGN.CENTER, font="Arial Black")
    return s

def body(s, items, x=Inches(0.6), y=Inches(1.45), w=Inches(12.1), size=17):
    """items: list of (level, lead, rest). level 0 = bullet, 1 = sub-dash.
       lead is bold; rest is normal. Either can be ''."""
    box = s.shapes.add_textbox(x, y, w, Inches(5.0)); tf = box.text_frame
    tf.word_wrap = True
    for i, (lvl, lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10 if lvl == 0 else 5)
        mark = "•  " if lvl == 0 else "      –  "
        sz = size if lvl == 0 else size - 2
        r0 = p.add_run(); r0.text = mark
        r0.font.size = Pt(sz); r0.font.color.rgb = (RED if lvl == 0 else GRAY)
        r0.font.name = "Calibri"; r0.font.bold = (lvl == 0)
        if lead:
            r1 = p.add_run(); r1.text = lead + (" " if rest else "")
            r1.font.size = Pt(sz); r1.font.bold = True
            r1.font.color.rgb = BLACK; r1.font.name = "Calibri"
        if rest:
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(sz); r2.font.color.rgb = BLACK; r2.font.name = "Calibri"
    return box

# =====================================================================
# 1  COVER
cover("ISZE – UTI Meeting",
      "Parts programme  ·  UIO  ·  Genuine Parts Share & PMS      |      June 2026")

# 2  AGENDA
s = content("Agenda")
body(s, [
    (0, "Parts programme (SPA)", "– aligning the basket and the flexibility we can offer."),
    (0, "BP 2027", "– led by UTI."),
    (0, "UIO reconciliation", "– confirming the active parc together."),
    (0, "Genuine Parts Share (GPS) & PMS", "– what they are and how we measure them."),
], y=Inches(1.7), size=19)

# 3  SPA
s = content("Parts programme (SPA)")
body(s, [
    (0, "", "We modelled 53 part numbers together to build a sustainable, competitive programme."),
    (0, "", "Current total volumes are lower than before, which makes the agreed targets harder to reach for both sides."),
    (0, "Flexibility from our side:", "we can run the programme over a 6-month period rather than a full year."),
    (0, "Two ways to work it:", "line by line, or on the total basket – whichever fits UTI's purchasing best."),
    (0, "", "The aim is a programme that stays competitive for UTI and viable for ISZE."),
], y=Inches(1.55), size=17)

# 4  LINES TO REVIEW
s = content("Two lines to review together")
body(s, [
    (0, "NOx sensor (8983528481):", "the largest single line in the basket."),
    (1, "", "A high share of this line moves through warranty rather than retail sales."),
    (1, "", "We propose reviewing the warranty volume together before fixing a programme price."),
    (0, "EGR cooler / valve:", "included specifically to help address parallel-import competition."),
    (1, "", "These are genuine-only parts, so the programme price is about defending against grey imports, not the aftermarket."),
], y=Inches(1.55), size=17)

# 5  UIO context + estimate
s = content("UIO reconciliation – our starting point")
body(s, [
    (0, "", "We are locking our units-in-operation (active parc) figures market by market, and would like to confirm UTI's directly before we close the numbers."),
    (0, "Our current model estimate for Israel, end-2025:", ""),
    (1, "N-Series:", "6,444 units"),
    (1, "F-Series:", "4,809 units"),
    (0, "", "These are derived from IML cumulative shipment data applied to a survival curve – not registration totals – so your local view may differ."),
], y=Inches(1.55), size=17)

# 6  UIO what we'd like
s = content("UIO reconciliation – what would help us")
body(s, [
    (0, "Active parc as of 31/12/2025", "– operating vehicles only, not cumulative registrations."),
    (0, "CV N- and F-series", "– split by tonnage level, if possible."),
    (0, "Source of the data", "– registration authority, your own records, or an estimate."),
    (0, "Any local scrap or de-registration rate", "you use internally."),
    (0, "Anything that makes your market diverge from our curve", "– used imports, exports, road conditions, registration practices."),
    (0, "Timing:", "information would be welcome by the beginning of next week."),
], y=Inches(1.5), size=16)

# 7  GPS
s = content("Genuine Parts Share (GPS)")
body(s, [
    (0, "The strategic question:", "of the total parts spend an Isuzu vehicle generates over its life, how much goes to ISZE genuine parts versus the aftermarket?"),
    (0, "", "It is hard to measure directly – we do not see aftermarket sales – so we start with what we can measure (PMS), then widen the lens."),
    (0, "We run it first at ISZE → distributor level;", "at a later stage it can extend to lower (dealer) levels."),
    (0, "", "UTI is welcome to run the same view at their own level if useful."),
], y=Inches(1.55), size=17)

# 8  PMS
s = content("PMS – the measurable entry point")
body(s, [
    (0, "PMS coverage:", "of the predictable annual maintenance demand each vehicle generates – oil, fuel and air filter, brake, clutch – how much does ISZE actually capture?"),
    (0, "How we measure it:", ""),
    (1, "", "UIO × replacement frequency = potential demand, compared with invoiced sales = coverage %."),
    (0, "What it shows:", "per category, where genuine sales are leaking to the aftermarket – so promotion and pricing target the real gap, not a gut feeling."),
    (0, "", "This is why accurate UIO matters: it is the base of the whole calculation."),
], y=Inches(1.55), size=17)

# 9  NEXT STEPS
s = content("Next steps")
body(s, [
    (0, "UIO:", "UTI to share active parc as of 31/12/2025, by the beginning of next week."),
    (0, "SPA:", "confirm the straightforward lines; review NOx warranty volume together; agree 6-month or full term, line-by-line or total."),
    (0, "GPS / PMS:", "confirm UTI's participation and the level to start from."),
], y=Inches(1.7), size=18)

# 10 CLOSING
closing()

import os
out = r"C:\Users\jerom\Desktop\claude\ISZE\Preparation Meetings\UTI\UTI_Meeting_June2026.pptx"
try:
    prs.save(out)
    print("saved", out)
except PermissionError:
    alt = out.replace(".pptx", "_v2.pptx")
    prs.save(alt)
    print("LOCKED - saved to", alt)
