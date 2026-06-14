# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

RED=RGBColor(0xCC,0x0E,0x1E)
doc=Document()
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)

def h(t,lvl=1):
    p=doc.add_heading(t,level=lvl)
    for r in p.runs: r.font.color.rgb=RED if lvl<=1 else RGBColor(0x40,0x40,0x40)
    return p
def para(t,bold=False,italic=False):
    p=doc.add_paragraph(); r=p.add_run(t); r.bold=bold; r.italic=italic; return p
def b(t):
    doc.add_paragraph(t,style='List Bullet')
def b2(t):
    p=doc.add_paragraph(t,style='List Bullet 2'); return p

doc.add_heading('UTI Meeting — Questions to Ask & Assumptions',level=0)
para('Internal preparation document · June 2026. Not for sharing as-is.',italic=True)
para('Two parts: (1) everything to ask UTI in the meeting, by topic; (2) all assumptions behind '
     'the figures used in the deck, so every number is defensible if challenged.')

# ============================ PART 1 ============================
h('Part 1 — What to ask UTI',1)

h('Parts programme (SPA)',2)
b('Which of the 53 lines can we confirm today as they stand?')
b('Do you prefer to work line-by-line, or fix the total basket first and only discuss exceptions?')
b('Are you open to a 6-month programme rather than a full year, so we can review and adjust?')
b('NOx sensor (8983528481): can we park this line and reconcile the warranty claim volume together '
  'before setting a programme price? (A large share of its volume is warranty, not retail.)')
b('EGR cooler / valve: confirm these are needed mainly to counter parallel / grey imports — is that '
  'still the priority on these lines?')

h('BP FY27',2)
para('Mostly UTI-led, but worth probing:')
b('FY27 started −17% vs budget and −31% vs last year (Apr–May). What is driving the softness — '
  'market, stock, end-customer demand?')
b('May order intake (€238K) was above sales — is the pipeline genuinely recovering?')
b('What support from ISZE would most help you close the gap to budget?')

h('UIO reconciliation',2)
para('This is what we already asked by email — confirm UTI will deliver it:')
b('Active parc as of 31/12/2025 — operating vehicles only, not cumulative registrations.')
b('CV N- and F-series, split by tonnage level if possible.')
b('Source of the data — registration authority, your own records, or an estimate.')
b('Any local scrap or de-registration rate you use internally.')
b('Anything that makes your market diverge from our survival curve — used imports, exports, '
  'road conditions, registration practices.')
b('Timing: can we have this by the beginning of next week?')

h('GPS & PMS',2)
b('Will UTI take part in the FY27 capture exercise?')
b('At which level do you want to start — ISZE→distributor, or would you also run it at your own level?')
b('Can you share your total category sales (oil/fuel/air filter, brake, clutch) so we can complete '
  'the capture calculation? (Today we only see what runs through ISZE.)')
b('Brakes specifically: how often do your 7.5t (N75) trucks actually go through front brake pads — '
  'and what annual mileage do they really run? (Our model assumes ~80,000 km / once every ~4 years '
  'at 20,000 km/yr — almost certainly too slack for urban distribution work. This is the single '
  'number that most distorts the brake result.)')
b('Do the rest of our working assumptions look right for Israel — share of parc in service, '
  'climate / duty cycle? (See Part 2 for the full list.)')

# ============================ PART 2 ============================
h('Part 2 — Assumptions behind the numbers',1)
para('Every figure in the deck rests on the assumptions below. The two that move the result most '
     'are annual mileage and the share of the parc actively in service.',bold=True)

h('A. UIO / active parc',2)
b('Our parc estimate = IML cumulative shipment data applied to the IMIT CV survival (scrap) curve — '
  'NOT registration totals.')
b('End-2025 model estimate: N-Series 6,444 + F-Series 4,809 = 11,253 units.')
b('The survival curve assumes a standard vehicle life. If Israel runs longer or shorter (used imports, '
  'exports, road conditions, registration practice), the real parc will diverge.')
b('Note a discrepancy to reconcile: the GPS/PMS capture model uses a UIO of ~10,668–11,019 units '
  '(internal "v13" model), while the shipment×survival estimate is 11,253. Different models / vintages — '
  'this is exactly why we are asking UTI to confirm.')

h('B. Genuine vs BVP — corrected from the catalog crossref (the part-dimension is broken)',2)
para('DATA-QUALITY WARNING. bvp_pn_dimension.csv — the table that labels each part genuine vs BVP — '
     'is unreliable: it mislabels 414 of 431 genuine catalog part numbers as "bvp", and a few BVP '
     'parts (e.g. oil-filter PN 1876102890) as "genuine". mix_genuine_bvp_pms.csv and '
     'bvp_downtrade_gp.csv are built on it, so their genuine/BVP splits are NOT trustworthy. Every '
     'figure below is re-derived directly from sales_fact + genuine_bvp_crossref.csv (the authoritative '
     'BVP catalog pairing, genuine_pn <-> bvp_pn).',bold=True)
b('Overall parts mix, 2025: UTI 88.9% genuine / 11.1% BVP; ITUK 97.3% genuine / 2.7% BVP. '
  'Both are genuine-dominant accounts — BVP is a small slice of total spend, NOT the headline.')
b('Within the contestable basket (only parts that HAVE a BVP alternative): UTI takes BVP 82% of '
  'units; ITUK only 21%. THIS is the real difference — UTI substitutes to the BVP economy line where '
  'ITUK keeps buying genuine.')
b('Filters are the gap. UTI buys air / fuel / oil filters ~100% as BVP. ITUK runs the same filters '
  '80% (fuel), 85% (air), 71% (oil) genuine — proof the genuine filter sale holds in a comparable '
  'network. Brake pads: both ~97–100% genuine already, so brakes are NOT a UTI gap.')
b('Opportunity (REVENUE, clean from sales): convert UTI BVP filters to genuine at network genuine '
  'prices — fuel filter +€54K (~11,400 units, BVP ~€8.7/u vs genuine ~€13.4/u), air filter +€11K. '
  'Fuel + air ≈ €65K/yr revenue. EXCLUDE oil filter: the genuine oil element is a low-price commodity '
  '(genuine ~€5/u < UTI BVP ~€10/u), so converting oil LOSES ~€23K. This is REVENUE; a gross-profit '
  'version needs a clean cost source — the old €45K GP and €221K gap both came from broken/soft files, '
  'do NOT quote them.')
b('This is measured from invoiced quantities and prices — no replacement-interval or mileage '
  'assumption needed. It is the defensible core of the GPS/PMS story; lead with the UTI-vs-ITUK '
  'filter gap.')

h('B-note. Deprecated numbers — do not quote',2)
b('Old N75 capture model: ~€221K gap and ~23% brake capture — built on an N75-only PN list and soft '
  'interval multipliers (its brake assumption implied a front-pad change only every ~80,000 km; a '
  '7.5t urban truck wears pads in ~30,000–60,000 km). Superseded; not defensible.')
b('Dimension-based figures ("oil 76% genuine", the "€45K GP", the oil-filter "trap", "UTI shifting to '
  'genuine"): all artifacts of the line_type mislabelling in Section B. Replaced by the crossref-based '
  'numbers above.')
b('The brake-interval insight still matters for ONE future question: sizing true aftermarket leakage '
  '(genuine demand captured in no form, not even BVP). That needs the real interval and share-of-parc '
  'in service — calibrate with UTI. Separate, softer exercise; not part of the firm filter story.')

h('C. GPS — spend per vehicle',2)
b('Spend per unit = total genuine + BVP parts sales ÷ UIO units.')
b('2023: €3.64M ÷ 11,019 = €331 · 2024: €2.65M ÷ 10,668 = €249 · 2025: €3.34M ÷ 10,668 (proxy) = €313.')
b('2025 figure inherits the proxy UIO, so treat the 2024→2025 move as indicative.')
b('The internal "green" traffic light is relative to comparable markets — a peer benchmark we do '
  'NOT share with UTI.')

h('D. BP FY27 performance',2)
b('These are actuals from the management report (Apr–May FY27) — no modelling assumptions.')
b('vs Budget and vs Prior year are simple ratios on the same figures.')

h('E. General caveats',2)
b('Data vintage: GPS/PMS figures cover 2023–2025; 2025 UIO is a proxy.')
b('Warranty data (used for the NOx point) is CV-2025 calendar; the SPA run-rate is its own window — '
  'so "about half of NOx volume is warranty" is indicative, not a precise overlap.')
b('No other distributor\'s figures are used anywhere in the UTI-facing deck (Part 3 below is '
  'internal network context only — not for sharing).')

# ============================ PART 3 ============================
h('Part 3 — Network context & strategic insights (internal, not for sharing)',1)
para('Sparring material. All figures re-derived from sales_fact + genuine_bvp_crossref.csv (2025), '
     'NOT the broken part-dimension. See Section B.',italic=True)

h('3.1  UTI in the network',2)
b('UTI is a mid-size PMS distributor. Overall parts mix 2025: 88.9% genuine / 11.1% BVP — '
  'it is NOT a BVP account; BVP is ~11% of spend.')
b('The signal is BEHAVIOUR inside the contestable basket: UTI takes BVP 82% of the time a BVP option '
  'exists; ITUK only 21%. UTI is the network\'s heaviest BVP-substituter on filters; ITUK among the '
  'lightest.')
b('So ITUK is the right benchmark — not because it sells more genuine overall (both are genuine-heavy '
  'overall), but because on the exact parts where UTI downtrades to BVP (filters), ITUK keeps genuine.')

h('3.2  The filter upsell — the clean prize',2)
b('UTI air / fuel / oil filters are bought ~100% BVP. Converting to genuine at network genuine prices: '
  'fuel +€54K, air +€11K = ~€65K/yr REVENUE. Oil excluded (genuine is the cheaper commodity).')
b('Fuel filter is the single prize: ~11,400 BVP units/yr at ~€8.7 vs genuine ~€13.4. ITUK already '
  'runs fuel 80% genuine — the conversion is demonstrably achievable in a comparable network.')
b('Brakes & clutch are NOT the UTI story: brake pads are already ~97% genuine; clutch kits have no '
  'genuine alternative being bought. The earlier "brakes & clutch under-captured" framing was an '
  'artifact of the broken dimension.')

h('3.3  Oil filter — leave it on BVP',2)
b('On oil filter the genuine element is a low-price commodity: network genuine ~€5/unit vs UTI BVP '
  '~€10/unit. Converting oil to genuine LOSES revenue (~€23K). Skip oil; focus fuel + air.')

h('3.4  The data bug that nearly shipped',2)
b('Two earlier versions of this analysis (incl. a "UTI 87% BVP / shifting to genuine / oil 76% '
  'genuine" story) were WRONG — both built on the part-dimension that mislabels genuine<->BVP. The '
  'single most misleading artifact: BVP oil-filter PN 1876102890 (3,793 units) tagged "genuine", '
  'inventing a genuine-oil surge that never happened (real genuine oil filters: ~12 units).')
b('Standing rule reinforced: classify genuine vs BVP from the catalog crossref (genuine_pn / bvp_pn), '
  'never from the dimension line_type, until the dimension is rebuilt. Verify any extreme single-line '
  'swing against the raw part number before believing it.')

h('3.5  AIM (aftermarket-risk overlay) — not usable yet',2)
b('AIM is doubly unusable: (a) it is built on the same broken dimension, and (b) even if fixed, it '
  'has no external aftermarket reference, so it only ever sees genuine<->BVP (both Isuzu) and cannot '
  'price independent-aftermarket risk. It flags ~€11K network — noise. Do not quote.')
b('To make AIM real: rebuild on the crossref, then add an external aftermarket price/brand feed and '
  'the missing 23 BVP wear categories (per BVP_FINDING.md the model covers only filters). Ask Snap-on '
  'for the BVP master file to get the per-category PNs.')
b('Until then, the defensible lever is the measured UTI-vs-ITUK filter gap (Section B / §3.2), not AIM.')

h('3.6  Data remediation needed (beyond this deck)',2)
b('Fix: realign bvp_pn_dimension.csv line_type to genuine_bvp_crossref.csv, then rebuild '
  'mix_genuine_bvp_pms.csv, bvp_downtrade_gp.csv and the aim_* files.')
b('This dimension also feeds the PRISM / GPS dashboards, so the same genuine-vs-BVP error is live '
  'there — the remediation is a network-data fix, not just a UTI-deck fix.')

out=r'C:\Users\jerom\Desktop\claude\ISZE\Preparation Meetings\UTI\UTI_Questions_and_Assumptions.docx'
try: doc.save(out); print('saved',out)
except PermissionError:
    alt=out.replace('.docx','_NEW.docx'); doc.save(alt); print('LOCKED - saved',alt)
