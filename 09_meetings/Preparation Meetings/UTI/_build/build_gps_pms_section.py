# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

RED       = RGBColor(0xCC, 0x0E, 0x1E)
LIGHT_RED = RGBColor(0xDC, 0x60, 0x6E)
VLRED     = RGBColor(0xF7, 0xDD, 0xE0)   # very light red, subtle motif
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
BLACK     = RGBColor(0x1A, 0x1A, 0x1A)
GRAY      = RGBColor(0x5A, 0x5A, 0x5A)
LGRAY     = RGBColor(0xF0, 0xF0, 0xF0)

SW, SH = Inches(13.33), Inches(7.5)
TEMPLATE = r"C:\Users\jerom\Desktop\claude\ISZE\Isuzu\isuzu-ppt-template\assets\template.pptx"

prs = Presentation(TEMPLATE)
while prs.slides:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId); del prs.slides._sldIdLst[0]

def set_white_bg(s):
    f = s.background.fill; f.solid(); f.fore_color.rgb = WHITE

def txt(s, text, left, top, width, height, size=14, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, font="Calibri", italic=False, wrap=True):
    box = s.shapes.add_textbox(left, top, width, height); tf = box.text_frame
    tf.word_wrap = wrap
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = font
    return box

def rect(s, left, top, width, height, fill_color):
    sh = s.shapes.add_shape(1, left, top, width, height)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def oval(s, left, top, width, height, fill_color):
    sh = s.shapes.add_shape(9, left, top, width, height)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def add_logo(s, left=Inches(0.28), top=Inches(6.80), size=18):
    return txt(s, "ISUZU", left, top, Inches(2.0), Inches(0.55),
               size=size, bold=True, color=RED, font="Arial Black")

def add_footer_bar(s):
    h = Inches(0.13); rect(s, 0, SH - h, SW, h, RED)

def content(title):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_white_bg(s)
    add_logo(s); add_footer_bar(s)
    txt(s, title, Inches(0.28), Inches(0.30), Inches(12.77), Inches(0.72),
        size=30, bold=True, color=BLACK)
    rect(s, Inches(0.28), Inches(1.02), Inches(12.77), Emu(38100), RED)
    return s

def body(s, items, x=Inches(0.6), y=Inches(1.6), w=Inches(12.1), size=18, big_gap=16):
    box = s.shapes.add_textbox(x, y, w, Inches(5.0)); tf = box.text_frame
    tf.word_wrap = True
    for i, (lvl, lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(big_gap if lvl == 0 else 5)
        mark = "•  " if lvl == 0 else "      –  "; sz = size if lvl == 0 else size-3
        r0 = p.add_run(); r0.text = mark
        r0.font.size = Pt(sz); r0.font.color.rgb = (RED if lvl==0 else GRAY)
        r0.font.bold = (lvl==0); r0.font.name="Calibri"
        if lead:
            r1=p.add_run(); r1.text=lead+(" " if rest else "")
            r1.font.size=Pt(sz); r1.font.bold=True; r1.font.color.rgb=BLACK; r1.font.name="Calibri"
        if rest:
            r2=p.add_run(); r2.text=rest
            r2.font.size=Pt(sz); r2.font.color.rgb=BLACK; r2.font.name="Calibri"
    return box

# =====================================================================
# 1  THE INVISIBLE MARKET  (recreated from explainer slide 2)
s = content("The invisible market")
body(s, [
    (0, "Every Isuzu in operation generates predictable parts demand.", ""),
    (0, "", "Today, neither ISZE nor our distributors can see how much of that demand we actually capture."),
    (0, "", "This is our first attempt to make it visible — and an invitation to refine it together."),
], y=Inches(1.7), size=19, big_gap=22)

# 2  GPS — elegant definition (three-beat, like the screen)
s = prs.slides.add_slide(prs.slide_layouts[6]); set_white_bg(s)
# subtle motif: two faint overlapping ovals bottom-right
oval(s, Inches(8.6), Inches(4.6), Inches(7.5), Inches(5.2), VLRED)
oval(s, Inches(10.2), Inches(5.6), Inches(6.5), Inches(4.5), LIGHT_RED)
add_logo(s); add_footer_bar(s)
txt(s, "GENUINE PARTS SHARE  (GPS)", Inches(0.6), Inches(1.05), Inches(9.0), Inches(0.4),
    size=14, bold=True, color=RED)
txt(s, "Every Isuzu on the road generates\na lifetime of parts demand.", Inches(0.6), Inches(1.75),
    Inches(9.3), Inches(1.8), size=34, bold=True, color=BLACK)
txt(s, "Genuine Parts Share is how much of that demand comes back to Isuzu —\nand how much leaks to the aftermarket.",
    Inches(0.6), Inches(3.85), Inches(9.0), Inches(1.1), size=19, color=GRAY)
txt(s, "What we don't capture, someone else supplies.", Inches(0.6), Inches(5.15),
    Inches(9.2), Inches(0.7), size=22, bold=True, italic=True, color=RED)

# 3  PMS — definition (clean)
s = content("PMS — the measurable entry point")
body(s, [
    (0, "PMS — Periodic Maintenance Service:", "the routine service parts every vehicle needs on a schedule — oil, fuel and air filters, brakes, clutch."),
    (0, "PMS coverage", "is the share of that demand we capture: UIO × replacement frequency, measured against what we actually invoice."),
    (0, "", "Maintenance demand is predictable — so it is the part of the market we can already see, and the clearest opportunity to act on."),
    (0, "It is the entry point to GPS:", "start with what is measurable, then widen the lens to the full parts basket."),
], y=Inches(1.6), size=18, big_gap=14)

# 4  FY27 ROLLOUT  (recreated from explainer slide 5)
s = content("FY27 rollout — across all ISZE markets")
body(s, [
    (0, "", "We are running the PMS coverage exercise for every ISZE market in FY27."),
    (0, "Two lenses, per market:", ""),
    (1, "PMS items comparison", "– within your market, which categories you capture best, and where the gap is."),
    (1, "Market comparison", "– an anonymised peer view: similar parc, different coverage — what moves the needle."),
    (0, "Phased:", "Q1 data & method · Q2 coverage & first analyses · Q3 action plans · Q4 implementation & feedback."),
    (0, "", "The exercise only works if every market takes part."),
], y=Inches(1.55), size=17, big_gap=11)

import os
out = r"C:\Users\jerom\Desktop\claude\ISZE\Preparation Meetings\UTI\GPS_PMS_Section.pptx"
try:
    prs.save(out); print("saved", out)
except PermissionError:
    alt = out.replace(".pptx","_v2.pptx"); prs.save(alt); print("LOCKED - saved", alt)
