# -*- coding: utf-8 -*-
"""Scan and produce fully clean EU Conf deck — remove ALL yellow/review markup."""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

SRC = Path(r"C:\Users\jerom\Downloads\EU Conf 26 ISZE Parts Sales 20260608 MH CHK.pptx")
OUT = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\Preparation Meetings\EU_Conf26"
    r"\EU Conf 26 ISZE Parts Sales 20260608 CLEAN.pptx"
)
OUT_DL = Path(r"C:\Users\jerom\Downloads\EU Conf 26 ISZE Parts Sales 20260608 CLEAN.pptx")

RED = RGBColor(0xD2, 0x00, 0x00)
GRAY = RGBColor(0x64, 0x64, 0x64)
BLACK = RGBColor(0x24, 0x24, 0x24)

YELLOW_HEX = {
    "FFFF00", "FFD01E", "FFC000", "FFEB3B", "FBBF24", "FFE135", "FFD700",
    "FFF176", "FFEE58", "FDD835", "FBC02D", "F9A825",
}


def remove_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def shape_has_yellow_stroke_or_fill(shape) -> bool:
    el = shape._element
    xml = el.xml
    for hx in YELLOW_HEX:
        if hx in xml.upper():
            return True
    return False


def is_review_shape(shape) -> bool:
    name = shape.name or ""
    if "Speech Bubble" in name:
        return True
    if "Arrow Connector" in name or name.startswith("Connector "):
        return True
    if name in ("Oval 42", "Oval 57"):
        return True
    if shape_has_yellow_stroke_or_fill(shape):
        return True
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t in ("VIN?", "Text overflow") or t.startswith("Use the same logic"):
            return True
    return False


def set_run_text(shape, text: str, *, size_pt: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def remove_all_review_shapes(prs) -> int:
    n = 0
    for slide in prs.slides:
        for sh in list(slide.shapes):
            if is_review_shape(sh):
                remove_shape(sh)
                n += 1
    return n


def apply_content_fixes(prs) -> None:
    replace_pairs = [
        ("Periodic Maintenance system", "Periodic Maintenance Service"),
        ("first?", "first"),
        ("Genuine Parts Spend per Unit", "Genuine Parts Share (GPS)"),
        ("Net Sales EUR  ÷  UIO  =  €/vehicle/year",
         "Actual genuine sales  ÷  Estimated parts demand  =  GPS %"),
    ]
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text
            for old, new in replace_pairs:
                if old in t:
                    if "Net Sales EUR" in old:
                        set_run_text(sh, new, size_pt=14, bold=True, color=RED)
                    elif "Genuine Parts Spend" in old:
                        set_run_text(sh, new, size_pt=14, bold=True, color=BLACK)
                    else:
                        set_run_text(sh, t.replace(old, new))

    s8 = prs.slides[7]
    for sh in s8.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Used to benchmark total parts"):
            set_run_text(
                sh,
                "Same logic as PMS: actual, estimate, coverage — applied to the full genuine parts basket.",
                size_pt=12, color=GRAY,
            )

    s9 = prs.slides[8]
    for sh in s9.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.replace("\x0b", "\n").strip()
        if t == "GPS — broad benchmark across all part numbers":
            set_run_text(sh, "GPS — same actual / estimate / coverage logic, full parts basket", size_pt=12)
        elif "All genuine PN" in t and "net sales" in t:
            set_run_text(sh, "Estimated\nparts demand", size_pt=11, bold=True)
        elif t.replace("\n", " ") == "UIO vehicle parc":
            set_run_text(sh, "Actual genuine\ninvoiced sales", size_pt=11, bold=True)
        elif t == "€/vehicle/year":
            set_run_text(sh, "GPS %\n(coverage)", size_pt=11, bold=True)
        elif t.startswith("PMS Coverage = Actual") and "GPS =" not in t:
            set_run_text(
                sh,
                "GPS = Actual genuine invoiced sales ÷ Estimated total parts demand    ·    "
                "PMS Coverage = Actual ISZE invoiced PMS sales ÷ Estimated potential PMS demand",
                size_pt=11, bold=True, color=RED,
            )

    s11 = prs.slides[10]
    for sh in s11.shapes:
        if sh.has_text_frame and "reporting for reporting" in sh.text_frame.text:
            set_run_text(sh, "The goal is real parts opportunities — not reporting for reporting's sake.", size_pt=12)


def scan_file(path: Path) -> list[str]:
    issues: list[str] = []
    prs = Presentation(str(path))
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if is_review_shape(sh):
                issues.append(f"slide {idx}: {sh.name}")
    with ZipFile(path) as z:
        for sn in sorted(n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)):
            xml = z.read(sn).decode("utf-8").upper()
            if "SPEECH BUBBLE" in xml:
                issues.append(f"{sn}: speech bubble xml")
            found = [y for y in YELLOW_HEX if y in xml]
            if found:
                issues.append(f"{sn}: yellow hex {found}")
    return issues


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source missing: {SRC}")

    backup = SRC.with_name(SRC.stem + f"_pre_clean_{datetime.now():%Y%m%d_%H%M}.pptx")
    shutil.copy2(SRC, backup)

    prs = Presentation(str(SRC))
    removed = remove_all_review_shapes(prs)
    apply_content_fixes(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    shutil.copy2(OUT, OUT_DL)
    # also overwrite original working name
    shutil.copy2(OUT, SRC)
    ws_copy = Path(
        r"C:\Users\jerom\Desktop\claude\ISZE\Preparation Meetings\EU_Conf26"
        r"\EU Conf 26 ISZE Parts Sales 20260608 MH CHK.pptx"
    )
    shutil.copy2(OUT, ws_copy)

    issues = scan_file(OUT)
    print(f"Removed {removed} review/yellow shapes")
    print(f"CLEAN deck: {OUT}")
    print(f"Downloads:  {OUT_DL}")
    print(f"Also updated: {SRC}")
    print("Scan issues:", issues or "NONE")


if __name__ == "__main__":
    main()
