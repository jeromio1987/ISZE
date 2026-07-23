# -*- coding: utf-8 -*-
"""Build ITUK and Imotors GPS/PMS decks using EU Conf 26 Isuzu master layout."""
from __future__ import annotations

import json
import sys
from copy import deepcopy
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[3]
PRISM = ROOT / "03_analytics" / "Isuzu" / "prism"
DERIVED = ROOT / "01_data" / "derived"
OUT_DIR = Path(__file__).resolve().parent
TEMPLATE = OUT_DIR / "EU Conf 26 ISZE Parts Sales 20260608 CLEAN.pptx"

RED = RGBColor(0xD2, 0x00, 0x00)
BLACK = RGBColor(0x24, 0x24, 0x24)
GRAY = RGBColor(0x64, 0x64, 0x64)
FONT = "Arial"

LY_INTRO = 1          # 1.2 Intro - Red
LY_CHAPTER = 10       # 3.1 Chapter - Red
LY_CONTENT = 24       # 4.11 Content - Title, Text, Red line
LY_CONTENT_IMG = 22   # 4.9 Content - Red, Title, Text, img.
LY_END = 30           # 5.1 END - Red symbol ENG

# Prototype slide indices in CLEAN deck (0-based)
PROTO_DEFINITIONS = 4   # slide 5 — GPS & PMS definitions
PROTO_BENEFITS = 7      # slide 8 — three-column benefit cards


def _load_prototype_shapes() -> dict[str, list]:
    prs = Presentation(str(TEMPLATE))
    return {
        "definitions": [deepcopy(sh.element) for sh in prs.slides[PROTO_DEFINITIONS].shapes],
        "benefits": [deepcopy(sh.element) for sh in prs.slides[PROTO_BENEFITS].shapes],
    }


PROTOTYPES = _load_prototype_shapes()
facts_global: dict = {}


def _load_facts() -> dict:
    with (DERIVED / "metrics_snapshot.json").open(encoding="utf-8") as f:
        snap = json.load(f)["metrics"]

    gps = pd.read_csv(DERIVED / "gps_spu_by_distributor.csv")
    pms_sum = pd.read_csv(PRISM / "evidence-app" / "data" / "derived" / "pms_summary_by_db.csv")
    pms_cap = pd.read_csv(PRISM / "evidence-app" / "data" / "derived" / "pms_capture_by_db_year.csv")
    parts_mix = pd.read_csv(PRISM / "evidence-app" / "data" / "derived" / "db_parts_mix_classification.csv")
    waterfall = pd.read_csv(PRISM / "evidence-app" / "data" / "derived" / "revenue_waterfall.csv")
    leakage = pd.read_csv(DERIVED / "dealer_leakage_ituk.csv")

    def gps_row(code: str) -> pd.Series:
        r = gps[(gps["distributor_code"] == code) & (gps["year"] == 2025)]
        return r.iloc[0]

    def pms_row(code: str) -> pd.Series:
        r = pms_sum[(pms_sum["distributor_code"] == code) & (pms_sum["year"] == 2025)]
        return r.iloc[0]

    def cap(code: str, cat: str, layer: str) -> float:
        r = pms_cap[
            (pms_cap["distributor_code"] == code)
            & (pms_cap["year"] == 2025)
            & (pms_cap["pms_category"] == cat)
            & (pms_cap["line_type"] == layer)
        ]
        return float(r["capture_rate"].iloc[0]) if len(r) else 0.0

    def rev(code: str, year: int) -> float:
        r = parts_mix[(parts_mix["distributor_code"] == code) & (parts_mix["year"] == year)]
        return float(r["total_eur"].iloc[0]) if len(r) else 0.0

    def wf_delta(code: str, family: str) -> float:
        r = waterfall[
            (waterfall["distributor_code"] == code)
            & (waterfall["year_to"] == 2025)
            & (waterfall["product_family"] == family)
        ]
        return float(r["total_delta_eur"].iloc[0]) if len(r) else 0.0

    ituk_g, pt_g = gps_row("ITUK"), gps_row("PORTUGAL")
    ituk_p, pt_p = pms_row("ITUK"), pms_row("PORTUGAL")
    wallet_ituk = float(ituk_g.get("max_spend_eur", 0) or 0)
    wallet_pt = float(pt_g.get("max_spend_eur", 0) or 0)

    return {
        "verify_ok": snap.get("verify.fail", {}).get("value", 1) == 0,
        "network_gps": snap["waterfall.network.gps_pct_2025"]["value"],
        "network_wallet_m": snap["waterfall.network.wallet_eur_2025"]["value"] / 1e6,
        "ituk": {
            "rev_24": rev("ITUK", 2024),
            "rev_25": rev("ITUK", 2025),
            "uio": int(ituk_g["uio_units"]),
            "spu": float(ituk_g["spend_per_unit_eur"]),
            "genuine_mix": snap["mix.ITUK.genuine_share_eur_pct_2025"]["value"],
            "wallet_m": wallet_ituk / 1e6,
            "wallet_genuine_pct": float(ituk_g.get("gps_capture_ex_warranty_vs_max", 0) or 0) * 100,
            "stretch_eur": wallet_ituk * 0.01,
            "pms_total": float(ituk_p["capture_rate_total_wtd"]) * 100,
            "pms_genuine": float(ituk_p["capture_rate_genuine_wtd"]) * 100,
            "pms_gap_m": float(ituk_p["revenue_gap_eur_sum"]) / 1e6,
            "top_gaps": str(ituk_p["top_gap_categories"]).replace("|", " · "),
            "oil_genuine": cap("ITUK", "oil_filter", "genuine") * 100,
            "oil_total": min(cap("ITUK", "oil_filter", "total") * 100, 100),
            "brake_shoe": cap("ITUK", "brake_shoe", "total") * 100,
            "pms_filters_delta": wf_delta("ITUK", "PMS_Filters"),
            "dealer_gap_k": float(snap["leakage.ituk.benchmark_gap_eur"]["value"]) / 1000,
            "dealers": leakage.nlargest(5, "gap_eur"),
            "has_dealer_data": True,
        },
        "imotors": {
            "rev_24": rev("PORTUGAL", 2024),
            "rev_25": rev("PORTUGAL", 2025),
            "uio": int(pt_g["uio_units"]),
            "spu": float(pt_g["spend_per_unit_eur"]),
            "wallet_m": wallet_pt / 1e6,
            "wallet_genuine_pct": float(pt_g.get("gps_capture_ex_warranty_vs_max", 0) or 0) * 100,
            "stretch_eur": wallet_pt * 0.01,
            "pms_total": float(pt_p["capture_rate_total_wtd"]) * 100,
            "pms_genuine": float(pt_p["capture_rate_genuine_wtd"]) * 100,
            "pms_gap_k": float(pt_p["revenue_gap_eur_sum"]) / 1000,
            "top_gaps": str(pt_p["top_gap_categories"]).replace("|", " · "),
            "oil_total": cap("PORTUGAL", "oil_filter", "total") * 100,
            "oil_genuine": cap("PORTUGAL", "oil_filter", "genuine") * 100,
            "air_total": cap("PORTUGAL", "air_filter", "total") * 100,
            "brakes_delta": wf_delta("PORTUGAL", "Brakes"),
            "has_dealer_data": False,
        },
    }


def _eur(v: float, decimals: int = 0) -> str:
    if abs(v) >= 1_000_000:
        return f"€{v / 1_000_000:.{1 if decimals else 0}f}M"
    if abs(v) >= 1000:
        return f"€{v / 1000:.{0 if decimals == 0 else 1}f}K"
    return f"€{v:,.0f}"


def _pct(v: float, decimals: int = 1) -> str:
    return f"{v:.{decimals}f}%"


def _yoy(cur: float, prev: float) -> str:
    if not prev:
        return "n/a"
    return f"{(cur / prev - 1) * 100:+.0f}%"


def _new_presentation() -> Presentation:
    prs = Presentation(str(TEMPLATE))
    while prs.slides:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]
    return prs


def _set_ph(slide, idx: int, text: str, *, size: int | None = None, bold: bool | None = None) -> None:
    ph = slide.placeholders[idx]
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _set_shape_text(shape, text: str, *, size: int = 14, bold: bool = False, color: RGBColor = BLACK) -> None:
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def _slide_from_prototype(prs: Presentation, proto_shapes: list, layout_idx: int = LY_CONTENT_IMG):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    for sh in list(slide.shapes):
        el = sh.element
        el.getparent().remove(el)
    for el in proto_shapes:
        slide.shapes._spTree.insert_element_before(deepcopy(el), "p:extLst")
    return slide


def _shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


class DeckBuilder:
    def __init__(self) -> None:
        self.prs = _new_presentation()

    def intro(self, title: str, subtitle: str) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LY_INTRO])
        _set_ph(slide, 0, title, size=36, bold=True)
        _set_ph(slide, 1, subtitle, size=18)

    def chapter(self, title: str) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LY_CHAPTER])
        _set_ph(slide, 0, title, size=40, bold=True)

    def content(self, title: str, body: str, *, body_size: int = 16) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LY_CONTENT])
        _set_ph(slide, 0, title, size=28, bold=True)
        _set_ph(slide, 18, body, size=body_size)

    def definitions_slide(self) -> None:
        _slide_from_prototype(self.prs, PROTOTYPES["definitions"])

    def _three_column_slide(
        self,
        title: str,
        subtitle: str,
        footer: str,
        columns: list[tuple[str, str, list[str]]],
    ) -> None:
        """Reuse EU Conf benefit-card layout (icons: col1–3)."""
        slide = _slide_from_prototype(self.prs, PROTOTYPES["benefits"])
        _set_shape_text(_shape_by_name(slide, "TextBox 1"), title, size=28, bold=True)
        _set_shape_text(_shape_by_name(slide, "TextBox 2"), subtitle, size=16, color=GRAY)
        _set_shape_text(_shape_by_name(slide, "TextBox 26"), footer, size=13, color=GRAY, bold=True)
        boxes = [
            ("TextBox 6", "TextBox 7", "TextBox 8", "TextBox 9", "TextBox 10"),
            ("TextBox 13", "TextBox 14", "TextBox 15", "TextBox 16", "TextBox 17"),
            ("TextBox 20", "TextBox 21", "TextBox 22", "TextBox 23", "TextBox 24"),
        ]
        for (icon_box, title_box, b1, b2, b3), (icon, col_title, bullets) in zip(boxes, columns):
            _set_shape_text(_shape_by_name(slide, icon_box), icon, size=36)
            _set_shape_text(_shape_by_name(slide, title_box), col_title, size=18, bold=True, color=RED)
            for box, line in zip((b1, b2, b3), bullets):
                _set_shape_text(_shape_by_name(slide, box), line, size=14)

    def dealer_data_slide(self, market_label: str, has_data: bool) -> None:
        if has_data:
            sub = (
                f"{market_label} shares dealer sell-out with ISZE — "
                "we see what leaves your warehouse and what reaches the workshop."
            )
            footer = (
                "Dealer-level data turns GPS/PMS from a distributor scorecard "
                "into a joint sales plan — not reporting for reporting's sake."
            )
            columns = [
                ("📦", "Replenishment", [
                    "• right parts in stock before the dealer calls",
                    "• filter & wear SKUs aligned to real bay demand",
                    "• fewer emergency orders and stock-outs",
                ]),
                ("📊", "Sales analysis", [
                    "• category mix vs network median by dealer",
                    "• spot under-performing accounts early",
                    "• target promotions where data shows a gap",
                ]),
                ("📍", "Capture lost sales", [
                    "• name the dealer + category + € opportunity",
                    "• PMS gaps become account actions, not averages",
                    "• track recovery quarter by quarter",
                ]),
            ]
        else:
            sub = (
                f"Today we model {market_label} from distributor billing and UIO. "
                "Dealer sell-out data would sharpen every number on this deck."
            )
            footer = (
                "Sharing dealer sales data is voluntary — but it is the fastest way "
                "to move from market averages to named opportunities."
            )
            columns = [
                ("📦", "Replenishment", [
                    "• forecast demand from workshop pull-through",
                    "• stock the SKUs dealers actually fit — not just what you buy",
                    "• reduce dead stock and urgent airfreight",
                ]),
                ("📊", "Sales analysis", [
                    "• compare dealer mix inside your network",
                    "• see which bays over-index on BVP or aftermarket",
                    "• build evidence-based dealer conversations",
                ]),
                ("📍", "Capture lost sales", [
                    "• quantify leakage per dealer and category",
                    "• prioritise field visits on € impact",
                    "• measure +1% GPS progress at dealer level",
                ]),
            ]
        slide = _slide_from_prototype(self.prs, PROTOTYPES["benefits"])
        _set_shape_text(
            _shape_by_name(slide, "TextBox 1"),
            "Sales to dealer data — why it matters",
            size=28,
            bold=True,
        )
        _set_shape_text(_shape_by_name(slide, "TextBox 2"), sub, size=16, color=GRAY)
        _set_shape_text(_shape_by_name(slide, "TextBox 26"), footer, size=13, color=GRAY, bold=True)
        boxes = [
            ("TextBox 6", "TextBox 7", "TextBox 8", "TextBox 9", "TextBox 10"),
            ("TextBox 13", "TextBox 14", "TextBox 15", "TextBox 16", "TextBox 17"),
            ("TextBox 20", "TextBox 21", "TextBox 22", "TextBox 23", "TextBox 24"),
        ]
        for (icon_box, title_box, b1, b2, b3), (icon, col_title, bullets) in zip(boxes, columns):
            _set_shape_text(_shape_by_name(slide, icon_box), icon, size=36)
            _set_shape_text(_shape_by_name(slide, title_box), col_title, size=18, bold=True, color=RED)
            for box, line in zip((b1, b2, b3), bullets):
                _set_shape_text(_shape_by_name(slide, box), line, size=14)

    def lubricants_slide(self, market: str) -> None:
        if market == "ituk":
            self._three_column_slide(
                "Centralised lubricants purchasing",
                "ISZE sources genuine BESCO lubricants centrally for Europe — "
                "ITUK distributes to dealers on one approved programme.",
                "Lubricants are a pull-through lever: every oil service is a genuine parts opportunity.",
                [
                    ("🛢️", "Centralised purchasing", [
                        "• ISZE negotiates volume with IML on behalf of the network",
                        "• one European price framework — less local sourcing friction",
                        "• simplified import, labelling and product compliance",
                    ]),
                    ("🔧", "For ITUK & dealers", [
                        "• genuine Isuzu-approved formulation — not generic CK-4",
                        "• displace Texaco Delo / aftermarket drums in the bay",
                        "• bundle lube + filters in a single call-off to dealers",
                    ]),
                    ("📈", "Parts & GPS synergy", [
                        "• oil change flags workshops still in the Isuzu ecosystem",
                        "• tie service intervals to filter & PMS capture",
                        "• lube volume signals where to focus lost-sales recovery",
                    ]),
                ],
            )
        else:
            self._three_column_slide(
                "Centralised lubricants purchasing",
                "Portugal joins the same European hub — Imotors sells genuine lubricants "
                "sourced by ISZE, alongside the parts basket.",
                "One supplier story for the workshop: genuine lube, genuine filters, genuine wear parts.",
                [
                    ("🛢️", "Centralised purchasing", [
                        "• scale with other ISZE markets on volume terms",
                        "• fewer parallel supplier relationships to manage",
                        "• ISZE holds the IML / BESCO supplier interface",
                    ]),
                    ("🔧", "For Imotors & workshops", [
                        "• predictable replenishment — aligned to service seasonality",
                        "• competitive vs local Castrol / Total alternatives",
                        "• workshop pack sizes matched to CV bay practice",
                    ]),
                    ("📈", "Parts & GPS synergy", [
                        "• every lube service is a filter upsell moment",
                        "• supports PMS fill-rate programmes already on this deck",
                        "• strengthens the genuine-only message at dealer level",
                    ]),
                ],
            )

    def end(self) -> None:
        self.prs.slides.add_slide(self.prs.slide_layouts[LY_END])

    def save(self, path: Path) -> Path:
        try:
            self.prs.save(path)
            return path
        except PermissionError:
            alt = path.with_name(path.stem + "_v2.pptx")
            self.prs.save(alt)
            return alt


def build_ituk_deck(f: dict) -> DeckBuilder:
    it = f["ituk"]
    d = DeckBuilder()

    d.intro(
        "ISZE – ITUK",
        "Genuine Parts Share & PMS  ·  FY2026 +1% GPS  ·  June 2026",
    )
    d.content(
        "Agenda",
        "• Why genuine now — Golden Window and FY2026 +1% GPS\n"
        "• GPS & PMS definitions\n"
        "• Sales to dealer data — replenishment, analysis, lost sales\n"
        "• Centralised lubricants purchasing — BESCO programme\n"
        "• Your performance — GPS snapshot, PMS gaps, dealer opportunities\n"
        "• BVP stance — protect genuine margin\n"
        "• FY2026 joint actions",
    )
    d.content(
        "Why genuine now — the Golden Window",
        "2025–2032 is peak parts-revenue potential for the ICE D-Max / N-series parc.\n\n"
        f"FY2026 stretch: +1 percentage point on GPS coverage "
        f"(≈ {_eur(it['stretch_eur'])} genuine on your wallet).\n\n"
        f"Network baseline CY2025: {_pct(f['network_gps'], 1)} GPS on "
        f"{_eur(f['network_wallet_m'] * 1e6)} addressable wallet.",
    )
    d.definitions_slide()
    d.dealer_data_slide("ITUK", has_data=True)
    d.lubricants_slide("ituk")
    d.content(
        "The +1% GPS stretch — FY2026",
        f"Addressable wallet: {_eur(it['wallet_m'] * 1e6)}\n"
        f"Current GPS wallet capture: {_pct(it['wallet_genuine_pct'])}\n"
        f"FY2026 target: {_pct(it['wallet_genuine_pct'] + 1)}\n"
        f"+1% value: ≈ {_eur(it['stretch_eur'])} incremental genuine\n\n"
        "Joint commitment — tracked quarterly from Apr 2026 baseline.",
    )
    d.content(
        "ITUK — GPS & revenue snapshot (CY2025)",
        f"Parts revenue: {_eur(it['rev_25'])} ({_yoy(it['rev_25'], it['rev_24'])} vs CY2024)\n"
        f"UIO: {it['uio']:,} (DVLA Mar 2026)  ·  SPU: {_eur(it['spu'])}\n"
        f"Contested-line genuine mix: {_pct(it['genuine_mix'], 1)}\n"
        f"PMS coverage: {_pct(it['pms_total'], 1)} total · {_pct(it['pms_genuine'], 1)} genuine\n"
        f"Modelled PMS gap: {_eur(it['pms_gap_m'] * 1e6)} · Top gaps: {it['top_gaps']}\n\n"
        f"PMS_Filters revenue {_eur(it['pms_filters_delta'])} YoY while total parts dipped — "
        "maintenance basket is resilient.",
    )
    d.content(
        "ITUK — PMS category focus",
        f"Oil filter (genuine): {_pct(min(it['oil_genuine'], 999), 0)} — mix win, demand saturated\n"
        f"Oil filter (total, capped): {_pct(it['oil_total'], 0)}\n"
        f"Brake shoe: {_pct(it['brake_shoe'], 1)} — #1 modelled gap\n"
        f"Dealer benchmark gap: {_eur(it['dealer_gap_k'] * 1000)} vs network-median mix\n\n"
        "With dealer sell-out we can name accounts — not just categories.",
    )
    dealers = it["dealers"]
    lines = ["Dealer · Category · Gap €", ""]
    for _, r in dealers.iterrows():
        lines.append(f"{str(r['dealer_name'])[:38]} · {r['category']} · €{r['gap_eur']:,.0f}")
    d.content("ITUK — dealer opportunities (top 5)", "\n".join(lines), body_size=14)
    d.content(
        "ITUK — BVP & contested sales",
        f"ITUK is a genuine stronghold: {_pct(it['genuine_mix'], 1)} contested-line genuine share.\n\n"
        "• Oil filter — genuine reclaiming share; do not push BVP at peak-window GP cost\n"
        "• Brakes — 100% genuine on contested lines; defend this position\n"
        "• Tier 1 BVP elsewhere — incremental only where genuine € is already minimal\n\n"
        "Maximising BVP share at ITUK would be value-destructive in the Golden Window.",
    )
    d.content(
        "ITUK — FY2026 joint actions",
        "1. Dealer PMS plan — top-10 accounts for brake shoe + urea filter\n"
        "2. Protect genuine on oil filter — no BVP push\n"
        "3. Quarterly GPS/PMS review — sell-out feed + ISZE billing; track +1% wallet capture\n"
        "4. Lubricants — agree BESCO pilot scope with top dealers Q1 FY2026",
    )
    d.content(
        "Benchmark — ITUK vs network",
        f"{'':22} {'ITUK':>10} {'Network':>10}\n"
        f"GPS wallet %{'':12} {it['wallet_genuine_pct']:>9.1f}% {f['network_gps']:>9.1f}%\n"
        f"PMS total %{'':13} {it['pms_total']:>9.1f}%\n"
        f"PMS genuine %{'':11} {it['pms_genuine']:>9.1f}%\n"
        f"SPU €{'':17} {it['spu']:>9,.0f}\n"
        f"Parts rev CY25{'':10} {_eur(it['rev_25']):>10}",
        body_size=15,
    )
    d.content(
        "Next steps",
        "• Agree dealer top-10 list and owner by end of June\n"
        "• Confirm quarterly GPS/PMS cadence from Apr 2026\n"
        "• Continue monthly dealer sell-out feed for replenishment & gap tracking",
    )
    d.end()
    return d


def build_imotors_deck(f: dict) -> DeckBuilder:
    im = f["imotors"]
    d = DeckBuilder()

    d.intro(
        "ISZE – Imotors",
        "Genuine Parts Share & PMS  ·  FY2026 +1% GPS  ·  June 2026",
    )
    d.content(
        "Agenda",
        "• Why genuine now — Golden Window and FY2026 +1% GPS\n"
        "• GPS & PMS definitions\n"
        "• Sales to dealer data — replenishment, analysis, lost sales\n"
        "• Centralised lubricants purchasing — European hub model\n"
        "• Your performance — GPS snapshot, PMS fill rate\n"
        "• BVP stance — genuine first on filters\n"
        "• FY2026 joint actions",
    )
    d.content(
        "Why genuine now — the Golden Window",
        "2025–2032 is peak parts-revenue potential for the ICE D-Max / N-series parc.\n\n"
        f"FY2026 stretch: +1 percentage point on GPS coverage "
        f"(≈ {_eur(im['stretch_eur'])} genuine on your wallet).\n\n"
        f"Network baseline CY2025: {_pct(f['network_gps'], 1)} GPS on "
        f"{_eur(f['network_wallet_m'] * 1e6)} addressable wallet.",
    )
    d.definitions_slide()
    d.dealer_data_slide("Imotors", has_data=False)
    d.lubricants_slide("imotors")
    d.content(
        "The +1% GPS stretch — FY2026",
        f"Addressable wallet: {_eur(im['wallet_m'] * 1e6)}\n"
        f"Current GPS wallet capture: {_pct(im['wallet_genuine_pct'])}\n"
        f"FY2026 target: {_pct(im['wallet_genuine_pct'] + 1)}\n"
        f"+1% value: ≈ {_eur(im['stretch_eur'])} incremental genuine\n\n"
        "Joint commitment — tracked quarterly from Apr 2026 baseline.",
    )
    d.content(
        "Imotors — GPS & revenue snapshot (CY2025)",
        f"Parts revenue: {_eur(im['rev_25'])} ({_yoy(im['rev_25'], im['rev_24'])} vs CY2024)\n"
        f"UIO: {im['uio']:,} (confirm with Imotors)  ·  SPU: {_eur(im['spu'])}\n"
        f"PMS coverage: {_pct(im['pms_total'], 1)} total · {_pct(im['pms_genuine'], 1)} genuine\n"
        f"Modelled PMS gap: {_eur(im['pms_gap_k'] * 1000)} · Top gaps: {im['top_gaps']}\n"
        f"Brakes YoY: {_eur(im['brakes_delta'])} — growth category\n\n"
        f"Genuine PMS capture {_pct(im['pms_genuine'], 1)} — largest upside lever vs network leaders.",
    )
    d.content(
        "Imotors — PMS category focus",
        f"Oil filter (total): {_pct(im['oil_total'], 1)} — genuine {_pct(im['oil_genuine'], 1)}\n"
        f"Air filter (total): {_pct(im['air_total'], 1)} — interval caveat applies\n"
        f"Top gaps: {im['top_gaps']}\n"
        f"+1% GPS value: {_eur(im['stretch_eur'])} on addressable wallet\n\n"
        "Dealer sell-out would sharpen these category gaps to named workshops.",
    )
    d.content(
        "Imotors — BVP & contested sales",
        "• Brakes — 35% genuine EUR on contested lines; revenue +€21K YoY\n"
        f"• Oil & air filter — grow genuine first (oil capture only {_pct(im['oil_total'], 1)})\n"
        "• Clutch kit — 68% BVP units; decide per account\n"
        "• Tier 1 BVP — suspension, cab, electrical where genuine € is minimal\n\n"
        "Strategy: maximise genuine on winnable PMS lines; BVP only where aftermarket already won.",
    )
    d.content(
        "Imotors — FY2026 joint actions",
        "1. Oil + air filter campaign — Q1 FY2026 genuine-first promotion\n"
        "2. Brake shoe programme — #1 modelled gap; bay-level targets\n"
        f"3. UIO reconciliation — confirm active parc (model: {im['uio']:,} units)\n"
        "4. Explore dealer sell-out sharing — unlock replenishment & lost-sales view\n"
        "5. Lubricants — confirm interest in centralised BESCO programme Q1 FY2026",
    )
    d.content(
        "Benchmark — Imotors vs network",
        f"{'':22} {'Imotors':>10} {'Network':>10}\n"
        f"GPS wallet %{'':12} {im['wallet_genuine_pct']:>9.1f}% {f['network_gps']:>9.1f}%\n"
        f"PMS total %{'':13} {im['pms_total']:>9.1f}%\n"
        f"PMS genuine %{'':11} {im['pms_genuine']:>9.1f}%\n"
        f"SPU €{'':17} {im['spu']:>9,.0f}\n"
        f"Parts rev CY25{'':10} {_eur(im['rev_25']):>10}",
        body_size=15,
    )
    d.content(
        "Next steps",
        "• Agree Q1 FY2026 filter campaign launch date\n"
        "• Imotors to confirm UIO by end of June\n"
        "• Discuss dealer data format for sell-out sharing (monthly CSV)",
    )
    d.end()
    return d


def main() -> int:
    global facts_global
    if not TEMPLATE.is_file():
        print(f"ERROR: template not found: {TEMPLATE}", file=sys.stderr)
        return 1

    facts_global = _load_facts()
    if not facts_global["verify_ok"]:
        print("WARNING: verification report not fully green", file=sys.stderr)

    for builder, name in (
        (build_ituk_deck(facts_global), "ITUK_GPS_PMS_June2026.pptx"),
        (build_imotors_deck(facts_global), "Imotors_GPS_PMS_June2026.pptx"),
    ):
        path = builder.save(OUT_DIR / name)
        print(f"saved {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
