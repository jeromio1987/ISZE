# -*- coding: utf-8 -*-
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
    r"C:\Users\jerom\Downloads\EU Conf 26 ISZE Parts Sales 20260608 CLEAN.pptx"
)
prs = Presentation(str(path))
slide = prs.slides[7]
print("SLIDE 8:", path.name)
print()

boxes = []
texts = []
for sh in slide.shapes:
    text = sh.text_frame.text.strip().replace("\n", " | ") if sh.has_text_frame else ""
    info = {
        "name": sh.name,
        "type": str(sh.shape_type),
        "left": sh.left,
        "top": sh.top,
        "width": sh.width,
        "height": sh.height,
        "right": sh.left + sh.width,
        "bottom": sh.top + sh.height,
        "text": text[:140],
    }
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        try:
            ast = str(sh.auto_shape_type)
            if "ROUNDED" in ast:
                boxes.append(info)
        except Exception:
            pass
    if sh.has_text_frame and info["text"]:
        texts.append(info)

print("=== ROUNDED BOXES ===")
for b in sorted(boxes, key=lambda x: x["left"]):
    print(
        f"{b['name']:20} L={b['left']:>8} T={b['top']:>8} "
        f"W={b['width']:>8} H={b['height']:>8} R={b['right']:>8} B={b['bottom']:>8}"
    )

print("\n=== TEXT SHAPES ===")
for t in sorted(texts, key=lambda x: (x["left"], x["top"])):
    print(
        f"{t['name']:20} L={t['left']:>8} T={t['top']:>8} "
        f"W={t['width']:>8} H={t['height']:>8} R={t['right']:>8} B={t['bottom']:>8}"
    )
    print(f"  TEXT: {t['text']}")
    for b in boxes:
        same_col = (t["left"] < 5000000 and b["left"] < 5000000) or (
            t["left"] >= 5000000 and b["left"] >= 5000000
        )
        if not same_col:
            continue
        over_r = t["right"] > b["right"]
        over_b = t["bottom"] > b["bottom"]
        over_l = t["left"] < b["left"]
        if over_r or over_b or over_l:
            print(
                f"  OVERFLOW vs {b['name']}: "
                f"left={over_l} right={over_r} bottom={over_b} "
                f"(dR={t['right']-b['right']}, dB={t['bottom']-b['bottom']})"
            )
    print()
