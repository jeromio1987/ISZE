"""New 17 Aug pack. Never writes How_We_Proceed_DRAFT or ORIGINAL."""
from __future__ import annotations

import os
import shutil
import subprocess
import sys

from PIL import Image, ImageOps
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.environ.get("ISZE_ROOT", r"C:\Users\jerom\Desktop\claude\ISZE")
TPL = os.path.join(ROOT, r"04_tools\ppt_template_v2")
sys.path.insert(0, TPL)

import isze_chart_theme as ct  # noqa: E402
import isze_deck_builder as db  # noqa: E402

OUT_FULL = os.path.join(ROOT, r"09_meetings\ISZE_PMS_Parts_Support_20260817_how_pct.pptx")
OUT_SHORT = os.path.join(ROOT, r"09_meetings\ISZE_PMS_Parts_Support_SHORT_20260817_how_pct.pptx")
ROOT_FULL = os.path.join(ROOT, "ISZE_PMS_Parts_Support_20260817_how_pct.pptx")
ROOT_SHORT = os.path.join(ROOT, "ISZE_PMS_Parts_Support_SHORT_20260817_how_pct.pptx")
CHART_DIR = os.path.join(HERE, "_charts_20260817")
os.makedirs(CHART_DIR, exist_ok=True)
PPT_EXE = r"C:\Program Files\Microsoft Office\root\Office16\POWERPNT.EXE"
PHOTO = os.path.join(HERE, "assets", "P700_Isuzu_Motors_Nseries_2023.jpg")

# Duncan register on content slides. Charts follow the same red.
INK = RGBColor(0x1B, 0x27, 0x33)
SLATE = RGBColor(0x34, 0x42, 0x51)
RED = RGBColor(0xED, 0x1C, 0x24)
WHITE = db.WHITE
NAVY = INK
PALE = RGBColor(0xF4, 0xF6, 0xF8)
FONT = "Calibri"

ct.RED = "#ED1C24"
ct.CHARCOAL = "#1B2733"
ct.GRAY = "#344251"
ct.SERIES_COLORS = ["#ED1C24", "#1B2733", "#5A6A7A", "#B7B7B7", "#F07A7E"]
ct.FONT_FAMILY = "Calibri"
import matplotlib.pyplot as plt  # noqa: E402
plt.rcParams["font.family"] = "Calibri"


def _guard(path: str) -> str:
    name = os.path.basename(path)
    if "How_We_Proceed" in name:
        raise SystemExit(f"Refusing to write {path}")
    return path


def _inches(x) -> float:
    return float(x.inches) if hasattr(x, "inches") else float(x)


def picture_cover(slide, src, left, top, width, height, cache_name):
    """Crop-fill a photo into the box (no letterbox)."""
    dpi = 180
    w_px = max(int(_inches(width) * dpi), 8)
    h_px = max(int(_inches(height) * dpi), 8)
    cache = os.path.join(CHART_DIR, cache_name)
    with Image.open(src) as im:
        fitted = ImageOps.fit(
            im.convert("RGB"), (w_px, h_px),
            method=Image.Resampling.LANCZOS, centering=(0.52, 0.42),
        )
        fitted.save(cache, quality=92)
    slide.shapes.add_picture(cache, left, top, width=width, height=height)


def add_chart(slide, path, left, top, width, height):
    """Insert a chart PNG stretched to the box so it is not a small island."""
    slide.shapes.add_picture(path, left, top, width=width, height=height)


def photo_panel(slide, cache_name):
    """Full-height right bleed. Caption is P700 only."""
    left = Inches(8.40)
    width = Inches(4.93)
    picture_cover(slide, PHOTO, left, 0, width, db.SH, cache_name)
    db.rect(slide, left, 0, Inches(0.06), db.SH, RED)
    db.rect(slide, left, Inches(6.96), width, Inches(0.44), INK)
    t(slide, "P700", left + Inches(0.16), Inches(6.96), Inches(4.50), Inches(0.44),
      size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def t(slide, text, left, top, width, height, size=12, bold=False, color=SLATE,
      align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    return db.txt(
        slide, text, left, top, width, height, size=size, bold=bold,
        color=color, align=align, font=FONT, anchor=anchor,
    )


def round_rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    return sh


def badge(slide, n, left, top):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.36), Inches(0.36))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RED
    sh.line.fill.background()
    t(slide, str(n), left, top, Inches(0.36), Inches(0.36), size=12, bold=True,
      color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def header(prs, title, page_no, kicker="Parts support  ·  scheduled PMS"):
    s = db.blank_slide(prs)
    db.set_white_bg(s)
    db.rect(s, 0, 0, Inches(0.18), db.SH, RED)
    t(s, kicker.upper(), Inches(0.48), Inches(0.18), Inches(11.4), Inches(0.26),
      size=11, bold=True, color=RED)
    t(s, title, Inches(0.48), Inches(0.44), Inches(12.2), Inches(0.50),
      size=26, bold=True, color=INK)
    db.rect(s, Inches(0.48), Inches(0.98), Inches(12.4), Emu(18000), RED)
    db.add_logo(s, left=Inches(0.48), top=Inches(6.92), height=Inches(0.26))
    db.add_footer_bar(s)
    db.add_page_number(s, page_no)
    return s


def cat_card(slide, left, top, width, height, title, tag, body, accent=RED):
    round_rect(slide, left, top, width, height, PALE)
    db.rect(slide, left, top, width, Inches(0.08), accent)
    t(slide, title, left + Inches(0.16), top + Inches(0.18), width - Inches(0.32),
      Inches(0.32), size=16, bold=True, color=INK)
    t(slide, tag, left + Inches(0.16), top + Inches(0.50), width - Inches(0.32),
      Inches(0.28), size=12, bold=True, color=RED)
    t(slide, body, left + Inches(0.16), top + Inches(0.82), width - Inches(0.32),
      height - Inches(0.96), size=12, color=SLATE)


def table(slide, headers, rows, left, top, width, col_w, row_h=0.38, head_h=0.36):
    head_h = Inches(head_h)
    row_h = Inches(row_h)
    db.rect(slide, left, top, width, head_h, INK)
    x = left
    for i, (label, w) in enumerate(zip(headers, col_w)):
        align = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.RIGHT
        t(slide, label, x + Inches(0.10), top, w - Inches(0.16), head_h,
          size=11, bold=True, color=WHITE, align=align, anchor=MSO_ANCHOR.MIDDLE)
        x += w
    y = top + head_h
    for ri, row in enumerate(rows):
        if ri % 2 == 1:
            db.rect(slide, left, y, width, row_h, PALE)
        x = left
        for i, (val, w) in enumerate(zip(row, col_w)):
            align = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.RIGHT
            t(slide, val, x + Inches(0.10), y, w - Inches(0.16), row_h,
              size=13, bold=(i == 0), color=INK if i == 0 else SLATE,
              align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += w
        y += row_h
    return y


def copy_pair(src, dest):
    os.makedirs(os.path.dirname(src), exist_ok=True)
    shutil.copy2(src, dest)


def close_locked(paths):
    try:
        import win32com.client  # type: ignore
        app = win32com.client.GetActiveObject("PowerPoint.Application")
        wanted = {os.path.normcase(os.path.abspath(p)) for p in paths}
        for pres in list(app.Presentations):
            try:
                fp = os.path.normcase(os.path.abspath(pres.FullName))
            except Exception:
                continue
            if "How_We_Proceed" in fp:
                continue
            if fp in wanted:
                pres.Save()
                pres.Close()
    except Exception:
        pass


def open_ppt(path):
    try:
        subprocess.Popen([PPT_EXE, path])
        return True
    except Exception:
        return False


def cover(prs):
    s = db.blank_slide(prs)
    db.set_white_bg(s)
    if os.path.isfile(PHOTO):
        photo_panel(s, "cover_bleed.jpg")
    db.rect(s, 0, 0, Inches(0.18), db.SH, RED)
    t(s, "ISUZU MOTORS EUROPE", Inches(0.55), Inches(1.55), Inches(7.5), Inches(0.32),
      size=12, bold=True, color=RED)
    t(s, "Parts support for scheduled PMS", Inches(0.55), Inches(2.05), Inches(7.5), Inches(1.15),
      size=32, bold=True, color=INK)
    t(s, "N75 first. Oil and fuel are the core. Air is the attach. "
         "Genuine first, unless a lower genuine SPA or a BVP is the better price.",
      Inches(0.55), Inches(3.35), Inches(7.4), Inches(1.15), size=15, color=SLATE)
    t(s, "18 August 2026  ·  complementary to the ITUK PMS pack",
      Inches(0.55), Inches(5.55), Inches(7.4), Inches(0.36), size=14, bold=True, color=INK)
    s.shapes.add_picture(db.LOGO_RED, Inches(0.55), Inches(6.48), height=Inches(0.32))
    db.add_footer_bar(s)
    return s


def contents(prs):
    s = header(prs, "Contents", 2)
    items = [
        ("01", "Where this sits", "Same N75 first, same 40,000 km service, our numbers on his parts map."),
        ("02", "The one percent", "All PMS parts, ITUK to dealer. 368 units. About EUR 7,530."),
        ("03", "Coverage", "Fuel is already there. Air is not. Next slide is how the percentages are built."),
        ("04", "Price route", "BVP where genuine is quiet. SPA if genuine can come in cheaper."),
        ("05", "Urea, brakes, clutch", "Urea stays. Brakes on condition. Clutch as a kit, manuals only."),
        ("06", "Aftermarket", "Same step: projected ITUK to dealers versus workshop purchase."),
    ]
    y = Inches(1.22)
    for num, title, body in items:
        db.rect(s, Inches(0.48), y, Inches(0.70), Inches(0.70), INK)
        t(s, num, Inches(0.48), y, Inches(0.70), Inches(0.70), size=14, bold=True,
          color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        t(s, title, Inches(1.38), y + Inches(0.04), Inches(11.2), Inches(0.30),
          size=16, bold=True, color=INK)
        t(s, body, Inches(1.38), y + Inches(0.34), Inches(11.2), Inches(0.32),
          size=13, color=SLATE)
        y += Inches(0.82)
    return s


def sits(prs, page_no=3):
    s = header(prs, "Same map as the ITUK pack. Our figures on it.", page_no)
    t(s, "Phase 1 is N75 fleet, 40,000 km, scheduled PMS only. Do not manufacture extra visits.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.36), size=14, color=SLATE)
    w = Inches(3.00)
    gap = Inches(0.16)
    x0 = Inches(0.48)
    cat_card(s, x0, Inches(1.62), w, Inches(4.55),
             "Oil + fuel", "Core pack",
             "Oil and fuel sit on every scheduled service. Fuel cover is already 95 percent to dealers. Oil BVP is already moving. Oil stays on the set. It is not a separate push.")
    cat_card(s, x0 + w + gap, Inches(1.62), w, Inches(4.55),
             "Air filter", "Primary attach",
             "Promote inclusion when the service menu requires it. Cover is 20 percent. That is the hole on the visit we already have.")
    cat_card(s, x0 + 2 * (w + gap), Inches(1.62), w, Inches(4.55),
             "Brakes", "Condition-based",
             "No volume push. Supply when the job needs it. Pad cover about 51 percent. No BVP on the R90 pad. SPA only if price is the barrier.")
    cat_card(s, x0 + 3 * (w + gap), Inches(1.62), w, Inches(4.55),
             "Clutch", "VIN / manual",
             "No broad push. Kit 5876102320, manuals only. Planning share 47 percent until a transmission file exists.")
    return s


def goal(prs, page_no=4):
    s = header(prs, "One extra point on all PMS parts", page_no)
    round_rect(s, Inches(0.48), Inches(1.22), Inches(6.05), Inches(2.55), PALE)
    t(s, "ITUK to dealer, calendar 2025", Inches(0.70), Inches(1.38), Inches(5.6), Inches(0.28),
      size=12, bold=True, color=RED)
    t(s, "+1%", Inches(0.70), Inches(1.70), Inches(5.6), Inches(0.85),
      size=54, bold=True, color=INK)
    t(s, "368 extra units  ·  about EUR 7,530 at ISZE to ITUK",
      Inches(0.70), Inches(2.60), Inches(5.6), Inches(0.70), size=16, color=SLATE)
    round_rect(s, Inches(6.75), Inches(1.22), Inches(6.15), Inches(2.55), INK)
    t(s, "Same mix as today. 56 part numbers. 36,843 units in the year.",
      Inches(6.95), Inches(1.45), Inches(5.75), Inches(0.70), size=14, color=WHITE)
    t(s, "Not 117 extra complete services. Not filters only. Brake is 42 percent of that euro.",
      Inches(6.95), Inches(2.25), Inches(5.75), Inches(1.15), size=14, color=WHITE)
    table(
        s,
        ["Family", "1% units", "1% EUR"],
        [
            ["Fuel", "180", "2,207"],
            ["Oil", "112", "774"],
            ["Brake", "36", "3,177"],
            ["Air", "36", "1,062"],
            ["Clutch", "4", "310"],
        ],
        Inches(0.48), Inches(4.00), Inches(12.42),
        [Inches(6.20), Inches(3.11), Inches(3.11)],
    )
    return s


def coverage(prs, page_no=5):
    s = header(prs, "N75 filter coverage, calendar 2025", page_no)
    t(s, "Sold quantity divided by the quantity expected from the IML schedule. "
         "9,359 vehicles. 1.25 services a year. Expected oil, fuel and air: 11,699. "
         "Fuel is already at the N75 service level. Air is not.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.50), size=14, color=SLATE)
    chart = os.path.join(CHART_DIR, "cover.png")
    ct.bar_chart(
        ["Oil", "Fuel", "Air"],
        {"ITUK to dealer": [46.2, 95.2, 19.6], "ISZE to ITUK": [39.9, 94.2, 19.0]},
        chart, value_fmt="{:.0f}%", figsize=(12.4, 3.05),
    )
    add_chart(s, chart, Inches(0.48), Inches(1.62), Inches(12.42), Inches(2.85))
    table(
        s,
        ["", "Expected", "ITUK to dealer", "Cover", "ISZE to ITUK", "Cover"],
        [
            ["Oil filter", "11,699", "5,404", "46.2%", "4,673", "39.9%"],
            ["Fuel filter", "11,699", "11,139", "95.2%", "11,021", "94.2%"],
            ["Air filter", "11,699", "2,289", "19.6%", "2,226", "19.0%"],
        ],
        Inches(0.48), Inches(4.55), Inches(12.42),
        [Inches(2.40), Inches(2.00), Inches(2.20), Inches(1.80), Inches(2.20), Inches(1.82)],
        row_h=0.42,
    )
    return s


def how_pct(prs, page_no=6):
    s = header(prs, "How the cover percentages are built", page_no)
    round_rect(s, Inches(0.48), Inches(1.16), Inches(12.42), Inches(0.72), INK)
    t(s, "Cover %  =  units sold in calendar 2025  /  units the service schedule expects",
      Inches(0.70), Inches(1.16), Inches(12.00), Inches(0.72),
      size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    w = Inches(4.00)
    gap = Inches(0.21)
    x0 = Inches(0.48)
    cat_card(s, x0, Inches(2.06), w, Inches(1.95),
             "1.  The car park", "9,359 N75",
             "United Kingdom N75 registered in 2024. Both engine generations. Not the 22,389 complete UK parc.")
    cat_card(s, x0 + w + gap, Inches(2.06), w, Inches(1.95),
             "2.  How often", "1.25 services / year",
             "Planning: 50,000 km a year, filter every 40,000 km. 50,000 / 40,000 = 1.25.")
    cat_card(s, x0 + 2 * (w + gap), Inches(2.06), w, Inches(1.95),
             "3.  Expected units", "11,699 filters",
             "9,359 x 1.25. Same expected quantity for oil, fuel and air. That is the denominator.")
    table(
        s,
        ["Worked example", "Sold 2025", "Expected", "Cover"],
        [
            ["Oil, ITUK to dealer", "5,404", "11,699", "46.2%"],
            ["Oil, ISZE to ITUK", "4,673", "11,699", "39.9%"],
            ["Fuel, ITUK to dealer", "11,139", "11,699", "95.2%"],
            ["Air, ITUK to dealer", "2,289", "11,699", "19.6%"],
        ],
        Inches(0.48), Inches(4.18), Inches(12.42),
        [Inches(4.20), Inches(2.74), Inches(2.74), Inches(2.74)],
        row_h=0.38,
    )
    t(s, "Same method on both layers. Pad cover uses 0.5 events a year, not 1.25. "
         "The +1% slide is different: it is 1% of actual 2025 mix, not 1% of expected services. "
         "Fuel 94% is versus N75 expected, not versus 22,389.",
      Inches(0.48), Inches(6.18), Inches(12.4), Inches(0.42), size=13, color=SLATE)
    return s


def price_route(prs, page_no=7):
    s = header(prs, "Price route. Pick the lower genuine or BVP.", page_no)
    t(s, "Agreed with ITUK 17 August. BVP where genuine is not moving, unless an SPA on genuine comes in cheaper. "
         "On the set: genuine first, unless genuine is already cheaper than BVP.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.55), size=14, color=SLATE)
    rows = [
        (1, "Fuel", "Genuine already moves.", "Keep genuine. BVP twin exists if needed."),
        (2, "Air", "Genuine cover 20 percent.", "BVP 5876102500, or SPA on genuine if that is lower."),
        (3, "Oil", "BVP already supplied.", "Stays on the set. Not a separate push."),
        (4, "Pad", "No BVP.", "SPA on genuine only, and only when the job needs the pad."),
    ]
    y = Inches(1.82)
    for n, title, a, b in rows:
        round_rect(s, Inches(0.48), y, Inches(12.42), Inches(1.00), PALE)
        badge(s, n, Inches(0.66), y + Inches(0.32))
        t(s, title, Inches(1.20), y + Inches(0.12), Inches(2.2), Inches(0.76),
          size=16, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        t(s, a, Inches(3.50), y + Inches(0.12), Inches(4.3), Inches(0.76),
          size=14, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        t(s, b, Inches(7.90), y + Inches(0.12), Inches(4.7), Inches(0.76),
          size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.10)
    return s


def filter_set(prs, page_no=8):
    s = header(prs, "One scheduled service. Three filters.", page_no)
    t(s, "Genuine set first. BVP only if genuine does not take, or if genuine is not the lower price. "
         "Air is the large piece of the bar. That is the attach on a visit that already happens.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.50), size=14, color=SLATE)
    chart = os.path.join(CHART_DIR, "set.png")
    ct.bar_chart(
        ["Genuine set", "Today (oil BVP in)", "Full BVP set"],
        {
            "ISZE to ITUK": [58.37, 53.88, 35.50],
            "Projected ITUK to dealers": [81.72, 75.43, 49.70],
        },
        chart, value_fmt="€{:.0f}", figsize=(12.4, 2.90),
    )
    add_chart(s, chart, Inches(0.48), Inches(1.62), Inches(12.42), Inches(2.62))
    table(
        s,
        ["", "Oil", "Fuel", "Air", "Set"],
        [
            ["Genuine PN", "8983282070", "8975425400", "8981772710", "three lines"],
            ["ISZE to ITUK EUR", "14.45", "12.56", "31.36", "58.37"],
            ["Projected to dealers EUR", "20.23", "17.58", "43.90", "81.72"],
            ["BVP PN", "1876102890", "1876102860", "5876102500", "if needed"],
        ],
        Inches(0.48), Inches(4.32), Inches(12.42),
        [Inches(3.10), Inches(2.33), Inches(2.33), Inches(2.33), Inches(2.33)],
        row_h=0.38,
    )
    return s


def urea(prs, page_no=9):
    s = header(prs, "Urea filter kit stays", page_no)
    t(s, "ITUK asked to keep it. Stays. Not a maybe. Not on the 56 PMS part numbers. Separate line. "
         "The item is the DEF pump filter kit, not bulk fluid.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.48), size=14, color=SLATE)
    cat_card(s, Inches(0.48), Inches(1.70), Inches(6.10), Inches(2.20),
             "Genuine  8983507161", "The live kit",
             "Keep genuine first. Projected ITUK to dealers about EUR 99 at 1.40.")
    cat_card(s, Inches(6.80), Inches(1.70), Inches(6.10), Inches(2.20),
             "BVP  1876102341", "If genuine does not take",
             "Use BVP if genuine does not take, unless a genuine SPA comes in lower.")
    table(
        s,
        ["", "Genuine 8983507161", "BVP 1876102341"],
        [
            ["ITUK to dealer qty", "541", "34"],
            ["ISZE to ITUK qty", "506", "46"],
            ["ISZE to ITUK EUR", "70.53", "30.72"],
            ["Projected to dealers EUR", "99", "43"],
        ],
        Inches(0.48), Inches(4.08), Inches(12.42),
        [Inches(4.14), Inches(4.14), Inches(4.14)],
        row_h=0.42,
    )
    t(s, "Projected to dealers = ISZE to ITUK x 1.40.",
      Inches(0.48), Inches(6.38), Inches(12.4), Inches(0.28), size=12, color=SLATE)
    return s


def brakes(prs, page_no=10):
    s = header(prs, "Brakes stay on condition. The pad is the price issue.", page_no)
    t(s, "No promotional trigger. When the pad is on the job, there is no BVP. SPA only. "
         "Compare projected ITUK to dealers with aftermarket workshop purchase. Same step.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.48), size=14, color=SLATE)
    chart = os.path.join(CHART_DIR, "pad.png")
    ct.bar_chart(
        ["ISZE to ITUK", "Projected ITUK to dealers", "Aftermarket workshop"],
        {"EUR": [89.11, 124.75, 46.12]},
        chart, value_fmt="€{:.0f}", figsize=(12.4, 2.35),
    )
    add_chart(s, chart, Inches(0.48), Inches(1.60), Inches(12.42), Inches(2.20))
    table(
        s,
        ["", "ITUK to dealer", "ISZE to ITUK", "Cover"],
        [
            ["Pad 8982447940", "2,367", "2,282", "51%  (0.5 / year)"],
            ["Disc rear 8981919560", "487", "521", "10%  (0.25 axle-set)"],
            ["Disc front 8982488951", "393", "520", "8%  (0.25 axle-set)"],
        ],
        Inches(0.48), Inches(3.92), Inches(12.42),
        [Inches(4.20), Inches(2.74), Inches(2.74), Inches(2.74)],
        row_h=0.44,
    )
    t(s, "Discs fit N75 NPR and carry over to P700. They are not on the ITUK PMS list. Aftermarket on the pad: Blue Print / Febi about EUR 46.",
      Inches(0.48), Inches(5.72), Inches(12.4), Inches(0.40), size=13, color=SLATE)
    t(s, "No BVP on the R90 pad. SPA on genuine only, and only when the job needs the pad.",
      Inches(0.48), Inches(6.12), Inches(12.4), Inches(0.32), size=13, bold=True, color=INK)
    return s


def clutch(prs, page_no=11):
    s = header(prs, "Clutch as a kit. Manuals only.", page_no)
    t(s, "No broad push. VIN / manual. 47 percent planning share until a transmission file exists. Do not mix disc 8981649181.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.40), size=14, color=SLATE)
    cat_card(s, Inches(0.48), Inches(1.68), Inches(6.10), Inches(2.55),
             "BVP kit  5876102320", "One number",
             "ITUK to dealer 63. ISZE to ITUK 78. ISZE to ITUK EUR 194.69. Projected to dealers about EUR 273.")
    cat_card(s, Inches(6.80), Inches(1.68), Inches(6.10), Inches(2.55),
             "Genuine three pieces", "If the workshop wants pieces",
             "8981649171 + 8970388314 + 1313100120. ISZE to ITUK EUR 282.08. Projected about EUR 395.")
    table(
        s,
        ["Piece", "PN", "ITUK to dealer", "ISZE to ITUK"],
        [
            ["Disc", "8981649171", "50", "47"],
            ["Pressure plate", "8970388314", "9", "54"],
            ["Release bearing", "1313100120", "38", "21"],
        ],
        Inches(0.48), Inches(4.45), Inches(12.42),
        [Inches(3.40), Inches(3.34), Inches(2.84), Inches(2.84)],
        row_h=0.46,
    )
    return s


def aftermarket(prs, page_no=12):
    s = header(prs, "The pad is well above aftermarket. Filters are close.", page_no)
    t(s, "Management line: the pad is the problem. Filters are close. The clutch is not. "
         "Compare projected ITUK price to dealers with what a workshop pays for the aftermarket part. Same step. ISZE to ITUK is shown so the chain is visible.",
      Inches(0.48), Inches(1.14), Inches(12.4), Inches(0.55), size=14, color=SLATE)
    table(
        s,
        ["Part", "ISZE to ITUK", "Projected to dealers", "Aftermarket workshop", "Read"],
        [
            ["Oil filter", "14.45", "20.23", "15.10", "Above"],
            ["Fuel filter", "12.56", "17.58", "14.45", "Above"],
            ["Air filter", "31.36", "43.90", "47.05", "In line"],
            ["Brake pad", "89.11", "124.75", "46.12", "Well above"],
            ["Clutch disc", "84.47", "118.26", "284", "Below"],
        ],
        Inches(0.48), Inches(1.82), Inches(12.42),
        [Inches(2.70), Inches(2.20), Inches(2.70), Inches(2.70), Inches(2.12)],
        row_h=0.46,
    )
    y = Inches(4.55)
    labels = [
        ("Above", "Oil and fuel, after 1.40"),
        ("In line", "Air"),
        ("Well above", "Pad. This is the SPA talk."),
    ]
    w = Inches(4.00)
    for i, (a, b) in enumerate(labels):
        x = Inches(0.48) + i * (w + Inches(0.18))
        fill = INK if i == 2 else PALE
        ink = WHITE if i == 2 else INK
        sl = WHITE if i == 2 else SLATE
        round_rect(s, x, y, w, Inches(1.55), fill)
        t(s, a, x + Inches(0.18), y + Inches(0.22), w - Inches(0.36), Inches(0.40),
          size=18, bold=True, color=ink)
        t(s, b, x + Inches(0.18), y + Inches(0.70), w - Inches(0.36), Inches(0.60),
          size=14, color=sl)
    t(s, "Projected to dealers = ISZE to ITUK x 1.40. Working mark-up. Confirm before IML. KS International: no price on these part numbers.",
      Inches(0.48), Inches(6.22), Inches(12.4), Inches(0.32), size=12, color=SLATE)
    return s


def close_slide(prs):
    s = db.blank_slide(prs)
    db.set_white_bg(s)
    if os.path.isfile(PHOTO):
        photo_panel(s, "close_bleed.jpg")
    db.rect(s, 0, 0, Inches(0.18), db.SH, RED)
    t(s, "October. N75 first.", Inches(0.55), Inches(2.35), Inches(7.5), Inches(0.70),
      size=32, bold=True, color=INK)
    t(s, "Scheduled PMS. Genuine first. BVP or SPA, whichever is the lower price. Urea stays.",
      Inches(0.55), Inches(3.20), Inches(7.4), Inches(1.00), size=16, color=SLATE)
    s.shapes.add_picture(db.LOGO_RED, Inches(0.55), Inches(6.48), height=Inches(0.32))
    db.add_footer_bar(s)
    return s


def build_full():
    prs = db.new_presentation()
    cover(prs)
    contents(prs)
    sits(prs)
    goal(prs)
    coverage(prs)
    how_pct(prs)
    price_route(prs)
    filter_set(prs)
    urea(prs)
    brakes(prs)
    clutch(prs)
    aftermarket(prs)
    close_slide(prs)
    path = _guard(OUT_FULL)
    prs.save(path)
    copy_pair(path, _guard(ROOT_FULL))
    return path, 13


def build_short():
    prs = db.new_presentation()
    cover(prs)
    sits(prs, 2)
    goal(prs, 3)
    coverage(prs, 4)
    how_pct(prs, 5)
    price_route(prs, 6)
    filter_set(prs, 7)
    aftermarket(prs, 8)
    close_slide(prs)
    path = _guard(OUT_SHORT)
    prs.save(path)
    copy_pair(path, _guard(ROOT_SHORT))
    return path, 9


def main():
    close_locked([OUT_FULL, OUT_SHORT, ROOT_FULL, ROOT_SHORT])
    full, n1 = build_full()
    short, n2 = build_short()
    showed = open_ppt(full)
    print("full", full, "slides", n1)
    print("short", short, "slides", n2)
    print("powerpoint", showed)
    print("did_not_touch DRAFT or ORIGINAL")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
