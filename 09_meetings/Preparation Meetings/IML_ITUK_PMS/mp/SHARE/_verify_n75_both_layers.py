from pathlib import Path
from openpyxl import load_workbook

PATHS = [
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\ISZE_PMS_N75_Coverage_Calculations.xlsx"),
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\ISZE_PMS_N75_Coverage_Calculations.xlsx"),
]
FORBIDDEN = ["IML v7", "Model B", "PRISM", "sold-to", "street", "\u2014", "\u2013"]

issues = []
wb = load_workbook(PATHS[0], data_only=False)
need = [
    "Cover", "Both_layers", "Parameters", "Expected", "Actual", "Coverage",
    "Worked_example_Fuel", "Worked_example_Air", "Worked_example_Oil", "What_this_is_not",
]
if wb.sheetnames != need:
    issues.append(f"order {wb.sheetnames}")
if PATHS[0].stat().st_size != PATHS[1].stat().st_size:
    issues.append("size mismatch")

# Both_layers formulas
bl = wb["Both_layers"]
for r, cat in ((6, "Fuel"), (7, "Air"), (8, "Oil")):
    for col in (3, 4, 5, 6, 7):
        v = bl.cell(r, col).value
        if not (isinstance(v, str) and v.startswith("=")):
            issues.append(f"Both_layers {cat} col{col} not formula: {v}")
    print(f"Both_layers r{r} {cat}:", [bl.cell(r, c).value for c in range(1, 8)])

# worked fuel/air both layers
for name, a, b in (
    ("Worked_example_Fuel", "FUEL_QTY_ISZE_ITUK", "FUEL_QTY_DEALER"),
    ("Worked_example_Air", "AIR_QTY_ISZE_ITUK", "AIR_QTY_DEALER"),
):
    sh = wb[name]
    print(name, "7a", sh["C11"].value, "8a", sh["C12"].value, "7b", sh["C13"].value, "8b", sh["C14"].value)
    if a not in str(sh["C11"].value):
        issues.append(f"{name} missing {a}")
    if b not in str(sh["C13"].value):
        issues.append(f"{name} missing {b}")
    if sh["C12"].value != "=D11/D10":
        issues.append(f"{name} 8a {sh['C12'].value}")
    if sh["C14"].value != "=D13/D10":
        issues.append(f"{name} 8b {sh['C14'].value}")

# params lock
lock = {
    "UIO_N75": 9359, "KM_PER_YEAR": 50000, "FILTER_INTERVAL_KM": 40000,
    "WEAR_INTERVAL_KM": 100000, "CLUTCH_MANUAL_SHARE": 0.47, "ADDRESSABLE_SHARE": 1.0,
    "OIL_QTY_ISZE_ITUK": 4673, "FUEL_QTY_ISZE_ITUK": 11021, "AIR_QTY_ISZE_ITUK": 2226,
    "OIL_QTY_DEALER": 5404, "FUEL_QTY_DEALER": 11139, "AIR_QTY_DEALER": 2289,
}
for name, expect in lock.items():
    sheet, coord = list(wb.defined_names[name].destinations)[0]
    val = wb[sheet][coord].value
    if val != expect:
        issues.append(f"{name}={val}")

# leaks
for name in wb.sheetnames:
    for row in wb[name].iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            for bad in FORBIDDEN:
                if bad in v:
                    issues.append(f"leak {name} {cell.coordinate} {bad}")

# computed table
uio, km, iv = 9359, 50000, 40000
exp = uio * (km / iv)
print("\nTABLE")
print(f"{'cat':8} {'expected':>10} {'ISZE->ITUK':>10} {'%':>7} {'ITUK->DLR':>10} {'%':>7}")
for cat, q1, q2 in (
    ("Fuel", 11021, 11139),
    ("Air", 2226, 2289),
    ("Oil", 4673, 5404),
):
    print(f"{cat:8} {exp:10.2f} {q1:10} {q1/exp:7.1%} {q2:10} {q2/exp:7.1%}")

# oil labelled
if "not an ISZE campaign" not in str(bl["B8"].value):
    issues.append("oil campaign label missing on Both_layers")

print("issues", issues)
if issues:
    raise SystemExit(1)
print("VERIFY OK")
