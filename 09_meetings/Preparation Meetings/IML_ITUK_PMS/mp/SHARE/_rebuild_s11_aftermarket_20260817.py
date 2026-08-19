"""Replace ONLY slide 11 on a copy of the 17 Aug pack.

Never writes How_We_Proceed_DRAFT or ORIGINAL.
Never rebuilds the rest of the deck.
"""
from __future__ import annotations

import importlib.util
import os
import shutil
import subprocess
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.environ.get("ISZE_ROOT", r"C:\Users\jerom\Desktop\claude\ISZE")
SRC = os.path.join(ROOT, r"09_meetings\ISZE_PMS_Parts_Support_20260817.pptx")
OUT = os.path.join(ROOT, r"09_meetings\ISZE_PMS_Parts_Support_20260817_s11.pptx")
OUT_ROOT = os.path.join(ROOT, "ISZE_PMS_Parts_Support_20260817_s11.pptx")
PPT_EXE = r"C:\Program Files\Microsoft Office\root\Office16\POWERPNT.EXE"
TPL = os.path.join(ROOT, r"04_tools\ppt_template_v2")
sys.path.insert(0, TPL)

import isze_deck_builder as db  # noqa: E402

spec = importlib.util.spec_from_file_location(
    "ps17", os.path.join(HERE, "_build_parts_support_20260817.py")
)
ps = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(ps)


def _guard(path: str) -> str:
    if "How_We_Proceed" in os.path.basename(path):
        raise SystemExit(f"Refusing to write {path}")
    return path


def wipe(slide) -> None:
    tree = slide.shapes._spTree
    for child in list(tree):
        tag = child.tag
        if tag.endswith("nvGrpSpPr") or tag.endswith("grpSpPr"):
            continue
        tree.remove(child)


def paint_header(slide, title: str, page_no: int) -> None:
    db.set_white_bg(slide)
    db.rect(slide, 0, 0, Inches(0.18), db.SH, ps.RED)
    ps.t(
        slide, "PARTS SUPPORT  ·  SCHEDULED PMS",
        Inches(0.48), Inches(0.14), Inches(11.4), Inches(0.22),
        size=11, bold=True, color=ps.RED,
    )
    ps.t(
        slide, title,
        Inches(0.48), Inches(0.34), Inches(12.2), Inches(0.40),
        size=22, bold=True, color=ps.INK,
    )
    db.rect(slide, Inches(0.48), Inches(0.76), Inches(12.4), Emu(18000), ps.RED)
    db.add_logo(slide, left=Inches(0.48), top=Inches(6.92), height=Inches(0.26))
    db.add_footer_bar(slide)
    db.add_page_number(slide, page_no)


def paint_slide11(slide) -> None:
    paint_header(slide, "Aftermarket compare. Projected figures only.", 11)

    # Big disclaimer. Sensitive. Must be unmissable.
    db.rect(slide, Inches(0.48), Inches(0.88), Inches(12.42), Inches(1.18), ps.RED)
    ps.t(
        slide, "DISCLAIMER",
        Inches(0.64), Inches(0.94), Inches(12.10), Inches(0.26),
        size=13, bold=True, color=ps.WHITE,
    )
    ps.t(
        slide,
        "We do not know the actual ISZE dealer price. We do not know the actual dealer sell-out. "
        "We work with a projected ITUK margin (x 1.40 on our list) and a projected dealer (DLR) margin. "
        "The 1.40 column is not a confirmed ITUK invoice and not a known dealer sell-out.",
        Inches(0.64), Inches(1.20), Inches(12.10), Inches(0.76),
        size=13, bold=True, color=ps.WHITE,
    )

    ps.t(
        slide,
        "Read only projected ITUK to dealers versus aftermarket workshop. Same step. Do not sign ISZE to ITUK against aftermarket.",
        Inches(0.48), Inches(2.12), Inches(12.42), Inches(0.28),
        size=13, color=ps.SLATE,
    )

    ps.table(
        slide,
        ["Part", "ISZE to ITUK list", "Projected ITUK to dealers", "Aftermarket workshop", "Read"],
        [
            ["Oil filter", "14.45", "20.23", "15.10", "Above"],
            ["Fuel filter", "12.56", "17.58", "14.45", "Above"],
            ["Air filter", "31.36", "43.90", "47.05", "In line"],
            ["Brake pad", "89.11", "124.75", "46.12", "Well above"],
            ["Clutch disc", "84.47", "118.26", "284", "Below"],
        ],
        Inches(0.48), Inches(2.42), Inches(12.42),
        [Inches(2.20), Inches(2.30), Inches(2.90), Inches(2.70), Inches(2.32)],
        row_h=0.34,
        head_h=0.32,
    )

    # Source + webshop
    ps.round_rect(slide, Inches(0.48), Inches(4.52), Inches(6.10), Inches(2.22), ps.PALE)
    db.rect(slide, Inches(0.48), Inches(4.52), Inches(6.10), Inches(0.08), ps.RED)
    ps.t(
        slide, "Source",
        Inches(0.64), Inches(4.66), Inches(5.78), Inches(0.26),
        size=14, bold=True, color=ps.INK,
    )
    ps.t(
        slide,
        "Aftermarket = Eucon UK IAM price survey 2023, HCV sheet, workshop purchase GBP, converted at 1.17. "
        "Not a 2026 live UK shop quote.\n\n"
        "KS International webshop: searched these exact part numbers, August 2026. No price. "
        "That is the live webshop check. It did not give a 2026 quote.",
        Inches(0.64), Inches(4.94), Inches(5.78), Inches(1.68),
        size=12, color=ps.SLATE,
    )

    # Brand quality
    ps.round_rect(slide, Inches(6.80), Inches(4.52), Inches(6.10), Inches(2.22), ps.PALE)
    db.rect(slide, Inches(6.80), Inches(4.52), Inches(6.10), Inches(0.08), ps.RED)
    ps.t(
        slide, "Brand quality",
        Inches(6.96), Inches(4.66), Inches(5.78), Inches(0.26),
        size=14, bold=True, color=ps.INK,
    )
    ps.t(
        slide,
        "IAM brands Eucon marked OE comparable (Blue Print, Febi, Bosch, Mahle). "
        "Same application, not the same part as genuine.\n\n"
        "Pad 46.12 is the cheap IAM cluster (Blue Print / Febi). Ferodo is higher. Textar is much higher. The median is not Ferodo.\n\n"
        "Clutch 284 is OE-comparable Blue Print only. Exedy and Sampa are marked lower quality.",
        Inches(6.96), Inches(4.94), Inches(5.78), Inches(1.68),
        size=12, color=ps.SLATE,
    )


def main() -> int:
    if not os.path.isfile(SRC):
        raise SystemExit(f"Missing source pack: {SRC}")
    _guard(OUT)
    _guard(OUT_ROOT)
    ps.close_locked([OUT, OUT_ROOT])
    shutil.copy2(SRC, OUT)
    prs = Presentation(OUT)
    n = len(prs.slides)
    if n != 12:
        raise SystemExit(f"Expected 12 slides in source copy, got {n}")
    slide = prs.slides[10]
    wipe(slide)
    paint_slide11(slide)
    prs.save(OUT)
    shutil.copy2(OUT, OUT_ROOT)
    showed = False
    try:
        subprocess.Popen([PPT_EXE, OUT])
        showed = True
    except Exception:
        showed = False
    print("out", OUT)
    print("root_copy", OUT_ROOT)
    print("slides", len(prs.slides))
    print("changed_slide", 11)
    print("powerpoint", showed)
    print("did_not_touch DRAFT or ORIGINAL or 20260817 full pack")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
