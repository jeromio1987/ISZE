"""Verify N75 coverage workbook: formulas present, params locked, no shared leaks."""
from pathlib import Path
from openpyxl import load_workbook

PATHS = [
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\ISZE_PMS_N75_Coverage_Calculations.xlsx"),
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\ISZE_PMS_N75_Coverage_Calculations.xlsx"),
]
FORBIDDEN = [
    "IML v7", "Model B", "PRISM", "sold-to", "street", "Interval Model",
    "Duncan", "Rene", "Jerome", "Masumi", "\u2014", "\u2013", "blank-%",
]
EXPECTED_SHEETS = [
    "Cover", "Parameters", "Expected", "Actual", "Coverage",
    "Worked_example_Oil", "What_this_is_not",
]

# Locked source values
LOCK = {
    "UIO_N75": 9359,
    "KM_PER_YEAR": 50000,
    "FILTER_INTERVAL_KM": 40000,
    "WEAR_INTERVAL_KM": 100000,
    "CLUTCH_MANUAL_SHARE": 0.47,
    "ADDRESSABLE_SHARE": 1.0,
    "OIL_QTY_ISZE_ITUK": 4673,
    "FUEL_QTY_ISZE_ITUK": 11021,
    "AIR_QTY_ISZE_ITUK": 2226,
    "BRAKE_QTY_ISZE_ITUK": 2902,
    "CLUTCH_QTY_ISZE_ITUK": 243,
    "OIL_QTY_DEALER": 5404,
    "FUEL_QTY_DEALER": 11139,
    "AIR_QTY_DEALER": 2289,
    "BRAKE_QTY_DEALER": 2923,
    "CLUTCH_QTY_DEALER": 340,
}

issues = []
wb0 = load_workbook(PATHS[0], data_only=False)
wb1 = load_workbook(PATHS[1], data_only=False)

if wb0.sheetnames != EXPECTED_SHEETS:
    issues.append(f"sheetnames {wb0.sheetnames}")
if PATHS[0].stat().st_size != PATHS[1].stat().st_size:
    issues.append(f"size mismatch {PATHS[0].stat().st_size} vs {PATHS[1].stat().st_size}")

# named range values
print("=== named ranges ===")
for name, expect in LOCK.items():
    dn = wb0.defined_names[name]
    dest = list(dn.destinations)[0]
    sheet, coord = dest
    val = wb0[sheet][coord].value
    ok = val == expect
    print(f"  {name:32} {val!r:16} expect {expect!r:16} {'OK' if ok else 'FAIL'}")
    if not ok:
        issues.append(f"{name}={val} != {expect}")

# derived formula cells
print("=== derived formulas on Parameters ===")
ws = wb0["Parameters"]
derived = {
    "C12": ("=C6/C7", "filter events"),
    "C13": ("=C6/C8", "wear events"),
    "C14": ("=C5*C9", "clutch UIO"),
    "C15": ("=C5*C12*C10", "expected filters"),
    "C16": ("=C5*C13*C10", "expected brake"),
    "C17": ("=C14*C13*C10", "expected clutch"),
    "C42": ("=C5+C39+C40+C41", "10684 checksum"),
}
for coord, (formula, lab) in derived.items():
    val = ws[coord].value
    ok = val == formula
    print(f"  {coord} {val!r:30} {lab:20} {'OK' if ok else 'FAIL'}")
    if not ok:
        issues.append(f"{coord} {val} != {formula}")

# coverage formulas must be formulas
print("=== Coverage D5:D9 and F5:F9 and H5:H9 ===")
cov = wb0["Coverage"]
for r in range(5, 10):
    for col in ("D", "F", "H"):
        v = cov[f"{col}{r}"].value
        is_f = isinstance(v, str) and v.startswith("=")
        print(f"  {col}{r} {v!r:40} {'OK' if is_f else 'FAIL'}")
        if not is_f:
            issues.append(f"Coverage {col}{r} not formula: {v}")

# count formulas
print("=== formula counts ===")
for name in EXPECTED_SHEETS:
    n = 0
    leaks = []
    for row in wb0[name].iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                n += 1
            if isinstance(v, str):
                for bad in FORBIDDEN:
                    if bad.lower() in v.lower() if bad.isascii() else bad in v:
                        leaks.append((cell.coordinate, bad, v[:80]))
    print(f"  {name:22} formulas={n:3} leaks={len(leaks)}")
    for item in leaks:
        issues.append(f"leak {name} {item}")
        print("   LEAK", item)

# python maths vs source display
uio, km, iv, wear, mt, addr = 9359, 50000, 40000, 100000, 0.47, 1.0
filt = km / iv
w = km / wear
exp_f = uio * filt * addr
exp_b = uio * w * addr
exp_c = uio * mt * w * addr
print("=== python lock check ===")
print(f"  filter events {filt}")
print(f"  wear events   {w}")
print(f"  expected filters {exp_f}  pack 11699")
print(f"  expected brake   {exp_b}  pack 4680")
print(f"  expected clutch  {exp_c}  pack 2199")
print(f"  oil DB   {4673/exp_f:.4%}  unrounded {4673.3438/exp_f:.4%}")
print(f"  fuel DB  {11021/exp_f:.4%}")
print(f"  air DB   {2226/exp_f:.4%}")
print(f"  brake DB {2902/exp_b:.4%}")
print(f"  clutch DB {243/exp_c:.4%}")
print(f"  oil DLR  {5404/exp_f:.4%}")
print(f"  mix checksum {9359+651+257+417}")

# freeze + wrap sample
for name in EXPECTED_SHEETS:
    sh = wb0[name]
    if sh.freeze_panes != "A5":
        issues.append(f"{name} freeze={sh.freeze_panes}")
    # notes col wrap
    # skip

print("=== issues", len(issues), "===")
for i in issues:
    print(" ", i)
if issues:
    raise SystemExit(1)
print("VERIFY OK")
