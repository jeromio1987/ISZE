# Rebuild IML-facing PMS study — clear N75 / total parc / IML compare / all DBs
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
import pandas as pd
import json

base = Path(r"C:\Users\jerom\Desktop\claude\ISZE\09_meetings\Preparation Meetings\IML_ITUK_PMS")
out = base / "mp" / "SHARE" / "ISZE_ITUK_UTI_PMS_Study_CY2025_for_IML.xlsx"
out_alt = base / "mp" / "SHARE" / "ISZE_ITUK_UTI_PMS_Study_CY2025_for_IML_CLEAR.xlsx"

wb = Workbook()

navy = PatternFill("solid", fgColor="1A2B4A")
paper = PatternFill("solid", fgColor="F7F4EE")
green = PatternFill("solid", fgColor="E6F5EF")
redf = PatternFill("solid", fgColor="F8E4E3")
yellow = PatternFill("solid", fgColor="FFF2CC")
grey = PatternFill("solid", fgColor="E7E6E6")
amber = PatternFill("solid", fgColor="FCE4D6")
hdr = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
title = Font(color="1A2B4A", bold=True, name="Calibri", size=16)
sec = Font(color="1A2B4A", bold=True, name="Calibri", size=12)
body = Font(name="Calibri", size=10)
muted = Font(name="Calibri", size=9, color="6B6256")
warn = Font(name="Calibri", size=10, color="C0552D", bold=True)
thin = Border(
    left=Side(style="thin", color="D9D3C8"),
    right=Side(style="thin", color="D9D3C8"),
    top=Side(style="thin", color="D9D3C8"),
    bottom=Side(style="thin", color="D9D3C8"),
)
pct, num, eur, dec = "0.0%", "#,##0", "#,##0", "0.00"


def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = navy
        cell.font = hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def border_row(ws, r, cols):
    for c in range(1, cols + 1):
        ws.cell(row=r, column=c).border = thin


# ---------- load all-DB PMS ----------
csv_path = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\pms_capture_by_db_year.csv")
df = pd.read_csv(csv_path)
d = df[(df.year == 2025) & (df.line_type == "total")].copy()

uti_path = Path(
    r"C:\Users\jerom\Desktop\claude\ISZE\03_analytics\Isuzu\pms-pilot\uti_iml_v7_results_2026-07-06.json"
)
uti_iml = json.loads(uti_path.read_text(encoding="utf-8"))["2025"]

# IML schedule freqs (same as Apr ITUK study)
FREQ_FILTER = 1.25  # 50000/40000
FREQ_WEAR = 0.5     # 50000/100000
CLUTCH_SHARE = 0.47

# Canonical ITUK N75 (Interval Model) + full DVLA (Apr study)
N75 = {
    "uio": 9359,
    "db": {
        "Oil filter": (4673, 7.74),
        "Fuel filter": (11021, 13.06),
        "Air filter": (2226, 30.30),
        "Clutch": (243, 80.72),
        "Brake": (2902, 88.63),
    },
    "dlr": {
        "Oil filter": (5404, 7.74),
        "Fuel filter": (11139, 13.06),
        "Air filter": (2289, 30.30),
        "Clutch": (340, 80.72),
        "Brake": (2923, 88.63),
    },
}
FULL = {
    "uio": 22389,
    "db": {
        "Oil filter": (9818, 7.74),
        "Fuel filter": (18228, 13.06),
        "Air filter": (3692, 30.30),
        "Clutch": (411, 80.72),
        "Brake": (3953, 88.63),
    },
    "dlr_oil_qty": 11353,
    "dlr_oil_price": 8.028,
}

# ========== 00 READ ME ==========
ws = wb.active
ws.title = "00_READ_ME_FIRST"
ws["A1"] = "HOW TO READ THIS WORKBOOK — read this page first"
ws["A1"].font = title
ws.merge_cells("A1:B1")

explain = [
    (3, "What this file is",
     "PMS progress pack for the IML / ITUK meeting. Three different layers — do not mix them."),
    (5, "LAYER 1 — ITUK N75 PILOT (sheet 01)",
     "This is the pilot study. UIO = 9,359 N75 vehicles only. Distributor (DB) AND dealer (DLR). "
     "Same service schedule as the Apr 2026 IML/ITUK analysis (50,000 km/yr · oil/fuel/air every 40,000 km)."),
    (7, "LAYER 2 — ITUK TOTAL CAR PARK (sheet 02)",
     "UIO = 22,389 = FULL DVLA registration (all models: N35 + N75 + other). "
     "This is NOT the pilot parc. It is the network ruler / Apr 2026 business-case total. "
     "Oil cover 35.1% · gap ~EUR 2.04M. Do NOT add this gap to the N75 gap."),
    (9, "LAYER 3 — COMPARISON with IML (sheet 03)",
     "Side-by-side: IML Mar kick-off table (20k km/yr · oil 20k) vs GIDSS workshop schedule (Rene) "
     "vs Apr IML/ITUK study (50k/40k) vs N75 cut. Same oil actuals → different % if assumptions change."),
    (11, "LAYER 4 — ALL DISTRIBUTORS with PMS data (sheet 04)",
     "Every DB where we have CY2025 PMS capture rows. Covers are REBASED to the same IML schedule "
     "(UIO × 1.25 for filters) so they can sit next to ITUK. "
     "ITUK on that sheet uses the Apr study numbers (35.1%), not the PRISM internal lane."),
    (13, "LAYER 5 — UTI detail on IML schedule (sheet 05)",
     "UTI only is fully rebased category-by-category like ITUK. Other DBs: oil/fuel/air/brake summary on sheet 04."),
    (15, "Yellow cells",
     "Editable inputs. Everything labelled Cover / Expected / Gap is a FORMULA."),
    (17, "External language",
     'Say "IML service schedule" and "Apr 2026 PMS analysis". No internal nicknames. No personal names on slides.'),
]
for row, head, text in explain:
    ws.cell(row=row, column=1, value=head).font = sec
    ws.cell(row=row + 1, column=1, value=text).font = body
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=2)
    ws.row_dimensions[row + 1].height = 48

ws["A19"] = "Sheet map"
ws["A19"].font = sec
for i, line in enumerate(
    [
        "01_ITUK_N75_STUDY          ← OPEN THIS FIRST (pilot)",
        "02_ITUK_TOTAL_CAR_PARK     ← 22,389 full DVLA (context)",
        "03_COMPARE_IML_vs_STUDY    ← kick-off / GIDSS / Apr / N75",
        "04_ALL_DBs_PMS_CY2025      ← all DBs with PMS data",
        "05_UTI_on_IML_schedule     ← UTI detail",
        "06_Assumptions             ← yellow inputs (mileage, intervals)",
    ],
    20,
):
    ws.cell(row=i, column=1, value=line).font = body

autosize(ws, [92, 20])

# ========== 06 Assumptions first (refs) ==========
ws = wb.create_sheet("06_Assumptions")
ws["A1"] = "ASSUMPTIONS (edit yellow) — used by sheets 01–05"
ws["A1"].font = title
ws["A3"] = "Parameter"
ws["B3"] = "Value"
ws["C3"] = "Meaning"
style_header_row(ws, 3, 3)
params = [
    ("Annual mileage km/yr", 50000, "Planning assumption (Apr study). Not a measured fleet average."),
    ("Oil interval km", 40000, "GIDSS NPR75 MY24 + Apr study"),
    ("Fuel interval km", 40000, "GIDSS NPR75 MY24"),
    ("Air interval km", 40000, "Planning (GIDSS MY24 air = indicator)"),
    ("Clutch interval km", 100000, "Wear assumption (disc/kit)"),
    ("Brake interval km", 100000, "Wear assumption"),
    ("Clutch MT share", 0.47, "Convention — not VIN-verified"),
    ("UIO N75 pilot", 9359, "N75 both engine gens — THE PILOT PARC"),
    ("UIO total car park", 22389, "Full DVLA all models — NOT the pilot"),
    ("UIO UTI", 10336, "UTI parc"),
]
for i, (a, b, c) in enumerate(params, 4):
    ws.cell(row=i, column=1, value=a)
    cell = ws.cell(row=i, column=2, value=b)
    cell.fill = yellow
    cell.border = thin
    cell.number_format = pct if "share" in a else num
    ws.cell(row=i, column=3, value=c).font = muted

ws["A15"] = "Derived freq = mileage / interval"
ws["A15"].font = sec
ws["A16"] = "Oil/Fuel/Air freq"
ws["B16"] = "=B4/B5"
ws["B16"].number_format = "0.00"
ws["A17"] = "Clutch/Brake freq"
ws["B17"] = "=B4/B8"
ws["B17"].number_format = "0.00"
ws["A18"] = "Clutch UIO N75"
ws["B18"] = "=B11*B10"
ws["B18"].number_format = num
ws["A19"] = "Clutch UIO total car park"
ws["B19"] = "=B12*B10"
ws["B19"].number_format = num
ws["A20"] = "Clutch UIO UTI"
ws["B20"] = "=B13*B10"
ws["B20"].number_format = num

ws["A22"] = "Kick-off reference (grey — comparison only)"
ws["A22"].font = sec
ws["A23"] = "Kick-off mileage example"
ws["B23"] = 20000
ws["B23"].fill = grey
ws["A24"] = "Kick-off oil interval"
ws["B24"] = 20000
ws["B24"].fill = grey
autosize(ws, [28, 14, 70])

# helper to write a cover block
def write_cover_block(ws, start_row, title_txt, uio_ref, clutch_uio_ref, cats_dict, note):
    ws.cell(row=start_row, column=1, value=title_txt).font = sec
    ws.cell(row=start_row + 1, column=1, value=note).font = muted
    headers = [
        "Category", "UIO", "Freq/yr", "Expected", "Actual qty", "Cover %",
        "Avg €", "Actual €", "Expected €", "Gap €",
    ]
    hr = start_row + 3
    for c, h in enumerate(headers, 1):
        ws.cell(row=hr, column=c, value=h)
    style_header_row(ws, hr, 10)

    order = ["Oil filter", "Fuel filter", "Air filter", "Clutch", "Brake"]
    freq_map = {
        "Oil filter": "='06_Assumptions'!B16",
        "Fuel filter": "='06_Assumptions'!B16",
        "Air filter": "='06_Assumptions'!B16",
        "Clutch": "='06_Assumptions'!B17",
        "Brake": "='06_Assumptions'!B17",
    }
    for i, cat in enumerate(order):
        r = hr + 1 + i
        act, price = cats_dict[cat]
        uio = clutch_uio_ref if cat == "Clutch" else uio_ref
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=uio).number_format = num
        ws.cell(row=r, column=3, value=freq_map[cat]).number_format = "0.00"
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = num
        a = ws.cell(row=r, column=5, value=act)
        a.fill = yellow
        a.number_format = num
        ws.cell(row=r, column=6, value=f"=IF(D{r}=0,0,E{r}/D{r})").number_format = pct
        p = ws.cell(row=r, column=7, value=price)
        p.fill = yellow
        p.number_format = dec
        ws.cell(row=r, column=8, value=f"=E{r}*G{r}").number_format = eur
        ws.cell(row=r, column=9, value=f"=D{r}*G{r}").number_format = eur
        ws.cell(row=r, column=10, value=f"=MAX(0,I{r}-H{r})").number_format = eur
        border_row(ws, r, 10)

    tot = hr + 6
    ws.cell(row=tot, column=1, value="TOTAL 5 cat").font = Font(bold=True)
    ws.cell(row=tot, column=5, value=f"=SUM(E{hr+1}:E{hr+5})").number_format = num
    ws.cell(row=tot, column=8, value=f"=SUM(H{hr+1}:H{hr+5})").number_format = eur
    ws.cell(row=tot, column=9, value=f"=SUM(I{hr+1}:I{hr+5})").number_format = eur
    ws.cell(row=tot, column=10, value=f"=SUM(J{hr+1}:J{hr+5})").number_format = eur
    ws.cell(row=tot, column=6, value=f"=IF(I{tot}=0,0,H{tot}/I{tot})").number_format = pct
    border_row(ws, tot, 10)
    return tot


# ========== 01 ITUK N75 STUDY ==========
ws = wb.create_sheet("01_ITUK_N75_STUDY", 1)
ws["A1"] = "ITUK N75 PILOT STUDY — this is the main pilot page"
ws["A1"].font = title
ws["A1"].fill = green
ws.merge_cells("A1:J1")
ws["A2"] = (
    "UIO = 9,359 = N75 vehicles only (both engine generations). "
    "NOT 22,389. Schedule = Apr 2026 IML/ITUK analysis (50k km/yr ÷ 40k oil). "
    "Phase-one = brake + air. Clutch deferred (VIN)."
)
ws["A2"].font = body
ws.merge_cells("A2:J2")
ws.row_dimensions[2].height = 36

ws["A3"] = "DISTRIBUTOR (ISZE → ITUK) — N75 cut"
end_db = write_cover_block(
    ws,
    4,
    "",
    "='06_Assumptions'!B11",
    "='06_Assumptions'!B18",
    N75["db"],
    "Actuals = Interval Model ASC-DB N-Series (share-scaled). Yellow = editable.",
)

ws.cell(row=end_db + 2, column=1, value="DEALER SELL-OUT (ITUK → dealers) — N75 cut").font = sec
end_dlr = write_cover_block(
    ws,
    end_db + 2,
    "",
    "='06_Assumptions'!B11",
    "='06_Assumptions'!B18",
    N75["dlr"],
    "Actuals = bottom-up PN split (exclude N35 + F-Series). Real workshop sell-out.",
)

ws.cell(row=end_dlr + 2, column=1, value="HEADLINE (at default assumptions)").font = sec
ws.cell(row=end_dlr + 3, column=1, value="Oil cover DB N75")
ws.cell(row=end_dlr + 3, column=2, value="=F9").number_format = pct
ws.cell(row=end_dlr + 3, column=2).fill = green
ws.cell(row=end_dlr + 4, column=1, value="Oil cover DLR N75")
ws.cell(row=end_dlr + 4, column=2, value=f"=F{end_db + 6}").number_format = pct
ws.cell(row=end_dlr + 4, column=2).fill = green
ws.cell(row=end_dlr + 5, column=1, value="5-cat gap DB N75 €")
ws.cell(row=end_dlr + 5, column=2, value="=J14").number_format = eur
ws.cell(row=end_dlr + 5, column=2).fill = redf
ws.cell(row=end_dlr + 6, column=1, value="Do not add N75 gap to the €2.04M total-car-park gap — N75 is a subset.")
ws.cell(row=end_dlr + 6, column=1).font = warn
ws.merge_cells(start_row=end_dlr + 6, start_column=1, end_row=end_dlr + 6, end_column=6)

autosize(ws, [18, 12, 10, 12, 12, 10, 10, 12, 12, 12])

# Fix: write_cover_block used blank title and start_row+3 for header — verify row numbers
# Looking at write_cover_block: start_row has empty title at start_row, note at start_row+1, header at start_row+3
# For first block start_row=4: header=7, data=8-12, total=13
# Wait I had title_txt="" at row 4, note at 5, header at 7, cats 8-12, total 13
# Headlines reference F9 and J14 — need to fix to actual oil row and total row.

# Recalculate: start_row=4 → hr=7, oil=8, total=13
# Second block start = 13+2=15 → hr=18, oil=19, total=24
# Fix headline formulas after save by reading — or rewrite sheet 01 more carefully without the helper bug.

# I'll rebuild sheet 01 explicitly without helper confusion.
wb.remove(ws)
ws = wb.create_sheet("01_ITUK_N75_STUDY", 1)
ws["A1"] = "ITUK N75 PILOT STUDY — open this first"
ws["A1"].font = title
ws["A1"].fill = green
ws.merge_cells("A1:J1")
ws["A2"] = (
    "PILOT PARC UIO = 9,359 (N75 only). "
    "22,389 is the TOTAL car park on sheet 02 — different layer. "
    "Schedule: 50,000 km/yr · filters every 40,000 km (Apr IML/ITUK study = GIDSS NPR75 oil/fuel)."
)
ws["A2"].font = body
ws.merge_cells("A2:J2")
ws.row_dimensions[2].height = 40

ws["A4"] = "A. DISTRIBUTOR LEVEL (ISZE → ITUK) — N75 only"
ws["A4"].font = sec
ws["A5"] = "Source: Interval Model ASC-DB N-Series. Yellow = inputs."
ws["A5"].font = muted

headers = ["Category", "UIO", "Freq/yr", "Expected", "Actual qty", "Cover %", "Avg €", "Actual €", "Expected €", "Gap €"]
for c, h in enumerate(headers, 1):
    ws.cell(row=6, column=c, value=h)
style_header_row(ws, 6, 10)

n75_db_rows = [
    ("Oil filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 4673, 7.74),
    ("Fuel filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 11021, 13.06),
    ("Air filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 2226, 30.30),
    ("Clutch", "='06_Assumptions'!B18", "='06_Assumptions'!B17", 243, 80.72),
    ("Brake", "='06_Assumptions'!B11", "='06_Assumptions'!B17", 2902, 88.63),
]
for i, (cat, uio, freq, act, price) in enumerate(n75_db_rows, 7):
    ws.cell(row=i, column=1, value=cat)
    ws.cell(row=i, column=2, value=uio).number_format = num
    ws.cell(row=i, column=3, value=freq).number_format = "0.00"
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").number_format = num
    a = ws.cell(row=i, column=5, value=act)
    a.fill = yellow
    a.number_format = num
    ws.cell(row=i, column=6, value=f"=IF(D{i}=0,0,E{i}/D{i})").number_format = pct
    p = ws.cell(row=i, column=7, value=price)
    p.fill = yellow
    p.number_format = dec
    ws.cell(row=i, column=8, value=f"=E{i}*G{i}").number_format = eur
    ws.cell(row=i, column=9, value=f"=D{i}*G{i}").number_format = eur
    ws.cell(row=i, column=10, value=f"=MAX(0,I{i}-H{i})").number_format = eur
    border_row(ws, i, 10)
ws["A12"] = "TOTAL N75 DB"
ws["A12"].font = Font(bold=True)
ws["E12"] = "=SUM(E7:E11)"
ws["H12"] = "=SUM(H7:H11)"
ws["I12"] = "=SUM(I7:I11)"
ws["J12"] = "=SUM(J7:J11)"
ws["F12"] = "=IF(I12=0,0,H12/I12)"
for col in "EFHIJ":
    ws[f"{col}12"].number_format = pct if col == "F" else (eur if col in "HIJ" else num)
border_row(ws, 12, 10)

ws["A14"] = "B. DEALER LEVEL (ITUK → dealers) — N75 only"
ws["A14"].font = sec
ws["A15"] = "Source: Interval Model DB-DLR N-Series (bottom-up PN). Yellow = inputs."
ws["A15"].font = muted
for c, h in enumerate(headers, 1):
    ws.cell(row=16, column=c, value=h)
style_header_row(ws, 16, 10)

n75_dlr_rows = [
    ("Oil filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 5404, 7.74),
    ("Fuel filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 11139, 13.06),
    ("Air filter", "='06_Assumptions'!B11", "='06_Assumptions'!B16", 2289, 30.30),
    ("Clutch", "='06_Assumptions'!B18", "='06_Assumptions'!B17", 340, 80.72),
    ("Brake", "='06_Assumptions'!B11", "='06_Assumptions'!B17", 2923, 88.63),
]
for i, (cat, uio, freq, act, price) in enumerate(n75_dlr_rows, 17):
    ws.cell(row=i, column=1, value=cat)
    ws.cell(row=i, column=2, value=uio).number_format = num
    ws.cell(row=i, column=3, value=freq).number_format = "0.00"
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").number_format = num
    a = ws.cell(row=i, column=5, value=act)
    a.fill = yellow
    a.number_format = num
    ws.cell(row=i, column=6, value=f"=IF(D{i}=0,0,E{i}/D{i})").number_format = pct
    p = ws.cell(row=i, column=7, value=price)
    p.fill = yellow
    p.number_format = dec
    ws.cell(row=i, column=8, value=f"=E{i}*G{i}").number_format = eur
    ws.cell(row=i, column=9, value=f"=D{i}*G{i}").number_format = eur
    ws.cell(row=i, column=10, value=f"=MAX(0,I{i}-H{i})").number_format = eur
    border_row(ws, i, 10)
ws["A22"] = "TOTAL N75 DLR"
ws["A22"].font = Font(bold=True)
ws["E22"] = "=SUM(E17:E21)"
ws["H22"] = "=SUM(H17:H21)"
ws["I22"] = "=SUM(I17:I21)"
ws["J22"] = "=SUM(J17:J21)"
ws["F22"] = "=IF(I22=0,0,H22/I22)"
for col in "EFHIJ":
    ws[f"{col}22"].number_format = pct if col == "F" else (eur if col in "HIJ" else num)
border_row(ws, 22, 10)

ws["A24"] = "C. N75 HEADLINES (formulas)"
ws["A24"].font = sec
ws["A25"] = "Oil cover — N75 distributor"
ws["B25"] = "=F7"
ws["B25"].number_format = pct
ws["B25"].fill = green
ws["C25"] = "← meeting pilot oil (DB)"
ws["A26"] = "Oil cover — N75 dealer"
ws["B26"] = "=F17"
ws["B26"].number_format = pct
ws["B26"].fill = green
ws["C26"] = "← meeting pilot oil (DLR)"
ws["A27"] = "5-cat revenue gap — N75 DB €"
ws["B27"] = "=J12"
ws["B27"].number_format = eur
ws["B27"].fill = redf
ws["A28"] = "WARNING"
ws["A28"].font = warn
ws["B28"] = "Do not add N75 € gap to the €2.04M on sheet 02. N75 is inside the total car park."
ws["B28"].font = warn
ws.merge_cells("B28:J28")
autosize(ws, [32, 12, 12, 12, 12, 10, 10, 12, 12, 12])

# ========== 02 TOTAL CAR PARK ==========
ws = wb.create_sheet("02_ITUK_TOTAL_CAR_PARK", 2)
ws["A1"] = "ITUK TOTAL CAR PARK — UIO 22,389 (full DVLA, all models)"
ws["A1"].font = title
ws["A1"].fill = amber
ws.merge_cells("A1:J1")
ws["A2"] = (
    "This is NOT the N75 pilot. This is the Apr 2026 IML/ITUK business-case total "
    "(N35 + N75 + other). Use for network context and the €2.04M gap. Pilot numbers = sheet 01."
)
ws["A2"].font = body
ws.merge_cells("A2:J2")
ws.row_dimensions[2].height = 40

ws["A4"] = "DISTRIBUTOR (ISZE → ITUK) — full DVLA"
ws["A4"].font = sec
for c, h in enumerate(headers, 1):
    ws.cell(row=5, column=c, value=h)
style_header_row(ws, 5, 10)

full_rows = [
    ("Oil filter", "='06_Assumptions'!B12", "='06_Assumptions'!B16", 9818, 7.74),
    ("Fuel filter", "='06_Assumptions'!B12", "='06_Assumptions'!B16", 18228, 13.06),
    ("Air filter", "='06_Assumptions'!B12", "='06_Assumptions'!B16", 3692, 30.30),
    ("Clutch", "='06_Assumptions'!B19", "='06_Assumptions'!B17", 411, 80.72),
    ("Brake", "='06_Assumptions'!B12", "='06_Assumptions'!B17", 3953, 88.63),
]
for i, (cat, uio, freq, act, price) in enumerate(full_rows, 6):
    ws.cell(row=i, column=1, value=cat)
    ws.cell(row=i, column=2, value=uio).number_format = num
    ws.cell(row=i, column=3, value=freq).number_format = "0.00"
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").number_format = num
    a = ws.cell(row=i, column=5, value=act)
    a.fill = yellow
    a.number_format = num
    ws.cell(row=i, column=6, value=f"=IF(D{i}=0,0,E{i}/D{i})").number_format = pct
    p = ws.cell(row=i, column=7, value=price)
    p.fill = yellow
    p.number_format = dec
    ws.cell(row=i, column=8, value=f"=E{i}*G{i}").number_format = eur
    ws.cell(row=i, column=9, value=f"=D{i}*G{i}").number_format = eur
    ws.cell(row=i, column=10, value=f"=MAX(0,I{i}-H{i})").number_format = eur
    border_row(ws, i, 10)
ws["A11"] = "TOTAL full DVLA"
ws["A11"].font = Font(bold=True)
ws["E11"] = "=SUM(E6:E10)"
ws["H11"] = "=SUM(H6:H10)"
ws["I11"] = "=SUM(I6:I10)"
ws["J11"] = "=SUM(J6:J10)"
ws["F11"] = "=IF(I11=0,0,H11/I11)"
for col in "EFHIJ":
    ws[f"{col}11"].number_format = pct if col == "F" else (eur if col in "HIJ" else num)
border_row(ws, 11, 10)

ws["A13"] = "Dealer oil (full parc) for reference"
ws["A13"].font = sec
ws["A14"] = "Oil actual DLR full parc"
ws["B14"] = 11353
ws["B14"].fill = yellow
ws["B14"].number_format = num
ws["A15"] = "Oil cover DLR full parc"
ws["B15"] = "=IF(D6=0,0,B14/D6)"
ws["B15"].number_format = pct
ws["C15"] = "Should read ~40.6% at defaults"

ws["A17"] = "HEADLINES — total car park"
ws["A17"].font = sec
ws["A18"] = "Oil cover DB (full DVLA)"
ws["B18"] = "=F6"
ws["B18"].number_format = pct
ws["B18"].fill = amber
ws["A19"] = "5-cat gap € (full DVLA)"
ws["B19"] = "=J11"
ws["B19"].number_format = eur
ws["B19"].fill = redf
ws["C19"] = "~EUR 2.04M at defaults"
autosize(ws, [32, 14, 14, 12, 12, 10, 10, 12, 12, 12])

# ========== 03 COMPARE ==========
ws = wb.create_sheet("03_COMPARE_IML_vs_STUDY", 3)
ws["A1"] = "COMPARISON — IML kick-off vs GIDSS vs Apr study vs N75 pilot"
ws["A1"].font = title
ws.merge_cells("A1:G1")
ws["A2"] = (
    "Same idea: Cover % = Actual ÷ Expected. Expected changes when mileage or interval changes. "
    "This sheet explains why numbers disagree."
)
ws["A2"].font = muted
ws.merge_cells("A2:G2")

for c, h in enumerate(
    ["Lens", "What it is", "UIO", "Mileage", "Oil interval", "Oil cover", "Use in meeting?"],
    1,
):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 7)

rows = [
    (
        "IML Mar kick-off table",
        "FIG-KEY potential example in mail",
        22389,
        20000,
        20000,
        "='06_Assumptions'!B12*('06_Assumptions'!B23/'06_Assumptions'!B24)",
        "Context only — oil interval ≠ current GIDSS NPR75",
    ),
    (
        "GIDSS NPR75 MY24 (Rene)",
        "OEM workshop schedule — oil/fuel R every 40k",
        "—",
        "not stated",
        40000,
        "— (schedule only)",
        "YES — confirms 40k oil/fuel for current N75",
    ),
    (
        "Apr IML/ITUK study — TOTAL car park",
        "Sheet 02 — full DVLA",
        "='06_Assumptions'!B12",
        "='06_Assumptions'!B4",
        "='06_Assumptions'!B5",
        "='02_ITUK_TOTAL_CAR_PARK'!F6",
        "YES — network / €2.04M context",
    ),
    (
        "ITUK N75 PILOT study",
        "Sheet 01 — pilot parc only",
        "='06_Assumptions'!B11",
        "='06_Assumptions'!B4",
        "='06_Assumptions'!B5",
        "='01_ITUK_N75_STUDY'!F7",
        "YES — lead with this for the pilot",
    ),
    (
        "TRAP — mix 50k mileage + 20k oil",
        "Inflates expected → ~17% oil",
        22389,
        50000,
        20000,
        None,
        "NO — never mix with Apr study",
    ),
]
for i, (lens, what, uio, mile, oil_iv, cover, use) in enumerate(rows, 5):
    ws.cell(row=i, column=1, value=lens)
    ws.cell(row=i, column=2, value=what).font = muted
    ws.cell(row=i, column=3, value=uio)
    if isinstance(uio, (int, float)) or (isinstance(uio, str) and uio.startswith("=")):
        ws.cell(row=i, column=3).number_format = num
    ws.cell(row=i, column=4, value=mile)
    if isinstance(mile, (int, float)) or (isinstance(mile, str) and str(mile).startswith("=")):
        ws.cell(row=i, column=4).number_format = num
    ws.cell(row=i, column=5, value=oil_iv)
    if isinstance(oil_iv, (int, float)) or (isinstance(oil_iv, str) and str(oil_iv).startswith("=")):
        ws.cell(row=i, column=5).number_format = num
    if cover is None:
        # trap: actual 9818 / (22389 * 50000/20000)
        ws.cell(row=i, column=6, value="=9818/(22389*50000/20000)")
    else:
        ws.cell(row=i, column=6, value=cover)
    if isinstance(ws.cell(row=i, column=6).value, str) and str(ws.cell(row=i, column=6).value).startswith("="):
        # kick-off row 5 is expected not cover — fix
        pass
    ws.cell(row=i, column=7, value=use)
    border_row(ws, i, 7)
    if "N75 PILOT" in lens:
        for c in range(1, 8):
            ws.cell(row=i, column=c).fill = green
    if "TRAP" in lens:
        for c in range(1, 8):
            ws.cell(row=i, column=c).fill = redf

# Fix row 5: kick-off cover = 9818 / (uio * mile/oil)
ws["F5"] = "=9818/(C5*D5/E5)"
ws["F5"].number_format = pct
ws["F7"].number_format = pct
ws["F8"].number_format = pct
ws["F9"].number_format = pct

ws["A11"] = "Meeting line (copy)"
ws["A11"].font = sec
ws["A12"] = (
    "Lead with N75 pilot (sheet 01): oil ~40% DB / ~46% DLR. "
    "Full DVLA 35.1% / €2.04M is total car park context (sheet 02). "
    "GIDSS confirms 40k oil for current NPR75 — kick-off 20k oil does not."
)
ws["A12"].font = body
ws.merge_cells("A12:G12")
ws.row_dimensions[12].height = 48
autosize(ws, [36, 40, 12, 12, 12, 12, 42])

# ========== 04 ALL DBs ==========
ws = wb.create_sheet("04_ALL_DBs_PMS_CY2025", 4)
ws["A1"] = "ALL DISTRIBUTORS with PMS data — CY2025 — IML schedule rebase"
ws["A1"].font = title
ws.merge_cells("A1:L1")
ws["A2"] = (
    "Every DB in pms_capture_by_db_year.csv for CY2025 (total line). "
    "Cover here = Actual ÷ (UIO × 1.25) for filters — SAME schedule as ITUK Apr study — so DBs are comparable. "
    "ITUK row uses Apr study oil 35.1% (total car park), not the PRISM internal lane. "
    "UTI also has a full IML detail sheet (05)."
)
ws["A2"].font = muted
ws.merge_cells("A2:L2")
ws.row_dimensions[2].height = 52

headers4 = [
    "Distributor",
    "UIO",
    "Oil actual",
    "Oil expected (IML)",
    "Oil cover (IML)",
    "Fuel cover (IML)",
    "Air cover (IML)",
    "Brake pad cover (IML)",
    "Oil gap € (PRISM SP)",
    "Lane note",
]
for c, h in enumerate(headers4, 1):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 10)

# pivot actuals
cats_need = ["oil_filter", "fuel_filter", "air_filter", "brake_pad"]
pivot = {}
for db in sorted(d.distributor_code.unique()):
    sub = d[d.distributor_code == db]
    uio = int(sub["uio_units"].mode().iloc[0]) if len(sub) else 0
    row = {"uio": uio}
    for cat in cats_need:
        r = sub[sub.pms_category == cat]
        if len(r):
            row[cat] = {
                "qty": float(r.iloc[0]["actual_qty"]),
                "gap": float(r.iloc[0]["revenue_gap_eur"]),
                "prism_cap": float(r.iloc[0]["capture_rate"]),
            }
        else:
            row[cat] = {"qty": 0.0, "gap": 0.0, "prism_cap": None}
    pivot[db] = row

r = 5
# Put ITUK first (canonical), then UTI, then others alpha
order_dbs = ["ITUK", "UTI"] + [x for x in sorted(pivot) if x not in ("ITUK", "UTI")]
for db in order_dbs:
    info = pivot[db]
    uio = info["uio"]
    ws.cell(row=r, column=1, value=db)
    ws.cell(row=r, column=2, value=uio).number_format = num

    if db == "ITUK":
        # Use Apr study numbers on total car park for meeting clarity
        ws.cell(row=r, column=2, value=22389).number_format = num
        ws.cell(row=r, column=3, value=9818).number_format = num
        ws.cell(row=r, column=4, value=27986).number_format = num
        ws.cell(row=r, column=5, value=0.351).number_format = pct
        ws.cell(row=r, column=6, value=0.651).number_format = pct
        ws.cell(row=r, column=7, value=0.132).number_format = pct
        ws.cell(row=r, column=8, value=0.353).number_format = pct  # brake family study
        ws.cell(row=r, column=9, value=140602).number_format = eur
        ws.cell(row=r, column=10, value="Apr IML/ITUK study TOTAL car park — see sheets 01–02 for N75")
        for c in range(1, 11):
            ws.cell(row=r, column=c).fill = amber
    elif db == "UTI":
        # IML-rebased from JSON
        u = uti_iml["uio"]
        oc = uti_iml["categories"]
        ws.cell(row=r, column=2, value=u).number_format = num
        ws.cell(row=r, column=3, value=oc["oil_filter"]["qty"]).number_format = num
        ws.cell(row=r, column=4, value=oc["oil_filter"]["exp"]).number_format = num
        ws.cell(row=r, column=5, value=oc["oil_filter"]["capture_pct"] / 100).number_format = pct
        ws.cell(row=r, column=6, value=oc["fuel_filter"]["capture_pct"] / 100).number_format = pct
        ws.cell(row=r, column=7, value=oc["air_filter"]["capture_pct"] / 100).number_format = pct
        ws.cell(row=r, column=8, value=oc["brake_pad"]["capture_pct"] / 100).number_format = pct
        ws.cell(row=r, column=9, value=oc["oil_filter"]["gap"]).number_format = eur
        ws.cell(row=r, column=10, value="IML schedule rebase (same ruler as ITUK) — detail on sheet 05")
        for c in range(1, 11):
            ws.cell(row=r, column=c).fill = green
    else:
        oil_qty = info["oil_filter"]["qty"]
        oil_exp = uio * FREQ_FILTER
        fuel_qty = info["fuel_filter"]["qty"]
        air_qty = info["air_filter"]["qty"]
        brake_qty = info["brake_pad"]["qty"]
        brake_exp = uio * FREQ_WEAR
        ws.cell(row=r, column=3, value=oil_qty).number_format = num
        ws.cell(row=r, column=4, value=oil_exp).number_format = num
        ws.cell(row=r, column=5, value=(oil_qty / oil_exp if oil_exp else 0)).number_format = pct
        ws.cell(row=r, column=6, value=(fuel_qty / oil_exp if oil_exp else 0)).number_format = pct
        ws.cell(row=r, column=7, value=(air_qty / oil_exp if oil_exp else 0)).number_format = pct
        ws.cell(row=r, column=8, value=(brake_qty / brake_exp if brake_exp else 0)).number_format = pct
        ws.cell(row=r, column=9, value=info["oil_filter"]["gap"]).number_format = eur
        ws.cell(
            row=r,
            column=10,
            value="IML-rebased from PRISM actuals (UIO×1.25). Oil € gap still PRISM SP.",
        ).font = muted
    border_row(ws, r, 10)
    r += 1

# Extra ITUK N75 row
ws.cell(row=r, column=1, value="ITUK N75 pilot")
ws.cell(row=r, column=2, value=9359).number_format = num
ws.cell(row=r, column=3, value=4673).number_format = num
ws.cell(row=r, column=4, value=11698.75).number_format = num
ws.cell(row=r, column=5, value=4673 / 11698.75).number_format = pct
ws.cell(row=r, column=6, value=11021 / 11698.75).number_format = pct
ws.cell(row=r, column=7, value=2226 / 11698.75).number_format = pct
ws.cell(row=r, column=8, value=2902 / 4679.5).number_format = pct
ws.cell(row=r, column=9, value=54377).number_format = eur
ws.cell(row=r, column=10, value="PILOT — sheet 01 (subset of ITUK total car park)")
for c in range(1, 11):
    ws.cell(row=r, column=c).fill = green
border_row(ws, r, 10)
r += 2

ws.cell(row=r, column=1, value="How to read ITUK on this sheet").font = sec
ws.cell(row=r + 1, column=1, value=(
    "Row ITUK = total car park 22,389 (Apr study). "
    "Row ITUK N75 pilot = 9,359. "
    "Other DBs = same IML filter frequency so oil/air/fuel % are comparable. "
    "UMI and UTI are separate Israel distributors — never blend."
)).font = body
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=10)
ws.row_dimensions[r + 1].height = 48
autosize(ws, [16, 10, 12, 16, 14, 14, 14, 16, 14, 55])

# ========== 05 UTI ==========
ws = wb.create_sheet("05_UTI_on_IML_schedule", 5)
ws["A1"] = "UTI — same IML service schedule as ITUK (comparable ruler)"
ws["A1"].font = title
ws.merge_cells("A1:H1")
ws["A2"] = (
    f"UIO {uti_iml['uio']} · FY/CY2025 rebased. Not an ITUK pilot contract. "
    "Clutch VIN-tagged. Do not blend with UMI."
)
ws["A2"].font = muted

for c, h in enumerate(
    ["Category", "UIO", "Expected", "Actual qty", "Cover %", "Gap €", "Note"],
    1,
):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 7)

uti_cats = [
    ("Oil filter", "oil_filter", ""),
    ("Fuel filter", "fuel_filter", ""),
    ("Air filter", "air_filter", ""),
    ("Brake pad", "brake_pad", "Main UTI EUR gap family"),
    ("Clutch disc", "clutch_disc", "VIN provisional"),
]
uio_u = uti_iml["uio"]
for i, (lab, key, note) in enumerate(uti_cats, 5):
    oc = uti_iml["categories"][key]
    ws.cell(row=i, column=1, value=lab)
    ws.cell(row=i, column=2, value=uio_u).number_format = num
    ws.cell(row=i, column=3, value=oc["exp"]).number_format = num
    a = ws.cell(row=i, column=4, value=oc["qty"])
    a.fill = yellow
    a.number_format = num
    ws.cell(row=i, column=5, value=oc["capture_pct"] / 100).number_format = pct
    ws.cell(row=i, column=6, value=oc["gap"]).number_format = eur
    ws.cell(row=i, column=7, value=note).font = muted
    border_row(ws, i, 7)

ws["A11"] = "vs ITUK N75 oil (sheet 01)"
ws["B11"] = "='01_ITUK_N75_STUDY'!F7"
ws["B11"].number_format = pct
ws["C11"] = "ITUK N75 DB oil"
ws["A12"] = "vs ITUK total car park oil (sheet 02)"
ws["B12"] = "='02_ITUK_TOTAL_CAR_PARK'!F6"
ws["B12"].number_format = pct
ws["C12"] = "ITUK full DVLA oil"
ws["A13"] = "UTI oil (this sheet)"
ws["B13"] = "=E5"
ws["B13"].number_format = pct
ws["B13"].fill = green
autosize(ws, [18, 12, 12, 12, 10, 12, 28])

# Move assumptions to end already at 06; ensure sheet order
# Reorder: 00, 01, 02, 03, 04, 05, 06
order = [
    "00_READ_ME_FIRST",
    "01_ITUK_N75_STUDY",
    "02_ITUK_TOTAL_CAR_PARK",
    "03_COMPARE_IML_vs_STUDY",
    "04_ALL_DBs_PMS_CY2025",
    "05_UTI_on_IML_schedule",
    "06_Assumptions",
]
for i, name in enumerate(order):
    wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

for name in order:
    if name != "00_READ_ME_FIRST":
        wb[name].freeze_panes = "A5"

try:
    wb.save(out)
    saved = out
except PermissionError:
    wb.save(out_alt)
    saved = out_alt

print("SAVED", saved)
print("sheets", wb.sheetnames)
print("DBs on sheet 04:", order_dbs)
print("size", saved.stat().st_size)
