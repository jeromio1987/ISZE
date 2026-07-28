# -*- coding: utf-8 -*-
"""Build V2 ~10-page deep study HTML + Excel for I Motors + UMI."""
from __future__ import annotations

import json
from pathlib import Path
from string import Template

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
DATA = json.loads((OUT / "study_data.json").read_text(encoding="utf-8"))
HTML = OUT / "ISZE_Imotors_UMI_Deep_Study_V2.html"
XLSX = OUT / "ISZE_Imotors_UMI_Deep_Study_V2.xlsx"

FILL = {
    "IMOTORS": {"avg_pct": 87.5, "source": "PBI SPI customer view · session 2026-07-26"},
    "UMI": {"avg_pct": None, "source": "Not isolated on fill page this session — use PN losers + ORDER gap"},
}


def eur(x):
    v = float(x or 0)
    if abs(v) >= 1_000_000:
        return f"€{v/1_000_000:.2f}M"
    return f"€{v/1000:.0f}K"


def get_dbs():
    imo = next(d for d in DATA["distributors"] if d["key"] == "IMOTORS")
    umi = next(d for d in DATA["distributors"] if d["key"] == "UMI")
    return imo, umi


def peers_gps():
    import csv

    path = ROOT / "01_data" / "00_derived" / "gps_spu_by_distributor.csv"
    rows = list(csv.DictReader(path.open(encoding="utf-8")))
    out = []
    for code in ["PORTUGAL", "UMI", "UTI", "ITUK", "IUK", "BELGIUM", "IRELAND"]:
        for r in rows:
            if r["distributor_code"] == code and r["year"] == "2025":
                out.append(
                    {
                        "code": code,
                        "spu": round(float(r["spend_per_unit_eur"] or 0), 1),
                        "uio": int(float(r["uio_units"] or 0)),
                        "gps": None
                        if r.get("gps_capture_vs_max") in (None, "")
                        else round(100 * float(r["gps_capture_vs_max"]), 1),
                        "sales": round(float(r["total_sales_eur"]), 0),
                        "tl": r.get("traffic_light"),
                    }
                )
    return out


def pn_table(rows):
    trs = []
    for r in rows:
        trs.append(
            f"<tr><td>{r['rank']}</td><td><code>{r['part_number']}</code></td>"
            f"<td>{r['description']}</td><td class='num'>{eur(r['sales_eur'])}</td>"
            f"<td class='num'>{r['share_pct']:.1f}%</td></tr>"
        )
    return "\n".join(trs)


def mover_table(rows, bad=False):
    cls = "bad" if bad else "good"
    trs = []
    for r in rows[:8]:
        trs.append(
            f"<tr><td><code>{r['part_number']}</code></td><td>{r['description']}</td>"
            f"<td class='num'>{eur(r['fy_a_eur'])}</td><td class='num'>{eur(r['fy_b_eur'])}</td>"
            f"<td class='num {cls}'>{eur(r['delta_eur'])}</td></tr>"
        )
    return "\n".join(trs)


def build_html():
    imo, umi = get_dbs()
    pbi = DATA["pbi_h1_2026"]
    drills, order = pbi["drills"], pbi["order"]
    peers = peers_gps()
    imo_pms, umi_pms = imo["pms_summary"][0], umi["pms_summary"][0]
    imo_gps = next(g for g in imo["gps"] if g["year"] == 2025)
    umi_gps = next(g for g in umi["gps"] if g["year"] == 2025)
    imo_fy24, imo_fy25 = imo["fy_totals"]["2024"]["sales_eur"], imo["fy_totals"]["2025"]["sales_eur"]
    umi_fy24, umi_fy25 = umi["fy_totals"]["2024"]["sales_eur"], umi["fy_totals"]["2025"]["sales_eur"]
    imo_yoy = round(100 * (imo_fy25 - imo_fy24) / imo_fy24, 1)
    umi_yoy = round(100 * (umi_fy25 - umi_fy24) / umi_fy24, 1)
    imo_pn = [x for x in imo["top_pn"] if x["fiscal_year"] == 2025][:12]
    umi_pn = [x for x in umi["top_pn"] if x["fiscal_year"] == 2025][:12]

    chart_data = {
        "imoMonthly": imo["monthly"],
        "umiMonthly": umi["monthly"],
        "imoSaty24": [x for x in imo["by_saty"] if x["fiscal_year"] == 2024],
        "imoSaty25": [x for x in imo["by_saty"] if x["fiscal_year"] == 2025],
        "umiSaty24": [x for x in umi["by_saty"] if x["fiscal_year"] == 2024],
        "umiSaty25": [x for x in umi["by_saty"] if x["fiscal_year"] == 2025],
        "imoFy": [imo_fy24, imo_fy25],
        "umiFy": [umi_fy24, umi_fy25],
        "orderSales": {
            "labels": ["I Motors", "UMI"],
            "sales": [drills["IMOTORS"]["sales_eur"], drills["UMI"]["sales_eur"]],
            "demand": [order["customers"]["IMOTORS"]["demand_eur"], order["customers"]["UMI"]["demand_eur"]],
        },
        "peers": peers,
        "imoPms": sorted(imo["pms_categories_total"], key=lambda r: -r["revenue_gap_eur"])[:10],
        "umiPms": sorted(umi["pms_categories_total"], key=lambda r: -r["revenue_gap_eur"])[:10],
        "imoCatH1": drills["IMOTORS"].get("category", []),
        "umiCatH1": drills["UMI"].get("category", []),
        "imoSrcH1": drills["IMOTORS"].get("source", []),
        "umiSrcH1": drills["UMI"].get("source", []),
    }

    # Body uses Template ($vars). Chart script uses @@DATA@@ replace (no brace wars).
    body = Template("""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>I Motors &amp; UMI — Deep Sales &amp; Order Study V2</title>
<style>
*{box-sizing:border-box}
html,body{margin:0;padding:0;-webkit-print-color-adjust:exact;print-color-adjust:exact}
body{font-family:"Segoe UI",Helvetica,Arial,sans-serif;color:#222;background:#f7f4ee;font-size:15px;line-height:1.5}
h1,h2,h3{font-family:Georgia,"Times New Roman",serif;color:#1a2b4a;margin:0}
.page{max-width:980px;margin:0 auto;padding:2.2rem 1.4rem 1rem;min-height:80vh;border-bottom:1px solid #d9d3c8}
@media print{.page{break-after:page;min-height:0;border:0}.no-print{display:none!important}body{background:#fff}}
.cover{background:linear-gradient(155deg,#15233c 0%,#203759 55%,#2b4a78 100%);color:#eef1f6;padding:3.5rem 1.6rem 3rem}
.cover .inner{max-width:980px;margin:0 auto}
.cover h1{color:#fff;font-size:2.5rem;line-height:1.12}
.cover .rule{width:4.2rem;height:3px;background:#c0552d;margin:1.1rem 0 1rem}
.chip{display:inline-block;background:rgba(255,255,255,.12);border:1px solid rgba(255,255,255,.22);padding:.2rem .65rem;border-radius:4px;font-size:.78rem;margin:.15rem .25rem .15rem 0}
.nav{position:sticky;top:0;z-index:20;background:#f7f4ee;border-bottom:1px solid #d9d3c8;padding:.45rem 1rem;display:flex;flex-wrap:wrap;gap:.35rem}
.nav a{font-size:.75rem;color:#1a2b4a;text-decoration:none;padding:.25rem .45rem;border:1px solid #d9d3c8;background:#fff}
h2.sec{font-size:1.45rem;margin:0 0 .7rem;padding-bottom:.35rem;border-bottom:2px solid #1a2b4a}
h2.sec .no{color:#c0552d}
h3{font-size:1.05rem;margin:1.1rem 0 .4rem}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:.7rem;margin:1rem 0}
.stat{background:#fff;border:1px solid #d9d3c8;padding:.85rem 1rem}
.stat .lbl{font-size:.7rem;text-transform:uppercase;letter-spacing:.04em;color:#6b6256}
.stat .val{font-size:1.3rem;font-weight:700;color:#1a2b4a;margin-top:.1rem}
.stat .sub{font-size:.82rem;color:#6b6256}
.good{color:#2e7d54!important}.bad{color:#b23b3b!important}.amber{color:#a67c00!important}
.tl{display:inline-block;width:.7rem;height:.7rem;border-radius:50%;margin-right:.35rem;vertical-align:middle}
.tl.g{background:#1d9e75}.tl.a{background:#eda100}.tl.r{background:#e24b4a}
.callout{border-left:3px solid #c0552d;background:#fff;padding:.85rem 1rem;margin:1rem 0}
.callout.warn{border-left-color:#eda100}
.chart{position:relative;width:100%;height:320px;background:#fff;border:1px solid #d9d3c8;padding:.5rem;margin:.8rem 0}
.chart.tall{height:420px}
.two{display:grid;grid-template-columns:1fr 1fr;gap:1rem}
@media(max-width:820px){.two{grid-template-columns:1fr}}
table{width:100%;border-collapse:collapse;font-size:.86rem;background:#fff;margin:.7rem 0}
th{background:#1a2b4a;color:#fff;text-align:left;padding:.5rem .6rem;font-size:.72rem;text-transform:uppercase}
td{padding:.45rem .6rem;border-bottom:1px solid #d9d3c8}
tbody tr:nth-child(even){background:#faf8f4}
td.num,.num{text-align:right;font-variant-numeric:tabular-nums}
code{font-family:Consolas,monospace;font-size:.84em}
.muted{color:#6b6256;font-size:.9rem}
.ask{background:#fff;border:1px solid #d9d3c8;padding:.7rem 1rem;margin:.4rem 0}
.footer{max-width:980px;margin:0 auto;padding:1rem 1.4rem 2.5rem;font-size:.75rem;color:#6b6256}
.heat{display:grid;grid-template-columns:1fr 1fr;gap:.45rem}
.heat .c{background:#fff;border:1px solid #d9d3c8;padding:.55rem .65rem;font-size:.8rem}
.heat .c .p{font-weight:700;font-size:1.05rem}
</style>
</head>
<body>
<div class="cover"><div class="inner">
<span class="chip">ISZE Parts · Meeting prep V2</span>
<span class="chip">Deep study · 10 pages</span>
<h1>I Motors &amp; UMI<br/>Sales, Order &amp; Capture</h1>
<div class="rule"></div>
<p style="max-width:40rem;opacity:.93">Two distributor stories — Portugal growth vs Israel LCV realisation gap. Hierarchy from fiscal sales through SaTy, GPS, PMS basket, and part number.</p>
<p style="color:#c9d2e0;font-size:.9rem">FY = Apr–Mar · Billing through $billing_max · PBI H1 CY2026 refreshed 2026-07-21 · Generated 2026-07-26</p>
</div></div>
<nav class="nav no-print">
<a href="#p1">01 Scope</a><a href="#p2">02 Exec</a><a href="#p3">03 Trend</a><a href="#p4">04 Order</a>
<a href="#p5">05 SaTy</a><a href="#p6">06 GPS</a><a href="#p7">07 PMS</a><a href="#p8">08 Family</a>
<a href="#p9">09 Fill</a><a href="#p10">10 Asks</a>
</nav>

<section class="page" id="p1">
<h2 class="sec"><span class="no">01</span> Scope &amp; periods</h2>
<div class="callout"><strong>Who:</strong> I Motors (Portugal, SAP 92/519) and UMI (Israel LCV, SAP 1797).<br/>
<strong>IDS Israel:</strong> covers the UMI dealer network only. Other Israel distributors, if shown, are separate context rows — do not combine with IDS Israel / PMS / GPS scores.</div>
<table>
<thead><tr><th>Layer</th><th>Period</th><th>Source</th></tr></thead>
<tbody>
<tr><td>Fiscal sales, SaTy, PN</td><td>FY2024 &amp; FY2025 (Apr–Mar)</td><td>ISZE billing</td></tr>
<tr><td>Sales tendency</td><td>24 fiscal months</td><td>ISZE billing</td></tr>
<tr><td>ORDER vs SALES</td><td>H1 CY2026</td><td>Power BI ORDER / SALES</td></tr>
<tr><td>GPS / PMS</td><td>CY2025</td><td>PRISM</td></tr>
<tr><td>Fill rate</td><td>PBI SPI</td><td>I Motors captured; UMI not isolated</td></tr>
</tbody></table>
<p class="muted">Primary lens = <strong>total</strong> sales / total PMS capture. Genuine is secondary.</p>
</section>

<section class="page" id="p2">
<h2 class="sec"><span class="no">02</span> Executive contrast</h2>
<div class="grid">
<div class="stat"><div class="lbl"><span class="tl g"></span>I Motors FY</div><div class="val good">+$imo_yoy%</div><div class="sub">$imo_fy24 → $imo_fy25</div></div>
<div class="stat"><div class="lbl"><span class="tl r"></span>UMI FY</div><div class="val bad">$umi_yoy%</div><div class="sub">$umi_fy24 → $umi_fy25</div></div>
<div class="stat"><div class="lbl"><span class="tl g"></span>I Motors H1 sales</div><div class="val good">+$imo_h1_yoy%</div><div class="sub">$imo_h1_sales</div></div>
<div class="stat"><div class="lbl"><span class="tl r"></span>UMI H1 sales</div><div class="val bad">$umi_h1_yoy%</div><div class="sub">$umi_h1_sales</div></div>
<div class="stat"><div class="lbl"><span class="tl a"></span>UMI H1 demand</div><div class="val">−0.4%</div><div class="sub">$umi_h1_demand · almost flat</div></div>
<div class="stat"><div class="lbl"><span class="tl r"></span>UMI gap</div><div class="val bad">$umi_gap</div><div class="sub">Demand ahead of billing</div></div>
</div>
<div class="chart"><canvas id="cFy" role="img" aria-label="FY compare"></canvas></div>
<div class="two">
<div class="callout"><strong>I Motors:</strong> Growing, diversified (top-10 $imo_top10% of FY25). Sales≈demand. PMS total $imo_pms%.</div>
<div class="callout warn"><strong>UMI:</strong> Invoice down, demand holds. Top-10 $umi_top10%. PMS total $umi_pms% · gap $umi_pms_gap.</div>
</div>
</section>

<section class="page" id="p3">
<h2 class="sec"><span class="no">03</span> Sales tendency — 24 months</h2>
<p class="muted">FY2024 months then FY2025 months.</p>
<h3>I Motors — increasing</h3>
<div class="chart"><canvas id="cImoM" role="img" aria-label="I Motors monthly"></canvas></div>
<h3>UMI — dropping / soft</h3>
<div class="chart"><canvas id="cUmiM" role="img" aria-label="UMI monthly"></canvas></div>
</section>

<section class="page" id="p4">
<h2 class="sec"><span class="no">04</span> ORDER vs SALES — lost realisation</h2>
<div class="chart"><canvas id="cOrd" role="img" aria-label="Order vs sales"></canvas></div>
<table>
<thead><tr><th>DB</th><th class="num">Sales</th><th class="num">YoY</th><th class="num">Demand</th><th class="num">YoY</th><th class="num">Gap</th></tr></thead>
<tbody>
<tr><td>I Motors</td><td class="num">$imo_h1_sales</td><td class="num good">+$imo_h1_yoy%</td><td class="num">$imo_h1_demand</td><td class="num good">+$imo_dem_yoy%</td><td class="num">$imo_gap</td></tr>
<tr><td>UMI</td><td class="num">$umi_h1_sales</td><td class="num bad">$umi_h1_yoy%</td><td class="num">$umi_h1_demand</td><td class="num">−0.4%</td><td class="num bad">$umi_gap</td></tr>
<tr><td>UTI (context)</td><td class="num">$uti_h1_sales</td><td class="num good">+15%</td><td class="num">$uti_h1_demand</td><td class="num good">+62%</td><td class="num muted">—</td></tr>
</tbody></table>
<div class="callout warn">Israel market demand +26% is largely driven by another Israel distributor (+62% demand). UMI −19% sales is UMI-specific — not “Israel soft”.</div>
<div class="two">
<div><h3>UMI H1 source</h3><div class="chart"><canvas id="cUmiSrc"></canvas></div></div>
<div><h3>I Motors H1 source</h3><div class="chart"><canvas id="cImoSrc"></canvas></div></div>
</div>
</section>

<section class="page" id="p5">
<h2 class="sec"><span class="no">05</span> SaTy evolution</h2>
<p class="muted">SOR = planned. RUSH/VOR = pressure. Share of €.</p>
<div class="two">
<div><h3>I Motors</h3><div class="chart"><canvas id="cImoSaty"></canvas></div></div>
<div><h3>UMI</h3><div class="chart"><canvas id="cUmiSaty"></canvas></div></div>
</div>
<div class="callout">I Motors SOR ~82% FY25. UMI SOR ~86% with material VOR; almost no RUSH coding.</div>
</section>

<section class="page" id="p6">
<h2 class="sec"><span class="no">06</span> GPS — SPU &amp; peers</h2>
<div class="grid">
<div class="stat"><div class="lbl">I Motors CY2025 SPU</div><div class="val">€$imo_spu</div><div class="sub">UIO $imo_uio · GPS $imo_gps% · $imo_tl</div></div>
<div class="stat"><div class="lbl">UMI CY2025 SPU</div><div class="val">€$umi_spu</div><div class="sub">UIO $umi_uio · GPS $umi_gps% · $umi_tl</div></div>
</div>
<div class="chart tall"><canvas id="cPeers"></canvas></div>
<div class="callout warn">UMI GPS green ≠ healthy basket — lubricants inflate wallet. Treat other Israel distributors as separate rows.</div>
</section>

<section class="page" id="p7">
<h2 class="sec"><span class="no">07</span> PMS — total capture first</h2>
<div class="grid">
<div class="stat"><div class="lbl">I Motors total</div><div class="val">$imo_pms%</div><div class="sub">Gap $imo_pms_gap · genuine $imo_gen%</div></div>
<div class="stat"><div class="lbl">UMI total</div><div class="val bad">$umi_pms%</div><div class="sub">Gap $umi_pms_gap · oil filter $umi_oil%</div></div>
</div>
<h3>I Motors gaps</h3><div class="chart tall"><canvas id="cImoPms"></canvas></div>
<h3>UMI gaps</h3><div class="chart tall"><canvas id="cUmiPms"></canvas></div>
<div class="heat">
<div class="c"><div class="muted">I Motors top gaps</div><div class="p">$imo_gaps</div></div>
<div class="c"><div class="muted">UMI top gaps</div><div class="p">$umi_gaps</div></div>
</div>
</section>

<section class="page" id="p8">
<h2 class="sec"><span class="no">08</span> Strong vs weak — H1 categories</h2>
<div class="two">
<div><h3>I Motors H1</h3><div class="chart tall"><canvas id="cImoCat"></canvas></div>
<p class="muted">Emission +87% · Engine +77% · Competitive −27%.</p></div>
<div><h3>UMI H1</h3><div class="chart tall"><canvas id="cUmiCat"></canvas></div>
<p class="muted">Others −17% · Emission −29% · Engine −62% · LOCAL −55%.</p></div>
</div>
</section>

<section class="page" id="p9">
<h2 class="sec"><span class="no">09</span> Fill &amp; where we lose sales</h2>
<div class="grid">
<div class="stat"><div class="lbl">I Motors fill</div><div class="val">$imo_fill%</div><div class="sub">$imo_fill_src</div></div>
<div class="stat"><div class="lbl">UMI fill</div><div class="val amber">n/a*</div><div class="sub">$umi_fill_src</div></div>
</div>
<div class="two">
<div><h3>I Motors PN losers</h3>
<table><thead><tr><th>PN</th><th>Desc</th><th class="num">FY24</th><th class="num">FY25</th><th class="num">Δ</th></tr></thead>
<tbody>$imo_losers</tbody></table></div>
<div><h3>UMI PN losers</h3>
<table><thead><tr><th>PN</th><th>Desc</th><th class="num">FY24</th><th class="num">FY25</th><th class="num">Δ</th></tr></thead>
<tbody>$umi_losers</tbody></table></div>
</div>
<div class="callout">Probe: (1) UMI ORDER→SALES gap $umi_gap (2) LOCAL/Shell −55% (3) clutch/brake PMS (4) I Motors competitive −27%.</div>
</section>

<section class="page" id="p10">
<h2 class="sec"><span class="no">10</span> Asks &amp; appendix</h2>
<h3>I Motors</h3>
<div class="ask"><strong>1.</strong> Competitive-parts H1 −27%: price, stock, or channel?</div>
<div class="ask"><strong>2.</strong> Dealer view on brake shoe / clutch / air filter.</div>
<div class="ask"><strong>3.</strong> Is ~11% RUSH intentional or planning miss?</div>
<h3>UMI</h3>
<div class="ask"><strong>1.</strong> What sits between ORDER and SALES in H1?</div>
<div class="ask"><strong>2.</strong> Rimula / LOCAL outlook for rest of FY.</div>
<div class="ask"><strong>3.</strong> Clutch &amp; brake — lead with total capture ($umi_pms%), not genuine-only.</div>
<div class="ask"><strong>4.</strong> IDS Israel covers the UMI network only — do not combine other Israel distributors with IDS Israel / PMS / GPS for this discussion.</div>
<h3>Appendix — Top PN FY2025</h3>
<div class="two">
<div><h3>I Motors</h3><table><thead><tr><th>#</th><th>PN</th><th>Desc</th><th class="num">€</th><th class="num">%</th></tr></thead><tbody>$imo_pn</tbody></table></div>
<div><h3>UMI</h3><table><thead><tr><th>#</th><th>PN</th><th>Desc</th><th class="num">€</th><th class="num">%</th></tr></thead><tbody>$umi_pn</tbody></table></div>
</div>
</section>
<div class="footer">Sources: ISZE billing FY2024–FY2025 · Power BI SALES/ORDER H1 CY2026 · PRISM GPS/PMS CY2025. Total capture primary. IDS Israel = UMI network only. Excel: ISZE_Imotors_UMI_Deep_Study_V2.xlsx</div>
<script src="assets/chart.umd.min.js"></script>
<script>
const D = @@DATA@@;
const ink='#52514e', grid='#e1e0d9', navy='#2a78d6', terra='#c0552d', good='#1d9e75', bad='#e24b4a', grey='#888780';
new Chart(cFy,{type:'bar',data:{labels:['FY2024','FY2025'],datasets:[
  {label:'I Motors',data:D.imoFy,backgroundColor:navy,borderRadius:4,barThickness:36},
  {label:'UMI',data:D.umiFy,backgroundColor:terra,borderRadius:4,barThickness:36}
]},options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},
  scales:{y:{ticks:{callback:v=>'€'+(v/1e6).toFixed(1)+'M',color:ink},grid:{color:grid}},x:{ticks:{color:ink},grid:{display:false}}}}});
function monthly(id,rows,color){
  const a=rows.filter(r=>r.fiscal_year===2024), b=rows.filter(r=>r.fiscal_year===2025);
  const labels=[...a.map(r=>r.ym.slice(5)),'|',...b.map(r=>r.ym.slice(5))];
  const data=[...a.map(r=>r.sales_eur),null,...b.map(r=>r.sales_eur)];
  new Chart(document.getElementById(id),{type:'bar',data:{labels,datasets:[{data,backgroundColor:data.map((v,i)=>i===a.length?'transparent':color),borderRadius:2,barThickness:9}]},
    options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false},title:{display:true,text:'FY2024 | FY2025',color:ink}},
      scales:{y:{ticks:{callback:v=>'€'+(v/1000).toFixed(0)+'K',color:ink},grid:{color:grid}},x:{ticks:{color:ink,maxRotation:90,font:{size:8}},grid:{display:false}}}}});
}
monthly('cImoM',D.imoMonthly,navy); monthly('cUmiM',D.umiMonthly,terra);
new Chart(cOrd,{type:'bar',data:{labels:D.orderSales.labels,datasets:[
  {label:'Sales H1',data:D.orderSales.sales,backgroundColor:navy,borderRadius:4,barThickness:28},
  {label:'Demand H1',data:D.orderSales.demand,backgroundColor:terra,borderRadius:4,barThickness:28}
]},options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},
  scales:{y:{ticks:{callback:v=>'€'+(v/1000).toFixed(0)+'K',color:ink},grid:{color:grid}},x:{ticks:{color:ink},grid:{display:false}}}}});
function doughnut(id,rows){const r=(rows||[]).filter(x=>x.eur>0);
  new Chart(document.getElementById(id),{type:'doughnut',data:{labels:r.map(x=>x.name+' '+x.pct.toFixed(0)+'%'),datasets:[{data:r.map(x=>x.eur),backgroundColor:[navy,terra,good,grey]}]},
    options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}}}});}
doughnut('cUmiSrc',D.umiSrcH1); doughnut('cImoSrc',D.imoSrcH1);
function saty(id,a,b){const keys=[...new Set([...a,...b].map(x=>x.saty))].filter(k=>k&&k!=='(blank)');
  new Chart(document.getElementById(id),{type:'bar',data:{labels:keys,datasets:[
    {label:'FY2024',data:keys.map(k=>(a.find(x=>x.saty===k)||{}).share_pct||0),backgroundColor:grey,barThickness:16},
    {label:'FY2025',data:keys.map(k=>(b.find(x=>x.saty===k)||{}).share_pct||0),backgroundColor:navy,barThickness:16}
  ]},options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},
    scales:{x:{max:100,ticks:{callback:v=>v+'%',color:ink},grid:{color:grid}},y:{ticks:{color:ink},grid:{display:false}}}}});}
saty('cImoSaty',D.imoSaty24,D.imoSaty25); saty('cUmiSaty',D.umiSaty24,D.umiSaty25);
new Chart(cPeers,{type:'bar',data:{labels:D.peers.map(p=>p.code),datasets:[{label:'SPU',data:D.peers.map(p=>p.spu),
  backgroundColor:D.peers.map(p=>p.code==='UMI'?terra:(p.code==='PORTUGAL'?navy:grey)),borderRadius:4,barThickness:22}]},
  options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false},title:{display:true,text:'Peer SPU — IDS Israel = UMI network',color:ink}},
    scales:{x:{ticks:{callback:v=>'€'+v,color:ink},grid:{color:grid}},y:{ticks:{color:ink},grid:{display:false}}}}});
function pms(id,rows){new Chart(document.getElementById(id),{type:'bar',data:{labels:rows.map(r=>r.category+' · '+(r.capture_rate*100).toFixed(0)+'%'),
  datasets:[{data:rows.map(r=>r.revenue_gap_eur),backgroundColor:bad,borderRadius:4,barThickness:16}]},
  options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false}},
    scales:{x:{ticks:{callback:v=>'€'+(v/1000).toFixed(0)+'K',color:ink},grid:{color:grid}},y:{ticks:{color:ink,font:{size:10}},grid:{display:false}}}}});}
pms('cImoPms',D.imoPms); pms('cUmiPms',D.umiPms);
function cats(id,rows,color){const r=(rows||[]).slice().sort((a,b)=>b.eur-a.eur);
  new Chart(document.getElementById(id),{type:'bar',data:{labels:r.map(x=>x.name),datasets:[{data:r.map(x=>x.eur),backgroundColor:r.map(x=>(x.yoy||0)<0?bad:color),borderRadius:4,barThickness:14}]},
  options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false}},
    scales:{x:{ticks:{callback:v=>'€'+(v/1000).toFixed(0)+'K',color:ink},grid:{color:grid}},y:{ticks:{color:ink,font:{size:10}},grid:{display:false}}}}});}
cats('cImoCat',D.imoCatH1,navy); cats('cUmiCat',D.umiCatH1,terra);
</script>
</body></html>""")

    html = body.safe_substitute(
        billing_max=DATA["local_billing_max"],
        imo_yoy=f"{imo_yoy:.1f}",
        umi_yoy=f"{umi_yoy:.1f}",
        imo_fy24=eur(imo_fy24),
        imo_fy25=eur(imo_fy25),
        umi_fy24=eur(umi_fy24),
        umi_fy25=eur(umi_fy25),
        imo_h1_yoy=f"{drills['IMOTORS']['yoy_pct']:.0f}",
        umi_h1_yoy=f"{drills['UMI']['yoy_pct']:.0f}",
        imo_h1_sales=eur(drills["IMOTORS"]["sales_eur"]),
        umi_h1_sales=eur(drills["UMI"]["sales_eur"]),
        imo_h1_demand=eur(order["customers"]["IMOTORS"]["demand_eur"]),
        umi_h1_demand=eur(order["customers"]["UMI"]["demand_eur"]),
        imo_dem_yoy=f"{order['customers']['IMOTORS']['yoy_pct']:.0f}",
        imo_gap=eur(order["sales_vs_demand"]["IMOTORS"]["gap_eur"]),
        umi_gap=eur(order["sales_vs_demand"]["UMI"]["gap_eur"]),
        uti_h1_sales=eur(pbi["customers"]["UTI"]["sales_eur"]),
        uti_h1_demand=eur(order["customers"]["UTI"]["demand_eur"]),
        imo_top10=f"{imo['concentration']['FY2025']['top10_share_pct']:.0f}",
        umi_top10=f"{umi['concentration']['FY2025']['top10_share_pct']:.0f}",
        imo_pms=f"{100*imo_pms['capture_rate_total_wtd']:.1f}",
        umi_pms=f"{100*umi_pms['capture_rate_total_wtd']:.1f}",
        imo_pms_gap=eur(imo_pms["revenue_gap_eur"]),
        umi_pms_gap=eur(umi_pms["revenue_gap_eur"]),
        imo_gen=f"{100*imo_pms['capture_rate_genuine_wtd']:.1f}",
        umi_oil=f"{100*umi_pms['oil_filter_capture']:.0f}",
        imo_spu=f"{imo_gps['spu']:.0f}",
        umi_spu=f"{umi_gps['spu']:.0f}",
        imo_uio=f"{imo_gps['uio']:,.0f}",
        umi_uio=f"{umi_gps['uio']:,.0f}",
        imo_gps=f"{100*(imo_gps['gps_capture_vs_max'] or 0):.0f}",
        umi_gps=f"{100*(umi_gps['gps_capture_vs_max'] or 0):.0f}",
        imo_tl=imo_gps["traffic_light"],
        umi_tl=umi_gps["traffic_light"],
        imo_gaps=imo_pms["top_gap_categories"].replace("|", " · "),
        umi_gaps=umi_pms["top_gap_categories"].replace("|", " · "),
        imo_fill=f"{FILL['IMOTORS']['avg_pct']:.1f}",
        imo_fill_src=FILL["IMOTORS"]["source"],
        umi_fill_src=FILL["UMI"]["source"],
        imo_losers=mover_table(imo["pn_movers_fy24_to_fy25"]["losers"], bad=True),
        umi_losers=mover_table(umi["pn_movers_fy24_to_fy25"]["losers"], bad=True),
        imo_pn=pn_table(imo_pn),
        umi_pn=pn_table(umi_pn),
    )
    html = html.replace("@@DATA@@", json.dumps(chart_data))
    HTML.write_text(html, encoding="utf-8")
    print("wrote", HTML)
    return {
        "imo": imo,
        "umi": umi,
        "drills": drills,
        "order": order,
        "peers": peers,
        "imo_fy24": imo_fy24,
        "imo_fy25": imo_fy25,
        "umi_fy24": umi_fy24,
        "umi_fy25": umi_fy25,
        "imo_yoy": imo_yoy,
        "umi_yoy": umi_yoy,
        "pbi": pbi,
    }


def build_excel(P):
    imo, umi = P["imo"], P["umi"]
    drills, order = P["drills"], P["order"]
    wb = Workbook()
    navy_f = PatternFill("solid", fgColor="1A2B4A")
    head = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title = Font(name="Calibri", bold=True, color="1A2B4A", size=14)
    body = Font(name="Calibri", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D3C8"),
        right=Side(style="thin", color="D9D3C8"),
        top=Side(style="thin", color="D9D3C8"),
        bottom=Side(style="thin", color="D9D3C8"),
    )
    zebra = PatternFill("solid", fgColor="F7F4EE")

    def hdr(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row, c)
            cell.fill = navy_f
            cell.font = head
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def table(ws, r0, headers, rows, money=()):
        for j, h in enumerate(headers, 1):
            ws.cell(r0, j, h)
        hdr(ws, r0, len(headers))
        for i, row in enumerate(rows):
            for j, val in enumerate(row, 1):
                cell = ws.cell(r0 + 1 + i, j, val)
                cell.font = body
                cell.border = thin
                cell.alignment = wrap
                if i % 2:
                    cell.fill = zebra
                if j in money and isinstance(val, (int, float)):
                    cell.number_format = "#,##0"

    def widths(ws, ws_):
        for i, w in enumerate(ws_, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(w, 28 if i == 2 else w)

    ws = wb.active
    ws.title = "00_README"
    ws["A1"] = "I Motors & UMI — Deep Study V2 (10 pages)"
    ws["A1"].font = title
    for i, t in enumerate(
        [
            "",
            "HTML: ISZE_Imotors_UMI_Deep_Study_V2.html",
            DATA["fiscal_definition"],
            "Primary = total PMS capture. IDS Israel = UMI network only.",
            "All-DB shallow packs exist; deepen later with this V2 template.",
            "Resume all-DB: python DB_Sales_Order_Study/_build_network_study.py --resume",
        ],
        2,
    ):
        ws.cell(i, 1, t).alignment = wrap
    ws.column_dimensions["A"].width = 100

    ws = wb.create_sheet("01_02_Executive")
    table(
        ws,
        1,
        ["Metric", "I Motors", "UMI"],
        [
            ["FY2024 €", P["imo_fy24"], P["umi_fy24"]],
            ["FY2025 €", P["imo_fy25"], P["umi_fy25"]],
            ["FY YoY %", P["imo_yoy"], P["umi_yoy"]],
            ["H1 sales €", drills["IMOTORS"]["sales_eur"], drills["UMI"]["sales_eur"]],
            ["H1 sales YoY %", drills["IMOTORS"]["yoy_pct"], drills["UMI"]["yoy_pct"]],
            ["H1 demand €", order["customers"]["IMOTORS"]["demand_eur"], order["customers"]["UMI"]["demand_eur"]],
            ["H1 gap €", order["sales_vs_demand"]["IMOTORS"]["gap_eur"], order["sales_vs_demand"]["UMI"]["gap_eur"]],
            ["PMS total %", round(100 * imo["pms_summary"][0]["capture_rate_total_wtd"], 1), round(100 * umi["pms_summary"][0]["capture_rate_total_wtd"], 1)],
            ["PMS gap €", imo["pms_summary"][0]["revenue_gap_eur"], umi["pms_summary"][0]["revenue_gap_eur"]],
            ["Fill %", FILL["IMOTORS"]["avg_pct"], "n/a"],
        ],
        money={2, 3},
    )
    widths(ws, [36, 18, 18])

    ws = wb.create_sheet("03_Monthly")
    rows = []
    for db, name in ((imo, "IMOTORS"), (umi, "UMI")):
        for m in db["monthly"]:
            rows.append([name, m["fiscal_year"], m["ym"], m["sales_eur"], m["qty"], m["lines"]])
    table(ws, 1, ["DB", "FY", "YM", "Sales €", "Qty", "Lines"], rows, money={4})
    widths(ws, [12, 10, 12, 14, 12, 10])

    ws = wb.create_sheet("04_Order_vs_Sales")
    table(
        ws,
        1,
        ["DB", "Sales", "Sales YoY", "Demand", "Demand YoY", "Gap"],
        [
            ["IMOTORS", drills["IMOTORS"]["sales_eur"], drills["IMOTORS"]["yoy_pct"], order["customers"]["IMOTORS"]["demand_eur"], order["customers"]["IMOTORS"]["yoy_pct"], order["sales_vs_demand"]["IMOTORS"]["gap_eur"]],
            ["UMI", drills["UMI"]["sales_eur"], drills["UMI"]["yoy_pct"], order["customers"]["UMI"]["demand_eur"], order["customers"]["UMI"]["yoy_pct"], order["sales_vs_demand"]["UMI"]["gap_eur"]],
            ["UTI", P["pbi"]["customers"]["UTI"]["sales_eur"], 15, order["customers"]["UTI"]["demand_eur"], order["customers"]["UTI"]["yoy_pct"], None],
        ],
        money={2, 4, 6},
    )
    widths(ws, [14, 14, 12, 14, 12, 12])

    ws = wb.create_sheet("05_SaTy")
    rows = []
    for db, name in ((imo, "IMOTORS"), (umi, "UMI")):
        for x in db["by_saty"]:
            if x["fiscal_year"] in (2024, 2025):
                rows.append([name, x["fiscal_year"], x["saty"], x["sales_eur"], x["share_pct"]])
    table(ws, 1, ["DB", "FY", "SaTy", "Sales €", "Share %"], rows, money={4})
    widths(ws, [12, 8, 12, 14, 12])

    ws = wb.create_sheet("06_GPS_Peers")
    table(ws, 1, ["Code", "SPU €", "UIO", "GPS %", "Sales", "TL"], [[p["code"], p["spu"], p["uio"], p["gps"], p["sales"], p["tl"]] for p in P["peers"]], money={2, 5})
    widths(ws, [12, 12, 12, 12, 14, 12])

    ws = wb.create_sheet("07_PMS")
    r = 1
    for db, name in ((imo, "IMOTORS"), (umi, "UMI")):
        ws.cell(r, 1, name).font = title
        cats = sorted(db["pms_categories_total"], key=lambda x: -x["revenue_gap_eur"])
        table(ws, r + 1, ["Category", "Capture", "Actual €", "Gap €"], [[c["category"], c["capture_rate"], c["actual_rev_eur"], c["revenue_gap_eur"]] for c in cats], money={3, 4})
        for i in range(len(cats)):
            ws.cell(r + 2 + i, 2).number_format = "0.0%"
        r += 3 + len(cats)
    widths(ws, [28, 12, 14, 14])

    ws = wb.create_sheet("08_H1_Category")
    r = 1
    for key, name in (("IMOTORS", "I Motors"), ("UMI", "UMI")):
        ws.cell(r, 1, name).font = title
        cats = drills[key].get("category", [])
        table(ws, r + 1, ["Category", "€", "Share %", "YoY"], [[c["name"], c["eur"], c["pct"], c.get("yoy")] for c in cats], money={2})
        r += 3 + len(cats)
    widths(ws, [28, 12, 12, 12])

    ws = wb.create_sheet("09_10_PN_Asks")
    r = 1
    for db, name in ((imo, "IMOTORS"), (umi, "UMI")):
        ws.cell(r, 1, f"{name} top PN FY25").font = title
        pn = [x for x in db["top_pn"] if x["fiscal_year"] == 2025][:20]
        table(ws, r + 1, ["Rank", "PN", "Desc", "€", "%"], [[x["rank"], x["part_number"], x["description"], x["sales_eur"], x["share_pct"]] for x in pn], money={4})
        r += 3 + len(pn)
        ws.cell(r, 1, f"{name} losers").font = Font(bold=True)
        losers = db["pn_movers_fy24_to_fy25"]["losers"][:15]
        table(ws, r + 1, ["PN", "Desc", "FY24", "FY25", "Delta"], [[x["part_number"], x["description"], x["fy_a_eur"], x["fy_b_eur"], x["delta_eur"]] for x in losers], money={3, 4, 5})
        r += 3 + len(losers)
    widths(ws, [10, 28, 40, 12, 12])

    for name in wb.sheetnames:
        wb[name].freeze_panes = "A2"
    wb.save(XLSX)
    print("wrote", XLSX)


def main():
    P = build_html()
    build_excel(P)


if __name__ == "__main__":
    main()
