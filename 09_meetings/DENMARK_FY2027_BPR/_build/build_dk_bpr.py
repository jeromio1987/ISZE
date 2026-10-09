"""Isuzu Danmark | FY2027 Business Performance Review (charcoal layout of the I Motors review, 24 Sep 2026).

Every figure is computed here from a named source. Typed in, with their source stated:
  - Power BI fill-rate points (individual fill rate, read 1 Oct 2026; perfect fill by order type, read 5 Oct 2026)
  - Isuzu Danmark's own VIN count by registration year (their pivot, shared 5 Oct 2026; excl. Faroe Islands and Greenland)
  - back-order lines on 29-09-2026 (8), as given by Jerome on 5 Oct 2026

Sources
  ISZE billing          : 01_data/00_derived/sales_fact.parquet, sold-to 2027 + 2751 (Isuzu Danmark), through 31 Aug 2026
  FY2027 BP (4+8)       : 01_data/90_INBOX/_bp_fy27_4plus8_2026-10-01/Full Package Plan VP+ACC FY2027-28 BP 4+8 V3.xlsx,
                          sheet 'VP FY27-28', row '2027 - ISUZU DANMARK' (block 'BP 2027 (FY revised 4+8)' and 'BP 2028 (FY)')
  Extended target       : 01_data/ISZE_GPS_Extended_Target_FY27_V3.1.xlsx, sheet 'Dashboard', row DENMARK (cols C, D, F, G, I)
  Sales to dealer, BO   : 01_data/11_planning/KPI Master File august 2026.xlsx, sheet 'Parts Handling'
  Registrations         : same file, sheet 'Vehicle Registration, RM, WE', 'No. of Vehicles Registered LCV'
  Genuine parts share   : 01_data/00_derived/gps_spu_by_distributor.csv (CY2025, all markets) and gps_master.csv (FY billing)
  Fleet (UIO), ISZE     : 01_data/00_derived/uio_by_market_type_year.csv (ISUZU_UIO_Overview_2026 '1.UIO Overview')
  Part classification   : 01_data/00_derived/aim_pn_classification_20260920.csv (catalog_l1 / l2), bvp_pn_dimension.csv
"""
import pathlib
import numpy as np
import openpyxl
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

NAMED = False   # other markets in the ranking slides: False = anonymous (distributor-facing), True = country names
MIN_PT = 9      # smallest font size anywhere in the deck

R = pathlib.Path(r"C:\Users\jerom\Desktop\claude\ISZE")
DER = R / "01_data" / "00_derived"
HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE.parent / "Isuzu_Danmark_FY2027_Business_Performance_Review_2026-10-05.pptx"
LOGO = R / "09_meetings" / "IMotors_FY2027_BPR" / "charcoal_first7" / "_build" / "isuzu_logo.png"
BP_FILE = R / "01_data" / "90_INBOX" / "_bp_fy27_4plus8_2026-10-01" / "Full Package Plan VP+ACC FY2027-28 BP 4+8 V3.xlsx"
EXT_FILE = R / "01_data" / "ISZE_GPS_Extended_Target_FY27_V3.1.xlsx"
KPI_FILE = R / "01_data" / "11_planning" / "KPI Master File august 2026.xlsx"

# ======================================================================= data
sf = pd.read_parquet(DER / "sales_fact.parquet")
sf["fy"] = np.where(sf.billing_date.dt.month >= 4, sf.billing_date.dt.year + 1, sf.billing_date.dt.year)
sf["ym"] = sf.billing_date.dt.to_period("M")
dk = sf[sf.sold_to_sap.isin([2027.0, 2751.0])].copy()
LAST = dk.billing_date.max()                                   # 2026-08-31
AA = [4, 5, 6, 7, 8]                                            # comparable Apr-Aug
FR = dk.part_number.eq("TRANSPORT")                             # freight line, billed with the parts

bill_m = dk.groupby("ym").net_sales_value.sum()
def ym(y, m): return pd.Period(f"{y}-{m:02d}", "M")
FY_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
def fy_series(fy):
    out = []
    for i in range(12):
        m = (i + 3) % 12 + 1; y = fy - 1 if m >= 4 else fy
        out.append(float(bill_m.get(ym(y, m), np.nan)))
    return out
b26, b27 = fy_series(2026), fy_series(2027)
ytd27 = sum(b27[:5]); ytd26 = sum(b26[:5]); fy26_tot = sum(b26)

# BP
ws = openpyxl.load_workbook(BP_FILE, read_only=True, data_only=True)["VP FY27-28"]
rows = list(ws.iter_rows(values_only=True))
def bp_row(start):
    for r in rows[start:start + 40]:
        if r[1] and "DANMARK" in str(r[1]).upper():
            return [float(v) for v in r[3:15]], float(r[17])
bp48, bp48_tot = bp_row(40)          # BP 2027 (FY revised 4+8): Apr-Jul actual, Aug-Mar plan
bp28, bp28_tot = bp_row(80)          # BP 2028 (FY)
bp_ytd = sum(bp48[:5]); bp_rest = sum(bp48[5:]); need_rest = bp48_tot - ytd27

# Extended target (business performance review): BP + stretch % of market potential
ext_row = next(r for r in openpyxl.load_workbook(EXT_FILE, read_only=True, data_only=True)["Dashboard"].iter_rows(values_only=True)
               if r[0] == "DENMARK")
EXT_BP, EXT_POT, EXT_PCT, EXT_STRETCH, EXT_TGT = float(ext_row[2]), float(ext_row[3]), float(ext_row[5]), float(ext_row[6]), float(ext_row[8])
assert abs(EXT_BP + EXT_STRETCH - EXT_TGT) < 1 and abs(EXT_POT * EXT_PCT - EXT_STRETCH) < 1
ext_rest = EXT_TGT - ytd27

# KPI master file
wb = openpyxl.load_workbook(KPI_FILE, read_only=True, data_only=True)
ph = list(wb["Parts Handling"].iter_rows(values_only=True))
HDR = ph[2]
def kpi_row(section_title):
    sec = None
    for r in ph:
        if r[0] and isinstance(r[0], str) and r[1] is None and r[2] is None: sec = r[0]
        if sec == section_title and r[0] == "Isuzu Danmark": return dict(zip(HDR[1:], r[1:]))
    raise KeyError(section_title)
SD = kpi_row("Parts Sales DB to DLR in €"); BO = kpi_row("Back Order Lines")
assert ph[10][0] == "Isuzu Danmark"
sd_fy26_keys = ["Apr 25", "Mai 25", "Jun 25", "Jul 25", "Aug 25", "Sep 25", "Oct 25", "Nov 25", "Dec 25", "Jan 26", "Feb 26", "Mar 26"]
sd_fy27_keys = ["Apr 26", "May 26", "Jun 26", "Jul 26"]
sd26 = [float(SD[k]) for k in sd_fy26_keys]; sd27 = [float(SD[k]) for k in sd_fy27_keys]
sd_fy26 = sum(sd26); sd_aj26 = sum(sd26[:4]); sd_aj27 = sum(sd27)
sd_r12 = sd_fy26 - sd_aj26 + sd_aj27
bo_keys = ["Aug 25", "Sep 25", "Oct 25", "Nov 25", "Dec 25", "Jan 26", "Feb 26", "Mar 26", "Apr 26", "May 26", "Jun 26", "Jul 26"]
bo = [BO[k] for k in bo_keys]
BO_NOW, BO_DATE = 8, "29-09-2026"
reg_rows = list(wb["Vehicle Registration, RM, WE"].iter_rows(values_only=True))
REG = dict(zip(reg_rows[2][1:], next(r for r in reg_rows if r[0] == "Isuzu Danmark")[1:]))
reg25 = sum(int(REG[k]) for k in ["Jan 25", "Feb 25", "Mrz 25", "Apr 25", "Mai 25", "Jun 25", "Jul 25", "Aug 25", "Sep 25", "Oct 25", "Nov 25", "Dec 25"])

# Fill rate, Power BI
FR_M = ["Sep 25", "Oct 25", "Nov 25", "Dec 25", "Jan 26", "Feb 26", "Mar 26", "Apr 26", "May 26", "Jun 26", "Jul 26", "Aug 26", "Sep 26"]
FR_V = [0.97, 0.94, 0.86, 0.91, 0.94, 0.92, 0.87, 0.83, 0.84, 0.97, 0.96, 0.94, 0.88]      # individual fill rate, labels as shown
FR_TARGET = 0.93
# perfect fill by order type, exact values from the report's data points (RUSH has no lines in Dec 2025 and Apr 2026)
FT_SOR = [97.33, 96.27, 93.46, 92.08, 97.50, 91.53, 92.31, 92.00, 91.14, 96.25, 97.33, 94.44, 90.00]
FT_VOR = [95.12, 84.62, 74.58, 89.47, 92.55, 91.49, 83.10, 69.70, 77.05, 98.08, 93.94, 92.31, 88.64]
FT_RUSH = [100.0, 100.0, 0.0, None, 100.0, 100.0, 83.33, None, 0.0, 100.0, 100.0, 100.0, 50.0]

# Order pattern (parts only: freight and uncoded lines excluded)
def urgent_share(df):
    c = df[df.saty.notna() & ~df.part_number.eq("TRANSPORT") & ~df.saty.eq("DR")]
    return c[c.saty.isin(["VOR", "RUSH", "SVOR"])].net_sales_value.sum() / c.net_sales_value.sum(), c.net_sales_value.sum()
u25, _ = urgent_share(dk[dk.fy == 2025]); u26, c26 = urgent_share(dk[dk.fy == 2026])
u_ja26, _ = urgent_share(dk[(dk.billing_date >= "2026-07-01") & (dk.billing_date <= "2026-08-31")])
u_ja25, _ = urgent_share(dk[(dk.billing_date >= "2025-07-01") & (dk.billing_date <= "2025-08-31")])
net26 = sf[(sf.fy == 2026) & sf.sold_to_sap.notna()]
u_net, _ = urgent_share(net26)
saty_cov27 = dk[(dk.fy == 2027) & ~FR & dk.saty.notna()].net_sales_value.sum() / dk[(dk.fy == 2027) & ~FR].net_sales_value.sum()
rush_lines26 = int((dk[(dk.fy == 2026)].saty == "RUSH").sum())

# Part-number pattern (FY2026, parts only)
p = dk[~FR]
p26 = p[p.fy == 2026]; p25 = p[p.fy == 2025]; p27 = p[p.fy == 2027]
g26 = p26.groupby("part_number").net_sales_value.sum().sort_values(ascending=False)
n_pn26 = len(g26); top20 = g26.head(20).sum() / g26.sum()
n80 = int((g26.cumsum() / g26.sum() < 0.8).sum() + 1)
months_per_pn = p26.groupby("part_number").ym.nunique()
once = int((months_per_pn == 1).sum()); six = int((months_per_pn >= 6).sum())
new27 = set(p27.part_number) - set(p[p.fy <= 2026].part_number)
new27_eur = p27[p27.part_number.isin(new27)].net_sales_value.sum()
lost = set(p25.part_number) - set(p26.part_number)
lost_eur = p25[p25.part_number.isin(lost)].net_sales_value.sum()
pn_aa = {fy: p[(p.fy == fy) & p.billing_date.dt.month.isin(AA)].part_number.nunique() for fy in (2025, 2026, 2027)}
top_pn = (p26.groupby(["part_number", "description"]).net_sales_value.sum().sort_values(ascending=False).head(8))

# Categories, comparable Apr-Aug
cls = pd.read_csv(DER / "aim_pn_classification_20260920.csv", dtype=str, usecols=["part_number", "catalog_l1", "catalog_l2"])
pc = p[p.billing_date.dt.month.isin(AA) & p.fy.isin([2026, 2027])].merge(cls, on="part_number", how="left")
pc["catalog_l1"] = pc.catalog_l1.fillna("Unclassified").replace({"Other / review": "Other"})
cat = pc.pivot_table(index="catalog_l1", columns="fy", values="net_sales_value", aggfunc="sum").fillna(0)
cat = cat[cat.max(axis=1) > 1000].sort_values(2027, ascending=False)
rep27 = pc[(pc.catalog_l1 == "Repair parts") & (pc.fy == 2027)]
rep27_elec = rep27[rep27.catalog_l2 == "Electrical sensors"].net_sales_value.sum()
bvp = pd.read_csv(DER / "bvp_pn_dimension.csv", dtype=str, usecols=["part_number", "line_type"]).drop_duplicates("part_number")
pb = p26.merge(bvp, on="part_number", how="left")
bvp26 = pb[pb.line_type == "bvp"].net_sales_value.sum()

# Genuine parts share, ISZE fleet
gm = pd.read_csv(DER / "gps_master.csv", index_col=0)
spu = pd.read_csv(DER / "gps_spu_by_distributor.csv")
r25 = spu[(spu.year == 2025) & spu.gps_capture_vs_max.notna()].sort_values("gps_capture_vs_max", ascending=False).reset_index(drop=True)
dk25 = r25[r25.distributor_code == "DENMARK"].iloc[0]
WALLET = float(dk25.max_spend_eur); UIO25 = float(dk25.uio_units); RATE = WALLET / UIO25
BILL25 = float(dk25.total_sales_eur)
dk_rank = int(r25.index[r25.distributor_code == "DENMARK"][0]) + 1; n_mkts = len(r25)
med = float(r25.gps_capture_vs_max.median())
net_w = float(r25.total_sales_eur.sum() / r25.max_spend_eur.sum())
nord = r25[r25.region == "Nordics"].reset_index(drop=True)
dk_nrank = int(nord.index[nord.distributor_code == "DENMARK"][0]) + 1
NET_FY = {fy: float(gm.loc["Danmark", f"net_FY{fy}"]) for fy in (2024, 2025, 2026)}
gfy = {fy: NET_FY[fy] / WALLET for fy in NET_FY}
gps25 = float(dk25.gps_capture_vs_max); gps25_xw = float(dk25.gps_capture_ex_warranty_vs_max)

# Fleet: ISZE (shipping data + IMIT scrap-rate logic) and Isuzu Danmark's own VIN count
u = pd.read_csv(DER / "uio_by_market_type_year.csv")
ud = u[u.distributor_code == "DENMARK"].groupby("year").uio_units.sum()
uio = {y: float(ud.loc[y]) for y in range(2019, 2027)}
ISZE25, ISZE26 = uio[2025], uio[2026]                                   # 851 (2025 estimate), 945 in the file
ISZE26_SAID = 944                                                       # LCV row of the same file, as quoted by ISZE
DK_VIN = {2008: 1, 2010: 1, 2014: 4, 2015: 28, 2016: 50, 2017: 45, 2018: 57, 2019: 69, 2020: 55, 2021: 65,
          2022: 72, 2023: 91, 2024: 88, 2025: 82, 2026: 33}
OWN26 = sum(DK_VIN.values()); OWN25 = OWN26 - DK_VIN[2026]
assert OWN26 == 741 and OWN25 == 708 and round(ISZE25) == 851
W_OWN = OWN25 * RATE
gps25_own = BILL25 / W_OWN
gfy_own = {fy: NET_FY[fy] / W_OWN for fy in NET_FY}
others = r25[r25.distributor_code != "DENMARK"]
rank_own = int((others.gps_capture_vs_max > gps25_own).sum()) + 1
nothers = nord[nord.distributor_code != "DENMARK"]
nrank_own = int((nothers.gps_capture_vs_max > gps25_own).sum()) + 1
def next_up(share, pool):
    above = pool[pool.gps_capture_vs_max > share].gps_capture_vs_max
    return float(above.min()) if len(above) else None

m16 = [pd.Period(f"2025-{m:02d}", "M") for m in range(4, 13)] + [pd.Period(f"2026-{m:02d}", "M") for m in range(1, 8)]
sd16 = sd26 + sd27
bp16 = [float(dk[(dk.ym == p_) & ~FR].net_sales_value.sum()) for p_ in m16]
ratio = sum(sd16) / sum(bp16)

# ======================================================================= style (I Motors charcoal layout)
BG, PANEL, BORDER, ROW2 = "262626", "303030", "6E6E6E", "2B2B2B"
WHITE, MUTED, GRID = "F2F2F2", "BDBDBD", "4D4D4D"
RED, AMBER, GREY, GREEN, LGREEN = "E30613", "F5A000", "B7B7B7", "31A354", "55C271"
FONT = "Aptos"
def rgb(h): return RGBColor.from_string(h)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def txt(slide, x, y, w, h, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    size = max(size, MIN_PT)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p_ = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_.alignment = align
        r = p_.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
    return tb

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def hline(slide, x, y, w, color, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)

PAGE = [1]
def base(section, title, sub, source):
    PAGE[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    txt(s, 0.25, 0.10, 4.5, 0.2, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 9, WHITE, True)
    txt(s, 0.25, 0.42, 8.0, 0.22, section.upper(), 9, RED, True)
    txt(s, 0.25, 0.68, 9.45, 0.66, title, 20 if len(title) <= 60 else 16.5, WHITE, True)
    txt(s, 0.25, 1.42, 9.45, 0.24, sub, 9, MUTED)
    hline(s, 0.25, 7.09, 12.85, "575757")
    txt(s, 0.25, 7.13, 11.9, 0.34, source, 9, "8C8C8C")
    txt(s, 12.4, 7.13, 0.7, 0.2, str(PAGE[0]), 9, MUTED, align=PP_ALIGN.RIGHT)
    return s

def panel(slide, points, x=9.86, y=1.46, w=3.25, h=5.55):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.28, y + 0.24, w - 0.5, 0.2, "OBSERVATIONS / QUESTIONS", 9, RED, True)
    yy = y + 0.66; step = (h - 0.85) / max(len(points), 1)
    for i, pnt in enumerate(points):
        rect(slide, x + 0.28, yy + 0.07, 0.09, 0.09, RED if i == 0 else AMBER, shape=MSO_SHAPE.OVAL)
        txt(slide, x + 0.52, yy, w - 0.75, step - 0.06, pnt, 9, WHITE)
        yy += step

def tile(slide, x, y, w, h, label, big, sub, color=WHITE, big_size=24):
    rect(slide, x, y, w, h, PANEL, BORDER)
    txt(slide, x + 0.16, y + 0.14, w - 0.26, 0.2, label, 9, RED, True)
    txt(slide, x + 0.16, y + 0.42, w - 0.26, 0.5, big, big_size, color, True)
    txt(slide, x + 0.16, y + 0.98, w - 0.26, h - 1.05, sub, 9, MUTED)

def style_chart(ch, num_fmt='#,##0', legend=True):
    ch.font.name = FONT; ch.font.size = Pt(MIN_PT); ch.font.color.rgb = rgb(MUTED); ch.has_title = False
    if legend:
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.color.rgb = rgb(MUTED); ch.legend.font.size = Pt(MIN_PT)
    else: ch.has_legend = False
    va = ch.value_axis; va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(GRID); va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.color.rgb = rgb(MUTED); va.tick_labels.font.size = Pt(MIN_PT)
    va.tick_labels.number_format = num_fmt; va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(GRID); ca.tick_labels.font.color.rgb = rgb(MUTED); ca.tick_labels.font.size = Pt(MIN_PT)

def labels(series, fmt, color=WHITE, pos=XL_LABEL_POSITION.OUTSIDE_END, size=9):
    dl = series.data_labels; dl.show_value = True
    dl.number_format = fmt; dl.number_format_is_linked = False
    dl.font.size = Pt(max(size, MIN_PT)); dl.font.color.rgb = rgb(color); dl.font.bold = True; dl.position = pos

def col_fill(series, c):
    series.format.fill.solid(); series.format.fill.fore_color.rgb = rgb(c); series.format.line.fill.background()

def point_fill(series, i, c):
    pt = series.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)

def line_fmt(series, c, width=2.25, marker=True, dash=False):
    series.smooth = False
    series.format.line.color.rgb = rgb(c); series.format.line.width = Pt(width)
    if dash: series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if marker:
        series.marker.style = XL_MARKER_STYLE.CIRCLE; series.marker.size = 6
        series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = rgb(c)
        series.marker.format.line.color.rgb = rgb(c)
    else: series.marker.style = XL_MARKER_STYLE.NONE

def chart(slide, kind, cd, x, y, w, h):
    return slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart

def bar_box(slide, text, color=WHITE, y=6.22, h=0.78, size=10.5):
    rect(slide, 0.4, y, 9.2, h, PANEL, BORDER)
    txt(slide, 0.62, y, 8.8, h, text, size, color, True, anchor=MSO_ANCHOR.MIDDLE)

def k(v): return f"€{v/1000:,.1f}K"
def pct(v, d=1): return f"{v*100:.{d}f}%"
def sgn(v, d=1): return f"{'+' if v >= 0 else ''}{v*100:.{d}f}%"
FLEET_TXT = f"ISZE fleet {UIO25:,.0f}  |  Isuzu Danmark VIN count {OWN25:,.0f}"

# ======================================================================= 1 cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
rect(s, 0, 0, 0.10, 7.5, RED)
txt(s, 0.45, 0.32, 5, 0.2, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", 9, WHITE, True)
txt(s, 0.45, 2.05, 10.5, 1.3, ["Isuzu Danmark", "FY2027 Business Performance Review"], 30, WHITE, True)
txt(s, 0.45, 3.55, 11.5, 0.3, "FY27 plan & extended target  ·  Sales to dealer  ·  Fill rate  ·  Order pattern  ·  Fleet  ·  Genuine parts share ranking", 11, MUTED)
txt(s, 0.45, 4.05, 11.5, 0.3, "ISZE billing through August 2026, fill rate through September 2026  |  Review draft, 5 October 2026", 10, MUTED)
s.shapes.add_picture(str(LOGO), Inches(0.45), Inches(6.6), height=Inches(0.36))

# ======================================================================= 2 agenda (with the data rule from the network review)
s = base("Agenda", "Review structure", "Isuzu Danmark FY2027 business performance review", "Fiscal year = April to March (FY2027 = April 2026 to March 2027).")
items = [("01", "Where we stand", "FY2027 plan, extended target and pace; sales to dealer"),
         ("02", "Service", "Fill rate from ISZE, in total and by order type; back orders and stock"),
         ("03", "What your purchases tell us", "Order pattern, part-number pattern, category movers"),
         ("04", "Fleet and genuine parts share", "ISZE fleet and your VIN count; share and ranking on both"),
         ("05", "Next steps", "Part-level sales to dealer, priorities, action points")]
for i, (n, h, d) in enumerate(items):
    y = 1.95 + i * 0.92
    rect(s, 0.4, y, 0.62, 0.62, RED); txt(s, 0.4, y, 0.62, 0.62, n, 16, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, 1.2, y + 0.02, 7.5, 0.3, h, 14, WHITE, True); txt(s, 1.2, y + 0.36, 7.5, 0.25, d, 10, MUTED)
rect(s, 9.86, 1.95, 3.25, 4.6, PANEL, BORDER)
txt(s, 10.14, 2.2, 2.8, 0.2, "DATA RULE FOR THIS REVIEW", 9, RED, True)
txt(s, 10.14, 2.6, 2.8, 3.8, ["Every figure is read from a named ISZE file or report on its build date, listed in the footer of each page and in the method page.",
                               "", "Where a measure is modelled, the page says so. Where Isuzu Danmark does not report a KPI, it is shown as not reported, never as zero.",
                               "", "Fiscal and calendar years are labelled on every page and never mixed in one chart."], 9, WHITE)

# ======================================================================= 3 executive view
s = base("Executive view", "Billing is on plan and dealer sales grow faster; fill rate on urgent orders is the gap to close",
         "Periods differ by KPI: ISZE billing Apr–Aug 2026; sales to dealer Apr–Jul 2026; fill rate Sep 2026; genuine parts share calendar 2025.",
         f"Sources: ISZE billing; FY2027 BP 4+8; Extended Target FY27 V3.1; KPI Master File Aug 2026; Power BI fill rate; back-order lines {BO_DATE}; ISZE fleet file and Isuzu Danmark VIN count.")
tiles = [
    ("ISZE BILLING VS BP", k(ytd27), [f"Apr–Aug 2026 vs {k(bp_ytd)} BP", f"{sgn(ytd27 / bp_ytd - 1)} vs plan, {sgn(ytd27 / ytd26 - 1)} vs last year"], LGREEN, 24),
    ("FY2027 EXTENDED TARGET", k(EXT_TGT), [f"BP {k(EXT_BP)} + {pct(EXT_PCT)} stretch", f"{k(ext_rest)} to go, {k(ext_rest / 7)} a month"], AMBER, 24),
    ("SALES TO DEALER", k(sd_aj27), [f"Apr–Jul 2026 vs {k(sd_aj26)}", f"{sgn(sd_aj27 / sd_aj26 - 1)} year on year"], LGREEN, 24),
    ("FILL RATE ISZE → DANMARK", pct(FR_V[-1], 0), [f"Sep 2026 vs {pct(FR_TARGET, 0)} target", f"urgent orders lowest: {FT_VOR[7]:.1f}% (Apr)"], RED, 24),
    ("URGENT ORDER SHARE", pct(u26, 0), ["of parts billing FY2026", f"network: {pct(u_net, 0)}"], AMBER, 24),
    ("BACK-ORDER LINES", f"{BO_NOW}", [f"on {BO_DATE}", "days in stock not reported"], WHITE, 24),
    ("GENUINE PARTS SHARE", f"{pct(gps25)} | {pct(gps25_own)}", ["calendar 2025, ISZE fleet | own VINs", f"rank #{dk_rank} | #{rank_own} of {n_mkts} markets"], AMBER, 18),
    ("FLEET 2025 (UIO)", f"{UIO25:,.0f} | {OWN25:,.0f}", ["ISZE fleet file | your VIN count", f"2026: {ISZE26_SAID} | {OWN26}"], AMBER, 22),
]
tw, th, gx, gy = 2.2, 1.6, 0.13, 0.15
for i, (lab, big, sub, c, bs) in enumerate(tiles):
    tile(s, 0.4 + (i % 4) * (tw + gx), 1.8 + (i // 4) * (th + gy), tw, th, lab, big, sub, c, bs)
rect(s, 0.4, 5.4, 4 * tw + 3 * gx, 1.6, PANEL, BORDER)
txt(s, 0.6, 5.55, 1.6, 0.25, "Management read", 10, RED, True)
txt(s, 2.25, 5.55, 7.15, 1.4, f"Isuzu Danmark is on plan: {k(ytd27)} billed against {k(bp_ytd)} in the BP, and dealer sales grow faster still. "
    f"The extended target of {k(EXT_TGT)} needs {k(ext_rest / 7)} a month from September. "
    f"What holds the business back is availability: about {pct(u26, 0)} of parts are bought as urgent orders, and that is where ISZE fill rate is weakest.", 10, WHITE)
panel(s, [
    f"The extended target adds {k(EXT_STRETCH)} ({pct(EXT_PCT)} of market potential) to the {k(EXT_BP)} BP. At the current pace of {k(ytd27 / 5)} a month it is within reach.",
    f"Dealer sales are {k(sd_r12)} over the last 12 months: €{ratio:.1f} sold to dealers for every €1 bought from ISZE since April 2025. We want to understand that difference.",
    f"Genuine parts share: {pct(gps25)} on the ISZE fleet ({UIO25:,.0f}), {pct(gps25_own)} on your own VIN count ({OWN25:,.0f}). One agreed fleet base would settle it.",
    "Question: can Isuzu Danmark share sales to dealer by part number and month, so stock can follow dealer demand?"])

# ======================================================================= 4 extended target (business performance review)
s = base("FY2027 extended target", f"FY2027 extended target: the BP plus a {pct(EXT_PCT)} stretch makes {k(EXT_TGT)}",
         "ISZE → Isuzu Danmark  |  business performance review  |  stretch = percentage of Denmark's market potential added on top of the BP",
         "Sources: ISZE_GPS_Extended_Target_FY27_V3.1.xlsx, Dashboard row DENMARK (BP, max spend, stretch %, stretch EUR, extended target); ISZE billing to 31 Aug 2026; FY2027 BP 4+8.")
cols = [(0.4, 5.6), (6.0, 2.0)]
y = 1.85
for (x, w), h in zip(cols, ["Step", "Amount"]):
    rect(s, x, y, w, 0.42, RED); txt(s, x + 0.15, y, w - 0.3, 0.42, h, 10, WHITE, True, PP_ALIGN.LEFT if x < 1 else PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
y += 0.46
steps = [("Business Plan FY2027", f"€{EXT_BP:,.0f}", False), ("Market potential FY2027 (maximum spend)", f"€{EXT_POT:,.0f}", False),
         (f"Stretch: +{pct(EXT_PCT)} of market potential", f"€{EXT_STRETCH:,.0f}", False), ("FY2027 EXTENDED TARGET", f"€{EXT_TGT:,.0f}", True)]
for i, (a_, b2, hl) in enumerate(steps):
    fillc = ROW2 if i % 2 == 0 else PANEL
    for (x, w), v in zip(cols, (a_, b2)):
        rect(s, x, y, w, 0.46, fillc)
        txt(s, x + 0.15, y, w - 0.3, 0.46, v, 11, RED if hl else WHITE, hl, PP_ALIGN.LEFT if x < 1 else PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    y += 0.49
txt(s, 0.4, 4.45, 4.0, 0.25, "Billed so far against the extended target", 10, WHITE, True)
txt(s, 5.6, 4.45, 2.4, 0.25, k(EXT_TGT), 10, WHITE, True, PP_ALIGN.RIGHT)
bw = 7.6; done = ytd27 / EXT_TGT
rect(s, 0.4, 4.75, bw, 0.55, "4A4A4A"); rect(s, 0.4, 4.75, bw * done, 0.55, RED)
txt(s, 0.55, 4.75, 3.0, 0.55, f"{k(ytd27)} billed (Apr–Aug)", 11, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.4 + bw * done + 0.15, 4.75, 4.2, 0.55, f"{k(ext_rest)} to go (Sep–Mar)", 11, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
t3w = (9.2 - 2 * 0.2) / 3
for i, (lab, big, sub, c) in enumerate([
        ("NEEDED PER MONTH", k(ext_rest / 7), ["Sep–Mar, to reach the extended target"], AMBER),
        ("BP PLAN PER MONTH", k(bp_rest / 7), ["Sep–Mar in the BP 4+8"], WHITE),
        ("BILLED PER MONTH", k(ytd27 / 5), ["Apr–Aug 2026 average"], LGREEN)]):
    tile(s, 0.4 + i * (t3w + 0.2), 5.5, t3w, 1.5, lab, big, sub, c, 24)
panel(s, [
    f"The business performance review sets the stretch at {pct(EXT_PCT)} of Denmark's market potential ({k(EXT_POT)}): {k(EXT_STRETCH)} on top of the {k(EXT_BP)} BP.",
    f"After five months Isuzu Danmark has billed {k(ytd27)}, {pct(done)} of the extended target; five of twelve months is {pct(5 / 12)}.",
    f"From September {k(ext_rest / 7)} a month is needed, against {k(bp_rest / 7)} in the BP and {k(ytd27 / 5)} billed per month so far.",
    f"Question: which actions (stock list, competitive parts, fill rate) does Isuzu Danmark see behind the extra {k(EXT_STRETCH)}?"])

# ======================================================================= 5 FY27 plan monthly
s = base("FY2027 plan", f"April to August billing is on the BP; the extended target needs {k(ext_rest / 7)} a month from September",
         "ISZE billing to Isuzu Danmark, € K per month  |  FY2027 actual to 31 August 2026 against the BP 4+8",
         "Sources: ISZE billing, sold-to 2027 and 2751, freight included as in the BP; Full Package Plan VP+ACC FY2027-28 BP 4+8 V3 (1 Oct 2026); Extended Target FY27 V3.1.")
cd = CategoryChartData(); cd.categories = FY_MONTHS
cd.add_series("Actual FY2027", [round(v / 1000, 1) if not np.isnan(v) else None for v in b27[:5]] + [None] * 7)
cd.add_series("BP 4+8", [round(v / 1000, 1) for v in bp48])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 4.25)
style_chart(ch, '€#,##0"K"'); ch.plots[0].gap_width = 70; ch.plots[0].overlap = -5
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 25
for se, c in zip(ch.plots[0].series, [RED, GREY]): col_fill(se, c); labels(se, '#,##0.0', MUTED)
bar_box(s, f"Apr–Aug {k(ytd27)} vs {k(bp_ytd)} BP ({sgn(ytd27 / bp_ytd - 1)})   |   FY2027 BP {k(bp48_tot)}   |   extended target {k(EXT_TGT)}   |   FY2028 BP {k(bp28_tot)}", WHITE, size=10)
panel(s, [
    f"Billing is {k(ytd27)} after five months against {k(bp_ytd)} in the BP: {sgn(ytd27 / bp_ytd - 1)}. July was the strongest month at {k(b27[3])}.",
    f"The BP for September to March is {k(bp_rest)}, {k(bp_rest / 7)} a month. The extended target of {k(EXT_TGT)} needs {k(ext_rest / 7)} a month.",
    f"The BP adds {sgn(bp48_tot / fy26_tot - 1)} on FY2026 ({k(fy26_tot)}); the extended target adds {sgn(EXT_TGT / fy26_tot - 1)}. FY2028 BP is {k(bp28_tot)}.",
    "Question: is there anything in the second half (fleet deals, campaigns, stock decisions) that changes the monthly pace?"])

# ======================================================================= 6 FY27 pace
cum26 = np.cumsum(b26); cum27 = np.cumsum(b27[:5]); cumbp = np.cumsum(bp48)
s = base("FY2027 pace", f"April–August billing is {sgn(ytd27 / ytd26 - 1)} on last year; a straight-line full year gives {pct(ytd27 / 5 * 12 / bp48_tot, 0)} of the BP",
         "Cumulative ISZE billing to Isuzu Danmark, € K  |  FY2026 against FY2027 to 31 August 2026 and the BP 4+8",
         "Sources: ISZE billing, sold-to 2027 and 2751, freight included; FY2027 BP 4+8; Extended Target FY27 V3.1. Straight line = Apr–Aug billing ÷ 5 × 12.")
cd = CategoryChartData(); cd.categories = FY_MONTHS
cd.add_series("FY2026", [round(v / 1000, 1) for v in cum26])
cd.add_series("FY2027", [round(v / 1000, 1) for v in cum27] + [None] * 7)
cd.add_series("BP 4+8", [round(v / 1000, 1) for v in cumbp])
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.8, 9.2, 3.7)
style_chart(ch, '€#,##0"K"'); ch.value_axis.minimum_scale = 0
a, b_, c_ = ch.plots[0].series
line_fmt(a, GREY, 2); line_fmt(b_, RED, 2.75); line_fmt(c_, AMBER, 1.75, marker=False, dash=True)
labels(b_, '#,##0.0', WHITE, XL_LABEL_POSITION.ABOVE)
for i, (lab, big, sub, c) in enumerate([
        ("APR–AUG FY2026", k(ytd26), ["Same five months last year"], WHITE),
        ("APR–AUG FY2027", k(ytd27), [f"{sgn(ytd27 / ytd26 - 1)} year on year"], LGREEN),
        ("STRAIGHT-LINE FY2027", k(ytd27 / 5 * 12), [f"{pct(ytd27 / 5 * 12 / bp48_tot)} of BP, {pct(ytd27 / 5 * 12 / EXT_TGT)} of extended target"], LGREEN)]):
    tile(s, 0.4 + i * (t3w + 0.2), 5.55, t3w, 1.45, lab, big, sub, c, 22)
panel(s, [
    f"FY2027 runs above FY2026 in every cumulative month so far; the gap is {k(ytd27 - ytd26)} after August.",
    f"FY2026 ended at {k(fy26_tot)}. The BP needs {k(need_rest / 7)} a month from September, the extended target {k(ext_rest / 7)}.",
    f"The straight line is a check, not a forecast: FY2026 had its strongest months in October and February ({k(b26[6])} and {k(b26[10])}).",
    "Question: which months does Isuzu Danmark expect to be strong this year, so ISZE can plan supply around them?"])

# ======================================================================= 7 sales to dealer
s = base("Sales to dealer", f"Dealer sales are up {sgn(sd_aj27 / sd_aj26 - 1, 0)} on April–July last year; the last 12 months reached {k(sd_r12)}",
         "Isuzu Danmark reported sales to dealer, € K per month  |  FY2026 (April 2025 to March 2026) and FY2027 to July 2026",
         "Source: KPI Master File august 2026.xlsx, Parts Handling – Parts Sales DB to DLR in €. August 2026 not yet reported.")
cd = CategoryChartData(); cd.categories = FY_MONTHS
cd.add_series("FY2026", [round(v / 1000, 1) for v in sd26])
cd.add_series("FY2027", [round(v / 1000, 1) for v in sd27] + [None] * 8)
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 3.7)
style_chart(ch, '€#,##0"K"'); ch.plots[0].gap_width = 70; ch.value_axis.minimum_scale = 0
for se, c in zip(ch.plots[0].series, [GREY, RED]): col_fill(se, c); labels(se, '#,##0', MUTED)
for i, (lab, big, sub, c) in enumerate([
        ("APR–JUL FY2026", k(sd_aj26), ["Reported sales to dealer"], WHITE),
        ("APR–JUL FY2027", k(sd_aj27), [f"{sgn(sd_aj27 / sd_aj26 - 1)} vs comparable period"], LGREEN),
        ("FY2026 TOTAL", k(sd_fy26), [f"Largest month: {k(max(sd26))} (Dec 2025)"], WHITE)]):
    tile(s, 0.4 + i * (t3w + 0.2), 5.55, t3w, 1.45, lab, big, sub, c, 22)
panel(s, [
    f"Every FY2027 month so far is above the same month last year; July 2026 was {k(sd27[3])} against {k(sd26[3])}.",
    f"FY2026 had a strong winter: November and December 2025 were {k(sd26[7])} and {k(sd26[8])}, more than double the spring months.",
    "These are Isuzu Danmark's sales to its dealers, at its own prices. They are not comparable in level with ISZE billing (next page).",
    "Question: what drove the November–December 2025 peak: fleet orders, a campaign, accessories, or tyres and lubricants?"])

# ======================================================================= 8 sales to dealer vs ISZE billing
lab16 = [p_.strftime("%b %y") for p_ in m16]
s = base("Sales to dealer  ·  ISZE billing", "Dealer sales and ISZE billing do not move together; the winter peak did not reach our billing",
         "€ K per month  |  Isuzu Danmark sales to dealer against ISZE parts billing to Isuzu Danmark (freight excluded)  |  April 2025 to July 2026",
         "Sources: KPI Master File Aug 2026 (sales to dealer); ISZE billing, sold-to 2027 and 2751, freight excluded. Different price levels: compare movements, not levels.")
cd = CategoryChartData(); cd.categories = lab16
cd.add_series("Sales to dealer", [round(v / 1000, 1) for v in sd16])
cd.add_series("ISZE parts billing", [round(v / 1000, 1) for v in bp16])
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.8, 9.2, 4.25)
style_chart(ch, '€#,##0"K"'); ch.value_axis.minimum_scale = 0
line_fmt(ch.plots[0].series[0], RED, 2.5); line_fmt(ch.plots[0].series[1], GREY, 2.25)
labels(ch.plots[0].series[0], '#,##0', WHITE, XL_LABEL_POSITION.ABOVE)
bar_box(s, f"Apr 2025 – Jul 2026:  sales to dealer {k(sum(sd16))}   |   ISZE parts billing {k(sum(bp16))}   |   €{ratio:.1f} of dealer sales per €1 bought from ISZE", size=10)
panel(s, [
    f"In November–December 2025 dealer sales rose to {k(sd16[7])} and {k(sd16[8])} while ISZE parts billing stayed near {k((bp16[7] + bp16[8]) / 2)} a month.",
    f"Over 16 months Isuzu Danmark sold €{ratio:.1f} to dealers for every €1 of parts bought from ISZE. Distributor margin alone does not explain a gap of that size.",
    "Possible reasons: the dealer figure includes accessories, lubricants, tyres or parts from other suppliers, or stock built before April 2025 was sold.",
    "Question: what does the sales-to-dealer figure include? Part-level data would answer this directly."])

# ======================================================================= 9 fill rate
s = base("Fill rate  ·  ISZE → Isuzu Danmark", f"ISZE fill rate to Isuzu Danmark recovered to 94–97% in summer, but fell back to {pct(FR_V[-1], 0)} in September",
         f"ISZE supplying Isuzu Danmark  |  individual fill rate  |  Sep 2025 – Sep 2026  |  latest point {pct(FR_V[-1], 0)}",
         f"Source: Power BI, SALES DATA VP NEW – INDIVIDUAL FILL RATE, customer ISZ DENMARK, engines excluded (data updated 1 Oct 2026). Target {pct(FR_TARGET, 0)}, the current ISZE internal target.")
cd = CategoryChartData(); cd.categories = FR_M
cd.add_series("Fill rate ISZE → Isuzu Danmark", FR_V); cd.add_series(f"Target {pct(FR_TARGET, 0)}", [FR_TARGET] * len(FR_V))
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '0%'); ch.value_axis.minimum_scale = 0.80; ch.value_axis.maximum_scale = 1.0; ch.value_axis.major_unit = 0.05
a, t = ch.plots[0].series
line_fmt(a, RED, 2.5); line_fmt(t, AMBER, 1.75, marker=False, dash=True)
labels(a, '0%', WHITE, XL_LABEL_POSITION.ABOVE, 10)
below = sum(v < FR_TARGET for v in FR_V)
bar_box(s, f"Latest {pct(FR_V[-1], 0)}   |   high {pct(max(FR_V), 0)}   |   low {pct(min(FR_V), 0)}   |   {below} of {len(FR_V)} months below the {pct(FR_TARGET, 0)} target", RED, y=6.25, h=0.75, size=11)
panel(s, [
    f"This is our own supply to Isuzu Danmark. {len(FR_V) - below} of the {len(FR_V)} months are at or above the {pct(FR_TARGET, 0)} target; the low point is {pct(min(FR_V), 0)} in April 2026.",
    "March to May 2026 sat at 83–87%, three months in a row below target. The next page shows that the dip sits in urgent orders.",
    f"June to August recovered to 94–97%; September dropped to {pct(FR_V[-1], 0)}. ISZE will check which part numbers caused the drop.",
    "Question for both sides: which part numbers does Isuzu Danmark miss most often, so ISZE can stock them for Denmark?"])

# ======================================================================= 10 fill rate by order type
avg = lambda v: sum(x for x in v if x is not None) / len([x for x in v if x is not None])
s = base("Fill rate  ·  by order type", f"Stock orders are filled at 90–98%; urgent orders dropped to {FT_VOR[7]:.0f}–{FT_VOR[8]:.0f}% in April and May",
         "ISZE supplying Isuzu Danmark  |  perfect fill by order type, %  |  Sep 2025 – Sep 2026  |  SOR = stock order, VOR = urgent order, RUSH = rush order",
         "Source: Power BI, SALES DATA VP NEW – FILL RATE BY ORDER TYPE, customer ISZ DENMARK, measure 'Perfect fill' (order types FD, PRD, RE, SVOR excluded by the report).")
cd = CategoryChartData(); cd.categories = FR_M
cd.add_series("Stock orders (SOR)", [v / 100 for v in FT_SOR])
cd.add_series("Urgent orders (VOR)", [v / 100 for v in FT_VOR])
cd.add_series("Rush orders (RUSH)", [None if v is None else v / 100 for v in FT_RUSH])
ch = chart(s, XL_CHART_TYPE.LINE_MARKERS, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '0%'); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.value_axis.major_unit = 0.2
sor, vor, rush = ch.plots[0].series
line_fmt(sor, GREY, 2.25); line_fmt(vor, RED, 2.75); line_fmt(rush, AMBER, 1.25, dash=True)
labels(vor, '0%', WHITE, XL_LABEL_POSITION.BELOW, 9)
bar_box(s, f"Average of the 13 months:  stock orders {avg(FT_SOR):.1f}%   |   urgent orders {avg(FT_VOR):.1f}%   |   lowest urgent month {min(FT_VOR):.1f}% (Apr 2026)", y=6.25, h=0.75, size=10.5)
panel(s, [
    f"Stock orders are filled at {min(FT_SOR):.0f}–{max(FT_SOR):.0f}% every month. The weak spot is urgent orders: {FT_VOR[2]:.1f}% in November 2025, {FT_VOR[7]:.1f}% and {FT_VOR[8]:.1f}% in April and May 2026.",
    f"Urgent orders are about {pct(u26, 0)} of Isuzu Danmark's parts value, so the weakest fill sits on a large part of the business.",
    f"Rush orders are very few ({rush_lines26} billed lines in FY2026); 0% in November and May means one or two lines not filled, not a trend.",
    "This page uses Power BI's perfect fill by order type, a different measure from the previous page. Question: which urgent part numbers were missed in April and May?"])

# ======================================================================= 11 stock & availability
s = base("Stock & availability  ·  Isuzu Danmark", "Back orders stay low; stock cover is not reported yet",
         "Isuzu Danmark back-order lines  |  KPI Master File  |  Aug 2025 – Jul 2026, and the count on 29-09-2026",
         f"Sources: KPI Master File august 2026.xlsx, Parts Handling – Back Order Lines and Days of Parts in Stock; back-order lines on {BO_DATE}.")
txt(s, 0.5, 1.85, 4.3, 0.22, "BACK-ORDER LINES", 9, RED, True)
cd = CategoryChartData(); cd.categories = bo_keys + [BO_DATE]; cd.add_series("Back-order lines", [int(v) for v in bo] + [BO_NOW])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 2.1, 5.6, 3.9)
style_chart(ch, '#,##0', legend=False); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 15; ch.plots[0].gap_width = 45
col_fill(ch.plots[0].series[0], GREY); labels(ch.plots[0].series[0], '#,##0', WHITE)
point_fill(ch.plots[0].series[0], len(bo), AMBER)
rect(s, 6.25, 2.1, 3.35, 3.9, PANEL, BORDER)
txt(s, 6.45, 2.3, 3.0, 0.22, "MONTHS OF STOCK (MOS)", 9, RED, True)
txt(s, 6.45, 2.7, 3.0, 0.6, "Not reported", 22, AMBER, True)
txt(s, 6.45, 3.45, 3.0, 2.4, ["'Days of Parts in Stock' is marked N/A (not tracked) for Isuzu Danmark in every month of the KPI file.",
                               "", "With MOS reported, we can see whether urgent orders come from low stock or from parts never stocked."], 9, WHITE)
bar_box(s, f"Back-order lines: {min(bo)}–{max(bo)} per month (Aug 2025 – Jul 2026)   |   {BO_NOW} on {BO_DATE}   |   MOS: not reported", size=10.5)
panel(s, [
    f"Back orders are small: between {min(bo)} and {max(bo)} lines a month, with peaks of 10 in October and November 2025, and {BO_NOW} on {BO_DATE}.",
    "Low back orders and a high urgent-order share can go together: urgent orders are often placed exactly to avoid a back order.",
    "Without days in stock, we cannot tell how much buffer Isuzu Danmark holds on its fast movers.",
    "Request: report days of parts in stock monthly in the KPI file from October 2026."])

# ======================================================================= 12 order pattern
s = base("What your purchases tell us  ·  Order pattern", f"About {pct(u26, 0)} of parts are bought as urgent orders, against {pct(u_net, 0)} across the network",
         "Share of ISZE parts billing by order type  |  fiscal years  |  freight lines and lines without order type excluded",
         "Source: ISZE billing, order type per line. Urgent = VOR, RUSH, SVOR. Apr–Jun 2026 billing has no order type in the SAP export, so FY2027 uses Jul–Aug. Network = all ISZE customers, FY2026.")
cats = ["FY2025", "FY2026", "Jul–Aug 2025", "Jul–Aug 2026", "Network FY2026"]
vals = [u25, u26, u_ja25, u_ja26, u_net]
cd = CategoryChartData(); cd.categories = cats
cd.add_series("Urgent orders", [round(v, 3) for v in vals]); cd.add_series("Stock and other orders", [round(1 - v, 3) for v in vals])
ch = chart(s, XL_CHART_TYPE.COLUMN_STACKED_100, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '0%'); ch.plots[0].gap_width = 60
col_fill(ch.plots[0].series[0], RED); col_fill(ch.plots[0].series[1], "4A4A4A")
labels(ch.plots[0].series[0], '0%', WHITE, XL_LABEL_POSITION.CENTER, 11); labels(ch.plots[0].series[1], '0%', MUTED, XL_LABEL_POSITION.CENTER, 10)
point_fill(ch.plots[0].series[0], 4, GREY)
bar_box(s, f"Urgent share:  FY2025 {pct(u25, 0)}   |   FY2026 {pct(u26, 0)}   |   Jul–Aug 2026 {pct(u_ja26, 0)}   |   network {pct(u_net, 0)}", y=6.25, h=0.75, size=11)
panel(s, [
    f"Roughly two in every five euros of parts reach Isuzu Danmark as urgent orders, about {u26 / u_net:.1f} times the network share.",
    f"Urgent orders are also where ISZE fill is weakest: {FT_VOR[7]:.1f}% in April 2026 against {FT_SOR[7]:.1f}% for stock orders.",
    "A higher stock-order share on fast movers would lift fill rate and lower freight cost per part, on both sides.",
    "Question: which parts does Isuzu Danmark keep on its own shelf today, and is there a list we could stock against together?"])

# ======================================================================= 13 part-number pattern
s = base("What your purchases tell us  ·  Part numbers", f"Two in three part numbers were bought in one month only; {n80} make 80% of the value",
         "ISZE parts billing to Isuzu Danmark  |  FY2026 (April 2025 to March 2026), freight excluded  |  new part numbers in FY2027 to August",
         "Source: ISZE billing, sold-to 2027 and 2751, by part number and billing month.")
t4w = (9.2 - 3 * 0.17) / 4
for i, (lab, big, sub, c) in enumerate([
        ("PART NUMBERS FY2026", f"{n_pn26}", [f"Top 20 = {pct(top20, 0)} of value", f"{n80} part numbers = 80%"], WHITE),
        ("ONE MONTH ONLY", f"{once}", [f"{pct(once / n_pn26, 0)} of all part numbers", f"only {six} bought in 6+ months"], AMBER),
        ("NEW IN FY2027", f"{len(new27)}", ["never bought before Apr 2026", f"{k(new27_eur)} so far"], WHITE),
        ("NOT REPEATED", f"{len(lost)}", ["bought in FY2025, not FY2026", f"{k(lost_eur)} in FY2025"], AMBER)]):
    tile(s, 0.4 + i * (t4w + 0.17), 1.8, t4w, 1.62, lab, big, sub, c, 24)
txt(s, 0.4, 3.62, 9.2, 0.22, "LARGEST PART NUMBERS, FY2026", 9, RED, True)
y = 3.92
for i, ((pn, desc), v) in enumerate(top_pn.items()):
    fillc = ROW2 if i % 2 == 0 else PANEL
    rect(s, 0.4, y, 9.2, 0.33, fillc)
    txt(s, 0.55, y, 1.6, 0.33, pn, 9.5, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 2.2, y, 5.2, 0.33, str(desc).title(), 9.5, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 7.6, y, 1.8, 0.33, f"€{v:,.0f}", 9.5, WHITE, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    y += 0.36
panel(s, [
    f"Isuzu Danmark buys a long tail: {pct(once / n_pn26, 0)} of part numbers appear in one month only. That is the pattern of ordering per job, not of restocking.",
    f"Only {six} part numbers were bought in six or more months. These are the natural candidates for a stock list.",
    f"The range is widening: {pn_aa[2027]} part numbers in April–August 2026 against {pn_aa[2026]} a year earlier, {len(new27)} of them never bought before.",
    f"Best Value Parts are {k(bvp26)} of FY2026 billing, almost all one front brake-pad kit. Question: is there interest in BVP for older D-Max vehicles?"])

# ======================================================================= 14 category movers
def cv(name, fy): return float(cat.loc[name, fy]) if name in cat.index else 0.0
s = base("What your purchases tell us  ·  Categories", "Repair parts doubled on last year, mostly electrical sensors; accident parts are lower",
         "ISZE parts billing to Isuzu Danmark by part group, € K  |  comparable April–August, FY2026 against FY2027",
         "Sources: ISZE billing, freight excluded; ISZE part classification (catalogue level 1 and 2). Competitive parts = parts the independent aftermarket also sells.")
cd = CategoryChartData(); cd.categories = list(cat.index)
cd.add_series("Apr–Aug FY2026", [round(v / 1000, 1) for v in cat[2026]]); cd.add_series("Apr–Aug FY2027", [round(v / 1000, 1) for v in cat[2027]])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '€#,##0"K"'); ch.plots[0].gap_width = 70; ch.value_axis.minimum_scale = 0
for se, c in zip(ch.plots[0].series, [GREY, RED]): col_fill(se, c); labels(se, '#,##0.0', MUTED)
bar_box(s, f"Repair parts {k(cv('Repair parts', 2026))} → {k(cv('Repair parts', 2027))}   |   Accident parts {k(cv('Accident parts', 2026))} → {k(cv('Accident parts', 2027))}   |   Competitive parts {k(cv('Competitive parts', 2026))} → {k(cv('Competitive parts', 2027))}", y=6.25, h=0.75, size=10)
panel(s, [
    f"Repair parts are parts replaced when something fails rather than at a service. They grew from {k(cv('Repair parts', 2026))} to {k(cv('Repair parts', 2027))}; {k(rep27_elec)} of that is electrical sensors, mainly particulate (PM) sensors.",
    f"Competitive parts (filters, brakes and other parts the aftermarket also sells) are the largest group at {k(cv('Competitive parts', 2027))}, about flat on last year.",
    f"Accident parts fell from {k(cv('Accident parts', 2026))} to {k(cv('Accident parts', 2027))}; that group follows claims and is hard to plan.",
    "Question: is Isuzu Danmark seeing repeat PM-sensor failures on D-Max, and do independent workshops service the older vehicles?"])

# ======================================================================= 15 fleet (UIO): ISZE against Isuzu Danmark
s = base("Fleet", f"Two fleet counts for Denmark: ISZE {UIO25:,.0f} and Isuzu Danmark {OWN25:,.0f} for 2025",
         "Isuzu vehicles in operation in Denmark  |  ISZE: shipping data and IMIT scrap-rate logic  |  Isuzu Danmark: VINs by registration year",
         "Sources: ISUZU_UIO_Overview_2026 (ISZE fleet file); Isuzu Danmark VIN count by registration year (shared Oct 2026), Faroe Islands and Greenland excluded; KPI Master File registrations.")
yrs = [y for y in DK_VIN if y >= 2014]
cd = CategoryChartData(); cd.categories = ["before 2014"] + [str(y) + (" YTD" if y == 2026 else "") for y in yrs]
cd.add_series("Isuzu Danmark VINs by registration year", [DK_VIN[2008] + DK_VIN[2010]] + [DK_VIN[y] for y in yrs])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 4.7, 4.3)
style_chart(ch, '#,##0', legend=False); ch.plots[0].gap_width = 45; ch.value_axis.minimum_scale = 0
col_fill(ch.plots[0].series[0], GREY); labels(ch.plots[0].series[0], '#,##0', WHITE)
point_fill(ch.plots[0].series[0], len(yrs), AMBER)
txt(s, 0.4, 1.75, 4.7, 0.25, "ISUZU DANMARK: VINs ON THE ROAD BY REGISTRATION YEAR", 9, RED, True)
cols = [(5.3, 1.55), (6.85, 1.35), (8.2, 1.4)]
y = 1.85
for (x, w), h in zip(cols, ["", "ISZE", "Isuzu Danmark"]):
    rect(s, x, y, w, 0.42, RED); txt(s, x + 0.1, y, w - 0.18, 0.42, h, 10, WHITE, True, PP_ALIGN.CENTER if h else PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
y += 0.46
trows = [("Fleet 2025", f"{UIO25:,.0f}", f"{OWN25:,.0f}"), ("Fleet 2026", f"{ISZE26_SAID:,.0f}", f"{OWN26:,.0f} (YTD)"),
         ("Difference 2025", "", f"{OWN25 - UIO25:,.0f} ({(OWN25 / UIO25 - 1) * 100:.0f}%)"),
         ("Basis", "shipping data + scrap rate", "registered VINs"),
         ("Fleet potential 2025", k(WALLET), k(W_OWN)),
         ("Genuine parts share 2025", pct(gps25), pct(gps25_own))]
for i, row in enumerate(trows):
    fillc = ROW2 if i % 2 == 0 else PANEL; last = i == len(trows) - 1
    for j, ((x, w), v) in enumerate(zip(cols, row)):
        rect(s, x, y, w, 0.56, fillc)
        txt(s, x + 0.1, y, w - 0.18, 0.56, v, 11 if last and j else 9.5, (AMBER if j == 2 else WHITE) if last else WHITE, last or j == 0,
            PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    y += 0.59
bar_box(s, f"Same ISZE billing in calendar 2025 ({k(BILL25)}), two fleets:  {pct(gps25)} on {UIO25:,.0f} vehicles   →   {pct(gps25_own)} on {OWN25:,.0f} vehicles", size=10.5)
panel(s, [
    f"The counts differ by {UIO25 - OWN25:,.0f} vehicles for 2025 ({(1 - OWN25 / UIO25) * 100:.0f}%). ISZE counts from shipping data with a scrap-rate model; Isuzu Danmark counts VINs registered today. 2026: {ISZE26_SAID} against {OWN26}.",
    "Faroe Islands and Greenland are not in Isuzu Danmark's count. Question: does Isuzu Danmark supply parts there? If yes, those vehicles belong in the fleet behind our billing.",
    f"Your KPI return shows {reg25} new LCV registrations in 2025; the VIN file holds {DK_VIN[2025]} from 2025. Question: what explains the difference (demo, export, Faroe Islands and Greenland)?",
    "Until one base is agreed, this review shows the genuine parts share on both fleets."])

# ======================================================================= 16 genuine parts share trend, both fleets
s = base("Genuine parts share", f"Genuine parts share FY2026: {pct(gfy[2026])} on the ISZE fleet, {pct(gfy_own[2026])} on your VIN count",
         f"ISZE billing to Isuzu Danmark (parts, warranty included) ÷ fleet potential  |  fiscal years on a fixed 2025 fleet  |  €{RATE:,.0f} modelled parts spend per vehicle",
         "Sources: ISZE genuine parts share files (gps_master FY billing; gps_spu_by_distributor 2025 fleet potential); Isuzu Danmark VIN count. Fleet potential = fleet × modelled parts spend, not measured competitor spend.")
cd = CategoryChartData(); cd.categories = ["FY2024", "FY2025", "FY2026"]
cd.add_series(f"ISZE fleet ({UIO25:,.0f})", [round(gfy[fy], 3) for fy in (2024, 2025, 2026)])
cd.add_series(f"Isuzu Danmark VIN count ({OWN25:,.0f})", [round(gfy_own[fy], 3) for fy in (2024, 2025, 2026)])
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 0.4, 1.8, 9.2, 4.3)
style_chart(ch, '0%'); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.plots[0].gap_width = 70; ch.plots[0].overlap = -5
col_fill(ch.plots[0].series[0], RED); col_fill(ch.plots[0].series[1], AMBER)
labels(ch.plots[0].series[0], '0.0%', WHITE, size=11); labels(ch.plots[0].series[1], '0.0%', WHITE, size=11)
bar_box(s, f"Calendar 2025 (rankings):  {pct(gps25)} ISZE fleet   |   {pct(gps25_own)} Isuzu Danmark VIN count   |   warranty excluded, ISZE fleet: {pct(gps25_xw)}", y=6.25, h=0.75, size=10)
panel(s, [
    f"FY2026 is the best year of the three on either fleet: {pct(gfy[2026])} on the ISZE fleet, {pct(gfy_own[2026])} on Isuzu Danmark's own count.",
    f"The fleet count moves the level by about {(gfy_own[2026] - gfy[2026]) * 100:.0f} points; the trend is the same. FY2027 billing to August is {sgn(ytd27 / ytd26 - 1)} on last year.",
    f"Both use the same billing and the same €{RATE:,.0f} per vehicle; only the number of vehicles differs ({UIO25:,.0f} against {OWN25:,.0f}).",
    "Question: which families does Isuzu Danmark think it loses most to the aftermarket: filters, brakes or body parts?"])

# ======================================================================= 17 ranking, all markets, both fleets
entries = [(r_.gps_capture_vs_max, r_.country, r_.distributor_code) for _, r_ in others.iterrows()]
entries += [(gps25, f"Denmark, ISZE fleet ({UIO25:,.0f})", "DK_ISZE"), (gps25_own, f"Denmark, own VINs ({OWN25:,.0f})", "DK_OWN")]
entries.sort(key=lambda e: -e[0])
names, vals_, kinds, n_o = [], [], [], 0
for v, c_, code in entries:
    if code.startswith("DK_"): names.append(c_)
    else:
        n_o += 1
        names.append((str(c_) if code not in ("IUK", "ITUK", "UMI", "UTI") else f"{c_} ({code})") if NAMED else f"Market {n_o}")
    vals_.append(float(v)); kinds.append(code)
s = base("Genuine parts share  ·  Network ranking", f"Denmark ranks #{dk_rank} of {n_mkts} on the ISZE fleet and #{rank_own} on its own VIN count",
         f"Genuine parts share by market  |  calendar 2025  |  ISZE billing ÷ fleet potential  |  Denmark shown on both fleets  |  {'markets named' if NAMED else 'other markets not named'}",
         "Source: gps_spu_by_distributor.csv (2025, full year): ISZE billing incl. warranty ÷ fleet × modelled parts spend per vehicle; Denmark also on Isuzu Danmark's VIN count. Markets without a fleet base are not ranked.")
cd = CategoryChartData(); cd.categories = names[::-1]; cd.add_series("Genuine parts share 2025", [round(v, 3) for v in vals_[::-1]])
ch = chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cd, 0.4, 1.75, 9.2, 5.25)
style_chart(ch, '0%', legend=False); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.plots[0].gap_width = 30
col_fill(ch.plots[0].series[0], "5A5A5A"); labels(ch.plots[0].series[0], '0.0%', MUTED)
for i, kd in enumerate(kinds[::-1]):
    if kd == "DK_ISZE": point_fill(ch.plots[0].series[0], i, RED)
    if kd == "DK_OWN": point_fill(ch.plots[0].series[0], i, AMBER)
panel(s, [
    f"On the ISZE fleet Denmark is #{dk_rank} of {n_mkts} ({pct(gps25)}); on its own VIN count it is #{rank_own} ({pct(gps25_own)}). Nobody is above 100%; the highest market is {pct(float(r25.gps_capture_vs_max.iloc[0]))}.",
    f"Network median {pct(med)}; all ranked markets together {pct(net_w)}. Denmark is above both on either fleet.",
    "The other markets are on ISZE fleet counts. If their own counts are also lower than ours, they would move up too.",
    "The same purchases put Denmark in the middle or near the top of the table. That is why one agreed fleet base matters."])

# ======================================================================= 18 Nordic ranking, both fleets
nent = [(r_.gps_capture_vs_max, r_.country, r_.distributor_code, r_.uio_units) for _, r_ in nothers.iterrows()]
nent += [(gps25, "Denmark, ISZE fleet", "DK_ISZE", UIO25), (gps25_own, "Denmark, own VINs", "DK_OWN", OWN25)]
nent.sort(key=lambda e: -e[0])
nnames, letter = [], 0
for v, c_, code, _u in nent:
    if code.startswith("DK_"): nnames.append(c_)
    else:
        nnames.append(str(c_) if NAMED else f"Nordic market {chr(65 + letter)}"); letter += 1
up_i = next_up(gps25, nothers); up_o = next_up(gps25_own, nothers)
s = base("Genuine parts share  ·  Nordic ranking", f"In the Nordics Denmark is #{dk_nrank} of {len(nord)} on the ISZE fleet and #{nrank_own} on its own VIN count",
         f"Genuine parts share, Nordic markets  |  calendar 2025  |  same method as the network ranking  |  {'markets named' if NAMED else 'other markets not named'}",
         "Source: gps_spu_by_distributor.csv (2025, full year), region Nordics; Denmark also on Isuzu Danmark's VIN count. Fleet potential = fleet × modelled parts spend per vehicle.")
cd = CategoryChartData(); cd.categories = nnames[::-1]; cd.add_series("Genuine parts share 2025", [round(float(e[0]), 3) for e in nent[::-1]])
ch = chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cd, 0.4, 1.8, 5.5, 4.3)
style_chart(ch, '0%', legend=False); ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.plots[0].gap_width = 40
col_fill(ch.plots[0].series[0], "5A5A5A"); labels(ch.plots[0].series[0], '0.0%', WHITE, size=10)
for i, e in enumerate(nent[::-1]):
    if e[2] == "DK_ISZE": point_fill(ch.plots[0].series[0], i, RED)
    if e[2] == "DK_OWN": point_fill(ch.plots[0].series[0], i, AMBER)
cols = [(6.1, 1.75), (7.85, 0.85), (8.7, 0.9)]
y = 1.85
for (x, w), h in zip(cols, ["Market", "Fleet", "Share"]):
    rect(s, x, y, w, 0.38, RED); txt(s, x + 0.08, y, w - 0.14, 0.38, h, 9.5, WHITE, True, PP_ALIGN.LEFT if x < 7 else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
y += 0.42
for i, (e, nm) in enumerate(zip(nent, nnames)):
    fillc = ROW2 if i % 2 == 0 else PANEL; hl = e[2].startswith("DK_"); col = RED if e[2] == "DK_ISZE" else AMBER if e[2] == "DK_OWN" else WHITE
    for (x, w), v in zip(cols, [nm, f"{e[3]:,.0f}", pct(float(e[0]))]):
        rect(s, x, y, w, 0.44, fillc)
        txt(s, x + 0.08, y, w - 0.14, 0.44, v, 9.5, col, hl, PP_ALIGN.LEFT if x < 7 else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    y += 0.47
bar_box(s, f"Next place up:  ISZE fleet +{(up_i - gps25) * 100:.1f} points (≈ {k((up_i - gps25) * WALLET)} a year)   |   own VINs +{(up_o - gps25_own) * 100:.1f} points (≈ {k((up_o - gps25_own) * W_OWN)} a year)", y=6.25, h=0.75, size=10.5)
panel(s, [
    f"On the ISZE fleet Denmark is #{dk_nrank} of {len(nord)}, {(up_i - gps25) * 100:.1f} points behind the next market. On its own VIN count it is #{nrank_own}, only {(up_o - gps25_own) * 100:.1f} points from second place.",
    f"The two leaders have fleets of {nord.uio_units.iloc[0]:,.0f} and {nord.uio_units.iloc[1]:,.0f} vehicles, the same size class as Denmark, so they are a fair benchmark.",
    "With a small fleet, one large fleet customer or one bulk order moves the share by several points in a year.",
    "Question: what would Isuzu Danmark need from ISZE (stock list, pricing on competitive parts, fill rate) to move up?"])

# ======================================================================= 19 the ask: part-level sales to dealer
s = base("Next steps  ·  Data request", "Part-level sales to dealer: the same report you send today, by part number",
         "What we ask, and what it gives Isuzu Danmark", "Request: sales to dealer by part number and month, quantities only, from 1 January 2025 onwards; monthly after that. No prices or margins needed.")
cards = [("1", "Already halfway there", ["Isuzu Danmark reports sales to dealer every month in the KPI file.", "The ask is the same figure split by part number: quantity per part per month. No prices, no margins, no dealer names."]),
         ("2", "Stock that follows your dealers", ["Your fast movers become visible, so ISZE can stock them for Denmark.", f"Fewer urgent orders (today {pct(u26, 0)} of parts value) and a higher fill rate on the first order."]),
         ("3", "A clear view of where demand goes", ["Part-level data separates a stock movement from a lost job.", f"It answers the €{ratio:.1f}-per-€1 question earlier in this review and makes service-part cover measurable."]),
         ("4", "A fair target and a fair ranking", ["Targets and the genuine parts share can be built on what dealers actually sell.", "Together with one agreed fleet count, both sides work from one set of numbers."])]
cw = (9.2 - 0.3) / 2
for i, (n, h, body) in enumerate(cards):
    x = 0.4 + (i % 2) * (cw + 0.3); y = 1.85 + (i // 2) * 2.2
    rect(s, x, y, cw, 2.0, PANEL, BORDER)
    rect(s, x + 0.2, y + 0.2, 0.45, 0.45, RED); txt(s, x + 0.2, y + 0.2, 0.45, 0.45, n, 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    txt(s, x + 0.8, y + 0.24, cw - 1.0, 0.35, h, 12, WHITE, True)
    txt(s, x + 0.8, y + 0.68, cw - 1.0, 1.25, body, 9.5, MUTED)
rect(s, 0.4, 6.3, 9.2, 0.7, RED)
txt(s, 0.65, 6.3, 8.8, 0.7, "What we use it for: stock planning, fill rate and joint category actions. It is not used to review Isuzu Danmark's pricing to dealers.", 10, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
panel(s, [
    "Other distributors already share part-level sales to dealer with ISZE. It is what made stock-build and campaign effects visible in their reviews.",
    "The file can be a simple export: invoice month, part number, quantity. ISZE does the mapping to families and Genuine / BVP.",
    "We share the results back: a monthly view of your dealer demand by family, next to what you buy from ISZE.",
    "Question: who at Isuzu Danmark can provide the first file, and can it start with January 2025 to July 2026?"])

# ======================================================================= 20 discussion priorities
s = base("Discussion priorities", "Three topics should drive the Isuzu Danmark review",
         "Built on ISZE billing, the BP and extended target, the KPI file, Power BI fill rate and both fleet counts", "Sources: as on the individual pages.")
pri = [("Extended target", f"FY2027 extended target {k(EXT_TGT)}: the {k(EXT_BP)} BP plus a {pct(EXT_PCT)} stretch ({k(EXT_STRETCH)}). {k(ytd27)} is billed; September to March needs {k(ext_rest / 7)} a month against {k(ytd27 / 5)} so far."),
       ("Availability", f"Bring fill rate back above {pct(FR_TARGET, 0)}, starting with urgent orders ({FT_VOR[7]:.0f}% in April), and move the {six} regularly bought part numbers onto stock orders. Urgent share from {pct(u26, 0)} towards the network {pct(u_net, 0)}."),
       ("One set of numbers", f"Agree the fleet base (ISZE {UIO25:,.0f} / {ISZE26_SAID}, Isuzu Danmark {OWN25:,.0f} / {OWN26}, Faroe Islands and Greenland) and what the sales-to-dealer figure includes. Share, ranking and targets then rest on one base.")]
for i, (h, body) in enumerate(pri):
    y = 1.9 + i * 1.65
    rect(s, 0.4, y, 9.2, 1.5, PANEL, BORDER)
    rect(s, 0.4, y, 0.12, 1.5, RED)
    txt(s, 0.75, y + 0.16, 8.6, 0.32, f"{i + 1}.  {h}", 13, WHITE, True)
    txt(s, 0.75, y + 0.58, 8.6, 0.9, body, 10, MUTED)
panel(s, [
    f"Isuzu Danmark is on plan and growing. The extended target is {k(EXT_STRETCH)} above the BP and reachable at today's pace.",
    "Fill rate and urgent orders are two sides of one problem: ISZE stocks what it sees, and today it mainly sees urgent demand.",
    "The fleet question is in Isuzu Danmark's interest as much as ours: a correct base makes the targets and the ranking fair.",
    "Question: which of the three does Isuzu Danmark want to start with?"])

# ======================================================================= 21 action points (with the 'one decision' box from the network review)
s = base("Action points", "Proposed output of the review for Isuzu Danmark and ISZE", "To agree in the meeting; owners per side", "Owners and dates to be confirmed in the meeting.")
acts = [("Extended target", f"Agree the monthly path to {k(EXT_TGT)} and the actions behind the {k(EXT_STRETCH)} stretch.", "Joint"),
        ("Part-level sales to dealer", "First file Jan 2025 – Jul 2026 (month, part number, quantity), then monthly.", "Isuzu Danmark"),
        ("Fleet base", f"Confirm the {OWN26} scope (Faroe Islands, Greenland, deregistered vehicles); ISZE checks its {ISZE26_SAID} shipping basis.", "Joint"),
        ("Days in stock", "Report days of parts in stock monthly in the KPI file.", "Isuzu Danmark"),
        ("Fill rate", "List the urgent part numbers behind the April–May and September 2026 dips; agree a Denmark stock list.", "ISZE"),
        ("Order pattern", f"Move regularly bought parts from urgent to stock orders; review the urgent share quarterly (today {pct(u26, 0)}).", "Joint"),
        ("Sales to dealer scope", "Confirm what the reported sales-to-dealer figure includes (accessories, lubricants, other suppliers).", "Isuzu Danmark")]
cols = [(0.4, 2.2), (2.6, 5.45), (8.05, 1.55)]
y = 1.85
for (x, w), h in zip(cols, ["Topic", "Action", "Owner"]):
    rect(s, x, y, w, 0.4, RED); txt(s, x + 0.1, y, w - 0.15, 0.4, h, 9.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
y += 0.44
for i, row in enumerate(acts):
    fillc = ROW2 if i % 2 == 0 else PANEL
    for j, ((x, w), v) in enumerate(zip(cols, row)):
        rect(s, x, y, w, 0.5, fillc)
        txt(s, x + 0.1, y, w - 0.18, 0.5, v, 9, WHITE, j == 0, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.53
rect(s, 0.4, y + 0.1, 9.2, 0.85, PANEL, BORDER)
txt(s, 0.6, y + 0.18, 3.0, 0.25, "THE ONE DECISION WE NEED", 9, RED, True)
txt(s, 0.6, y + 0.45, 8.8, 0.45, "Whether part-level sales to dealer becomes a standing part of the Isuzu Danmark review. With it, stock, fill rate and targets follow dealer demand.", 10, WHITE, True)
panel(s, [
    "Suggested rhythm: a short monthly check on billing against the extended target, fill rate, urgent share and sales to dealer; this full review twice a year.",
    "The first part-level file lets ISZE draft a Denmark stock list within a month.",
    "The fleet count and the sales-to-dealer scope are one-off questions; once answered, every later number gets easier to read.",
    "Next review: after the first part-level data, with FY2027 half-year billing."])

# ======================================================================= 22 appendix: method
s = base("Appendix  ·  Method", "How these figures were built", "Data used, what was assumed, and what was not shown",
         "All files are ISZE internal unless stated; build date 5 October 2026.")
blocks = [
    ("ISZE billing", f"SAP billing to sold-to 2027 and 2751 (both Isuzu Danmark), to {LAST:%d %B %Y}. Plan pages include the freight line, as the BP does; all other analysis excludes it. Fiscal year April–March."),
    ("BP and extended target", f"BP 4+8 (1 Oct 2026): FY2027 {k(bp48_tot)}, FY2028 {k(bp28_tot)}. Extended Target FY27 V3.1: BP {k(EXT_BP)} + {pct(EXT_PCT)} of market potential {k(EXT_POT)} = {k(EXT_TGT)}."),
    ("KPI file and back orders", f"KPI Master File August 2026: sales to dealer (€), back-order lines, days in stock (not tracked for Denmark), new LCV registrations. Back-order lines on {BO_DATE}: {BO_NOW}."),
    ("Fill rate", "Power BI SALES DATA VP NEW: INDIVIDUAL FILL RATE (engines excluded, labels as shown) and FILL RATE BY ORDER TYPE (measure 'Perfect fill'), customer ISZ DENMARK."),
    ("Order type", f"SAP order type per billing line. Urgent = VOR, RUSH, SVOR. April–June 2026 has no order type in the export ({pct(1 - saty_cov27, 0)} of FY2027 parts value uncoded), so FY2027 uses July–August."),
    ("Genuine parts share", f"ISZE billing incl. warranty ÷ (fleet × €{RATE:,.0f} modelled parts spend per vehicle). Ranking: calendar 2025, {n_mkts} markets. Trend: fiscal billing on the fixed 2025 fleet."),
    ("Fleet", f"ISZE: fleet file, 2025 estimate {UIO25:,.0f}, 2026 {ISZE26_SAID} (shipping data, IMIT scrap-rate logic). Isuzu Danmark: {OWN26} VINs excl. Faroe Islands and Greenland; {OWN25} registered up to end 2025."),
    ("Not shown", "Service-part (PMS) cover: on a fleet of this size single bulk orders push several families above 100% in one year, so the measure is not reliable without sales-to-dealer data. No dealer prices or margins were used.")]
cw = (12.7 - 0.3) / 2
for i, (h, body) in enumerate(blocks):
    x = 0.4 + (i % 2) * (cw + 0.3); y = 1.85 + (i // 2) * 1.3
    rect(s, x, y, cw, 1.18, PANEL, BORDER)
    txt(s, x + 0.2, y + 0.12, cw - 0.4, 0.25, h.upper(), 9, RED, True)
    txt(s, x + 0.2, y + 0.4, cw - 0.4, 0.75, body, 9.5, WHITE)

prs.save(OUT)
print("saved", OUT)
print({"ytd27": ytd27, "ext": (EXT_BP, EXT_POT, EXT_PCT, EXT_STRETCH, EXT_TGT), "ext_rest_pm": ext_rest / 7, "gps25": gps25, "gps25_own": gps25_own,
       "rank": (dk_rank, rank_own, n_mkts), "nrank": (dk_nrank, nrank_own, len(nord)), "gfy": gfy, "gfy_own": gfy_own, "W_OWN": W_OWN,
       "up": (up_i, up_o), "rep27_elec": rep27_elec, "rush_lines26": rush_lines26, "ISZE26_file": ISZE26})
