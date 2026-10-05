"""Render PNG previews of the deck straight from the PPTX geometry.

No PowerPoint or LibreOffice is available in this session, so this reads the
slide XML (shape positions, text runs, pictures) and composites each slide with
PIL. It is a composition check: it proves where every element sits and that the
images are there, at the deck's real coordinates. Text metrics are approximate,
so treat font fitting as indicative rather than exact.
"""
from __future__ import annotations

import io
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECK = HERE / "ISZE_Business_Performance_Review_2026.pptx"
OUT = HERE / "_qa" / "slides"
OUT.mkdir(parents=True, exist_ok=True)

EMU_PER_IN = 914400
DPI = 110
FONT_PATHS = [
    r"C:\Windows\Fonts\calibri.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\arial.ttf",
]
FONT_BOLD = [r"C:\Windows\Fonts\calibrib.ttf", r"C:\Windows\Fonts\seguisb.ttf",
             r"C:\Windows\Fonts\arialbd.ttf"]


def _font(paths, size):
    for p in paths:
        if Path(p).exists():
            try:
                return ImageFont.truetype(p, size)
            except OSError:
                continue
    return ImageFont.load_default()


prs = Presentation(DECK)
SW = int(prs.slide_width / EMU_PER_IN * DPI)
SH = int(prs.slide_height / EMU_PER_IN * DPI)

report = []
for idx, slide in enumerate(prs.slides, start=1):
    img = Image.new("RGB", (SW, SH), "white")
    dr = ImageDraw.Draw(img)
    pics = 0
    txts = 0
    crowded = []

    for shape in slide.shapes:
        l = int((shape.left or 0) / EMU_PER_IN * DPI)
        t = int((shape.top or 0) / EMU_PER_IN * DPI)
        w = int((shape.width or 0) / EMU_PER_IN * DPI)
        h = int((shape.height or 0) / EMU_PER_IN * DPI)

        # solid-fill autoshapes (accent bars, cards, the cream divider block)
        if shape.shape_type == 1 or shape.__class__.__name__ == "AutoShape":
            try:
                if shape.fill.type is not None and shape.fill.type == 1:
                    rgb = shape.fill.fore_color.rgb
                    dr.rectangle([l, t, l + w, t + h], fill=(rgb[0], rgb[1], rgb[2]))
            except Exception:  # noqa: BLE001
                pass

        if shape.__class__.__name__ == "Picture":
            pics += 1
            try:
                blob = shape.image.blob
                im = Image.open(io.BytesIO(blob)).convert("RGB")
                im = im.resize((max(1, w), max(1, h)))
                img.paste(im, (l, t))
            except Exception as exc:  # noqa: BLE001
                dr.rectangle([l, t, l + w, t + h], outline="red", width=3)
                crowded.append(f"picture failed: {exc}")
            continue

        if shape.has_text_frame and shape.text_frame.text.strip():
            txts += 1
            y = t
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                if not text:
                    continue
                sizes = [r.font.size.pt for r in para.runs if r.font.size]
                bold = any(r.font.bold for r in para.runs)
                size_pt = max(sizes) if sizes else 14
                px = max(8, int(size_pt * DPI / 72))
                f = _font(FONT_BOLD if bold else FONT_PATHS, px)
                colour = (0x1A, 0x1A, 0x1A)
                try:
                    if para.runs and para.runs[0].font.color and para.runs[0].font.color.rgb:
                        c = para.runs[0].font.color.rgb
                        colour = (c[0], c[1], c[2])
                except Exception:  # noqa: BLE001
                    pass
                # naive wrap within the box
                words = text.split()
                line = ""
                lines = []
                for word in words:
                    probe = (line + " " + word).strip()
                    if dr.textlength(probe, font=f) > max(20, w) and line:
                        lines.append(line)
                        line = word
                    else:
                        line = probe
                if line:
                    lines.append(line)
                for ln in lines[:14]:
                    dr.text((l, y), ln, font=f, fill=colour)
                    y += int(px * 1.28)
                y += int(px * 0.3)
            if y > t + h + px:  # text ran past its frame
                crowded.append(f"text overflow ~{y - (t + h)}px in box at {l},{t}")

    p = OUT / f"slide_{idx:02d}.png"
    img.save(p)

    # blankness check on the composed slide
    import numpy as np
    arr = np.array(img.convert("L"))
    ink = float((arr < 245).mean() * 100)
    report.append({
        "slide": idx, "png": str(p), "pics": pics, "texts": txts,
        "ink_pct": round(ink, 2), "warnings": crowded,
    })

lines = []
lines.append(f"{len(report)} slides previewed at {SW}x{SH}px -> {OUT}")
lines.append("")
bad = []
for r in report:
    flag = ""
    if r["ink_pct"] < 1.0:
        flag = "  <-- NEARLY BLANK"
        bad.append(r["slide"])
    warn = ("  WARN: " + "; ".join(r["warnings"])) if r["warnings"] else ""
    lines.append(f"slide {r['slide']:02d}  pics {r['pics']}  text-boxes {r['texts']:2d}  "
                 f"ink {r['ink_pct']:5.2f}%{flag}{warn}")
lines.append("")
lines.append("blank slides: " + (", ".join(map(str, bad)) if bad else "none"))
text = "\n".join(lines)
print(text)
(HERE / "_qa" / "preview_report.txt").write_text(text, encoding="utf-8")
