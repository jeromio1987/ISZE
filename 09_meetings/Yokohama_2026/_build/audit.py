"""Full audit of ISZE_GPS_Extended_Target_FY27_V3.3.xlsx.

A. inputs vs repo sources   B. formula pattern per tab   C. independent recompute of every output
D. cross-sheet ties          E. every formula cell carries a value, no errors
Usage: python -I audit.py <formulas.xlsx> <recalculated.xlsx> <repo root>
"""
import csv, json, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L

FX, RC, ROOT = sys.argv[1:4]
wf = openpyxl.load_workbook(FX)
wv = openpyxl.load_workbook(RC, data_only=True)
res = []  # (section, check, ok, detail)


def chk(sec, name, ok, detail=""):
    res.append((sec, name, bool(ok), detail))


def close(a, b, tol=0.5):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


LCV = ["IUK", "UMI", "NORWAY", "SWEDEN", "HUNGARY", "FINLAND", "ROMANIA", "DENMARK", "ICELAND"]
MIX = ["GERMANY", "BELGIUM", "PORTUGAL", "IRELAND", "GREECE", "CYPRUS", "MALTA"]
uio = {}
for r in csv.DictReader(open(f"{ROOT}/01_data/00_derived/uio_by_market_type_year.csv", encoding="utf-8-sig")):
    k = (r["distributor_code"], r["market"], r["model_type"], int(r["year"]))
    uio[k] = uio.get(k, 0) + float(r["uio_units"] or 0)


def U(code, m, cy):
    mk = "BELGIUM" if code == "BELGIUM" else None
    return round(sum(v for (c, x, mm, y), v in uio.items() if c == code and mm == m and y == cy and (mk is None or x == mk)))


bill = {(int(r["FY"]), r["canonical_name"]): float(r["net_gps"])
        for r in csv.DictReader(open(f"{ROOT}/01_data/00_derived/gps_billing_fy_ALLCUSTOMERS.csv", encoding="utf-8-sig"))}
BN = {"IUK": "IUK", "UMI": "UMI", "NORWAY": "Norway", "SWEDEN": "Sverige", "HUNGARY": "AutoWallis", "FINLAND": "Nordic FI",
      "ROMANIA": "ROM Auto", "DENMARK": "Danmark", "ICELAND": "BL Iceland", "GERMANY": "ISD", "BELGIUM": "IBX", "PORTUGAL": "Imotors",
      "IRELAND": "Ireland", "GREECE": "RSA", "CYPRUS": "Loutsios", "MALTA": "GasanZammit"}
bp = json.load(open(f"{ROOT}/01_data/91_build/bp_fy27_reviewed.json"))["by_gps_code"]
COLS = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I"}

# ------------------------------------------------------------------ A. inputs vs sources
for t in LCV:
    ws = wf[t]
    for fy, c in COLS.items():
        chk("A inputs", f"{t} fleet LCV FY{fy % 100} (V19 CY{fy - 1})", ws[f"{c}9"].value == U(t, "LCV", fy - 1), f"{ws[f'{c}9'].value} vs {U(t, 'LCV', fy - 1)}")
    if t == "DENMARK":
        chk("A inputs", "DENMARK has 1 LD truck not counted (LCV-only tab, as in V3.2)", U("DENMARK", "LD", 2026) <= 1, f"LD CY2026 = {U('DENMARK', 'LD', 2026)}")
    for c, fy in (("C", 2025), ("D", 2026)):
        chk("A inputs", f"{t} ISZE billing FY{fy % 100}", close(ws[f"{c}20"].value, bill[(fy, BN[t])], 1), f"{ws[f'{c}20'].value} vs {bill[(fy, BN[t])]:.2f}")
    chk("A inputs", f"{t} BP FY27", close(ws["E20"].value, bp[t], 1), f"{ws['E20'].value} vs {bp[t]}")
for t in MIX:
    ws = wf[t]
    for r, m in ((9, "LD"), (10, "MD"), (11, "F-Series"), (12, "LCV")):
        for fy, c in COLS.items():
            chk("A inputs", f"{t} fleet {m} FY{fy % 100}", ws[f"{c}{r}"].value == U(t, m, fy - 1), f"{ws[f'{c}{r}'].value} vs {U(t, m, fy - 1)}")
    for c, fy in (("C", 2025), ("D", 2026)):
        chk("A inputs", f"{t} ISZE billing FY{fy % 100}", close(ws[f"{c}26"].value, bill[(fy, BN[t])], 1), f"{ws[f'{c}26'].value} vs {bill[(fy, BN[t])]:.2f}")
    chk("A inputs", f"{t} BP FY27", close(ws["E26"].value, bp[t], 1), f"{ws['E26'].value} vs {bp[t]}")
D = wv["Dashboard"]
chk("A inputs", "ITUK BP FY27 (Dashboard C5)", close(D["C5"].value, bp["ITUK"], 1), f"{D['C5'].value} vs {bp['ITUK']}")
chk("A inputs", "UTI BP FY27 (Dashboard C6)", close(D["C6"].value, bp["UTI"], 1), f"{D['C6'].value} vs {bp['UTI']}")
I = wv["Inputs"]
for r, code, m in ((26, "ITUK", "LD"), (27, "ITUK", "F-Series"), (29, "UTI", "LD"), (30, "UTI", "MD")):
    for fy in range(2026, 2032):
        c = L(CI("C") + fy - 2026)
        chk("A inputs", f"Inputs {code} {m} FY{fy % 100}", I[f"{c}{r}"].value == U(code, m, fy - 1), f"{I[f'{c}{r}'].value} vs {U(code, m, fy - 1)}")
for c, v in (("D", 35284444.2378), ("E", 37530000), ("F", 41530410), ("G", 46385705), ("H", 52096000)):
    chk("A inputs", f"Inputs network BP {c}21", close(I[f"{c}21"].value, v, 1), "PL plan row 186 as recorded in gps_full_run.py / handover")
chk("A inputs", "Inputs: LCV 250 / LD 739 / MD 887 / F 887", [I[f"B{r}"].value for r in (4, 5, 6, 7)] == [250, 739, 887, 887])

# ------------------------------------------------------------------ B. formula pattern per tab
def norm(f, t):
    return f.replace(f"{t}!", "TAB!") if isinstance(f, str) else f


for group, ref in ((LCV, "IUK"), (MIX, "GERMANY")):
    rws = wf[ref]
    for t in group:
        ws = wf[t]; diffs = []
        for row in rws.iter_rows():
            for c in row:
                a = c.value; b = ws[c.coordinate].value
                if isinstance(a, str) and a.startswith("=") or isinstance(b, str) and str(b).startswith("="):
                    if norm(a, ref) != norm(b, t):
                        diffs.append(c.coordinate)
        chk("B formulas", f"{t}: every formula identical to {ref}", not diffs, ",".join(diffs[:8]))

# ------------------------------------------------------------------ C. independent recompute
price = {2026: 1.0, 2027: 1.05}
for fy in (2028, 2029, 2030, 2031):
    price[fy] = price[fy - 1] * 1.03
RATE = {"LCV": 250, "LD": 739, "MD": 887, "F-Series": 887}
DR = {"LCV": 448, "LD": 1120, "MD": 1344, "F-Series": 1344}
D2 = wv["Dashboard (2)"]
BASE = {2026: "K", 2027: "U", 2028: "AE", 2029: "AO", 2030: "AY", 2031: "BI"}
ROW = {"IUK": 11, "UMI": 12, "NORWAY": 13, "SWEDEN": 14, "HUNGARY": 15, "FINLAND": 16, "ROMANIA": 17, "DENMARK": 18, "ICELAND": 19,
       "GERMANY": 23, "BELGIUM": 26, "PORTUGAL": 30, "IRELAND": 34, "GREECE": 37, "CYPRUS": 40, "MALTA": 43}
DROW = {"ITUK": 5, "UTI": 6, "IUK": 7, "UMI": 8, "NORWAY": 9, "SWEDEN": 10, "HUNGARY": 11, "FINLAND": 12, "ROMANIA": 13, "DENMARK": 14,
        "ICELAND": 15, "GERMANY": 16, "BELGIUM": 17, "PORTUGAL": 18, "IRELAND": 19, "GREECE": 20, "CYPRUS": 21, "MALTA": 22}


def cell(r, b, off):
    return D2.cell(r, CI(BASE[b]) + off).value if isinstance(b, int) else None


models = lambda t: ["LCV"] if t in LCV else ["LD", "MD", "F-Series", "LCV"]
S27, SL = 0.018, 0.010
CEIL = wv["Inputs"]["B38"].value  # 9.99 when the cap is off
model = {}
for t in LCV + MIX:
    pot = {fy: sum(U(t, m, fy - 1) * RATE[m] * price[fy] for m in models(t)) for fy in range(2026, 2032)}
    dpot = {fy: sum(U(t, m, fy - 1) * DR[m] * price[fy] for m in models(t)) for fy in range(2026, 2032)}
    sh27 = bp[t] / pot[2027]
    st = max(0, min(S27, CEIL - sh27))
    tg = {2027: bp[t] + pot[2027] * st}
    sh = {2027: tg[2027] / pot[2027]}
    for fy in range(2028, 2032):
        p = sh[fy - 1]; sh[fy] = p if p >= CEIL else min(p + SL, CEIL); tg[fy] = sh[fy] * pot[fy]
    s2d = D[f"Q{DROW[t]}"].value
    dt = {}
    if isinstance(s2d, (int, float)):
        dt[2027] = s2d + pot[2027] * st / 0.65
        for fy in range(2028, 2032):
            dt[fy] = dt[fy - 1] * tg[fy] / tg[fy - 1]
    model[t] = (pot, tg, sh, dpot, dt, st)
    r = ROW[t]
    for fy in range(2026, 2032):
        tol = 900 if t == "DENMARK" else 1
        chk("C recompute", f"{t} potential FY{fy % 100}", close(cell(r, fy, 3), pot[fy], tol), f"{cell(r, fy, 3)} vs {pot[fy]:.0f}")
        chk("C recompute", f"{t} dealer potential FY{fy % 100}", close(cell(r, fy, 7), dpot[fy], tol * 2), f"{cell(r, fy, 7)} vs {dpot[fy]:.0f}")
        if fy >= 2027:
            chk("C recompute", f"{t} target FY{fy % 100}", close(cell(r, fy, 5), tg[fy], tol), f"{cell(r, fy, 5)} vs {tg[fy]:.0f}")
            chk("C recompute", f"{t} share FY{fy % 100}", close(cell(r, fy, 4), sh[fy], 0.0005 if t != "DENMARK" else 0.01), f"{cell(r, fy, 4)} vs {sh[fy]:.4f}")
            if dt:
                chk("C recompute", f"{t} dealer target FY{fy % 100}", close(cell(r, fy, 8), dt[fy], tol * 2), f"{cell(r, fy, 8)} vs {dt[fy]:.0f}")
            else:
                chk("C recompute", f"{t} dealer target FY{fy % 100} = 'no KPI data'", cell(r, fy, 8) == "no KPI data", str(cell(r, fy, 8)))
    dr = DROW[t]
    chk("C recompute", f"{t} Dashboard stretch EUR", close(D[f"G{dr}"].value, pot[2027] * st, 1), f"{D[f'G{dr}'].value} vs {pot[2027] * st:.0f}")
    chk("C recompute", f"{t} Dashboard ext target", close(D[f"I{dr}"].value, tg[2027], 1))
    chk("C recompute", f"{t} Dashboard max spend", close(D[f"D{dr}"].value, pot[2027], 900 if t == "DENMARK" else 1))
# ITUK / UTI as presented FY27, extended after
ituk27 = D2["Z9"].value; pot27 = D2["X9"].value
chk("C recompute", "ITUK FY27 target = Dashboard I5 (as presented)", close(ituk27, D["I5"].value, 1))
w6, w7 = D2["W6"].value, D2["W7"].value
sh = D2["Y9"].value
for fy in range(2028, 2032):
    ld = U("ITUK", "LD", fy - 1); f = U("ITUK", "F-Series", fy - 1)
    n35 = round(ld * w6 / (w6 + w7)); n75 = ld - n35
    k = price[fy] / 1.05
    pot = n35 * D2["V6"].value * k + n75 * D2["V7"].value * k + f * D2["V8"].value * k
    sh = sh if sh >= CEIL else min(sh + SL, CEIL)
    chk("C recompute", f"ITUK potential FY{fy % 100}", close(cell(9, fy, 3), pot, 1), f"{cell(9, fy, 3)} vs {pot:.0f}")
    chk("C recompute", f"ITUK target FY{fy % 100}", close(cell(9, fy, 5), pot * sh, 1), f"{cell(9, fy, 5)} vs {pot * sh:.0f}")
sh = D2["Y10"].value
for fy in range(2028, 2032):
    u = U("UTI", "LD", fy - 1) + U("UTI", "MD", fy - 1)
    pot = u * D2["V10"].value * price[fy] / 1.05
    sh = sh if sh >= CEIL else min(sh + SL, CEIL)
    chk("C recompute", f"UTI target FY{fy % 100}", close(cell(10, fy, 5), pot * sh, 1), f"{cell(10, fy, 5)} vs {pot * sh:.0f}")

# ------------------------------------------------------------------ D. ties
W = wv["Waterfall VP"]
chk("D ties", "Waterfall still closes to plan (T5 = 35,284,444)", close(W["T5"].value, 35284444.2378, 1), str(W["T5"].value))
chk("D ties", "Waterfall GPS step P4 = Dashboard stretch total G23", close(W["P4"].value, D["G23"].value, 1))
chk("D ties", "Waterfall Q4 = 952,037 - GPS step", close(W["Q4"].value, 952037.4225 - D["G23"].value, 1), f"{W['Q4'].value}")
chk("D ties", "Dashboard C23 = sum of 18 BP (bp json)", close(D["C23"].value, sum(bp.values()), 2), f"{D['C23'].value} vs {sum(bp.values())}")
chk("D ties", "Dashboard (2) Z44 = Dashboard I23 (FY27 ext targets)", close(D2["Z44"].value, D["I23"].value, 1))
chk("D ties", "Dashboard (2) C..I = Dashboard C..I (18 rows)", all(close(D2[f"{c}{r2}"].value, D[f"{c}{r1}"].value, 0.01)
    for (r1, r2) in zip(DROW.values(), [6, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 24, 27, 31, 35, 38, 41]) for c in "CDEGI"))
chk("D ties", "Inputs D23 outside-18 = 35.28M - sum BP", close(I["D23"].value, 35284444.2378 - sum(bp.values()), 2))
for t in LCV + MIX:
    chk("D ties", f"{t} share <= ceiling after stretch (or flat if already above)", (D[f"H{DROW[t]}"].value <= CEIL + 1e-7) or D[f"F{DROW[t]}"].value == 0)
# known difference, reported not hidden
chk("D ties", "ITUK max spend on MH data on both dashboards", close(D["D5"].value, D2["X9"].value, 1), f"{D['D5'].value:,.0f} vs {D2['X9'].value:,.0f}")
chk("D ties", "UTI max spend on MH data on both dashboards", close(D["D6"].value, D2["X10"].value, 1), f"{D['D6'].value:,.0f} vs {D2['X10'].value:,.0f}")

# ------------------------------------------------------------------ E. values present, no errors
empty, errs, nform = [], [], 0
for ws in wf:
    wsv = wv[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                nform += 1
                v = wsv[c.coordinate].value
                if v is None and ws.title != "Waterfall VP (original)":
                    empty.append(f"{ws.title}!{c.coordinate}")
                if isinstance(v, str) and v.startswith(("#", "Err")):
                    errs.append(f"{ws.title}!{c.coordinate}={v}")
chk("E values", f"all {nform} formula cells carry a value", not empty, ", ".join(empty[:10]))
chk("E values", "no #REF / #DIV / #VALUE / #NAME", not errs, ", ".join(errs[:10]))
chk("E values", "no external links", not wf._external_links)

# ------------------------------------------------------------------ report
by = {}
for s, n, ok, d in res:
    by.setdefault(s, [0, 0]); by[s][0 if ok else 1] += 1
for s, (p, f) in by.items():
    print(f"{s:14} pass {p:4}  fail {f}")
print("\nFAILS / INFO:")
for s, n, ok, d in res:
    if not ok:
        print(f" - [{s}] {n}: {d}")
json.dump([dict(section=s, check=n, ok=ok, detail=d) for s, n, ok, d in res], open("audit_result.json", "w"), indent=1)
