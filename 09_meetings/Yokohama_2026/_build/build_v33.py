"""Build ISZE_GPS_Extended_Target_FY27_V3.3.xlsx from Jerome's V3.2.

Changes (Jerome, 23-10-2026 answers):
  * LCV rate: EUR 250 ISZE-net FY26 base -> x1.05 FY27 (262.50) -> x1.03/yr FY28-31.
    Replaces MH's 291.2 x 1.16 x 0.65 = 219.56 chain (FX and margin applied twice).
  * UIO stays V19 detail rows, shifted one column (FY = CY-1). FY28..FY31 <- CY2027F..CY2030F.
  * Ceiling per distributor: stretch never pushes GPS share above the ceiling (100%).
    A DB already at/above it gets no stretch; its share is held flat.
  * ITUK / UTI kept as presented (V4 basis) for FY27; extended FY28-31 on the same rule.
  * FY25/FY26 ISZE actuals refreshed from 01_data/00_derived/gps_billing_fy_ALLCUSTOMERS.csv.
  * Dashboard (2): external [1] links made internal, Malta LCV rate bug fixed, Greece/Cyprus
    MD vehicles added back, FY2028-FY2031 blocks filled, network BP comparison added.

Usage: python -I build_v33.py <in_v32.xlsx> <uio_by_market_type_year.csv> <gps_billing_fy.csv> <out.xlsx>
"""
import csv
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

SRC, UIO_CSV, BILL_CSV, OUT = sys.argv[1:5]

LCV_TABS = ["IUK", "UMI", "NORWAY", "SWEDEN", "HUNGARY", "FINLAND", "ROMANIA", "DENMARK", "ICELAND"]
MIX_TABS = ["GERMANY", "BELGIUM", "PORTUGAL", "IRELAND", "GREECE", "CYPRUS", "MALTA"]
# tab -> (uio distributor_code, market filter) ; Belgium = Belgian parc only (scope note kept)
UIO_KEY = {t: (t, None) for t in LCV_TABS + MIX_TABS}
UIO_KEY["BELGIUM"] = ("BELGIUM", "BELGIUM")
BILL_NAME = {"IUK": "IUK", "UMI": "UMI", "NORWAY": "Norway", "SWEDEN": "Sverige", "HUNGARY": "AutoWallis",
             "FINLAND": "Nordic FI", "ROMANIA": "ROM Auto", "DENMARK": "Danmark", "ICELAND": "BL Iceland",
             "GERMANY": "ISD", "BELGIUM": "IBX", "PORTUGAL": "Imotors", "IRELAND": "Ireland",
             "GREECE": "RSA",  # payer 66 = Greece (labelled RSA in the billing file)
             "CYPRUS": "Loutsios", "MALTA": "GasanZammit", "ITUK": "ITUK", "UTI": "UTI"}

BLUE = Font(color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(italic=True, color="595959", size=9)

# ---------------------------------------------------------------- source data
uio = {}  # (code, market or None, model_type, year) -> units
with open(UIO_CSV, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        y = int(r["year"])
        if 2024 <= y <= 2030:
            for mk in (None, r["market"]):
                k = (r["distributor_code"], mk, r["model_type"], y)
                uio[k] = uio.get(k, 0) + float(r["uio_units"] or 0)


def U(tab, model, cy):
    code, mk = UIO_KEY[tab]
    return int(round(uio.get((code, mk, model, cy), 0)))


bill = {}
with open(BILL_CSV, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        bill[(int(r["FY"]), r["canonical_name"])] = float(r["net_gps"])

wb = openpyxl.load_workbook(SRC)
log = []

# ---------------------------------------------------------------- inputs (Dashboard (2) row 1)
D2 = wb["Dashboard (2)"]
INP = {"stretch": "N1", "ceiling": "Q1", "price": "T1"}
for lab_cell, lab, cell, val, fmt in [("K1", "GPS stretch FY27 (pt)", "N1", 0.018, "0.0%"),
                                      ("K2", "GPS stretch FY28-31 / yr (pt)", "N2", 0.010, "0.0%"),
                                      ("O1", "Share ceiling", "Q1", 1.0, "0%"),
                                      ("R1", "Price FY28-31 / yr", "T1", 0.03, "0.0%")]:
    D2[lab_cell] = lab
    D2[lab_cell].font = Font(bold=True)
    D2[cell] = val
    D2[cell].number_format = fmt
    D2[cell].font = BLUE
    D2[cell].fill = INPUT_FILL
D2["V2"] = "FY27 stretch must stay >= ~1.1% to cover the EUR 952K the waterfall needs (Dashboard C25)."
D2["V2"].font = NOTE
D2["V1"] = ("Blue/yellow = input. Stretch is capped per distributor so GPS share never passes the ceiling; "
            "a distributor already above it keeps its share flat (no stretch).")
D2["V1"].font = NOTE
D2.row_dimensions[2].hidden = False  # row 2 holds the FY28-31 input
STRETCH = "'Dashboard (2)'!$N$1"
STRETCH_LATER = "'Dashboard (2)'!$N$2"
CEIL = "'Dashboard (2)'!$Q$1"
PRICE = "'Dashboard (2)'!$T$1"

YEARS = [2028, 2029, 2030, 2031]   # FY
YCOL = {2028: "F", 2029: "G", 2030: "H", 2031: "I"}  # tab columns for FY28..FY31


def style_like(dst, src):
    dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format


# ---------------------------------------------------------------- LCV tabs
for t in LCV_TABS:
    ws = wb[t]
    # rate chain -> EUR 250 base
    ws["A14"] = "ISZE-net EUR/unit (= DB cost)"
    ws["B13"] = "FY26 base"
    ws["B14"] = 250
    ws["B14"].font = BLUE
    ws["C13"] = None; ws["C14"] = None
    ws["D13"] = "FY26"
    ws["D14"] = "=B14"
    ws["E14"] = "=D14*1.05"
    for i, fy in enumerate(YEARS):
        c = YCOL[fy]; prev = L(CI(c) - 1)
        ws[f"{c}13"] = f"FY{fy % 100} (+price)"
        ws[f"{c}14"] = f"={prev}14*(1+{PRICE})"
        style_like(ws[f"{c}13"], ws["E13"]); style_like(ws[f"{c}14"], ws["E14"])
    ws["A15"] = ("LCV rate: EUR 250 ISZE-net FY26 base (same rate as the fleet study), +5% FY27 = 262.50, "
                 "then + price input per year. Replaces MH chain 291.2 x 1.16 x 0.65 = 219.56 (FX and margin applied twice).")
    # UIO FY28..31
    for fy in YEARS:
        c = YCOL[fy]
        ws[f"{c}8"] = f"FY{fy % 100} (CY{fy - 1}F)"
        style_like(ws[f"{c}8"], ws["E8"])
        ws[f"{c}9"] = U(t, "LCV", fy - 1)
        style_like(ws[f"{c}9"], ws["E9"])
        ws[f"{c}10"] = f"={c}9"
        style_like(ws[f"{c}10"], ws["E10"])
    chk = U(t, "LCV", 2026)
    if ws["E9"].value != chk:
        log.append(f"WARN {t} E9 {ws['E9'].value} != UIO CY2026 {chk}")
    # max spend FY28..31
    for fy in YEARS:
        c = YCOL[fy]
        ws[f"{c}18"] = f"FY{fy % 100}"
        style_like(ws[f"{c}18"], ws["E18"])
        ws[f"{c}19"] = f"={c}9*{c}14"
        style_like(ws[f"{c}19"], ws["E19"])
    ws["F20"] = None
    ws["A27"] = "FY28-FY31 max spend (row 19, cols F:I) feeds Dashboard (2). FY27 actual cell E20 = BP reviewed 24-04-2026."
    ws["A27"].font = NOTE
    # actuals FY25 / FY26 from GPS-audited billing
    for col, fy in (("C", 2025), ("D", 2026)):
        v = bill.get((fy, BILL_NAME[t]))
        if v is not None:
            old = ws[f"{col}20"].value
            if old is None or abs(float(old) - v) > 1:
                log.append(f"{t} FY{fy % 100} actual {old} -> {v:,.0f}")
            ws[f"{col}20"] = round(v, 2)
    # capped stretch
    ws["E22"] = f"=MAX(0,MIN({STRETCH},{CEIL}-E21))"
    ws["A22"] = "Stretch (%) = global input, capped at ceiling"

# ---------------------------------------------------------------- mixed tabs
for t in MIX_TABS:
    ws = wb[t]
    ws["C20"] = 250
    ws["C20"].font = BLUE
    ws["A20"] = "LCV / D-Max  · EUR 250 ISZE-net base"
    ws["A21"] = ("CV rates held at v9 (739 / 887) +5% FY27, per Jerome. LCV: EUR 250 base +5% (262.50). "
                 "FY28-FY31 = previous year x (1 + price input).")
    for i, fy in enumerate(YEARS):
        c = L(CI("E") + i)  # rates FY28..31 in E..H, rows 16-20
        prev = L(CI(c) - 1)
        ws[f"{c}16"] = f"FY{fy % 100}"
        style_like(ws[f"{c}16"], ws["D16"])
        for r in range(17, 21):
            ws[f"{c}{r}"] = f"={prev}{r}*(1+{PRICE})"
            style_like(ws[f"{c}{r}"], ws[f"D{r}"])
    models = {9: "LD", 10: "MD", 11: "F-Series", 12: "LCV"}
    for r, m in models.items():
        if (ws[f"E{r}"].value or 0) != U(t, m, 2026):
            log.append(f"WARN {t} E{r} {ws[f'E{r}'].value} != UIO CY2026 {U(t, m, 2026)}")
    for fy in YEARS:
        c = YCOL[fy]
        ws[f"{c}8"] = f"FY{fy % 100} (CY{fy - 1}F)"
        style_like(ws[f"{c}8"], ws["E8"])
        for r, m in models.items():
            ws[f"{c}{r}"] = U(t, m, fy - 1)
            style_like(ws[f"{c}{r}"], ws[f"E{r}"])
        ws[f"{c}13"] = f"=SUM({c}9:{c}12)"
        style_like(ws[f"{c}13"], ws["E13"])
    # max spend FY28..31 in F25:I25 (rates in E..H)
    for i, fy in enumerate(YEARS):
        c = YCOL[fy]; rc = L(CI("E") + i)
        ws[f"{c}24"] = f"FY{fy % 100}"
        style_like(ws[f"{c}24"], ws["E24"])
        ws[f"{c}25"] = f"={c}9*{rc}17+{c}10*{rc}18+{c}11*{rc}19+{c}12*{rc}20"
        style_like(ws[f"{c}25"], ws["E25"])
    ws["F26"] = None
    ws["A33"] = "FY28-FY31 max spend (row 25, cols F:I) feeds Dashboard (2). FY27 actual cell E26 = BP reviewed 24-04-2026."
    ws["A33"].font = NOTE
    for col, fy in (("C", 2025), ("D", 2026)):
        v = bill.get((fy, BILL_NAME[t]))
        if v is not None:
            old = ws[f"{col}26"].value
            if old is None or abs(float(old) - v) > 1:
                log.append(f"{t} FY{fy % 100} actual {old} -> {v:,.0f}")
            ws[f"{col}26"] = round(v, 2)
    ws["E28"] = f"=MAX(0,MIN({STRETCH},{CEIL}-E27))"
    ws["A28"] = "Stretch (%) = global input, capped at ceiling"

# ---------------------------------------------------------------- Dashboard (English): stretch label
DB = wb["Dashboard"]
DB["J12"] = "Above ceiling at BP — no stretch"
DB["J19"] = "Stretch capped at 100% share"
DB["J17"] = "Belgian vehicles only (IBX also serves NL/PL/LU)"

# ---------------------------------------------------------------- Dashboard (2)
# 1. internal refs, link C..I to Dashboard
for row in D2.iter_rows():
    for c in row:
        if isinstance(c.value, str) and "[1]" in c.value:
            c.value = c.value.replace("[1]", "")
wb._external_links = []

DASH_ROW = {"ITUK": 5, "UTI": 6, "IUK": 7, "UMI": 8, "NORWAY": 9, "SWEDEN": 10, "HUNGARY": 11, "FINLAND": 12,
            "ROMANIA": 13, "DENMARK": 14, "ICELAND": 15, "GERMANY": 16, "BELGIUM": 17, "PORTUGAL": 18,
            "IRELAND": 19, "GREECE": 20, "CYPRUS": 21, "MALTA": 22}
D2_ROW = {"ITUK": 6, "UTI": 10, "IUK": 11, "UMI": 12, "NORWAY": 13, "SWEDEN": 14, "HUNGARY": 15, "FINLAND": 16,
          "ROMANIA": 17, "DENMARK": 18, "ICELAND": 19, "GERMANY": 20, "BELGIUM": 24, "PORTUGAL": 27,
          "IRELAND": 31, "GREECE": 35, "CYPRUS": 38, "MALTA": 41}
for k, r in D2_ROW.items():
    for col in "CDEFGHI":
        D2[f"{col}{r}"] = f"=Dashboard!{col}{DASH_ROW[k]}"

# 2. segment rows of mixed DBs: list of (tab rate row, tab uio row)
SEG = {
    "GERMANY": [(20, [(17, 9)]), (21, [(18, 10)]), (22, [(20, 12)])],
    "BELGIUM": [(24, [(17, 9)]), (25, [(20, 12)])],
    "PORTUGAL": [(27, [(17, 9)]), (28, [(18, 10), (19, 11)]), (29, [(20, 12)])],
    "IRELAND": [(31, [(17, 9)]), (32, [(18, 10), (19, 11)]), (33, [(20, 12)])],
    "GREECE": [(35, [(17, 9), (18, 10)]), (36, [(20, 12)])],
    "CYPRUS": [(38, [(17, 9), (18, 10)]), (39, [(20, 12)])],
    "MALTA": [(41, [(17, 9)]), (42, [(20, 12)])],
}
TOTAL_ROW = {"GERMANY": 23, "BELGIUM": 26, "PORTUGAL": 30, "IRELAND": 34, "GREECE": 37, "CYPRUS": 40, "MALTA": 43}
LCV_ROW = {t: D2_ROW[t] for t in LCV_TABS}

# block base columns: FY26 K, FY27 U, FY28 AE, FY29 AO, FY30 AY, FY31 BI
BASE = {2026: "K", 2027: "U", 2028: "AE", 2029: "AO", 2030: "AY", 2031: "BI"}
# mixed-tab rate column and UIO column per FY
MIX_RATE = {2026: "C", 2027: "D", 2028: "E", 2029: "F", 2030: "G", 2031: "H"}
MIX_UIO = {2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I"}
LCV_RATE = {2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I"}
LCV_UIO = MIX_UIO


def off(fy, n):
    return L(CI(BASE[fy]) + n)


def fmt_from(dst, src_cell):
    style_like(D2[dst], D2[src_cell])


# template formats from FY27 block
T_RATE, T_UIO, T_POT, T_SHARE, T_TGT = "V11", "W11", "X11", "Y11", "Z11"

for fy in [2026, 2027] + YEARS:
    lab, rate, uioc, pot = off(fy, 0), off(fy, 1), off(fy, 2), off(fy, 3)
    share, tgt = off(fy, 4), off(fy, 5)
    # LCV single-row DBs
    for t, r in LCV_ROW.items():
        D2[f"{lab}{r}"] = "LCV"
        D2[f"{rate}{r}"] = f"={t}!{LCV_RATE[fy]}14"
        D2[f"{uioc}{r}"] = f"={t}!{LCV_UIO[fy]}9"
        D2[f"{pot}{r}"] = f"={rate}{r}*{uioc}{r}"
        for c, tmpl in ((rate, T_RATE), (uioc, T_UIO), (pot, T_POT)):
            fmt_from(f"{c}{r}", tmpl)
    # mixed DBs, segment rows
    for t, segs in SEG.items():
        for r, parts in segs:
            if len(parts) == 1:
                rr, ur = parts[0]
                D2[f"{rate}{r}"] = f"={t}!{MIX_RATE[fy]}{rr}"
                D2[f"{uioc}{r}"] = f"={t}!{MIX_UIO[fy]}{ur}"
                D2[f"{pot}{r}"] = f"={rate}{r}*{uioc}{r}"
            else:
                D2[f"{uioc}{r}"] = "=" + "+".join(f"{t}!{MIX_UIO[fy]}{ur}" for _, ur in parts)
                D2[f"{pot}{r}"] = "=" + "+".join(f"{t}!{MIX_RATE[fy]}{rr}*{t}!{MIX_UIO[fy]}{ur}" for rr, ur in parts)
                D2[f"{rate}{r}"] = f"=IFERROR({pot}{r}/{uioc}{r},0)"
            if fy >= 2028:
                D2[f"{lab}{r}"] = D2[f"U{r}"].value
            for c, tmpl in ((rate, T_RATE), (uioc, T_UIO), (pot, T_POT)):
                fmt_from(f"{c}{r}", tmpl)
        tr = TOTAL_ROW[t]
        first, last = segs[0][0], segs[-1][0]
        D2[f"{uioc}{tr}"] = f"=SUM({uioc}{first}:{uioc}{last})"
        D2[f"{pot}{tr}"] = f"=SUM({pot}{first}:{pot}{last})"
        for c, tmpl in ((uioc, T_UIO), (pot, T_POT)):
            fmt_from(f"{c}{tr}", tmpl)
D2["K35"] = "CV N+MD"; D2["U35"] = "CV N+MD"; D2["K38"] = "CV N+MD"; D2["U38"] = "CV N+MD"
for b in ("AE", "AO", "AY", "BI"):
    D2[f"{b}35"] = "CV N+MD"; D2[f"{b}38"] = "CV N+MD"

# total-level row per DB (where share / target sit)
SHARE_ROW = dict(LCV_ROW)
SHARE_ROW.update(TOTAL_ROW)

# FY26 actual (O) and FY27 target (Z) live
for t in LCV_TABS:
    r = SHARE_ROW[t]
    D2[f"O{r}"] = f"={t}!D20"
    D2[f"P{r}"] = f"=O{r}/N{r}"
    D2[f"Z{r}"] = f"={t}!E24"
    D2[f"Y{r}"] = f"=Z{r}/X{r}"
for t in MIX_TABS:
    r = SHARE_ROW[t]
    D2[f"O{r}"] = f"={t}!D26"
    D2[f"P{r}"] = f"=O{r}/N{r}"
    D2[f"Z{r}"] = f"={t}!E30"
    D2[f"Y{r}"] = f"=Z{r}/X{r}"
    for c, tmpl in (("O", "O11"), ("P", "P11"), ("Y", T_SHARE), ("Z", T_TGT)):
        fmt_from(f"{c}{r}", tmpl)
D2["O9"] = round(bill[(2026, "ITUK")], 2)
D2["O10"] = round(bill[(2026, "UTI")], 2)
D2["Z9"] = "=Dashboard!I5"
D2["Z10"] = "=Dashboard!I6"

# ITUK (rows 6-8 N35/N75/F, total 9) and UTI (row 10): FY28+ on V4 basis, CY = FY-1
ituk_ld = {fy: U_ for fy, U_ in ((fy, int(round(uio.get(("ITUK", None, "LD", fy - 1), 0)))) for fy in YEARS)}
ituk_f = {fy: int(round(uio.get(("ITUK", None, "F-Series", fy - 1), 0))) for fy in YEARS}
uti = {fy: int(round(uio.get(("UTI", None, "LD", fy - 1), 0) + uio.get(("UTI", None, "MD", fy - 1), 0))) for fy in YEARS}
for fy in YEARS:
    prevb = BASE[fy - 1]
    lab, rate, uioc, pot = off(fy, 0), off(fy, 1), off(fy, 2), off(fy, 3)
    pr = L(CI(prevb) + 1)
    drate, dpot, dtgt, dsh = off(fy, 6), off(fy, 7), off(fy, 8), off(fy, 9)
    pdrate, pdtgt, pdsh = L(CI(prevb) + 6), L(CI(prevb) + 8), L(CI(prevb) + 9)
    n35_share = "($W$6/($W$6+$W$7))"
    for r, lbl in ((6, "N35"), (7, "N75"), (8, "F")):
        D2[f"{lab}{r}"] = lbl
        D2[f"{rate}{r}"] = f"={pr}{r}*(1+{PRICE})"
        if r == 6:
            D2[f"{uioc}{r}"] = f"=ROUND({ituk_ld[fy]}*{n35_share},0)"
        elif r == 7:
            D2[f"{uioc}{r}"] = f"={ituk_ld[fy]}-{uioc}6"
        else:
            D2[f"{uioc}{r}"] = ituk_f[fy]
        D2[f"{pot}{r}"] = f"={rate}{r}*{uioc}{r}"
        D2[f"{drate}{r}"] = f"={pdrate}{r}*(1+{PRICE})"
        D2[f"{dpot}{r}"] = f"={drate}{r}*{uioc}{r}"
        for c, tmpl in ((rate, "V6"), (uioc, "W6"), (pot, "X6"), (drate, "AA6"), (dpot, "AB6")):
            fmt_from(f"{c}{r}", tmpl)
    D2[f"{uioc}9"] = f"=SUM({uioc}6:{uioc}8)"
    D2[f"{pot}9"] = f"=SUM({pot}6:{pot}8)"
    D2[f"{dpot}9"] = f"=SUM({dpot}6:{dpot}8)"
    psh_d = f"{pdsh}9"
    D2[f"{dsh}9"] = f"=IF({psh_d}>={CEIL},{psh_d},MIN({psh_d}+{STRETCH_LATER},{CEIL}))"
    D2[f"{dtgt}9"] = f"={dpot}9*{dsh}9"
    for c, tmpl in ((uioc, "W9"), (pot, "X9"), (dpot, "AB9"), (dsh, "AD9"), (dtgt, "AC9")):
        fmt_from(f"{c}9", tmpl)
    # UTI single row
    D2[f"{lab}10"] = "CV"
    D2[f"{rate}10"] = f"={pr}10*(1+{PRICE})"
    D2[f"{uioc}10"] = uti[fy]
    D2[f"{pot}10"] = f"={rate}10*{uioc}10"
    for c, tmpl in ((rate, "V10"), (uioc, "W10"), (pot, "X10")):
        fmt_from(f"{c}10", tmpl)
# ITUK FY27 dealer share was 0.786+0.018: make it follow the inputs
D2["AC9"] = f"=AB9*AD9"
D2["AD9"] = f"=IF(T9>={CEIL},T9,MIN(T9+{STRETCH},{CEIL}))"

# share / target FY28+ for every DB (total-level rows incl. ITUK row 9, UTI row 10)
ALL_SHARE_ROWS = [9, 10] + [SHARE_ROW[t] for t in LCV_TABS + MIX_TABS]
for fy in YEARS:
    prevb = BASE[fy - 1]
    psh = L(CI(prevb) + 4)
    pot, share, tgt = off(fy, 3), off(fy, 4), off(fy, 5)
    for r in ALL_SHARE_ROWS:
        D2[f"{share}{r}"] = f"=IF({psh}{r}>={CEIL},{psh}{r},MIN({psh}{r}+{STRETCH_LATER},{CEIL}))"
        D2[f"{tgt}{r}"] = f"={pot}{r}*{share}{r}"
        fmt_from(f"{share}{r}", T_SHARE)
        fmt_from(f"{tgt}{r}", T_TGT)

# 3. TOTAL row 44 + network BP comparison (rows 45-47) per year
sum_rows = ALL_SHARE_ROWS
BP_NET = {2027: 35284444.2378, 2028: 37530000, 2029: 41530410, 2030: 46385705, 2031: 52096000}
OUTSIDE = 35284444.2378 - 33222361.956  # CIS / MIDI / KIMA / IMG: BP but no UIO -> held flat
for fy in [2026, 2027] + YEARS:
    lab, uioc, pot, share, tgt = off(fy, 0), off(fy, 2), off(fy, 3), off(fy, 4), off(fy, 5)
    D2[f"{lab}44"] = "TOTAL 18"
    D2[f"{lab}44"].font = Font(bold=True)
    D2[f"{uioc}44"] = "=" + "+".join(f"{uioc}{r}" for r in sum_rows)
    D2[f"{pot}44"] = "=" + "+".join(f"{pot}{r}" for r in sum_rows)
    for c in (uioc, pot):
        fmt_from(f"{c}44", T_POT); D2[f"{c}44"].font = Font(bold=True)
    if fy == 2026:
        D2["O44"] = "=" + "+".join(f"O{r}" for r in sum_rows)
        D2["P44"] = "=O44/N44"
        for c, tmpl in (("O", "O11"), ("P", "P11")):
            fmt_from(f"{c}44", tmpl); D2[f"{c}44"].font = Font(bold=True)
        D2["L44"] = None
        D2["N45"] = None
        continue
    D2[f"{tgt}44"] = "=" + "+".join(f"{tgt}{r}" for r in sum_rows)
    D2[f"{share}44"] = f"={tgt}44/{pot}44"
    fmt_from(f"{share}44", T_SHARE); fmt_from(f"{tgt}44", T_TGT)
    D2[f"{share}44"].font = Font(bold=True); D2[f"{tgt}44"].font = Font(bold=True)
    D2[f"{lab}45"] = "+ outside 18 (CIS/MIDI/KIMA/IMG, flat)"
    D2[f"{tgt}45"] = round(OUTSIDE, 2)
    D2[f"{lab}46"] = "Network BP (PL plan row 186)"
    D2[f"{tgt}46"] = BP_NET[fy]
    D2[f"{lab}47"] = "Gap: target - BP (+ = above plan)"
    D2[f"{tgt}47"] = f"={tgt}44+{tgt}45-{tgt}46"
    for rr in (45, 46, 47):
        fmt_from(f"{tgt}{rr}", T_TGT)
        D2[f"{lab}{rr}"].font = NOTE if rr != 47 else Font(bold=True)
    D2[f"{tgt}47"].font = Font(bold=True)
D2["N44"] = "=" + "+".join(f"N{r}" for r in sum_rows)
D2["M44"] = "=" + "+".join(f"M{r}" for r in sum_rows)

# headers for FY28..31 blocks already exist (row 3-5); copy widths
for fy in YEARS:
    for n in range(10):
        src = L(CI("U") + n); dst = off(fy, n)
        D2.column_dimensions[dst].width = D2.column_dimensions[src].width or 12

# WATCH notes refreshed
D2["A54"] = ("• CEILING (v3.3): stretch is capped per distributor at the share ceiling (input Q1). "
             "FINLAND is above 100% at BP (104.9%) — no stretch, share held flat. It is the BP, not the model.")
D2["A53"] = ("• V4 IUK max spend = 50,202 (CY2027F) x EUR250 = 12,550,500. This file (v3.3): 49,280 (CY2026F) x 262.50 = 12,936,000.")
DB["A33"] = D2["A54"].value
DB["A32"] = D2["A53"].value

# ---------------------------------------------------------------- Method changelog
M = wb["Method"]
lines = [
    "",
    "VERSION 3.3 — 23-10-2026 (Jerome's answers in the Claude session)",
    "1. LCV rate: EUR 250 ISZE-net FY26 base (fleet-study rate) -> x1.05 FY27 = 262.50 -> x(1+price) FY28-FY31. "
    "MH's 291.2 x 1.16 x 0.65 = 219.56 chain retired (FX and margin applied twice).",
    "2. UIO: V19 detail rows kept, one-column shift (FY = CY-1). FY28..FY31 <- CY2027F..CY2030F "
    "(01_data/00_derived/uio_by_market_type_year.csv). Corrected CY2026 file (30-07, ~9% lower LCV) NOT applied — V20 not adopted.",
    "3. Ceiling per distributor (input Q1 on Dashboard (2), 100%): stretch = MAX(0, MIN(stretch input, ceiling - GPS share)). "
    "Share already at/above the ceiling is held flat. FY28-31: share = previous share + FY28-31 stretch input (N2, 1.0%), capped the same way.",
    "   Why two inputs: FY27 at 1.0% gives EUR 923K stretch, EUR 29K short of the EUR 952K the waterfall needs, so FY27 stays 1.8%. "
    "FY28-31 at 1.8% puts the 18 + outside customers EUR 3.3-4.2M a year above the network plan and drives Iceland to 99.8% / Ireland to the ceiling; "
    "1.0% stays above plan every year (FY31 +0.9M); 0.8% gives FY31 +0.2M only.",
    "4. ITUK / UTI: FY27 kept as presented (V4 basis). FY28-31 extended with the same share rule on their own fleet rows.",
    "5. FY25 / FY26 ISZE actuals refreshed from 00_derived/gps_billing_fy_ALLCUSTOMERS.csv (net_gps). Greece = payer 66 ('RSA' label).",
    "6. Dashboard (2): links to the old Sept workbook made internal; Malta LCV row used the F-Series rate (887) — fixed; "
    "Greece and Cyprus medium-duty vehicles were left out of the potential — added; FY2028-FY2031 filled; network BP comparison rows 45-47.",
    "7. Network BP FY28-FY31 = PL plan row 186: 37.53M / 41.53M / 46.39M / 52.10M. Customers outside the 18 (EUR 2.06M) held flat.",
    "8. Dealer level (販社) only kept for ITUK, as before (dealer block dropped 17-07-2026).",
]
start = M.max_row + 2
for i, s in enumerate(lines):
    M[f"A{start + i}"] = s
M[f"A{start + 1}"].font = Font(bold=True)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("\n".join(log))
print("saved", OUT)
