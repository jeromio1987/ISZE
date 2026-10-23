"""Independent recompute of the 'UIO V20' sheet from the V20 change list, V19 fleet and model rules."""
import csv, sys, openpyxl
from openpyxl.utils import column_index_from_string as CI
RC, ROOT = sys.argv[1:3]
wv = openpyxl.load_workbook(RC, data_only=True)
V = wv["UIO V20"]; D2 = wv["Dashboard (2)"]
last = {}
for r in csv.DictReader(open(f"{ROOT}/06_uio/_CURRENT/UIO_V20_CHANGELOG_2026-09-22c.csv", encoding="utf-8-sig")):
    last[(r["country"], r["model_type"], int(r["year"]))] = float(r["new_value_v20"])
u19 = {}
for r in csv.DictReader(open(f"{ROOT}/01_data/00_derived/uio_by_market_type_year.csv", encoding="utf-8-sig")):
    if r["distributor_code"] == "BELGIUM" and r["market"] != "BELGIUM":
        continue
    k = (r["distributor_code"], r["model_type"], int(r["year"]))
    u19[k] = u19.get(k, 0) + float(r["uio_units"] or 0)
C = {"IUK": "U KING", "UMI": "ISRAEL", "PORTUGAL": "PORTUGL"}
LCV = ["IUK", "UMI", "NORWAY", "SWEDEN", "HUNGARY", "FINLAND", "ROMANIA", "DENMARK", "ICELAND"]
MIX = ["GERMANY", "BELGIUM", "PORTUGAL", "IRELAND", "GREECE", "CYPRUS", "MALTA"]
RATE = {"LCV": 250, "LD": 739, "MD": 887, "F-Series": 887}
price = {2026: 1.0, 2027: 1.05}
for fy in range(2028, 2032):
    price[fy] = price[fy - 1] * 1.03
ROW2 = {"IUK": 11, "UMI": 12, "NORWAY": 13, "SWEDEN": 14, "HUNGARY": 15, "FINLAND": 16, "ROMANIA": 17, "DENMARK": 18, "ICELAND": 19,
        "GERMANY": 23, "BELGIUM": 26, "PORTUGAL": 30, "IRELAND": 34, "GREECE": 37, "CYPRUS": 40, "MALTA": 43}


def fleet20(t, m, cy):
    if cy == 2025:
        return round(u19.get((t, m, cy), 0))
    return last.get((C.get(t, t), m, cy), round(u19.get((t, m, cy), 0)))


# locate rows on the sheet
rows = {V.cell(r, 1).value: r for r in range(60, 125) if V.cell(r, 1).value in LCV + MIX + ["ITUK", "UTI"]}
pot_rows, share_rows, tgt_rows = {}, {}, {}
for r in range(55, 125):
    n = V.cell(r, 1).value
    if n in LCV + MIX + ["ITUK", "UTI"]:
        (pot_rows if n not in pot_rows else share_rows if n not in share_rows else tgt_rows)[n] = r
ok = bad = 0
tot = {fy: 0 for fy in range(2027, 2032)}
for t in LCV + MIX:
    models = ["LCV"] if t in LCV else ["LD", "MD", "F-Series", "LCV"]
    pot = {fy: sum(fleet20(t, m, fy - 1) * RATE[m] for m in models) * price[fy] for fy in range(2026, 2032)}
    tg27 = D2.cell(ROW2[t], CI("Z")).value
    sh = tg27 / pot[2027]; tg = {2027: tg27}
    for fy in range(2028, 2032):
        sh = min(sh + 0.010, 9.99) if sh < 9.99 else sh
        tg[fy] = sh * pot[fy]
    for fy in range(2026, 2032):
        cell = V.cell(pot_rows[t], 3 + fy - 2026).value
        good = abs(cell - pot[fy]) < 1
        ok += good; bad += not good
        if not good: print("POT", t, fy, cell, round(pot[fy]))
    for fy in range(2027, 2032):
        cell = V.cell(tgt_rows[t], 3 + fy - 2026).value
        good = abs(cell - tg[fy]) < 1
        ok += good; bad += not good
        if not good: print("TGT", t, fy, cell, round(tg[fy]))
        tot[fy] += tg[fy]
for n, col in (("ITUK", 9), ("UTI", 10)):
    for fy in range(2027, 2032):
        tot[fy] += D2.cell(col, CI({2027: "U", 2028: "AE", 2029: "AO", 2030: "AY", 2031: "BI"}[fy]) + 5).value
plan = {2027: 35284444.2378, 2028: 37530000, 2029: 41530410, 2030: 46385705, 2031: 52096000}
outside = 35284444.2378 - 33222361.956
summ = {V.cell(r, 1).value: r for r in range(120, V.max_row + 1)}
r_ag = summ["Against plan, UIO V20"]
for fy in range(2027, 2032):
    exp = tot[fy] + outside - plan[fy]; cell = V.cell(r_ag, 3 + fy - 2026).value
    good = abs(cell - exp) < 2
    ok += good; bad += not good
    print(f"FY{fy % 100} against plan V20: sheet {cell:,.0f}  independent {exp:,.0f}  {'OK' if good else 'DIFF'}")
print("UIO V20 recompute: ok", ok, "bad", bad)
