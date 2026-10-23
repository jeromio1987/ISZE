import csv, openpyxl
from openpyxl.utils import column_index_from_string as CI
uio = {}
for r in csv.DictReader(open('/home/user/ISZE/01_data/00_derived/uio_by_market_type_year.csv', encoding='utf-8-sig')):
    k = (r['distributor_code'], r['market'], r['model_type'], int(r['year']))
    uio[k] = uio.get(k, 0) + float(r['uio_units'] or 0)
def U(code, m, cy, market=None):
    return sum(v for (c, mk, mm, y), v in uio.items() if c == code and mm == m and y == cy and (market is None or mk == market))
rate = {'LCV': 250, 'LD': 739, 'MD': 887, 'F-Series': 887}
price = {2026: 1, 2027: 1.05}
for fy in (2028, 2029, 2030, 2031): price[fy] = price[fy - 1] * 1.03
dbs = ['IUK','UMI','NORWAY','SWEDEN','HUNGARY','FINLAND','ROMANIA','DENMARK','ICELAND','GERMANY','BELGIUM','PORTUGAL','IRELAND','GREECE','CYPRUS','MALTA']
wb = openpyxl.load_workbook('out/recalc/ISZE_GPS_Extended_Target_FY27_V3.3.xlsx', data_only=True)
D = wb['Dashboard']; D2 = wb['Dashboard (2)']
bp = {D[f'A{r}'].value: D[f'C{r}'].value for r in range(5, 23)}
base = {'K': 2026, 'U': 2027, 'AE': 2028, 'AO': 2029, 'AY': 2030, 'BI': 2031}
pot = {}; share = {}; tgt = {}
for d in dbs:
    for fy in range(2026, 2032):
        mk = 'BELGIUM' if d == 'BELGIUM' else None
        pot[(d, fy)] = sum(U(d, m, fy - 1, mk) * rate[m] * price[fy] for m in rate)
    s27 = bp[d] / pot[(d, 2027)]
    st = max(0, min(0.018, 1 - s27)); share[(d, 2027)] = s27 + st; tgt[(d, 2027)] = bp[d] + pot[(d, 2027)] * st
    for fy in range(2028, 2032):
        p = share[(d, fy - 1)]; share[(d, fy)] = p if p >= 1 else min(p + 0.01, 1); tgt[(d, fy)] = share[(d, fy)] * pot[(d, fy)]
print('FY  potential16(py)  potential16(xlsx-ITUK-UTI)  target16(py)  target16(xlsx)')
for b, fy in base.items():
    px = D2.cell(44, CI(b) + 3).value - D2.cell(9, CI(b) + 3).value - D2.cell(10, CI(b) + 3).value
    pp = sum(pot[(d, fy)] for d in dbs)
    line = f'{fy} {pp:14,.0f} {px:14,.0f}'
    if fy >= 2027:
        tx = D2.cell(44, CI(b) + 5).value - D2.cell(9, CI(b) + 5).value - D2.cell(10, CI(b) + 5).value
        line += f' {sum(tgt[(d, fy)] for d in dbs):14,.0f} {tx:14,.0f}'
    print(line)
