"""
AIM Class 0 PN Detector
=======================
Scans the AIM workbook for parts with Class 0 (no margin factor assigned).
Run after any AIM refresh to catch newly unclassified Source J parts.

Usage:
    python aim_class0_detector.py
    python aim_class0_detector.py --workbook "path/to/AIM_CY2025_Analysis.xlsx"

Output:
    - Console summary
    - class0_report_YYYYMMDD.csv  (same folder as workbook)
"""

import sys
import os
import argparse
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

DEFAULT_WORKBOOK = r"C:\Users\jerom\Desktop\claude\ISZE\Master Dashboard\AIM_CY2025_Analysis.xlsx"

# Columns to look for in the Parts Master sheet
# Adjust these if the AIM workbook structure changes
SHEET_NAMES = ["Parts Master", "PartsMaster", "parts_master", "AIM Parts"]
PN_COL_NAMES = ["PN", "Part Number", "Part_Number", "pn", "part_number"]
CLASS_COL_NAMES = ["Class", "class", "ABC Class", "abc_class", "AIM_Class", "Pricing Class"]
SOURCE_COL_NAMES = ["Source", "source", "PN Source", "pn_source"]
DESC_COL_NAMES = ["Description", "desc", "Part Description", "DESCRIPTION"]
REV_COL_NAMES = ["Revenue", "revenue", "Rev", "CY2025 Revenue", "Rev_EUR"]


def find_col(headers, candidates):
    """Return 0-based column index of first matching candidate, or None."""
    for name in candidates:
        for i, h in enumerate(headers):
            if h and str(h).strip().lower() == name.lower():
                return i
    return None


def detect_class0(workbook_path: str) -> list[dict]:
    """Return list of Class 0 part records from the AIM workbook."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)

    # Find the right sheet
    sheet = None
    for name in SHEET_NAMES:
        if name in wb.sheetnames:
            sheet = wb[name]
            break
    if sheet is None:
        # Fall back: look for any sheet with "part" in name
        for sname in wb.sheetnames:
            if "part" in sname.lower() or "master" in sname.lower():
                sheet = wb[sname]
                break
    if sheet is None:
        print(f"ERROR: Could not find Parts Master sheet. Available sheets: {wb.sheetnames}")
        return []

    # Read header row (row 1)
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("ERROR: Sheet is empty.")
        return []

    headers = [str(h) if h is not None else "" for h in rows[0]]

    col_pn    = find_col(headers, PN_COL_NAMES)
    col_class = find_col(headers, CLASS_COL_NAMES)
    col_src   = find_col(headers, SOURCE_COL_NAMES)
    col_desc  = find_col(headers, DESC_COL_NAMES)
    col_rev   = find_col(headers, REV_COL_NAMES)

    if col_pn is None or col_class is None:
        print(f"WARNING: Could not map required columns. Headers found: {headers[:15]}")
        print("         Adjust PN_COL_NAMES / CLASS_COL_NAMES in the script.")
        return []

    class0_parts = []
    for i, row in enumerate(rows[1:], start=2):  # row 2 onwards
        if row[col_pn] is None:
            continue
        raw_class = str(row[col_class]).strip() if row[col_class] is not None else ""
        # Detect Class 0: zero, "0", empty, "None", blank
        is_class0 = raw_class in ("0", "0.0", "", "None", "N/A", "-")
        if is_class0:
            class0_parts.append({
                "row": i,
                "pn": str(row[col_pn]).strip(),
                "class": raw_class or "(blank)",
                "source": str(row[col_src]).strip() if col_src and row[col_src] else "",
                "description": str(row[col_desc]).strip() if col_desc and row[col_desc] else "",
                "revenue": row[col_rev] if col_rev else None,
            })
    return class0_parts


def write_csv(records: list[dict], out_path: str):
    import csv
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["row", "pn", "class", "source", "description", "revenue"])
        writer.writeheader()
        writer.writerows(records)


def main():
    parser = argparse.ArgumentParser(description="Detect Class 0 PNs in AIM workbook")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK, help="Path to AIM Excel workbook")
    args = parser.parse_args()

    wb_path = args.workbook
    if not os.path.exists(wb_path):
        print(f"ERROR: Workbook not found: {wb_path}")
        sys.exit(1)

    print(f"Scanning: {wb_path}")
    records = detect_class0(wb_path)

    if not records:
        print("✅ No Class 0 PNs detected.")
        return

    total_rev = sum(r["revenue"] for r in records if isinstance(r["revenue"], (int, float)))
    source_j  = [r for r in records if r["source"].upper() in ("J", "SOURCE J", "JP")]

    print(f"\n⚠️  CLASS 0 PNs DETECTED: {len(records)}")
    print(f"   Total revenue at risk: €{total_rev:,.0f}")
    print(f"   Source J count: {len(source_j)} / {len(records)}")
    print()

    # Print table
    print(f"{'PN':<16} {'Source':<10} {'Class':<8} {'Revenue':>10}  Description")
    print("-" * 80)
    for r in sorted(records, key=lambda x: -(x["revenue"] or 0)):
        rev_str = f"€{r['revenue']:,.0f}" if isinstance(r["revenue"], (int, float)) else "n/a"
        desc = (r["description"][:38] + "…") if len(r["description"]) > 39 else r["description"]
        print(f"{r['pn']:<16} {r['source']:<10} {r['class']:<8} {rev_str:>10}  {desc}")

    # Write CSV
    ts = datetime.now().strftime("%Y%m%d")
    out_dir = os.path.dirname(wb_path)
    csv_path = os.path.join(out_dir, f"class0_report_{ts}.csv")
    write_csv(records, csv_path)
    print(f"\nCSV written: {csv_path}")
    print("\nAction: Assign pricing class to each PN above, or exclude from AIM scoring.")
    print("        Source J = Japan-sourced — check IML margin table for correct class factor.")


if __name__ == "__main__":
    main()
