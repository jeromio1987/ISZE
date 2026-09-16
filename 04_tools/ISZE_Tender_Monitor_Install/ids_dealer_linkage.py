"""IDS dealer / distributor linkage for tender opportunities.

Commercial MAP: country → GPS distributor (yaml) + best-effort IDS dealer
from the real IDS 2026 ranking workbook. Never invents dealers.

Match methods:
  - geo_nearest     — buyer/title/city token hits a dealer name or address city
  - country_primary — IDS rank #1 for that country (no geo hit)
  - unmapped        — no yaml distributor and/or no IDS dealers for the country

Hard rules: null ≠ 0; no fake dealers; no expected revenue.
"""
from __future__ import annotations

import csv
import json
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parent
ISZE_ROOT = ROOT.parents[1]
MAP_YAML = ROOT / "country_distributor_map.yaml"
DIST_DIM = ISZE_ROOT / "01_data" / "00_derived" / "distributor_dimension.csv"
IDS_RANKINGS_XLSX = (
    ISZE_ROOT
    / "01_data"
    / "90_INBOX"
    / "IDS_2026"
    / "rankings"
    / "IDS_2026_all_country_rankings.xlsx"
)
IDS_DERIVED_CSV = ISZE_ROOT / "01_data" / "00_derived" / "ids_dealer_network.csv"
IDS_MODULE_STATUS = ROOT / "_ids_dealer_status.json"

# Tender country name → IDS workbook tab(s). Order = preference for primary.
# Germany / Austria / … absent on purpose — no IDS dealer sheet on disk.
TENDER_TO_IDS_TABS: Dict[str, List[str]] = {
    "Belgium": ["Belgium"],
    "Netherlands": ["Netherlands"],
    "Luxembourg": ["Luxembourg"],
    "Denmark": ["Denmark"],
    "Poland": ["Poland"],
    "Greece": ["Greece"],
    "Norway": ["Norway"],
    "United Kingdom": ["UK CV", "UK LCV"],  # yaml primary = ITUK (CV)
    "France": ["France"],
    "Italy": ["Italy"],
    "Switzerland": ["Switzerland"],
    "Spain": ["Spain"],
    "Israel": ["Israel LCV"],
    "Romania": ["Romania"],
    "Portugal": ["Portugal"],
    "Hungary": ["Hungary"],
    "Malta": ["Malta"],
    "Cyprus": ["Cyprus"],
}

# Soft city tokens sometimes embedded in IDS dealer names (not inventing dealers —
# only matching substrings already present in the ranking Name column).
_CITY_ALIASES: Dict[str, List[str]] = {
    "bucuresti": ["bucharest", "bucuresti", "bucureșți"],
    "bucharest": ["bucharest", "bucuresti"],
    "oradea": ["oradea"],
    "prahova": ["prahova", "ploiesti", "ploiești"],
    "barneveld": ["barneveld"],
    "venlo": ["venlo"],
    "namur": ["namur"],
    "huy": ["huy"],
    "marche": ["marche"],
    "waddinxveen": ["waddinxveen"],
    "zwolle": ["zwolle"],
    "veenendaal": ["veenendaal"],
}


def _load_yaml_map() -> Dict[str, Any]:
    text = MAP_YAML.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore

        return yaml.safe_load(text) or {}
    except Exception:
        country: Dict[str, Optional[str]] = {}
        iso3: Dict[str, str] = {}
        section = None
        for line in text.splitlines():
            s = line.strip()
            if not s or s.startswith("#"):
                continue
            if s.endswith(":") and not s.startswith(" "):
                section = s[:-1]
                continue
            if ":" in s:
                k, v = s.split(":", 1)
                k, v = k.strip(), v.split("#")[0].strip()
                val: Optional[str] = None if v == "null" else v
                if section == "country_to_distributor":
                    country[k] = val
                elif section == "iso3_to_country":
                    iso3[k] = val or ""
        return {"country_to_distributor": country, "iso3_to_country": iso3}


@lru_cache(maxsize=1)
def _maps() -> Tuple[Dict[str, Optional[str]], Dict[str, str]]:
    data = _load_yaml_map()
    return (
        dict(data.get("country_to_distributor") or {}),
        dict(data.get("iso3_to_country") or {}),
    )


def resolve_distributor(country: str, country_code: str = "") -> Optional[str]:
    cmap, iso = _maps()
    name = (country or "").strip()
    code = (country_code or "").strip().upper()
    if not name and code:
        name = iso.get(code, "")
    if name in cmap:
        return cmap.get(name)
    if code in iso:
        return cmap.get(iso[code])
    return None


@lru_cache(maxsize=1)
def _distributor_labels() -> Dict[str, str]:
    """distributor_code → commercial label from distributor_dimension.csv."""
    out: Dict[str, str] = {}
    if not DIST_DIM.exists():
        return out
    with DIST_DIM.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            code = (row.get("distributor_code") or "").strip()
            if not code:
                continue
            name = (row.get("customer_name") or "").strip()
            label = (row.get("gps_db_label") or "").strip()
            candidate = name or label or code
            prev = out.get(code)
            # Prefer a real company name over a short GPS label when multiple SAP rows share a code
            if not prev:
                out[code] = candidate
            elif name and (not prev or prev == label or len(name) > len(prev)):
                out[code] = name
    return out


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").casefold()).strip()


def _city_from_name(name: str) -> Optional[str]:
    """Best-effort city token already present in the IDS dealer name."""
    n = _norm(name)
    # Prefer explicit place words after a dash / at end
    for sep in (" - ", " – ", " — "):
        if sep in name:
            tail = name.split(sep)[-1].strip()
            if len(tail) >= 3 and not tail.isupper():
                return tail
    # Known Romanian / Benelux tokens in the ranking names
    for token in (
        "BUCURESTI",
        "ORADEA",
        "PRAHOVA",
        "Barneveld",
        "Venlo",
        "Namur",
        "Huy",
        "Marche",
        "Waddinxveen",
        "Zwolle",
        "Veenendaal",
        "SintTruiden",
    ):
        if token.casefold() in n:
            return token.replace("SintTruiden", "Sint-Truiden")
    return None


def _load_dealers_from_xlsx() -> Dict[str, List[Dict[str, Any]]]:
    if not IDS_RANKINGS_XLSX.exists():
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}
    by_tab: Dict[str, List[Dict[str, Any]]] = {}
    wb = openpyxl.load_workbook(IDS_RANKINGS_XLSX, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if sheet_name == "Index":
                continue
            ws = wb[sheet_name]
            rows: List[Dict[str, Any]] = []
            header = None
            for row in ws.iter_rows(values_only=True):
                if not row or row[0] is None:
                    continue
                if header is None:
                    header = [str(c or "").strip() for c in row]
                    continue
                rank_raw, name_raw = row[0], row[1] if len(row) > 1 else None
                name = str(name_raw or "").strip()
                if not name:
                    continue
                try:
                    rank = int(rank_raw)
                except (TypeError, ValueError):
                    continue
                pts = None
                if len(row) > 2 and row[2] is not None:
                    try:
                        pts = float(row[2])
                    except (TypeError, ValueError):
                        pts = None
                rows.append(
                    {
                        "ids_tab": sheet_name,
                        "rank": rank,
                        "dealer_name": name,
                        "points_pct": pts,
                        "city": _city_from_name(name),
                        "address": None,
                        "source_file": IDS_RANKINGS_XLSX.name,
                    }
                )
            if rows:
                rows.sort(key=lambda r: r["rank"])
                by_tab[sheet_name] = rows
    finally:
        wb.close()
    return by_tab


def _load_dealers_from_csv() -> Dict[str, List[Dict[str, Any]]]:
    if not IDS_DERIVED_CSV.exists():
        return {}
    by_tab: Dict[str, List[Dict[str, Any]]] = {}
    with IDS_DERIVED_CSV.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            tab = (row.get("ids_tab") or "").strip()
            name = (row.get("dealer_name") or "").strip()
            if not tab or not name:
                continue
            try:
                rank = int(row.get("rank") or 0)
            except ValueError:
                continue
            city = (row.get("city") or "").strip() or _city_from_name(name)
            pts_raw = (row.get("points_pct") or "").strip()
            pts = float(pts_raw) if pts_raw else None
            by_tab.setdefault(tab, []).append(
                {
                    "ids_tab": tab,
                    "rank": rank,
                    "dealer_name": name,
                    "points_pct": pts,
                    "city": city or None,
                    "address": (row.get("address") or "").strip() or None,
                    "source_file": (row.get("source_file") or IDS_DERIVED_CSV.name),
                }
            )
    for tab in by_tab:
        by_tab[tab].sort(key=lambda r: r["rank"])
    return by_tab


@lru_cache(maxsize=1)
def load_ids_dealers_by_tab() -> Dict[str, List[Dict[str, Any]]]:
    """IDS tab → ranked dealer rows. Prefer live xlsx; fall back to derived CSV."""
    from_xlsx = _load_dealers_from_xlsx()
    if from_xlsx:
        return from_xlsx
    return _load_dealers_from_csv()


def export_derived_csv(path: Optional[Path] = None) -> Path:
    """Write flat IDS dealer network CSV under 01_data/00_derived (real rankings only)."""
    out = path or IDS_DERIVED_CSV
    by_tab = _load_dealers_from_xlsx() or _load_dealers_from_csv()
    out.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "ids_tab",
        "rank",
        "dealer_name",
        "points_pct",
        "city",
        "address",
        "source_file",
    ]
    with out.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for tab in sorted(by_tab.keys()):
            for row in by_tab[tab]:
                w.writerow({k: row.get(k) for k in fields})
    return out


def ids_source_status() -> Dict[str, Any]:
    by_tab = load_ids_dealers_by_tab()
    xlsx_ok = IDS_RANKINGS_XLSX.exists()
    csv_ok = IDS_DERIVED_CSV.exists()
    blocked_countries = sorted(
        c
        for c, dist in _maps()[0].items()
        if dist and c not in TENDER_TO_IDS_TABS
    )
    return {
        "ids_rankings_xlsx": str(IDS_RANKINGS_XLSX),
        "ids_rankings_xlsx_exists": xlsx_ok,
        "ids_derived_csv": str(IDS_DERIVED_CSV),
        "ids_derived_csv_exists": csv_ok,
        "tabs_loaded": sorted(by_tab.keys()),
        "dealer_count": sum(len(v) for v in by_tab.values()),
        "tender_countries_with_ids_tabs": sorted(TENDER_TO_IDS_TABS.keys()),
        "gps_mapped_countries_without_ids_dealers": blocked_countries,
        "blocked_note": (
            None
            if xlsx_ok or csv_ok
            else (
                "BLOCKED — drop IDS_2026_all_country_rankings.xlsx "
                f"at {IDS_RANKINGS_XLSX} (or per-country *dealer_ranking.xlsx)."
            )
        ),
    }


def _haystack(row: Dict[str, Any]) -> str:
    parts = [
        row.get("buyer") or row.get("Buyer") or "",
        row.get("title") or row.get("Title") or "",
        row.get("city") or row.get("City") or "",
        row.get("region") or row.get("Region") or "",
        row.get("locality") or "",
        row.get("NUTS") or row.get("nuts") or "",
    ]
    return _norm(" ".join(str(p) for p in parts if p))


_STOP_TOKENS = frozenset(
    {
        "open",
        "auto",
        "center",
        "centre",
        "motor",
        "motors",
        "garage",
        "group",
        "groep",
        "ltd",
        "limited",
        "gmbh",
        "bv",
        "sa",
        "srl",
        "truck",
        "trucks",
        "commercial",
        "commercials",
        "vehicle",
        "vehicles",
        "service",
        "services",
        "parts",
        "isuzu",
        "sud",  # too short / common in BE names when alone
    }
)


def _geo_score(dealer: Dict[str, Any], hay: str) -> float:
    if not hay:
        return 0.0
    name = _norm(dealer.get("dealer_name") or "")
    city = _norm(dealer.get("city") or "")
    addr = _norm(dealer.get("address") or "")
    score = 0.0
    if name and len(name) >= 5 and name in hay:
        score = max(score, 0.95)
    # Token overlap — skip short / generic brand words (avoid "auto"/"center" false hits)
    for tok in re.findall(r"[a-z0-9]{5,}", name):
        if tok in _STOP_TOKENS:
            continue
        if tok in hay:
            score = max(score, 0.8)
    if city and len(city) >= 4 and city in hay:
        score = max(score, 0.85)
    if addr:
        for tok in re.findall(r"[a-z0-9]{5,}", addr):
            if tok in _STOP_TOKENS:
                continue
            if tok in hay:
                score = max(score, 0.75)
    # Alias expansion for city tokens
    for key, aliases in _CITY_ALIASES.items():
        if key in name or key in city:
            if any(a in hay for a in aliases):
                score = max(score, 0.82)
    return score


def _dealers_for_country(country: str, distributor: Optional[str]) -> List[Dict[str, Any]]:
    tabs = list(TENDER_TO_IDS_TABS.get((country or "").strip()) or [])
    # Prefer CV tab when UK + ITUK
    if country == "United Kingdom" and distributor == "ITUK":
        tabs = ["UK CV", "UK LCV"]
    elif country == "United Kingdom" and distributor == "IUK":
        tabs = ["UK LCV", "UK CV"]
    by_tab = load_ids_dealers_by_tab()
    out: List[Dict[str, Any]] = []
    for tab in tabs:
        out.extend(by_tab.get(tab) or [])
    return out


def attach_ids_dealer(row: Dict[str, Any]) -> Dict[str, Any]:
    """Attach ids_* commercial fields. Mutates and returns row. Never invents dealers."""
    row.setdefault("ids_distributor", None)
    row.setdefault("ids_distributor_label", None)
    row.setdefault("ids_dealer_name", None)
    row.setdefault("ids_dealer_city", None)
    row.setdefault("ids_dealer_rank", None)
    row.setdefault("ids_match_method", "unmapped")
    row.setdefault("ids_match_confidence", None)
    row.setdefault("ids_tab", None)

    country = str(row.get("Country") or row.get("country") or "").strip()
    ccode = str(row.get("Country Code") or row.get("country_code") or "").strip()
    dist = resolve_distributor(country, ccode)
    labels = _distributor_labels()

    row["ids_distributor"] = dist
    row["ids_distributor_label"] = labels.get(dist) if dist else None
    # Keep legacy GPS field in sync when empty (commercial ownership ≠ fleet sizing)
    if dist and not row.get("distributor_code"):
        row["distributor_code"] = dist

    dealers = _dealers_for_country(country, dist)
    if not dealers:
        row["ids_match_method"] = "unmapped"
        row["ids_match_confidence"] = None
        return row

    hay = _haystack(row)
    best: Optional[Dict[str, Any]] = None
    best_score = 0.0
    for d in dealers:
        sc = _geo_score(d, hay)
        if sc > best_score:
            best_score = sc
            best = d
        elif sc == best_score and sc > 0 and best is not None:
            if int(d["rank"]) < int(best["rank"]):
                best = d

    if best is not None and best_score >= 0.75:
        row["ids_dealer_name"] = best["dealer_name"]
        row["ids_dealer_city"] = best.get("city")
        row["ids_dealer_rank"] = best.get("rank")
        row["ids_tab"] = best.get("ids_tab")
        row["ids_match_method"] = "geo_nearest"
        row["ids_match_confidence"] = round(min(0.95, best_score), 2)
        return row

    primary = dealers[0]  # already sorted by rank within first preferred tab
    # If multiple tabs, prefer lowest rank within the first tab only
    first_tab = primary.get("ids_tab")
    tab_primaries = [d for d in dealers if d.get("ids_tab") == first_tab]
    primary = min(tab_primaries, key=lambda d: int(d["rank"]))

    row["ids_dealer_name"] = primary["dealer_name"]
    row["ids_dealer_city"] = primary.get("city")
    row["ids_dealer_rank"] = primary.get("rank")
    row["ids_tab"] = primary.get("ids_tab")
    row["ids_match_method"] = "country_primary"
    row["ids_match_confidence"] = 0.5
    return row


def enrich_payload_ids(payload: Dict[str, Any]) -> Dict[str, Any]:
    for n in payload.get("notices") or []:
        attach_ids_dealer(n)
    for o in payload.get("opportunities") or []:
        attach_ids_dealer(o)
    meta = payload.setdefault("meta", {})
    meta["ids_dealer"] = ids_source_status()
    try:
        IDS_MODULE_STATUS.write_text(
            json.dumps(meta["ids_dealer"], indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    except OSError:
        pass
    return payload


if __name__ == "__main__":
    status = ids_source_status()
    print(json.dumps(status, indent=2, ensure_ascii=False))
    if status.get("ids_rankings_xlsx_exists"):
        path = export_derived_csv()
        print(f"Wrote {path}")
    for sample in (
        {"country": "Romania", "buyer": "Primaria Bucuresti", "title": "Pickup fleet"},
        {"country": "Romania", "buyer": "CJ Cluj", "title": "Vehicles"},
        {"country": "Germany", "buyer": "Stadt Köln", "title": "LKW"},
        {"country": "Belgium", "buyer": "Commune de Namur", "title": "Utilitaires"},
    ):
        print(json.dumps(attach_ids_dealer(dict(sample)), indent=2, ensure_ascii=False))
