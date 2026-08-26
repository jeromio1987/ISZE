#!/usr/bin/env python3
"""Merge gated Austria USP notices into primary isuzu_tenders_v3.xlsx.

Preserves Isuzu + not-for-us gate. Does not invent notices. Does not touch Nordics keys / RO sheet.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
MAIN = HERE / "isuzu_tenders_v3.xlsx"
SIDE = HERE / "isuzu_tenders_v3_at.xlsx"

SHEET_MAP = {
    "Austria USP": ("Austria USP", "AT-USP"),
}


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _row_keys_from_dict(row: Dict[str, Any]) -> Tuple[str, str]:
    url = _norm(row.get("URL") or row.get("url")).lower()
    nid = _norm(row.get("Source Notice ID") or row.get("source_notice_id")).lower()
    return url, nid


def _sheet_raw_dicts(ws: Worksheet) -> Tuple[List[str], List[Dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_norm(h) for h in rows[0]]
    out: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if not any(v not in (None, "") for v in raw):
            continue
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        src = _norm(row.get("Source"))
        title = _norm(row.get("Title"))
        if not title and src.lower().startswith("no "):
            continue
        if src.startswith("No results") or src.startswith("No AT"):
            continue
        if title.lower().startswith("no "):
            continue
        out.append(row)
    return headers, out


def _collect_other_sheet_keys(wb: openpyxl.Workbook, skip: Set[str]) -> Tuple[Set[str], Set[str]]:
    urls: Set[str] = set()
    ids: Set[str] = set()
    for name in wb.sheetnames:
        if name in skip:
            continue
        headers, data = _sheet_raw_dicts(wb[name])
        if not headers:
            continue
        for row in data:
            u, n = _row_keys_from_dict(row)
            if u:
                urls.add(u)
            if n:
                ids.add(n)
    return urls, ids


def _export_headers_from_norm(norm: Dict[str, Any], raw: Dict[str, Any], headers: List[str]) -> Dict[str, Any]:
    key_map = {
        "Source": ("Source", "source"),
        "Source Notice ID": ("Source Notice ID", "source_notice_id"),
        "Lifecycle": ("Lifecycle", "lifecycle"),
        "Notice Type": ("Notice Type", "notice_type"),
        "Procedure Type": ("Procedure Type", "procedure_type"),
        "Buyer Type": ("Buyer Type", "buyer_type"),
        "Score": ("Score", "score"),
        "FleetSize": ("FleetSize", "fleet_size"),
        "Title": ("Title", "title"),
        "Buyer": ("Buyer", "buyer"),
        "Country": ("Country", "country"),
        "Country Code": ("Country Code", "country_code"),
        "Published": ("Published", "published"),
        "Deadline": ("Deadline", "deadline"),
        "Value": ("Value", "value"),
        "Currency": ("Currency", "currency"),
        "CPV Code": ("CPV Code", "cpv"),
        "CPV Desc": ("CPV Desc", "cpv_desc"),
        "URL": ("URL", "url"),
        "Description": ("Description", "description"),
        "Fit Score": ("Fit Score", "fit_score"),
        "Win Score": ("Win Score", "win_score"),
        "Econ Score": ("Econ Score", "econ_score"),
        "Decision": ("Decision", "decision"),
        "Evidence Score": ("Evidence Score", "evidence_score"),
        "Evidence Label": ("Evidence Label", "evidence_label"),
        "Suggested Owner": ("Suggested Owner", "owner_suggested"),
        "Suggested Status": ("Suggested Status", "status_suggested"),
        "Days to Deadline": ("Days to Deadline", "days_to_deadline"),
        "Urgency": ("Urgency", "urgency"),
        "Reasons": ("Reasons", "reasons"),
    }
    out: Dict[str, Any] = {}
    for h in headers:
        aliases = key_map.get(h, (h,))
        val = None
        for a in aliases:
            if a in raw and raw[a] not in (None, ""):
                val = raw[a]
                break
            if a in norm and norm[a] not in (None, ""):
                val = norm[a]
                break
        out[h] = "" if val is None else val
    return out


def _gate_rows(raw_rows: List[Dict[str, Any]], source_label: str) -> Tuple[List[Dict[str, Any]], int]:
    from isze_core_v3 import normalize_raw_row, row_passes_isuzu_selection_gate
    from tender_scraper_v3 import EXPORT_HEADERS

    kept: List[Dict[str, Any]] = []
    dropped = 0
    for raw in raw_rows:
        raw = dict(raw)
        raw.setdefault("_source", source_label)
        norm = normalize_raw_row(raw, source_label)
        if not row_passes_isuzu_selection_gate(norm):
            dropped += 1
            continue
        kept.append(_export_headers_from_norm(norm, raw, list(EXPORT_HEADERS)))
    return kept, dropped


def _rewrite_sheet(ws: Worksheet, headers: List[str], rows: List[Dict[str, Any]]) -> None:
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def _ensure_sheet(wb: openpyxl.Workbook, name: str) -> Worksheet:
    if name in wb.sheetnames:
        return wb[name]
    return wb.create_sheet(name)


def merge() -> Dict[str, Any]:
    from tender_scraper_v3 import EXPORT_HEADERS

    if not MAIN.exists():
        raise SystemExit(f"Main workbook missing: {MAIN}")
    if not SIDE.exists():
        raise SystemExit(f"Side workbook missing: {SIDE}")

    main_wb = openpyxl.load_workbook(MAIN)
    side_wb = openpyxl.load_workbook(SIDE)

    skip_names = {main_name for main_name, _ in SHEET_MAP.values()}
    other_urls, other_ids = _collect_other_sheet_keys(main_wb, skip_names)

    stats: Dict[str, Any] = {"sheets": {}, "gate": {}}

    for side_name, (main_name, source_label) in SHEET_MAP.items():
        if side_name not in side_wb.sheetnames:
            headers = list(EXPORT_HEADERS)
            ws = _ensure_sheet(main_wb, main_name)
            _rewrite_sheet(ws, headers, [])
            stats["sheets"][main_name] = {"note": "side sheet missing — wrote empty"}
            continue

        headers, side_raw = _sheet_raw_dicts(side_wb[side_name])
        headers = headers or list(EXPORT_HEADERS)
        gated, dropped = _gate_rows(side_raw, source_label)

        merged: List[Dict[str, Any]] = []
        seen_u: Set[str] = set()
        seen_n: Set[str] = set()
        skipped_cross = 0
        skipped_dup = 0

        for row in gated:
            u, n = _row_keys_from_dict(row)
            if u and u in seen_u:
                skipped_dup += 1
                continue
            if n and n in seen_n:
                skipped_dup += 1
                continue
            if u and u in other_urls:
                skipped_cross += 1
                continue
            if n and n in other_ids:
                skipped_cross += 1
                continue
            if u:
                seen_u.add(u)
            if n:
                seen_n.add(n)
            merged.append(row)

        ws = _ensure_sheet(main_wb, main_name)
        _rewrite_sheet(ws, list(EXPORT_HEADERS), merged)
        stats["sheets"][main_name] = {
            "side_raw": len(side_raw),
            "gated_kept": len(gated),
            "gate_dropped": dropped,
            "merged_into_main": len(merged),
            "skipped_cross_sheet_dup": skipped_cross,
            "skipped_internal_dup": skipped_dup,
        }
        stats["gate"][main_name] = {"raw": len(side_raw), "kept": len(gated), "dropped": dropped}
        print(
            f"[{main_name}] side_raw={len(side_raw)} gated={len(gated)} "
            f"(drop {dropped}) merged={len(merged)} cross_dup={skipped_cross}"
        )

    main_wb.save(MAIN)
    stats["main"] = str(MAIN)
    stats["side"] = str(SIDE)
    out = HERE / "_phase1_at_merge_stats.json"
    out.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {out}")
    print(f"Saved {MAIN}")
    return stats


if __name__ == "__main__":
    merge()
