#!/usr/bin/env python3
"""Build ISZE Tender OS v4 — GPS / PMS addressable-opportunity dashboard.

Self-contained HTML (DATA fully embedded) so file:// works offline. No CDN.
Does not modify isze_dashboard_v2.html.

Prefer: import load_input_rows / build_payload from build_dashboard_v2, or load an
existing tender_data.json / isze_dashboard_v2_data.json via --json-input.

Outcomes are stored in tender_outcomes.json for future analysis but are NOT
editable in the v4 UI (outcome dropdown removed).
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, Optional

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if sys.stderr.encoding and sys.stderr.encoding.lower() not in ("utf-8", "utf8"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
ISZE_ROOT = HERE.parents[1]
V2_JSON = HERE / "isze_dashboard_v2_data.json"
V4_JSON = HERE / "isze_dashboard_v4_data.json"
V4_HTML = HERE / "isze_dashboard_v4_opportunity.html"
V4_XLSX = HERE / "isze_dashboard_v4_opportunities.xlsx"
XLSX = HERE / "isuzu_tenders_v3.xlsx"
SOURCE_HEALTH = HERE / "_source_health.json"
BP_FY27 = ISZE_ROOT / "01_data" / "91_build" / "bp_fy27_reviewed.json"
MAP_YAML = HERE / "country_distributor_map.yaml"

# Extra GPS codes in BP that are not the primary country_to_distributor entry
GPS_CODE_ALIASES = {
    "IUK": "United Kingdom",
    "ITUK": "United Kingdom",
    "UMI": "Israel",
    "UTI": "Israel",
}


def load_payload(input_path: Optional[Path] = None, json_input: Optional[Path] = None) -> Dict[str, Any]:
    if json_input:
        if not json_input.exists():
            raise SystemExit(f"JSON input not found: {json_input}")
        print(f"Loading {json_input}")
        payload = json.loads(json_input.read_text(encoding="utf-8"))
    elif input_path and input_path.suffix.lower() in (".xlsx", ".html") and input_path.exists():
        print(f"Building from {input_path.name}")
        from build_dashboard_v2 import build_payload, load_input_rows

        payload = build_payload(load_input_rows(input_path))
    elif V2_JSON.exists():
        print(f"Loading {V2_JSON.name}")
        payload = json.loads(V2_JSON.read_text(encoding="utf-8"))
    else:
        from build_dashboard_v2 import build_payload, load_input_rows

        if not XLSX.exists():
            raise SystemExit(f"Neither {V2_JSON.name} nor {XLSX.name} found")
        print(f"{V2_JSON.name} missing — building from {XLSX.name}")
        payload = build_payload(load_input_rows(XLSX))

    # Ensure GPS/PMS + outcomes are applied (idempotent; build_payload already does this)
    try:
        from gps_pms_linkage import enrich_payload

        enrich_payload(payload)
    except Exception as exc:
        print(f"[GPS/PMS] enrich skipped: {exc}")
    try:
        from tender_outcomes import merge_outcomes_into_payload

        merge_outcomes_into_payload(payload)
    except Exception as exc:
        print(f"[Outcomes] merge skipped: {exc}")

    # Re-apply Isuzu + not-for-us commercial gate when loading a pre-gate JSON snapshot
    try:
        from isze_core_v3 import row_passes_isuzu_selection_gate

        before_n = len(payload.get("notices") or [])
        before_o = len(payload.get("opportunities") or [])
        # Tag aftersales before gate so real fleet R&M can survive not-for-us
        try:
            from isze_core_v3 import attach_aftersales

            for n in payload.get("notices") or []:
                attach_aftersales(n)
            for o in payload.get("opportunities") or []:
                attach_aftersales(o)
        except Exception as exc_as:
            print(f"[Aftersales] pre-gate tag skipped: {exc_as}")

        payload["notices"] = [
            n for n in (payload.get("notices") or []) if row_passes_isuzu_selection_gate(n)
        ]
        payload["opportunities"] = [
            o for o in (payload.get("opportunities") or []) if row_passes_isuzu_selection_gate(o)
        ]
        dropped_n = before_n - len(payload["notices"])
        dropped_o = before_o - len(payload["opportunities"])
        if dropped_n or dropped_o:
            print(
                f"[Isuzu+not-for-us gate] filtered snapshot: notices {before_n}->{len(payload['notices'])} "
                f"(-{dropped_n}), opportunities {before_o}->{len(payload['opportunities'])} (-{dropped_o})"
            )
            meta = payload.setdefault("meta", {})
            stats = meta.setdefault("stats", {})
            stats["raw_notices"] = len(payload["notices"])
            stats["opportunities"] = len(payload["opportunities"])
            stats["isuzu_gate_dropped_notices"] = dropped_n
            stats["isuzu_gate_dropped_opportunities"] = dropped_o
        else:
            meta = payload.setdefault("meta", {})
            stats = meta.setdefault("stats", {})
            stats.setdefault("isuzu_gate_dropped_notices", 0)
            stats.setdefault("isuzu_gate_dropped_opportunities", 0)
    except Exception as exc:
        print(f"[Isuzu gate] filter skipped: {exc}")

    # Aftersales tags (service contract / parts / warranty) — refresh counts post-gate
    try:
        from isze_core_v3 import attach_aftersales

        for n in payload.get("notices") or []:
            attach_aftersales(n)
        tagged = 0
        for o in payload.get("opportunities") or []:
            attach_aftersales(o)
            if o.get("aftersales_tags"):
                tagged += 1
        meta = payload.setdefault("meta", {})
        stats = meta.setdefault("stats", {})
        stats["aftersales_tagged"] = tagged
        print(f"[Aftersales] tagged {tagged} of {len(payload.get('opportunities') or [])} opportunities")
    except Exception as exc:
        print(f"[Aftersales] enrich skipped: {exc}")

    return payload


def load_source_health() -> Optional[Dict[str, Any]]:
    if not SOURCE_HEALTH.exists():
        return None
    try:
        return json.loads(SOURCE_HEALTH.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"[Source health] read failed: {exc}")
        return None


def load_workbook_sheets(xlsx_path: Optional[Path] = None) -> list:
    """List sheet names + data-row counts from isuzu_tenders_v3.xlsx (or override)."""
    path = xlsx_path or XLSX
    if not path.exists():
        return []
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for name in wb.sheetnames:
            ws = wb[name]
            data_rows = max(0, (ws.max_row or 1) - 1)
            sheets.append(
                {
                    "name": name,
                    "data_rows": data_rows,
                    "cols": ws.max_column or 0,
                }
            )
        wb.close()
        return sheets
    except Exception as exc:
        print(f"[Workbook sheets] read failed: {exc}")
        return []


def _country_to_distributor() -> Dict[str, Optional[str]]:
    if not MAP_YAML.exists():
        return {}
    text = MAP_YAML.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore

        data = yaml.safe_load(text) or {}
        return dict(data.get("country_to_distributor") or {})
    except Exception:
        out: Dict[str, Optional[str]] = {}
        section = None
        for line in text.splitlines():
            s = line.strip()
            if not s or s.startswith("#"):
                continue
            if s.endswith(":") and s.count(":") == 1:
                section = s[:-1]
                continue
            if section == "country_to_distributor" and ":" in s:
                k, v = s.split(":", 1)
                k = k.strip()
                v = v.split("#")[0].strip()
                out[k] = None if v == "null" else v
        return out


def load_bp_revenue_shares() -> Optional[Dict[str, Any]]:
    """Per-country network revenue shares from FY27 reviewed BP, or None if unavailable."""
    if not BP_FY27.exists():
        return None
    try:
        bp = json.loads(BP_FY27.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"[BP FY27] read failed: {exc}")
        return None
    by_code = bp.get("by_gps_code") or {}
    if not by_code:
        return None
    cmap = _country_to_distributor()
    code_to_country = {code: country for country, code in cmap.items() if code}
    for code, country in GPS_CODE_ALIASES.items():
        code_to_country.setdefault(code, country)
    by_country: Dict[str, float] = {}
    for code, eur in by_code.items():
        try:
            val = float(eur)
        except (TypeError, ValueError):
            continue
        country = code_to_country.get(code)
        if not country:
            continue
        by_country[country] = by_country.get(country, 0.0) + val
    total = float(bp.get("vp_network_total_eur") or sum(by_country.values()))
    if total <= 0 or not by_country:
        return None
    return {
        "available": True,
        "snapshot": bp.get("reviewed") or bp.get("generated"),
        "network_total_eur": total,
        "by_country": {
            c: {"revenue_eur": round(v, 2), "revenue_share": round(v / total, 6)}
            for c, v in by_country.items()
        },
        "label": (
            "Network revenue share from FY27 reviewed BP vs pipeline opportunity share. "
            "Diagnostic coverage only — not a forecast."
        ),
    }


def _has_fleet(o: Dict[str, Any]) -> bool:
    fs = o.get("fleet_size")
    if fs in (None, "", 0, "0"):
        fs = o.get("FleetSize")
    try:
        return int(fs) > 0
    except (TypeError, ValueError):
        return False


def _fleet_int(o: Dict[str, Any]) -> Optional[int]:
    """Return extracted fleet count, or None when missing. Never coerce null→0."""
    fs = o.get("fleet_size")
    if fs in (None, "", 0, "0"):
        fs = o.get("FleetSize")
    try:
        n = int(fs)
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


# Product-fit segments eligible for GPS € (not van). Stated-fleet sum ≠ CV tender count ≠ network sales.
GPS_ELIGIBLE_SEGMENTS = frozenset({"pickup", "truck"})


def addressable_vehicle_stats(payload: Dict[str, Any]) -> Dict[str, Any]:
    """Sum stated fleet units from extract. Never invent fleet. Not market potential.

    vehicles_pickup_truck — sum fleet where product_segment in pickup/truck (layer c)
    vehicles_all_stated   — sum fleet wherever extracted (incl van/LCV/weak)
    Do not label either as addressable market / annual demand. CV tender count is separate (layer b).
    """
    opps = payload.get("opportunities") or []
    vehicles_pickup_truck = 0
    vehicles_all_stated = 0
    rows_with_fleet = 0
    rows_pickup_truck_with_fleet = 0
    by_country: Dict[str, Dict[str, int]] = {}
    for o in opps:
        n = _fleet_int(o)
        if n is None:
            continue
        rows_with_fleet += 1
        vehicles_all_stated += n
        seg = o.get("product_segment") or "weak"
        c = o.get("country") or "—"
        bucket = by_country.setdefault(
            c,
            {
                "vehicles_pickup_truck": 0,
                "vehicles_all_stated": 0,
                "rows_with_fleet": 0,
                "rows_pickup_truck_with_fleet": 0,
            },
        )
        bucket["vehicles_all_stated"] += n
        bucket["rows_with_fleet"] += 1
        if seg in GPS_ELIGIBLE_SEGMENTS:
            vehicles_pickup_truck += n
            rows_pickup_truck_with_fleet += 1
            bucket["vehicles_pickup_truck"] += n
            bucket["rows_pickup_truck_with_fleet"] += 1
    return {
        "vehicles_pickup_truck": vehicles_pickup_truck,
        "vehicles_all_stated": vehicles_all_stated,
        "rows_with_fleet": rows_with_fleet,
        "rows_pickup_truck_with_fleet": rows_pickup_truck_with_fleet,
        "by_country": by_country,
        "definition": (
            "vehicles_pickup_truck = sum(fleet_size) where fleet extracted AND "
            "product_segment in {pickup, truck}. "
            "vehicles_all_stated = sum(fleet_size) wherever fleet extracted "
            "(incl van/LCV/weak) — labeled 'stated in notices (incl van/LCV)'. "
            "Null fleet rows omitted from sums; per-row display never coerces null→0. "
            "Independent of GPS € (product-fit still blocks van GPS wallet)."
        ),
    }


def load_gps_map_meta() -> Dict[str, Any]:
    """Country → {distributor, has_spu}. Real pack only — never invent SPU."""
    from gps_pms_linkage import _gps_latest

    cmap = _country_to_distributor()
    latest = _gps_latest()
    by_country: Dict[str, Dict[str, Any]] = {}
    for country, dist in cmap.items():
        has_spu = bool(dist and dist in latest)
        by_country[country] = {
            "distributor": dist,
            "has_spu": has_spu,
            "can_size_gps": has_spu,
        }
    # Also mark CSV countries that appear under a different name
    for code, row in latest.items():
        cname = (row.get("country") or "").strip()
        if cname and cname not in by_country:
            by_country[cname] = {
                "distributor": code,
                "has_spu": True,
                "can_size_gps": True,
            }
    sized_ok = sorted(c for c, m in by_country.items() if m.get("can_size_gps"))
    no_spu = sorted(c for c, m in by_country.items() if not m.get("can_size_gps"))
    return {
        "by_country": by_country,
        "countries_with_spu": sized_ok,
        "countries_without_spu": no_spu,
        "note": (
            "GPS € only where notice has fleet count, country has SPU in "
            "gps_spu_by_distributor.csv, and product segment is pickup/truck. "
            "Null ≠ 0. No invented SPU."
        ),
    }


def sized_stats(payload: Dict[str, Any]) -> Dict[str, Any]:
    opps = payload.get("opportunities") or []
    sized = [o for o in opps if o.get("gps_addressable_per_year_eur") is not None]
    unsized = len(opps) - len(sized)
    with_fleet = [o for o in opps if _has_fleet(o)]
    no_fleet = len(opps) - len(with_fleet)
    fleet_no_gps = [o for o in with_fleet if o.get("gps_addressable_per_year_eur") is None]
    total_gps = sum(float(o["gps_addressable_per_year_eur"]) for o in sized)
    seg_counts: Dict[str, int] = {}
    for o in opps:
        seg = o.get("product_segment") or "weak"
        seg_counts[seg] = seg_counts.get(seg, 0) + 1
    veh = addressable_vehicle_stats(payload)
    return {
        "opportunities": len(opps),
        "sized": len(sized),
        "unsized": unsized,
        "with_fleet": len(with_fleet),
        "no_fleet": no_fleet,
        "fleet_but_no_gps_map": len(fleet_no_gps),
        "gps_addressable_sum_eur": round(total_gps, 2),
        "product_segment_counts": seg_counts,
        "vehicles_pickup_truck": veh["vehicles_pickup_truck"],
        "vehicles_all_stated": veh["vehicles_all_stated"],
        "rows_with_fleet": veh["rows_with_fleet"],
        "rows_pickup_truck_with_fleet": veh["rows_pickup_truck_with_fleet"],
        "addressable_vehicles_definition": veh["definition"],
    }


HTML_TEMPLATE = r'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta name="color-scheme" content="only light">
<title>ISZE Tender OS v4 — Addressable Opportunity</title>
<style>
  :root {
    color-scheme: only light;
    --bg: #f7f4ee;
    --ink: #15233c;
    --accent: #c0552d;
    --panel: #ffffff;
    --muted: #5a6578;
    --line: #d8d2c6;
    --navy: #15233c;
    --paper: #f7f4ee;
    --ok: #1a6b3c;
    --warn: #9a6400;
    --bad: #a33a1a;
  }
  * { box-sizing: border-box; }
  html {
    color-scheme: only light;
    background: var(--bg);
    color: var(--ink);
  }
  body {
    margin: 0;
    font-family: Georgia, "Times New Roman", serif;
    font-size: 14px;
    line-height: 1.45;
    background: var(--bg);
    color: var(--ink);
    min-height: 100vh;
  }
  h1, h2, h3 {
    font-family: Georgia, "Times New Roman", serif;
    color: var(--ink);
    font-weight: 700;
  }
  a { color: var(--accent); }
  header {
    background: var(--navy);
    color: var(--paper);
    padding: 22px 28px 18px;
  }
  header h1 {
    margin: 0;
    font-size: 1.55rem;
    color: var(--paper);
  }
  header .sub {
    margin: 6px 0 0;
    opacity: 0.85;
    font-size: 0.92rem;
    color: var(--paper);
  }
  .wrap { padding: 18px 28px 40px; max-width: 1400px; margin: 0 auto; }
  .banner {
    margin: 14px 0 0;
    padding: 10px 14px;
    border-radius: 8px;
    font-size: 0.9rem;
    background: #eef3f9;
    color: var(--ink);
    border: 1px solid var(--line);
  }
  .banner.ok { background: #e8f5ee; border-color: #b7dcc6; }
  .banner.warn { background: #fff6db; border-color: #e8d48a; }
  .banner.bad { background: #fdece6; border-color: #e8b5a0; }
  .banner strong { color: var(--ink); }
  .cards {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
    gap: 14px;
    margin: 18px 0;
  }
  .card {
    background: var(--panel);
    border: 1px solid var(--line);
    border-radius: 10px;
    padding: 16px 18px;
    color: var(--ink);
  }
  .card .eyebrow {
    font-size: 0.72rem;
    text-transform: uppercase;
    letter-spacing: 0.06em;
    color: var(--muted);
    font-family: Georgia, serif;
  }
  .card .big {
    margin-top: 6px;
    font-size: 1.85rem;
    font-weight: 700;
    color: var(--ink);
  }
  .card .big.accent { color: var(--accent); }
  .card .sub {
    margin-top: 6px;
    font-size: 0.86rem;
    color: var(--muted);
  }
  .card.hero-vehicles {
    border: 2px solid var(--accent);
    background: #fff8f4;
    grid-column: span 1;
  }
  .card.hero-vehicles .eyebrow {
    color: var(--accent);
    font-weight: 700;
  }
  .card.hero-vehicles .big {
    font-size: 2.35rem;
    color: var(--accent);
  }
  .card.sanity-scale {
    background: #15233c;
    color: #f7f4ee;
    border: 1px solid #1a2b4a;
    border-left: 3px solid #c0552d;
    grid-column: 1 / -1;
  }
  .card.sanity-scale .eyebrow,
  .card.sanity-scale .eyebrow .kpi-tip {
    color: #c8d0dc;
    border-bottom-color: #c0552d;
  }
  .card.sanity-scale .big {
    font-size: 1.15rem;
    font-weight: 700;
    color: #f7f4ee;
    margin-top: 8px;
    line-height: 1.35;
  }
  .card.sanity-scale .sub { color: #c8d0dc; }
  .fleet-disclaimer {
    background: #fff6db;
    border: 1px solid #e8d48a;
    border-left: 4px solid #c0552d;
    border-radius: 8px;
    padding: 10px 14px;
    margin: -6px 0 14px;
    font-size: 0.88rem;
    color: var(--ink);
    line-height: 1.4;
  }
  .fleet-disclaimer strong { color: var(--navy); }
  .card.unsized-emphasis {
    border: 1px solid var(--line);
    background: #fff;
  }
  .card.unsized-emphasis .big.accent { color: var(--ink); }
  /* Coverage vs dark — from _coverage_chip_snippet.html / _MARKET_COVERAGE_DARK_MAP.md */
  .card.coverage-dark {
    background: #15233c;
    color: #f7f4ee;
    border: 1px solid #1a2b4a;
    border-left: 3px solid #c0552d;
  }
  .card.coverage-dark .eyebrow,
  .card.coverage-dark .eyebrow .kpi-tip {
    color: #c8d0dc;
    border-bottom-color: #c0552d;
  }
  .card.coverage-dark .big {
    color: #f7f4ee;
  }
  .card.coverage-dark .sub {
    color: #c8d0dc;
  }
  .gps-explain {
    background: var(--paper);
    border: 1px solid var(--line);
    border-left: 4px solid var(--accent);
    border-radius: 8px;
    padding: 12px 16px;
    margin: 0 0 16px;
    color: var(--ink);
    font-size: 0.92rem;
    line-height: 1.45;
  }
  .gps-explain strong { color: var(--navy); font-family: Georgia, serif; }
  .money-signal {
    display: inline-block;
    margin-top: 4px;
    padding: 2px 8px;
    border-radius: 4px;
    font-size: 0.72rem;
    font-weight: 700;
    background: #eef3f9;
    color: var(--navy);
    border: 1px solid #c5d0e0;
  }
  .gps-empty {
    color: var(--muted);
    font-size: 0.78rem;
    font-weight: 500;
  }
  .panel {
    background: var(--panel);
    border: 1px solid var(--line);
    border-radius: 10px;
    padding: 16px 18px;
    margin-bottom: 16px;
    color: var(--ink);
  }
  .panel h2 {
    margin: 0 0 4px;
    font-size: 1.15rem;
    color: var(--ink);
  }
  .panel .subtle {
    color: var(--muted);
    font-size: 0.88rem;
    margin-bottom: 12px;
  }
  .grid-2 {
    display: grid;
    grid-template-columns: 1.2fr 0.8fr;
    gap: 16px;
  }
  @media (max-width: 980px) {
    .grid-2 { grid-template-columns: 1fr; }
    .wrap, header { padding-left: 14px; padding-right: 14px; }
  }
  .table-wrap {
    overflow: auto;
    border: 1px solid var(--line);
    border-radius: 8px;
    background: #fff;
  }
  table {
    width: 100%;
    border-collapse: collapse;
    font-size: 0.9rem;
    color: var(--ink);
    background: #fff;
  }
  thead th {
    position: sticky;
    top: 0;
    z-index: 2;
    background: var(--navy);
    color: var(--paper);
    text-align: left;
    padding: 10px 12px;
    font-size: 0.72rem;
    text-transform: uppercase;
    letter-spacing: 0.04em;
    white-space: nowrap;
    cursor: pointer;
    user-select: none;
  }
  thead th:hover { filter: brightness(1.08); }
  thead th .sort-ind { opacity: 0.7; margin-left: 4px; }
  td {
    padding: 10px 12px;
    border-bottom: 1px solid #ebe6dc;
    vertical-align: top;
    color: var(--ink);
    background: #fff;
  }
  tbody tr { cursor: pointer; }
  tbody tr:hover td { background: #f0ebe3; }
  tbody tr.selected td { background: #e8eef7; }
  .muted { color: var(--muted); }
  .num { text-align: right; font-variant-numeric: tabular-nums; }
  .pill {
    display: inline-block;
    padding: 2px 8px;
    border-radius: 999px;
    font-size: 0.74rem;
    font-weight: 700;
    white-space: nowrap;
  }
  .pill.Attack { background: #e7f6ed; color: #126936; }
  .pill.Selective { background: #fff6db; color: #9a6400; }
  .pill.Watch { background: #eaf2ff; color: #1f5ca8; }
  .pill.Ignore { background: #f0f0f0; color: #5a6578; }
  .pill.rm { background: #fef3c7; color: #713f12; }
  .pill.ok { background: #e7f6ed; color: #126936; }
  .pill.warn { background: #fff6db; color: #9a6400; }
  .pill.bad { background: #fdece6; color: #a33a1a; }
  .pill.seg-pickup { background: #e7f6ed; color: #126936; }
  .pill.seg-truck { background: #eaf2ff; color: #1f5ca8; }
  .pill.seg-lcv_van { background: #fff0e8; color: #a33a1a; }
  .pill.seg-rm_only { background: #fef3c7; color: #713f12; }
  .pill.seg-weak { background: #f0f0f0; color: #5a6578; }
  .why-chips { display: flex; flex-wrap: wrap; gap: 4px; max-width: 280px; }
  .pill.why-nofleet { background: #eef3f9; color: #15233c; }
  .pill.why-nomap { background: #fdece6; color: #a33a1a; }
  .pill.why-van { background: #fff0e8; color: #a33a1a; }
  .pill.why-nordics { background: #fff6db; color: #9a6400; }
  .pill.why-pack { background: #f0f0f0; color: #5a6578; }
  .pill.ids-dealer { background: #eaf2ff; color: #15233c; display: inline-block; line-height: 1.25; }
  .pill.ids-dist { background: #eef3f9; color: #15233c; display: inline-block; line-height: 1.25; }

  .as-chips { display: flex; flex-wrap: wrap; gap: 4px; max-width: 220px; align-items: center; }
  .pill.as-service_contract { background: #15233c; color: #f7f4ee; }
  .pill.as-repair_maintenance { background: #c0552d; color: #f7f4ee; }
  .pill.as-parts { background: #1a4d3e; color: #f7f4ee; }
  .pill.as-extended_warranty { background: #1f5ca8; color: #f7f4ee; }
  .pill.as-full_service_lease { background: #5a3d2b; color: #f7f4ee; }
  .pill.as-empty { background: #eef3f9; color: #5a6578; font-weight: 600; }
  .seg-mix { font-size: 0.78rem; color: var(--muted); white-space: nowrap; }
  .gps-cell-main { font-weight: 700; color: var(--ink); font-size: 0.95rem; }
  .gps-cell-n { font-size: 0.78rem; color: var(--muted); margin-top: 2px; }
  .veh-hero {
    font-size: 1.35rem;
    font-weight: 700;
    color: var(--accent);
    font-variant-numeric: tabular-nums;
    line-height: 1.2;
  }
  .veh-hero .kpi-tip { color: var(--accent); border-bottom-color: var(--accent); }
  .veh-hero-sub {
    font-size: 0.72rem;
    color: var(--muted);
    margin-top: 3px;
    font-weight: 500;
  }
  th.th-veh {
    background: #1a2b4a;
    color: #f7f4ee;
    font-size: 0.78rem;
    letter-spacing: 0.05em;
  }
  th.th-gps-sec {
    color: #c8d0dc;
    font-weight: 600;
  }
  .market-toolbar {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
    align-items: center;
    margin-bottom: 10px;
  }
  .market-toolbar label {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    font-size: 0.86rem;
    color: var(--muted);
    cursor: pointer;
    user-select: none;
  }
  .market-toolbar input[type="checkbox"] {
    width: 16px;
    height: 16px;
    accent-color: var(--accent);
  }
  td.fleet-emph .kpi-tip,
  td.fleet-emph { font-weight: 700; color: var(--navy); }
  .detail-col { /* toggled via body class */ }
  body:not(.show-market-detail) .detail-col { display: none; }
  button, select {
    font: inherit;
    font-size: 0.9rem;
  }
  .btn {
    background: var(--navy);
    color: var(--paper);
    border: none;
    border-radius: 8px;
    padding: 9px 14px;
    cursor: pointer;
    font-weight: 700;
  }
  .btn:hover { filter: brightness(1.08); }
  .btn.secondary {
    background: #fff;
    color: var(--navy);
    border: 1px solid var(--line);
  }
  .filters {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
    margin-bottom: 12px;
    align-items: center;
  }
  .filters input, .filters select {
    padding: 8px 10px;
    border: 1px solid var(--line);
    border-radius: 8px;
    background: #fff;
    color: var(--ink);
    min-width: 140px;
  }
  .filters input { min-width: 220px; flex: 1; }
  /* KPI tooltips via ::after (not only title=) */
  .kpi-tip {
    position: relative;
    display: inline-block;
    cursor: help;
    border-bottom: 1px dotted var(--accent);
  }
  .kpi-tip::after {
    content: attr(data-tip);
    position: absolute;
    left: 0;
    bottom: calc(100% + 8px);
    width: min(320px, 70vw);
    padding: 10px 12px;
    background: #15233c;
    color: #f7f4ee;
    border-left: 3px solid #c0552d;
    border-radius: 6px;
    font-size: 0.8rem;
    line-height: 1.35;
    font-family: Georgia, serif;
    opacity: 0;
    pointer-events: none;
    transform: translateY(4px);
    transition: opacity 0.12s ease, transform 0.12s ease;
    z-index: 40;
    white-space: normal;
  }
  .kpi-tip:hover::after {
    opacity: 1;
    transform: translateY(0);
  }
  .cov-row {
    display: grid;
    grid-template-columns: 110px 1fr;
    gap: 10px;
    align-items: center;
    margin-bottom: 10px;
    font-size: 0.86rem;
  }
  .cov-bars { display: grid; gap: 4px; }
  .cov-track {
    height: 10px;
    background: #ebe6dc;
    border-radius: 999px;
    overflow: hidden;
  }
  .cov-fill-opp { height: 100%; background: var(--navy); }
  .cov-fill-gps { height: 100%; background: var(--accent); }
  .cov-legend {
    display: flex;
    gap: 16px;
    font-size: 0.82rem;
    color: var(--muted);
    margin-bottom: 12px;
  }
  .cov-legend span::before {
    content: "";
    display: inline-block;
    width: 12px;
    height: 10px;
    border-radius: 2px;
    margin-right: 6px;
    vertical-align: middle;
  }
  .cov-legend .l-opp::before { background: var(--navy); }
  .cov-legend .l-gps::before { background: var(--accent); }
  .cov-legend .l-rev::before { background: #8b4513; }
  .cov-fill-rev { height: 100%; background: #8b4513; }
  #pmsPanel { display: none; }
  #pmsPanel.open { display: block; }
  .pms-head {
    display: flex;
    justify-content: space-between;
    gap: 12px;
    align-items: flex-start;
    flex-wrap: wrap;
  }
  .kv-grid {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(160px, 1fr));
    gap: 10px;
    margin: 12px 0;
  }
  .kv {
    border: 1px solid var(--line);
    border-radius: 8px;
    padding: 10px;
    background: var(--paper);
    color: var(--ink);
  }
  .kv .k { font-size: 0.72rem; text-transform: uppercase; letter-spacing: 0.05em; color: var(--muted); }
  .kv .v { margin-top: 4px; font-weight: 700; color: var(--ink); }
  /* Top nav tabs (restored from v2/v3 pattern) */
  .tabs {
    display: flex;
    gap: 8px;
    flex-wrap: wrap;
    margin: 0 0 16px;
    padding: 0;
  }
  .tab {
    background: #fff;
    color: var(--navy);
    border: 1px solid var(--line);
    border-radius: 8px;
    padding: 10px 16px;
    cursor: pointer;
    font-weight: 700;
    font-family: Georgia, "Times New Roman", serif;
    font-size: 0.9rem;
  }
  .tab:hover { border-color: var(--accent); color: var(--accent); }
  .tab.active {
    background: var(--navy);
    color: var(--paper);
    border-color: var(--navy);
  }
  .view { display: none; }
  .view.active { display: block; }
  .explain-grid {
    display: grid;
    gap: 14px;
    max-width: 920px;
  }
  .explain-card {
    background: #fff;
    border: 1px solid var(--line);
    border-radius: 10px;
    padding: 16px 18px;
    color: var(--ink);
  }
  .explain-card h3 {
    margin: 0 0 8px;
    font-size: 1.05rem;
    color: var(--navy);
  }
  .explain-card p { margin: 0 0 8px; color: var(--ink); }
  .explain-card p:last-child { margin-bottom: 0; }
  .explain-card .callout {
    margin-top: 10px;
    padding: 10px 12px;
    background: #eef3f9;
    border-left: 3px solid var(--accent);
    border-radius: 6px;
    color: var(--ink);
    font-size: 0.9rem;
  }
  .score-badge {
    display: inline-block;
    font-weight: 800;
    padding: 4px 12px;
    border-radius: 6px;
    font-size: 0.85rem;
    margin-right: 8px;
  }
  .score-badge.fit { background: #eaf2ff; color: #1f5ca8; }
  .score-badge.win { background: #e7f6ed; color: #126936; }
  .score-badge.econ { background: #fff6db; color: #9a6400; }
  .method-table {
    width: 100%;
    border-collapse: collapse;
    font-size: 0.85rem;
    margin-top: 8px;
  }
  .method-table th {
    background: var(--navy);
    color: var(--paper);
    text-align: left;
    padding: 8px 10px;
  }
  .method-table td {
    padding: 8px 10px;
    border-bottom: 1px solid #ebe6dc;
    color: var(--ink);
    background: #fff;
  }
  .method-table td.pts { text-align: center; font-weight: 700; color: var(--accent); }
  .decision-grid {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(160px, 1fr));
    gap: 10px;
    margin-top: 10px;
  }
  .decision-box {
    border-radius: 8px;
    padding: 12px;
    border: 1px solid var(--line);
    color: var(--ink);
    background: var(--paper);
  }
  .decision-box strong { display: block; margin-bottom: 6px; color: var(--navy); }
</style>
</head>
<body>
<header>
  <h1>ISZE Tender OS v4 — Addressable Opportunity</h1>
  <p class="sub">GPS / PMS framing of the live tender pipeline. Figures are addressable opportunity if won — not expected revenue. Snapshot embeds all data for offline use.</p>
  <div id="sourceBanner" class="banner"></div>
</header>

<div class="wrap">
  <nav class="tabs" role="tablist" aria-label="Dashboard sections">
    <button type="button" class="tab active" data-view="opportunities" onclick="switchView('opportunities')">Opportunities</button>
    <button type="button" class="tab" data-view="howTenders" onclick="switchView('howTenders')">How tenders work</button>
    <button type="button" class="tab" data-view="byMarket" onclick="switchView('byMarket')">By market</button>
    <button type="button" class="tab" data-view="sources" onclick="switchView('sources')">Sources / health</button>
    <button type="button" class="tab" data-view="method" onclick="switchView('method')">Method / scoring</button>
    <button type="button" class="tab" data-view="dataSheets" onclick="switchView('dataSheets')">Export / Data sheets</button>
  </nav>

  <section class="view active" id="view-opportunities">
    <div class="cards" id="headlineCards"></div>
    <div class="fleet-disclaimer" id="fleetDisclaimer"></div>
    <div class="gps-explain" id="gpsExplainBanner"></div>

    <div class="panel" id="pmsPanel">
      <div class="pms-head">
        <div>
          <h2 style="margin:0">PMS breakdown</h2>
          <p class="subtle" id="pmsTitle" style="margin:6px 0 0"></p>
        </div>
        <button class="btn secondary" type="button" onclick="closePms()">Close</button>
      </div>
      <div id="pmsBody"></div>
    </div>

    <div class="panel">
      <h2>Opportunity table</h2>
      <p class="subtle" id="oppTableSubtle"></p>
      <div class="filters">
        <input id="qSearch" type="search" placeholder="Search buyer, title, country… (type CV = pickup+truck)" oninput="renderOppTable()">
        <select id="fCountry" onchange="renderOppTable()"><option value="">All markets</option></select>
        <select id="fDecision" onchange="renderOppTable()">
          <option value="">All decisions</option>
          <option value="Attack">Attack</option>
          <option value="Selective">Selective</option>
          <option value="Watch">Watch</option>
          <option value="Ignore">Ignore</option>
        </select>
        <select id="fSized" onchange="renderOppTable()">
          <option value="">All (with or without GPS €)</option>
          <option value="1">Has GPS € only</option>
          <option value="0">No GPS € (fleet/map/product-fit)</option>
          <option value="fleet0">No fleet count in notice</option>
          <option value="val_nofleet">Contract value but no fleet</option>
        </select>
        <select id="fSegment" onchange="renderOppTable()">
          <option value="cv" selected>CV (pickup + truck)</option>
          <option value="">All product segments</option>
          <option value="pickup">Pickup / D-Max-shaped</option>
          <option value="truck">Truck / chassis</option>
          <option value="lcv_van">Postal van / LCV (rare survivors)</option>
          <option value="rm_only">R&amp;M only</option>
          <option value="weak">Generic / weak fit</option>
        </select>
        <select id="fRm" onchange="renderOppTable()">
          <option value="">All contracts</option>
          <option value="1">R&amp;M only</option>
          <option value="aftersales">Any aftersales hook</option>
          <option value="service_contract">Service contract</option>
          <option value="parts">Parts</option>
          <option value="extended_warranty">Ext. warranty</option>
          <option value="full_service_lease">Full-service lease</option>
        </select>
      </div>
      <div class="table-wrap" style="max-height:620px"><table id="oppTable"></table></div>
      <p class="subtle" id="oppTableFooter" style="margin-top:10px"></p>
    </div>
  </section>

  <section class="view" id="view-howTenders">
    <div class="explain-grid" id="howTendersBody"></div>
  </section>

  <section class="view" id="view-byMarket">
    <div class="grid-2">
      <div class="panel">
        <h2>By market</h2>
        <p class="subtle">Hero column: <strong>stated fleet in captured notices (pickup/truck)</strong> — extract-subset only, not market demand. Addressable GPS € is secondary when sized. Low fleet count → coverage gap (portal map P0), not a tiny market. Detail columns stay hidden until toggled.</p>
        <div class="market-toolbar">
          <label><input type="checkbox" id="marketDetailToggle" onchange="toggleMarketDetail(this.checked)"> Show detail columns</label>
        </div>
        <div class="table-wrap"><table id="marketTable"></table></div>
      </div>
      <div class="panel">
        <h2>Coverage vs revenue</h2>
        <p class="subtle" id="coverageSubtle"></p>
        <div class="cov-legend" id="coverageLegend"></div>
        <div id="coveragePanel"></div>
      </div>
    </div>
  </section>

  <section class="view" id="view-sources">
    <div class="panel">
      <h2>Sources / health</h2>
      <p class="subtle">Live status from <code>_source_health.json</code> — OK / SUSPECT / BROKEN / AWAITING_KEY / CODE_READY / RETIRED / BLOCKED. Empty sheets are honest, not fake “no results” notices. Country coverage paper: <code>_COUNTRY_TENDER_PORTALS_MAP.md</code> — national LIVE ~55% (16/29; includes NO/FI/DK/PT/BG/CZ). Wire next: ES WAF. LT CODE_READY. HU/IE/SE BLOCKED with unblock paths. <strong>Do not</strong> ask Jerome for Nordics API keys or open those portals again.</p>
      <div class="table-wrap"><table id="sourcesTable"></table></div>
    </div>
  </section>

  <section class="view" id="view-method">
    <div class="explain-grid" id="methodBody"></div>
  </section>

  <section class="view" id="view-dataSheets">
    <div class="panel">
      <h2>Export / Data sheets</h2>
      <p class="subtle">Raw workbook worksheets in <code>isuzu_tenders_v3.xlsx</code> (source portals) plus the v4 opportunities export. Row counts are data rows excluding the header.</p>
      <div class="table-wrap"><table id="dataSheetsTable"></table></div>
      <p class="subtle" style="margin-top:12px" id="dataSheetsFooter"></p>
    </div>
  </section>
</div>

<script>
const DATA = __DATA__;

let marketSort = { key: 'vehicles_pickup_truck', dir: -1 };
let oppSort = { key: 'fleet_size', dir: -1 };
let selectedId = null;
let showMarketDetail = false;

function toggleMarketDetail(on) {
  showMarketDetail = !!on;
  document.body.classList.toggle('show-market-detail', showMarketDetail);
  renderMarketTable();
}

function esc(v) {
  return String(v ?? '').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
}
function n(v) { return Number(v || 0); }
function fmtCount(v) { return new Intl.NumberFormat('en-GB').format(v || 0); }
function fmtMoney(v) {
  if (v === null || v === undefined || v === '') return '—';
  return new Intl.NumberFormat('en-GB', { style: 'currency', currency: 'EUR', maximumFractionDigits: 0 }).format(v);
}
function fmtPct(v) {
  if (v === null || v === undefined || v === '') return '—';
  return (Number(v) * 100).toFixed(1) + '%';
}
function dash(v) {
  if (v === null || v === undefined || v === '') return '—';
  return v;
}

function countryHasGpsMap(c) {
  const gm = (DATA.gps_map && DATA.gps_map.by_country && DATA.gps_map.by_country[c]) || null;
  return !!(gm && gm.has_spu && gm.distributor);
}

function nordicsKeyAwaiting(c) {
  if (c !== 'Norway' && c !== 'Finland') return false;
  const sh = DATA.source_health || {};
  const needle = c === 'Norway' ? 'Norway' : 'Finland';
  return Object.entries(sh).some(([k, v]) =>
    String(k).indexOf(needle) >= 0 && v && v.status === 'AWAITING_KEY'
  );
}

function blankReasonForOpp(o) {
  if (hasGps(o)) return null;
  if (!hasFleet(o)) return 'no_fleet';
  const seg = o.product_segment || 'weak';
  if (seg === 'lcv_van' || seg === 'rm_only' || seg === 'weak') return 'van_only';
  return 'no_map';
}

function marketRows() {
  const map = {};
  (DATA.opportunities || []).forEach(o => {
    const c = o.country || '—';
    if (!map[c]) {
      map[c] = {
        market: c,
        notices: 0,
        opportunities: 0,
        action: 0,
        gps: 0,
        sized: 0,
        with_fleet: 0,
        no_fleet: 0,
        no_map: 0,
        van_only: 0,
        pickup: 0,
        truck: 0,
        lcv_van: 0,
        rm_count: 0,
        rm_share: null,
        capture: null,
        has_gps_map: countryHasGpsMap(c),
        awaiting_nordics: nordicsKeyAwaiting(c),
        _distSet: {},
        distributors_touched: [],
        dealers_assigned: 0,
        ids_unmapped: 0,
        vehicles_pickup_truck: 0,
        vehicles_all_stated: 0,
      };
    }
    const m = map[c];
    m.opportunities += 1;
    m.notices += Number(o.notice_count || 1);
    if (o.decision === 'Attack' || o.decision === 'Selective') m.action += 1;
    if (o.rm_contract) m.rm_count += 1;
    const seg = o.product_segment || 'weak';
    if (seg === 'pickup') m.pickup += 1;
    else if (seg === 'truck') m.truck += 1;
    else if (seg === 'lcv_van') m.lcv_van += 1;
    if (hasFleet(o)) {
      m.with_fleet += 1;
      const fn = fleetNum(o);
      if (fn !== null) {
        m.vehicles_all_stated += fn;
        if (seg === 'pickup' || seg === 'truck') m.vehicles_pickup_truck += fn;
      }
    }
    const dist = o.ids_distributor || o.distributor_code;
    if (dist) m._distSet[dist] = true;
    if (o.ids_dealer_name) m.dealers_assigned += 1;
    if ((o.ids_match_method || 'unmapped') === 'unmapped') m.ids_unmapped += 1;
    const g = o.gps_addressable_per_year_eur;
    if (g !== null && g !== undefined) {
      m.gps += Number(g);
      m.sized += 1;
    } else {
      const reason = blankReasonForOpp(o);
      if (reason === 'no_fleet') m.no_fleet += 1;
      else if (reason === 'no_map') m.no_map += 1;
      else if (reason === 'van_only') m.van_only += 1;
    }
    if (o.gps_capture_rate !== null && o.gps_capture_rate !== undefined) {
      const cap = Number(o.gps_capture_rate);
      if (cap <= 1.000001) m.capture = cap;
    }
  });
  Object.values(map).forEach(m => {
    m.rm_share = m.opportunities ? m.rm_count / m.opportunities : null;
    if (!m.sized) m.capture = null;
    m.distributors_touched = Object.keys(m._distSet).sort();
    delete m._distSet;
  });
  return Object.values(map);
}

function hasFleet(o) {
  return fleetNum(o) !== null;
}
function fleetNum(o) {
  const fs = o.fleet_size;
  if (fs === null || fs === undefined || fs === '') return null;
  const n = Number(fs);
  if (!Number.isFinite(n) || n <= 0) return null;
  return n;
}
function hasPubValue(o) {
  return o.value_amount !== null && o.value_amount !== undefined && Number(o.value_amount) > 0;
}
function hasGps(o) {
  return o.gps_addressable_per_year_eur !== null && o.gps_addressable_per_year_eur !== undefined;
}

function addressableVehicleTotals(opps) {
  let vehicles_pickup_truck = 0;
  let vehicles_all_stated = 0;
  let rows_with_fleet = 0;
  let rows_pickup_truck_with_fleet = 0;
  (opps || []).forEach(o => {
    const n = fleetNum(o);
    if (n === null) return;
    rows_with_fleet += 1;
    vehicles_all_stated += n;
    const seg = o.product_segment || 'weak';
    if (seg === 'pickup' || seg === 'truck') {
      vehicles_pickup_truck += n;
      rows_pickup_truck_with_fleet += 1;
    }
  });
  return { vehicles_pickup_truck, vehicles_all_stated, rows_with_fleet, rows_pickup_truck_with_fleet };
}

function whyChipsHtml(r) {
  const chips = [];
  if (r.no_fleet > 0) {
    chips.push(`<span class="pill why-nofleet kpi-tip" data-tip="Notice never states a vehicle count. We do not invent fleet size — GPS € stays null (not €0).">${fmtCount(r.no_fleet)} no fleet</span>`);
  }
  if (r.no_map > 0) {
    chips.push(`<span class="pill why-nomap kpi-tip" data-tip="Fleet + pickup/truck exist, but this country has no SPU row in gps_spu_by_distributor.csv (no GPS map in source pack). Null, not invented.">${fmtCount(r.no_map)} no GPS map</span>`);
  }
  if (r.van_only > 0) {
    chips.push(`<span class="pill why-van kpi-tip" data-tip="Fleet known but product-fit gate blocks GPS € (van/LCV, R&M-only, or weak fit). Pickup/network SPU is not applied. Null ≠ 0.">${fmtCount(r.van_only)} van-only (no GPS)</span>`);
  }
  if (r.awaiting_nordics) {
    chips.push(`<span class="pill why-nordics kpi-tip" data-tip="National NO/FI portal status is AWAITING_KEY in this build's health snapshot — diagnose nordics_api_keys.env + verify_nordics_keys.py silently. Keys are already registered; do not ask Jerome to re-paste the setup sheet. TED above-threshold rows may still appear.">awaiting Nordics key</span>`);
  }
  if (!r.has_gps_map && r.sized === 0 && r.no_map === 0) {
    chips.push(`<span class="pill why-pack kpi-tip" data-tip="This country has no distributor SPU in the GPS source pack. Even if a fleet count appears later, GPS € stays null until a real SPU row exists — never invented.">no GPS map in source pack</span>`);
  }
  if (!chips.length && r.sized > 0) {
    chips.push(`<span class="pill ok kpi-tip" data-tip="At least one opportunity has fleet + GPS map + pickup/truck product-fit.">GPS sized</span>`);
  }
  return chips.length ? `<div class="why-chips">${chips.join('')}</div>` : '<span class="muted">—</span>';
}

function segMixHtml(r) {
  const parts = [];
  if (r.pickup) parts.push(`${r.pickup} pickup`);
  if (r.truck) parts.push(`${r.truck} truck`);
  if (r.lcv_van) parts.push(`${r.lcv_van} van`);
  if (!parts.length) return '<span class="muted">—</span>';
  return `<span class="seg-mix kpi-tip" data-tip="Product-fit segment counts in this market. GPS € only applies to pickup/truck when fleet + map exist.">${esc(parts.join(' · '))}</span>`;
}

function reasonSummary(r) {
  const bits = [];
  if (r.no_fleet) bits.push(`${r.no_fleet} no fleet`);
  if (r.no_map) bits.push(`${r.no_map} no GPS map`);
  if (r.van_only) bits.push(`${r.van_only} van-only`);
  if (r.awaiting_nordics) bits.push('awaiting Nordics key');
  if (!r.has_gps_map && !r.sized) bits.push('no GPS map in source pack');
  return bits.join('; ') || (r.sized ? 'GPS sized' : '—');
}

function renderHeadline() {
  const opps = DATA.opportunities || [];
  const sized = opps.filter(hasGps);
  const withFleet = opps.filter(hasFleet);
  const noFleet = opps.length - withFleet.length;
  const withVal = opps.filter(hasPubValue);
  const valNoFleet = withVal.filter(o => !hasFleet(o));
  const sum = sized.length ? sized.reduce((a, o) => a + Number(o.gps_addressable_per_year_eur), 0) : null;
  const valSum = withVal.length ? withVal.reduce((a, o) => a + Number(o.value_amount), 0) : null;
  const gen = (DATA.meta && DATA.meta.generated_at) ? DATA.meta.generated_at : 'embedded snapshot';
  const fleetPct = opps.length ? (100 * withFleet.length / opps.length).toFixed(1) : '0';
  const veh = addressableVehicleTotals(opps);
  const asTagged = opps.filter(o => o.aftersales_tags && o.aftersales_tags.length).length;
  const cvOpps = opps.filter(isCvSegment).length;
  document.getElementById('headlineCards').innerHTML = `
    <div class="card hero-vehicles">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Sum of extracted fleet_size units on pickup/truck rows — only from the subset of tenders we currently scrape and extract. NOT tender count. NOT annual market demand. NOT sales potential. Three layers: (a) network ~20k LCV / ~78k CV/yr · (b) CV tenders in feed (${fmtCount(cvOpps)} this snapshot) · (c) this stated-fleet sum. Null fleet omitted. Snapshot: ${esc(gen)}.">Stated fleet in captured notices (pickup/truck)</span></div>
      <div class="big accent">${fmtCount(veh.vehicles_pickup_truck)}</div>
      <div class="sub">${fmtCount(veh.rows_pickup_truck_with_fleet)} notices with extract · ≠ ${fmtCount(cvOpps)} CV tenders · ≠ ~78k CV/yr</div>
    </div>
    <div class="card">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Sum of GPS €/year only on rows with a published vehicle count, a GPS market wallet rate, AND a product segment that matches the GPS mix (pickup / truck). Formula: fleet × (max_spend_eur ÷ uio_units). Supporting figure when sized — secondary to stated fleet in captured notices. Snapshot: ${esc(gen)}.">Addressable GPS / year</span></div>
      <div class="big">${fmtMoney(sum)}</div>
      <div class="sub">${fmtCount(sized.length)} sized of ${fmtCount(opps.length)} · money when map + fit</div>
    </div>
    <div class="card">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Sum of extracted fleet_size wherever the notice states a count — including van/LCV and weak-fit rows. Still an extract-subset of captured notices — not network sales volume. Null rows omitted. Snapshot: ${esc(gen)}.">Stated fleet all segments</span></div>
      <div class="big">${fmtCount(veh.vehicles_all_stated)}</div>
      <div class="sub">${fmtCount(veh.rows_with_fleet)} notices · stated in notices (incl van/LCV)</div>
    </div>
    <div class="card">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Count of opportunities where notice text or CPV matched aftersales language (service contract, R&M, parts, extended warranty, full-service lease). Never invented — empty when the notice is vehicles-only. See _AFTERSALES_TAGS.md. Snapshot: ${esc(gen)}.">Aftersales hooks</span></div>
      <div class="big accent">${fmtCount(asTagged)}</div>
      <div class="sub">${fmtCount(asTagged)} of ${fmtCount(opps.length)} mention service / parts / warranty</div>
    </div>
    <div class="card">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Sum of estimated contract values the notice itself published. Separate from GPS parts € — never converted into fleet or GPS. Missing value stays out of this sum (not treated as €0).">Published contract value</span></div>
      <div class="big">${fmtMoney(valSum)}</div>
      <div class="sub">${fmtCount(withVal.length)} publish a value · ${fmtCount(valNoFleet.length)} value but no fleet</div>
    </div>
    <div class="card">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Count of opportunity families whose notice never states how many vehicles. We do not invent a fleet size. Without a count there is no GPS €/year. These rows stay in the pipeline as leads (buyer, Fit, deadline).">No fleet count in notice</span></div>
      <div class="big">${fmtCount(noFleet)}</div>
      <div class="sub">${fleetPct}% of pipeline · ${fmtCount(withFleet.length)} state a count</div>
    </div>
    <div class="card coverage-dark">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Share of Isuzu-relevant public notices we already ingest (TED ± national) — not GPS-sized share, not network sales. Band Medium (2026-08-25). Country portal map: national LIVE ~55% of 29 footprint countries; ~38% DARK/BLOCKED; CODE_READY ES+LT — primary reason stated-fleet hero stays extract-thin. Wire P0: ES WAF. Nordics keys DONE. See _COUNTRY_TENDER_PORTALS_MAP.md.">Coverage vs dark</span></div>
      <div class="big">~35–55%</div>
      <div class="sub">~55% national LIVE · ~38% DARK/BLOCKED · ES/LT CODE_READY · see portal map</div>
    </div>
    <div class="card sanity-scale">
      <div class="eyebrow"><span class="kpi-tip" data-tip="Indicative Jerome / ISZE network annual volumes for order-of-magnitude sanity only — not from this scrape. Ask if stale. Mandatory layers: (a) network sales scale ~20k LCV / ~78k CV/yr · (b) CV tenders captured in feed · (c) fleet units stated in those tenders. ${fmtCount(veh.vehicles_pickup_truck)} ≠ ${fmtCount(cvOpps)} ≠ ~78k. If (c) &lt;&lt; (a) by &gt;10× without coverage caveats → FAIL delivery. Low fleet → portal map P0.">Sanity context — three layers</span></div>
      <div class="big">Jerome network scale (indicative): ~20k LCV / ~78k CV per year — tender radar coverage is a fraction; do not equate.</div>
      <div class="sub">${fmtCount(veh.vehicles_pickup_truck)} stated fleet units ≠ ${fmtCount(cvOpps)} CV tenders ≠ ~78k CV/yr · national LIVE ~55% · DARK/BLOCKED ~38% · wire P0: ES WAF</div>
    </div>`;
  document.getElementById('fleetDisclaimer').innerHTML =
    `<strong>Not annual market demand. Not sales potential.</strong> ` +
    `Only vehicle counts printed in the subset of tenders we currently scrape + extract. ` +
    `<strong>${fmtCount(veh.vehicles_pickup_truck)} stated fleet units ≠ ${fmtCount(cvOpps)} CV tenders ≠ ~78k CV annual sales.</strong> ` +
    `National LIVE ~55% of footprint; ~38% DARK/BLOCKED; CODE_READY ES+LT — why the extract stays thin. ` +
    `Coverage P0: ES WAF. Nordics keys DONE. See <code>_COUNTRY_TENDER_PORTALS_MAP.md</code>.`;
  document.getElementById('gpsExplainBanner').innerHTML =
    `<strong>${fmtCount(veh.vehicles_pickup_truck)} stated fleet units</strong> (pickup/truck extract)` +
    ` · <strong>${fmtCount(cvOpps)} CV tenders</strong> in feed (filter CV)` +
    ` · network ~78k CV/yr indicative` +
    ` · portal footprint: ~55% national LIVE / ~38% DARK-BLOCKED / ES+LT CODE_READY` +
    ` · GPS € ${fmtMoney(sum)} where sized (${fmtCount(sized.length)} of ${fmtCount(opps.length)})` +
    ` · three layers · null ≠ 0.`;
  document.getElementById('oppTableSubtle').innerHTML =
    `Default sort: <strong>Fleet count</strong> descending — vehicle volume first. ` +
    `Product filter <strong>CV (pickup + truck)</strong> = Isuzu commercial shortlist (${fmtCount((DATA.opportunities||[]).filter(isCvSegment).length)} rows this snapshot). ` +
    `Search <code>CV</code> does the same (not LCV vans). ` +
    `Fleet + Aftersales stay prominent. Empty GPS = no published count, no market map, or van/LCV gate. ` +
    `Click a row for detail. <span class="kpi-tip" data-tip="GPS € = fleet × max wallet per unit for pickup/truck only. Contract value = buyer-stated deal size. Neither invents the other. Fleet column shows extracted counts only — null stays —.">Fleet + two money signals stay separate</span>.`;
}

function renderSourceBanner() {
  const el = document.getElementById('sourceBanner');
  const sh = DATA.source_health;
  if (!sh || !Object.keys(sh).length) {
    el.className = 'banner warn';
    el.innerHTML = '<strong>Source health:</strong> <code>_source_health.json</code> not present — run a scrape with health tracking to populate this banner.';
    return;
  }
  const entries = Object.entries(sh);
  const broken = entries.filter(([, v]) => v && v.status === 'BROKEN');
  const awaiting = entries.filter(([, v]) => v && v.status === 'AWAITING_KEY');
  const codeReady = entries.filter(([, v]) => v && v.status === 'CODE_READY');
  const suspect = entries.filter(([, v]) => v && v.status === 'SUSPECT');
  // RETIRED / WONTFIX / SKIPPED are intentional — show as muted notes, never as BROKEN.
  const retired = entries.filter(([, v]) => v && (v.status === 'RETIRED' || v.status === 'WONTFIX' || v.status === 'SKIPPED'));
  const ok = entries.filter(([, v]) => v && v.status === 'OK');
  const fmt = (pairs) => pairs.map(([k, v]) => esc(k) + (v.last_error ? ' (' + esc(v.last_error) + ')' : '')).join('; ');
  if (broken.length) {
    el.className = 'banner bad';
    el.innerHTML = `<strong>Source health — BROKEN:</strong> ${fmt(broken)}` +
      (awaiting.length ? ` · <strong>AWAITING_KEY:</strong> ${awaiting.map(([k]) => esc(k)).join(', ')}` : '') +
      (codeReady.length ? ` · <strong>CODE_READY:</strong> ${codeReady.map(([k]) => esc(k)).join(', ')}` : '') +
      (suspect.length ? ` · <strong>SUSPECT:</strong> ${suspect.map(([k]) => esc(k)).join(', ')}` : '') +
      ` · ${ok.length} OK`;
  } else if (awaiting.length || codeReady.length || suspect.length) {
    el.className = 'banner warn';
    const parts = [];
    if (awaiting.length) parts.push(`<strong>AWAITING_KEY:</strong> ${fmt(awaiting)}`);
    if (codeReady.length) parts.push(`<strong>CODE_READY:</strong> ${fmt(codeReady)}`);
    if (suspect.length) parts.push(`<strong>SUSPECT (secondary):</strong> ${suspect.map(([k]) => esc(k)).join(', ')}`);
    if (retired.length) parts.push(`RETIRED/WONTFIX: ${retired.map(([k]) => esc(k)).join(', ')}`);
    parts.push(`${ok.length} OK`);
    el.innerHTML = `<strong>Source health:</strong> ${parts.join(' · ')}`;
  } else {
    el.className = 'banner ok';
    const note = retired.length
      ? ` · ${retired.length} retired/wontfix (ignored)`
      : '';
    el.innerHTML = `<strong>Source health:</strong> all ${ok.length} active sources OK${note}.`;
  }
}

function sortMarkets(rows) {
  const { key, dir } = marketSort;
  return [...rows].sort((a, b) => {
    const av = a[key], bv = b[key];
    if (av === null || av === undefined) return 1;
    if (bv === null || bv === undefined) return -1;
    if (typeof av === 'string') return dir * String(av).localeCompare(String(bv));
    return dir * (n(av) - n(bv));
  });
}

function setMarketSort(key) {
  if (marketSort.key === key) marketSort.dir *= -1;
  else { marketSort.key = key; marketSort.dir = key === 'market' ? 1 : -1; }
  renderMarketTable();
}

function idsDistRollupHtml(r) {
  const dists = r.distributors_touched || [];
  const tip = 'Distributors from country_distributor_map.yaml on opportunities in this market. Dealers from IDS 2026 rankings when a country sheet exists — never invented. Germany and other GPS-mapped markets without an IDS tab stay distributor-only / unmapped at dealer level.';
  if (!dists.length) {
    return `<span class="gps-empty kpi-tip" data-tip="${esc(tip)}">unmapped</span>`;
  }
  const dealerBit = r.dealers_assigned
    ? `${fmtCount(r.dealers_assigned)} with IDS dealer`
    : 'no IDS dealer sheet / unmapped dealers';
  return `<span class="kpi-tip" data-tip="${esc(tip)}"><strong>${esc(dists.join(', '))}</strong><div class="muted" style="font-size:0.75rem;margin-top:2px">${esc(dealerBit)}</div></span>`;
}

function renderMarketTable() {
  const rows = sortMarkets(marketRows());
  const ind = (k) => marketSort.key === k ? (marketSort.dir < 0 ? '▼' : '▲') : '';
  const head = `<thead><tr>
    <th onclick="setMarketSort('market')">Market<span class="sort-ind">${ind('market')}</span></th>
    <th class="th-veh" onclick="setMarketSort('vehicles_pickup_truck')">Stated fleet (pickup/truck)<span class="sort-ind">${ind('vehicles_pickup_truck')}</span></th>
    <th class="th-gps-sec" onclick="setMarketSort('gps')">Addressable GPS / yr<span class="sort-ind">${ind('gps')}</span></th>
    <th class="detail-col" onclick="setMarketSort('notices')">Notices<span class="sort-ind">${ind('notices')}</span></th>
    <th class="detail-col" onclick="setMarketSort('with_fleet')">With fleet<span class="sort-ind">${ind('with_fleet')}</span></th>
    <th class="detail-col" onclick="setMarketSort('sized')">GPS-sized<span class="sort-ind">${ind('sized')}</span></th>
    <th class="detail-col">Why blank</th>
    <th class="detail-col" onclick="setMarketSort('capture')">Capture<span class="sort-ind">${ind('capture')}</span></th>
    <th class="detail-col">Segment mix</th>
    <th class="detail-col">Distributors</th>
  </tr></thead>`;
  const body = rows.map(r => {
    const gpsCell = r.sized
      ? `<div class="gps-cell-main">${fmtMoney(r.gps)}</div><div class="gps-cell-n">${fmtCount(r.sized)} sized</div>`
      : `<span class="gps-empty kpi-tip" data-tip="${esc(reasonSummary(r))}">null</span>`;
    const capCell = r.sized
      ? `<span class="kpi-tip" data-tip="gps_capture_vs_max from latest GPS SPU row for the mapped distributor. Not a win rate. Never shown above 100%.">${fmtPct(r.capture)}</span>`
      : `<span class="gps-empty kpi-tip" data-tip="${esc('Capture rate only when GPS-sized > 0. ' + reasonSummary(r))}">—</span>`;
    const vehTip = `Stated fleet in captured notices = sum of extracted fleet_size on pickup+truck rows in this market (nulls omitted). Extract-subset only — NOT market potential, NOT annual demand. Stated all segments = sum wherever fleet extracted (incl van/LCV). Low totals → coverage gap, not a tiny market.`;
    const vehCell = (r.vehicles_pickup_truck || r.vehicles_all_stated)
      ? `<div class="veh-hero"><span class="kpi-tip" data-tip="${esc(vehTip)}">${fmtCount(r.vehicles_pickup_truck)}</span></div>` +
        (r.vehicles_all_stated !== r.vehicles_pickup_truck
          ? `<div class="veh-hero-sub">${fmtCount(r.vehicles_all_stated)} stated (incl van)</div>`
          : `<div class="veh-hero-sub">pickup/truck · extract-subset</div>`)
      : `<span class="gps-empty kpi-tip" data-tip="${esc(vehTip)}">—</span>`;
    return `<tr>
    <td>${esc(r.market)}</td>
    <td class="num">${vehCell}</td>
    <td class="num"><span class="kpi-tip" data-tip="Sum of gps_addressable_per_year_eur for sized rows. Secondary to stated fleet in captured notices. Null inputs excluded — never coerced to 0.">${gpsCell}</span></td>
    <td class="num detail-col"><span class="kpi-tip" data-tip="Sum of notice_count for opportunity families in this market.">${fmtCount(r.notices)}</span></td>
    <td class="num detail-col"><span class="kpi-tip" data-tip="Opportunities in this market whose notice states a vehicle count.">${fmtCount(r.with_fleet)}</span></td>
    <td class="num detail-col"><span class="kpi-tip" data-tip="Rows with fleet + GPS map + pickup/truck product-fit — the only rows that contribute GPS €.">${fmtCount(r.sized)}</span></td>
    <td class="detail-col">${whyChipsHtml(r)}</td>
    <td class="num detail-col">${capCell}</td>
    <td class="detail-col">${segMixHtml(r)}</td>
    <td class="detail-col">${idsDistRollupHtml(r)}</td>
  </tr>`;
  }).join('');
  document.getElementById('marketTable').innerHTML = head + `<tbody>${body}</tbody>`;
}

function renderCoverage() {
  const rows = marketRows();
  const totalOpp = rows.reduce((a, r) => a + r.opportunities, 0) || 1;
  const totalGps = rows.reduce((a, r) => a + r.gps, 0);
  const bp = DATA.revenue_shares;
  const hasRev = !!(bp && bp.available && bp.by_country);
  const subtle = document.getElementById('coverageSubtle');
  const legend = document.getElementById('coverageLegend');
  if (hasRev) {
    subtle.innerHTML = `<span class="kpi-tip" data-tip="${esc(bp.label || '')}">Pipeline opportunity share vs network revenue share</span> by market. BP snapshot: ${esc(bp.snapshot || '—')}. When no sized GPS, reason counts replace “no sized GPS”. Diagnostic only — not a forecast.`;
    legend.innerHTML = `<span class="l-opp">Share of opportunities</span><span class="l-rev">Share of network revenue</span><span class="l-gps">Share of addressable GPS</span>`;
  } else {
    subtle.innerHTML = `<span class="kpi-tip" data-tip="FY27 network revenue file not loaded — showing pipeline opportunity-count share vs sized GPS addressable share only. No revenue shares invented.">Share of opportunities vs share of addressable GPS</span> by market. Pipeline-share only (no network revenue file in this build).`;
    legend.innerHTML = `<span class="l-opp">Share of opportunities</span><span class="l-gps">Share of addressable GPS</span>`;
  }
  const ranked = [...rows].sort((a, b) => b.opportunities - a.opportunities).slice(0, 16);
  const gpsDenom = totalGps > 0 ? totalGps : 1;
  document.getElementById('coveragePanel').innerHTML = ranked.map(r => {
    const so = 100 * r.opportunities / totalOpp;
    const sg = totalGps > 0 && r.gps ? (100 * r.gps / gpsDenom) : 0;
    const revShare = hasRev && bp.by_country[r.market] ? Number(bp.by_country[r.market].revenue_share) : null;
    const sr = revShare != null ? (100 * revShare) : null;
    const revBar = hasRev
      ? `<div class="cov-track"><div class="cov-fill-rev" style="width:${sr != null ? Math.max(2, sr) : 0}%"></div></div>`
      : '';
    const revTxt = hasRev
      ? (sr != null ? ` · ${sr.toFixed(1)}% of network revenue` : ' · no revenue share mapped')
      : '';
    const gpsTxt = r.sized
      ? `${sg.toFixed(1)}% of addressable GPS`
      : reasonSummary(r);
    return `<div class="cov-row">
      <div>${esc(r.market)}</div>
      <div class="cov-bars">
        <div class="cov-track"><div class="cov-fill-opp" style="width:${Math.max(2, so)}%"></div></div>
        ${revBar}
        <div class="cov-track"><div class="cov-fill-gps" style="width:${r.gps ? Math.max(2, sg) : 0}%"></div></div>
        <div class="muted" style="font-size:0.75rem">${so.toFixed(1)}% of opps${revTxt} · ${esc(gpsTxt)}</div>
        ${r.sized ? '' : whyChipsHtml(r)}
      </div>
    </div>`;
  }).join('') || '<div class="muted">No markets in snapshot.</div>';
}

function isCvSegment(o) {
  const seg = o.product_segment || 'weak';
  return seg === 'pickup' || seg === 'truck';
}

function filteredOpps() {
  const q = (document.getElementById('qSearch').value || '').trim().toLowerCase();
  const fc = document.getElementById('fCountry').value;
  const fd = document.getElementById('fDecision').value;
  const fs = document.getElementById('fSized').value;
  const fseg = document.getElementById('fSegment').value;
  const fr = document.getElementById('fRm').value;
  // "cv" / "commercial vehicle" = Isuzu CV shortlist (pickup+truck), not LCV substring
  const qIsCv = (q === 'cv' || q === 'cvs' || q === 'commercial vehicle' || q === 'commercial vehicles');
  return (DATA.opportunities || []).filter(o => {
    if (fc && o.country !== fc) return false;
    if (fd && o.decision !== fd) return false;
    if (fseg === 'cv') {
      if (!isCvSegment(o)) return false;
    } else if (fseg && (o.product_segment || 'weak') !== fseg) {
      return false;
    }
    const sized = hasGps(o);
    const fleet = hasFleet(o);
    const val = hasPubValue(o);
    if (fs === '1' && !sized) return false;
    if (fs === '0' && sized) return false;
    if (fs === 'fleet0' && fleet) return false;
    if (fs === 'val_nofleet' && !(val && !fleet)) return false;
    if (fr === '1' && !o.rm_contract) return false;
    if (fr === 'aftersales' && !(o.aftersales_tags && o.aftersales_tags.length)) return false;
    if (fr && fr !== '1' && fr !== 'aftersales' && fr !== '') {
      const tags = o.aftersales_tags || [];
      if (!tags.includes(fr)) return false;
    }
    if (q) {
      if (qIsCv) return isCvSegment(o);
      const hay = [o.title, o.buyer, o.country, o.decision, o.distributor_code, o.ids_distributor, o.ids_dealer_name, o.ids_dealer_city, o.product_segment, o.product_segment_label, o.aftersales_summary, (o.aftersales_tags||[]).join(' ')].join(' ').toLowerCase();
      if (!hay.includes(q)) return false;
    }
    return true;
  });
}

function sortOpps(rows) {
  const { key, dir } = oppSort;
  return [...rows].sort((a, b) => {
    let av = a[key], bv = b[key];
    const aNull = av === null || av === undefined || av === '';
    const bNull = bv === null || bv === undefined || bv === '';
    if (aNull && bNull) return 0;
    if (aNull) return 1;  // nulls last when DESC
    if (bNull) return -1;
    if (typeof av === 'string' || typeof bv === 'string') return dir * String(av).localeCompare(String(bv));
    return dir * (n(av) - n(bv));
  });
}

function setOppSort(key) {
  if (oppSort.key === key) oppSort.dir *= -1;
  else { oppSort.key = key; oppSort.dir = (key === 'title' || key === 'buyer' || key === 'country' || key === 'product_segment' || key === 'ids_dealer_name') ? 1 : -1; }
  renderOppTable();
}

function segChip(o) {
  const code = o.product_segment || 'weak';
  const label = o.product_segment_label || code;
  return `<span class="pill seg-${esc(code)} kpi-tip" data-tip="Product segment from title/CPV — drives whether GPS € can apply. GPS wallet is pickup/network UIO; pure van/LCV stays null (no N-Series SPU invented).">${esc(label)}</span>`;
}

function idsDealerTip(o) {
  const method = o.ids_match_method || 'unmapped';
  const conf = o.ids_match_confidence != null ? ` · confidence ${o.ids_match_confidence}` : '';
  const label = o.ids_distributor_label ? ` (${o.ids_distributor_label})` : '';
  if (method === 'geo_nearest') {
    return `IDS match: geo_nearest${conf}. Buyer/title hit dealer name or city token already in IDS rankings. Distributor ${o.ids_distributor || '—'}${label}. Rank #${o.ids_dealer_rank || '—'} on tab ${o.ids_tab || '—'}. Never invented.`;
  }
  if (method === 'country_primary') {
    return `IDS match: country_primary${conf}. No city hit in buyer/title — assigned IDS rank #1 for this country (${o.ids_tab || '—'}). Distributor ${o.ids_distributor || '—'}${label}. Not a geo claim.`;
  }
  if (o.ids_distributor) {
    return `IDS match: unmapped at dealer level. Distributor ${o.ids_distributor}${label} from yaml — no IDS dealer sheet for this country (e.g. Germany). Drop a *dealer_ranking.xlsx / DE tab to unlock.`;
  }
  return 'IDS match: unmapped. No country→distributor in yaml and no IDS dealers. Null, not invented.';
}

function idsDealerChip(o) {
  const tip = idsDealerTip(o);
  const dist = o.ids_distributor || o.distributor_code;
  if (o.ids_dealer_name) {
    const city = o.ids_dealer_city ? ` · ${o.ids_dealer_city}` : '';
    return `<span class="pill ids-dealer kpi-tip" data-tip="${esc(tip)}">${esc(o.ids_dealer_name)}${esc(city)}<div class="muted" style="font-size:0.72rem;margin-top:2px">${esc(dist || '')}</div></span>`;
  }
  if (dist) {
    return `<span class="pill ids-dist kpi-tip" data-tip="${esc(tip)}">${esc(dist)}<div class="muted" style="font-size:0.72rem;margin-top:2px">no IDS dealer</div></span>`;
  }
  return `<span class="gps-empty kpi-tip" data-tip="${esc(tip)}">unmapped</span>`;
}


const AS_LABELS = {
  service_contract: 'Service contract',
  repair_maintenance: 'R&M',
  parts: 'Parts',
  extended_warranty: 'Ext. warranty',
  full_service_lease: 'Full-service lease'
};

function aftersalesChips(o) {
  const tags = o.aftersales_tags || [];
  const tip = o.aftersales_match_tip
    || (tags.length
      ? ('Matched in notice — ' + (o.aftersales_summary || tags.join(', ')))
      : 'No service/parts/warranty language found in notice');
  if (!tags.length) {
    return `<span class="pill as-empty kpi-tip" data-tip="${esc(tip)}">vehicles only</span>`;
  }
  return `<div class="as-chips kpi-tip" data-tip="${esc(tip)}">${tags.map(t =>
    `<span class="pill as-${esc(t)}">${esc(AS_LABELS[t] || t)}</span>`
  ).join('')}</div>`;
}

function renderOppTable() {
  const all = DATA.opportunities || [];
  const withFleet = all.filter(hasFleet).length;
  const rows = sortOpps(filteredOpps());
  const ind = (k) => oppSort.key === k ? (oppSort.dir < 0 ? '▼' : '▲') : '';
  const head = `<thead><tr>
    <th onclick="setOppSort('decision')">Decision<span class="sort-ind">${ind('decision')}</span></th>
    <th onclick="setOppSort('product_segment')">Product<span class="sort-ind">${ind('product_segment')}</span></th>
    <th onclick="setOppSort('buyer')">Buyer<span class="sort-ind">${ind('buyer')}</span></th>
    <th onclick="setOppSort('country')">Country<span class="sort-ind">${ind('country')}</span></th>
    <th class="th-veh" onclick="setOppSort('fleet_size')">Fleet<span class="sort-ind">${ind('fleet_size')}</span></th>
    <th onclick="setOppSort('aftersales_summary')">Aftersales<span class="sort-ind">${ind('aftersales_summary')}</span></th>
    <th onclick="setOppSort('ids_dealer_name')">Dealer / IDS<span class="sort-ind">${ind('ids_dealer_name')}</span></th>
    <th onclick="setOppSort('deadline_date')">Deadline<span class="sort-ind">${ind('deadline_date')}</span></th>
    <th onclick="setOppSort('value_amount')">Contract value<span class="sort-ind">${ind('value_amount')}</span></th>
    <th class="th-gps-sec" onclick="setOppSort('gps_addressable_per_year_eur')">GPS € / yr<span class="sort-ind">${ind('gps_addressable_per_year_eur')}</span></th>
  </tr></thead>`;
  const body = rows.map(o => {
    const fleet = hasFleet(o);
    const gps = hasGps(o);
    const val = hasPubValue(o);
    const seg = o.product_segment || 'weak';
    let tip, gpsCell;
    if (gps) {
      tip = o.opportunity_basis
        ? esc(o.opportunity_basis)
        : 'fleet × (max_spend ÷ UIO) for this market. Addressable if won — not expected revenue.';
      gpsCell = `<span class="kpi-tip" data-tip="${tip}">${fmtMoney(o.gps_addressable_per_year_eur)}</span>`;
    } else if (fleet && (seg === 'lcv_van' || seg === 'rm_only' || seg === 'weak')) {
      tip = o.opportunity_basis
        ? esc(o.opportunity_basis)
        : 'Fleet known but product segment is not pickup/truck GPS-eligible — GPS € null (not invented).';
      const short = seg === 'lcv_van' ? '— van/LCV not GPS' : (seg === 'rm_only' ? '— R&M only' : '— weak fit');
      gpsCell = `<span class="gps-empty kpi-tip" data-tip="${tip}">${short}</span>`;
    } else if (fleet) {
      tip = o.opportunity_basis
        ? esc(o.opportunity_basis)
        : 'Fleet count known, but this country has no GPS wallet mapping — GPS € stays blank (not invented).';
      gpsCell = `<span class="gps-empty kpi-tip" data-tip="${tip}">— no GPS map</span>`;
    } else {
      tip = 'Notice does not state how many vehicles. GPS € only when a count is published. Not €0 — unknown.';
      gpsCell = `<span class="gps-empty kpi-tip" data-tip="${tip}">— no fleet in notice</span>`;
    }
    let valCell;
    if (val) {
      const badge = (!fleet)
        ? `<div class="money-signal kpi-tip" data-tip="Buyer-published contract value. Not GPS parts € and not a stand-in for fleet size.">contract € (no fleet)</div>`
        : '';
      valCell = `${fmtMoney(o.value_amount)}${badge}`;
    } else {
      valCell = '—';
    }
    const fleetCell = fleet
      ? `<span class="kpi-tip" data-tip="Extracted fleet_size from notice text. Never invented. Primary volume signal per opportunity.">${fmtCount(o.fleet_size)}</span>`
      : `<span class="gps-empty kpi-tip" data-tip="Vehicle count not published in the notice text we have. We never invent it. Null ≠ 0.">—</span>`;
    return `<tr class="${selectedId === o.opportunity_id ? 'selected' : ''}" onclick="openPms('${esc(o.opportunity_id)}')">
      <td><span class="pill ${esc(o.decision || '')}">${esc(o.decision || '—')}</span>
        <div class="muted" style="font-size:0.78rem;margin-top:4px;max-width:220px">${esc((o.title || '').slice(0, 80))}${(o.title || '').length > 80 ? '…' : ''}</div></td>
      <td>${segChip(o)}</td>
      <td>${esc(o.buyer || '—')}</td>
      <td>${esc(o.country || '—')}</td>
      <td class="num fleet-emph">${fleetCell}</td>
      <td>${aftersalesChips(o)}</td>
      <td>${idsDealerChip(o)}</td>
      <td>${esc(dash(o.deadline_date))}</td>
      <td class="num">${valCell}</td>
      <td class="num">${gpsCell}</td>
    </tr>`;
  }).join('') || `<tr><td colspan="10"><div class="muted" style="padding:12px">No rows match filters.</div></td></tr>`;
  document.getElementById('oppTable').innerHTML = head + `<tbody>${body}</tbody>`;
  const veh = addressableVehicleTotals(all);
  document.getElementById('oppTableFooter').innerHTML =
    `Showing ${fmtCount(rows.length)} of ${fmtCount(all.length)} · ` +
    `<strong>${fmtCount(veh.vehicles_pickup_truck)}</strong> stated fleet pickup/truck (captured notices) · ` +
    `<strong>${fmtCount(veh.vehicles_all_stated)}</strong> stated (incl van/LCV) · ` +
    `not market potential · GPS € only when the notice states how many vehicles — ` +
    `<strong>${fmtCount(withFleet)} of ${fmtCount(all.length)}</strong> have a count · product-fit gate strips van/LCV GPS · Dealer/IDS from real rankings (null = unmapped).`;
}

function openPms(id) {
  const o = (DATA.opportunities || []).find(x => x.opportunity_id === id);
  if (!o) return;
  selectedId = id;
  const panel = document.getElementById('pmsPanel');
  panel.classList.add('open');
  document.getElementById('pmsTitle').textContent = (o.buyer || '—') + ' · ' + (o.country || '—') + ' · ' + (o.title || '');
  const cats = o.pms_demand_by_category || [];
  let table = '';
  if (cats.length) {
    table = `<div class="table-wrap"><table>
      <thead><tr><th>Category</th><th>Qty demand</th><th>Avg SP €</th><th>Addressable €</th></tr></thead>
      <tbody>${cats.map(c => `<tr>
        <td>${esc(c.category_label || c.pms_category || '—')}</td>
        <td class="num">${c.expected_qty != null ? Number(c.expected_qty).toFixed(2) : '—'}</td>
        <td class="num">${c.avg_sp_eur != null ? Number(c.avg_sp_eur).toFixed(2) : '—'}</td>
        <td class="num"><span class="kpi-tip" data-tip="fleet_size × expected_qty_per_uio × avg_sp_eur for this PMS category (genuine line, latest year). Addressable opportunity.">${c.addressable_eur != null ? fmtMoney(c.addressable_eur) : '—'}</span></td>
      </tr>`).join('')}</tbody></table></div>`;
  } else {
    table = `<div class="muted" style="padding:8px 0">No PMS category demand — no published fleet count, or this distributor has no PMS rows.</div>`;
  }
  document.getElementById('pmsBody').innerHTML = `
    <div class="kv-grid">
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="Vehicle count only when the notice text states it. Never invented.">Fleet count in notice</span></div><div class="v">${hasFleet(o) ? fmtCount(o.fleet_size) : '— not published'}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="Product reality from title/CPV. GPS € only for pickup/truck when fleet + map exist. Pure van/LCV → null (no N-Series SPU invented).">Product segment</span></div><div class="v">${segChip(o)}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="${esc(o.opportunity_basis || 'Null — needs published fleet count, GPS market wallet, and product-fit (pickup/truck). Not included: guessed fleets, win rates, van GPS fantasy.')}">GPS € / yr</span></div><div class="v">${hasGps(o) ? fmtMoney(o.gps_addressable_per_year_eur) : (hasFleet(o) ? ((o.product_segment === 'lcv_van' || o.product_segment === 'rm_only' || o.product_segment === 'weak') ? '— product-fit gate' : '— no GPS map for this country') : '— no fleet count')}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="fleet × actual spend_per_unit_eur (today's capture intensity). Only when GPS € exists.">At current capture</span></div><div class="v">${o.capture_at_current_rate_eur != null ? fmtMoney(o.capture_at_current_rate_eur) : '—'}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="Full GPS addressable wallet if the fleet sat under authorised R&M. Same inputs as GPS €.">Under R&amp;M ceiling</span></div><div class="v">${o.capture_under_rm_eur != null ? fmtMoney(o.capture_under_rm_eur) : '—'}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="Buyer-published contract value from the notice. Separate from GPS parts € — neither invents the other.">Published contract value</span></div><div class="v">${hasPubValue(o) ? fmtMoney(o.value_amount) + (!hasFleet(o) ? ' <span class="money-signal">contract € only</span>' : '') : '—'}</div></div>
      <div class="kv"><div class="k">Distributor</div><div class="v">${esc(o.ids_distributor || o.distributor_code || '—')}${o.ids_distributor_label ? ' · ' + esc(o.ids_distributor_label) : ''}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="${esc(idsDealerTip(o))}">Dealer / IDS</span></div><div class="v">${o.ids_dealer_name ? esc(o.ids_dealer_name) + (o.ids_dealer_city ? ' · ' + esc(o.ids_dealer_city) : '') + ' · ' + esc(o.ids_match_method || '') : (o.ids_distributor ? 'distributor only · ' + esc(o.ids_match_method || 'unmapped') : 'unmapped')}</div></div>
      <div class="kv"><div class="k"><span class="kpi-tip" data-tip="${esc(o.aftersales_match_tip || 'No service/parts/warranty language found in notice')}">Aftersales</span></div><div class="v">${aftersalesChips(o)}</div></div>
    </div>
    ${!hasFleet(o) ? `<p class="subtle" style="margin-top:12px">This is still a real tender lead (buyer / Fit / deadline). Without a published vehicle count we cannot size GPS parts € — and we will not invent one.</p>` : ''}
    ${o.url ? `<p style="margin-top:12px"><a class="btn secondary" href="${esc(o.url)}" target="_blank" rel="noopener">Open notice</a></p>` : ''}
    ${o.opportunity_basis ? `<p class="subtle" style="margin-top:12px">${esc(o.opportunity_basis)}</p>` : ''}
    <h3 style="margin:16px 0 8px;font-size:1rem;color:var(--ink)">PMS demand by category</h3>
    ${table}`;
  renderOppTable();
  panel.scrollIntoView({ behavior: 'smooth', block: 'nearest' });
}

function closePms() {
  selectedId = null;
  document.getElementById('pmsPanel').classList.remove('open');
  renderOppTable();
}

function populateFilters() {
  const countries = [...new Set((DATA.opportunities || []).map(o => o.country).filter(Boolean))].sort();
  const sel = document.getElementById('fCountry');
  countries.forEach(c => {
    const opt = document.createElement('option');
    opt.value = c;
    opt.textContent = c;
    sel.appendChild(opt);
  });
}

function switchView(view) {
  document.querySelectorAll('.tab').forEach(el => el.classList.toggle('active', el.dataset.view === view));
  document.querySelectorAll('.view').forEach(el => el.classList.toggle('active', el.id === 'view-' + view));
  if (view === 'howTenders') renderHowTenders();
  if (view === 'byMarket') { renderMarketTable(); renderCoverage(); }
  if (view === 'sources') renderSourcesTable();
  if (view === 'method') renderMethod();
  if (view === 'dataSheets') renderDataSheets();
  if (view === 'opportunities') renderOppTable();
}

function liveSnapshotBits() {
  const opps = DATA.opportunities || [];
  const stats = (DATA.meta && DATA.meta.stats) || {};
  const veh = addressableVehicleTotals(opps);
  const asTagged = opps.filter(o => o.aftersales_tags && o.aftersales_tags.length).length;
  const withFleet = opps.filter(o => o.fleet_size != null && Number(o.fleet_size) > 0).length;
  const sized = opps.filter(o => o.gps_addressable_per_year_eur != null).length;
  let gpsSum = 0;
  opps.forEach(o => {
    const v = o.gps_addressable_per_year_eur;
    if (v != null && !isNaN(Number(v))) gpsSum += Number(v);
  });
  return {
    nOpps: opps.length || stats.opportunities || 0,
    vehPickup: veh.vehicles_pickup_truck,
    vehAll: veh.vehicles_all_stated,
    asTagged,
    withFleet,
    sized,
    gpsSum,
    gen: (DATA.meta && DATA.meta.generated_at) || 'snapshot',
  };
}

function renderHowTenders() {
  const s = liveSnapshotBits();
  const pctFleet = s.nOpps ? Math.round(1000 * s.withFleet / s.nOpps) / 10 : 0;
  const cvN = (DATA.opportunities || []).filter(isCvSegment).length;
  document.getElementById('howTendersBody').innerHTML = `
    <div class="explain-card">
      <h3>Three layers (never conflate)</h3>
      <p><strong>1. Network scale (indicative):</strong> ~20k LCV / ~78k CV sold per year across Jerome’s footprint — order-of-magnitude sanity only; ask if stale.</p>
      <p><strong>2. CV tenders captured:</strong> <strong>${fmtCount(cvN)}</strong> pickup+truck of <strong>${fmtCount(s.nOpps)}</strong> opportunities this snapshot (product filter <strong>CV</strong>). CV <em>is</em> in the feed.</p>
      <p><strong>3. Stated fleet units:</strong> <strong>${fmtCount(s.vehPickup)}</strong> = sum of extracted fleet counts on pickup/truck notices that publish a number — <em>not</em> market potential, <em>not</em> tender count.</p>
      <div class="callout"><strong>${fmtCount(s.vehPickup)} ≠ ${fmtCount(cvN)} ≠ ~78k.</strong> Tiny fleet extract is expected while national LIVE is only ~55% of footprint countries and ~38% stay DARK/BLOCKED (ES/LT CODE_READY). Diagnosis = <strong>coverage P0</strong>, not a tiny Isuzu market. Paper: <code>_COUNTRY_TENDER_PORTALS_MAP.md</code>. Wire order: ES WAF (Nordics/DK/PT/BG/CZ DONE).</div>
    </div>
    <div class="explain-card">
      <h3>What you are looking at</h3>
      <p>Public buyers publish <strong>tender notices</strong> when they want to buy vehicles, service, or parts. This dashboard turns those notices into <strong>opportunity families</strong> (same buyer + title + country collapsed together) and scores them for Isuzu Europe (ISZE).</p>
      <p>Live: GPS floor <strong>${fmtMoney(s.gpsSum)}</strong> where fleet + map + product-fit allow · Aftersales tagged <strong>${fmtCount(s.asTagged)}</strong> of ${fmtCount(s.nOpps)} · six tabs stay: Opportunities · How tenders · By market · Sources · Method · Export.</p>
    </div>
    <div class="explain-card">
      <h3>Notice, CPV, deadline</h3>
      <p><strong>Notice</strong> = the public advert (TED, BOAMP, BZP, Contracts Finder…). It has a buyer, a title, often a link to the full PDF.</p>
      <p><strong>CPV</strong> = Common Procurement Vocabulary — a product code. Vehicle buys sit in 341xxxx; repair &amp; maintenance in 501xxxx. We use CPV plus title language to decide product segment (pickup / truck / van / R&amp;M).</p>
      <p><strong>Deadline</strong> = last day to bid. Many TED list rows omit it — missing deadline is a data gap, not “closed.”</p>
    </div>
    <div class="explain-card">
      <h3>Fleet count vs GPS €</h3>
      <p><strong>Fleet</strong> = how many vehicles the notice literally states. We never invent a count. On this snapshot only <strong>${fmtCount(s.withFleet)}</strong> of ${fmtCount(s.nOpps)} (${pctFleet}%) publish a number.</p>
      <p><strong>GPS € / year</strong> = addressable parts wallet if won: fleet × (max GPS spend ÷ UIO units) for that market — and only when product-fit is pickup or truck. Postal/courier panel vans are hard-dropped from the desk (not GPS-null clutter). Live floor is ~${fmtMoney(s.gpsSum)}, not a historical €30k figure.</p>
      <div class="callout"><strong>Unsized rows are still real leads.</strong> Empty GPS means “no published count / no map / van blocked,” not “worthless tender.”</div>
    </div>
    <div class="explain-card">
      <h3>Product segments</h3>
      <p><strong>Pickup / truck</strong> feed the stated-fleet-in-captured-notices sum and can carry GPS €. That sum is <em>not</em> market potential. Default product filter is <strong>CV (pickup + truck)</strong>. Postal/courier panel vans are <strong>hard-dropped</strong> (see <code>_NOT_FOR_US_GATE.md</code>) — not kept “just in case.” <strong>R&amp;M only</strong> is aftersales approach, not a new vehicle buy.</p>
    </div>
    <div class="explain-card">
      <h3>Aftersales play (ISZE)</h3>
      <p>If the notice mentions a <strong>service contract</strong>, repair &amp; maintenance, <strong>parts</strong>, extended <strong>warranty</strong>, or full-service lease, the Aftersales column shows chips. That is an ISZE aftersales play — even when the row is vehicles-led.</p>
      <p>Empty Aftersales = <em>vehicles only</em> (no service/parts/warranty language found). Live: <strong>${fmtCount(s.asTagged)}</strong> of ${fmtCount(s.nOpps)} tagged. See <code>_AFTERSALES_TAGS.md</code>.</p>
    </div>
    <div class="explain-card">
      <h3>Where to go next</h3>
      <p><strong>Opportunities</strong> — Fleet + Aftersales; filter <strong>CV</strong>. <strong>By market</strong> — stated fleet first; low fleet → coverage P0. <strong>Sources / health</strong> — live scrapers + portal map. <strong>Method</strong> · <strong>Export</strong>. Full country map: <code>_COUNTRY_TENDER_PORTALS_MAP.md</code>.</p>
    </div>`;
}

function renderSourcesTable() {
  const sh = DATA.source_health || {};
  const keys = Object.keys(sh).sort();
  const statusClass = (st) => {
    const u = String(st || '').toUpperCase();
    if (u === 'OK') return 'ok';
    if (u === 'SUSPECT' || u === 'AWAITING_KEY' || u === 'CODE_READY') return 'warn';
    if (u === 'BROKEN' || u === 'RETIRED') return 'bad';
    return '';
  };
  const head = `<thead><tr><th>Source</th><th>Status</th><th>Last count</th><th>Zero runs</th><th>Note / last error</th></tr></thead>`;
  let body;
  if (!keys.length) {
    body = `<tbody><tr><td colspan="5" class="muted">No <code>_source_health.json</code> embedded — run a scrape with health tracking.</td></tr></tbody>`;
  } else {
    body = `<tbody>${keys.map(k => {
      const row = sh[k] || {};
      const st = row.status || '—';
      const note = row.note || row.last_error || '—';
      return `<tr style="cursor:default">
        <td>${esc(k)}</td>
        <td><span class="pill ${statusClass(st)}">${esc(st)}</span></td>
        <td class="num">${row.last_count != null ? fmtCount(row.last_count) : '—'}</td>
        <td class="num">${row.zero_runs != null ? fmtCount(row.zero_runs) : '—'}</td>
        <td class="muted">${esc(String(note))}</td>
      </tr>`;
    }).join('')}</tbody>`;
  }
  document.getElementById('sourcesTable').innerHTML = head + body;
}

function renderMethod() {
  document.getElementById('methodBody').innerHTML = `
    <div class="explain-card">
      <h3>How tenders are scored</h3>
      <p>Every opportunity gets three scores. They drive Attack / Selective / Watch / Ignore and pipeline order. Ported from the v2 Scoring Guide — weights unchanged.</p>
      <div style="font-family:Consolas,monospace;background:#15233c;color:#f7f4ee;padding:12px 14px;border-radius:8px;margin-top:8px;border-left:3px solid #c0552d">
        Opportunity Score = Fit × 0.4 + Win × 0.3 + Econ × 0.3
      </div>
    </div>
    <div class="explain-card">
      <span class="score-badge fit">FIT</span><strong>Strategic Fit 0–100</strong>
      <p style="margin-top:8px">How well the notice matches Isuzu products (D-Max, N/F-Series, genuine parts &amp; service).</p>
      <table class="method-table">
        <thead><tr><th>Signal</th><th>Pts</th><th>What triggers it</th></tr></thead>
        <tbody>
          <tr><td>ISZE brand signal</td><td class="pts">+40</td><td>isuzu / d-max / n-series / f-series in title or description</td></tr>
          <tr><td>Service / parts signal</td><td class="pts">+25</td><td>CPV 501xx or service/parts terminology</td></tr>
          <tr><td>Target CPV match</td><td class="pts">+20</td><td>Exact hit on ISZE CPV list (pickups, vans, R&amp;M…)</td></tr>
          <tr><td>Vehicle procurement</td><td class="pts">+15</td><td>Generic vehicle CPV <em>plus</em> brand or service qualifier</td></tr>
          <tr><td>Pickup / light truck</td><td class="pts">+12</td><td>Pickup vocabulary (pickup, pritschenwagen…)</td></tr>
          <tr><td>Framework route</td><td class="pts">+8</td><td>Multi-lot framework — recurring call-offs</td></tr>
        </tbody>
      </table>
      <div class="callout">Design intent: “vehicles” alone without a second qualifier scores ≤5 → Ignore. Two signals needed for Watch or above.</div>
    </div>
    <div class="explain-card">
      <span class="score-badge win">WIN</span><strong>Win Feasibility 0–100</strong>
      <p style="margin-top:8px">Can the distributor act before close? Timing, completeness, stage.</p>
      <table class="method-table">
        <thead><tr><th>Factor</th><th>Pts</th><th>Notes</th></tr></thead>
        <tbody>
          <tr><td>Base</td><td class="pts">20</td><td>Every tender starts here</td></tr>
          <tr><td>Deadline &gt; 21 days</td><td class="pts">+25</td><td>Good runway</td></tr>
          <tr><td>Deadline 7–21 days</td><td class="pts">+15</td><td>Window tightening</td></tr>
          <tr><td>Deadline &lt; 7 days</td><td class="pts">+5</td><td>Urgent</td></tr>
          <tr><td>Deadline missing</td><td class="pts">±0</td><td>TED often omits — check full notice</td></tr>
          <tr><td>Deadline passed</td><td class="pts">−30</td><td>Closed — award / R&amp;M approach only</td></tr>
          <tr><td>Buyer / country / evidence</td><td class="pts">+8–20</td><td>Known buyer, country, complete fields</td></tr>
          <tr><td>Award / result stage</td><td class="pts">−15</td><td>Already awarded</td></tr>
        </tbody>
      </table>
    </div>
    <div class="explain-card">
      <span class="score-badge econ">ECON</span><strong>Economic Attractiveness 0–100</strong>
      <p style="margin-top:8px">Contract size, recurring parts, fleet scale.</p>
      <table class="method-table">
        <thead><tr><th>Factor</th><th>Pts</th><th>Notes</th></tr></thead>
        <tbody>
          <tr><td>Base</td><td class="pts">10</td><td>Every tender starts here</td></tr>
          <tr><td>Large ≥ €500k</td><td class="pts">+40</td><td>Major fleet or long maintenance</td></tr>
          <tr><td>Medium €100k–€500k</td><td class="pts">+25</td><td>Worth distributor time</td></tr>
          <tr><td>Service / R&amp;M signal</td><td class="pts">+24</td><td>Recurring parts lifetime value</td></tr>
          <tr><td>Framework</td><td class="pts">+20</td><td>Multi-year call-offs</td></tr>
          <tr><td>Fleet ≥ 20 / 5–19</td><td class="pts">+15 / +10</td><td>Material parts demand</td></tr>
        </tbody>
      </table>
    </div>
    <div class="explain-card">
      <h3>Decision thresholds</h3>
      <div class="decision-grid">
        <div class="decision-box"><strong>Attack</strong>Fit ≥ 60 · Win ≥ 55 · Econ ≥ 45 (all three).</div>
        <div class="decision-box"><strong>Selective</strong>Fit ≥ 52 and (Win ≥ 45 or Econ ≥ 50).</div>
        <div class="decision-box"><strong>Watch</strong>Fit ≥ 38 or Econ ≥ 32 — monitor / low effort.</div>
        <div class="decision-box"><strong>Ignore</strong>Below thresholds — wrong product or closed noise.</div>
      </div>
      <div class="callout"><strong>Special case:</strong> R&amp;M and awarded “result” notices stay Watch — approach the contract winner, do not treat as a fresh vehicle bid.</div>
    </div>
    <div class="explain-card">
      <h3>GPS product-fit + not-for-us (v4)</h3>
      <p>GPS € only on pickup/truck with published fleet and a GPS map row. Postal/courier panel vans are hard-dropped from opportunities (<code>_NOT_FOR_US_GATE.md</code>) — not kept as GPS-null clutter. Null ≠ €0. See <code>02_METHOD_GPS_PRODUCT_FIT.md</code>.</p>
    </div>`;
}

function renderDataSheets() {
  const sheets = DATA.workbook_sheets || [];
  const head = `<thead><tr><th>Workbook sheet</th><th>Data rows</th><th>Columns</th><th>Role</th></tr></thead>`;
  const roleHint = (name) => {
    const n = String(name).toLowerCase();
    if (n.includes('ted') || n.includes('eu')) return 'EU TED Europa notices';
    if (n.includes('boamp') || n.includes('france')) return 'France national (BOAMP)';
    if (n.includes('bzp') || n.includes('poland')) return 'Poland BZP';
    if (n.includes('contracts finder')) return 'UK Contracts Finder';
    if (n.includes('uk tenders') || n.includes('find a tender')) return 'UK Find a Tender';
    if (n.includes('belgium')) return 'Belgium public procurement';
    if (n.includes('israel')) return 'Israel tenders';
    if (n.includes('netherlands') || n.includes('tenderned')) return 'Netherlands TenderNed';
    if (n.includes('norway') || n.includes('doffin')) return 'Norway Doffin (key may be pending)';
    if (n.includes('finland') || n.includes('hilma')) return 'Finland HILMA (key may be pending)';
    if (n.includes('germany')) return 'Germany national / TED DE';
    return 'Source portal worksheet';
  };
  let body;
  if (!sheets.length) {
    body = `<tbody><tr><td colspan="4" class="muted">Workbook sheet list not embedded — rebuild from isuzu_tenders_v3.xlsx.</td></tr></tbody>`;
  } else {
    body = `<tbody>${sheets.map(s => `<tr style="cursor:default">
      <td>${esc(s.name)}</td>
      <td class="num">${fmtCount(s.data_rows)}</td>
      <td class="num">${fmtCount(s.cols)}</td>
      <td class="muted">${esc(roleHint(s.name))}</td>
    </tr>`).join('')}</tbody>`;
  }
  document.getElementById('dataSheetsTable').innerHTML = head + body;
  const total = sheets.reduce((a, s) => a + (Number(s.data_rows) || 0), 0);
  const s = liveSnapshotBits();
  document.getElementById('dataSheetsFooter').innerHTML =
    `Raw workbook: <code>isuzu_tenders_v3.xlsx</code> · ${fmtCount(sheets.length)} sheets · ${fmtCount(total)} data rows (pre-gate). ` +
    `Filtered opportunity export: <code>isze_dashboard_v4_opportunities.xlsx</code> (${fmtCount(s.nOpps)} rows, includes Aftersales). ` +
    `Double-click <code>OPEN_isze_dashboard_v4_opportunities_xlsx.cmd</code> to open the export.`;
}

(function init() {
  populateFilters();
  renderHeadline();
  renderSourceBanner();
  renderMarketTable();
  renderCoverage();
  renderOppTable();
  renderHowTenders();
  renderSourcesTable();
  renderMethod();
  renderDataSheets();
})();
</script>
</body>
</html>
'''



def write_opportunities_xlsx(payload: Dict[str, Any], out_path: Path) -> None:
    """Sheet for Jerome: opportunities + Aftersales column (rebuild path)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("[Aftersales xlsx] openpyxl missing — skipped")
        return
    headers = [
        "Decision", "Product", "Buyer", "Country", "Dealer / IDS", "Distributor",
        "Deadline", "Fleet count", "Contract value", "GPS EUR/yr",
        "Aftersales", "Aftersales tags", "R&M contract", "Title", "URL",
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Opportunities"
    header_fill = PatternFill("solid", fgColor="15233C")
    header_font = Font(color="F7F4EE", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True)
    for o in payload.get("opportunities") or []:
        dealer = o.get("ids_dealer_name") or ""
        if o.get("ids_dealer_city"):
            dealer = f"{dealer} · {o['ids_dealer_city']}" if dealer else o["ids_dealer_city"]
        if not dealer:
            dealer = o.get("ids_distributor") or o.get("distributor_code") or "unmapped"
        tags = o.get("aftersales_tags") or []
        ws.append([
            o.get("decision") or "",
            o.get("product_segment_label") or o.get("product_segment") or "",
            o.get("buyer") or "",
            o.get("country") or "",
            dealer,
            o.get("ids_distributor") or o.get("distributor_code") or "",
            o.get("deadline_date") or "",
            o.get("fleet_size") if o.get("fleet_size") else "",
            o.get("value_amount") if o.get("value_amount") is not None else "",
            o.get("gps_addressable_per_year_eur") if o.get("gps_addressable_per_year_eur") is not None else "",
            o.get("aftersales_summary") or "vehicles only",
            ", ".join(tags),
            "YES" if o.get("rm_contract") else "",
            o.get("title") or "",
            o.get("url") or "",
        ])
    ws.freeze_panes = "A2"
    wb.save(out_path)
    print(f"Wrote {out_path}")


def build_html(
    payload: Dict[str, Any],
    source_health: Optional[Dict[str, Any]],
    revenue_shares: Optional[Dict[str, Any]],
    gps_map: Optional[Dict[str, Any]] = None,
    workbook_sheets: Optional[list] = None,
) -> str:
    # Slim notices out of HTML to keep file manageable; opportunities carry GPS/PMS.
    embed = {
        "meta": payload.get("meta") or {},
        "opportunities": payload.get("opportunities") or [],
        "source_health": source_health or {},
        "revenue_shares": revenue_shares,
        "gps_map": gps_map or {},
        "addressable_vehicles": payload.get("addressable_vehicles") or {},
        "workbook_sheets": workbook_sheets or payload.get("workbook_sheets") or [],
    }
    data_json = json.dumps(embed, ensure_ascii=False, separators=(",", ":"))
    # Prevent </script> breakouts in titles
    data_json = data_json.replace("</", "<\\/")
    return HTML_TEMPLATE.replace("__DATA__", data_json)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build ISZE Tender OS v4 addressable-opportunity dashboard"
    )
    parser.add_argument(
        "--input",
        default=str(XLSX),
        help="Input XLSX (or legacy HTML). Used when --json-input is not set and v2 JSON missing.",
    )
    parser.add_argument(
        "--json-input",
        default="",
        help="Load existing tender_data.json / isze_dashboard_v2_data.json instead of XLSX",
    )
    parser.add_argument("--output", default=str(V4_HTML), help="Output HTML path")
    parser.add_argument("--json-output", default=str(V4_JSON), help="Output JSON path")
    args = parser.parse_args()

    os.chdir(HERE)
    json_in = Path(args.json_input) if args.json_input else None
    input_path = Path(args.input)
    payload = load_payload(input_path=input_path, json_input=json_in)
    source_health = load_source_health()
    revenue_shares = load_bp_revenue_shares()
    workbook_sheets = load_workbook_sheets(XLSX if XLSX.exists() else input_path)
    try:
        gps_map = load_gps_map_meta()
    except Exception as exc:
        print(f"[GPS map] meta skipped: {exc}")
        gps_map = {}

    out_json = Path(args.json_output)
    out_html = Path(args.output)
    # Full payload JSON for tooling; HTML embed is slimmed
    stats = sized_stats(payload)
    veh = addressable_vehicle_stats(payload)
    meta = payload.setdefault("meta", {})
    meta_stats = meta.setdefault("stats", {})
    meta_stats.update(
        {
            "vehicles_pickup_truck": veh["vehicles_pickup_truck"],
            "vehicles_all_stated": veh["vehicles_all_stated"],
            "rows_with_fleet": veh["rows_with_fleet"],
            "rows_pickup_truck_with_fleet": veh["rows_pickup_truck_with_fleet"],
            "addressable_vehicles_definition": veh["definition"],
            "vehicles_by_country": veh["by_country"],
        }
    )
    payload["addressable_vehicles"] = {
        "vehicles_pickup_truck": veh["vehicles_pickup_truck"],
        "vehicles_all_stated": veh["vehicles_all_stated"],
        "rows_with_fleet": veh["rows_with_fleet"],
        "rows_pickup_truck_with_fleet": veh["rows_pickup_truck_with_fleet"],
        "by_country": veh["by_country"],
        "definition": veh["definition"],
    }
    as_tagged = sum(1 for o in (payload.get("opportunities") or []) if o.get("aftersales_tags"))
    meta_stats["aftersales_tagged"] = as_tagged
    out_payload = dict(payload)
    out_payload["source_health"] = source_health or {}
    out_payload["revenue_shares"] = revenue_shares
    out_payload["gps_map"] = gps_map
    out_payload["workbook_sheets"] = workbook_sheets
    out_json.write_text(json.dumps(out_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    html = build_html(
        payload, source_health, revenue_shares, gps_map, workbook_sheets
    )
    out_html.write_text(html, encoding="utf-8")
    write_opportunities_xlsx(payload, V4_XLSX)

    print(f"Built {out_html}")
    print(f"Wrote {out_json}")
    print(f"  Opportunities: {stats['opportunities']}")
    print(f"  Sized (GPS €): {stats['sized']}")
    print(f"  Unsized (no GPS €): {stats['unsized']}")
    print(f"  With fleet count: {stats['with_fleet']}")
    print(f"  No fleet in notice: {stats['no_fleet']}")
    print(f"  Fleet but no GPS € (map or product-fit): {stats['fleet_but_no_gps_map']}")
    print(f"  Addressable GPS sum EUR: {stats['gps_addressable_sum_eur']}")
    print(f"  Stated fleet in captured notices (pickup/truck): {veh['vehicles_pickup_truck']} ({veh['rows_pickup_truck_with_fleet']} rows) — NOT market potential")
    print(f"  Stated fleet all segments: {veh['vehicles_all_stated']} ({veh['rows_with_fleet']} rows)")
    print(f"  Aftersales tagged: {as_tagged} of {stats['opportunities']}")
    segs = stats.get("product_segment_counts") or {}
    if segs:
        print(f"  Product segments: {', '.join(f'{k}={v}' for k, v in sorted(segs.items()))}")
    print(f"  Source health: {'present' if source_health else 'missing'}")
    print(f"  Workbook sheets: {len(workbook_sheets)} from {XLSX.name if XLSX.exists() else 'n/a'}")
    print(f"  Revenue shares: {'FY27 BP loaded' if revenue_shares else 'pipeline-only'}")
    ids_meta = (payload.get("meta") or {}).get("ids_dealer")
    if ids_meta:
        print(f"  IDS dealer: preserved ({ids_meta.get('status') or ids_meta.get('source') or 'linked'})")


if __name__ == "__main__":
    main()
