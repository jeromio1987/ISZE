"""Shared loaders for the DB sales-review toolchain.

Everything that used to be re-discovered by hand per distributor lives either in
db_profiles.yaml (constants) or in this module (the loaders that read the
canonical files). Nothing here is DB-specific.

See ..\\ppt_template_v2\\DB_SALES_REVIEW_BLUEPRINT.md for the method and the
tone rules; this module only implements the data side of it.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
import yaml

# --------------------------------------------------------------------- paths --
ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
DERIVED = ISZE / "01_data" / "00_derived"
PRICING = ISZE / "01_data" / "01_Pricing"
HERE = Path(__file__).resolve().parent
PROFILES = HERE / "db_profiles.yaml"
WARRANTY_GRAIN = DERIVED / "warranty_raw_grain.csv"

# profile key → exact db_name values in warranty_raw_grain.csv
# (ERP has no warranty doc type — claim grain is the PN-level exclusion source)
WARRANTY_GRAIN_DB: dict[str, list[str]] = {
    "IUK": ["ISUZU (UK)"],
    "PORTUGAL": ["IMOTORS LDA", "PORTUGAL"],
    "ISD": ["ISUZU SALES DEUTSCHLAND - ISD", "ISD"],
    # IBX payer bills BE/NL/PL — include all Benelux claim rows
    "IBX": [
        "ISUZU BENELUX - BELGIUM",
        "ISUZU BENELUX - NETHERLANDS",
        "Isuzu Benelux NV (Poland)",
    ],
    "PETROS": ["PETROS PETROPOULOS"],
    "IRELAND": ["Isuzu Ireland 2020 Unlimited Company", "IRELAND"],
    "UMI": ["UMI"],
    "NORWAY": ["RSA"],
    "SWEDEN": ["Isuzu Sverige AB"],
    "HUNGARY": ["WAE LTD.", "WAE"],
    "FINLAND": ["Isuzu Motor Finland OY"],
    "ROMANIA": ["ISUZU ROM AUTO SRL"],
    "DENMARK": ["ISUZU Danmark"],
    "ICELAND": ["BL ehf."],
    "CYPRUS": ["A. LOUTSIOS & SONS", "CYPRUS"],
    "MALTA": ["GASANZAMMIT MOTORS LIMITED", "MALTA"],
}

# Months of claim-assessment lag before the prior window (claims trail billing)
WARRANTY_CLAIM_LAG_MONTHS = 3

_WARRANTY_TEXT_RE = re.compile(
    r"\bSENSOR\W*PM\b|\bPM\W*SENSOR\b|\bWARRANTY\b|\bCAMPAIGN\b|\bRECALL\b",
    re.IGNORECASE,
)

UIO_FILE = ISZE / "06_uio" / "_CURRENT" / "ISUZU_UIO_Overview_2026.xlsx"
FY27_FILE = ISZE / "01_data" / "ISZE_GPS_Extended_Target_FY27_V3.1.xlsx"
IDS_FILE = (ISZE / "01_data" / "90_INBOX" / "IDS_2026" / "rankings"
            / "IDS_2026_TOTAL_ranking_all.xlsx")
BUDGET_FILE = ISZE / "01_data" / "11_planning" / "budget_per_DB_reviewed_24042026.xlsx"
CARPARK_FILE = DERIVED / "UIO_carpark_for_GPS_recalc_2026-07-30.csv"

OUT_ROOT = ISZE / "09_meetings"

# profile key -> distributor_DB name in the carpark CSV (30 Jul 2026 pack)
CARPARK_DB = {
    "IUK": "IUK",
    "PORTUGAL": "Imotors",
    "ISD": "ISD-Lang",
    "IBX": "IBX",
    "PETROS": "Petros Petropoulos",
    "IRELAND": "Isuzu Ireland",
    "UMI": "UMI",
    "NORWAY": "RSA",
    "SWEDEN": "Isuzu Sverige",
    "HUNGARY": "AutoWallis",
    "FINLAND": "Nordic Auto.",
    "ROMANIA": "Isuzu Rom Auto",
    "DENMARK": "Isuzu Denmark",
    "ICELAND": "BL",
    "CYPRUS": "Loutsios & Sons",
    "MALTA": "GasanZammit",
}
_CARPARK_MODEL_TO_SEG = {
    "LD": "ld",
    "MD": "md",
    "F-Series": "f_series",
    "LCV": "lcv",
}


# ------------------------------------------------------------------ profiles --
def load_profile(key: str) -> dict:
    """Return one distributor's profile with `defaults` merged in under 'defaults'."""
    cfg = yaml.safe_load(PROFILES.read_text(encoding="utf-8"))
    key = key.upper()
    if key not in cfg["distributors"]:
        raise SystemExit(
            f"Unknown DB '{key}'. Known: {', '.join(sorted(cfg['distributors']))}"
        )
    p = dict(cfg["distributors"][key])
    p["key"] = key
    p["defaults"] = cfg["defaults"]
    return p


def list_profiles() -> list[str]:
    cfg = yaml.safe_load(PROFILES.read_text(encoding="utf-8"))
    return sorted(
        key for key, profile in cfg["distributors"].items()
        if profile.get("build_enabled", True)
    )


# ------------------------------------------------------------------- helpers --
def norm_pn(x) -> str | None:
    """Normalise a part number cell. Returns None for blanks/NaN."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return None
    s = str(x).strip()
    if not s or s.lower() in {"nan", "none"}:
        return None
    try:
        if "." in s:
            f = float(s)
            if f == int(f):
                s = str(int(f))
    except ValueError:
        pass
    return s


def _dates(s):
    return pd.to_datetime(s, errors="coerce")


# --------------------------------------------------------------------- sales --
def load_sales(profile: dict) -> pd.DataFrame:
    """Line-level billing for one DB across the prior and current window.

    Adds three columns the rest of the toolchain depends on:
      period     window_label_prior / window_label_current
      is_credit  True for BillT G2 or SaTy CR
      sap_desc   the free-text description (the raw exports hold it in the
                 column *labelled* 'Material Number' — the two are swapped)

    Credit memos carry a POSITIVE Net value in the raw exports but are
    economically negative; they are signed here so a rebate credit cannot read
    as a gain. Verified against net_value_gps_signed in the tagged extract.

    CY2026 current window spanning past 31 Mar pulls Apr–Jun from
    q1fy2027_billing_tagged.csv (same payer / Net value / credit convention).
    """
    d = profile["defaults"]
    payer = profile["payer"]
    L25, L26 = d["window_label_prior"], d["window_label_current"]
    w25, w26 = d["window_prior"], d["window_current"]
    frames = []

    usecols_cur = ["Billing Date", "Payer.1", "Material", "Material Number",
                   "Billed Quantity", "Net value", "SaTy", "BillT"]

    def _slice_cur(df: pd.DataFrame, start: str, end: str) -> pd.DataFrame:
        out = df[df["Payer.1"] == payer].copy()
        out["Billing Date"] = _dates(out["Billing Date"])
        out = out[(out["Billing Date"] >= start) & (out["Billing Date"] < end)]
        out["period"] = L26
        return out.rename(columns={"Billed Quantity": "qty", "Net value": "net_eur"})

    # CY current: Jan–Mar from cy2026_q1, Apr+ from q1fy2027 tagged extract
    cur_q1 = pd.read_csv(DERIVED / "cy2026_q1.csv", low_memory=False, usecols=usecols_cur)
    frames.append(_slice_cur(cur_q1, w26[0], min(w26[1], "2026-04-01")))

    if w26[1] > "2026-04-01":
        cur_q2 = pd.read_csv(DERIVED / "q1fy2027_billing_tagged.csv", low_memory=False,
                             usecols=usecols_cur)
        frames.append(_slice_cur(cur_q2, max(w26[0], "2026-04-01"), w26[1]))

    pri = pd.read_csv(DERIVED / "cy2025_raw.csv", low_memory=False,
                      usecols=["Billing Date", "DB", "Material", "Material Number",
                               "Billed Quantity", "Net value", "SaTy"])
    pri = pri[pri["DB"] == payer].copy()
    pri["Billing Date"] = _dates(pri["Billing Date"])
    pri = pri[(pri["Billing Date"] >= w25[0]) & (pri["Billing Date"] < w25[1])]
    pri["period"] = L25
    pri["BillT"] = None
    frames.append(pri.rename(columns={"Billed Quantity": "qty", "Net value": "net_eur"}))

    all_ = pd.concat(frames, ignore_index=True)
    if all_.empty:
        raise SystemExit(f"No billing rows found for payer '{payer}'. Check the profile.")

    all_["sap_desc"] = all_["Material Number"].astype(str).str.strip()
    all_["Material"] = all_["Material"].map(norm_pn)
    all_["qty"] = pd.to_numeric(all_["qty"], errors="coerce").fillna(0.0)
    all_["net_eur"] = pd.to_numeric(all_["net_eur"], errors="coerce").fillna(0.0)
    all_["SaTy"] = all_["SaTy"].astype(str).str.strip().str.upper()
    all_["is_credit"] = (all_["BillT"].astype(str).str.upper() == "G2") | (all_["SaTy"] == "CR")
    all_.loc[all_["is_credit"], "net_eur"] = -all_.loc[all_["is_credit"], "net_eur"].abs()
    all_.loc[all_["is_credit"], "qty"] = -all_.loc[all_["is_credit"], "qty"].abs()
    return all_


def fy_billing_derived(profile: dict, start="2025-04-01", end="2026-04-01") -> float:
    """Sum a fiscal year straight from the CY files — used by preflight to prove
    the source files are complete for this DB before anything is built on them."""
    payer = profile["payer"]
    tot = 0.0
    for f, col in [("cy2026_q1.csv", "Payer.1"), ("cy2025_raw.csv", "DB")]:
        df = pd.read_csv(DERIVED / f, low_memory=False, usecols=[col, "Billing Date", "Net value"])
        df = df[df[col] == payer]
        df["Billing Date"] = _dates(df["Billing Date"])
        df = df[(df["Billing Date"] >= start) & (df["Billing Date"] < end)]
        tot += float(pd.to_numeric(df["Net value"], errors="coerce").fillna(0).sum())
    return tot


# ----------------------------------------------------------------- warranty --
def is_warranty_campaign_text(pn_value, description) -> bool:
    """PN/description heuristic for campaign / recall / PM-sensor style lines."""
    text = f"{pn_value or ''} {description or ''}"
    return bool(_WARRANTY_TEXT_RE.search(text))


def _yyyymm(ts) -> int:
    ts = pd.Timestamp(ts)
    return int(ts.year) * 100 + int(ts.month)


def _shift_yyyymm(yyyymm: int, months: int) -> int:
    y, m = divmod(int(yyyymm), 100)
    idx = y * 12 + (m - 1) + months
    return (idx // 12) * 100 + (idx % 12) + 1


def load_warranty_claim_pns(profile: dict,
                            window_prior: tuple[str, str] | None = None,
                            window_current: tuple[str, str] | None = None,
                            lag_months: int = WARRANTY_CLAIM_LAG_MONTHS,
                            ) -> dict:
    """PN set from warranty_raw_grain for this DB over the analysis windows.

    ERP billing has no warranty document type, so commercial H1 movers exclude
    PNs that appear on claim lines for this distributor. Claim assessment dates
    lag billing — the prior window is opened ``lag_months`` earlier.

    Returns dict with keys: pns (set[str]), n_claim_lines, claim_eur, db_names,
    assess_month_from, assess_month_to, source, available (bool).
    """
    d = profile["defaults"]
    w25 = window_prior or tuple(d["window_prior"])
    w26 = window_current or tuple(d["window_current"])
    names = list(profile.get("warranty_grain_db") or WARRANTY_GRAIN_DB.get(profile["key"]) or [])
    empty = {
        "pns": set(), "n_claim_lines": 0, "claim_eur": 0.0, "db_names": names,
        "assess_month_from": None, "assess_month_to": None,
        "source": str(WARRANTY_GRAIN), "available": False,
        "method": "claim_pn_match+text_heuristic",
    }
    if not names or not WARRANTY_GRAIN.exists():
        return empty

    grain = pd.read_csv(WARRANTY_GRAIN, low_memory=False,
                        usecols=["assess_month", "db_name", "part_no", "claim_eur"])
    hit = grain[grain["db_name"].isin(names)].copy()
    if hit.empty:
        return empty

    start = _shift_yyyymm(_yyyymm(w25[0]), -lag_months)
    end = _yyyymm(w26[1])  # exclusive end date → last month before end
    # window end is exclusive (e.g. 2026-07-01) → include through prior month
    end_inclusive = _shift_yyyymm(end, -1)
    am = pd.to_numeric(hit["assess_month"], errors="coerce")
    hit = hit[(am >= start) & (am <= end_inclusive)]
    pns = {p for p in hit["part_no"].map(norm_pn) if p}
    return {
        "pns": pns,
        "n_claim_lines": int(len(hit)),
        "claim_eur": float(pd.to_numeric(hit["claim_eur"], errors="coerce").fillna(0).sum()),
        "db_names": names,
        "assess_month_from": int(start),
        "assess_month_to": int(end_inclusive),
        "source": str(WARRANTY_GRAIN),
        "available": True,
        "method": "claim_pn_match+text_heuristic",
    }


def is_warranty_pn(pn_value, description, claim_pns: set[str]) -> bool:
    """True when PN is on a claim for this DB or matches the campaign heuristic."""
    pn = norm_pn(pn_value)
    if pn and pn in claim_pns:
        return True
    return is_warranty_campaign_text(pn_value, description)


# ------------------------------------------------------------------ ITC map --
def load_itc_map():
    """P/N -> Newest P/N with chains resolved transitively, plus descriptions."""
    cache = DERIVED / "price_file_2026_pn_newest.csv"
    if cache.exists():
        df = pd.read_csv(cache, dtype=str)
    else:
        df = pd.read_excel(PRICING / "Price file 2026_full.xlsx", sheet_name="Sheet1",
                           usecols=["P/N", "Newest P/N", "Description"], dtype=str)
        df.to_csv(cache, index=False)

    df = df.rename(columns={"P/N": "pn_raw", "Newest P/N": "newest_raw", "Description": "desc"})
    df["pn"] = df["pn_raw"].map(norm_pn)
    df["newest"] = df["newest_raw"].map(norm_pn)
    df = df.dropna(subset=["pn"]).drop_duplicates(subset=["pn"], keep="last")
    df["newest"] = df["newest"].fillna(df["pn"])
    df.loc[df["newest"] == "", "newest"] = df.loc[df["newest"] == "", "pn"]

    direct = dict(zip(df["pn"], df["newest"]))

    def resolve(pn):
        seen, cur = set(), pn
        for _ in range(20):
            nxt = direct.get(cur, cur)
            if nxt == cur or nxt in seen:
                return nxt or cur
            seen.add(cur)
            cur = nxt
        return cur

    resolved = {pn: resolve(pn) for pn in direct}
    desc_by_pn = dict(zip(df["pn"], df["desc"].fillna("")))
    desc_by_newest: dict[str, str] = {}
    for pn, newest in resolved.items():
        if not desc_by_newest.get(newest):
            desc_by_newest[newest] = desc_by_pn.get(newest) or desc_by_pn.get(pn) or ""
    return resolved, desc_by_newest


def load_catalog() -> pd.DataFrame:
    cat = pd.read_csv(DERIVED / "parts_catalog_dimension.csv", dtype=str)
    cat["part_number"] = cat["part_number"].map(norm_pn)
    return cat[["part_number", "pr00_group", "pms_category"]].drop_duplicates("part_number")


def load_bvp_pns() -> set[str]:
    bvp = pd.read_csv(DERIVED / "bvp_pn_dimension.csv", dtype=str)
    return set(bvp["part_number"].dropna().map(norm_pn).dropna())


# ------------------------------------------------------------- reference data --
def load_gps_master(profile: dict) -> dict:
    g = pd.read_csv(DERIVED / "gps_master.csv")
    hit = g[g["Unnamed: 0"] == profile["gps_key"]]
    if hit.empty:
        raise SystemExit(f"gps_key '{profile['gps_key']}' not in gps_master.csv")
    row = hit.iloc[0]
    keys = ["bill_FY2024", "bill_FY2025", "bill_FY2026", "net_FY2024", "net_FY2025",
            "net_FY2026", "uio_2025", "rate", "max_spend",
            "gps_pct_FY2024", "gps_pct_FY2025", "gps_pct_FY2026"]
    return {k: (None if pd.isna(row[k]) else float(row[k])) for k in keys}


def load_carpark(profile: dict) -> dict | None:
    """Segmented UIO + target wallet from the 30 Jul 2026 carpark pack.

    Returns None when this profile has no row in the pack. Numerators are NOT
    in the pack — pair with load_gps_master / load_segmented_gps.
    """
    cdb = profile.get("carpark_db") or CARPARK_DB.get(profile["key"])
    if not cdb or not CARPARK_FILE.exists():
        return None
    df = pd.read_csv(CARPARK_FILE)
    df = df[(df["distributor_DB"] == cdb) & (~df["is_MIDI_exclude"].astype(bool))]
    if df.empty:
        return None

    exclude = set(profile.get("exclude_uio_segments") or [])
    by_seg: dict[str, dict] = {}
    for _, r in df.iterrows():
        seg = _CARPARK_MODEL_TO_SEG.get(str(r["model_type"]))
        if not seg or seg in exclude:
            continue
        slot = by_seg.setdefault(seg, {
            "uio_2025": 0, "uio_2026": 0,
            "wallet_2025": 0.0, "wallet_2026": 0.0,
            "rate": float(r["target_eur_per_vehicle_yr"]),
            "gps_target_group": str(r["gps_target_group"]),
        })
        slot["uio_2025"] += int(r["uio_CY2025_unchanged"] or 0)
        slot["uio_2026"] += int(r["uio_CY2026_corrected"] or 0)
        slot["wallet_2025"] += float(r["wallet_CY2025_eur"] or 0)
        slot["wallet_2026"] += float(r["wallet_CY2026_corrected_eur"] or 0)

    if not by_seg:
        return None

    statuses = sorted({str(s) for s in df["status"].dropna().unique()})
    return {
        "carpark_db": cdb,
        "by_segment": by_seg,
        "uio_2025": int(sum(v["uio_2025"] for v in by_seg.values())),
        "uio_2026": int(sum(v["uio_2026"] for v in by_seg.values())),
        "wallet_2025": float(sum(v["wallet_2025"] for v in by_seg.values())),
        "wallet_2026": float(sum(v["wallet_2026"] for v in by_seg.values())),
        "status": statuses,
        "source": CARPARK_FILE.name,
    }


def load_segmented_gps(profile: dict) -> dict | None:
    """GPS = gps_master net sales ÷ carpark segmented wallet (CY2025 frozen denom).

    Headline for decks is FY2026. If capture exceeds 100%, show_pct is False —
    Mode A slides must use the amount variant (Ireland-class cases).
    """
    park = load_carpark(profile)
    if not park or not park["wallet_2025"]:
        return None
    gm = load_gps_master(profile)
    net25, net26 = gm.get("net_FY2025"), gm.get("net_FY2026")
    w25, w26 = park["wallet_2025"], park["wallet_2026"]

    def _pct(net, wallet):
        if net is None or not wallet:
            return None
        return net / wallet

    gps25 = _pct(net25, w25)
    gps26 = _pct(net26, w25)          # frozen CY2025 denom (gps_master convention)
    gps26_w26 = _pct(net26, w26)      # matching-year view for FY27 planning

    show_pct = gps26 is not None and 0 < gps26 <= 1.0
    return {
        "carpark_db": park["carpark_db"],
        "uio_2025": park["uio_2025"],
        "uio_2026": park["uio_2026"],
        "wallet_2025": w25,
        "wallet_2026": w26,
        "net_FY2025": net25,
        "net_FY2026": net26,
        "gps_FY2025": gps25,
        "gps_FY2026": gps26,
        "gps_FY2026_on_cy2026_wallet": gps26_w26,
        "show_pct": show_pct,
        "status": park["status"],
        "by_segment": park["by_segment"],
        "source": park["source"],
        # flat-rate reference from gps_master (never put the recon on a Mode A slide)
        "gps_master_flat_FY2026": gm.get("gps_pct_FY2026"),
        "gps_master_max_spend": gm.get("max_spend"),
    }


def latest_complete_pms_year(years) -> int:
    """Latest full calendar year in a PMS series.

    Mid-year rows (e.g. CY2026 YTD in Aug 2026) keep a full-year expected_qty
    denominator. Pairing that with a partial numerator crushes capture %
    (Hungary brake rotor: 5.1% on 2026 YTD vs 36.1% on CY2025 full year — same
    file, same model). Decks and peer bars must use a complete year only.
    """
    from datetime import date
    ys = sorted({int(y) for y in years if pd.notna(y)})
    if not ys:
        raise SystemExit("No PMS years available")
    cy = date.today().year
    complete = [y for y in ys if y < cy]
    return max(complete) if complete else max(ys)


def load_pms(profile: dict, year: int | None = None):
    p = pd.read_csv(DERIVED / "pms_capture_by_db_year.csv")
    p = p[(p["distributor_code"] == profile["pms_code"]) & (p["line_type"] == "total")]
    if p.empty:
        raise SystemExit(f"pms_code '{profile['pms_code']}' not in pms_capture_by_db_year.csv")
    year = year or latest_complete_pms_year(p["year"])
    return year, p[p["year"] == year].reset_index(drop=True)


def pms_vs_network(profile: dict, categories=None, year: int | None = None) -> dict:
    """This DB's capture against the rest of the network, on the same expected-
    quantity model — which normalises away the interval assumptions and is the
    only defensible way to show these rates to a distributor."""
    categories = categories or ["oil_filter", "fuel_filter", "air_filter",
                                "brake_pad", "brake_rotor"]
    p = pd.read_csv(DERIVED / "pms_capture_by_db_year.csv")
    p = p[p["line_type"] == "total"]
    year = year or latest_complete_pms_year(p["year"])
    p = p[p["year"] == year]
    code = profile["pms_code"]
    out = {}
    for cat in categories:
        sub = p[p["pms_category"] == cat]
        me = sub[sub["distributor_code"] == code]
        peers = sub[sub["distributor_code"] != code]
        if me.empty or peers["expected_qty"].sum() == 0:
            continue
        out[cat] = {
            "db": float(me["capture_rate"].iloc[0]) * 100,
            "network": float(peers["actual_qty"].sum() / peers["expected_qty"].sum()) * 100,
            "best": float(peers["capture_rate"].max()) * 100,
        }
    return out


def load_uio(profile: dict) -> dict:
    """{segment: {year: units}} straight from 06_uio/_CURRENT. Years are the
    literal header labels with the E/F suffix stripped (2025E -> 2025).

    For deck purposes only, profiles can request a corrected CY2026 estimate.
    It backs registrations out of the workbook's 3% roll-forward and reapplies
    the historical survival rates (10.2% CV / 12.3% LCV). The source workbook
    is never changed. IUK's locked LCV estimate is handled as a fixed override.

    When the 30 Jul 2026 carpark pack has a row for this DB, CY2025 and CY2026
    are taken from that pack (authoritative for GPS + fleet slides).
    """
    import openpyxl
    wb = openpyxl.load_workbook(UIO_FILE, data_only=True, read_only=True)
    ws = wb["1.UIO Overview"]
    rows = list(ws.iter_rows(min_row=1, max_row=130, values_only=True))
    wb.close()
    hdr = rows[5]
    ycols = {}
    for c in range(6, len(hdr)):
        lab = str(hdr[c]) if hdr[c] is not None else ""
        m = re.match(r"^(\d{4})[EF]?$", lab)
        if m:
            ycols[int(m.group(1))] = c

    out: dict[str, dict[int, int]] = {}
    for seg, rowno in (profile["uio_rows"] or {}).items():
        if not rowno or seg in set(profile.get("exclude_uio_segments") or []):
            continue
        r = rows[rowno - 1]
        out[seg] = {y: int(r[c]) for y, c in sorted(ycols.items())
                    if c < len(r) and isinstance(r[c], (int, float)) and r[c]}
    if not out:
        raise SystemExit(f"No UIO rows configured for {profile['key']}")

    if profile.get("uio_cy2026_method") == "corrected_survival":
        for seg, vals in out.items():
            prior, workbook_2026 = vals.get(2025), vals.get(2026)
            if not prior or workbook_2026 is None:
                continue
            scrap = 0.123 if seg == "lcv" else 0.102
            implied_regs = workbook_2026 - prior * 0.97
            vals[2026] = int(round(prior * (1 - scrap) + implied_regs))

    for seg, value in (profile.get("uio_2026_fixed") or {}).items():
        if seg in out:
            out[seg][2026] = int(value)

    # Carpark pack wins for 2025/2026 when present (GPS wallet + fleet slides).
    park = load_carpark(profile)
    if park:
        for seg, slot in park["by_segment"].items():
            if seg not in out:
                out[seg] = {}
            out[seg][2025] = int(slot["uio_2025"])
            out[seg][2026] = int(slot["uio_2026"])

    total = {}
    for seg_vals in out.values():
        for y, v in seg_vals.items():
            total[y] = total.get(y, 0) + v
    out["total"] = dict(sorted(total.items()))
    return out


def load_s2d_target(profile: dict) -> dict:
    """Load Dashboard P:U and return the locked Sales-to-Dealer target chain.

    FY2027 target = FY2026 S2D × (1 + UIO growth) × 1.05 × 1.018.
    When the KPI S2D base is unavailable, FY2026 ISZE billing is grossed up at
    a 65% ISZE-to-dealer factor and explicitly marked as imputed.
    """
    import openpyxl
    wb = openpyxl.load_workbook(FY27_FILE, data_only=True, read_only=True)
    ws = wb["Dashboard"]
    dashboard_key = str(profile.get("dashboard_key") or profile["fy27_tab"]).strip().upper()
    hit = None
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row and str(row[0] or "").strip().upper() == dashboard_key:
            hit = row
            break
    wb.close()
    if hit is None:
        raise SystemExit(f"dashboard_key '{dashboard_key}' not in {FY27_FILE.name} Dashboard")

    def num(v):
        return float(v) if isinstance(v, (int, float)) and not pd.isna(v) else None

    dashboard_billing = num(hit[13]) if len(hit) > 13 else None  # N
    s2d_fy25 = num(hit[15]) if len(hit) > 15 else None           # P
    s2d_fy26 = num(hit[16]) if len(hit) > 16 else None           # Q
    dashboard_growth = num(hit[17]) if len(hit) > 17 else None   # R
    price = num(hit[18]) if len(hit) > 18 else None              # S
    additional = num(hit[19]) if len(hit) > 19 else None         # T
    dashboard_target = num(hit[20]) if len(hit) > 20 else None   # U

    imputed = s2d_fy26 is None
    if imputed:
        billing = load_gps_master(profile).get("bill_FY2026") or dashboard_billing
        if billing is None:
            raise SystemExit(f"No FY2026 billing available to impute S2D for {profile['key']}")
        s2d_fy26 = billing / 0.65

    growth = dashboard_growth
    growth_overridden = bool(
        profile.get("uio_cy2026_method") or profile.get("uio_2026_fixed")
    )
    if growth_overridden:
        uio = load_uio(profile)["total"]
        if not uio.get(2025):
            raise SystemExit(f"No CY2025 UIO available to calculate S2D growth for {profile['key']}")
        growth = uio.get(2026, 0) / uio[2025] - 1
    if growth is None:
        raise SystemExit(f"No UIO growth available for S2D target for {profile['key']}")

    price = 0.05 if price is None else price
    additional = 0.018 if additional is None else additional
    target = s2d_fy26 * (1 + growth) * (1 + price) * (1 + additional)
    return {
        "dashboard_key": dashboard_key,
        "fy25": s2d_fy25,
        "fy26": s2d_fy26,
        "uio_growth": growth,
        "dashboard_uio_growth": dashboard_growth,
        "price": price,
        "additional": additional,
        "fy27": round(target),
        "dashboard_fy27": dashboard_target,
        "s2d_imputed": imputed,
        "growth_overridden": growth_overridden,
        "billing_imputation_base": billing if imputed else None,
    }


_FY27_LABELS = {
    "max_spend": "Max spend / potential",
    "actual": "Actual (ISZE billing)",
    "gps": "Share (GPS",
    "stretch_pct": "Stretch idea",
    "stretch_eur": "Stretch amount (= max spend",
    "ext_target": "TTL amount (= actual + stretch)",
    "gps_new": "Share (new)",
    "s2d": "Parts Sales DB-to-DLR",
}
_FY27_SEGMENTS = {
    "ld": "LD (N75", "md": "MD", "f_series": "F-Series", "lcv": "LCV",
}


def load_fy27(profile: dict) -> dict:
    """Read the DB's tab in ISZE_GPS_Extended_Target_FY27_V3.1.xlsx by LABEL, not
    by fixed row — the LCV-only tabs (IUK) and the mixed tabs have different
    layouts. Columns C/D/E are FY25 / FY26 / FY27 throughout."""
    import openpyxl
    wb = openpyxl.load_workbook(FY27_FILE, data_only=True)
    tab = profile["fy27_tab"]
    if tab not in wb.sheetnames:
        raise SystemExit(f"fy27_tab '{tab}' not in {FY27_FILE.name}")
    ws = wb[tab]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]
    wb.close()

    def num(v):
        return float(v) if isinstance(v, (int, float)) else None

    out: dict = {"tab": tab, "segments": {}}
    for r in rows:
        lab = str(r[0]).strip() if r and r[0] else ""
        if not lab:
            continue
        for key, seg_lab in _FY27_SEGMENTS.items():
            if lab.startswith(seg_lab) and key not in out["segments"]:
                out["segments"][key] = {
                    "fy25": num(r[2]) if len(r) > 2 else None,
                    "fy26": num(r[3]) if len(r) > 3 else None,
                    "fy27": num(r[4]) if len(r) > 4 else None,
                }
        for key, needle in _FY27_LABELS.items():
            if lab.startswith(needle) and key not in out:
                out[key] = {
                    "fy25": num(r[2]) if len(r) > 2 else None,
                    "fy26": num(r[3]) if len(r) > 3 else None,
                    "fy27": num(r[4]) if len(r) > 4 else None,
                }
    missing = [k for k in ("max_spend", "actual", "ext_target") if k not in out]
    if missing:
        raise SystemExit(f"FY27 tab '{tab}' is missing rows for: {missing}")
    out["s2d_target"] = load_s2d_target(profile)
    return out


def load_ids(profile: dict):
    """(this DB's dealers, network average, DB average) or None when the DB has
    no IDS 2026 assessment."""
    if not profile.get("ids_code"):
        return None
    df = pd.read_excel(IDS_FILE, sheet_name=0)
    mine = df[df["Country"] == profile["ids_code"]].copy()
    if mine.empty:
        return None
    mine = mine.sort_values("Rank (country)")
    return mine, float(df["Points achieved %"].mean()), float(mine["Points achieved %"].mean())
