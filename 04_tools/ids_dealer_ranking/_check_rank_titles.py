#!/usr/bin/env python3
from pathlib import Path
from openpyxl import load_workbook

r = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\IDS_2026\rankings")
for f in [
    "IUK_dealer_ranking.xlsx",
    "ITUK_dealer_ranking.xlsx",
    "Israel_dealer_ranking.xlsx",
    "RSA_dealer_ranking.xlsx",
    "Petros_dealer_ranking.xlsx",
    "BE_dealer_ranking.xlsx",
    "NL_dealer_ranking.xlsx",
    "LUX_dealer_ranking.xlsx",
    "Switzerland_dealer_ranking.xlsx",
    "IDS_2026_all_country_rankings.xlsx",
]:
    p = r / f
    wb = load_workbook(p, read_only=True)
    print(f, "->", wb.sheetnames[:8])
    wb.close()
