"""Approximate dealer geo for IDS 2026 map pins.

Priority: dealer override → MIDI NETWORK lat/lon (FR/IT/ES/CH) →
city token in address / MIDI town → city token in name →
country centre with small deterministic jitter (last resort). Not surveyed
street addresses.
"""

from __future__ import annotations

import hashlib
import math
import re
from collections import defaultdict
from functools import lru_cache
from pathlib import Path
from typing import Any

# --- City centres (name / address token → lat, lon, label) ---
_CITY_COORDS: dict[str, tuple[float, float, str]] = {
    # Portugal
    "BEJA": (38.0151, -7.8632, "Beja"),
    "EVORA": (38.5667, -7.9000, "Évora"),
    "PORTALEGRE": (39.2938, -7.4312, "Portalegre"),
    "CASCAIS": (38.6979, -9.4215, "Cascais"),
    "LISBOA": (38.7223, -9.1393, "Lisboa"),
    "LEIRIA": (39.7436, -8.8071, "Leiria"),
    "POVOA": (41.3833, -8.7667, "Póvoa de Varzim"),
    "VIANA": (41.6918, -8.8344, "Viana do Castelo"),
    "PACOS": (41.2760, -8.3760, "Paços de Ferreira"),
    "PAÇOS": (41.2760, -8.3760, "Paços de Ferreira"),
    "PENAFIEL": (41.2080, -8.2820, "Penafiel"),
    "CALDAS": (39.4070, -9.1350, "Caldas da Rainha"),
    # Belgium / NL place tokens
    "LOCHRISTI": (51.0964, 3.8314, "Lochristi"),
    "SINT-TRUIDEN": (50.8167, 5.1833, "Sint-Truiden"),
    "HUY": (50.5187, 5.2328, "Huy"),
    "NAMUR": (50.4674, 4.8720, "Namur"),
    "PLOMBIERES": (50.7360, 5.9660, "Plombières"),
    "GROOT-BIJGAARDEN": (50.8730, 4.2660, "Groot-Bijgaarden"),
    "AMSTELVEEN": (52.3081, 4.8694, "Amstelveen"),
    "VENLO": (51.3704, 6.1724, "Venlo"),
    "BORN": (51.0320, 5.8100, "Born"),
    "DELFGAUW": (51.9940, 4.3950, "Delfgauw"),
    "BARNEVELD": (52.1400, 5.5840, "Barneveld"),
    "WADDINXVEEN": (52.0450, 4.6510, "Waddinxveen"),
    "ESCH": (49.4958, 5.9806, "Esch-sur-Alzette"),
    # Denmark
    "AALBORG": (57.0488, 9.9217, "Aalborg"),
    "FREDERIKSHAVN": (57.4407, 10.5366, "Frederikshavn"),
    "HJORRING": (57.4642, 9.9823, "Hjørring"),
    "HJØRRING": (57.4642, 9.9823, "Hjørring"),
    "AABENRAA": (55.0443, 9.4178, "Aabenraa"),
    "HOLSTEBRO": (56.3601, 8.6161, "Holstebro"),
    "RINGKOBING": (56.0901, 8.2440, "Ringkøbing"),
    "RINGKØBING": (56.0901, 8.2440, "Ringkøbing"),
    "KOLDING": (55.4904, 9.4722, "Kolding"),
    # France (name tokens)
    "CAVAILLON": (43.8375, 5.0381, "Cavaillon"),
    "LORRAINE": (48.6921, 6.1844, "Nancy (Lorraine)"),
    "BEZIERS": (43.3442, 3.2158, "Béziers"),
    "BÉZIERS": (43.3442, 3.2158, "Béziers"),
    "BONNEVILLE": (46.0775, 6.4090, "Bonneville"),
    "CHAUMONT": (48.1112, 5.1390, "Chaumont"),
    "ESSONNE": (48.6300, 2.4400, "Évry (Essonne)"),
    "CAMARGUE": (43.5333, 4.5667, "Camargue"),
    "ALSACE": (48.5734, 7.7521, "Strasbourg (Alsace)"),
    "LUISANT": (48.4300, 1.4750, "Luisant"),
    "MANTES": (48.9908, 1.7167, "Mantes-la-Jolie"),
    "MONTELIMAR": (44.5569, 4.7495, "Montélimar"),
    "MONTÉLIMAR": (44.5569, 4.7495, "Montélimar"),
    "MONTPELLIER": (43.6108, 3.8767, "Montpellier"),
    "BELFORT": (47.6397, 6.8638, "Belfort"),
    "STRASBOURG": (48.5734, 7.7521, "Strasbourg"),
    "ROUEN": (49.4431, 1.0993, "Rouen"),
    "TOURAINE": (47.3941, 0.6848, "Tours (Touraine)"),
    "PYRENEES": (43.2951, -0.3708, "Pau (Pyrénées)"),
    "PYRÉNÉES": (43.2951, -0.3708, "Pau (Pyrénées)"),
    "GARLABAN": (43.2965, 5.5381, "Aubagne (Garlaban)"),
    # Israel place tokens
    "JERUSALEM": (31.7683, 35.2137, "Jerusalem"),
    "ASHDOD": (31.8044, 34.6553, "Ashdod"),
    "HERZELIA": (32.1624, 34.8447, "Herzliya"),
    "HERZLIYA": (32.1624, 34.8447, "Herzliya"),
    "REHOVOT": (31.8928, 34.8113, "Rehovot"),
    "TVERIA": (32.7922, 35.5312, "Tiberias"),
    "RISHON": (31.9730, 34.7925, "Rishon LeZion"),
    "ASHKELON": (31.6688, 34.5743, "Ashkelon"),
    "NAHARIA": (33.0059, 35.0941, "Nahariya"),
    "BEER SHEVA": (31.2530, 34.7915, "Be'er Sheva"),
    "HAIFA": (32.7940, 34.9896, "Haifa"),
    "KIRYAT ATA": (32.8000, 35.1060, "Kiryat Ata"),
    "MODI'IN": (31.9077, 35.0076, "Modi'in"),
    "TLV": (32.0853, 34.7818, "Tel Aviv"),
    "RAANANA": (32.1848, 34.8708, "Ra'anana"),
    "NETIVOT": (31.4220, 34.5886, "Netivot"),
    "PETACH TIKVA": (32.0840, 34.8878, "Petah Tikva"),
    "HADERA": (32.4340, 34.9196, "Hadera"),
    "BEIT SHEMESH": (31.7470, 34.9881, "Beit Shemesh"),
    "EMEK HEFER": (32.3800, 34.9200, "Emek Hefer"),
    "NOF HAGALIL": (32.7100, 35.3200, "Nof HaGalil"),
    # Italy
    "CUNEO": (44.3841, 7.5420, "Cuneo"),
    "MIRANDOLA": (44.8867, 11.0660, "Mirandola"),
    "PADOVA": (45.4064, 11.8768, "Padova"),
    "BIELLES": (45.5620, 8.0580, "Biella"),
    # Norway
    "EVJE": (58.5850, 7.8060, "Evje"),
    "ALTA": (69.9689, 23.2717, "Alta"),
    "FINNSNES": (69.2296, 17.9812, "Finnsnes"),
    "LARVIK": (59.0533, 10.0353, "Larvik"),
    "GOL": (60.7010, 8.9420, "Gol"),
    "HARESTUA": (60.2070, 10.7210, "Harestua"),
    "HOKKSUND": (59.7680, 9.9100, "Hokksund"),
    "BRYNE": (58.7354, 5.6477, "Bryne"),
    "HONEFOSS": (60.1680, 10.2560, "Hønefoss"),
    "HØNEFOSS": (60.1680, 10.2560, "Hønefoss"),
    "LAKSELV": (70.0510, 24.9720, "Lakselv"),
    "HAMAR": (60.7945, 11.0680, "Hamar"),
    "KONGSVINGER": (60.1905, 12.0012, "Kongsvinger"),
    "TRONDHEIM": (63.4305, 10.3951, "Trondheim"),
    "ARENDAL": (58.4618, 8.7723, "Arendal"),
    "DRAMMEN": (59.7440, 10.2045, "Drammen"),
    "FORUS": (58.8930, 5.7330, "Forus / Stavanger"),
    "FREDRIKSTAD": (59.2181, 10.9298, "Fredrikstad"),
    "OSLO": (59.9139, 10.7522, "Oslo"),
    "TROMSO": (69.6492, 18.9553, "Tromsø"),
    "TROMSØ": (69.6492, 18.9553, "Tromsø"),
    "SARPSBORG": (59.2839, 11.1096, "Sarpsborg"),
    "SELJORD": (59.4830, 8.6300, "Seljord"),
    "ELVERUM": (60.8810, 11.5630, "Elverum"),
    "NOTODDEN": (59.5594, 9.2590, "Notodden"),
    "KVAM": (60.3700, 6.1400, "Kvam"),
    # Romania
    "ORADEA": (47.0465, 21.9189, "Oradea"),
    "BUCURESTI": (44.4268, 26.1025, "Bucharest"),
    "PRAHOVA": (44.9500, 26.0200, "Ploiești (Prahova)"),
    "BISTRITA": (47.1332, 24.5001, "Bistrița"),
    "CRAIOVA": (44.3302, 23.7949, "Craiova"),
    # Hungary
    "PECS": (46.0727, 18.2328, "Pécs"),
    "PÉCS": (46.0727, 18.2328, "Pécs"),
    "SZEKSZARD": (46.3470, 18.7060, "Szekszárd"),
    "SZEKSZÁRD": (46.3470, 18.7060, "Szekszárd"),
    "BP.": (47.4979, 19.0402, "Budapest"),
    # Switzerland
    "KERZERS": (46.9750, 7.1950, "Kerzers"),
    "SIERRE": (46.2919, 7.5356, "Sierre"),
    "UZWIL": (47.4360, 9.1330, "Uzwil"),
    "OBERBUREN": (47.4500, 9.1600, "Oberbüren"),
    "OBERBÜREN": (47.4500, 9.1600, "Oberbüren"),
    # Spain
    "TENERIFE": (28.2916, -16.6291, "Tenerife"),
    "BILBO": (43.2630, -2.9350, "Bilbao"),
    "PISUERGA": (41.6523, -4.7245, "Valladolid (Pisuerga)"),
    # UK place tokens
    "KIRRIEMUIR": (56.6730, -3.0040, "Kirriemuir"),
    "GRIMSBY": (53.5670, -0.0800, "Grimsby"),
    "HULL": (53.7676, -0.3274, "Hull"),
    "KINGS LANGLEY": (51.7130, -0.4500, "Kings Langley"),
    "MANCHESTER": (53.4808, -2.2426, "Manchester"),
    "WARRINGTON": (53.3900, -2.5970, "Warrington"),
    "YORK": (53.9591, -1.0815, "York"),
    "ASHBURTON": (50.5160, -3.7560, "Ashburton"),
    "STOKE": (53.0027, -2.1794, "Stoke-on-Trent"),
    "GLOUCESTER": (51.8642, -2.2382, "Gloucester"),
    "TIPTON": (52.5290, -2.0670, "Tipton"),
    "WORCESTER": (52.1936, -2.2215, "Worcester"),
    "ARLESEY": (52.0070, -0.2650, "Arlesey"),
    "BICESTER": (51.9000, -1.1540, "Bicester"),
    "HEATHROW": (51.4700, -0.4543, "Heathrow"),
    "CUMBRIA": (54.4609, -3.0886, "Cumbria"),
    "ELGIN": (57.6494, -3.3186, "Elgin"),
    "INVERNESS": (57.4778, -4.2247, "Inverness"),
    "FERNDOWN": (50.8000, -1.9000, "Ferndown"),
    "IVYBRIDGE": (50.3900, -3.9200, "Ivybridge"),
    "EXETER": (50.7184, -3.5339, "Exeter"),
    "BLACKBURN": (53.7480, -2.4820, "Blackburn"),
    "LEEDS": (53.8008, -1.5491, "Leeds"),
    "BURGESS HILL": (50.9550, -0.1320, "Burgess Hill"),
    "MAIDSTONE": (51.2704, 0.5227, "Maidstone"),
    "WEST THURROCK": (51.4780, 0.2750, "West Thurrock"),
    "NORWICH": (52.6309, 1.2974, "Norwich"),
    "LEICESTER": (52.6369, -1.1398, "Leicester"),
    "NORTHAMPTON": (52.2405, -0.9027, "Northampton"),
    "TUXFORD": (53.2300, -0.9000, "Tuxford"),
    "PETERBOROUGH": (52.5695, -0.2405, "Peterborough"),
    "LEADENHAM": (53.0600, -0.5800, "Leadenham"),
    "LIVINGSTON": (55.8860, -3.5230, "Livingston"),
    # Poland (IDS addresses + official Isuzu PL network cities)
    "TORUN": (53.0138, 18.5984, "Toruń"),
    "TORUŃ": (53.0138, 18.5984, "Toruń"),
    "OLSZTYN": (53.7784, 20.4801, "Olsztyn"),
    "SLUPSK": (54.4641, 17.0285, "Słupsk"),
    "SŁUPSK": (54.4641, 17.0285, "Słupsk"),
    "SWILCZA": (50.0710, 21.9660, "Świlcza (Rzeszów)"),
    "ŚWILCZA": (50.0710, 21.9660, "Świlcza (Rzeszów)"),
    "RZESZOW": (50.0412, 21.9991, "Rzeszów"),
    "RZESZÓW": (50.0412, 21.9991, "Rzeszów"),
    "PLOCK": (52.5463, 19.7065, "Płock"),
    "PŁOCK": (52.5463, 19.7065, "Płock"),
    "DABROWA GORNICZA": (50.3217, 19.1876, "Dąbrowa Górnicza"),
    "DĄBROWA GÓRNICZA": (50.3217, 19.1876, "Dąbrowa Górnicza"),
    "POLKOWICE": (51.5022, 16.0726, "Polkowice"),
    "GDYNIA": (54.5189, 18.5305, "Gdynia"),
    "KOMORNIKI": (52.3380, 16.8070, "Komorniki"),
    "LUBLIN": (51.2465, 22.5684, "Lublin"),
    "WARSZAWA": (52.2297, 21.0122, "Warszawa"),
    "WARSAW": (52.2297, 21.0122, "Warszawa"),
    "BIELSKO-BIALA": (49.8225, 19.0444, "Bielsko-Biała"),
    "BIELSKO-BIAŁA": (49.8225, 19.0444, "Bielsko-Biała"),
    "BIELSKO": (49.8225, 19.0444, "Bielsko-Biała"),
    "KRAKOW": (50.0647, 19.9450, "Kraków"),
    "KRAKÓW": (50.0647, 19.9450, "Kraków"),
    "POZNAN": (52.4064, 16.9252, "Poznań"),
    "POZNAŃ": (52.4064, 16.9252, "Poznań"),
    "MASLOW": (50.9060, 20.7260, "Masłów (Kielce)"),
    "MASŁÓW": (50.9060, 20.7260, "Masłów (Kielce)"),
    "KIELCE": (50.8661, 20.6270, "Kielce"),
    "BIALYSTOK": (53.1325, 23.1688, "Białystok"),
    "BIAŁYSTOK": (53.1325, 23.1688, "Białystok"),
    "KRUPNIKI": (53.1325, 23.1688, "Białystok (Krupniki)"),
    "BLONIE": (52.1950, 20.6170, "Błonie"),
    "BŁONIE": (52.1950, 20.6170, "Błonie"),
    "JELENIA GORA": (50.9044, 15.7194, "Jelenia Góra"),
    "JELENIA GÓRA": (50.9044, 15.7194, "Jelenia Góra"),
    "OPOLE": (50.6751, 17.9213, "Opole"),
    "WEJHEROWO": (54.6056, 18.2335, "Wejherowo"),
    "WROCLAW": (51.1079, 17.0385, "Wrocław"),
    "WROCŁAW": (51.1079, 17.0385, "Wrocław"),
    "GDANSK": (54.3520, 18.6466, "Gdańsk"),
    "GDAŃSK": (54.3520, 18.6466, "Gdańsk"),
    "LODZ": (51.7592, 19.4560, "Łódź"),
    "ŁÓDŹ": (51.7592, 19.4560, "Łódź"),
    "SZCZECIN": (53.4285, 14.5528, "Szczecin"),
    "KATOWICE": (50.2649, 19.0238, "Katowice"),
    "RADOM": (51.4027, 21.1471, "Radom"),
    "CZĘSTOCHOWA": (50.8118, 19.1203, "Częstochowa"),
    "CZESTOCHOWA": (50.8118, 19.1203, "Częstochowa"),
    # --- Hungary cities (from evaluation-sheet addresses) ---
    "BUDAPEST": (47.4979, 19.0402, "Budapest"),
    "NYIREGYHAZA": (47.9553, 21.7167, "Nyíregyháza"),
    "NYÍREGYHÁZA": (47.9553, 21.7167, "Nyíregyháza"),
    "SZEKESFEHERVAR": (47.1890, 18.4100, "Székesfehérvár"),
    "SZÉKESFEHÉRVÁR": (47.1890, 18.4100, "Székesfehérvár"),
    "KESZTHELY": (46.7680, 17.2430, "Keszthely"),
    "BEKESCSABA": (46.6800, 21.1000, "Békéscsaba"),
    "BÉKÉSCSABA": (46.6800, 21.1000, "Békéscsaba"),
    "SZEGED": (46.2530, 20.1414, "Szeged"),
    "VESZPREM": (47.0920, 17.9110, "Veszprém"),
    "VESZPRÉM": (47.0920, 17.9110, "Veszprém"),
    "VERESEGYHAZ": (47.6570, 19.1700, "Veresegyház"),
    "VERESEGYHÁZ": (47.6570, 19.1700, "Veresegyház"),
    "EGER": (47.9020, 20.3770, "Eger"),
    "DEBRECEN": (47.5316, 21.6273, "Debrecen"),
    "KECSKEMET": (46.9060, 19.6910, "Kecskemét"),
    "KECSKEMÉT": (46.9060, 19.6910, "Kecskemét"),
    "MISKOLC": (48.1030, 20.7780, "Miskolc"),
    "CUNEOTRE": (44.3841, 7.5420, "Cuneo"),
    "BIELLESI": (45.5620, 8.0580, "Biella"),
    "SOVECA 43": (45.0430, 3.8850, "Le Puy-en-Velay area"),
    "SOVECA 63": (45.7772, 3.0820, "Clermont-Ferrand area"),

    # --- Belgium (IDS Address row) ---
    "HAMME": (51.0914, 4.1370, "Hamme"),
    "GEEL": (51.1659, 4.9890, "Geel"),
    "ANS": (50.6623, 5.5203, "Ans"),
    "LA LOUVIERE": (50.4768, 4.1879, "La Louvière"),
    "LOUVIERE": (50.4768, 4.1879, "La Louvière"),
    "OOSTERZELE": (50.9490, 3.8030, "Oosterzele"),
    "SCHELDEWINDEKE": (50.9490, 3.8030, "Oosterzele"),
    "SINT-GILLIS-WAAS": (51.2190, 4.1240, "Sint-Gillis-Waas"),
    "ONHAYE": (50.2410, 4.8400, "Onhaye"),
    "PHILIPPEVILLE": (50.1960, 4.5440, "Philippeville"),
    "IZEGEM": (50.9139, 3.2138, "Izegem"),
    "DOUR": (50.3960, 3.7780, "Dour"),
    "NINOVE": (50.8276, 4.0175, "Ninove"),
    "LIBRAMONT": (49.9197, 5.3794, "Libramont"),
    "STREE": (50.2800, 4.3000, "Strée"),
    "STRÉE": (50.2800, 4.3000, "Strée"),
    "OUDSBERGEN": (51.0860, 5.5550, "Oudsbergen"),
    "BRECHT": (51.3500, 4.6400, "Brecht"),
    "WAANRODE": (50.9400, 4.9800, "Waanrode"),
    "KAPELLEN": (51.3130, 4.4350, "Kapellen"),
    "VEURNE": (51.0723, 2.6620, "Veurne"),
    "HOMBEEK": (51.0000, 4.4500, "Hombeek"),
    "ARLON": (49.6833, 5.8167, "Arlon"),
    "SOUMAGNE": (50.6130, 5.7410, "Soumagne"),
    "IEPER": (50.8510, 2.8856, "Ieper"),
    "YPRES": (50.8510, 2.8856, "Ieper"),
    "JABBEKE": (51.1810, 3.0930, "Jabbeke"),
    "BERINGEN": (51.0490, 5.2260, "Beringen"),
    # --- Netherlands ---
    "SPIJKENISSE": (51.8450, 4.3290, "Spijkenisse"),
    "WINTERSWIJK": (51.9725, 6.7190, "Winterswijk"),
    "ZWOLLE": (52.5168, 6.0830, "Zwolle"),
    "VEENENDAAL": (52.0286, 5.5580, "Veenendaal"),
    # --- Denmark ---
    "BOGENSE": (55.5669, 10.0889, "Bogense"),
    "SKAERBAEK": (55.1570, 8.7680, "Skærbæk"),
    "SKÆRBÆK": (55.1570, 8.7680, "Skærbæk"),
    "SVENDBORG": (55.0598, 10.6068, "Svendborg"),
    "AUNING": (56.4310, 10.3780, "Auning"),
    "GRENAA": (56.4158, 10.8783, "Grenå"),
    "GRENÅ": (56.4158, 10.8783, "Grenå"),
    "BALLERUP": (55.7316, 12.3633, "Ballerup"),
    "GALTEN": (56.1560, 9.9070, "Galten"),
    "SJELLE": (56.1560, 9.9070, "Galten"),
    "THISTED": (56.9552, 8.6944, "Thisted"),
    "TARP": (55.4660, 8.5500, "Esbjerg (Tarp)"),
    "ESBJERG": (55.4765, 8.4594, "Esbjerg"),
    "ODENSE": (55.4038, 10.4024, "Odense"),
    # --- Greece (IDS Address row — approx city centres) ---
    "SPARTI": (37.0733, 22.4297, "Sparti"),
    "SPARTA": (37.0733, 22.4297, "Sparti"),
    "MYTILINI": (39.1100, 26.5547, "Mytilini"),
    "MYTILENE": (39.1100, 26.5547, "Mytilini"),
    "KORINTHOS": (37.9386, 22.9322, "Korinthos"),
    "CORINTH": (37.9386, 22.9322, "Korinthos"),
    "HAIDARI": (38.0114, 23.6658, "Haidari"),
    "CHAIDARI": (38.0114, 23.6658, "Haidari"),
    "PATRA": (38.2466, 21.7346, "Patra"),
    "PATRAS": (38.2466, 21.7346, "Patra"),
    "NEA MOUDANIA": (40.2430, 23.2850, "Nea Moudania"),
    "MOUDANIA": (40.2430, 23.2850, "Nea Moudania"),
    "AVLONAS": (38.2510, 23.6940, "Avlonas"),
    "MOIRES": (35.0510, 24.8730, "Moires"),
    "LARISA": (39.6390, 22.4191, "Larisa"),
    "LARISSA": (39.6390, 22.4191, "Larisa"),
    "THESSALONIKI": (40.6401, 22.9444, "Thessaloniki"),
    "RETHYMNO": (35.3656, 24.4821, "Rethymno"),
    "RETHIMNOU": (35.3656, 24.4821, "Rethymno"),
    "RETHYMNOU": (35.3656, 24.4821, "Rethymno"),
    "ATHENS": (37.9838, 23.7275, "Athens"),
    "ATHINA": (37.9838, 23.7275, "Athens"),
    "LAMIA": (38.9000, 22.4333, "Lamia"),
    "TRIKALA": (39.5556, 21.7678, "Trikala"),
    "CHIOS": (38.3678, 26.1358, "Chios"),
    "PIRGOS": (37.6750, 21.4410, "Pyrgos"),
    "PYRGOS": (37.6750, 21.4410, "Pyrgos"),
    "KOZANI": (40.3007, 21.7880, "Kozani"),
    "SKYDRAS": (40.7670, 22.1520, "Skydra"),
    "SKYDRA": (40.7670, 22.1520, "Skydra"),
    "NAFPLIO": (37.5670, 22.8053, "Nafplio"),
    "KALAMATA": (37.0389, 22.1142, "Kalamata"),
    "CHANIA": (35.5138, 24.0180, "Chania"),
    "KARDITSA": (39.3656, 21.9217, "Karditsa"),
    "ALEXANDROUPOLI": (40.8457, 25.8730, "Alexandroupoli"),
    "XANTHI": (41.1347, 24.8880, "Xanthi"),
    "LIVADIA": (38.4362, 22.8747, "Livadia"),
    "IOANNINA": (39.6650, 20.8537, "Ioannina"),
    "AGRINIO": (38.6214, 21.4076, "Agrinio"),
    "TRIPOLI": (37.5101, 22.3796, "Tripoli"),
    "VOLOS": (39.3610, 22.9422, "Volos"),
    "PREVEZA": (38.9592, 20.7510, "Preveza"),
    "SOFIA": (42.6977, 23.3219, "Sofia"),
    # --- Norway (extra from Address) ---
    "ANKENES": (68.4220, 17.3750, "Ankenes"),
    "EGERSUND": (58.4514, 5.9997, "Egersund"),
    "ALVDAL": (62.1080, 10.6320, "Alvdal"),
    "STORD": (59.7790, 5.5000, "Stord"),
    "STJORDAL": (63.4709, 10.9260, "Stjørdal"),
    "STJØRDAL": (63.4709, 10.9260, "Stjørdal"),
    "RISSA": (63.5880, 9.9480, "Rissa"),
    "MALMEFJORDEN": (62.9170, 7.1500, "Malmefjorden"),
    "KARMSUND": (59.3660, 5.3000, "Karmsund / Haugesund"),
    "SKIEN": (59.2085, 9.6090, "Skien"),
    "HUNNDALEN": (60.8000, 10.7000, "Hunndalen / Gjøvik"),
    "HESSENG": (69.6870, 29.9970, "Hesseng"),
    "HELLANDSJO": (63.4500, 8.7000, "Hellandsjøen"),
    # --- UK LCV / IUK (Address) ---
    "GALASHIELS": (55.6140, -2.8070, "Galashiels"),
    "TWEEDBANK": (55.6050, -2.7600, "Tweedbank"),
    "BURTON ON TRENT": (52.8060, -1.6370, "Burton upon Trent"),
    "CHICHESTER": (50.8365, -0.7792, "Chichester"),
    "CHELMSFORD": (51.7356, 0.4685, "Chelmsford"),
    "SOHAM": (52.3330, 0.3360, "Soham"),
    "ELY": (52.3990, 0.2620, "Ely"),
    "BOSTON": (52.9780, -0.0260, "Boston"),
    "ROBERTSBRIDGE": (50.9850, 0.4740, "Robertsbridge"),
    "EAGLESHAM": (55.7410, -4.2750, "Eaglesham"),
    "GLASGOW": (55.8642, -4.2518, "Glasgow"),
    "MAYDOWN": (55.0400, -7.2800, "Maydown / Derry"),
    "WINCANTON": (51.0560, -2.4100, "Wincanton"),
    "HALESWORTH": (52.3440, 1.5030, "Halesworth"),
    "MACCLESFIELD": (53.2587, -2.1231, "Macclesfield"),
    "MARTON": (53.2100, -2.2200, "Marton / Macclesfield"),
    "ANTRIM": (54.7170, -6.2070, "Antrim"),
    "IPSWICH": (52.0567, 1.1482, "Ipswich"),
    "LISNASKEA": (54.2530, -7.4420, "Lisnaskea"),
    "TINGEWICK": (51.9900, -1.0500, "Tingewick"),
    "KENDAL": (54.3280, -2.7470, "Kendal"),
    "LEVENS": (54.2700, -2.7800, "Levens / Kendal"),
    "MENAI BRIDGE": (53.2260, -4.1630, "Menai Bridge"),
    "CARLISLE": (54.8951, -2.9382, "Carlisle"),
    "SHEFFIELD": (53.3811, -1.4701, "Sheffield"),
    "DARLINGTON": (54.5270, -1.5528, "Darlington"),
    # --- UK CV / ITUK name tokens ---
    "CYMRU": (52.1307, -3.7837, "Wales (approx)"),
    "NORFOLK": (52.6309, 1.2974, "Norwich (Norfolk)"),
    # --- France name tokens (extra) ---
    "ANDRES": (50.8670, 1.9200, "Andres"),
    "ALPES": (45.1885, 5.7245, "Grenoble (Alpes)"),
    "SARTHOIS": (48.0061, 0.1996, "Le Mans (Sarthe)"),
    "GARD": (43.8367, 4.3601, "Nîmes (Gard)"),
    "LEREENNE": (47.3320, 2.8700, "Léré"),
    # --- Spain name tokens ---
    "BILBOTRUCK": (43.2630, -2.9350, "Bilbao"),
    "VIGATANA": (41.9300, 2.2540, "Vic"),
    "ILERDAUTO": (41.6170, 0.6200, "Lleida"),
    "LEOMOTOR": (42.5987, -5.5671, "León"),
    "ROFERVIGO": (42.2406, -8.7207, "Vigo"),
    "TALAUTO": (40.4168, -3.7038, "Madrid (Talauto)"),
    "GAURSA": (42.8467, -2.6716, "Vitoria-Gasteiz"),
    "ARABAUTO": (42.8467, -2.6716, "Álava"),
    # --- Italy name tokens ---
    "ADRIATICA": (43.6158, 13.5189, "Ancona (Adriatica)"),
    "ERZELLI": (44.4056, 8.9463, "Genova (Erzelli)"),
    "BURANO": (45.4850, 12.4170, "Burano / Venezia"),
    "MONTELLO": (45.8000, 11.9500, "Montello / Treviso"),
    "MITTELCAR": (46.4983, 11.3548, "Bolzano area"),
    # --- Hungary ---
    "SZFV": (47.1890, 18.4100, "Székesfehérvár"),
    "SZEKESFEHERVAR": (47.1890, 18.4100, "Székesfehérvár"),
    "SZEKESFEHÉRVAR": (47.1890, 18.4100, "Székesfehérvár"),
    "KOROS": (46.6800, 21.1000, "Körös area"),
    "KÖRÖS": (46.6800, 21.1000, "Körös area"),
    # --- Switzerland (MIDI NETWORK Town when lat/lon blank) ---
    "ALOUETTES": (46.5197, 6.6323, "Lausanne area"),
    "BIRRER": (47.5530, 7.9830, "Sisseln"),
    "EMIL FREY": (46.8500, 9.5320, "Chur"),
    "SETTELEN": (47.5596, 7.5886, "Basel"),
    "VUD": (46.7020, 6.5700, "Chavornay"),
    "ROMANDIE": (46.7020, 6.5700, "Chavornay"),
    "GEHRIG": (47.3700, 8.3800, "Rudolfstetten"),
    "TRAMELAN": (47.2250, 7.1000, "Tramelan"),
    "COLDRERIO": (45.8540, 9.0050, "Coldrerio"),
    "ENNENDA": (47.0330, 9.0830, "Ennenda"),
    "SISSELN": (47.5530, 7.9830, "Sisseln"),
    "BIASCA": (46.3600, 8.9700, "Biasca"),
    "CHUR": (46.8500, 9.5320, "Chur"),
    "GERLAFINGEN": (47.1700, 7.5700, "Gerlafingen"),
    "FAULENSEE": (46.6750, 7.7000, "Faulensee"),
    "REINACH": (47.2550, 8.1800, "Reinach AG"),
    "KEMPTTHAL": (47.4550, 8.7050, "Kemptthal"),
    "LAMBOING": (47.1850, 7.2100, "Lamboing"),
    "MUNCHALTORF": (47.2820, 8.6500, "Mönchaltorf"),
    "MÜNCHALTORF": (47.2820, 8.6500, "Mönchaltorf"),
    "MONCHALTORF": (47.2820, 8.6500, "Mönchaltorf"),
    "MÖNCHALTORF": (47.2820, 8.6500, "Mönchaltorf"),
    "CHAM": (47.1820, 8.4630, "Cham"),
    "HOLSTEIN": (47.4300, 7.7700, "Hölstein"),
    "HÖLSTEIN": (47.4300, 7.7700, "Hölstein"),
    "OBERDORF": (46.9600, 8.3900, "Oberdorf NW"),
    "BASEL": (47.5596, 7.5886, "Basel"),
    "CHAVORNAY": (46.7020, 6.5700, "Chavornay"),
    "RUDOLFSTETTEN": (47.3700, 8.3800, "Rudolfstetten"),
    "ALTSTATTEN": (47.3780, 9.5550, "Altstätten"),
    "ALTSTÄTTEN": (47.3780, 9.5550, "Altstätten"),
    # --- Italy towns (MIDI NETWORK without lat) ---
    "TEGGIANO": (40.3800, 15.5400, "Teggiano"),
    "MODUGNO": (41.0850, 16.7830, "Modugno"),
    # --- Spain towns (MIDI NETWORK without lat / name tokens) ---
    "ALMENDRALEJO": (38.6830, -6.4080, "Almendralejo"),
    "ZARAGOZA": (41.6488, -0.8891, "Zaragoza"),
    "VITORIA": (42.8467, -2.6716, "Vitoria-Gasteiz"),
    "GASTEIZ": (42.8467, -2.6716, "Vitoria-Gasteiz"),
    # --- UK CV / ITUK cities ---
    "BRISTOL": (51.4545, -2.5879, "Bristol"),
    "TEWKESBURY": (51.9920, -2.1560, "Tewkesbury"),
    "BASINGSTOKE": (51.2667, -1.0876, "Basingstoke"),
    "SOUTHAMPTON": (50.9097, -1.4044, "Southampton"),
    "MARCHWOOD": (50.8900, -1.4540, "Marchwood"),
    "TREDEGAR": (51.7200, -3.2500, "Tredegar"),
    "REDRUTH": (50.2330, -5.2240, "Redruth"),
    "TRURO": (50.2630, -5.0510, "Truro"),
    "TIVERTON": (50.9020, -3.4860, "Tiverton"),
    "CARDIFF": (51.4816, -3.1791, "Cardiff"),
    "BARKING": (51.5360, 0.0800, "Barking"),
    "DAGENHAM": (51.5400, 0.1400, "Dagenham"),
    "TAKELEY": (51.8680, 0.2660, "Takeley"),
    "PARK ROYAL": (51.5280, -0.2700, "Park Royal"),
    # --- Romania extras from filenames / names ---
    "HINODE": (44.4268, 26.1025, "Bucharest area"),
    "MARSHAL": (45.9432, 24.9668, "Romania (Marshal)"),
    "RODAV": (45.6500, 25.6000, "Brașov area"),
    "OPEN AUTO": (44.4268, 26.1025, "Bucharest area"),
    "PREMIUM CARS": (44.4268, 26.1025, "Bucharest area"),
    "LUK MOTORS": (45.7489, 21.2257, "Timișoara area"),
    "MIT MOTORS": (46.7712, 23.6236, "Cluj area"),
    "AUTO KINO": (45.7489, 21.2257, "Timișoara area"),
    "AUTOCOMO": (44.3302, 23.7949, "Craiova area"),
    "IMG": (44.4268, 26.1025, "Bucharest area"),

}

_DEALER_GEO_OVERRIDE: dict[str, tuple[float, float, str]] = {
    # Portugal public directories
    "ROLMAX": (41.1239, -8.6118, "Vila Nova de Gaia"),
    "AUTO IMPERIAL": (41.8061, -6.7567, "Bragança"),
    "MSCAR": (37.0870, -8.0307, "Almancil"),
    "SOTERMÁQUINAS": (38.6593, -27.2194, "Angra do Heroísmo"),
    "SOTERMAQUINAS": (38.6593, -27.2194, "Angra do Heroísmo"),
    "XANAUTO": (38.8309, -9.1681, "Loures"),
    # Single-dealer countries / clear HQ
    "A. LOUTSIOS AND SONS LTD": (34.7071, 33.0226, "Limassol"),
    "GASANZAMMIT MOTORS LIMITED": (35.8997, 14.5146, "Msida / Malta"),
    "AUTOPOLIS ESCH": (49.4958, 5.9806, "Esch-sur-Alzette"),
    # Israel without clear token
    "NEEMAN": (32.0853, 34.7818, "Tel Aviv area"),
    "KFIR YAUD": (32.0800, 34.8900, "Central Israel"),
    "LEV HAGALIL": (32.9200, 35.3000, "Upper Galilee"),

    # NL without clear city token in short address
    "AUTOBEDRIJF TIGGELOVEN B.V.": (51.9725, 6.7190, "Winterswijk"),
    "KOEN DIJKERS": (51.9200, 4.4500, "Rotterdam area"),
    # Belgium incomplete address
    "AUTO.IT": (50.6326, 5.5797, "Liège area"),
    # Denmark without city in short address
    "KJERULFF A/S": (55.4038, 10.4024, "Odense area"),
    "ANDERS HANSEN AUTOMOBILER": (55.4660, 8.5500, "Esbjerg (Tarp)"),
    "BILBOEL": (55.2560, 10.5980, "Ørbæk"),
    "JENS E BILER": (55.6761, 12.5683, "Copenhagen area"),
    # Norway short address
    "GJERMUNDSEN BILSENTER A/S": (59.2675, 10.4076, "Tønsberg area"),
    "BLAKSTAD AUTO AS": (63.4500, 8.7000, "Hellandsjøen"),
    "TUNGEN BIL AS": (63.4305, 10.3951, "Trondheim area"),
    # UK LCV / IUK no/weak address
    "ADAM": (54.8951, -2.9382, "Carlisle area"),
    "CROSSKEYS": (51.4800, -3.1800, "Cardiff / Crosskeys"),
    "FUSSELLWADMAN": (51.3800, -2.3600, "Bath / Somerset area"),
    "GARETHHUGHS": (53.2280, -4.1280, "North Wales"),
    "PATKIRK": (54.5970, -5.9300, "Northern Ireland"),
    "TANNERS": (51.7520, -1.2577, "Oxford area"),
    "THREEWAYS": (53.2280, -4.1280, "North Wales"),
    "M M BELLINGER AND SONS LTD": (51.6100, -1.4200, "Grove / Oxfordshire"),
    "COMMERCIAL MOTORS (WALES) LTD": (52.1307, -3.7837, "Wales (approx)"),
    "AUTO NART AG": (47.0330, 9.0830, "Ennenda"),
    # Poland — short / ambiguous address or name without city token
    "TOP AUTO ANDRZEJ AREFIEW": (53.1325, 23.1688, "Białystok"),
    # Hungary — report names → cities from evaluation-sheet addresses
    "AUTÓ CARO": (47.5070, 19.0820, "Budapest"),
    "AUTO CARO": (47.5070, 19.0820, "Budapest"),
    "CSŐSZ": (47.9553, 21.7167, "Nyíregyháza"),
    "CSOSZ": (47.9553, 21.7167, "Nyíregyháza"),
    "FÁBIÁN BP.ISUZU MÁRKSZERVIZ SZABVÁNYOK 2026 MAGYAR": (47.4979, 19.0402, "Budapest"),
    "FABIAN BP.ISUZU MARKSZERVIZ SZABVANYOK 2026 MAGYAR": (47.4979, 19.0402, "Budapest"),
    "GADÁCSI": (46.7680, 17.2430, "Keszthely"),
    "GADACSI": (46.7680, 17.2430, "Keszthely"),
    "KÖRÖS": (46.6800, 21.1000, "Békéscsaba"),
    "KOROS": (46.6800, 21.1000, "Békéscsaba"),
    "LINARTECH": (46.2530, 20.1414, "Szeged"),
    "MARSAL": (47.0920, 17.9110, "Veszprém"),
    "PÁSZTOR": (47.6570, 19.1700, "Veresegyház"),
    "PASZTOR": (47.6570, 19.1700, "Veresegyház"),
    "SHICK-R": (47.9020, 20.3770, "Eger"),
    "TIPP-AUTÓ": (47.5316, 21.6273, "Debrecen"),
    "TIPP-AUTO": (47.5316, 21.6273, "Debrecen"),
    "TORMÁSI ISUZU MÁRKASZERVIZ SZABVÁNYOK 2026 MAGYAR": (46.9060, 19.6910, "Kecskemét"),
    "TORMASI ISUZU MARKASZERVIZ SZABVANYOK 2026 MAGYAR": (46.9060, 19.6910, "Kecskemét"),
    "UNIVER-CAR": (48.1030, 20.7780, "Miskolc"),
    "WMP": (47.4979, 19.0402, "Budapest"),
    # Italy
    "CUNEOTRE S.P.A.": (44.3841, 7.5420, "Cuneo"),
    "NUOVA VIB VEICOLI INDUSTRIALI BIELLESI S.R.L.": (45.5620, 8.0580, "Biella"),
    # UK CV / ITUK — Isuzu Truck UK store pages / public dealer addresses
    "NORSCOT TRUCK AND VAN LTD": (57.4778, -4.2247, "Inverness area"),
    "CORDWALLIS COMMERCIALS": (51.4700, -0.4543, "Heathrow / Cordwallis"),
    "PAYNES GARAGES LTD": (52.6369, -1.1398, "Leicester area"),
    "PENTAGON GROUP": (52.9225, -1.4746, "Derby area"),
    "A K COMMERCIALS LTD": (51.4545, -2.5879, "Bristol"),
    "ARTHUR SPRIGGS LTD": (51.9920, -2.1560, "Tewkesbury"),
    "D T TRUCKS LTD": (51.5360, 0.0800, "Barking"),
    "MANCHETTS": (52.5695, -0.2405, "Peterborough"),
    "PR FLEET SERVICES LTD": (51.5280, -0.2700, "Park Royal"),
    "SAPPHIRE VEHICLE SERVICES": (51.2667, -1.0876, "Basingstoke"),
    "SHEPPARD COMMERCIAL SERVICES LTD": (50.8900, -1.4540, "Marchwood"),
    "SW COMMERCIALS LTD": (50.2630, -5.0510, "Truro area"),
    "T & M COMMERCIALS LTD": (50.9020, -3.4860, "Tiverton"),
    "TANNERS ELECTRICS LTD": (51.4816, -3.1791, "Cardiff"),
    "WALDRON COMMERCIALS LTD": (51.7200, -3.2500, "Tredegar"),
    "T R AUTOS 2003": (51.8680, 0.2660, "Takeley"),
    "T.I.M COMMERCIALS": (52.9710, -1.3090, "Ilkeston"),
    "TIM COMMERCIALS": (52.9710, -1.3090, "Ilkeston"),
}

# Mainland / core IDS bounding boxes — used to reject absurd city matches and
# to keep the map focused on Europe/Israel (not DOM-TOM / wrong continent).
# (min_lat, max_lat, min_lon, max_lon) — generous for islands that belong.
_COUNTRY_BOUNDS: dict[str, tuple[float, float, float, float]] = {
    "Belgium": (49.4, 51.6, 2.4, 6.5),
    "Benelux": (49.4, 53.7, 2.4, 7.3),
    "Cyprus": (34.5, 35.8, 32.0, 34.8),
    "Denmark": (54.5, 57.9, 8.0, 15.3),
    "France": (41.0, 51.5, -5.5, 10.0),  # metro only — DOM-TOM verified separately
    "Greece": (34.7, 41.8, 19.3, 29.7),
    "Hungary": (45.7, 48.6, 16.0, 23.0),
    "Israel LCV": (29.4, 33.4, 34.2, 35.9),
    "Italy": (36.0, 47.5, 6.5, 19.0),
    "Luxembourg": (49.4, 50.2, 5.7, 6.6),
    "Malta": (35.7, 36.1, 14.1, 14.6),
    "Netherlands": (50.7, 53.7, 3.2, 7.3),
    "Norway": (57.8, 71.5, 4.5, 31.5),
    "Poland": (49.0, 55.0, 14.0, 24.5),
    "Portugal": (36.8, 42.3, -9.6, -6.0),  # mainland; Açores kept via allowlist
    "Romania": (43.5, 48.3, 20.0, 30.0),
    "Spain": (35.9, 44.0, -9.5, 4.5),  # mainland + Balearics; Canaries allowlisted
    "Switzerland": (45.7, 47.9, 5.8, 10.6),
    "UK CV": (49.8, 61.0, -8.5, 2.0),
    "UK LCV": (49.8, 61.0, -8.5, 2.0),
}

# Cities that exist in more than one country — only match when country matches.
_CITY_COUNTRY_SCOPE: dict[str, frozenset[str]] = {
    "SOFIA": frozenset({"Bulgaria"}),  # never Greece/IDS — "Sofia France Auto" false hit
}

# Verified French overseas (DOM-TOM) markers — stay on the map; fitBounds stays Europe.
# Prefer MIDI "FRANCE DOM TOM" + hard lat/lon far-box over loose place substrings
# (Saint-Pierre-Du-Mont / Saint-Pierre-En-Faucigny are metro France).
_OVERSEAS_PLACE_MARKERS = (
    "FORT-DE-FRANCE",
    "MAMOUDZOU",
    "GUADELOUPE",
    "MARTINIQUE",
    "MAYOTTE",
    "GUYANE",
    "NOUMEA",
    "TAHITI",
    "PAPEETE",
    "CAYENNE",
    "POINTE-A-PITRE",
    "SAINT-PIERRE CEDEX",  # Réunion — not Saint-Pierre-Du-Mont / not Madagascar
)

# Clear shared place labels so Réunion is never read as Madagascar.
# Longer / specific markers first — do not add a bare "REUNION" entry that
# would rewrite "Saint-Pierre (Réunion)" down to "Réunion" on a second pass.
_OVERSEAS_PLACE_LABELS: tuple[tuple[str, str], ...] = (
    ("FORT-DE-FRANCE", "Fort-de-France (Martinique)"),
    ("MAMOUDZOU", "Mamoudzou (Mayotte)"),
    ("MAYOTTE", "Mamoudzou (Mayotte)"),
    ("SAINT-PIERRE CEDEX", "Saint-Pierre (Réunion)"),
    ("SAINT-PIERRE (RÉUNION)", "Saint-Pierre (Réunion)"),
    ("SAINT-PIERRE (REUNION)", "Saint-Pierre (Réunion)"),
)

# Country centre used when no city token / override (jittered per dealer).
_COUNTRY_CENTRE: dict[str, tuple[float, float, str]] = {
    "Belgium": (50.8503, 4.3517, "Belgium (approx)"),
    "Benelux": (51.0, 4.5, "Benelux (approx)"),
    "Cyprus": (34.7071, 33.0226, "Cyprus (approx)"),
    "Denmark": (56.2639, 9.5018, "Denmark (approx)"),
    "France": (46.6034, 1.8883, "France (approx)"),
    "Greece": (39.0742, 21.8243, "Greece (approx)"),
    "Hungary": (47.1625, 19.5033, "Hungary (approx)"),
    "Israel LCV": (31.5, 34.9, "Israel (approx)"),
    "Italy": (42.5, 12.5, "Italy (approx)"),
    "Luxembourg": (49.8153, 6.1296, "Luxembourg (approx)"),
    "Malta": (35.8997, 14.5146, "Malta (approx)"),
    "Netherlands": (52.1326, 5.2913, "Netherlands (approx)"),
    "Norway": (64.0, 12.0, "Norway (approx)"),
    "Poland": (52.0, 19.0, "Poland (approx)"),
    "Portugal": (39.5, -8.0, "Portugal (approx)"),
    "Romania": (45.9, 24.9, "Romania (approx)"),
    "Spain": (40.2, -3.7, "Spain (approx)"),
    "Switzerland": (46.8, 8.2, "Switzerland (approx)"),
    # Centroids follow the population, not the label: IUK (= UK LCV) skews
    # north/Scotland, ITUK (= UK CV) truck centres sit mid-England.
    "UK LCV": (54.0, -2.5, "UK (approx)"),
    "UK CV": (52.5, -1.5, "UK (approx)"),
}

# Sanity-check thresholds (always run on rebuild)
_CLUSTER_RADIUS_KM = 10.0
_CLUSTER_SHARE_FAIL = 0.50
_MIN_DEALERS_FOR_GEO_FAIL = 5

# Countries where we expect strong city placement (addresses / name tokens).
# Regression below the floor fails the dashboard build.
_EXPECTED_CITY_FLOOR: dict[str, float] = {
    "Poland": 0.80,
    "Portugal": 0.80,
    "Greece": 0.70,
    "Belgium": 0.60,
    "Netherlands": 0.60,
    "Israel LCV": 0.70,
    "Cyprus": 0.50,
    "Malta": 0.50,
    "Luxembourg": 0.50,
    # MIDI NETWORK lat/lon (or Town) covers FR/IT/ES/CH
    "France": 0.85,
    "Italy": 0.85,
    "Spain": 0.85,
    "Switzerland": 0.80,
}

_MIDI_COUNTRY_MAP = {
    "FRANCE": "France",
    "FRANCE DOM TOM": "France",
    "ITALY": "Italy",
    "SPAIN": "Spain",
    "SWISS": "Switzerland",
    "SWITZERLAND": "Switzerland",
}


def _midi_file_path() -> Path | None:
    try:
        from ids_paths import MIDI_FILE

        p = Path(MIDI_FILE)
        return p if p.is_file() else None
    except Exception:
        return None


@lru_cache(maxsize=1)
def load_midi_network_geo() -> dict[str, dict[str, Any]]:
    """Map normalised DLRName → {lat, lon, town, address, postal, country}.

    Source: MIDI_FR.xlsx NETWORK sheet (Town + Latitude/Longitude). Exact name
    match against IDS MIDI result-report dealers is 1:1 for FR/IT/ES/CH.
    """
    path = _midi_file_path()
    if path is None:
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    out: dict[str, dict[str, Any]] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "NETWORK" not in wb.sheetnames:
            return {}
        ws = wb["NETWORK"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[0] if row else None
            if not name:
                continue
            raw_country = str(row[5] or "").strip().upper()
            country = _MIDI_COUNTRY_MAP.get(raw_country)
            if not country:
                continue
            lat = lon = None
            try:
                if row[6] not in (None, ""):
                    lat = float(row[6])
                if row[7] not in (None, ""):
                    lon = float(row[7])
            except (TypeError, ValueError):
                lat = lon = None
            town = str(row[4] or "").strip()
            address = str(row[1] or "").strip()
            postal = str(row[2] or "").strip()
            full = str(row[13] or "").strip() if len(row) > 13 else ""
            key = _norm_dealer_key(str(name))
            out[key] = {
                "name": str(name).strip(),
                "country": country,
                "midi_country": raw_country,
                "town": town,
                "address": address,
                "postal": postal,
                "full_address": full,
                "lat": lat,
                "lon": lon,
            }
    finally:
        wb.close()
    return out


def midi_lookup(name: str) -> dict[str, Any] | None:
    """Exact normalised-name lookup into MIDI NETWORK."""
    if not (name or "").strip():
        return None
    return load_midi_network_geo().get(_norm_dealer_key(name))


def enrich_address_from_midi(dealers: list[dict[str, Any]]) -> int:
    """Fill empty dealer address from MIDI NETWORK Town / street. Returns count."""
    n = 0
    for d in dealers:
        if (d.get("address") or "").strip():
            continue
        hit = midi_lookup(d.get("name") or "")
        if not hit:
            continue
        bits = [
            hit.get("address") or "",
            " ".join(
                x for x in [(hit.get("postal") or ""), (hit.get("town") or "")] if x
            ).strip(),
        ]
        addr = ", ".join(b for b in bits if b).strip()
        if not addr and hit.get("full_address"):
            addr = str(hit["full_address"]).strip()
        if not addr and hit.get("town"):
            addr = str(hit["town"]).strip()
        if addr:
            d["address"] = addr
            n += 1
    return n


def _norm_dealer_key(name: str) -> str:
    upper = (name or "").upper().strip()
    for a, b in (
        ("Á", "A"),
        ("À", "A"),
        ("Ã", "A"),
        ("Â", "A"),
        ("É", "E"),
        ("Ê", "E"),
        ("Í", "I"),
        ("Ó", "O"),
        ("Õ", "O"),
        ("Ô", "O"),
        ("Ú", "U"),
        ("Ç", "C"),
        ("Ø", "O"),
        ("Å", "A"),
        ("Æ", "AE"),
        ("Ő", "O"),
        ("Ű", "U"),
    ):
        upper = upper.replace(a, b)
    return " ".join(upper.split())


def _jitter(seed: str, lat: float, lon: float, *, km: float = 28.0) -> tuple[float, float]:
    """Deterministic offset so country-centre pins do not stack exactly."""
    h = hashlib.md5(seed.encode("utf-8")).hexdigest()
    a = int(h[:8], 16) / 0xFFFFFFFF
    b = int(h[8:16], 16) / 0xFFFFFFFF
    ang = a * 2 * math.pi
    dist = (0.25 + 0.75 * b) * km
    dlat = (dist / 111.0) * math.cos(ang)
    dlon = (dist / (111.0 * max(0.2, math.cos(math.radians(lat))))) * math.sin(ang)
    return lat + dlat, lon + dlon


def _match_city_token(
    text: str, *, country: str | None = None
) -> tuple[float, float, str] | None:
    """Longest city-token match in free text (address or dealer name).

    Uses alphanumeric word boundaries so short tokens (e.g. BILBO) do not
    false-hit longer names (Bilboel). Country-scoped tokens (e.g. SOFIA)
    only match when the dealer country is in the allow-list.
    """
    upper = (text or "").upper()
    if not upper.strip():
        return None
    for key, (lat, lon, label) in sorted(
        _CITY_COORDS.items(), key=lambda kv: -len(kv[0])
    ):
        scope = _CITY_COUNTRY_SCOPE.get(key)
        if scope is not None and (country or "") not in scope:
            continue
        pat = r"(?<![A-Z0-9])" + re.escape(key) + r"(?![A-Z0-9])"
        if re.search(pat, upper):
            if country and not _coords_in_country_bounds(lat, lon, country):
                continue
            return lat, lon, label
    return None


def _coords_in_country_bounds(lat: float, lon: float, country: str) -> bool:
    bounds = _COUNTRY_BOUNDS.get(country)
    if not bounds:
        return True
    lo_la, hi_la, lo_lo, hi_lo = bounds
    return lo_la <= lat <= hi_la and lo_lo <= lon <= hi_lo


def _is_verified_dom_tom(
    *,
    place: str | None,
    lat: float | None,
    lon: float | None,
    midi_country: str | None = None,
) -> bool:
    """True for known French DOM-TOM (Martinique / Réunion / Mayotte / …).

    Réunion (lon ~55.5) is east of Madagascar — not Madagascar itself.
    """
    if midi_country and "DOM" in str(midi_country).upper():
        return True
    place_u = (place or "").upper()
    if any(m in place_u for m in _OVERSEAS_PLACE_MARKERS):
        return True
    if lat is None or lon is None:
        return False
    # Hard boxes for the three MIDI DOM-TOM pins (sanity / anti-Madagascar).
    if 14.3 <= lat <= 14.9 and -61.3 <= lon <= -60.8:
        return True  # Martinique / Fort-de-France
    if -21.5 <= lat <= -20.8 and 55.2 <= lon <= 55.9:
        return True  # Réunion / Saint-Pierre (NOT Madagascar ~lon 43–50)
    if -13.1 <= lat <= -12.5 and 45.0 <= lon <= 45.4:
        return True  # Mayotte / Mamoudzou
    return False


def _overseas_display_place(
    place: str | None,
    *,
    lat: float | None = None,
    lon: float | None = None,
) -> str | None:
    """Rewrite DOM-TOM town labels so Réunion is not confused with Madagascar."""
    place_u = (place or "").upper()
    for marker, label in _OVERSEAS_PLACE_LABELS:
        if marker in place_u:
            return label
    if lat is not None and lon is not None:
        if 14.3 <= lat <= 14.9 and -61.3 <= lon <= -60.8:
            return "Fort-de-France (Martinique)"
        if -21.5 <= lat <= -20.8 and 55.2 <= lon <= 55.9:
            return "Saint-Pierre (Réunion)"
        if -13.1 <= lat <= -12.5 and 45.0 <= lon <= 45.4:
            return "Mamoudzou (Mayotte)"
    return place


def _is_overseas_pin(
    *,
    place: str | None,
    lat: float | None,
    lon: float | None,
    country: str | None,
    midi_country: str | None = None,
) -> bool:
    """True for DOM-TOM / far overseas (outside Europe-core fitBounds)."""
    if _is_verified_dom_tom(
        place=place, lat=lat, lon=lon, midi_country=midi_country
    ):
        return True
    if lat is None or lon is None:
        return False
    # Far outside Europe/Israel box used by IDS network map
    if lon < -20 or lon > 40 or lat < 28 or lat > 72:
        # Açores / Madeira / Canaries / Cyprus already in or near box; flag rest
        if country == "Portugal" and -32 <= lon <= -6 and 32 <= lat <= 40:
            return False  # Açores / Madeira — keep on map, fitBounds excludes
        if country == "Spain" and -19 <= lon <= -13 and 27 <= lat <= 30:
            return False  # Canaries — keep, fitBounds excludes
        if country in {"Cyprus", "Malta", "Israel LCV"}:
            return False
        return True
    # Metro France / Europe coords are never overseas via place-name alone.
    return False


def _should_map_exclude_overseas(
    *,
    place: str | None,
    lat: float | None,
    lon: float | None,
    midi_country: str | None = None,
) -> bool:
    """Hide mystery far pins; keep verified French DOM-TOM on the map."""
    if _is_verified_dom_tom(
        place=place, lat=lat, lon=lon, midi_country=midi_country
    ):
        return False
    return True


def _haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(min(1.0, math.sqrt(a)))


def place_for_dealer(
    name: str,
    country: str | None = None,
    *,
    address: str | None = None,
    allow_country_centre: bool = True,
) -> tuple[float | None, float | None, str | None, str]:
    """Return (lat, lon, place, geo_source_label)."""
    raw = (name or "").strip()
    upper = raw.upper()
    upper = upper.replace("PAÃOS", "PAÇOS").replace("PACOS", "PAÇOS")
    key_norm = _norm_dealer_key(raw)

    for override_key, (lat, lon, label) in _DEALER_GEO_OVERRIDE.items():
        if key_norm == _norm_dealer_key(override_key) or upper == override_key.upper():
            return lat, lon, label, "Approx city (directory / known place)"

    midi = midi_lookup(raw)
    if midi and midi.get("lat") is not None and midi.get("lon") is not None:
        town = (midi.get("town") or "").strip() or "MIDI network"
        place = town.title() if town == town.upper() else town
        lat, lon = float(midi["lat"]), float(midi["lon"])
        # Prefer MIDI coords only when they sit in the dealer's country bbox
        # (or are explicitly overseas DOM-TOM — still returned; map may show them).
        midi_cty = (midi.get("country") or country or "")
        midi_raw = str(midi.get("midi_country") or "")
        overseas = _is_overseas_pin(
            place=place, lat=lat, lon=lon, country=country, midi_country=midi_raw
        )
        if overseas or not country or _coords_in_country_bounds(lat, lon, country) or midi_cty == country:
            if overseas:
                place = _overseas_display_place(place, lat=lat, lon=lon) or place
                verified = _is_verified_dom_tom(
                    place=place, lat=lat, lon=lon, midi_country=midi_raw
                )
                src = (
                    "Overseas territory (French DOM-TOM)"
                    if verified
                    else "Overseas territory (far from Europe core)"
                )
            else:
                src = "Approx city (MIDI network)"
            return lat, lon, place, src

    # Prefer MIDI town text when result reports have no Address row
    midi_town = (midi or {}).get("town") if midi else None
    addr_text = " ".join(
        x for x in [(address or "").strip(), (midi_town or "").strip()] if x
    )
    addr_hit = _match_city_token(addr_text, country=country)
    if addr_hit:
        lat, lon, label = addr_hit
        src = (
            "Approx city from MIDI town"
            if midi_town and not (address or "").strip()
            else "Approx city from address"
        )
        return lat, lon, label, src

    name_hit = _match_city_token(upper, country=country)
    if name_hit:
        lat, lon, label = name_hit
        return lat, lon, label, "Approx city from dealer name"

    if allow_country_centre and country and country in _COUNTRY_CENTRE:
        clat, clon, clabel = _COUNTRY_CENTRE[country]
        jlat, jlon = _jitter(f"{country}::{raw}", clat, clon)
        return jlat, jlon, clabel, "Approx country centre (no city in name)"

    return None, None, None, "No geo"


def attach_geo_to_dealers(
    dealers: list[dict[str, Any]],
    *,
    allow_country_centre: bool = True,
) -> dict[str, int]:
    """Mutate dealers in place: add geo dict (or null). Returns coverage counts."""
    enrich_address_from_midi(dealers)
    ok = 0
    city = 0
    country_fb = 0
    missing = 0
    for d in dealers:
        lat, lon, place, label = place_for_dealer(
            d.get("name") or "",
            d.get("country"),
            address=d.get("address"),
            allow_country_centre=allow_country_centre,
        )
        if lat is None:
            d["geo"] = None
            missing += 1
            continue
        midi_raw = (midi_lookup(d.get("name") or "") or {}).get("midi_country")
        overseas = _is_overseas_pin(
            place=place,
            lat=lat,
            lon=lon,
            country=d.get("country"),
            midi_country=midi_raw,
        )
        if overseas:
            place = _overseas_display_place(place, lat=lat, lon=lon) or place
            verified = _is_verified_dom_tom(
                place=place, lat=lat, lon=lon, midi_country=midi_raw
            )
            if verified:
                label = "Overseas territory (French DOM-TOM)"
            elif "Overseas" not in (label or ""):
                label = "Overseas territory (far from Europe core)"
        map_exclude = bool(
            overseas
            and _should_map_exclude_overseas(
                place=place, lat=lat, lon=lon, midi_country=midi_raw
            )
        )
        d["geo"] = {
            "lat": round(lat, 5),
            "lon": round(lon, 5),
            "place": place,
            "geo_label": label,
            "map_exclude": map_exclude,
        }
        ok += 1
        if "country centre" in label:
            country_fb += 1
        else:
            city += 1
    return {
        "geo_ok": ok,
        "geo_city": city,
        "geo_country_fallback": country_fb,
        "geo_missing": missing,
        "dealer_count": len(dealers),
    }


def map_points_from_dealers(
    dealers: list[dict[str, Any]],
    *,
    country_avg_pct: float | None = None,
    skip_benelux_duplicates: bool = True,
    allow_country_centre: bool = True,
) -> list[dict[str, Any]]:
    """Ranking-style map points with dealer %, country avg %, delta pp."""
    # Prefer non-Benelux rows when the same name exists under BE/NL/LUX
    skip_names: set[str] = set()
    if skip_benelux_duplicates:
        for d in dealers:
            if d.get("country") != "Benelux":
                skip_names.add((d.get("name") or "").strip().casefold())

    ranked = sorted(
        [d for d in dealers if d.get("total_pct") is not None],
        key=lambda d: (-(d.get("total_pct") or 0), d.get("name") or ""),
    )
    points: list[dict[str, Any]] = []
    for i, d in enumerate(ranked, 1):
        ctry = d.get("country") or ""
        name = (d.get("name") or "").strip()
        if (
            skip_benelux_duplicates
            and ctry == "Benelux"
            and name.casefold() in skip_names
        ):
            # Still count in tables; omit duplicate pin
            lat = lon = place = None
            geo_label = "Skipped (shown under Belgium / Netherlands / Luxembourg)"
        else:
            geo = d.get("geo")
            if geo and geo.get("lat") is not None:
                lat, lon = geo["lat"], geo["lon"]
                place = geo.get("place")
                geo_label = geo.get("geo_label") or ""
                if geo.get("map_exclude"):
                    # Keep geo_label for transparency; do not draw pin on Europe map
                    lat = lon = None
                    if "Overseas" not in geo_label:
                        geo_label = "Overseas territory (hidden on Europe map)"
            else:
                lat, lon, place, geo_label = place_for_dealer(
                    name,
                    ctry,
                    address=d.get("address"),
                    allow_country_centre=allow_country_centre,
                )
                if lat is not None and _is_overseas_pin(
                    place=place, lat=lat, lon=lon, country=ctry
                ):
                    if _should_map_exclude_overseas(
                        place=place, lat=lat, lon=lon
                    ):
                        lat = lon = None
                        geo_label = "Overseas territory (hidden on Europe map)"
                    else:
                        place = _overseas_display_place(place, lat=lat, lon=lon) or place
                        geo_label = "Overseas territory (French DOM-TOM)"

        pct = d.get("total_pct")
        cavg = d.get("country_avg_pct")
        if cavg is None:
            cavg = country_avg_pct
        pct_disp = round(pct * 100, 1) if pct is not None else None
        cavg_disp = round(cavg * 100, 1) if cavg is not None else None
        delta_pp = (
            round((pct - cavg) * 100, 1)
            if pct is not None and cavg is not None
            else None
        )
        points.append(
            {
                "id": d.get("id"),
                "rank": i,
                "name": d.get("name"),
                "country": ctry,
                "pct": pct_disp,
                "country_avg": cavg_disp,
                "delta_pp": delta_pp,
                "lat": lat,
                "lon": lon,
                "place": place,
                "geo_label": geo_label,
            }
        )
    return points


def geo_coverage_summary(dealers: list[dict[str, Any]]) -> dict[str, Any]:
    """Per-country pin coverage for build logs."""
    by: dict[str, dict[str, int]] = {}
    for d in dealers:
        c = d.get("country") or "?"
        slot = by.setdefault(c, {"n": 0, "pinned": 0, "city": 0, "country_fb": 0})
        slot["n"] += 1
        g = d.get("geo")
        if g and g.get("lat") is not None:
            slot["pinned"] += 1
            if "country centre" in (g.get("geo_label") or ""):
                slot["country_fb"] += 1
            else:
                slot["city"] += 1
    return by


def sanity_check_geo_clustering(
    dealers: list[dict[str, Any]],
    *,
    radius_km: float = _CLUSTER_RADIUS_KM,
    share_threshold: float = _CLUSTER_SHARE_FAIL,
    min_dealers: int = _MIN_DEALERS_FOR_GEO_FAIL,
) -> dict[str, Any]:
    """Detect country-centre blobs / over-clustered pins.

    Returns:
      {
        "ok": bool,  # False when any hard fail
        "errors": [str, ...],
        "warnings": [str, ...],
        "countries": {country: {n, pinned, city, country_fb, cluster_max, cluster_share, unique_places}},
      }
    """
    by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for d in dealers:
        by[d.get("country") or "?"].append(d)

    errors: list[str] = []
    warnings: list[str] = []
    countries: dict[str, dict[str, Any]] = {}

    for country in sorted(by.keys()):
        rows = by[country]
        pinned = [
            d
            for d in rows
            if (d.get("geo") or {}).get("lat") is not None
            and (d.get("geo") or {}).get("lon") is not None
        ]
        city_n = sum(
            1
            for d in pinned
            if "country centre" not in ((d.get("geo") or {}).get("geo_label") or "")
        )
        fb_n = len(pinned) - city_n
        places = {
            ((d.get("geo") or {}).get("place") or "").strip()
            for d in pinned
            if (d.get("geo") or {}).get("place")
        }
        places.discard("")

        cluster_max = 0
        coords = [
            (float(d["geo"]["lat"]), float(d["geo"]["lon"]))  # type: ignore[index]
            for d in pinned
        ]
        n = len(coords)
        if n:
            for i, (la, lo) in enumerate(coords):
                near = sum(
                    1
                    for j, (lb, lob) in enumerate(coords)
                    if _haversine_km(la, lo, lb, lob) <= radius_km
                )
                if near > cluster_max:
                    cluster_max = near
        cluster_share = (cluster_max / n) if n else 0.0

        slot = {
            "n": len(rows),
            "pinned": n,
            "city": city_n,
            "country_fb": fb_n,
            "unique_places": len(places),
            "cluster_max": cluster_max,
            "cluster_share": round(cluster_share, 3),
            "cluster_radius_km": radius_km,
        }
        countries[country] = slot

        if n < min_dealers:
            continue

        fb_share = fb_n / n
        city_share = (city_n / n) if n else 0.0
        has_address = any((d.get("address") or "").strip() for d in rows)
        centre = _COUNTRY_CENTRE.get(country)
        near_centre = 0
        if centre and coords:
            clat, clon, _ = centre
            near_centre = sum(
                1 for la, lo in coords if _haversine_km(la, lo, clat, clon) <= 40.0
            )
        near_share = (near_centre / n) if n else 0.0
        blob = fb_share >= share_threshold and near_share >= share_threshold

        # Out-of-bounds on map = hard fail for non-overseas pins.
        # Verified French DOM-TOM may sit on the map (Réunion / Antilles / Mayotte);
        # Europe fitBounds still ignores them. Mystery far pins must stay map_exclude.
        oob_on_map = 0
        for d in pinned:
            g = d.get("geo") or {}
            if g.get("map_exclude"):
                continue
            la, lo = float(g["lat"]), float(g["lon"])
            ctry = d.get("country") or country
            midi_raw = (midi_lookup(d.get("name") or "") or {}).get("midi_country")
            if _is_verified_dom_tom(
                place=g.get("place"), lat=la, lon=lo, midi_country=midi_raw
            ):
                continue
            if _is_overseas_pin(
                place=g.get("place"), lat=la, lon=lo, country=ctry, midi_country=midi_raw
            ) or (
                ctry in _COUNTRY_BOUNDS and not _coords_in_country_bounds(la, lo, ctry)
            ):
                # Açores / Madeira / Canaries allowed as pins but rare
                if ctry == "Portugal" and -32 <= lo <= -6 and 32 <= la <= 40:
                    continue
                if ctry == "Spain" and -19 <= lo <= -13 and 27 <= la <= 30:
                    continue
                oob_on_map += 1
        if oob_on_map:
            errors.append(
                f"{country}: {oob_on_map} map pin(s) outside country/Europe bounds "
                f"(ocean / wrong continent) — fix tokens or map_exclude"
            )

        floor = _EXPECTED_CITY_FLOOR.get(country)
        if floor is not None and city_share < floor:
            errors.append(
                f"{country}: city placement {city_n}/{n} ({city_share:.0%}) "
                f"below expected floor {floor:.0%} — geo regression"
            )
        elif blob and has_address and city_share < 0.15:
            # Addresses present but almost nothing mapped — loud warn (expand city tokens)
            warnings.append(
                f"{country}: {fb_n}/{n} country-centre pins "
                f"({near_centre} within 40 km of centre) despite Address data — "
                f"add city tokens from addresses"
            )
        elif blob:
            warnings.append(
                f"{country}: {fb_n}/{n} pins ({fb_share:.0%}) are country-centre "
                f"fallback clustered near centre — add address/city tokens"
            )
        elif fb_share >= share_threshold:
            warnings.append(
                f"{country}: {fb_n}/{n} pins ({fb_share:.0%}) use country-centre "
                f"fallback — prefer city from address/name"
            )
        elif cluster_share >= share_threshold and len(places) <= 2:
            warnings.append(
                f"{country}: {cluster_max}/{n} pins ({cluster_share:.0%}) within "
                f"{radius_km:.0f} km across {len(places)} place label(s) — review"
            )

    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "countries": countries,
    }


def format_geo_sanity_report(result: dict[str, Any]) -> str:
    """Human-readable multi-line report for build logs."""
    lines = ["Geo sanity check (cluster / country-centre):"]
    countries: dict[str, dict[str, Any]] = result.get("countries") or {}
    for c, s in countries.items():
        lines.append(
            f"  {c}: n={s['n']} pinned={s['pinned']} "
            f"city={s['city']} centre_fb={s['country_fb']} "
            f"places={s['unique_places']} "
            f"cluster={s['cluster_max']}/{s['pinned'] or 0} "
            f"({s['cluster_share']:.0%} <={s['cluster_radius_km']:.0f}km)"
        )
    for w in result.get("warnings") or []:
        lines.append(f"  WARN: {w}")
    for e in result.get("errors") or []:
        lines.append(f"  FAIL: {e}")
    if result.get("ok"):
        lines.append("  RESULT: PASS")
    else:
        lines.append("  RESULT: FAIL")
    return "\n".join(lines)
