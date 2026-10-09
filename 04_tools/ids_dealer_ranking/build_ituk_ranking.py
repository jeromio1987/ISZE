#!/usr/bin/env python3
"""Build ITUK dealer ranking from consolidated IDS ITUK self evaluations workbook.

Sheet2 has two dealer blocks on the same tab:
  Block 1 — names row 1,  percentage row 23, points row 24
  Block 2 — names row 27, percentage row 49, points row 50

Skip dealers with 0 points (did not participate). Duplicate names keep the non-zero score.

Section rows A.1–H.1 sit between the name row and the total-% row. Letter groups
(A–H) are rolled up for the UK CV dashboard (score_basis UKCV-AH). This is not
the continental A–E / 280-point book — do not overlay network A–E on these scores.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from build_country_ranking import _coerce_percent, write_ranking_workbook

from ids_paths import ITUK_INPUT, ITUK_OUTPUT, country_label

DEFAULT_INPUT = ITUK_INPUT
DEFAULT_OUTPUT = ITUK_OUTPUT

# name_row, percent_row, points_row
ITUK_BLOCKS = (
    (1, 23, 24),
    (27, 49, 50),
)

SUBSEC_RE = re.compile(r"^([A-H])\.(\d+)\s*$", re.I)

# Sheet2 letter rollups — English labels for shared UK CV HTML
UKCV_SECTION_DEFS: list[dict[str, str]] = [
    {"id": "A", "label": "A. Premises, stationery and uniforms"},
    {"id": "B", "label": "B. Organisation and management"},
    {"id": "C", "label": "C. Workshop and service operation"},
    {"id": "D", "label": "D. Parts facility and operations"},
    {"id": "E", "label": "E. New vehicle sales"},
    {"id": "F", "label": "F. Marketing"},
    {"id": "G", "label": "G. Standards review"},
    {"id": "H", "label": "H. Key performance indicators"},
]
UKCV_SECTION_IDS = tuple(s["id"] for s in UKCV_SECTION_DEFS)
UKCV_SCORE_BASIS = "UKCV-AH"


def _resolve_sheet(wb, sheet_arg: str | None):
    if sheet_arg:
        if sheet_arg.isdigit():
            return wb.worksheets[int(sheet_arg)]
        return wb[sheet_arg]
    if "Sheet2" in wb.sheetnames:
        return wb["Sheet2"]
    return wb.worksheets[1]


def _section_rows_for_block(ws, *, name_row: int, percent_row: int) -> list[tuple[int, str, float]]:
    """Return (sheet_row, letter, max_pts) for each A.1–H.1 line in the block."""
    out: list[tuple[int, str, float]] = []
    for r in range(name_row + 1, percent_row):
        ref = ws.cell(r, 1).value
        if not ref:
            continue
        m = SUBSEC_RE.match(str(ref).strip())
        if not m:
            continue
        letter = m.group(1).upper()
        max_raw = ws.cell(r, 3).value
        try:
            max_pts = float(max_raw) if max_raw is not None else 0.0
        except (TypeError, ValueError):
            max_pts = 0.0
        out.append((r, letter, max_pts))
    return out


def _roll_letter_sections(
    ws,
    *,
    col: int,
    section_rows: list[tuple[int, str, float]],
) -> dict[str, dict]:
    """Sum Sheet2 sub-rows into A–H {score, max, pct} blocks."""
    score_sum: dict[str, float] = {sid: 0.0 for sid in UKCV_SECTION_IDS}
    max_sum: dict[str, float] = {sid: 0.0 for sid in UKCV_SECTION_IDS}
    seen: dict[str, bool] = {sid: False for sid in UKCV_SECTION_IDS}

    for r, letter, max_pts in section_rows:
        if letter not in score_sum:
            continue
        raw = ws.cell(r, col).value
        try:
            pts = float(raw) if raw is not None else 0.0
        except (TypeError, ValueError):
            pts = 0.0
        # Zero-max lines (e.g. E.1 on Parts & Service block) do not contribute.
        if max_pts <= 0:
            continue
        score_sum[letter] += pts
        max_sum[letter] += max_pts
        seen[letter] = True

    sections: dict[str, dict] = {}
    for sid in UKCV_SECTION_IDS:
        mx = max_sum[sid]
        if not seen[sid] or mx <= 0:
            sections[sid] = {
                "score": None,
                "max": mx if mx > 0 else None,
                "pct": None,
                "fill": "not_scored",
                "country_avg_score": None,
                "country_avg_pct": None,
            }
            continue
        sc = score_sum[sid]
        sections[sid] = {
            "score": sc,
            "max": mx,
            "pct": sc / mx,
            "fill": "scored",
            "country_avg_score": None,
            "country_avg_pct": None,
        }
    return sections


def _read_block(
    ws,
    *,
    name_row: int,
    percent_row: int,
    points_row: int,
    first_dealer_col: int = 5,
    source_file: str,
) -> list[dict]:
    max_pts = ws.cell(points_row, 3).value
    if not isinstance(max_pts, (int, float)) or max_pts <= 0:
        max_pts = 234

    section_rows = _section_rows_for_block(ws, name_row=name_row, percent_row=percent_row)

    rows: list[dict] = []
    col = first_dealer_col
    while col < ws.max_column:
        name = ws.cell(name_row, col).value
        if not name or not str(name).strip():
            col += 1
            continue
        name = str(name).strip()
        if name.startswith(",") or "Section" in name or name == "Parts & Service":
            col += 1
            continue

        points_raw = ws.cell(points_row, col).value
        pct_raw = ws.cell(percent_row, col + 1).value

        try:
            points_val = float(points_raw) if points_raw is not None else 0.0
        except (TypeError, ValueError):
            points_val = 0.0

        if points_val == 0:
            col += 2
            continue

        pct = _coerce_percent(pct_raw)
        if pct is None and points_val > 0:
            pct = round(points_val / float(max_pts) * 100, 2)
        items_pct = round(points_val / float(max_pts) * 100, 2) if points_val > 0 else None

        rows.append(
            {
                "name": name,
                "points": pct,
                "items": items_pct,
                "raw_points": points_val,
                "max_pts": float(max_pts),
                "file": source_file,
                "error": None,
                "sections": _roll_letter_sections(
                    ws, col=col, section_rows=section_rows
                ),
                "score_basis": UKCV_SCORE_BASIS,
            }
        )
        col += 2

    return rows


def _merge_dealer_rows(blocks: list[list[dict]]) -> list[dict]:
    """One row per dealer; prefer higher assessed points when name appears twice."""
    by_name: dict[str, dict] = {}
    for block_rows in blocks:
        for row in block_rows:
            key = row["name"].casefold()
            existing = by_name.get(key)
            if existing is None or (row["items"] or 0) > (existing["items"] or 0):
                by_name[key] = row

    merged = list(by_name.values())
    merged.sort(
        key=lambda r: (-(r["points"] or -1), -(r["items"] or -1), r["name"].lower()),
    )
    return merged


def read_ituk_dealers(path: Path, *, sheet_name: str | None = None) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = _resolve_sheet(wb, sheet_name)

    block_rows = [
        _read_block(
            ws,
            name_row=name_row,
            percent_row=percent_row,
            points_row=points_row,
            source_file=path.name,
        )
        for name_row, percent_row, points_row in ITUK_BLOCKS
    ]
    wb.close()
    return _merge_dealer_rows(block_rows)


def ituk_sections_by_name(path: Path | None = None) -> dict[str, dict]:
    """casefold(name) → dealer row with rolled A–H sections (from Sheet2)."""
    src = path or ITUK_INPUT
    if not src.is_file():
        return {}
    return {r["name"].casefold(): r for r in read_ituk_dealers(src)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sheet", default=None, help="Sheet name or index (default: Sheet2)")
    args = parser.parse_args()

    if not args.input.is_file():
        print(f"Input not found: {args.input}", file=sys.stderr)
        return 1

    rows = read_ituk_dealers(args.input, sheet_name=args.sheet)
    if not rows:
        print("No participating dealers found (all points were 0).", file=sys.stderr)
        return 2

    write_ranking_workbook(
        country_label("ITUK"),
        rows,
        args.output.resolve(),
    )

    print(f"Wrote {args.output}")
    print(f"  dealers ranked: {len(rows)}")
    print(f"  top: {rows[0]['name']} ({rows[0]['points']}%)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
