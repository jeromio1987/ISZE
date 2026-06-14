"""One-off extractor for dealer standards audit."""
from ids_paths import RSA_TEMPLATE
from openpyxl import load_workbook
from collections import Counter

SAMPLE = RSA_TEMPLATE


def main():
    wb = load_workbook(SAMPLE, data_only=True)
    ev = wb["Evaluation Sheet_English"]
    # find header row
    header_row = None
    for r in range(1, 15):
        vals = [str(ev.cell(r, c).value or "").lower() for c in range(1, 17)]
        if "question" in " ".join(vals):
            header_row = r
            print("HEADER ROW", r, [ev.cell(r, c).value for c in range(1, 17)])
            break

    questions = []
    for r in range((header_row or 4) + 1, ev.max_row + 1):
        no = ev.cell(r, 1).value
        upd = ev.cell(r, 2).value
        q = ev.cell(r, 3).value
        if not q or not isinstance(q, str) or len(q) < 20:
            continue
        if str(q).strip().lower() in ("question", "total"):
            continue
        row = {
            "no": no,
            "update": upd,
            "question": q.strip(),
            "priority": ev.cell(r, 4).value,
            "method": ev.cell(r, 5).value,
            "point": ev.cell(r, 6).value,
            "self_eval": ev.cell(r, 11).value if ev.max_column >= 11 else None,
        }
        questions.append(row)

    print(f"\nTotal questions: {len(questions)}")
    print("Priorities:", Counter(q["priority"] for q in questions))
    print("Methods:", Counter(q["method"] for q in questions))
    print("Points:", Counter(q["point"] for q in questions))
    print("\n--- First 8 ---")
    for q in questions[:8]:
        print(f"{q['no']} | P{q['priority']} | {q['point']}pt | {q['method']}")
        print(f"  {q['question'][:200]}")
    print("\n--- Long / compound (>=140 chars) ---")
    for q in questions:
        if len(q["question"]) >= 140:
            print(f"{q['no']}: {q['question'][:180]}...")

    cov = wb["Cover Page"]
    print("\n--- Cover key text ---")
    for r in range(1, 25):
        for c in range(1, 10):
            v = cov.cell(r, c).value
            if v and isinstance(v, str) and len(v) > 15:
                print(f"  [{r},{c}] {v[:120]}")

    sc = wb["Score Card"]
    print("\n--- Score Card headers ---")
    for r in range(1, 6):
        print(r, [sc.cell(r, c).value for c in range(1, 14)])

    wb.close()


if __name__ == "__main__":
    main()
