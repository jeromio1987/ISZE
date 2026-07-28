#!/usr/bin/env python3
"""Combine per-country dealer ranking workbooks into one multi-tab Excel file."""

from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ids_paths import RANKINGS_DIR
OUTPUT = RANKINGS_DIR / "IDS_2026_all_country_rankings.xlsx"

# (source filename, English tab name) — main total ranking sheet only
SOURCES = (
    ("BE_dealer_ranking.xlsx", "Belgium"),
    ("NL_dealer_ranking.xlsx", "Netherlands"),
    ("LUX_dealer_ranking.xlsx", "Luxembourg"),
    ("Denmark_dealer_ranking.xlsx", "Denmark"),
    ("Poland_dealer_ranking.xlsx", "Poland"),
    ("Petros_dealer_ranking.xlsx", "Greece"),
    ("RSA_dealer_ranking.xlsx", "Norway"),
    ("IUK_dealer_ranking.xlsx", "UK CV"),
    ("ITUK_dealer_ranking.xlsx", "UK LCV"),
    ("France_dealer_ranking.xlsx", "France"),
    ("Italy_dealer_ranking.xlsx", "Italy"),
    ("Switzerland_dealer_ranking.xlsx", "Switzerland"),
    ("Spain_dealer_ranking.xlsx", "Spain"),
    ("MIDI_FR_dealer_ranking.xlsx", "France"),
    ("MIDI_IT_dealer_ranking.xlsx", "Italy"),
    ("MIDI_CH_dealer_ranking.xlsx", "Switzerland"),
    ("MIDI_ES_dealer_ranking.xlsx", "Spain"),
    ("Israel_dealer_ranking.xlsx", "Israel LCV"),
    ("Romania_dealer_ranking.xlsx", "Romania"),
    ("Portugal_dealer_ranking.xlsx", "Portugal"),
    ("Hungary_dealer_ranking.xlsx", "Hungary"),
    ("Malta_dealer_ranking.xlsx", "Malta"),
    ("Cyprus_dealer_ranking.xlsx", "Cyprus"),
)


def _copy_sheet(source_ws, target_ws) -> None:
    for row in source_ws.iter_rows():
        for cell in row:
            t = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                t.font = copy(cell.font)
                t.border = copy(cell.border)
                t.fill = copy(cell.fill)
                t.number_format = copy(cell.number_format)
                t.protection = copy(cell.protection)
                t.alignment = copy(cell.alignment)
    for col, dim in source_ws.column_dimensions.items():
        target_ws.column_dimensions[col].width = dim.width
    for row, dim in source_ws.row_dimensions.items():
        target_ws.row_dimensions[row].height = dim.height
    target_ws.freeze_panes = source_ws.freeze_panes


def _add_index_sheet(wb, summary: list[tuple[str, str, int]]) -> None:
    ws = wb.create_sheet("Index", 0)
    ws.append(["Tab", "Source file", "Dealers ranked"])
    for c in range(1, 4):
        ws.cell(1, c).font = Font(bold=True)
    for tab, src, n in summary:
        ws.append([tab, src, n])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16


def main() -> int:
    RANKINGS_DIR.mkdir(parents=True, exist_ok=True)
    out_wb = Workbook()
    out_wb.remove(out_wb.active)

    summary: list[tuple[str, str, int]] = []
    missing: list[str] = []

    for filename, tab_name in SOURCES:
        path = RANKINGS_DIR / filename
        if not path.is_file():
            missing.append(filename)
            continue

        src_wb = load_workbook(path, data_only=False)
        src_name = tab_name if tab_name in src_wb.sheetnames else src_wb.sheetnames[0]
        src_ws = src_wb[src_name]

        if tab_name in out_wb.sheetnames:
            continue
        tgt_ws = out_wb.create_sheet(tab_name)
        _copy_sheet(src_ws, tgt_ws)

        dealers = max(0, tgt_ws.max_row - 1)
        summary.append((tab_name, filename, dealers))
        src_wb.close()

    if not summary:
        print("No ranking files found.", file=sys.stderr)
        return 1

    _add_index_sheet(out_wb, summary)
    out_wb.save(OUTPUT)

    print(f"Wrote {OUTPUT}")
    print(f"  tabs: Index + {', '.join(t for t, _, _ in summary)}")
    if missing:
        print("  missing:", ", ".join(missing), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
