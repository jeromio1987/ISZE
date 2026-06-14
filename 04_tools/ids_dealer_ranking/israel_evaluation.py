"""Read totals from Israel UMI IDS templates (Evaluation Sheet_English)."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

TOP_SECTION_RE = re.compile(r"^[A-F]\. [A-Za-z]")
B1_SECTION_RE = re.compile(r"^B\.1\s+[A-Za-z]", re.IGNORECASE)


def find_israel_eval_sheet(sheetnames: list[str]) -> str | None:
    for name in sheetnames:
        low = name.strip().lower()
        if "evaluation sheet" in low and "english" in low:
            return name
    for name in sheetnames:
        low = name.strip().lower()
        if "formular" in low and "evaluare" in low:
            return name
    for name in sheetnames:
        low = name.strip().lower()
        if "evaluation" in low and "comparison" not in low:
            return name
    for name in sheetnames:
        low = name.strip().lower()
        if "evaluare" in low and "standarde" in low:
            return name
    return None


def _float_cell(value) -> float:
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def read_metrics_from_ws(ws) -> tuple[float | None, float | None]:
    """Overall points % and items % (same on this template)."""
    max_total = 0.0
    assessed_total = 0.0
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str):
            continue
        if not TOP_SECTION_RE.match(label.strip()):
            continue
        max_total += _float_cell(row[6] if len(row) > 6 else None)
        assessed_total += _float_cell(row[11] if len(row) > 11 else None)

    if max_total <= 0:
        return None, None
    pct = round(assessed_total / max_total * 100, 2)
    return pct, pct


def read_b1_from_ws(ws) -> tuple[float | None, float | None, float | None, float | None]:
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str):
            continue
        if not B1_SECTION_RE.match(label.strip()):
            continue
        max_val = _float_cell(row[6] if len(row) > 6 else None)
        assessed_val = _float_cell(row[11] if len(row) > 11 else None)
        if max_val <= 0:
            return None, None, assessed_val, max_val
        pct = round(assessed_val / max_val * 100, 2)
        return pct, pct, assessed_val, max_val
    return None, None, None, None


def read_israel_total(path: Path) -> tuple[float | None, float | None]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = find_israel_eval_sheet(wb.sheetnames)
    if not sheet:
        wb.close()
        raise ValueError(f"no evaluation sheet in {path.name}")
    pts, items = read_metrics_from_ws(wb[sheet])
    wb.close()
    return pts, items


def read_israel_b1(path: Path) -> tuple[float | None, float | None, float | None, float | None]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = find_israel_eval_sheet(wb.sheetnames)
    if not sheet:
        wb.close()
        raise ValueError(f"no evaluation sheet in {path.name}")
    result = read_b1_from_ws(wb[sheet])
    wb.close()
    if result[0] is None and result[1] is None:
        raise ValueError(f"B.1 section row not found on evaluation sheet in {path.name}")
    return result
