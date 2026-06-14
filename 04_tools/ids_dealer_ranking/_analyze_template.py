"""Analyze evaluation template for duplicates and structure."""
from ids_paths import RSA_TEMPLATE
from openpyxl import load_workbook
from collections import defaultdict

f = RSA_TEMPLATE
wb = load_workbook(f, data_only=True)
ev = wb["Evaluation Sheet_English"]

questions = []
for r in range(5, ev.max_row + 1):
    no = ev.cell(r, 2).value
    q = ev.cell(r, 4).value
    pt = ev.cell(r, 7).value
    if q and isinstance(q, str) and len(q) > 20 and isinstance(pt, (int, float)):
        questions.append((r, str(no), q.strip(), pt))

# normalize for duplicate detection
def norm(q):
    return " ".join(q.lower().split())[:120]

by_norm = defaultdict(list)
for item in questions:
    by_norm[norm(item[2])].append(item)

dups = {k: v for k, v in by_norm.items() if len(v) > 1}
print("total questions", len(questions))
print("duplicate groups", len(dups))
for k, items in list(dups.items())[:15]:
    print("---")
    for row, no, q, pt in items:
        print(f"  row {row} {no} pts={pt} | {q[:90]}")

wb.close()
