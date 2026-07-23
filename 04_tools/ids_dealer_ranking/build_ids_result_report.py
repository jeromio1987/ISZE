#!/usr/bin/env python3
"""Build IDS Evaluation Result Report (wide matrix) per country — DK .xls clone.

Clones ``2026 IDS Evaluation Result Report_DK.xls`` so layout, merges, and colours
match exactly. Dealers start column D; column C is max-points only (never dealers).

Outputs (under ``90_INBOX/IDS_2026/reports/evaluation_result/``):
  - ``2026 IDS Evaluation Result Report_{CC}.xls``  — **DK layout, colours, merges** (open this)
  - ``_delta/2026 IDS Evaluation Result Report_{CC}_Delta.xlsx``  — plain numbers only (optional)

DK template: ``IDS_2026/reports/_template/2026 IDS Evaluation Result Report_DK.xls``
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import Any

import xlrd
from openpyxl import Workbook, load_workbook
from xlutils.copy import copy
from xlwt import Formula
from xlwt.Formatting import Protection
from xlwt.Style import XFStyle

from build_country_ranking import clean_dealer_name, dealer_display_name, dealer_name_from_filename
from ids_paths import DATA_ROOT, resolve_dk_template, resolve_input_dir, result_report_delta_path, result_report_path

# 0-based column indices — match DK .xls (A=0, B=1, C=max pts, D=first dealer)
COL_LABEL = 0
COL_TEXT = 1
COL_MAX = 2
COL_FIRST_DEALER = 3
ROW_COUNTRY = 1
ROW_META_START = 2
ROW_SCORE_ORIG = 12
ROW_SCORE_PCT = 13
ROW_DATA_START = 14
ROW_AVG_LABEL = 11  # Excel row 12 — DK: "Total Average" / "Ave. %"
AVG_HEADER_SCORE = "Total Average"
AVG_HEADER_PCT = "Ave. %"

# xlwt style cache keyed by (workbook id, xf_idx)
_XFSTYLE_CACHE: dict[tuple[int, int], XFStyle] = {}


def _xf_idx_to_xfstyle(wb: Any, xf_idx: int) -> XFStyle:
    """Rebuild xlwt XFStyle from copied workbook so writes keep DK row colours."""
    key = (id(wb), xf_idx)
    cached = _XFSTYLE_CACHE.get(key)
    if cached is not None:
        return cached
    sc = wb._Workbook__styles  # xlwt internal
    xf_tuple = sc._xf_x2id.get(xf_idx)
    if not xf_tuple:
        return wb.get_default_style()
    font_idx, _num_fmt_idx, alignment, borders, pattern, protection = xf_tuple
    st = XFStyle()
    st.alignment = alignment
    st.borders = borders
    st.pattern = pattern
    st.protection = protection
    st.font = sc._font_x2id[font_idx]
    _XFSTYLE_CACHE[key] = st
    return st


_PERCENT_STYLE_CACHE: dict[int, XFStyle] = {}


def _percent_style(base: XFStyle, fmt: str = "0.0%") -> XFStyle:
    """Variant of a row style formatted as a percentage and left unlocked (user-editable)."""
    cached = _PERCENT_STYLE_CACHE.get(id(base))
    if cached is not None:
        return cached
    st = XFStyle()
    st.alignment = base.alignment
    st.borders = base.borders
    st.pattern = base.pattern
    st.font = base.font
    st.num_format_str = fmt
    st.protection = Protection()
    st.protection.cell_locked = False
    st.protection.formula_hidden = base.protection.formula_hidden
    _PERCENT_STYLE_CACHE[id(base)] = st
    return st


def _snapshot_row_styles(ws: Any, wb: Any, nrows: int) -> dict[int, XFStyle]:
    """Capture per-row XF from cloned template (cols A–D) before overwriting dealers."""
    styles: dict[int, XFStyle] = {}
    default = wb.get_default_style()
    for r in range(nrows):
        row_obj = ws._Worksheet__rows.get(r)
        if not row_obj:
            styles[r] = default
            continue
        found = default
        for c in (COL_FIRST_DEALER, COL_LABEL, COL_TEXT, COL_MAX):
            cell = row_obj._Row__cells.get(c)
            if cell is not None:
                found = _xf_idx_to_xfstyle(wb, cell.xf_idx)
                break
        styles[r] = found
    return styles


def _styled_write(
    ws: Any,
    row_styles: dict[int, XFStyle],
    r: int,
    c: int,
    val: Any,
    *,
    default: XFStyle,
) -> None:
    ws.write(r, c, val, row_styles.get(r, default))


def _styled_write_pct(
    ws: Any,
    row_styles: dict[int, XFStyle],
    r: int,
    c: int,
    val: Any,
    *,
    default: XFStyle,
) -> None:
    """Like _styled_write, but renders the value as a percentage (e.g. 73.0%) and unlocks the cell."""
    ws.write(r, c, val, _percent_style(row_styles.get(r, default)))


def _excel_col(col: int) -> str:
    return xlrd.colname(col)


def _dealer_range(r: int, n_dealers: int) -> str:
    """A1 range for dealer columns on one row (Excel 1-based row)."""
    first = _excel_col(COL_FIRST_DEALER)
    last = _excel_col(COL_FIRST_DEALER + n_dealers - 1)
    row = r + 1
    return f"{first}{row}:{last}{row}"


def _avg_score_formula(r: int, n_dealers: int) -> Formula:
    return Formula(f"AVERAGE({_dealer_range(r, n_dealers)})")


def _avg_pct_formula(r: int, avg_col: int) -> Formula:
    row = r + 1
    avg_cell = f"${_excel_col(avg_col)}${row}"
    max_cell = f"$C${row}"
    return Formula(f"IF({max_cell}=0,\"\",MIN({avg_cell}/{max_cell},1))")


def _child_question_sheet_rows(matrix_rows: list[TemplateRow], trow: TemplateRow) -> list[int]:
    rows: list[int] = []
    started = False
    for row in matrix_rows:
        if row.sheet_row == trow.sheet_row:
            started = True
            continue
        if not started:
            continue
        if row.kind in ("category", "subcategory"):
            break
        if row.kind == "question" and row.sheet_row is not None:
            rows.append(row.sheet_row)
    return rows


def _child_subcategory_sheet_rows(matrix_rows: list[TemplateRow], trow: TemplateRow) -> list[int]:
    rows: list[int] = []
    started = False
    for row in matrix_rows:
        if row.sheet_row == trow.sheet_row:
            started = True
            continue
        if not started:
            continue
        if row.kind == "category":
            break
        if row.kind == "subcategory" and row.sheet_row is not None:
            rows.append(row.sheet_row)
    return rows


def _sum_range_formula(dealer_col: int, sheet_rows: list[int]) -> Formula | None:
    """DK-style =SUM(D17:D38) for consecutive question rows under a subcategory."""
    if not sheet_rows:
        return None
    col = _excel_col(dealer_col)
    excel_rows = sorted(r + 1 for r in sheet_rows)
    if len(excel_rows) == 1:
        return Formula(f"{col}{excel_rows[0]}")
    if excel_rows[-1] - excel_rows[0] + 1 == len(excel_rows):
        return Formula(f"SUM({col}{excel_rows[0]}:{col}{excel_rows[-1]})")
    refs = ",".join(f"{col}{r}" for r in excel_rows)
    return Formula(f"SUM({refs})")


def _add_rows_formula(dealer_col: int, sheet_rows: list[int]) -> Formula | None:
    """DK-style =D16+D39+D44 for top-level category rows."""
    if not sheet_rows:
        return None
    col = _excel_col(dealer_col)
    parts = [f"{col}{r + 1}" for r in sorted(sheet_rows)]
    if len(parts) == 1:
        return Formula(parts[0])
    return Formula("+".join(parts))


def _section_formula(
    trow: TemplateRow, matrix_rows: list[TemplateRow], dealer_col: int
) -> Formula | None:
    if trow.kind == "subcategory":
        return _sum_range_formula(dealer_col, _child_question_sheet_rows(matrix_rows, trow))
    if trow.kind == "category":
        sub_rows = _child_subcategory_sheet_rows(matrix_rows, trow)
        if sub_rows:
            return _add_rows_formula(dealer_col, sub_rows)
        return _sum_range_formula(dealer_col, _child_question_sheet_rows(matrix_rows, trow))
    return None


def _category_sheet_rows(matrix_rows: list[TemplateRow]) -> list[int]:
    return [t.sheet_row for t in matrix_rows if t.kind == "category" and t.sheet_row is not None]

QUESTION_RE = re.compile(r"^[A-H][.\-]\d+[-.]\d+", re.IGNORECASE)
SUBCAT_RE = re.compile(r"^[A-H]\.\d+\s", re.IGNORECASE)
TOPCAT_RE = re.compile(r"^[A-H]\.\s", re.IGNORECASE)

META_ROWS = (
    "Dealer Name (*Write official company name for IDS certification)",
    "Evaluation Date",
    "Address",
    "Evaluator",
    "Contact Person of the Dealer and his position",
    "Available Models",
    "Other Vehicle Brand",
    "Opening Hours",
    "The number of Servicing Vehicle  per month",
)

COVER_LABELS: dict[str, tuple[str, ...]] = {
    "dealer_name": (
        "dealer name",
        "naam dealer",
        "nom du concessionnaire",
        "nom du revendeur",
        "nom dealer",
    ),
    "evaluation_date": ("evaluation date", "evaluatie datum", "date d'évaluation", "date d'evaluation"),
    "address": ("address", "adres", "adresse"),
    "evaluator": ("evaluator", "beoordelaar", "évaluateur", "evaluateur"),
    "contact": (
        "contact person",
        "verkoop personeel",
        "personne de contact",
        "contact person of the dealer",
    ),
    "models": ("available models", "beschikbare modellen", "modèles disponibles", "modeles disponibles"),
    "other_brands": ("other vehicle brand", "andere automerken", "autres marques"),
    "hours": ("opening hours", "openingsuren", "heures d'ouverture"),
    "servicing": (
        "servicing vehicle",
        "number of servicing",
        "aantal voertuigen",
        "véhicules entretenus",
    ),
}


@dataclass
class TemplateRow:
    kind: str  # meta_total | category | subcategory | question | other
    ref: str
    text: str
    max_points: float | None
    question_ids: list[str] = field(default_factory=list)
    sheet_row: int | None = None


@dataclass
class StaffRow:
    name: str = ""
    title: str = ""
    nationality: str = ""
    age: str = ""
    isuzu_years: str = ""
    mechanic_years: str = ""


@dataclass
class DealerData:
    name: str
    file: str
    meta: dict[str, Any]
    question_scores: dict[str, float]
    mechanics: list[StaffRow] = field(default_factory=list)
    advisors: list[StaffRow] = field(default_factory=list)
    total_override: float | None = None
    # 0–1 fraction when set (ITUK consolidated: do not derive as total/280).
    total_pct_override: float | None = None
    error: str | None = None


def _norm_ref(ref: str) -> str:
    ref = str(ref).strip().replace("–", "-").replace("—", "-")
    m = re.match(r"^([A-H])[-.](\d+)[-.](\d+)$", ref, re.I)
    if m:
        return f"{m.group(1).upper()}.{m.group(2)}-{m.group(3)}"
    return ref


def _classify_ref(ref: str) -> str:
    ref = str(ref).strip()
    if QUESTION_RE.match(ref) or re.match(r"^[A-H][-.]\d+[-.]\d+", ref, re.I):
        return "question"
    if SUBCAT_RE.match(ref):
        return "subcategory"
    if TOPCAT_RE.match(ref):
        return "category"
    return "other"


def load_dk_template(path: Path | None = None, *, sheet: str = "Score") -> list[TemplateRow]:
    if path is None:
        path = resolve_dk_template()
    wb = xlrd.open_workbook(str(path))
    sh = wb.sheet_by_name(sheet)
    rows: list[TemplateRow] = []
    current_questions: list[str] = []

    def flush_questions(target: TemplateRow) -> None:
        target.question_ids = list(current_questions)

    for r in range(14, sh.nrows):
        ref = str(sh.cell_value(r, 0)).strip()
        text = str(sh.cell_value(r, 1) or "").strip()
        mx = sh.cell_value(r, 2)
        max_pts = float(mx) if isinstance(mx, (int, float)) and mx != "" else None
        if not ref and max_pts == 280:
            rows.append(TemplateRow("meta_total", "", "", 280.0))
            continue
        if not ref and not text and max_pts is None:
            continue
        kind = _classify_ref(ref)
        row = TemplateRow(kind, ref, text, max_pts)
        if kind == "question":
            qid = _norm_ref(ref)
            current_questions.append(qid)
            row.question_ids = [qid]
        elif kind in ("category", "subcategory"):
            current_questions = []
            row.question_ids = []
        rows.append(row)

    # assign question ranges to category/subcategory rows
    questions: list[tuple[str, float]] = []
    for row in rows:
        if row.kind == "question" and row.ref:
            questions.append((_norm_ref(row.ref), row.max_points or 0))

    # rebuild hierarchy: walk rows, accumulate question ids until next same/higher level
    idx = 0
    while idx < len(rows):
        row = rows[idx]
        if row.kind in ("category", "subcategory"):
            level = 0 if row.kind == "category" else 1
            qids: list[str] = []
            j = idx + 1
            while j < len(rows):
                nxt = rows[j]
                if nxt.kind == "category":
                    break
                if level == 1 and nxt.kind == "subcategory":
                    break
                if nxt.kind == "question":
                    qids.extend(nxt.question_ids)
                j += 1
            row.question_ids = qids
        idx += 1
    return rows


def find_eval_sheet(sheetnames: list[str]) -> str | None:
    skip = ("comparison", "vergelijk", "comparaison", "score card", "scorecard", "scorekaart", "eredmény", "eredmeny", "fiche de r")
    patterns = (
        "evaluation sheet",
        "evaluatieformulier",
        "feuille d",
        "ficha de aval",
        "értékelő",
        "ertekelo",
        "formular de evaluare",
        "evaluare",
    )
    for name in sheetnames:
        low = name.strip().lower()
        if any(x in low for x in skip):
            continue
        if any(k in low for k in patterns):
            return name
    for name in sheetnames:
        low = name.strip().lower()
        if any(x in low for x in skip):
            continue
        if "evaluat" in low or "evaluare" in low or "avalia" in low:
            return name
    return None


def find_cover_sheet(sheetnames: list[str]) -> str | None:
    for name in sheetnames:
        low = name.lower()
        if any(k in low for k in ("cover page", "voorpagina", "page de couverture", "fedlap", "capa")):
            return name
    return None


def _dealer_total(d: DealerData) -> float:
    if d.total_override is not None:
        return d.total_override
    return sum(d.question_scores.values())


def _dealer_ok(d: DealerData) -> bool:
    return not d.error and (bool(d.question_scores) or d.total_override is not None)


def _hungarian_rubric_score(row: tuple) -> float | None:
    """HU template: max col I (8); dealer score col N (13); J–M are fixed rubric labels.

    Rubric cells often contain static ``nem``/``igen``/option text identical across
    dealers — do not treat those as the self-eval mark. Prefer the numeric score
    in column N (same value the Eredmény scorecard uses).
    """
    if len(row) < 9:
        return None
    mx = row[8]
    if not isinstance(mx, (int, float)):
        return None
    max_pt = float(mx)
    if len(row) > 13 and isinstance(row[13], (int, float)):
        return min(float(row[13]), max_pt)
    # Fallback when N is blank: Point 0–3 rubric mark (rare on current HU files).
    yes = {"igen", "yes", "ja", "oui", "igen."}
    for rubric_pts, cell in zip((0, 1, 2, 3), row[9:13]):
        if cell is None:
            continue
        if str(cell).strip().lower() in yes:
            return float(rubric_pts) if rubric_pts <= max_pt else max_pt
    return _coerce_score(row[12] if len(row) > 12 else None, max_pt)


_YES_MARKS = frozenset({"yes", "ja", "oui", "da", "sim", "y", "igen"})


def _rubric_mark_score(row: tuple, max_pt: float) -> float | None:
    """Infer points from Point 0/1/2/3 rubric columns (H–K) when self-eval cell is wrong."""
    for rubric_pts, cell in zip((0, 1, 2, 3), row[7:11] if len(row) > 10 else ()):
        if cell is None:
            continue
        s = str(cell).strip().lower()
        if s in _YES_MARKS:
            return float(min(rubric_pts, max_pt))
    return None


def _standard_eval_score(row: tuple, max_pt: float) -> float | None:
    """Self-eval col L, with rubric fallback and cap at max points."""
    raw = _coerce_score(row[11] if len(row) > 11 else None, max_pt)
    rubric = _rubric_mark_score(row, max_pt)
    if raw is not None and raw <= max_pt:
        return raw
    if rubric is not None:
        return rubric
    if raw is not None:
        return min(raw, max_pt)
    return None


def _parse_eval_rows(rows_iter, *, hungarian: bool = False) -> dict[str, float]:
    scores: dict[str, float] = {}
    for row in rows_iter:
        if not row or len(row) < 8:
            continue
        no = row[1] if len(row) > 1 else None
        if not no:
            continue
        ref = str(no).strip()
        if not (QUESTION_RE.match(ref) or re.match(r"^[A-H][-.]\d+[-.]\d+", ref, re.I)):
            continue
        if hungarian:
            sc = _hungarian_rubric_score(row)
        else:
            mx = row[6] if len(row) > 6 else None
            if not isinstance(mx, (int, float)):
                continue
            sc = _standard_eval_score(row, float(mx))
        if sc is not None:
            scores[_norm_ref(ref)] = sc
    return scores


def parse_dealer_workbook(path: Path) -> DealerData:
    name = dealer_display_name(path)
    sheetnames, eval_sheet, rows_iter, open_err = _open_eval_rows(path)
    if open_err:
        return DealerData(name=name, file=path.name, meta={}, question_scores={}, error=open_err)

    meta: dict[str, Any] = {}
    mechanics: list[StaffRow] = []
    advisors: list[StaffRow] = []

    if path.suffix.lower() != ".xls":
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            cover = find_cover_sheet(wb.sheetnames)
            if cover:
                cws = wb[cover]
                meta = _read_cover_field_map(cws)
                cleaned = clean_dealer_name(
                    str(meta["dealer_name"]) if meta.get("dealer_name") is not None else None
                )
                if cleaned:
                    name = cleaned
                    meta["dealer_name"] = cleaned
                else:
                    # Cover still has the certificate placeholder — use filename stem.
                    name = dealer_name_from_filename(path)
                    meta["dealer_name"] = name
                mechanics = _parse_staff_block(cws, "mechanic")
                if not mechanics:
                    mechanics = _parse_staff_block(cws, "monteur")
                advisors = _parse_staff_block(cws, "service advisor")
                if not advisors:
                    advisors = _parse_staff_block(cws, "service adviseur")
            wb.close()
        except Exception:
            pass

    hungarian = eval_sheet and (
        "értékel" in eval_sheet.lower() or "ertekel" in eval_sheet.lower()
    )
    scores = _parse_eval_rows(rows_iter(), hungarian=bool(hungarian))

    if not scores:
        return DealerData(
            name=name,
            file=path.name,
            meta=meta,
            question_scores={},
            mechanics=mechanics,
            advisors=advisors,
            error="no question scores parsed",
        )
    return DealerData(
        name=name,
        file=path.name,
        meta=meta,
        question_scores=scores,
        mechanics=mechanics,
        advisors=advisors,
    )


def _open_eval_rows(path: Path):
    """Open eval sheet rows. Returns (sheetnames, eval_sheet, rows_iter, error)."""
    if path.suffix.lower() == ".xls":
        try:
            rb = xlrd.open_workbook(str(path))
        except Exception as exc:
            return [], None, None, str(exc)
        eval_sheet = find_eval_sheet(rb.sheet_names())
        if not eval_sheet:
            return rb.sheet_names(), None, None, f"no evaluation sheet in {rb.sheet_names()}"
        sh = rb.sheet_by_name(eval_sheet)

        def rows():
            for r in range(4, sh.nrows):
                yield tuple(sh.cell_value(r, c) for c in range(sh.ncols))

        return rb.sheet_names(), eval_sheet, rows, None

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return [], None, None, str(exc)
    eval_sheet = find_eval_sheet(wb.sheetnames)
    if not eval_sheet:
        names = list(wb.sheetnames)
        wb.close()
        return names, None, None, f"no evaluation sheet in {names}"
    ws = wb[eval_sheet]

    def rows():
        try:
            for row in ws.iter_rows(min_row=5, values_only=True):
                yield row
        finally:
            wb.close()

    return wb.sheetnames, eval_sheet, rows, None


def parse_ituk_dealers(path: Path) -> list[DealerData]:
    """ITUK consolidated workbook — totals only (no per-question breakdown)."""
    from build_ituk_ranking import ITUK_BLOCKS, _merge_dealer_rows, _read_block

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb.worksheets[1]
        block_rows = [
            _read_block(
                ws,
                name_row=nr,
                percent_row=pr,
                points_row=pt,
                source_file=path.name,
            )
            for nr, pr, pt in ITUK_BLOCKS
        ]
        wb.close()
    except Exception as exc:
        return [DealerData("ITUK", path.name, {}, {}, error=str(exc))]

    dealers: list[DealerData] = []
    for row in _merge_dealer_rows(block_rows):
        raw = row.get("raw_points")
        if raw is None or float(raw) <= 0:
            continue
        # Prefer sheet %; else points/max. Store as 0–1 for Score (pct) row.
        pct_100 = row.get("points")
        if pct_100 is None:
            pct_100 = row.get("items")
        pct_frac = (float(pct_100) / 100.0) if pct_100 is not None else None
        dealers.append(
            DealerData(
                name=row["name"],
                file=path.name,
                meta={},
                question_scores={},
                total_override=float(raw),
                total_pct_override=pct_frac,
            )
        )
    return dealers


def _coerce_score(value: Any, max_pt: float) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return max_pt if value else 0.0
    if isinstance(value, (int, float)):
        v = float(value)
        # Only values strictly between 0 and 1 are %-of-max (e.g. 0.5 on max 2 → 1 pt).
        # Integer 1 on max 3 is one point, not 100% — fixes NL/BE Zelfevaluatie column.
        if 0 < v < 1 and max_pt > 1:
            return round(v * max_pt, 4)
        return v
    s = str(value).strip().upper()
    if s in {"YES", "JA", "OUI", "DA", "SIM", "Y"}:
        return max_pt
    if s in {"NO", "NEE", "NON", "NU", "N"}:
        return 0.0
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def _label_match(label: str, keys: tuple[str, ...]) -> bool:
    low = label.strip().lower()
    return any(k in low for k in keys)


def _read_cover_field_map(ws) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=1, max_row=80, values_only=True):
        if not row:
            continue
        for col in range(len(row) - 1):
            label = row[col]
            if not isinstance(label, str):
                continue
            val = row[col + 1] if col + 1 < len(row) else None
            # also check label col B value col E pattern
            if _label_match(label, COVER_LABELS["dealer_name"]):
                out["dealer_name"] = val
            elif _label_match(label, COVER_LABELS["evaluation_date"]):
                out["evaluation_date"] = val
            elif _label_match(label, COVER_LABELS["address"]):
                out["address"] = val
            elif _label_match(label, COVER_LABELS["evaluator"]):
                out["evaluator"] = val
            elif _label_match(label, COVER_LABELS["contact"]):
                out["contact"] = val
            elif _label_match(label, COVER_LABELS["models"]):
                out["models"] = val
            elif _label_match(label, COVER_LABELS["other_brands"]):
                out["other_brands"] = val
            elif _label_match(label, COVER_LABELS["hours"]):
                out["hours"] = val
            elif _label_match(label, COVER_LABELS["servicing"]):
                out["servicing"] = val
        # B/E layout: label col 2, value col 5
        if len(row) > 4:
            label = row[1]
            val = row[4]
            if isinstance(label, str):
                if _label_match(label, COVER_LABELS["dealer_name"]):
                    out["dealer_name"] = val
                elif _label_match(label, COVER_LABELS["evaluation_date"]):
                    out["evaluation_date"] = val
                elif _label_match(label, COVER_LABELS["address"]):
                    out["address"] = val
                elif _label_match(label, COVER_LABELS["evaluator"]):
                    out["evaluator"] = val
                elif _label_match(label, COVER_LABELS["contact"]):
                    out["contact"] = val
                elif _label_match(label, COVER_LABELS["models"]):
                    out["models"] = val
                elif _label_match(label, COVER_LABELS["other_brands"]):
                    out["other_brands"] = val
                elif _label_match(label, COVER_LABELS["hours"]):
                    out["hours"] = val
                elif _label_match(label, COVER_LABELS["servicing"]):
                    out["servicing"] = val
    return out


def _parse_staff_block(ws, start_label: str) -> list[StaffRow]:
    """Best-effort mechanic / advisor table on cover page."""
    rows: list[StaffRow] = []
    in_block = False
    header_map: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=120, values_only=True):
        cells = list(row)
        texts = [str(c).strip() if c is not None else "" for c in cells]
        joined = " ".join(texts).lower()
        if start_label.lower() in joined:
            in_block = True
            header_map = {}
            continue
        if not in_block:
            continue
        if any(
            x in joined
            for x in (
                "serviceadviseur",
                "service advisor",
                "basisinformatie",
                "evaluation information",
            )
        ) and start_label.lower() not in joined:
            break
        # header row with Name / Title
        if "name" in joined and ("title" in joined or "functie" in joined or "titre" in joined):
            for i, t in enumerate(texts):
                tl = t.lower()
                if t in {"Name", "Naam", "Nom"}:
                    header_map["name"] = i
                elif "title" in tl or "functie" in tl or "titre" in tl:
                    header_map["title"] = i
                elif "national" in tl:
                    header_map["nationality"] = i
                elif tl == "age" or "leeftijd" in tl:
                    header_map["age"] = i
                elif "isuzu" in tl:
                    header_map["isuzu_years"] = i
                elif "mechanic" in tl or "monteur" in tl:
                    header_map["mechanic_years"] = i
            continue
        if not header_map:
            continue
        name = texts[header_map.get("name", 0)] if header_map else ""
        if not name or name.lower() in {"name", "naam", "nr."}:
            continue
        if re.fullmatch(r"\d+", name):
            continue
        rows.append(
            StaffRow(
                name=name,
                title=texts[header_map["title"]] if "title" in header_map else "",
                nationality=texts[header_map["nationality"]] if "nationality" in header_map else "",
                age=texts[header_map["age"]] if "age" in header_map else "",
                isuzu_years=texts[header_map["isuzu_years"]] if "isuzu_years" in header_map else "",
                mechanic_years=texts[header_map["mechanic_years"]] if "mechanic_years" in header_map else "",
            )
        )
    return rows


def load_matrix_rows_from_sheet(sh) -> list[TemplateRow]:
    """Row specs from an existing DK-layout main sheet (DK26 / BE26)."""
    rows: list[TemplateRow] = []
    for r in range(ROW_DATA_START, sh.nrows):
        ref = str(sh.cell_value(r, COL_LABEL)).strip()
        text = str(sh.cell_value(r, COL_TEXT) or "").strip()
        mx = sh.cell_value(r, COL_MAX)
        max_pts = float(mx) if isinstance(mx, (int, float)) and mx != "" else None
        if not ref and not text and max_pts is None:
            continue
        kind = _classify_ref(ref) if ref else "other"
        rows.append(TemplateRow(kind, ref, text, max_pts, sheet_row=r))
    # attach question id lists to category rows
    idx = 0
    while idx < len(rows):
        row = rows[idx]
        if row.kind in ("category", "subcategory"):
            level = 0 if row.kind == "category" else 1
            qids: list[str] = []
            j = idx + 1
            while j < len(rows):
                nxt = rows[j]
                if nxt.kind == "category":
                    break
                if level == 1 and nxt.kind == "subcategory":
                    break
                if nxt.kind == "question":
                    qids.append(_norm_ref(nxt.ref))
                j += 1
            row.question_ids = qids
        elif row.kind == "question":
            row.question_ids = [_norm_ref(row.ref)]
        idx += 1
    return rows


def _row_value(dealer: DealerData, trow: TemplateRow) -> float | None:
    if trow.kind == "question":
        qid = _norm_ref(trow.ref)
        return dealer.question_scores.get(qid)
    if trow.kind in ("category", "subcategory"):
        vals = [dealer.question_scores.get(q) for q in trow.question_ids]
        vals = [v for v in vals if v is not None]
        return sum(vals) if vals else None
    if trow.kind == "meta_total":
        return _dealer_total(dealer)
    return None


def _pct(score: float | None, max_pts: float | None) -> float | None:
    if score is None or not max_pts:
        return None
    return score / max_pts


def _fmt_meta(val: Any) -> Any:
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.isoformat()
    return val


def _write_staff_sheet(ws, country: str, dealers: list[DealerData], attr: str) -> None:
    """Populate Mechanic Info / Service Advisor Info (headers already in DK template)."""
    row = 2  # 0-based: row 3 is header in template; data from row 3 (index 2) after header
    # DK template row index 1 = headers, data starts index 2
    row = 2
    for d in dealers:
        staff_list = getattr(d, attr)
        if staff_list:
            for s in staff_list:
                ws.write(row, 0, country.upper())
                ws.write(row, 1, d.name)
                ws.write(row, 2, s.name)
                ws.write(row, 3, s.title)
                ws.write(row, 4, s.nationality)
                ws.write(row, 5, s.age)
                ws.write(row, 6, s.isuzu_years)
                ws.write(row, 7, s.mechanic_years)
                row += 1
        else:
            ws.write(row, 0, country.upper())
            ws.write(row, 1, d.name)
            row += 1


def build_delta_workbook(
    country: str,
    dealers: list[DealerData],
    matrix_rows: list[TemplateRow],
    averages: dict[int, float],
    output: Path,
) -> None:
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = f"{country}26_Delta"[:31]
    ws["A1"] = "Delta vs country average (dealer score - country avg)"
    ws["A2"] = f"Country  {country.upper()}"
    ws["A1"].font = Font(bold=True)
    ws["A2"].font = Font(bold=True)

    # Usability: keep refs/labels visible while scrolling dealers.
    ws.freeze_panes = "D4"

    # Column sizing: make this readable without manual formatting.
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 2

    wrap_left = Alignment(wrap_text=True, vertical="top")
    first_col = 4
    for i, d in enumerate(dealers):
        ws.cell(row=3, column=first_col + i, value=d.name)
    for trow in matrix_rows:
        r = trow.sheet_row + 1 if trow.sheet_row is not None else None
        if r is None:
            continue
        avg = averages.get(trow.sheet_row)
        if avg is None:
            continue
        ws.cell(row=r, column=1, value=trow.ref)
        c2 = ws.cell(row=r, column=2, value=trow.text)
        c2.alignment = wrap_left
        for i, d in enumerate(dealers):
            val = _row_value(d, trow)
            if val is not None:
                ws.cell(row=r, column=first_col + i, value=round(float(val) - float(avg), 4))

    # Dealer columns: consistent width for readability
    for i in range(len(dealers)):
        col_letter = ws.cell(row=3, column=first_col + i).column_letter
        ws.column_dimensions[col_letter].width = 18
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def build_report(
    country: str,
    dealers: list[DealerData],
    template_path: Path,
    output: Path,
) -> None:
    ok = [d for d in dealers if _dealer_ok(d)]
    ok.sort(key=lambda d: d.name.lower())
    if not ok:
        raise ValueError("no valid dealers to include")

    rb = xlrd.open_workbook(str(template_path), formatting_info=True)
    sh_t = rb.sheet_by_index(0)
    matrix_rows = load_matrix_rows_from_sheet(sh_t)

    n = len(ok)
    avg_col = COL_FIRST_DEALER + n
    avg_pct_col = avg_col + 1
    clear_to = max(sh_t.ncols, avg_pct_col + 1)

    wb = copy(rb)
    main = wb.get_sheet(0)
    try:
        main.name = f"{country}26"[:31]
    except Exception:
        pass

    row_styles = _snapshot_row_styles(main, wb, sh_t.nrows)
    default_style = wb.get_default_style()

    # Country label (keep DK merges / styles in cols A-C)
    main.write(ROW_COUNTRY, COL_LABEL, f"Country  {country.upper()}")

    # Clear old dealer + average columns only (never col C except existing template maxima)
    for r in range(sh_t.nrows):
        for c in range(COL_FIRST_DEALER, clear_to):
            _styled_write(main, row_styles, r, c, "", default=default_style)

    # Dealer index numbers row 2
    for i in range(n):
        _styled_write(main, row_styles, ROW_COUNTRY, COL_FIRST_DEALER + i, i + 1, default=default_style)

    meta_keys = (
        "dealer_name",
        "evaluation_date",
        "address",
        "evaluator",
        "contact",
        "models",
        "other_brands",
        "hours",
        "servicing",
    )
    for mi, key in enumerate(meta_keys):
        r = ROW_META_START + mi
        for i, d in enumerate(ok):
            val = d.name if key == "dealer_name" else d.meta.get(key)
            _styled_write(main, row_styles, r, COL_FIRST_DEALER + i, _fmt_meta(val), default=default_style)

    totals = [_dealer_total(d) for d in ok]
    use_detail_formulas = any(d.question_scores for d in ok)
    # Always write cached numeric section/total scores (not only Excel formulas).
    # Reviewers opening via xlrd / non-Excel tools otherwise see blanks vs self-eval
    # (Romania_tb discrepancy: B.5 / B.7 subcategory rows).

    for i, total in enumerate(totals):
        _styled_write(
            main, row_styles, ROW_SCORE_ORIG, COL_FIRST_DEALER + i, round(total, 4), default=default_style
        )
        pct_override = ok[i].total_pct_override
        if pct_override is not None:
            _styled_write_pct(
                main, row_styles, ROW_SCORE_PCT, COL_FIRST_DEALER + i, round(float(pct_override), 6), default=default_style
            )
        else:
            _styled_write_pct(
                main, row_styles, ROW_SCORE_PCT, COL_FIRST_DEALER + i, round(total / 280, 6), default=default_style
            )

    # DK row 12: average column headers (row 2 has no avg labels)
    _styled_write(main, row_styles, ROW_AVG_LABEL, avg_col, AVG_HEADER_SCORE, default=default_style)
    _styled_write(main, row_styles, ROW_AVG_LABEL, avg_pct_col, AVG_HEADER_PCT, default=default_style)

    avg_total = mean(totals)
    _styled_write(
        main, row_styles, ROW_SCORE_ORIG, avg_col, _avg_score_formula(ROW_SCORE_ORIG, n), default=default_style
    )
    _styled_write_pct(
        main, row_styles, ROW_SCORE_ORIG, avg_pct_col, _avg_pct_formula(ROW_SCORE_ORIG, avg_col), default=default_style
    )

    averages: dict[int, float] = {
        ROW_SCORE_ORIG: avg_total,
    }

    for trow in matrix_rows:
        r = trow.sheet_row
        if r is None:
            continue

        if trow.kind in ("category", "subcategory") and use_detail_formulas:
            for i, d in enumerate(ok):
                val = _row_value(d, trow)
                if val is not None:
                    if trow.max_points:
                        val = min(float(val), float(trow.max_points))
                    _styled_write(
                        main, row_styles, r, COL_FIRST_DEALER + i, round(float(val), 4), default=default_style
                    )
            vals = [v for v in (_row_value(d, trow) for d in ok) if v is not None]
            if vals:
                averages[r] = mean(vals)
            _styled_write(main, row_styles, r, avg_col, _avg_score_formula(r, n), default=default_style)
            if trow.max_points:
                _styled_write_pct(
                    main,
                    row_styles,
                    r,
                    avg_pct_col,
                    _avg_pct_formula(r, avg_col),
                    default=default_style,
                )
            continue

        vals: list[float] = []
        for i, d in enumerate(ok):
            val = _row_value(d, trow)
            if val is not None:
                if trow.max_points:
                    val = min(float(val), float(trow.max_points))
                _styled_write(main, row_styles, r, COL_FIRST_DEALER + i, round(val, 4), default=default_style)
                vals.append(float(val))
        if vals:
            avg = mean(vals)
            averages[r] = avg
            _styled_write(main, row_styles, r, avg_col, _avg_score_formula(r, n), default=default_style)
            if trow.max_points:
                _styled_write_pct(
                    main,
                    row_styles,
                    r,
                    avg_pct_col,
                    _avg_pct_formula(r, avg_col),
                    default=default_style,
                )

    # Staff sheets (indices 1 and 2 in DK master)
    _write_staff_sheet(wb.get_sheet(1), country, ok, "mechanics")
    _write_staff_sheet(wb.get_sheet(2), country, ok, "advisors")
    # Sheet 3 = Score reference — left unchanged from DK template

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))

    delta_path = result_report_delta_path(country)
    build_delta_workbook(country, ok, matrix_rows, averages, delta_path)


def parse_midi_dealers(cod_div: str, path: Path | None = None) -> list[DealerData]:
    """Build DealerData rows from MIDI_FR.xlsx DLR SELF EVALUATION for one CodDiv."""
    from midi_fr_ranking import MIDI_FILE, load_network_names

    midi_path = path or MIDI_FILE
    if not midi_path.is_file():
        return [DealerData("MIDI", str(midi_path), {}, {}, error=f"missing {midi_path}")]

    code = cod_div.strip().upper()
    names = load_network_names(midi_path)

    max_by_ref: dict[str, float] = {}
    wb = load_workbook(midi_path, read_only=True, data_only=True)
    if "Evaluation Sheet_English" in wb.sheetnames:
        for row in wb["Evaluation Sheet_English"].iter_rows(min_row=4, values_only=True):
            no = row[1] if len(row) > 1 else None
            if not no:
                continue
            ref = str(no).strip()
            if not QUESTION_RE.match(ref) and not re.match(r"^[A-H][-.]\d+[-.]\d+", ref, re.I):
                continue
            mx = row[6] if len(row) > 6 else None
            if isinstance(mx, (int, float)):
                max_by_ref[_norm_ref(ref)] = float(mx)

    scores_by_ana: dict[str, dict[str, float]] = {}
    for row in wb["DLR SELF EVALUATION"].iter_rows(min_row=8, values_only=True):
        if not row or len(row) < 7:
            continue
        if str(row[2] or "").strip().upper() != code:
            continue
        ref_raw, ana = row[1], row[3]
        if not ref_raw or not ana:
            continue
        ref = _norm_ref(str(ref_raw))
        if not QUESTION_RE.match(ref) and not re.match(r"^[A-H]\.\d+-\d+", ref, re.I):
            continue
        try:
            val = float(row[6]) if row[6] is not None else 0.0
        except (TypeError, ValueError):
            val = 0.0
        mx = max_by_ref.get(ref)
        if mx is not None:
            val = min(val, mx)
        scores_by_ana.setdefault(str(ana).strip(), {})[ref] = val
    wb.close()

    dealers: list[DealerData] = []
    for ana, scores in sorted(
        scores_by_ana.items(), key=lambda kv: (names.get(kv[0]) or names.get(kv[0].lstrip("0")) or kv[0]).lower()
    ):
        name = names.get(ana) or names.get(ana.lstrip("0")) or ana
        dealers.append(
            DealerData(
                name=str(name).strip(),
                file=f"{midi_path.name}:{code}:{ana}",
                meta={"cod_ana": ana, "cod_div": code},
                question_scores=scores,
            )
        )
    return dealers


def collect_dealer_files(input_dir: Path) -> list[Path]:
    skip_names = ("met grafiek", "ranking", "dealers romania 2025")
    files: list[Path] = []
    for p in sorted(input_dir.rglob("*")):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".xlsx", ".xls"}:
            continue
        if p.name.startswith("~$"):
            continue
        low = p.name.lower()
        if "ranking" in low or any(s in low for s in skip_names):
            continue
        files.append(p)
    return files


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--country", default="BE", help="Country code/label (default BE)")
    parser.add_argument(
        "--input-dir",
        type=Path,
        help="Folder with dealer workbooks (default: IDS_2026/{country})",
    )
    parser.add_argument(
        "--extra-input-dir",
        type=Path,
        action="append",
        default=[],
        help="Additional folder(s) of dealer workbooks (repeatable; used for IBX Benelux)",
    )
    parser.add_argument(
        "--midi",
        metavar="COD",
        help="Build from MIDI_FR.xlsx for CodDiv (FR/IT/ES/CH) instead of per-dealer folders",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="Output .xls path",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=None,
        help="DK .xls template path (default: IDS_2026/reports/_template/)",
    )
    parser.add_argument("--qa", action="store_true", help="Run layout QA vs DK after build")
    args = parser.parse_args()

    country = args.country
    template_path = args.template or resolve_dk_template()
    if not template_path.is_file():
        print(f"DK template not found: {template_path}", file=sys.stderr)
        return 1

    if args.midi:
        dealers = parse_midi_dealers(args.midi)
    elif country.upper() == "ITUK":
        from ids_paths import ITUK_INPUT

        if not ITUK_INPUT.is_file():
            print(f"ITUK input not found: {ITUK_INPUT}", file=sys.stderr)
            return 1
        dealers = parse_ituk_dealers(ITUK_INPUT)
    else:
        input_dir = args.input_dir or resolve_input_dir(country, DATA_ROOT)
        dirs = [input_dir, *args.extra_input_dir]
        files: list[Path] = []
        for d in dirs:
            if not d.is_dir():
                print(f"Input directory not found: {d}", file=sys.stderr)
                return 1
            files.extend(collect_dealer_files(d))
        # de-dupe by resolved path
        seen: set[Path] = set()
        unique: list[Path] = []
        for f in files:
            key = f.resolve()
            if key in seen:
                continue
            seen.add(key)
            unique.append(f)
        if not unique:
            print(f"No workbooks in {dirs}", file=sys.stderr)
            return 1
        dealers = [parse_dealer_workbook(p) for p in unique]

    output = args.output or result_report_path(country)
    build_report(country, dealers, template_path, output.resolve())

    ok = sum(1 for d in dealers if _dealer_ok(d))
    bad = len(dealers) - ok
    print(f"Wrote {output}")
    print(f"  dealers included: {ok}")
    print(f"  delta (plain): {result_report_delta_path(country)}")
    if bad:
        print(f"  skipped/errors: {bad}")
        for d in dealers:
            if d.error:
                print(f"    {d.file}: {d.error}")

    if args.qa:
        from ids_report_qa import compare_to_dk

        result = compare_to_dk(output.resolve())
        print(f"  QA ok={result['ok']} issues={len(result['issues'])}")
        for issue in result["issues"][:15]:
            print(f"    - {issue}")
        if not result["ok"]:
            return 3
    return 0 if ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
