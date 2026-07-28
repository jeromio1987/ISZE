"""MIDI consolidated rankings from MIDI_FR.xlsx (multi-country DLR SELF EVALUATION)."""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from ids_paths import DATA_ROOT, MIDI_FILE, RANKINGS_DIR

MIDI_FILE = MIDI_FILE
RANKINGS_DIR = RANKINGS_DIR

MAX_POINTS_TOTAL = 280
QUESTION_REF = re.compile(r"^[A-H]\.\d", re.IGNORECASE)
B1_REF = re.compile(r"^B\.1", re.IGNORECASE)

# CodDiv column values in DLR SELF EVALUATION -> workbook tab / country label
MIDI_COUNTRY_LABELS: dict[str, str] = {
    "FR": "France",
    "IT": "Italy",
    "CH": "Switzerland",
    "ES": "Spain",
}


def midi_label(cod_div: str) -> str:
    code = cod_div.strip().upper()
    return MIDI_COUNTRY_LABELS.get(code, f"MIDI_{code}")


def discover_midi_cod_divs(path: Path = MIDI_FILE) -> list[str]:
    """Country codes present in DLR SELF EVALUATION (e.g. FR, IT, CH, ES)."""
    if not path.is_file():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["DLR SELF EVALUATION"]
    codes: set[str] = set()
    for row in ws.iter_rows(values_only=True):
        raw = row[2]
        if not raw:
            continue
        code = str(raw).strip().upper()
        if len(code) == 2 and code.isalpha() and code != "CO":
            codes.add(code)
    wb.close()
    order = {c: i for i, c in enumerate(MIDI_COUNTRY_LABELS)}
    return sorted(codes, key=lambda c: (order.get(c, 99), c))


def load_network_names(path: Path) -> dict[str, str]:
    """Map CUSTOMER_CODE (column M) -> DLRName."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["NETWORK"]
    mapping: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]
        code = row[12] if len(row) > 12 else None
        if name and code:
            mapping[str(code).strip()] = str(name).strip()
    wb.close()
    return mapping


def load_b1_max_points(path: Path) -> float:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Evaluation Sheet_English"]
    total = 0.0
    for row in ws.iter_rows(min_row=4, values_only=True):
        no = row[1] if len(row) > 1 else None
        if not no or not re.match(r"^B\.1-\d", str(no).strip(), re.IGNORECASE):
            continue
        try:
            total += float(row[6] if len(row) > 6 else 0 or 0)
        except (TypeError, ValueError):
            pass
    wb.close()
    return total or 12.0


def _resolve_name(names: dict[str, str], ana: str) -> str:
    name = names.get(ana) or names.get(ana.lstrip("0"))
    return name or ana


def _collect_midi_rows(
    path: Path,
    cod_div: str,
    *,
    question_rx: re.Pattern[str],
    max_points: float,
    include_zero_scores: bool = False,
) -> list[dict]:
    code = cod_div.strip().upper()
    names = load_network_names(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["DLR SELF EVALUATION"]

    stats: dict[str, dict] = defaultdict(
        lambda: {"points_sum": 0.0, "items_scored": 0, "items_total": 0}
    )

    for row in ws.iter_rows(values_only=True):
        if row[2] != code or not row[3]:
            continue
        ref = row[1]
        if not ref or ref == "ReferenceNoInIDS":
            continue
        if not question_rx.match(str(ref).strip()):
            continue

        ana = str(row[3]).strip()
        score = row[6]
        try:
            val = float(score) if score is not None else 0.0
        except (TypeError, ValueError):
            val = 0.0

        stats[ana]["items_total"] += 1
        if val > 0:
            stats[ana]["items_scored"] += 1
        stats[ana]["points_sum"] += val

    wb.close()

    rows: list[dict] = []
    for ana, s in stats.items():
        if s["items_total"] == 0:
            continue
        if not include_zero_scores and s["points_sum"] == 0:
            continue

        points_pct = round(s["points_sum"] / max_points * 100, 2)
        items_pct = round(s["items_scored"] / s["items_total"] * 100, 2)

        rows.append(
            {
                "name": _resolve_name(names, ana),
                "points": points_pct,
                "items": items_pct,
                "cod_ana": ana,
                "cod_div": code,
                "assessed_points": round(s["points_sum"], 2),
                "max_points": max_points,
                "items_count": s["items_total"],
                "file": path.name,
                "error": None,
            }
        )

    rows.sort(key=lambda r: (-(r["points"] or -1), -(r["items"] or -1), r["name"].lower()))
    return rows


def collect_midi_total_dealers(cod_div: str, path: Path = MIDI_FILE) -> list[dict]:
    # Keep zero-score shells so MIDI DLR count matches the source NETWORK /
    # self-eval set (Tomoe 27 Jul: count must not drop; scores may be 0).
    return _collect_midi_rows(
        path,
        cod_div,
        question_rx=QUESTION_REF,
        max_points=MAX_POINTS_TOTAL,
        include_zero_scores=True,
    )


def collect_midi_b1_dealers(cod_div: str, path: Path = MIDI_FILE) -> list[dict]:
    b1_max = load_b1_max_points(path)
    total_codes = {r["cod_ana"] for r in collect_midi_total_dealers(cod_div, path)}
    rows = _collect_midi_rows(
        path,
        cod_div,
        question_rx=B1_REF,
        max_points=b1_max,
        include_zero_scores=True,
    )
    by_ana = {r["cod_ana"]: r for r in rows}
    aligned: list[dict] = []
    for total_row in collect_midi_total_dealers(cod_div, path):
        ana = total_row["cod_ana"]
        if ana in by_ana:
            aligned.append(by_ana[ana])
        else:
            aligned.append(
                {
                    "name": total_row["name"],
                    "points": 0.0,
                    "items": 0.0,
                    "cod_ana": ana,
                    "cod_div": cod_div.strip().upper(),
                    "assessed_points": 0.0,
                    "max_points": b1_max,
                    "items_count": 0,
                    "file": path.name,
                    "error": None,
                }
            )
    aligned.sort(key=lambda r: (-(r["points"] or -1), -(r["items"] or -1), r["name"].lower()))
    return aligned


def collect_fr_total_dealers(path: Path = MIDI_FILE) -> list[dict]:
    return collect_midi_total_dealers("FR", path)


def collect_fr_b1_dealers(path: Path = MIDI_FILE) -> list[dict]:
    return collect_midi_b1_dealers("FR", path)


def total_output_path(cod_div: str) -> Path:
    return RANKINGS_DIR / f"{midi_label(cod_div)}_dealer_ranking.xlsx"


def b1_output_path(cod_div: str) -> Path:
    return RANKINGS_DIR / f"{midi_label(cod_div)}_B1_dealer_ranking.xlsx"


def write_total_workbook(rows: list[dict], output: Path, *, sheet_title: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers = [
        "Rank",
        "Name",
        "CodAna",
        "Points achieved %",
        "Items scored %",
        "Points (scored)",
        "Questions",
        "Source",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for rank, row in enumerate(rows, start=1):
        ws.append(
            [
                rank,
                row["name"],
                row["cod_ana"],
                row["points"],
                row["items"],
                row["assessed_points"],
                row["items_count"],
                row["file"],
            ]
        )

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 20

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def write_b1_workbook(rows: list[dict], output: Path, *, sheet_title: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers = [
        "Rank",
        "Name",
        "CodAna",
        "B.1 Points achieved %",
        "B.1 Items scored %",
        "B.1 Points (scored)",
        "B.1 Points (max)",
        "Source file",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for rank, row in enumerate(rows, start=1):
        ws.append(
            [
                rank,
                row["name"],
                row["cod_ana"],
                row["points"],
                row["items"],
                row["assessed_points"],
                row["max_points"],
                row["file"],
            ]
        )

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 20

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def build_all_midi_rankings(path: Path = MIDI_FILE) -> list[tuple[str, int, int]]:
    """Build total + B.1 workbooks for every CodDiv in the MIDI file. Returns summary tuples."""
    summary: list[tuple[str, int, int]] = []
    for cod in discover_midi_cod_divs(path):
        label = midi_label(cod)
        total_rows = collect_midi_total_dealers(cod, path)
        if total_rows:
            write_total_workbook(
                total_rows, total_output_path(cod), sheet_title=label
            )
        b1_rows = collect_midi_b1_dealers(cod, path)
        if b1_rows:
            write_b1_workbook(b1_rows, b1_output_path(cod), sheet_title=label)
        summary.append((label, len(total_rows), len(b1_rows)))
    return summary
