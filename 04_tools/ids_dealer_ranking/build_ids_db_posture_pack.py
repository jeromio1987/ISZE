#!/usr/bin/env python3
"""IDS DB Parts posture one-pagers + within-country D vs C CAP split.

Outputs under IDS_2026/reports/kaizen_db_posture/:
  - IDS_DB_Parts_Posture_2026.html
  - IDS_DB_Parts_Posture_Register.xlsx
  - study_payload.json
"""

from __future__ import annotations

import json
import math
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from build_ids_parts_kaizen_study import (
    COUNTRY_MAP,
    build_join_rows,
    correlations,
    ids_country_section_d,
    load_genuine_share_2025,
    load_gps_2025,
    load_kpi_fill_mos,
    load_pms_db_2025,
)
from ids_kaizen_hygiene import (
    dealers_for_stats,
    hygienic_network,
    is_rollup_country,
    section_pct,
)
from ids_park_age import attach_park_age, load_park_age_proxy
from ids_paths import DATA_ROOT, IDS_RESULT_REPORTS_DIR
from ids_report_extract import build_dashboard_payload

OUT_DIR = DATA_ROOT / "reports" / "kaizen_db_posture"
CHART_JS = Path(__file__).resolve().parents[2] / "09_meetings" / "Preparation Meetings" / "PMS_Top10" / "assets" / "chart.umd.min.js"
_TEMPLATE = Path(__file__).resolve().parent / "ids_db_posture_template.html"
_APP_JS = Path(__file__).resolve().parent / "ids_db_posture_app.js"

# Absolute floors used only when a country has too few dealers for quartile logic
ABS_WEAK = 0.65

GLOSSARY: dict[str, str] = {
    "pearson": (
        "Pearson r measures how two series move together across markets, from −1 to +1.\n"
        "+1 = perfect lockstep\n"
        "0 = no linear link\n"
        "−1 = perfect opposite\n\n"
        "Here each point is ONE distributor (not one dealer).\n"
        "So this is a cross-market prior — NOT a correlation inside IBX or IUK."
    ),
    "section_d": (
        "IDS Section D = Parts standards in the dealer self-evaluation "
        "(stock, counter, genuine display, targets, fill tracking, etc.).\n"
        "Shown as % of maximum points for that section.\n"
        "D avg = average across scored dealers in the country.\n"
        "It is a PROCESS / STANDARDS score — not € sales."
    ),
    "d_min": (
        "Lowest Section D % among scored dealers in that country.\n"
        "A big gap between D avg and D min = 'floor problem': "
        "the network average looks fine but a few dealers are very weak."
    ),
    "genuine_mix": (
        "Genuine parts share (€ mix) = Genuine € ÷ (Genuine € + BVP €).\n\n"
        "YES — it is how high genuine is compared with BVP inside that basket.\n"
        "Example: 80% means €80 genuine and €20 BVP out of €100 in that mix.\n\n"
        "It is NOT 'genuine vs all independent aftermarket'.\n"
        "It is NOT PMS capture %.\n"
        "High mix (e.g. 85–98%) = mostly genuine in the Genuine+BVP euro mix.\n"
        "Low mix (e.g. <45–50%) = BVP takes a large € share of that basket.\n\n"
        "Israel/UMI can show a very low mix — that is a mix/policy story, not fill."
    ),
    "pms_total": (
        "PMS total capture = actual parts sales ÷ expected wallet (all PMS lines).\n"
        "PRIMARY commercial lens — sales is sales (includes BVP/total basket).\n"
        "High IDS D does NOT automatically mean high PMS capture "
        "(in this study the cross-market link is even slightly negative)."
    ),
    "pms_genuine": (
        "PMS genuine capture = same idea as total capture, but only the genuine line type.\n"
        "Extra lens — do not replace total capture with this as the headline."
    ),
    "fill": (
        "KPI first fill rate = share of parts demand lines shipped from stock "
        "on first request (KPI Master).\n"
        "Higher is usually better ops (fewer backorders).\n"
        "Rough guide: well above ~90% looks strong; ~60% (e.g. Petros) is soft.\n"
        "Always read with MOS and park age — older shrinking parks often need more "
        "odd/slow-moving parts, which can drag fill even with good processes."
    ),
    "mos": (
        "MOS = Months Of Supply (days in KPI Master, averaged).\n"
        "Roughly: how many days of stock you are carrying vs sales pace.\n\n"
        "What is 'high' MOS?\n"
        "• Lean: often ~40–70 days (e.g. IBX ~47 in this extract)\n"
        "• Typical/elevated: ~90–120 days\n"
        "• High: ~140+ days (capital tied in stock; e.g. UMI ~147)\n"
        "• Extreme: 200–300+ days (e.g. some MIDI readings) — investigate\n\n"
        "High MOS is not 'good' by itself. Too high = overstock / slow movers. "
        "Too low with poor fill = stockouts. Target bands are market-specific."
    ),
    "gps_spu": (
        "GPS €/UIO (PRISM) = parts € intensity per Unit In Operation (parc vehicle).\n\n"
        "IMPORTANT — warranty: the headline figure used here is spend_per_unit_eur, "
        "which INCLUDES warranty parts (warranty is embedded in ERP genuine sales).\n"
        "An ex-warranty €/UIO also exists and is shown on cards as '€/UIO ex-w'.\n\n"
        "This is the better-positioned '€ per unit' metric vs sparse KPI Parts Sales per Unit.\n"
        "Different from genuine mix % and from IDS Section D.\n"
        "Traffic light (green/amber) is on the INCLUSIVE €/UIO unless labelled ex-w."
    ),
    "gps_spu_ex": (
        "GPS €/UIO excluding warranty parts.\n"
        "Headline GPS on cards still includes warranty; use this to see how much warranty "
        "inflates €/UIO (large gaps e.g. on UMI)."
    ),
    "gps_light": (
        "PRISM traffic light on GPS €/UIO intensity for that distributor "
        "(on the INCLUSIVE / with-warranty figure unless labelled ex-w).\n"
        "Green = intensity in a healthier band vs peers/targets; amber = soft. "
        "Not an IDS result."
    ),
    "park_age": (
        "Park ageing proxy from UIO.xlsx stock time series (not a true mean vehicle age).\n"
        "5y net park change % = (UIO_2024 − UIO_2019) / UIO_2024.\n"
        "Negative = shrinking parc → remaining fleet tends to be older "
        "(harder repairs, more odd parts, lower first-fill risk).\n"
        "Strongly positive = younger / expanding parc."
    ),
    "kpi_spu": (
        "KPI Master row: Parts Sales per Unit in € (€ parts ÷ unit base).\n"
        "Same FAMILY of idea as GPS €/UIO (intensity per vehicle), but sparsely filled.\n"
        "When blank: prefer GPS €/UIO — it is better positioned for this comparison.\n"
        "See §05. The ~0.98 correlation uses just 3 markets — ignore that bar."
    ),
    "cap_d_only": (
        "COUNT of dealers in this country tagged 'D-only' for Kaizen.\n\n"
        "Meaning: weak on IDS Parts (Section D), but Service (Section C) is OK.\n"
        "Example: CAP D-only = 4 → four dealers need a Parts/counter CAP, not a workshop CAP.\n\n"
        "This is NOT question D.1 or D.4 from the Excel — it is a headcount."
    ),
    "cap_d_plus_c": (
        "COUNT of dealers in this country tagged 'D+C' for Kaizen.\n\n"
        "Meaning: weak on BOTH Parts (D) and Service (C).\n"
        "Example: CAP D+C = 1 → one dealer needs a combined workshop + parts CAP "
        "(RO attach / pull-through), not only a counter fix.\n\n"
        "This is NOT an IDS question number — it is a headcount."
    ),
    "cap_c_only": (
        "COUNT of dealers weak on Service (C) only, while Parts (D) looks OK.\n"
        "Again a headcount — not an Excel question code."
    ),
    "benelux": (
        "Benelux result-report tab = BE + NL dealers combined.\n"
        "If we also keep BE and NL tabs, summing countries double-counts those dealers.\n"
        "Hygiene drops the Benelux rollup from network averages."
    ),
    "zero_shells": (
        "Dealers with 0% / blank totals — often empty MIDI self-eval shells still in the file.\n"
        "They pull averages down unfairly, so hygiene excludes them from scored stats."
    ),
    "raw_rows": "Every dealer row in the dashboard extract before hygiene filters.",
    "clean_rows": "Dealer rows after dropping Benelux rollup and 0% unscored shells.",
}


def build_kpi_spu_coverage(kpi: dict[str, dict[str, float | None]]) -> dict[str, Any]:
    """Who has Parts Sales per Unit filled in KPI Master vs IDS map."""
    ids_by_kpi: dict[str, list[str]] = {}
    for country, meta in COUNTRY_MAP.items():
        kn = meta.get("kpi")
        if not kn:
            continue
        ids_by_kpi.setdefault(str(kn), []).append(country)

    rows: list[dict[str, Any]] = []
    # Prefer IDS-mapped names first, then other KPI names with any parts metric
    names = sorted(set(kpi.keys()) | set(ids_by_kpi.keys()))
    for name in names:
        vals = kpi.get(name) or {}
        # skip empty ghosts
        if not any(vals.get(k) is not None for k in ("kpi_fill_rate", "kpi_mos_days", "kpi_spu_eur")):
            if name not in ids_by_kpi:
                continue
        rows.append(
            {
                "kpi_name": name,
                "ids_countries": ", ".join(ids_by_kpi.get(name, [])) or "— (not in IDS join map)",
                "kpi_fill_rate": vals.get("kpi_fill_rate"),
                "kpi_mos_days": vals.get("kpi_mos_days"),
                "kpi_spu_eur": vals.get("kpi_spu_eur"),
            }
        )

    # Sort: IDS-linked first, then SPU filled, then name
    rows.sort(
        key=lambda r: (
            0 if r["ids_countries"].startswith("—") else 0 if "—" not in r["ids_countries"][:1] else 1,
            0 if "—" not in str(r["ids_countries"]) else 1,
            0 if r.get("kpi_spu_eur") is not None else 1,
            str(r["kpi_name"]),
        )
    )

    filled = [r for r in rows if r.get("kpi_spu_eur") is not None and "—" not in r["ids_countries"]]
    blank = [r for r in rows if r.get("kpi_spu_eur") is None and "—" not in r["ids_countries"]]
    return {
        "rows": rows,
        "filled": [{"kpi_name": r["kpi_name"], "ids": r["ids_countries"]} for r in filled],
        "blank_on_ids_map": [{"kpi_name": r["kpi_name"], "ids": r["ids_countries"]} for r in blank],
    }


def _pct(v: Any) -> float | None:
    if v is None:
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    if x > 1.5:
        return round(x, 1)
    return round(x * 100, 1)


def classify_cap_type(d_pct: float, c_pct: float | None, d_cut: float, c_cut: float | None) -> str:
    """D-only = parts/counter process; D+C = workshop pull-through; other buckets for context."""
    d_weak = d_pct < d_cut
    c_weak = c_pct is not None and c_cut is not None and c_pct < c_cut
    if d_weak and c_weak:
        return "D+C"
    if d_weak and not c_weak:
        return "D-only"
    if (not d_weak) and c_weak:
        return "C-only"
    return "OK"


def within_country_cap_split(payload: dict[str, Any]) -> dict[str, Any]:
    """Per-country: dealers weak on D alone vs weak on D and C."""
    by_c: dict[str, list[dict[str, Any]]] = {}
    for d in dealers_for_stats(payload):
        if d.get("totals_only"):
            continue
        by_c.setdefault(str(d["country"]), []).append(d)

    countries: list[dict[str, Any]] = []
    examples: list[dict[str, Any]] = []
    totals = {"D-only": 0, "D+C": 0, "C-only": 0, "OK": 0, "unclassified": 0}

    for country, dealers in sorted(by_c.items()):
        if is_rollup_country(country):
            continue
        d_vals = [v for d in dealers if (v := section_pct(d, "D")) is not None]
        c_vals = [v for d in dealers if (v := section_pct(d, "C")) is not None]
        if len(d_vals) < 2:
            continue
        d_cut = median(d_vals)
        c_cut = median(c_vals) if c_vals else None
        # Prefer absolute floor when median is still high (network is strong)
        d_cut_eff = min(d_cut, ABS_WEAK) if d_cut is not None else ABS_WEAK
        c_cut_eff = min(c_cut, ABS_WEAK) if c_cut is not None else None

        counts = {"D-only": 0, "D+C": 0, "C-only": 0, "OK": 0}
        for d in dealers:
            dp = section_pct(d, "D")
            cp = section_pct(d, "C")
            if dp is None:
                totals["unclassified"] += 1
                continue
            kind = classify_cap_type(dp, cp, d_cut_eff, c_cut_eff)
            counts[kind] += 1
            totals[kind] += 1
            if kind in ("D-only", "D+C") and len([e for e in examples if e["country"] == country]) < 3:
                examples.append(
                    {
                        "country": country,
                        "name": d["name"],
                        "cap_type": kind,
                        "d_pct": dp,
                        "c_pct": cp,
                        "total_pct": d.get("total_pct"),
                        "ask": (
                            "Counter / inventory / genuine display gaps — Parts process CAP"
                            if kind == "D-only"
                            else "Parts + Service together — workshop pull-through & RO attach CAP"
                        ),
                    }
                )

        countries.append(
            {
                "country": country,
                "n": len(dealers),
                "d_cut": d_cut_eff,
                "c_cut": c_cut_eff,
                "counts": counts,
            }
        )

    return {"countries": countries, "examples": examples, "totals": totals}


def db_thesis(row: dict[str, Any]) -> dict[str, str]:
    """Short posture thesis + two DB asks — PMS excluded from conclusions."""
    c = row["ids_country"]
    d_avg = row.get("ids_d_avg") or 0
    d_min = row.get("ids_d_min")
    fill = row.get("kpi_fill_rate")
    mix = row.get("genuine_share_eur")
    gps = row.get("gps_spu_eur")
    light = (row.get("gps_light") or "").lower()
    ageing = row.get("ageing_risk")
    ch5 = row.get("park_change_5y_pct")

    spread = None
    if d_avg is not None and d_min is not None:
        spread = (d_avg - d_min) * 100

    archetype = "Balanced"
    thesis = (
        "IDS Parts standards and ops metrics are not loudly mismatched — "
        "still close floor dealers and watch park age vs fill."
    )
    asks = [
        "Which of your bottom-3 IDS Section D dealers also rank bottom-third for parts purchases?",
        "Can you share dealer-level first-fill or genuine-mix for those names (even rough)?",
    ]

    if c == "IUK" or (spread is not None and spread >= 35):
        archetype = "Floor problem"
        thesis = (
            "Network D average looks fine, but the floor is weak — Kaizen the outliers "
            "(counter process / targets), not a DB-wide Parts redesign."
        )
        asks = [
            "Confirm York Van / Adam (or your worst D names) rank for parts purchases vs network.",
            "Do they miss monthly parts targets and first-fill tracking (D.2-9 / D.2-10), or is it display/space only?",
        ]
    elif fill is not None and fill < 0.7:
        archetype = "Ops + standards"
        if ageing in {"high", "elevated"} and ch5 is not None and ch5 < 0:
            thesis = (
                f"Section D and first-fill are both soft, and the parc is shrinking "
                f"({ch5*100:.0f}% net over 5y) — older remaining vehicles often need odd/slow-moving parts, "
                f"which can drag fill. Treat as ops hygiene + age-aware stocking, not fill-blaming alone."
            )
            asks = [
                "How much of the fill gap is odd/obsolete parts for older vehicles vs process (ordering/min-max)?",
                "Which dealers drag fill most, and do they sit in older local fleets?",
            ]
        else:
            thesis = (
                "Section D and first-fill are both soft — treat as ops hygiene (fill, MOS, ordering) "
                "alongside IDS Excel gaps."
            )
            asks = [
                "What drives the low first-fill — stock policy, supplier lead time, or dealer ordering behaviour?",
                "Which dealers drag fill most, and do they match your weakest IDS Section D names?",
            ]
    elif mix is not None and mix < 0.45:
        archetype = "Mix gap"
        thesis = (
            "Genuine € mix is low relative to IDS D — validate whether Section D is paper compliance "
            "vs real counter mix (genuine default / BVP placement)."
        )
        asks = [
            "Is genuine the default recommendation at the counter, or is BVP / non-genuine pushed on price?",
            "Any dealer-level genuine share for the weakest IDS D dealers?",
        ]
    elif light == "amber" or (gps is not None and gps < 120):
        archetype = "Low €/UIO intensity"
        thesis = (
            "GPS €/UIO (incl. warranty) is soft while IDS may look OK — standards and spend intensity diverge. "
            "Check €/UIO ex-warranty too; use UIO for priority weight, not as an IDS causal story."
        )
        asks = [
            "Is low €/UIO age-mix, channel leakage, or under-worked parc?",
            "Which dealers cover the largest UIO with the weakest Parts standards?",
        ]
    elif ageing == "high" and fill is not None and fill < 0.9:
        archetype = "Ageing parc pressure"
        thesis = (
            f"Parc is shrinking ({(ch5 or 0)*100:.0f}% / 5y) — remaining fleet likely older. "
            "Even with middling IDS D, expect harder repairs and fill pressure; age-band stocking matters."
        )
        asks = [
            "Do you stock by age-band / slow-movers for older vehicles, or only by recent demand?",
            "Where does first-fill fail most — old VIN cohorts or fast movers?",
        ]

    if c == "ISRAEL":
        thesis += " IDS Israel = UMI only — never blend UTI."
        if mix is not None and mix < 0.25:
            archetype = "Mix gap (UMI)"

    if c in ("BE", "NL"):
        thesis += " IBX shares one GPS pool — dealer CAP on Degroote-class names."

    if c == "HUNGARY" and row.get("kpi_name") is None:
        thesis += " No KPI Master Parts row — posture relies on IDS + PRISM only."

    if c == "PETROS" and ageing in {"high", "elevated"}:
        thesis += (
            " Petros specifically: park decline supports the fill story — "
            "do not conclude fill is only dealer discipline."
        )

    return {"archetype": archetype, "thesis": thesis, "ask_1": asks[0], "ask_2": asks[1]}


def build_db_cards(join_rows: list[dict[str, Any]], cap_split: dict[str, Any]) -> list[dict[str, Any]]:
    cap_by_c = {r["country"]: r for r in cap_split.get("countries") or []}
    cards: list[dict[str, Any]] = []
    for row in join_rows:
        if is_rollup_country(row["ids_country"]):
            continue
        thesis = db_thesis(row)
        cap = cap_by_c.get(row["ids_country"], {})
        counts = cap.get("counts") or {}
        clean_row = {k: v for k, v in row.items() if not str(k).startswith("pms_")}
        cards.append(
            {
                **clean_row,
                **thesis,
                "cap_d_only": counts.get("D-only", 0),
                "cap_d_plus_c": counts.get("D+C", 0),
                "cap_c_only": counts.get("C-only", 0),
                "cap_ok": counts.get("OK", 0),
                "ids_d_avg_pct": _pct(row.get("ids_d_avg")),
                "ids_d_min_pct": _pct(row.get("ids_d_min")),
                "ids_total_avg_pct": _pct(row.get("ids_total_avg")),
                "genuine_share_pct": _pct(row.get("genuine_share_eur")),
                "fill_pct": _pct(row.get("kpi_fill_rate")),
                "kpi_mos_days": row.get("kpi_mos_days"),
                "park_change_5y_pct_disp": (
                    round(float(row["park_change_5y_pct"]) * 100, 1)
                    if row.get("park_change_5y_pct") is not None
                    else None
                ),
            }
        )
    order = {
        "Floor problem": 0,
        "Ops + standards": 1,
        "Ageing parc pressure": 2,
        "Mix gap": 3,
        "Mix gap (UMI)": 3,
        "Low €/UIO intensity": 4,
        "Balanced": 5,
    }
    cards.sort(key=lambda c: (order.get(c["archetype"], 9), -(c.get("ids_d_avg") or 0)))
    return cards


def build_park_fill_analysis(cards: list[dict[str, Any]]) -> dict[str, Any]:
    from build_ids_parts_kaizen_study import _pearson

    pairs = [
        c
        for c in cards
        if c.get("park_change_5y_pct") is not None and c.get("kpi_fill_rate") is not None
    ]
    r = None
    if len(pairs) >= 3:
        r = _pearson(
            [float(c["park_change_5y_pct"]) for c in pairs],
            [float(c["kpi_fill_rate"]) for c in pairs],
        )
    rows = []
    for c in sorted(
        cards,
        key=lambda x: (
            0 if x.get("ageing_risk") == "high" else 1 if x.get("ageing_risk") == "elevated" else 2,
            x.get("park_change_5y_pct") if x.get("park_change_5y_pct") is not None else 9,
        ),
    ):
        if c.get("park_change_5y_pct") is None and c.get("kpi_fill_rate") is None:
            continue
        note = ""
        if c.get("ageing_risk") in {"high", "elevated"} and (c.get("fill_pct") or 100) < 85:
            note = "Shrinking park + soft fill — age/odd-parts hypothesis fits"
        elif (c.get("park_change_5y_pct") or 0) > 0.1:
            note = "Growing park — fill issues less likely driven by old VIN mix alone"
        rows.append(
            {
                "country": c["ids_country"],
                "uio_name": c.get("uio_park_name"),
                "uio_units": c.get("uio_units_park"),
                "park_change_5y_pct": c.get("park_change_5y_pct"),
                "ageing_label": c.get("ageing_label"),
                "ageing_risk": c.get("ageing_risk"),
                "fill_pct": c.get("fill_pct"),
                "ids_d_avg_pct": c.get("ids_d_avg_pct"),
                "note": note,
            }
        )
    return {
        "r_park_change_vs_fill": r,
        "n": len(pairs),
        "rows": rows,
        "method_note": (
            "Ageing proxy = 5-year net UIO change from UIO.xlsx (latest vs −5y), not mean vehicle age. "
            "Hypothesis: shrinking parks leave older vehicles → harder repairs → lower fill. "
            "Positive r supports younger/growing ↔ higher fill."
        ),
    }


def hygiene_snapshot(payload: dict[str, Any]) -> dict[str, Any]:
    raw_n = len(payload.get("dealers") or [])
    raw_net = payload.get("network") or {}
    clean = hygienic_network(payload)
    zero_shells = sum(
        1
        for d in payload.get("dealers") or []
        if not is_rollup_country(str(d.get("country") or ""))
        and (d.get("total_pct") is None or float(d.get("total_pct") or 0) <= 0)
    )
    rollup_n = sum(1 for d in payload.get("dealers") or [] if is_rollup_country(str(d.get("country") or "")))
    return {
        "raw_dealer_rows": raw_n,
        "raw_network_d_pct": _pct((raw_net.get("sections_pct") or {}).get("D")),
        "raw_network_total_pct": _pct(raw_net.get("total_pct")),
        "clean_dealer_rows": clean.get("dealer_count"),
        "clean_network_d_pct": _pct((clean.get("sections_pct") or {}).get("D")),
        "clean_network_total_pct": _pct(clean.get("total_pct")),
        "dropped_benelux_rollup": rollup_n,
        "dropped_zero_shells": zero_shells,
        "note": "Clean stats exclude Benelux rollup (BE+NL already counted) and 0% unscored shells (MIDI empties).",
    }


def write_register(
    path: Path,
    cards: list[dict[str, Any]],
    cap_split: dict[str, Any],
    park_fill: dict[str, Any] | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DB_Posture"
    headers = [
        "Country",
        "PRISM DB",
        "Archetype",
        "Thesis",
        "Ask 1",
        "Ask 2",
        "Dealers scored",
        "IDS D avg %",
        "IDS D min %",
        "Genuine share %",
        "GPS €/UIO incl warranty",
        "GPS €/UIO ex warranty",
        "GPS light",
        "Fill %",
        "MOS days",
        "Park change 5y %",
        "Ageing label",
        "CAP D-only",
        "CAP D+C",
        "CAP C-only",
        "CAP OK",
        "Note",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1A2B4A")
    for c in cards:
        ws.append(
            [
                c["ids_country"],
                c.get("prism_db"),
                c["archetype"],
                c["thesis"],
                c["ask_1"],
                c["ask_2"],
                c.get("dealer_count"),
                c.get("ids_d_avg_pct"),
                c.get("ids_d_min_pct"),
                c.get("genuine_share_pct"),
                c.get("gps_spu_eur"),
                c.get("gps_spu_ex_warranty_eur"),
                c.get("gps_light"),
                c.get("fill_pct"),
                c.get("kpi_mos_days"),
                c.get("park_change_5y_pct_disp"),
                c.get("ageing_label"),
                c.get("cap_d_only"),
                c.get("cap_d_plus_c"),
                c.get("cap_c_only"),
                c.get("cap_ok"),
                c.get("note"),
            ]
        )

    ws2 = wb.create_sheet("CAP_Split_by_Country")
    ws2.append(["Country", "N scored", "D cut", "C cut", "D-only", "D+C", "C-only", "OK"])
    for cell in ws2[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1A2B4A")
    for r in cap_split.get("countries") or []:
        ws2.append(
            [
                r["country"],
                r["n"],
                round(r["d_cut"] * 100, 1) if r.get("d_cut") is not None else None,
                round(r["c_cut"] * 100, 1) if r.get("c_cut") is not None else None,
                r["counts"].get("D-only"),
                r["counts"].get("D+C"),
                r["counts"].get("C-only"),
                r["counts"].get("OK"),
            ]
        )

    ws3 = wb.create_sheet("CAP_Examples")
    ws3.append(["Country", "Dealer", "CAP type", "D %", "C %", "Total %", "Ask"])
    for cell in ws3[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="C0552D")
    for e in cap_split.get("examples") or []:
        ws3.append(
            [
                e["country"],
                e["name"],
                e["cap_type"],
                _pct(e.get("d_pct")),
                _pct(e.get("c_pct")),
                _pct(e.get("total_pct")),
                e.get("ask"),
            ]
        )

    if park_fill:
        ws4 = wb.create_sheet("Park_vs_Fill")
        ws4.append(
            [
                "Country",
                "UIO name",
                "UIO units",
                "Park 5y change %",
                "Ageing label",
                "Ageing risk",
                "Fill %",
                "D avg %",
                "Note",
            ]
        )
        for cell in ws4[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1A2B4A")
        for r in park_fill.get("rows") or []:
            ch = r.get("park_change_5y_pct")
            ws4.append(
                [
                    r.get("country"),
                    r.get("uio_name"),
                    r.get("uio_units"),
                    round(ch * 100, 1) if ch is not None else None,
                    r.get("ageing_label"),
                    r.get("ageing_risk"),
                    r.get("fill_pct"),
                    r.get("ids_d_avg_pct"),
                    r.get("note"),
                ]
            )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _json_safe(payload: dict) -> str:
    def scrub(o: Any) -> Any:
        if isinstance(o, float) and (math.isnan(o) or math.isinf(o)):
            return None
        if isinstance(o, dict):
            return {k: scrub(v) for k, v in o.items()}
        if isinstance(o, list):
            return [scrub(v) for v in o]
        return o

    return json.dumps(scrub(payload), ensure_ascii=False, allow_nan=False)


def render_html(payload: dict, chart_js: str) -> str:
    return (
        _TEMPLATE.read_text(encoding="utf-8")
        .replace("__BUILT_AT__", payload.get("built_at", ""))
        .replace("__CHART_JS__", chart_js)
        .replace("__DATA_JSON__", _json_safe(payload))
        .replace("__APP_JS__", _APP_JS.read_text(encoding="utf-8"))
    )


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    payload_ids = build_dashboard_payload(IDS_RESULT_REPORTS_DIR)
    hygiene = hygiene_snapshot(payload_ids)
    ids_d = ids_country_section_d(payload_ids)
    kpi = load_kpi_fill_mos()
    join_rows = build_join_rows(
        ids_d,
        load_pms_db_2025(),
        load_gps_2025(),
        load_genuine_share_2025(),
        kpi,
    )
    # PMS still loaded for optional offline use, but stripped from conclusions/cards
    attach_park_age(join_rows, load_park_age_proxy())
    corr = correlations(join_rows)
    cap_split = within_country_cap_split(payload_ids)
    cards = build_db_cards(join_rows, cap_split)
    spu_cov = build_kpi_spu_coverage(kpi)
    park_fill = build_park_fill_analysis(cards)

    study = {
        "built_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "hygiene": hygiene,
        "correlations": corr,
        "correlations_chart": {
            k: corr[k]
            for k in (
                "genuine_share_eur",
                "gps_spu_eur",
                "kpi_fill_rate",
                "kpi_mos_days",
            )
            if k in corr
        },
        "db_cards": cards,
        "cap_split": cap_split,
        "park_fill": park_fill,
        "glossary": GLOSSARY,
        "kpi_spu_coverage": spu_cov,
        "gps_warranty_note": (
            "Headline GPS €/UIO in this pack is spend_per_unit_eur and INCLUDES warranty parts. "
            "Cards also show €/UIO ex-warranty where available."
        ),
        "metric_notes": {
            "pearson": "Network r across DBs — not a within-DB Pearson. Used only as prior.",
            "cap_split": "Within country: weak D vs median∩65% floor; D+C if Service (C) also weak.",
            "chart_excludes": "Correlation chart and DB conclusions exclude PMS and KPI Parts Sales per Unit.",
            "israel": "IDS Israel = UMI only.",
            "park_age": park_fill.get("method_note"),
        },
    }

    (OUT_DIR / "study_payload.json").write_text(
        json.dumps(study, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )
    xlsx_path = OUT_DIR / "IDS_DB_Parts_Posture_Register.xlsx"
    write_register(xlsx_path, cards, cap_split, park_fill)

    if not CHART_JS.is_file():
        print(f"Chart.js missing: {CHART_JS}", file=sys.stderr)
        return 2
    chart_js = re.sub(r"\n//# sourceMappingURL=.*$", "", CHART_JS.read_text(encoding="utf-8"))
    html = render_html(study, chart_js)
    html_path = OUT_DIR / "IDS_DB_Parts_Posture_2026.html"
    html_path.write_text(html, encoding="utf-8")

    downloads = Path.home() / "Downloads"
    if downloads.is_dir():
        (downloads / html_path.name).write_text(html, encoding="utf-8")
        import shutil

        shutil.copy2(xlsx_path, downloads / xlsx_path.name)

    print("Wrote", html_path)
    print("Wrote", xlsx_path)
    print("Hygiene raw→clean dealers:", hygiene["raw_dealer_rows"], "→", hygiene["clean_dealer_rows"])
    print("CAP totals:", cap_split["totals"])
    print("DB cards:", len(cards))
    for c in cards[:6]:
        print(f"  {c['ids_country']:10} {c['archetype']:22} D={c.get('ids_d_avg_pct')}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
