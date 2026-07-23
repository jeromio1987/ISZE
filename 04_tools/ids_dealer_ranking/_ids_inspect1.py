import sys
sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
from openpyxl import load_workbook
import xlrd

# 1) 26 summary.xlsx
p = Path(r"C:\Users\jerom\Downloads\26 summary.xlsx")
wb = load_workbook(p, data_only=True)
print("=== 26 summary sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n--- {sn} dims {ws.dimensions} ---")
    for i, row in enumerate(ws.iter_rows(max_row=25, max_col=12, values_only=True), 1):
        if any(c is not None and str(c).strip() != "" for c in row):
            print(i, row)
wb.close()
