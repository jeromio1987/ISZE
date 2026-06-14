#!/usr/bin/env python3
"""Fill DB Evaluation for Awards 26.xlsx from IDS rankings and evaluation files."""

from __future__ import annotations

import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from ids_paths import DATA_ROOT, RANKINGS_DIR

DATA_ROOT = DATA_ROOT
RANKINGS_DIR = DATA_ROOT / "rankings"
AWARDS_PATH = RANKINGS_DIR / "DB Evaluation for Awards 26.xlsx"
def _latest_b1_workbook() -> Path | None:
    candidates = sorted(RANKINGS_DIR.glob("B1_Business_Plan_ranking*.xlsx"), reverse=True)
    return candidates[0] if candidates else None


B1_PATH = _latest_b1_workbook() or (RANKINGS_DIR / "B1_Business_Plan_ranking.xlsx")
TOTAL_RANKING = RANKINGS_DIR / "IDS_2026_TOTAL_ranking_all.xlsx"

# Import shared readers from sibling module
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_country_ranking import read_scorecard_metrics  # noqa: E402

TOTAL_LABELS = frozenset({"totaal", "total", "gesamt", "totale", "sum", "somme"})
MAX_IDS_POINTS = 280
ITUK_MAX_POINTS = 234


@dataclass
class CountryMetrics:
    n_ranked: int
    avg_pct: float  # 0–1
    hi_pts: int | None
    lo_pts: int | None
    all_ge_50_pct: bool
    b1_avg_pts: float | None


def _iter_xlsx(folder: Path):
    for path in sorted(folder.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        yield path


def _with_zip_xlsx(folder: Path):
    """Yield xlsx paths, extracting zip bundles to a temp dir when needed."""
    yielded = False
    for path in _iter_xlsx(folder):
        yielded = True
        yield path
    if yielded:
        return
    for zpath in sorted(folder.glob("*.zip")):
        with TemporaryDirectory() as tmp:
            with zipfile.ZipFile(zpath) as zf:
                zf.extractall(tmp)
            for path in _iter_xlsx(Path(tmp)):
                yield path


def _pick_achieved_points(g: object, h: object, j: object) -> float | None:
    """Pick auditor/self total points from scorecard total row (cols G/H/J)."""
    candidates: list[float] = []
    for v in (h, g, j):
        try:
            f = float(v) if v is not None else None
        except (TypeError, ValueError):
            f = None
        if f is None or f <= 0 or f > MAX_IDS_POINTS:
            continue
        candidates.append(f)
    return max(candidates) if candidates else None


def _read_total_row_points(path: Path) -> tuple[float | None, float | None]:
    """Return (achieved_points, pct_0_1) from scorecard total row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    for sn in wb.sheetnames:
        low = sn.strip().lower()
        if low not in ("scorecard", "scorekaart") and "score card" not in low and "scorekaart" not in low:
            continue
        ws = wb[sn]
        for r in range(15, 35):
            label = ws.cell(r, 2).value
            if not label:
                continue
            if str(label).strip().lower() not in TOTAL_LABELS:
                continue
            g = ws.cell(r, 7).value
            h = ws.cell(r, 8).value
            j = ws.cell(r, 10).value
            pct = ws.cell(r, 11).value
            wb.close()
            pts = _pick_achieved_points(g, h, j)
            try:
                p = float(pct) if pct is not None else None
            except (TypeError, ValueError):
                p = None
            if p is not None and p > 1:
                p = p / 100
            return pts, p
        wb.close()
        return None, None
    wb.close()
    return None, None


def _pts_from_pcts(pcts: list[float], max_pts: float = MAX_IDS_POINTS) -> tuple[int, int]:
    scaled = [p * max_pts for p in pcts]
    return int(round(max(scaled))), int(round(min(scaled)))


def _is_eval_workbook(path: Path) -> bool:
    upper = path.name.upper()
    if upper.startswith("~$"):
        return False
    if "DEALERS" in upper and "EVALUATION" not in upper:
        return False
    return True


def _find_eval_sheet(sheetnames: list[str]) -> str | None:
    from israel_evaluation import find_israel_eval_sheet

    found = find_israel_eval_sheet(sheetnames)
    if found:
        return found
    for name in sheetnames:
        low = name.strip().lower()
        if "formular" in low and "evaluare" in low:
            return name
        if "evaluare" in low and "standarde" in low:
            return name
    return None


def _read_eval_sheet_pct(path: Path) -> float | None:
    from israel_evaluation import read_metrics_from_ws

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = _find_eval_sheet(wb.sheetnames)
    if not sheet:
        wb.close()
        return None
    pct, _ = read_metrics_from_ws(wb[sheet])
    wb.close()
    if pct is None:
        return None
    return pct / 100 if pct > 1 else pct


def metrics_from_folder(folder: Path) -> CountryMetrics | None:
    if not folder.is_dir():
        return None
    pcts: list[float] = []
    pts: list[float] = []
    for path in _iter_xlsx(folder):
        if not _is_eval_workbook(path):
            continue
        try:
            pct_val, _ = read_scorecard_metrics(path)
        except Exception:
            pct_val = None
        raw_pts, raw_pct = _read_total_row_points(path)
        if pct_val is not None:
            pcts.append(pct_val / 100 if pct_val > 1 else pct_val)
        elif raw_pct is not None:
            pcts.append(raw_pct)
        else:
            eval_pct = _read_eval_sheet_pct(path)
            if eval_pct is not None:
                pcts.append(eval_pct)
        if raw_pts is not None:
            pts.append(raw_pts)
    if not pcts:
        return None
    if pts:
        hi_pts, lo_pts = int(round(max(pts))), int(round(min(pts)))
    else:
        hi_pts, lo_pts = _pts_from_pcts(pcts)
    return CountryMetrics(
        n_ranked=len(pcts),
        avg_pct=sum(pcts) / len(pcts),
        hi_pts=hi_pts,
        lo_pts=lo_pts,
        all_ge_50_pct=all(p >= 0.5 for p in pcts),
        b1_avg_pts=None,
    )


def metrics_from_ranking_code(code: str) -> CountryMetrics | None:
    if not TOTAL_RANKING.is_file():
        return None
    wb = load_workbook(TOTAL_RANKING, read_only=True, data_only=True)
    ws = wb["All dealers"]
    pcts: list[float] = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 2).value) != code:
            continue
        v = ws.cell(r, 5).value
        if v is None:
            continue
        try:
            pcts.append(float(v) / 100)
        except (TypeError, ValueError):
            pass
    wb.close()
    if not pcts:
        return None
    hi_pts, lo_pts = _pts_from_pcts(pcts)
    return CountryMetrics(
        n_ranked=len(pcts),
        avg_pct=sum(pcts) / len(pcts),
        hi_pts=hi_pts,
        lo_pts=lo_pts,
        all_ge_50_pct=all(p >= 0.5 for p in pcts),
        b1_avg_pts=None,
    )


def metrics_from_ituk() -> CountryMetrics | None:
    path = DATA_ROOT / "ITUK" / "IDS ITUK self evaluations.xlsx"
    if not path.is_file():
        return None
    from build_ituk_ranking import ITUK_BLOCKS, _read_block, _resolve_sheet

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = _resolve_sheet(wb, None)
    rows: list[dict] = []
    for name_row, pct_row, pts_row in ITUK_BLOCKS:
        rows.extend(
            _read_block(
                ws,
                name_row=name_row,
                percent_row=pct_row,
                points_row=pts_row,
                source_file=path.name,
            )
        )
    wb.close()
    if not rows:
        return None
    pcts = [r["points"] / 100 for r in rows if r.get("points") is not None]
    raw = [float(r["items"]) for r in rows if r.get("items") is not None]
    hi_pts, lo_pts = (
        (int(round(max(raw))), int(round(min(raw))))
        if raw
        else _pts_from_pcts(pcts, ITUK_MAX_POINTS)
    )
    return CountryMetrics(
        n_ranked=len(pcts),
        avg_pct=sum(pcts) / len(pcts),
        hi_pts=hi_pts,
        lo_pts=lo_pts,
        all_ge_50_pct=all(p >= 0.5 for p in pcts),
        b1_avg_pts=None,
    )


def metrics_from_israel_folder() -> CountryMetrics | None:
    folder = DATA_ROOT / "Israel"
    if not folder.is_dir():
        return None
    from israel_evaluation import find_israel_eval_sheet, read_metrics_from_ws

    pcts: list[float] = []
    assessed: list[float] = []
    for path in _iter_xlsx(folder):
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet = find_israel_eval_sheet(wb.sheetnames)
        if not sheet:
            wb.close()
            continue
        pts, _ = read_metrics_from_ws(wb[sheet])
        wb.close()
        if pts is None:
            continue
        pcts.append(pts / 100 if pts > 1 else pts)
        assessed.append(pts / 100 * MAX_IDS_POINTS if pts > 1 else pts * MAX_IDS_POINTS)
    if not pcts:
        return None
    hi_pts, lo_pts = _pts_from_pcts(pcts)
    if assessed:
        hi_pts, lo_pts = int(round(max(assessed))), int(round(min(assessed)))
    return CountryMetrics(
        n_ranked=len(pcts),
        avg_pct=sum(pcts) / len(pcts),
        hi_pts=hi_pts,
        lo_pts=lo_pts,
        all_ge_50_pct=all(p >= 0.5 for p in pcts),
        b1_avg_pts=None,
    )


def b1_avg_from_folder(folder: Path) -> float | None:
    from build_b1_ranking import read_b1_from_workbook

    vals: list[float] = []
    for path in _iter_xlsx(folder):
        if not _is_eval_workbook(path):
            continue
        try:
            _, _, assessed, _ = read_b1_from_workbook(path)
        except Exception:
            continue
        if assessed is not None:
            try:
                vals.append(float(assessed))
            except (TypeError, ValueError):
                pass
    if not vals:
        return None
    return sum(vals) / len(vals)


def load_b1_averages() -> dict[str, float]:
    if not B1_PATH.is_file():
        return {}
    wb = load_workbook(B1_PATH, read_only=True, data_only=True)
    out: dict[str, float] = {}
    for sn in wb.sheetnames:
        if sn.startswith("Map_") or sn.endswith("_err"):
            continue
        ws = wb[sn]
        vals: list[float] = []
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 5).value
            if v is None:
                continue
            try:
                vals.append(float(v))
            except (TypeError, ValueError):
                pass
        if vals:
            out[sn] = sum(vals) / len(vals)
    wb.close()
    return out


def resolve_input_dir(country: str) -> Path:
    from ids_paths import resolve_input_dir as _resolve

    return _resolve(country, DATA_ROOT)


def fmt_b1(value: float) -> str:
    s = f"{value:.2f}".replace(".", ",")
    return s.rstrip("0").rstrip(",") if "," in s else s


# Awards sheet row -> data source
ROW_SOURCES: dict[int, tuple[str, Path | None]] = {
    7: ("Portugal", DATA_ROOT / "Portugal"),
    8: ("Poland", DATA_ROOT / "Poland"),
    9: ("BE", resolve_input_dir("BE")),
    10: ("BE_LU", resolve_input_dir("BE")),  # Luxembourg — IBX / Belgium network
    11: ("NL", DATA_ROOT / "NL"),
    20: ("Romania", resolve_input_dir("Romania")),
    23: ("ITUK", None),
    24: ("IUK", DATA_ROOT / "IUK"),
    27: ("MIDI_FR", None),
    28: ("MIDI_IT", None),
    29: ("MIDI_ES", None),
    30: ("MIDI_CH", None),
    35: ("BG", DATA_ROOT / "Petros"),  # Bulgaria — Petros distributor
    36: ("Petros", DATA_ROOT / "Petros"),
    37: ("RSA", DATA_ROOT / "RSA"),
    38: ("Israel", DATA_ROOT / "Israel"),
}

# Rows already manually completed — do not overwrite score columns
SKIP_ROWS = frozenset({12, 14, 22, 26, 29})

B1_SHEET_MAP = {
    "Portugal": "Portugal",
    "Poland": "Poland",
    "BE": "BE",
    "BE_LU": "BE",
    "NL": "NL",
    "Romania": "Romania",
    "Portugal": "Portugal",
    "Hungary": "Hungary",
    "BG": "Petros",
    "ITUK": "ITUK",
    "IUK": "IUK",
    "MIDI_FR": "MIDI_FR",
    "MIDI_IT": "MIDI_IT",
    "MIDI_ES": "MIDI_ES",
    "MIDI_CH": "MIDI_CH",
    "Petros": "Petros",
    "RSA": "RSA",
    "Israel": "Israel",
}


def get_metrics(key: str, folder: Path | None) -> CountryMetrics | None:
    if key == "ITUK":
        return metrics_from_ituk()
    if key == "Israel":
        return metrics_from_israel_folder()
    if key == "BE_LU" and folder is not None and folder.is_dir():
        return metrics_from_folder(folder)
    if key == "BG" and folder is not None and folder.is_dir():
        return metrics_from_folder(folder)
    if folder is not None and folder.is_dir():
        m = metrics_from_folder(folder)
        if m:
            return m
    ranking_codes = {
        "BE": "BE",
        "Poland": "Poland",
        "NL": "NL",
        "Petros": "Petros",
        "BG": "Petros",
        "RSA": "RSA",
        "IUK": "IUK",
        "MIDI_FR": "MIDI_FR",
        "MIDI_IT": "MIDI_IT",
        "MIDI_ES": "MIDI_ES",
        "MIDI_CH": "MIDI_CH",
        "Portugal": "Portugal",
        "Hungary": "Hungary",
    }
    code = ranking_codes.get(key)
    if code:
        return metrics_from_ranking_code(code)
    return None


def refresh_awards_b1(dry_run: bool = False) -> int:
    """Update only column O (B.1 average points) from the latest B1 ranking workbook."""
    b1_avgs = load_b1_averages()
    wb = load_workbook(AWARDS_PATH)
    ws = wb["Awards"]
    updated = 0

    for row, key in {r: k for r, (k, _) in ROW_SOURCES.items()}.items():
        b1_key = B1_SHEET_MAP.get(key)
        avg = b1_avgs.get(b1_key) if b1_key else None
        if avg is None or avg <= 0:
            continue
        ws.cell(row, 15).value = fmt_b1(avg)
        print(f"  row {row} {key}: B1={fmt_b1(avg)}")
        updated += 1

    if dry_run:
        print(f"Dry run — would update {updated} B1 cells")
        return 0

    wb.save(AWARDS_PATH)
    print(f"Saved B1 column in {AWARDS_PATH}")
    return 0


def fill_awards(dry_run: bool = False) -> int:
    b1_avgs = load_b1_averages()
    wb = load_workbook(AWARDS_PATH)
    ws = wb["Awards"]
    updated = 0

    for row, (key, folder) in ROW_SOURCES.items():
        if row in SKIP_ROWS:
            continue
        subs = ws.cell(row, 6).value
        if not subs or subs == 0:
            continue
        if ws.cell(row, 10).value not in (None, ""):
            continue

        m = get_metrics(key, folder)
        if m is None:
            print(f"  SKIP row {row} {key}: no metrics")
            continue

        b1_key = B1_SHEET_MAP.get(key)
        if b1_key and b1_key in b1_avgs:
            m.b1_avg_pts = b1_avgs[b1_key]
        elif folder is not None and folder.is_dir():
            m.b1_avg_pts = b1_avg_from_folder(folder)

        ws.cell(row, 10).value = round(m.avg_pct, 4)
        if m.hi_pts is not None:
            ws.cell(row, 11).value = m.hi_pts
        if m.lo_pts is not None:
            ws.cell(row, 12).value = m.lo_pts
        ws.cell(row, 14).value = "Yes" if m.all_ge_50_pct else "No"
        if m.b1_avg_pts is not None and m.b1_avg_pts > 0:
            ws.cell(row, 15).value = fmt_b1(m.b1_avg_pts)

        print(
            f"  row {row} {key}: avg={m.avg_pct:.2%} hi={m.hi_pts} lo={m.lo_pts} "
            f"all>=50%={m.all_ge_50_pct} b1={m.b1_avg_pts}"
        )
        updated += 1

    if dry_run:
        print(f"Dry run — would update {updated} rows")
        return 0

    wb.save(AWARDS_PATH)
    print(f"Saved {AWARDS_PATH} ({updated} rows updated)")
    return 0


if __name__ == "__main__":
    dry = "--dry-run" in sys.argv
    if "--b1-only" in sys.argv:
        raise SystemExit(refresh_awards_b1(dry_run=dry))
    raise SystemExit(fill_awards(dry_run=dry))
