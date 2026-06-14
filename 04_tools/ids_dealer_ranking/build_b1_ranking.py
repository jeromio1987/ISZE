#!/usr/bin/env python3
"""Build B.1 Business Plan rankings — one workbook tab per country, no charts.

Also writes a dated "Map_YYYY-MM-DD" summary tab with input coverage.
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from build_country_ranking import (
    PREFIXES,
    _coerce_percent,
    dealer_display_name,
    dealer_name_from_filename,
    read_scorecard_metrics,
)

from ids_paths import DATA_ROOT, RANKINGS_DIR

DATA_ROOT = DATA_ROOT
RUN_DATE = date.today()
OUTPUT = DATA_ROOT / "rankings" / f"B1_Business_Plan_ranking_{RUN_DATE:%Y-%m-%d}.xlsx"

SKIP_DIRS = {"rankings", "reports", "__pycache__", "MIDI"}


def _find_score_sheet_name(sheetnames: list[str]) -> str | None:
    for name in sheetnames:
        if name.strip().lower() == "scorecard":
            return name
    for name in sheetnames:
        low = name.strip().lower()
        if any(
            k in low
            for k in (
                "scorekaart",
                "score card",
                "fiche de r",
                "results",
                "eredmény",
                "eredmeny",
            )
        ):
            return name
    return None


def _is_b1_subcategory(label) -> bool:
    if not label:
        return False
    s = str(label).strip()
    return bool(re.match(r"^B\.1(\b|[\s\-])", s, re.IGNORECASE))


def _find_eval_sheet(wb) -> str | None:
    """Heuristic: pick an evaluation sheet when Score Card is missing."""
    from israel_evaluation import find_israel_eval_sheet

    found = find_israel_eval_sheet(wb.sheetnames)
    if found:
        return found
    for name in wb.sheetnames:
        low = name.strip().lower()
        if "evaluation" in low and "comparison" not in low:
            return name
    for name in wb.sheetnames:
        low = name.strip().lower()
        if any(
            k in low
            for k in ("evaluare", "formulier", "standard", "standarde", "értékel", "ertekel")
        ):
            return name
    return None


def read_b1_from_eval_sheet(
    path: Path,
) -> tuple[float | None, float | None, float | None, float | None]:
    """Fallback: read B.1 section totals from the Evaluation sheet.

    Many templates encode section totals on rows like "B.1 Business Plan" with:
      - max points in column G (7)
      - assessed points in column L (12)
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_name = _find_eval_sheet(wb)
    if not sheet_name:
        wb.close()
        raise ValueError(f"no evaluation sheet in {path.name}")

    from israel_evaluation import find_israel_eval_sheet, read_b1_from_ws

    ws = wb[sheet_name]
    if find_israel_eval_sheet([sheet_name]):
        pts_pct, items_pct, assessed_val, max_val = read_b1_from_ws(ws)
        wb.close()
        if pts_pct is None and items_pct is None:
            raise ValueError(f"B.1 section row not found on evaluation sheet in {path.name}")
        return pts_pct, items_pct, assessed_val, max_val

    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str):
            continue
        s = label.strip()
        if not _is_b1_subcategory(s):
            continue
        max_pts = row[6] if len(row) > 6 else None
        assessed = row[11] if len(row) > 11 else None
        if max_pts in (None, "") and len(row) > 8:
            max_pts = row[8]
        if assessed in (None, "") and len(row) > 7:
            assessed = row[7]
        try:
            max_val = float(max_pts) if max_pts is not None else None
        except (TypeError, ValueError):
            max_val = None
        try:
            assessed_val = float(assessed) if assessed is not None else None
        except (TypeError, ValueError):
            assessed_val = None

        pts_pct = None
        if assessed_val is not None and max_val not in (None, 0):
            try:
                pts_pct = round(assessed_val / max_val * 100, 2)
            except ZeroDivisionError:
                pts_pct = None

        wb.close()
        return pts_pct, pts_pct, assessed_val, max_val

    wb.close()
    raise ValueError(f"B.1 section row not found on evaluation sheet in {path.name}")


def _float_cell(value) -> float | None:
    try:
        return float(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _read_b1_from_scorecard_row(ws, row: int) -> tuple[float | None, float | None, float | None, float | None]:
    """Read B.1 row — percentage from col K is authoritative on English templates."""
    max_pts = _float_cell(ws.cell(row, 6).value)
    pts_pct = _coerce_percent(ws.cell(row, 11).value)
    items_pct = _coerce_percent(ws.cell(row, 12).value)

    dealer_pts: float | None = None
    if pts_pct is not None and max_pts:
        dealer_pts = round(pts_pct / 100 * max_pts, 2)
    else:
        for candidate in (
            _float_cell(ws.cell(row, 7).value),
            _float_cell(ws.cell(row, 10).value),
            _float_cell(ws.cell(row, 8).value),
        ):
            if candidate is None or candidate < 0:
                continue
            if max_pts and candidate > max_pts:
                continue
            dealer_pts = candidate
            break
        if dealer_pts is not None and max_pts:
            pts_pct = round(dealer_pts / max_pts * 100, 2)

    if items_pct is None:
        items_pct = pts_pct

    return pts_pct, items_pct, dealer_pts, max_pts


def read_b1_from_workbook(path: Path) -> tuple[float | None, float | None, float | None, float | None]:
    """Return points_%, items_%, dealer_points, max_points for B.1 subcategory."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sc_name = _find_score_sheet_name(wb.sheetnames)
    if not sc_name:
        wb.close()
        # Fallback for templates without Score Card / Scorecard tabs.
        return read_b1_from_eval_sheet(path)

    ws = wb[sc_name]
    for r in range(1, min(ws.max_row + 1, 40)):
        sub = ws.cell(r, 5).value
        cat = ws.cell(r, 2).value
        if _is_b1_subcategory(sub) or (
            cat and _is_b1_subcategory(cat) and sub is None
        ):
            result = _read_b1_from_scorecard_row(ws, r)
            wb.close()
            return result

    wb.close()
    raise ValueError(f"B.1 subcategory row not found in {path.name}")


ITUK_NAME_ROWS = (1, 27)


def _read_b1_ituk_block(
    ws,
    *,
    name_row: int,
    b1_row: int,
    source_file: str,
) -> list[dict]:
    max_pts = ws.cell(b1_row, 3).value
    rows: list[dict] = []
    col = 5
    while col < ws.max_column:
        name = ws.cell(name_row, col).value
        if not name or not str(name).strip():
            col += 1
            continue
        name = str(name).strip()
        if name.startswith(",") or "Section" in name or name == "Parts & Service":
            col += 1
            continue

        assessed = ws.cell(b1_row, col).value
        section_pct = ws.cell(b1_row, col + 1).value
        try:
            assessed_val = float(assessed) if assessed is not None else 0.0
        except (TypeError, ValueError):
            assessed_val = 0.0

        pts_pct = _coerce_percent(section_pct)
        if pts_pct is None and max_pts:
            try:
                pts_pct = round(assessed_val / float(max_pts) * 100, 2)
            except (TypeError, ValueError, ZeroDivisionError):
                pts_pct = None
        if pts_pct is None:
            pts_pct = 0.0

        rows.append(
            {
                "name": name,
                "points": pts_pct,
                "items": pts_pct,
                "assessed_points": assessed_val,
                "max_points": max_pts,
                "file": source_file,
            }
        )
        col += 2
    return rows


def _merge_b1_ituk_rows(blocks: list[list[dict]]) -> list[dict]:
    """One row per dealer; prefer higher B.1 assessed points when name appears twice."""
    by_name: dict[str, dict] = {}
    for block_rows in blocks:
        for row in block_rows:
            key = row["name"].casefold()
            existing = by_name.get(key)
            if existing is None or (row["assessed_points"] or 0) > (
                existing["assessed_points"] or 0
            ):
                by_name[key] = row
    merged = list(by_name.values())
    merged.sort(key=lambda r: (-(r["points"] or -1), r["name"].lower()))
    return merged


def read_b1_ituk(path: Path) -> list[dict]:
    """B.1 from ITUK Sheet2 — both dealer blocks (rows 1 and 27), incl. 0% scores."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb.worksheets[1]

    b1_row = None
    for r in range(1, min(ws.max_row + 1, 60)):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a and str(a).strip().upper() == "B.1" and b:
            b1_row = r
            break

    if b1_row is None:
        wb.close()
        raise ValueError("B.1 row not found on ITUK Sheet2")

    blocks = [
        _read_b1_ituk_block(ws, name_row=name_row, b1_row=b1_row, source_file=path.name)
        for name_row in ITUK_NAME_ROWS
    ]
    wb.close()
    merged = _merge_b1_ituk_rows(blocks)
    by_name = {r["name"].casefold(): r for r in merged}

    from build_ituk_ranking import read_ituk_dealers

    aligned: list[dict] = []
    max_pts = merged[0]["max_points"] if merged else 12
    for total_row in read_ituk_dealers(path):
        key = total_row["name"].casefold()
        if key in by_name:
            aligned.append(by_name[key])
        else:
            aligned.append(
                {
                    "name": total_row["name"],
                    "points": 0.0,
                    "items": 0.0,
                    "assessed_points": 0.0,
                    "max_points": max_pts,
                    "file": path.name,
                }
            )
    aligned.sort(key=lambda r: (-(r["points"] or -1), r["name"].lower()))
    return aligned


from ids_paths import discover_country_folders, resolve_input_dir


def discover_countries() -> list[str]:
    """All country-like folders under DATA_ROOT with at least one .xlsx (incl. nested)."""
    return discover_country_folders(DATA_ROOT)


def collect_country(country: str) -> list[dict]:
    if country == "ITUK":
        path = DATA_ROOT / "ITUK" / "IDS ITUK self evaluations.xlsx"
        if not path.is_file():
            return []
        return read_b1_ituk(path)

    input_dir = resolve_input_dir(country)
    if not input_dir.is_dir():
        return []

    rows: list[dict] = []
    for path in sorted(input_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            total_pts, _ = read_scorecard_metrics(path)
        except Exception:
            total_pts = None
        if total_pts is None or total_pts <= 0:
            continue
        try:
            pts_pct, items_pct, dealer_pts, max_pts = read_b1_from_workbook(path)
        except Exception as exc:
            rows.append(
                {
                    "name": dealer_display_name(path),
                    "points": None,
                    "items": None,
                    "assessed_points": None,
                    "max_points": None,
                    "file": path.name,
                    "error": str(exc),
                }
            )
            continue
        if pts_pct is None and items_pct is None and dealer_pts is None:
            continue
        rows.append(
            {
                "name": dealer_display_name(path),
                "points": 0.0 if pts_pct is None else pts_pct,
                "items": 0.0 if items_pct is None else items_pct,
                "assessed_points": dealer_pts,
                "max_points": max_pts,
                "file": path.name,
                "error": None,
            }
        )

    rows.sort(key=lambda r: (-(r["points"] or -1), -(r["items"] or -1), r["name"].lower()))
    return rows


def write_country_sheet(wb: Workbook, country: str, rows: list[dict]) -> None:
    title = country[:31]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)

    headers = [
        "Rank",
        "Name",
        "B.1 Points achieved %",
        "B.1 Items scored %",
        "B.1 Points (scored)",
        "B.1 Points (max)",
        "Source file",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    rank = 0
    for row in rows:
        if row.get("error"):
            continue
        if row["points"] is None and row["items"] is None:
            continue
        rank += 1
        ws.append(
            [
                rank,
                row["name"],
                row["points"],
                row["items"],
                row.get("assessed_points"),
                row.get("max_points"),
                row["file"],
            ]
        )

    err = [r for r in rows if r.get("error")]
    if err:
        es = wb.create_sheet(f"{title[:28]}_err"[:31])
        es.append(["File", "Dealer", "Error"])
        for r in err:
            es.append([r["file"], r["name"], r["error"]])

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 40


def write_map_sheet(
    wb: Workbook,
    *,
    run_date: date,
    summary_rows: list[dict],
) -> None:
    title = f"Map_{run_date:%Y-%m-%d}"[:31]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 0)

    headers = ["Country", "Input folder", "XLSX files found", "Dealers ranked", "Errors"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for r in summary_rows:
        ws.append(
            [
                r["country"],
                str(r["input_dir"]),
                r["xlsx_files"],
                r["ranked_dealers"],
                r["errors"],
            ]
        )

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 10
    ws.freeze_panes = "A2"


def main() -> int:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    countries = discover_countries()
    if not countries:
        print(f"No country folders with xlsx found under {DATA_ROOT}", file=sys.stderr)
        return 1

    summary_rows: list[dict] = []
    for country in countries:
        input_dir = (
            (DATA_ROOT / "ITUK") if country == "ITUK" else resolve_input_dir(country)
        )
        if country == "ITUK":
            xlsx_files = 1 if (DATA_ROOT / "ITUK" / "IDS ITUK self evaluations.xlsx").is_file() else 0
        else:
            try:
                xlsx_files = sum(
                    1
                    for p in input_dir.glob("*.xlsx")
                    if p.is_file() and not p.name.startswith("~$")
                )
            except Exception:
                xlsx_files = 0

        rows = collect_country(country)
        if not rows:
            print(f"SKIP {country}: no data")
            continue
        write_country_sheet(wb, country, rows)
        ok = sum(1 for r in rows if not r.get("error"))
        err = sum(1 for r in rows if r.get("error"))
        summary_rows.append(
            {
                "country": country,
                "input_dir": input_dir,
                "xlsx_files": xlsx_files,
                "ranked_dealers": ok,
                "errors": err,
            }
        )
        print(f"{country}: {ok} dealers ranked")

    from midi_fr_ranking import (
        MIDI_FILE,
        collect_midi_b1_dealers,
        discover_midi_cod_divs,
        midi_label,
    )

    if MIDI_FILE.is_file():
        for cod in discover_midi_cod_divs(MIDI_FILE):
            label = midi_label(cod)
            midi_rows = [
                {
                    "name": r["name"],
                    "points": r["points"],
                    "items": r["items"],
                    "assessed_points": r["assessed_points"],
                    "max_points": r["max_points"],
                    "file": r["file"],
                    "error": None,
                }
                for r in collect_midi_b1_dealers(cod, MIDI_FILE)
            ]
            if not midi_rows:
                print(f"SKIP {label}: no B.1 data")
                continue
            write_country_sheet(wb, label, midi_rows)
            summary_rows.append(
                {
                    "country": label,
                    "input_dir": MIDI_FILE.parent,
                    "xlsx_files": 1,
                    "ranked_dealers": len(midi_rows),
                    "errors": 0,
                }
            )
            print(f"{label}: {len(midi_rows)} dealers ranked (MIDI_FR.xlsx / {cod})")

    if not wb.sheetnames:
        print("No rankings produced.", file=sys.stderr)
        return 1

    if summary_rows:
        write_map_sheet(wb, run_date=RUN_DATE, summary_rows=summary_rows)

    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
