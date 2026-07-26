#!/usr/bin/env python3
"""Build two single-file global rankings: total IDS and B.1 across all countries."""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from ids_paths import RANKINGS_DIR
RUN_DATE = date.today()

OUTPUT_TOTAL = RANKINGS_DIR / "IDS_2026_TOTAL_ranking_all.xlsx"
OUTPUT_B1 = RANKINGS_DIR / "IDS_2026_B1_ranking_all.xlsx"


def _safe_save(wb: Workbook, path: Path) -> Path:
    """Save workbook; if target is locked, write a dated sibling file instead."""
    RANKINGS_DIR.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        tmp.replace(path)
        return path
    except OSError:
        alt = path.with_name(f"{path.stem}_{RUN_DATE}{path.suffix}")
        tmp.replace(alt)
        print(f"  (target locked — wrote {alt.name} instead)", file=sys.stderr)
        return alt

TOTAL_SOURCES = (
    ("BE_dealer_ranking.xlsx", "BE"),
    ("NL_dealer_ranking.xlsx", "NL"),
    ("Denmark_dealer_ranking.xlsx", "Denmark"),
    ("Poland_dealer_ranking.xlsx", "Poland"),
    ("Petros_dealer_ranking.xlsx", "Petros"),
    ("RSA_dealer_ranking.xlsx", "RSA"),
    ("IUK_dealer_ranking.xlsx", "IUK"),
    ("ITUK_dealer_ranking.xlsx", "ITUK"),
    ("MIDI_FR_dealer_ranking.xlsx", "MIDI_FR"),
    ("MIDI_IT_dealer_ranking.xlsx", "MIDI_IT"),
    ("MIDI_CH_dealer_ranking.xlsx", "MIDI_CH"),
    ("MIDI_ES_dealer_ranking.xlsx", "MIDI_ES"),
    ("Israel_dealer_ranking.xlsx", "Israel"),
    ("Romania_dealer_ranking.xlsx", "Romania"),
    ("Portugal_dealer_ranking.xlsx", "Portugal"),
    ("Hungary_dealer_ranking.xlsx", "Hungary"),
    ("Malta_dealer_ranking.xlsx", "Malta"),
    ("Cyprus_dealer_ranking.xlsx", "Cyprus"),
)

B1_HEADERS = (
    "Rank",
    "Name",
    "B.1 Points achieved %",
    "B.1 Items scored %",
    "B.1 Points (scored)",
    "B.1 Points (max)",
    "Source file",
)

SKIP_B1_TABS = re.compile(r"^(Map_|Index$|.*_err$)", re.IGNORECASE)
MIDI_B1_TAB_PREFIX = re.compile(r"^MIDI_", re.IGNORECASE)


def _header_row(ws) -> list[str]:
    return [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]


def _col_for(headers: list[str], label: str, default: int) -> int:
    if label in headers:
        return headers.index(label) + 1
    return default


def _col_for_any(headers: list[str], labels: tuple[str, ...], default: int) -> int:
    for label in labels:
        if label in headers:
            return headers.index(label) + 1
    return default


ITUK_ITEMS_MAX = 234


def _coerce_items_pct(
    value,
    *,
    header: str,
    points_pct: float | None,
) -> float | None:
    """Normalise Items scored % — some legacy ITUK sheets store raw counts."""
    if value is None:
        return points_pct
    try:
        items_f = float(value)
    except (TypeError, ValueError):
        return points_pct
    if items_f <= 100.001:
        return round(items_f, 2)
    if header == "Assessed points" or items_f > 150:
        return round(items_f / ITUK_ITEMS_MAX * 100, 2)
    return round(items_f, 2)


def _read_total_country(path: Path, country: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = country if country in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    headers = _header_row(ws)
    rank_col = _col_for(headers, "Rank", 1)
    name_col = _col_for(headers, "Name", 2)
    pts_col = _col_for_any(headers, ("Points achieved %", "Total percentage"), 3)
    items_col = _col_for_any(headers, ("Items scored %", "Assessed points"), 4)
    items_header = headers[items_col - 1] if items_col <= len(headers) else ""
    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if not name or str(name).strip().startswith("("):
            continue
        pts = ws.cell(r, pts_col).value
        items = ws.cell(r, items_col).value if items_col <= ws.max_column else pts
        try:
            pts_f = float(pts) if pts is not None else None
        except (TypeError, ValueError):
            pts_f = None
        if pts_f is None or pts_f > 100:
            continue
        items_f = _coerce_items_pct(items, header=items_header, points_pct=pts_f)
        rows.append(
            {
                "country": country,
                "country_rank": ws.cell(r, rank_col).value,
                "name": str(name).strip(),
                "points": pts_f,
                "items": items_f,
            }
        )
    wb.close()
    return rows


def _read_b1_country(ws, country: str) -> list[dict]:
    headers = _header_row(ws)
    name_col = headers.index("Name") + 1 if "Name" in headers else 2
    pts_col = (
        headers.index("B.1 Points achieved %") + 1
        if "B.1 Points achieved %" in headers
        else 3
    )
    items_col = (
        headers.index("B.1 Items scored %") + 1
        if "B.1 Items scored %" in headers
        else 4
    )
    scored_col = (
        headers.index("B.1 Points (scored)") + 1
        if "B.1 Points (scored)" in headers
        else 5
    )
    max_col = (
        headers.index("B.1 Points (max)") + 1 if "B.1 Points (max)" in headers else 6
    )
    file_col = headers.index("Source file") + 1 if "Source file" in headers else 7

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if not name:
            continue
        pts = ws.cell(r, pts_col).value
        try:
            pts_f = float(pts) if pts is not None else None
        except (TypeError, ValueError):
            pts_f = None
        if pts_f is None:
            continue
        rows.append(
            {
                "country": country,
                "country_rank": ws.cell(r, 1).value,
                "name": str(name).strip(),
                "points": pts_f,
                "items": ws.cell(r, items_col).value,
                "scored": ws.cell(r, scored_col).value,
                "max_pts": ws.cell(r, max_col).value,
                "file": ws.cell(r, file_col).value,
            }
        )
    return rows


def _find_latest_b1_workbook() -> Path | None:
    candidates = sorted(RANKINGS_DIR.glob("B1_Business_Plan_ranking*.xlsx"), reverse=True)
    return candidates[0] if candidates else None


def _write_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    data_rows: list[list],
    *,
    col_widths: tuple[float, ...] = (),
) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 0)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in data_rows:
        ws.append(row)
    from openpyxl.utils import get_column_letter

    widths = col_widths or (8, 12, 10, 36, 18, 18)
    for i, w in enumerate(widths, start=1):
        if i <= len(headers):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_total_file() -> int:
    all_rows: list[dict] = []
    missing: list[str] = []
    for filename, country in TOTAL_SOURCES:
        path = RANKINGS_DIR / filename
        if not path.is_file():
            missing.append(filename)
            continue
        all_rows.extend(_read_total_country(path, country))

    if not all_rows:
        print("No total ranking sources found.", file=sys.stderr)
        return 1

    all_rows.sort(key=lambda r: (-r["points"], r["country"].casefold(), r["name"].casefold()))

    out = Workbook()
    out.remove(out.active)
    headers = [
        "Global rank",
        "Country",
        "Rank (country)",
        "Dealer",
        "Points achieved %",
        "Items scored %",
    ]
    data = []
    for i, r in enumerate(all_rows, start=1):
        data.append(
            [
                i,
                r["country"],
                r["country_rank"],
                r["name"],
                r["points"],
                r["items"],
            ]
        )
    _write_sheet(
        out,
        "All dealers",
        headers,
        data,
        col_widths=(12, 14, 14, 40, 20, 18),
    )

    # Summary by country
    ws_sum = out.create_sheet("By country")
    ws_sum.append(["Country", "Dealers", "Avg points %", "Max points %"])
    for c in range(1, 5):
        ws_sum.cell(1, c).font = Font(bold=True)
    by_country: dict[str, list[float]] = {}
    for r in all_rows:
        by_country.setdefault(r["country"], []).append(r["points"])
    for country in sorted(by_country, key=str.casefold):
        vals = by_country[country]
        ws_sum.append(
            [
                country,
                len(vals),
                round(sum(vals) / len(vals), 2),
                max(vals),
            ]
        )
    ws_sum.column_dimensions["A"].width = 14
    ws_sum.column_dimensions["B"].width = 10
    ws_sum.column_dimensions["C"].width = 16
    ws_sum.column_dimensions["D"].width = 16

    meta = out.create_sheet("Info")
    meta.append(["Generated", str(RUN_DATE)])
    meta.append(["Dealers", len(all_rows)])
    meta.append(["Countries", len(by_country)])

    saved = _safe_save(out, OUTPUT_TOTAL)
    print(f"Wrote {saved}")
    print(f"  {len(all_rows)} dealers across {len(by_country)} countries")
    if missing:
        print("  missing sources:", ", ".join(missing), file=sys.stderr)
    return 0 if not missing else 1


def _read_all_midi_b1() -> list[dict]:
    from midi_fr_ranking import collect_midi_b1_dealers, discover_midi_cod_divs, midi_label

    rows: list[dict] = []
    for cod in discover_midi_cod_divs():
        label = midi_label(cod)
        for rank, row in enumerate(collect_midi_b1_dealers(cod), start=1):
            rows.append(
                {
                    "country": label,
                    "country_rank": rank,
                    "name": row["name"],
                    "points": row["points"],
                    "items": row["items"],
                    "scored": row["assessed_points"],
                    "max_pts": row["max_points"],
                    "file": row["file"],
                }
            )
    return rows


def build_b1_file() -> int:
    all_rows: list[dict] = []

    b1_path = _find_latest_b1_workbook()
    if b1_path:
        wb = load_workbook(b1_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            if SKIP_B1_TABS.match(sheet_name) or MIDI_B1_TAB_PREFIX.match(sheet_name):
                continue
            all_rows.extend(_read_b1_country(wb[sheet_name], sheet_name))
        wb.close()

    all_rows.extend(_read_all_midi_b1())

    if not all_rows:
        print(
            "No B.1 rows found. Run build_b1_ranking.py and build_midi_fr_ranking.py.",
            file=sys.stderr,
        )
        return 1

    all_rows.sort(
        key=lambda r: (-r["points"], r["country"].casefold(), r["name"].casefold()),
    )

    out = Workbook()
    out.remove(out.active)
    headers = [
        "Global rank",
        "Country",
        "Rank (country)",
        "Dealer",
        *B1_HEADERS[2:],
    ]
    data = []
    for i, r in enumerate(all_rows, start=1):
        data.append(
            [
                i,
                r["country"],
                r["country_rank"],
                r["name"],
                r["points"],
                r["items"],
                r["scored"],
                r["max_pts"],
                r["file"],
            ]
        )
    _write_sheet(
        out,
        "All dealers",
        headers,
        data,
        col_widths=(12, 14, 14, 40, 20, 18, 14, 14, 40),
    )

    ws_sum = out.create_sheet("By country")
    ws_sum.append(["Country", "Dealers", "Avg B.1 %", "Avg B.1 pts", "Max dealer B.1 %"])
    for c in range(1, 6):
        ws_sum.cell(1, c).font = Font(bold=True)
    by_country: dict[str, list[dict]] = {}
    for r in all_rows:
        by_country.setdefault(r["country"], []).append(r)
    for country in sorted(by_country, key=str.casefold):
        rows_c = by_country[country]
        pcts = [float(x["points"]) for x in rows_c]
        scored = [float(x["scored"]) for x in rows_c if x.get("scored") is not None]
        ws_sum.append(
            [
                country,
                len(rows_c),
                round(sum(pcts) / len(pcts), 2),
                round(sum(scored) / len(scored), 2) if scored else None,
                max(pcts),
            ]
        )

    meta = out.create_sheet("Info")
    meta.append(["Generated", str(RUN_DATE)])
    meta.append(["Source B1 workbook", b1_path.name])
    meta.append(["Dealers", len(all_rows)])
    meta.append(["Countries", len(by_country)])

    saved = _safe_save(out, OUTPUT_B1)
    print(f"Wrote {saved}")
    print(f"  {len(all_rows)} dealers across {len(by_country)} countries")
    return 0


def main() -> int:
    code = build_total_file()
    code_b1 = build_b1_file()
    return code if code else code_b1


if __name__ == "__main__":
    raise SystemExit(main())
