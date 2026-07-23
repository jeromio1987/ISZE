#!/usr/bin/env python3
"""Full IDS Kaizen study — dealers per country + country×section compare.

Primary focus: IDS scores and Excel question gaps (not GPS/PMS correlation).
Outputs under IDS_2026/reports/kaizen_ids/
"""

from __future__ import annotations

import json
import math
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Any

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_ids_result_report import (
    COL_FIRST_DEALER,
    COL_MAX,
    COL_TEXT,
    QUESTION_RE,
    ROW_META_START,
    load_matrix_rows_from_sheet,
)
from ids_paths import DATA_ROOT, IDS_RESULT_REPORTS_DIR
from ids_report_extract import build_dashboard_payload

ISZE = Path(__file__).resolve().parents[2]
CHART_JS = ISZE / "09_meetings" / "Preparation Meetings" / "PMS_Top10" / "assets" / "chart.umd.min.js"
OUT_DIR = DATA_ROOT / "reports" / "kaizen_ids"
_TEMPLATE = Path(__file__).resolve().parent / "ids_kaizen_full_template.html"
_APP_JS = Path(__file__).resolve().parent / "ids_kaizen_full_app.js"

DEALERS_PER_COUNTRY = 3
SECTIONS = ("A", "B", "C", "D", "E")
SECTION_LABELS = {
    "A": "A Premises",
    "B": "B Organisation",
    "C": "C Service",
    "D": "D Parts",
    "E": "E Vehicle sales",
}
RANKINGS_DIR = DATA_ROOT / "rankings"
TOTAL_RANKING = RANKINGS_DIR / "IDS_2026_TOTAL_ranking_all.xlsx"
B1_RANKING = RANKINGS_DIR / "IDS_2026_B1_ranking_all.xlsx"



def _f(x: Any) -> float | None:
    if x is None or x == "":
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    if math.isnan(v):
        return None
    return v


def _json_safe(payload: dict) -> str:
    def scrub(o: Any) -> Any:
        if isinstance(o, float) and (math.isnan(o) or math.isinf(o)):
            return None
        if isinstance(o, dict):
            return {k: scrub(v) for k, v in o.items()}
        if isinstance(o, list):
            return [scrub(v) for v in o]
        return o

    return json.dumps(scrub(payload), ensure_ascii=False, allow_nan=False)


def _report_path(country: str) -> Path | None:
    for p in IDS_RESULT_REPORTS_DIR.glob("2026 IDS Evaluation Result Report_*.xls"):
        stem = p.stem.replace("2026 IDS Evaluation Result Report_", "")
        if stem.upper() == country.upper() or country.lower() in stem.lower():
            return p
    return None


def extract_worst_questions(
    country: str, dealer_name: str, *, top_n: int = 10, section: str | None = None
) -> list[dict[str, Any]]:
    path = _report_path(country)
    if path is None:
        return [{"error": f"no report for {country}"}]
    rb = xlrd.open_workbook(str(path))
    sh = rb.sheet_by_index(0)
    col = None
    for c in range(COL_FIRST_DEALER, sh.ncols):
        name = str(sh.cell_value(ROW_META_START, c) or "").strip()
        if name.casefold() == dealer_name.casefold():
            col = c
            break
    if col is None:
        for c in range(COL_FIRST_DEALER, sh.ncols):
            name = str(sh.cell_value(ROW_META_START, c) or "").strip()
            if dealer_name.casefold() in name.casefold() or name.casefold() in dealer_name.casefold():
                col = c
                break
    if col is None:
        return [{"error": f"dealer not found: {dealer_name}"}]

    matrix = load_matrix_rows_from_sheet(sh)
    current_sec = None
    gaps: list[dict[str, Any]] = []
    for row in matrix:
        ref = (row.ref or "").strip()
        if row.kind == "category":
            m = re.match(r"^([A-E])\.", ref, re.I)
            current_sec = m.group(1).upper() if m else None
            continue
        if row.kind != "question" or row.sheet_row is None:
            continue
        if not QUESTION_RE.match(ref):
            continue
        if section and current_sec != section.upper():
            continue
        max_pts = row.max_points or _f(sh.cell_value(row.sheet_row, COL_MAX)) or 0.0
        score = _f(sh.cell_value(row.sheet_row, col))
        if score is None or max_pts <= 0:
            continue
        pct = score / max_pts
        text = row.text or str(sh.cell_value(row.sheet_row, COL_TEXT) or "")
        gaps.append(
            {
                "ref": ref,
                "section": current_sec,
                "text": text[:200],
                "score": score,
                "max": max_pts,
                "pct": pct,
                "gap_pts": max_pts - score,
            }
        )
    gaps.sort(key=lambda g: (g["pct"], -g["gap_pts"]))
    return gaps[:top_n]


# Rollup tabs that duplicate BE+NL dealers — skip for Kaizen focus picks.
SKIP_FOCUS_COUNTRIES = frozenset({"BENELUX"})


def pick_dealers_per_country(payload: dict, n: int = DEALERS_PER_COUNTRY) -> list[dict[str, Any]]:
    """Bottom-n dealers by total % per country (include totals_only for ITUK).

    Skips Benelux rollup (already covered by BE+NL) and unscored 0% shells
    (common in MIDI empty workbooks) so focus lists are actionable.
    """
    by_c: dict[str, list[dict]] = {}
    for d in payload["dealers"]:
        by_c.setdefault(d["country"], []).append(d)

    picked: list[dict[str, Any]] = []
    for country, dealers in sorted(by_c.items()):
        if country.upper() in SKIP_FOCUS_COUNTRIES:
            continue
        # Prefer real scores: drop blank/0% shells unless nothing else remains.
        scored = [
            d
            for d in dealers
            if d.get("total_pct") is not None and float(d.get("total_pct") or 0) > 0
        ]
        pool = scored if scored else dealers
        ranked = sorted(pool, key=lambda d: (d.get("total_pct") is None, d.get("total_pct") or 0))
        take = ranked[: min(n, len(ranked))]
        country_avg = payload["country_profiles"].get(country, {}).get("total_pct")
        for rank_i, d in enumerate(take, 1):
            # weakest section for this dealer (if available)
            weak_sec = None
            weak_pct = None
            for s in SECTIONS:
                sp = d.get("sections", {}).get(s, {}).get("pct")
                if sp is None:
                    continue
                if weak_pct is None or sp < weak_pct:
                    weak_pct = sp
                    weak_sec = s
            picked.append(
                {
                    "country": country,
                    "name": d["name"],
                    "id": d["id"],
                    "rank_in_country": rank_i,
                    "ids_total_pct": d.get("total_pct"),
                    "country_avg_pct": country_avg,
                    "vs_country_pp": (
                        None
                        if d.get("total_pct") is None or country_avg is None
                        else (d["total_pct"] - country_avg) * 100
                    ),
                    "totals_only": bool(d.get("totals_only")),
                    "sections": {
                        s: d.get("sections", {}).get(s, {}).get("pct") for s in SECTIONS
                    },
                    "weakest_section": weak_sec,
                    "weakest_section_pct": weak_pct,
                }
            )
    return picked


def country_section_matrix(payload: dict) -> list[dict[str, Any]]:
    rows = []
    network = payload.get("network", {}).get("sections_pct", {})
    for country, prof in sorted(payload["country_profiles"].items()):
        sec = prof.get("sections_pct") or {}
        row = {
            "country": country,
            "dealer_count": prof.get("dealer_count"),
            "total_pct": prof.get("total_pct"),
            "sections": {},
            "vs_network_pp": {},
            "weakest_section": None,
            "weakest_pct": None,
        }
        for s in SECTIONS:
            v = sec.get(s)
            row["sections"][s] = v
            nv = network.get(s)
            if v is not None and nv is not None:
                row["vs_network_pp"][s] = (v - nv) * 100
            if v is not None and (row["weakest_pct"] is None or v < row["weakest_pct"]):
                row["weakest_pct"] = v
                row["weakest_section"] = s
        rows.append(row)
    return rows


def country_underperform_flags(matrix: list[dict[str, Any]], threshold_pp: float = -8.0) -> list[dict[str, Any]]:
    """Side quest: countries where a section is clearly below network."""
    flags = []
    for row in matrix:
        for s in SECTIONS:
            pp = row.get("vs_network_pp", {}).get(s)
            if pp is not None and pp <= threshold_pp:
                flags.append(
                    {
                        "country": row["country"],
                        "section": s,
                        "section_label": SECTION_LABELS[s],
                        "country_pct": row["sections"].get(s),
                        "vs_network_pp": pp,
                        "ask": f"Why is {row['country']} underperforming on {SECTION_LABELS[s]} vs network?",
                    }
                )
    flags.sort(key=lambda x: x["vs_network_pp"])
    return flags


def _norm_name(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().casefold())


def _norm_country(s: str) -> str:
    return (s or "").strip().upper().replace(" ", "")


def load_ranking_lookup(path: Path, *, pct_col: int = 4) -> dict[tuple[str, str], dict[str, Any]]:
    """Map (country, dealer_name) → rank fields from TOTAL or B1 workbook."""
    if not path.is_file():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["All dealers"] if "All dealers" in wb.sheetnames else wb.active
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[3] is None:
            continue
        country = _norm_country(str(row[1] or ""))
        name = _norm_name(str(row[3]))
        if not country or not name:
            continue
        pct_raw = _f(row[pct_col])
        if pct_raw is not None and pct_raw > 1.5:
            pct_raw = pct_raw / 100.0
        out[(country, name)] = {
            "global_rank": row[0],
            "country_rank": row[2],
            "pct": pct_raw,
        }
    wb.close()
    return out


def _lookup_rank(
    lu: dict[tuple[str, str], dict[str, Any]], country: str, name: str
) -> dict[str, Any] | None:
    """Exact (country, name) only — no substring fuzzy (avoids Adam→Adam Purves)."""
    return lu.get((_norm_country(country), _norm_name(name)))


def attach_ranks(
    focus: list[dict[str, Any]],
    total_lu: dict[tuple[str, str], dict[str, Any]],
    b1_lu: dict[tuple[str, str], dict[str, Any]],
) -> None:
    for f in focus:
        t = _lookup_rank(total_lu, f["country"], f["name"])
        b = _lookup_rank(b1_lu, f["country"], f["name"])
        f["ids_global_rank"] = t.get("global_rank") if t else None
        f["ids_country_rank_total"] = t.get("country_rank") if t else f.get("rank_in_country")
        f["b1_country_rank"] = b.get("country_rank") if b else None
        f["b1_pct"] = b.get("pct") if b else None


def build_db_questions(f: dict[str, Any], n_in_country: int | None = None) -> list[dict[str, str]]:
    """Structured interview pack for the distributor about one weak dealer."""
    name = f["name"]
    country = f["country"]
    weak = f.get("weakest_section")
    weak_label = SECTION_LABELS.get(weak or "", weak or "n/a")
    ids_pct = _pct(f.get("ids_total_pct"))
    cr = f.get("ids_country_rank_total") or f.get("rank_in_country")
    n = n_in_country or "?"
    b1r = f.get("b1_country_rank")
    b1p = _pct(f.get("b1_pct")) if f.get("b1_pct") is not None else None

    qs: list[dict[str, str]] = [
        {
            "id": "AS_RANK",
            "theme": "Aftersales",
            "question": (
                f"Where does {name} rank in your {country} network for aftersales "
                f"(parts + labour revenue) last FY and YTD — top / mid / bottom third?"
            ),
        },
        {
            "id": "SALES_RANK",
            "theme": "Vehicle sales",
            "question": (
                f"Where does {name} rank for new (and used, if tracked) vehicle sales "
                f"in {country}? Are they a sales leader with weak standards, or weak on both?"
            ),
        },
        {
            "id": "PARTS_BUY",
            "theme": "Parts",
            "question": (
                f"Where does {name} rank for parts purchases from you (genuine + BVP)? "
                f"Any dealer-level first-fill / backorder / genuine-mix data you can share?"
            ),
        },
        {
            "id": "WORKSHOP",
            "theme": "Workshop",
            "question": (
                f"How does {name} compare on workshop throughput (ROs / vehicles serviced / month) "
                f"and retention vs your other dealers?"
            ),
        },
        {
            "id": "IDS_CONTEXT",
            "theme": "IDS",
            "question": (
                f"{name} is IDS Kaizen pick #{cr} of ~{n} in {country} "
                f"(total {ids_pct}% if scored; weakest section {weak_label}). "
                f"Do you agree this ranking matches commercial reality?"
            ),
        },
        {
            "id": "B1_PLAN",
            "theme": "Business plan",
            "question": (
                f"Does {name} have an agreed aftersales / parts business plan (IDS B.1) "
                f"with monthly review?"
                + (
                    f" (IDS B.1 country rank {b1r}, score {b1p}%.)"
                    if b1r is not None
                    else " What is their B.1 status?"
                )
            ),
        },
        {
            "id": "GAP_WHY",
            "theme": "Root cause",
            "question": (
                f"For the Excel gaps on {name} (see Gap_Detail): is the cause people, process, "
                f"stock, multi-brand priority, or data/self-score quality?"
            ),
        },
        {
            "id": "COMMERCIAL_VS_IDS",
            "theme": "Priority",
            "question": (
                f"If {name} is commercially important but IDS-weak: what support will you give "
                f"in 90 days? If commercially weak and IDS-weak: development, probation, or exit?"
            ),
        },
        {
            "id": "OWNER_CAP",
            "theme": "Ownership",
            "question": (
                f"Who owns the Kaizen CAP for {name} at dealer and at DB, and what is the "
                f"go/no-go check at day 90?"
            ),
        },
    ]
    if weak:
        qs.insert(
            5,
            {
                "id": "SECTION_FOCUS",
                "theme": f"Section {weak}",
                "question": (
                    f"Specifically on {weak_label}: why is {name} soft here, and what is the "
                    f"one countermeasure you will run in the next 90 days?"
                ),
            },
        )
    return qs


def write_register(path: Path, focus: list[dict], matrix: list[dict], flags: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kaizen_Dealers"
    headers = [
        "Country",
        "Rank in country (Kaizen)",
        "Dealer",
        "IDS total %",
        "IDS country rank (TOTAL file)",
        "IDS global rank",
        "B.1 country rank",
        "B.1 %",
        "Country avg %",
        "vs country pp",
        "Weakest section",
        "Weakest %",
        "A %",
        "B %",
        "C %",
        "D %",
        "E %",
        "Top Excel gaps",
        "Totals only?",
        "DB owner",
        "Due",
        "Status",
    ]
    hf = PatternFill("solid", fgColor="1A2B4A")
    hfont = Font(color="FFFFFF", bold=True)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.fill = hf
        c.font = hfont
        c.alignment = Alignment(wrap_text=True)

    for i, f in enumerate(focus, 2):
        gaps = f.get("gaps") or []
        gap_txt = "; ".join(
            f"{g['ref']} {g['score']:g}/{g['max']:g}" for g in gaps[:5] if "ref" in g
        )
        ws.cell(i, 1, f["country"])
        ws.cell(i, 2, f["rank_in_country"])
        ws.cell(i, 3, f["name"])
        ws.cell(i, 4, _pct(f.get("ids_total_pct")))
        ws.cell(i, 5, f.get("ids_country_rank_total"))
        ws.cell(i, 6, f.get("ids_global_rank"))
        ws.cell(i, 7, f.get("b1_country_rank"))
        ws.cell(i, 8, _pct(f.get("b1_pct")) if f.get("b1_pct") is not None else None)
        ws.cell(i, 9, _pct(f.get("country_avg_pct")))
        ws.cell(i, 10, round(f["vs_country_pp"], 1) if f.get("vs_country_pp") is not None else None)
        ws.cell(i, 11, f.get("weakest_section"))
        ws.cell(i, 12, _pct(f.get("weakest_section_pct")))
        for j, s in enumerate(SECTIONS):
            ws.cell(i, 13 + j, _pct(f.get("sections", {}).get(s)))
        ws.cell(i, 18, gap_txt)
        ws.cell(i, 19, "Y" if f.get("totals_only") else "N")
        ws.cell(i, 22, "Open")

    ws2 = wb.create_sheet("Country_Sections")
    ws2.append(
        ["Country", "Dealers", "Total %", "A %", "B %", "C %", "D %", "E %", "Weakest", "vs net A", "vs net B", "vs net C", "vs net D", "vs net E"]
    )
    for r in matrix:
        ws2.append(
            [
                r["country"],
                r["dealer_count"],
                _pct(r.get("total_pct")),
                *[_pct(r["sections"].get(s)) for s in SECTIONS],
                r.get("weakest_section"),
                *[
                    round(r["vs_network_pp"][s], 1) if s in r.get("vs_network_pp", {}) else None
                    for s in SECTIONS
                ],
            ]
        )

    ws3 = wb.create_sheet("Country_Ask_List")
    ws3.append(["Country", "Section", "Country %", "vs network pp", "Question for DB"])
    for fl in flags:
        ws3.append(
            [
                fl["country"],
                fl["section_label"],
                _pct(fl.get("country_pct")),
                round(fl["vs_network_pp"], 1),
                fl["ask"],
            ]
        )

    ws_db = wb.create_sheet("DB_Questions")
    ws_db.append(
        [
            "Country",
            "Dealer",
            "Q id",
            "Theme",
            "Question for DB",
            "DB answer",
            "Evidence / rank given",
            "Follow-up",
        ]
    )
    for f in focus:
        for q in f.get("db_questions") or []:
            ws_db.append(
                [
                    f["country"],
                    f["name"],
                    q.get("id"),
                    q.get("theme"),
                    q.get("question"),
                    "",
                    "",
                    "",
                ]
            )

    ws4 = wb.create_sheet("Gap_Detail")
    ws4.append(["Country", "Dealer", "Section", "Ref", "Score", "Max", "Pct %", "Question"])
    for f in focus:
        for g in f.get("gaps") or []:
            if "ref" not in g:
                continue
            ws4.append(
                [
                    f["country"],
                    f["name"],
                    g.get("section"),
                    g["ref"],
                    g["score"],
                    g["max"],
                    round(g["pct"] * 100, 1),
                    g.get("text", ""),
                ]
            )

    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 14
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _pct(v: Any) -> float | None:
    if v is None:
        return None
    return round(float(v) * 100, 1)


def render_html(payload: dict, chart_js: str) -> str:
    return (
        _TEMPLATE.read_text(encoding="utf-8")
        .replace("__BUILT_AT__", payload.get("built_at", ""))
        .replace("__CHART_JS__", chart_js)
        .replace("__DATA_JSON__", _json_safe(payload))
        .replace("__APP_JS__", _APP_JS.read_text(encoding="utf-8"))
    )


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    ids = build_dashboard_payload(IDS_RESULT_REPORTS_DIR)
    focus = pick_dealers_per_country(ids, DEALERS_PER_COUNTRY)
    matrix = country_section_matrix(ids)
    flags = country_underperform_flags(matrix, threshold_pp=-8.0)

    total_lu = load_ranking_lookup(TOTAL_RANKING, pct_col=4)
    b1_lu = load_ranking_lookup(B1_RANKING, pct_col=4)
    attach_ranks(focus, total_lu, b1_lu)

    dealers_per_country = {
        c: int(ids["country_profiles"][c].get("dealer_count") or 0)
        for c in ids["country_profiles"]
    }

    for f in focus:
        if f.get("totals_only"):
            f["gaps"] = [{"note": "ITUK totals-only — no question rows in result report"}]
            f["suggested"] = "Use consolidated ITUK workbook + local audit for Kaizen themes"
        else:
            sec = f.get("weakest_section")
            gaps = extract_worst_questions(f["country"], f["name"], top_n=8, section=sec)
            if not gaps or (len(gaps) == 1 and "error" in gaps[0]):
                gaps = extract_worst_questions(f["country"], f["name"], top_n=8)
            f["gaps"] = gaps
            f["suggested"] = (
                f"Fix weakest section {sec}: " + ", ".join(g["ref"] for g in gaps[:4] if "ref" in g)
                if sec
                else "Close top Excel gaps"
            )
        f["db_questions"] = build_db_questions(
            f, n_in_country=dealers_per_country.get(f["country"])
        )

    study = {
        "built_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "dealers_per_country": DEALERS_PER_COUNTRY,
        "scored_dealers": sum(1 for d in ids["dealers"] if not d.get("totals_only")),
        "total_dealers": len(ids["dealers"]),
        "countries": sorted(ids["country_profiles"].keys()),
        "section_labels": SECTION_LABELS,
        "network_sections": ids.get("network", {}).get("sections_pct", {}),
        "focus_dealers": focus,
        "country_matrix": matrix,
        "country_flags": flags,
    }

    (OUT_DIR / "study_payload.json").write_text(
        json.dumps(study, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )
    write_register(OUT_DIR / "IDS_Kaizen_Full_Register.xlsx", focus, matrix, flags)

    chart_js = CHART_JS.read_text(encoding="utf-8")
    chart_js = re.sub(r"\n//# sourceMappingURL=.*$", "", chart_js)
    html = render_html(study, chart_js)
    html_path = OUT_DIR / "IDS_Kaizen_Full_Study_2026.html"
    html_path.write_text(html, encoding="utf-8")

    downloads = Path.home() / "Downloads"
    if downloads.is_dir():
        (downloads / html_path.name).write_text(html, encoding="utf-8")
        import shutil

        shutil.copy2(OUT_DIR / "IDS_Kaizen_Full_Register.xlsx", downloads / "IDS_Kaizen_Full_Register.xlsx")

    print("Wrote", html_path)
    print("Wrote", OUT_DIR / "IDS_Kaizen_Full_Register.xlsx")
    print(f"Focus dealers: {len(focus)} ({DEALERS_PER_COUNTRY}/country × {len(ids['country_profiles'])} countries)")
    print(f"Country section flags (≤ -8pp vs network): {len(flags)}")
    for fl in flags[:8]:
        print(f"  {fl['country']} {fl['section']} {fl['vs_network_pp']:+.1f}pp")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
