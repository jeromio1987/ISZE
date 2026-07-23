"""
GPS Full Run — Clean methodology across all available sales data
================================================================
Exclusions (GPS only):
  - Material: TRANSPORT, LATE PAYMENT INTEREST, SERVICE FEE, COST DEMURRAGE
  - Material Number: 550047xxx / 550070xxx (Shell Rimula lubricants)
  - BillT: G2 (credit memo), RE (return), S1, S2 (pro-forma/intercompany)
  - SaTy: CR (credit order), ZDR (special debit memo)

Surcharge neutralisation (strip urgency premium from net value):
  - SOR  0%  → no change
  - RUSH +3% → net / 1.03
  - VOR  +7% → net / 1.07
  - SVOR +11%→ net / 1.11
  - PRD  -3% → net / 0.97

Fiscal Year = Apr–Mar (ISZE convention)
Data sources:
  - CY2023-2024.XLSX  → FY2024 (complete) + FY2025 Apr-Dec portion
  - CY2025 parquet    → FY2025 Jan-Mar portion + FY2026 Apr-Dec portion
  - sales_2026tillApril.XLSX → FY2026 Jan-Mar portion

DB identification:
  - Raw files: use Payer field; UMI/UTI split via Sold-to pt
  - Parquet: use sold_to_sap → mapped to payer via xref built from raw
"""

import pandas as pd
import numpy as np
from pathlib import Path
import sys
import openpyxl, csv, warnings
warnings.filterwarnings('ignore')

_TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(_TOOLS))
from _shared.paths import DATA_ROOT, DERIVED_DIR, RAW_SALES_DIR  # noqa: E402

BASE   = _TOOLS.parent
RAW_23 = RAW_SALES_DIR / "CY2023-2024.XLSX"
PAR_25 = DERIVED_DIR / "sales_fact.parquet"
RAW_26 = DATA_ROOT / "sales_2026tillApril.XLSX"
UIO_F  = DATA_ROOT / "06_uio" / "_CURRENT" / "ISUZU_UIO_Overview_2026.xlsx"
OUT    = DERIVED_DIR / "gps_clean_all_fy.csv"

SURCHARGE = {'SOR':1.00,'RUSH':1.03,'VOR':1.07,'SVOR':1.11,'PRD':0.97}

EXCL_MATERIAL  = {'TRANSPORT','LATE PAYMENT INTEREST'}
EXCL_MAT_PART  = {'SERVICE FEE','COST DEMURRAGE'}  # contains match
EXCL_SHELL_PFX = ('550047','550070')
EXCL_BILLT     = {'G2','RE','S1','S2'}
EXCL_SATY      = {'CR','ZDR'}

BP = {2027:35.28e6, 2028:37.53e6, 2029:41_530_410, 2030:46_385_705}  # FY29/30 = live PL Plan row 186 (was 42.56/48.27e6 pre-revision phantom)

# UIO max spend rates EUR/unit/yr
RATES = {
    89:250.0,    # IUK LCV
    2194:739.20, # ITUK CV/LD
    93:400.0,    # ISD-Lang Germany (blended LCV+LD)
    92:250.0,    # Imotors (Belgium/other)
    90:250.0,
    1797:250.0,  # UMI Israel LCV
    1799:739.20, # UTI Israel CV
}

# ── HELPERS ──────────────────────────────────────────────────────────────
def apply_exclusions(df, mat_col='Material', matnum_col='Material Number',
                     billt_col='BillT', saty_col='SaTy'):
    mat  = df[mat_col].astype(str).str.upper().str.strip()
    mnum = df[matnum_col].astype(str).str.strip() if matnum_col in df.columns else pd.Series('', index=df.index)
    bt   = df[billt_col].astype(str).str.upper().str.strip() if billt_col in df.columns else pd.Series('', index=df.index)
    st   = df[saty_col].astype(str).str.upper().str.strip()  if saty_col  in df.columns else pd.Series('', index=df.index)

    mask = (
        mat.isin(EXCL_MATERIAL) |
        mat.str.contains('SERVICE FEE|COST DEMURRAGE|LATE PAYMENT', na=False) |
        mnum.str.startswith(EXCL_SHELL_PFX) |
        bt.isin(EXCL_BILLT) |
        st.isin(EXCL_SATY)
    )
    return df[~mask].copy()

def neutralise_surcharge(df, net_col='net_value', saty_col='SaTy'):
    saty = df[saty_col].astype(str).str.upper().str.strip() if saty_col in df.columns else 'SOR'
    factor = saty.map(SURCHARGE).fillna(1.0)
    df = df.copy()
    df[net_col] = df[net_col] / factor
    return df

def add_fy(df, date_col='billing_date'):
    dt = pd.to_datetime(df[date_col], errors='coerce')
    # FY = Apr-Mar: if month >= 4 → FY = year+1, else FY = year
    df = df.copy()
    df['FY'] = dt.dt.year.where(dt.dt.month < 4, dt.dt.year + 1)
    return df

# ── 1. READ CY2023-2024 RAW ──────────────────────────────────────────────
print("Reading CY2023-2024 raw...")
wb = openpyxl.load_workbook(RAW_23, read_only=True, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
wb.close()

hdr = list(rows[0])
# Fix duplicate column names
seen = {}
clean_hdr = []
for h in hdr:
    h = str(h) if h else 'Unknown'
    if h in seen:
        seen[h] += 1
        clean_hdr.append(f"{h}_{seen[h]}")
    else:
        seen[h] = 0
        clean_hdr.append(h)

raw23 = pd.DataFrame(rows[1:], columns=clean_hdr)
raw23['billing_date'] = pd.to_datetime(raw23['Billing Date'], errors='coerce')
raw23['net_value']    = pd.to_numeric(raw23['Net value'], errors='coerce').fillna(0)
raw23['Payer']        = pd.to_numeric(raw23['Payer'], errors='coerce')
raw23['Sold-to pt']   = pd.to_numeric(raw23['Sold-to pt'], errors='coerce')

# UMI/UTI split: Payer=1797 + Sold-to=1799 → assign to UTI
raw23['db_payer'] = raw23.apply(
    lambda r: 1799 if (r['Payer']==1797 and r['Sold-to pt']==1799) else r['Payer'], axis=1
)

# Build xref: sold_to → payer (for parquet mapping)
xref = raw23.dropna(subset=['Sold-to pt','Payer']).groupby('Sold-to pt')['Payer'].agg(
    lambda x: x.mode()[0]).to_dict()
print(f"  xref built: {len(xref)} sold-to codes")

raw23 = apply_exclusions(raw23, mat_col='Material', matnum_col='Material Number',
                         billt_col='BillT', saty_col='SaTy')
raw23 = neutralise_surcharge(raw23, net_col='net_value', saty_col='SaTy')
raw23 = add_fy(raw23)
raw23 = raw23[['billing_date','FY','db_payer','net_value']].rename(columns={'db_payer':'db'})
print(f"  Clean rows: {len(raw23):,} | FYs: {sorted(raw23.FY.dropna().unique())}")

# ── 2. READ CY2025 PARQUET ──────────────────────────────────────────────
print("Reading CY2025 parquet...")
par = pd.read_parquet(PAR_25)
par25 = par[par['year']==2025].copy()
par25['billing_date'] = pd.to_datetime(par25['billing_date'], errors='coerce')
par25['net_value']    = pd.to_numeric(par25['net_sales_value'], errors='coerce').fillna(0)

# Map sold_to_sap → payer using xref from raw
par25['db'] = par25['sold_to_sap'].map(xref)

# Rename columns for exclusion function
par25['Material']          = par25['material'].astype(str)
par25['Material Number']   = par25['part_number'].astype(str)
par25['BillT']             = ''   # not in parquet — already excluded upstream for CY2025
par25['SaTy']              = par25['saty'].astype(str) if 'saty' in par25.columns else 'SOR'

par25 = apply_exclusions(par25)
par25 = neutralise_surcharge(par25, net_col='net_value', saty_col='SaTy')
par25 = add_fy(par25)
par25 = par25[['billing_date','FY','db','net_value']]
print(f"  Clean rows: {len(par25):,} | FYs: {sorted(par25.FY.dropna().unique())}")

# ── 3. READ FY2026 RAW (Jan-Mar 2026) ───────────────────────────────────
print("Reading sales_2026tillApril...")
wb2 = openpyxl.load_workbook(RAW_26, read_only=True, data_only=True)
ws2 = wb2.active
rows2 = list(ws2.iter_rows(values_only=True))
wb2.close()

hdr2 = [str(h) if h else f'Col{i}' for i,h in enumerate(rows2[0])]
raw26 = pd.DataFrame(rows2[1:], columns=hdr2)
raw26['billing_date'] = pd.to_datetime(raw26.get('Billing Date', raw26.iloc[:,0]), errors='coerce')
raw26['net_value']    = pd.to_numeric(raw26.get('Net value', raw26['Net value']), errors='coerce').fillna(0)

# Identify payer column
payer_col  = 'Payer'   if 'Payer'     in raw26.columns else [c for c in raw26.columns if 'ayer' in c][0]
soldto_col = 'Sold-to pt' if 'Sold-to pt' in raw26.columns else [c for c in raw26.columns if 'old' in c.lower()][0]
mat_col    = 'Material'   if 'Material'   in raw26.columns else 'Material'
billt_col  = 'BillT'     if 'BillT'       in raw26.columns else ''
saty_col   = 'SaTy'      if 'SaTy'        in raw26.columns else ''

raw26[payer_col]  = pd.to_numeric(raw26[payer_col], errors='coerce')
raw26[soldto_col] = pd.to_numeric(raw26[soldto_col], errors='coerce')
raw26['db_payer'] = raw26.apply(
    lambda r: 1799 if (r[payer_col]==1797 and r[soldto_col]==1799) else r[payer_col], axis=1
)

raw26_clean = apply_exclusions(raw26, mat_col=mat_col,
                                matnum_col='Material Number' if 'Material Number' in raw26.columns else mat_col,
                                billt_col=billt_col, saty_col=saty_col)
raw26_clean = neutralise_surcharge(raw26_clean, net_col='net_value', saty_col=saty_col)
raw26_clean = add_fy(raw26_clean)
raw26_clean = raw26_clean[['billing_date','FY','db_payer','net_value']].rename(columns={'db_payer':'db'})
raw26_clean = raw26_clean[raw26_clean['FY']==2026]  # Jan-Mar 2026 only
print(f"  Clean rows: {len(raw26_clean):,} | FYs: {sorted(raw26_clean.FY.dropna().unique())}")

# ── 4. COMBINE ──────────────────────────────────────────────────────────
print("Combining all sources...")
all_data = pd.concat([raw23, par25, raw26_clean], ignore_index=True)
all_data = all_data.dropna(subset=['FY','db'])
all_data['FY']  = all_data['FY'].astype(int)
all_data['db']  = all_data['db'].astype(int)

# GPS per DB per FY
gps_by_db_fy = all_data.groupby(['FY','db'])['net_value'].sum().reset_index()
gps_by_db_fy.columns = ['FY','db','clean_gps_eur']

print("\n=== Clean GPS per FY (all DBs combined) ===")
print(gps_by_db_fy.groupby('FY')['clean_gps_eur'].sum().apply(lambda x: f"€{x/1e6:.3f}M").to_string())

# ── 5. UIO MAX SPEND ────────────────────────────────────────────────────
print("\nLoading UIO v18...")

# UIO summary from memory (v18 key DBs 2025E-2030F)
# Payer code → UIO per year
uio_data = {
    2194: {2025:23068,2026:24100,2027:25200,2028:26400,2029:27300,2030:28209},  # ITUK
    89:   {2025:47783,2026:48900,2027:50000,2028:51000,2029:51900,2030:52804},  # IUK
    93:   {2025:13780,2026:14900,2027:16100,2028:17300,2029:18500,2030:19700},  # Germany
    1797: {2025:18542,2026:19000,2027:19400,2028:19800,2029:20100,2030:20339}, # UMI
    1799: {2025:11253,2026:10800,2027:10200,2028:9700, 2029:8500, 2030:7337},  # UTI
}

# For FY2024/2025 max spend use 2025E UIO (closest available)
def get_uio(db, fy):
    yr = min(max(fy, 2025), 2030)
    return uio_data.get(db, {}).get(yr, None)

def get_rate(db):
    return RATES.get(db, 250.0)

# ── 6. GPS% PER DB PER FY ───────────────────────────────────────────────
gps_by_db_fy['uio']      = gps_by_db_fy.apply(lambda r: get_uio(r['db'], r['FY']), axis=1)
gps_by_db_fy['rate']     = gps_by_db_fy['db'].map(get_rate)
gps_by_db_fy['max_spend'] = gps_by_db_fy['uio'] * gps_by_db_fy['rate']
gps_by_db_fy['gps_pct']  = gps_by_db_fy['clean_gps_eur'] / gps_by_db_fy['max_spend']

# Focus on FY2024/2025/2026
for fy in [2024,2025,2026]:
    sub = gps_by_db_fy[gps_by_db_fy.FY==fy].sort_values('clean_gps_eur', ascending=False)
    total = sub['clean_gps_eur'].sum()
    print(f"\n=== FY{fy} — Total: €{total/1e6:.3f}M ===")
    for _,r in sub.head(10).iterrows():
        gps_str = f"{r['gps_pct']*100:.1f}%" if pd.notna(r['gps_pct']) else "no UIO"
        print(f"  DB {int(r['db'])}: €{r['clean_gps_eur']/1e3:,.0f}K | GPS {gps_str}")

# ── 7. BASELINE: AVG GPS% FY2025+FY2026 ────────────────────────────────
print("\n=== Baseline: 2-year avg GPS% FY2025+FY2026 ===")
base = gps_by_db_fy[gps_by_db_fy['FY'].isin([2025,2026])].copy()
baseline = base.groupby('db').agg(
    total_sales=('clean_gps_eur','sum'),
    total_max  =('max_spend','sum')
).reset_index()
baseline['gps_baseline'] = (baseline['total_sales'] / baseline['total_max']).clip(upper=0.95)
baseline = baseline.dropna(subset=['gps_baseline'])
print(baseline[['db','total_sales','gps_baseline']].sort_values('total_sales',ascending=False).to_string())

# ── 8. OPTION A UPLIFT FY2027-2030 ─────────────────────────────────────
cis_proj = {2027:1.70e6, 2028:1.80e6, 2029:1.90e6, 2030:2.00e6}

print("\n=== Option A uplift per FY ===")
proj_rows = []
for yr in [2027,2028,2029,2030]:
    baseline[f'uio_{yr}']  = baseline['db'].apply(lambda d: get_uio(d, yr) or 0)
    baseline[f'ms_{yr}']   = baseline[f'uio_{yr}'] * baseline['db'].map(get_rate)
    max_t   = baseline[f'ms_{yr}'].sum()
    target  = BP[yr] - cis_proj[yr]
    wbase   = (baseline[f'ms_{yr}'] * baseline['gps_baseline']).sum()
    X       = (target - wbase) / max_t
    baseline[f'gps_{yr}']  = (baseline['gps_baseline'] + X).clip(upper=0.90)
    baseline[f'sales_{yr}']= baseline[f'ms_{yr}'] * baseline[f'gps_{yr}']
    total   = baseline[f'sales_{yr}'].sum() + cis_proj[yr]
    print(f"  FY{yr}: uplift X={X*100:+.2f}pp | projected=€{total/1e6:.2f}M | BP=€{BP[yr]/1e6:.2f}M | gap=€{(total-BP[yr])/1e3:+.0f}K")

# ── 9. SAVE ─────────────────────────────────────────────────────────────
gps_by_db_fy.to_csv(OUT, index=False)
baseline.to_csv(str(OUT).replace('.csv','_projection.csv'), index=False)
print(f"\nSaved: {OUT}")
print(f"Saved: {str(OUT).replace('.csv','_projection.csv')}")
