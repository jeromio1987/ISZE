#!/usr/bin/env python3
"""Run pre-delivery Excel QA across ISZE deliverable folders.

See 00_docs/EXCEL_PRE_DELIVERY_CHECK.md
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
ISZE = TOOLS.parent
IDS = TOOLS / "ids_dealer_ranking"

SKIP_PARTS = re.compile(
    r"(_archive|_Archive|~\$| - Copy|\.backup_|\\BE\\|\\NL\\|\\Poland\\|\\RSA\\|\\IUK\\|\\Israel\\|\\Hungary\\|\\Portugal\\|\\Romania\\|\\Petros\\|\\Gasanzammit\\|\\Loutsios\\|\\ITUK\\|\\Denmark\\|\\MIDI\\)",
    re.I,
)
DATED_SUFFIX = re.compile(r"_20\d{2}-\d{2}-\d{2}(?:_\w+)?\.xlsx$", re.I)


@dataclass
class SuiteResult:
    name: str
    kind: str = "data"  # data | format
    passed: int = 0
    failed: int = 0
    skipped: int = 0
    failures: list[str] = field(default_factory=list)


def _discover_xlsx(paths: list[Path], *, canonical_only: bool = False) -> list[Path]:
    out: list[Path] = []
    for base in paths:
        if not base.is_dir():
            if base.is_file() and base.suffix.lower() in {".xlsx", ".xlsm"}:
                out.append(base)
            continue
        for p in sorted(base.rglob("*.xlsx")):
            if SKIP_PARTS.search(str(p)):
                continue
            if canonical_only and DATED_SUFFIX.search(p.name):
                continue
            out.append(p)
        for p in sorted(base.rglob("*.xlsm")):
            if SKIP_PARTS.search(str(p)):
                continue
            if canonical_only and DATED_SUFFIX.search(p.name):
                continue
            out.append(p)
    return sorted(set(out))


def _run_pre_delivery(paths: list[Path], name: str) -> SuiteResult:
    res = SuiteResult(name, kind="format")
    if not paths:
        return res
    cmd = [sys.executable, str(TOOLS / "excel_pre_delivery_check.py"), "--json"]
    cmd.extend(str(p) for p in paths)
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode not in (0, 1) or not proc.stdout.strip():
        res.failed = len(paths)
        res.failures.append(proc.stderr or proc.stdout or "excel_pre_delivery_check failed")
        return res
    payload = json.loads(proc.stdout)
    for r in payload.get("results", []):
        if r.get("ok"):
            res.passed += 1
        else:
            res.failed += 1
            name = Path(r["path"]).name
            issues = r.get("sheets", [])
            detail = "; ".join(
                f"{i['check']}: {i['detail']}"
                for sh in issues
                for i in sh.get("issues", [])[:2]
            )
            res.failures.append(f"{name}: {detail or 'issues'}")
    return res


def _run_ids_logic() -> SuiteResult:
    res = SuiteResult("ids_logic_check", kind="data")
    proc = subprocess.run(
        [sys.executable, str(IDS / "ids_logic_check.py"), "--excel", "--json"],
        cwd=IDS,
        capture_output=True,
        text=True,
    )
    if proc.returncode not in (0, 1) or not proc.stdout.strip():
        res.failed = 1
        res.failures.append(proc.stderr or proc.stdout or "ids_logic_check failed")
        return res
    payload = json.loads(proc.stdout)
    for r in payload.get("results", []):
        if r.get("ok"):
            res.passed += 1
        else:
            res.failed += 1
            res.failures.append(
                f"{r['country']}: " + "; ".join(i["detail"] for i in r.get("issues", [])[:3])
            )
    return res


def _is_pct_header(label: str) -> bool:
    low = label.strip().lower()
    if "%" not in low:
        return False
    # Raw counts mislabelled in legacy ITUK exports.
    if "assessed points" in low or "(scored)" in low or "(max)" in low:
        return False
    return True


def _check_ranking_pct(paths: list[Path]) -> SuiteResult:
    """Flag % columns in ranking workbooks where values exceed 100%."""
    from openpyxl import load_workbook

    res = SuiteResult("ids_ranking_pct", kind="data")
    for path in paths:
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            res.failed += 1
            res.failures.append(f"{path.name}: cannot open ({exc})")
            continue
        bad: list[str] = []
        for ws in wb.worksheets:
            header_row = next(
                ws.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            )
            if not header_row:
                continue
            pct_cols: list[tuple[int, str]] = []
            for ci, val in enumerate(header_row):
                if isinstance(val, str) and _is_pct_header(val):
                    pct_cols.append((ci, val.strip()))
            if not pct_cols:
                continue
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row or 0, 500), values_only=True):
                for ci, label in pct_cols:
                    if ci >= len(row):
                        continue
                    val = row[ci]
                    if not isinstance(val, (int, float)):
                        continue
                    v = float(val)
                    if v <= 1.5:
                        if v > 1.001:
                            bad.append(f"{label}={v}")
                    elif v > 100.001:
                        bad.append(f"{label}={v}")
        wb.close()
        if bad:
            res.failed += 1
            res.failures.append(f"{path.name}: % column >100% sample {bad[:3]}")
        else:
            res.passed += 1
    return res


def main() -> int:
    inbox = ISZE / "01_data" / "90_INBOX" / "IDS_2026"

    suites: list[SuiteResult] = []

    # IDS country reports — full logic + layout + Excel formulas
    suites.append(_run_ids_logic())

    delta = _discover_xlsx([inbox / "reports" / "evaluation_result" / "_delta"])
    suites.append(_run_pre_delivery(delta, "ids_delta_format"))

    rankings = _discover_xlsx([inbox / "rankings"], canonical_only=True)
    suites.append(_run_pre_delivery(rankings, "ids_rankings_format"))
    suites.append(_check_ranking_pct(rankings))

    uio = _discover_xlsx(
        [
            ISZE / "06_uio" / "_CURRENT",
            ISZE / "06_uio" / "Ireland_SIMI_Comparison",
        ]
    )
    suites.append(_run_pre_delivery(uio, "uio_format"))

    autodoc = [
        p
        for p in _discover_xlsx([ISZE / "04_tools" / "autodoc_scraper"])
        if "backup" not in p.name.lower() and "_HANDOFF" not in str(p)
    ]
    suites.append(_run_pre_delivery(autodoc, "autodoc_format"))

    pms = [
        p
        for p in _discover_xlsx([ISZE / "03_analytics" / "Isuzu" / "pms-pilot"], canonical_only=True)
        if p.name.startswith(("ITUK_PMS", "ISZE_EU_PMS", "ISZE_ITUK"))
    ]
    suites.append(_run_pre_delivery(pms, "pms_format"))

    meetings = _discover_xlsx([ISZE / "09_meetings" / "Preparation Meetings"])
    suites.append(_run_pre_delivery(meetings, "meetings_format"))

    gps = list((ISZE / "04_tools").glob("ISZE_*GPS*.xlsx")) + list(
        (ISZE / "04_tools").glob("ISZE_SPA*.xlsx")
    )
    suites.append(_run_pre_delivery([p for p in gps if p.is_file()], "gps_spa_format"))

    data_suites = [s for s in suites if s.kind == "data"]
    fmt_suites = [s for s in suites if s.kind == "format"]
    data_fail = sum(s.failed for s in data_suites)
    fmt_fail = sum(s.failed for s in fmt_suites)

    print("ISZE Excel QA — full sweep")
    print("=" * 60)
    print("DATA LOGIC (must pass)")
    for s in data_suites:
        status = "PASS" if s.failed == 0 else "FAIL"
        print(f"  {s.name:26} {status:4}  passed={s.passed}  failed={s.failed}")
        for f in s.failures[:5]:
            print(f"    - {f}")
    print()
    print("FORMATTING (merge/wrap/extend — informational)")
    for s in fmt_suites:
        status = "PASS" if s.failed == 0 else "WARN"
        print(f"  {s.name:26} {status:4}  passed={s.passed}  failed={s.failed}")
        for f in s.failures[:3]:
            print(f"    - {f}")
        if len(s.failures) > 3:
            print(f"    ... +{len(s.failures) - 3} more")
    print("=" * 60)
    print(f"DATA:   {sum(s.passed for s in data_suites)} passed, {data_fail} failed")
    print(f"FORMAT: {sum(s.passed for s in fmt_suites)} passed, {fmt_fail} warnings")

    report = ISZE / "01_data" / "00_derived" / "excel_qa_report.json"
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(
        json.dumps(
            {
                "ok": data_fail == 0,
                "data_failed": data_fail,
                "format_warnings": fmt_fail,
                "passed": sum(s.passed for s in suites),
                "failed": sum(s.failed for s in suites),
                "suites": [
                    {
                        "name": s.name,
                        "passed": s.passed,
                        "failed": s.failed,
                        "failures": s.failures,
                    }
                    for s in suites
                ],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Report: {report}")
    return 1 if data_fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
