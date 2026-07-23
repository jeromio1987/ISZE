"""
ISZE branded deck builder - v2

What changed vs. the old isuzu-ppt-template skill:
  1. Real Isuzu logo (assets/isuzu_logo_red.png) instead of Arial-Black text
     pretending to be a logo.
  2. Charts are matplotlib PNGs from isze_chart_theme.py, not native pptx
     charts - inserted as pictures. Direct labels, brand colors, no chart junk.
  3. Added a "chart + takeaways" two-column layout, because a chart alone
     without a stated so-what is the other half of what makes decks look
     generic. Every content slide can carry 2-4 bullet takeaways next to
     the visual.
  4. Kept the same brand colors / dimensions as the original so it's a
     drop-in upgrade, not a redesign.

Usage: see build_demo_deck.py in this folder for a full worked example.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

RED       = RGBColor(0xCC, 0x0E, 0x1E)
LIGHT_RED = RGBColor(0xDC, 0x60, 0x6E)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
BLACK     = RGBColor(0x1A, 0x1A, 0x1A)
GRAY      = RGBColor(0x5A, 0x5A, 0x5A)
GRAY_BG   = RGBColor(0xEC, 0xEC, 0xEC)

SW = Inches(13.33)
SH = Inches(7.5)

LOGO_RED   = os.path.join(ASSETS, "isuzu_logo_red.png")          # 930x162, transparent
LOGO_GROUP = os.path.join(ASSETS, "isuzu_group_wordmark_black.png")  # 2500x310, transparent


# --------------------------------------------------------------- helpers --
def new_presentation():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_white_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def txt(slide, text, left, top, width, height, size=14, bold=False,
        color=BLACK, align=PP_ALIGN.LEFT, font="Calibri", wrap=True,
        anchor=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    if anchor:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return box


def bullet_list(slide, items, left, top, width, height, size=13,
                 color=BLACK, bullet_color=RED, line_spacing=1.25):
    """items: list[str]. Each rendered as '— text' with a red dash marker,
    matching the no-native-bullet-glyph look used in clean decks."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = line_spacing
        r1 = p.add_run()
        r1.text = "—  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet_color
        r1.font.name = "Calibri"
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.font.name = "Calibri"
    return box


def rect(slide, left, top, width, height, fill_color, line=False):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    if not line:
        s.line.fill.background()
    return s


def add_logo(slide, left=Inches(0.28), top=Inches(6.85), height=Inches(0.32)):
    """Real Isuzu logo image, aspect-locked."""
    slide.shapes.add_picture(LOGO_RED, left, top, height=height)


def add_footer_bar(slide):
    h = Inches(0.10)
    rect(slide, 0, SH - h, SW, h, RED)


def add_page_number(slide, n):
    txt(slide, str(n), SW - Inches(0.65), SH - Inches(0.42), Inches(0.4),
        Inches(0.3), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def add_picture_fit(slide, path, left, top, width, height):
    """Insert an image, preserving aspect ratio, centered within the box."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio = width / height
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        new_w = width
        new_h = int(width / img_ratio)
    else:
        new_h = height
        new_w = int(height * img_ratio)
    off_l = left + (width - new_w) // 2
    off_t = top + (height - new_h) // 2
    slide.shapes.add_picture(path, off_l, off_t, width=new_w, height=new_h)


# --------------------------------------------------------------- layouts --
def build_cover(prs, title, subtitle, eyebrow="ISZE  ·  VEHICLE PARTS & ACCESSORIES"):
    s = blank_slide(prs)
    set_white_bg(s)
    # accent block instead of the old fake-oval "wave"
    rect(s, 0, 0, Inches(0.22), SH, RED)
    txt(s, eyebrow, Inches(0.55), Inches(0.55), Inches(9), Inches(0.4),
        size=12, bold=True, color=RED, font="Calibri")
    txt(s, title, Inches(0.55), Inches(2.55), Inches(10.5), Inches(2.1),
        size=40, bold=True, color=BLACK)
    txt(s, subtitle, Inches(0.55), Inches(4.75), Inches(8.0), Inches(0.5),
        size=16, color=GRAY)
    s.shapes.add_picture(LOGO_RED, Inches(0.55), SH - Inches(1.05), height=Inches(0.42))
    add_footer_bar(s)
    return s


def build_content(prs, title, eyebrow=None, page_no=None):
    s = blank_slide(prs)
    set_white_bg(s)
    add_logo(s)
    add_footer_bar(s)
    y = Inches(0.22)
    if eyebrow:
        txt(s, eyebrow.upper(), Inches(0.32), y, Inches(12.5), Inches(0.3),
            size=11, bold=True, color=RED, font="Calibri")
        y = Inches(0.52)
    txt(s, title, Inches(0.32), y, Inches(12.7), Inches(0.62),
        size=25, bold=True, color=BLACK)
    rect(s, Inches(0.32), Inches(1.10), Inches(12.7), Emu(28000), RED)
    if page_no:
        add_page_number(s, page_no)
    return s


def build_two_col(prs, title, chart_path, takeaways, eyebrow=None, page_no=None,
                   takeaway_heading="Takeaways"):
    """Chart on the left (~62% width), 2-4 bullet takeaways on the right.
    This is the layout that turns 'here's a chart' into 'here's the story'."""
    s = build_content(prs, title, eyebrow=eyebrow, page_no=page_no)
    content_top = Inches(1.30)
    content_h = Inches(5.35)
    chart_w = Inches(7.85)
    add_picture_fit(s, chart_path, Inches(0.32), content_top, chart_w, content_h)

    col_left = Inches(8.45)
    col_w = Inches(4.55)
    rect(s, col_left, content_top, Emu(22000), content_h, GRAY_BG)
    txt(s, takeaway_heading.upper(), col_left + Inches(0.22), content_top + Inches(0.05),
        col_w - Inches(0.3), Inches(0.35), size=12, bold=True, color=RED)
    bullet_list(s, takeaways, col_left + Inches(0.22), content_top + Inches(0.55),
                col_w - Inches(0.4), content_h - Inches(0.6), size=12.5)
    return s


def build_full_chart(prs, title, chart_path, eyebrow=None, page_no=None):
    s = build_content(prs, title, eyebrow=eyebrow, page_no=page_no)
    add_picture_fit(s, chart_path, Inches(0.32), Inches(1.30), Inches(12.7), Inches(5.35))
    return s


def build_closing(prs):
    s = blank_slide(prs)
    set_white_bg(s)
    rect(s, 0, 0, Inches(0.22), SH, RED)
    txt(s, "Moving the World – for You", Inches(2.5), Inches(3.35),
        Inches(8.5), Inches(0.5), size=16, color=GRAY, align=PP_ALIGN.CENTER)
    s.shapes.add_picture(LOGO_RED, (SW - Inches(3.2)) // 2, Inches(2.7), width=Inches(3.2))
    return s
