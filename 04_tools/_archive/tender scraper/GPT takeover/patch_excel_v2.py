"""
Normalize an existing ISZE tender workbook into the v2 schema.

Run:
  python patch_excel_v2.py /path/to/isuzu_tenders.xlsx

Output:
  /path/to/isuzu_tenders_v2_patched.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl

from isze_core import normalize_raw_row, source_label_from_sheet
from tender_scraper_v2 import EXPORT_HEADERS, flatten_for_export


def patch_workbook(input_path: Path, output_path: Path) -> None:
    wb = openpyxl.load_workbook(input_path, data_only=True)
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    total_rows = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        values = list(ws.iter_rows(values_only=True))
        out_ws = out_wb.create_sheet(sheet_name)
        out_ws.append(EXPORT_HEADERS)
        if not values:
            continue
        headers = [str(h).strip() if h is not None else "" for h in values[0]]
        source_label = source_label_from_sheet(sheet_name)
        fixed = 0
        for raw_values in values[1:]:
            if not any(raw_values):
                continue
            raw: Dict[str, Any] = {headers[i]: raw_values[i] for i in range(min(len(headers), len(raw_values)))}
            if len(headers) == 1 and isinstance(raw_values[0], str) and raw_values[0].lower().startswith("no "):
                continue
            if str(raw.get("Title", "")).lower().startswith("no "):
                continue
            raw["_source"] = source_label
            row = flatten_for_export(normalize_raw_row(raw, source_label))
            out_ws.append([row.get(h, "") for h in EXPORT_HEADERS])
            fixed += 1
        total_rows += fixed
        print(f"  {sheet_name}: {fixed} rows normalized")

    out_wb.save(output_path)
    print(f"\n✓ Saved {output_path}")
    print(f"  Total normalized rows: {total_rows}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python patch_excel_v2.py /path/to/isuzu_tenders.xlsx [output.xlsx]")
    input_path = Path(sys.argv[1]).expanduser().resolve()
    if not input_path.exists():
        raise SystemExit(f"Input not found: {input_path}")
    if len(sys.argv) >= 3:
        output_path = Path(sys.argv[2]).expanduser().resolve()
    else:
        output_path = input_path.with_name(input_path.stem + "_v2_patched.xlsx")
    patch_workbook(input_path, output_path)


if __name__ == "__main__":
    main()
