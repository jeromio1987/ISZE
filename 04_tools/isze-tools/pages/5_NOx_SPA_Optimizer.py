import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
import numpy as np
from io import BytesIO
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
import os, sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

st.set_page_config(page_title="NOx SPA Optimizer", layout="wide", page_icon="🔧")

st.title("🔧 NOx Sensor SPA — Revenue Optimizer")
st.caption(
    "ITUK Special Price Agreement · Compare L12 actual vs. SPA request · "
    "Entity margin split (ISZE / IML / IMIT) · Breakeven qty & price per PN"
)

# ──────────────────────────────────────────────────────────────────────────────
# Data loading
# ──────────────────────────────────────────────────────────────────────────────
# __file__ = .../claude/Prive/jarvis/pages/5_NOx_SPA_Optimizer.py
# 4 x dirname up → .../claude/
_HERE = os.path.abspath(__file__)
_CLAUDE_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(_HERE))))
DATA_PATH = os.path.join(_CLAUDE_ROOT, "Isuzu", "data", "Nox Sensor SPA Requests_.xlsx")


def parse_eur(v):
    """Handle European decimal comma: '102,65' → 102.65"""
    if isinstance(v, str):
        return float(v.replace(",", ".").replace(" ", ""))
    return float(v) if v is not None else 0.0


@st.cache_data
def load_data(path: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path)
    ws = wb["Sheet4"]
    rows = []
    for r in ws.iter_rows(min_row=3, max_row=14, values_only=True):
        if r[0] is None:
            continue
        rows.append(
            {
                "PN":          str(int(r[0])),
                "Description": (r[1] or "").strip(),
                "Buy_Price":   parse_eur(r[2]),   # C – ISZE cost/unit (from IML/IMIT)
                "Sell_Price":  parse_eur(r[3]),   # D – ISZE IE40 sell price to ITUK
                "SPA_Qty":     int(r[5]) if r[5] else 0,    # F – ITUK requested SPA volume
                "SPA_Price":   float(r[6]) if r[6] else 0.0, # G – ITUK requested SPA price
                "L12_Qty":     int(r[9]) if r[9] else 0,    # J – actual last-12-month qty
                "L12_Cost":    float(r[15]) if r[15] else 0.0, # P – actual L12 total cost
                "L12_TO":      float(r[16]) if r[16] else 0.0, # Q – actual L12 total turnover
            }
        )
    df = pd.DataFrame(rows)
    df["L12_Profit"] = df["L12_TO"] - df["L12_Cost"]
    df["L12_GM_pct"] = df["L12_Profit"] / df["L12_TO"]
    df["L12_Cost_pu"] = df["L12_Cost"] / df["L12_Qty"]
    df["L12_Sell_pu"] = df["L12_TO"] / df["L12_Qty"]
    return df


try:
    base_df = load_data(DATA_PATH)
except FileNotFoundError:
    st.error(f"Data file not found: `{DATA_PATH}`")
    st.stop()
except Exception as exc:
    st.error(f"Could not load data: {exc}")
    st.stop()

# ──────────────────────────────────────────────────────────────────────────────
# Sidebar
# ──────────────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.header("⚙️ Parameters")

    # ── Entity support ──────────────────────────────────────────────────────
    st.markdown("### 🏢 Entity Margin Support")
    st.caption(
        "Each % reduces ISZE's buy price.  "
        "IML = Source J parts · IMIT = Source T parts.  "
        "Applied globally here; per-PN override available below."
    )

    iml_pct = st.slider("IML (Japan) BP reduction %", 0, 20, 0, 1,
                         help="IML reduces base price in JPY → ISZE EUR cost drops") / 100
    imit_pct = st.slider("IMIT (Thailand) BP reduction %", 0, 20, 0, 1,
                          help="IMIT reduces their transfer price → ISZE EUR cost drops") / 100
    total_supp = iml_pct + imit_pct

    if total_supp > 0:
        isze_carries = max(0.0, 1.0 - total_supp)
        st.info(
            f"**Entity support: {total_supp*100:.0f}%** of buy price  \n"
            f"IML {iml_pct*100:.0f}% + IMIT {imit_pct*100:.0f}%  \n"
            f"ISZE still absorbs: {isze_carries*100:.0f}% of original cost"
        )

    # ── Breakeven target ────────────────────────────────────────────────────
    st.markdown("---")
    st.markdown("### 🎯 Profit Target")
    tgt_mode = st.radio("Breakeven target", ["Match L12 profit", "Custom portfolio (€)"])
    custom_tgt = None
    if tgt_mode == "Custom portfolio (€)":
        custom_tgt = st.number_input(
            "Target portfolio profit (€)",
            value=int(base_df["L12_Profit"].sum()),
            step=500,
        )

    # ── Display ─────────────────────────────────────────────────────────────
    st.markdown("---")
    st.markdown("### 📊 Display")
    show_min_supp = st.checkbox("Show min entity support needed", value=True)
    show_raw = st.checkbox("Show raw Excel data", value=False)
    st.markdown("---")
    st.caption(f"Source: `Nox Sensor SPA Requests_.xlsx`")


# ──────────────────────────────────────────────────────────────────────────────
# Editable SPA proposal (st.data_editor)
# ──────────────────────────────────────────────────────────────────────────────
st.subheader("📥 SPA Proposal — Edit to Simulate")
st.caption(
    "SPA Price and SPA Qty are editable. "
    "Change them per PN and all results update instantly. "
    "Optional: set a custom entity support % per PN in the last column."
)

edit_input = base_df[
    ["PN", "Description", "Buy_Price", "Sell_Price", "SPA_Price", "SPA_Qty",
     "L12_Qty", "L12_TO", "L12_Profit"]
].copy()
edit_input["Entity_Override_pct"] = round(total_supp * 100, 0)  # default = sidebar value

col_cfg = {
    "PN":                 st.column_config.TextColumn("Part No.", disabled=True, width=120),
    "Description":        st.column_config.TextColumn("Description", disabled=True, width=220),
    "Buy_Price":          st.column_config.NumberColumn("Buy Price €", format="€%.2f", disabled=True),
    "Sell_Price":         st.column_config.NumberColumn("IE40 Price €", format="€%.2f", disabled=True),
    "SPA_Price":          st.column_config.NumberColumn("▶ SPA Price €", format="€%.2f", min_value=0.0, step=1.0),
    "SPA_Qty":            st.column_config.NumberColumn("▶ SPA Qty", format="%d", min_value=0, step=10),
    "L12_Qty":            st.column_config.NumberColumn("L12 Qty", disabled=True),
    "L12_TO":             st.column_config.NumberColumn("L12 Turnover €", format="€%.0f", disabled=True),
    "L12_Profit":         st.column_config.NumberColumn("L12 Profit €", format="€%.0f", disabled=True),
    "Entity_Override_pct": st.column_config.NumberColumn(
        "Entity % (override)",
        format="%.0f%%",
        min_value=0, max_value=40, step=1,
        help="Total entity BP reduction % for this PN. Overrides sidebar value.",
    ),
}

edited = st.data_editor(
    edit_input, column_config=col_cfg, use_container_width=True,
    num_rows="fixed", key="spa_editor", height=430
)

# ──────────────────────────────────────────────────────────────────────────────
# Core calculations
# ──────────────────────────────────────────────────────────────────────────────
df = edited.copy()

# Effective cost after per-PN entity support
df["Eff_Supp"] = df["Entity_Override_pct"] / 100
df["Eff_Cost"] = df["Buy_Price"] * (1 - df["Eff_Supp"])

# SPA scenario
df["SPA_TO"]     = df["SPA_Price"] * df["SPA_Qty"]
df["SPA_Cost"]   = df["Eff_Cost"]  * df["SPA_Qty"]
df["SPA_Profit"] = df["SPA_TO"]    - df["SPA_Cost"]
df["SPA_GM_pct"] = np.where(df["SPA_TO"] > 0, df["SPA_Profit"] / df["SPA_TO"], 0.0)
df["L12_GM_pct"] = df["L12_Profit"] / df["L12_TO"]
df["Disc_pct"]   = (df["Sell_Price"] - df["SPA_Price"]) / df["Sell_Price"]

# Deltas vs L12
df["Δ_TO"]     = df["SPA_TO"]     - df["L12_TO"]
df["Δ_Profit"] = df["SPA_Profit"] - df["L12_Profit"]

# Per-PN profit targets
if tgt_mode == "Match L12 profit":
    df["Target"] = df["L12_Profit"]
else:
    _share = df["L12_Profit"] / df["L12_Profit"].sum()
    df["Target"] = _share * custom_tgt

# Breakeven QTY  → at SPA price, how many units to hit target?
# Target = (SPA_Price - Eff_Cost) × q_be  →  q_be = Target / unit_margin
df["Unit_Margin"] = df["SPA_Price"] - df["Eff_Cost"]
df["BE_Qty"] = np.where(
    df["Unit_Margin"] > 0.01,
    np.ceil(df["Target"] / df["Unit_Margin"]),
    np.inf,
)

# Breakeven PRICE → at SPA qty, what price hits target?
# Target = (p_be - Eff_Cost) × SPA_Qty  →  p_be = Target/SPA_Qty + Eff_Cost
df["BE_Price"] = np.where(
    df["SPA_Qty"] > 0,
    df["Target"] / df["SPA_Qty"] + df["Eff_Cost"],
    np.inf,
)

# Min entity support (% of buy price) to make each PN profit-neutral at SPA qty — no support assumed
# (SPA_Price - Buy_Price*(1-x)) * SPA_Qty = L12_Profit
# x = 1 - (SPA_Price - L12_Profit/SPA_Qty) / Buy_Price
df["Min_Supp_Needed"] = np.where(
    (df["SPA_Price"] - df["Buy_Price"]) * df["SPA_Qty"] >= df["L12_Profit"],
    0.0,
    np.clip(
        1 - (df["SPA_Price"] - df["L12_Profit"] / df["SPA_Qty"]) / df["Buy_Price"],
        0, 0.999,
    ),
)

# Revenue already positive at SPA qty + zero support?
df["No_Supp_OK"] = (df["SPA_Price"] - df["Buy_Price"]) * df["SPA_Qty"] >= df["L12_Profit"]


def get_status(r):
    if r["SPA_Price"] <= r["Eff_Cost"]:
        return "❌ Below Cost"
    if r["Δ_Profit"] >= 0:
        return "✅ Profit ↑"
    if r["Δ_TO"] >= 0:
        return "⚠️ TO ↑ / Profit ↓"
    return "🔴 Both ↓"


df["Status"] = df.apply(get_status, axis=1)

# ──────────────────────────────────────────────────────────────────────────────
# Portfolio summary metrics
# ──────────────────────────────────────────────────────────────────────────────
st.subheader("📦 Portfolio Summary")

l12_to  = df["L12_TO"].sum()
l12_p   = df["L12_Profit"].sum()
spa_to  = df["SPA_TO"].sum()
spa_p   = df["SPA_Profit"].sum()
d_to    = spa_to - l12_to
d_p     = spa_p  - l12_p
n_pos   = int((df["Δ_Profit"] >= 0).sum())
max_supp_needed = df["Min_Supp_Needed"].max()
worst_pn = df.loc[df["Min_Supp_Needed"].idxmax(), "PN"]

c1, c2, c3, c4, c5, c6 = st.columns(6)
c1.metric("L12 Turnover",      f"€{l12_to:,.0f}")
c2.metric("L12 Profit",        f"€{l12_p:,.0f}",  f"GM {l12_p/l12_to*100:.1f}%")
c3.metric("SPA Turnover",      f"€{spa_to:,.0f}",  f"{d_to:+,.0f}",
          delta_color="normal" if d_to >= 0 else "inverse")
c4.metric("SPA Profit",        f"€{spa_p:,.0f}",   f"{d_p:+,.0f}",
          delta_color="normal" if d_p >= 0 else "inverse")
c5.metric("PNs profit-positive", f"{n_pos} / {len(df)}")
c6.metric("Max supp. needed",  f"{max_supp_needed*100:.1f}%",
          f"PN {worst_pn} (hardest)", delta_color="off")

# ──────────────────────────────────────────────────────────────────────────────
# Detailed results table
# ──────────────────────────────────────────────────────────────────────────────
st.subheader("📋 Full PN Results")

_STATUS_COLORS = {
    "✅ Profit ↑":       "background-color:#C6EFCE;color:#276221",
    "⚠️ TO ↑ / Profit ↓": "background-color:#FFEB9C;color:#9C6500",
    "🔴 Both ↓":         "background-color:#FFC7CE;color:#9C0006",
    "❌ Below Cost":     "background-color:#FF0000;color:#FFFFFF",
}

result_cols = [
    "PN", "Description",
    "Buy_Price", "Eff_Cost", "Sell_Price", "Disc_pct",
    "L12_Qty",  "L12_TO",     "L12_Profit",  "L12_GM_pct",
    "SPA_Qty",  "SPA_Price",  "SPA_TO",      "SPA_Profit", "SPA_GM_pct",
    "Δ_TO",     "Δ_Profit",
    "BE_Qty",   "BE_Price",
    "Min_Supp_Needed",
    "Status",
]

_fmt = {
    "Buy_Price":       "€{:.2f}", "Eff_Cost":    "€{:.2f}", "Sell_Price":  "€{:.2f}",
    "SPA_Price":       "€{:.2f}", "BE_Price":    "€{:.2f}",
    "Disc_pct":        "{:.1%}",  "L12_GM_pct":  "{:.1%}",  "SPA_GM_pct":  "{:.1%}",
    "Min_Supp_Needed": "{:.1%}",
    "L12_TO":          "€{:,.0f}","L12_Profit":  "€{:,.0f}",
    "SPA_TO":          "€{:,.0f}","SPA_Profit":  "€{:,.0f}",
    "Δ_TO":            "€{:+,.0f}","Δ_Profit":   "€{:+,.0f}",
    "BE_Qty":          "{:,.0f}",
}


def _style_status(val):
    return _STATUS_COLORS.get(val, "")


def _style_delta(val):
    if isinstance(val, (int, float)) and not np.isinf(val):
        return "color:green;font-weight:bold" if val > 0 else "color:red" if val < 0 else ""
    return ""


styled_result = (
    df[result_cols]
    .style
    .applymap(_style_status, subset=["Status"])
    .applymap(_style_delta,  subset=["Δ_TO", "Δ_Profit"])
    .format(_fmt, na_rep="—")
)

st.dataframe(styled_result, use_container_width=True, height=430)

# ──────────────────────────────────────────────────────────────────────────────
# Tabs: Profit Curves / Breakeven / Entity Impact
# ──────────────────────────────────────────────────────────────────────────────
tab1, tab2, tab3 = st.tabs(
    ["📈 Profit Curves", "🎯 Breakeven — Qty & Price", "🏢 Entity Support Impact"]
)

# ── Tab 1 — Profit curves (per PN deep-dive) ───────────────────────────────
with tab1:
    sel_pn = st.selectbox(
        "Select PN to inspect",
        df["PN"].tolist(),
        format_func=lambda x: f"{x}  —  {df.loc[df['PN']==x, 'Description'].iloc[0][:35]}",
    )
    row = df[df["PN"] == sel_pn].iloc[0]

    q_be   = row["BE_Qty"]   if row["BE_Qty"]   != np.inf else 0
    q_max  = max(row["L12_Qty"], row["SPA_Qty"], q_be) * 1.55
    if q_max <= 0:
        q_max = 500

    qs = np.linspace(0, q_max, 400)

    # SPA profit curve (with entity support)
    spa_profits_curve = (row["SPA_Price"] - row["Eff_Cost"]) * qs
    # IE40 profit curve (no discount, ISZE cost = Buy_Price)
    ie40_profits_curve = (row["Sell_Price"] - row["Buy_Price"]) * qs

    fig1 = go.Figure()

    # IE40 reference
    fig1.add_trace(go.Scatter(
        x=qs, y=ie40_profits_curve,
        name=f"IE40 price (€{row['Sell_Price']:.2f}) — no discount",
        line=dict(color="#5B9BD5", width=2, dash="dot"),
    ))
    # SPA curve
    fig1.add_trace(go.Scatter(
        x=qs, y=spa_profits_curve,
        name=f"SPA price (€{row['SPA_Price']:.2f}) + entity support {row['Eff_Supp']*100:.0f}%",
        line=dict(color="#ED7D31", width=2.5),
        fill="tozeroy",
        fillcolor="rgba(237,125,49,0.08)",
    ))

    # L12 profit target (horizontal)
    fig1.add_hline(
        y=row["L12_Profit"],
        line=dict(color="#1F497D", dash="dash", width=2),
        annotation_text=f"L12 Profit Target: €{row['L12_Profit']:,.0f}",
        annotation_position="bottom right",
    )
    fig1.add_hline(y=0, line=dict(color="black", width=1))

    # L12 actual point
    fig1.add_trace(go.Scatter(
        x=[row["L12_Qty"]], y=[row["L12_Profit"]],
        mode="markers+text",
        marker=dict(size=13, color="#1F497D", symbol="circle"),
        text=["L12 actual"], textposition="top right",
        name=f"L12 actual ({row['L12_Qty']} units, €{row['L12_Profit']:,.0f})",
    ))

    # SPA requested point
    fig1.add_trace(go.Scatter(
        x=[row["SPA_Qty"]], y=[row["SPA_Profit"]],
        mode="markers+text",
        marker=dict(size=14, color="#ED7D31", symbol="star"),
        text=["SPA request"], textposition="top right",
        name=f"SPA request ({row['SPA_Qty']} units, €{row['SPA_Profit']:,.0f})",
    ))

    # Breakeven point
    if row["BE_Qty"] != np.inf and row["BE_Qty"] <= q_max:
        fig1.add_trace(go.Scatter(
            x=[row["BE_Qty"]], y=[row["Target"]],
            mode="markers+text",
            marker=dict(size=14, color="red", symbol="x"),
            text=["Breakeven"], textposition="top right",
            name=f"BE qty = {row['BE_Qty']:,.0f} units",
        ))
        fig1.add_vline(
            x=row["BE_Qty"],
            line=dict(color="red", dash="dash", width=1),
            annotation_text=f"BE: {row['BE_Qty']:,.0f}",
        )

    fig1.update_layout(
        title=dict(
            text=f"{row['PN']} — {row['Description']}",
            font=dict(size=14),
        ),
        xaxis_title="Quantity (units / year)",
        yaxis_title="Profit (€)",
        height=460,
        legend=dict(orientation="h", yanchor="bottom", y=1.02),
        hovermode="x unified",
    )
    st.plotly_chart(fig1, use_container_width=True)

    # Key metrics for selected PN
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("SPA Profit",   f"€{row['SPA_Profit']:,.0f}",
              f"{row['Δ_Profit']:+,.0f} vs L12")
    m2.metric("SPA GM%",      f"{row['SPA_GM_pct']*100:.1f}%",
              f"{(row['SPA_GM_pct']-row['L12_GM_pct'])*100:+.1f}pp vs L12")
    m3.metric("SPA Discount", f"{row['Disc_pct']*100:.1f}%",
              f"€{row['Sell_Price']:.2f} → €{row['SPA_Price']:.2f}")
    m4.metric("BE Qty needed",
              f"{row['BE_Qty']:,.0f}" if row["BE_Qty"] != np.inf else "—",
              f"vs {row['SPA_Qty']} committed ({'+' if row['BE_Qty']<=row['SPA_Qty'] else ''}"
              f"{row['SPA_Qty']-row['BE_Qty']:+.0f})"
              if row["BE_Qty"] != np.inf else "negative margin")
    m5.metric("BE Price needed", f"€{row['BE_Price']:.2f}",
              f"vs €{row['SPA_Price']:.2f} ({row['BE_Price']-row['SPA_Price']:+.2f})")

# ── Tab 2 — Breakeven: qty & price ────────────────────────────────────────
with tab2:
    st.markdown("#### How far is each PN from being revenue-positive at the requested deal?")

    col_a, col_b = st.columns(2)

    with col_a:
        st.markdown("**Quantity — what ITUK must commit vs. what makes us whole**")
        valid = df[df["BE_Qty"] != np.inf].copy()
        valid_sorted = valid.sort_values("BE_Qty_factor" if "BE_Qty_factor" in valid.columns else "PN")

        # Compute factor for ordering
        valid["BE_factor"] = valid["BE_Qty"] / valid["SPA_Qty"]

        fig_qty = go.Figure()
        fig_qty.add_trace(go.Bar(
            name="L12 Actual Qty",
            x=valid["PN"], y=valid["L12_Qty"],
            marker_color="#5B9BD5", opacity=0.7,
        ))
        fig_qty.add_trace(go.Bar(
            name="SPA Committed Qty",
            x=valid["PN"], y=valid["SPA_Qty"],
            marker_color="#ED7D31", opacity=0.9,
        ))
        fig_qty.add_trace(go.Bar(
            name="BE Qty (needed to match L12 profit)",
            x=valid["PN"], y=valid["BE_Qty"],
            marker_color="#FF4444", opacity=0.55,
        ))
        fig_qty.update_layout(
            barmode="group", height=360,
            yaxis_title="Units / year",
            legend=dict(orientation="h", yanchor="bottom", y=1.0),
            margin=dict(t=20),
        )
        st.plotly_chart(fig_qty, use_container_width=True)

        # Gap table
        gap = valid[["PN", "L12_Qty", "SPA_Qty", "BE_Qty"]].copy()
        gap["Gap"] = gap["BE_Qty"] - gap["SPA_Qty"]
        gap["Verdict"] = gap["Gap"].apply(
            lambda g: "✅ Committed qty sufficient"
            if g <= 0
            else f"❌ Need +{g:.0f} more units ({g/gap.loc[gap['Gap']==g, 'SPA_Qty'].iloc[0]*100:.0f}%)"
        )
        st.dataframe(
            gap[["PN", "L12_Qty", "SPA_Qty", "BE_Qty", "Gap", "Verdict"]]
            .style.format({"BE_Qty": "{:.0f}", "Gap": "{:+.0f}"}),
            use_container_width=True, height=350,
        )

    with col_b:
        st.markdown("**Price — what price makes the deal neutral at SPA committed qty**")
        fig_price = go.Figure()
        fig_price.add_trace(go.Bar(
            name="IE40 Price (current)",
            x=df["PN"], y=df["Sell_Price"],
            marker_color="#5B9BD5", opacity=0.7,
        ))
        fig_price.add_trace(go.Bar(
            name="BE Price needed",
            x=df["PN"],
            y=df["BE_Price"].replace(np.inf, np.nan).clip(upper=df["Sell_Price"] * 1.6),
            marker_color="#FF4444", opacity=0.6,
        ))
        fig_price.add_trace(go.Bar(
            name="SPA Price requested",
            x=df["PN"], y=df["SPA_Price"],
            marker_color="#ED7D31", opacity=0.95,
        ))
        fig_price.add_trace(go.Scatter(
            name="Effective ISZE cost",
            x=df["PN"], y=df["Eff_Cost"],
            mode="lines+markers",
            line=dict(color="black", dash="dot", width=2),
            marker=dict(symbol="diamond", size=8),
        ))
        fig_price.update_layout(
            barmode="group", height=360,
            yaxis_title="Price (€)",
            legend=dict(orientation="h", yanchor="bottom", y=1.0),
            margin=dict(t=20),
        )
        st.plotly_chart(fig_price, use_container_width=True)

        price_gap = df[["PN", "Eff_Cost", "SPA_Price", "BE_Price"]].copy()
        price_gap["Gap €"] = price_gap["BE_Price"].replace(np.inf, np.nan) - price_gap["SPA_Price"]
        price_gap["Verdict"] = price_gap["Gap €"].apply(
            lambda g: "✅ SPA price sufficient"
            if isinstance(g, float) and g <= 0
            else (f"❌ Need +€{g:.2f}/unit" if isinstance(g, float) else "❌ Selling below cost")
        )
        st.dataframe(
            price_gap[["PN", "Eff_Cost", "SPA_Price", "BE_Price", "Gap €", "Verdict"]]
            .style.format({
                "Eff_Cost":  "€{:.2f}", "SPA_Price": "€{:.2f}",
                "BE_Price":  "€{:.2f}", "Gap €":     "€{:+.2f}",
            }, na_rep="—"),
            use_container_width=True, height=350,
        )

# ── Tab 3 — Entity support impact ─────────────────────────────────────────
with tab3:
    st.markdown(
        "#### Minimum BP reduction ISZE needs from IML / IMIT to make each PN profit-neutral at SPA qty"
    )

    col_c, col_d = st.columns(2)

    with col_c:
        st.markdown("**Per-PN minimum support needed (% of buy price)**")
        min_s = df[["PN", "Min_Supp_Needed", "No_Supp_OK"]].copy().sort_values("Min_Supp_Needed")

        bar_colors = []
        for _, r in min_s.iterrows():
            if r["No_Supp_OK"]:
                bar_colors.append("#C6EFCE")
            elif r["Min_Supp_Needed"] <= total_supp:
                bar_colors.append("#FFEB9C")
            else:
                bar_colors.append("#FFC7CE")

        fig_min = go.Figure()
        fig_min.add_trace(go.Bar(
            x=min_s["Min_Supp_Needed"] * 100,
            y=min_s["PN"],
            orientation="h",
            marker_color=bar_colors,
            marker_line_color="gray", marker_line_width=0.5,
            text=[f"{v*100:.1f}%" for v in min_s["Min_Supp_Needed"]],
            textposition="outside",
        ))
        if total_supp > 0:
            fig_min.add_vline(
                x=total_supp * 100,
                line=dict(color="#1F497D", dash="dash", width=2),
                annotation_text=f"Current support {total_supp*100:.0f}%",
                annotation_position="top right",
            )
        fig_min.update_layout(
            height=400,
            xaxis_title="Buy price reduction needed (%)",
            yaxis_title="Part Number",
            showlegend=False,
            margin=dict(l=90, r=60),
        )
        st.plotly_chart(fig_min, use_container_width=True)

        # Legend
        st.markdown(
            "🟩 Already OK without support  "
            "🟨 Covered by current support  "
            "🟥 Still not covered — needs more"
        )

    with col_d:
        st.markdown("**Portfolio profit sensitivity to total entity BP reduction**")
        st.caption("How does total SPA portfolio profit change as IML+IMIT collectively lower their prices?")

        supp_range = np.linspace(0, 0.30, 61)
        portfolio_profits = [
            float(((df["SPA_Price"] - df["Buy_Price"] * (1 - s)) * df["SPA_Qty"]).sum())
            for s in supp_range
        ]
        l12_total_p = float(df["L12_Profit"].sum())

        fig_sens = go.Figure()
        fig_sens.add_trace(go.Scatter(
            x=supp_range * 100,
            y=portfolio_profits,
            name="SPA portfolio profit",
            line=dict(color="#ED7D31", width=2.5),
            fill="tozeroy",
            fillcolor="rgba(237,125,49,0.10)",
        ))
        fig_sens.add_hline(
            y=l12_total_p,
            line=dict(color="#1F497D", dash="dash", width=2.5),
            annotation_text=f"L12 profit benchmark: €{l12_total_p:,.0f}",
            annotation_position="bottom right",
        )
        fig_sens.add_hline(y=0, line=dict(color="black", width=1))

        # Current setting marker
        curr_p = float(((df["SPA_Price"] - df["Buy_Price"] * (1 - total_supp)) * df["SPA_Qty"]).sum())
        fig_sens.add_trace(go.Scatter(
            x=[total_supp * 100], y=[curr_p],
            mode="markers",
            marker=dict(size=13, color="red", symbol="diamond"),
            name=f"Current ({total_supp*100:.0f}%): €{curr_p:,.0f}",
        ))

        fig_sens.update_layout(
            height=400,
            xaxis_title="Total entity BP reduction (%)",
            yaxis_title="Portfolio profit (€)",
            legend=dict(orientation="h", yanchor="bottom", y=1.02),
        )
        st.plotly_chart(fig_sens, use_container_width=True)

        # Find exact % where portfolio profit = L12 profit
        try:
            be_supp_pct = float(np.interp(l12_total_p, portfolio_profits, supp_range * 100))
        except Exception:
            be_supp_pct = float("nan")

        if not np.isnan(be_supp_pct):
            st.metric(
                "Entity support to match L12 portfolio profit",
                f"{be_supp_pct:.1f}%",
                f"e.g. IML {be_supp_pct/2:.1f}% + IMIT {be_supp_pct/2:.1f}% (50/50)",
                delta_color="off",
            )
        else:
            st.info("Portfolio already matches L12 profit without entity support.")

        # No-support baseline
        no_supp_p = float(((df["SPA_Price"] - df["Buy_Price"]) * df["SPA_Qty"]).sum())
        st.metric(
            "SPA profit with ZERO entity support",
            f"€{no_supp_p:,.0f}",
            f"{no_supp_p - l12_total_p:+,.0f} vs L12",
            delta_color="normal" if no_supp_p >= l12_total_p else "inverse",
        )

# ──────────────────────────────────────────────────────────────────────────────
# Optional: raw data
# ──────────────────────────────────────────────────────────────────────────────
if show_raw:
    st.subheader("🗂️ Raw Excel Data")
    st.dataframe(base_df, use_container_width=True)

# ──────────────────────────────────────────────────────────────────────────────
# Excel export
# ──────────────────────────────────────────────────────────────────────────────
st.subheader("⬇️ Export")


def build_excel(df_in: pd.DataFrame, iml: float, imit: float) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SPA Analysis"

    ISZE_BLUE = PatternFill("solid", fgColor="003087")
    H_L12     = PatternFill("solid", fgColor="4472C4")
    H_SPA     = PatternFill("solid", fgColor="70AD47")
    H_DELTA   = PatternFill("solid", fgColor="ED7D31")
    H_REF     = PatternFill("solid", fgColor="1F497D")
    W_FONT    = Font(color="FFFFFF", bold=True)
    BOLD      = Font(bold=True)
    CENTER    = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Row 1 – title
    ws.merge_cells("A1:V1")
    ws["A1"] = (
        f"NOx Sensor SPA — Revenue Analysis  |  ITUK  |  "
        f"IML support: {iml*100:.0f}%  |  IMIT support: {imit*100:.0f}%  |  "
        f"Effective total: {(iml+imit)*100:.0f}%"
    )
    ws["A1"].font  = Font(bold=True, size=12, color="FFFFFF")
    ws["A1"].fill  = ISZE_BLUE
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 22

    # Row 2 – section labels (merged)
    section_map = {
        "A2:B2": ("REF", H_REF),
        "C2:F2": ("ISZE Cost Reference", H_REF),
        "G2:J2": ("L12 Actual", H_L12),
        "K2:O2": ("SPA Scenario (with entity support)", H_SPA),
        "P2:Q2": ("Delta", H_DELTA),
        "R2:S2": ("Breakeven", H_DELTA),
        "T2:U2": ("Min Entity Support", H_DELTA),
        "V2:V2": ("Status", H_REF),
    }
    for rng, (label, fill) in section_map.items():
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value     = label
        c.fill      = fill
        c.font      = W_FONT
        c.alignment = CENTER

    # Row 3 – column headers
    headers = [
        "Part No.", "Description",
        "Buy Price €", "Eff. Cost €", "IE40 Price €", "Discount %",
        "L12 Qty", "L12 Turnover €", "L12 Profit €", "L12 GM%",
        "SPA Qty", "SPA Price €", "SPA Turnover €", "SPA Profit €", "SPA GM%",
        "Δ Turnover €", "Δ Profit €",
        "BE Qty", "BE Price €",
        "Min Supp %", "No Supp OK?",
        "Status",
    ]
    col_fills = (
        [H_REF]*2 + [H_REF]*4 +
        [H_L12]*4 + [H_SPA]*5 +
        [H_DELTA]*2 + [H_DELTA]*2 + [H_DELTA]*2 + [H_REF]
    )
    for ci, (hdr, fill) in enumerate(zip(headers, col_fills), 1):
        cell = ws.cell(3, ci, hdr)
        cell.fill      = fill
        cell.font      = W_FONT
        cell.alignment = CENTER
    ws.row_dimensions[3].height = 30

    # Data rows
    fill_status = {
        "✅ Profit ↑":         PatternFill("solid", fgColor="C6EFCE"),
        "⚠️ TO ↑ / Profit ↓":  PatternFill("solid", fgColor="FFEB9C"),
        "🔴 Both ↓":           PatternFill("solid", fgColor="FFC7CE"),
        "❌ Below Cost":        PatternFill("solid", fgColor="FF0000"),
    }
    data_cols_map = [
        ("PN", None), ("Description", None),
        ("Buy_Price", "€#,##0.00"), ("Eff_Cost", "€#,##0.00"),
        ("Sell_Price", "€#,##0.00"), ("Disc_pct", "0.0%"),
        ("L12_Qty", "#,##0"), ("L12_TO", "€#,##0"), ("L12_Profit", "€#,##0"), ("L12_GM_pct", "0.0%"),
        ("SPA_Qty", "#,##0"), ("SPA_Price", "€#,##0.00"),
        ("SPA_TO", "€#,##0"), ("SPA_Profit", "€#,##0"), ("SPA_GM_pct", "0.0%"),
        ("Δ_TO", "€#,##0"), ("Δ_Profit", "€#,##0"),
        ("BE_Qty", "#,##0"), ("BE_Price", "€#,##0.00"),
        ("Min_Supp_Needed", "0.0%"), ("No_Supp_OK", None),
        ("Status", None),
    ]

    for ri, (_, row_data) in enumerate(df_in.iterrows(), 4):
        s_fill = fill_status.get(str(row_data.get("Status", "")))
        for ci, (col, fmt) in enumerate(data_cols_map, 1):
            raw = row_data.get(col, "")
            val = None if (isinstance(raw, float) and np.isinf(raw)) else raw
            if isinstance(val, bool):
                val = "Yes" if val else "No"
            cell = ws.cell(ri, ci, val)
            if fmt:
                cell.number_format = fmt
            if s_fill and ci > 6:
                cell.fill = s_fill

    # Totals row
    last = len(df_in) + 4
    ws.cell(last, 1, "TOTAL").font = BOLD
    sums = {
        8:  df_in["L12_TO"].sum(),
        9:  df_in["L12_Profit"].sum(),
        13: df_in["SPA_TO"].sum(),
        14: df_in["SPA_Profit"].sum(),
        16: df_in["Δ_TO"].sum(),
        17: df_in["Δ_Profit"].sum(),
    }
    for ci, val in sums.items():
        c = ws.cell(last, ci, round(val, 0))
        c.font = BOLD
        c.number_format = "€#,##0"

    # Sheet 2 – Entity sensitivity table
    ws2 = wb.create_sheet("Entity Sensitivity")
    ws2.append(["Entity BP reduction %", "Portfolio Profit €", "vs L12 Profit €", "Covers L12?"])
    for cell in ws2[1]:
        cell.font = BOLD
    l12_p_total = float(df_in["L12_Profit"].sum())
    for s in np.linspace(0, 0.30, 31):
        p = float(((df_in["SPA_Price"] - df_in["Buy_Price"] * (1 - s)) * df_in["SPA_Qty"]).sum())
        ws2.append([
            f"{s*100:.0f}%",
            round(p, 0),
            round(p - l12_p_total, 0),
            "✅" if p >= l12_p_total else "❌",
        ])

    # Column widths
    for sheet in [ws, ws2]:
        for col_cells in sheet.columns:
            max_len = max(
                (len(str(c.value)) for c in col_cells if c.value is not None), default=8
            )
            sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_len + 2, 28)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(data_cols_map))}3"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


exp1, exp2 = st.columns(2)
with exp1:
    xlsx_data = build_excel(df, iml_pct, imit_pct)
    fname = f"NOx_SPA_Analysis_IML{int(iml_pct*100)}pct_IMIT{int(imit_pct*100)}pct.xlsx"
    st.download_button(
        "📥 Download Excel (full analysis + sensitivity sheet)",
        data=xlsx_data, file_name=fname,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )
with exp2:
    csv_data = df[result_cols].to_csv(index=False).encode()
    st.download_button(
        "📄 Download CSV",
        data=csv_data, file_name="NOx_SPA_Analysis.csv",
        mime="text/csv", use_container_width=True,
    )
