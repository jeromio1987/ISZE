# -*- coding: utf-8 -*-
"""Read-only: page and slide counts for the documents the index will describe."""
import os
import fitz
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

PDFS = [
    "INDEX.md",
    "ISUZU_UNIVERSE_WHITE_PAPER_2026-09-23.pdf",
    "ISUZU_UNIVERSE_3_0_2026-09-23.pdf",
    "ISUZU_UNIVERSE_3_0_ANNEX_2026-09-23.pdf",
    "Universe_Yellow_Paper_v0.5_2026-09-22.pdf",
    "GPT_HANDOVER_2026-09-23/UNIVERSE_3_0_CRITIQUE_2026-09-23.pdf",
    "ISUZU_UNIVERSE_COMMERCIAL_VISION_3_0.pdf",
    "ISUZU_UNIVERSE_COMMERCIAL_TOOLS_WHITE_PAPER.pdf",
    "01_Strategy/AIM_PRISM_Five_Action_Studies.pdf",
]

for f in PDFS:
    p = os.path.join(ROOT, f)
    if not os.path.exists(p):
        print("MISSING  %s" % f)
        continue
    if f.lower().endswith(".md"):
        print("MD       %s  %d bytes" % (f, os.path.getsize(p)))
        continue
    d = fitz.open(p)
    print("PDF      %s  pages=%d  bytes=%d" % (f, d.page_count, os.path.getsize(p)))
    d.close()

pp = os.path.join(ROOT, "ISUZU_UNIVERSE_3_0_DECK_2026-09-23.pptx")
pr = Presentation(pp)
slides = list(pr.slides)
notes = sum(1 for s in slides
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
print("PPTX     ISUZU_UNIVERSE_3_0_DECK_2026-09-23.pptx  slides=%d  notes_on=%d"
      % (len(slides), notes))
