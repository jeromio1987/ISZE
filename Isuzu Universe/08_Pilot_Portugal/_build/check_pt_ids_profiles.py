"""Verify the Portugal Dealer Standards 2026 profile pages against the U14 brief."""
from __future__ import annotations

import html as html_lib
import json
import re
import sys
from pathlib import Path

import openpyxl

BUILD = Path(__file__).resolve().parent
ROOT = BUILD.parent
PAGES = ROOT / "pages"
sys.path.insert(0, str(BUILD))
from build_pt_ids_profiles import (  # noqa: E402
    DEALERS,
    INPUT_DIR,
    TOTAL_QUESTIONS,
    pages_dir_for,
    read_dealer,
    resolve_file,
)

PAGE = "Dealer_Standards_2026.html"
OPEN = "OPEN_Dealer_Standards_2026.cmd"

EXPECTED = {
    "PT1": (171, 12, 7, 4, 0, 194),
    "PT2": (185, 0, 7, 2, 0, 194),
    "PT3": (147, 25, 13, 9, 0, 194),
}
KEYS = ("highest", "lower", "lowest", "matches_no_option", "no_answer", "total")

ACRONYMS = ["AIM", "PRISM", "DLOD", "IDS", "GPS", "BVP", "IML", "IMIT", "ISZE"]
WORDS = ["campaign", "sell-out", "sellout", "wallet", "score", "rank"]
# "rank" allowed only inside "No total or ranking"
RANK_OK = "No total or ranking"


def strip_blocks(raw: str) -> str:
    raw = re.sub(r"<script\b[^>]*>.*?</script>", " ", raw, flags=re.I | re.S)
    return re.sub(r"<style\b[^>]*>.*?</style>", " ", raw, flags=re.I | re.S)


def to_text(fragment: str) -> str:
    text = re.sub(r"<[^>]+>", " ", fragment)
    return re.sub(r"\s+", " ", html_lib.unescape(text))


def visible_text(raw: str) -> str:
    return to_text(strip_blocks(raw))


def visible_outside_verbatim(raw: str) -> str:
    body = strip_blocks(raw)
    body = re.sub(r'<[^>]*class="[^"]*\bverbatim\b[^"]*"[^>]*>.*?</[^>]+>', " ", body, flags=re.S)
    return to_text(body)


def other_portugal_dealer_names(keep_files: set[str]) -> list[str]:
    """Names from the other 13 Portugal workbooks (cover sheet)."""
    names: list[str] = []
    for path in sorted(INPUT_DIR.glob("*.xlsx")):
        if path.name in keep_files:
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws0 = wb[wb.sheetnames[0]]
        found = None
        for row in ws0.iter_rows(values_only=True):
            if not row:
                continue
            cells = list(row)
            for j, v in enumerate(cells):
                if v is None:
                    continue
                if "Nome do Concession" in str(v):
                    for k in range(j + 1, len(cells)):
                        if cells[k] is not None and str(cells[k]).strip():
                            found = str(cells[k]).strip().lstrip(". ").strip()
                            break
            if found:
                break
        wb.close()
        if found:
            names.append(found)
    return names


def main() -> int:
    errors: list[str] = []

    man_path = PAGES / "ids_manifest.json"
    if not man_path.exists():
        print("FAIL missing ids_manifest.json")
        return 1
    man = json.loads(man_path.read_text(encoding="utf-8"))
    by_code = {d["code"]: d for d in man["dealers"]}
    if set(by_code) != set(EXPECTED):
        errors.append(f"manifest codes {sorted(by_code)} != {sorted(EXPECTED)}")
    for code, exp in EXPECTED.items():
        got = by_code.get(code, {})
        for k, v in zip(KEYS, exp):
            if got.get(k) != v:
                errors.append(f"{code}.{k}: {got.get(k)} != {v}")

    keep_files: set[str] = set()
    dealer_answers: dict[str, list[dict]] = {}
    for dealer in DEALERS:
        path = resolve_file(dealer)
        keep_files.add(path.name)
        answers, _, _ = read_dealer(path)
        dealer_answers[str(dealer["code"])] = answers

    other_names = other_portugal_dealer_names(keep_files)
    own_names = {str(d["page_name"]) for d in DEALERS} | {str(d["form_name"]) for d in DEALERS}

    ok = 0
    for dealer in DEALERS:
        code = str(dealer["code"])
        before = len(errors)
        folder = pages_dir_for(dealer)
        page, cmd = folder / PAGE, folder / OPEN
        if not page.exists() or not cmd.exists():
            errors.append(f"{code}: missing page or OPEN file")
            continue
        raw = page.read_text(encoding="utf-8")
        answers = dealer_answers[code]

        q_rows = re.findall(r'<tr class="q">(.*?)</tr>', raw, flags=re.S)
        if len(q_rows) != TOTAL_QUESTIONS:
            errors.append(f"{code}: {len(q_rows)} question rows, expected {TOTAL_QUESTIONS}")
        for q, row in zip(answers, q_rows):
            cells = re.findall(r"<td[^>]*>(.*?)</td>", row, flags=re.S)
            if len(cells) < 3 or cells[0] != html_lib.escape(q["number"]):
                errors.append(f"{code}: row order differs at {q['number']}")
            elif cells[1] != html_lib.escape(q["wording"]):
                errors.append(f"{code}: wording differs for {q['number']}")
            elif cells[2] != html_lib.escape(q["answer_raw"]):
                errors.append(f"{code}: answer differs for {q['number']}")

        vis = visible_text(raw)
        outside = visible_outside_verbatim(raw)
        for term in ACRONYMS:
            if re.search(rf"(?<![A-Za-z0-9]){term}(?![A-Za-z0-9])", outside):
                errors.append(f"{code}: banned term {term!r}")
        for term in WORDS:
            if term == "rank":
                cleaned = outside.replace(RANK_OK, " ")
                if re.search(r"(?<![A-Za-z])rank(?![A-Za-z])", cleaned, flags=re.I):
                    errors.append(f"{code}: banned term 'rank'")
                continue
            if term == "score":
                if re.search(r"(?<![A-Za-z])score(?![A-Za-z])", outside, flags=re.I):
                    errors.append(f"{code}: banned term 'score'")
                continue
            if term in outside.lower():
                errors.append(f"{code}: banned term {term!r}")

        if "%" in outside:
            errors.append(f"{code}: % outside verbatim cells")

        for name in other_names:
            if name in own_names:
                continue
            if name and (name in raw or html_lib.escape(name) in raw):
                # skip short fragments that collide with English words
                if len(name) < 4:
                    continue
                errors.append(f"{code}: other Portugal dealer name present: {name}")

        for m in re.finditer(r"""\b(?:src|href)\s*=\s*["']([^"']*)["']""", raw, flags=re.I):
            if "http" in m.group(1).lower():
                errors.append(f"{code}: http in {m.group(0)}")

        for bad in ("NaN", "undefined", "null"):
            if re.search(rf"(?<![A-Za-z0-9]){bad}(?![A-Za-z0-9])", vis):
                errors.append(f"{code}: {bad!r} in visible text")

        cmd_bytes = cmd.read_bytes()
        if b"\r\n" not in cmd_bytes or re.search(rb"(?<!\r)\n", cmd_bytes):
            errors.append(f"{code}: OPEN file is not CRLF")
        if PAGE.encode() not in cmd_bytes:
            errors.append(f"{code}: OPEN file does not name {PAGE}")

        if len(errors) == before:
            ok += 1

    if errors:
        print("FAIL")
        for e in errors[:60]:
            print(" -", e)
        if len(errors) > 60:
            print(f" ... +{len(errors) - 60} more")
        return 1
    print(f"PT_IDS_PROFILES_OK {ok}/3")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
