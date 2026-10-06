"""Build one Dealer Standards 2026 self-evaluation page per Portugal visit dealer."""
from __future__ import annotations

import hashlib
import html
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ISZE = Path(__file__).resolve().parents[3]
PAGES = ROOT / "pages"
INPUT_DIR = (
    ISZE
    / "01_data"
    / "08_Distributor"
    / "IDS_2026"
    / "Portugal"
    / "IDS Evaluation 2026 - PORTUGAL - FINAL"
)

QUESTION_RE = re.compile(r"^[A-E]\.\d+-\d+")
SUBSECTION_RE = re.compile(r"^[A-E]\.\d+\s")
SECTION_RE = re.compile(r"^[A-E]\.\s")
TOTAL_QUESTIONS = 194

SECTIONS = {
    "A": "Premises, equipment and clothing",
    "B": "Organisation and management",
    "C": "Service",
    "D": "Parts and accessories",
    "E": "Vehicle sales",
}

HIGHEST = "Highest answer chosen"
LOWER = "A lower answer chosen"
LOWEST = "Lowest answer chosen"
NOMATCH = "Recorded points match no answer option"
NONE = "No answer recorded"
CLASSES = [HIGHEST, LOWER, LOWEST, NOMATCH, NONE]
CLASS_KEY = {
    HIGHEST: "highest",
    LOWER: "lower",
    LOWEST: "lowest",
    NOMATCH: "matches_no_option",
    NONE: "no_answer",
}
CLASS_CSS = {
    HIGHEST: "c-highest",
    LOWER: "c-lower",
    LOWEST: "c-lowest",
    NOMATCH: "c-nomatch",
    NONE: "c-none",
}
CLASS_TIP = {
    HIGHEST: "The dealer picked the highest points option that the form offers for this question.",
    LOWER: "The dealer picked an option that scores more than zero but less than the highest points the form offers.",
    LOWEST: "The dealer picked the zero-points option for a question that has a higher option available.",
    NOMATCH: "The form records a points value that does not match any filled answer option for this question. "
    "The points are shown in words; they are not repaired.",
    NONE: "The self-evaluation cell for this question is empty on the form.",
}

# Exact mapping from the U14 brief — do not reopen.
DEALERS: list[dict[str, object]] = [
    {
        "code": "PT1",
        "form_name": "Rolmax",
        "page_name": "Rolmax",
        "file_glob": "*ROLMAX*",
        "hash_prefix": "856b3c4baf5a831b",
    },
    {
        "code": "PT2",
        "form_name": "Sociedade Comercial de Automóveis Reno, SA",
        "page_name": "Auto Reno Paços Ferreira",
        "file_glob": "*AUTO RENO PA*OS FERREIRA*",
        "hash_prefix": "c078cb6f2d608c2f",
    },
    {
        "code": "PT3",
        "form_name": "SOCARROS SA",
        "page_name": "Socarros Caldas",
        "file_glob": "*SOCARROS CALDAS*",
        "hash_prefix": "3d181b2c1d4a37f6",
    },
]


def sha256_full(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def folder_slug(page_name: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch == " " else " " for ch in page_name)
    return "_".join(cleaned.split())


def pages_dir_for(dealer: dict[str, object]) -> Path:
    return PAGES / f"{dealer['code']}_{folder_slug(str(dealer['page_name']))}"


def resolve_file(dealer: dict[str, object]) -> Path:
    matches = sorted(INPUT_DIR.glob(str(dealer["file_glob"])))
    if len(matches) != 1:
        raise SystemExit(f"{dealer['code']}: expected 1 file for {dealer['file_glob']!r}, found {len(matches)}")
    return matches[0]


def points_in_words(n: int) -> str:
    if n == 1:
        return "1 point"
    return f"{n} points"


def answer_class(picked, max_points, option_text: str) -> str:
    """Check order: empty first, then no-match, then highest, then zero, then lower."""
    if picked is None:
        return NONE
    if isinstance(picked, str) and picked.strip() == "":
        return NONE
    try:
        s = int(picked)
    except (TypeError, ValueError):
        s = int(float(picked))
    try:
        max_p = int(max_points) if max_points is not None else 0
    except (TypeError, ValueError):
        max_p = int(float(max_points))
    opt = (option_text or "").strip()
    if opt == "" or s > max_p:
        return NOMATCH
    if s == max_p:
        return HIGHEST
    if s == 0 and max_p > 0:
        return LOWEST
    if 0 < s < max_p:
        return LOWER
    return NOMATCH


def display_answer(cls: str, option_text: str, picked) -> str:
    if cls == NONE:
        return ""
    if cls == NOMATCH:
        try:
            n = int(picked)
        except (TypeError, ValueError):
            n = int(float(picked))
        return (
            f"the form records {points_in_words(n)}; "
            f"no option on the form carries {points_in_words(n)}"
        )
    return option_text


def parse_eval_date(ws0) -> datetime:
    for row in ws0.iter_rows(values_only=True):
        if not row:
            continue
        cells = list(row)
        for j, v in enumerate(cells):
            if v is None:
                continue
            if "Data da Avalia" in str(v):
                for k in range(j + 1, len(cells)):
                    if cells[k] is None:
                        continue
                    val = cells[k]
                    if isinstance(val, datetime):
                        return val
                    if isinstance(val, str) and val.strip():
                        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                            try:
                                return datetime.strptime(val.strip()[:10], fmt)
                            except ValueError:
                                continue
                    raise SystemExit(f"cannot parse evaluation date: {val!r}")
    raise SystemExit("evaluation date label not found on cover sheet")


def format_date(d: datetime) -> str:
    return f"{d.day} {d.strftime('%B %Y')}"


def read_dealer(path: Path) -> tuple[list[dict], datetime, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws0 = wb[wb.sheetnames[0]]
    ws1 = wb[wb.sheetnames[1]]
    eval_date = parse_eval_date(ws0)

    questions: list[dict] = []
    section_letter = None
    section_pt = ""
    sub_code, sub_title = None, None
    for row in ws1.iter_rows(values_only=True):
        if not row or len(row) < 12 or row[1] is None:
            continue
        first = str(row[1]).strip()
        if SECTION_RE.match(first):
            section_letter = first[0]
            section_pt = first
            continue
        if SUBSECTION_RE.match(first):
            # title is the rest after the code token
            parts = first.split(None, 1)
            sub_code = parts[0]
            sub_title = parts[1] if len(parts) > 1 else ""
            continue
        if QUESTION_RE.match(first):
            max_p = row[6]
            opts = {0: row[7], 1: row[8], 2: row[9], 3: row[10]}
            picked = row[11]
            if picked is None or (isinstance(picked, str) and str(picked).strip() == ""):
                s = None
                option = ""
            else:
                s = int(float(picked))
                option = "" if opts.get(s) is None else str(opts[s])
            cls = answer_class(picked, max_p, option)
            questions.append(
                {
                    "number": first,
                    "section": first[0],
                    "section_pt": section_pt,
                    "sub_code": sub_code,
                    "sub_title": sub_title,
                    "wording": "" if row[3] is None else str(row[3]),
                    "max_points": max_p,
                    "option": option,
                    "picked": picked if s is None else s,
                    "answer_raw": display_answer(cls, option, picked if s is None else s),
                    "cls": cls,
                    "reviewed": "" if len(row) <= 13 or row[13] is None else str(row[13]),
                }
            )
    wb.close()
    if len(questions) != TOTAL_QUESTIONS:
        raise SystemExit(f"{path.name}: expected {TOTAL_QUESTIONS} questions, found {len(questions)}")
    return questions, eval_date, path.name


def tip(text: str, tip_text: str) -> str:
    return (
        f'<span class="kpi-tip" tabindex="0">{html.escape(text)}'
        f'<span class="tip">{html.escape(tip_text)}</span></span>'
    )


def stacked_bar(counts: dict[str, int], total: int) -> str:
    parts = [
        '<svg class="stack" viewBox="0 0 1000 18" preserveAspectRatio="none" '
        'role="img" aria-hidden="true">'
    ]
    x = 0.0
    for c in CLASSES:
        n = counts[c]
        if not n:
            continue
        w = 1000.0 * n / total
        parts.append(f'<rect class="{CLASS_CSS[c]}" x="{x:.2f}" y="0" width="{w:.2f}" height="18"/>')
        x += w
    parts.append("</svg>")
    return "".join(parts)


def count_line(counts: dict[str, int]) -> str:
    bits = [
        f'<span class="cl"><i class="sw {CLASS_CSS[c]}"></i>{html.escape(c)} {counts[c]}</span>'
        for c in CLASSES
        if counts[c]
    ]
    return " ".join(bits)


def question_label(n: int) -> str:
    return f"{n} question" if n == 1 else f"{n} questions"


def render_page(dealer: dict, answers: list[dict], counts: dict[str, int], date_text: str) -> str:
    page_name = str(dealer["page_name"])

    subs: dict[str, dict] = {}
    for a in answers:
        key = a["sub_code"]
        if key not in subs:
            subs[key] = {"title": a["sub_title"], "counts": {c: 0 for c in CLASSES}, "total": 0}
        subs[key]["counts"][a["cls"]] += 1
        subs[key]["total"] += 1

    kpis = "".join(
        f'<div class="kpi {CLASS_CSS[c]}-edge"><p class="label">{html.escape(c)}</p>'
        f'<p class="val">{tip(f"{counts[c]} of {TOTAL_QUESTIONS}", CLASS_TIP[c])}</p></div>'
        for c in CLASSES
    )

    legend = " ".join(
        f'<span class="cl"><i class="sw {CLASS_CSS[c]}"></i>{html.escape(c)}</span>' for c in CLASSES
    )
    overview_rows = []
    for code, s in subs.items():
        overview_rows.append(
            '<div class="ov-row">'
            f'<p class="ov-label"><b>{html.escape(code)}</b> {html.escape(s["title"])} '
            f'<span class="muted">{question_label(s["total"])}</span></p>'
            f'{stacked_bar(s["counts"], s["total"])}'
            f'<p class="ov-counts">{count_line(s["counts"])}</p>'
            "</div>"
        )

    high_items = "".join(
        f'<li><b>{html.escape(code)}</b> {html.escape(s["title"])}: '
        f'{s["counts"][HIGHEST]} of {s["total"]}</li>'
        for code, s in subs.items()
        if s["counts"][HIGHEST]
    )
    gap_list = [a for a in answers if a["cls"] in (LOWER, LOWEST, NOMATCH, NONE)]
    gap_items = "".join(
        f'<li><span class="badge {CLASS_CSS[a["cls"]]}">{html.escape(a["cls"])}</span> '
        f'<b>{html.escape(a["number"])}</b> '
        f'<span class="verbatim">{html.escape(a["wording"])}</span> — '
        f'<span class="verbatim">{html.escape(a["answer_raw"])}</span></li>'
        for a in gap_list
    )
    if not gap_items:
        gap_items = '<li class="muted">None. Every question has the highest answer chosen.</li>'

    section_blocks = []
    for letter, sec_name in SECTIONS.items():
        sec = [a for a in answers if a["section"] == letter]
        if not sec:
            continue
        sec_counts = {c: sum(1 for a in sec if a["cls"] == c) for c in CLASSES}
        rows = []
        last_sub = None
        for a in sec:
            if a["sub_code"] != last_sub:
                rows.append(
                    f'<tr class="sub"><th colspan="4">{html.escape(a["sub_code"])} '
                    f'{html.escape(a["sub_title"])}</th></tr>'
                )
                last_sub = a["sub_code"]
            rows.append(
                '<tr class="q">'
                f'<td class="qn" data-l="No.">{html.escape(a["number"])}</td>'
                f'<td class="verbatim" data-l="Question">{html.escape(a["wording"])}</td>'
                f'<td class="verbatim" data-l="Answer chosen">{html.escape(a["answer_raw"])}</td>'
                f'<td data-l="Reading"><span class="badge {CLASS_CSS[a["cls"]]}">{html.escape(a["cls"])}</span></td>'
                "</tr>"
            )
        section_blocks.append(
            '<details class="sec">'
            f'<summary>{letter}. {html.escape(sec_name)} '
            f'<span class="sum">{question_label(len(sec))}: {count_line(sec_counts)}</span></summary>'
            '<div class="table-wrap"><table><thead><tr>'
            "<th>No.</th><th>Question</th><th>Answer chosen</th><th>Reading</th>"
            "</tr></thead><tbody>"
            + "".join(rows)
            + "</tbody></table></div></details>"
        )

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>Isuzu Dealer Standards 2026: your self-evaluation — {html.escape(page_name)}</title>
<style>
:root {{
  --navy:#15233c; --paper:#f7f4ee; --ink:#222; --terra:#c0552d; --line:#d5cfc4;
  --lower:#b08a2e; --lowest:#7d8796; --nomatch:#8a4b2f; --none:#e7e1d6;
}}
* {{ box-sizing:border-box; }}
html,body {{ margin:0; padding:0; background:var(--paper); color:var(--ink);
  font-family: Georgia, "Times New Roman", serif; }}
body {{ padding:24px 20px 64px; max-width:1100px; margin:0 auto; overflow-x:hidden; }}
h1 {{ font-size:1.65rem; color:var(--navy); margin:0 0 4px; line-height:1.25; }}
h2 {{ font-size:1.15rem; color:var(--navy); margin:0 0 12px; }}
.dealer {{ font-size:1.2rem; color:var(--navy); margin:0 0 8px; }}
.source {{ color:#445; font-size:0.95rem; margin:0 0 22px; max-width:820px; }}
.muted {{ color:#556; font-weight:400; }}
.kpis {{ display:grid; grid-template-columns:repeat(5,minmax(0,1fr)); gap:12px; margin:0 0 28px; }}
.kpi {{ background:#fff; border:1px solid var(--line); border-left:4px solid var(--navy);
  padding:14px 14px 12px; min-width:0; }}
.kpi.c-lower-edge {{ border-left-color:var(--lower); }}
.kpi.c-lowest-edge {{ border-left-color:var(--lowest); }}
.kpi.c-nomatch-edge {{ border-left-color:var(--nomatch); }}
.kpi.c-none-edge {{ border-left-color:#b9b0a1; }}
.kpi .label {{ font-size:0.72rem; text-transform:uppercase; letter-spacing:0.03em; color:#445; margin:0 0 6px; line-height:1.25; }}
.kpi .val {{ font-size:1.25rem; color:var(--navy); font-weight:700; margin:0; }}
.kpi-tip {{ position:relative; cursor:help; border-bottom:1px dotted var(--terra); }}
.kpi-tip .tip {{
  display:none; position:absolute; left:0; top:calc(100% + 8px); z-index:20;
  width:min(300px, 80vw); background:var(--navy); color:var(--paper);
  padding:10px 12px; font-size:0.82rem; line-height:1.35; border-top:3px solid var(--terra);
  box-shadow:0 8px 18px rgba(21,35,60,0.25); font-weight:400;
}}
.kpi:nth-child(n+4) .tip {{ left:auto; right:0; }}
.kpi-tip:hover .tip, .kpi-tip:focus .tip {{ display:block; }}
.box {{ background:#fff; border:1px solid var(--line); padding:16px 18px; margin:0 0 28px; }}
.legend {{ margin:0 0 14px; font-size:0.88rem; color:#334; }}
.cl {{ display:inline-block; margin:0 14px 4px 0; white-space:nowrap; }}
.sw {{ display:inline-block; width:12px; height:12px; margin-right:6px; vertical-align:-1px; }}
.c-highest {{ background:var(--navy); fill:var(--navy); }}
.c-lower {{ background:var(--lower); fill:var(--lower); }}
.c-lowest {{ background:var(--lowest); fill:var(--lowest); }}
.c-nomatch {{ background:var(--nomatch); fill:var(--nomatch); }}
.c-none {{ background:var(--none); fill:var(--none); }}
i.sw.c-none {{ outline:1px solid #a69d8e; }}
.ov-row {{ margin:0 0 14px; }}
.ov-label {{ margin:0 0 4px; color:var(--navy); font-size:0.95rem; }}
.stack {{ display:block; width:100%; height:18px; background:var(--none); outline:1px solid var(--line); }}
.ov-counts {{ margin:4px 0 0; font-size:0.82rem; color:#334; }}
.lists {{ display:grid; grid-template-columns:1fr 1fr; gap:16px; margin:0 0 28px; }}
.lists .box {{ margin:0; }}
.lists ul {{ margin:0; padding-left:1.1rem; }}
.lists li {{ margin:0 0 8px; line-height:1.4; }}
.badge {{ display:inline-block; padding:1px 7px; font-size:0.74rem; border-radius:2px; white-space:nowrap; }}
.badge.c-highest {{ color:var(--paper); }}
.badge.c-lower {{ color:#fff; }}
.badge.c-lowest {{ color:#fff; }}
.badge.c-nomatch {{ color:#fff; }}
.badge.c-none {{ color:var(--ink); outline:1px solid #a69d8e; }}
details.sec {{ border:1px solid var(--line); background:#fff; margin:0 0 8px; }}
summary {{ cursor:pointer; padding:10px 12px; color:var(--navy); font-weight:700; list-style:none; }}
summary::-webkit-details-marker {{ display:none; }}
summary::before {{ content:"+ "; color:var(--terra); }}
details[open] > summary::before {{ content:"− "; }}
.sum {{ display:block; font-weight:400; color:#445; font-size:0.85rem; margin-top:4px; }}
.table-wrap {{ margin:0 10px 12px; border:1px solid var(--line); }}
table {{ border-collapse:collapse; width:100%; font-size:0.84rem; background:#fff; table-layout:fixed; }}
th, td {{ border-bottom:1px solid var(--line); padding:7px 8px; text-align:left; vertical-align:top;
  color:var(--ink); overflow-wrap:anywhere; }}
thead th {{ background:var(--navy); color:var(--paper); }}
thead th:nth-child(1) {{ width:72px; }} thead th:nth-child(2) {{ width:48%; }}
thead th:nth-child(4) {{ width:170px; }}
tr.sub th {{ background:#eef0f4; color:var(--navy); font-weight:700; }}
td.qn {{ font-family: Consolas, "Courier New", monospace; white-space:nowrap; }}
.limits ul {{ margin:8px 0 0; padding-left:1.2rem; }}
.limits li {{ margin:4px 0; }}
@media (max-width:900px) {{
  .kpis {{ grid-template-columns:1fr 1fr; }}
  .kpi:nth-child(n+4) .tip {{ left:0; right:auto; }}
  .kpi:nth-child(even) .tip {{ left:auto; right:0; }}
  .lists {{ grid-template-columns:1fr; }}
}}
@media (max-width:700px) {{
  body {{ padding:16px 12px 48px; }}
  h1 {{ font-size:1.35rem; }}
  table, thead, tbody, tr, td, th {{ display:block; width:auto !important; }}
  thead {{ display:none; }}
  tr.q {{ border-bottom:2px solid var(--line); padding:6px 0; }}
  tr.q td {{ border:0; padding:3px 8px; }}
  tr.q td::before {{ content:attr(data-l); display:block; font-size:0.72rem; text-transform:uppercase;
    letter-spacing:0.04em; color:#556; }}
}}
@media (max-width:420px) {{
  .kpis {{ grid-template-columns:1fr; }}
  .kpi .tip, .kpi:nth-child(even) .tip {{ left:0; right:auto; }}
}}
</style>
</head>
<body>
<header>
  <h1>Isuzu Dealer Standards 2026: your self-evaluation</h1>
  <p class="dealer">{html.escape(page_name)}</p>
  <p class="source">As recorded by the dealer on the 2026 self-evaluation form, dated {html.escape(date_text)}.</p>
</header>

<section class="kpis" aria-label="Answers by type">
{kpis}
</section>

<section class="box" aria-label="Overview by subsection">
  <h2>Your answers by subsection</h2>
  <p class="legend">{legend}</p>
  {''.join(overview_rows)}
</section>

<section class="lists">
  <div class="box">
    <h2>Where the highest answer was chosen</h2>
    <ul>{high_items}</ul>
  </div>
  <div class="box">
    <h2>Where it was not, or the points match no option</h2>
    <ul>{gap_items}</ul>
  </div>
</section>

<section aria-label="Full questionnaire">
  <h2>The full questionnaire</h2>
  {''.join(section_blocks)}
</section>

<section class="box limits">
  <h2>What this page does not show</h2>
  <ul>
    <li>No total or ranking. The page counts your answers by class; it does not grade them.</li>
    <li>No other dealer's answers.</li>
    <li>No review by Isuzu. The review column on the form is empty for every question.</li>
    <li>Nothing about parts sales.</li>
  </ul>
</section>
</body>
</html>
"""


def main() -> int:
    if not INPUT_DIR.is_dir():
        raise SystemExit(f"missing input folder: {INPUT_DIR}")

    PAGES.mkdir(parents=True, exist_ok=True)
    manifest = {
        "built_at": datetime.now(timezone.utc).isoformat(),
        "input_hashes": {},
        "dealers": [],
    }

    for dealer in DEALERS:
        path = resolve_file(dealer)
        full_hash = sha256_full(path)
        if not full_hash.startswith(str(dealer["hash_prefix"])):
            raise SystemExit(
                f"{dealer['code']}: hash {full_hash[:16]} != {dealer['hash_prefix']}"
            )
        answers, eval_date, fname = read_dealer(path)
        counts = {c: sum(1 for a in answers if a["cls"] == c) for c in CLASSES}
        date_text = format_date(eval_date)

        out_dir = pages_dir_for(dealer)
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / "Dealer_Standards_2026.html").write_text(
            render_page(dealer, answers, counts, date_text), encoding="utf-8"
        )
        (out_dir / "OPEN_Dealer_Standards_2026.cmd").write_bytes(
            b'@echo off\r\nstart "" "%~dp0Dealer_Standards_2026.html"\r\n'
        )
        entry = {
            "code": dealer["code"],
            "page_name": dealer["page_name"],
            "form_name": dealer["form_name"],
            "folder": out_dir.name,
            "eval_date": eval_date.date().isoformat(),
        }
        entry.update({CLASS_KEY[c]: counts[c] for c in CLASSES})
        entry["total"] = sum(counts.values())
        manifest["dealers"].append(entry)
        manifest["input_hashes"][fname] = full_hash
        print(
            f"{dealer['code']}: "
            + " ".join(f"{CLASS_KEY[c]}={counts[c]}" for c in CLASSES)
            + f" total={entry['total']} date={entry['eval_date']}"
        )

    man_path = PAGES / "ids_manifest.json"
    man_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {man_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
