"""List Eucon 2023 COMPARABLE OEs not yet in oem_peer_basket (Mode B hygiene).

Writes 01_data/00_derived/eucon_comparable_oem_gap_20260810.csv
Usage: python 03_analytics/isze_pricing/scan_eucon_comparable_oem_gap.py
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

RAW = Path(__file__).resolve().parents[2] / "01_data" / "01_Pricing" / "_raw"
BASKET = Path(__file__).resolve().parent / "oem_peer_basket.csv"
UK = RAW / "Eucon_PriceSurvey_UK_2023_OEM_rawdata_20240109.xlsx"
ITUK = RAW / "Eucon_PriceSurvey_ITUK_2023_OEM_rawdata.xlsx"
OUT = Path(__file__).resolve().parents[2] / "01_data" / "00_derived" / "eucon_comparable_oem_gap_20260810.csv"


def pn_norm(x) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    s = str(x).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    s = re.sub(r"[^0-9A-Za-z]", "", s)
    if s.isdigit() and len(s) <= 10:
        return s.zfill(10)
    return s


def load_sheet(path: Path, sheet: str | None = None) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sn = sheet or wb.sheetnames[0]
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c is not None else f"c{i}" for i, c in enumerate(rows[0])]
    seen: dict[str, int] = {}
    nh = []
    for h in hdr:
        if h in seen:
            seen[h] += 1
            nh.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 0
            nh.append(h)
    df = pd.DataFrame(rows[1:], columns=nh)
    wb.close()
    return df


def extract(df: pd.DataFrame, src: str) -> pd.DataFrame:
    ccol = [c for c in df.columns if "comparab" in c.lower()][0]
    desc = (
        "CustomerProductDescription"
        if "CustomerProductDescription" in df.columns
        else "CustomerPartProperties"
    )
    m = df[df[ccol].astype(str).str.upper().str.contains("COMPARABLE", na=False)].copy()
    rows = []
    for _, r in m.iterrows():
        rows.append(
            {
                "src": src,
                "isze_pn": pn_norm(r.get("CustomerPartNumber")),
                "peer_oe_pn": pn_norm(r.get("CompetitorPartNumber")),
                "brand": r.get("CompetitorBrand"),
                "vehicle": r.get("CompetitorVehicle"),
                "veh_desc": r.get("CompetitorVehicleDescription"),
                "desc": r.get(desc),
                "comparability": r.get(ccol),
            }
        )
    return pd.DataFrame(rows)


def main() -> None:
    uk = load_sheet(UK, "Isuzu UK")
    ituk = load_sheet(ITUK)
    comp = pd.concat([extract(uk, "UK"), extract(ituk, "ITUK")], ignore_index=True)
    comp = comp[comp["peer_oe_pn"].astype(bool) & comp["isze_pn"].astype(bool)]

    basket = pd.read_csv(BASKET, dtype=str)
    mapped = {pn_norm(x) for x in basket["peer_oe_pn"].dropna()} - {""}
    basket_pairs = {
        (pn_norm(a), pn_norm(b))
        for a, b in zip(basket["isze_pn"], basket["peer_oe_pn"])
        if pn_norm(b)
    }

    comp["already_mapped_oe"] = comp["peer_oe_pn"].isin(mapped)
    comp["already_mapped_pair"] = [
        (a, b) in basket_pairs for a, b in zip(comp["isze_pn"], comp["peer_oe_pn"])
    ]

    print("COMPARABLE with OE:", len(comp))
    print("by brand:")
    print(comp["brand"].value_counts().to_string())
    print("\nUnique OEs total:", comp["peer_oe_pn"].nunique())
    print("Already mapped unique OEs:", len(mapped))
    new_oe = comp[~comp["already_mapped_oe"]]
    print("Rows with OE not in basket:", len(new_oe), "unique OEs", new_oe["peer_oe_pn"].nunique())
    print(new_oe.groupby("brand").size().to_string())
    print("\nUnmapped brand x vehicle:")
    print(new_oe.groupby(["brand", "vehicle"]).size().sort_values(ascending=False).to_string())

    # Focus peers Jerome cares about
    focus = new_oe[
        new_oe["brand"]
        .astype(str)
        .str.lower()
        .str.contains(
            "toyota|ford|mitsubishi|daf|mercedes|man|vw|volkswagen|nissan|iveco",
            regex=True,
            na=False,
        )
    ].copy()
    print("\nFocus-peer unmapped rows:", len(focus))
    for _, r in focus.sort_values(["brand", "vehicle", "desc"]).iterrows():
        print(
            f"  {r['brand']}|{r['vehicle']}|OE={r['peer_oe_pn']}|ISZE={r['isze_pn']}|"
            f"{str(r['desc'])[:55]}"
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    comp.to_csv(OUT, index=False)
    print("\nWrote", OUT)


if __name__ == "__main__":
    main()
