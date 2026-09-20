#!/usr/bin/env python3
"""
Build pr00_classification.csv from IE40 PR00 price list.

Exports per PN: family, group, fig_key (Fig No. + Key No.), level3 type.

Output: Data/derived/pr00_classification.csv

Run:  python scripts/build_pr00_classification.py
      python scripts/build_pr00_classification.py "C:/path/to/PR00.xlsx"
"""

from __future__ import annotations

import csv
import glob
import os
import re
import sys
from pathlib import Path

import openpyxl

PRISM_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PRISM_ROOT))

from config import DERIVED_DIR, PRISM_DATA_DIR, RAW_REFERENCE_DIR  # noqa: E402

HOME = Path(os.path.expanduser("~"))

SEARCH_DIRS = [
    Path(RAW_REFERENCE_DIR),
    Path(PRISM_DATA_DIR),
    HOME / "Desktop" / "claude" / "ISZE",
    HOME / "Desktop",
    HOME / "Downloads",
]
PATTERNS = ["*PR00*classified*.xlsx", "IE40*PR00*.xlsx", "*PR00*3*level*.xlsx"]

OUT = Path(DERIVED_DIR) / "pr00_classification.csv"


def resolve_pr00_path() -> Path:
    if len(sys.argv) > 1:
        p = Path(sys.argv[1])
        if not p.exists():
            raise FileNotFoundError(f"PR00 path given but not found: {p}")
        return p
    for d in SEARCH_DIRS:
        for pat in PATTERNS:
            hits = sorted(glob.glob(str(d / pat)))
            if hits:
                return Path(hits[-1])
    raise FileNotFoundError(
        "PR00 workbook not found. Searched:\n  "
        + "\n  ".join(str(d) for d in SEARCH_DIRS)
        + "\nPass the path explicitly: python scripts/build_pr00_classification.py <path>"
    )


def _norm(pn) -> str | None:
    """Normalize to the same 10-char zero-padded key used by aim / sales / the gate.

    Previously used lstrip('0'), which made every leading-zero part number invisible
    to the gate's PR00 join (audit 2026-09-17: 1,616 keys never checked).
    """
    if pn is None:
        return None
    s = re.sub(r"\.0$", "", str(pn).strip().upper())
    if not s:
        return None
    if s.isdigit():
        return s.zfill(10)
    return s


def _fig_key_from_parts(fig_no, key_no) -> str | None:
    if fig_no is None or key_no is None:
        return None
    fs = str(fig_no).strip().replace(".0", "")
    ks = str(key_no).strip().replace(".0", "")
    if not fs or not ks or fs.lower() in ("nan", "none") or ks.lower() in ("nan", "none"):
        return None
    if not fs.isdigit() or not ks.isdigit():
        return None
    return f"{fs.zfill(3)}-{ks.zfill(3)}"


def _header_map(header_row: tuple) -> dict[str, int]:
    m: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        key = str(cell).strip().lower()
        m[key] = i
    return m


def build() -> int:
    src = resolve_pr00_path()
    print(f"[pr00] source: {src.name}", flush=True)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]

    header: dict[str, int] | None = None
    out: dict[str, dict] = {}

    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            header = _header_map(row)
            continue
        if not header:
            continue

        def col(name: str) -> int | None:
            return header.get(name.lower())

        i_pn = col("p/n")
        i_new = col("newest p/n")
        i_fam = col("level 1 family")
        i_grp = col("level 2 group")
        i_l3 = col("level 3 type")
        i_fig = col("fig no.")
        i_key = col("key no.")

        if i_grp is None:
            continue

        family = row[i_fam] if i_fam is not None else None
        group = row[i_grp] if i_grp is not None else None
        if not group:
            continue

        l3 = row[i_l3] if i_l3 is not None else None
        fig_no = row[i_fig] if i_fig is not None else None
        key_no = row[i_key] if i_key is not None else None
        fk = _fig_key_from_parts(fig_no, key_no)

        # Own P/N always wins. Newest P/N is only a fallback when that key has no
        # own row yet (audit 2026-09-17: 1,083 keys inherited a different node's
        # description because first-writer-wins let a predecessor fill the successor).
        own_n = _norm(row[i_pn] if i_pn is not None and i_pn < len(row) else None)
        new_n = _norm(row[i_new] if i_new is not None and i_new < len(row) else None)
        rec = {
            "family": str(family or ""),
            "group": str(group),
            "level3_type": str(l3 or ""),
            "fig_key": fk or "",
            "fig_no": str(fig_no or "").replace(".0", "") if fig_no is not None else "",
            "key_no": str(key_no or "").replace(".0", "") if key_no is not None else "",
        }
        if own_n:
            out[own_n] = rec  # overwrite — own row is authoritative
        if new_n and new_n not in out:
            out[new_n] = dict(rec)

    wb.close()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["pn_norm", "family", "group", "level3_type", "fig_key", "fig_no", "key_no"])
        for n, rec in out.items():
            w.writerow(
                [
                    n,
                    rec["family"],
                    rec["group"],
                    rec["level3_type"],
                    rec["fig_key"],
                    rec["fig_no"],
                    rec["key_no"],
                ]
            )

    with_fk = sum(1 for r in out.values() if r["fig_key"])
    print(f"[pr00] wrote {OUT} | {len(out):,} parts | fig_key on {with_fk:,}")
    return len(out)


if __name__ == "__main__":
    try:
        n = build()
        raise SystemExit(0 if n else 1)
    except FileNotFoundError as e:
        if OUT.exists():
            print(f"[pr00] SKIP — {e}; using existing {OUT.name}")
            raise SystemExit(0)
        raise
