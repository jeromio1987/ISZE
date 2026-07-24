#!/usr/bin/env python3
"""
Revenue Runway — the Golden Window as a model, not a thesis.

Projects the parts wallet per distributor through 2030 using the UIO v17
workbook's own parc projections (sheet '1.UIO Overview', 2025E-2030F columns,
MIDI rows excluded per ISZE KPI policy), valued at the official GPS max-spend
targets per model group.

Baseline capture scenario: distributor's CY2025 GPS % held flat — a SCENARIO,
not a forecast. Output flags the peak wallet year and the cumulative window.

Age-spend curve: NOT applied (decision_model_assumptions.yaml age_spend_curve
is null until IML supplies the genuine-EUR-by-age curve). Wallet is therefore
age-flat — known limitation, same as the GPS engine today.

Output: Data/derived/revenue_runway.csv
Run from prism/.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

PRISM_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PRISM_ROOT))

from config import DERIVED_DIR, resolve_uio_path  # noqa: E402
from modules.data_prep import SEGMENT_ONLY, V13_DIST_TO_CODE  # noqa: E402
from modules.gps_model_targets import (  # noqa: E402
    map_uio_row_to_model_group,
    targets_by_group,
)

OUT_CSV = DERIVED_DIR / "revenue_runway.csv"
ASSUMPTIONS = PRISM_ROOT / "data" / "decision_model_assumptions.yaml"
WATERFALL_CSV = DERIVED_DIR / "wallet_waterfall.csv"
ITUK_PROJ_CACHE = PRISM_ROOT / "data" / "ituk_uio_projection_cache.csv"


def _ituk_forecast_rows() -> pd.DataFrame:
    """ITUK 2025-2030 from UIO v17 sheet 4 Section C.

    The v17 workbook was saved without recalculation, so sheet 1's ITUK
    forecast formulas have no cached values and openpyxl reads None. We read
    the computed values once via Excel COM (read-only, never saved) and cache
    them to prism/data/ituk_uio_projection_cache.csv for Excel-free reruns.
    Aggregation mirrors sheet 1's own formulas:
      CV LD       = N35 + N50/55 + N62/65 + N75 + UNSPEC   (Section C rows 23-27)
      CV F-Series = F110/120 + F135 + F150/180             (Section C rows 20-22)
    """
    if ITUK_PROJ_CACHE.exists():
        return pd.read_csv(ITUK_PROJ_CACHE)

    try:
        import pythoncom
        import win32com.client as wc
    except ImportError:
        print("WARN: pywin32 unavailable and no ITUK projection cache — ITUK forecast skipped")
        return pd.DataFrame()

    pythoncom.CoInitialize()
    xl = wc.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(str(resolve_uio_path()), ReadOnly=True, UpdateLinks=0)
        ws = wb.Worksheets("4. ITUK UIO Projection")
        vals = ws.Range("B19:H28").Value
        wb.Close(False)
    finally:
        xl.Quit()

    header = [str(v) for v in vals[0]]  # Model, 2025..2030
    years = [int(y) for y in header[1:]]
    by_model = {str(row[0]): row[1:] for row in vals[1:]}

    ld_models = ["N35", "N50/55", "N62/65", "N75", "UNSPEC"]
    f_models = ["F110/120", "F135", "F150/180"]
    rows = []
    for i, yr in enumerate(years):
        ld = sum(float(by_model[m][i] or 0) for m in ld_models if m in by_model)
        fs = sum(float(by_model[m][i] or 0) for m in f_models if m in by_model)
        rows.append({"distributor_code": "ITUK", "year": yr, "uio_units": round(ld),
                     "cv_lcv": "CV", "model": "LD", "is_forecast": True})
        rows.append({"distributor_code": "ITUK", "year": yr, "uio_units": round(fs),
                     "cv_lcv": "CV", "model": "F-Series", "is_forecast": True})
    df = pd.DataFrame(rows)
    df.to_csv(ITUK_PROJ_CACHE, index=False)
    print(f"runway: ITUK projection read via Excel COM, cached -> {ITUK_PROJ_CACHE.name}")
    return df


def load_uio_all_years(path: str, year_min: int, year_max: int) -> pd.DataFrame:
    """UIO Overview incl. forecast columns ('2025E', '2026F', ...). MIDI excluded."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["1.UIO Overview"]
    header = list(ws.iter_rows(min_row=6, max_row=6, values_only=True))[0]

    year_cols: dict[int, int] = {}
    for i, h in enumerate(header):
        token = str(h).strip().rstrip("EF") if h is not None else ""
        try:
            yr = int(token)
        except ValueError:
            continue
        if year_min <= yr <= year_max:
            year_cols[yr] = i

    rows = []
    for row in ws.iter_rows(min_row=7, max_row=500, values_only=True):
        if row[0] is None:
            continue
        dist_v13 = str(row[5]).strip() if row[5] else ""
        if not dist_v13 or dist_v13 == "None" or "MIDI" in dist_v13.upper():
            continue
        dist_code = V13_DIST_TO_CODE.get(dist_v13)
        if not dist_code:
            continue
        cv_lcv = str(row[2]).strip() if row[2] else ""
        if dist_code in SEGMENT_ONLY and cv_lcv != SEGMENT_ONLY[dist_code]:
            continue
        model = str(row[3]).strip() if row[3] else ""
        for yr, ci in year_cols.items():
            val = row[ci]
            if val is None:
                continue
            try:
                units = float(val)
            except (TypeError, ValueError):
                continue
            rows.append(
                {
                    "distributor_code": dist_code,
                    "year": yr,
                    "uio_units": units,
                    "cv_lcv": cv_lcv,
                    "model": model,
                    "is_forecast": yr >= 2025,
                }
            )
    return pd.DataFrame(rows)


def main() -> int:
    cfg = yaml.safe_load(ASSUMPTIONS.read_text(encoding="utf-8"))["revenue_runway"]
    y0, y1 = int(cfg["start_year"]), int(cfg["end_year"])
    basis_year = int(cfg["capture_basis_year"])
    if cfg.get("age_spend_curve"):
        print("NOTE: age_spend_curve set but v1 builder ignores it — extend before use.")

    uio = load_uio_all_years(str(resolve_uio_path()), y0, y1)
    if uio.empty:
        print("ERROR: no UIO rows parsed")
        return 1

    # ITUK forecast cells are uncached formulas in v17 — backfill from sheet 4
    ituk_fc = uio[(uio["distributor_code"] == "ITUK") & (uio["year"] >= 2025)]
    if ituk_fc.empty:
        extra = _ituk_forecast_rows()
        if not extra.empty:
            extra = extra[(extra["year"] >= y0) & (extra["year"] <= y1)]
            uio = pd.concat([uio, extra], ignore_index=True)
    targets = targets_by_group()
    uio["model_group"] = uio.apply(
        lambda r: map_uio_row_to_model_group(r["model"], r["cv_lcv"]), axis=1
    )
    uio["target_eur"] = uio["model_group"].map(targets).fillna(targets.get("N75_N50_N62", 1120))
    uio["wallet_eur"] = uio["uio_units"] * uio["target_eur"]

    proj = (
        uio.groupby(["distributor_code", "year", "is_forecast"], as_index=False)
        .agg(uio_units=("uio_units", "sum"), wallet_eur=("wallet_eur", "sum"))
    )

    # CY2025 GPS % per distributor from the waterfall (same identity)
    gps_basis = {}
    if WATERFALL_CSV.exists():
        wf = pd.read_csv(WATERFALL_CSV)
        wfb = wf[wf["year"] == basis_year]
        gps_basis = dict(zip(wfb["distributor_code"], wfb["gps_pct"]))
    else:
        print("WARN: wallet_waterfall.csv missing — run build_wallet_waterfall.py first")

    proj["gps_pct_basis"] = proj["distributor_code"].map(gps_basis)
    proj["projected_capture_eur"] = (
        proj["wallet_eur"] * proj["gps_pct_basis"] / 100.0
    ).round(2)
    proj["scenario_note"] = (
        f"capture = CY{basis_year} GPS% held flat (BASELINE SCENARIO); "
        "wallet = UIO v17 parc projection x GPS max targets, age-flat "
        "(age-spend curve pending from IML)"
    )

    # Peak year + window metrics per distributor
    peaks = []
    for dist, grp in proj.groupby("distributor_code"):
        grp = grp.sort_values("year")
        peak_row = grp.loc[grp["wallet_eur"].idxmax()]
        fc = grp[grp["year"] >= 2026]
        last = grp[grp["year"] == y1]["wallet_eur"]
        first = grp[grp["year"] == y0]["wallet_eur"]
        peaks.append(
            {
                "distributor_code": dist,
                "peak_wallet_year": int(peak_row["year"]),
                "peak_wallet_eur": round(float(peak_row["wallet_eur"]), 2),
                "wallet_2030_vs_2024_pct": round(
                    (float(last.iloc[0]) / float(first.iloc[0]) - 1) * 100, 1
                )
                if len(last) and len(first) and float(first.iloc[0]) > 0
                else None,
                "cum_wallet_2026_2030_eur": round(float(fc["wallet_eur"].sum()), 2),
                "cum_projected_capture_2026_2030_eur": round(
                    float(fc["projected_capture_eur"].sum()), 2
                )
                if fc["projected_capture_eur"].notna().any()
                else None,
            }
        )
    out = proj.merge(pd.DataFrame(peaks), on="distributor_code", how="left")
    out = out.sort_values(["distributor_code", "year"])
    out.to_csv(OUT_CSV, index=False)

    net = out.groupby("year")["wallet_eur"].sum()
    peak_year = int(net.idxmax())
    print(
        f"runway: {len(out)} rows -> {OUT_CSV.name} | network wallet peaks {peak_year} "
        f"(EUR{net.max()/1e6:.1f}M); 2030 = EUR{net.get(2030, float('nan'))/1e6:.1f}M"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
