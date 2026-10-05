#!/usr/bin/env python3
"""
Real per-distributor UIO split by model-type bucket (CV-standard [LD+MD], F-Series, LCV).

WHY cv_std STAYS COLLAPSED (reviewed 2026-09-17 — do not "fix" this without reading):

`cv_std` merges N35 and N75 deliberately. This file feeds the PMS capture engine
(`expected_by_category_bucketed`), where what matters is SERVICES PER YEAR, and the pipeline
applies one rate to both bands (0.5 svc/yr). Splitting them here would yield an identical
expected_qty, so the bucket stays collapsed rather than implying a distinction the service
model does not make.

The GPS side DOES need the split, because the rate card prices N35 at EUR 296 and N75 at
EUR 739 — a 2.5x difference. That split is applied downstream in
`modules/data_prep._split_collapsed_ld_rows()`, which reads the workbook's own ITUK carpark
breakdown (sheet '2. Vehicle Volume' section B) and emits explicit N35 / N75 rows before the
rate mapper sees them. See `modules/gps_model_targets.map_uio_row_to_model_group`.

Net effect for ITUK: cv_std 22,190.9 is still N35 11,206 + N75 10,984. Same total either way;
only the GPS ceiling moves.

Fixes the PMS capture pipeline's hardcoded 2-market scope allowlist
(pms_scope_for_distributor's LCV_DISTRIBUTORS={"IUK","UMI"} / CV_TRUCK_DISTRIBUTORS=
{"ITUK","UTI"}, everything else silently defaulted to N75-truck scope). 14 of 18
markets were being scored against the wrong vehicle class as a result — this reads
the real per-country model-type UIO already carried in the UIO workbook's
"1.UIO Overview" sheet instead.

Country -> reporting distributor_code join uses distributor_master_mapping.csv's
uio_match column, except United Kingdom and Israel, which the UIO sheet itself
already splits row-by-row into ITUK/IUK and UTI/UMI (trusted directly). The Benelux
group (Belgium/Netherlands/Luxembourg/Poland) rolls up to the single "BELGIUM"
reporting code, matching how pms_capture_by_db_year.csv already reports it (real
business structure per distributor_master_mapping.csv's notes column: "Covers
BE/NL/PL/LUX"), not an assumption.

Output: data/uio_by_model_bucket.csv (distributor_code, cv_std_uio, f_series_uio, lcv_uio, source_year)
Run whenever the UIO workbook is refreshed.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

PRISM_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PRISM_ROOT))

from config import resolve_uio_path, resolve_mapping_path  # noqa: E402

OUTPUT_CSV = PRISM_ROOT / "data" / "uio_by_model_bucket.csv"

# UK and Israel: the UIO sheet already splits these row-by-row (CV rows tagged
# ITUK/UTI, the LCV row tagged IUK/UMI) — trust that column directly rather than
# the master-mapping join, which only has one uio_match code per country.
DIRECT_DISTRIBUTOR_COUNTRIES = {"U KING", "ISRAEL"}

# Known one-to-many reporting rollups (real business structure, not a guess —
# see distributor_master_mapping.csv notes column for each).
COUNTRY_ROLLUP = {
    "NETHLDS": "BELGIUM",
    "LUXMBRG": "BELGIUM",
    "POLAND": "BELGIUM",
}

# Real gaps/inconsistencies found by diffing the UIO sheet's actual "Country"
# codes against distributor_master_mapping.csv's uio_match column (2026-07-15):
# - Romania has no row in the master mapping at all, even though it's a live
#   PMS/GPS reporting distributor_code; the UIO sheet's own country label
#   already equals it verbatim.
# - Portugal's master-mapping uio_match is spelled "PORTUGAL", but the UIO
#   sheet's actual country code is the truncated "PORTUGL" (and "AZORS" for
#   the Azores territory) — a spelling mismatch in the master file, not a
#   business-structure choice. Confirmed no other ISZE-reporting country is
#   affected (the rest of the mismatches found are MIDI territories with no
#   real UIO data, already excluded, or parser artifacts from subtotal rows).
EXTRA_DIRECT_COUNTRY_CODES = {
    "ROMANIA": "ROMANIA",
    "PORTUGL": "PORTUGAL",
    "AZORS": "PORTUGAL",
}

# MIDI Europe (uio_match="ITALY") is explicitly excluded from GPS/PMS — no ISZE-
# direct sales data available, see distributor_master_mapping.csv notes and
# project_midi_not_customer.
EXCLUDED_DISTRIBUTOR_CODES = {"MIDI"}


def _country_to_distributor_map() -> dict[str, str]:
    mapping = pd.read_csv(resolve_mapping_path())
    out: dict[str, str] = {}
    for _, row in mapping.iterrows():
        code = COUNTRY_ROLLUP.get(row["uio_match"], row["distributor_code"])
        out[row["uio_match"]] = code
    return out


def build_uio_by_model_bucket() -> pd.DataFrame:
    wb = openpyxl.load_workbook(resolve_uio_path(), read_only=True, data_only=True)
    ws = wb["1.UIO Overview"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header = rows[5]
    # Pick the latest actuals/estimate column ("...E") rather than the last column,
    # which is a forward forecast ("...F") several years out.
    estimate_cols = [i for i, v in enumerate(header) if v and str(v).endswith("E")]
    if not estimate_cols:
        raise ValueError("No '...E' (estimate) year column found in UIO Overview header")
    latest_col = max(estimate_cols)
    latest_year_label = header[latest_col]

    country_map = _country_to_distributor_map()

    acc: dict[str, dict[str, float]] = {}
    for r in rows[6:]:
        if r[0] is None:
            continue
        country, cv_lcv, model_type, uio_distributor = r[1], r[2], r[3], r[5]
        val = r[latest_col] or 0
        if not val:
            continue

        if country in DIRECT_DISTRIBUTOR_COUNTRIES:
            code = uio_distributor
        elif country in EXTRA_DIRECT_COUNTRY_CODES:
            code = EXTRA_DIRECT_COUNTRY_CODES[country]
        else:
            code = country_map.get(country)
        if not code or code in EXCLUDED_DISTRIBUTOR_CODES:
            continue

        bucket = acc.setdefault(code, {"cv_std_uio": 0.0, "f_series_uio": 0.0, "lcv_uio": 0.0})
        if cv_lcv == "LCV":
            bucket["lcv_uio"] += val
        elif model_type == "F-Series":
            bucket["f_series_uio"] += val
        else:  # LD, MD
            bucket["cv_std_uio"] += val

    out_rows = [
        {"distributor_code": code, **vals, "source_year": latest_year_label}
        for code, vals in sorted(acc.items())
    ]
    return pd.DataFrame(out_rows)

if __name__ == "__main__":
    df = build_uio_by_model_bucket()
    df.to_csv(OUTPUT_CSV, index=False)
    print(f"Wrote {len(df)} distributors to {OUTPUT_CSV}")
    print(df.to_string(index=False))
