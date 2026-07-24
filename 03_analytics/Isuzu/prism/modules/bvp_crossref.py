"""
Genuine ↔ BVP part-number cross-reference from catalog rows + PMS registry.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from modules.bvp_catalog_io import (
    SKIP_SHEETS,
    CatalogColumns,
    detect_catalog_columns,
    normalize_pn,
    resolve_sheet_meta,
)
from modules.pms_engine import CATEGORY_LABELS, PMS_PN_REGISTRY


def _family_for_category(cat: str) -> str:
    if cat in ("oil_filter", "fuel_filter", "air_filter", "urea_filter", "blower_filter"):
        return "PMS_Filters"
    if "brake" in cat:
        return "Brakes"
    if cat in ("clutch_disc", "clutch_pressure_plate", "clutch_kit", "release_bearing", "pilot_bearing", "clutch_booster"):
        return "Clutch_Driveline"
    return "Other"


def extract_pairs_from_sheet(
    sheet_name: str,
    rows: list[tuple],
    cols: CatalogColumns,
) -> list[dict]:
    family, subgroup, pms_flag, pms_cat = resolve_sheet_meta(sheet_name)
    out: list[dict] = []
    last_bvp: str | None = None
    last_group: str | None = None
    carry_vehicle = carry_engine = carry_type = carry_origin = None

    for row in rows[cols.data_start_row :]:
        gen = normalize_pn(_row_val(row, cols.col_genuine))
        bvp = normalize_pn(_row_val(row, cols.col_bvp))
        line_no = _row_val(row, cols.col_no)
        part_type = _clean_attr(_row_val(row, cols.col_type))
        vehicle = _clean_attr(_row_val(row, cols.col_vehicle))
        engine = _clean_attr(_row_val(row, cols.col_engine))
        origin = _clean_attr(_row_val(row, cols.col_origin))

        if not any([gen, bvp, vehicle, engine, part_type, line_no]):
            continue

        if vehicle:
            carry_vehicle = vehicle
        if engine:
            carry_engine = engine
        if part_type:
            carry_type = part_type
        if origin:
            carry_origin = origin

        if bvp:
            last_bvp = bvp
            last_group = f"{sheet_name}::{line_no if line_no is not None else bvp}"

        if not gen or not last_bvp:
            continue

        pair_source = "catalog_paired" if bvp else "catalog_continuation"
        out.append(
            {
                "genuine_pn": gen,
                "bvp_pn": bvp or last_bvp,
                "product_family": family,
                "product_subgroup": subgroup,
                "pms_flag": pms_flag,
                "pms_category": pms_cat or "",
                "category_label": CATEGORY_LABELS.get(pms_cat, "") if pms_cat else "",
                "sheet_source": sheet_name,
                "catalog_line_no": line_no,
                "catalog_type": carry_type or "",
                "vehicle_model": carry_vehicle or "",
                "engine_type": carry_engine or "",
                "origin": carry_origin or "",
                "pair_source": pair_source,
                "crossref_group": last_group or "",
                "registry_model": "",
            }
        )

    return out


def _row_val(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _clean_attr(val) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    return s if s else None


def build_crossref_from_catalog(catalog_path: Path) -> pd.DataFrame:
    import openpyxl

    rows_all: list[dict] = []
    wb = openpyxl.load_workbook(catalog_path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        sheet_rows = list(ws.iter_rows(values_only=True))
        cols = detect_catalog_columns(sheet_rows)
        if cols is None:
            continue
        rows_all.extend(extract_pairs_from_sheet(sheet_name, sheet_rows, cols))
    wb.close()
    return pd.DataFrame(rows_all)


def build_crossref_from_registry(existing: pd.DataFrame | None = None) -> pd.DataFrame:
    """Add model×category pairs from PMS_PN_REGISTRY not already in catalog crossref."""
    existing_keys: set[tuple[str, str]] = set()
    if existing is not None and not existing.empty:
        existing_keys = set(
            zip(existing["genuine_pn"].astype(str), existing["bvp_pn"].astype(str))
        )

    buckets: dict[tuple[str, str], dict[str, list[str]]] = {}
    for pn, meta in PMS_PN_REGISTRY.items():
        pn_n = normalize_pn(pn)
        if not pn_n:
            continue
        key = (meta["model"], meta["category"])
        buckets.setdefault(key, {"genuine": [], "bvp": []})
        buckets[key][meta["type"]].append(pn_n)

    rows: list[dict] = []
    for (model, cat), parts in buckets.items():
        gens = sorted(set(parts["genuine"]))
        bvps = sorted(set(parts["bvp"]))
        if not gens or not bvps:
            continue
        family = _family_for_category(cat)
        for g in gens:
            for b in bvps:
                if (g, b) in existing_keys:
                    continue
                rows.append(
                    {
                        "genuine_pn": g,
                        "bvp_pn": b,
                        "product_family": family,
                        "product_subgroup": cat,
                        "pms_flag": True,
                        "pms_category": cat,
                        "category_label": CATEGORY_LABELS.get(cat, cat),
                        "sheet_source": "PMS_PN_REGISTRY",
                        "catalog_line_no": None,
                        "catalog_type": "",
                        "vehicle_model": "",
                        "engine_type": "",
                        "origin": "",
                        "pair_source": "registry_model_category",
                        "crossref_group": f"{model}::{cat}",
                        "registry_model": model,
                    }
                )
    return pd.DataFrame(rows)


def build_genuine_bvp_crossref(catalog_path: Path) -> pd.DataFrame:
    df_cat = build_crossref_from_catalog(catalog_path)
    df_reg = build_crossref_from_registry(df_cat)
    df = pd.concat([df_cat, df_reg], ignore_index=True)
    if df.empty:
        return df

    df = df.drop_duplicates(subset=["genuine_pn", "bvp_pn"], keep="first")
    df = df.sort_values(
        ["product_family", "product_subgroup", "sheet_source", "genuine_pn", "bvp_pn"]
    ).reset_index(drop=True)
    df.insert(0, "crossref_id", range(1, len(df) + 1))
    return df


def load_crossref(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(
        path,
        dtype={"genuine_pn": str, "bvp_pn": str, "crossref_id": int},
    )


def tag_sales_with_crossref(
    df_sales: pd.DataFrame, df_crossref: pd.DataFrame
) -> pd.DataFrame:
    """
    Tag PMS sales lines with crossref_id, role (genuine|bvp), and counterpart PN.
    One sale row can match multiple pairs if a PN maps to several BVP/genuine lines.
    """
    from modules.bvp_analytics import ensure_part_number

    if df_crossref.empty:
        return df_sales.copy()

    df = ensure_part_number(df_sales)
    xref_cols = [
        "crossref_id",
        "genuine_pn",
        "bvp_pn",
        "product_subgroup",
        "product_family",
        "pair_source",
        "vehicle_model",
        "engine_type",
    ]
    xref = df_crossref[[c for c in xref_cols if c in df_crossref.columns]].copy()

    g = df.merge(xref, left_on="part_number", right_on="genuine_pn", how="inner")
    g["crossref_role"] = "genuine"
    g["counterpart_pn"] = g["bvp_pn"]

    b = df.merge(xref, left_on="part_number", right_on="bvp_pn", how="inner")
    b["crossref_role"] = "bvp"
    b["counterpart_pn"] = b["genuine_pn"]

    return pd.concat([g, b], ignore_index=True)
