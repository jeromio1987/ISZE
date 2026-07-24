"""
Build flat bvp_pn_dimension from BVP PARTS CATALOG.xlsx + PRISM PMS PN registry (genuine).
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from modules.bvp_catalog_io import (
    SKIP_SHEETS,
    coalesce_pms_category,
    normalize_pn,
    resolve_sheet_meta,
)
from modules.bvp_product_nav import attach_navigation_columns
from modules.pms_engine import CATEGORY_LABELS, PMS_PN_REGISTRY


def _extract_pns_from_sheet(ws) -> list[str]:
    found: set[str] = set()
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            pn = normalize_pn(cell)
            if pn:
                found.add(pn)
    return sorted(found)


def build_dimension_from_catalog(catalog_path: Path) -> pd.DataFrame:
    import openpyxl

    rows: list[dict] = []
    wb = openpyxl.load_workbook(catalog_path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in SKIP_SHEETS:
            continue
        family, subgroup, pms_flag, pms_cat = resolve_sheet_meta(sheet_name)
        pms_cat = coalesce_pms_category(pms_cat, subgroup, pms_flag)
        ws = wb[sheet_name]
        for pn in _extract_pns_from_sheet(ws):
            rows.append(
                {
                    "part_number": pn,
                    "product_family": family,
                    "product_subgroup": subgroup,
                    "sheet_source": sheet_name,
                    "line_type": "bvp",
                    "in_bvp_catalog": True,
                    "pms_flag": pms_flag,
                    "pms_category": pms_cat,
                    "category_label": CATEGORY_LABELS.get(pms_cat, "") if pms_cat else "",
                }
            )
    wb.close()

    df = pd.DataFrame(rows)
    if df.empty:
        return df

    gen_rows: list[dict] = []
    for pn, meta in PMS_PN_REGISTRY.items():
        pn_n = normalize_pn(pn)
        if not pn_n:
            continue
        cat = meta["category"]
        reg_type = meta.get("type", "genuine")
        if reg_type not in ("genuine", "bvp"):
            reg_type = "genuine"
        gen_rows.append(
            {
                "part_number": pn_n,
                "product_family": "PMS_Filters"
                if cat in ("oil_filter", "fuel_filter", "air_filter")
                else "Brakes"
                if "brake" in cat
                else "Clutch_Driveline",
                "product_subgroup": cat,
                "sheet_source": "PMS_PN_REGISTRY",
                "line_type": reg_type,
                "in_bvp_catalog": reg_type == "bvp",
                "pms_flag": True,
                "pms_category": cat,
                "category_label": CATEGORY_LABELS.get(cat, cat),
            }
        )

    df_gen = pd.DataFrame(gen_rows)
    df = df.drop_duplicates("part_number", keep="first")
    df_gen = df_gen.drop_duplicates("part_number", keep="first")

    out = df.set_index("part_number")
    for pn, row in df_gen.set_index("part_number").iterrows():
        if pn not in out.index:
            out.loc[pn] = row
        else:
            out.loc[pn, "pms_flag"] = bool(out.loc[pn, "pms_flag"]) or bool(row["pms_flag"])
            if pd.isna(out.loc[pn, "pms_category"]) and pd.notna(row["pms_category"]):
                out.loc[pn, "pms_category"] = row["pms_category"]
                out.loc[pn, "category_label"] = row["category_label"]
            # Registry is authoritative for line_type when the PN is not in the BVP catalog extract.
            if row["sheet_source"] == "PMS_PN_REGISTRY" and not bool(out.loc[pn, "in_bvp_catalog"]):
                out.loc[pn, "line_type"] = row["line_type"]
                out.loc[pn, "in_bvp_catalog"] = bool(row["in_bvp_catalog"])

    out = out.reset_index()
    out = _realign_line_type_from_crossref(out, catalog_path)
    out = attach_navigation_columns(out)
    return out.sort_values(
        ["product_group", "product_subgroup", "part_number"]
    ).reset_index(drop=True)


def _realign_line_type_from_crossref(
    out: pd.DataFrame, catalog_path: Path
) -> pd.DataFrame:
    """Authoritative genuine/BVP split from the catalog cross-reference.

    The raw sheet extract tags *every* PN it sees as ``bvp`` because the BVP
    catalog lists genuine ↔ BVP pairs side by side. That mislabels the genuine
    column (414/431 genuine catalog PNs) as BVP. The cross-reference builder
    reads the two columns separately, so its ``genuine_pn``/``bvp_pn`` sets are
    the source of truth: a PN in ``bvp_pn`` IS bvp, a PN in ``genuine_pn`` IS
    genuine. Realign ``line_type``/``in_bvp_catalog`` against them.
    """
    from modules.bvp_crossref import build_genuine_bvp_crossref

    xref = build_genuine_bvp_crossref(catalog_path)
    if xref is None or xref.empty:
        return out

    bvp_pns = set(xref["bvp_pn"].dropna().astype(str).str.strip())
    gen_pns = set(xref["genuine_pn"].dropna().astype(str).str.strip())

    pn = out["part_number"].astype(str).str.strip()
    is_bvp = pn.isin(bvp_pns)
    is_gen = pn.isin(gen_pns) & ~is_bvp  # bvp_pn wins any (empirically empty) overlap

    out.loc[is_bvp, "line_type"] = "bvp"
    out.loc[is_bvp, "in_bvp_catalog"] = True
    out.loc[is_gen, "line_type"] = "genuine"
    out.loc[is_gen, "in_bvp_catalog"] = False
    return out


def load_dimension(path: Path) -> pd.DataFrame:
    if path.suffix == ".parquet":
        return pd.read_parquet(path)
    return pd.read_csv(path, dtype={"part_number": str})
