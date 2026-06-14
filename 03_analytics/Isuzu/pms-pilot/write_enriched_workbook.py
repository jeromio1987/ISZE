"""
Append EPC enrichment columns to the FIG-KEY Parts sheet of the v5 workbook.

Source: epc_enriched.json (parser output)
Target: ITUK_PMS_Analysis_Apr2026_v5.xlsx (writes v6 with new sheet variant)
"""
import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ENRICH_COLUMNS = [
    ('EPC_Found', 'epc_found'),
    ('EPC_Total', 'epc_total_results'),
    ('EPC_Editions', 'epc_editions'),
    ('EPC_AvgPositions', 'epc_avg_positions'),
    ('EPC_MultiPosFlag', 'epc_multi_position_flag'),
    ('EPC_Vehicle_Families', 'epc_vehicle_families'),
    ('EPC_Buckets', 'epc_buckets'),
    ('EPC_Model_Series', 'epc_model_series'),
    ('EPC_GVW_Class', 'epc_gvw_classes'),
    ('EPC_Markets', 'epc_markets'),
    ('EPC_Emissions', 'epc_emissions'),
    ('EPC_Year_Min', 'epc_year_min'),
    ('EPC_Year_Max', 'epc_year_max'),
    ('EPC_Partial', 'epc_partial'),
    ('EPC_Error', 'epc_error'),
]


def main():
    src = 'ITUK_PMS_Analysis_Apr2026_v5.xlsx'
    out = 'ITUK_PMS_Analysis_Apr2026_v6_EPC_enriched.xlsx'

    with open('epc_enriched.json', encoding='utf-8') as f:
        enriched_list = json.load(f)
    by_pn = {e['pn']: e for e in enriched_list}

    wb = openpyxl.load_workbook(src)
    ws = wb['FIG-KEY Parts']

    # Find the column index for "Newest P/N" (col B in v5)
    # Find header row by looking for cell that says 'Newest P/N'
    header_row = None
    newest_col = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.strip() == 'Newest P/N':
                header_row = cell.row
                newest_col = cell.column
                break
        if header_row:
            break
    if not header_row:
        raise SystemExit('Could not locate Newest P/N column.')

    # Find rightmost existing column on header row
    last_col = ws.max_column
    while last_col > 0 and ws.cell(row=header_row, column=last_col).value is None:
        last_col -= 1
    start_col = last_col + 1

    # Write enrichment headers
    bold = Font(bold=True)
    hdr_fill = PatternFill('solid', fgColor='305496')
    hdr_font = Font(bold=True, color='FFFFFF')
    for offset, (col_name, _) in enumerate(ENRICH_COLUMNS):
        c = ws.cell(row=header_row, column=start_col + offset, value=col_name)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center', wrap_text=True)

    # Fill data rows
    written = 0
    not_found_fill = PatternFill('solid', fgColor='FCE4D6')  # light orange for N
    multi_pos_fill = PatternFill('solid', fgColor='FFF2CC')  # light yellow

    for r in range(header_row + 1, ws.max_row + 1):
        pn_cell = ws.cell(row=r, column=newest_col)
        pn = pn_cell.value
        if not pn:
            continue
        pn = str(pn).strip()
        data = by_pn.get(pn)
        if not data:
            # Mark pending
            ws.cell(row=r, column=start_col, value='(pending)')
            continue
        written += 1
        for offset, (col_name, key) in enumerate(ENRICH_COLUMNS):
            v = data.get(key)
            cell = ws.cell(row=r, column=start_col + offset, value=v if v not in ('', None) else None)
            if col_name == 'EPC_Found' and v == 'N':
                cell.fill = not_found_fill
            if col_name == 'EPC_MultiPosFlag' and v == 'Y':
                cell.fill = multi_pos_fill

    # Save
    wb.save(out)
    print(f'Wrote {out}')
    print(f'  Rows enriched: {written}')
    print(f'  Enrichment columns added: {len(ENRICH_COLUMNS)} starting at column {start_col}')


if __name__ == '__main__':
    main()
