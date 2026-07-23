import openpyxl
from collections import defaultdict
from datetime import datetime
import json

EXCL_BILLT = {'S1','G2','RE','ZFD','S2'}
EXCL_SATY = {'CR','RE'}

def fy(dt):
    return dt.year + 1 if dt.month >= 4 else dt.year

PAYER_DB = {
    '2194': 'ITUK', '89': 'IUK', '90': 'IBX',
    '92': 'Imotors', '93': 'ISD', '66': 'RSA', '2181': 'Sverige',
    '76': 'Ireland', '393': 'Loutsios', '86': 'Norway', '396': 'ROM Auto',
    '537': 'AutoWallis', '1759': 'GasanZammit', '2679': 'Nordic FI',
    '2027': 'Danmark', '1576': 'BL Iceland', '88': 'MIDI',
    '105': 'IMG', '2294': 'CIS-KTK', '2731': 'GA', '2732': 'CBC',
    '2372': 'KIMA', '77': 'TR', '2929': 'Sverige2'
}

results = defaultdict(lambda: defaultdict(float))

for fpath in [
    '/sessions/beautiful-kind-bell/mnt/ISZE/01_data/raw/sales/CY2023-2024.XLSX',
    '/sessions/beautiful-kind-bell/mnt/ISZE/01_data/CY2024-2025.XLSX',
    '/sessions/beautiful-kind-bell/mnt/ISZE/01_data/sales_2026tillApril.XLSX',
]:
    fname = fpath.split('/')[-1]
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    if not h[0]:
        print(f'{fname}: SKIP pivot'); continue
    billt_i = list(h).index('BillT')
    saty_i = list(h).index('SaTy')
    netval_i = list(h).index('Net value')
    kept = excl = 0
    for r in rows[1:]:
        if not isinstance(r[0], datetime): continue
        if r[billt_i] in EXCL_BILLT or r[saty_i] in EXCL_SATY:
            excl += 1; continue
        if str(r[4]) in ('1797', '1799'): continue
        db = PAYER_DB.get(str(r[4]), f'UNK_{r[4]}')
        results[db][fy(r[0])] += r[netval_i] or 0
        kept += 1
    print(f'{fname}: kept={kept:,} excl={excl:,}')
    wb.close()

wb = openpyxl.load_workbook("/sessions/beautiful-kind-bell/mnt/ISZE/01_data/Sales_israel_1799&_1797.XLSX", read_only=True, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
h = rows[0]
billt_i = list(h).index('BillT'); saty_i = list(h).index('SaTy'); netval_i = list(h).index('Net value')
kept = excl = 0
for r in rows[1:]:
    if not isinstance(r[0], datetime): continue
    if r[billt_i] in EXCL_BILLT or r[saty_i] in EXCL_SATY:
        excl += 1; continue
    db = 'UTI' if str(r[5]) == '1799' else 'UMI'
    results[db][fy(r[0])] += r[netval_i] or 0
    kept += 1
print(f'Israel: kept={kept:,} excl={excl:,}')
wb.close()

output = {db: {str(y): round(v, 2) for y, v in fydata.items()} for db, fydata in results.items()}
with open('/sessions/beautiful-kind-bell/mnt/ISZE/01_data/derived/gps_billing_clean.json', 'w') as f:
    json.dump(output, f, indent=2)

print(f'\n{"DB":<15} {"FY2024":>12} {"FY2025":>12} {"FY2026*":>12}')
print('-'*55)
for db in sorted(results, key=lambda d: -results[d].get(2025, 0)):
    r = results[db]
    print(f'{db:<15} {r.get(2024,0)/1000:>10.1f}k  {r.get(2025,0)/1000:>10.1f}k  {r.get(2026,0)/1000:>10.1f}k')
print('\nDone.')
