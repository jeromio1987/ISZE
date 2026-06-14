"""Load UIO v18 + ITUK DVLA car park + gps_results CY2025 for GPS v9."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
UIO_V18 = ROOT / "Data" / "ISUZU_UIO_Overview_2026_v18.xlsx"
CARPARK_XLSX = ROOT / "Isuzu" / "UIO" / "Script" / "Parts Sales Calculation FY26.xlsx"
GPS_RESULTS = ROOT / "Data" / "_build" / "gps_results.json"
OUT_JSON = ROOT / "Data" / "_build" / "uio_v18_gps.json"

RATE_N35 = 295.68
RATE_N75 = 739.20
RATE_FMD = 887.04
RATE_LCV = 250.0

ITUK_GP = 0.34
ITUK_FX = 1.12

BP_ITUK = {
    2026: {"retail_gbp": 13_129_200, "isze_eur": 6_974_170.05},
    2027: {"retail_gbp": 13_785_660, "isze_eur": 7_425_511.43},
}
KOSHA_VP_EUR = {2027: 35_300_000}

DIST_MAP = {
    ("BELGIUM", "IBX"): "BELGIUM",
    ("CYPRUS", "Loutsios & Sons"): "CYPRUS",
    ("DENMARK", "Isuzu Denmark"): "DENMARK",
    ("FINLAND", "Nordic Auto."): "FINLAND",
    ("GERMANY", "ISD-Lang"): "GERMANY",
    ("GREECE", "Petros Petropoulos"): "GREECE",
    ("HUNGARY", "AutoWallis"): "HUNGARY",
    ("ICELAND", "BL"): "ICELAND",
    ("IRELAND", "Isuzu Ireland"): "IRELAND",
    ("ISRAEL", "UMI"): "UMI",
    ("ITALY", "MIDI Europe"): "UTI",
    ("MALTA", "GasanZammit"): "MALTA",
    ("NORWAY", "RSA"): "NORWAY",
    ("PORTUGL", "Imotors"): "PORTUGAL",
    ("ROMANIA", "Isuzu Rom Auto"): "ROMANIA",
    ("SWEDEN", "Isuzu Sverige"): "SWEDEN",
    ("U KING", "IUK"): "IUK",
    ("U KING", "ITUK"): "ITUK",
}

GPS_DBS = [
    "ITUK", "IUK", "UMI", "UTI", "GERMANY", "BELGIUM", "IRELAND", "NORWAY",
    "SWEDEN", "DENMARK", "FINLAND", "ICELAND", "CYPRUS", "GREECE", "MALTA",
    "ROMANIA", "HUNGARY", "PORTUGAL",
]

# Which UIO buckets count toward max spend per distributor
SCOPE = {
    "IUK": ("LCV",),
    "UMI": ("LCV",),
    "UTI": ("LD", "FMD"),
    "NORWAY": ("LCV",),
    "SWEDEN": ("LCV",),
    "DENMARK": ("LCV",),
    "FINLAND": ("LCV",),
    "ICELAND": ("LCV",),
    "HUNGARY": ("LCV",),
    "ROMANIA": ("LCV",),
}
DEFAULT_SCOPE = ("LD", "FMD", "LCV")


def _bucket(mtype: str) -> str:
    mt = str(mtype or "").upper()
    if "LCV" in mt or "D-MAX" in mt or "DMAX" in mt:
        return "LCV"
    if mt in ("MD", "F-SERIES", "F SERIES") or mt == "F":
        return "FMD"
    return "LD"


def _scope_buckets(db: str) -> tuple[str, ...]:
    if db == "ITUK":
        return ("N35", "N75_LD", "F_MD")
    return SCOPE.get(db, DEFAULT_SCOPE)


def load_ituk_carpark() -> dict[str, int]:
    wb = openpyxl.load_workbook(CARPARK_XLSX, read_only=True, data_only=True)
    ws = wb["UIO 24"]
    totals: dict[str, int] = defaultdict(int)
    for row in ws.iter_rows(min_row=2, values_only=True):
        short = str(row[3] or "").strip()
        n = int(row[4] or 0)
        if short and n > 0:
            totals[short] += n
    wb.close()
    return dict(totals)


def load_ituk_projection_models() -> dict[str, list[float]]:
    wb = openpyxl.load_workbook(UIO_V18, read_only=True, data_only=True)
    ws = wb["4. ITUK UIO Projection"]
    models: dict[str, list[float]] = {}
    for row in ws.iter_rows(min_row=20, max_row=27, values_only=True):
        name = str(row[1] or "").strip()
        if name and name != "Total ITUK UIO":
            models[name] = [float(row[i] or 0) for i in (2, 3, 4)]
    wb.close()
    return models


def _ituk_buckets_from_models(models: dict[str, list[float]], idx: int) -> dict[str, float]:
    return {
        "N35": models["N35"][idx],
        "N75_LD": (
            models["N75"][idx]
            + models["N50/55"][idx]
            + models["N62/65"][idx]
            + models["UNSPEC"][idx]
        ),
        "F_MD": (
            models["F110/120"][idx]
            + models["F135"][idx]
            + models["F150/180"][idx]
        ),
    }


def _ituk_buckets_from_carpark(cp: dict[str, int]) -> dict[str, float]:
    return {
        "N35": float(cp.get("N35", 0)),
        "N75_LD": float(
            cp.get("N75", 0)
            + cp.get("N50/55", 0)
            + cp.get("N62/65", 0)
            + cp.get("UNSPECIFIED", 0)
        ),
        "F_MD": float(
            cp.get("F110/120", 0) + cp.get("F135", 0) + cp.get("F150/180", 0)
        ),
    }


def load_network_uio_v18() -> dict[str, dict[int, dict[str, float]]]:
    wb = openpyxl.load_workbook(UIO_V18, read_only=True, data_only=True)
    ws = wb["1.UIO Overview"]
    all_rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(all_rows) if r[0] is not None and "Ctry" in str(r[0]))
    hdr = all_rows[hdr_i]
    ycols = {
        int(m.group(1)): j
        for j, h in enumerate(hdr)
        if h and (m := re.match(r"^\s*(\d{4})[EF]?\s*$", str(h)))
        and 2024 <= int(m.group(1)) <= 2030
    }
    agg: dict[str, dict[int, dict[str, float]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(float))
    )
    for r in all_rows[hdr_i + 1 :]:
        if r[1] is None or r[5] is None:
            continue
        db = DIST_MAP.get((str(r[1]).strip(), str(r[5]).strip()))
        if not db or db == "ITUK":
            continue
        b = _bucket(r[3])
        for y, j in ycols.items():
            if isinstance(r[j], (int, float)):
                agg[db][y][b] += float(r[j])
    wb.close()
    return {k: {yy: dict(buckets) for yy, buckets in v.items()} for k, v in agg.items()}


def load_gps_results_2025() -> dict[str, dict[str, float]]:
    rows = json.load(open(GPS_RESULTS, encoding="utf-8-sig"))
    out: dict[str, dict[str, float]] = {}
    for r in rows:
        if r.get("year") != 2025:
            continue
        db = r["db"]
        if db not in GPS_DBS or db == "ITUK":
            continue
        out[db] = {
            "LD": float(r.get("uio_cv_ld") or 0),
            "FMD": float((r.get("uio_f") or 0) + (r.get("uio_md") or 0)),
            "LCV": float(r.get("uio_lcv") or 0),
        }
    return out


def _scoped_total(buckets: dict[str, float], db: str) -> float:
    scope = _scope_buckets(db)
    if db == "ITUK":
        return sum(buckets.get(k, 0) for k in scope)
    return sum(buckets.get(k, 0) for k in scope if k in buckets)


def max_spend_from_buckets(buckets: dict[str, float], db: str) -> float:
    if db == "ITUK":
        return (
            buckets.get("N35", 0) * RATE_N35
            + buckets.get("N75_LD", 0) * RATE_N75
            + buckets.get("F_MD", 0) * RATE_FMD
        )
    return (
        buckets.get("LD", 0) * RATE_N75
        + buckets.get("FMD", 0) * RATE_FMD
        + buckets.get("LCV", 0) * RATE_LCV
    )


def _scale_buckets(base: dict[str, float], factor: float, db: str) -> dict[str, float]:
    scope = _scope_buckets(db)
    return {k: round(v * factor, 1) if k in scope else 0.0 for k, v in base.items()}


def build_gps_uio_package() -> dict:
    carpark = load_ituk_carpark()
    ituk_models = load_ituk_projection_models()
    v18 = load_network_uio_v18()
    gps25 = load_gps_results_2025()

    ituk_uio: dict[int, dict] = {
        2025: _ituk_buckets_from_carpark(carpark),
        2026: _ituk_buckets_from_models(ituk_models, 1),
        2027: _ituk_buckets_from_models(ituk_models, 2),
    }

    distributors: dict[str, dict] = {
        "ITUK": {
            "scope": "CV — DVLA car park CY2025 · v18 model 2026–27",
            "uio": ituk_uio,
            "max_spend": {y: max_spend_from_buckets(b, "ITUK") for y, b in ituk_uio.items()},
            "rates_note": "N35 €296 · N75/LD €739 · F/MD €887",
        }
    }

    for db in GPS_DBS:
        if db == "ITUK":
            continue
        base = gps25[db]
        v18_25 = v18.get(db, {}).get(2025, {})
        v18_26 = v18.get(db, {}).get(2026, {})
        v18_27 = v18.get(db, {}).get(2027, {})

        def _growth(target: dict[str, float]) -> float:
            t25 = _scoped_total(v18_25, db)
            t_tgt = sum(target.get(k, 0) for k in _scope_buckets(db) if k in target)
            return (t_tgt / t25) if t25 > 0 else 1.0

        uio_by_year = {
            2025: {k: int(round(v)) for k, v in base.items()},
            2026: {
                k: int(round(v))
                for k, v in _scale_buckets(base, _growth(v18_26), db).items()
            },
            2027: {
                k: int(round(v))
                for k, v in _scale_buckets(base, _growth(v18_27), db).items()
            },
        }
        distributors[db] = {
            "scope": (
                "LCV only"
                if SCOPE.get(db) == ("LCV",)
                else ("CV only" if db == "UTI" else "CV + LCV")
            ),
            "uio": uio_by_year,
            "max_spend": {
                y: max_spend_from_buckets(
                    {k: float(v) for k, v in b.items()}, db
                )
                for y, b in uio_by_year.items()
            },
            "rates_note": "LD €739 · F/MD €887 · LCV €250",
        }

    return {
        "source_uio": str(UIO_V18),
        "source_uio_cy2025": str(GPS_RESULTS),
        "source_ituk_carpark": str(CARPARK_XLSX),
        "rates": {
            "N35": RATE_N35,
            "N75_LD": RATE_N75,
            "F_MD": RATE_FMD,
            "LCV": RATE_LCV,
            "gp": ITUK_GP,
            "fx": ITUK_FX,
        },
        "ituk_carpark_raw": carpark,
        "ituk_models_v18": ituk_models,
        "distributors": distributors,
        "bp_ituk": BP_ITUK,
        "kosha_vp_eur": KOSHA_VP_EUR,
    }


def main() -> None:
    pkg = build_gps_uio_package()
    OUT_JSON.write_text(json.dumps(pkg, indent=2), encoding="utf-8")
    print("wrote", OUT_JSON)
    for db in ("ITUK", "UTI", "GERMANY"):
        d = pkg["distributors"][db]
        print(db, d["uio"].get(2025), "max25", d["max_spend"][2025])


if __name__ == "__main__":
    main()
