"""Build ISZE_GPS_Calculator_v9.xlsx — UIO v18 + DVLA car park + study max spend rates."""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gps_v9_loaders import (
    BP_ITUK,
    GPS_DBS,
    KOSHA_VP_EUR,
    RATE_FMD,
    RATE_LCV,
    RATE_N35,
    RATE_N75,
    build_gps_uio_package,
)

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
OUT = ROOT / "Data" / "ISZE_GPS_Calculator_v9.xlsx"

BLUE = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
H1 = Font(name="Arial", size=13, bold=True, color="CC0E1E")
SMALL = Font(name="Arial", size=8, italic=True, color="808080")
HDR = Font(name="Arial", size=9, bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", start_color="CC0E1E")
YELLOW = PatternFill("solid", start_color="FFFF00")
GREEN = Font(name="Arial", size=10, color="008000")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# CY2025 ISZE→DB sales (ERP / sales_fact) — unchanged
SALES_25 = {
    "ITUK": 5_769_867.82,
    "IUK": 6_221_920.41,
    "UMI": 3_388_801.28,
    "UTI": 3_402_116.94,
    "GERMANY": 2_577_765.90,
    "BELGIUM": 2_143_951.90,
    "IRELAND": 1_213_660.33,
    "NORWAY": 1_438_472.63,
    "SWEDEN": 1_289_481.04,
    "DENMARK": 150_218.49,
    "FINLAND": 285_912.53,
    "ICELAND": 87_259.89,
    "CYPRUS": 429_512.03,
    "GREECE": 528_390.82,
    "MALTA": 206_502.27,
    "ROMANIA": 177_276.35,
    "HUNGARY": 330_030.58,
    "PORTUGAL": 1_712_790.96,
}

WARRANTY = {
    "ITUK": 374_426.05,
    "IUK": 1_054_395,
    "UMI": 0,
    "UTI": 255_967.77,
    "GERMANY": 259_675,
    "BELGIUM": 74_677,
    "IRELAND": 34_857,
    "NORWAY": 290_712,
    "SWEDEN": 189_618,
    "DENMARK": 9_115,
    "FINLAND": 48_676,
    "ICELAND": 7_035,
    "CYPRUS": 28_041,
    "GREECE": 56_159,
    "MALTA": 18_219,
    "ROMANIA": 19_377,
    "HUNGARY": 77_434,
    "PORTUGAL": 130_588,
}

pkg = build_gps_uio_package()
DIST = pkg["distributors"]


def _uio_total(db: str, year: int) -> int:
    u = DIST[db]["uio"][year]
    return int(round(sum(v for v in u.values() if isinstance(v, (int, float)))))


def _uio_detail_str(db: str, year: int) -> str:
    u = DIST[db]["uio"][year]
    if db == "ITUK":
        return f"N35 {u['N35']:,.0f} · N75/LD {u['N75_LD']:,.0f} · F/MD {u['F_MD']:,.0f}"
    parts = []
    if u.get("LD"):
        parts.append(f"LD {u['LD']:,}")
    if u.get("FMD"):
        parts.append(f"F/MD {u['FMD']:,}")
    if u.get("LCV"):
        parts.append(f"LCV {u['LCV']:,}")
    return " · ".join(parts) if parts else "—"


base = [
    (
        db,
        DIST[db]["scope"],
        _uio_total(db, 2025),
        DIST[db]["max_spend"][2025],
        SALES_25[db],
    )
    for db in GPS_DBS
]

hist = [
    r
    for r in json.load(open(ROOT / "Data/_build/gps_results.json", encoding="utf-8-sig"))
    if r["db"] != "Distributor" and isinstance(r["year"], int)
]

wb = openpyxl.Workbook()

# ============ Read Me
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
lines = [
    ("ISZE GPS Calculator v9", H1),
    ("", BLACK),
    ("Management language", BOLD),
    ("• Say: max spend per unit (€/vehicle/yr at ISZE billing).", BLACK),
    ("• Do NOT say: wallet.", BLACK),
    ("• GPS % = ISZE→DB parts sales ÷ max spend (UIO × model rates). ISZE billing only.", BLACK),
    ("", BLACK),
    ("Rates (all CV distributors)", BOLD),
    ("• N35 Grafter €296 · N75/LD CV €739 · F/MD €887 · LCV/D-Max €250.", BLACK),
    ("• Formula: retail £ × (1−GP 34%) × FX 1.12 — replaces legacy €1,120 / €1,344.", BLACK),
    ("", BLACK),
    ("UIO sources", BOLD),
    ("• ITUK CY2025: DVLA car park (Parts Sales Calculation FY26) — 22,389 CV.", BLACK),
    ("• ITUK 2026–27: UIO v18 sheet 4 model projection.", BLACK),
    ("• Network: ISUZU_UIO_Overview_2026_v18.xlsx (LD / F+MD / LCV).", BLACK),
    ("", BLACK),
    ("Tabs", BOLD),
    ("1 Max Spend Rates — standard CV/LCV rates + ITUK car park build", BLACK),
    ("2 GPS 2025 — actuals, warranty incl/ex (no blended €/unit)", BLACK),
    ("3 Projection FY26", BLACK),
    ("4 BP FY2027 — GPS pp gap to business plan", BLACK),
    ("5 Assumptions", BLACK),
]
for i, (txt, font) in enumerate(lines, 1):
    ws.cell(i, 1, txt).font = font
ws.column_dimensions["A"].width = 105

# ============ Max Spend Rates
ws = wb.create_sheet("Max Spend Rates")
ws.sheet_view.showGridLines = False
ws["A1"] = "MAX SPEND PER UNIT — ISZE billing (€/vehicle/yr). Blue = adjustable."
ws["A1"].font = H1
ws["A2"] = "All CV distributors use study-derived rates (GP 34%, FX 1.12). Legacy €1,120/€1,344 retired."
ws["A2"].font = SMALL

hdrs = ["Segment", "Retail £/veh/yr", "GP %", "ISZE max spend €/veh/yr", "Applied to"]
for j, h in enumerate(hdrs, 1):
    c = ws.cell(4, j, h)
    c.font = HDR
    c.fill = HDRFILL
    c.border = BORDER

std_rates = [
    ("N35 Grafter", 400, RATE_N35, "ITUK car park N35"),
    ("N75 / N50 / N62 (LD CV)", 1000, RATE_N75, "ITUK N75+ · network LD"),
    ("F-series / MD", 1200, RATE_FMD, "ITUK F · network F/MD"),
    ("LCV / D-Max", None, RATE_LCV, "IUK · UMI · mixed markets"),
]
for i, (seg, retail, rate, note) in enumerate(std_rates):
    r = 5 + i
    ws.cell(r, 1, seg).font = BOLD
    if retail is not None:
        rc = ws.cell(r, 2, retail)
        rc.font = BLUE
        rc.border = BORDER
        rc.number_format = "£#,##0"
        ws.cell(r, 3, 0.34).number_format = "0%"
    else:
        ws.cell(r, 2, "—")
        ws.cell(r, 3, "—")
    c = ws.cell(r, 4, rate)
    c.font = BLUE
    c.border = BORDER
    c.number_format = "€#,##0"
    ws.cell(r, 5, note).font = SMALL

# ITUK car park block
cp = pkg["ituk_carpark_raw"]
b = 5 + len(std_rates) + 2
ws.cell(b, 1, "ITUK — DVLA car park CY2025 (max spend build)").font = H1
ws.cell(b + 1, 1, f"Source: Parts Sales Calculation FY26 · UIO v18 for 2026–27 · GP 34% · FX 1.12").font = SMALL
mh = ["Bucket", "UIO 2025", "€/veh/yr", "Max spend 2025 €", "UIO 2026F", "Max spend 2026F €"]
for j, h in enumerate(mh, 1):
    c = ws.cell(b + 3, j, h)
    c.font = HDR
    c.fill = HDRFILL
    c.border = BORDER

ituk_rows = [
    ("N35 Grafter", cp.get("N35", 0), RATE_N35, DIST["ITUK"]["uio"][2026]["N35"]),
    (
        "N75 / N50 / N62 / unspec",
        cp.get("N75", 0) + cp.get("N50/55", 0) + cp.get("N62/65", 0) + cp.get("UNSPECIFIED", 0),
        RATE_N75,
        DIST["ITUK"]["uio"][2026]["N75_LD"],
    ),
    (
        "F-series",
        cp.get("F110/120", 0) + cp.get("F135", 0) + cp.get("F150/180", 0),
        RATE_FMD,
        DIST["ITUK"]["uio"][2026]["F_MD"],
    ),
]
ITUK_RATE_START = b + 4
for i, (label, u25, rate, u26) in enumerate(ituk_rows):
    r = ITUK_RATE_START + i
    ws.cell(r, 1, label).font = BOLD
    c25 = ws.cell(r, 2, u25)
    c25.font = BLUE
    c25.border = BORDER
    c25.number_format = "#,##0"
    ws.cell(r, 3, rate).number_format = "€#,##0"
    ws.cell(r, 4, f"=B{r}*C{r}").number_format = "€#,##0"
    c26 = ws.cell(r, 5, round(u26))
    c26.font = BLUE
    c26.border = BORDER
    c26.number_format = "#,##0"
    ws.cell(r, 6, f"=E{r}*C{r}").number_format = "€#,##0"

tr = ITUK_RATE_START + len(ituk_rows)
ws.cell(tr, 1, "ITUK CV total").font = BOLD
for col in (2, 4, 5, 6):
    letter = get_column_letter(col)
    ws.cell(tr, col, f"=SUM({letter}{ITUK_RATE_START}:{letter}{tr-1})").font = BOLD
    ws.cell(tr, col).number_format = "€#,##0" if col in (4, 6) else "#,##0"
ITUK_MAX_25_CELL = f"D{tr}"
ITUK_MAX_26_CELL = f"F{tr}"

for j, w in enumerate([22, 12, 8, 14, 12, 14], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ============ GPS 2025
ws = wb.create_sheet("GPS 2025")
ws.sheet_view.showGridLines = False
ws["A1"] = "GPS 2025 — ISZE→DB actuals vs max spend (UIO v18 / ITUK car park)"
ws["A1"].font = H1
ws["A2"] = "No blended €/unit column — mixed CV/LCV parc makes SPU misleading. Use GPS % only."
ws["A2"].font = SMALL
hdrs = [
    "Distributor", "Scope", "UIO 2025", "Max spend (€)", "ISZE sales incl. warr (€)",
    "Warranty (€) [INPUT]", "Sales ex. warr (€)", "GPS % incl.", "GPS % ex.",
]
for j, h in enumerate(hdrs, 1):
    c = ws.cell(4, j, h)
    c.font = HDR
    c.fill = HDRFILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")

for i, (db, scope, uio, maxsp, sales) in enumerate(base):
    r = 5 + i
    ws.cell(r, 1, db).font = BOLD
    ws.cell(r, 2, scope).font = BLACK
    ws.cell(r, 3, uio).number_format = "#,##0"
    if db == "ITUK":
        c = ws.cell(r, 4, f"='Max Spend Rates'!{ITUK_MAX_25_CELL}")
        c.font = GREEN
    else:
        ws.cell(r, 4, maxsp).number_format = "€#,##0"
    ws.cell(r, 5, sales).number_format = "€#,##0"
    cw = ws.cell(r, 6, WARRANTY.get(db, 0))
    cw.font = BLUE
    cw.border = BORDER
    cw.number_format = "€#,##0"
    ws.cell(r, 7, f"=E{r}-F{r}").number_format = "€#,##0"
    ws.cell(r, 8, f"=E{r}/D{r}").number_format = "0.0%"
    ws.cell(r, 9, f"=G{r}/D{r}").number_format = "0.0%"

tr = 5 + len(base)
ws.cell(tr, 1, "TOTAL").font = BOLD
for col, letter in ((3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G")):
    c = ws.cell(tr, col, f"=SUM({letter}5:{letter}{tr-1})")
    c.font = BOLD
    c.number_format = "#,##0" if col == 3 else "€#,##0"
ws.cell(tr, 8, f"=E{tr}/D{tr}").font = BOLD
ws.cell(tr, 8).number_format = "0.0%"
ws.cell(tr, 9, f"=G{tr}/D{tr}").font = BOLD
ws.cell(tr, 9).number_format = "0.0%"

BASE_ROW = {db: 5 + i for i, (db, *_r) in enumerate(base)}

for j, w in enumerate([12, 18, 9, 13, 15, 12, 14, 9, 9], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ============ Projection FY26
ws = wb.create_sheet("Projection FY26")
ws.sheet_view.showGridLines = False
ws["A1"] = "FY2026 projection — GPS uplift vs 2025 (pp)"
ws["A1"].font = H1
ws["A4"] = "Uplift (pp):"
ws["A4"].font = BOLD
up = ws.cell(4, 3, 0.01)
up.font = BLUE
up.fill = YELLOW
up.border = BORDER
up.number_format = "0.0%"
UP = "$C$4"

hdrs = [
    "Distributor", "UIO 26 (detail)", "Max spend FY26 (€)", "GPS 2025", "GPS FY26 proj",
    "Proj. sales €", "Sales 2025 €", "Δ vs 2025 €",
]
for j, h in enumerate(hdrs, 1):
    c = ws.cell(6, j, h)
    c.font = HDR
    c.fill = HDRFILL
    c.border = BORDER

r0 = 7
for i, (db, scope, _uio, _maxsp, sales) in enumerate(base):
    r = r0 + i
    br = BASE_ROW[db]
    ws.cell(r, 1, db).font = BOLD
    ws.cell(r, 2, _uio_detail_str(db, 2026)).font = SMALL
    if db == "ITUK":
        c = ws.cell(r, 3, f"='Max Spend Rates'!{ITUK_MAX_26_CELL}")
    else:
        c = ws.cell(r, 3, DIST[db]["max_spend"][2026])
    c.font = GREEN
    c.number_format = "€#,##0"
    c = ws.cell(r, 4, f"='GPS 2025'!H{br}")
    c.font = GREEN
    c.number_format = "0.0%"
    ws.cell(r, 5, f"=D{r}+{UP}").number_format = "0.0%"
    ws.cell(r, 6, f"=E{r}*C{r}").number_format = "€#,##0"
    c = ws.cell(r, 7, f"='GPS 2025'!E{br}")
    c.font = GREEN
    c.number_format = "€#,##0"
    ws.cell(r, 8, f"=F{r}-G{r}").number_format = "€#,##0;(€#,##0);-"

tr = r0 + len(base)
ws.cell(tr, 1, "TOTAL").font = BOLD
for col in (3, 6, 7, 8):
    letter = get_column_letter(col)
    ws.cell(tr, col, f"=SUM({letter}{r0}:{letter}{tr-1})").font = BOLD
    ws.cell(tr, col).number_format = "€#,##0"
ws.cell(tr, 5, f"=F{tr}/C{tr}").font = BOLD
ws.cell(tr, 5).number_format = "0.0%"

# ============ BP FY2027 GPS gap
ws = wb.create_sheet("BP FY2027")
ws.sheet_view.showGridLines = False
ws["A1"] = "FY2027 — GPS increase needed to reach business plan"
ws["A1"].font = H1
ws["A2"] = (
    "ITUK: BP ISZE plan (€7.43M). Other distributors: Kosha VP €35.3M allocated by CY2025 sales share."
)
ws["A2"].font = SMALL

hdrs = [
    "Distributor", "Max spend FY27 (€)", "GPS 2025 ex.", "Sales 2025 ex. (€)",
    "BP / plan FY27 (€)", "GPS needed", "Δ GPS pp", "Add'l sales €", "€ per +1pp",
]
for j, h in enumerate(hdrs, 1):
    c = ws.cell(4, j, h)
    c.font = HDR
    c.fill = HDRFILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")

total_sales_25 = sum(SALES_25.values())
kosha_27 = KOSHA_VP_EUR[2027]
bp_targets = {db: kosha_27 * (SALES_25[db] / total_sales_25) for db in GPS_DBS}
bp_targets["ITUK"] = BP_ITUK[2027]["isze_eur"]

r0 = 5
for i, db in enumerate(GPS_DBS):
    r = r0 + i
    br = BASE_ROW[db]
    max27 = DIST[db]["max_spend"][2027]
    ws.cell(r, 1, db).font = BOLD
    if db == "ITUK":
        c = ws.cell(r, 2, max27)
    else:
        c = ws.cell(r, 2, max27)
    c.number_format = "€#,##0"
    ws.cell(r, 3, f"='GPS 2025'!I{br}").number_format = "0.0%"
    ws.cell(r, 4, f"='GPS 2025'!G{br}").number_format = "€#,##0"
    if db == "ITUK":
        c = ws.cell(r, 5, BP_ITUK[2027]["isze_eur"])
        c.font = BLUE
        c.border = BORDER
    else:
        ws.cell(r, 5, bp_targets[db]).number_format = "€#,##0"
    ws.cell(r, 5).number_format = "€#,##0"
    ws.cell(r, 6, f"=E{r}/B{r}").number_format = "0.0%"
    ws.cell(r, 7, f"=F{r}-C{r}").number_format = "0.0%"
    ws.cell(r, 8, f"=E{r}-D{r}").number_format = "€#,##0;(€#,##0);-"
    ws.cell(r, 9, f"=B{r}*0.01").number_format = "€#,##0"

tr = r0 + len(GPS_DBS)
ws.cell(tr, 1, "NETWORK TOTAL").font = BOLD
for col in (2, 4, 5, 8, 9):
    letter = get_column_letter(col)
    ws.cell(tr, col, f"=SUM({letter}{r0}:{letter}{tr-1})").font = BOLD
    ws.cell(tr, col).number_format = "€#,##0"
ws.cell(tr, 6, f"=E{tr}/B{tr}").font = BOLD
ws.cell(tr, 6).number_format = "0.0%"
ws.cell(tr, 3, f"='GPS 2025'!I{5 + len(base)}").font = BOLD
ws.cell(tr, 3).number_format = "0.0%"
ws.cell(tr, 7, f"=F{tr}-C{tr}").font = BOLD
ws.cell(tr, 7).number_format = "0.0%"

note = tr + 2
ws.cell(note, 1, "ITUK detail").font = BOLD
ws.cell(note + 1, 1, f"BP retail FY27 £{BP_ITUK[2027]['retail_gbp']:,} · ISZE plan €{BP_ITUK[2027]['isze_eur']:,.0f}")
ws.cell(note + 2, 1, f"UIO FY27 (v18): {_uio_detail_str('ITUK', 2027)}")
ws.cell(note + 3, 1, "Δ GPS pp = GPS needed for plan − CY2025 GPS ex-warranty")

for j, w in enumerate([12, 14, 10, 14, 14, 10, 10, 12, 10], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ============ Assumptions
ws = wb.create_sheet("Assumptions")
ws.sheet_view.showGridLines = False
ws["A1"] = "Assumptions & data sources (v9 — Jun 2026 refresh)"
ws["A1"].font = H1
assumptions = [
    "GPS scope = ISZE → distributor billing only.",
    "Max spend rates: £400/£1,000/£1,200 retail × (1−34% GP) × FX 1.12 → €296/€739/€887. LCV €250.",
    "Legacy network €1,120 / €1,344 retired (skipped GP haircut).",
    f"UIO workbook: {pkg['source_uio']}.",
    f"ITUK CY2025 car park: {pkg['source_ituk_carpark']} (DVLA, 22,389 CV).",
    "ITUK 2026–27 UIO: v18 sheet 4 model projection.",
    "ERP sales CY2025: sales_fact / PRISM net billing.",
    "Warranty: Data/Warranty_CY2025_by_DB.md — claim price basis.",
    "BP FY2027 ITUK: ITUK BUSINESSPLAN + Target Study ISZE plan row.",
    f"BP FY2027 network proxy: Kosha VP €{kosha_27/1e6:.1f}M allocated by CY2025 sales share.",
]
for i, txt in enumerate(assumptions, 3):
    ws.cell(i, 1, f"{i-2}. {txt}").font = BLACK
ws.column_dimensions["A"].width = 110

try:
    wb.save(OUT)
    saved = OUT
except PermissionError:
    alt = OUT.with_name(OUT.stem + "_new.xlsx")
    wb.save(alt)
    saved = alt
    print("WARN: could not overwrite (file open?) — wrote", alt)

# Export snapshot for dashboard build
snap = {
    "base25": [[db, _uio_total(db, 2025), DIST[db]["max_spend"][2025], SALES_25[db]] for db in GPS_DBS],
    "warranty": WARRANTY,
    "uio26": {db: DIST[db]["uio"][2026] for db in GPS_DBS},
    "uio27_max": {db: DIST[db]["max_spend"][2027] for db in GPS_DBS},
    "bp27": bp_targets,
    "rates": [RATE_N35, RATE_N75, RATE_FMD],
    "bp_ituk": BP_ITUK,
}
(ROOT / "Data" / "_build" / "gps_v9_snapshot.json").write_text(
    json.dumps(snap, indent=2), encoding="utf-8"
)

ituk_max = DIST["ITUK"]["max_spend"][2025]
ituk_gps = SALES_25["ITUK"] / ituk_max
print("saved", saved)
print(f"ITUK GPS incl: {ituk_gps:.1%} | max spend €{ituk_max:,.0f} | car park UIO {_uio_total('ITUK', 2025):,}")
