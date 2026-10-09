# -*- coding: utf-8 -*-
"""Don't Lose the Deal - other distributors study. Numbers builder (2026-10-08).

Reads only derived data plus the ITUK deal register workbook. Writes payload.json
next to this file. Every figure that reaches the study comes from here.
Windows: FY2026 = 1 Apr 2025 - 31 Mar 2026. GBP->EUR 1.17 (conversion only).
"""
from __future__ import annotations

import glob
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
D = ROOT / "01_data" / "00_derived"
OUT = Path(__file__).parent
FX = 1.17
FY0, FY1 = "2025-04-01", "2026-03-31"

BOOK = ["ITUK", "IUK", "UTI", "UMI", "GERMANY", "BELGIUM", "PORTUGAL", "NORWAY", "SWEDEN",
        "IRELAND", "GREECE", "CYPRUS", "HUNGARY", "FINLAND", "MALTA", "ROMANIA", "DENMARK", "ICELAND"]
WARR_MAP = {"ITUK": "ITUK", "IUK": "IUK", "UTI": "UTI", "UMI": "UMI", "ISD": "GERMANY", "RSA": "NORWAY",
            "Sverige": "SWEDEN", "IBX": "BELGIUM", "Imotors": "PORTUGAL", "Ireland": "IRELAND",
            "Greece": "GREECE", "AutoWallis": "HUNGARY", "ROM Auto": "ROMANIA", "Loutsios": "CYPRUS",
            "GasanZammit": "MALTA", "Nordic FI": "FINLAND", "Danmark": "DENMARK", "BL Iceland": "ICELAND"}
IDS_TAB = {"BELGIUM": "Belgium", "GREECE": "Greece", "HUNGARY": "Hungary", "NORWAY": "Norway",
           "PORTUGAL": "Portugal", "ROMANIA": "Romania", "DENMARK": "Denmark", "UMI": "Israel LCV",
           "IUK": "UK LCV", "ITUK": "UK CV", "CYPRUS": "Cyprus", "MALTA": "Malta"}
STD_FILES = {"ITUK": "named dealers CY2025 (64) + network", "IUK": "network by part, CY2025 + 2026 YTD",
             "UTI": "network, Aug 2025 - Mar 2026", "GREECE": "network, Apr - Dec 2025",
             "IRELAND": "network, CY2025 + 2026 YTD", "PORTUGAL": "network, CY2025 + 2026 YTD"}


def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


# ---------------------------------------------------------------- 1. register
def load_register():
    p = glob.glob(str(ROOT / "01_data/90_INBOX/ituk_std_mail_hunt_20260903/pulled_attachments/*DLAD*2026*"))[0]
    wb = openpyxl.load_workbook(p, data_only=True)
    rows = []
    for name in ["North", "South West", "South East"]:
        for r in wb[name].iter_rows(min_row=3, values_only=True):
            if num(r[8]) is None or num(r[10]) is None:
                continue
            rows.append(dict(region=name, date=r[0], dealer=str(r[2]).strip(), cust=str(r[3]).strip() if r[3] else None,
                             reg=str(r[4]).strip().upper().replace(" ", "") if r[4] else None,
                             pn=str(r[5]).strip().split(".")[0] if r[5] is not None else None,
                             desc=str(r[6]).strip(), dn=num(r[8]), pur=num(r[9]), off=num(r[10]),
                             qty=num(r[16]) or 1, sup=num(r[18]) or 0, comp=r[13], cmt=r[15]))
    reg = pd.DataFrame(rows)
    reg["date"] = pd.to_datetime(reg["date"], errors="coerce")
    # SW has month names only (no year). File is the 2026 workbook, 3 months: JAN/FEB/MARCH 2026 (assumption, stated)
    mm = {"JAN": 1, "FEB": 2, "MARCH": 3, "MAR": 3}
    for i, r in reg.iterrows():
        if pd.isna(r["date"]) and r["region"] == "South West":
            m = mm.get(str(r["dealer"]).upper(), None)
        # month text lives in column A; re-read below
    return reg, wb


def sw_months(wb):
    out = []
    for r in wb["South West"].iter_rows(min_row=3, values_only=True):
        if num(r[8]) is None or num(r[10]) is None:
            continue
        out.append(str(r[0]).strip().upper())
    return out


def fleet_block(wb):
    ws = wb["Fleet Support"]
    rows = [r[:16] for r in ws.iter_rows(min_row=7, values_only=True)]
    df = pd.DataFrame(rows, columns=["cust", "pn", "desc", "qty", "nf_terms", "nf_revised", "support", "qty_total",
                                     "support_total", "month", "dealer", "revised_dp", "dealer_margin", "ituk_landed",
                                     "ituk_margin", "order_ref"])
    df = df[df["cust"].notna()]
    df["support_total"] = pd.to_numeric(df["support_total"], errors="coerce")
    return df.groupby(df["cust"].astype(str).str.strip()).support_total.sum().sort_values(ascending=False)


# ---------------------------------------------------------------- 2. billing pools
def billing_pools():
    dim = pd.read_csv(D / "distributor_dimension.csv")
    dim = dim[dim.network_role == "network_db"][["sold_to_sap", "distributor_code", "scope"]]
    s = pd.read_parquet(D / "sales_fact.parquet").merge(dim, on="sold_to_sap", how="inner")
    fy = s[(s.billing_date >= FY0) & (s.billing_date <= FY1) & (~s.saty.isin(["CR", "RE", "FD"]))].copy()
    pc = pd.read_parquet(D / "parts_catalog_dimension.parquet", columns=["part_number", "pms_registry", "in_bvp_catalog"]).drop_duplicates("part_number")
    cl = pd.read_csv(D / "aim_pn_classification_20260920.csv", usecols=["part_number", "catalog_l1", "aim_l1_fixed"], dtype=str).drop_duplicates("part_number")
    fy = fy.merge(pc, on="part_number", how="left").merge(cl, on="part_number", how="left")
    fy["fam"] = fy.aim_l1_fixed.fillna(fy.catalog_l1).fillna("Unclassified")
    g = fy.groupby(["distributor_code", "part_number"]).agg(val=("net_sales_value", "sum"), qty=("billed_qty", "sum")).reset_index()
    g["asp"] = np.where(g.qty > 0, g.val / g.qty, np.nan)
    fy = fy.merge(g[["distributor_code", "part_number", "asp"]], on=["distributor_code", "part_number"], how="left")
    fy["pms"] = fy.pms_registry.astype("boolean").fillna(False).astype(bool)
    fy["bvp"] = fy.in_bvp_catalog.astype("boolean").fillna(False).astype(bool)
    fy["hard"] = (~fy.pms) & (~fy.bvp)
    fy["profile"] = fy.hard & (fy.asp >= 100)
    fy["profile250"] = fy.hard & (fy.asp >= 250)
    fy["profile500"] = fy.hard & (fy.asp >= 500)
    rows = {}
    for d, x in fy.groupby("distributor_code"):
        pr = x[x.profile]
        rows[d] = dict(
            scope=x.scope.iloc[0],
            billed=float(x.net_sales_value.sum()),
            pms=float(x.net_sales_value[x.pms].sum()),
            bvp=float(x.net_sales_value[x.bvp].sum()),
            hard=float(x.net_sales_value[x.hard].sum()),
            profile=float(pr.net_sales_value.sum()),
            profile250=float(x.net_sales_value[x.profile250].sum()),
            profile500=float(x.net_sales_value[x.profile500].sum()),
            profile_units=float(pr.billed_qty.sum()),
            profile_avg_price=float(pr.net_sales_value.sum() / pr.billed_qty.sum()) if pr.billed_qty.sum() > 0 else None,
            profile_pns=int(pr.part_number.nunique()),
        )
    fams = fy[fy.profile].groupby(["distributor_code", "fam"]).net_sales_value.sum().reset_index()
    return rows, fams, fy


def warranty_fy26():
    w = pd.read_csv(D / "distributor_warranty_expense_fiscal.csv")
    return {WARR_MAP[r["db"]]: float(r["warr_FY2026"]) for _, r in w.iterrows() if r["db"] in WARR_MAP}



def warranty_profile_share():
    """Share of warranty claim euros that sit on hard parts (not service, not value line) priced >= EUR 100 a unit.
    Measured on warranty_raw_grain.csv, FY2025-FY2026 claims, unit price = claim euro / quantity."""
    w = pd.read_csv(D / "warranty_raw_grain.csv", low_memory=False)
    w["am"] = w.assess_month.astype(str)
    w["fy"] = np.where(w.am.str[4:6].astype(int) >= 4, w.am.str[:4].astype(int) + 1, w.am.str[:4].astype(int))
    w = w[w.fy.isin([2025, 2026])].copy()
    w["pn"] = w.part_no.astype(str).str.replace(r"\.0$", "", regex=True)
    pc = pd.read_parquet(D / "parts_catalog_dimension.parquet", columns=["part_number", "pms_registry", "in_bvp_catalog"]).drop_duplicates("part_number")
    w = w.merge(pc, left_on="pn", right_on="part_number", how="left")
    w["hard"] = ~(w.pms_registry.astype("boolean").fillna(False).astype(bool)) & ~(w.in_bvp_catalog.astype("boolean").fillna(False).astype(bool))
    w["unit"] = np.where(w.qty > 0, w.claim_eur / w.qty, np.nan)
    w["prof"] = w.hard & (w.unit >= 100)
    m = {"ISUZU (UK)": "IUK", "UMI": "UMI", "RSA": "NORWAY", "Isuzu Sverige AB": "SWEDEN", "ISUZU SALES DEUTSCHLAND - ISD": "GERMANY",
         "ISD": "GERMANY", "ISUZU BENELUX - BELGIUM": "BELGIUM", "IMOTORS LDA": "PORTUGAL", "PORTUGAL": "PORTUGAL",
         "Isuzu Ireland 2020 Unlimited Company": "IRELAND", "IRELAND": "IRELAND", "PETROS PETROPOULOS": "GREECE",
         "WAE LTD.": "HUNGARY", "ISUZU ROM AUTO SRL": "ROMANIA", "A. LOUTSIOS & SONS": "CYPRUS", "CYPRUS": "CYPRUS",
         "GASANZAMMIT MOTORS LIMITED": "MALTA", "MALTA": "MALTA", "Isuzu Motor Finland OY": "FINLAND", "ISUZU Danmark": "DENMARK",
         "BL ehf.": "ICELAND", "ITUK": "ITUK", "UTI": "UTI"}
    w["d"] = w.db_name.map(m)
    a = w.dropna(subset=["d"]).groupby("d").apply(lambda x: x.claim_eur[x.prof].sum() / x.claim_eur.sum(), include_groups=False)
    return a.to_dict()


def rebates():
    """Year-end 'Rebate Material' credit documents (SaTy CR) per distributor and fiscal year, from sales_fact.
    Sold-to 2194 is ITUK's payer account and is added to ITUK."""
    s = pd.read_parquet(D / "sales_fact.parquet")
    dim = pd.read_csv(D / "distributor_dimension.csv")
    m = dim.set_index("sold_to_sap").distributor_code.to_dict()
    m[2194] = "ITUK"
    c = s[s.part_number.astype(str).str.upper().str.contains("REBATE")].copy()
    c["fy"] = c.billing_date.apply(lambda d: d.year + 1 if d.month >= 4 else d.year)
    c["d"] = c.sold_to_sap.map(m)
    t = c.dropna(subset=["d"]).pivot_table(index="d", columns="fy", values="net_sales_value", aggfunc="sum").fillna(0.0)
    docs = c[(c.fy == 2026)].dropna(subset=["d"]).groupby("d").apply(lambda x: [dict(date=str(r.billing_date.date()), eur=float(r.net_sales_value)) for r in x.itertuples()], include_groups=False)
    return t.to_dict(orient="index"), docs.to_dict()


def contest_evidence(fy):
    """Competitor price evidence that touches hard parts (not service, not value line, >= EUR 100 a unit).
    delta = (independent price - ISZE price) / ISZE price; negative = independent is cheaper.
    Weights = FY2026 hard-part billing on that part number in the stated markets."""
    prof = fy[fy.profile]
    def rev(ds):
        x = prof if ds is None else prof[prof.distributor_code.isin(ds)]
        return x.groupby("part_number").net_sales_value.sum()
    def st(name, df, dcol, w, markets):
        df = df.copy()
        df["pn"] = df.pn.astype(str).str.replace(r"\.0$", "", regex=True)
        df = df.drop_duplicates("pn")
        df["w"] = df.pn.map(w).fillna(0)
        d = df[df.w > 0]
        x, tw = d[dcol], d.w.sum()
        return dict(source=name, markets=markets, pns_priced=int(len(df)), pns_hard_in_pool=int(len(d)), covered_eur=float(tw),
                    pool_eur=float(w.sum()), coverage=float(tw / w.sum()),
                    share_below10=float(d.w[x <= -.10].sum() / tw), share_below25=float(d.w[x <= -.25].sum() / tw),
                    share_above=float(d.w[x > 0].sum() / tw), wtd_delta=float((x * d.w).sum() / tw))
    e = pd.read_csv(D / "eucon_price_gap.csv")
    u = pd.read_csv(D / "eucon_uk_survey_2023_price_gap.csv")
    k = pd.read_csv(D / "iam_ks_price_gap.csv", low_memory=False)
    return [
        st("Online independent listings, Germany and Norway, Mar 2026", e, "Avg_Delta_pct", rev(["GERMANY", "NORWAY"]), "Germany + Norway"),
        st("Online independent listings, Germany and Norway, Mar 2026", e, "Avg_Delta_pct", rev(None), "all 18 distributors"),
        st("UK independent workshop prices, 2023 survey", u, "delta_pct", rev(["ITUK"]), "ITUK"),
        st("UK independent workshop prices, 2023 survey", u, "delta_pct", rev(["IUK"]), "IUK"),
        st("UK trade-shop listings, Aug 2026", k, "delta_pct", rev(["ITUK"]), "ITUK"),
        st("UK trade-shop listings, Aug 2026", k, "delta_pct", rev(["IUK"]), "IUK"),
    ]


def rebate_labels():
    """FY2026 rebate credit documents with the label the order file carries (Purchase Order Number), matched on sold-to, date, euro."""
    s = pd.read_parquet(D / "sales_fact.parquet")
    dim = pd.read_csv(D / "distributor_dimension.csv")
    m = dim.set_index("sold_to_sap").distributor_code.to_dict()
    m[2194] = "ITUK"
    c = s[s.part_number.astype(str).str.upper().str.contains("REBATE")].copy()
    c = c[(c.billing_date >= FY0) & (c.billing_date <= FY1)]
    c["d"] = c.billing_date.dt.date.astype(str)
    c["eur_r"] = c.net_sales_value.round(0)
    o = pd.read_csv(D / "order_headers_from_Order_Detail.csv", low_memory=False)
    o = o[o.SaTy == "CR"].copy()
    o["d"] = pd.to_datetime(o.created).dt.date.astype(str)
    o["eur_r"] = o.eur.round(0)
    j = c.merge(o[["Sold-to pt", "d", "eur_r", "Purchase Order Number"]], left_on=["sold_to_sap", "d", "eur_r"],
                right_on=["Sold-to pt", "d", "eur_r"], how="left")
    def klass(x):
        t = str(x).lower() if isinstance(x, str) else ""
        if "kickback" in t: return "kickback"
        if "dlod" in t: return "DLOD"
        if "scrap" in t: return "scrap support"
        if "sales support" in t: return "sales support"
        if "correction" in t: return "correction"
        return "no label in order file"
    j["klass"] = j["Purchase Order Number"].map(klass)
    j["dist"] = j.sold_to_sap.map(m)
    out = {}
    for d, g in j.groupby("dist"):
        out[d] = [dict(date=str(r.d), eur=float(r.net_sales_value), label=(r._asdict().get("_9") if False else None), klass=r.klass) for r in g.itertuples()]
        for rec, (_, r) in zip(out[d], g.iterrows()):
            rec["label"] = r["Purchase Order Number"] if isinstance(r["Purchase Order Number"], str) else None
    tot = j.groupby("klass").net_sales_value.sum().to_dict()
    return out, {k: float(v) for k, v in tot.items()}

def margin_by_dist():
    o = pd.read_csv(D / "order_headers_from_Order_Detail.csv", low_memory=False)
    dim = pd.read_csv(D / "distributor_dimension.csv")
    dim = dim[dim.network_role == "network_db"][["sold_to_sap", "distributor_code"]]
    o = o.merge(dim, left_on="Sold-to pt", right_on="sold_to_sap", how="left")
    f = o[(o.created >= FY0) & (o.created <= FY1) & (~o.SaTy.isin(["CR", "G2", "RE", "FD"])) & o.distributor_code.notna() & (o.cost > 0)]
    a = f.groupby("distributor_code").agg(eur=("eur", "sum"), gm=("gm_eur", "sum"))
    return (a.gm / a.eur).to_dict(), float(f.gm_eur.sum() / f.eur.sum())


def main():
    reg, wb = load_register()
    swm = sw_months(wb)
    # SW months -> dates in 2026 (assumption)
    sw_idx = reg.index[reg.region == "South West"]
    mm = {"JAN": "2026-01-15", "FEB": "2026-02-15", "MARCH": "2026-03-15", "MAR": "2026-03-15"}
    for i, m in zip(sw_idx, swm):
        reg.loc[i, "date"] = pd.Timestamp(mm.get(m, "2026-02-15"))
    reg["sup_eur"] = reg.sup * FX
    reg["dn_eur"] = reg.dn * FX
    reg["pur_eur"] = reg.pur * FX
    fleet = fleet_block(wb)

    # classification + ISZE price validation
    pc = pd.read_parquet(D / "parts_catalog_dimension.parquet", columns=["part_number", "pms_registry", "in_bvp_catalog"]).drop_duplicates("part_number")
    cl = pd.read_csv(D / "aim_pn_classification_20260920.csv", usecols=["part_number", "catalog_l1", "aim_l1_fixed"], dtype=str).drop_duplicates("part_number")
    reg = reg.merge(pc, left_on="pn", right_on="part_number", how="left").merge(cl, on="part_number", how="left", suffixes=("", "_c"))
    reg["fam"] = reg.aim_l1_fixed.fillna(reg.catalog_l1).fillna("Unclassified")
    s = pd.read_parquet(D / "sales_fact.parquet")
    it = s[(s.sold_to_sap == 2575) & (s.billing_date >= "2025-01-01") & (~s.saty.isin(["CR", "RE", "FD"])) & (s.billed_qty > 0)]
    asp = (it.groupby("part_number").net_sales_value.sum() / it.groupby("part_number").billed_qty.sum()).rename("isze_asp")
    reg = reg.merge(asp, left_on="pn", right_index=True, how="left")
    matched = reg.isze_asp.notna()
    ratio = (reg.isze_asp / reg.pur_eur)[matched]
    reg["isze_unit"] = reg.isze_asp.fillna(reg.pur_eur * float(ratio.median()))
    reg["isze_bill"] = reg.isze_unit * reg.qty

    # register facts
    tot_sup = reg.sup_eur.sum()
    regn = {
        "lines": int(len(reg)),
        "support_gbp": float(reg.sup.sum()),
        "dn_gbp": float((reg.dn * reg.qty).sum()),
        "pur_gbp": float((reg.pur * reg.qty).sum()),
        "offer_gbp": float((reg.off * reg.qty).sum()),
        "matched_lines": int(matched.sum()),
        "pur_vs_isze_ratio_median": float(ratio.median()),
        "pur_vs_isze_ratio_p25": float(ratio.quantile(.25)),
        "pur_vs_isze_ratio_p75": float(ratio.quantile(.75)),
        "non_pms_lines": int((reg.pms_registry.astype("boolean").fillna(False) == False).sum()),
        "bvp_lines": int(reg.in_bvp_catalog.astype("boolean").fillna(False).sum()),
        "pms_lines": int(reg.pms_registry.astype("boolean").fillna(False).sum()),
        "comment_blank": int(reg.cmt.isna().sum()),
        "comp_price_filled": int(reg.comp.notna().sum()),
        "warranty_comment": int(reg.cmt.astype(str).str.lower().str.contains("warrant").sum()),
        "secondhand_comment": int(reg.cmt.astype(str).str.lower().str.contains("second").sum()),
        "internet_comment": int(reg.cmt.astype(str).str.lower().str.contains("internet").sum()),
    }
    regn["support_over_dn"] = regn["support_gbp"] / regn["dn_gbp"]
    regn["support_over_isze_bill"] = float(reg.sup_eur.sum() / reg.isze_bill.sum())
    # Euro-weighted pool margins (total cost ÷ total price) — useful for funding arithmetic.
    regn["normal_margin_on_dn"] = 1 - regn["pur_gbp"] / regn["dn_gbp"]
    regn["retained_margin_on_offer"] = 1 - regn["pur_gbp"] / regn["offer_gbp"]
    # Line-level margins (Duncan Retained Margin % style) — board narrative uses these.
    line_normal = 1 - reg.pur / reg.dn
    line_retained = 1 - reg.pur / reg.off
    regn["median_normal_margin_on_dn"] = float(line_normal.median())
    regn["median_retained_margin_on_offer"] = float(line_retained.median())
    regn["avg_retained_margin_on_offer"] = float(line_retained.mean())
    regn["lines_below_cost"] = int((reg.off < reg.pur).sum())
    regn["dn_price_over_cost"] = regn["dn_gbp"] / regn["pur_gbp"]
    regn["isze_bill_eur"] = float(reg.isze_bill.sum())
    srt = reg.sort_values("sup_eur", ascending=False)
    n10 = int(round(len(reg) * .10))
    regn["top10pct_lines_support_share"] = float(srt.sup_eur.head(n10).sum() / tot_sup)
    regn["lines_ge500_support_gbp"] = float(reg.sup[reg.sup >= 500].sum())
    regn["lines_ge500_n"] = int((reg.sup >= 500).sum())
    regn["lines_lt150_n"] = int((reg.sup < 150).sum())
    regn["lines_lt150_support_gbp"] = float(reg.sup[reg.sup < 150].sum())
    regn["median_support_per_line_gbp"] = float(reg.sup.median())
    dsup = reg.groupby("dealer").sup.sum().sort_values(ascending=False)
    regn["dealers"] = int(len(dsup))
    regn["top3_dealer_share"] = float(dsup.head(3).sum() / dsup.sum())
    regn["top5_dealer_share"] = float(dsup.head(5).sum() / dsup.sum())
    # repeat
    j = reg[reg.reg.notna()]
    vv = j.groupby("reg").agg(dates=("date", "nunique"), dealers=("dealer", "nunique"), lines=("pn", "size"))
    regn["vehicles"] = int(len(vv))
    regn["vehicles_2plus_dates"] = int((vv.dates >= 2).sum())
    regn["vehicles_multi_line"] = int((vv.lines > 1).sum())
    regn["customers"] = int(reg.cust.nunique())
    # family mix of support
    fam = reg.groupby("fam").agg(lines=("sup", "size"), support_eur=("sup_eur", "sum"), isze_bill_eur=("isze_bill", "sum")).sort_values("support_eur", ascending=False)
    fam["share"] = fam.support_eur / fam.support_eur.sum()
    regn["families"] = [dict(fam=k, lines=int(v.lines), support_eur=float(v.support_eur), share=float(v.share)) for k, v in fam.head(10).iterrows()]
    # ISZE price thresholds
    for T in (100, 150, 250):
        regn[f"support_share_isze_price_ge_{T}"] = float(reg.sup_eur[reg.isze_unit >= T].sum() / tot_sup)
    # windows
    def win(a, b, regions=None):
        m = (reg.date >= a) & (reg.date <= b)
        if regions:
            m &= reg.region.isin(regions)
        x = reg[m]
        return dict(lines=int(len(x)), support_eur=float(x.sup_eur.sum()), isze_bill_eur=float(x.isze_bill.sum()), vehicles=int(x.reg.nunique()))
    windows = {
        "north_cy2025": win("2025-01-01", "2025-12-31", ["North"]),
        "fy26_visible": win(FY0, FY1),
        "q1_2026_sw_se": win("2026-01-01", "2026-03-31", ["South West", "South East"]),
        "north_full": win("2024-11-01", "2025-12-31", ["North"]),
    }
    fleet_total_gbp = float(fleet.sum())
    fleet_top = fleet.head(3).to_dict()

    # distributors
    pools, fams, fy = billing_pools()
    warr = warranty_fy26()
    fprof = warranty_profile_share()
    reb, reb_docs = rebates()
    reb_lab, reb_tot = rebate_labels()
    contest = contest_evidence(fy)
    gm, gm_net = margin_by_dist()
    spu = pd.read_csv(D / "gps_spu_by_distributor.csv")
    spu25 = spu[spu.year == 2025].set_index("distributor_code")
    risk = pd.read_csv(D / "db_aftersales_risk.csv").set_index("distributor_code")
    mix = pd.read_csv(D / "db_parts_mix_classification.csv")
    mix25 = mix[mix.year == 2025].set_index("distributor_code")
    ids = pd.read_csv(D / "ids_dealer_network_cleaned_2026-09-20.csv").groupby("ids_tab").size().to_dict()
    ow = pd.read_csv(D / "wallet_waterfall.csv")
    ow25 = ow[ow.year == 2025].set_index("distributor_code")
    dists = []
    for d in BOOK:
        p = pools[d]
        w = warr.get(d, 0.0)
        f = fprof.get(d, 1.0)
        profile_low = max(0.0, p["profile"] - w * f)
        row = dict(code=d, scope=p["scope"], billed=p["billed"], warranty=w, commercial=p["billed"] - w,
                   warranty_share=w / p["billed"], warranty_in_pool_share=f, warranty_in_pool=w * f, hard=p["hard"], pms=p["pms"], bvp=p["bvp"],
                   profile=p["profile"], profile_low=profile_low, profile250=p["profile250"], profile500=p["profile500"],
                   profile_share_billed=p["profile"] / p["billed"], profile_avg_price=p["profile_avg_price"],
                   profile_pns=p["profile_pns"], gm=gm.get(d, gm_net),
                   uio=float(spu25.loc[d, "uio_units"]) if d in spu25.index else None,
                   gps_capture_ex_w=float(spu25.loc[d, "gps_capture_ex_warranty_vs_max"]) if d in spu25.index else None,
                   spend_ex_w=float(spu25.loc[d, "spend_per_unit_ex_warranty_eur"]) if d in spu25.index else None,
                   dealers=(int(ids[IDS_TAB[d]]) if (d in IDS_TAB and IDS_TAB[d] in ids and ids[IDS_TAB[d]] >= 3) else None),
                   std_file=STD_FILES.get(d),
                   contested_eur=float(mix25.loc[d, "contested_eur"]) if d in mix25.index else None,
                   contested_bvp_capture=float(mix25.loc[d, "contested_bvp_capture"]) if d in mix25.index else None,
                   eucon_undercut_pct=float(risk.loc[d, "eucon_undercut_pct"]) if d in risk.index else None,
                   n_eucon=float(risk.loc[d, "n_eucon"]) if d in risk.index else None,
                   risk_score=float(risk.loc[d, "risk_score"]) if d in risk.index else None,
                   vor_urgent_share=float(risk.loc[d, "vor_urgent_share"]) if d in risk.index else None,
                   wallet_gap=float(ow25.loc[d, "addressable_uncaptured_eur"]) if d in ow25.index else None,
                   rebate_fy26=float(reb.get(d, {}).get(2026, 0.0)), rebate_fy25=float(reb.get(d, {}).get(2025, 0.0)), rebate_docs_fy26=reb_docs.get(d, []), rebate_labelled=reb_lab.get(d, []))
        row["commercial_per_vehicle"] = row["commercial"] / row["uio"] if row["uio"] else None
        row["margin_per_vehicle"] = row["commercial_per_vehicle"] * row["gm"] if row["commercial_per_vehicle"] else None
        dists.append(row)

    # ITUK usage intensity (touch rate) from the register
    it_d = next(x for x in dists if x["code"] == "ITUK")
    P = it_d["profile_low"]
    u = {
        "pool_low": P,
        "u_north_cy2025": windows["north_cy2025"]["isze_bill_eur"] / P,
        "u_fy26_visible": windows["fy26_visible"]["isze_bill_eur"] / P,
        "u_q1_annualised": 4 * windows["q1_2026_sw_se"]["isze_bill_eur"] / P,
    }
    rate = {
        "cases_per_1000_uio_fy26_visible": windows["fy26_visible"]["lines"] / it_d["uio"] * 1000,
        "vehicles_per_1000_uio_fy26_visible": windows["fy26_visible"]["vehicles"] / it_d["uio"] * 1000,
        "avg_isze_bill_per_line": regn["isze_bill_eur"] / regn["lines"],
    }

    payload = dict(as_of="2026-10-08", fx=FX, register=regn, windows=windows, fleet_total_gbp=fleet_total_gbp, fleet_top=fleet_top,
                   dists=dists, intensity=u, rate=rate, network_gm=gm_net, contest=contest, rebate_class_totals=reb_tot)
    (OUT / "payload.json").write_text(json.dumps(payload, indent=1, default=float), encoding="utf-8")
    fams.to_csv(OUT / "profile_pool_by_family.csv", index=False)
    pd.DataFrame(dists).to_csv(OUT / "distributor_table.csv", index=False)
    reg.drop(columns=["part_number", "part_number_c"], errors="ignore").to_csv(OUT / "register_lines_classified.csv", index=False)
    print("ok", OUT / "payload.json")


if __name__ == "__main__":
    main()
