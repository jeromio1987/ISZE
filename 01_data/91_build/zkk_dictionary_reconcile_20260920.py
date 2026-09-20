# -*- coding: utf-8 -*-
"""
ZKK dictionary -- two reconciliations before the document is written.

1. The earlier audit reports ZKK Japanese-description purity 12.32%; my pass gives 12.33%.
   Test whether the gap is the split-character set (the benchmark script splits on ; , ( ).
2. The settled position records 27 three-digit prefixes in the 2022 HKK workbook; my
   extract found 26. List them per sheet.

Read-only. Nothing written to disk.

Run: python 01_data/91_build/zkk_dictionary_reconcile_20260920.py
"""
from __future__ import annotations

import collections
import csv
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "01_data" / "00_derived" / "jp_codes_cache_20260920.csv"
SPINE = ROOT / "01_data" / "00_derived" / "aim_master_v2.csv"
HKK2022 = ROOT / "01_data" / "HKK Code 20221028.xlsx"


def s(v) -> str:
    if v is None:
        return ""
    t = str(v).strip()
    if t.endswith(".0") and t[:-2].replace("-", "").isdigit():
        t = t[:-2]
    return t


def nfkc(v) -> str:
    return unicodedata.normalize("NFKC", str(v or "")).strip()


def make_noun(pattern: str):
    def f(desc: str) -> str:
        d = re.sub(r"\s+", " ", (desc or "").upper()).strip()
        h = re.sub(r"\s+", " ", re.split(pattern, d)[0]).strip(" .*/,")
        return h or "(blank)"
    return f


def main() -> int:
    P = print
    cache = {}
    with CACHE.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            pn = s(r.get("pn"))
            if pn:
                cache[pn] = r
    base = {}
    with SPINE.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            pn = s(r.get("part_number"))
            if not pn:
                continue
            c = cache.get(pn)
            if not c:
                continue
            z = s(c.get("zkk"))
            if not z:
                continue
            base[pn] = {"desc": s(c.get("jp_description")), "zkk": z,
                        "hkk": s(c.get("hkk")), "zuk": s(c.get("zuk"))}

    P("=" * 96)
    P("N1. PURITY UNDER DIFFERENT LEADING-NOUN DEFINITIONS (same 139,524-row base)")
    P("=" * 96)
    defs = {
        "split on ; , ( )   [benchmark script]": r"[;,(]",
        "split on ; ,       [this document]": r"[;,](?!$)",
        "split on ; only": r"[;]",
        "split on ; , ( ) / [widest]": r"[;,(/]",
    }
    P(f"  {'definition':<40} {'ZKK':>8} {'HKK':>8} {'ZUK':>8}")
    for label, pat in defs.items():
        nf = make_noun(pat)
        line = f"  {label:<40}"
        for key in ("zkk", "hkk", "zuk"):
            mp = collections.defaultdict(collections.Counter)
            for v in base.values():
                mp[v[key]][nf(v["desc"])] += 1
            tot = sum(sum(c.values()) for c in mp.values())
            dom = sum(c.most_common(1)[0][1] for c in mp.values())
            line += f" {100*dom/tot:>7.2f}%"
        P(line)

    P("")
    P("=" * 96)
    P("N2. THE 2022 HKK WORKBOOK -- prefixes per sheet")
    P("=" * 96)
    wb = openpyxl.load_workbook(HKK2022, read_only=True, data_only=True)
    allpfx = collections.defaultdict(set)
    for sh in wb.sheetnames:
        ws = wb[sh]
        for row in ws.iter_rows(values_only=True):
            cells = [nfkc(c) for c in row]
            for i in range(len(cells) - 2):
                if re.fullmatch(r"\d{3}", cells[i]) and re.fullmatch(r"[A-Za-z0-9]", cells[i + 1]):
                    name = cells[i + 2]
                    if len(name) > 1 and not name.startswith("("):
                        allpfx[sh].add(cells[i])
    wb.close()
    tot = set()
    for sh, pf in allpfx.items():
        P(f"  sheet {sh:<12} {len(pf):>3} prefixes: {sorted(pf)}")
        tot |= pf
    P(f"  union across sheets: {len(tot)} -> {sorted(tot)}")

    zkk_on_file = sorted({s(r.get("zkk")) for r in cache.values() if s(r.get("zkk"))})
    P(f"  ZKK values that are a 2022 prefix: {len(set(zkk_on_file) & tot)} -> "
      f"{sorted(set(zkk_on_file) & tot)}")
    P("")
    P("  the union above includes headers of the form '3-digit + slot' only; a prefix that")
    P("  appears solely inside a bracketed column heading is not counted. Settled position")
    P("  records 27 (Parts 13, Accessories 3, Lubricant 11); this extract gives the counts above.")
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    raise SystemExit(main())
