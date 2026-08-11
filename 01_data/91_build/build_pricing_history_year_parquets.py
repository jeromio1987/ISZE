"""Build year-stamped PR00 / ZPR1 history parquets from 01_Pricing/history/.

Matches existing 2023 schema:
  pricing_pr00_ie40_YYYY.parquet  → Material, Price_EUR, list_year
  pricing_zpr1_YYYY.parquet       → Customer, Material, Price, Currency,
                                    Valid_From, Valid_To, sheet, list_year

Current (2026) engine tables stay:
  pricing_pr00_ie40.parquet / pricing_zpr1.parquet
  (built by build_pricing_enquiry_tables.py — do not overwrite here)
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

ISZE = Path(__file__).resolve().parents[2]
HIST = ISZE / "01_data" / "01_Pricing" / "history"
DERIVED = ISZE / "01_data" / "00_derived"

PR00_SOURCES = {
    2023: HIST / "IE40_PR00_2023-12-31.xlsx",
    2024: HIST / "IE40_PR00_2024-12-31.xlsx",
    2025: HIST / "IE40_PR00_2025-12-31.xlsx",
}
ZPR1_SOURCES = {
    2023: HIST / "ZPR1_2023-01-03.xlsx",
    2024: HIST / "ZPR1_2024-01-03.xlsx",
    2025: HIST / "ZPR1_2025-02-01.xlsx",
}


def _parse_price(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return pd.NA
    if isinstance(v, str):
        s = v.strip().replace(",", ".")
        if not s or s.upper() in {"#VALUE!", "#REF!", "#N/A", "NAN"}:
            return pd.NA
        v = s
    return pd.to_numeric(v, errors="coerce")


def build_pr00(year: int) -> tuple[Path, int]:
    path = PR00_SOURCES[year]
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Prefer sheet named as year, else first sheet
    name = str(year) if str(year) in wb.sheetnames else wb.sheetnames[0]
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    header = [str(x).strip() if x is not None else "" for x in next(it)]
    # Material / Price_EUR (history convention) or Material / Price
    hmap = {h.lower(): i for i, h in enumerate(header)}
    i_mat = hmap.get("material", 0)
    i_price = hmap.get("price_eur", hmap.get("price", 1))
    mats, prices = [], []
    for r in it:
        if r is None or i_mat >= len(r):
            continue
        mat = r[i_mat]
        if mat is None or (isinstance(mat, str) and not mat.strip()):
            continue
        mats.append(mat)
        prices.append(_parse_price(r[i_price] if i_price < len(r) else None))
    wb.close()
    df = pd.DataFrame({"Material": mats, "Price_EUR": prices, "list_year": year})
    out = DERIVED / f"pricing_pr00_ie40_{year}.parquet"
    df.to_parquet(out, index=False)
    print(f"{out.name}: {len(df):,} ← {path.name} sheet={name}")
    return out, len(df)


def build_zpr1(year: int) -> tuple[Path, int]:
    path = ZPR1_SOURCES[year]
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    frames = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(x).strip() if x is not None else "" for x in next(it)]
        except StopIteration:
            continue
        hmap = {h.lower().replace(" ", "_"): i for i, h in enumerate(header)}
        # tolerate "Valid From" / "Valid_From"
        def col(*names):
            for n in names:
                if n in hmap:
                    return hmap[n]
            return None

        i_cust = col("customer")
        i_mat = col("material")
        i_price = col("price")
        i_cur = col("currency")
        i_vf = col("valid_from", "validfrom")
        i_vt = col("valid_to", "validto")
        if i_cust is None or i_mat is None or i_price is None:
            print(f"  skip sheet {sheet}: unexpected header {header}")
            continue
        rows = []
        for r in it:
            if r is None or i_mat >= len(r):
                continue
            mat = r[i_mat]
            if mat is None or (isinstance(mat, str) and not str(mat).strip()):
                continue
            rows.append(
                {
                    "Customer": r[i_cust],
                    "Material": mat,
                    "Price": _parse_price(r[i_price] if i_price < len(r) else None),
                    "Currency": (
                        None
                        if i_cur is None or i_cur >= len(r) or r[i_cur] is None
                        else str(r[i_cur]).strip()
                    ),
                    "Valid_From": None if i_vf is None or i_vf >= len(r) else r[i_vf],
                    "Valid_To": None if i_vt is None or i_vt >= len(r) else r[i_vt],
                    "sheet": sheet,
                    "list_year": year,
                }
            )
        if rows:
            frames.append(pd.DataFrame(rows))
            print(f"  zpr1 {year} sheet {sheet}: {len(rows):,}")
    wb.close()
    if not frames:
        raise RuntimeError(f"No ZPR1 rows for {year} from {path}")
    df = pd.concat(frames, ignore_index=True)
    # Align dtypes with 2023 extract
    df["Valid_From"] = pd.to_datetime(df["Valid_From"], errors="coerce")
    df["Valid_To"] = pd.to_datetime(df["Valid_To"], errors="coerce")
    out = DERIVED / f"pricing_zpr1_{year}.parquet"
    df.to_parquet(out, index=False)
    print(f"{out.name}: {len(df):,} ← {path.name}")
    return out, len(df)


def main(years: list[int] | None = None) -> None:
    years = years or [2023, 2024, 2025]
    DERIVED.mkdir(parents=True, exist_ok=True)
    meta = {
        "built_at_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "builder": "01_data/91_build/build_pricing_history_year_parquets.py",
        "note": "History extracts only; current 2026 = pricing_pr00_ie40.parquet / pricing_zpr1.parquet",
        "outputs": {},
    }
    for y in years:
        p_path, p_n = build_pr00(y)
        z_path, z_n = build_zpr1(y)
        meta["outputs"][str(y)] = {
            "pr00": {"path": str(p_path.relative_to(ISZE)).replace("\\", "/"), "rows": p_n,
                     "source": PR00_SOURCES[y].name},
            "zpr1": {"path": str(z_path.relative_to(ISZE)).replace("\\", "/"), "rows": z_n,
                     "source": ZPR1_SOURCES[y].name},
        }
    meta_path = DERIVED / "pricing_history_years.meta.json"
    meta_path.write_text(json.dumps(meta, indent=2) + "\n", encoding="utf-8")
    print("meta →", meta_path)


if __name__ == "__main__":
    import sys

    ys = [int(a) for a in sys.argv[1:]] if len(sys.argv) > 1 else [2024, 2025]
    main(ys)
