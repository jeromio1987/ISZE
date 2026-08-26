"""Corrected (v3) fill for Bart's VPACC action plan — Jerome's exact rules.

1. My added comments: Arial 9 black, non-bold, non-italic (matches sheet body font).
2. No grey fill/font on my additions.
3. Never touch a cell that already has colleague text, and never touch a row
   whose PIC contains 'BS' (case-insensitive).
4. Green = my (empty) comment cell ONLY, and ONLY where the row's JUL column
   has an 'x'/'X' cross. No green on Progress/month cells, no green on future
   months.

Geometry:
  VP:  JUL month = col I (9);  comment col = S (19)  "Comments July";  PIC = E (5)
  ACC: JUL month = col J (10); report col  = BC (55) "July Activity Report"; PIC = F (6)
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

SRC = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\20260820__Action plan FY2027 VPACC_July update.xlsx")
OUT_DIR = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\11_planning")
OUT = OUT_DIR / "Action plan FY2027 VPACC_August update.xlsx"

GREEN = PatternFill("solid", start_color="92D050", end_color="92D050")
BODY_FONT = Font(name="Arial", size=9, bold=False, italic=False, color="FF000000")

VP_JUL = 9
VP_COMMENT = 19
VP_PIC = 5
ACC_JUL = 10
ACC_REPORT = 55
ACC_PIC = 6


def is_cross(v) -> bool:
    return isinstance(v, str) and v.strip().upper() == "X"


def pic_has_bs(pic) -> bool:
    return isinstance(pic, str) and "bs" in pic.lower()


def is_empty(v) -> bool:
    return v is None or str(v).strip() == ""


# ---------- notes (full intended set; unsafe rows are auto-skipped) ----------
vp_notes = {
    6:  "Aug: ITUK DLOD draft for management review in preparation.",
    7:  "Aug: all MOU's signed (carried from July).",
    15: "Aug: tailor-made IUK brake SPA (6-month) approved.",
    19: "Aug: stock clearance campaign live (3-month validity from July).",
    21: "Aug: IUK 6-month brake SPA approved; spot campaigns running.",
    23: "Aug: next training = IPBM foundation after EDI (IDS deferred while ISZE/IMG rework it).",
    26: "Aug: EDI training #1 done 29/07; #2 on 26/08 (3 sessions Jul/Aug/Sep).",
    27: "Aug: parts competition roadmap awaiting management approval.",
    34: "GPS recalculated for all customers (GPS Extended Target FY27 V3.1).",
    35: ("Jul/Aug GPS reviews: Portugal (Imotors) capture 63.3%; ISD ~65.9% "
         "(ISZJ <=5yr fleet method); IUK 10yr car park 46,828 for global part-share recalc; "
         "Greece below network avg; Ireland benchmark = Portugal."),
    37: "ITUK: F2F meeting with IML held to introduce the promo initiative.",
    38: "ITUK target set.",
    42: "DB selection (Petros + IUK) done.",
    43: "Petros/IUK target selection in progress.",
    47: "Promotional method (which product) being decided with DBs.",
    52: "Aug: KPI's received from most DBs; several August meetings to push improvement; SharePoint auto-receiving ~50% done.",
    53: "Aug: automation ticket for KPI requests raised + discussed with IT.",
    54: "Aug: Allianz meeting end-August re Portugal development.",
    58: "Aug: Portugal agrees to visit Sept/Oct; new Imotors contact discussion Sept.",
    62: "Aug: R&M proposal to DBs without R&M (presentation 4 Jun done).",
    63: "Aug: awaiting Uematsu-san start for new dealer-visit schedule.",
    64: "Aug: market intelligence — 2nd potential 3rd-party provider contacted to compare.",
    67: "Aug: IDS audit schedule changed by IMG.",
    69: "Aug: IMG meeting to discuss follow-up after Uematsu-san start; visited DBs must verify Kaizen by end Nov.",
    73: "Aug: follow-up meeting with IML; gathering questionnaires.",
}

acc_notes = {
    6:  "Aug: ISV over-fender delay resolved via weekly allocation updates; IBX forecast + firm order for Sep-Nov (TTL 190x ~ EUR 136K).",
    7:  "Aug: FY27 target adjusted to EUR 2,809K on strong V-Cross/B-liner sales; YTD EUR 1.6M well ahead of plan.",
    8:  "Aug: 3-party meeting (7 Aug) for IAST to explain ISV engine-hood-protector claim outcome.",
    9:  "Aug: DB survey input combined and summary shared to IAST.",
    10: "Aug: Navco bracket modification needed after trial; next trial August.",
    12: "Aug: 300x rubber mats invoiced (EUR 16K); Kenya interest forwarded to IEA.",
    13: "Aug: IUK replenishment order EUR 146K invoiced (incl earlier BO).",
    15: "Aug: Petros tow-bar fit issue on MY26 regular cab — IBX assisting GDW investigation.",
    16: "Aug: KIMA fleet-deal accessory order complete; awaiting full container to invoice.",
    39: "Aug: LHD+RHD BSIS bracket promoted; first sales expected during August.",
    41: "Aug: DB's notified all accessory info is on SharePoint; 3-month DB evaluation to follow Sept.",
    61: "Aug: Navco LED bulb trial fitment at Dealer De Witte during August.",
    62: "Aug: 23/07 trial showed MY25/MY26 tailgate lock difference; minor bracket mod; next trial Aug.",
    69: "Aug: GZ & LS to order BSIS; LS order 20x received; IMOT ordered 55x from Midi.",
}


def apply(ws, notes, pic_col, jul_col, comment_col, sheet_name):
    written = 0
    green = 0
    skipped = []
    for row, note in notes.items():
        pic = ws.cell(row, pic_col).value
        if pic_has_bs(pic):
            skipped.append((row, "BS owner"))
            continue
        cell = ws.cell(row, comment_col)
        if not is_empty(cell.value):
            skipped.append((row, "cell already has colleague text"))
            continue
        cell.value = note
        cell.font = BODY_FONT
        if is_cross(ws.cell(row, jul_col).value):
            cell.fill = GREEN
            green += 1
        else:
            # ensure no stray fill on non-green cells
            cell.fill = PatternFill(fill_type=None)
        written += 1
    print(f"{sheet_name}: written={written} green={green} skipped={len(skipped)}")
    for s in skipped:
        print(f"  skip row {s[0]}: {s[1]}")
    return written, green


def main() -> int:
    wb = openpyxl.load_workbook(str(SRC))
    vp = wb["Action Plan VP"]
    acc = wb["Action Plan ACC"]

    vp_w, vp_g = apply(vp, vp_notes, VP_PIC, VP_JUL, VP_COMMENT, "VP")
    acc_w, acc_g = apply(acc, acc_notes, ACC_PIC, ACC_JUL, ACC_REPORT, "ACC")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUT))
    print(f"saved {OUT}")
    print(f"TOTAL green = {vp_g + acc_g}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
