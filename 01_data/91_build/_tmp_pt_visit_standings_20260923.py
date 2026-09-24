#!/usr/bin/env python3
"""Extract Portugal dealer standings: country rank, global rank of 518, section A-E."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DASH = ROOT / "01_data/08_Distributor/IDS_2026/reports/dashboard/IDS_2026_dealer_compare.html"
TOTAL = ROOT / "01_data/08_Distributor/IDS_2026/rankings/IDS_2026_TOTAL_ranking_all.xlsx"
B1 = ROOT / "01_data/08_Distributor/IDS_2026/rankings/IDS_2026_B1_ranking_all.xlsx"

t = DASH.read_text(encoding="utf-8", errors="replace")
start = t.find('"dealers":')
seg = t[start:]
dec = json.JSONDecoder()
payload, _ = dec.raw_decode(seg[seg.find("[") :])
dealers = [d for d in payload if d.get("country") == "Portugal"]
print("Portugal dealers in dashboard:", len(dealers))
print("keys:", sorted(dealers[0].keys()))

# global totals
wb = openpyxl.load_workbook(TOTAL, data_only=True)
ws = wb["All dealers"]
rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
vals = [r[4] for r in rows if isinstance(r[4], (int, float))]
n_all = len(vals)
net_avg = sum(vals) / n_all

wb2 = openpyxl.load_workbook(B1, data_only=True)
rows2 = [r for r in wb2["All dealers"].iter_rows(min_row=2, values_only=True)]
b1 = {(r[1], r[3]): r for r in rows2}


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip().upper()


b1n = {norm(f"{k[0]}|{k[1]}"): v for k, v in b1.items()}

out = []
for d in dealers:
    nm = d["name"]
    gr = next((r[0] for r in rows if norm(r[3]) == norm(nm) and r[1] == "Portugal"), None)
    b = b1n.get(norm(f"Portugal|{nm}"))
    pct = d["total_pct"]
    top_pct = 100 * sum(1 for v in vals if v >= pct * 100) / n_all
    secs = d.get("sections") or {}
    out.append(
        {
            "name": nm,
            "country_rank": d.get("country_rank"),
            "global_rank": gr,
            "pct": round(pct * 100, 2),
            "top_share_of_network": round(top_pct, 1),
            "vs_network_pp": round((pct * 100) - net_avg, 2),
            "vs_country_pp": round((pct * 100) - d["country_avg_pct"] * 100, 2),
            "sections": {
                s: {
                    "pct": round(v["pct"] * 100, 1),
                    "ctry_pp": round((v["pct"] - v["country_avg_pct"]) * 100, 1),
                    "net_pp": round((v["pct"] - v["network_avg_pct"]) * 100, 1),
                }
                for s, v in secs.items()
            },
            "b1_global_rank": b[0] if b else None,
            "b1_pct": b[4] if b else None,
        }
    )

out.sort(key=lambda r: r["country_rank"] or 99)
print()
print(f"network avg {net_avg:.2f}%  n={n_all}")
print(f"Portugal avg {dealers[0]['country_avg_pct']*100:.2f}%")
print()
hdr = f"{'#':>3} {'dealer':<28} {'tot%':>6} {'glob':>5} {'top%':>5} {'vsNet':>6} {'vsPt':>6} | A     B     C     D     E   | B1"
print(hdr)
for r in out:
    s = r["sections"]
    sec = " ".join(f"{s[k]['pct']:>5.1f}" if k in s else "  -  " for k in "ABCDE")
    print(
        f"{(r['country_rank'] or 0):>3} {r['name'][:28]:<28} {r['pct']:>6.2f} {(r['global_rank'] or 0):>5} "
        f"{r['top_share_of_network']:>5.1f} {r['vs_network_pp']:>6.1f} {r['vs_country_pp']:>6.1f} | {sec} | "
        f"{r['b1_global_rank']} ({r['b1_pct']})"
    )

visits = ["ROLMAX", "AUTO RENO PAÇOS FERREIRA", "AUTO INDUSTRIAL LEIRIA"]
print()
print("=== visit dealers ===")
for r in out:
    if norm(r["name"]) in [norm(v) for v in visits]:
        print(json.dumps(r, indent=1, ensure_ascii=False))
