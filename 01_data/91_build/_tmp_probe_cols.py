"""Check whether openpyxl collapses identical adjacent column widths when reading."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BOOK = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing"
            r"\Price_increase_by_part_group_per_customer.xlsx")

z = zipfile.ZipFile(BOOK)
wb = load_workbook(BOOK)
for i, name in enumerate(wb.sheetnames, start=1):
    xml = z.read(f"xl/worksheets/sheet{i}.xml").decode()
    m = re.search(r"<cols>.*?</cols>", xml, re.S)
    print("=" * 70)
    print(name)
    print("  raw <cols>:", m.group(0)[:600] if m else "(none)")
    dims = {k: (v.width, v.min, v.max) for k, v in wb[name].column_dimensions.items()}
    print("  openpyxl dims:", dims)
