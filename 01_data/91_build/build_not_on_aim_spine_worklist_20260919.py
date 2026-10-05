# -*- coding: utf-8 -*-
"""Work the Not-on-AIM residual the same way as the Sep-18 gap pass.

Does NOT patch aim_master_v2.csv. Writes Excel for Jerome.
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# Reuse the proven description rules from the gap classifier
import importlib.util

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
D = ISZE / "01_data" / "00_derived"
PRICING = ISZE / "01_data" / "01_Pricing"
OUT_XLSX = PRICING / "Not_on_AIM_spine_WORKLIST_classify_20260919.xlsx"
OUT_CSV = D / "pricing_not_on_aim_spine_worklist_20260919.csv"
OPEN_CMD = PRICING / "OPEN_Not_on_AIM_spine_WORKLIST_20260919.cmd"

spec = importlib.util.spec_from_file_location(
    "gapcls",
    ISZE / "01_data" / "91_build" / "classify_aim_gap_auto_assign_20260918.py",
)
gapcls = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(gapcls)
DESC_RULES = gapcls.DESC_RULES


def rule_match(desc: str):
    d = (desc or "").strip()
    for rx, l1, l2, l3, meth in DESC_RULES:
        if rx.search(d):
            return l1, l2, l3, meth
    return None


def main() -> None:
    resid = pd.read_csv(
        D / "pricing_not_on_aim_spine_residual_after_tools_asm_20260919.csv",
        dtype=str,
    )
    resid["pn"] = resid["pn"].str.replace(r"\D", "", regex=True).str.zfill(10)
    for c in ("billed_eur", "step_eur", "qty", "prior", "new", "weighted_pct"):
        if c in resid.columns:
            resid[c] = pd.to_numeric(resid[c], errors="coerce")

    assign = pd.read_csv(D / "aim_gap_auto_assign_20260918.csv", dtype=str, low_memory=False)
    assign["pn"] = assign["part_number"].str.replace(r"\D", "", regex=True).str.zfill(10)

    m = resid.merge(
        assign[
            [
                "pn",
                "confidence",
                "method",
                "proposed_aim_l1",
                "proposed_aim_l2",
                "proposed_aim_l3",
                "notes",
                "hs6",
            ]
        ],
        on="pn",
        how="left",
    )

    # Re-fire description rules (many unassigned rows already match clear rules)
    upgrades = []
    for i, r in m.iterrows():
        hit = rule_match(str(r.get("description") or ""))
        if not hit:
            continue
        l1, l2, l3, meth = hit
        conf0 = str(r.get("confidence") or "none").lower()
        has = bool(str(r.get("proposed_aim_l1") or "").strip())
        # Upgrade unassigned / none, or fill blanks; leave existing medium/low proposal
        # but stamp rule_high when rule is clear and prior was none/unassigned
        if conf0 in ("none", "nan", "") or not has:
            m.at[i, "proposed_aim_l1"] = l1
            m.at[i, "proposed_aim_l2"] = l2
            m.at[i, "proposed_aim_l3"] = l3
            m.at[i, "confidence"] = "high"
            m.at[i, "method"] = f"desc_rule_rerun:{meth}"
            m.at[i, "notes"] = (
                str(r.get("notes") or "")
                + "; upgraded_20260919_desc_rule"
            ).strip("; ")
            upgrades.append(r["pn"])
        elif conf0 in ("medium", "low") and str(r.get("proposed_aim_l1")) == l1:
            # Same L1 as prior proposal — promote to high when hard rule matches
            m.at[i, "confidence"] = "high"
            m.at[i, "method"] = f"desc_rule_promote:{meth}|was_{conf0}"
            m.at[i, "proposed_aim_l2"] = l2
            m.at[i, "proposed_aim_l3"] = l3
            m.at[i, "notes"] = (
                str(r.get("notes") or "")
                + "; promoted_20260919_desc_rule"
            ).strip("; ")
            upgrades.append(r["pn"])

    m["confidence"] = m["confidence"].fillna("none").replace({"": "none", "nan": "none"})
    m["work_status"] = m.apply(
        lambda r: (
            "ready_high_apply"
            if str(r["confidence"]).lower() == "high"
            and str(r.get("proposed_aim_l1") or "").strip()
            else (
                "review_medium"
                if str(r["confidence"]).lower() == "medium"
                else (
                    "review_low"
                    if str(r["confidence"]).lower() == "low"
                    else "needs_eq_or_manual"
                )
            )
        ),
        axis=1,
    )

    m = m.sort_values("billed_eur", ascending=False)
    m.to_csv(OUT_CSV, index=False)

    # ---- Excel
    wb = Workbook()
    navy = Font(bold=True, color="FFFFFF")
    head = PatternFill("solid", fgColor="15233C")
    title = Font(bold=True, size=14, color="15233C")

    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "Not on AIM spine — classify worklist (same method as Sep-18 gap pass)"
    ov["A1"].font = title
    ov["A2"] = (
        "Why this was still open: the Sep-18 auto-assign already scored these PNs, but only "
        "high-confidence rows were written into aim_master_v2. This residual is almost all "
        "medium / low / unassigned — so the pricing cube still shows Not on AIM spine. "
        "Tools / MX2 already stripped. ENGINE ASM / CAB ASM were never in this residual. "
        "This workbook does NOT patch the master — say apply when you want high rows written."
    )
    ov.merge_cells("A2:G2")
    ov["A2"].alignment = Alignment(wrap_text=True)
    ov.row_dimensions[2].height = 64

    def conf_sum(df):
        return (
            df.groupby("work_status", dropna=False)
            .agg(part_numbers=("pn", "nunique"), billed_eur=("billed_eur", "sum"), step_eur=("step_eur", "sum"))
            .reset_index()
            .sort_values("billed_eur", ascending=False)
        )

    ov["A4"] = "Work status"
    ov["B4"] = "Part numbers"
    ov["C4"] = "Parts billed EUR"
    ov["D4"] = "Step EUR"
    for c in ov[4]:
        c.fill = head
        c.font = navy
    summary = conf_sum(m)
    for i, r in enumerate(summary.itertuples(index=False), 5):
        ov.cell(i, 1, r.work_status)
        ov.cell(i, 2, int(r.part_numbers))
        c = ov.cell(i, 3, float(r.billed_eur))
        c.number_format = "#,##0.00"
        c2 = ov.cell(i, 4, float(r.step_eur or 0))
        c2.number_format = "#,##0.00"
    ov.cell(12, 1, "Total residual")
    ov.cell(12, 2, int(m["pn"].nunique()))
    t = ov.cell(12, 3, float(m["billed_eur"].sum()))
    t.number_format = "#,##0.00"
    ov.cell(14, 1, "Promoted / filled by description rules this pass")
    ov.cell(14, 2, len(set(upgrades)))
    ov.column_dimensions["A"].width = 28
    ov.column_dimensions["B"].width = 14
    ov.column_dimensions["C"].width = 18
    ov.column_dimensions["D"].width = 14

    cols = [
        "pn",
        "description",
        "billed_eur",
        "step_eur",
        "work_status",
        "confidence",
        "proposed_aim_l1",
        "proposed_aim_l2",
        "proposed_aim_l3",
        "method",
        "notes",
        "price_class",
        "dbs",
        "priced_both",
    ]
    rename = {
        "pn": "Part number",
        "description": "Description",
        "billed_eur": "Parts billed EUR",
        "step_eur": "Step EUR",
        "work_status": "Work status",
        "confidence": "Confidence",
        "proposed_aim_l1": "Proposed AIM L1",
        "proposed_aim_l2": "Proposed AIM L2",
        "proposed_aim_l3": "Proposed AIM L3",
        "method": "Method",
        "notes": "Notes",
        "price_class": "Price class",
        "dbs": "Distributors",
        "priced_both": "Priced both sides",
    }

    def sheet(name: str, df: pd.DataFrame) -> None:
        w = wb.create_sheet(name)
        out = df[[c for c in cols if c in df.columns]].rename(
            columns={k: v for k, v in rename.items() if k in df.columns}
        )
        for r in dataframe_to_rows(out, index=False, header=True):
            w.append(r)
        for cell in w[1]:
            cell.fill = head
            cell.font = navy
        headers = [c.value for c in w[1]]
        for col_i, h in enumerate(headers, 1):
            for row in range(2, w.max_row + 1):
                cell = w.cell(row, col_i)
                if h in ("Parts billed EUR", "Step EUR") and isinstance(
                    cell.value, (int, float)
                ):
                    cell.number_format = "#,##0.00"
            w.column_dimensions[get_column_letter(col_i)].width = min(
                32, max(12, len(str(h)) + 2)
            )
        w.auto_filter.ref = w.dimensions
        w.freeze_panes = "A2"

    sheet("1 Ready high - apply", m[m["work_status"] == "ready_high_apply"])
    sheet("2 Review medium", m[m["work_status"] == "review_medium"])
    sheet("3 Review low", m[m["work_status"] == "review_low"])
    sheet("4 Needs EQ or manual", m[m["work_status"] == "needs_eq_or_manual"])
    sheet("All residual", m)

    # L1 mix for ready + medium
    mix = wb.create_sheet("Proposed L1 mix")
    mix["A1"] = "Proposed AIM L1 — ready high + medium (by billed euro)"
    mix["A1"].font = title
    sub = m[m["work_status"].isin(["ready_high_apply", "review_medium"])]
    g = (
        sub.groupby("proposed_aim_l1", dropna=False)
        .agg(pn=("pn", "nunique"), billed=("billed_eur", "sum"), step=("step_eur", "sum"))
        .reset_index()
        .sort_values("billed", ascending=False)
    )
    mix["A3"] = "Proposed AIM L1"
    mix["B3"] = "Part numbers"
    mix["C3"] = "Parts billed EUR"
    mix["D3"] = "Step EUR"
    for c in mix[3]:
        c.fill = head
        c.font = navy
    for i, r in enumerate(g.itertuples(index=False), 4):
        mix.cell(i, 1, r.proposed_aim_l1 or "(blank)")
        mix.cell(i, 2, int(r.pn))
        c = mix.cell(i, 3, float(r.billed))
        c.number_format = "#,##0.00"
        c2 = mix.cell(i, 4, float(r.step or 0))
        c2.number_format = "#,##0.00"
    mix.column_dimensions["A"].width = 28
    mix.column_dimensions["B"].width = 14
    mix.column_dimensions["C"].width = 18
    mix.column_dimensions["D"].width = 14

    wb.save(OUT_XLSX)
    OPEN_CMD.write_text(
        '@echo off\r\nstart "" "%~dp0Not_on_AIM_spine_WORKLIST_classify_20260919.xlsx"\r\n',
        encoding="ascii",
        newline="",
    )

    print("OUT", OUT_XLSX)
    print(summary.to_string(index=False))
    print("upgraded_or_promoted", len(set(upgrades)))
    print(
        "ready_high billed",
        float(m.loc[m["work_status"] == "ready_high_apply", "billed_eur"].sum()),
    )


if __name__ == "__main__":
    main()
