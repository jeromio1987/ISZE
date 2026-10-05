"""Capture the workbook's figures BEFORE the rebuild, for the before/after report."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BOOK = ISZE / "01_data" / "01_Pricing" / "Price_increase_by_part_group_per_customer.xlsx"
OUT = ISZE / "01_data" / "91_build" / "_price_inc_before.json"

wb = load_workbook(BOOK)
report: dict = {"path": str(BOOK), "sheets": wb.sheetnames}

ws = wb["By part group (network)"]
rows = list(ws.iter_rows(values_only=True))
head = rows[0]
step_i = head.index("Step (EUR)")
billed_i = head.index("L12M billed (EUR)")
qty_i = head.index("L12M qty")
rate_i = head.index("Increase %")
pns_i = head.index("Part numbers")
data = rows[1:]
report["network_sheet"] = {
    "header": list(head),
    "rows": len(data),
    "step_sum_eur": round(sum(r[step_i] for r in data), 2),
    "billed_sum_eur": round(sum(r[billed_i] for r in data), 2),
    "qty_sum": sum(r[qty_i] for r in data),
    "part_numbers_sum": sum(r[pns_i] for r in data),
    "rate_min_pct": min(r[rate_i] for r in data),
    "rate_max_pct": max(r[rate_i] for r in data),
    "top_by_rate": data[0][1],
}

ws2 = wb["Most and least up per customer"]
rows2 = list(ws2.iter_rows(values_only=True))
nan_cells = []
for row in ws2.iter_rows():
    for c in row:
        if isinstance(c.value, str) and "nan" in c.value.lower():
            nan_cells.append(c.coordinate)
report["most_least"] = {
    "header": list(rows2[0]),
    "rows": len(rows2) - 1,
    "nan_cells": nan_cells,
}

ws3 = wb["Every customer x part group"]
rows3 = list(ws3.iter_rows(values_only=True))
h3 = rows3[0]
s3 = h3.index("Step (EUR)")
b3 = h3.index("L12M billed (EUR)")
d3 = rows3[1:]
report["customer_group"] = {
    "header": list(h3),
    "rows": len(d3),
    "step_sum_eur": round(sum(r[s3] for r in d3), 2),
    "billed_sum_eur": round(sum(r[b3] for r in d3), 2),
    "blank_per_truck": sum(1 for r in d3 if r[h3.index("EUR per truck")] is None),
}

ws4 = wb["Increase % grid"]
rows4 = list(ws4.iter_rows(values_only=True))
report["grid"] = {"header": list(rows4[0]), "rows": len(rows4) - 1,
                  "cols": len(rows4[0])}

# formatting defects
defects = {}
for name in wb.sheetnames:
    s = wb[name]
    defects[name] = {
        "freeze_panes": s.freeze_panes,
        "autofilter": s.auto_filter.ref,
        "default_width_cols": sum(
            1 for i in range(1, s.max_column + 1)
            if s.column_dimensions[s.cell(1, i).column_letter].width in (None, 8.43)),
        "cols": s.max_column,
        "formula_cells": sum(
            1 for row in s.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")),
        "conditional_formats": len(list(s.conditional_formatting)),
        "fills": sum(1 for row in s.iter_rows() for c in row
                     if c.fill is not None and c.fill.fill_type == "solid"),
        "number_formats": sorted({c.number_format for row in s.iter_rows() for c in row}),
    }
report["defects_before"] = defects

OUT.write_text(json.dumps(report, indent=2), encoding="utf-8")
print(json.dumps(report, indent=2))
