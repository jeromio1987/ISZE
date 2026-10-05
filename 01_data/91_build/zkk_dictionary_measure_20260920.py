# -*- coding: utf-8 -*-
"""
ZKK code dictionary -- measurement pass.

Read-only. Nothing is rebuilt. Sources:
  01_data/00_derived/jp_codes_cache_20260920.csv   (536,854 rows, per-Pno Japan codes)
  01_data/00_derived/aim_master_v2.csv             (165,213 rows, the spine: euro + AIM + jp_*)
  01_data/HKK Code 20221028.xlsx                   (2022 sales-category workbook)

Outputs nothing to disk except stdout. Every table is printed as markdown-ready rows.

Run: python 01_data/91_build/zkk_dictionary_measure_20260920.py
"""
from __future__ import annotations

import collections
import csv
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "01_data" / "00_derived" / "jp_codes_cache_20260920.csv"
SPINE = ROOT / "01_data" / "00_derived" / "aim_master_v2.csv"
HKK2022 = ROOT / "01_data" / "HKK Code 20221028.xlsx"
ZKK_COUNTS = ROOT / "01_data" / "90_INBOX" / "_zkk_code_counts_20260920.json"

BLANKISH = {"", "-", "--", "N/A", "NA", "NONE", "UNKNOWN", "(BLANK)"}


def s(v) -> str:
    if v is None:
        return ""
    t = str(v).strip()
    if t.endswith(".0") and t[:-2].replace("-", "").isdigit():
        t = t[:-2]
    return t


def collapse(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip()


def noun_task(desc: str) -> str:
    """Task definition: first token before any ';' or ','."""
    d = collapse((desc or "").upper())
    head = re.split(r"[;,]", d)[0]
    head = collapse(head).strip(" .-*/")
    return head if head and head not in BLANKISH else "(blank)"


def noun_wide(desc: str) -> str:
    """Wider definition: also splits on '(' and '/'. Benchmark used ';' ',' '('."""
    d = collapse((desc or "").upper())
    head = re.split(r"[;,]", d)[0]
    head = re.split(r"[(\[/]", head)[0]
    head = collapse(head).strip(" .-*")
    return head if head and head not in BLANKISH else "(blank)"


def nfkc(t) -> str:
    return unicodedata.normalize("NFKC", str(t or "")).strip()


def main() -> int:
    out = []
    P = out.append

    # ---------------- load cache ----------------
    rows_by_pn: dict[str, list[dict]] = collections.defaultdict(list)
    zkk_rows = collections.Counter()
    hkk_rows = collections.Counter()
    zuk_rows = collections.Counter()
    jp_desc_blank_rows = 0
    n_cache = 0
    with CACHE.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            pn = s(r.get("pn"))
            if not pn:
                continue
            n_cache += 1
            zkk = s(r.get("zkk"))
            hkk = s(r.get("hkk"))
            zuk = s(r.get("zuk"))
            desc = s(r.get("jp_description"))
            zkk_rows[zkk] += 1
            hkk_rows[hkk] += 1
            zuk_rows[zuk] += 1
            if collapse(desc).upper() in BLANKISH:
                jp_desc_blank_rows += 1
            rows_by_pn[pn].append({"zkk": zkk, "hkk": hkk, "zuk": zuk, "desc": desc})

    pn_zkk_multi = sum(
        1 for pn, rr in rows_by_pn.items() if len({x["zkk"] for x in rr}) > 1
    )
    pn_hkk_multi = sum(
        1 for pn, rr in rows_by_pn.items() if len({x["hkk"] for x in rr}) > 1
    )

    # ---------------- load spine ----------------
    spine = {}
    total_eur = 0.0
    with SPINE.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            pn = s(r.get("part_number"))
            if not pn:
                continue
            e = float(r.get("l12m_billed_eur") or 0)
            spine[pn] = {
                "eur": e,
                "desc": s(r.get("description")),
                "aim_l1": s(r.get("aim_l1_fixed")) or s(r.get("aim_l1")) or "(blank)",
                "aim_l2": s(r.get("aim_l2_fixed")) or s(r.get("aim_l2")) or "(blank)",
            }
            total_eur += e

    # ---------------- common base ----------------
    # one ZKK per PN is required for a clean base; take the modal code where a PN shows more than one
    base: dict[str, dict] = {}
    for pn, rr in rows_by_pn.items():
        if pn not in spine:
            continue
        modal = collections.Counter(x["zkk"] for x in rr).most_common(1)[0][0]
        if not modal:
            continue
        base[pn] = {
            "eur": spine[pn]["eur"],
            "desc": rr[0]["desc"],
            "zkk": modal,
            "hkk": collections.Counter(x["hkk"] for x in rr).most_common(1)[0][0],
            "zuk": collections.Counter(x["zuk"] for x in rr).most_common(1)[0][0],
            "aim_l1": spine[pn]["aim_l1"],
            "aim_l2": spine[pn]["aim_l2"],
        }
    base_eur = sum(v["eur"] for v in base.values())

    P(f"### SECTION 0 - base")
    P(f"cache rows            : {n_cache:,}")
    P(f"cache distinct PNs    : {len(rows_by_pn):,}")
    P(f"cache PNs with >1 ZKK : {pn_zkk_multi:,}")
    P(f"cache PNs with >1 HKK : {pn_hkk_multi:,}")
    P(f"cache jp_desc blank   : {jp_desc_blank_rows:,} ({100*jp_desc_blank_rows/max(n_cache,1):.2f}% of rows)")
    P(f"distinct ZKK values   : {len([k for k in zkk_rows if k])}  (blank rows: {zkk_rows.get('',0)})")
    P(f"distinct HKK values   : {len([k for k in hkk_rows if k])}  (blank rows: {hkk_rows.get('',0)})")
    P(f"distinct ZUK values   : {len([k for k in zuk_rows if k])}  (blank rows: {zuk_rows.get('',0)})")
    P(f"spine rows            : {len(spine):,}   total EUR {total_eur:,.2f}")
    P(f"COMMON BASE PNs       : {len(base):,}   EUR {base_eur:,.2f}  ({100*base_eur/total_eur:.2f}% of book)")

    # ---------------- SECTION 1 - structural ----------------
    P("")
    P("### SECTION 1 - structural, ZKK vs HKK")
    hkk_vals = sorted(k for k in hkk_rows if k)
    hkk_pfx = sorted({k[:3] for k in hkk_vals})
    zkk_vals = sorted(k for k in zkk_rows if k)
    zkk_set = set(zkk_vals)
    full_match = [k for k in hkk_vals if k in zkk_set]
    pfx_match = sorted(p for p in hkk_pfx if p in zkk_set)
    P(f"HKK distinct values                     : {len(hkk_vals)}")
    P(f"HKK distinct 3-char prefixes            : {len(hkk_pfx)}")
    P(f"HKK 3-char prefixes that are a ZKK value: {len(pfx_match)}  -> {pfx_match}")
    P(f"HKK full 4-char values that are a ZKK   : {len(full_match)} -> {full_match}")
    P(f"HKK 3-char prefixes NOT a ZKK value     : {sorted(set(hkk_pfx)-zkk_set)}")
    P(f"HKK length set                          : {sorted({len(k) for k in hkk_vals})}")
    P(f"ZKK length set                          : {sorted({len(k) for k in zkk_vals})}")
    P(f"ZUK length set                          : {sorted({len(k) for k in zuk_rows if k})}")
    P(f"ZUK middle-position chars               : {sorted({k[2] for k in zuk_rows if k})}")
    P(f"ZUK chars per position                  : "
      + str([sorted({k[i] for k in zuk_rows if k}) for i in range(5)]))

    # 2022 workbook prefixes
    wb_pfx = set()
    try:
        wb = openpyxl.load_workbook(HKK2022, read_only=True, data_only=True)
        for sh in wb.sheetnames:
            ws = wb[sh]
            for row in ws.iter_rows(values_only=True):
                cells = [nfkc(c) for c in row]
                for i, c in enumerate(cells):
                    if re.fullmatch(r"\d{3}", c):
                        if i + 1 < len(cells) and re.fullmatch(r"[A-Za-z0-9]", cells[i + 1] or ""):
                            wb_pfx.add(c)
        wb.close()
    except Exception as exc:  # pragma: no cover
        P(f"2022 workbook read failed: {exc}")
    P(f"2022 workbook prefixes (my extract)     : {len(wb_pfx)} -> {sorted(wb_pfx)}")
    P(f"ZKK values that are a 2022 prefix       : {len(zkk_set & wb_pfx)} -> {sorted(zkk_set & wb_pfx)}")

    # ---------------- SECTION 2 - per ZKK code table ----------------
    per: dict[str, dict] = {}
    for pn, v in base.items():
        c = per.setdefault(v["zkk"], {"pns": 0, "eur": 0.0, "nouns": collections.Counter(),
                                      "nouns_wide": collections.Counter(),
                                      "aim": collections.Counter(), "aim_eur": collections.Counter()})
        c["pns"] += 1
        c["eur"] += v["eur"]
        c["nouns"][noun_task(v["desc"])] += 1
        c["nouns_wide"][noun_wide(v["desc"])] += 1
        c["aim"][v["aim_l1"]] += 1
        c["aim_eur"][v["aim_l1"]] += v["eur"]

    def purity(per_map, key):
        """row-weighted dominant-noun share, and single-noun-code count"""
        tot = 0
        dom_rows = 0
        single = 0
        for c, v in per_map.items():
            n = sum(v[key].values())
            if not n:
                continue
            tot += n
            dom_rows += v[key].most_common(1)[0][1]
            if len(v[key]) == 1:
                single += 1
        return (100 * dom_rows / tot if tot else 0.0), single

    P("")
    P("### SECTION 2 - per ZKK code (billed book, common base)")
    P("| ZKK | part numbers | EUR | file rows | top 1 noun | share | top 2 noun | share | top 3 noun | share |")
    P("|---|---:|---:|---:|---|---:|---|---:|---|---:|")
    for c, v in sorted(per.items(), key=lambda kv: -kv[1]["eur"]):
        t3 = v["nouns"].most_common(3)
        while len(t3) < 3:
            t3.append(("(none)", 0))
        P("| `{c}` | {n:,} | {e:,.0f} | {fr:,} | {a} | {pa:.1f}% | {b} | {pb:.1f}% | {d} | {pd:.1f}% |".format(
            c=c, n=v["pns"], e=v["eur"], fr=zkk_rows.get(c, 0),
            a=t3[0][0], pa=100 * t3[0][1] / v["pns"],
            b=t3[1][0], pb=100 * t3[1][1] / v["pns"],
            d=t3[2][0], pd=100 * t3[2][1] / v["pns"]))

    # ---------------- SECTION 3 - coherence ----------------
    P("")
    P("### SECTION 3 - coherence")
    rows = []
    for c, v in per.items():
        top_n, top_cnt = v["nouns"].most_common(1)[0]
        top_wide_n, top_wide_cnt = v["nouns_wide"].most_common(1)[0]
        rows.append({
            "zkk": c,
            "pns": v["pns"],
            "eur": v["eur"],
            "top": top_n,
            "share": 100 * top_cnt / v["pns"],
            "top_wide": top_wide_n,
            "share_wide": 100 * top_wide_cnt / v["pns"],
            "distinct_nouns": len(v["nouns"]),
            "top3_share": 100 * sum(x for _, x in v["nouns"].most_common(3)) / v["pns"],
        })
    coop_rows = [r for r in rows if r["pns"] >= 200 and r["share"] >= 80.0]
    coop_rows_wide = [r for r in rows if r["pns"] >= 200 and r["share_wide"] >= 80.0]
    near = [r for r in rows if r["pns"] >= 200 and 60.0 <= r["share"] < 80.0]
    near_wide = [r for r in rows if r["pns"] >= 200 and 60.0 <= r["share_wide"] < 80.0]
    frag = [r for r in rows if r["pns"] >= 200 and r["share"] < 40.0]

    def tab(rs):
        return " ".join(f"`{r['zkk']}`({r['pns']:,}pn,{r['share']:.0f}%,{r['top']})" for r in sorted(rs, key=lambda x: -x["pns"]))

    def tabw(rs):
        return " ".join(f"`{r['zkk']}`({r['pns']:,}pn,{r['share_wide']:.0f}%,{r['top_wide']})" for r in sorted(rs, key=lambda x: -x["pns"]))

    P(f"codes in billed book                              : {len(rows)}")
    P(f"codes with >=200 PNs                              : {sum(1 for r in rows if r['pns']>=200)}")
    P(f"COHERENT  (>=200 PNs, dominant noun >=80%)        : {len(coop_rows)}")
    P(f"   -> {tab(coop_rows)}")
    P(f"COHERENT (wide noun def, >=200 PNs, >=80%)        : {len(coop_rows_wide)}")
    P(f"   -> {tabw(coop_rows_wide)}")
    P(f"NEAR      (>=200 PNs, 60-80%)                     : {len(near)}")
    P(f"   -> {tab(near)}")
    P(f"NEAR wide (>=200 PNs, 60-80%)                     : {len(near_wide)}")
    P(f"   -> {tabw(near_wide)}")
    P(f"FRAGMENTED(>=200 PNs, dominant noun <40%)         : {len(frag)}")
    P(f"   -> {tab(frag)}")
    P("")
    P("full ranking by dominant-noun share (PN>=200 first):")
    P("| ZKK | PNs | EUR | dom noun | share | share(wide) | distinct nouns | top3 |")
    P("|---|---:|---:|---|---:|---:|---:|---:|")
    for r in sorted([x for x in rows if x["pns"] >= 200], key=lambda x: -x["share"]):
        P(f"| `{r['zkk']}` | {r['pns']:,} | {r['eur']:,.0f} | {r['top']} | {r['share']:.1f}% | "
          f"{r['share_wide']:.1f}% | {r['distinct_nouns']:,} | {r['top3_share']:.1f}% |")
    P("")
    P("small codes (<200 PNs), listed:")
    P("| ZKK | PNs | EUR | dom noun | share | distinct nouns |")
    P("|---|---:|---:|---|---:|---:|")
    for r in sorted([x for x in rows if x["pns"] < 200], key=lambda x: -x["eur"]):
        P(f"| `{r['zkk']}` | {r['pns']:,} | {r['eur']:,.0f} | {r['top']} | {r['share']:.1f}% | {r['distinct_nouns']:,} |")

    # ---------------- SECTION 4 - ZKK vs AIM ----------------
    P("")
    P("### SECTION 4 - ZKK vs AIM family")
    fam_tot = collections.Counter()
    for pn, v in base.items():
        fam_tot[v["aim_l1"]] += v["eur"]
    big = [r for r in rows if r["eur"] >= 0.005 * base_eur]
    P(f"ZKK codes >=0.5% of base EUR: {len(big)}")
    P("| ZKK | EUR | share of book | dominant AIM L1 | share of code EUR | 2nd AIM L1 | share |")
    P("|---|---:|---:|---|---:|---|---:|")
    detail = []
    for r in sorted(big, key=lambda x: -x["eur"]):
        c = r["zkk"]
        ae = per[c]["aim_eur"]
        top = ae.most_common(2)
        while len(top) < 2:
            top.append(("(none)", 0.0))
        detail.append((c, r["eur"], 100 * r["eur"] / base_eur, top[0][0], 100 * top[0][1] / r["eur"],
                       top[1][0], 100 * top[1][1] / r["eur"]))
        P(f"| `{c}` | {r['eur']:,.0f} | {100*r['eur']/base_eur:.2f}% | {top[0][0]} | "
          f"{100*top[0][1]/r['eur']:.1f}% | {top[1][0]} | {100*top[1][1]/r['eur']:.1f}% |")
    P("")
    P("Reverse view - how many ZKK codes inside each AIM L1, and top code share:")
    fam_zkk = collections.defaultdict(collections.Counter)
    for pn, v in base.items():
        fam_zkk[v["aim_l1"]][v["zkk"]] += v["eur"]
    P("| AIM L1 | EUR | ZKK codes >=1% of family | top ZKK | top ZKK share of family |")
    P("|---|---:|---:|---|---:|")
    for fam, tot in fam_tot.most_common(15):
        cc = fam_zkk[fam]
        nz = sum(1 for _, e in cc.items() if e >= 0.01 * tot)
        t = cc.most_common(1)[0]
        P(f"| {fam} | {tot:,.0f} | {nz} | `{t[0]}` | {100*t[1]/tot:.1f}% |")
    P("")
    txt = []
    for (c, e, sh, f1, p1, f2, p2) in detail[:8]:
        txt.append(f"`{c}` EUR{e:,.0f} ({sh:.2f}% book): {f1} {p1:.1f}%, {f2} {p2:.1f}%")
    P("largest codes, dominance detail: " + " | ".join(txt))

    # ---------------- SECTION 5 - HKK vs ZKK as a level ----------------
    P("")
    P("### SECTION 5 - HKK vs ZKK as a level (same base, same method)")
    def purity_of(key_fn, noun_fn):
        mp = collections.defaultdict(collections.Counter)
        for pn, v in base.items():
            mp[key_fn(v)][noun_fn(v["desc"])] += 1
        tot = sum(sum(c.values()) for c in mp.values())
        dom = sum(c.most_common(1)[0][1] for c in mp.values())
        single = sum(1 for c in mp.values() if len(c) == 1)
        return len(mp), (100 * dom / tot if tot else 0.0), single, tot

    for label, kf in (("ZKK (3 char)", lambda v: v["zkk"]),
                      ("HKK (4 char)", lambda v: v["hkk"]),
                      ("HKK first 3 char", lambda v: v["hkk"][:3]),
                      ("AIM L1", lambda v: v["aim_l1"]),
                      ("AIM L1>L2", lambda v: v["aim_l1"] + " > " + v["aim_l2"])):
        for nlabel, nf in (("task-noun", noun_task), ("wide-noun", noun_wide)):
            ncodes, pur, single, tot = purity_of(kf, nf)
            P(f"{label:<18} {nlabel:<10} codes={ncodes:<4} purity={pur:6.2f}%  single-noun-codes={single:<4} rows={tot:,}")
    P("")
    P("HKK code-level detail, billed book, sorted by EUR:")
    hkk_agg = collections.defaultdict(lambda: {"pns": 0, "eur": 0.0, "nouns": collections.Counter(), "zkk": collections.Counter()})
    for pn, v in base.items():
        h = hkk_agg[v["hkk"]]
        h["pns"] += 1
        h["eur"] += v["eur"]
        h["nouns"][noun_task(v["desc"])] += 1
        h["zkk"][v["zkk"]] += 1
    P("| HKK | PNs | EUR | dom noun | share | dom ZKK inside | distinct ZKK |")
    P("|---|---:|---:|---|---:|---|---:|")
    for c, v in sorted(hkk_agg.items(), key=lambda kv: -kv[1]["eur"]):
        if v["pns"] < 2:
            continue
        t = v["nouns"].most_common(1)[0]
        z = v["zkk"].most_common(1)[0]
        P(f"| `{c}` | {v['pns']:,} | {v['eur']:,.0f} | {t[0]} | {100*t[1]/v['pns']:.1f}% | `{z[0]}` | {len(v['zkk'])} |")

    # ---------------- SECTION 6 - ZUK ----------------
    P("")
    P("### SECTION 6 - ZUK shape")
    P(f"distinct ZUK values in cache      : {len([k for k in zuk_rows if k])}")
    P(f"distinct ZUK values on billed book: {len({v['zuk'] for v in base.values()})}")
    zuk_agg = collections.defaultdict(lambda: {"pns": 0, "eur": 0.0})
    for pn, v in base.items():
        zuk_agg[v["zuk"]]["pns"] += 1
        zuk_agg[v["zuk"]]["eur"] += v["eur"]
    top = sorted(zuk_agg.items(), key=lambda kv: -kv[1]["eur"])[:15]
    for c, v in top:
        P(f"  ZUK `{c}` PNs={v['pns']:,} EUR={v['eur']:,.0f}")
    P(f"top 10 ZUK share of base EUR      : {100*sum(v['eur'] for _,v in top[:10])/base_eur:.2f}%")
    P(f"ZUK codes >=0.5% of base EUR      : {sum(1 for _,v in zuk_agg.items() if v['eur']>=0.005*base_eur)}")
    # does ZUK separate what HKK/ZKK do not?
    pair_zkk_shared = sum(v["pns"] for v in zuk_agg.values())
    P(f"ZUK 3rd char distribution (billed) : {collections.Counter(c[2] for c in zuk_agg for _ in range(zuk_agg[c]['pns'])).most_common()}")

    print("\n".join(out))
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    raise SystemExit(main())
