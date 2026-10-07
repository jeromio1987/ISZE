"""Build the KPI panel from the per-distributor check sheets (read-only inputs).

Reads the 80 submission files in
  01_data/90_INBOX/_bo_scrape_3m_20261007/mailbox_extra
and writes one tidy long table to
  01_data/90_INBOX/_kpi_panel_20261007/kpi_panel.csv

Columns: period | report_month_source | distributor | metric | value | unit |
         sheet_name | source_file

Design notes
------------
* Everything is read as text; numbers are converted deliberately so that a
  percentage like 0.97 is never turned into a date and 'N/D' stays a string.
* A cell that openpyxl returns as a datetime (a date-formatted numeric cell) is
  converted back to its Excel serial number, not reformatted as a date.
* 'N/D' / 'N/A' and free-text answers are kept as literal values (a distributor
  that does not report is a finding, not a blank).
* Two layouts exist so far:
    layout "flat"    -> month columns D..O, one row per metric
    layout "country" -> an extra Country column D, then month columns E..P,
                        several rows can share one metric (one per country)
  The layout is detected per file; nothing is assumed to fit one template.
"""
from __future__ import annotations

import base64
import csv
import io
import json
import re
import warnings
from datetime import date, datetime, time
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

SRC = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_bo_scrape_3m_20261007\mailbox_extra")
OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_kpi_panel_20261007")
OUT.mkdir(parents=True, exist_ok=True)

MONTHS = ["January", "Feburary", "February", "March", "April", "May", "June",
          "July", "August", "September", "October", "November", "December"]
MONTH_NUM = {}
for i, m in enumerate(["January", "February", "March", "April", "May", "June",
                       "July", "August", "September", "October", "November", "December"], 1):
    MONTH_NUM[m.lower()] = i
MONTH_NUM["feburary"] = 2  # sheet typo

# --- metric recognition -----------------------------------------------------
def canonical_metric(label: str) -> tuple[str, str] | None:
    """Return (metric_name, unit) for a row label, or None if not a KPI row."""
    s = (label or "").lower()
    if "days or months of parts" in s or "(mos)" in s:
        return "Days or Months of Parts in Stock (MOS)", "months"
    if "first fill rate" in s:
        return "First Fill Rate", "ratio (as written: 0.97 = 97%)"
    if "b/o lines" in s or "back order lines" in s:
        return "Back Order Lines", "lines (count)"
    if "parts sales distributor to all dealers" in s or "parts sales db to dlr" in s:
        return "Parts Sales DB to DLR", "EUR (as written)"
    if "inventory" in s and "opening" in s:
        return "Parts Inventory (Opening)", "EUR (as written)"
    if "inventory" in s and "closing" in s:
        return "Parts Inventory (Closing)", "EUR (as written)"
    if "inventory" in s:
        return "Parts Inventory", "EUR (as written)"
    if "recall campaign rate" in s:
        return "Recall Campaign Rate", "ratio (as written)"
    if "vehicles maintenance" in s:
        return "No. of Vehicles Maintenance", "count"
    if "vehicles visits" in s:
        return "No. of Vehicles Visits", "count"
    if "warranty cases" in s:
        return "No. of Warranty Cases", "count"
    if "customer satisfaction" in s:
        return "Customer Satisfaction", "ratio (as written)"
    if "customer complaints" in s or "customer complaint" in s:
        return "No. of Customer Complaints", "count"
    if "dealer network efficiency" in s:
        return "Dealer Network Efficiency Rate", "ratio (as written)"
    if "registered lcv" in s:
        return "Vehicles Registered LCV", "count"
    if "registered cv" in s:
        return "Vehicles Registered CV", "count"
    if "service contracts sold lcv" in s:
        return "Service Contracts Sold LCV", "count"
    if "service contracts sold cv" in s:
        return "Service Contracts Sold CV", "count"
    if "warranty extensions sold lcv" in s:
        return "Warranty Extensions Sold LCV", "count"
    if "warranty extensions sold cv" in s:
        return "Warranty Extensions Sold CV", "count"
    if "service contract" in s:
        return "Service Contracts Sold (unspecified)", "count"
    if "warranty extension" in s:
        return "Warranty Extensions Sold (unspecified)", "count"
    if "aftersales campaign" in s:
        return "Aftersales Campaigns Info", "text"
    return None


def cell_text(v) -> str | None:
    """Convert a raw openpyxl cell to text without inventing dates."""
    if v is None:
        return None
    if isinstance(v, datetime):
        # date-formatted numeric cell: return its Excel serial, not a date
        base = datetime(1899, 12, 30)
        delta = v - base
        return f"{delta.days + delta.seconds / 86400:.6f}".rstrip("0").rstrip(".")
    if isinstance(v, (date, time)):
        return str(v)
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    return str(v).strip()


def load_wb(path: Path):
    raw = path.read_bytes()
    if raw[:2] == b"PK":
        return openpyxl.load_workbook(path, data_only=True)
    # base64-encoded xlsx (as some Outlook saves can be)
    return openpyxl.load_workbook(io.BytesIO(base64.b64decode(raw)), data_only=True)


def is_check_sheet(ws) -> bool:
    title = cell_text(ws.cell(1, 3).value) or ""
    if "distributor kpi" in title.lower():
        return True
    for r in range(1, min(ws.max_row, 15) + 1):
        v = cell_text(ws.cell(r, 3).value) or ""
        if "first fill rate" in v.lower():
            return True
    return False


def check_sheets(wb):
    """Every sheet in the workbook that is a monthly check sheet."""
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if is_check_sheet(ws):
            out.append((sn, ws))
    return out


def header_info(ws):
    """Find the header row and the month columns (and country column if any)."""
    for r in range(1, min(ws.max_row, 12) + 1):
        row = [cell_text(ws.cell(r, c).value) or "" for c in range(1, ws.max_column + 1)]
        if any("january" in x.lower() for x in row):
            date_cols = []
            for c, x in enumerate(row, 1):
                xl = x.lower()
                if xl in MONTH_NUM:
                    date_cols.append((MONTH_NUM[xl], c))
            country_col = None
            for c, x in enumerate(row, 1):
                if x.strip().lower() == "country":
                    country_col = c
            return r, date_cols, country_col
    return None, [], None


def year_of(ws) -> int | None:
    t = cell_text(ws.cell(2, 3).value) or ""
    m = re.search(r"(20\d\d)", t)
    return int(m.group(1)) if m else None


def distributor_from_title(ws) -> str | None:
    t = cell_text(ws.cell(2, 3).value) or ""
    t = re.sub(r"(?i)^year\s*20\d\d", "", t)
    t = re.sub(r"(?i)\brevised\b", "", t)
    t = t.replace("&", "&").strip(" -_")
    return t or None


def title_names(ws) -> list[str]:
    t2 = cell_text(ws.cell(2, 3).value) or ""
    t2 = re.sub(r"(?i)^year\s*20\d\d", "", t2)
    if "\n" in t2 or "\r" in t2:
        parts = re.split(r"[\r\n]+", t2)
    else:
        parts = re.split(r"\s{2,}", t2)
    out = []
    for part in parts:
        part = re.sub(r"(?i)\brevised\b", "", part)
        part = part.replace("\u3000", " ").strip(" -_\t")
        if part:
            out.append(ALIASES.get(part, part))
    return out


def tail_names(ws) -> list[str]:
    """Distributor name carried at the end of the C1 title, if any."""
    t1 = cell_text(ws.cell(1, 3).value) or ""
    for pre in ("Check Sheet Monthly", "Check Sheet", "Checklist", "KPI", "SPI"):
        i = t1.rfind(pre)
        if i != -1:
            tail = t1[i + len(pre):].strip(" -_&\t")
            low = tail.lower()
            if tail and not re.search(r"\b(check|sheet|spi|kpi|img|isze|distributor)\b", low):
                return [ALIASES.get(tail, tail)]
    return []


def distributor_from_filename(name: str) -> str | None:
    """Best-effort distributor token from the file name (fallback only)."""
    stem = Path(name).stem
    stem = re.sub(r"(?i)^copy of\b", "", stem).strip()
    # text between the yyyymmdd stamp and the boilerplate that follows it
    m = re.search(r"\d{8}[_ ]+(.+?)\s*(?:\bKPI\b|\bSPI\b|\bCheck\b|$)", stem, re.I)
    if m:
        token = m.group(1).strip(" -–_")
        token = re.sub(r"(?i)^(isuzu)\s+", "", token).strip(" -–_")
        token = re.sub(r"(?i)\b(monthly|revised|2026|26)\b", "", token).strip(" -–_")
        if token:
            return token
    return None


# light, disclosed alias clean-up so the same source is not split in two
ALIASES = {
    "\u3000\u3000BL": "BL",
    "BL ": "BL",
    "Ro": "Romania",
    "RO": "Romania",
    "Isuzu Rom Auto": "Romania",
    "Isuzu RO": "Romania",
    "Poland": "Isuzu Polska",
    "UPPF_RT": "Iceland (UPPF_RT)",
    "IM Nordic": "Isuzu Nordic (IM Nordic)",
    "FINLAND & BALTICS": "Finland & Baltics",
    "Finland& Baltics": "Finland & Baltics",
    "Loutsios ": "Loutsios",
    "Isuzu Ireland ": "Isuzu Ireland",
    "Isuzu Polska ": "Isuzu Polska",
}

# Explicit, verified file -> distributor overrides, used only when the sheet
# itself carries no distributor name. Each was confirmed from the workbook or
# from a byte-identical twin (see report):
#   CharlotteMarston_*.xlsx            == Tomoe..._26_ITUK.xlsx (same md5)
#   AaronSingh-Chutti_*IUK*.xlsx       = IUK (distributor's own file)
#   BavoJacobs_*Isuzu OE*.xlsx         = Isuzu OE (Benelux/NL/Poland shared stock)
#   NunoMachado_*Portugal*.xlsx        = Portugal
#   Sigur*UPPF_RT*.xlsx                = Iceland (UPPF_RT)
#   ion.david_isuzu-romani_*.xlsx      = Romania
FILE_DIST = [
    (re.compile(r"charlottemarston", re.I), "ITUK"),
    (re.compile(r"aaron", re.I), "IUK"),
    (re.compile(r"bavojacobs", re.I), "Isuzu OE"),
    (re.compile(r"nunomachado", re.I), "Portugal"),
    (re.compile(r"upPf.rt|sigur", re.I), "Iceland (UPPF_RT)"),
    (re.compile(r"isuzu-romani|ION\.DAVID", re.I), "Romania"),
]

# Known distributor tokens carried in the file name. Used only as a last resort
# when the sheet itself (C2 and C1) has no name; reported as such.
NAME_TOKENS = [
    ("IBX", "IBX"), ("ITUK", "ITUK"), ("IUK", "IUK"), ("KIMA", "KIMA"),
    ("UTI", "UTI"), ("UMI", "UMI"), ("RSA", "RSA"), ("MIDI", "MIDI"),
    ("AIOS", "AIOS"), ("ISD", "ISD"), ("AWD", "AWD"),
    ("Nordic", "Isuzu Nordic (IM Nordic)"), ("Denmark", "Denmark"),
    ("Ireland", "Isuzu Ireland"), ("Poland", "Isuzu Polska"),
    ("Sweden", "Isuzu Sverige"), ("Hungary", "Hungary"),
    ("Romania", "Romania"), ("Greece", "Greece"), ("Malta", "Malta"),
    ("Cyprus", "Cyprus"), ("Portugal", "Portugal"),
    ("Finland", "Finland & Baltics"), ("Loutsios", "Loutsios"),
]


def dist_from_file(name: str) -> str | None:
    for rx, dist in FILE_DIST:
        if rx.search(name):
            return dist
    for token, dist in NAME_TOKENS:
        if re.search(r"(?<![A-Za-z])" + re.escape(token) + r"(?![A-Za-z])", name, re.I):
            return dist
    # trailing tokens with no word boundary after (e.g. "...Sheet26_BL.xlsx")
    for token, dist in (("BL", "BL"), ("IUK", "IUK")):
        if re.search(re.escape(token), name, re.I):
            return dist
    return None


def distributors_in_sheet(ws, title_dist, file_dist, date_cols) -> list[str]:
    """Names actually present in the sheet (col C2 line 2), or file fallback."""
    found = title_names(ws)
    if not found and file_dist:
        found.append(file_dist)
    if not found:
        found.append("(unstated)")
    return found


def is_blank_sheet(ws, date_cols, country_col) -> bool:
    """True if no KPI row in this sheet has any month value at all."""
    for r in range(1, ws.max_row + 1):
        if canonical_metric(cell_text(ws.cell(r, 3).value)) is None:
            continue
        for _, c in date_cols:
            v = cell_text(ws.cell(r, c).value)
            if v:
                return False
    return True


def main() -> int:
    rows: list[dict] = []
    layouts: dict[str, list[str]] = {}
    skipped: list[str] = []
    empty_sheets: list[str] = []

    for path in sorted(SRC.glob("*.xlsx")):
        name = path.name
        if "Fill Rate details" in name or "Graph Fill Rate" in name:
            skipped.append(f"{name} (fill-rate detail workbook, not a check sheet)")
            continue
        try:
            wb = load_wb(path)
        except Exception as exc:
            skipped.append(f"{name} (unreadable: {exc})")
            continue

        sheets = check_sheets(wb)
        if not sheets:
            skipped.append(f"{name} (no check-sheet sheet)")
            continue

        src_date = re.search(r"_(\d{8})_", name)
        report_src = src_date.group(1) if src_date else ""
        file_dist = distributor_from_filename(name)
        if file_dist:
            file_dist = ALIASES.get(file_dist, file_dist)

        for sn, ws in sheets:
            hdr, date_cols, country_col = header_info(ws)
            if not date_cols:
                skipped.append(f"{name} [{sn}] (no month header row)")
                continue
            if is_blank_sheet(ws, date_cols, country_col):
                empty_sheets.append(f"{name} [{sn}]")
                continue
            year = year_of(ws)
            if year is None:
                ym = re.search(r"(20\d\d)", name)
                year = int(ym.group(1)) if ym else None
            layout = "country" if country_col else "flat"
            layouts.setdefault(layout, []).append(f"{name} [{sn}]")

            # in-sheet title wins; C1 tail, explicit file map, then filename
            names = title_names(ws)
            if not names:
                names = tail_names(ws)
            if not names:
                fd = dist_from_file(name)
                if fd:
                    names = [fd]
            if not names and file_dist:
                names = [file_dist]
            if not names:
                names = ["(unstated)"]

            for r in range(hdr + 1, ws.max_row + 1):
                label = cell_text(ws.cell(r, 3).value)
                if not label:
                    continue
                cm = canonical_metric(label)
                if cm is None:
                    continue
                metric, unit = cm
                country = cell_text(ws.cell(r, country_col).value) if country_col else None

                if country_col:
                    dist = " / ".join(names[:1] + [country]) if country else names[0]
                else:
                    dist = names[0]

                for mnum, c in date_cols:
                    val = cell_text(ws.cell(r, c).value)
                    if val is None or val == "":
                        continue
                    period = f"{year}-{mnum:02d}" if year else f"??-{mnum:02d}"
                    rows.append({
                        "period": period,
                        "report_month_source": report_src,
                        "distributor": dist,
                        "metric": metric,
                        "value": val,
                        "unit": unit,
                        "sheet_name": sn,
                        "source_file": name,
                    })

    cols = ["period", "report_month_source", "distributor", "metric",
            "value", "unit", "sheet_name", "source_file"]
    with (OUT / "kpi_panel.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    meta = {
        "rows": len(rows),
        "files_parsed": sum(len(v) for v in layouts.values()),
        "layouts": {k: {"count": len(v), "files": v} for k, v in layouts.items()},
        "skipped": skipped,
        "empty_sheets": empty_sheets,
        "distributors": sorted({r["distributor"] for r in rows}),
        "periods": sorted({r["period"] for r in rows}),
    }
    (OUT / "kpi_panel_meta.json").write_text(
        json.dumps(meta, indent=1, ensure_ascii=False), encoding="utf-8")

    print(f"rows={len(rows)}  files_parsed={meta['files_parsed']}  skipped={len(skipped)}")
    for k, v in layouts.items():
        print(f"layout {k}: {len(v)} files/sheets")
    print("empty sheets (blank):", empty_sheets)
    print("skipped:", skipped)
    print("distributors:", meta["distributors"])
    print("periods:", meta["periods"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
