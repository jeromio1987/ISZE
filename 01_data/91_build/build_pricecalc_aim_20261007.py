"""Build AIM-aligned PriceCalc workbook from v3 + live config.yaml.

Outputs:
  01_data/01_Pricing/ISZE_PriceCalc_AIM_20261007.xlsx
  01_data/01_Pricing/aim_category_target_gm_20261007.csv
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pandas as pd
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "91_legacy/Isuzu/Divers/ISZE_PriceCalc_2026_v3.xlsx"
OUT_DIR = ROOT / "01_data/01_Pricing"
OUT_XLSX = OUT_DIR / "ISZE_PriceCalc_AIM_20261007.xlsx"
OUT_CSV = OUT_DIR / "aim_category_target_gm_20261007.csv"
CFG = ROOT / "03_analytics/isze_pricing/config.yaml"
AIM = ROOT / "Pricing revision/AIM_Master_PN_Categories_Figkey_20260921.csv"
LIST = ROOT / "01_data/00_derived/pricing_list_pn.parquet"

ROLE = {
    "Maintenance Service": "wear_pms",
    "Lubrication": "wear_pms",
    "Fluids Chemicals": "wear_pms",
    "Brakes": "wear_pms",
    "Clutch": "wear_pms",
    "Cooling": "wear_pms",
    "Wheels Hub": "wear_pms",
    "Intake Exhaust Emission": "captive",
    "Fuel System": "captive",
    "Electrical": "captive",
    "Engine": "captive",
    "Transmission": "captive",
    "Driveline Axle": "commercial",
    "Steering": "commercial",
    "Suspension": "commercial",
    "Lighting": "commercial",
    "Cabin Interior": "commercial",
    "Chassis Frame": "body_legal",
    "Accident & Body": "body_legal",
    "Literature Marketing": "unclassified",
    "Logistics Packaging": "unclassified",
    "Tools Equipment": "unclassified",
    "Unknown": "unclassified",
    "Unknown Review": "unclassified",
}

ROLE_RULE = {
    "wear_pms": "Price below class-implied to protect dealer habit (Spec launch wear).",
    "captive": "Price to margin, at or above class-implied (Spec launch captive).",
    "commercial": "Hold near class; competition can match (Dial COMMERCIAL).",
    "body_legal": "Hold near class; crash/legal (Dial BODY / LEGAL).",
    "unclassified": "No AIM family — keep class letter until classified.",
}

# ASSUMPTION until Jerome locks targets (percentage points vs billed book class GM)
PP_ADJ = {
    "wear_pms": -5.0,
    "captive": +5.0,
    "commercial": 0.0,
    "body_legal": 0.0,
    "unclassified": 0.0,
}


def main() -> None:
    with open(CFG, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    fx = float(cfg["formula"]["fx_rate"])
    src_j = float(cfg["formula"]["source_discount"]["J"])
    src_t = float(cfg["formula"]["source_discount"]["T"])
    class_factor = {k: float(v) for k, v in cfg["formula"]["class_factor"].items()}
    cis_factor = float(cfg["markets"]["CIS"]["factor_flat"])
    cis_overlap = float(cfg["markets"]["CIS"]["qingling_overlap_gm"])
    cis_no = float(cfg["markets"]["CIS"]["no_overlap_gm"])

    aim = pd.read_csv(
        AIM,
        usecols=["part_number", "aim_level_1", "l12m_billed_eur"],
        dtype={"part_number": str},
        low_memory=False,
    )
    lst = pd.read_parquet(LIST, columns=["part_number", "price_class", "gm"])
    lst["part_number"] = lst["part_number"].astype(str)
    aim["pn"] = aim["part_number"].astype(str).str.replace(r"\.0$", "", regex=True)
    lst["pn"] = lst["part_number"].astype(str).str.replace(r"\.0$", "", regex=True)

    def class_gm(letter):
        if letter is None or (isinstance(letter, float) and pd.isna(letter)):
            return None
        s = str(letter).strip().upper()
        for ch in s:
            if ch in class_factor:
                return 1.0 - class_factor[ch]
        return None

    lst["gm_class"] = lst["price_class"].map(class_gm)
    lst["gm_use"] = lst["gm"].where(lst["gm"].notna(), lst["gm_class"])

    j = aim.merge(lst[["pn", "price_class", "gm_use"]], on="pn", how="left")
    j["w"] = j["l12m_billed_eur"].fillna(0).clip(lower=0)
    j = j[j["aim_level_1"].notna() & j["gm_use"].notna()]

    rows = []
    for l1, g in j.groupby("aim_level_1"):
        w = g["w"].sum()
        book_gm = float((g["gm_use"] * g["w"]).sum() / w) if w > 0 else float(g["gm_use"].median())
        role = ROLE.get(l1, "unclassified")
        adj = PP_ADJ[role]
        proposed = max(0.30, min(0.65, book_gm + adj / 100.0))
        rows.append(
            {
                "aim_level_1": l1,
                "pricing_role": role,
                "role_rule": ROLE_RULE[role],
                "book_gm_weighted": round(book_gm, 4),
                "book_gm_pct": round(book_gm * 100, 1),
                "pp_vs_book": adj,
                "target_gm": round(proposed, 4),
                "target_gm_pct": round(proposed * 100, 1),
                "class_factor_for_sap": round(1.0 - proposed, 4),
                "billed_eur_weight": round(float(w), 2),
                "n_pn_with_class": int(g["pn"].nunique()),
                "use_mode": "CHECK",
                "status_note": (
                    "CHECK against class until Jerome confirms ISZE can set "
                    "new-PN margin by AIM category (Spec Q1)."
                ),
                "assumption": (
                    "target_gm = book_gm + role pp (wear -5, captive +5, else 0), "
                    "clipped 30–65%. Not locked."
                ),
                "as_of": "2026-10-07",
                "config_fx": fx,
                "config_source_j": src_j,
                "config_source_t": src_t,
            }
        )

    seen = {r["aim_level_1"] for r in rows}
    for l1 in sorted(aim["aim_level_1"].dropna().unique()):
        if l1 in seen:
            continue
        role = ROLE.get(l1, "unclassified")
        book_gm = 0.53
        adj = PP_ADJ[role]
        proposed = max(0.30, min(0.65, book_gm + adj / 100.0))
        rows.append(
            {
                "aim_level_1": l1,
                "pricing_role": role,
                "role_rule": ROLE_RULE[role],
                "book_gm_weighted": None,
                "book_gm_pct": None,
                "pp_vs_book": adj,
                "target_gm": round(proposed, 4),
                "target_gm_pct": round(proposed * 100, 1),
                "class_factor_for_sap": round(1.0 - proposed, 4),
                "billed_eur_weight": 0.0,
                "n_pn_with_class": 0,
                "use_mode": "CHECK",
                "status_note": "No billed×class join — defaulted book to class C 53% GM. CHECK only.",
                "assumption": "target_gm = book_gm + role pp, clipped 30–65%. Not locked.",
                "as_of": "2026-10-07",
                "config_fx": fx,
                "config_source_j": src_j,
                "config_source_t": src_t,
            }
        )

    lookup = (
        pd.DataFrame(rows)
        .sort_values(["pricing_role", "aim_level_1"])
        .reset_index(drop=True)
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    lookup.to_csv(OUT_CSV, index=False)

    if not SRC.exists():
        raise SystemExit(f"Missing source workbook: {SRC}")

    shutil.copy2(SRC, OUT_XLSX)
    wb = openpyxl.load_workbook(OUT_XLSX)
    ws = wb["Config"]

    ws["A1"] = "ISZE — PRICE CALCULATION — CONFIG (AIM-aligned 2026-10-07)"
    ws["A4"] = "FX Rate JPY/EUR (live)"
    ws["B4"] = fx
    ws["A5"] = "Source J Discount (Japan/IML)"
    ws["B5"] = src_j
    ws["A6"] = "Source T Discount (Thailand/IMIT)"
    ws["B6"] = src_t
    ws["A7"] = "CIS Factor (GA + CBC)"
    ws["B7"] = cis_factor
    ws["B8"] = "01.10.2026"
    ws["B9"] = "31.12.2999"

    letter_rows = {}
    for r in range(12, 22):
        letter = ws.cell(r, 1).value
        if letter:
            letter_rows[str(letter).strip().upper()] = r

    for letter, factor in class_factor.items():
        if letter in letter_rows:
            ws.cell(letter_rows[letter], 2).value = factor

    ws["A23"] = (
        "Blue cells = inputs. Black = formulas. Rate authority: "
        "03_analytics/isze_pricing/config.yaml — refresh this sheet when that file moves."
    )
    ws["A24"] = (
        f"Aligned 2026-10-07 from config.yaml: FX {fx}, J {src_j}, T {src_t}, "
        f"class T factor {class_factor['T']} (40% GM). CIS flat {cis_factor} (60% GM)."
    )
    ws["A25"] = (
        "AIM-led margin: use sheet 'AIM Margin Lookup' as CHECK against class "
        "until Spec Q1 is answered. SAP sheets stay here."
    )
    ws["A26"] = (
        f"CIS overlap targets in config (not auto-applied): Qingling-overlap "
        f"{cis_overlap:.0%} GM; no-overlap {cis_no:.0%} GM."
    )

    if "AIM Margin Lookup" in wb.sheetnames:
        del wb["AIM Margin Lookup"]
    wsL = wb.create_sheet("AIM Margin Lookup", 1)

    headers = [
        "aim_level_1",
        "pricing_role",
        "role_rule",
        "book_gm_pct",
        "pp_vs_book",
        "target_gm_pct",
        "class_factor_for_sap",
        "use_mode",
        "billed_eur_weight",
        "n_pn_with_class",
        "status_note",
    ]
    header_fill = PatternFill("solid", fgColor="15233C")
    header_font = Font(color="F7F4EE", bold=True)
    for i, h in enumerate(headers, 1):
        cell = wsL.cell(1, i, h)
        cell.fill = header_fill
        cell.font = header_font

    for i, row in enumerate(lookup.itertuples(index=False), 2):
        wsL.cell(i, 1, row.aim_level_1)
        wsL.cell(i, 2, row.pricing_role)
        wsL.cell(i, 3, row.role_rule)
        wsL.cell(i, 4, row.book_gm_pct)
        wsL.cell(i, 5, row.pp_vs_book)
        wsL.cell(i, 6, row.target_gm_pct)
        wsL.cell(i, 7, row.class_factor_for_sap)
        wsL.cell(i, 8, row.use_mode)
        wsL.cell(i, 9, row.billed_eur_weight)
        wsL.cell(i, 10, row.n_pn_with_class)
        wsL.cell(i, 11, row.status_note)

    last = len(lookup) + 1
    wsL["A28"] = "How to use"
    wsL["A29"] = "1) Find the part's AIM level-1 family (from AIM Master table)."
    wsL["A30"] = "2) Read target_gm_pct and class_factor_for_sap."
    wsL["A31"] = (
        "3) Until Spec Q1 = yes: keep Pricing Class from IML (col P on New Parts); "
        "compare GM% to target — this is CHECK mode."
    )
    wsL["A32"] = (
        "4) When Spec Q1 = yes: use class_factor_for_sap as the PR00 divisor for that line."
    )
    wsL["A33"] = (
        "5) Supersessions: inherit predecessor AIM family; do not re-pick role. "
        "New PNs: assign AIM family first, then read this table."
    )
    wsL["A34"] = f"CSV twin: {OUT_CSV}"
    wsL["A35"] = (
        "ASSUMPTION: wear −5pp / captive +5pp vs billed book class GM. "
        "Not a locked policy — Jerome sets final targets."
    )

    for col in range(1, 12):
        wsL.column_dimensions[get_column_letter(col)].width = (
            48 if col in (3, 11) else 18
        )

    if "Ops Rules" in wb.sheetnames:
        del wb["Ops Rules"]
    wsO = wb.create_sheet("Ops Rules", 2)
    ops = [
        ("ISZE PriceCalc — AIM-led operating rules", ""),
        ("As of", "2026-10-07"),
        ("", ""),
        (
            "What this file is for",
            "Day-to-day: paste IML line → FOB € → sales price € → SAP Fast Entry "
            "(PR00 / ISR 1797 / ISR 1799 / CIS).",
        ),
        (
            "What the revision pack is for",
            "Portfolio decisions: base-price step, FX, stack, AIM-by-source. "
            "It does not mint SAP upload rows.",
        ),
        ("", ""),
        (
            "Rate authority",
            "03_analytics/isze_pricing/config.yaml. Refresh Config when FX or "
            "discounts move. Do not keep a second living rate book.",
        ),
        (
            "Live dials on Config",
            f"FX {fx} · Source J {src_j} · Source T {src_t} · class T factor "
            f"{class_factor['T']} (40% GM) · CIS {cis_factor} (60% GM)",
        ),
        ("", ""),
        (
            "Supersession",
            "Old PN → newest PN: inherit predecessor AIM category. Keep price "
            "continuity unless the yen BP step forces a move. Newest Part No. "
            "drives Material for SAP. Do not re-open the category debate.",
        ),
        (
            "Truly new PN",
            "Assign AIM level-1 first. Read AIM Margin Lookup for target %GM. "
            "FOB still = BP_JPY / FX × source discount. Only the margin step is "
            "AIM-led once Spec Q1 is yes.",
        ),
        (
            "Until Spec Q1 is yes",
            "SAP still follows the IML class letter. Use AIM target as a CHECK "
            "only (class implies X% GM; AIM family wanted Y%).",
        ),
        (
            "Israel / CIS",
            "Still use the SAP Output sheets in this workbook. Revision pack can "
            "advise CIS margin policy; this file builds the paste rows.",
        ),
        ("", ""),
        (
            "Do not",
            "Retire this file because AIM Source HTML exists. Do not price new "
            "PNs from the 47k-line simulation workbook. Do not leave Config on "
            "163.7 / 0.62.",
        ),
    ]
    for i, (a, b) in enumerate(ops, 1):
        wsO.cell(i, 1, a)
        wsO.cell(i, 2, b)
        if b == "" and a:
            wsO.cell(i, 1).font = Font(bold=True, color="15233C")
    wsO.column_dimensions["A"].width = 28
    wsO.column_dimensions["B"].width = 110

    wsN = wb["➕ New Parts"]
    wsN.cell(3, 41, "AIM L1\n(paste)")
    wsN.cell(3, 42, "AIM target\n%GM (lookup)")
    wsN.cell(3, 43, "AIM check\n(class GM − target)")
    wsN.cell(3, 44, "Mode")
    wsN.cell(
        1,
        41,
        "AIM helpers → paste L1 in col AO; target/% check fill from AIM Margin Lookup "
        "(CHECK mode until Spec Q1)",
    )

    n_lookup = last
    for r in range(5, 220):
        wsN.cell(
            r,
            42,
            f'=IF(AO{r}="","",IFERROR(VLOOKUP(AO{r},\'AIM Margin Lookup\'!$A$2:$G${n_lookup},6,FALSE),""))',
        )
        wsN.cell(
            r,
            43,
            f'=IF(OR(T{r}="",AP{r}=""),"",ROUND(T{r}*100-AP{r},1))',
        )
        wsN.cell(
            r,
            44,
            f'=IF(AO{r}="","",IFERROR(VLOOKUP(AO{r},\'AIM Margin Lookup\'!$A$2:$H${n_lookup},8,FALSE),"CHECK"))',
        )

    wb.save(OUT_XLSX)
    print("WROTE", OUT_XLSX)
    print("WROTE", OUT_CSV)
    print("lookup_rows", len(lookup))
    print(
        lookup[
            ["aim_level_1", "pricing_role", "book_gm_pct", "target_gm_pct", "use_mode"]
        ].to_string(index=False)
    )


if __name__ == "__main__":
    main()
