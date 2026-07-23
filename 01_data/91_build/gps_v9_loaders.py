"""Load UIO v19 (canonical) + ITUK DVLA car park + gps_results CY2025 for GPS v9/v6."""
from __future__ import annotations

import json
import re
import shutil
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
DATA = ROOT / "01_data"
BUILD = DATA / "91_build"
UIO_CURRENT = ROOT / "06_uio" / "_CURRENT" / "ISUZU_UIO_Overview_2026.xlsx"
UIO_FALLBACKS = [
    DATA / "ISUZU_UIO_Overview_2026_v19.xlsx",
    ROOT / "06_uio" / "_Archive" / "ISUZU_UIO_Overview_2026_v19.xlsx",
]
CARPARK_XLSX = ROOT / "91_legacy" / "Isuzu" / "UIO" / "Script" / "Parts Sales Calculation FY26.xlsx"
GPS_RESULTS = BUILD / "gps_results.json"
OUT_JSON_V19 = BUILD / "uio_v19_gps.json"
OUT_JSON_LEGACY = BUILD / "uio_v18_gps.json"  # alias copy for older scripts until retired

UIO_VERSION = 19

RATE_N35 = 296.0
RATE_N75 = 739.0
RATE_FMD = 887.0
RATE_LCV = 250.0

ITUK_GP = 0.34
ITUK_FX = 1.12

BP_ITUK = {
    2026: {"retail_gbp": 13_129_200, "isze_eur": 6_974_170.05},
    2027: {"retail_gbp": 13_129_200, "isze_eur": 6_974_170.05},
}
KOSHA_VP_EUR = {2027: 35_284_400}

DIST_MAP = {
    ("BELGIUM", "IBX"): "BELGIUM",
    ("CYPRUS", "Loutsios & Sons"): "CYPRUS",
    ("DENMARK", "Isuzu Denmark"): "DENMARK",
    ("FINLAND", "Nordic Auto."): "FINLAND",
    ("GERMANY", "ISD-Lang"): "GERMANY",
    ("GREECE", "Petros Petropoulos"): "GREECE",
    ("HUNGARY", "AutoWallis"): "HUNGARY",
    ("ICELAND", "BL"): "ICELAND",
    ("IRELAND", "Isuzu Ireland"): "IRELAND",
    ("ISRAEL", "UMI"): "UMI",
    ("ISRAEL", "UTI"): "UTI",
    ("MALTA", "GasanZammit"): "MALTA",
    ("NORWAY", "RSA"): "NORWAY",
    ("PORTUGL", "Imotors"): "PORTUGAL",
    ("ROMANIA", "Isuzu Rom Auto"): "ROMANIA",
    ("SWEDEN", "Isuzu Sverige"): "SWEDEN",
    ("U KING", "IUK"): "IUK",
    ("U KING", "ITUK"): "ITUK",
}

GPS_DBS = [
    "ITUK", "IUK", "UMI", "UTI", "GERMANY", "BELGIUM", "IRELAND", "NORWAY",
    "SWEDEN", "DENMARK", "FINLAND", "ICELAND", "CYPRUS", "GREECE", "MALTA",
    "ROMANIA", "HUNGARY", "PORTUGAL",
]

SCOPE = {
    "IUK": ("LCV",),
    "UMI": ("LCV",),
    "UTI": ("LD", "FMD"),
    "NORWAY": ("LCV",),
    "SWEDEN": ("LCV",),
    "DENMARK": ("LCV",),
    "FINLAND": ("LCV",),
    "ICELAND": ("LCV",),
    "HUNGARY": ("LCV",),
    "ROMANIA": ("LCV",),
}
DEFAULT_SCOPE = ("LD", "FMD", "LCV")


def resolve_uio_workbook() -> Path:
    for p in [UIO_CURRENT, *UIO_FALLBACKS]:
        if p.is_file():
            return p
    raise FileNotFoundError(
        f"No UIO workbook found. Expected canonical: {UIO_CURRENT}"
    )


def _bucket(mtype: str) -> str:
    mt = str(mtype or "").upper()
    if "LCV" in mt or "D-MAX" in mt or "DMAX" in mt:
        return "LCV"
    if mt in ("MD", "F-SERIES", "F SERIES") or mt == "F":
        return "FMD"
    return "LD"


def _scope_buckets(db: str) -> tuple[str, ...]:
    if db == "ITUK":
        return ("N35", "N75_LD", "F_MD")
    return SCOPE.get(db, DEFAULT_SCOPE)


def load_ituk_carpark() -> dict[str, int]:
    wb = openpyxl.load_workbook(CARPARK_XLSX, read_only=True, data_only=True)
    ws = wb["UIO 24"]
    totals: dict[str, int] = defaultdict(int)
    for row in ws.iter_rows(min_row=2, values_only=True):
        short = str(row[3] or "").strip()
        n = int(row[4] or 0)
        if short and n > 0:
            totals[short] += n
    wb.close()
    return dict(totals)


def load_ituk_projection_models(uio_path: Path) -> dict[str, list[float]]:
    wb = openpyxl.load_workbook(uio_path, read_only=True, data_only=True)
    ws = wb["4. ITUK UIO Projection"]
    models: dict[str, list[float]] = {}
    for row in ws.iter_rows(min_row=20, max_row=35, values_only=True):
        name = str(row[1] or "").strip()
        if not name or name.startswith("Total") or name == "(":
            continue
        try:
            vals = [float(row[i] or 0) for i in range(2, 8)]  # 2025–2030
        except (TypeError, ValueError):
            continue
        if any(vals):
            models[name] = vals
    wb.close()
    if not models:
        raise ValueError(f"No ITUK model rows in {uio_path} sheet '4. ITUK UIO Projection'")
    return models


def _ituk_buckets_from_models(models: dict[str, list[float]], idx: int) -> dict[str, float]:
    return {
        "N35": models.get("N35", [0, 0, 0])[idx],
        "N75_LD": (
            models.get("N75", [0, 0, 0])[idx]
            + models.get("N50/55", [0, 0, 0])[idx]
            + models.get("N62/65", [0, 0, 0])[idx]
            + models.get("UNSPEC", [0, 0, 0])[idx]
        ),
        "F_MD": (
            models.get("F110/120", [0, 0, 0])[idx]
            + models.get("F135", [0, 0, 0])[idx]
            + models.get("F150/180", [0, 0, 0])[idx]
        ),
    }


def _ituk_buckets_from_carpark(cp: dict[str, int]) -> dict[str, float]:
    return {
        "N35": float(cp.get("N35", 0)),
        "N75_LD": float(
            cp.get("N75", 0)
            + cp.get("N50/55", 0)
            + cp.get("N62/65", 0)
            + cp.get("UNSPECIFIED", 0)
        ),
        "F_MD": float(
            cp.get("F110/120", 0) + cp.get("F135", 0) + cp.get("F150/180", 0)
        ),
    }


def load_network_uio(uio_path: Path) -> dict[str, dict[int, dict[str, float]]]:
    wb = openpyxl.load_workbook(uio_path, read_only=True, data_only=True)
    ws = wb["1.UIO Overview"]
    all_rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(all_rows) if r[0] is not None and "Ctry" in str(r[0]))
    hdr = all_rows[hdr_i]
    ycols = {
        int(m.group(1)): j
        for j, h in enumerate(hdr)
        if h and (m := re.match(r"^\s*(\d{4})[EF]?\s*$", str(h)))
        and 2024 <= int(m.group(1)) <= 2030
    }
    agg: dict[str, dict[int, dict[str, float]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(float))
    )
    for r in all_rows[hdr_i + 1 :]:
        if r[1] is None or r[5] is None:
            continue
        db = DIST_MAP.get((str(r[1]).strip(), str(r[5]).strip()))
        if not db or db == "ITUK":
            continue
        b = _bucket(r[3])
        for y, j in ycols.items():
            if isinstance(r[j], (int, float)):
                agg[db][y][b] += float(r[j])
    wb.close()
    return {k: {yy: dict(buckets) for yy, buckets in v.items()} for k, v in agg.items()}


# Backward-compatible alias
load_network_uio_v18 = load_network_uio


def load_gps_results_2025() -> dict[str, dict[str, float]]:
    rows = json.load(open(GPS_RESULTS, encoding="utf-8-sig"))
    out: dict[str, dict[str, float]] = {}
    for r in rows:
        if r.get("year") != 2025:
            continue
        db = r["db"]
        if db not in GPS_DBS or db == "ITUK":
            continue
        out[db] = {
            "LD": float(r.get("uio_cv_ld") or 0),
            "FMD": float((r.get("uio_f") or 0) + (r.get("uio_md") or 0)),
            "LCV": float(r.get("uio_lcv") or 0),
        }
    return out


def _scoped_total(buckets: dict[str, float], db: str) -> float:
    scope = _scope_buckets(db)
    if db == "ITUK":
        return sum(buckets.get(k, 0) for k in scope)
    return sum(buckets.get(k, 0) for k in scope if k in buckets)


def max_spend_from_buckets(buckets: dict[str, float], db: str) -> float:
    if db == "ITUK":
        return (
            buckets.get("N35", 0) * RATE_N35
            + buckets.get("N75_LD", 0) * RATE_N75
            + buckets.get("F_MD", 0) * RATE_FMD
        )
    return (
        buckets.get("LD", 0) * RATE_N75
        + buckets.get("FMD", 0) * RATE_FMD
        + buckets.get("LCV", 0) * RATE_LCV
    )


def _scale_buckets(base: dict[str, float], factor: float, db: str) -> dict[str, float]:
    scope = _scope_buckets(db)
    return {k: round(v * factor, 1) if k in scope else 0.0 for k, v in base.items()}


def build_gps_uio_package() -> dict:
    uio_path = resolve_uio_workbook()
    carpark = load_ituk_carpark()
    ituk_models = load_ituk_projection_models(uio_path)
    network_uio = load_network_uio(uio_path)
    gps25 = load_gps_results_2025()

    ituk_uio: dict[int, dict] = {
        2025: _ituk_buckets_from_carpark(carpark),
        2026: _ituk_buckets_from_models(ituk_models, 1),
        2027: _ituk_buckets_from_models(ituk_models, 2),
        2028: _ituk_buckets_from_models(ituk_models, 3),
        2029: _ituk_buckets_from_models(ituk_models, 4),
        2030: _ituk_buckets_from_models(ituk_models, 5),
    }

    distributors: dict[str, dict] = {
        "ITUK": {
            "scope": f"CV — DVLA car park CY2025 · UIO v{UIO_VERSION} model 2026–27",
            "uio": ituk_uio,
            "max_spend": {y: max_spend_from_buckets(b, "ITUK") for y, b in ituk_uio.items()},
            "rates_note": "N35 €296 · N75/LD €739 · F/MD €887",
        }
    }

    for db in GPS_DBS:
        if db == "ITUK":
            continue
        base = gps25[db]
        uio_25 = network_uio.get(db, {}).get(2025, {})
        uio_26 = network_uio.get(db, {}).get(2026, {})
        uio_27 = network_uio.get(db, {}).get(2027, {})

        def _growth(target: dict[str, float]) -> float:
            t25 = _scoped_total(uio_25, db)
            t_tgt = sum(target.get(k, 0) for k in _scope_buckets(db) if k in target)
            return (t_tgt / t25) if t25 > 0 else 1.0

        uio_by_year = {
            2025: {k: int(round(v)) for k, v in base.items()},
            2026: {
                k: int(round(v))
                for k, v in _scale_buckets(base, _growth(uio_26), db).items()
            },
            2027: {
                k: int(round(v))
                for k, v in _scale_buckets(base, _growth(uio_27), db).items()
            },
        }
        distributors[db] = {
            "scope": (
                "LCV only"
                if SCOPE.get(db) == ("LCV",)
                else ("CV only" if db == "UTI" else "CV + LCV")
            ),
            "uio": uio_by_year,
            "max_spend": {
                y: max_spend_from_buckets(
                    {k: float(v) for k, v in b.items()}, db
                )
                for y, b in uio_by_year.items()
            },
            "rates_note": "LD €739 · F/MD €887 · LCV €250",
        }

    return {
        "uio_version": UIO_VERSION,
        "source_uio": str(uio_path),
        "source_uio_cy2025": str(GPS_RESULTS),
        "source_ituk_carpark": str(CARPARK_XLSX),
        "rates": {
            "N35": RATE_N35,
            "N75_LD": RATE_N75,
            "F_MD": RATE_FMD,
            "LCV": RATE_LCV,
            "gp": ITUK_GP,
            "fx": ITUK_FX,
        },
        "ituk_carpark_raw": carpark,
        "ituk_models": ituk_models,
        "ituk_models_v18": ituk_models,  # legacy key for dashboards reading old field name
        "distributors": distributors,
        "bp_ituk": BP_ITUK,
        "kosha_vp_eur": KOSHA_VP_EUR,
    }


def write_uio_gps_json(pkg: dict | None = None) -> Path:
    pkg = pkg or build_gps_uio_package()
    text = json.dumps(pkg, indent=2)
    OUT_JSON_V19.write_text(text, encoding="utf-8")
    shutil.copy2(OUT_JSON_V19, OUT_JSON_LEGACY)
    return OUT_JSON_V19


def main() -> None:
    pkg = build_gps_uio_package()
    path = write_uio_gps_json(pkg)
    print("wrote", path, "(and legacy alias uio_v18_gps.json)")
    print("source", pkg["source_uio"])
    for db in ("ITUK", "UTI", "GERMANY"):
        d = pkg["distributors"][db]
        print(db, "max25", f"{d['max_spend'][2025]:,.0f}", "max27", f"{d['max_spend'][2027]:,.0f}")


if __name__ == "__main__":
    main()
