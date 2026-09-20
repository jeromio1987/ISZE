"""Verify the rebuilt workbook: formats, caches, checks, and the specific defects."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BOOK = ISZE / "01_data" / "01_Pricing" / "Price_increase_by_part_group_per_customer.xlsx"

wf = load_workbook(BOOK, data_only=False)
wv = load_workbook(BOOK, data_only=True)

print("SHEETS")
for i, n in enumerate(wf.sheetnames):
    print(f"  {i} {n}  (opens on: {wf.active.title})")
print()

# ---- 1. nan / #### strings anywhere -------------------------------------
print("TEXT DEFECTS")
bad = []
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                low = c.value.lower()
                if re.search(r"\bnan\b|nan%|\bnone\b|#REF|#####|inf\b|Traceback", low):
                    bad.append((ws.title, c.coordinate, c.value[:70]))
print(f"  nan/none/#-style text: {len(bad)}")
for b in bad[:10]:
    print("   ", b)
print()

# ---- 2. formulas and their cached values --------------------------------
print("FORMULA CACHE")
totals = []
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                cached = wv[ws.title][c.coordinate].value
                totals.append((ws.title, c.coordinate, c.value, cached))
print(f"  formula cells: {len(totals)}")
blank = [t for t in totals if t[3] is None]
print(f"  formula cells with no cached value: {len(blank)}")
for t in blank[:12]:
    print("   BLANK", t[:3])
print()

# ---- 3. the Checks sheet ------------------------------------------------
print("CHECKS SHEET")
ck = wf["Checks"]
ckv = wv["Checks"]
hdr = None
for r in range(1, ck.max_row + 1):
    v = ck.cell(r, 1).value
    if v == "Check":
        hdr = r
        break
print("  header row:", hdr)
rows = []
for r in range(hdr + 1, ck.max_row + 1):
    name = ck.cell(r, 1).value
    if not name:
        continue
    res = ckv.cell(r, 6).value
    rows.append((name, ck.cell(r, 2).value, ckv.cell(r, 2).value,
                 ckv.cell(r, 3).value, ckv.cell(r, 4).value, ckv.cell(r, 5).value, res))
for name, f, cv, ev, dv, tv, res in rows:
    flag = "  " if res == "OK" else "**"
    print(f"  {flag}{res!s:6s} {name[:56]:58s} here={cv} exp={ev} diff={dv} tol={tv}")
print()
notok = [r for r in rows if r[6] != "OK"]
print(f"  rows not reading OK: {len(notok)}")
print()

# ---- 4. formatting coverage --------------------------------------------
print("FORMATTING")
for ws in wf.worksheets:
    widths = {}
    for i in range(1, ws.max_column + 1):
        letter = get_column_letter(i)
        d = ws.column_dimensions.get(letter)
        widths[letter] = None if d is None or d.width is None else round(d.width, 2)
    default_like = [k for k, v in widths.items() if v is None or v == 8.43]
    fills = sum(1 for row in ws.iter_rows() for c in row
                if c.fill is not None and c.fill.fill_type == "solid")
    fmts = sorted({c.number_format for row in ws.iter_rows() for c in row
                   if c.number_format and c.number_format != "General"})
    print(f"  {ws.title:32s} cols={ws.max_column:3d} freeze={ws.freeze_panes!s:6s} "
          f"filter={ws.auto_filter.ref!s:22s}")
    print(f"      default-width cols: {len(default_like)} {default_like}")
    print(f"      solid fills: {fills} | cf rules: {len(list(ws.conditional_formatting))}"
          f" | tab={ws.sheet_properties.tabColor}")
    print(f"      non-General formats: {fmts}")
    print(f"      merges: {len(list(ws.merged_cells.ranges))}")
print()

# ---- 5. Overview figures ------------------------------------------------
print("OVERVIEW")
ov = wv["Overview"]
for r in range(1, 13):
    lab = ov.cell(r, 1).value
    val = ov.cell(r, 2).value
    if lab:
        print(f"  r{r}: {str(lab)[:60]:62s} {val}")
print()

# ---- 6. the acceptance numbers as stored --------------------------------
d = wv["Calculation detail"]
tot = None
for r in range(1, d.max_row + 1):
    if d.cell(r, 1).value == "Total":
        tot = r
print("detail total row", tot)
print("  step      ", d.cell(tot, 6).value)
print("  billed    ", d.cell(tot, 5).value)
print("  basket    ", d.cell(tot, 7).value)
print("  movement  ", d.cell(tot, 8).value)
print("  coverage  ", d.cell(tot, 9).value)

# ---- 7. row counts and the totals on each sheet -------------------------
for name, tcol in [("Part groups, network-wide", 6), ("Customers by part group", 7)]:
    ws = wf[name]
    tr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "Total":
            tr = r
    print(f"{name}: total row {tr}, value col {tcol} = {wv[name].cell(tr, tcol).value}")

# ---- 8. rank sheet: the Denmark case and the count column ---------------
rk = wv["Highest and lowest by customer"]
print()
print("RANK SHEET (customer, groups, high, low)")
for r in range(5, 5 + rk.max_row):
    cust = rk.cell(r, 1).value
    if not cust:
        break
    n = rk.cell(r, 6).value
    low = rk.cell(r, 8).value
    print(f"  {str(cust)[:28]:30s} n={n!s:4s} low={'BLANK' if not low else low[:44]}")
