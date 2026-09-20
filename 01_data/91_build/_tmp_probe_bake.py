"""Probe: does pywin32 drive Excel recalc reliably in this session?"""
from __future__ import annotations

import shutil
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

TMP = Path(__file__).resolve().parent / "_tmp_bake_probe.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "P"
ws["A1"] = 2
ws["A2"] = 3
ws["A3"] = "=SUM(A1:A2)"
ws["A4"] = "=A3*2"
wb.save(TMP)
print("wrote", TMP)

import win32com.client  # noqa: E402

app = None
t0 = time.time()
try:
    app = win32com.client.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    book = app.Workbooks.Open(str(TMP))
    app.CalculateFullRebuild()
    book.Save()
    book.Close(SaveChanges=False)
    print("com seconds", round(time.time() - t0, 1))
finally:
    if app is not None:
        app.Quit()

wb2 = load_workbook(TMP, data_only=True)
print("cached A3", wb2["P"]["A3"].value, "A4", wb2["P"]["A4"].value)
wb3 = load_workbook(TMP, data_only=False)
print("formula A3", wb3["P"]["A3"].value)
TMP.unlink(missing_ok=True)
