"""One-off: flatten IE30/40/50 contacts overview + compare to DB_email_22_DBs."""
from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
NEW_P = ROOT / r"01_data\08_Distributor\overview_IE30_IE40_IE50_contacts.xlsx"
OLD_P = ROOT / r"01_data\08_Distributor\DB_email_22_DBs.xlsx"
OUT_MD = ROOT / r"01_data\08_Distributor\CONTACTS_REGISTER.md"
OUT_CSV = ROOT / r"01_data\00_derived\distributor_daily_contacts_ie30_ie40_ie50.csv"

EMAIL_RE = re.compile(r"[\w.+'-]+@[\w.-]+\.[A-Za-z]{2,}", re.I)


def main() -> None:
    wb = openpyxl.load_workbook(NEW_P, read_only=True, data_only=True)
    rows_out: list[dict] = []
    emails_new: set[str] = set()
    customers: set[str] = set()
    rows = list(wb["All customers"].iter_rows(values_only=True))
    cur_cust = cur_branch = None
    for r in rows[1:]:
        vals = list(r) + [None] * 8
        if vals[0] not in (None, ""):
            cur_cust = vals[0]
        if vals[1] not in (None, ""):
            cur_branch = vals[1]
        cust, branch = cur_cust, cur_branch
        if cust:
            customers.add(str(cust).strip())
        for line, name, email in (
            ("IE40", vals[2], vals[3]),
            ("IE30", vals[4], vals[5]),
            ("IE50", vals[6], vals[7]),
        ):
            if not name and not email:
                continue
            em = str(email).strip() if email else ""
            nm = str(name).strip() if name else ""
            for m in EMAIL_RE.findall(em):
                emails_new.add(m.lower())
            rows_out.append(
                {
                    "customer": str(cust).strip() if cust else "",
                    "branch": str(branch).strip() if branch else "",
                    "product_line": line,
                    "contact_name": nm,
                    "email": em.lower() if em else "",
                }
            )

    wb2 = openpyxl.load_workbook(OLD_P, read_only=True, data_only=True)
    emails_old: set[str] = set()
    for r in list(wb2["22 DBs"].iter_rows(values_only=True))[1:]:
        cell = r[4] if len(r) > 4 else None
        if cell:
            emails_old.update(m.lower() for m in EMAIL_RE.findall(str(cell)))

    only_new = sorted(emails_new - emails_old)
    only_old = sorted(emails_old - emails_new)
    both = sorted(emails_new & emails_old)
    n_rows = sum(1 for r in rows_out if r["contact_name"] or r["email"])

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["customer", "branch", "product_line", "contact_name", "email"],
        )
        w.writeheader()
        for row in rows_out:
            if row["contact_name"] or row["email"]:
                w.writerow(row)

    lines = [
        "# Distributor contacts register",
        "",
        "**Canonical daily-contacts workbook (by product line):**",
        "`01_data/08_Distributor/overview_IE30_IE40_IE50_contacts.xlsx`",
        "",
        "**Email allow-list master (MAPI collector / domains):**",
        "`01_data/08_Distributor/DB_email_22_DBs.xlsx` → builds `04_tools/interactions/db_domains.yaml`",
        "",
        "**Derived flat list:**",
        "`01_data/00_derived/distributor_daily_contacts_ie30_ie40_ie50.csv`",
        "",
        "## Source (2026-09-03)",
        "",
        "Jerome self-sent from Sent Items (`jerome.van.der.pluym@isuzu-e.com` → himself),",
        "blank subject, attachment `overview IE30-IE40-IE50 contacts.xlsx`.",
        "Intake: `01_data/90_INBOX/contacts_self_sent_20260903/`.",
        "",
        "## What is in the workbook",
        "",
        "Sheets: `All customers`, `IE30`, `IE40`, `IE50`.",
        "Columns: customer, branch, daily contact name + email per IE30 / IE40 / IE50.",
        "",
        f"- Customer blocks on All customers: **{len(customers)}**",
        f"- Contact rows exported (name or email present): **{n_rows}**",
        f"- Unique emails in overview: **{len(emails_new)}**",
        f"- Unique emails in DB_email_22_DBs: **{len(emails_old)}**",
        f"- Overlap: **{len(both)}**",
        f"- In overview but not in DB_email master: **{len(only_new)}**",
        f"- In DB_email master but not in overview: **{len(only_old)}**",
        "",
        "## Practice (do not replace blindly)",
        "",
        "- Keep **DB_email_22_DBs.xlsx** as the domain / allow-list source for mail collection.",
        "- Keep **overview_IE30_IE40_IE50_contacts.xlsx** as the operational “who do I call per product line” register.",
        "- Enrich DB_email from overview only after a human pass (shared Astara domains etc.).",
        "",
        "### Emails in overview not yet in DB_email_22_DBs",
        "",
    ]
    if only_new:
        lines.extend(f"- `{e}`" for e in only_new)
    else:
        lines.append("- (none)")
    lines += ["", "### Customers (All customers)", ""]
    lines.extend(f"- {c}" for c in sorted(customers))
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT_MD}")
    print(f"wrote {OUT_CSV}")
    print(
        f"customers={len(customers)} emails_new={len(emails_new)} "
        f"only_new={len(only_new)} only_old={len(only_old)}"
    )
    for e in only_new:
        print("ONLY_NEW", e)


if __name__ == "__main__":
    main()
