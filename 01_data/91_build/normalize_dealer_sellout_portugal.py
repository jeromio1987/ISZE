"""Normalize Portugal (Imotors / Astara) sales-to-dealer XT -> canonical long CSV.

Native shape (two sheets, one line per invoice line):
  '2025'  CY2025 full         | Data da Fatura | Referencia | Sum of Quantidade
  '2026'  CY2026 YTD Aug      | Data da Fatura | Referencia | Sum of Quantidade

Rules:
  - distributor_code = PORTUGAL (Imotors LDA / Astara Portugal) - this is DB->dealer
    sell-in, NOT ISZE->DB billing. Never load into sales_fact.
  - PN = Referencia (Isuzu-style part numbers)
  - qty = Sum of Quantidade (skip zero lines)
  - period from Data da Fatura (YYYY-MM)
  - grain = network_pn (no dealer dimension in source)
  - qty_only (no EUR in source); 2026 sheet is partial_year

Outputs:
  - dealer_sellout_long_PORTUGAL_CY2025.csv
  - dealer_sellout_long_PORTUGAL_CY2026YTD.csv

Target schema: 01_data/12_raw/dealer_sales/DEALER_SELL_OUT.md
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

FIELDS = [
    "distributor_code",
    "source_file",
    "grain",
    "year",
    "period",
    "period_grain",
    "part_number",
    "part_number_raw",
    "description",
    "qty",
    "dealer_name",
    "vehicle_scope",
    "flags",
]


def normalize_pn(raw) -> tuple[str | None, str]:
    s = str(raw).strip() if raw is not None else ""
    digits = re.sub(r"\D", "", s)
    if len(digits) < 8:
        return None, s
    return digits.zfill(10), s


def parse_date(v) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def pick(idx: dict[str, int], *names: str) -> int | None:
    """First matching header, exact then case-insensitive-substring."""
    for n in names:
        if n in idx:
            return idx[n]
    low = {k.lower(): v for k, v in idx.items()}
    for n in names:
        for k, v in low.items():
            if n.lower() in k:
                return v
    return None


def normalize_portugal(src: Path, out_dir: Path) -> dict[int, Path]:
    import openpyxl

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    sheet_years = {"2025": 2025, "2026": 2026}
    by_year: dict[int, list[dict]] = {2025: [], 2026: []}
    year_totals: dict[tuple[int, str], float] = defaultdict(float)
    raw_first: dict[str, str] = {}
    stats: dict[int, dict] = {}

    for sheet_name, year in sheet_years.items():
        if sheet_name not in wb.sheetnames:
            print(f"WARN: sheet {sheet_name!r} not found (have {wb.sheetnames}) - skipped")
            continue
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
        d_i = pick(idx, "Data da Fatura", "Data")
        pn_i = pick(idx, "Referencia", "Referência", "Part")
        q_i = pick(idx, "Sum of Quantidade", "Quantidade", "Qty")
        if d_i is None or pn_i is None or q_i is None:
            raise SystemExit(f"Sheet {sheet_name}: cannot map columns from {list(idx)}")

        agg: dict[tuple[int, str, str], float] = defaultdict(float)
        kept = skipped_qty = skipped_pn = skipped_date = 0
        for r in it:
            d = parse_date(r[d_i])
            if d is None:
                skipped_date += 1
                continue
            try:
                qty = float(r[q_i] or 0)
            except (TypeError, ValueError):
                skipped_qty += 1
                continue
            if qty == 0:
                skipped_qty += 1
                continue
            pn, raw = normalize_pn(r[pn_i])
            if not pn:
                skipped_pn += 1
                continue
            agg[(d.year, d.strftime("%Y-%m"), pn)] += qty
            raw_first.setdefault(pn, raw)
            kept += 1

        flags = "qty_only|network_pn" if year == 2025 else "qty_only|network_pn|partial_year"
        for (y, period, pn), q in sorted(agg.items()):
            by_year[year].append(
                dict(
                    distributor_code="PORTUGAL",
                    source_file=src.name,
                    grain="network_pn",
                    year=y,
                    period=period,
                    period_grain="month",
                    part_number=pn,
                    part_number_raw=raw_first.get(pn, pn),
                    description="",
                    qty=q,
                    dealer_name="",
                    vehicle_scope="unknown",
                    flags=flags,
                )
            )
            year_totals[(year, pn)] += q

        stats[year] = dict(lines_kept=kept, skipped_date=skipped_date,
                           skipped_qty=skipped_qty, skipped_pn=skipped_pn,
                           months=len({p for (_, p, _) in agg}), pns=len({p for (_, _, p) in agg}))

    wb.close()

    for (year, pn), q in sorted(year_totals.items()):
        flags = "qty_only|network_pn" if year == 2025 else "qty_only|network_pn|partial_year"
        by_year[year].append(
            dict(
                distributor_code="PORTUGAL",
                source_file=src.name,
                grain="network_pn",
                year=year,
                period=str(year) if year == 2025 else f"{year}-YTD",
                period_grain="year" if year == 2025 else "ytd",
                part_number=pn,
                part_number_raw=raw_first.get(pn, pn),
                description="",
                qty=q,
                dealer_name="",
                vehicle_scope="unknown",
                flags=flags,
            )
        )

    out_dir.mkdir(parents=True, exist_ok=True)
    outs = {
        2025: out_dir / "dealer_sellout_long_PORTUGAL_CY2025.csv",
        2026: out_dir / "dealer_sellout_long_PORTUGAL_CY2026YTD.csv",
    }
    for year, path in outs.items():
        rows = by_year.get(year, [])
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS)
            w.writeheader()
            w.writerows(rows)
        n_month = sum(1 for r in rows if r["period_grain"] == "month")
        n_roll = sum(1 for r in rows if r["period_grain"] in ("year", "ytd"))
        qty_roll = sum(r["qty"] for r in rows if r["period_grain"] in ("year", "ytd"))
        s = stats.get(year, {})
        print(f"{path.name}: rows={len(rows)} month={n_month} rollup={n_roll} "
              f"rollup_qty={qty_roll:,.0f} pns={s.get('pns')} months={s.get('months')} "
              f"lines_kept={s.get('lines_kept')} skipped(date/qty/pn)="
              f"{s.get('skipped_date')}/{s.get('skipped_qty')}/{s.get('skipped_pn')}")
    return outs


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--src",
        type=Path,
        default=Path("01_data/12_raw/dealer_sales/PORTUGAL/PORTUGAL_Sales_to_Dealer_2025_2026YTD_Aug.xlsx"),
    )
    ap.add_argument("--out-dir", type=Path, default=Path("01_data/00_derived"))
    args = ap.parse_args()
    src, out_dir = args.src, args.out_dir
    if not src.is_absolute():
        root = Path(__file__).resolve().parents[2]
        src = root / src
        out_dir = root / out_dir
    normalize_portugal(src, out_dir)


if __name__ == "__main__":
    main()
