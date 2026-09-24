# -*- coding: utf-8 -*-
"""Read-only: the 44 in the source classified workbook, and the ENGINE ASM leaf there."""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
A = ROOT / "01_data" / "91_build" / "_tmp_apply_export_pr00_20260924"
WB = ROOT / "01_data" / "01_Pricing" / "IE40 Price list as of 13.03.2026 PR00_3_level_classified_GPT.xlsx"

pns = {l.strip() for l in (A / "carried_out_not_in_master.txt").read_text().splitlines() if l.strip()}
wb = openpyxl.load_workbook(WB, read_only=True)
ws = wb["Sheet1"]
rows = ws.iter_rows(values_only=True)
hdr = next(rows)
idx = {h: i for i, h in enumerate(hdr)}
hit = {}
eng = []
for r in rows:
    pn = str(r[idx["P/N"]] or "").strip()
    if pn in pns:
        hit[pn] = r
    if str(r[idx["Description"]] or "").strip().upper().startswith("ENGINE ASM"):
        eng.append(r)

print(f"of the 44 present in the classified source workbook: {len(hit)}")
print("missing:", sorted(pns - set(hit)))
for pn in sorted(hit):
    r = hit[pn]
    print(f"  {pn} | {r[idx['Description']]} | L1={r[idx['Level 1 Family']]} | L2={r[idx['Level 2 Group']]} "
          f"| L3={r[idx['Level 3 Type']]} | fig={r[idx['Fig No.']]}-{r[idx['Key No.']]}")

print(f"\nrows whose Description starts with ENGINE ASM: {len(eng)}")
print("by full path:", Counter(
    (r[idx["Level 1 Family"]], r[idx["Level 2 Group"]], r[idx["Level 3 Type"]]) for r in eng
).most_common(10))
print("by description:", Counter(str(r[idx["Description"]]).strip() for r in eng).most_common(10))

print("\nrows on the exact leaf ('Repair parts','Engine parts','Engine assembly'):")
leaf = [r for r in eng if (r[idx["Level 1 Family"]], r[idx["Level 2 Group"]], r[idx["Level 3 Type"]])
        == ("Repair parts", "Engine parts", "Engine assembly")]
print("count:", len(leaf))
print("sample:", [(str(r[idx["P/N"]]), str(r[idx["Description"]])) for r in leaf[:6]])
print("\nall rows in the whole workbook on that leaf:")
n = 0
leaf_all = []
wb2 = openpyxl.load_workbook(WB, read_only=True)
ws2 = wb2["Sheet1"]
it = ws2.iter_rows(values_only=True)
next(it)
for r in it:
    if (str(r[idx["Level 1 Family"]]).strip(), str(r[idx["Level 2 Group"]]).strip(),
            str(r[idx["Level 3 Type"]]).strip()) == ("Repair parts", "Engine parts", "Engine assembly"):
        n += 1
        leaf_all.append(str(r[idx["P/N"]]).strip())
print("count:", n)
print("is 8981422000 on the leaf?", "8981422000" in leaf_all)
print("is any of the 44 on the leaf?", sorted(pns & set(leaf_all)))
