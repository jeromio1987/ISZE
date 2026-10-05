import inspect
import io
import zipfile

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBar, DataBarRule

print(inspect.getsource(DataBar))
print("=" * 60)
try:
    from openpyxl.formatting.rule import DataBarRule as _R

    src = inspect.getsource(_R.__init__)
    print(src)
except Exception as e:  # pragma: no cover
    print("no src", e)

print("=" * 60)


def dump(rule, label):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws.conditional_formatting.add("A1:A2", rule)
    b = io.BytesIO()
    wb.save(b)
    z = zipfile.ZipFile(b)
    xml = z.read("xl/worksheets/sheet1.xml").decode()
    start = xml.find("<conditionalFormatting")
    print(label, "->", xml[start : start + 700])
    print()


r1 = DataBarRule(start_type="num", start_value=0, end_type="max", color="C0552D", showValue=True)
dump(r1, "plain")

r2 = DataBarRule(start_type="num", start_value=0, end_type="max", color="C0552D", showValue=True)
r2.dataBar.gradient = False
dump(r2, "gradient False on dataBar")

r3 = DataBarRule(start_type="num", start_value=0, end_type="max", color="3DC0552D", showValue=True)
dump(r3, "color with alpha 3D")
