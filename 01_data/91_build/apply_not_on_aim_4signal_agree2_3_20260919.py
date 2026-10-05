# -*- coding: utf-8 -*-
"""Apply agree3_high + agree2_review from 4-signal score to aim_master_v2.

Jerome: Apply 2026-09-19 — both tiers.

Does NOT apply agree1_*, conflict_*, needs_eq, eq_last_resort.
Pattern mirrors apply_not_on_aim_spine_high_20260919.py.
"""
from __future__ import annotations

import shutil
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
D = ISZE / "01_data" / "00_derived"
PRICING = ISZE / "01_data" / "01_Pricing"
MASTER = D / "aim_master_v2.csv"
SCORED = D / "pricing_not_on_aim_4signal_20260919.csv"
PATCH = D / "aim_gap_4signal_agree2_3_patch_20260919.csv"
OUT_XLSX = PRICING / "Not_on_AIM_spine_4SIGNAL_APPLIED_agree2_3_20260919.xlsx"
OPEN_CMD = PRICING / "OPEN_Not_on_AIM_spine_4SIGNAL_APPLIED_20260919.cmd"

APPLY_STATUSES = ("agree3_high", "agree2_review")
STATUS_MARKER = {
    "agree3_high": "4signal_agree3",
    "agree2_review": "4signal_agree2",
}
PLACEHOLDER = {"", "Unknown", "Unknown Review", "Not on AIM spine"}
STAMP = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def norm_pn(s: object) -> str:
    import re

    return re.sub(r"\D", "", str(s or "")).zfill(10)


def write_applied_xlsx(patch: pd.DataFrame, stats: dict) -> None:
    wb = Workbook()
    navy = Font(bold=True, color="FFFFFF")
    head = PatternFill("solid", fgColor="15233C")
    title = Font(bold=True, size=14, color="15233C")

    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "Not on AIM — APPLIED (agree3_high + agree2_review)"
    ov["A1"].font = title
    ov["A2"] = (
        "Written into aim_master_v2 on this run. Labels taken only from proposed_aim_l1/l2/l3 "
        "in the 4-signal score file. Weak / conflict / EQ tiers were NOT applied. "
        "Tools / MX2 already excluded from the scored residual."
    )
    ov.merge_cells("A2:F2")
    ov["A2"].alignment = Alignment(wrap_text=True)
    ov.row_dimensions[2].height = 48

    rows = [
        ("agree3_high PNs", stats["agree3_n"]),
        ("agree3_high billed EUR", stats["agree3_eur"]),
        ("agree2_review PNs", stats["agree2_n"]),
        ("agree2_review billed EUR", stats["agree2_eur"]),
        ("Total applied PNs", stats["total_n"]),
        ("Total applied billed EUR", stats["total_eur"]),
        ("Master rows updated", stats["upd"]),
        ("Master rows added", stats["add"]),
        ("Backup", stats["bak"]),
        ("Patch CSV", str(PATCH)),
        ("Master", str(MASTER)),
    ]
    ov["A4"] = "Item"
    ov["B4"] = "Value"
    for c in ov[4]:
        c.fill = head
        c.font = navy
    for i, (k, v) in enumerate(rows, 5):
        ov.cell(i, 1, k)
        cell = ov.cell(i, 2, v)
        if "EUR" in k and isinstance(v, (int, float)):
            cell.number_format = "#,##0.00"
    ov.column_dimensions["A"].width = 28
    ov.column_dimensions["B"].width = 72

    cols = [
        "part_number",
        "description",
        "billed_eur",
        "status",
        "aim_l1_fixed",
        "aim_l2_fixed",
        "aim_l3",
        "method_stamp",
        "score_method",
        "notes",
    ]
    rename = {
        "part_number": "Part number",
        "description": "Description",
        "billed_eur": "Parts billed EUR",
        "status": "Status",
        "aim_l1_fixed": "AIM L1 written",
        "aim_l2_fixed": "AIM L2 written",
        "aim_l3": "AIM L3 written",
        "method_stamp": "Method stamp",
        "score_method": "Score method",
        "notes": "Notes",
    }

    def sheet(name: str, df: pd.DataFrame) -> None:
        w = wb.create_sheet(name[:31])
        use = df[[c for c in cols if c in df.columns]].rename(columns=rename)
        for row in dataframe_to_rows(use, index=False, header=True):
            w.append(row)
        for cell in w[1]:
            cell.fill = head
            cell.font = navy
        headers = [c.value for c in w[1]]
        for col_i, h in enumerate(headers, 1):
            for row in range(2, w.max_row + 1):
                cell = w.cell(row, col_i)
                if h == "Parts billed EUR" and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
            w.column_dimensions[get_column_letter(col_i)].width = min(
                32, max(12, len(str(h)) + 2)
            )
        w.auto_filter.ref = w.dimensions
        w.freeze_panes = "A2"

    sheet("APPLIED", patch.sort_values("billed_eur", ascending=False))
    sheet("agree3_high", patch[patch["status"] == "agree3_high"])
    sheet("agree2_review", patch[patch["status"] == "agree2_review"])

    wb.save(OUT_XLSX)
    OPEN_CMD.write_text(
        '@echo off\r\nstart "" "%~dp0Not_on_AIM_spine_4SIGNAL_APPLIED_agree2_3_20260919.xlsx"\r\n',
        encoding="ascii",
        newline="",
    )


def main() -> None:
    scored = pd.read_csv(SCORED, dtype=str)
    scored["billed_eur"] = pd.to_numeric(scored["billed_eur"], errors="coerce")
    apply = scored[scored["status"].isin(APPLY_STATUSES)].copy()
    apply["pn"] = apply["pn"].map(norm_pn)
    assert len(apply) > 0, "no agree3/agree2 rows"

    # Refuse empty proposed labels — never invent AIM
    bad = apply["proposed_aim_l1"].fillna("").str.strip().isin(["", "nan", "None"])
    if bad.any():
        raise SystemExit(f"{int(bad.sum())} rows missing proposed_aim_l1 — abort")

    patch = pd.DataFrame(
        {
            "part_number": apply["pn"],
            "description": apply["description"].fillna(""),
            "billed_eur": apply["billed_eur"],
            "status": apply["status"],
            "aim_l1_fixed": apply["proposed_aim_l1"].fillna("").astype(str).str.strip(),
            "aim_l2_fixed": apply["proposed_aim_l2"].fillna("").astype(str).str.strip(),
            "aim_l3": apply["proposed_aim_l3"].fillna("").astype(str).str.strip(),
            "method_stamp": apply["status"].map(STATUS_MARKER),
            "score_method": apply["method"].fillna(""),
            "notes": apply["notes"].fillna(""),
        }
    )
    patch.to_csv(PATCH, index=False)

    bak = D / f"aim_master_v2.csv.bak_pre_4signal_agree2_3_{STAMP}"
    shutil.copy2(MASTER, bak)
    print("backup", bak)

    m = pd.read_csv(MASTER, dtype=str, low_memory=False)
    m["pn_key"] = m["part_number"].map(norm_pn)
    patch["pn_key"] = patch["part_number"].map(norm_pn)

    before_unknown = m["aim_l1_fixed"].fillna("").isin(PLACEHOLDER).sum()

    upd = 0
    add = 0
    for _, r in patch.iterrows():
        hit = m["pn_key"] == r["pn_key"]
        n = int(hit.sum())
        l1 = str(r["aim_l1_fixed"])
        l2 = str(r["aim_l2_fixed"])
        l3 = str(r["aim_l3"])
        marker = str(r["method_stamp"])
        if n == 0:
            row = {c: "" for c in m.columns}
            row["part_number"] = r["part_number"]
            row["pn_key"] = r["pn_key"]
            row["description"] = r.get("description") or ""
            row["aim_l1_fixed"] = l1
            row["aim_l2_fixed"] = l2
            row["aim_l3"] = l3
            if "aim_l1" in m.columns:
                row["aim_l1"] = l1
            if "aim_l2" in m.columns:
                row["aim_l2"] = l2
            if "has_aim" in m.columns:
                row["has_aim"] = "1"
            row["aim_changed"] = marker
            row["aim_source"] = marker
            m = pd.concat([m, pd.DataFrame([row])], ignore_index=True)
            add += 1
            continue
        m.loc[hit, "aim_l1_fixed"] = l1
        m.loc[hit, "aim_l2_fixed"] = l2
        m.loc[hit, "aim_l3"] = l3
        if "aim_l1" in m.columns:
            m.loc[hit, "aim_l1"] = l1
        if "aim_l2" in m.columns:
            m.loc[hit, "aim_l2"] = l2
        if "has_aim" in m.columns:
            m.loc[hit, "has_aim"] = "1"
        prev = m.loc[hit, "aim_changed"].fillna("").astype(str)
        m.loc[hit, "aim_changed"] = [
            marker if not x.strip() else (x if marker in x else f"{x}|{marker}")
            for x in prev
        ]
        src = m.loc[hit, "aim_source"].fillna("").astype(str)
        m.loc[hit, "aim_source"] = [
            marker if not x.strip() else (x if marker in x else f"{x}|{marker}")
            for x in src
        ]
        upd += n

    m = m.drop(columns=["pn_key"])
    m.to_csv(MASTER, index=False)

    after_unknown = m["aim_l1_fixed"].fillna("").isin(PLACEHOLDER).sum()

    a3 = patch[patch["status"] == "agree3_high"]
    a2 = patch[patch["status"] == "agree2_review"]
    stats = {
        "agree3_n": len(a3),
        "agree3_eur": float(a3["billed_eur"].sum()),
        "agree2_n": len(a2),
        "agree2_eur": float(a2["billed_eur"].sum()),
        "total_n": len(patch),
        "total_eur": float(patch["billed_eur"].sum()),
        "upd": upd,
        "add": add,
        "bak": str(bak),
    }
    write_applied_xlsx(patch, stats)

    print("patch_rows", len(patch))
    print("agree3_high", stats["agree3_n"], stats["agree3_eur"])
    print("agree2_review", stats["agree2_n"], stats["agree2_eur"])
    print("master_rows_updated", upd)
    print("master_rows_added", add)
    print("unknown_l1_before", int(before_unknown), "after", int(after_unknown))
    print("billed_applied", stats["total_eur"])
    print("PATCH", PATCH)
    print("MASTER", MASTER)
    print("XLSX", OUT_XLSX)
    print("OPEN", OPEN_CMD)
    print("BAK", bak)


if __name__ == "__main__":
    main()
