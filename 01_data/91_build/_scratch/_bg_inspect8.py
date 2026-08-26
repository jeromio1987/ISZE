import openpyxl, sys, csv, os
sys.stdout.reconfigure(encoding='utf-8')

def dump(path, sheet, nrows=50, ncols=None, start=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    print("="*100)
    print(f"SHEET: {sheet}  dims={ws.max_row}x{ws.max_column}")
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        if cnt < start:
            cnt += 1; continue
        vals = list(row)
        if ncols:
            vals = vals[:ncols]
        print(cnt, vals)
        cnt += 1
        if cnt >= start+nrows:
            break

p = r"c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx"
dump(p, "DG Material Pareto", 50, start=3)

# model_dimension.csv header
md = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\07_UIO\vin\model_dimension.csv"
print("\nmodel_dimension.csv exists:", os.path.exists(md))
if os.path.exists(md):
    with open(md, newline='', encoding='utf-8-sig') as fh:
        rd = csv.reader(fh)
        for i, row in enumerate(rd):
            print(i, row)
            if i >= 5:
                break
else:
    # search for it
    for root, dirs, files in os.walk(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\07_UIO"):
        for f in files:
            if 'model' in f.lower():
                print("  found:", os.path.join(root, f))
