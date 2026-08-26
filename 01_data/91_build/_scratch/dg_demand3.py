import sys, io, os
import pandas as pd
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

P = r'c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx'
BASE = r'C:\Users\jerom\Desktop\claude\ISZE\01_data'

dg = pd.read_excel(P, sheet_name='DG list (deduped)', header=3)
dg.columns = [str(c).strip() for c in dg.columns]
def norm(x):
    if x is None: return None
    s = str(x).strip()
    return s.zfill(10)
dg['_mat'] = dg['Material'].map(norm)
dg['_new'] = dg['Newest material no.'].map(norm)
dg_319 = set(dg['_mat'])
dg_union = set(dg['_mat']) | set(dg['_new'])
desc = dict(zip(dg['_mat'], dg['Description'])) | dict(zip(dg['_new'], dg['Description']))
def fam(d):
    x = str(d).upper()
    if 'AIR BAG' in x or 'AIRBAG' in x: return 'Airbag'
    if 'BELT' in x: return 'Seatbelt pretensioner'
    if 'CONTROL UNIT' in x or 'SRS' in x: return 'SRS control unit'
    if any(k in x for k in ['OIL','LUB','FLUID','CLEAN','ATF','GREASE','DEF','ADBLUE','BATTERY']): return 'Bulk lubricant'
    return 'Other'

def summarize(name, pn_qty, period):
    pn_qty = {norm(k): v for k, v in pn_qty.items() if norm(k)}
    tot_pn = len(pn_qty)
    tot_q = sum(pn_qty.values())
    hits = {p: q for p, q in pn_qty.items() if p in dg_union}
    fam_q = {}
    for p, q in hits.items():
        f = fam(desc.get(p,''))
        fam_q[f] = fam_q.get(f,0)+q
    print(f"{name:20s} period={period:18s} totalPN={tot_pn:6d} totalQty={tot_q:9.0f}  DG_PN={len(hits):3d}  DG_qty={sum(hits.values()):7.1f}  fams={fam_q}")
    return hits

print("DB                      period            totalPN totalQty  DG_PN  DG_qty   family-qty")

# IUK / IRELAND / GREECE from long CSV, MONTH rows only (avoid year-agg double count)
for db in ['IUK','IRELAND']:
    for yr in ['CY2025','CY2026YTD']:
        f = os.path.join(BASE, '00_derived', f'dealer_sellout_long_{db}_{yr}.csv')
        df = pd.read_csv(f, dtype={'part_number':str})
        m = df[df['period_grain']=='month'] if 'period_grain' in df.columns else df
        if len(m)==0: m = df
        agg = m.groupby('part_number')['qty'].sum().to_dict()
        summarize(f"{db}_{yr}", agg, yr)

f = r'C:\Users\jerom\Desktop\claude\ISZE\02_dashboards\Data\00_derived\dealer_sellout_long_GREECE_CY2025.csv'
df = pd.read_csv(f, dtype={'part_number':str})
m = df[df['period_grain']=='month'] if 'period_grain' in df.columns else df
if len(m)==0: m = df
agg = m.groupby('part_number')['qty'].sum().to_dict()
summarize('GREECE', agg, 'CY2025')

# ITUK
f = r'C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\ITUK\Sales to Dealers_ITUK_CY2025.xlsx'
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
ws = wb['Sheet1']
agg = {}
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 5: continue
    item = r[1]
    if item is None: continue
    s = str(item).strip()
    pn = s.split(' - ')[0].split(' ')[0].strip()
    if pn == 'Item Number' or not pn.isdigit(): continue
    qty = 0
    for c in r[3:15]:
        if c is not None:
            try: qty += float(c)
            except: pass
    agg[pn] = agg.get(pn,0)+qty
summarize('ITUK', agg, 'CY2025')

# UTI: sum qty Month across all monthly files (note Feb-Jul 2025 gap)
utidir = r'C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\UTI'
u_qty = {}
months_present = []
for fn in sorted(os.listdir(utidir)):
    if not fn.endswith('.xlsx'): continue
    wb = openpyxl.load_workbook(os.path.join(utidir, fn), read_only=True, data_only=True)
    ws = wb['UTIPartSales']
    months_present.append(fn)
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 6: continue
        if not r or len(r)==0: continue
        pn = norm(r[0])
        if not pn or not pn.isdigit(): continue
        qm = r[1] if len(r)>1 and r[1] is not None else 0
        try: qm = float(qm)
        except: qm = 0
        u_qty[pn] = u_qty.get(pn,0)+qm
summarize('UTI (all months avail)', u_qty, '2025+2026 partial')
print("  UTI monthly files present:", len(months_present), "(missing CY2025 Feb-Jul)")

print("\nDG union size (319 mat + newest):", len(dg_union))
