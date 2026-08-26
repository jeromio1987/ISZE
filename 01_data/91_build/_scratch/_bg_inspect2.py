import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

def dump(path, sheet, nrows=8, ncols=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    print("="*100)
    print(f"SHEET: {sheet}  dims={ws.max_row}x{ws.max_column}")
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if ncols:
            vals = vals[:ncols]
        print(cnt, vals)
        cnt += 1
        if cnt >= nrows:
            break

p = r"c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx"
dump(p, "Data (joined)", 6)
dump(p, "Orders", 6)
dump(p, "DG list (deduped)", 8)
dump(p, "Master data", 8)
