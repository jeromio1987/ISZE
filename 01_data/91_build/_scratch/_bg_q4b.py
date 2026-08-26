import pandas as pd, numpy as np, sys
sys.stdout.reconfigure(encoding='utf-8')

D = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX"
m = pd.read_csv(f"{D}\\_part_stats.csv", dtype=str)
m['air_bookings'] = m['air_bookings'].astype(int)

# boundary analysis
vc = m['air_bookings'].value_counts().sort_index(ascending=False)
print("air bookings distribution (>=8):")
print(vc[vc.index>=8])

# how many with >=9, and how many exactly 8
n9 = (m['air_bookings']>=9).sum()
n8 = (m['air_bookings']==8).sum()
print(f"\nparts with >=9 bookings: {n9}, parts with exactly 8: {n8}")
print("So top 80 boundary is at 8 bookings; need first {0} of the {1} tied parts".format(80-n9, n8))

# check pms_registry among top 80
top80 = m.sort_values('air_bookings', ascending=False).head(80)
pc = pd.read_csv(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\parts_catalog_dimension.csv", dtype=str)
pc['pms_category'] = pc['pms_category'].fillna('').str.strip()
pc['pms_registry'] = pc['pms_registry'].fillna('').str.strip()

t80 = top80.merge(pc[['part_number','pms_category','pms_registry','classification_source','fig_key_label']],
                  left_on='Material', right_on='part_number', how='left')
print("\n80 top flyers with any PMS signal:")
print(t80[t80['pms_category']!=''][['Material','Description','pms_category','pms_registry','classification_source','air_bookings']].to_string(index=False))

# also: how many of top 80 have pms_registry True
print("\npms_registry True in top80:", (t80['pms_registry']=='True').sum())

# search v4 workbook for 'PMS' mentions
import openpyxl
v4 = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\UMIUTI_Logistics_Study_v4_enhanced.xlsx"
wb = openpyxl.load_workbook(v4, read_only=True, data_only=True)
print("\n===== v4 PMS mentions =====")
for s in wb.sheetnames:
    ws = wb[s]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and 'PMS' in cell:
                print(f"[{s}]", cell[:200])
                break
