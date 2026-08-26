import openpyxl, sys, glob, os
sys.stdout.reconfigure(encoding='utf-8')

def dump(path, sheet, nrows=8, ncols=None, start=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    print("="*100)
    print(f"SHEET: {sheet}  dims={ws.max_row}x{ws.max_column}")
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        if cnt < start:
            cnt += 1; continue
        vals = list(row)
        if ncols:
            vals = vals[:ncols]
        print(cnt, vals)
        cnt += 1
        if cnt >= start+nrows:
            break

p4 = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\UMIUTI_Logistics_Study_v4_enhanced.xlsx"
dump(p4, "Back-order join", 12, start=6)

# backorder file
bo = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\08_Distributor\BO OVERVIEW - UNIVERSAL TRUCKS ISRAEL LTD 01-07-2026 (1).xlsx"
wb = openpyxl.load_workbook(bo, read_only=True, data_only=True)
print("BO sheets:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print("   ", s, ws.max_row, "x", ws.max_column)

# dealer sellout files
d = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\UTI"
print("\nDealer sellout files:")
for f in sorted(os.listdir(d)):
    print("   ", f, os.path.getsize(os.path.join(d,f)))
