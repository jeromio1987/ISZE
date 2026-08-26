import sys, io, os
import pandas as pd
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

P = r'c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx'
BASE = r'C:\Users\jerom\Desktop\claude\ISZE\01_data'

dg = pd.read_excel(P, sheet_name='DG list (deduped)', header=3)
dg.columns = [str(c).strip() for c in dg.columns]
def norm(x):
    if x is None: return None
    s = str(x).strip()
    return s.zfill(10)
dg['_mat'] = dg['Material'].map(norm)
dg['_new'] = dg['Newest material no.'].map(norm)
dg_319 = set(dg['_mat'])
dg_union = set(dg['_mat']) | set(dg['_new'])
desc = dict(zip(dg['_mat'], dg['Description'])) | dict(zip(dg['_new'], dg['Description']))

def fam(desc):
    x = str(desc).upper()
    if 'AIR BAG' in x or 'AIRBAG' in x: return 'Airbag'
    if 'BELT' in x: return 'Seatbelt pretensioner'
    if 'CONTROL UNIT' in x or 'SRS' in x: return 'SRS control unit'
    if any(k in x for k in ['OIL','LUB','FLUID','CLEAN','ATF','GREASE','DEF','ADBLUE','BATTERY']): return 'Bulk lubricant'
    return 'Other'

def show(name, pn_qty, min_qty=0):
    pn_qty = {norm(k): v for k, v in pn_qty.items() if norm(k)}
    hits = {p: q for p, q in pn_qty.items() if p in dg_union}
    hits = {p: q for p, q in hits.items() if q >= min_qty}
    print(f"\n### {name}: {len(hits)} DG PNs, qty={round(sum(hits.values()),1)}")
    for p, q in sorted(hits.items(), key=lambda x: -x[1]):
        print(f"   {p}  {str(desc.get(p,'?'))[:38]:38s}  {fam(desc.get(p,'')):22s}  qty={q:g}".replace('g', 'g'))
    return hits

# UTI full union
utidir = r'C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\UTI'
u_pn = set()
u_qty = {}
for fn in sorted(os.listdir(utidir)):
    if not fn.endswith('.xlsx'): continue
    wb = openpyxl.load_workbook(os.path.join(utidir, fn), read_only=True, data_only=True)
    ws = wb['UTIPartSales']
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 6: continue
        if not r or len(r)==0: continue
        pn = norm(r[0])
        if not pn or not pn.isdigit(): continue
        u_pn.add(pn)
print("UTI union distinct PNs:", len(u_pn))
print("UTI union DG PNs (319):", len(u_pn & dg_319))
print("UTI union DG PNs (337):", len(u_pn & dg_union))
print("UTI DG PNs:", sorted(u_pn & dg_union))
for p in sorted(u_pn & dg_union):
    print("   ", p, desc.get(p,'?'))

# ITUK supersession extra
print("\nITUK DG (union) full list:")
f = r'C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\ITUK\Sales to Dealers_ITUK_CY2025.xlsx'
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
ws = wb['Sheet1']
agg = {}
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 5: continue
    item = r[1]
    if item is None: continue
    s = str(item).strip()
    pn = s.split(' - ')[0].split(' ')[0].strip()
    if pn == 'Item Number' or not pn.isdigit(): continue
    qty = 0
    for c in r[3:15]:
        if c is not None:
            try: qty += float(c)
            except: pass
    agg[pn] = agg.get(pn,0)+qty
show("ITUK", agg)
