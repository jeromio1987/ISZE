import openpyxl
import io, sys

OUT = io.open(r'C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\dg_explore.txt', 'w', encoding='utf-8')
def p(*a):
    print(*a, file=OUT)

wb = openpyxl.load_workbook(r'c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx', read_only=True)
for name in ['Data (joined)','Orders','DG list (deduped)','Dangerous goods','DG x MoT Summary','DG Material Pareto','Executive Summary','Air Booking Economics','Savings Scenarios','DG Material Pareto']:
    ws = wb[name]
    p('=====', name, ' rows=', ws.max_row, 'cols=', ws.max_column, '=====')
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True)):
        p(i+1, row)
    p()
OUT.close()
print('done')
