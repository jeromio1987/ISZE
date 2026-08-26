"""Verify v3 action plan against Jerome's exact rules."""
import openpyxl

SRC = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\11_planning\Action plan FY2027 VPACC_August update.xlsx"
OUT = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_ap_verify3.txt"

GREEN = "0092D050"
wb = openpyxl.load_workbook(SRC)
vp = wb["Action Plan VP"]
acc = wb["Action Plan ACC"]

VP_JUL = 9
VP_COMMENT = 19
ACC_JUL = 10
ACC_REPORT = 55

lines = []

def fill_rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None

def font_desc(cell):
    f = cell.font
    return dict(name=f.name, size=f.size, bold=f.bold, italic=f.italic, color=(f.color.rgb if f and f.color and isinstance(f.color, openpyxl.styles.Color) else f.color))

# --- green census ---
def census(ws, name, jul_col, report_col, month_cols):
    total = 0
    by_col = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if fill_rgb(ws.cell(r, c)) == GREEN:
                total += 1
                by_col[c] = by_col.get(c, 0) + 1
    lines.append(f"{name}: total green cells = {total}")
    lines.append(f"  green by column: {by_col}")
    # green only in report column
    bad_cols = [c for c in by_col if c != report_col]
    lines.append(f"  green outside report col (should be none): {bad_cols}")
    # every green report cell has July x
    nox = []
    for r in range(1, ws.max_row + 1):
        if fill_rgb(ws.cell(r, report_col)) == GREEN:
            jv = ws.cell(r, jul_col).value
            if not (isinstance(jv, str) and jv.strip().upper() == "X"):
                nox.append(r)
    lines.append(f"  green report cells WITHOUT July x (should be none): {nox}")
    # no green in future months (Aug..Mar)
    fut = [c for c in month_cols if c > jul_col]
    futg = []
    for r in range(1, ws.max_row + 1):
        for c in fut:
            if fill_rgb(ws.cell(r, c)) == GREEN:
                futg.append((r, c))
    lines.append(f"  green in future months (should be none): {futg}")
    return total

lines.append("=== GREEN CENSUS ===")
t1 = census(vp, "VP", VP_JUL, VP_COMMENT, list(range(6, 18)))
t2 = census(acc, "ACC", ACC_JUL, ACC_REPORT, list(range(7, 19)))
lines.append(f"TOTAL GREEN = {t1 + t2}")

# --- font check: my comment cells ---
lines.append("")
lines.append("=== FONT CHECK (my comment cells) ===")
vp_rows = [6, 7, 15, 19, 21, 34, 35, 37, 38, 42, 43, 47, 62, 67, 73]
acc_rows = [9]
for r in vp_rows:
    lines.append(f"  VP S{r}: {font_desc(vp.cell(r, VP_COMMENT))}")
for r in acc_rows:
    lines.append(f"  ACC BC{r}: {font_desc(acc.cell(r, ACC_REPORT))}")

# --- grey check: any grey fill/font on my cells ---
lines.append("")
lines.append("=== GREY CHECK (my comment cells) ===")
GREY_FILLS = {"FF808080", "FFA6A6A6", "FFD9D9D9", "FFBFBFBF", "FFD3D3D3", "FFC0C0C0", "FFEDEDED", "FFF2F2F2"}
grey_found = []
for name, ws, col, rows in [("VP", vp, VP_COMMENT, vp_rows), ("ACC", acc, ACC_REPORT, acc_rows)]:
    for r in rows:
        c = ws.cell(r, col)
        frgb = fill_rgb(c)
        if frgb and frgb.upper() in GREY_FILLS:
            grey_found.append(f"{name} row{r} fill {frgb}")
        if c.font and c.font.color and isinstance(c.font.color, openpyxl.styles.Color):
            crgb = c.font.color.rgb
            if crgb and crgb.upper() in GREY_FILLS:
                grey_found.append(f"{name} row{r} fontcolor {crgb}")
lines.append(f"  grey on my cells (should be none): {grey_found}")

# --- BS cells unchanged: compare against source ---
lines.append("")
lines.append("=== BS CELLS UNCHANGED ===")
SRC2 = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\20260820__Action plan FY2027 VPACC_July update.xlsx"
wsrc = openpyxl.load_workbook(SRC2)
vp_s = wsrc["Action Plan VP"]
acc_s = wsrc["Action Plan ACC"]
bs_vp_rows = [23, 25, 26, 27, 29, 30, 31, 32, 33, 63, 64, 65]
changed = []
for r in bs_vp_rows:
    for c in range(1, vp.max_column + 1):
        a = vp.cell(r, c).value
        b = vp_s.cell(r, c).value
        if str(a) != str(b):
            changed.append(f"VP row{r} col{c}")
lines.append(f"  BS VP rows changed (should be none): {changed}")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("\n".join(lines))
