"""Verify corrected action plan against Jerome's exact rules."""
import openpyxl

SRC = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\11_planning\Action plan FY2027 VPACC_August update.xlsx"
OUT = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_ap_verify2.txt"

GREEN = "0092D050"
wb = openpyxl.load_workbook(SRC)
vp = wb["Action Plan VP"]
acc = wb["Action Plan ACC"]

VP_JUL = 9
VP_COMMENT = 19
ACC_JUL = 10
ACC_REPORT = 55

lines = []

def fill_rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None

# 1. green cell census across whole used range (both sheets), grouped by column
def green_census(ws, name, month_cols, jul_col, report_col):
    total = 0
    by_col = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if fill_rgb(ws.cell(r, c)) == GREEN:
                total += 1
                by_col[c] = by_col.get(c, 0) + 1
    lines.append(f"{name}: total green cells = {total}")
    lines.append(f"  by column: {by_col}")
    # 2. every green must be in the report/comment column only
    bad = [c for c in by_col if c != report_col]
    lines.append(f"  green in non-report columns (should be none): {bad}")
    # 3. every green report cell must have a July x on its row
    problems = []
    for r in range(1, ws.max_row + 1):
        if fill_rgb(ws.cell(r, report_col)) == GREEN:
            jv = ws.cell(r, jul_col).value
            if not (isinstance(jv, str) and jv.strip().upper() == "X"):
                problems.append((r, jv))
    lines.append(f"  green report cells WITHOUT July x (should be none): {problems}")
    # 4. July x rows WITHOUT green (should exist only where I added no comment)
    #    just report count of July-x rows for context
    july_x_rows = [r for r in range(1, ws.max_row + 1)
                   if isinstance(ws.cell(r, jul_col).value, str) and ws.cell(r, jul_col).value.strip().upper() == "X"]
    lines.append(f"  July-x rows total: {len(july_x_rows)}")
    # 5. no green in Aug/Sep/Oct (future) columns
    fut = [c for c in month_cols if c > jul_col]
    fut_green = []
    for r in range(1, ws.max_row + 1):
        for c in fut:
            if fill_rgb(ws.cell(r, c)) == GREEN:
                fut_green.append((r, c))
    lines.append(f"  green in future-month cols (should be none): {fut_green}")
    return total

lines.append("=== VP ===")
t1 = green_census(vp, "VP", list(range(6, 18)), VP_JUL, VP_COMMENT)
lines.append("=== ACC ===")
t2 = green_census(acc, "ACC", list(range(7, 19)), ACC_JUL, ACC_REPORT)
lines.append(f"TOTAL GREEN = {t1 + t2}")

# 6. italic check: my notes are italic rich-text runs
from openpyxl.cell.rich_text import CellRichText
lines.append("")
lines.append("=== Italic check (sample of my comment cells) ===")
for name, ws, col, rows in [
    ("VP", vp, VP_COMMENT, [6, 34, 35, 52, 7]),
    ("ACC", acc, ACC_REPORT, [7, 62, 69]),
]:
    for r in rows:
        v = ws.cell(r, col).value
        if isinstance(v, CellRichText):
            italic_runs = []
            for tb in v:
                if hasattr(tb, "font") and tb.font is not None:
                    italic_runs.append(bool(tb.font.i))
            lines.append(f"  {name} row{r}: rich text, italic flags per block = {italic_runs}")
        else:
            lines.append(f"  {name} row{r}: NOT rich text -> {type(v).__name__}")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("\n".join(lines))
