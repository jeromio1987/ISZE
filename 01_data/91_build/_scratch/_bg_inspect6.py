import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

def dump(path, sheet, nrows=40, ncols=None, start=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    print("="*100)
    print(f"SHEET: {sheet}  dims={ws.max_row}x{ws.max_column}")
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        if cnt < start:
            cnt += 1; continue
        vals = list(row)
        if ncols:
            vals = vals[:ncols]
        # print compactly
        print(cnt, vals)
        cnt += 1
        if cnt >= start+nrows:
            break

p = r"c:\Users\jerom\Downloads\UMIUTI_Logistics_Study_v2.xlsx"
for s in ["Executive Summary", "Air Booking Economics", "Stocking Candidates", "Strange Behaviours"]:
    dump(p, s, 45)
