"""Verify filled action plan: green fill counts + sample note cells."""
import openpyxl

SRC = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\11_planning\Action plan FY2027 VPACC_August update.xlsx"
OUT = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_ap_verify.txt"

wb = openpyxl.load_workbook(SRC)
lines = []

vp = wb["Action Plan VP"]
acc = wb["Action Plan ACC"]

# green fill count
vp_green = 0
acc_green = 0
for r in range(6, vp.max_row + 1):
    for c in range(6, 19):
        f = vp.cell(r, c).fill
        if f and f.fgColor and f.fgColor.rgb == "0092D050":
            vp_green += 1
for r in range(6, acc.max_row + 1):
    for c in range(7, 20):
        f = acc.cell(r, c).fill
        if f and f.fgColor and f.fgColor.rgb == "0092D050":
            acc_green += 1

lines.append(f"VP green cells: {vp_green}")
lines.append(f"ACC green cells: {acc_green}")

lines.append("")
lines.append("=== VP sample: row 6 (progress + July comment) ===")
lines.append(f"  progress R6 = {vp.cell(6,18).value!r} fill={vp.cell(6,18).fill.fgColor.rgb if vp.cell(6,18).fill and vp.cell(6,18).fill.fgColor else None}")
lines.append(f"  month F6 = {vp.cell(6,6).value!r} fill={vp.cell(6,6).fill.fgColor.rgb if vp.cell(6,6).fill and vp.cell(6,6).fill.fgColor else None}")
lines.append(f"  month G6 = {vp.cell(6,7).value!r} fill={vp.cell(6,7).fill.fgColor.rgb if vp.cell(6,7).fill and vp.cell(6,7).fill.fgColor else None}")
lines.append(f"  S6 comment = {vp.cell(6,19).value!r}")

lines.append("")
lines.append("=== VP GPS rows 34-35 (S col) ===")
lines.append(f"  S34 = {vp.cell(34,19).value!r}")
lines.append(f"  S35 = {vp.cell(35,19).value!r}")

lines.append("")
lines.append("=== ACC sample: row 7 (progress + July report) ===")
lines.append(f"  progress S7 = {acc.cell(7,19).value!r}")
lines.append(f"  month G7 = {acc.cell(7,7).value!r} fill={acc.cell(7,7).fill.fgColor.rgb if acc.cell(7,7).fill and acc.cell(7,7).fill.fgColor else None}")
lines.append(f"  month K7 = {acc.cell(7,11).value!r} fill={acc.cell(7,11).fill.fgColor.rgb if acc.cell(7,11).fill and acc.cell(7,11).fill.fgColor else None}")
lines.append(f"  BC7 = {acc.cell(7,55).value!r}")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("\n".join(lines))
