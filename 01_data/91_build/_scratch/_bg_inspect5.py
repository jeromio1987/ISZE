import openpyxl, sys, csv, os
sys.stdout.reconfigure(encoding='utf-8')

def dump(path, sheet, nrows=8, ncols=None, start=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet is None:
        sheet = wb.sheetnames[0]
    ws = wb[sheet]
    print("="*100)
    print(f"SHEET: {sheet}  dims={ws.max_row}x{ws.max_column}")
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        if cnt < start:
            cnt += 1; continue
        vals = list(row)
        if ncols:
            vals = vals[:ncols]
        print(cnt, vals)
        cnt += 1
        if cnt >= start+nrows:
            break

d = r"C:\Users\jerom\Desktop\claude\ISZE\01_data\12_raw\dealer_sales\UTI"
f0 = os.path.join(d, "UTIPartSales_2025-01.xlsx")
wb = openpyxl.load_workbook(f0, read_only=True, data_only=True)
print("sellout sheets:", wb.sheetnames)
dump(f0, None, 8)

# CSVs
for c in [r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\parts_catalog_dimension.csv",
          r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\bvp_pn_dimension.csv",
          r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\distributor_dimension.csv",
          r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\gps_master.csv"]:
    print("="*100)
    print("CSV:", os.path.basename(c))
    with open(c, newline='', encoding='utf-8-sig') as fh:
        rd = csv.reader(fh)
        for i, row in enumerate(rd):
            print(i, row)
            if i >= 3:
                break
