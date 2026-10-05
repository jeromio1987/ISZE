"""Audit the freeze panes, the layout, and the Source column against Jerome's file.

The defect: hdr_row set freeze_panes on every call, and the market sheets call it again
far down the sheet for the pass-through grid. That pinned almost the whole sheet - PR00
froze at A10990 of 10,994 rows, leaving 5 scrollable. His file freezes once, under the
header, and the rest of his sheets do not freeze at all.
"""
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

OURS = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing"
            r"\ISZE_Pricing_Stack_Simulation_2026-09-19.xlsx")
HIS = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\_raw"
           r"\pricing_simulation_2026_5pct.xlsx")
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
fails = []


def say(m=""):
    print(m)


def check(lab, ok):
    print(f"  {'PASS' if ok else 'FAIL'}  {lab}")
    if not ok:
        fails.append(lab)


# ---------------------------------------------------------------- freeze panes
say("=" * 100)
say("FREEZE PANES AND SCROLLABLE ROWS")
say("=" * 100)


def panes(path, want_names=None):
    z = zipfile.ZipFile(path)
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    names = re.findall(r"<sheet[^>]*\bname=\"([^\"]+)\"[^>]*r:id=\"([^\"]+)\"", wbxml)
    rel_map = {}
    for m in re.finditer(r"<Relationship\b([^>]*)/?>", rels):
        at = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        rel_map[at.get("Id")] = at.get("Target", "")
    out = {}
    for nm, rid in names:
        tgt = rel_map.get(rid, "")
        part = tgt.lstrip("/") if tgt.startswith("/") else (
            tgt if tgt.startswith("xl/") else "xl/" + tgt)
        if part not in z.namelist():
            continue
        root = ElementTree.fromstring(z.read(part))
        pane = root.find(f".//{NS}pane")
        top = pane.get("topLeftCell") if pane is not None else None
        ysplit = int(pane.get("ySplit") or 0) if pane is not None else 0
        dim = root.find(f"{NS}dimension")
        ref = dim.get("ref") if dim is not None else ""
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ref or "")
        max_row = int(m.group(4)) if m else 0
        out[nm] = {"topLeftCell": top, "ySplit": ysplit, "max_row": max_row,
                   "scrollable": max_row - ysplit if top else max_row}
    return out


A = panes(HIS)
B = panes(OURS)
say()
say(f"  {'sheet':<10}{'source':<9}{'topLeftCell':>13}{'ySplit':>8}{'max_row':>10}"
    f"{'scrollable':>12}")
for src, tag in ((A, "HIS"), (B, "OURS")):
    for nm, d in src.items():
        say(f"  {nm[:9]:<10}{tag:<9}{str(d['topLeftCell']):>13}{d['ySplit']:>8}"
            f"{d['max_row']:>10,}{d['scrollable']:>12,}")
    say()

say("  ASSERTIONS")
for nm in ("PR00", "UTI", "UMI", "CIS"):
    d = B.get(nm, {})
    check(f"{nm}: freeze pane topLeftCell is A2 (got {d.get('topLeftCell')!r})",
          d.get("topLeftCell") == "A2")
    check(f"{nm}: ySplit is 1 (got {d.get('ySplit')})", d.get("ySplit") == 1)
    sc = d.get("scrollable", 0)
    check(f"{nm}: scrollable rows {sc:,} in the thousands, never 5", sc > 1000)
d = B.get("Legends", {})
check(f"Legends: no freeze pane (got {d.get('topLeftCell')!r})",
      d.get("topLeftCell") is None)

# ---------------------------------------------------------------- layout
say()
say("=" * 100)
say("LAYOUT: WRAPPING, ROW HEIGHTS, COLUMN WIDTH")
say("=" * 100)
wb = load_workbook(OURS)
say(f"  {'sheet':<10}{'rows':>9}{'cols':>6}{'wrapped':>9}{'rows w/ht':>10}"
    f"{'max col width':>14}{'freeze':>9}")
for nm in wb.sheetnames:
    ws = wb[nm]
    wrapped = sum(1 for row in ws.iter_rows(min_row=2)
                  for c in row
                  if c.alignment is not None and c.alignment.wrap_text)
    heights = sum(1 for r in range(2, ws.max_row + 1)
                  if ws.row_dimensions.get(r) is not None
                  and ws.row_dimensions[r].height is not None)
    widths = [d.width for d in ws.column_dimensions.values() if d.width]
    say(f"  {nm:<10}{ws.max_row:>9,}{ws.max_column:>6}{wrapped:>9,}{heights:>10}"
        f"{max(widths) if widths else 0:>14.1f}{str(ws.freeze_panes):>9}")
    if nm in ("PR00", "UTI", "UMI", "CIS"):
        check(f"{nm}: zero wrapped cells in the data ({wrapped})", wrapped == 0)
        check(f"{nm}: no explicit row height on data rows ({heights})", heights == 0)
        check(f"{nm}: max column width {max(widths) if widths else 0:.1f} <= 42",
              (max(widths) if widths else 0) <= 42)

# ---------------------------------------------------------------- the Source column
say()
say("=" * 100)
say("THE SOURCE COLUMN: HIS CODES J / T / BLANK ONLY")
say("=" * 100)
for nm in ("PR00", "UTI", "UMI", "CIS"):
    ws = wb[nm]
    vals = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 4).value
        if ws.cell(r, 1).value == "TOTAL" or v is None and ws.cell(r, 1).value is None:
            if ws.cell(r, 1).value in ("RESULT", "Pass-through dial") or v is None:
                pass
        key = v if v not in (None, "") else "(blank)"
        vals[key] = vals.get(key, 0) + 1
    # restrict to the data rows: down to the TOTAL marker
    tot = next((r for r in range(2, ws.max_row + 1)
                if ws.cell(r, 1).value == "TOTAL"), ws.max_row)
    vals = {}
    for r in range(2, tot):
        v = ws.cell(r, 4).value
        key = v if v not in (None, "") else "(blank)"
        vals[key] = vals.get(key, 0) + 1
    say(f"  {nm:<6} {vals}")
    unknown = [k for k in vals if k not in ("J", "T", "(blank)")]
    check(f"{nm}: Source holds only J, T or blank (found {unknown})", not unknown)
check("no invented provenance code anywhere in the workbook",
      not any(isinstance(c.value, str) and c.value in ("J99", "T99", "X99")
              for ws in wb.worksheets for row in ws.iter_rows() for c in row))

# ---------------------------------------------------------------- formulas
say()
say("=" * 100)
say("FORMULAS: PARSING ONLY <f> ELEMENTS")
say("=" * 100)
z = zipfile.ZipFile(OURS)
bad = {}
hist = 0
for n in z.namelist():
    if not re.match(r"xl/worksheets/sheet\d+\.xml$", n):
        continue
    root = ElementTree.fromstring(z.read(n))
    txt = " ".join((f.text or "") for f in root.iter(f"{NS}f"))
    for fn in ("VLOOKUP", "SUMPRODUCT", "SUMIF", "COUNTIF", "INDEX", "MATCH",
               "HLOOKUP", "XLOOKUP"):
        c = txt.count(fn + "(")
        if c:
            bad[n] = bad.get(n, 0) + c
    for rate in ("163.7", "172.6"):
        for m in re.finditer(re.escape(rate), txt):
            before = txt[m.start() - 1] if m.start() else ""
            if not (before.isdigit() or before == "."):
                hist += 1
check(f"zero lookup/array functions in <f> only ({bad})", not bad)
check(f"zero formulas containing a historical rate ({hist})", hist == 0)

# ---------------------------------------------------------------- container
say()
say("=" * 100)
say("CONTAINER")
say("=" * 100)
check("zip integrity", z.testzip() is None)
xml_bad = 0
for n in z.namelist():
    if n.endswith(".xml") or n.endswith(".rels"):
        try:
            ElementTree.fromstring(z.read(n))
        except ElementTree.ParseError:
            xml_bad += 1
check("every XML part well-formed", xml_bad == 0)
sheets = re.findall(r"<sheet[^>]*\bname=\"([^\"]+)\"", z.read("xl/workbook.xml").decode())
check(f"five sheets declared ({sheets})",
      sheets == ["Legends", "PR00", "UTI", "UMI", "CIS"])
check("no macro part", not any("vbaProject" in n for n in z.namelist()))
check("no external link", not any("externalLink" in n for n in z.namelist()))
check("no pivot cache", not any("pivotCache" in n or "pivotTable" in n
                                for n in z.namelist()))
largest = max((z.getinfo(n).file_size for n in z.namelist()
               if re.match(r"xl/worksheets/sheet\d+\.xml$", n)), default=0)
check(f"largest sheet part {largest/1e6:.2f} MB, his was 14.72 MB",
      largest < 14_723_392)
heis = zipfile.ZipFile(HIS)
say()
say(f"  his largest sheet part : "
    f"{max(heis.getinfo(n).file_size for n in heis.namelist() if re.match(r'xl/worksheets/sheet.d+.xml', n) or 'worksheets' in n)/1e6:.2f} MB")
say(f"  ours                   : {largest/1e6:.2f} MB")

say()
if fails:
    say(f"FREEZE_AND_SOURCE_AUDIT_FAIL ({len(fails)})")
    for f in fails:
        say(f"  - {f}")
else:
    say("FREEZE_AND_SOURCE_AUDIT_OK")
