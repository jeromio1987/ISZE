# -*- coding: utf-8 -*-
"""AIM class vs Not-on-AIM overview after this morning's applies."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
D = ISZE / "01_data" / "00_derived"
P = ISZE / "01_data" / "01_Pricing"
OUT = P / "AIM_class_status_overview_20260919.xlsx"
CMD = P / "OPEN_AIM_class_status_overview_20260919.cmd"


def main() -> None:
    book_pn, book_eur = 15872, 34993609.4
    has_pn, has_eur = 13329, 34004042.0
    not_pn, not_eur = 2543, 989567.3
    open_pn, open_eur = 2531, 881466.3
    tools_pn, tools_eur = 11, 3919.2
    mx2_pn, mx2_eur = 1, 104181.9
    high_pn, high_eur = 299, 427228.9
    a23_pn, a23_eur = 275, 244132.2
    start_pn, start_eur = 3105, 1552827.4

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    navy = PatternFill("solid", fgColor="15233C")
    head = Font(bold=True, color="FFFFFF")
    title = Font(bold=True, size=14, color="15233C")
    money = '€#,##0'

    ws["A1"] = "AIM class status — pricing revision book"
    ws["A1"].font = title
    ws.merge_cells("A1:D1")
    ws["A2"] = (
        "As of 2026-09-19 morning. Window = price revision lines. "
        "Tools + MX2 parked out of the open work residual."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 32

    ws["A4"] = "1. Book now (after this morning)"
    ws["A4"].font = Font(bold=True, color="15233C")
    for i, h in enumerate(["Bucket", "PNs", "Billed €", "Share of book billed"], 1):
        c = ws.cell(5, i, h)
        c.fill = navy
        c.font = head
    rows1 = [
        ("HAS an AIM class (L1)", has_pn, has_eur, has_eur / book_eur),
        ("Still Not on AIM spine — total", not_pn, not_eur, not_eur / book_eur),
        ("  of which: open residual (work left)", open_pn, open_eur, open_eur / book_eur),
        ("  of which: special tools (parked)", tools_pn, tools_eur, tools_eur / book_eur),
        ("  of which: MX2 (parked)", mx2_pn, mx2_eur, mx2_eur / book_eur),
        ("BOOK TOTAL", book_pn, book_eur, 1.0),
    ]
    for r, row in enumerate(rows1, 6):
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1])
        ws.cell(r, 3, row[2]).number_format = money
        ws.cell(r, 4, row[3]).number_format = "0.0%"
        if row[0] == "BOOK TOTAL":
            for c in range(1, 5):
                ws.cell(r, c).font = Font(bold=True)

    ws["A14"] = "2. Classified this morning (Not-on-AIM open residual)"
    ws["A14"].font = Font(bold=True, color="15233C")
    for i, h in enumerate(["Step", "PNs", "Billed €", "Notes"], 1):
        c = ws.cell(15, i, h)
        c.fill = navy
        c.font = head
    rows2 = [
        (
            "Start — open residual after tools/MX2 out",
            start_pn,
            start_eur,
            "Before any apply today",
        ),
        (
            "Applied — high confidence (desc rules)",
            high_pn,
            high_eur,
            "First apply this morning",
        ),
        (
            "Applied — 5-signal agree3 + agree2",
            a23_pn,
            a23_eur,
            "agree3 28 / EUR 18.7k + agree2 247 / EUR 225.4k",
        ),
        (
            "Still open residual now",
            open_pn,
            open_eur,
            "Needs EQ / weak / conflict / strong-desc alone",
        ),
        (
            "Total classified this morning",
            high_pn + a23_pn,
            high_eur + a23_eur,
            "574 PNs onto AIM",
        ),
    ]
    for r, row in enumerate(rows2, 16):
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1])
        ws.cell(r, 3, row[2]).number_format = money
        ws.cell(r, 4, row[3])
        if "Total classified" in row[0] or "Still open" in row[0]:
            for c in range(1, 5):
                ws.cell(r, c).font = Font(bold=True)

    s4 = pd.read_csv(D / "pricing_not_on_aim_4signal_20260919.csv")
    left = s4[~s4["status"].isin(["agree3_high", "agree2_review"])].copy()
    ws["A23"] = "3. Still open — by 5-signal status (after agree2/3 apply)"
    ws["A23"].font = Font(bold=True, color="15233C")
    for i, h in enumerate(["Status", "PNs", "Billed €", "What it means"], 1):
        c = ws.cell(24, i, h)
        c.fill = navy
        c.font = head
    meaning = {
        "agree1_desc_strong": "One strong description vote only — not applied",
        "agree1_weak": "Soft / single weak vote (often dimension alone)",
        "conflict_needs_eq": "Signals disagree — EQ queue",
        "needs_eq": "No usable vote — EQ queue",
        "eq_last_resort": "Old EQ hit on disk — re-check tonight",
    }
    r = 25
    for status, g in left.groupby("status"):
        ws.cell(r, 1, status)
        ws.cell(r, 2, len(g))
        ws.cell(r, 3, g.billed_eur.sum()).number_format = money
        ws.cell(r, 4, meaning.get(status, ""))
        r += 1
    ws.cell(r, 1, "OPEN TOTAL")
    ws.cell(r, 2, len(left))
    ws.cell(r, 3, left.billed_eur.sum()).number_format = money
    for c in range(1, 4):
        ws.cell(r, c).font = Font(bold=True)

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 55
    ws.freeze_panes = "A6"
    wb.save(OUT)
    CMD.write_bytes(
        b'@echo off\r\nstart "" "%~dp0AIM_class_status_overview_20260919.xlsx"\r\n'
    )
    print(f"wrote {OUT}")
    print(f"open left {len(left)} / EUR {left.billed_eur.sum():,.1f}")


if __name__ == "__main__":
    main()
