# -*- coding: utf-8 -*-
"""
Build a PN -> Japanese-code map from the ISZJ price CD and measure what it covers
of the ISZE billed book.

Source : 01_data/【CY2026Price Revison】PRICE CD.xlsx   sheet 'PRICE CD 250912'
         cols: Pno, Parts name, New B/P, New Part Class, Current B/P, Current Part Class,
               EPLE, FIGKEY, Example Model, ZUK, HKK, ZKK
Spine  : 01_data/00_derived/aim_master_v2.csv  (has l12m_billed_eur)

Outputs:
  01_data/00_derived/jp_codes_cache_20260920.csv        one row per PN
  01_data/90_INBOX/_jp_codes_match_report_20260920.json coverage, by code

Read-only on the source. Writes only the two files above.

Run: python 01_data/91_build/extract_jp_classification_codes_20260920.py
"""
from __future__ import annotations

import collections
import csv
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "01_data" / "【CY2026Price Revison】PRICE CD.xlsx"
CACHE = ROOT / "01_data" / "00_derived" / "jp_codes_cache_20260920.csv"
REPORT = ROOT / "01_data" / "90_INBOX" / "_jp_codes_match_report_20260920.json"
SPINE = ROOT / "01_data" / "00_derived" / "aim_master_v2.csv"

# source column -> cache column
WANT = {
    "Pno": "pn",
    "Parts name": "jp_description",
    "FIGKEY": "jp_fig_key",
    "EPLE": "eple",
    "Example Model": "example_model",
    "ZUK": "zuk",
    "HKK": "hkk",
    "ZKK": "zkk",
    "New Part Class": "class_new",
    "Current Part Class": "class_current",
}
OUT_COLS = list(WANT.values())


def s(v) -> str:
    if v is None:
        return ""
    t = str(v).strip()
    return t[:-2] if t.endswith(".0") else t


def main() -> int:
    if not SRC.exists():
        print(f"MISSING source: {SRC}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [s(h) for h in next(it)]
    idx = {h: i for i, h in enumerate(hdr) if h}
    missing = [k for k in WANT if k not in idx]
    if missing:
        print(f"MISSING columns: {missing}", file=sys.stderr)
        wb.close()
        return 3

    CACHE.parent.mkdir(parents=True, exist_ok=True)
    blank = collections.Counter()
    n = 0
    with CACHE.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(OUT_COLS)
        for row in it:
            pn = s(row[idx["Pno"]])
            if not pn:
                continue
            n += 1
            out = []
            for src_col, _ in WANT.items():
                val = s(row[idx[src_col]])
                if not val:
                    blank[src_col] += 1
                out.append(val)
            w.writerow(out)
    wb.close()
    print(f"cached {n} PNs -> {CACHE.name}")

    # ---- the spine, and the billed book -----------------------------------------
    total_eur = 0.0
    spine: dict[str, float] = {}
    spine_rows = 0
    with SPINE.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            spine_rows += 1
            pn = s(r.get("part_number"))
            if not pn:
                continue
            e = float(r.get("l12m_billed_eur") or 0)
            spine[pn] = spine.get(pn, 0.0) + e
            total_eur += e
    print(f"spine {spine_rows} rows / {len(spine)} PNs / EUR {total_eur:,.0f} L12M")

    # ---- join --------------------------------------------------------------------
    def by(code_field: str):
        buckets: dict[str, list] = collections.defaultdict(lambda: [0, 0.0])
        cov_rows = cov_eur = 0
        seen = 0
        with CACHE.open(encoding="utf-8", newline="") as fh:
            for r in csv.DictReader(fh):
                seen += 1
                code = r[code_field]
                if not code or code == "9999":
                    continue
                if r["pn"] not in spine:
                    continue
                e = spine[r["pn"]]
                cov_rows += 1
                cov_eur += e
                d = buckets[code]
                d[0] += 1
                d[1] += e
        top = sorted(
            ({"code": k, "pns": v[0], "eur": round(v[1], 2)} for k, v in buckets.items()),
            key=lambda x: -x["eur"],
        )
        return {
            "field": code_field,
            "codes_seen": len(buckets),
            "spine_pns_covered": cov_rows,
            "spine_eur_covered": round(cov_eur, 2),
            "eur_share_pct": round(100 * cov_eur / total_eur, 2) if total_eur else 0.0,
            "pn_share_pct": round(100 * cov_rows / len(spine), 2) if spine else 0.0,
            "top_codes": top[:30],
            "cache_rows": seen,
        }

    report = {
        "source": str(SRC),
        "cache": str(CACHE),
        "cache_rows": n,
        "blank_counts": dict(blank),
        "spine": {"rows": spine_rows, "pns": len(spine), "l12m_eur": round(total_eur, 2)},
        "coverage": {f: by(f) for f in ("zkk", "zuk", "eple", "hkk", "jp_fig_key", "class_new")},
    }
    REPORT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print()
    print(f"{'field':<12} {'codes':>7} {'PNs hit':>8} {'EUR hit':>16} {'% of EUR':>9} {'% of PNs':>9}")
    for f, v in report["coverage"].items():
        print(
            f"{f:<12} {v['codes_seen']:>7} {v['spine_pns_covered']:>8} "
            f"{v['spine_eur_covered']:>16,.0f} {v['eur_share_pct']:>8.2f}% {v['pn_share_pct']:>8.2f}%"
        )
    print(f"\nreport -> {REPORT.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
