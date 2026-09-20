"""Normalize Romania (Isuzu Rom Auto SRL) sales-to-dealer monthly files -> canonical long CSV.

Native shape (one file per calendar month, 2026-01 .. 2026-08):
  - Title cell in column B: 'Parts and accessories invoiced in <Month> <Year>'
  - Header row (variable offset, 5-7): B = 'Part number', C = 'Parts and accessories invoiced/by quantity'
  - One row per part number with the invoiced quantity for that month

Rules:
  - distributor_code = ROMANIA (Isuzu Rom Auto SRL, SAP sold_to 396) - this is DB->dealer
    sell-in, NOT ISZE->DB billing. Never load into sales_fact.
  - period = month from the FILENAME (the native file carries no invoice date column)
  - qty = quantity in column C; skip zero/blank lines
  - grain = network_pn (no dealer dimension in source)
  - qty_only (no EUR in source); 2026 is a partial year (Jan-Aug)
  - PN = column B as-is. Numeric codes are zero-padded to 10 digits. Non-numeric
    short codes (accessory / service codes such as EAGTB2911B) cannot be mapped to the
    10-digit canonical key and are NOT silently dropped: they are written to a sidecar
    file so the excluded volume stays visible.

Outputs (into 01_data/00_derived):
  - dealer_sellout_long_ROMANIA_CY2026YTD.csv
  - dealer_sellout_unmapped_codes_ROMANIA_CY2026YTD.csv   (sidecar, excluded lines)

Note on the rollup grain: the YTD row is written with `period_grain="year"`, matching the
CY2025 files and the loader contract in modules/dealer_sales_loader.py, which selects rollup
rows with `period_grain == "year"` only. The PORTUGAL / IRELAND / IUK CY2026YTD adapters use
`"ytd"` instead, which that loader does not recognise (it then adds month + ytd together).
That is a pre-existing latent defect reported to Jerome on 2026-09-17; this adapter does not
rely on it. The period string still carries "-YTD" and the `partial_year` flag is set.
The sidecar is deliberately NOT named `dealer_sellout_long_*` so that it is not picked up by
the `dealer_sellout_long_*.csv` globs in the leakage tooling.

Target schema: 01_data/12_raw/dealer_sales/DEALER_SELL_OUT.md
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

FIELDS = [
    "distributor_code",
    "source_file",
    "grain",
    "year",
    "period",
    "period_grain",
    "part_number",
    "part_number_raw",
    "description",
    "qty",
    "dealer_name",
    "vehicle_scope",
    "flags",
]

UNMAPPED_FIELDS = ["source_file", "period", "part_number_raw", "qty", "reason"]

MONTH_RE = re.compile(r"(\d{4})-(\d{2})")
DISTRIBUTOR_CODE = "ROMANIA"
YEAR = 2026


def normalize_pn(raw) -> tuple[str | None, str]:
    """Canonical 10-digit key from a numeric Isuzu part number; else (None, raw)."""
    s = str(raw).strip() if raw is not None else ""
    digits = re.sub(r"\D", "", s)
    if not digits or len(digits) < 8:
        return None, s
    return digits.zfill(10), s


def _find_header(rows: list[tuple]) -> int:
    """Header row index: the row whose column B reads 'Part number'."""
    for i, r in enumerate(rows):
        for cell in r:
            if cell is not None and str(cell).strip().lower() == "part number":
                return i
    raise SystemExit("Could not locate 'Part number' header row")


def _month_from_name(path: Path) -> str:
    m = MONTH_RE.search(path.stem)
    if not m:
        raise SystemExit(f"Cannot read YYYY-MM from filename: {path.name}")
    y, mm = int(m.group(1)), int(m.group(2))
    if y != YEAR:
        raise SystemExit(f"Unexpected year {y} in {path.name} (expected {YEAR})")
    if not 1 <= mm <= 12:
        raise SystemExit(f"Bad month {mm} in {path.name}")
    return f"{y}-{mm:02d}"


def normalize_romania(src_dir: Path, out_dir: Path) -> dict[str, Path]:
    import openpyxl

    files = sorted(src_dir.glob(f"*Sales_to_Dealer_{YEAR}-*.xlsx"))
    if not files:
        raise SystemExit(f"No Romania monthly files found in {src_dir}")

    # (period, pn) -> qty ; keep first raw spelling per PN
    agg: dict[tuple[str, str], float] = defaultdict(float)
    raw_first: dict[str, str] = {}
    unmapped: list[dict] = []
    per_month: dict[str, dict] = {}

    for path in files:
        period = _month_from_name(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        hdr = _find_header(rows)
        # column B = part number, column C = quantity (native layout is fixed)
        kept = skipped_blank = skipped_qty = 0
        for r in rows[hdr + 1:]:
            if len(r) < 3:
                continue
            pn_cell, qty_cell = r[1], r[2]
            if pn_cell is None or str(pn_cell).strip() == "":
                skipped_blank += 1
                continue
            try:
                qty = float(qty_cell or 0)
            except (TypeError, ValueError):
                skipped_qty += 1
                continue
            if qty == 0:
                skipped_qty += 1
                continue
            pn, raw = normalize_pn(pn_cell)
            if not pn:
                unmapped.append(
                    dict(
                        source_file=path.name,
                        period=period,
                        part_number_raw=raw,
                        qty=qty,
                        reason="not a 10-digit Isuzu part number (accessory/service code)",
                    )
                )
                continue
            agg[(period, pn)] += qty
            raw_first.setdefault(pn, raw)
            kept += 1

        per_month[period] = dict(lines_kept=kept, skipped_blank=skipped_blank,
                                 skipped_qty=skipped_qty)

    # ---- main long rows: month grain + YTD rollup
    flags = "qty_only|network_pn|partial_year"
    out_rows: list[dict] = []
    for (period, pn), q in sorted(agg.items()):
        out_rows.append(
            dict(
                distributor_code=DISTRIBUTOR_CODE,
                source_file=f"{DISTRIBUTOR_CODE}_Sales_to_Dealer_{period}.xlsx",
                grain="network_pn",
                year=YEAR,
                period=period,
                period_grain="month",
                part_number=pn,
                part_number_raw=raw_first.get(pn, pn),
                description="",
                qty=q,
                dealer_name="",
                vehicle_scope="unknown",
                flags=flags,
            )
        )

    ytd_totals: dict[str, float] = defaultdict(float)
    for (_, pn), q in agg.items():
        ytd_totals[pn] += q
    for pn, q in sorted(ytd_totals.items()):
        out_rows.append(
            dict(
                distributor_code=DISTRIBUTOR_CODE,
                source_file=f"{DISTRIBUTOR_CODE}_Sales_to_Dealer_{YEAR}YTD_Jan-Aug.xlsx",
                grain="network_pn",
                year=YEAR,
                period=f"{YEAR}-YTD",
                # "year" (not "ytd"): the canonical loader selects rollup rows with
                # period_grain == "year" only. See module docstring.
                period_grain="year",
                part_number=pn,
                part_number_raw=raw_first.get(pn, pn),
                description="",
                qty=q,
                dealer_name="",
                vehicle_scope="unknown",
                flags=flags,
            )
        )

    out_dir.mkdir(parents=True, exist_ok=True)
    out_main = out_dir / f"dealer_sellout_long_{DISTRIBUTOR_CODE}_CY{YEAR}YTD.csv"
    with out_main.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(out_rows)

    # Sidecar name is deliberately NOT 'dealer_sellout_long_*' so the wildcard globs in the
    # leakage tooling do not mistake it for another distributor's long file.
    out_unmapped = out_dir / f"dealer_sellout_unmapped_codes_{DISTRIBUTOR_CODE}_CY{YEAR}YTD.csv"
    with out_unmapped.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=UNMAPPED_FIELDS)
        w.writeheader()
        w.writerows(sorted(unmapped, key=lambda r: (r["period"], r["part_number_raw"])))

    # ---- report
    n_month = sum(1 for r in out_rows if r["period_grain"] == "month")
    n_ytd = sum(1 for r in out_rows if r["period_grain"] == "year")
    qty_month = sum(r["qty"] for r in out_rows if r["period_grain"] == "month")
    qty_unmapped = sum(r["qty"] for r in unmapped)
    months = sorted(per_month)
    print(f"months read      : {len(months)} ({months[0]} .. {months[-1]})")
    for m in months:
        s = per_month[m]
        mq = sum(q for (p, _), q in agg.items() if p == m)
        print(f"  {m}: lines={s['lines_kept']:4d}  qty={mq:6,.0f}")
    print(f"main CSV         : {out_main.name}")
    print(f"  month rows     : {n_month}")
    print(f"  ytd rows       : {n_ytd}")
    print(f"  unique PNs     : {len(ytd_totals):,}")
    print(f"  total qty      : {qty_month:,.0f}")
    print(f"unmapped sidecar : {out_unmapped.name}")
    print(f"  excluded lines : {len(unmapped)}  qty={qty_unmapped:,.0f} "
          f"({100 * qty_unmapped / (qty_month + qty_unmapped):.2f}% of source qty)")
    return {"main": out_main, "unmapped": out_unmapped}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--src-dir",
        type=Path,
        default=Path("01_data/12_raw/dealer_sales/ROMANIA"),
    )
    ap.add_argument("--out-dir", type=Path, default=Path("01_data/00_derived"))
    args = ap.parse_args()
    src, out_dir = args.src_dir, args.out_dir
    if not src.is_absolute():
        root = Path(__file__).resolve().parents[2]
        src = root / src
        out_dir = root / out_dir
    normalize_romania(src, out_dir)


if __name__ == "__main__":
    main()
