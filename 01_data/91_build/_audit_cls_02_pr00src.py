#!/usr/bin/env python3
"""Audit step 2: open the PR00 source workbooks that the classification came from."""
from __future__ import annotations
from pathlib import Path
import openpyxl, sys

P = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing")
targets = [
    "IE40 Price list as of 13.03.2026 PR00_3_level_classified_GPT.xlsx",
    "Price_file_PR00_IE40_20260810.xlsx",
]
for name in targets:
    hits = list(P.glob(f"**/{name}"))
    if not hits:
        print(f"!! not found: {name}")
        continue
    p = hits[0]
    print("=" * 78)
    print(p)
    print(f"{p.stat().st_size/1e6:.1f} MB")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"\n  sheet {ws.title!r}  dims={ws.calculate_dimension()}  "
              f"max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True)):
            cells = ["" if c is None else str(c)[:34] for c in row[:24]]
            print(f"    r{i+1}: {cells}")
    wb.close()
