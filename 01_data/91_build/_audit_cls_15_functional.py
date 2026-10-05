#!/usr/bin/env python3
"""Audit step 15: there is already a functional taxonomy on disk. What is in it?"""
from __future__ import annotations
from pathlib import Path
import openpyxl

P = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\ISZE_Parts_Sales_Classification_Mar2026_v2.xlsx")
print(P, f"{P.stat().st_size/1e6:.1f} MB")
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== {ws.title!r}  max_row={ws.max_row} max_col={ws.max_column} ===")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True)):
        cells = ["" if c is None else str(c)[:30] for c in row[:10]]
        if any(c.strip() for c in cells):
            print(f"  r{i+1}: {cells}")
wb.close()
