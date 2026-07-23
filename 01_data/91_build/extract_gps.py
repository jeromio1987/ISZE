import openpyxl, sys, json
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

gps = r"C:\Users\jerom\Desktop\claude\ISZE\Isuzu\Divers\ISZE_GPS_Calculator_v3.xlsx"
wb = openpyxl.load_workbook(gps, data_only=True)

ws = wb["GPS_Results_2026"]
rows = []
for r in ws.iter_rows(min_row=5, values_only=True):
    if r[0] is None:
        continue
    rows.append({"db": r[0], "year": r[1], "uio_cv_ld": r[2], "uio_f": r[3], "uio_md": r[4],
                 "uio_lcv": r[5], "uio_total": r[6], "wallet": r[7], "sales": r[8],
                 "spu": r[9], "gps": r[10]})
print(json.dumps(rows))
print("---EXWARR---")
ws2 = wb["GPS_ex_Warranty"]
for r in ws2.iter_rows(min_row=1, max_row=min(ws2.max_row, 15), max_col=12, values_only=True):
    vals = ["" if v is None else str(v)[:26] for v in r]
    if any(v != "" for v in vals):
        print(" | ".join(vals))
for r in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=12, values_only=True):
    if r[0] and "ITUK" in str(r[0]):
        print(r)
