# -*- coding: utf-8 -*-
"""Complete PN list: old BP missing / new BP missing for the price simulation.

Source: price_bp_weighted_compare_pn.csv (L12M billing window, in-scope network).
  old BP = bp_prior_revised  (CY2026 revision PRICE CD 250912)
  new BP = bp_new_jpy        (PRICE_CD_20260910)

Outputs one workbook with Overview + three lists (old missing, new missing,
complete gap union). Jerome fills / requests the missing base prices for the
simulation file.

Run: python 01_data/91_build/build_bp_missing_for_simulation_20260920.py
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
D = ISZE / "01_data" / "00_derived"
P = ISZE / "01_data" / "01_Pricing"
SRC = D / "price_bp_weighted_compare_pn.csv"
OUT = P / "BP_missing_for_price_simulation_20260920.xlsx"
CMD = P / "OPEN_BP_missing_for_price_simulation_20260920.cmd"
CSV_OLD = D / "bp_missing_old_for_simulation_20260920.csv"
CSV_NEW = D / "bp_missing_new_for_simulation_20260920.csv"
CSV_ALL = D / "bp_missing_complete_for_simulation_20260920.csv"

NAVY = "15233C"
MONEY = '€#,##0.00'


def _sheet(ws, df: pd.DataFrame, title: str, subtitle: str) -> None:
    navy = PatternFill("solid", fgColor=NAVY)
    head = Font(bold=True, color="FFFFFF")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(8, max(1, len(df.columns))))
    ws["A2"] = subtitle
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(8, max(1, len(df.columns))))
    ws.row_dimensions[2].height = 36

    start = 4
    for r_i, row in enumerate(dataframe_to_rows(df, index=False, header=True), start):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            if r_i == start:
                cell.fill = navy
                cell.font = head
            elif isinstance(val, float) and df.columns[c_i - 1] in (
                "billed_eur",
                "billed_qty",
                "bp_new_jpy",
                "bp_prior_revised",
            ):
                if "eur" in df.columns[c_i - 1]:
                    cell.number_format = MONEY
                elif "bp_" in df.columns[c_i - 1]:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0.00"

    widths = {
        "part_number": 14,
        "description": 42,
        "gap": 18,
        "unpriced_reason": 28,
        "billed_eur": 12,
        "billed_qty": 12,
        "bp_new_jpy": 12,
        "bp_prior_revised": 16,
        "n_customers": 12,
        "distributors": 40,
        "what_is_needed": 55,
    }
    for i, col in enumerate(df.columns, 1):
        letter = chr(64 + i) if i <= 26 else "A"
        # openpyxl column letter for >26
        from openpyxl.utils import get_column_letter

        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(df.columns))}{start + len(df)}"


def main() -> None:
    raw = pd.read_csv(SRC, low_memory=False)
    ins = raw[raw["in_scope"] == True].copy()  # noqa: E712

    # Prefer the most common unpriced_reason per PN (by euro)
    reason = (
        ins.groupby(["part_number", "unpriced_reason"], dropna=False)["net_sales_value"]
        .sum()
        .reset_index()
        .sort_values("net_sales_value", ascending=False)
        .drop_duplicates("part_number")
        [["part_number", "unpriced_reason"]]
    )

    g = (
        ins.groupby("part_number", as_index=False)
        .agg(
            description=("billed_description", "first"),
            billed_eur=("net_sales_value", "sum"),
            billed_qty=("billed_qty", "sum"),
            bp_new_jpy=("bp_new_jpy", "max"),
            bp_prior_revised=("bp_prior_revised", "max"),
            n_customers=("customer_code", "nunique"),
            distributors=(
                "gps_db_label",
                lambda s: ", ".join(sorted({str(x) for x in s if pd.notna(x) and str(x).strip()})[:10]),
            ),
        )
    )
    g = g.merge(reason, on="part_number", how="left")

    miss_old = g["bp_prior_revised"].isna()
    miss_new = g["bp_new_jpy"].isna()
    both_ok = (~miss_old) & (~miss_new)

    def gap_label(r) -> str:
        o, n = pd.isna(r["bp_prior_revised"]), pd.isna(r["bp_new_jpy"])
        if o and n:
            return "BOTH missing"
        if o:
            return "OLD BP missing"
        if n:
            return "NEW BP missing"
        return "OK"

    def need(r) -> str:
        o, n = pd.isna(r["bp_prior_revised"]), pd.isna(r["bp_new_jpy"])
        if o and n:
            return "Need prior (old) base price JPY AND new PRICE CD base price JPY"
        if o:
            return "Need prior (old) base price JPY — new BP already present"
        if n:
            return "Need new PRICE CD base price JPY — old BP already present"
        return ""

    gaps = g[miss_old | miss_new].copy()
    gaps["gap"] = gaps.apply(gap_label, axis=1)
    gaps["what_is_needed"] = gaps.apply(need, axis=1)
    gaps = gaps.sort_values(["gap", "billed_eur"], ascending=[True, False])

    old_list = gaps[gaps["gap"].isin(["OLD BP missing", "BOTH missing"])].copy()
    new_list = gaps[gaps["gap"].isin(["NEW BP missing", "BOTH missing"])].copy()

    cols = [
        "part_number",
        "description",
        "gap",
        "what_is_needed",
        "billed_eur",
        "billed_qty",
        "bp_prior_revised",
        "bp_new_jpy",
        "unpriced_reason",
        "n_customers",
        "distributors",
    ]
    old_list = old_list[cols]
    new_list = new_list[cols]
    complete = gaps[cols]

    old_list.to_csv(CSV_OLD, index=False)
    new_list.to_csv(CSV_NEW, index=False)
    complete.to_csv(CSV_ALL, index=False)

    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "Base prices missing for the price simulation"
    ov["A1"].font = Font(bold=True, size=14, color=NAVY)
    ov["A2"] = (
        "In-scope pricing network, L12M billing 1 Jul 2025 – 30 Jun 2026. "
        "Old BP = prior revision (PRICE CD 250912). New BP = PRICE CD 20260910. "
        "One row per part number. Fill the missing JPY base price(s) into the simulation file."
    )
    ov["A2"].alignment = Alignment(wrap_text=True)
    ov.merge_cells("A2:D2")
    ov.row_dimensions[2].height = 48

    navy = PatternFill("solid", fgColor=NAVY)
    head = Font(bold=True, color="FFFFFF")
    for i, h in enumerate(["Bucket", "PNs", "Billed €", "Note"], 1):
        c = ov.cell(4, i, h)
        c.fill = navy
        c.font = head

    n_ok = int(both_ok.sum())
    eur_ok = float(g.loc[both_ok, "billed_eur"].sum())
    rows = [
        ("Book in scope (unique PN)", len(g), float(g.billed_eur.sum()), "All in-scope billed part numbers"),
        ("OK — old and new BP both present", n_ok, eur_ok, "Simulation pair exists"),
        (
            "OLD BP missing (incl. both)",
            len(old_list),
            float(old_list.billed_eur.sum()),
            "Need prior / old base price JPY",
        ),
        (
            "NEW BP missing (incl. both)",
            len(new_list),
            float(new_list.billed_eur.sum()),
            "Need new PRICE CD base price JPY",
        ),
        (
            "BOTH missing",
            int((gaps["gap"] == "BOTH missing").sum()),
            float(gaps.loc[gaps["gap"] == "BOTH missing", "billed_eur"].sum()),
            "Need both sides",
        ),
        (
            "COMPLETE gap list (union)",
            len(complete),
            float(complete.billed_eur.sum()),
            "Every PN missing at least one side",
        ),
    ]
    for r_i, row in enumerate(rows, 5):
        ov.cell(r_i, 1, row[0])
        ov.cell(r_i, 2, row[1])
        ov.cell(r_i, 3, row[2]).number_format = MONEY
        ov.cell(r_i, 4, row[3])

    ov["A12"] = "By gap type (complete list)"
    ov["A12"].font = Font(bold=True, color=NAVY)
    for i, h in enumerate(["Gap", "PNs", "Billed €"], 1):
        c = ov.cell(13, i, h)
        c.fill = navy
        c.font = head
    r = 14
    for gap, sub in complete.groupby("gap"):
        ov.cell(r, 1, gap)
        ov.cell(r, 2, len(sub))
        ov.cell(r, 3, float(sub.billed_eur.sum())).number_format = MONEY
        r += 1

    ov["A18"] = "By unpriced reason (complete list)"
    ov["A18"].font = Font(bold=True, color=NAVY)
    for i, h in enumerate(["Reason", "PNs", "Billed €"], 1):
        c = ov.cell(19, i, h)
        c.fill = navy
        c.font = head
    r = 20
    for reason, sub in complete.groupby(complete["unpriced_reason"].fillna("(blank)")):
        ov.cell(r, 1, reason)
        ov.cell(r, 2, len(sub))
        ov.cell(r, 3, float(sub.billed_eur.sum())).number_format = MONEY
        r += 1

    ov.column_dimensions["A"].width = 42
    ov.column_dimensions["B"].width = 10
    ov.column_dimensions["C"].width = 14
    ov.column_dimensions["D"].width = 48

    ws_old = wb.create_sheet("Missing_old_BP")
    _sheet(
        ws_old,
        old_list,
        "Part numbers with OLD (prior) base price missing",
        "Includes PNs that are also missing the new BP (gap = BOTH missing). Sorted by billed euro.",
    )
    ws_new = wb.create_sheet("Missing_new_BP")
    _sheet(
        ws_new,
        new_list,
        "Part numbers with NEW base price missing",
        "Includes PNs that are also missing the old BP (gap = BOTH missing). Sorted by billed euro.",
    )
    ws_all = wb.create_sheet("Complete_gaps")
    _sheet(
        ws_all,
        complete,
        "Complete list — every PN missing old BP, new BP, or both",
        "Use the gap column to filter. what_is_needed says exactly which JPY base price to put in the simulation file.",
    )

    wb._sheets = [ov, ws_old, ws_new, ws_all]
    wb.save(OUT)
    CMD.write_bytes(
        b'@echo off\r\nstart "" "%~dp0BP_missing_for_price_simulation_20260920.xlsx"\r\n'
    )
    print(f"wrote {OUT}")
    print(f"old missing {len(old_list)} / EUR {old_list.billed_eur.sum():,.1f}")
    print(f"new missing {len(new_list)} / EUR {new_list.billed_eur.sum():,.1f}")
    print(f"complete    {len(complete)} / EUR {complete.billed_eur.sum():,.1f}")
    print(complete.groupby("gap").agg(n=("part_number", "count"), eur=("billed_eur", "sum")).to_string())


if __name__ == "__main__":
    main()
