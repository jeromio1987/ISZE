import openpyxl, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ITUK UIO projection sheet from UIO v17
wb = openpyxl.load_workbook(r"C:\Users\jerom\Desktop\claude\ISZE\Data\ISUZU_UIO_Overview_2026_v17.xlsx", data_only=True, read_only=True)
ws = wb["4. ITUK UIO Projection"]
rows = list(ws.iter_rows(values_only=True))
start = next(i for i, r in enumerate(rows) if any("Section C" in str(c) for c in r if c))
for r in rows[start:start+16]:
    vals = ["" if v is None else str(v)[:24] for v in r[:9]]
    if any(v != "" for v in vals):
        print(" | ".join(vals))
wb.close()

# Warranty sums
wb = openpyxl.load_workbook(r"C:\Users\jerom\Desktop\claude\ISZE\Data\Warranty_CV_2025.xlsx", read_only=True)
for nm, qcol, pcol in [("ITUK", 2, 3), ("UTI", 2, 3)]:
    ws = wb[nm]
    tot = n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        q, p = r[qcol], r[pcol]
        if isinstance(q, (int, float)) and isinstance(p, (int, float)):
            tot += q * p; n += 1
    print(f"{nm}: {n} PNs, warranty parts value EUR {tot:,.2f}")

# Sheet4 per-country
ws = wb["Sheet4"]
import re
cur = None; sums = {}
for r in ws.iter_rows(min_row=2, max_col=3, values_only=True):
    lbl, cnt, val = r[0], r[1], r[2]
    if lbl is None: continue
    s = str(lbl).strip()
    if not re.match(r"^\d{8,}", s) and not s.lower().startswith("grand"):
        cur = s; sums.setdefault(cur, 0.0); continue
    if cur and cnt is not None:
        v = val
        if isinstance(v, str):
            v = v.replace(" ", "").replace(" ", "").replace(",", ".")
            try: v = float(v)
            except ValueError: v = None
        if isinstance(v, (int, float)):
            try: c = float(cnt)
            except (TypeError, ValueError): c = 1
            sums[cur] += c * v
for k, v in sums.items():
    print(f"Sheet4 {k}: {v:,.2f}")
wb.close()
