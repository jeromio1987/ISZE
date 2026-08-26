"""Harvest cover evaluation dates for NL/DK/LUX incoming + existing Huberty."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\IDS_2026")
STAGE = ROOT / "_INCOMING_2026-07-24"


def cover_date(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    dealer = None
    date_v = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, max_col=8, values_only=True), 1):
        label = str(row[1] or "").strip().lower()
        val = row[4]
        if any(k in label for k in ("dealer", "concessionnaire", "naam dealer")):
            dealer = str(val).strip() if val is not None else None
        if any(k in label for k in ("date", "datum")):
            date_v = val
    wb.close()
    return dealer, date_v


def show(label: str, paths: list[Path]) -> None:
    print(f"\n=== {label} ({len(paths)}) ===")
    rows = []
    for p in paths:
        dealer, date_v = cover_date(p)
        rows.append((str(date_v), dealer or "?", p.name))
    for r in sorted(rows, key=lambda x: (str(x[0]), x[2])):
        print(f"  {r[0]!s:28}  {r[1]!s:40}  {r[2]}")


def main() -> None:
    nl = sorted((STAGE / "NL").rglob("*.xlsx"))
    dk = sorted(
        p
        for p in (STAGE / "Denmark").rglob("*.xlsx")
        if "grafiek" not in p.name.lower()
    )
    lux = sorted((STAGE / "LUX").rglob("*.xlsx"))
    huberty = list((ROOT / "BE").rglob("*Huberty*.xlsx"))
    show("NL incoming", nl)
    show("DK dealers incoming", dk)
    show("LUX incoming", lux)
    show("Existing Huberty (BE)", huberty)


if __name__ == "__main__":
    main()
