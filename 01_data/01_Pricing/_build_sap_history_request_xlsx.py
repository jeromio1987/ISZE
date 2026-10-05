"""Build SAP_Pricing_History_Data_Request_by_Year.xlsx for Jerome's dated SAP pulls.

Three validity windows (European dates) × PR00 + ZPR1.
Parts_To_Pull = one row per Spec C4 missing PN with Y/N pull flags per window.

shared. Uses derived parquet/CSV only (never raw sales XLSX).
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def sanitize_xlsx_for_excel_desktop(path: Path) -> None:
    """Patch openpyxl OOXML quirks that make Excel Desktop prompt Repair / hang.

    Fixes seen vs Excel SaveAs repair:
    - empty <workbookProtection/>
    - <patternFill/> without patternType=\"none\"
    - rgb=\"00RRGGBB\" (alpha 00 = transparent) -> FF
    - absolute Target=\"/xl/...\" relationships -> relative
    - spurious pivotButton/quotePrefix on xf
    """
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    wb_xml = files["xl/workbook.xml"].decode("utf-8")
    wb_xml = re.sub(r"<workbookProtection\b[^/]*/>", "", wb_xml)
    wb_xml = re.sub(
        r"<workbookProtection\b[^>]*>.*?</workbookProtection>",
        "",
        wb_xml,
        flags=re.DOTALL,
    )
    files["xl/workbook.xml"] = wb_xml.encode("utf-8")

    styles = files["xl/styles.xml"].decode("utf-8")
    styles = styles.replace("<patternFill/>", '<patternFill patternType="none"/>')
    styles = styles.replace("<patternFill />", '<patternFill patternType="none"/>')
    styles = re.sub(r'rgb="00([0-9A-Fa-f]{6})"', r'rgb="FF\1"', styles)
    styles = re.sub(r'\s+pivotButton="0"', "", styles)
    styles = re.sub(r'\s+quotePrefix="0"', "", styles)
    files["xl/styles.xml"] = styles.encode("utf-8")

    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rels = rels.replace('Target="/xl/worksheets/', 'Target="worksheets/')
    rels = rels.replace('Target="/xl/theme/', 'Target="theme/')
    rels = rels.replace('Target="/xl/styles.xml"', 'Target="styles.xml"')
    rels = rels.replace('Target="/xl/sharedStrings.xml"', 'Target="sharedStrings.xml"')
    files["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    tmp = path.with_suffix(".sanitize.tmp.xlsx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)

ROOT = Path(__file__).resolve().parents[2]
PRICING = Path(__file__).resolve().parent
DER = ROOT / "01_data" / "00_derived"
OUT = PRICING / "SAP_Pricing_History_Data_Request_by_Year.xlsx"
RANKED_CSV = DER / "no_pr00_gap_cy2025_ranked.csv"
RANKED_META = DER / "no_pr00_gap_cy2025_ranked.meta.json"

# Jerome's windows (European DD/MM/YYYY). Prefer year-end for Window_2023;
# literal alternate 31/01/2023 recorded in How_To_Pull (typo ambiguity on 310122023).
WINDOWS = [
    {
        "label": "Window_2023",
        "valid_from": "01/01/2023",
        "valid_to": "31/12/2023",
        "valid_to_alternate": "31/01/2023",
        "note": (
            "Prefer Valid To 31/12/2023 (year-end). "
            "Jerome wrote 310122023 (9 digits) — also record alternate 31/01/2023 if that was intended."
        ),
        "year_match": 2023,
    },
    {
        "label": "Window_2024",
        "valid_from": "01/01/2024",
        "valid_to": "31/01/2025",
        "valid_to_alternate": None,
        "note": "As written 01012024-31012025.",
        "year_match": 2024,
    },
    {
        "label": "Window_2025",
        "valid_from": "01/02/2025",
        "valid_to": "17/03/2026",
        "valid_to_alternate": None,
        "note": "As written 01022025-17032026 (~ Mar 2026 IE40 snapshot era). Primary / P0 window.",
        "year_match": 2025,
    },
]

CONDITIONS = ("PR00", "ZPR1")

thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
title_font = Font(bold=True, name="Calibri", size=14, color="1F4E79")
subtitle_font = Font(name="Calibri", size=10, italic=True, color="555555")
wrap = Alignment(wrap_text=True, vertical="top")
p0_fill = PatternFill("solid", fgColor="FCE4D6")
p1_fill = PatternFill("solid", fgColor="FFF2CC")
p2_fill = PatternFill("solid", fgColor="E2EFDA")
y_fill = PatternFill("solid", fgColor="C6EFCE")
n_fill = PatternFill("solid", fgColor="F2F2F2")
have_fill = PatternFill("solid", fgColor="C6EFCE")
missing_fill = PatternFill("solid", fgColor="F8CBAD")
partial_fill = PatternFill("solid", fgColor="FCE4D6")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_rows(ws, start_row, headers, rows, row_height=28, yn_cols=None):
    yn_cols = set(yn_cols or [])
    for j, h in enumerate(headers, 1):
        ws.cell(start_row, j, h)
    style_header(ws, start_row, len(headers))
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for j, val in enumerate(row, 1):
            cell = ws.cell(r, j, val)
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="Calibri", size=10)
            h = headers[j - 1]
            if h in yn_cols:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if val == "Y":
                    cell.fill = y_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="006100")
                elif val == "N":
                    cell.fill = n_fill
            if h == "Priority":
                if val == "P0":
                    cell.fill = p0_fill
                elif val == "P1":
                    cell.fill = p1_fill
                elif val == "P2":
                    cell.fill = p2_fill
            if h == "Status":
                if val == "missing":
                    cell.fill = missing_fill
                elif val == "partial":
                    cell.fill = partial_fill
                elif val == "have":
                    cell.fill = have_fill
                elif val == "to_pull":
                    cell.fill = missing_fill
        ws.row_dimensions[r].height = row_height
    return start_row + 1 + len(rows)


def norm_pn(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, int):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    digits = re.sub(r"\D", "", s)
    return digits.lstrip("0") or ("0" if digits else "")


def format_pn_10(pn: str) -> str:
    """Pad numeric PN to 10 digits for SAP Material display when short."""
    if not pn or not str(pn).isdigit():
        return str(pn)
    s = str(pn)
    return s.zfill(10) if len(s) <= 10 else s


def priority_from_eur(eur: float) -> str:
    if eur >= 10_000:
        return "P0"
    if eur >= 1_000:
        return "P1"
    return "P2"


def load_pr00_sets(year: int) -> tuple[set[str], set[str], pd.DataFrame | None]:
    if year >= 2026:
        df = pd.read_parquet(DER / "pricing_pr00_ie40.parquet")
        df = df.rename(columns={"part_number": "Material", "sp_eur": "Price_EUR"})
        df["pn"] = df["Material"].map(norm_pn)
        df["Price_EUR"] = pd.to_numeric(df["Price_EUR"], errors="coerce")
        on_list = set(df.loc[df["pn"] != "", "pn"])
        if "is_superseded_row" in df.columns:
            active = df[~df["is_superseded_row"].fillna(False)].copy()
        else:
            active = df
        usable = set(
            active.loc[
                active["Price_EUR"].notna() & (active["Price_EUR"] > 0) & (active["pn"] != ""),
                "pn",
            ]
        )
        return on_list, usable, df

    path = DER / f"pricing_pr00_ie40_{year}.parquet"
    df = pd.read_parquet(path)
    df["pn"] = df["Material"].map(norm_pn)
    df["Price_EUR"] = pd.to_numeric(df["Price_EUR"], errors="coerce")
    on_list = set(df.loc[df["pn"] != "", "pn"])
    usable = set(
        df.loc[df["Price_EUR"].notna() & (df["Price_EUR"] > 0) & (df["pn"] != ""), "pn"]
    )
    return on_list, usable, None


def classify_vs_current(pn: str, pr_cur: pd.DataFrame, usable: set[str]) -> str:
    rows = pr_cur[pr_cur["pn"] == pn]
    if rows.empty:
        return "absent_from_pr00"
    if pn in usable:
        return "HAS_USABLE"
    if "is_superseded_row" in rows.columns:
        has_active = (~rows["is_superseded_row"].fillna(False)).any()
        if not has_active:
            return "superseded_only"
    return "on_list_no_usable_price"


def year_matched_missing_sets() -> dict[int, set[str]]:
    """PNs billed in year with no usable year-matched PR00."""
    sf = pd.read_parquet(
        DER / "sales_fact.parquet",
        columns=["part_number", "year", "net_sales_value"],
    )
    sf["pn"] = sf["part_number"].map(norm_pn)
    sf = sf[sf["pn"] != ""]
    sf["net_sales_value"] = pd.to_numeric(sf["net_sales_value"], errors="coerce").fillna(0.0)
    out: dict[int, set[str]] = {}
    for year in (2023, 2024, 2025):
        _, usable, _ = load_pr00_sets(year)
        y = sf[sf["year"] == year]
        g = y.groupby("pn", as_index=False).agg(eur=("net_sales_value", "sum"))
        g = g[g["eur"] > 0]
        out[year] = set(g.loc[~g["pn"].isin(usable), "pn"])
    return out


def build_parts_frame() -> tuple[pd.DataFrame, dict]:
    ranked = pd.read_csv(RANKED_CSV)
    meta = json.loads(RANKED_META.read_text(encoding="utf-8")) if RANKED_META.exists() else {}
    on26, us26, pr_cur = load_pr00_sets(2026)
    assert pr_cur is not None

    ranked = ranked.copy()
    ranked["pn"] = ranked["part_number"].map(norm_pn)
    ranked["reason"] = ranked["pn"].map(lambda p: classify_vs_current(p, pr_cur, us26))
    ranked["EUR_CY2025"] = pd.to_numeric(ranked["cy2025_eur"], errors="coerce").fillna(0.0)
    ranked["Priority"] = ranked["EUR_CY2025"].map(priority_from_eur)
    ranked["part_number_10"] = ranked["pn"].map(format_pn_10)

    ym = year_matched_missing_sets()

    # Window_2025: all Spec C4 → Y for PR00 and ZPR1
    # Window_2023/2024: Y if year-matched gap OR top € (P0 ≥ €10k)
    def flag_historic(pn: str, year: int, eur: float) -> str:
        if pn in ym.get(year, set()):
            return "Y"
        if eur >= 10_000:
            return "Y"
        return "N"

    ranked["Pull_PR00_2023"] = [
        flag_historic(p, 2023, e) for p, e in zip(ranked["pn"], ranked["EUR_CY2025"])
    ]
    ranked["Pull_ZPR1_2023"] = ranked["Pull_PR00_2023"]
    ranked["Pull_PR00_2024"] = [
        flag_historic(p, 2024, e) for p, e in zip(ranked["pn"], ranked["EUR_CY2025"])
    ]
    ranked["Pull_ZPR1_2024"] = ranked["Pull_PR00_2024"]
    ranked["Pull_PR00_2025"] = "Y"
    ranked["Pull_ZPR1_2025"] = "Y"

    def why_row(reason: str, p23: str, p24: str) -> str:
        base = {
            "absent_from_pr00": "Absent from current IE40 PR00 — need usable list price",
            "on_list_no_usable_price": "On list but no usable price (blank/0) — need Amount>0",
            "superseded_only": "Only superseded on current list — need active/successor price",
        }.get(reason, reason)
        bits = [base, "Window_2025 = P0 (all Spec C4)"]
        if p23 == "Y":
            bits.append("also year-gap or top-€ for Window_2023")
        if p24 == "Y":
            bits.append("also year-gap or top-€ for Window_2024")
        return "; ".join(bits)

    ranked["Reason"] = [
        why_row(r, a, b)
        for r, a, b in zip(
            ranked["reason"], ranked["Pull_PR00_2023"], ranked["Pull_PR00_2024"]
        )
    ]
    ranked["In_year_gap_2023"] = ranked["pn"].map(lambda p: "Y" if p in ym[2023] else "N")
    ranked["In_year_gap_2024"] = ranked["pn"].map(lambda p: "Y" if p in ym[2024] else "N")
    ranked["In_year_gap_2025"] = ranked["pn"].map(lambda p: "Y" if p in ym[2025] else "N")

    ranked = ranked.sort_values(["Priority", "EUR_CY2025"], ascending=[True, False])

    stats = {
        "ranked_pn_count": int(len(ranked)),
        "ranked_eur": round(float(ranked["EUR_CY2025"].sum()), 2),
        "by_reason": ranked["reason"].value_counts().to_dict(),
        "published_522_note": meta.get("published_2026_08_01", {}),
        "ym_counts": {y: len(s) for y, s in ym.items()},
        "pull_counts": {
            "Window_2023_PR00": int((ranked["Pull_PR00_2023"] == "Y").sum()),
            "Window_2023_ZPR1": int((ranked["Pull_ZPR1_2023"] == "Y").sum()),
            "Window_2024_PR00": int((ranked["Pull_PR00_2024"] == "Y").sum()),
            "Window_2024_ZPR1": int((ranked["Pull_ZPR1_2024"] == "Y").sum()),
            "Window_2025_PR00": int((ranked["Pull_PR00_2025"] == "Y").sum()),
            "Window_2025_ZPR1": int((ranked["Pull_ZPR1_2025"] == "Y").sum()),
        },
        "any_pull_2023": int((ranked["Pull_PR00_2023"] == "Y").sum()),
        "any_pull_2024": int((ranked["Pull_PR00_2024"] == "Y").sum()),
        "any_pull_2025": int((ranked["Pull_PR00_2025"] == "Y").sum()),
        "on26_usable": len(us26),
        "on26_any": len(on26),
    }
    return ranked, stats


def main() -> None:
    parts, stats = build_parts_frame()

    wb = Workbook()

    # ── Parts_To_Pull (main — first tab) ────────────────────────────────
    ws = wb.active
    ws.title = "Parts_To_Pull"
    ws["A1"] = "Parts to pull - one row per part number (Spec C4 missing list)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:N1")
    ws["A2"] = (
        f"shared | {stats['ranked_pn_count']} PNs | EUR {stats['ranked_eur']:,.0f} CY2025 billed without usable current PR00. "
        "Instruction: for each column Pull_* = Y, export that PN price where Valid From/To overlaps the matching window "
        "(see Pull_Windows). Window_2025 = all rows Y (P0). Window_2023/2024 = Y only if year-matched gap OR EUR>=10k. "
        "Both PR00 and ZPR1 when Y."
    )
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 52

    p_headers = [
        "part_number",
        "Description",
        "Reason",
        "EUR_CY2025",
        "Priority",
        "Pull_PR00_2023",
        "Pull_ZPR1_2023",
        "Pull_PR00_2024",
        "Pull_ZPR1_2024",
        "Pull_PR00_2025",
        "Pull_ZPR1_2025",
        "In_year_gap_2023",
        "In_year_gap_2024",
        "gap_reason_code",
    ]
    yn_cols = {
        "Pull_PR00_2023",
        "Pull_ZPR1_2023",
        "Pull_PR00_2024",
        "Pull_ZPR1_2024",
        "Pull_PR00_2025",
        "Pull_ZPR1_2025",
        "In_year_gap_2023",
        "In_year_gap_2024",
    }
    p_rows = []
    for rec in parts.itertuples(index=False):
        p_rows.append(
            [
                rec.part_number_10,
                rec.description if pd.notna(rec.description) else "",
                rec.Reason,
                round(float(rec.EUR_CY2025), 2),
                rec.Priority,
                rec.Pull_PR00_2023,
                rec.Pull_ZPR1_2023,
                rec.Pull_PR00_2024,
                rec.Pull_ZPR1_2024,
                rec.Pull_PR00_2025,
                rec.Pull_ZPR1_2025,
                rec.In_year_gap_2023,
                rec.In_year_gap_2024,
                rec.reason,
            ]
        )
    last = write_rows(ws, 4, p_headers, p_rows, row_height=32, yn_cols=yn_cols)
    ws.auto_filter.ref = f"A4:N{last - 1}"
    ws.freeze_panes = "A5"
    autosize(ws, [14, 28, 48, 12, 10, 14, 14, 14, 14, 14, 14, 14, 14, 22])

    # ── Pull_Windows (6 pulls) ──────────────────────────────────────────
    ws0 = wb.create_sheet("Pull_Windows", 1)
    ws0["A1"] = "SAP pull windows - 3 date ranges x PR00 + ZPR1 (6 extracts)"
    ws0["A1"].font = title_font
    ws0.merge_cells("A1:H1")
    ws0["A2"] = (
        "shared | Built 2026-08-12 | European dates DD/MM/YYYY. "
        "For each window, export BOTH network list (PR00 / IE40) and customer conditions (ZPR1). "
        "Export every condition row whose Valid From-Valid To overlaps the window. "
        "Parts list = sheet Parts_To_Pull (Y = pull that PN in that window)."
    )
    ws0["A2"].font = subtitle_font
    ws0["A2"].alignment = wrap
    ws0.merge_cells("A2:H2")
    ws0.row_dimensions[2].height = 48

    pw_headers = [
        "Pull_ID",
        "Window",
        "Condition",
        "Valid_From",
        "Valid_To",
        "Valid_To_alternate",
        "Status",
        "Notes",
    ]
    pw_rows = []
    for w in WINDOWS:
        for cond in CONDITIONS:
            pid = f"{w['label']}_{cond}"
            n_pn = stats["pull_counts"][f"{w['label']}_{cond}"]
            pw_rows.append(
                [
                    pid,
                    w["label"],
                    cond,
                    w["valid_from"],
                    w["valid_to"],
                    w["valid_to_alternate"] or "-",
                    "to_pull",
                    f"{w['note']} | Parts_To_Pull flag Y for {n_pn} PNs (Spec C4 list).",
                ]
            )
    write_rows(ws0, 4, pw_headers, pw_rows, row_height=40)
    ws0.auto_filter.ref = f"A4:H{4 + len(pw_rows)}"
    ws0.freeze_panes = "A5"
    autosize(ws0, [22, 28, 12, 14, 14, 18, 12, 70])

    ws0.cell(12, 1, "How many PNs flagged Y per window (unique PN; PR00=ZPR1 flags match)").font = Font(
        bold=True, name="Calibri", size=11
    )
    sum_h = ["Window", "Valid_From", "Valid_To", "PNs_flagged_Y", "Conditions"]
    sum_r = [
        [
            "Window_2023",
            "01/01/2023",
            "31/12/2023 (alt 31/01/2023)",
            stats["any_pull_2023"],
            "PR00 + ZPR1",
        ],
        [
            "Window_2024",
            "01/01/2024",
            "31/01/2025",
            stats["any_pull_2024"],
            "PR00 + ZPR1",
        ],
        [
            "Window_2025",
            "01/02/2025",
            "17/03/2026",
            stats["any_pull_2025"],
            "PR00 + ZPR1 (all Spec C4)",
        ],
    ]
    write_rows(ws0, 13, sum_h, sum_r, row_height=22)

    # ── How_To_Pull ─────────────────────────────────────────────────────
    ws3 = wb.create_sheet("How_To_Pull", 2)
    ws3["A1"] = "How to pull from SAP — plain English"
    ws3["A1"].font = title_font
    ws3.merge_cells("A1:B1")
    instructions = [
        ("What this file is for", "Six SAP extracts: three date windows x two condition types (PR00 and ZPR1)."),
        (
            "Date windows (European)",
            "Window_2023: 01/01/2023 to 31/12/2023 (preferred). "
            "Alternate if typo meant month-end Jan: 31/01/2023 (Jerome wrote 310122023). "
            "Window_2024: 01/01/2024 to 31/01/2025. "
            "Window_2025: 01/02/2025 to 17/03/2026.",
        ),
        (
            "Clarify once",
            "Please confirm Window_2023 Valid To = 31/12/2023 (year-end) vs 31/01/2023 (literal 9-digit string).",
        ),
        ("Sales organization", "IE40 for network list (PR00)."),
        (
            "Conditions",
            "PR00 = network list price. ZPR1 = customer-specific list. Pull BOTH for every window. Keep UTI and UMI customers separate.",
        ),
        (
            "Overlap rule",
            "Export every condition record for the Materials on Parts_To_Pull where Valid From / Valid To overlaps the window dates. "
            "Do not collapse to a single year-end price without validity columns.",
        ),
        (
            "Minimum columns",
            "Material (part number); Condition type (PR00 or ZPR1); Sales Organization (PR00); Customer (ZPR1); "
            "Amount; Currency; Valid From; Valid To; Deletion indicator if available.",
        ),
        (
            "Which parts",
            "Sheet Parts_To_Pull. For each Y in Pull_PR00_YYYY / Pull_ZPR1_YYYY, include that part in that window extract. "
            f"All {stats['ranked_pn_count']} rows are Y for Window_2025 (most relevant to current missing list).",
        ),
        (
            "Priority filter (optional)",
            "Filter Priority=P0 first (EUR>=10k CY2025). Then P1, then rest.",
        ),
        (
            "Reason codes",
            "absent_from_pr00 = not on current list. on_list_no_usable_price = blank/zero amount. "
            "superseded_only = only old superseded row - price the active/successor part too if known.",
        ),
        (
            "File naming",
            "IE40_PR00_Window_2023.xlsx, IE40_ZPR1_Window_2023.xlsx, ... same for 2024 and 2025. "
            "Land in 01_data/90_INBOX/ then sort to 01_Pricing/_raw/.",
        ),
        (
            "Drop zone",
            "C:\\Users\\jerom\\Desktop\\claude\\ISZE\\01_data\\90_INBOX\\",
        ),
        (
            "Already on file (do not re-pull blindly)",
            "See Have_Already. We already have year-end Material+EUR sheets and some ZPR1 with dates - "
            "this request is for dated condition rows overlapping the three windows for the listed PNs.",
        ),
        (
            "Rebuild this workbook",
            r"python C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\_build_sap_history_request_xlsx.py",
        ),
    ]
    ws3["A3"] = "Step"
    ws3["B3"] = "Detail"
    style_header(ws3, 3, 2)
    for i, (topic, text) in enumerate(instructions):
        r = 4 + i
        ws3.cell(r, 1, topic).font = Font(bold=True, name="Calibri", size=10)
        ws3.cell(r, 1).alignment = wrap
        ws3.cell(r, 1).border = thin
        ws3.cell(r, 2, text).alignment = wrap
        ws3.cell(r, 2).border = thin
        ws3.cell(r, 2).font = Font(name="Calibri", size=10)
        ws3.row_dimensions[r].height = 44
    autosize(ws3, [28, 110])
    ws3.freeze_panes = "A4"

    # ── Have_Already ────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Have_Already", 3)
    ws2["A1"] = "Already filed - do not treat as completing the 6 window pulls"
    ws2["A1"].font = title_font
    ws2.merge_cells("A1:E1")
    ws2["A2"] = (
        "Reference only. Year-end PR00 sheets are Material + EUR without full Valid From/To grain for these windows. "
        "ZPR1 history covers key customers for some vintages — still pull window-overlap extracts for Parts_To_Pull."
    )
    ws2["A2"].font = subtitle_font
    ws2["A2"].alignment = wrap
    ws2.merge_cells("A2:E2")
    ws2.row_dimensions[2].height = 40

    h2 = ["Category", "Vintage", "File_path", "Role", "Gap_vs_this_request"]
    have_rows = [
        [
            "PR00 IE40",
            "2023 YE stamp",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\IE40_PR00_2023-12-31.xlsx",
            "Material + Price_EUR",
            "Not a Window_2023 Valid From/To extract",
        ],
        [
            "PR00 IE40",
            "2024 YE stamp",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\IE40_PR00_2024-12-31.xlsx",
            "Material + Price_EUR",
            "Not Window_2024 dated overlap extract",
        ],
        [
            "PR00 IE40",
            "2025 YE stamp",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\IE40_PR00_2025-12-31.xlsx",
            "Material + Price_EUR",
            "Not Window_2025 dated overlap extract",
        ],
        [
            "PR00 IE40",
            "2026-03-13 snapshot",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\IE40 Price list as of 13.03.2026 PR00_3_level_classified_GPT.xlsx",
            "Point-in-time list",
            "Near Window_2025 end; not full window history",
        ],
        [
            "PR00 IE40",
            "2026-08-10 current",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\Price_file_PR00_IE40_20260810.xlsx",
            "Current list (defines Spec C4 gap)",
            "After Window_2025 end date",
        ],
        [
            "ZPR1",
            "2023 key customers",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\ZPR1_2023-01-03.xlsx",
            "Has Valid From/To for 1799/1797/2372/2733/2734",
            "Still pull Window_2023 for Parts_To_Pull Materials",
        ],
        [
            "ZPR1",
            "2024 key customers",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\ZPR1_2024-01-03.xlsx",
            "Key customers + Valid From",
            "Still pull Window_2024 for listed PNs",
        ],
        [
            "ZPR1",
            "2025 key customers",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\history\ZPR1_2025-02-01.xlsx",
            "1799/1797/2372",
            "Still pull Window_2025 for listed PNs",
        ],
        [
            "ZPR1 / CIS",
            "2026-08-10",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\EXPORT_ZPR1_20260810.xlsx",
            "Full current customer export",
            "After Window_2025",
        ],
        [
            "Spec C4 ranked list",
            "CY2025 vs current PR00",
            r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived\no_pr00_gap_cy2025_ranked.csv",
            f"{stats['ranked_pn_count']} PNs / €{stats['ranked_eur']:,.0f}",
            "Source for Parts_To_Pull",
        ],
    ]
    write_rows(ws2, 4, h2, have_rows, row_height=36)
    ws2.auto_filter.ref = f"A4:E{4 + len(have_rows)}"
    ws2.freeze_panes = "A5"
    autosize(ws2, [18, 28, 95, 36, 40])

    # ── Legend (short) ──────────────────────────────────────────────────
    ws4 = wb.create_sheet("Legend", 4)
    ws4["A1"] = "Legend"
    ws4["A1"].font = title_font
    leg = [
        ["Code", "Meaning"],
        ["Y", "Pull this PN for that window × condition"],
        ["N", "Skip for that window (still pull Window_2025 if Y there)"],
        ["P0", "€≥10k CY2025 — do first"],
        ["P1", "€≥1k"],
        ["P2", "Remainder"],
        ["to_pull", "Extract not yet received for this window"],
        ["Window_2025", "Primary — all Spec C4 PNs"],
        ["Window_2023/2024", "Y if year-matched missing on filed PR00 OR top-€ (P0)"],
    ]
    for i, row in enumerate(leg):
        for j, v in enumerate(row, 1):
            cell = ws4.cell(2 + i, j, v)
            cell.border = thin
            cell.alignment = wrap
            if i == 0:
                cell.fill = header_fill
                cell.font = header_font
            elif v == "Y":
                cell.fill = y_fill
            elif v == "N":
                cell.fill = n_fill
            elif v == "P0":
                cell.fill = p0_fill
            elif v == "P1":
                cell.fill = p1_fill
            elif v == "P2":
                cell.fill = p2_fill
            elif v == "to_pull":
                cell.fill = missing_fill
    autosize(ws4, [22, 85])
    ws4.freeze_panes = "A3"

    wb.save(OUT)
    sanitize_xlsx_for_excel_desktop(OUT)
    print("Wrote", OUT)
    print("Spec C4 PNs:", stats["ranked_pn_count"], "eur", stats["ranked_eur"])
    print("By reason:", stats["by_reason"])
    print("Year-matched gap sizes (all billed, not only Spec):", stats["ym_counts"])
    print("PNs flagged Y per window:")
    for k, v in stats["pull_counts"].items():
        print(f"  {k}: {v}")
    print(
        "Unique PN counts:",
        "2023=",
        stats["any_pull_2023"],
        "2024=",
        stats["any_pull_2024"],
        "2025=",
        stats["any_pull_2025"],
    )


if __name__ == "__main__":
    main()
