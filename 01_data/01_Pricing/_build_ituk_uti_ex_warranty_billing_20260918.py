"""ITUK / UTI price revision side study — three warranty cuts (2026-09-18).

Cut A — billing SaTy codes (what Jerome asked): exclude warranty/FOC/claim codes.
Cut B — claim-qty haircut (intent-aligned proxy when Cut A is empty).
Cut C — PN-ever-on-claim flag (earlier study; reference only).

Does NOT change the published network 5.9616% or rebuild AIM HTML.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
WARR = ROOT / "01_data" / "CV_Warranty_01-08_2026.xlsx"
SALES = ROOT / "01_data" / "00_derived" / "sales_fact.parquet"
PN_CSV = ROOT / "01_data" / "00_derived" / "price_bp_weighted_compare_pn.csv"
OUT_DIR = ROOT / "01_data" / "01_Pricing"
OUT_XLSX = OUT_DIR / "Pricing_Revision_ITUK_UTI_ex_warranty_billing_20260918.xlsx"
OUT_MD = OUT_DIR / "Pricing_Revision_ITUK_UTI_ex_warranty_billing_20260918_SUMMARY.md"
OUT_CMD = OUT_DIR / "OPEN_Pricing_Revision_ITUK_UTI_ex_warranty_billing_20260918.cmd"
TMP_JSON = OUT_DIR / "_tmp_ituk_uti_ex_warranty_billing_result.json"

L12M_START = pd.Timestamp("2025-07-01")
L12M_END = pd.Timestamp("2026-06-30")
CLAIM_WIN_LO = 20250701
CLAIM_WIN_HI = 20260630
FX_JPY_PER_EUR = 163.7

# sold_to_sap → DB label (ITUK ≠ IUK; UTI ≠ UMI)
DBS = {
    2575: "ITUK",
    1799: "UTI",
}

HEADER_FILL = PatternFill("solid", fgColor="15233C")
HEADER_FONT = Font(bold=True, color="F7F4EE")

# SaTy tokens that would mean warranty / FOC / claim if they existed in billing
WARR_SATY_RE = re.compile(
    r"WARR|WTY|CLM|FOC|CLAIM|WARRANT|FREE.?OF.?CHARGE|GOODWILL",
    re.IGNORECASE,
)


def norm_pn(value) -> str | None:
    """Same key as build_price_bp_weighted_compare (10-char digit PN)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        digits = str(int(value))
    else:
        digits = re.sub(r"\D", "", str(value).strip())
    if len(digits) < 8:
        return None
    return digits.zfill(10)


def weighted_metrics(qty: pd.Series, bp_prior: pd.Series, bp_new: pd.Series, billed_eur: pd.Series) -> dict:
    q = pd.to_numeric(qty, errors="coerce").fillna(0.0)
    prior_jpy = float((q * bp_prior).sum())
    new_jpy = float((q * bp_new).sum())
    step_eur = float((q * (bp_new - bp_prior)).sum() / FX_JPY_PER_EUR)
    wtd = (100.0 * (new_jpy / prior_jpy - 1.0)) if prior_jpy else float("nan")
    return {
        "billed_eur": float(pd.to_numeric(billed_eur, errors="coerce").fillna(0).sum()),
        "prior_eur": prior_jpy / FX_JPY_PER_EUR,
        "new_eur": new_jpy / FX_JPY_PER_EUR,
        "step_eur": step_eur,
        "weighted_pct": wtd,
        "n_pn": int((q > 0).sum()) if hasattr(q, "sum") else int(len(q)),
        "n_rows": int(len(q)),
        "qty_sum": float(q.sum()),
    }


def style_header(ws, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def autosize(ws, cols: list[str], max_w: int = 40) -> None:
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(max_w, len(c) + 2))


def main() -> None:
    # --- Price bases (published priced-both, per customer × PN) ---
    px = pd.read_csv(PN_CSV, dtype={"part_number": str, "customer_code": str})
    px["customer_code"] = px["customer_code"].astype(str).str.replace(r"\.0$", "", regex=True)
    px["pn"] = px["part_number"].map(norm_pn)
    px = px[
        px["customer_code"].isin({str(k) for k in DBS})
        & (px["in_scope"] == True)  # noqa: E712
        & px["bp_new_jpy"].notna()
        & px["bp_prior_revised"].notna()
        & px["pn"].notna()
    ].copy()
    px["db"] = px["customer_code"].map(lambda c: DBS[int(c)])

    # --- Sales fact L12M (SaTy audit + qty/eur by PN) ---
    sf = pd.read_parquet(
        SALES,
        columns=[
            "billing_date",
            "sold_to_sap",
            "saty",
            "billed_qty",
            "net_sales_value",
            "part_number",
        ],
    )
    sf["sold_to_sap"] = pd.to_numeric(sf["sold_to_sap"], errors="coerce")
    sf["billing_date"] = pd.to_datetime(sf["billing_date"])
    sf = sf[
        sf["sold_to_sap"].isin(DBS.keys())
        & (sf["billing_date"] >= L12M_START)
        & (sf["billing_date"] <= L12M_END)
    ].copy()
    sf["db"] = sf["sold_to_sap"].map(DBS)
    sf["pn"] = sf["part_number"].map(norm_pn)
    sf["saty_raw"] = sf["saty"]
    sf["saty_key"] = (
        sf["saty"]
        .fillna("(blank)")
        .astype(str)
        .str.strip()
        .replace({"": "(blank)", "nan": "(blank)", "None": "(blank)", "NaT": "(blank)"})
    )
    sf.loc[sf["saty_key"].str.len() == 0, "saty_key"] = "(blank)"

    saty_counts = (
        sf.groupby(["db", "saty_key"], as_index=False)
        .agg(rows=("saty_key", "size"), billed_eur=("net_sales_value", "sum"), billed_qty=("billed_qty", "sum"))
        .sort_values(["db", "rows"], ascending=[True, False])
    )
    sf["is_warranty_saty"] = sf["saty_key"].map(lambda s: bool(WARR_SATY_RE.search(str(s))))
    # Credits / returns — document, not treat as warranty
    sf["is_credit_saty"] = sf["saty_key"].str.upper().isin({"CR", "RE", "DR"})

    sales_pn = (
        sf.dropna(subset=["pn"])
        .groupby(["db", "sold_to_sap", "pn"], as_index=False)
        .agg(billed_qty=("billed_qty", "sum"), billed_eur=("net_sales_value", "sum"))
    )
    sales_pn["asp"] = np.where(
        sales_pn["billed_qty"] != 0,
        sales_pn["billed_eur"] / sales_pn["billed_qty"],
        np.nan,
    )

    # --- Warranty claims ---
    w = pd.read_excel(WARR, sheet_name="Sheet1")
    w["_pn1"] = w["PARTNBR"].map(norm_pn)
    w["_pn2"] = w["PARTNBR2"].map(norm_pn)
    w["QTY1"] = pd.to_numeric(w["QTY1"], errors="coerce").fillna(0.0)
    w["QTY2"] = pd.to_numeric(w["QTY2"], errors="coerce").fillna(0.0)
    w["BILLYMD"] = pd.to_numeric(w["BILLYMD"], errors="coerce")
    bill_min = int(w["BILLYMD"].min()) if w["BILLYMD"].notna().any() else None
    bill_max = int(w["BILLYMD"].max()) if w["BILLYMD"].notna().any() else None

    def claim_qty_table(sub: pd.DataFrame) -> pd.DataFrame:
        """Sum claim qty by PN (PARTNBR→QTY1 + PARTNBR2→QTY2)."""
        rows = []
        for _, r in sub.iterrows():
            if r["_pn1"]:
                rows.append({"pn": r["_pn1"], "claim_qty": float(r["QTY1"])})
            if r["_pn2"]:
                rows.append({"pn": r["_pn2"], "claim_qty": float(r["QTY2"])})
        if not rows:
            return pd.DataFrame(columns=["pn", "claim_qty"])
        return pd.DataFrame(rows).groupby("pn", as_index=False)["claim_qty"].sum()

    claim_meta = []
    claim_by_db_window: dict[str, pd.DataFrame] = {}
    claim_by_db_full: dict[str, pd.DataFrame] = {}
    warranty_pn_sets: dict[str, set] = {}

    for sap, db in DBS.items():
        sub = w[w["DST_Name"] == db].copy()
        win = sub[
            sub["BILLYMD"].notna()
            & (sub["BILLYMD"] >= CLAIM_WIN_LO)
            & (sub["BILLYMD"] <= CLAIM_WIN_HI)
        ]
        cq_win = claim_qty_table(win)
        cq_full = claim_qty_table(sub)
        claim_by_db_window[db] = cq_win
        claim_by_db_full[db] = cq_full

        pns: set = set()
        for col in ("_pn1", "_pn2"):
            pns.update(x for x in sub[col].dropna().unique())
        warranty_pn_sets[db] = pns

        billy_sub = sub["BILLYMD"]
        claim_meta.append(
            {
                "DST_Name": db,
                "claim_lines_full": len(sub),
                "claim_lines_window_202507_202606": len(win),
                "unique_PN_full": len(pns),
                "unique_PN_with_qty_window": int(cq_win["pn"].nunique()) if len(cq_win) else 0,
                "unique_PN_with_qty_full": int(cq_full["pn"].nunique()) if len(cq_full) else 0,
                "claim_qty_sum_window": float(cq_win["claim_qty"].sum()) if len(cq_win) else 0.0,
                "claim_qty_sum_full": float(cq_full["claim_qty"].sum()) if len(cq_full) else 0.0,
                "BILLYMD_min": int(billy_sub.min()) if billy_sub.notna().any() else None,
                "BILLYMD_max": int(billy_sub.max()) if billy_sub.notna().any() else None,
            }
        )

    summary_rows = []
    cut_a_detail = []
    cut_b_frames = []
    cut_c_frames = []

    for sap, db in DBS.items():
        sp = sales_pn[sales_pn["db"] == db][["pn", "billed_qty", "billed_eur", "asp"]].rename(
            columns={"billed_qty": "sf_qty", "billed_eur": "sf_eur"}
        )
        base = px[px["db"] == db].copy()
        base["bp_qty"] = base["billed_qty"]
        base = base.merge(sp, on="pn", how="left")
        assert (base["sf_qty"] - base["bp_qty"]).abs().fillna(0).max() < 1e-6, (
            f"{db}: sales_fact qty diverges from price_bp"
        )

        bp_prior = base["bp_prior_revised"]
        bp_new = base["bp_new_jpy"]
        published = weighted_metrics(base["bp_qty"], bp_prior, bp_new, base["sf_eur"])

        # ===== Cut A: SaTy warranty codes =====
        # On the priced-both PN set, exclude sales lines with warranty SaTy, then
        # re-aggregate qty/eur by PN and reweight. Expected: zero warranty SaTy rows.
        sf_db = sf[sf["db"] == db]
        n_warr_saty = int(sf_db["is_warranty_saty"].sum())
        n_credit = int(sf_db["is_credit_saty"].sum())
        # Re-aggregate excluding warranty SaTy (no-op if none)
        sf_a = sf_db[~sf_db["is_warranty_saty"]].dropna(subset=["pn"])
        agg_a = sf_a.groupby("pn", as_index=False).agg(
            qty_a=("billed_qty", "sum"), eur_a=("net_sales_value", "sum")
        )
        m_a = base.merge(agg_a, on="pn", how="left")
        # PNs that only had warranty-coded lines would drop to 0
        m_a["qty_a"] = m_a["qty_a"].fillna(0.0)
        m_a["eur_a"] = m_a["eur_a"].fillna(0.0)
        cut_a = weighted_metrics(m_a["qty_a"], m_a["bp_prior_revised"], m_a["bp_new_jpy"], m_a["eur_a"])
        cut_a_detail.append(
            {
                "db": db,
                "l12m_billing_rows": len(sf_db),
                "warranty_coded_saty_rows": n_warr_saty,
                "credit_saty_rows_CR_RE_DR": n_credit,
                "saty_values_observed": ", ".join(
                    sorted(sf_db["saty_key"].unique().tolist())
                ),
                "published_weighted_pct": published["weighted_pct"],
                "cut_a_weighted_pct": cut_a["weighted_pct"],
                "delta_pts": cut_a["weighted_pct"] - published["weighted_pct"],
                "note": (
                    "No WARR/WTY/CLM/FOC SaTy in billing — Cut A cannot move the number."
                    if n_warr_saty == 0
                    else "Warranty SaTy codes found — Cut A applied."
                ),
            }
        )
        summary_rows.append({"db": db, "cut": "published_priced_both", **published})
        summary_rows.append(
            {
                "db": db,
                "cut": "A_billing_saty_ex_warranty",
                **cut_a,
                "warranty_coded_rows": n_warr_saty,
            }
        )

        # ===== Cut B: claim-qty haircut (window primary; full-file sensitivity) =====
        for claim_label, cq in (
            ("B_claim_qty_haircut_window", claim_by_db_window[db]),
            ("B_claim_qty_haircut_fullfile_sens", claim_by_db_full[db]),
        ):
            m = base.merge(cq.rename(columns={"claim_qty": "claim_qty"}), on="pn", how="left")
            m["claim_qty"] = m["claim_qty"].fillna(0.0)
            m["commercial_qty"] = np.maximum(0.0, m["bp_qty"] - m["claim_qty"])
            m["commercial_eur"] = np.where(
                m["bp_qty"] > 0,
                m["sf_eur"] * (m["commercial_qty"] / m["bp_qty"]),
                0.0,
            )
            m["qty_haircut"] = m["bp_qty"] - m["commercial_qty"]
            mets = weighted_metrics(
                m["commercial_qty"], m["bp_prior_revised"], m["bp_new_jpy"], m["commercial_eur"]
            )
            summary_rows.append(
                {
                    "db": db,
                    "cut": claim_label,
                    **mets,
                    "claim_qty_applied_sum": float(m["claim_qty"].sum()),
                    "qty_haircut_sum": float(m["qty_haircut"].sum()),
                    "pns_with_claim_overlap": int((m["claim_qty"] > 0).sum()),
                    "pns_fully_zeroed": int(((m["bp_qty"] > 0) & (m["commercial_qty"] == 0)).sum()),
                }
            )
            if claim_label == "B_claim_qty_haircut_window":
                det = m[m["claim_qty"] > 0][
                    [
                        "db",
                        "pn",
                        "billed_description",
                        "bp_new_class",
                        "bp_qty",
                        "claim_qty",
                        "commercial_qty",
                        "sf_eur",
                        "commercial_eur",
                        "bp_prior_revised",
                        "bp_new_jpy",
                    ]
                ].copy()
                det["prior_eur_commercial"] = (
                    det["commercial_qty"] * det["bp_prior_revised"] / FX_JPY_PER_EUR
                )
                det["new_eur_commercial"] = (
                    det["commercial_qty"] * det["bp_new_jpy"] / FX_JPY_PER_EUR
                )
                det["step_eur_commercial"] = (
                    det["new_eur_commercial"] - det["prior_eur_commercial"]
                )
                det["delta_pct"] = np.where(
                    det["bp_prior_revised"] != 0,
                    100.0 * (det["bp_new_jpy"] / det["bp_prior_revised"] - 1.0),
                    np.nan,
                )
                det = det.sort_values("step_eur_commercial", ascending=False)
                cut_b_frames.append(det)

        # ===== Cut C: PN-flag (earlier study; reference) =====
        war_set = warranty_pn_sets[db]
        m_c = base.copy()
        m_c["is_warranty_pn"] = m_c["pn"].isin(war_set)
        ex_c = m_c[~m_c["is_warranty_pn"]]
        war_c = m_c[m_c["is_warranty_pn"]]
        cut_c_ex = weighted_metrics(
            ex_c["bp_qty"], ex_c["bp_prior_revised"], ex_c["bp_new_jpy"], ex_c["sf_eur"]
        )
        cut_c_war = weighted_metrics(
            war_c["bp_qty"], war_c["bp_prior_revised"], war_c["bp_new_jpy"], war_c["sf_eur"]
        )
        summary_rows.append(
            {
                "db": db,
                "cut": "C_pn_flag_ex_warranty_REF",
                **cut_c_ex,
                "warranty_pn_set_size": len(war_set),
            }
        )
        summary_rows.append(
            {
                "db": db,
                "cut": "C_pn_flag_warranty_only_REF",
                **cut_c_war,
                "warranty_pn_set_size": len(war_set),
            }
        )
        if len(war_c):
            g = (
                war_c.groupby("pn", as_index=False)
                .agg(
                    description=("billed_description", "first"),
                    price_class=("bp_new_class", "first"),
                    qty=("bp_qty", "sum"),
                    billed_eur=("sf_eur", "sum"),
                    prior_jpy=("bp_prior_revised", "first"),
                    new_jpy=("bp_new_jpy", "first"),
                )
            )
            # Recompute step at PN level with summed qty (one row per PN already)
            g["prior_eur"] = g["qty"] * g["prior_jpy"] / FX_JPY_PER_EUR
            g["new_eur"] = g["qty"] * g["new_jpy"] / FX_JPY_PER_EUR
            g["step_eur"] = g["new_eur"] - g["prior_eur"]
            g["delta_pct"] = np.where(
                g["prior_jpy"] != 0, 100.0 * (g["new_jpy"] / g["prior_jpy"] - 1.0), np.nan
            )
            g.insert(0, "db", db)
            g = g.sort_values("step_eur", ascending=False)
            cut_c_frames.append(g)

    summary = pd.DataFrame(summary_rows)
    cut_a_df = pd.DataFrame(cut_a_detail)
    cut_b_detail = pd.concat(cut_b_frames, ignore_index=True) if cut_b_frames else pd.DataFrame()
    cut_c_detail = pd.concat(cut_c_frames, ignore_index=True) if cut_c_frames else pd.DataFrame()

    # Console
    for db in DBS.values():
        s = summary[summary["db"] == db]
        pub = s[s["cut"] == "published_priced_both"].iloc[0]
        a = s[s["cut"] == "A_billing_saty_ex_warranty"].iloc[0]
        b = s[s["cut"] == "B_claim_qty_haircut_window"].iloc[0]
        bfull = s[s["cut"] == "B_claim_qty_haircut_fullfile_sens"].iloc[0]
        c = s[s["cut"] == "C_pn_flag_ex_warranty_REF"].iloc[0]
        print(
            f"{db}: pub {pub['weighted_pct']:.4f}% step {pub['step_eur']:,.2f} | "
            f"A {a['weighted_pct']:.4f}% (d{a['weighted_pct']-pub['weighted_pct']:+.4f}) | "
            f"B-win {b['weighted_pct']:.4f}% (d{b['weighted_pct']-pub['weighted_pct']:+.4f}) "
            f"step {b['step_eur']:,.2f} | "
            f"B-full {bfull['weighted_pct']:.4f}% | "
            f"C {c['weighted_pct']:.4f}% (d{c['weighted_pct']-pub['weighted_pct']:+.4f})"
        )

    # ========== Excel ==========
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme = [
        "ITUK / UTI price revision — ex-warranty by billing code vs claim-qty haircut",
        "",
        "Side study only. Published network weighted increase 5.9616% "
        "(step EUR 1,652,140.71) is NOT changed. No AIM HTML rebuild.",
        "",
        "Entities: ITUK (sold_to_sap 2575) ≠ IUK; UTI (sold_to_sap 1799) ≠ UMI.",
        "",
        "Cut A — Billing SaTy (what was asked)",
        "Exclude lines whose SaTy looks like warranty / FOC / claim "
        "(WARR, WTY, CLM, FOC, CLAIM, …). CR / RE / DR are credits / debit memos, "
        "not warranty codes — counted separately, not treated as warranty.",
        "Expected here: zero warranty-coded rows → weighted % unchanged vs published.",
        "",
        "Cut B — Claim-qty haircut (intent-aligned proxy)",
        "Because Cut A finds no warranty SaTy in sales_fact, this is the closest "
        "commercial definition: for each PN, commercial_qty = max(0, billed_qty − claim_qty). "
        "Weighted % = 100 × (sum(commercial_qty × new JPY) / sum(commercial_qty × prior JPY) − 1), "
        "same formula as the published exercise. Billed euro scaled by commercial_qty/billed_qty.",
        "Primary claim window: BILLYMD 20250701–20260630 (aligned to L12M). "
        "Full-file claim qty is a sensitivity only.",
        "",
        "Cut C — PN-flag reference (earlier study)",
        "Drop every PN that ever appears on a claim line for that DST in the warranty file. "
        "Shown for comparison only — not the primary answer.",
        "",
        "Sources",
        f"sales_fact.parquet — L12M {L12M_START.date()} … {L12M_END.date()}, sold_to_sap 2575 / 1799",
        f"price_bp_weighted_compare_pn.csv — in_scope priced-both bases (prior = bp_prior_revised)",
        f"CV warranty: {WARR.name}",
        f"FX for euro step only: {FX_JPY_PER_EUR} JPY/EUR (percentage is JPY-pure)",
        "",
        f"Built: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
    ]
    for i, line in enumerate(readme, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 118

    # Summary
    ws2 = wb.create_sheet("Summary")
    sum_cols = [
        "db",
        "cut",
        "weighted_pct",
        "step_eur",
        "billed_eur",
        "prior_eur",
        "new_eur",
        "qty_sum",
        "n_pn",
        "n_rows",
        "warranty_coded_rows",
        "claim_qty_applied_sum",
        "qty_haircut_sum",
        "pns_with_claim_overlap",
        "pns_fully_zeroed",
        "warranty_pn_set_size",
    ]
    ws2.append(sum_cols)
    for _, r in summary.iterrows():
        ws2.append([r.get(c, None) for c in sum_cols])
    style_header(ws2, len(sum_cols))
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=len(sum_cols)):
        for idx, cell in enumerate(row, 1):
            if idx == 3 and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"
            if idx in (4, 5, 6, 7, 12, 13) and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
    autosize(ws2, sum_cols, 28)

    # Cut A audit
    ws3 = wb.create_sheet("Cut A SaTy audit")
    ws3.append(
        [
            "db",
            "l12m_billing_rows",
            "warranty_coded_saty_rows",
            "credit_saty_rows_CR_RE_DR",
            "saty_values_observed",
            "published_weighted_pct",
            "cut_a_weighted_pct",
            "delta_pts",
            "note",
        ]
    )
    for _, r in cut_a_df.iterrows():
        ws3.append(
            [
                r["db"],
                r["l12m_billing_rows"],
                r["warranty_coded_saty_rows"],
                r["credit_saty_rows_CR_RE_DR"],
                r["saty_values_observed"],
                r["published_weighted_pct"],
                r["cut_a_weighted_pct"],
                r["delta_pts"],
                r["note"],
            ]
        )
    style_header(ws3, 9)
    ws3.append([])
    ws3.append(["SaTy value counts (L12M ITUK + UTI)"])
    ws3.append(["db", "saty", "rows", "billed_eur", "billed_qty"])
    for _, r in saty_counts.iterrows():
        ws3.append([r["db"], r["saty_key"], r["rows"], r["billed_eur"], r["billed_qty"]])
    # header for counts block
    for c in range(1, 6):
        cell = ws3.cell(ws3.max_row - len(saty_counts), c)
        # find the counts header row
        pass
    # Restyle the counts header row we just wrote before the data — locate it
    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, max_col=5):
        if row[0].value == "db" and row[1].value == "saty":
            for cell in row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
    autosize(ws3, ["db", "l12m_billing_rows", "warranty_coded", "credit", "saty_values", "pub", "cuta", "delta", "note"], 36)

    # Cut B detail
    ws4 = wb.create_sheet("Cut B PN haircut")
    if len(cut_b_detail):
        bcols = list(cut_b_detail.columns)
        ws4.append(bcols)
        for _, r in cut_b_detail.iterrows():
            ws4.append([r[c] for c in bcols])
        style_header(ws4, len(bcols))
        for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, max_col=len(bcols)):
            for idx, cell in enumerate(row, 1):
                name = bcols[idx - 1]
                if name.endswith("_eur") or name in (
                    "sf_eur",
                    "commercial_eur",
                    "prior_eur_commercial",
                    "new_eur_commercial",
                    "step_eur_commercial",
                ):
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
                if name in ("bp_qty", "claim_qty", "commercial_qty", "bp_prior_revised", "bp_new_jpy"):
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
                if name == "delta_pct" and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        autosize(ws4, bcols)
    else:
        ws4["A1"] = "No claim-qty overlap with priced-both PNs in the window."

    # Cut C reference
    ws5 = wb.create_sheet("Cut C PN flag REF")
    if len(cut_c_detail):
        ccols = list(cut_c_detail.columns)
        ws5.append(ccols)
        for _, r in cut_c_detail.iterrows():
            ws5.append([r[c] for c in ccols])
        style_header(ws5, len(ccols))
        autosize(ws5, ccols)
    else:
        ws5["A1"] = "No warranty PN overlap (Cut C)."

    # Warranty meta
    ws6 = wb.create_sheet("Warranty claim meta")
    ws6["A1"] = "Warranty file"
    ws6["B1"] = str(WARR)
    ws6["A2"] = "BILLYMD min (all file)"
    ws6["B2"] = bill_min
    ws6["A3"] = "BILLYMD max (all file)"
    ws6["B3"] = bill_max
    ws6["A4"] = "Preferred claim window"
    ws6["B4"] = f"{CLAIM_WIN_LO}–{CLAIM_WIN_HI}"
    ws6["A5"] = "Note"
    ws6["B5"] = (
        "File starts mid-window (first BILLYMD often Dec 2025). "
        "Window cut is honest but partial; full-file sensitivity is on Summary."
    )
    ws6["A1"].font = Font(bold=True)
    meta_cols = list(claim_meta[0].keys())
    for col, name in enumerate(meta_cols, 1):
        cell = ws6.cell(8, col, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for ri, row in enumerate(claim_meta, 9):
        for ci, name in enumerate(meta_cols, 1):
            cell = ws6.cell(ri, ci, row[name])
            if "qty" in name or "sum" in name:
                if isinstance(row[name], (int, float)):
                    cell.number_format = "#,##0.00"
    for i, name in enumerate(meta_cols, 1):
        ws6.column_dimensions[get_column_letter(i)].width = max(14, min(36, len(name) + 2))

    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")

    # ========== SUMMARY.md ==========
    lines = [
        "# ITUK / UTI — ex-warranty by billing code vs claim-qty haircut",
        "",
        "Side study only. Published network **5.9616%** (step EUR 1,652,140.71) is unchanged.",
        "",
        "## Verdict",
        "",
    ]

    # Cut A verdict
    lines.append(
        "**Cut A (billing SaTy):** there is no warranty / FOC / claim SaTy on ITUK or UTI "
        "billing in the L12M window. Observed codes are only SOR, blank, PRD, VOR, RUSH, CR, DR. "
        "Warranty parts still bill as normal genuine lines; PRISM subtracts warranty expense at "
        "fiscal GPS level, not via a SaTy flag. So Cut A removes **zero** rows and the weighted "
        "% is **unchanged** vs the published priced-both basis. This cut cannot move the number."
    )
    lines.append("")

    for db in DBS.values():
        s = summary[summary["db"] == db]
        pub = s[s["cut"] == "published_priced_both"].iloc[0]
        b = s[s["cut"] == "B_claim_qty_haircut_window"].iloc[0]
        bfull = s[s["cut"] == "B_claim_qty_haircut_fullfile_sens"].iloc[0]
        c = s[s["cut"] == "C_pn_flag_ex_warranty_REF"].iloc[0]
        d_b = b["weighted_pct"] - pub["weighted_pct"]
        d_c = c["weighted_pct"] - pub["weighted_pct"]
        direction_b = "raises" if d_b > 0 else ("lowers" if d_b < 0 else "does not change")
        direction_c = "raises" if d_c > 0 else ("lowers" if d_c < 0 else "does not change")
        lines.append(
            f"**Cut B ({db}, claim-qty haircut, window):** {direction_b} the weighted increase "
            f"from {pub['weighted_pct']:.4f}% to {b['weighted_pct']:.4f}% "
            f"({d_b:+.4f} pts). Step euro moves from EUR {pub['step_eur']:,.2f} to "
            f"EUR {b['step_eur']:,.2f}. "
            f"Full-file claim sensitivity: {bfull['weighted_pct']:.4f}% "
            f"(step EUR {bfull['step_eur']:,.2f})."
        )
        lines.append("")
        lines.append(
            f"**Cut C ({db}, PN-flag reference):** {direction_c} from "
            f"{pub['weighted_pct']:.4f}% to {c['weighted_pct']:.4f}% ({d_c:+.4f} pts). "
            f"Reference only — not the primary answer."
        )
        lines.append("")

    lines.extend(
        [
            "### Why Cut B can move the other way from Cut C",
            "",
            "Cut C drops the **entire** billed quantity of any PN that ever appears on a claim. "
            "That removes high-step (or low-step) parts completely, so the remaining basket’s "
            "weighted % can fall even when commercial reality is only a partial warranty share.",
            "",
            "Cut B only subtracts **claim quantity** from billed quantity. The PN stays in the "
            "basket at its residual commercial volume. If the haircut removes relatively more "
            "volume from parts with a **below-average** price step (or leaves parts with an "
            "above-average step), the weighted % can rise — Jerome’s intuition that “strip out "
            "warranty and the commercial increase looks different” — even when the blunt PN-flag "
            "cut moved the other way.",
            "",
            "## Numbers",
            "",
            "| DB | Cut | Weighted % | Step € | Billed € (weight basis) | Qty |",
            "|---|---|---:|---:|---:|---:|",
        ]
    )
    for _, r in summary.iterrows():
        lines.append(
            f"| {r['db']} | {r['cut']} | {r['weighted_pct']:.4f}% | "
            f"{r['step_eur']:,.2f} | {r['billed_eur']:,.2f} | {r['qty_sum']:,.1f} |"
        )

    lines.extend(
        [
            "",
            "## SaTy observed (L12M, sold_to_sap 2575 / 1799)",
            "",
            "| DB | SaTy | Rows |",
            "|---|---|---:|",
        ]
    )
    for _, r in saty_counts.iterrows():
        lines.append(f"| {r['db']} | {r['saty_key']} | {int(r['rows']):,} |")

    lines.extend(["", "## Warranty file", ""])
    lines.append(f"- File: `{WARR}`")
    lines.append(f"- BILLYMD all file: {bill_min} → {bill_max}")
    lines.append(f"- Preferred window for Cut B: {CLAIM_WIN_LO}–{CLAIM_WIN_HI}")
    for r in claim_meta:
        lines.append(
            f"- **{r['DST_Name']}:** {r['claim_lines_full']:,} claim lines "
            f"({r['claim_lines_window_202507_202606']:,} in window); "
            f"claim qty sum window {r['claim_qty_sum_window']:,.1f} / "
            f"full file {r['claim_qty_sum_full']:,.1f}; "
            f"BILLYMD {r['BILLYMD_min']} → {r['BILLYMD_max']}."
        )

    lines.extend(
        [
            "",
            "## Caveats",
            "",
            "- Cut A is empty by construction of the billing extract — not a modelling choice.",
            "- Warranty file CV claims 01–08 2026; first BILLYMD is mid-L12M, so the window "
            "haircut is partial vs a full Jul25–Jun26 claim extract.",
            "- Cut C is the earlier “PN ever on a claim” study, kept as reference only.",
            "- ITUK ≠ IUK; UTI ≠ UMI.",
            "",
            "## Outputs",
            "",
            f"- Excel: `{OUT_XLSX}`",
            f"- Open: `{OUT_CMD}`",
            "",
        ]
    )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT_MD}")

    cmd_body = "@echo off\r\n" + f'start "" "%~dp0{OUT_XLSX.name}"\r\n'
    OUT_CMD.write_bytes(cmd_body.encode("ascii"))
    print(f"Wrote {OUT_CMD}")

    TMP_JSON.write_text(
        json.dumps(
            {
                "summary": summary.replace({np.nan: None}).to_dict(orient="records"),
                "cut_a": cut_a_df.replace({np.nan: None}).to_dict(orient="records"),
                "claim_meta": claim_meta,
                "saty_counts": saty_counts.replace({np.nan: None}).to_dict(orient="records"),
            },
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    print("OK")


if __name__ == "__main__":
    main()
