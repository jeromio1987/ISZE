"""Side study: ITUK / UTI weighted price increase excluding warranty PNs."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
WARR = ROOT / "01_data" / "CV_Warranty_01-08_2026.xlsx"
PARQ = ROOT / "01_data" / "00_derived" / "pricing_aim_source_lines.parquet"
OUT_DIR = ROOT / "01_data" / "01_Pricing"
OUT_XLSX = OUT_DIR / "Pricing_Revision_ITUK_UTI_ex_warranty_20260918.xlsx"
OUT_MD = OUT_DIR / "Pricing_Revision_ITUK_UTI_ex_warranty_20260918_SUMMARY.md"
OUT_CMD = OUT_DIR / "OPEN_Pricing_Revision_ITUK_UTI_ex_warranty_20260918.cmd"
TMP_JSON = OUT_DIR / "_tmp_ituk_uti_ex_warranty_result.json"

DBS = ["ITUK", "UTI"]
HEADER_FILL = PatternFill("solid", fgColor="15233C")
HEADER_FONT = Font(bold=True, color="F7F4EE")


def norm_pn(s):
    if pd.isna(s):
        return None
    t = str(s).strip().upper()
    if not t or t in ("0", "0.0", "NAN", "NONE", "NULL"):
        return None
    if t.endswith(".0") and t[:-2].isdigit():
        t = t[:-2]
    return t


def metrics(df: pd.DataFrame) -> dict:
    billed = float(df["billed_eur"].sum())
    prior = float(df["prior_eur"].sum())
    new = float(df["new_eur"].sum())
    step = float(df["step_eur"].sum())
    n_pn = int(df["part_number"].nunique())
    n_rows = len(df)
    wtd = (100.0 * (new / prior - 1.0)) if prior else float("nan")
    return {
        "billed_eur": billed,
        "prior_eur": prior,
        "new_eur": new,
        "step_eur": step,
        "weighted_pct": wtd,
        "n_pn": n_pn,
        "n_rows": n_rows,
    }


def main() -> None:
    w = pd.read_excel(WARR, sheet_name="Sheet1")
    w["_pn1"] = w["PARTNBR"].map(norm_pn)
    w["_pn2"] = w["PARTNBR2"].map(norm_pn)
    billy = pd.to_numeric(w["BILLYMD"], errors="coerce")
    bill_min = int(billy.min()) if billy.notna().any() else None
    bill_max = int(billy.max()) if billy.notna().any() else None
    w["_part_eur"] = (
        pd.to_numeric(w["PARTAMT1"], errors="coerce").fillna(0)
        + pd.to_numeric(w["PARTAMT2"], errors="coerce").fillna(0)
    )

    warranty_sets: dict[str, set] = {}
    warranty_meta_rows = []
    for db in DBS:
        sub = w[w["DST_Name"] == db].copy()
        pns: set = set()
        for col in ("_pn1", "_pn2"):
            pns.update(x for x in sub[col].dropna().unique())
        warranty_sets[db] = pns
        billy_sub = pd.to_numeric(sub["BILLYMD"], errors="coerce")
        warranty_meta_rows.append(
            {
                "DST_Name": db,
                "claim_lines": len(sub),
                "unique_PN_PARTNBR": int(sub["_pn1"].nunique(dropna=True)),
                "unique_PN_PARTNBR2": int(sub["_pn2"].nunique(dropna=True)),
                "unique_PN_combined": len(pns),
                "BILLYMD_min": int(billy_sub.min()) if billy_sub.notna().any() else None,
                "BILLYMD_max": int(billy_sub.max()) if billy_sub.notna().any() else None,
                "part_amt_eur_sum": float(sub["_part_eur"].sum()),
            }
        )

    dst_counts = {str(k): int(v) for k, v in w["DST_Name"].value_counts().items()}

    px = pd.read_parquet(PARQ)
    px = px[px["db"].isin(DBS) & (px["priced_both"] == True)].copy()  # noqa: E712
    px["part_number"] = px["part_number"].map(norm_pn)

    summary_rows = []
    detail_frames = []

    for db in DBS:
        d = px[px["db"] == db].copy()
        war_set = warranty_sets[db]
        d["is_warranty_pn"] = d["part_number"].isin(war_set)
        all_m = metrics(d)
        ex_m = metrics(d[~d["is_warranty_pn"]])
        war = d[d["is_warranty_pn"]]
        war_m = metrics(war)

        for label, m in (
            ("all_in_scope_priced", all_m),
            ("ex_warranty", ex_m),
            ("warranty_only", war_m),
        ):
            summary_rows.append(
                {
                    "db": db,
                    "slice": label,
                    **m,
                    "warranty_pn_set_size": len(war_set),
                    "step_share_of_db_pct": (
                        100.0 * m["step_eur"] / all_m["step_eur"] if all_m["step_eur"] else float("nan")
                    ),
                    "billed_share_of_db_pct": (
                        100.0 * m["billed_eur"] / all_m["billed_eur"]
                        if all_m["billed_eur"]
                        else float("nan")
                    ),
                }
            )

        if len(war):
            g = (
                war.groupby("part_number", as_index=False)
                .agg(
                    description=("description", "first"),
                    price_class=("price_class", "first"),
                    qty=("qty", "sum"),
                    billed_eur=("billed_eur", "sum"),
                    prior_eur=("prior_eur", "sum"),
                    new_eur=("new_eur", "sum"),
                    step_eur=("step_eur", "sum"),
                    n_rows=("part_number", "size"),
                )
            )
            g["delta_pct"] = np.where(
                g["prior_eur"] != 0,
                100.0 * (g["new_eur"] / g["prior_eur"] - 1.0),
                np.nan,
            )
            g.insert(0, "db", db)
            g = g.sort_values("step_eur", ascending=False)
            detail_frames.append(g)

    summary = pd.DataFrame(summary_rows)
    detail = (
        pd.concat(detail_frames, ignore_index=True) if detail_frames else pd.DataFrame()
    )

    for db in DBS:
        s = summary[summary["db"] == db]
        all_r = s[s["slice"] == "all_in_scope_priced"].iloc[0]
        ex_r = s[s["slice"] == "ex_warranty"].iloc[0]
        war_r = s[s["slice"] == "warranty_only"].iloc[0]
        delta_pts = ex_r["weighted_pct"] - all_r["weighted_pct"]
        print(
            f"{db}: all {all_r['weighted_pct']:.4f}% step {all_r['step_eur']:,.2f} | "
            f"ex {ex_r['weighted_pct']:.4f}% ({delta_pts:+.4f} pts) | "
            f"war step {war_r['step_eur']:,.2f} ({war_r['step_share_of_db_pct']:.2f}% of step) "
            f"billed share {war_r['billed_share_of_db_pct']:.2f}% | "
            f"war PNs in exercise {war_r['n_pn']} of set {int(all_r['warranty_pn_set_size'])}"
        )

    # --- Excel ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme = [
        "ITUK / UTI price revision — exclude warranty parts (side study)",
        "",
        "Purpose",
        "Separate exercise for Jerome: recompute the qty-weighted price increase for ITUK and UTI "
        "after dropping part numbers that appear on CV warranty claim lines for that distributor. "
        "The published network figure is NOT changed.",
        "",
        "Published network (unchanged)",
        "Step EUR 1,652,140.71 / weighted 5.9616% on in-scope priced-both qty-weighted bases "
        "(full AIM source cube).",
        "",
        "Sources",
        f"Warranty: {WARR.name} (Sheet1) — CV warranty claim lines 01–08 2026.",
        "Pricing lines: 01_data/00_derived/pricing_aim_source_lines.parquet — "
        "filter db in {ITUK, UTI} and priced_both == True.",
        "",
        "Method",
        "1. Normalize part numbers from PARTNBR and PARTNBR2 (strip, upper, drop blanks/zeros). "
        "PARTNBR2 is included.",
        "2. Build warranty PN sets per DST_Name exact match: ITUK, UTI.",
        "3. For each DB, take all in-scope priced-both rows; then drop rows whose part_number "
        "sits in that DB's warranty PN set (ex-warranty); and show the warranty-only slice of "
        "the price exercise.",
        "4. Weighted % = 100 * (sum new_eur / sum prior_eur − 1), same basis as the published "
        "weighted %.",
        "",
        "Caveats",
        "• Warranty file is CV warranty 01–08 2026 claim lines. Excluding means "
        "“PN ever on a claim for that DST in this file,” not a perfect match to the "
        "Jul25–Jun26 sales window used in the price exercise.",
        "• PARTNBR2 is included in the warranty PN set.",
        "• This is a side study. The published 5.9616% network figure is unchanged.",
        "• UTI and ITUK are never blended with UMI or IUK.",
        "",
        f"Built: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
    ]
    for i, line in enumerate(readme, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    ws2 = wb.create_sheet("Summary")
    sum_cols = [
        "db",
        "slice",
        "billed_eur",
        "prior_eur",
        "new_eur",
        "step_eur",
        "weighted_pct",
        "n_pn",
        "n_rows",
        "warranty_pn_set_size",
        "step_share_of_db_pct",
        "billed_share_of_db_pct",
    ]
    ws2.append(sum_cols)
    for _, r in summary.iterrows():
        ws2.append([r[c] for c in sum_cols])
    for c in range(1, len(sum_cols) + 1):
        cell = ws2.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    money_cols = {3, 4, 5, 6}
    pct_cols = {7, 11, 12}
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=len(sum_cols)):
        for idx, cell in enumerate(row, 1):
            if idx in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            if idx in pct_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"
    for i, c in enumerate(sum_cols, 1):
        ws2.column_dimensions[get_column_letter(i)].width = max(14, len(c) + 2)

    ws3 = wb.create_sheet("PN detail")
    if len(detail):
        det_cols = [
            "db",
            "part_number",
            "description",
            "price_class",
            "qty",
            "billed_eur",
            "prior_eur",
            "new_eur",
            "step_eur",
            "delta_pct",
            "n_rows",
        ]
        ws3.append(det_cols)
        for _, r in detail.iterrows():
            ws3.append([r[c] for c in det_cols])
        for c in range(1, len(det_cols) + 1):
            cell = ws3.cell(1, c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for row in ws3.iter_rows(
            min_row=2, max_row=ws3.max_row, min_col=1, max_col=len(det_cols)
        ):
            for idx, cell in enumerate(row, 1):
                if idx in (5, 6, 7, 8, 9) and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                if idx == 10 and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        for i, c in enumerate(det_cols, 1):
            ws3.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c) + 2))
        ws3.column_dimensions["C"].width = 40
    else:
        ws3["A1"] = (
            "No warranty PNs overlapped the priced-both price exercise for ITUK/UTI."
        )

    ws4 = wb.create_sheet("Warranty file meta")
    ws4["A1"] = "Warranty file"
    ws4["B1"] = str(WARR)
    ws4["A2"] = "Sheet"
    ws4["B2"] = "Sheet1"
    ws4["A3"] = "Total claim lines"
    ws4["B3"] = len(w)
    ws4["A4"] = "BILLYMD min (all)"
    ws4["B4"] = bill_min
    ws4["A5"] = "BILLYMD max (all)"
    ws4["B5"] = bill_max
    ws4["A6"] = "Note"
    ws4["B6"] = (
        "Excluding = PN ever on a claim for that DST in this file "
        "(PARTNBR + PARTNBR2). Not a perfect match to the sales window of the price exercise."
    )
    ws4["A1"].font = Font(bold=True)

    ws4["A8"] = "Per DST (ITUK / UTI)"
    ws4["A8"].font = Font(bold=True)
    meta_cols = list(warranty_meta_rows[0].keys())
    for col, name in enumerate(meta_cols, 1):
        cell = ws4.cell(9, col, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for ri, row in enumerate(warranty_meta_rows, 10):
        for ci, name in enumerate(meta_cols, 1):
            val = row[name]
            cell = ws4.cell(ri, ci, val)
            if name == "part_amt_eur_sum" and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"

    ws4["A13"] = "DST_Name line counts (full file)"
    ws4["A13"].font = Font(bold=True)
    ws4["A14"] = "DST_Name"
    ws4["B14"] = "claim_lines"
    ws4["A14"].fill = HEADER_FILL
    ws4["A14"].font = HEADER_FONT
    ws4["B14"].fill = HEADER_FILL
    ws4["B14"].font = HEADER_FONT
    for i, (dst, cnt) in enumerate(sorted(dst_counts.items(), key=lambda x: -x[1]), 15):
        ws4.cell(i, 1, dst)
        ws4.cell(i, 2, int(cnt))
    for letter, width in (
        ("A", 28),
        ("B", 22),
        ("C", 22),
        ("D", 22),
        ("E", 22),
        ("F", 18),
        ("G", 18),
        ("H", 18),
    ):
        ws4.column_dimensions[letter].width = width

    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")

    # --- SUMMARY.md ---
    lines = [
        "# ITUK / UTI price revision — exclude warranty parts",
        "",
        "**Side study only.** The published network weighted increase (**5.9616%**, step "
        "**EUR 1,652,140.71**) is unchanged.",
        "",
        "## Verdict",
        "",
    ]
    for db in DBS:
        s = summary[summary["db"] == db]
        all_r = s[s["slice"] == "all_in_scope_priced"].iloc[0]
        ex_r = s[s["slice"] == "ex_warranty"].iloc[0]
        war_r = s[s["slice"] == "warranty_only"].iloc[0]
        delta_pts = ex_r["weighted_pct"] - all_r["weighted_pct"]
        if delta_pts > 0:
            direction = "raises"
        elif delta_pts < 0:
            direction = "lowers"
        else:
            direction = "does not change"
        lines.append(
            f"**{db}:** excluding warranty parts {direction} the weighted increase from "
            f"{all_r['weighted_pct']:.4f}% to {ex_r['weighted_pct']:.4f}% "
            f"({delta_pts:+.4f} percentage points). "
            f"Warranty PNs carry EUR {war_r['step_eur']:,.2f} of that DB’s priced step "
            f"({war_r['step_share_of_db_pct']:.1f}% of step; "
            f"{war_r['billed_share_of_db_pct']:.1f}% of billed)."
        )
        lines.append("")

    lines.extend(
        [
            "## Numbers",
            "",
            "| DB | Slice | Weighted % | Step € | Billed € | PNs |",
            "|---|---|---:|---:|---:|---:|",
        ]
    )
    for _, r in summary.iterrows():
        lines.append(
            f"| {r['db']} | {r['slice']} | {r['weighted_pct']:.4f}% | "
            f"{r['step_eur']:,.2f} | {r['billed_eur']:,.2f} | {int(r['n_pn'])} |"
        )
    lines.append("")
    lines.extend(
        [
            "## Warranty file used",
            "",
            f"- File: `{WARR}`",
            f"- Claim window (BILLYMD all file): {bill_min} → {bill_max}",
        ]
    )
    for r in warranty_meta_rows:
        lines.append(
            f"- **{r['DST_Name']}:** {r['claim_lines']:,} claim lines; "
            f"{r['unique_PN_combined']:,} unique PNs (PARTNBR + PARTNBR2); "
            f"BILLYMD {r['BILLYMD_min']} → {r['BILLYMD_max']}; "
            f"part amount sum EUR {r['part_amt_eur_sum']:,.2f}."
        )
    lines.extend(
        [
            "",
            "## Caveats",
            "",
            "- Warranty file is CV warranty **01–08 2026** claim lines. Excluding means "
            "“PN ever on a claim for that DST in this file,” not a perfect match to the "
            "Jul25–Jun26 sales window of the price exercise.",
            "- **PARTNBR2** is included in the warranty PN set.",
            "- Separate exercise; published **5.9616%** network figure unchanged.",
            "- ITUK ≠ IUK; UTI ≠ UMI — never blended.",
            "",
            "## Outputs",
            "",
            f"- Excel: `{OUT_XLSX}`",
            f"- Open: `{OUT_CMD}`",
            "",
        ]
    )
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT_MD}")

    cmd_body = "@echo off\r\n" + f'start "" "%~dp0{OUT_XLSX.name}"\r\n'
    OUT_CMD.write_bytes(cmd_body.encode("ascii"))
    print(f"Wrote {OUT_CMD}")

    TMP_JSON.write_text(
        json.dumps(
            {"summary": summary.to_dict(orient="records"), "warranty_meta": warranty_meta_rows},
            indent=2,
        ),
        encoding="utf-8",
    )
    print("OK")


if __name__ == "__main__":
    main()
