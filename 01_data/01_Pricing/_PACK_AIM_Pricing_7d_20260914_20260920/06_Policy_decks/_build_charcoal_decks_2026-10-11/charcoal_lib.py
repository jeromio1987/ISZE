"""Charcoal deck library: the I-Motors FY2027 review look, on python-pptx.

Look taken from I_Motors_FY2027_Business_Performance_Review_CHARCOAL.pptx (24 Sep 2026): charcoal 262626
background, 303030 cards, Isuzu red E30613, amber F5A000, Aptos, brand strip top left, red kicker, white title,
grey subtitle, a right-hand panel, a rule above the source line and the page number bottom right.
Sizes are larger than that deck's (its panel text is 8.6 pt) because these decks are read on screen by management.
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Aptos"


def rgb(h):
    return RGBColor.from_string(h)


BG = rgb("262626")
CARD = rgb("303030")
CARD2 = rgb("2B2B2B")
HEAD = rgb("3A3A3A")
LINE = rgb("4A4A4A")
BORDER = rgb("6A6A6A")
RED = rgb("E30613")
DRED = rgb("8A0A10")
AMBER = rgb("F5A000")
GREEN = rgb("31A354")
GREEN2 = rgb("55C271")
TXT = rgb("F2F2F2")
SUB = rgb("BDBDBD")
MUTE = rgb("8C8C8C")
BAR = rgb("B8B8B8")  # neutral chart grey, as in the I-Motors deck
WHITE = rgb("FFFFFF")

W, H = 13.333, 7.5
PANEL_X, PANEL_Y, PANEL_W, PANEL_H = 9.86, 1.5, 3.25, 5.55


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def blank(prs, bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


# ---------------------------------------------------------------- primitives
def text(s, x, y, w, h, paras, size=12, color=TXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name=None, spacing=None, line=None):
    """paras: str | list of paragraphs; a paragraph is a str or a list of (text, opts) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        sp = spacing[i] if isinstance(spacing, (list, tuple)) else spacing
        if sp is not None and i > 0:
            p.space_before = Pt(sp)
        if line:
            p.line_spacing = line
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
            if o.get("strike"):
                r.font._rPr.set("strike", "sngStrike")
    return tb


def rect(s, x, y, w, h, fill=CARD, line=None, name=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    return sh


def hline(s, x1, y1, x2, y2, color=BORDER, lw=0.75, name=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if name:
        c.name = name
    return c


def picture(s, path, x, y, w=None, h=None, name=None):
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if name:
        p.name = name
    return p


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- slide frame
def frame(s, kicker, title, subtitle=None, source=None, page=None, panel=None, wide=False, panel_label="WHAT THIS MEANS"):
    """Brand strip, red kicker, title, subtitle, rule, source, page number and (optionally) the right-hand panel."""
    text(s, 0.25, 0.10, 4.5, 0.18, "ISZE  ·  VEHICLE PARTS & ACCESSORIES", size=7, bold=True, name="brand")
    text(s, 0.25, 0.42, 7.0, 0.22, kicker.upper(), size=9, bold=True, color=RED, name="kicker")
    tw = 12.8 if wide else 9.35
    text(s, 0.25, 0.68, tw, 0.82, title, size=24, bold=True, name="title", line=0.92)
    if subtitle:
        text(s, 0.25, 1.52, tw, 0.28, subtitle, size=11, color=SUB, name="subtitle")
    hline(s, 0.25, 7.18, 13.08, 7.18, color=BORDER, name="footer_rule")
    if source:
        text(s, 0.25, 7.23, 11.9, 0.2, source, size=8, color=MUTE, name="source")
    if page is not None:
        text(s, 12.55, 7.23, 0.53, 0.2, str(page), size=8, color=SUB, align=PP_ALIGN.RIGHT, name="page")
    if panel:
        right_panel(s, panel, panel_label)


def right_panel(s, items, label="WHAT THIS MEANS", x=PANEL_X, y=PANEL_Y, w=PANEL_W, h=PANEL_H, size=11.5):
    """items: list of (text | runs, dot colour). Slots are even, top aligned, like the I-Motors panel."""
    rect(s, x, y, w, h, fill=CARD, line=BORDER, name="panel")
    text(s, x + 0.28, y + 0.24, w - 0.5, 0.22, label, size=8.5, bold=True, color=RED, name="panel_label")
    n = len(items)
    top, avail = y + 0.72, h - 0.95
    slot = avail / n
    for i, (t, col) in enumerate(items):
        yy = top + i * slot
        rect(s, x + 0.28, yy + 0.07, 0.1, 0.1, fill=col, shape=MSO_SHAPE.OVAL, name=f"panel_dot{i}")
        paras = [t] if isinstance(t, str) else [t]
        text(s, x + 0.52, yy, w - 0.78, slot - 0.12, paras, size=size, name=f"panel_t{i}")


def band(s, x, y, w, h, paras, size=13, name="band"):
    rect(s, x, y, w, h, fill=CARD, line=BORDER, name=name + "_bg")
    text(s, x + 0.25, y, w - 0.5, h, paras, size=size, anchor=MSO_ANCHOR.MIDDLE, name=name)


def kpi(s, x, y, w, big, small, color=TXT, big_size=28, small_size=10, h=1.15, name="kpi"):
    text(s, x, y, w, 0.5, big, size=big_size, bold=True, color=color, name=name + "_v")
    text(s, x, y + 0.55, w, h - 0.55, small, size=small_size, color=SUB, name=name + "_l")


def card(s, x, y, w, h, kicker, big, body, big_color=TXT, kick_color=RED, big_size=26, body_size=10.5, name="card"):
    rect(s, x, y, w, h, fill=CARD, line=BORDER, name=name + "_bg")
    text(s, x + 0.2, y + 0.15, w - 0.4, 0.22, kicker, size=8.5, bold=True, color=kick_color, name=name + "_k")
    text(s, x + 0.2, y + 0.45, w - 0.4, 0.55, big, size=big_size, bold=True, color=big_color, name=name + "_v")
    text(s, x + 0.2, y + 1.08, w - 0.4, h - 1.2, body, size=body_size, color=SUB, name=name + "_b")


def logo(s, x, y, w=2.07):
    return picture(s, Path(__file__).parent / "assets" / "isuzu_logo.png", x, y, w=w, name="isuzu_logo")


def cover(prs, kicker, title_lines, sub, small, tag=None):
    """Cover as in the I-Motors deck: red bar on the left, brand strip, big title, grey lines, logo."""
    s = blank(prs)
    rect(s, 0, 0, 0.1, H, fill=RED, name="cover_bar")
    text(s, 0.45, 0.32, 8, 0.2, kicker, size=8, bold=True, name="cover_kicker")
    text(s, 0.45, 1.9, 11.4, 1.9, title_lines, size=36, bold=True, name="cover_title", line=0.95, spacing=6)
    text(s, 0.45, 4.2, 10.8, 0.9, sub, size=15, color=SUB, name="cover_sub", line=1.05)
    text(s, 0.45, 5.35, 10.8, 0.5, small, size=10.5, color=SUB, name="cover_small")
    if tag:
        # Status / disclaimer on cover must stay readable on charcoal (MUTE was too faint).
        text(s, 0.45, 5.85, 10.8, 0.35, tag, size=11, color=SUB, name="cover_tag")
    logo(s, 0.45, 6.6)
    return s


# ---------------------------------------------------------------- charts
def _nofill(chart):
    cs = chart._chartSpace
    ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"

    def mk():
        sp = etree.SubElement(etree.Element("dummy"), f"{{{ns}}}spPr")
        etree.SubElement(sp, f"{{{a}}}noFill")
        ln = etree.SubElement(sp, f"{{{a}}}ln")
        etree.SubElement(ln, f"{{{a}}}noFill")
        return sp

    old = cs.find(qn("c:spPr"))
    if old is not None:
        cs.remove(old)
    cs.chart.addnext(mk())
    pa = cs.chart.plotArea
    old = pa.find(qn("c:spPr"))
    if old is not None:
        pa.remove(old)
    pa.append(mk())


def _base(chart, size):
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = TXT


def _label(point, txt, size, color, pos):
    lab = point.data_label
    lab.position = pos
    tf = lab.text_frame
    tf.text = txt
    for r in tf.paragraphs[0].runs:
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = color


def _no_invert(series, i):
    dpt = series._element.xpath(f"./c:dPt[c:idx/@val='{i}']")[0]
    if not dpt.xpath("./c:invertIfNegative"):
        inv = dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
        dpt.find(qn("c:idx")).addnext(inv)


def hbar(s, x, y, w, h, cats, vals, colors, fmt, name, size=11, vmax=None, vmin=None, gap=45, low_labels=False,
         label_colors=None):
    """Horizontal bars, first category on top. fmt: callable value -> label text."""
    cd = CategoryChartData()
    cd.categories = list(reversed(cats))
    cd.add_series("v", list(reversed(vals)))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    _base(ch, size)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    if vmax is not None:
        va.maximum_scale = vmax
    if vmin is not None:
        va.minimum_scale = vmin
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = TXT
    if low_labels:
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, (c, v) in enumerate(zip(reversed(colors), reversed(vals))):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
        pt.format.line.fill.background()
        _no_invert(ser, i)
        lc = TXT
        if label_colors:
            lc = list(reversed(label_colors))[i]
        _label(pt, fmt(v), size, lc, XL_LABEL_POSITION.OUTSIDE_END)
    _nofill(ch)
    return gf


def clustered(s, x, y, w, h, cats, series, fmt, name, size=10.5, horizontal=False, vmax=None, vmin=None, gap=60,
              overlap=-5, legend=True, label_size=None):
    """Clustered bars or columns. series: list of (name, values, colour). fmt: value -> label text."""
    cd = CategoryChartData()
    cats_ = list(reversed(cats)) if horizontal else list(cats)
    cd.categories = cats_
    for nm, vals, _ in series:
        cd.add_series(nm, list(reversed(vals)) if horizontal else list(vals))
    ct = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = s.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    _base(ch, size)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    if vmax is not None:
        va.maximum_scale = vmax
    if vmin is not None:
        va.minimum_scale = vmin
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = TXT
    sers = list(plot.series)
    if horizontal:
        sers = sers  # series order is kept; legend order follows series order
    for sr, (nm, vals, col) in zip(sers, series):
        sr.invert_if_negative = False
        sr.format.fill.solid()
        sr.format.fill.fore_color.rgb = col
        sr.format.line.fill.background()
        vv = list(reversed(vals)) if horizontal else list(vals)
        for i, v in enumerate(vv):
            pt = sr.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = col
            pt.format.line.fill.background()
            _no_invert(sr, i)
            _label(pt, fmt(v), label_size or size, TXT, XL_LABEL_POSITION.OUTSIDE_END)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
        ch.legend.font.name = FONT
        ch.legend.font.color.rgb = SUB
    _nofill(ch)
    return gf


def waterfall(s, x, y, w, h, steps, fmt, name, size=10.5, vmax=None, gap=35):
    """steps: list of (label, value, kind, colour); kind 'total' draws a full bar, 'delta' floats from the running
    total. Built as stacked columns with an invisible base."""
    cats, base, vis, cols, lab = [], [], [], [], []
    run = 0.0
    for label, v, kind, col in steps:
        cats.append(label)
        if kind == "total":
            base.append(0.0)
            vis.append(v)
            run = v
        else:
            new = run + v
            base.append(min(run, new))
            vis.append(abs(v))
            run = new
        cols.append(col)
        lab.append((v, kind))
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("base", base)
    cd.add_series("step", vis)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    _base(ch, size)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = 100
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = TXT
    sb, sv = plot.series[0], plot.series[1]
    sb.format.fill.background()
    sb.format.line.fill.background()
    for i, c in enumerate(cols):
        pt = sv.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
        pt.format.line.fill.background()
        if lab[i][1] == "total":
            _label(pt, fmt(*lab[i]), size, TXT, XL_LABEL_POSITION.INSIDE_END)
        else:
            # a step can be a thin sliver: label it on the invisible base, just under the sliver
            _label(sb.points[i], fmt(*lab[i]), size, TXT, XL_LABEL_POSITION.INSIDE_END)
    _nofill(ch)
    return gf


# ---------------------------------------------------------------- tables
def table(s, x, y, w, rows, col_w, row_h=0.36, head_h=0.4, size=10.5, head_size=9, aligns=None, name="table",
          colors=None, bold_first=True):
    """rows[0] is the header. colors: {(r, c): RGB} text colour overrides. Native table, dark, no grid."""
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(head_h + row_h * (nr - 1)))
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if a in tblPr.attrib:
            del tblPr.attrib[a]
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(w * cw / tot)
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, nr):
        tbl.rows[r].height = Inches(row_h)
    aligns = aligns or ["l"] + ["r"] * (nc - 1)
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD if r == 0 else (CARD if r % 2 else CARD2)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[aligns[c]]
            run = p.add_run()
            run.text = str(rows[r][c])
            f = run.font
            f.name = FONT
            f.size = Pt(head_size if r == 0 else size)
            f.bold = (r == 0) or (bold_first and c == 0)
            col = SUB if r == 0 else TXT
            if colors and (r, c) in colors:
                col = colors[(r, c)]
            f.color.rgb = col
    return gf


def eur_m(v, d=1):
    return f"€{v / 1e6:.{d}f}M"


def eur_k(v):
    return f"€{v / 1e3:,.0f}k"
