"""Download candidate photographs, validate native resolution, then embed reviewed assets.

This does not run or change anything automatically. The supplied input deck stays intact.
Run --download first. After checking the images, run --apply-reviewed-images.
No credentials, macros, external image links or private-data uploads are used.
"""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
from urllib.request import Request, urlopen
import warnings
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'photos'
SOURCES = ROOT / 'image_sources.json'
INPUT = ROOT / 'AIM_checked_input.pptx'
OUTPUT = ROOT / 'AIM_with_reviewed_photos.pptx'
KEYS = ('element', 'housing', 'kit', 'truck')
MAX_BYTES = 35 * 1024 * 1024


def validate(data: bytes, info: dict) -> tuple[Image.Image, dict]:
    if len(data) > MAX_BYTES:
        raise ValueError('Image exceeds the 35 MB limit.')
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as original:
            original.load()
            im = ImageOps.exif_transpose(original).copy()
    w, h = im.size
    if max(w, h) < info['minimum_long_edge'] or min(w, h) < info['minimum_short_edge']:
        raise ValueError(f'Native image is {w} x {h}; below the required resolution. No upscaling is used.')
    if im.mode == 'RGBA':
        bg = Image.new('RGB', im.size, 'white'); bg.paste(im, mask=im.getchannel('A')); im = bg
    else:
        im = im.convert('RGB')
    return im, {'width': w, 'height': h, 'source_sha256': hashlib.sha256(data).hexdigest()}


def local_image(key: str) -> Path | None:
    # Users may place an authorised original here instead of using the candidate URL.
    for ext in ('.png', '.jpg', '.jpeg', '.webp'):
        candidate = ASSETS / (key + ext)
        if candidate.is_file():
            return candidate
    return None


def prepare(sources: dict, download: bool) -> dict:
    ASSETS.mkdir(exist_ok=True)
    records = {}
    errors = []
    for key in KEYS:
        info = sources[key]
        path = local_image(key)
        data = None
        used_url = None
        try:
            if path:
                data = path.read_bytes()
            elif download:
                attempts = []
                for url in info['image_urls']:
                    try:
                        req = Request(url, headers={'User-Agent': 'Mozilla/5.0', 'Accept': 'image/*'})
                        with urlopen(req, timeout=35) as response:
                            data = response.read(MAX_BYTES + 1)
                        im, details = validate(data, info)
                        used_url = url
                        break
                    except Exception as exc:
                        attempts.append(f'{url}: {exc}')
                        data = None
                if data is None:
                    raise ValueError('No valid high-resolution download.\n' + '\n'.join(attempts))
            else:
                raise FileNotFoundError(f'Place an authorised original at photos/{key}.jpg or run --download.')
            im, details = validate(data, info)
            destination = ASSETS / f'{key}.png'
            im.save(destination, 'PNG', optimize=True)
            records[key] = {
                **details, 'path': str(destination), 'source_page': info['page'],
                'image_url': used_url, 'note': info['note'],
                'provided_locally': path is not None,
            }
            print(f'{key}: {details["width"]} x {details["height"]}, saved.')
        except Exception as exc:
            errors.append(f'{key}: {exc}')
    (ROOT / 'photo_download_results.json').write_text(json.dumps({'images': records, 'errors': errors}, indent=2), encoding='utf-8')
    if errors:
        raise RuntimeError('\n\n'.join(errors) + '\n\nThe presentation has NOT been changed.')
    sheet = Image.new('RGB', (1600, 1220), 'white')
    draw = ImageDraw.Draw(sheet)
    for i, key in enumerate(KEYS):
        with Image.open(records[key]['path']) as source:
            preview = source.copy(); preview.thumbnail((760, 535))
        x = (i % 2) * 800; y = (i // 2) * 610
        draw.text((x + 20, y + 10), f'{key}: {records[key]["width"]} x {records[key]["height"]}', fill='black')
        sheet.paste(preview, (x + (800-preview.width)//2, y + 45 + (535-preview.height)//2))
    sheet.save(ROOT / 'REVIEW_IMAGES_BEFORE_APPLYING.png')
    return records


def contain(slide, image_path: str, x: float, y: float, w: float, h: float, name: str):
    with Image.open(image_path) as im:
        aspect = im.width / im.height
    iw, ih = (w, w / aspect) if w / h < aspect else (h * aspect, h)
    pic = slide.shapes.add_picture(image_path, Inches(x + (w-iw)/2), Inches(y + (h-ih)/2), width=Inches(iw), height=Inches(ih))
    pic.name = name
    pic._element.nvPicPr.cNvPr.set('descr', name + '. Representative photograph; application not inferred.')
    return pic


def words(slide) -> Counter:
    return Counter(re.findall(r'\w+', ' '.join(s.text for s in slide.shapes if s.has_text_frame).lower()))


def apply(records: dict):
    p = Presentation(INPUT)
    if len(p.slides) != 24:
        raise ValueError('Expected the supplied 24-slide checked input deck.')
    before_words = [words(s) for s in p.slides]
    before_xml = [s._element.xml for s in p.slides]
    s8 = p.slides[7]
    if len(s8.shapes) != 32 or 'An element, a housing and a kit' not in s8.shapes[2].text:
        raise ValueError('Slide 8 differs from the expected checked input. Nothing was saved.')
    # Delete drawings only; all labels, explanatory text, headings and footers remain.
    for sh in list(s8.shapes)[8:23]:
        if sh.has_text_frame and sh.text.strip():
            raise ValueError('A pictogram position unexpectedly contains text. Nothing was saved.')
        sh._element.getparent().remove(sh._element)
    for key, x in [('element', .82), ('housing', 5.02), ('kit', 9.25)]:
        contain(s8, records[key]['path'], x, 2.68, 3.25, 1.72, 'AIM photograph: ' + key)
    s24 = p.slides[23]
    statement = s24.shapes[2]
    if 'From "Others"' not in statement.text:
        raise ValueError('Closing statement differs from the expected checked input.')
    # Preserve the exact message, reserve its original left-hand panel for the photograph.
    statement.left = Inches(.8); statement.top = Inches(1.60)
    statement.width = Inches(6.55); statement.height = Inches(1.42)
    tf = statement.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    par = tf.paragraphs[0]
    style = par.runs[0].font if par.runs else None
    # Keep the four original lines, with explicit spacing for reliable PowerPoint fit.
    for par in tf.paragraphs:
        par.line_spacing = Pt(24); par.space_before = Pt(0); par.space_after = Pt(0)
        par.font.size = Pt(22)
        for run in par.runs:
            run.font.size = Pt(22)
    contain(s24, records['truck']['path'], .8, 3.16, 6.42, 3.0, 'AIM photograph: Australian MY25 Isuzu N Series')
    for sn, keys in [(8, KEYS[:3]), (24, ('truck',))]:
        note = p.slides[sn-1].notes_slide.notes_text_frame
        if sn == 8:
            for par in note.paragraphs:
                if 'Conceptual silhouettes illustrate product form' in par.text:
                    par.text = par.text.replace('Conceptual silhouettes illustrate product form; they are not technical fitment drawings.', 'Photographs illustrate product form; they do not establish technical fitment.')
            # Remove the obsolete drawing-only status from notes, retaining all other notes.
            paragraphs = [par for par in note.paragraphs if 'Compact-edition image status:' in par.text]
            for par in paragraphs: par._p.getparent().remove(par._p)
        suffix = '\n[Sources: product photographs]\n'
        for key in keys:
            r = records[key]
            suffix += f'{key}: {r["source_page"]}\n{r["note"]}\nNative pixels: {r["width"]} x {r["height"]}.\n'
        suffix += 'Illustrations do not establish common fitment or interchangeability.\n[/Sources]\n'
        note.add_paragraph().text = suffix
    assert [words(s) for s in p.slides] == before_words, 'Slide wording unexpectedly changed.'
    for i, slide in enumerate(p.slides):
        if i not in (7, 23):
            assert slide._element.xml == before_xml[i], f'Slide {i+1} changed unexpectedly.'
    if OUTPUT.exists():
        raise FileExistsError(f'{OUTPUT.name} already exists; rename it before creating another version.')
    tmp = ROOT / 'AIM_with_reviewed_photos.tmp.pptx'
    p.save(tmp); tmp.replace(OUTPUT)
    validation = {
        'slide_count': len(p.slides), 'only_photo_slides_modified': [8,24],
        'all_slide_words_retained': True, 'native_image_dimensions': records,
        'other_slide_xml_unchanged': True,
        'render_check_required': 'Open in PowerPoint and inspect slides 8 and 24, then export PDF and check every slide.',
    }
    (ROOT / 'photo_embedding_results.json').write_text(json.dumps(validation, indent=2), encoding='utf-8')
    print('Saved:', OUTPUT)
    print('Images are embedded, not linked. Inspect slides 8 and 24 in PowerPoint before use.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download', action='store_true', help='Try candidate public image URLs; validate pixels; make a preview.')
    parser.add_argument('--apply-reviewed-images', action='store_true', help='Embed the four images after checking identity and permission to use them.')
    args = parser.parse_args()
    sources = json.loads(SOURCES.read_text(encoding='utf-8'))
    records = prepare(sources, args.download)
    if args.apply_reviewed_images:
        apply(records)
    else:
        print('Inspect REVIEW_IMAGES_BEFORE_APPLYING.png and source pages. Then run --apply-reviewed-images.')


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        print('STOP:', exc, file=sys.stderr)
        sys.exit(1)
