"""UTI weekly back-order lists WK2-WK40 2026 -> tidy trend table + copy-price comparison.

Read-only on every source. Writes t13*/t14* CSVs + _numbers_bo_trend.json beside this file.
Run: python uti_backorder_trend.py
Sources:
  weekly BO    01_data/90_INBOX/UTI_backorders_2026/BO OVERVIEW - UNIVERSAL TRUCKS ISRAEL LTD dd-mm-2026.xlsx
               (39 ISZE mails "BO overview UNIVERSAL TRUCKS ISRAEL LTD (WKnn)"; see manifest.csv there)
  Eran 9 Jun   ERAN_20260611_1436_BO OVERVIEW ... 09-06-2026.xlsx (UTI's annotated copy, mail "body parts")
  AM samples   AM parts samples.xlsx, sheet 'part no & price' (Eran, 14 Jul 2026)
  categories   APPROVED AIM release via 04_tools/_shared/classification_paths.approved_release()
  list price   2026 PR00 (pricing_pr00_ie40.parquet via response_lib.load_lists), today's list, not the
               list in force on each snapshot date
Definitions:
  order date   'Req.Dlv.Dt' column (requested delivery date on the order line; the lists carry no
               separate order-creation date, so this is a proxy)
  days open    snapshot date (date in the file name) - Req.Dlv.Dt
  value        Open Qty x 2026 PR00 list (EUR); the lists carry no price
  body/panel   release L1 'Accident & Body' (Panels = L2 'Panels'); lines the release does not file as body
               but whose description is a body word are kept and flagged body_by_description_only
"""
import csv
import json
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
INB = ROOT / "01_data" / "90_INBOX" / "UTI_backorders_2026"
HERE = Path(__file__).parent
sys.path.insert(0, str(ROOT / "04_tools" / "_shared"))
sys.path.insert(0, str(HERE.parent))
import classification_paths as cp  # noqa: E402
from response_lib import load_lists  # noqa: E402

z = lambda s: s.astype(str).str.replace(r"\.0$", "", regex=True).str.replace("-", "").str.strip().str.zfill(10)
N = {}

rel_path, rel = cp.approved_release()
N["release"] = {"id": rel["release_id"], "status": rel["status"]}
R = pd.read_csv(rel_path, dtype=str, usecols=["part_number", "aim_l1_fixed", "aim_l2_fixed", "aim_l3_fixed"])
R["pn"] = z(R.part_number)
R = R.drop_duplicates("pn").set_index("pn")
lists, _, _ = load_lists()
lp = lists[2026]

COLS = {"Sales Doc.": "sales_doc", "Customer Reference": "cust_ref", "Material": "material",
        "Material description": "description", "SaTy": "order_type", "Req.Dlv.Dt": "order_date",
        "Open Qty": "open_qty", "Transit qty": "transit_qty", "ETA": "eta", "Priority": "priority",
        "Next ETA Supplier": "next_eta_supplier",
        "Tentative future ETA ISZE (1st incoming shipment)": "tentative_eta_isze",
        "Confirmed Qty": "confirmed_qty", "ORDER WK": "order_wk_eran", "TOTAL WEEKS": "total_weeks_eran",
        "Sold-to n°": "sold_to"}


def read_bo(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["DATA"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rows[3]]
    recs = [dict(zip(hdr, r)) for r in rows[4:] if r[4] not in (None, "")]
    df = pd.DataFrame(recs).rename(columns=COLS)
    return df[[c for c in COLS.values() if c in df.columns]], tuple(hdr)


man = list(csv.DictReader(open(INB / "manifest.csv", encoding="utf-8")))
weekly = [m for m in man if m["kind"] == "weekly_BO"]
frames, headers = [], {}
for m in weekly:
    wk = int(re.search(r"\(WK(\d+)\)", m["mail_subject"]).group(1))
    snap = pd.to_datetime(re.search(r"(\d\d-\d\d-\d{4})\.xlsx$", m["file"]).group(1), format="%d-%m-%Y")
    df, hdr = read_bo(INB / m["file"])
    headers.setdefault(hdr, []).append(wk)
    df.insert(0, "week", wk)
    df.insert(1, "snapshot_date", snap)
    frames.append(df)
N["header_variants"] = {" | ".join(h): w for h, w in headers.items()}
bo = pd.concat(frames, ignore_index=True)
N["sold_to_values"] = sorted(bo.sold_to.astype(str).unique().tolist())
bo["pn"] = z(bo.material)
bo["order_date"] = pd.to_datetime(bo.order_date, errors="coerce")
bo["open_qty"] = pd.to_numeric(bo.open_qty, errors="coerce")
bo["days_open"] = (bo.snapshot_date - bo.order_date).dt.days
bo["list_eur_2026"] = bo.pn.map(lp)
bo["value_eur_at_list"] = bo.open_qty * bo.list_eur_2026
bo["l1"] = bo.pn.map(R.aim_l1_fixed).fillna("Unclassified")
bo["l2"] = bo.pn.map(R.aim_l2_fixed).fillna("")
bo["l3"] = bo.pn.map(R.aim_l3_fixed).fillna("")
bo["seg"] = np.where(bo.l1 == "Lighting", "Lighting", np.where(bo.l1 != "Accident & Body", "Not body",
                     np.where(bo.l2 == "Panels", "Panels", "Other body")))
body_words = r"PANEL|BUMPER|LID;|DOOR|GRILLE|FENDER|HOOD|ROOF|COWL|STEP|MEMBER; CAB|MOLDING|DEFLECTOR"
not_body_words = r"HARNESS|OIL TANK|PARCEL|STEERING|TRAY"   # interior / non-body hits of the word list (checked by hand)
_d = bo.description.astype(str).str.upper()
bo["body_by_description_only"] = bo.seg.eq("Not body") & _d.str.contains(body_words, regex=True) & ~_d.str.contains(not_body_words, regex=True)
N["missing_list_price_lines"] = int(bo.list_eur_2026.isna().sum())
N["missing_order_date_lines"] = int(bo.order_date.isna().sum())

# ---- tidy trend table: body/panel lines
body = bo[bo.seg.isin(["Panels", "Other body"]) | bo.body_by_description_only].copy()
tidy_cols = ["week", "snapshot_date", "pn", "description", "open_qty", "order_date", "order_type", "days_open",
             "value_eur_at_list", "list_eur_2026", "seg", "l3", "body_by_description_only", "sales_doc",
             "cust_ref", "transit_qty", "eta", "next_eta_supplier", "tentative_eta_isze"]
body[tidy_cols].sort_values(["week", "seg", "pn"]).to_csv(HERE / "t13_uti_bo_body_lines_wk2_wk40.csv", index=False)


def wk_summary(d, label):
    g = d.groupby("week").agg(snapshot=("snapshot_date", "first"), lines=("pn", "size"), parts=("pn", "nunique"),
                              units=("open_qty", "sum"), value_eur_at_list=("value_eur_at_list", "sum"),
                              median_days_open=("days_open", "median"), max_days_open=("days_open", "max"),
                              lines_over_60d=("days_open", lambda s: int((s > 60).sum())),
                              vor_rush_lines=("order_type", lambda s: int(s.isin(["VOR", "RUSH"]).sum())))
    g.insert(0, "scope", label)
    return g.reindex(sorted(bo.week.unique())).assign(scope=label).fillna({"lines": 0, "units": 0, "value_eur_at_list": 0})


summ = pd.concat([wk_summary(bo, "All lines"),
                  wk_summary(bo[bo.seg.isin(["Panels", "Other body"]) | bo.body_by_description_only], "Body incl. description-only"),
                  wk_summary(bo[bo.seg == "Panels"], "Panels (release)")])
summ.reset_index().to_csv(HERE / "t13b_uti_bo_weekly_summary.csv", index=False)

# ---- recurring part numbers (body)
nweeks = bo.week.nunique()
rec = body.groupby("pn").agg(description=("description", "first"), seg=("seg", "first"),
                             weeks_on_list=("week", "nunique"), first_week=("week", "min"), last_week=("week", "max"),
                             distinct_orders=("sales_doc", "nunique"), max_days_open=("days_open", "max"),
                             max_open_units=("open_qty", "max"),
                             order_types=("order_type", lambda s: ",".join(sorted(set(map(str, s))))),
                             list_eur_2026=("list_eur_2026", "first")).reset_index()
rec = rec.sort_values(["weeks_on_list", "distinct_orders"], ascending=False)
rec.to_csv(HERE / "t13c_uti_bo_body_recurring_parts.csv", index=False)

# ---- Eran's 9 Jun copy vs ISZE's 9 Jun list
er, er_hdr = read_bo(INB / "ERAN_20260611_1436_BO OVERVIEW - UNIVERSAL TRUCKS ISRAEL LTD 09-06-2026.xlsx")
er["pn"] = z(er.material)
er["seg"] = er.pn.map(R.aim_l1_fixed).fillna("Unclassified")
N["eran_9jun"] = {"lines": int(len(er)), "header": list(er_hdr),
                  "lines_with_order_wk": int(er.order_wk_eran.notna().sum()) if "order_wk_eran" in er else 0,
                  "lines_with_total_weeks": int(pd.to_numeric(er.get("total_weeks_eran"), errors="coerce").notna().sum()),
                  "total_weeks_values": pd.to_numeric(er.get("total_weeks_eran"), errors="coerce").dropna().describe().round(1).to_dict(),
                  "isze_wk24_lines": int((bo.week == 24).sum())}
erm = er[pd.to_numeric(er.get("total_weeks_eran"), errors="coerce").notna()][
    ["sales_doc", "pn", "description", "order_type", "order_date", "order_wk_eran", "open_qty", "total_weeks_eran", "eta"]].copy()
erm["l1"] = erm.pn.map(R.aim_l1_fixed).fillna("Unclassified")
erm["l2"] = erm.pn.map(R.aim_l2_fixed).fillna("")
erm.to_csv(HERE / "t13d_eran_9jun_marked_lines.csv", index=False)

# ---- copy prices (AM parts samples.xlsx)
ws = openpyxl.load_workbook(INB / "AM parts samples.xlsx", read_only=True, data_only=True)["part no & price"]
am = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] is not None]
am = pd.DataFrame([r[:5] for r in am], columns=["pn_raw", "uti_desc", "copy_price_local_supplier_eur",
                                                "genuine_fob_isze_per_uti_eur", "copy_qty_sold_per_month"])
am["pn"] = z(am.pn_raw)
am["list_eur_2026"] = am.pn.map(lp)
am["reference_dealer_price_eur_x2_19_provisional"] = am.list_eur_2026 * 2.19
am["copy_vs_list"] = am.copy_price_local_supplier_eur / am.list_eur_2026 - 1
am["copy_vs_uti_fob"] = am.copy_price_local_supplier_eur / am.genuine_fob_isze_per_uti_eur - 1
am["copy_vs_ref_dealer_price"] = am.copy_price_local_supplier_eur / am.reference_dealer_price_eur_x2_19_provisional - 1
am["uti_fob_vs_list"] = am.genuine_fob_isze_per_uti_eur / am.list_eur_2026 - 1
am["release_l1"] = am.pn.map(R.aim_l1_fixed).fillna("Unclassified")
am["release_l2"] = am.pn.map(R.aim_l2_fixed).fillna("")
am["duplicate_row_in_source"] = am.duplicated("pn", keep="first")
am.to_csv(HERE / "t14_uti_copy_prices_vs_list.csv", index=False)
N["am_rows"] = int(len(am))
N["am_missing_list"] = am.loc[am.list_eur_2026.isna(), "pn"].tolist()

json.dump(N, open(HERE / "_numbers_bo_trend.json", "w"), indent=1, default=str)
pd.set_option("display.width", 250, "display.max_columns", 30, "display.max_rows", 200)
print(json.dumps(N, indent=1, default=str)[:3000])
print(summ.reset_index().to_string())
print(rec.head(30).to_string())
print(am.drop(columns=["pn_raw"]).round(2).to_string())
