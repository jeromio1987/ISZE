"""+5% price proposal by part number: factory-led, adjusted only where evidence is solid.

Reads the in-scope revision base (price_bp_weighted_compare_pn.csv, in_scope and movable:
47,617 lines, EUR 32,595,852.63, Jul 2025 - Jun 2026), attaches the captured-demand and
independent-price evidence of the 2026-09-20 specification workbook, and writes a live-formula
workbook. Nothing published; no price file or pricing configuration touched.

Run after dealer_response.py:
    python build_price_5pct_proposal_20260929.py
"""
import json
import os
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from evidence_lib import attach_evidence, load_base

HERE = Path(__file__).parent
ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
OUT_XLSX = Path(os.environ.get("PROPOSAL_OUT",
                               ISZE / "01_data" / "01_Pricing" / "Price_Increase_5pct_Proposal_2026-09-29.xlsx"))
OUT_JSON = HERE / "price_5pct_proposal_summary_2026-09-29.json"
RESP = json.loads((HERE / "dealer_price_response_2026-09-29.json").read_text(encoding="utf-8"))
CFG = yaml.safe_load((ISZE / "03_analytics" / "isze_pricing" / "config.yaml").read_text(encoding="utf-8"))

TARGET = 0.05
LOW, HIGH_BAND, FAR = 0.95, 1.15, 1.60          # spec-basis bands; config theta_veto premium_above 0.60
TIER_LIMIT = {"T1": CFG["allocation"]["tier_limit"]["T1_ks_observed"],
              "T2": CFG["allocation"]["tier_limit"]["T2_ks_derived"]}
CLASS_MARGIN = {k: 1 - v for k, v in CFG["formula"]["class_factor"].items()}
MARKETS = [("PR00", "PR00 (EU distributors)"), ("UK (ITUK + IUK)", "UK (ITUK trucks + IUK D-Max)"),
           ("UTI", "UTI Israel (1799)"), ("UMI", "UMI Israel (1797)"), ("CIS", "CIS (Armenia, Kazakhstan)")]
SRC = {"J": "Japan (IML)", "T": "Thailand (IMIT)"}
EXPECT = {"lines": 47617, "billed": 32595852.63, "recipe": 0.059616}

# ------------------------------------------------------------------ data
ins = load_base()
L = attach_evidence(ins[ins.movable].copy())
assert len(L) == EXPECT["lines"], len(L)
assert abs(L.net_sales_value.sum() - EXPECT["billed"]) < 0.01
L["c"] = L.bp_new_jpy / L.bp_prior_revised - 1
ks = L.evidence_tier.isin(["T1", "T2"])
L["idx_ks"] = np.where(ks, L.ref_dealer_price_eur / L.independent_price_eur, np.nan)
L["lost"] = (L.zone_basis == "Measured") & (L.zone == "Lost")
L["room_line"] = ks & (L.idx_ks < LOW) & ~L.lost
L["far_line"] = ks & (L.idx_ks > FAR) & ~L.lost
L["room"] = np.where(L.room_line, np.minimum(L.evidence_tier.map(TIER_LIMIT).fillna(0),
                                             np.clip(HIGH_BAND / L.idx_ks - 1, 0, None)), 0.0)
w = L.net_sales_value


def pn_text(d):
    bits = []
    lost = (d[d.lost].sort_values("net_sales_value", ascending=False)
            .drop_duplicates(["distributor_code", "service_category"]))
    if len(lost):
        parts = [f"{r.distributor_code} {str(r.service_category).replace('_', ' ')} {r.capture_rate:.0%}"
                 for r in lost.sort_values("net_sales_value", ascending=False).head(3).itertuples()]
        bits.append("Distributor sells under 45% of its trucks' need: " + ", ".join(parts))
    kk = d[d.idx_ks.notna()]
    if len(kk):
        ix = np.average(kk.idx_ks, weights=kk.net_sales_value) if kk.net_sales_value.sum() > 0 else kk.idx_ks.mean()
        bits.append(f"UK shop (KS), same part: our dealer price = {ix:.2f}x shop price")
    if (d.evidence_tier == "T3").any():
        bits.append("Eucon competitor price on file (not used: unclear if trade or retail)")
    return "; ".join(bits)


g = L.groupby("pn")
P = pd.DataFrame({
    "billed": g.net_sales_value.sum(),
    "qty": g.billed_qty.sum(),
    "c": g.c.first(),
    "bp0": g.bp_prior_revised.first(),
    "bp1": g.bp_new_jpy.first(),
    "cls": g.bp_new_class.first(),
    "src": g.supply_source.first().map(SRC).fillna("No source label"),
    "list": g.sp_eur_pr00_ie40.first(),
})
top_desc = L.sort_values("net_sales_value", ascending=False).drop_duplicates("pn").set_index("pn").billed_description
P["desc"] = top_desc.reindex(P.index).fillna("").astype(str).str.slice(0, 80)
# Jerome 2026-09-30: "you should leave the pad kits at same price" -> brake pad kits held at today's price
P["padkit"] = top_desc.reindex(P.index).fillna("").astype(str).str.upper().str.contains(r"PAD\s*KIT").astype(int)
P["h_lost"] = L[L.lost].groupby("pn").net_sales_value.sum().reindex(P.index).fillna(0) / P.billed
P["h_far"] = L[L.far_line].groupby("pn").net_sales_value.sum().reindex(P.index).fillna(0) / P.billed
P["r"] = (L.room * w).groupby(L.pn).sum().reindex(P.index).fillna(0) / P.billed
kk = L[L.idx_ks.notna()]
P["idx_ks"] = (kk.idx_ks * kk.net_sales_value).groupby(kk.pn).sum() / kk.groupby("pn").net_sales_value.sum()
P["evidence"] = g.apply(pn_text, include_groups=False)
P["tier"] = g.evidence_tier.min()
mk = L.pivot_table(index="pn", columns="market", values="net_sales_value", aggfunc="sum", fill_value=0)
for code, _ in MARKETS:
    P[f"m_{code}"] = mk[code].reindex(P.index).fillna(0) if code in mk else 0.0
P["main_market"] = mk.idxmax(axis=1).reindex(P.index)
P["margin"] = P.cls.map(CLASS_MARGIN)
P = P.sort_values("billed", ascending=False)
P.index.name = "pn"
P["h_lost"] = P.h_lost.clip(0, 1).round(6)
P["h_far"] = P.h_far.clip(0, 1).round(6)
P["r"] = P.r.round(6)

# expected values with default switches (lost hold on, room on, far hold off, pad kits kept) - the verifier
# compares Excel with these
cpos = P.c.clip(lower=0)
keep = 1 - P.padkit
k = (TARGET * P.billed.sum() - (P.billed * P.r * keep).sum()) / (P.billed * cpos * (1 - P.h_lost) * keep).sum()
P["delta"] = keep * (cpos * k * (1 - P.h_lost) + P.r)
pads = P[P.padkit == 1]
pad_margin_after = 1 - (1 - pads.margin) * (1 + pads.c)
k0 = TARGET * P.billed.sum() / (P.billed * cpos).sum()
exp = {
    "k": k, "k_factory_only": k0,
    "achieved": float(np.average(P.delta, weights=P.billed)),
    "step_eur": float((P.billed * P.delta).sum()),
    "parts": int(len(P)), "parts_at_zero": int((P.delta <= 1e-5).sum()),
    "eur_at_zero": float(P.billed[P.delta <= 1e-5].sum()),
    "median_move": float(P.delta.median()), "p90_move": float(P.delta.quantile(0.9)), "max_move": float(P.delta.max()),
    "recipe_check": float((P.qty * (P.bp1 - P.bp0)).sum() / (P.qty * P.bp0).sum()),
    "pad_kits": int(len(pads)), "pad_kits_billed": float(pads.billed.sum()),
    "pad_kits_pns": pads.index.tolist(),
    "by_market": {c: float(np.average(P.delta, weights=P[f"m_{c}"])) for c, _ in MARKETS},
}
mm = P.margin.notna()
exp["margin_pts_proposal"] = float(np.average(((1 - P.margin) * (1 - (1 + P.c) / (1 + P.delta)))[mm], weights=P.billed[mm]))
exp["margin_pts_flat"] = float(np.average(((1 - P.margin) * (1 - (1 + P.c) / 1.05))[mm], weights=P.billed[mm]))
d0 = cpos * k0
exp["margin_pts_factory_only"] = float(np.average(((1 - P.margin) * (1 - (1 + P.c) / (1 + d0)))[mm], weights=P.billed[mm]))
for lab, dd in [("flat", 0.05), ("factory_only", d0), ("proposal", P.delta)]:
    pts = (1 - P.margin) * (1 - (1 + P.c) / (1 + dd))
    exp[f"eur_lose_1pt_{lab}"] = float(P.billed[mm & (pts < -0.01)].sum())
for name, e in [("none", 0.0), ("central", RESP["scenario_elasticities"]["central"]),
                ("stress", RESP["scenario_elasticities"]["stress"])]:
    exp[f"revenue_{name}"] = float((P.billed * ((1 + P.delta) * (1 + e * P.delta) - 1)).sum())
print(json.dumps(exp, indent=1, default=float))

# ------------------------------------------------------------------ not in proposal
un = ins[~ins.movable].copy()
desc = un.billed_description.fillna("").str.upper()
lubr = un.pn.str.startswith("05") & desc.str.contains("RIM|TELLUS|SHELL|HELIX|SPIRAX|OMALA|GADUS")
nonpart = ~un.pn.str.fullmatch(r"\d+")
un["why"] = np.select(
    [lubr, nonpart, un.bp_new_jpy.notna() & un.bp_prior_revised.isna(),
     un.bp_new_jpy.isna() & un.bp_prior_revised.notna()],
    ["Lubricant (separate business, out of scope)", "Not a part: rebate, transport or service line",
     "New factory price only: needs a first price", "No new factory price: not in the new price file"],
    default="No factory price on either side")

# ------------------------------------------------------------------ workbook
NAVY, CREAM = "15233C", "F7F4EE"
F = "Arial"
HFILL = PatternFill("solid", fgColor=NAVY)
HFONT = Font(name=F, bold=True, color=CREAM, size=10)
TITLE = Font(name=F, bold=True, size=14, color=NAVY)
SUB = Font(name=F, italic=True, size=9, color="555555")
SEC = Font(name=F, bold=True, size=11, color=NAVY)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF", bold=True)
LINK = Font(name=F, size=10, color="008000")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
WRAP = Alignment(wrap_text=True, vertical="top")
PCT1, PCT2, EUR0, EUR2 = "0.0%;-0.0%;-", "0.00%;-0.00%;-", "#,##0;(#,##0);-", "#,##0.00;(#,##0.00);-"

wb = Workbook()
ws_read = wb.active
ws_read.title = "Read me"
ws_sum = wb.create_sheet("Summary")
ws_p = wb.create_sheet("Proposal")
ws_rev = wb.create_sheet("Review lists")
ws_dr = wb.create_sheet("Dealer response")
ws_out = wb.create_sheet("Not in proposal")
ws_ck = wb.create_sheet("Checks")
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, col=1, height=30):
    for i, h in enumerate(labels):
        c = ws.cell(row, col + i, h)
        c.fill, c.font = HFILL, HFONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def put(ws, row, col, value, font=BODY, fmt=None, align=None, fill=None):
    c = ws.cell(row, col, value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if fill:
        c.fill = fill
    return c


# ---- Proposal sheet ---------------------------------------------------------
N = len(P)
first, last = 2, N + 1
COLS = ["Part number", "Description", "Price class", "Class margin", "Source of supply", "Main market",
        "Billed EUR (Jul 25-Jun 26)", "Quantity", "Current list EUR (PR00)", "Factory base price before (JPY)",
        "Factory base price after (JPY)", "Factory move", "Share held: distributor sells under 45% of need",
        "Share held: far above UK shop price", "Extra room below UK shop price", "Proposed move",
        "New list EUR (PR00)", "Extra EUR on last 12 months", "Class margin change (points)", "Rule applied",
        "Evidence on file", "Our dealer price / UK shop price"] + \
       [f"Billed EUR: {lab}" for _, lab in MARKETS] + ["Factory-led only move (comparison)", "Check flag",
                                                        "Brake pad kit: kept at today's price (1 = yes)"]
header(ws_p, 1, COLS, height=54)
S = "Summary!"
for i, (pn, r) in enumerate(P.iterrows()):
    rr = i + 2
    ws_p.cell(rr, 1, str(pn))
    ws_p.cell(rr, 2, r.desc)
    ws_p.cell(rr, 3, None if pd.isna(r.cls) else str(r.cls))
    ws_p.cell(rr, 4, None if pd.isna(r.margin) else round(float(r.margin), 4))
    ws_p.cell(rr, 5, r.src)
    ws_p.cell(rr, 6, None if pd.isna(r.main_market) else str(r.main_market))
    ws_p.cell(rr, 7, round(float(r.billed), 2))
    ws_p.cell(rr, 8, float(r.qty))
    ws_p.cell(rr, 9, None if pd.isna(r.list) or r.list <= 0 else round(float(r.list), 2))
    ws_p.cell(rr, 10, float(r.bp0))
    ws_p.cell(rr, 11, float(r.bp1))
    ws_p.cell(rr, 12, f"=K{rr}/J{rr}-1")
    ws_p.cell(rr, 13, float(r.h_lost))
    ws_p.cell(rr, 14, float(r.h_far))
    ws_p.cell(rr, 15, float(r.r))
    ws_p.cell(rr, 16, f"=(1-{S}$B$9*AD{rr})*(MAX(L{rr},0)*{S}$B$10*(1-{S}$B$6*M{rr}-{S}$B$8*N{rr})+{S}$B$7*O{rr})")
    ws_p.cell(rr, 17, f'=IF(I{rr}>0,ROUND(I{rr}*(1+P{rr}),2),"")')
    ws_p.cell(rr, 18, f"=G{rr}*P{rr}")
    ws_p.cell(rr, 19, f'=IF(D{rr}="","",(1-D{rr})*(1-(1+L{rr})/(1+P{rr})))')
    ws_p.cell(rr, 20, (f'=IF(AND({S}$B$9=1,AD{rr}=1),"Held: pad kit, price kept",'
                       f'IF(P{rr}<=0.00001,IF(L{rr}<=0,"No factory increase: held","Held: "&'
                       f'IF(AND({S}$B$6=1,M{rr}>0),"distributor sells under 45% of need","far above UK shop price")),'
                       f'IF(AND({S}$B$7=1,O{rr}>0),"Factory move + room below UK shop price",'
                       f'IF(OR(AND({S}$B$6=1,M{rr}>0),AND({S}$B$8=1,N{rr}>0)),"Factory move, partly held","Factory move"))))'))
    ws_p.cell(rr, 21, r.evidence)
    ws_p.cell(rr, 22, None if pd.isna(r.idx_ks) else round(float(r.idx_ks), 3))
    for j, (code, _) in enumerate(MARKETS):
        ws_p.cell(rr, 23 + j, round(float(r[f"m_{code}"]), 2))
    ws_p.cell(rr, 28, f"=MAX(L{rr},0)*{S}$B$26")
    ws_p.cell(rr, 29, f'=IF(L{rr}>0.15,"Factory move above 15%: check","")')
    ws_p.cell(rr, 30, int(r.padkit))

fmt = {4: PCT1, 7: EUR0, 8: "#,##0", 9: EUR2, 10: "#,##0", 11: "#,##0", 12: PCT2, 13: "0%", 14: "0%",
       15: PCT2, 16: PCT2, 17: EUR2, 18: EUR0, 19: PCT2, 22: "0.00", 23: EUR0, 24: EUR0, 25: EUR0, 26: EUR0,
       27: EUR0, 28: PCT2}
for row in ws_p.iter_rows(min_row=2, max_row=last):
    for c in row:
        c.font = BODY
        if c.column in fmt:
            c.number_format = fmt[c.column]
    row[15].font = BOLD
    row[16].font = BOLD
widths = [13, 34, 7, 8, 15, 16, 13, 10, 12, 12, 12, 9, 12, 12, 11, 10, 12, 12, 11, 34, 60, 12,
          13, 13, 13, 13, 13, 12, 26, 12]
for i, wd in enumerate(widths, start=1):
    ws_p.column_dimensions[get_column_letter(i)].width = wd
ws_p.freeze_panes = "C2"
ws_p.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"
rng = lambda col: f"Proposal!${col}${first}:${col}${last}"
G, Lr, Mr, Nr, Or, Pr, Rr, Sr, Tr, Dr, Cr, Er = (rng(x) for x in "GLMNOPRSTDCE")
ADr = rng("AD")

# ---- Summary sheet ------------------------------------------------------------
ws = ws_sum
put(ws, 1, 1, "Price proposal +5%: best allocation by part number", TITLE)
put(ws, 2, 1, "A proposal on paper. Nothing published, no price file, price list or pricing configuration changed. "
              "Built 29 September 2026 on what distributors bought July 2025 to June 2026.", SUB)
put(ws, 4, 1, "Settings - change the yellow cells, every figure below recalculates", SEC)
settings = [
    (5, "Target: weighted increase on what distributors bought in the last 12 months", TARGET, PCT1),
    (6, "Hold at 0% where the distributor sells under 45% of what its trucks need in that category (1 = on, 0 = off)", 1, "0"),
    (7, "Extra where the UK shop (KS) sells the same part number above our reference dealer price (1 = on, 0 = off)", 1, "0"),
    (8, "Hold at 0% where our reference dealer price is more than 1.6x the UK shop price (1 = on, 0 = off)", 0, "0"),
    (9, "Keep brake pad kits at today's price (1 = on, 0 = off) - Jerome, 30 September 2026", 1, "0"),
]
for r, lab, v, f in settings:
    put(ws, r, 1, lab, BODY, align=WRAP)
    put(ws, r, 2, v, INPUT, f, fill=YELLOW)
put(ws, 10, 1, "Scaling applied to each part's own factory move (calculated so the book lands on the target)", BODY, align=WRAP)
put(ws, 10, 2, (f"=($B$5*SUM({G})-$B$7*SUMPRODUCT({G},{Or},1-$B$9*{ADr}))/"
                f"SUMPRODUCT({G},({Lr}>0)*{Lr},1-$B$6*{Mr}-$B$8*{Nr},1-$B$9*{ADr})"), BODY, "0.0000")
put(ws, 10, 3, "e.g. 0.90 = each part moves 90% of its own factory move", SUB)

put(ws, 11, 1, "Result", SEC)
res = [
    (12, "Weighted increase achieved", f"=SUMPRODUCT({G},{Pr})/SUM({G})", PCT2),
    (13, "Extra EUR on the last 12 months' volume (before any dealer reaction)", f"=SUM({Rr})", EUR0),
    (14, "Billed base, EUR", f"=SUM({G})", EUR0),
    (15, "Part numbers priced", f"=COUNTA(Proposal!$A${first}:$A${last})", "#,##0"),
    # comparisons inside SUMPRODUCT, never decimal criteria strings: those break on comma-decimal Excel
    (16, "Part numbers held at 0%", f"=SUMPRODUCT(--({Pr}<=0.00001))", "#,##0"),
    (17, "Billed EUR on parts held at 0%", f"=SUMPRODUCT(({Pr}<=0.00001)*{G})", EUR0),
    (18, "Typical part (median move)", f"=MEDIAN({Pr})", PCT2),
    (19, "90% of parts move less than", f"=PERCENTILE({Pr},0.9)", PCT2),
    (20, "Largest single move (see Review lists: factory moves above 15%)", f"=MAX({Pr})", PCT1),
    (21, "Class margin change, billed-weighted, in points", f'=SUMPRODUCT({G},{Sr})/SUMIF({Dr},">0",{G})', PCT2),
]
for r, lab, f_, nf in res:
    put(ws, r, 1, lab, BODY, align=WRAP)
    put(ws, r, 2, f_, BOLD, nf)
put(ws, 22, 1, "Margin change holds exchange rate and purchase discount where they are: it shows only what the factory "
              "move and our price move do to each class margin.", SUB)

put(ws, 23, 1, "Three ways to reach the same +5%", SEC)
header(ws, 24, ["Option", "Scaling", "Extra EUR", "Parts at 0%", "Typical part move", "Class margin change (points)",
                 "Billed EUR on parts losing more than 1 margin point", "How it works"], height=44)
k0_formula = f"=$B$5*SUM({G})/SUMPRODUCT({G},({Lr}>0)*{Lr})"
AB = rng("AB")
opts = [
    ("Flat +5% on every part", None, f"=$B$5*SUM({G})", 0, "=$B$5",
     f'=SUMPRODUCT({G},({Dr}>0)*(1-{Dr})*(1-(1+{Lr})/(1+$B$5)))/SUMIF({Dr},">0",{G})',
     f'=SUMPRODUCT({G},({Dr}>0)*((1-{Dr})*(1-(1+{Lr})/(1+$B$5))<-0.01))',
     "Every part +5%. Parts with a big factory increase lose margin, parts with none gain it."),
    ("Factory-led only", k0_formula, f"=$B$5*SUM({G})", f'=COUNTIF({Lr},"<=0")', f"=MEDIAN({AB})",
     f'=SUMPRODUCT({G},({Dr}>0)*(1-{Dr})*(1-(1+{Lr})/(1+{AB})))/SUMIF({Dr},">0",{G})',
     f'=SUMPRODUCT({G},({Dr}>0)*((1-{Dr})*(1-(1+{Lr})/(1+{AB}))<-0.01))',
     "Each part moves in line with its own factory move, scaled to +5% overall."),
    ("This proposal", "=$B$10", "=$B$13", "=$B$16", "=$B$18", "=$B$21", f"=SUMPRODUCT({G},--({Sr}<-0.01))",
     "Factory-led; pad kits kept at today's price; held where the distributor already loses the category; extra where the UK shop sells above us."),
]
for i, (lab, sc, eur, zero, typ, mar, lose, how) in enumerate(opts):
    r = 25 + i
    put(ws, r, 1, lab, BOLD)
    put(ws, r, 2, sc, BODY, "0.0000")
    put(ws, r, 3, eur, BODY, EUR0)
    put(ws, r, 4, zero, BODY, "#,##0")
    put(ws, r, 5, typ, BODY, PCT2)
    put(ws, r, 6, mar, BODY, PCT2)
    put(ws, r, 7, lose, BOLD, EUR0)
    put(ws, r, 8, how, BODY)
    for cc in range(1, 9):
        ws.cell(r, cc).alignment = Alignment(vertical="center", wrap_text=(cc == 8))
    ws.row_dimensions[r].height = 30

put(ws, 29, 1, "What if dealers buy less after the increase", SEC)
put(ws, 30, 1, "Dealer response = % change in dealer purchases for +1% on our price. Measured on sales-to-dealer data, see 'Dealer response'.", SUB)
header(ws, 31, ["Scenario", "Dealer response", "Dealer volume change", "Extra EUR after volume effect", "As % of billed base", "Where the response figure comes from"], height=44)
ws.merge_cells("F31:H31")
sc = RESP["scenario_elasticities"]
scen = [("No reaction", sc["none"], "What the April 2026 increase showed on the UK truck network: volume up 3.7% March-August"),
        ("Central", sc["central"], "Four markets pooled, April 2026 list change (UK trucks, UK D-Max, Ireland, Portugal)"),
        ("Stress", sc["stress"], "Lowest 95% bound across the pooled measurements: a pessimistic case, not a forecast")]
for i, (lab, e, src) in enumerate(scen):
    r = 32 + i
    put(ws, r, 1, lab, BOLD)
    put(ws, r, 2, e, INPUT, "0.00", fill=YELLOW)
    put(ws, r, 3, f"=SUMPRODUCT({G},B{r}*{Pr})/SUM({G})", BODY, PCT2)
    put(ws, r, 4, f"=SUMPRODUCT({G},(1+{Pr})*(1+B{r}*{Pr}))-SUM({G})", BOLD, EUR0)
    put(ws, r, 5, f"=D{r}/SUM({G})", BODY, PCT2)
    put(ws, r, 6, src, SUB)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    for cc in range(1, 7):
        ws.cell(r, cc).alignment = Alignment(vertical="center", wrap_text=(cc == 6))
    ws.row_dimensions[r].height = 22

put(ws, 36, 1, "By market (one list price per part; Israel and CIS lists move by the same % on their own price)", SEC)
header(ws, 37, ["Market", "Billed EUR", "Proposed weighted increase", "Extra EUR", "Factory move, same weights"], height=44)
for i, (code, lab) in enumerate(MARKETS):
    r, col = 38 + i, get_column_letter(23 + i)
    M_ = rng(col)
    put(ws, r, 1, lab)
    put(ws, r, 2, f"=SUM({M_})", BODY, EUR0)
    put(ws, r, 3, f"=SUMPRODUCT({M_},{Pr})/SUM({M_})", BOLD, PCT2)
    put(ws, r, 4, f"=SUMPRODUCT({M_},{Pr})", BODY, EUR0)
    put(ws, r, 5, f"=SUMPRODUCT({M_},{Lr})/SUM({M_})", BODY, PCT2)
put(ws, 43, 1, "All markets", BOLD)
put(ws, 43, 2, "=SUM(B38:B42)", BOLD, EUR0)
put(ws, 43, 3, "=D43/B43", BOLD, PCT2)
put(ws, 43, 4, "=SUM(D38:D42)", BOLD, EUR0)
put(ws, 43, 5, f"=SUMPRODUCT({G},{Lr})/SUM({G})", BOLD, PCT2)

put(ws, 45, 1, "By source of supply", SEC)
header(ws, 46, ["Source", "Billed EUR", "Proposed weighted increase", "Extra EUR", "Factory move, same weights"], height=44)
for i, s in enumerate(["Japan (IML)", "Thailand (IMIT)", "No source label"]):
    r = 47 + i
    put(ws, r, 1, s)
    put(ws, r, 2, f'=SUMIF({Er},A{r},{G})', BODY, EUR0)
    put(ws, r, 3, f"=D{r}/B{r}", BOLD, PCT2)
    put(ws, r, 4, f'=SUMPRODUCT(({Er}=A{r})*{G},{Pr})', BODY, EUR0)
    put(ws, r, 5, f'=SUMPRODUCT(({Er}=A{r})*{G},{Lr})/B{r}', BODY, PCT2)

put(ws, 51, 1, "By price class (margin as issued by IML, config.yaml)", SEC)
header(ws, 52, ["Class", "Class margin", "Billed EUR", "Proposed weighted increase", "Factory move, same weights", "Class margin change (points)"], height=44)
classes = [c for c in ["C", "D", "F", "J", "M", "P", "V"] if (P.cls == c).any()]
for i, c in enumerate(classes):
    r = 53 + i
    put(ws, r, 1, c, BOLD)
    put(ws, r, 2, CLASS_MARGIN[c], BODY, PCT1)
    put(ws, r, 3, f'=SUMIF({Cr},A{r},{G})', BODY, EUR0)
    put(ws, r, 4, f'=SUMPRODUCT(({Cr}=A{r})*{G},{Pr})/C{r}', BOLD, PCT2)
    put(ws, r, 5, f'=SUMPRODUCT(({Cr}=A{r})*{G},{Lr})/C{r}', BODY, PCT2)
    put(ws, r, 6, f'=SUMPRODUCT(({Cr}=A{r})*{G},{Sr})/C{r}', BODY, PCT2)
rc = 53 + len(classes) + 1

put(ws, rc, 1, "By rule applied", SEC)
header(ws, rc + 1, ["Rule", "Part numbers", "Billed EUR", "Extra EUR", "Average move"])
rules = ["Factory move", "Factory move, partly held", "Factory move + room below UK shop price",
         "Held: pad kit, price kept", "Held: distributor sells under 45% of need", "Held: far above UK shop price",
         "No factory increase: held"]
for i, rl in enumerate(rules):
    r = rc + 2 + i
    put(ws, r, 1, rl)
    put(ws, r, 2, f'=COUNTIF({Tr},A{r})', BODY, "#,##0")
    put(ws, r, 3, f'=SUMIF({Tr},A{r},{G})', BODY, EUR0)
    put(ws, r, 4, f'=SUMIF({Tr},A{r},{Rr})', BODY, EUR0)
    put(ws, r, 5, f'=IF(C{r}=0,0,D{r}/C{r})', BODY, PCT2)
rt = rc + 2 + len(rules)
put(ws, rt, 1, "Total", BOLD)
for col in "BCD":
    put(ws, rt, "ABCD".index(col) + 1, f"=SUM({col}{rc + 2}:{col}{rt - 1})", BOLD, EUR0 if col != "B" else "#,##0")
put(ws, rt + 2, 1, ("Sources: revision base 01_data/00_derived/price_bp_weighted_compare_pn.csv (in scope, both factory prices "
                    "present); current list pricing_pn_enquiry.parquet; distributor share of need pms_capture_by_db_year.csv "
                    "(CY2025); UK shop prices iam_ks_price_gap.csv; class margins 03_analytics/isze_pricing/config.yaml; "
                    "dealer response 02_dashboards/pricing/_build/price_5pct_proposal_2026-09-29/."), SUB, align=WRAP)
ws.merge_cells(start_row=rt + 2, start_column=1, end_row=rt + 2, end_column=8)
ws.row_dimensions[rt + 2].height = 40
for col, wd in zip("ABCDEFGH", [62, 14, 14, 16, 16, 18, 20, 60]):
    ws.column_dimensions[col].width = wd
for r in range(5, 11):
    ws.row_dimensions[r].height = 28
ws.freeze_panes = "A4"

# ---- Review lists ---------------------------------------------------------------
ws = ws_rev
put(ws, 1, 1, "Review lists - the parts worth a look before anything goes out", TITLE)
put(ws, 2, 1, "Proposed move and new list price read live from the Proposal sheet.", SUB)
row = 4
look = lambda col, r: f"=INDEX(Proposal!${col}${first}:${col}${last},MATCH($A{r},Proposal!$A${first}:$A${last},0))"


def section(title, note, frame, col8=("Rule applied", "T", None)):
    global row
    put(ws, row, 1, f"{title} ({len(frame)} part numbers)", SEC)
    put(ws, row + 1, 1, note, SUB)
    header(ws, row + 2, ["Part number", "Description", "Billed EUR", "Factory move", "Proposed move",
                         "Current list EUR", "New list EUR", col8[0], "Evidence on file"])
    row += 3
    for pn, r in frame.iterrows():
        put(ws, row, 1, str(pn))
        put(ws, row, 2, r.desc)
        put(ws, row, 3, round(float(r.billed), 2), BODY, EUR0)
        put(ws, row, 4, look("L", row), LINK, PCT1)
        put(ws, row, 5, look("P", row), LINK, PCT2)
        put(ws, row, 6, look("I", row), LINK, EUR2)
        put(ws, row, 7, look("Q", row), LINK, EUR2)
        put(ws, row, 8, look(col8[1], row), LINK, col8[2])
        put(ws, row, 9, r.evidence, BODY)
        row += 1
    row += 1


section("Top 30 part numbers by billed euro", "Together these carry the largest share of the extra euro.", P.head(30))
section("Brake pad kits kept at today's price (Jerome, 30 September 2026)",
        "Proposed move 0% while Summary B9 is on. Column H shows what that does to each kit's class margin, holding "
        "exchange rate and purchase discount where they are.", P[P.padkit == 1],
        col8=("Class margin change (points)", "S", PCT2))
section("Held: the distributor sells under 45% of what its trucks need in that category",
        "Measured, CY2025: billed demand against the demand the fleet should generate. Part is held on that "
        "distributor's share of its volume.", P[P.h_lost > 0])
section("Extra: the UK shop (KS) sells the same part number above our reference dealer price",
        "Room is capped at the evidence limit in config.yaml (10 points same shop, 6 points derived) and at "
        "1.15x the shop price.", P[P.r > 0])
section("Far above the UK shop price (more than 1.6x) - switch off by default",
        "Not held by default: the dealer data shows these parts did not lose more volume after the April 2026 "
        "increase than comparable parts. Turn on Summary B8 to hold them.", P[P.h_far > 0])
section("Factory move above 15% - passed through, please check",
        "The factory base price moved more than 15% (some doubled). Passed through at the same scaling; "
        "no per-part ceiling (removed 20 September 2026).", P[P.c > 0.15])
for col, wd in zip("ABCDEFGHI", [13, 36, 12, 10, 10, 12, 12, 34, 80]):
    ws.column_dimensions[col].width = wd
ws.freeze_panes = "A4"

# ---- Dealer response ------------------------------------------------------------
ws = ws_dr
put(ws, 1, 1, "Dealer response: do dealers buy less of a part when our list price for it rises?", TITLE)
put(ws, 2, 1, ("Measured on sales-to-dealer data: UK trucks (ITUK) monthly Jan 2023 - Aug 2026, and UK D-Max (IUK), "
               "Ireland and Portugal monthly 2025 - 2026, against our dated list prices."), SUB)
verdict = [
    "Verdict: at part level, dealer purchases did not measurably react to our list price.",
    ("The April 2026 list change raised UK-relevant prices by a median 11.7% (confirmed on UK invoices: 11.6%). "
     "Parts that got more of it did not lose more dealer volume than similar parts that got less, in any of the four markets "
     "taken together. The pooled figure is about -0.16 (a 10% price rise costs about 1.6% of volume), with a range "
     "that runs from roughly -0.6 to +0.3."),
    ("A placebo test (does the April change 'predict' volume in the two months before it happened?) gives a figure "
     "of the same size, so the measured response is at the level of noise. It is not strong enough to set one part "
     "higher than another, so no part price in the proposal is set from it. It is used only for the volume "
     "scenarios on the Summary."),
    ("What it cannot see: a whole family switching to the aftermarket (the comparison is within a family), and "
     "slow erosion over several years. Dealer invoice prices are not in any file, so how each distributor passed "
     "our increase to its dealers is not measured; dealers stocking up in March 2026 (+14% on March 2025, UK trucks) "
     "suggests the UK distributor did pass it on."),
]
r = 4
for i, t in enumerate(verdict):
    put(ws, r, 1, t, BOLD if i == 0 else BODY, align=WRAP)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 18 if i == 0 else 44
    r += 1
r += 1
put(ws, r, 1, "1. UK truck distributor (ITUK), three list changes", SEC)
r += 1
header(ws, r, ["List change", "Parts", "Weighted list change", "Spread within family", "Dealer response",
               "Uncertainty (+/-)", "95% range low", "95% range high"])
r += 1
for ev, v in RESP["ituk_three_list_changes"]["events"].items():
    vals = [ev, v["parts"], v["qty_weighted_list_change"], v["within_family_spread_sd"], v["beta_dlp"], v["se_dlp"],
            f"=E{r}-1.96*F{r}", f"=E{r}+1.96*F{r}"]
    for j, x in enumerate(vals):
        put(ws, r, j + 1, x, BODY, [None, "#,##0", PCT1, PCT1, "0.00", "0.00", "0.00", "0.00"][j])
    r += 1
for lab, key in [("All three pooled", "pooled"), ("Pooled, part numbers not superseded", "pooled_excl_superseded"),
                 ("Pooled, parts with 10+ units", "pooled_min10_units")]:
    v = RESP["ituk_three_list_changes"][key]
    for j, x in enumerate([lab, v["n"], None, None, v["beta_dlp"], v["se_dlp"], f"=E{r}-1.96*F{r}", f"=E{r}+1.96*F{r}"]):
        put(ws, r, j + 1, x, BOLD if j == 0 else BODY, [None, "#,##0", None, None, "0.00", "0.00", "0.00", "0.00"][j])
    r += 1
for ev, v in RESP["ituk_three_list_changes"]["placebo_previous_year_trend"].items():
    for j, x in enumerate([f"Placebo: {ev} change vs previous year's trend", v["n"], None, None, v["beta_dlp"], v["se_dlp"],
                           f"=E{r}-1.96*F{r}", f"=E{r}+1.96*F{r}"]):
        put(ws, r, j + 1, x, SUB if j == 0 else BODY, [None, "#,##0", None, None, "0.00", "0.00", "0.00", "0.00"][j])
    r += 1
r += 1
put(ws, r, 1, "2. April 2026 list change, four markets", SEC)
r += 1
header(ws, r, ["Market", "Parts", "Weighted list change", "Dealer units before", "Dealer units after", "Volume change",
               "Dealer response", "Uncertainty (+/-)"])
r += 1
four = RESP["four_markets_april_2026"]
for mkt, v in four["by_market"].items():
    u = four["dealer_units"][mkt]
    for j, x in enumerate([f"{mkt} ({u['window']})", v["n"], v["qty_weighted_list_change"], u["units_before"],
                           u["units_after"], u["change"], v["beta_dlp"], v["se_dlp"]]):
        put(ws, r, j + 1, x, BODY, [None, "#,##0", PCT1, "#,##0", "#,##0", PCT1, "0.00", "0.00"][j])
    r += 1
for lab, key in [("Four markets pooled", "pooled"), ("Pooled, part numbers not superseded", "pooled_excl_superseded"),
                 ("Pooled, parts with 10+ units", "pooled_min10_units"),
                 ("Placebo: April change vs Jan-Feb growth (before it happened)", "placebo_jan_feb")]:
    v = four[key]
    for j, x in enumerate([lab, v["n"], None, None, None, None, v["beta_dlp"], v["se_dlp"]]):
        put(ws, r, j + 1, x, BOLD if j == 0 else BODY, [None, "#,##0", None, None, None, None, "0.00", "0.00"][j])
    r += 1
put(ws, r, 1, "Ireland's January-March 2026 file is incomplete (Jan-Feb -26%, March -70% on 2025); April-July looks complete.", SUB)
r += 2
put(ws, r, 1, "3. Do some kinds of part react more? (extra response on top of the pooled figure)", SEC)
r += 1
header(ws, r, ["Kind of part", "Share of parts", "Extra response", "Uncertainty (+/-)"])
r += 1
for lab, v in four["segments"].items():
    for j, x in enumerate([lab, v["share_of_parts"], v["extra_response"], v["extra_response_se"]]):
        put(ws, r, j + 1, x, BODY, [None, PCT1, "0.00", "0.00"][j])
    r += 1
put(ws, r, 1, "None differs measurably. Parts with a competitor price on file reacted slightly less, not more.", SUB)
r += 2
put(ws, r, 1, "4. UK truck dealer purchases by month (units)", SEC)
r += 1
header(ws, r, ["Month", "2023", "2024", "2025", "2026", "2026 vs 2025"])
r0 = r + 1
for mo in range(1, 13):
    r += 1
    v = RESP["ituk_dealer_units_by_month"][str(mo)]
    put(ws, r, 1, ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][mo - 1])
    for j, y in enumerate(["2023", "2024", "2025", "2026"]):
        put(ws, r, j + 2, v.get(y), BODY, "#,##0")
    put(ws, r, 6, f'=IF(E{r}="","",E{r}/D{r}-1)', BODY, PCT1)
ch = LineChart()
ch.title = "UK truck dealer purchases by month (units)"
ch.height, ch.width = 7.5, 15
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=r0 - 1, max_row=r), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r))
for s_ in ch.series:
    s_.smooth = False
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.number_format = "#,##0"
ch.legend.position = "b"
ws.add_chart(ch, f"A{r + 2}")
r += 17
r += 2
put(ws, r, 1, "5. Checks behind the method", SEC)
r += 1
ic = RESP["invoice_check_ituk"]
checks_txt = [
    f"Our list change matches what ITUK was invoiced: April 2026 list {ic['April 2026']['median_list_change']:.1%} vs "
    f"invoice {ic['April 2026']['median_invoice_change']:.1%}; January 2024 list {ic['January 2024']['median_list_change']:.1%} "
    f"vs invoice {ic['January 2024']['median_invoice_change']:.1%}. ITUK pays about 97.4% of list in every year.",
    f"The 2026 factory revision (PRICE CD 250912) weighs {RESP['cy2026_cd_weighted_base_price_move']['value']:.2%} on base price, "
    "matching IML's 'around 3%' of 10 November 2025. The 5.96% is the next revision (PRICE CD 20260910): two different "
    "revisions, not a conflict.",
    "List change dates read from the current list (valid from 13 March 2026) and ITUK invoices month by month: January 2024, "
    "February 2025, April 2026 (first full month).",
]
for t in checks_txt:
    put(ws, r, 1, t, BODY, align=WRAP)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 30
    r += 1
put(ws, r + 1, 1, ("Method: two-period Poisson with part-number fixed effects and family x event fixed effects, standard errors "
                   "clustered by part number; parts with 3+ dealer units before. Build: dealer_response.py, results "
                   "dealer_price_response_2026-09-29.json (02_dashboards/pricing/_build/price_5pct_proposal_2026-09-29/)."), SUB, align=WRAP)
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
ws.row_dimensions[r + 1].height = 30
for col, wd in zip("ABCDEFGH", [60, 10, 14, 14, 14, 13, 13, 13]):
    ws.column_dimensions[col].width = wd

# ---- Not in proposal ------------------------------------------------------------
ws = ws_out
put(ws, 1, 1, "Not in the proposal: in-scope lines without a factory price on both sides", TITLE)
put(ws, 2, 1, "Named, not dropped. They keep today's price until someone prices them.", SUB)
header(ws, 4, ["Why", "Lines", "Part numbers", "Billed EUR"])
grp = un.groupby("why").agg(lines=("pn", "size"), pns=("pn", "nunique"), eur=("net_sales_value", "sum")).sort_values("eur", ascending=False)
r = 5
for why, v in grp.iterrows():
    put(ws, r, 1, why)
    put(ws, r, 2, int(v.lines), BODY, "#,##0")
    put(ws, r, 3, int(v.pns), BODY, "#,##0")
    put(ws, r, 4, round(float(v.eur), 2), BODY, EUR0)
    r += 1
put(ws, r, 1, "Total", BOLD)
put(ws, r, 2, f"=SUM(B5:B{r - 1})", BOLD, "#,##0")
put(ws, r, 4, f"=SUM(D5:D{r - 1})", BOLD, EUR0)
r += 2
put(ws, r, 1, "Top 40 lines by billed euro", SEC)
r += 1
header(ws, r, ["Why", "Part number", "Description", "Customer", "Billed EUR"])
for _, v in un.sort_values("net_sales_value", ascending=False).head(40).iterrows():
    r += 1
    put(ws, r, 1, v.why)
    put(ws, r, 2, str(v.pn))
    put(ws, r, 3, str(v.billed_description)[:60])
    put(ws, r, 4, str(v.distributor_code if pd.notna(v.distributor_code) else v.gps_db_label))
    put(ws, r, 5, round(float(v.net_sales_value), 2), BODY, EUR0)
for col, wd in zip("ABCDE", [48, 14, 40, 14, 14]):
    ws.column_dimensions[col].width = wd

# ---- Checks ---------------------------------------------------------------------
ws = ws_ck
put(ws, 1, 1, "Checks - every line should read OK", TITLE)
header(ws, 3, ["Check", "Value", "Expected", "Result"])
ck = [
    ("Billed base equals the published in-scope priced base (EUR)", f"=SUM({G})", EXPECT["billed"], EUR2,
     "=IF(ABS(B4-C4)<0.05,\"OK\",\"CHECK\")"),
    ("Published recipe reproduces: quantity x factory move", f"=SUMPRODUCT(Proposal!$H${first}:$H${last},Proposal!$K${first}:$K${last}-Proposal!$J${first}:$J${last})/SUMPRODUCT(Proposal!$H${first}:$H${last},Proposal!$J${first}:$J${last})",
     EXPECT["recipe"], "0.0000%", "=IF(ABS(B5-C5)<0.000001,\"OK\",\"CHECK\")"),
    ("Weighted increase achieved equals the target", "=Summary!B12", "=Summary!B5", PCT2,
     "=IF(ABS(B6-C6)<0.0000001,\"OK\",\"CHECK\")"),
    ("No part proposed below 0%", f'=COUNTIF({Pr},"<0")', 0, "0", "=IF(B7=C7,\"OK\",\"CHECK\")"),
    ("Market columns add up to the base", "=SUM(Summary!B38:B42)", f"=SUM({G})", EUR2, "=IF(ABS(B8-C8)<0.05,\"OK\",\"CHECK\")"),
    ("Every part has a rule", f"=COUNTA({Tr})", f"=COUNTA(Proposal!$A${first}:$A${last})", "0", "=IF(B9=C9,\"OK\",\"CHECK\")"),
    ("By-rule table adds up to the extra euro", f"=Summary!D{rt}", "=Summary!B13", EUR2, "=IF(ABS(B10-C10)<0.05,\"OK\",\"CHECK\")"),
    ("In-scope lines in the base (count written by the builder)", EXPECT["lines"], 47617, "#,##0", "=IF(B11=C11,\"OK\",\"CHECK\")"),
    ("Brake pad kits at 0% while the switch is on (sum of their proposed moves)",
     f"=SUMPRODUCT({ADr},{Pr})*Summary!B9", 0, "0.0000%", "=IF(ABS(B12-C12)<0.0000001,\"OK\",\"CHECK\")"),
]
for i, (lab, v, e, nf, res_) in enumerate(ck):
    r = 4 + i
    put(ws, r, 1, lab, BODY, align=WRAP)
    put(ws, r, 2, v, BODY, nf)
    put(ws, r, 3, e, BODY, nf)
    put(ws, r, 4, res_, BOLD)
put(ws, 14, 1, "Parts without a current PR00 list price get a proposed % but no new list price: "
               f"{int(P.list.isna().sum())} part numbers, EUR {P.billed[P.list.isna()].sum():,.0f} billed.", SUB)
for col, wd in zip("ABCD", [62, 18, 18, 10]):
    ws.column_dimensions[col].width = wd

# ---- Read me ------------------------------------------------------------------
ws = ws_read
lines = [
    ("Price proposal +5%: best allocation by part number", TITLE),
    ("29 September 2026 - a proposal on paper. Nothing published or mailed; no price list, price file or pricing configuration changed.", SUB),
    ("", BODY),
    ("What it is", SEC),
    (f"A new list price for each of {len(P):,} part numbers that together average exactly +5.00% on what the distributors "
     f"bought from July 2025 to June 2026 (EUR {P.billed.sum():,.0f}). That is EUR {exp['step_eur']:,.0f} a year on the same volume.", BODY),
    ("", BODY),
    ("How each price is set", SEC),
    ("1. Each part follows its own factory move. The new factory base prices (PRICE CD 20260910) move each part by a different "
     "amount: 61% of the euro sits at +6.5% to +8%, some parts do not move, a few double. Each part gets its own move, scaled "
     f"down to {exp['k']:.0%} of it so the book lands on +5%. This keeps every part close to the margin its class sets, instead of "
     "a flat +5% that squeezes parts whose cost rose most and fattens parts whose cost did not move.", BODY),
    ("2. Held at 0% where the distributor already loses the category. Where a distributor sells under 45% of what its trucks "
     "need in that category (measured, 2025), that distributor's share of the part is held. Mostly IUK filters.", BODY),
    ("3. A little extra where an independent UK shop sells the same part number above our reference dealer price "
     "(5 parts; capped at the evidence limits already in the pricing configuration).", BODY),
    (f"4. Brake pad kits kept at today's price (your call, 30 September 2026): {len(pads)} kits, EUR {pads.billed.sum():,.0f} "
     f"billed. The factory raised {int((pads.c > 0.15).sum())} of them by more than 15%, so at an unchanged price their "
     f"class margin falls to {pad_margin_after[pads.c > 0.15].min():.0%}-{pad_margin_after[pads.c > 0.15].max():.0%}; "
     "the other parts carry the difference so the book still lands on +5%. Switch on the Summary (B9).", BODY),
    ("", BODY),
    ("What the dealer data says", SEC),
    ("On four markets, parts that got a bigger share of the April 2026 increase did not lose measurably more dealer volume. "
     "The measured response is small (about -0.16) and at the level of noise, so it cannot pick one part over another. "
     "It is used only for the volume scenarios. Detail on the 'Dealer response' sheet.", BODY),
    ("", BODY),
    ("What is deliberately not used", SEC),
    ("Eucon competitor prices: it is not clear whether they are trade or retail. That decides whether they compare with our list "
     "or with our reference dealer price (our list x 1.40 distributor margin x 1.275 dealer margin), a 1.785x difference. "
     "Category medians: they compare a part with a whole category. 'Far above the UK shop price' is a switch, off by default, "
     "because those parts did not react more in the dealer data.", BODY),
    ("", BODY),
    ("How to use it", SEC),
    ("Summary: change the yellow cells (target, four switches, the dealer-response scenarios); everything recalculates. "
     "Proposal: one row per part number with the reason and the evidence. Review lists: top parts, the pad kits, held parts, "
     "the extra-room parts and every factory move above 15%. Checks: every line should read OK.", BODY),
    ("", BODY),
    ("Not covered", SEC),
    ("Lines without a factory price on both sides (EUR 2.4M, see 'Not in proposal'). The CIS pricing basis question stays open; "
     "CIS moves by the same % on its current basis. Exchange rate and purchase discount are held where they are.", BODY),
]
for i, (t, f_) in enumerate(lines, start=1):
    c = put(ws, i, 1, t, f_, align=WRAP)
    if f_ is BODY:
        ws.row_dimensions[i].height = 6 if not t else 13.5 * -(-len(t) // 140) + 4
ws.column_dimensions["A"].width = 120

# ---- layout pass: freeze panes everywhere, readable column B, wrap long text in B and D
for ws in wb.worksheets:
    if ws.freeze_panes is None:
        ws.freeze_panes = "A2"
    if ws.title != "Read me" and (ws.column_dimensions["B"].width or 0) < 28:
        ws.column_dimensions["B"].width = 28
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), min_col=2, max_col=4):
        for c in row:
            if c.column in (2, 4) and isinstance(c.value, str) and len(c.value) > 40:
                a = c.alignment
                c.alignment = Alignment(wrap_text=True, vertical=a.vertical or "top", horizontal=a.horizontal)

OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT_XLSX)
exp["xlsx"] = str(OUT_XLSX)
exp["rules_default"] = {"hold_lost": 1, "room": 1, "hold_far": 0, "pad_kits_kept": 1}
exp["top30"] = P.head(30).index.tolist()
OUT_JSON.write_text(json.dumps(exp, indent=2, default=float), encoding="utf-8")
(OUT_XLSX.parent / "OPEN_Price_Increase_5pct_Proposal_2026-09-29.cmd").write_text(
    '@echo off\r\nstart "" "%~dp0Price_Increase_5pct_Proposal_2026-09-29.xlsx"\r\n', encoding="ascii")
print("written", OUT_XLSX, "rows", N)
