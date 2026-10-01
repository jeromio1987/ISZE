"""Verify the recalculated proposal workbook against the builder's own Python figures."""
import json
import re
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).parent
exp = json.loads((HERE / "price_5pct_proposal_summary_2026-09-29.json").read_text(encoding="utf-8"))
wb = load_workbook(exp["xlsx"], data_only=True)
fails = []


def check(label, got, want, tol):
    ok = got is not None and abs(float(got) - float(want)) <= tol
    print(f"{'OK  ' if ok else 'FAIL'} {label}: excel={got} python={want}")
    if not ok:
        fails.append(label)


s = wb["Summary"]
check("pad-kit switch on", s["B9"].value, 1, 0)
check("scaling k", s["B10"].value, exp["k"], 1e-9)
check("achieved", s["B12"].value, 0.05, 1e-9)
check("extra EUR", s["B13"].value, exp["step_eur"], 0.01)
check("parts", s["B15"].value, exp["parts"], 0)
check("parts at 0%", s["B16"].value, exp["parts_at_zero"], 0)
check("EUR at 0%", s["B17"].value, exp["eur_at_zero"], 0.05)
check("median move", s["B18"].value, exp["median_move"], 1e-9)
check("max move", s["B20"].value, exp["max_move"], 1e-9)
check("margin pts proposal", s["B21"].value, exp["margin_pts_proposal"], 1e-9)
check("margin pts flat", s["F25"].value, exp["margin_pts_flat"], 1e-9)
check("margin pts factory-only", s["F26"].value, exp["margin_pts_factory_only"], 1e-9)
check("k factory-only", s["B26"].value, exp["k_factory_only"], 1e-9)
check("lose >1pt flat", s["G25"].value, exp["eur_lose_1pt_flat"], 0.05)
check("lose >1pt factory-only", s["G26"].value, exp["eur_lose_1pt_factory_only"], 0.05)
check("lose >1pt proposal", s["G27"].value, exp["eur_lose_1pt_proposal"], 0.05)
for r, key in [(32, "revenue_none"), (33, "revenue_central"), (34, "revenue_stress")]:
    check(f"revenue {key}", s[f"D{r}"].value, exp[key], 0.05)
for i, (code, v) in enumerate(exp["by_market"].items()):
    check(f"market {code}", s[f"C{38 + i}"].value, v, 1e-9)

ck = wb["Checks"]
for r in range(4, 13):
    v = ck[f"D{r}"].value
    print(f"{'OK  ' if v == 'OK' else 'FAIL'} Checks row {r}: {ck[f'A{r}'].value} -> {v}")
    if v != "OK":
        fails.append(f"Checks row {r}")

errs = 0
banned = re.compile(r"\b(wallet|street|sell-?out|sellout|book b|mode a|mode b|class factor)\b", re.I)
hits = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str):
                if v.startswith("#") and v.rstrip("!?/0").upper() in ("#VALUE", "#NAME", "#DIV", "#REF", "#N/A", "#NUM", "#NULL"):
                    errs += 1
                if banned.search(v):
                    hits.append((ws.title, c.coordinate, v[:80]))
            if v in ("#VALUE!", "#NAME?", "#DIV/0!", "#REF!", "#N/A", "#NUM!", "#NULL!"):
                errs += 1
print(("OK  " if errs == 0 else "FAIL") + f" formula error cells: {errs}")
print(("OK  " if not hits else "FAIL") + f" banned-label hits: {hits[:5]}")
if errs:
    fails.append("formula errors")
if hits:
    fails.append("banned labels")

p = wb["Proposal"]
vals = np.array([row[15].value for row in p.iter_rows(min_row=2)], dtype=float)
print(("OK  " if np.isfinite(vals).all() and (vals >= 0).all() else "FAIL") + f" proposal column numeric and >= 0 ({len(vals)} rows)")
pad_rows = [(row[0].value, row[15].value, row[16].value, row[8].value, row[19].value)
            for row in p.iter_rows(min_row=2) if row[29].value == 1]
pad_ok = (len(pad_rows) == exp["pad_kits"] and sorted(r[0] for r in pad_rows) == sorted(exp["pad_kits_pns"])
          and all(abs(r[1]) < 1e-12 and r[4] == "Held: pad kit, price kept" for r in pad_rows)
          and all(r[3] is None or abs(r[2] - r[3]) < 0.005 for r in pad_rows))
print(("OK  " if pad_ok else "FAIL") + f" pad kits: {len(pad_rows)} rows, proposed move 0, new list = current list, rule text set")
if not pad_ok:
    fails.append("pad kits")
rev = wb["Review lists"]
na = sum(1 for row in rev.iter_rows(min_row=4) for c in row[3:8] if c.value in ("#N/A",))
print(("OK  " if na == 0 else "FAIL") + f" review-list lookups resolved (#N/A: {na})")
print("\nRESULT:", "ALL OK" if not fails else f"FAILED {fails}")
sys.exit(1 if fails else 0)
