"""Simulate HTML rollup for Not held — find why all_billed collapses to ~1329."""
import json
from pathlib import Path

d = json.loads(
    Path(r"C:\Users\jerom\Desktop\claude\ISZE\02_dashboards\pricing\_build\pricing_dashboard_data.json")
    .read_text(encoding="utf-8")
)
nhi = d["pool"]["src"].index("Not held")
gs = [r for r in d["groups"] if r[4] == nhi]
print("nh groups", len(gs))
print("sum all", sum(r[12] for r in gs))
print("sum billed", sum(r[8] for r in gs))

# simulate totalsNow: rollup by db index 0
from collections import defaultdict
m = defaultdict(lambda: {"all": 0, "billed": 0, "step": 0, "prior": 0, "new": 0})
for r in gs:
    o = m[r[0]]
    o["billed"] += r[8]
    o["prior"] += r[9]
    o["new"] += r[10]
    o["step"] += r[11]
    o["all"] += r[12]
t_all = sum(v["all"] for v in m.values())
t_billed = sum(v["billed"] for v in m.values())
t_step = sum(v["step"] for v in m.values())
print("rollup all", round(t_all, 2), "billed", round(t_billed, 2), "step", round(t_step, 2))
print("dbs in nh", len(m))

# Checks sheet cell map
import openpyxl
wb = openpyxl.load_workbook(
    r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\Pricing_Revision_AIM_Source_Analysis.xlsx",
    data_only=False,
)
ws = wb["Checks"]
for r in range(1, 30):
    a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
    if a or b:
        print(f"Checks R{r}: {a!r:.80} | {b!r:.60}" if a else f"Checks R{r}: None | {b!r:.60}")
