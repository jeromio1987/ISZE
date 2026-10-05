# -*- coding: utf-8 -*-
"""Static scan of the delivered workbook: formula count and language check."""
import io
import re
import sys

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

P = r"02_dashboards\pricing\ISZE_Pricing_Simulation_Spec_Basis.xlsx"
wb = openpyxl.load_workbook(P, data_only=False)

print("=== sheets and formula counts")
total_formula = 0
total_cells = 0
for ws in wb.worksheets:
    f = 0
    c = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            c += 1
            if isinstance(cell.value, str) and cell.value.startswith("="):
                f += 1
    total_formula += f
    total_cells += c
    print(f"  {ws.title:<24} rows={ws.max_row:>7,} cols={ws.max_column:>3} "
          f"formulas={f:>8,} non-empty={c:>8,}")
print(f"  {'TOTAL':<24} {'':>14} formulas={total_formula:>8,} non-empty={total_cells:>8,}")

print()
print("=== the input dials (shaded cells a reader may change)")
ws = wb["Legend"]
for rr in range(1, 40):
    lab = ws.cell(rr, 1).value
    v = ws.cell(rr, 2).value
    if lab is None or v is None:
        continue
    filled = ws.cell(rr, 2).fill
    is_dial = filled is not None and filled.fgColor is not None and \
        filled.fgColor.rgb not in (None, "00000000")
    if is_dial:
        print(f"  Legend!B{rr}  {lab} = {v}")

print()
print("=== banned / risky wording scan of every visible cell")
BANNED = ["wallet", "street", "sell-out", "sellout", "sell out", "book b",
          "ship-gate", "shipgate", "harness", "spar", "canonical",
          "iml v7", "model b", "prism", "yaml", ".md", ".py", ".json", ".csv",
          "00_docs", "01_data", "02_dashboards", "03_analytics", "c:\\",
          "pms", "gps", "ids", "aim", "bvp", "spu", "uio", "spa"]
hits = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or v.startswith("="):
                continue
            # Column B of PN data is the raw part description carried from the
            # billing extract. It is quoted material, not our copy.
            if ws.title == "PN data" and cell.column == 2:
                continue
            low = v.lower()
            for token in BANNED:
                if token in low:
                    hits.append((ws.title, cell.coordinate, token, v[:160]))
if hits:
    print(f"  {len(hits)} hits")
    for h in hits:
        print(f"    {h[0]}!{h[1]}  [{h[2]}]  {h[3]}")
else:
    print("  none")

print()
print("=== sheet names")
print(" ", wb.sheetnames)
