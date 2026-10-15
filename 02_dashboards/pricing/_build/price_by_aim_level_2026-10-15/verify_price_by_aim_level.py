"""Check the AIM-level price workbook in Excel itself, against an independent Python calculation.

1. Recalculates the delivered workbook in Excel and compares every part's move on every list, the totals and the
   Checks sheet with the builder's own model (shipped settings: target 5.30%, pad kits at 0%).
2. Neutral copy: every setting cleared, target 5.00% -> each part's move must equal the 'Factory-led only move' of
   the price revision workbook of October 2026.
3. Scenario copy: settings on L1, L2, L3 and a part, one target per list, mixed grammar (F+2, S, 8%, F-1, 0%) ->
   every part's move must equal the Python model.

    python verify_price_by_aim_level.py            (all three)
    python verify_price_by_aim_level.py --quick    (only 1)
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import build_price_by_aim_level as B  # noqa: E402  (loads the data and the model; does not rebuild)

RECALC = B.PROP / "recalc_excel.py"
EXP = json.loads((HERE / "expected_price_by_aim_level.json").read_text(encoding="utf-8"))
TOL = 1e-9
fails = []


def recalc(path):
    r = subprocess.run([sys.executable, str(RECALC), str(path)], capture_output=True, text=True)
    print("   ", r.stdout.strip().replace("\n", " | "))
    if r.returncode:
        raise SystemExit(r.stderr)


def read_parts(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Parts"]
    rows = ws.iter_rows(min_row=B.HDR_ROW, values_only=True)
    hdr = list(next(rows))
    data = [r for r in rows if r[0] is not None]
    df = pd.DataFrame(data, columns=hdr).set_index("Part number")
    st = {r[0]: r[1:6] for r in wb["Settings"].iter_rows(min_row=1, max_row=60, max_col=6, values_only=True) if r[0]}
    ck = [r for r in wb["Checks"].iter_rows(min_row=4, max_row=12, max_col=4, values_only=True) if r[0]]
    wb.close()
    return df, st, ck


def mv_col(k):
    return {"PR00": "PR00: MOVE", "UTI": "UTI: MOVE", "UMI": "UMI: MOVE",
            "OTH": "Other lists: MOVE (all-lists settings only)"}[k]


def compare(df, M, label):
    worst = 0.0
    for k in B.LISTS + ["OTH"]:
        got = pd.to_numeric(df[mv_col(k)], errors="coerce").reindex(B.P.index)
        d = float((got - M[f"mv_{k}"]).abs().max())
        worst = max(worst, d)
        if not np.isfinite(d) or d > TOL:
            fails.append(f"{label}: {k} moves differ by up to {d}")
    print(f"   {label}: largest difference between Excel and Python on {len(B.P):,} parts x 4 lists = {worst:.1e}")


def set_inputs(src, dst, target, mode=1, list_targets=None, l1=None, l2=None, l3=None, part=None, clear=True):
    shutil.copy(src, dst)
    wb = load_workbook(dst)
    s = wb["Settings"]
    s["B5"], s["B6"] = target, mode
    for k, cell in (("PR00", "B7"), ("UTI", "B8"), ("UMI", "B9"), ("OTH", "B10")):
        s[cell] = (list_targets or {}).get(k, target)
    for sheet, nlev, sets in (("By L1", 1, l1 or {}), ("By L2", 2, l2 or {}), ("By L3", 3, l3 or {})):
        ws = wb[sheet]
        vc = B.vcols(nlev)
        cols = {"ALL": vc["s_ALL"], "PR00": vc["s_PR00"], "UTI": vc["s_UTI"], "UMI": vc["s_UMI"]}
        for r in range(B.FIRST, ws.max_row + 1):
            key = tuple(ws.cell(r, j + 1).value for j in range(nlev))
            if key[0] is None or key[0] == "Total":
                continue
            if clear:
                for c in cols.values():
                    ws.cell(r, c).value = None
            k1 = key[0] if nlev == 1 else key
            for lst, v in sets.get(k1, {}).items():
                ws.cell(r, cols[lst]).value = v
    ws = wb["Parts"]
    hdr = [c.value for c in ws[B.HDR_ROW]]
    pcols = {k: hdr.index(f"SETTING on this part: {n}") + 1 for k, n in
             (("ALL", "all lists"), ("PR00", "PR00"), ("UTI", "UTI 1799"), ("UMI", "UMI 1797"))}
    for r in range(B.FIRST, ws.max_row + 1):
        pn = ws.cell(r, 1).value
        if clear:
            for c in pcols.values():
                ws.cell(r, c).value = None
        for lst, v in (part or {}).get(pn, {}).items():
            ws.cell(r, pcols[lst]).value = v
    wb.save(dst)


# ------------------------------------------------------------------ 1. delivered workbook
print("1. delivered workbook, shipped settings")
recalc(B.OUT)
df, st, ck = read_parts(B.OUT)
M0, O0 = B.model()
compare(df, M0, "shipped")
book_extra = st["Extra EUR a year on the same volume"][4]
print(f"   whole book: extra EUR {book_extra:,.2f} (Python {sum(O0['extra'].values()):,.2f}), "
      f"achieved {st['Move achieved, weighted'][4]:.6%}, status {st['Status'][4]}")
if abs(book_extra - sum(O0["extra"].values())) > 0.01 or st["Status"][4] != "OK":
    fails.append("shipped totals")
for i, k in enumerate(B.LISTS + ["OTH"]):
    if abs(st["Extra EUR a year on the same volume"][i] - O0["extra"][k]) > 0.01:
        fails.append(f"extra {k}")
    if abs(st["Margin EUR change (extra EUR less cost increase, parts with a known cost)"][i] - O0["mchg"][k]) > 0.01:
        fails.append(f"margin {k}")
for r in ck:
    print(f"   check: {r[3]:5} {r[0][:90]}")
    if r[3] != "OK":
        fails.append(f"check: {r[0]}")
res = {k: st[k] for k in ("Billed last 12 months, EUR", "Own factory move, weighted", "Move achieved, weighted",
                          "Extra EUR a year on the same volume", "Margin today (on billed, parts with a known cost)",
                          "Margin after", "Margin change, points",
                          "Margin EUR change (extra EUR less cost increase, parts with a known cost)",
                          "Parts billed on this list with no current price on it (they get a move, no new price)",
                          "Billed EUR on those parts")}
(HERE / "verify_result_shipped.json").write_text(json.dumps(res, default=str, indent=1), encoding="utf-8")

if "--quick" not in sys.argv:
    tmp = Path(tempfile.gettempdir()) / "isze_aim_level_verify"
    tmp.mkdir(exist_ok=True)
    # -------------------------------------------------------------- 2. neutral
    print("2. neutral copy: no settings, target 5.00%")
    dst = tmp / "neutral.xlsx"
    set_inputs(B.OUT, dst, 0.05)
    recalc(dst)
    dfn, stn, _ = read_parts(dst)
    Mn, _ = B.model(target=0.05, set_l3={}, set_part={})
    compare(dfn, Mn, "neutral")
    ref = pd.Series(EXP["neutral_5pct_moves"]).reindex(B.P.index)
    got = pd.to_numeric(dfn[mv_col("PR00")], errors="coerce").reindex(B.P.index)
    d = float((got - ref).abs().max())
    print(f"   vs 'Factory-led only move' of the price revision workbook: largest difference {d:.1e}")
    if d > TOL:
        fails.append("neutral vs revision workbook")
    # -------------------------------------------------------------- 3. scenario
    print("3. scenario copy: settings on every level, one target per list")
    g2 = B.G2[B.G2.l1 == "Brakes"].iloc[0]
    g3 = B.G3[B.G3.l3 == "PM Sensor"].iloc[0]
    l1 = {"Brakes": {"ALL": "F+2"}, "Intake Exhaust Emission": {"UMI": 0.0}}
    l2 = {(g2.l1, g2.l2): {"PR00": "S"}}
    l3 = {(r.l1, r.l2, r.l3): {"ALL": 0.0} for r in B.G3.itertuples() if r.l3 == "Brake Pad Kit"}
    l3[(g3.l1, g3.l2, g3.l3)] = {"UTI": 0.08}
    part = {pn: {"ALL": 0.0} for pn in B.PAD_KITS_MISLABELLED}
    part["8983767550"] = {"PR00": "F-1"}
    lt = {"PR00": 0.055, "UTI": 0.04, "UMI": 0.05, "OTH": 0.053}
    dst = tmp / "scenario.xlsx"
    set_inputs(B.OUT, dst, 0.053, mode=2, list_targets=lt, l1=l1, l2=l2, l3=l3, part=part)
    recalc(dst)
    dfs, sts, cks = read_parts(dst)
    Ms, Os = B.model(target=0.053, mode=2, set_l1=l1, set_l2=l2, set_l3=l3, set_part=part, list_targets=lt)
    compare(dfs, Ms, "scenario")
    for i, k in enumerate(B.LISTS + ["OTH"]):
        a = sts["Move achieved, weighted"][i]
        print(f"   {k}: target {lt[k]:.2%}, achieved {a:.4%}, status {sts['Status'][i]}")
        if abs(a - lt[k]) > 5e-7:
            fails.append(f"scenario target {k}")
    rotor = dfs.loc["8983767550"]
    print(f"   rotor 8983767550 on PR00: setting from '{rotor['PR00: setting comes from']}', "
          f"used '{rotor['PR00: setting used']}', move {rotor['PR00: MOVE']:.4%} (own factory move {rotor['Own factory move (IML/IMIT base price)']:.4%})")

print("\nRESULT:", "ALL OK" if not fails else "FAIL")
for f in fails:
    print("  -", f)
sys.exit(1 if fails else 0)
