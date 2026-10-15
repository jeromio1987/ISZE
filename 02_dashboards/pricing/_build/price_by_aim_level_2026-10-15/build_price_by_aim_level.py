"""Price revision by AIM level: one live workbook to set a price move on any AIM level (L1, L2, L3) or on a
single part, separately for the three price lists that matter: PR00 (EU and UK distributors), UTI Israel
trucks (customer 1799) and UMI Israel D-Max (customer 1797).

Base: the same in-scope revision base as the +5% proposal (47,617 lines, EUR 32,595,852.63, Jul 2025 - Jun 2026).
Parts without a setting follow their own factory move, scaled so the book (or each list) lands on the target.
Nothing published; no price file, price list or pricing configuration touched.

    python build_price_by_aim_level.py
"""
import json
import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
ISZE = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
PROP = ISZE / "02_dashboards" / "pricing" / "_build" / "price_5pct_proposal_2026-09-29"
sys.path.insert(0, str(PROP))
from evidence_lib import load_base  # noqa: E402

DER = ISZE / "01_data" / "00_derived"
OUT = Path(os.environ.get("AIM_LEVEL_OUT", ISZE / "01_data" / "01_Pricing" /
                          "ISZE_Price_Revision_by_AIM_Level_2026-10-15.xlsx"))
EXP_JSON = HERE / "expected_price_by_aim_level.json"
REL = DER / "_releases" / "aim_pn_classification_release_2026-10-11.csv"
MASUMI_WB = (ISZE / "01_data" / "01_Pricing" / "_PACK_AIM_Pricing_7d_20260914_20260920" /
             "08_Draft_to_MH_2026-10-11" / "ISZE_Parts_Price_Revision_Workbook_Oct2026.xlsx")
CFG = yaml.safe_load((ISZE / "03_analytics" / "isze_pricing" / "config.yaml").read_text(encoding="utf-8"))

EXPECT = {"lines": 47617, "billed": 32595852.63, "parts": 15273}
TARGET, MODE, CEIL, FLOOR = 0.053, 1, 0.05, 0.0
FX = float(CFG["formula"]["fx_rate"])
DISC = {k: float(v) for k, v in CFG["formula"]["source_discount"].items()}
LISTS = ["PR00", "UTI", "UMI"]
LIST_NAME = {"PR00": "PR00: EU and UK distributors", "UTI": "UTI Israel trucks (1799)",
             "UMI": "UMI Israel D-Max (1797)", "OTH": "Other own lists (2372, 2733, 2734)"}
OWN_LIST = {"1799": "UTI", "1797": "UMI", "2372": "OTH", "2733": "OTH", "2734": "OTH"}
# decided 30 Sep 2026: brake pad kits stay at today's price in every market. The AIM release of 11 Oct labels
# five of the 29 kits "Brake Disc", so they are held on the part itself as well as on L3 "Brake Pad Kit".
PAD_KITS_MISLABELLED = ["8978279610", "5876150640", "8978279590", "5876150650", "8978279601"]

# ------------------------------------------------------------------ data
L = load_base()
L = L[L.movable].copy()
assert len(L) == EXPECT["lines"], len(L)
assert abs(L.net_sales_value.sum() - EXPECT["billed"]) < 0.01
cust = L.customer_code.astype(str).str.strip().str.lstrip("0")
L["lst"] = cust.map(OWN_LIST).fillna("PR00")
L["c"] = L.bp_new_jpy / L.bp_prior_revised - 1

rel = pd.read_csv(REL, dtype=str, usecols=["part_number", "aim_l1_fixed", "aim_l2_fixed", "aim_l3_fixed"])
rel["part_number"] = rel.part_number.str.strip()
rel = rel.drop_duplicates("part_number").set_index("part_number")
enq = pd.read_parquet(DER / "pricing_pn_enquiry.parquet", columns=["part_number", "source", "sp_eur_pr00_ie40"])
enq["pn"] = enq.part_number.astype(str).str.strip()
enq = enq.drop_duplicates("pn").set_index("pn")
kw = pd.read_parquet(DER / "pricing_zpr1_key_wide.parquet", columns=["part_number", "zpr1_sp_UTI", "zpr1_sp_UMI"])
kw["part_number"] = kw.part_number.astype(str).str.strip()
kw = kw.drop_duplicates("part_number").set_index("part_number")

g = L.groupby("pn")
P = pd.DataFrame({"c": g.c.first(), "bp0": g.bp_prior_revised.first(), "bp1": g.bp_new_jpy.first(),
                  "cls": g.bp_new_class.first()})
for k in LISTS + ["OTH"]:
    d = L[L.lst == k].groupby("pn")
    P[f"b_{k}"] = d.net_sales_value.sum().reindex(P.index).fillna(0.0)
    P[f"q_{k}"] = d.billed_qty.sum().reindex(P.index).fillna(0.0)
P["billed"] = P[[f"b_{k}" for k in LISTS + ["OTH"]]].sum(axis=1)
P["desc"] = (L.sort_values("net_sales_value", ascending=False).drop_duplicates("pn").set_index("pn")
             .billed_description.reindex(P.index).fillna("").astype(str).str.slice(0, 60))
for lv in ("l1", "l2", "l3"):
    P[lv] = rel[f"aim_{lv}_fixed"].reindex(P.index).fillna("Not classified").astype(str)
P["src"] = enq.source.reindex(P.index)
P.loc[~P.src.isin(["J", "T"]), "src"] = ""
P["p_PR00"] = pd.to_numeric(enq.sp_eur_pr00_ie40.reindex(P.index), errors="coerce")
P["p_UTI"] = pd.to_numeric(kw.zpr1_sp_UTI.reindex(P.index), errors="coerce")
P["p_UMI"] = pd.to_numeric(kw.zpr1_sp_UMI.reindex(P.index), errors="coerce")
for k in LISTS:
    P.loc[~(P[f"p_{k}"] > 0), f"p_{k}"] = np.nan
unit = P.b_PR00 / P.q_PR00.where(P.q_PR00 > 0)
P["chk_pr00"] = unit / P.p_PR00
P = P.sort_values("billed", ascending=False)
P.index.name = "pn"
assert len(P) == EXPECT["parts"], len(P)
P["cp"] = P.c.clip(lower=0)
P["fob"] = np.where(P.src == "J", P.bp0 / FX * DISC["J"], np.where(P.src == "T", P.bp0 / FX * DISC["T"], np.nan))

# group tables, biggest billed first
def group_table(keys):
    t = P.groupby(keys).agg(parts=("billed", "size"), **{f"b_{k}": (f"b_{k}", "sum") for k in LISTS + ["OTH"]},
                            billed=("billed", "sum")).reset_index()
    wc = (P.billed * P.c).groupby([P[k] for k in keys]).sum().reset_index(name="wc")
    t = t.merge(wc, on=keys)
    t["fm"] = np.where(t.billed > 0, t.wc / t.billed, np.nan)
    return t.sort_values("billed", ascending=False).reset_index(drop=True)


G1, G2, G3 = group_table(["l1"]), group_table(["l1", "l2"]), group_table(["l1", "l2", "l3"])
HDR_ROW, FIRST = 5, 6
for t in (G1, G2, G3):
    t["row"] = np.arange(len(t)) + FIRST
P["r1"] = P.l1.map(dict(zip(G1.l1, G1.row)))
P["r2"] = pd.MultiIndex.from_frame(P[["l1", "l2"]]).map(dict(zip(zip(G2.l1, G2.l2), G2.row)))
P["r3"] = pd.MultiIndex.from_frame(P[["l1", "l2", "l3"]]).map(dict(zip(zip(G3.l1, G3.l2, G3.l3), G3.row)))
assert P[["r1", "r2", "r3"]].notna().all().all()

# shipped settings (blank everywhere else)
SET_L3 = {(r.l1, r.l2, r.l3): {"ALL": 0.0} for r in G3.itertuples() if r.l3 == "Brake Pad Kit"}
SET_PART = {pn: {"ALL": 0.0} for pn in PAD_KITS_MISLABELLED}
assert len(SET_L3) == 1 and all(pn in P.index for pn in PAD_KITS_MISLABELLED)

# ------------------------------------------------------------------ python model (the verifier compares Excel with it)
def parse(v, cp):
    """Return (scaled?, move). Same grammar as the workbook."""
    if v is None or (isinstance(v, str) and v.strip().upper() in ("", "S")):
        return True, None
    if isinstance(v, (int, float)):
        return False, float(v)
    s = v.strip().upper()
    if s.startswith("F"):
        rest = s[1:].replace(",", ".").replace("%", "").strip()
        pts = float(rest) if rest else 0.0
        return False, max(FLOOR, cp + pts / 100)
    return True, None


def resolve(pn, row, k, set_l1=None, set_l2=None, set_l3=None, set_part=None):
    set_l1, set_l2, set_l3 = set_l1 or {}, set_l2 or {}, set_l3 or {}
    set_part = SET_PART if set_part is None else set_part
    slots = [(set_part.get(pn, {}), "Part"), (set_l3.get((row.l1, row.l2, row.l3), {}), "L3"),
             (set_l2.get((row.l1, row.l2), {}), "L2"), (set_l1.get(row.l1, {}), "L1")]
    for d, _ in slots:
        for key in (k, "ALL"):
            if key in d and d[key] is not None and d[key] != "":
                return d[key]
    return None


def model(target=TARGET, mode=MODE, set_l3=None, set_part=None, set_l1=None, set_l2=None, list_targets=None):
    set_l3 = SET_L3 if set_l3 is None else set_l3
    M = pd.DataFrame(index=P.index)
    for k in LISTS + ["OTH"]:
        sc, mv = zip(*[parse(resolve(pn, r, k, set_l1=set_l1, set_l2=set_l2, set_l3=set_l3,
                                             set_part=set_part), r.cp)
                       for pn, r in P.iterrows()])
        M[f"sc_{k}"] = np.array(sc)
        M[f"sm_{k}"] = np.array([np.nan if m is None else m for m in mv], dtype=float)
    base = {k: P[f"b_{k}"].sum() for k in LISTS + ["OTH"]}
    F = {k: (P[f"b_{k}"] * M[f"sm_{k}"].fillna(0)).sum() for k in LISTS + ["OTH"]}
    B1 = {k: (P[f"b_{k}"] * P.cp * M[f"sc_{k}"]).sum() for k in LISTS + ["OTH"]}
    tot = sum(base.values())
    if mode == 1:
        s1 = (target * tot - sum(F.values())) / sum(B1.values())
        s1 = {k: s1 for k in base}
    else:
        lt = list_targets or {}
        s1 = {k: (lt.get(k, target) * base[k] - F[k]) / B1[k] for k in base}
    cap, capeur, unc = {}, {}, {}
    for k in base:
        cap[k] = M[f"sc_{k}"] & (P.cp * s1[k] > P.cp + CEIL)
        capeur[k] = (P[f"b_{k}"] * cap[k] * (P.cp + CEIL)).sum()
        unc[k] = B1[k] - (P[f"b_{k}"] * cap[k] * P.cp).sum()
    if mode == 1:
        s2 = (target * tot - sum(F.values()) - sum(capeur.values())) / sum(unc.values())
        s2 = {k: s2 for k in base}
    else:
        s2 = {k: (lt.get(k, target) * base[k] - F[k] - capeur[k]) / unc[k] for k in base}
    out = {"s1": s1, "s2": s2, "base": base}
    for k in base:
        scaled = np.maximum(FLOOR, np.minimum(s2[k] * P.cp, P.cp + CEIL))
        M[f"mv_{k}"] = np.where(M[f"sc_{k}"], scaled, M[f"sm_{k}"])
        M[f"ex_{k}"] = P[f"b_{k}"] * M[f"mv_{k}"]
        known = P.fob.notna()
        M[f"mc_{k}"] = np.where(known, P[f"b_{k}"] * M[f"mv_{k}"] - P[f"q_{k}"] * P.fob * P.c, np.nan)
    out["extra"] = {k: float(M[f"ex_{k}"].sum()) for k in base}
    out["mchg"] = {k: float(np.nansum(M[f"mc_{k}"])) for k in base}
    out["achieved"] = sum(out["extra"].values()) / tot
    return M, out


# view sheet layout: name columns, then fixed blocks
def vcols(nlev):
    b = nlev
    return {"parts": b + 1, "b_PR00": b + 2, "b_UTI": b + 3, "b_UMI": b + 4, "b_OTH": b + 5, "fm": b + 6,
            "s_ALL": b + 7, "s_PR00": b + 8, "s_UTI": b + 9, "s_UMI": b + 10,
            "m_PR00": b + 11, "m_UTI": b + 12, "m_UMI": b + 13, "m_OTH": b + 14,
            "extra": b + 15, "mchg": b + 16, "m0": b + 17, "m1": b + 18, "chk": b + 19}


def main():
    M0, OUT0 = model()
    print("shipped settings:", json.dumps({k: OUT0[k] for k in ("extra", "achieved", "s2")}, default=float))

    # neutral settings at 5.00% must equal the 'Factory-led only move' of the price revision workbook, part by part
    Mn, On = model(target=0.05, set_l3={}, set_part={})
    wb_m = load_workbook(MASUMI_WB, read_only=True, data_only=True)
    ref = {str(r[0]).strip(): r[27] for r in wb_m["Proposal"].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    wb_m.close()
    cmp = pd.Series(ref).reindex(P.index).astype(float)
    NEUTRAL_MAXDIFF = float((Mn.mv_PR00 - cmp).abs().max())
    assert cmp.notna().all() and NEUTRAL_MAXDIFF < 1e-9, NEUTRAL_MAXDIFF
    print(f"neutral reconciliation vs price revision workbook: max difference {NEUTRAL_MAXDIFF:.2e} on {len(cmp)} parts")

    EXP = {"target": TARGET, "mode": MODE, "billed": float(P.billed.sum()), "parts": len(P),
           "base": {k: float(v) for k, v in OUT0["base"].items()},
           "extra": OUT0["extra"], "mchg": OUT0["mchg"], "achieved": OUT0["achieved"],
           "s1": {k: float(v) for k, v in OUT0["s1"].items()}, "s2": {k: float(v) for k, v in OUT0["s2"].items()},
           "moves": {k: dict(zip(P.index, M0[f"mv_{k}"].round(12))) for k in LISTS + ["OTH"]},
           "neutral_5pct_moves": dict(zip(P.index, Mn.mv_PR00.round(12))),
           "neutral_maxdiff_vs_revision_workbook": NEUTRAL_MAXDIFF,
           "groups": {"l1": len(G1), "l2": len(G2), "l3": len(G3)}}
    EXP_JSON.write_text(json.dumps(EXP, default=float), encoding="utf-8")

    # ------------------------------------------------------------------ workbook styles
    NAVY, PAPER, TERRA = "15233C", "F7F4EE", "C0552D"
    F = "Arial"
    HFILL = PatternFill("solid", fgColor=NAVY)
    HFONT = Font(name=F, bold=True, color=PAPER, size=10)
    GFILL = PatternFill("solid", fgColor="D9D9D9")
    GFONT = Font(name=F, bold=True, color="222222", size=9)
    TITLE = Font(name=F, bold=True, size=14, color=NAVY)
    SUB = Font(name=F, italic=True, size=9, color="555555")
    SEC = Font(name=F, bold=True, size=11, color=NAVY)
    BODY = Font(name=F, size=10, color="222222")
    BOLD = Font(name=F, size=10, bold=True, color="222222")
    INPUT = Font(name=F, size=10, color="0000FF", bold=True)
    YELLOW = PatternFill("solid", fgColor="FFF2A8")
    TFILL = PatternFill(start_color="FBE3D8", end_color="FBE3D8", fill_type="solid")
    WRAP = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    PCT1, PCT2, EUR0, EUR2 = "0.0%", "0.00%", "#,##0;-#,##0;-", "#,##0.00"
    THIN = Border(bottom=Side(style="thin", color="BBBBBB"))

    wb = Workbook()
    ws_rm = wb.active
    ws_rm.title = "Read me"
    ws_s = wb.create_sheet("Settings")
    ws_1 = wb.create_sheet("By L1")
    ws_2 = wb.create_sheet("By L2")
    ws_3 = wb.create_sheet("By L3")
    ws_p = wb.create_sheet("Parts")
    ws_c = wb.create_sheet("Checks")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False


    def put(ws, r, c, v, font=BODY, fmt=None, al=None, fill=None):
        x = ws.cell(r, c, v)
        x.font = font
        if fmt:
            x.number_format = fmt
        if al:
            x.alignment = al
        if fill:
            x.fill = fill
        return x


    def header(ws, r, labels, col=1, grey=(), height=58):
        for i, h in enumerate(labels):
            x = ws.cell(r, col + i, h)
            isg = (col + i) in grey
            x.fill, x.font, x.alignment = (GFILL if isg else HFILL), (GFONT if isg else HFONT), CENTER
        ws.row_dimensions[r].height = height


    def okf(ref):
        """True when a setting cell is empty or uses the grammar."""
        return (f'OR({ref}="",AND(ISNUMBER({ref}),ABS({ref})<=0.5),UPPER(TRIM({ref}))="S",'
                f'LEFT(UPPER(TRIM({ref})),1)="F")')


    GRAMMAR = ("0% = keep today's price.  4% = move exactly +4%.  F = the part's own factory move.  "
               "F+2 = own factory move plus 2 points (F-1 = minus 1 point).  S = follow the scaled factory move "
               "(cancels a setting from a higher level).  Empty = take the setting from the level above.")
    ORDER = ("Which setting wins: the part itself, then L3, then L2, then L1. On each level a setting for one price list "
             "beats the 'all lists' setting. Parts with no setting anywhere follow their own factory move, scaled so the "
             "total lands on the target.")


    def add_validation(ws, rng, first_ref):
        dv = DataValidation(type="custom", formula1=okf(first_ref), allow_blank=True, showErrorMessage=True,
                            errorTitle="Setting not understood",
                            error="Type 0% (keep), 4% (exact move), F (own factory move), F+2 or F-1 (factory move "
                                  "plus or minus points) or S (follow the scaled factory move). Leave empty to inherit.")
        dv.add(rng)
        ws.add_data_validation(dv)


    # ------------------------------------------------------------------ Parts sheet layout
    N = len(P)
    R0, R1 = FIRST, FIRST + N - 1
    col = {}
    LAB = {"pn": "Part number", "desc": "Description", "l1": "AIM L1", "l2": "AIM L2", "l3": "AIM L3",
           "in_ALL": "SETTING on this part: all lists", "in_PR00": "SETTING on this part: PR00",
           "in_UTI": "SETTING on this part: UTI 1799", "in_UMI": "SETTING on this part: UMI 1797",
           "mv_PR00": "PR00: MOVE", "mv_UTI": "UTI: MOVE", "mv_UMI": "UMI: MOVE",
           "mv_OTH": "Other lists: MOVE (all-lists settings only)",
           "np_PR00": "PR00: NEW PRICE EUR", "np_UTI": "UTI: NEW PRICE EUR", "np_UMI": "UMI: NEW PRICE EUR",
           "p_PR00": "Current price PR00 (IE40 list 13 Mar 2026)",
           "p_UTI": "Current price UTI 1799 (own list, 10 Aug 2026)",
           "p_UMI": "Current price UMI 1797 (own list, 10 Aug 2026)",
           "at_PR00": "PR00: setting comes from", "at_UTI": "UTI: setting comes from", "at_UMI": "UMI: setting comes from",
           "at_OTH": "Other lists: setting comes from", "su_OTH": "Other lists: setting used",
           "sm_OTH": "Other lists: move from the setting",
           "c": "Own factory move (IML/IMIT base price)", "flag": "Check", "cls": "Price class",
           "src": "Source (J Japan IML, T Thailand IMIT)",
           "b_PR00": "Billed EUR PR00", "b_UTI": "Billed EUR UTI 1799", "b_UMI": "Billed EUR UMI 1797",
           "b_OTH": "Billed EUR other own lists",
           "ex_PR00": "PR00: extra EUR a year", "ex_UTI": "UTI: extra EUR a year", "ex_UMI": "UMI: extra EUR a year",
           "ex_OTH": "Other lists: extra EUR a year",
           "mc_PR00": "PR00: margin EUR change", "mc_UTI": "UTI: margin EUR change", "mc_UMI": "UMI: margin EUR change",
           "mc_OTH": "Other lists: margin EUR change",
           "fob": "Our cost today per piece, EUR (base price / FX x purchase discount)",
           "chk_pr00": "Billed average price / PR00 list (far from 1 = list or billing odd)",
           "cp": "Factory move used (a fall counts as 0)",
           "q_PR00": "Qty PR00", "q_UTI": "Qty UTI", "q_UMI": "Qty UMI", "q_OTH": "Qty other",
           "bp0": "Factory base price before (JPY)", "bp1": "Factory base price after (JPY)",
           "r1": "Row on By L1", "r2": "Row on By L2", "r3": "Row on By L3",
           "su_PR00": "PR00: setting used", "su_UTI": "UTI: setting used", "su_UMI": "UMI: setting used",
           "sm_PR00": "PR00: move from the setting (empty = scaled)", "sm_UTI": "UTI: move from the setting",
           "sm_UMI": "UMI: move from the setting",
           "cap_PR00": "PR00: at ceiling, first pass", "cap_UTI": "UTI: at ceiling, first pass",
           "cap_UMI": "UMI: at ceiling, first pass", "cap_OTH": "Other lists: at ceiling, first pass",
           "ct": "All lists: cost today EUR (known cost)", "ci": "All lists: cost increase EUR",
           "bk": "All lists: billed EUR with a known cost", "ek": "All lists: extra EUR with a known cost"}
    ORDER_COLS = ["pn", "desc", "l1", "l2", "l3", "in_ALL", "in_PR00", "in_UTI", "in_UMI",
                  "mv_PR00", "mv_UTI", "mv_UMI", "mv_OTH", "np_PR00", "np_UTI", "np_UMI", "p_PR00", "p_UTI", "p_UMI",
                  "at_PR00", "at_UTI", "at_UMI", "at_OTH", "c", "flag", "cls", "src", "b_PR00", "b_UTI", "b_UMI", "b_OTH",
                  "ex_PR00", "ex_UTI", "ex_UMI", "ex_OTH", "mc_PR00", "mc_UTI", "mc_UMI", "mc_OTH", "fob", "chk_pr00",
                  "cp", "q_PR00", "q_UTI", "q_UMI", "q_OTH", "bp0", "bp1", "r1", "r2", "r3",
                  "su_PR00", "su_UTI", "su_UMI", "su_OTH", "sm_PR00", "sm_UTI", "sm_UMI", "sm_OTH",
                  "cap_PR00", "cap_UTI", "cap_UMI", "cap_OTH", "ct", "ci", "bk", "ek"]
    assert set(ORDER_COLS) == set(LAB) and len(ORDER_COLS) == len(LAB)
    names = [(k, LAB[k]) for k in ORDER_COLS]
    for i, k in enumerate(ORDER_COLS):
        col[k] = i + 1
    nxt = len(ORDER_COLS) + 1

    def A(k, r=None, abs_=True):
        c = CL(col[k])
        if r is None:
            return f"Parts!${c}${R0}:${c}${R1}"
        return f"${c}{r}" if abs_ else f"{c}{r}"


    def RNG(k):
        return A(k)



    VIEWS = [(ws_1, G1, ["l1"], "By L1"), (ws_2, G2, ["l1", "l2"], "By L2"), (ws_3, G3, ["l1", "l2", "l3"], "By L3")]
    VMAX = max(len(G1), len(G2), len(G3)) + FIRST + 5


    def vref(sheet, nlev, key):
        c = CL(vcols(nlev)[key])
        return f"'{sheet}'!${c}$1:${c}${VMAX}"


    SL = "Settings"
    S_COL = {"PR00": "B", "UTI": "C", "UMI": "D", "OTH": "E", "BOOK": "F"}
    RS = {"base": 22, "fm": 23, "tgt": 24, "ach": 25, "stat": 26, "extra": 27, "cost": 28, "mchg": 29, "m0": 30,
          "m1": 31, "mpts": 32, "known": 33, "scale": 34, "nop": 35, "nopeur": 36,
          "F": 40, "B1": 41, "s1": 42, "cap": 43, "unc": 44, "s2": 45, "ct": 46, "bk": 47, "ek": 48}
    TGT_CELL = {"PR00": "$B$7", "UTI": "$B$8", "UMI": "$B$9", "OTH": "$B$10"}

    # ------------------------------------------------------------------ Parts sheet
    ws = ws_p
    put(ws, 1, 1, "Parts: one row per part number, its setting, its move and its new price on each list", TITLE)
    put(ws, 2, 1, "Blue on yellow = a setting you can type on this part (it beats every level). Grey headers = helper "
                  "columns. Black = values from files or formulas. " + GRAMMAR, SUB)
    put(ws, 3, 1, ORDER, SUB)
    grey = {col[k] for k in ORDER_COLS[ORDER_COLS.index("cp"):]} | {col[k] for k in ("at_PR00", "at_UTI", "at_UMI", "at_OTH")}
    header(ws, HDR_ROW, [n for _, n in names], grey=grey, height=80)
    lvl_lab = ["Part", "Part, all lists", "L3", "L3, all lists", "L2", "L2, all lists", "L1", "L1, all lists"]
    LAB_ARR = "{" + ",".join(f'"{x}"' for x in lvl_lab) + "}"
    NONE_LAB = "None (scaled)"
    T, I_ = CL(col["fob"]), CL(col["cp"])
    for i, (pn, r) in enumerate(P.iterrows()):
        rr = R0 + i
        vals = {"pn": pn, "desc": r.desc, "l1": r.l1, "l2": r.l2, "l3": r.l3, "cls": r.cls if pd.notna(r.cls) else "",
                "src": r.src, "c": float(r.c), "b_PR00": r.b_PR00, "b_UTI": r.b_UTI, "b_UMI": r.b_UMI, "b_OTH": r.b_OTH,
                "q_PR00": r.q_PR00, "q_UTI": r.q_UTI, "q_UMI": r.q_UMI, "q_OTH": r.q_OTH, "bp0": r.bp0, "bp1": r.bp1,
                "p_PR00": None if pd.isna(r.p_PR00) else float(r.p_PR00),
                "p_UTI": None if pd.isna(r.p_UTI) else float(r.p_UTI),
                "p_UMI": None if pd.isna(r.p_UMI) else float(r.p_UMI),
                "chk_pr00": None if pd.isna(r.chk_pr00) else round(float(r.chk_pr00), 3),
                "flag": "Factory move above 15%" if r.c > 0.15 else "",
                "r1": int(r.r1), "r2": int(r.r2), "r3": int(r.r3)}
        for k, v in vals.items():
            ws.cell(rr, col[k], v)
        ws.cell(rr, col["cp"], f"=MAX({CL(col['c'])}{rr},0)")
        bp0 = CL(col["bp0"])
        ws.cell(rr, col["fob"], f'=IF({CL(col["src"])}{rr}="J",{bp0}{rr}/{SL}!$B$13*{SL}!$B$14,'
                                f'IF({CL(col["src"])}{rr}="T",{bp0}{rr}/{SL}!$B$13*{SL}!$B$15,""))')
        if pn in SET_PART:
            ws.cell(rr, col["in_ALL"], SET_PART[pn]["ALL"])
        for k in LISTS + ["OTH"]:
            pa = f"{CL(col['in_ALL'])}{rr}"
            r1, r2, r3 = f"${CL(col['r1'])}{rr}", f"${CL(col['r2'])}{rr}", f"${CL(col['r3'])}{rr}"
            a3, a2, a1 = (f"INDEX({vref('By L3', 3, 's_ALL')},{r3})", f"INDEX({vref('By L2', 2, 's_ALL')},{r2})",
                          f"INDEX({vref('By L1', 1, 's_ALL')},{r1})")
            if k == "OTH":   # the other own lists take only the 'all lists' settings
                order, labs = [pa, a3, a2, a1], [lvl_lab[1], lvl_lab[3], lvl_lab[5], lvl_lab[7]]
            else:
                pk = f"{CL(col['in_' + k])}{rr}"
                order = [pk, pa, f"INDEX({vref('By L3', 3, 's_' + k)},{r3})", a3,
                         f"INDEX({vref('By L2', 2, 's_' + k)},{r2})", a2, f"INDEX({vref('By L1', 1, 's_' + k)},{r1})", a1]
                labs = lvl_lab
            f_at = f'"{NONE_LAB}"'
            for ref, lab in reversed(list(zip(order, labs))):
                f_at = f'IF({ref}<>"","{lab}",{f_at})'
            lab_arr = "{" + ",".join(f'"{x}"' for x in labs) + "}"
            at = f"{CL(col['at_' + k])}{rr}"
            su = f"{CL(col['su_' + k])}{rr}"
            sm = f"{CL(col['sm_' + k])}{rr}"
            ws.cell(rr, col[f"at_{k}"], "=" + f_at)
            ws.cell(rr, col[f"su_{k}"], f'=IF({at}="{NONE_LAB}","",CHOOSE(MATCH({at},{lab_arr},0),{",".join(order)}))')
            pts = (f'IFERROR(_xlfn.NUMBERVALUE(SUBSTITUTE(SUBSTITUTE(MID(TRIM({su}),2,20),",","."),"%",""),".","\'"),0)')
            ws.cell(rr, col[f"sm_{k}"],
                    f'=IF(OR({su}="",UPPER(TRIM({su}))="S"),"",IF(ISNUMBER({su}),{su},'
                    f'IF(LEFT(UPPER(TRIM({su})),1)="F",MAX({SL}!$B$12,${I_}{rr}+{pts}/100),"")))')
            sc = S_COL[k]
            ws.cell(rr, col[f"cap_{k}"], f'=IF(AND({sm}="",${I_}{rr}*{SL}!${sc}${RS["s1"]}>${I_}{rr}+{SL}!$B$11),1,0)')
            mv = f"{CL(col['mv_' + k])}{rr}"
            ws.cell(rr, col[f"mv_{k}"], f'=IF({sm}="",MAX({SL}!$B$12,MIN({SL}!${sc}${RS["s2"]}*${I_}{rr},'
                                        f'${I_}{rr}+{SL}!$B$11)),{sm})')
            if k != "OTH":
                pr = f"{CL(col['p_' + k])}{rr}"
                ws.cell(rr, col[f"np_{k}"], f'=IF(N({pr})>0,ROUND({pr}*(1+{mv}),2),"")')
            b, q = f"{CL(col['b_' + k])}{rr}", f"{CL(col['q_' + k])}{rr}"
            ws.cell(rr, col[f"ex_{k}"], f"={b}*{mv}")
            ws.cell(rr, col[f"mc_{k}"], f'=IF(ISNUMBER(${T}{rr}),{b}*{mv}-{q}*${T}{rr}*${CL(col["c"])}{rr},"")')
        qs = "+".join(f"{CL(col['q_' + k])}{rr}" for k in LISTS + ["OTH"])
        bs = "+".join(f"{CL(col['b_' + k])}{rr}" for k in LISTS + ["OTH"])
        es = "+".join(f"{CL(col['ex_' + k])}{rr}" for k in LISTS + ["OTH"])
        ws.cell(rr, col["ct"], f'=IF(ISNUMBER(${T}{rr}),({qs})*${T}{rr},"")')
        ws.cell(rr, col["ci"], f'=IF(ISNUMBER(${T}{rr}),({qs})*${T}{rr}*${CL(col["c"])}{rr},"")')
        ws.cell(rr, col["bk"], f'=IF(ISNUMBER(${T}{rr}),{bs},"")')
        ws.cell(rr, col["ek"], f'=IF(ISNUMBER(${T}{rr}),{es},"")')

    fmt = {"c": PCT2, "cp": PCT2, "fob": EUR2, "p_PR00": EUR2, "p_UTI": EUR2, "p_UMI": EUR2, "chk_pr00": "0.00",
           "ct": EUR0, "ci": EUR0, "bk": EUR0, "ek": EUR0, "bp0": "#,##0", "bp1": "#,##0"}
    for k in LISTS + ["OTH"]:
        fmt.update({f"b_{k}": EUR0, f"q_{k}": "#,##0", f"mv_{k}": PCT2, f"np_{k}": EUR2, f"ex_{k}": EUR0, f"mc_{k}": EUR0,
                    f"sm_{k}": PCT2, f"su_{k}": PCT1})
    for k, f_ in fmt.items():
        if k in col:
            for (cell,) in ws.iter_rows(min_row=R0, max_row=R1, min_col=col[k], max_col=col[k]):
                cell.number_format = f_
    for k in ("in_ALL", "in_PR00", "in_UTI", "in_UMI"):
        for (cell,) in ws.iter_rows(min_row=R0, max_row=R1, min_col=col[k], max_col=col[k]):
            cell.fill, cell.font, cell.number_format = YELLOW, INPUT, PCT1
        ws.column_dimensions[CL(col[k])].width = 11
    add_validation(ws, f"{CL(col['in_ALL'])}{R0}:{CL(col['in_UMI'])}{R1}", f"{CL(col['in_ALL'])}{R0}")
    widths = {"pn": 13, "desc": 30, "l1": 18, "l2": 20, "l3": 22, "cls": 6, "src": 8, "flag": 22}
    for i in range(1, nxt):
        ws.column_dimensions[CL(i)].width = 12
    for k, w_ in widths.items():
        ws.column_dimensions[CL(col[k])].width = w_
    for k in LISTS + ["OTH"]:
        ws.column_dimensions[CL(col[f"at_{k}"])].width = 15
    ws.freeze_panes = ws.cell(R0, col["l1"])
    ws.auto_filter.ref = f"A{HDR_ROW}:{CL(nxt - 1)}{R1}"
    for k in LISTS:
        for (cell,) in ws.iter_rows(min_row=R0, max_row=R1, min_col=col[f"mv_{k}"], max_col=col[f"mv_{k}"]):
            cell.font = BOLD

    # ------------------------------------------------------------------ view sheets (By L1 / L2 / L3)
    for ws, Gt, keys, sheet in VIEWS:
        nlev = len(keys)
        vc = vcols(nlev)
        lvname = {1: "AIM L1", 2: "AIM L2", 3: "AIM L3"}
        put(ws, 1, 1, f"{sheet}: set a price move per {lvname[nlev][4:]} group, per price list", TITLE)
        put(ws, 2, 1, GRAMMAR, SUB)
        put(ws, 3, 1, ORDER + " Biggest billed first; use the filter arrows to narrow by L1 or L2.", SUB)
        labels = [lvname[j + 1] for j in range(nlev)] + [
            "Parts", "Billed EUR PR00", "Billed EUR UTI 1799", "Billed EUR UMI 1797", "Billed EUR other own lists",
            "Own factory move, weighted", "SETTING all lists", "SETTING PR00", "SETTING UTI 1799", "SETTING UMI 1797",
            "Move PR00", "Move UTI 1799", "Move UMI 1797", "Move other lists", "Extra EUR a year (all lists)",
            "Margin EUR change (all lists)", "Margin today (known cost)", "Margin after", "Setting check"]
        header(ws, HDR_ROW, labels)
        for j in range(4):
            ws.cell(HDR_ROW, vc["s_ALL"] + j).fill = PatternFill("solid", fgColor=TERRA)
        for i, g_ in Gt.iterrows():
            rr = int(g_.row)
            for j, kk in enumerate(keys):
                put(ws, rr, j + 1, g_[kk], BOLD if j == nlev - 1 else BODY)
            put(ws, rr, vc["parts"], int(g_.parts), fmt="#,##0")
            for k in LISTS + ["OTH"]:
                put(ws, rr, vc[f"b_{k}"], float(g_[f"b_{k}"]), fmt=EUR0)
            put(ws, rr, vc["fm"], None if pd.isna(g_.fm) else float(g_.fm), fmt=PCT2)
            crit = ",".join(f"Parts!${CL(col[kk])}${R0}:${CL(col[kk])}${R1},${CL(j + 1)}{rr}" for j, kk in enumerate(keys))
            for k in LISTS + ["OTH"]:
                bc = f"{CL(vc['b_' + k])}{rr}"
                pos = f'{RNG("b_" + k)},">0"'
                put(ws, rr, vc[f"m_{k}"], f'=IFERROR(SUMIFS({RNG("ex_" + k)},{crit},{pos})/'
                                          f'SUMIFS({RNG("b_" + k)},{crit},{pos}),"")',
                    BOLD if k != "OTH" else BODY, fmt=PCT2)
            put(ws, rr, vc["extra"], "=" + "+".join(f"SUMIFS({RNG('ex_' + k)},{crit})" for k in LISTS + ["OTH"]), fmt=EUR0)
            put(ws, rr, vc["mchg"], "=" + "+".join(f"SUMIFS({RNG('mc_' + k)},{crit})" for k in LISTS + ["OTH"]), fmt=EUR0)
            ct, ci, bk, ek = (f"SUMIFS({RNG(x)},{crit})" for x in ("ct", "ci", "bk", "ek"))
            put(ws, rr, vc["m0"], f'=IFERROR(1-{ct}/{bk},"")', fmt=PCT1)
            put(ws, rr, vc["m1"], f'=IFERROR(1-({ct}+{ci})/({bk}+{ek}),"")', fmt=PCT1)
            ins = [f"{CL(vc[s])}{rr}" for s in ("s_ALL", "s_PR00", "s_UTI", "s_UMI")]
            put(ws, rr, vc["chk"], f'=IF(AND({",".join(okf(x) for x in ins)}),"","Not understood: type 0%, 4%, F, F+2 or S")',
                Font(name=F, size=9, bold=True, color=TERRA))
            for s in ("s_ALL", "s_PR00", "s_UTI", "s_UMI"):
                put(ws, rr, vc[s], None, INPUT, PCT1, fill=YELLOW)
            if nlev == 3 and (g_.l1, g_.l2, g_.l3) in SET_L3:
                ws.cell(rr, vc["s_ALL"], SET_L3[(g_.l1, g_.l2, g_.l3)]["ALL"])
        last = FIRST + len(Gt) - 1
        add_validation(ws, f"{CL(vc['s_ALL'])}{FIRST}:{CL(vc['s_UMI'])}{last}", f"{CL(vc['s_ALL'])}{FIRST}")
        # light terracotta on rows that carry a setting
        rng = f"A{FIRST}:{CL(vc['chk'])}{last}"
        first_in = f"${CL(vc['s_ALL'])}{FIRST}:${CL(vc['s_UMI'])}{FIRST}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"COUNTA({first_in})>0"], fill=TFILL))
        for j in range(nlev):
            ws.column_dimensions[CL(j + 1)].width = 24
        for c_ in range(nlev + 1, vc["chk"] + 1):
            ws.column_dimensions[CL(c_)].width = 11.5
        ws.column_dimensions[CL(vc["chk"])].width = 30
        ws.freeze_panes = ws.cell(FIRST, nlev + 1)
        ws.auto_filter.ref = f"A{HDR_ROW}:{CL(vc['chk'])}{last}"
        tr = last + 2
        put(ws, tr, 1, "Total", BOLD)
        for k in LISTS + ["OTH"]:
            c_ = CL(vc[f"b_{k}"])
            put(ws, tr, vc[f"b_{k}"], f"=SUM({c_}{FIRST}:{c_}{last})", BOLD, EUR0)
        for k in ("extra", "mchg"):
            c_ = CL(vc[k])
            put(ws, tr, vc[k], f"=SUM({c_}{FIRST}:{c_}{last})", BOLD, EUR0)

    # ------------------------------------------------------------------ Settings
    ws = ws_s
    put(ws, 1, 1, "Price revision by AIM level: PR00, UTI (1799) and UMI (1797)", TITLE)
    put(ws, 2, 1, "A tool on paper. Nothing published, no SAP price or price list changed. Built 15 October 2026 on what "
                  "distributors bought from July 2025 to June 2026: EUR 32,595,853 on 15,273 part numbers, the same base "
                  "as the price revision workbook of October 2026.", SUB)
    put(ws, 4, 1, "Settings - change the yellow cells; every figure in the workbook recalculates", SEC)
    rows = [(5, "Target: weighted increase on the last 12 months' billed euros", TARGET, PCT2),
            (6, "Target applies to: 1 = the whole book (one scaling for every list), 2 = each list on its own target", MODE,
             "0"),
            (7, "Own target, PR00 (used only when B6 = 2)", TARGET, PCT2),
            (8, "Own target, UTI 1799 (used only when B6 = 2)", TARGET, PCT2),
            (9, "Own target, UMI 1797 (used only when B6 = 2)", TARGET, PCT2),
            (10, "Own target, other own lists: Palestine 2372, Armenia 2733, Kazakhstan 2734 (used only when B6 = 2)",
             TARGET, PCT2),
            (11, "Ceiling for a part on the scaled factory move: its own factory move plus this many points "
                 "(decided 30 Sep and 2 Oct 2026)", CEIL, PCT1),
            (12, "Lowest move for a part on the factory move (a % you type is used exactly as typed)", FLOOR, PCT1),
            (13, "Exchange rate, JPY per EUR (cost basis for the margin; config of 1 Oct 2026)", FX, "0.0"),
            (14, "Purchase discount, Japan (IML)", DISC["J"], "0.000"),
            (15, "Purchase discount, Thailand (IMIT)", DISC["T"], "0.000")]
    for r, lab, v, f_ in rows:
        put(ws, r, 1, lab, BODY, al=WRAP)
        put(ws, r, 2, v, INPUT, f_, fill=YELLOW)
    dvm = DataValidation(type="whole", operator="between", formula1="1", formula2="2", showErrorMessage=True,
                         error="Type 1 (whole book) or 2 (each list).")
    dvm.add("B6")
    ws.add_data_validation(dvm)
    put(ws, 17, 1, "How to type a setting (on By L1, By L2, By L3, or on a single part on Parts)", SEC)
    put(ws, 18, 1, GRAMMAR, BODY, al=WRAP)
    put(ws, 19, 1, ORDER, BODY, al=WRAP)
    for r_ in (18, 19):
        ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=6)
        ws.row_dimensions[r_].height = 30
    header(ws, 21, ["Result", LIST_NAME["PR00"], LIST_NAME["UTI"], LIST_NAME["UMI"], LIST_NAME["OTH"], "Whole book"])
    labels = {"base": "Billed last 12 months, EUR", "fm": "Own factory move, weighted",
              "tgt": "Target", "ach": "Move achieved, weighted", "stat": "Status",
              "extra": "Extra EUR a year on the same volume", "cost": "Cost increase EUR (parts with a known cost)",
              "mchg": "Margin EUR change (extra EUR less cost increase, parts with a known cost)",
              "m0": "Margin today (on billed, parts with a known cost)", "m1": "Margin after",
              "mpts": "Margin change, points", "known": "Share of billed with a known cost",
              "scale": "Scaling applied to each factory move (1.00 = the full factory move)",
              "nop": "Parts billed on this list with no current price on it (they get a move, no new price)",
              "nopeur": "Billed EUR on those parts",
              "F": "Helper: EUR from parts with a setting", "B1": "Helper: billed x factory move, parts without a setting",
              "s1": "Helper: scaling, first pass", "cap": "Helper: EUR of parts held at their ceiling (first pass)",
              "unc": "Helper: billed x factory move, parts without a setting and below the ceiling",
              "s2": "Helper: scaling, second pass (used)", "ct": "Helper: cost today, EUR (known cost)",
              "bk": "Helper: billed EUR with a known cost", "ek": "Helper: extra EUR with a known cost"}
    fmts = {"base": EUR0, "fm": PCT2, "tgt": PCT2, "ach": PCT2, "extra": EUR0, "cost": EUR0, "mchg": EUR0, "m0": PCT1,
            "m1": PCT1, "mpts": PCT2, "known": PCT1, "scale": "0.0000", "nop": "#,##0", "nopeur": EUR0, "F": EUR0,
            "B1": EUR0, "s1": "0.000000", "cap": EUR0, "unc": EUR0, "s2": "0.000000", "ct": EUR0, "bk": EUR0, "ek": EUR0}
    put(ws, 39, 1, "Helpers (how the scaling is solved; no need to touch)", Font(name=F, bold=True, size=9, color="777777"))
    for key, r in RS.items():
        put(ws, r, 1, labels[key], BODY if r < 39 else Font(name=F, size=9, color="777777"), al=WRAP)
    B = lambda k: f"$B${RS[k]}"  # noqa: E731
    for k in LISTS + ["OTH"]:
        c = S_COL[k]
        b, q = RNG(f"b_{k}"), RNG(f"q_{k}")
        fobr, cr, cpr = RNG("fob"), RNG("c"), RNG("cp")
        f = {}
        f["base"] = f"=SUM({b})"
        f["fm"] = f"=SUMPRODUCT({b},{cr})/{c}{RS['base']}"
        f["tgt"] = f"=IF($B$6=1,$B$5,{TGT_CELL[k]})"
        f["ach"] = f"={c}{RS['extra']}/{c}{RS['base']}"
        f["stat"] = (f'=IF($B$6=1,"On the book target",IF(ABS({c}{RS["ach"]}-{c}{RS["tgt"]})<0.00005,"OK",'
                     f'IF({c}{RS["ach"]}<{c}{RS["tgt"]},"Not reached: ","Above target: ")&'
                     f'ROUND(({c}{RS["ach"]}-{c}{RS["tgt"]})*100,2)&" points"))')
        f["extra"] = f"=SUM({RNG('ex_' + k)})"
        f["cost"] = f"=SUMPRODUCT({q},{fobr},{cr})"
        f["mchg"] = f"=SUM({RNG('mc_' + k)})"
        f["m0"] = f"=IFERROR(1-{c}{RS['ct']}/{c}{RS['bk']},\"\")"
        f["m1"] = f"=IFERROR(1-({c}{RS['ct']}+{c}{RS['cost']})/({c}{RS['bk']}+{c}{RS['ek']}),\"\")"
        f["mpts"] = f"=IFERROR({c}{RS['m1']}-{c}{RS['m0']},\"\")"
        f["known"] = f"={c}{RS['bk']}/{c}{RS['base']}"
        f["scale"] = f"={c}{RS['s2']}"
        if k in LISTS:
            f["nop"] = f'=COUNTIFS({b},">0",{RNG("p_" + k)},"")'
            f["nopeur"] = f'=SUMIFS({b},{RNG("p_" + k)},"")'
            f["F"] = f"=SUMPRODUCT({b},{RNG('sm_' + k)})"
            f["B1"] = f'=SUMPRODUCT({b},{cpr},--({RNG("sm_" + k)}=""))'
        else:
            f["nop"] = "not printed"
            f["nopeur"] = "not printed"
            f["F"] = f"=SUMPRODUCT({b},{RNG('sm_OTH')})"
            f["B1"] = f'=SUMPRODUCT({b},{cpr},--({RNG("sm_OTH")}=""))'
        f["s1"] = (f"=IF($B$6=1,$F${RS['s1']},IFERROR(({c}{RS['tgt']}*{c}{RS['base']}-{c}{RS['F']})/{c}{RS['B1']},0))")
        f["cap"] = f"=SUMPRODUCT({b},{RNG('cap_' + k)},{cpr}+$B$11)"
        f["unc"] = f"={c}{RS['B1']}-SUMPRODUCT({b},{RNG('cap_' + k)},{cpr})"
        f["s2"] = (f"=IF($B$6=1,$F${RS['s2']},IFERROR(({c}{RS['tgt']}*{c}{RS['base']}-{c}{RS['F']}-{c}{RS['cap']})/"
                   f"{c}{RS['unc']},0))")
        f["ct"] = f"=SUMPRODUCT({q},{fobr})"
        f["bk"] = f"=SUMPRODUCT({b},--ISNUMBER({fobr}))"
        f["ek"] = f"=SUMPRODUCT({b},{RNG('mv_' + k)},--ISNUMBER({fobr}))"
        for key, v in f.items():
            put(ws, RS[key], ord(c) - 64, v, BOLD if key in ("ach", "extra", "mchg", "stat") else BODY, fmts.get(key))
    # whole book column F
    fb = {}
    sumrow = lambda key: f"=SUM(B{RS[key]}:E{RS[key]})"  # noqa: E731
    for key in ("base", "extra", "cost", "mchg", "F", "B1", "cap", "unc", "ct", "bk", "ek", "nop", "nopeur"):
        fb[key] = sumrow(key)
    fb["nop"] = f"=SUM(B{RS['nop']}:D{RS['nop']})"
    fb["nopeur"] = f"=SUM(B{RS['nopeur']}:D{RS['nopeur']})"
    fb["fm"] = f"=SUMPRODUCT(B{RS['base']}:E{RS['base']},B{RS['fm']}:E{RS['fm']})/F{RS['base']}"
    fb["tgt"] = f"=IF($B$6=1,$B$5,SUMPRODUCT(B{RS['base']}:E{RS['base']},B{RS['tgt']}:E{RS['tgt']})/F{RS['base']})"
    fb["ach"] = f"=F{RS['extra']}/F{RS['base']}"
    fb["stat"] = (f'=IF(ABS(F{RS["ach"]}-F{RS["tgt"]})<0.00005,"OK",IF(F{RS["ach"]}<F{RS["tgt"]},"Not reached: ",'
                  f'"Above target: ")&ROUND((F{RS["ach"]}-F{RS["tgt"]})*100,2)&" points")')
    fb["m0"] = f"=IFERROR(1-F{RS['ct']}/F{RS['bk']},\"\")"
    fb["m1"] = f"=IFERROR(1-(F{RS['ct']}+F{RS['cost']})/(F{RS['bk']}+F{RS['ek']}),\"\")"
    fb["mpts"] = f"=IFERROR(F{RS['m1']}-F{RS['m0']},\"\")"
    fb["known"] = f"=F{RS['bk']}/F{RS['base']}"
    fb["scale"] = f'=IF($B$6=1,F{RS["s2"]},"per list")'
    fb["s1"] = f"=IFERROR(($B$5*F{RS['base']}-F{RS['F']})/F{RS['B1']},0)"
    fb["s2"] = f"=IFERROR(($B$5*F{RS['base']}-F{RS['F']}-F{RS['cap']})/F{RS['unc']},0)"
    for key, v in fb.items():
        put(ws, RS[key], 6, v, BOLD if key in ("ach", "extra", "mchg", "stat") else BODY, fmts.get(key))
    for r in range(RS["F"], RS["ek"] + 1):
        for c_ in range(2, 7):
            ws.cell(r, c_).font = Font(name=F, size=9, color="777777")
    ws.conditional_formatting.add(f"B{RS['stat']}:F{RS['stat']}",
                                  FormulaRule(formula=[f'LEFT(B{RS["stat"]},3)="Not"'], font=Font(color=TERRA, bold=True)))
    ws.column_dimensions["A"].width = 62
    for c_ in "BCDEF":
        ws.column_dimensions[c_].width = 21
    ws.freeze_panes = "B5"
    put(ws, 50, 1, "Where the settings are now", SEC)
    cnt = [("By L1", ws_1, 1, G1), ("By L2", ws_2, 2, G2), ("By L3", ws_3, 3, G3)]
    for i, (lab, _, nlev, Gt) in enumerate(cnt):
        vc = vcols(nlev)
        rg = f"'{lab}'!{CL(vc['s_ALL'])}{FIRST}:{CL(vc['s_UMI'])}{FIRST + len(Gt) - 1}"
        put(ws, 51 + i, 1, f"Settings typed on {lab}")
        put(ws, 51 + i, 2, f"=COUNTA({rg})", fmt="#,##0")
    put(ws, 54, 1, "Settings typed on single parts (Parts sheet)")
    put(ws, 54, 2, f"=COUNTA(Parts!${CL(col['in_ALL'])}${R0}:${CL(col['in_UMI'])}${R1})", fmt="#,##0")
    put(ws, 55, 1, "Settings not understood anywhere (must be 0)")
    bad = []
    for lab, _, nlev, Gt in cnt:
        vc = vcols(nlev)
        bad.append(f"COUNTIF('{lab}'!{CL(vc['chk'])}{FIRST}:{CL(vc['chk'])}{FIRST + len(Gt) - 1},\"Not*\")")
    put(ws, 55, 2, "=" + "+".join(bad), BOLD, "#,##0")

    # chart: move by L1 per list against the factory move
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "r"
    ch.y_axis.crosses = "max"
    ch.y_axis.number_format = "0%"
    vc1 = vcols(1)
    last1 = FIRST + len(G1) - 1
    cats = Reference(ws_1, min_col=1, min_row=FIRST, max_row=last1)
    for key in ("fm", "m_PR00", "m_UTI", "m_UMI"):
        data = Reference(ws_1, min_col=vc1[key], min_row=HDR_ROW, max_row=last1)
        ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.height, ch.width = 16, 22
    for s, colr in zip(ch.series, ["9E9E9E", NAVY, TERRA, "6B8FB8"]):
        s.graphicalProperties.solidFill = colr
        s.graphicalProperties.line.solidFill = colr
    ch.x_axis.scaling.orientation = "maxMin"
    put(ws, 4, 8, "Which parts families carry the increase, per price list? Move by AIM L1 group against its own "
                  "factory move (grey).", SEC)
    ws.add_chart(ch, "H5")

    # ------------------------------------------------------------------ Checks
    ws = ws_c
    put(ws, 1, 1, "Checks - every line should read OK", TITLE)
    header(ws, 3, ["Check", "Value", "Expected", "Result"])
    chk = [
        ("Billed base equals the published in-scope priced base (EUR)", f"={SL}!F{RS['base']}", EXPECT["billed"], "abs1"),
        ("Part numbers in the tool", f"=COUNTA(Parts!$A${R0}:$A${R1})", EXPECT["parts"], "eq"),
        ("The four lists add up to the base (EUR)", f"={SL}!B{RS['base']}+{SL}!C{RS['base']}+{SL}!D{RS['base']}+"
                                                  f"{SL}!E{RS['base']}", EXPECT["billed"], "abs1"),
        ("Parts with no row in the AIM release of 11 Oct 2026 (grouped as 'Not classified'; they still get a move)",
         f'=COUNTIF(Parts!$C${R0}:$C${R1},"Not classified")', int((P.l1 == "Not classified").sum()), "eq"),
        ("Target reached (whole book, or each list when B6 = 2)", f'=IF({SL}!$B$6=1,'
                                                                  f'{SL}!F{RS["stat"]},IF(COUNTIF({SL}!B{RS["stat"]}:'
                                                                  f'E{RS["stat"]},"OK")=4,"OK","not on every list"))',
         "OK", "txt"),
        ("Settings not understood (must be 0)", f"={SL}!B55", 0, "eq"),
        ("Extra EUR on the parts equals the extra EUR on the settings sheet",
         f"=SUM('By L1'!{CL(vcols(1)['extra'])}{FIRST}:{CL(vcols(1)['extra'])}{last1})", f"={SL}!F{RS['extra']}", "absf"),
        ("No part on the scaled factory move moved below the lowest move",
         "=" + "+".join(f'COUNTIFS({RNG("sm_" + k)},"",{RNG("mv_" + k)},"<"&{SL}!$B$12)' for k in LISTS), 0, "eq"),
    ]
    for i, (lab, v, e, how) in enumerate(chk):
        r = 4 + i
        put(ws, r, 1, lab, al=WRAP)
        put(ws, r, 2, v, fmt=EUR2 if how.startswith("abs") else None)
        put(ws, r, 3, e, fmt=EUR2 if how.startswith("abs") else None)
        if how == "abs1":
            put(ws, r, 4, f'=IF(ABS(B{r}-C{r})<0.01,"OK","CHECK")', BOLD)
        elif how == "absf":
            put(ws, r, 4, f'=IF(ABS(B{r}-C{r})<0.01,"OK","CHECK")', BOLD)
        elif how == "eq":
            put(ws, r, 4, f'=IF(B{r}=C{r},"OK","CHECK")', BOLD)
        else:
            put(ws, r, 4, f'=IF(B{r}=C{r},"OK","CHECK")', BOLD)
    r = 4 + len(chk) + 1
    put(ws, r, 1, "Checked when the workbook was built (Python, not a live formula)", SEC)
    built = [
        f"With every setting empty and the target at 5.00%, each part's move equals the 'Factory-led only move' column "
        f"of the price revision workbook of October 2026, on all {len(cmp):,} part numbers (largest difference "
        f"{NEUTRAL_MAXDIFF:.1e}).",
        f"The five pad kits labelled 'Brake Disc' in the AIM release of 11 Oct 2026 ({', '.join(PAD_KITS_MISLABELLED)}) "
        f"carry their own 0% on the Parts sheet, so the 30 Sep 2026 decision to keep pad kits at today's price holds even "
        f"though their L3 label is wrong.",
        "Prices of the two Israel lists come from the price export of 10 Aug 2026, the newest on file. UTI: 89.2% of its "
        "billed euros has a price there; UMI: 83.7%. Parts without a price get a move but no new price.",
    ]
    neg = [(pn, k, float(P.at[pn, f"b_{k}"])) for k in LISTS + ["OTH"] for pn in P.index[P[f"b_{k}"] < 0]]
    if neg:
        built.append("Parts whose credits in the 12 months exceed their sales on a list (net billed below zero): "
                     + "; ".join(f"{pn} on {k}, EUR {v:,.0f}" for pn, k, v in neg)
                     + ". They stay in the base, as in the price revision workbook; the group moves on the By sheets "
                       "are weighted on parts with positive billing so one credit cannot distort a group average.")
    hi = P[P.c > 0.15]
    built.append(f"{len(hi)} part numbers have an own factory move above 15% (EUR {hi.billed.sum():,.0f} billed; "
                 f"{int((P.c > 0.5).sum())} of them above 50%, EUR {P.billed[P.c > 0.5].sum():,.0f}). Without a setting "
                 f"they follow that move, scaled and capped at their own move plus the ceiling. They are marked "
                 f"'Factory move above 15%' in the Check column on Parts.")
    for j, t in enumerate(built):
        put(ws, r + 1 + j, 1, t, al=WRAP)
        ws.merge_cells(start_row=r + 1 + j, start_column=1, end_row=r + 1 + j, end_column=4)
        ws.row_dimensions[r + 1 + j].height = 16 * (len(t) // 140 + 1) + 2
    ws.column_dimensions["A"].width = 70
    for c_ in "BC":
        ws.column_dimensions[c_].width = 18
    ws.column_dimensions["D"].width = 10
    ws.conditional_formatting.add(f"D4:D{3 + len(chk)}", FormulaRule(formula=['D4<>"OK"'], font=Font(color=TERRA, bold=True)))

    # ------------------------------------------------------------------ Read me
    ws = ws_rm
    txt = [
        ("Price revision by AIM level", TITLE),
        ("A tool on paper, built 15 October 2026. Nothing published, no SAP price, price list or pricing configuration "
         "changed.", SUB),
        ("", None),
        ("What it does", SEC),
        ("It sets a price move for every part on three price lists: PR00 (the EU and UK distributors), UTI Israel trucks "
         "(customer 1799) and UMI Israel D-Max (customer 1797). The two Israel distributors each have their own price list "
         "and are never blended. You can set the move on any AIM level - L1, L2 or L3 - or on a single part, for all "
         "three lists at once or for one list only. Every part without a setting follows its own factory move (the IML or "
         "IMIT base price change), scaled so that the whole book, or each list, lands exactly on the target.", BODY),
        ("", None),
        ("How to use it", SEC),
        ("1. Settings: set the target (default +5.30%, the figure in the current discussion with IML) and choose one "
         "target for the whole book or one per list.", BODY),
        ("2. By L1, By L2, By L3: type a setting in the yellow cells of a group. " + GRAMMAR, BODY),
        ("3. " + ORDER, BODY),
        ("4. Read the result on Settings (per list: move, extra euros, margin today and after) and on each By sheet. "
         "New prices per list are on Parts.", BODY),
        ("", None),
        ("What is already set", SEC),
        ("Brake pad kits at today's price on every list (decided 30 September 2026): on By L3, 'Brake Pad Kit' = 0%, and "
         "on Parts the five pad kits that the AIM release of 11 October labels 'Brake Disc' carry their own 0%. Nothing "
         "else is set.", BODY),
        ("", None),
        ("What it does not do", SEC),
        ("Distributor conditions inside PR00 are not here: the service-part holds in eight markets, the holds where a "
         "distributor sells under 45% of its trucks' need and the extra where a UK shop sells above our reference dealer "
         "price. Those stay in the price revision workbook of October 2026; they change what a distributor pays, not the "
         "list price. The total at a given target is the same in both files; the split across parts is not.", BODY),
        ("It covers the 15,273 part numbers billed in the 12 months; parts not billed in that window are not repriced. "
         "Palestine (2372), Armenia (2733) and Kazakhstan (2734) have their own price lists too; they follow the scaled "
         "factory move, take only the 'all lists' settings (so the pad kits stay at today's price there too) and count in the whole-book target; their new prices are not printed. "
         "Lubricants and lines without a factory price on both sides are outside the base.", BODY),
        ("", None),
        ("How the margin is counted", SEC),
        ("Cost per piece = factory base price in yen / exchange rate x purchase discount (Japan 0.673, Thailand 0.611, "
         "exchange rate 178.5, all yellow on Settings). Margin = 1 - cost / billed euros, on parts with a known source "
         "(about 95% of the billed euros). The margin euro change is the extra euros less the cost increase at the same "
         "volume. Exchange rate and discounts stay where they are; only the factory move and our price move change.", BODY),
        ("", None),
        ("Sources", SEC),
        ("Base: 01_data/00_derived/price_bp_weighted_compare_pn.csv (in scope and priced on both sides). AIM levels: "
         "01_data/00_derived/_releases/aim_pn_classification_release_2026-10-11.csv. PR00 list: IE40 list of 13 Mar 2026 "
         "(pricing_pn_enquiry.parquet). Israel lists: price export of 10 Aug 2026 (pricing_zpr1_key_wide.parquet). "
         "Rates: 03_analytics/isze_pricing/config.yaml. Builder: 02_dashboards/pricing/_build/price_by_aim_level_2026-10-15/"
         "build_price_by_aim_level.py.", SUB),
    ]
    for i, (t, f_) in enumerate(txt):
        if f_ is None:
            continue
        put(ws, i + 1, 1, t, f_, al=WRAP)
        if f_ is BODY or f_ is SUB:
            ws.row_dimensions[i + 1].height = max(15, 15 * (len(t) // 120 + 1))
    ws.column_dimensions["A"].width = 130

    wb.active = 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB), parts {N}, groups L1 {len(G1)} L2 {len(G2)} L3 {len(G3)}")


if __name__ == "__main__":
    main()
