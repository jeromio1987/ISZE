"""Two lists for Jerome, 19 Oct 2026: parts going up more than 15%, and parts bought but with no price on their list."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
H = Path(__file__).parent
OUT = H.parents[3] / "01_data/01_Pricing/ISZE_Price_Revision_lists_2026-10-19.xlsx"
p = pd.read_pickle(H / "parts_v2_all_brakes.pkl")
wb = Workbook()
def sheet(ws, title, cols, rows):
    ws["A1"] = title; ws["A1"].font = Font(name="Georgia", size=13, bold=True, color="15233C")
    for j, (h, w, f) in enumerate(cols, 1):
        c = ws.cell(3, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="15233C")
        c.alignment = Alignment(wrap_text=True, vertical="center"); ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[3].height = 32
    for i, r in enumerate(rows, 4):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, None if (isinstance(v, float) and pd.isna(v)) else v)
            if cols[j - 1][2]: c.number_format = cols[j - 1][2]
    ws.freeze_panes = "B4"; ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{3 + len(rows)}"
big = p[(p.pr00_m > 0.15) & (p.b_tot > 0)].sort_values("pr00_m", ascending=False)
ws = wb.active; ws.title = "Up more than 15%"
sheet(ws, f"{len(big)} parts that go up more than 15% because their base price jumped (billed July 2025 - June 2026)",
      [("Part number", 13, None), ("Description", 30, None), ("Family level 1", 20, None), ("Family level 3", 22, None), ("Made in", 15, None),
       ("Base price before (yen)", 11, "#,##0"), ("Base price after (yen)", 11, "#,##0"), ("Base price change", 10, "0.0%"),
       ("Proposed increase", 10, "0.0%"), ("PR00 price today (EUR)", 11, "#,##0.00"), ("PR00 new price (EUR)", 11, "#,##0.00"),
       ("Billed, all lists (EUR)", 12, "#,##0"), ("Extra per year (EUR)", 11, "#,##0")],
      [[x.pn, x.desc, x.l1, x.l3, x.src, x.bp0, x.bp1, x.bp_step, x.pr00_m, x.pr00_p,
        (round(x.pr00_p * (1 + x.pr00_m), 2) if pd.notna(x.pr00_p) else None), x.b_tot, x.x_tot] for x in big.itertuples()])
rows = []
for nm, lab, pc in (("pr00", "PR00", "pr00_p"), ("uti", "UTI 1799", "uti_p"), ("umi", "UMI 1797", "umi_p")):
    q = p[(p[nm + "_b"] > 0) & p[pc].isna()].sort_values(nm + "_b", ascending=False)
    rows += [[lab, x.pn, x.newest or None, x.desc, x.l1, x.l3, getattr(x, nm + "_b"), getattr(x, nm + "_m")] for x in q.itertuples()]
ws2 = wb.create_sheet("No price on the list")
sheet(ws2, f"{len(rows)} parts that were billed in July 2025 - June 2026 but have no price on that list in the SAP export of 10 August 2026",
      [("Price list", 10, None), ("Part number", 13, None), ("Newer part number in Power BI", 14, None), ("Description", 30, None),
       ("Family level 1", 20, None), ("Family level 3", 22, None), ("Billed on this list (EUR)", 12, "#,##0"), ("Proposed increase", 10, "0.0%")], rows)
wb.save(OUT); print(OUT, len(big), len(rows))
print(big[["pn", "desc", "src", "bp_step", "pr00_m", "b_tot"]].head(12).to_string())
print(pd.DataFrame(rows).groupby(0)[6].agg(["count", "sum"]))
