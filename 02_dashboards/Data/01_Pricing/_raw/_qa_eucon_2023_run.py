"""Hard QA for Eucon 2023 Price Survey rawdata — Mode B scratch; not a deliverable."""
from __future__ import annotations

import csv
import json
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

RAW = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\01_Pricing\_raw")
DER = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\00_derived")
BASKET = Path(r"C:\Users\jerom\Desktop\claude\ISZE\03_analytics\isze_pricing\oem_peer_basket.csv")
OUT = RAW / "_qa_eucon_2023_scratch.json"

UK_OEM = RAW / "Eucon_PriceSurvey_UK_2023_OEM_rawdata_20240109.xlsx"
ITUK_OEM = RAW / "Eucon_PriceSurvey_ITUK_2023_OEM_rawdata.xlsx"
UK_IAM = RAW / "Eucon_PriceSurvey_UK_2023_IAM_rawdata_20240109.xlsx"


def pn_norm(x) -> str:
    if x is None:
        return ""
    s = str(x).strip()
    if not s or s.lower() in ("nan", "none"):
        return ""
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    s = re.sub(r"[^0-9A-Za-z]", "", s)
    if s.isdigit() and len(s) <= 10:
        return s.zfill(10)
    return s


def to_float(x):
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", ""))
    except ValueError:
        return None


def load_sheet(path: Path, sheet: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return rows


def brand_match(survey_brand: str, peer_oem: str) -> bool:
    sb = (survey_brand or "").lower()
    po = (peer_oem or "").lower()
    if not sb or not po:
        return False
    if sb in po or po in sb:
        return True
    if "mercedes" in sb and "mercedes" in po:
        return True
    if sb == "ford" and "ford" in po:
        return True
    if sb == "toyota" and "toyota" in po:
        return True
    if sb == "daf" and "daf" in po:
        return True
    return False


def dup_analysis(data):
    keys = Counter((d["isuzu_pn"], d["peer_oe_pn"], d["peer_brand"]) for d in data)
    dups = [(k, v) for k, v in keys.items() if v > 1]
    samples = []
    for k, v in sorted(dups, key=lambda x: -x[1])[:5]:
        rows = [d for d in data if (d["isuzu_pn"], d["peer_oe_pn"], d["peer_brand"]) == k]
        vehicles = sorted(
            set(d.get("peer_vehicle", "") + "|" + d.get("vehicle", "") for d in rows)
        )
        samples.append(
            {
                "key": list(k),
                "n": v,
                "vehicles": vehicles[:6],
                "comparabilities": dict(Counter(d["comparability"] for d in rows)),
            }
        )
    return {
        "dup_key_count": len(dups),
        "dup_row_extra": sum(v - 1 for _, v in dups),
        "samples": samples,
    }


def fix(o):
    if isinstance(o, Counter):
        return dict(o)
    if isinstance(o, dict):
        return {str(k): fix(v) for k, v in o.items()}
    if isinstance(o, list):
        return [fix(x) for x in o]
    if isinstance(o, tuple):
        return list(o)
    return o


def main():
    rows = load_sheet(UK_OEM, "Isuzu UK")
    uk = []
    for r in rows[1:]:
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        uk.append(
            {
                "market": "UK",
                "vehicle": str(r[1] or ""),
                "prod": str(r[3] or ""),
                "desc": str(r[4] or ""),
                "isuzu_pn": pn_norm(r[5]),
                "peer_brand": str(r[8] or "").strip(),
                "peer_vehicle": str(r[9] or "").strip(),
                "peer_oe_pn": pn_norm(r[11]),
                "isze_price_main": to_float(r[15]),
                "comp_price_gbp": to_float(r[16]),
                "currency": str(r[17] or "").strip(),
                "comparability": str(r[21] or "").strip().upper(),
            }
        )

    helper_rows = load_sheet(UK_OEM, "Sheet1")
    helper = {}
    for r in helper_rows[1:]:
        if r[0] is None:
            continue
        pn = pn_norm(r[0])
        helper[pn] = {
            "isze_helper": to_float(r[2]),
            "comp_eur_helper": to_float(r[3]),
            "ratio_compEUR_over_isze": to_float(r[4]),
            "comp_gbp_helper": to_float(r[5]),
        }

    for d in uk:
        h = helper.get(d["isuzu_pn"], {})
        d["isze_helper"] = h.get("isze_helper")
        d["comp_eur_helper"] = h.get("comp_eur_helper")

    rows2 = load_sheet(ITUK_OEM, "Sheet1")
    ituk = []
    for r in rows2[1:]:
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        ituk.append(
            {
                "market": "ITUK",
                "vehicle": "",
                "isuzu_pn": pn_norm(r[0]),
                "peer_brand": str(r[3] or "").strip(),
                "peer_vehicle": str(r[4] or "").strip(),
                "peer_oe_pn": pn_norm(r[6]),
                "isze_price_main": to_float(r[10]),
                "comp_price_gbp": to_float(r[11]),
                "currency": str(r[12] or "").strip(),
                "comparability": str(r[16] or "").strip().upper(),
            }
        )

    def parse_iam(seg):
        out = []
        for r in load_sheet(UK_IAM, seg)[1:]:
            if not any(c is not None and str(c).strip() != "" for c in r):
                continue
            out.append(
                {
                    "segment": seg,
                    "isuzu_pn": pn_norm(r[1]),
                    "iam_brand": str(r[4] or "").strip(),
                    "iam_pn": pn_norm(r[5]) if r[5] not in (None, "") else "",
                    "list_price": to_float(r[6]),
                    "workshop_price": to_float(r[7]),
                    "comment": str(r[10] or "").strip(),
                }
            )
        return out

    iam = parse_iam("LCV") + parse_iam("HCV")

    pr00 = pd.read_parquet(DER / "pricing_pr00_ie40.parquet")
    pr00["pn"] = pr00["part_number"].map(pn_norm)
    pr00_active = pr00[~pr00["is_superseded_row"]].copy()
    pr00_map = pr00_active.groupby("pn")["sp_eur"].first().to_dict()
    pr00_active["newest"] = pr00_active["newest_pn"].map(pn_norm)
    newest_map = pr00_active.groupby("newest")["sp_eur"].first().to_dict()

    list_pn = pd.read_parquet(DER / "pricing_list_pn.parquet")
    list_pn["pn"] = list_pn["part_number"].map(pn_norm)
    list_map = (
        list_pn[~list_pn["is_superseded_row"]].groupby("pn")["sp_eur"].first().to_dict()
    )

    gap = pd.read_csv(DER / "eucon_price_gap.csv")
    gap["pn"] = gap["pn"].map(pn_norm)
    gap_map = gap.set_index("pn")["ISZE_IE40_Price_EUR"].to_dict()

    basket_pns = set()
    with BASKET.open(encoding="utf-8") as f:
        basket_rows = list(csv.DictReader(f))
        for row in basket_rows:
            basket_pns.add(pn_norm(row["isze_pn"]))

    uk_unique = {}
    for d in uk:
        if d["isuzu_pn"] not in uk_unique:
            uk_unique[d["isuzu_pn"]] = d

    candidates = sorted((basket_pns & set(uk_unique)) | (set(gap_map) & set(uk_unique)))
    for d in uk:
        if d["isze_price_main"] is not None and d["isuzu_pn"] not in candidates:
            candidates.append(d["isuzu_pn"])
    candidates = candidates[:20]

    spot = []
    for pn in candidates:
        d = uk_unique[pn]
        isze = d["isze_helper"] if d["isze_helper"] is not None else d["isze_price_main"]
        pr = pr00_map.get(pn) or newest_map.get(pn)
        lp = list_map.get(pn)
        gp = gap_map.get(pn)
        row = {
            "pn": pn,
            "survey_isze": isze,
            "pr00_sp_eur": pr,
            "list_sp_eur": lp,
            "gap_ie40": gp,
            "vs_pr00_pct": None
            if (isze is None or not pr)
            else (isze / pr - 1) * 100,
            "vs_street1785_pct": None
            if (isze is None or not pr)
            else (isze / (pr * 1.785) - 1) * 100,
            "survey_looks_like": None,
        }
        if isze is not None and pr:
            err_list = abs(isze / pr - 1)
            err_street = abs(isze / (pr * 1.785) - 1)
            if err_list <= 0.15:
                row["survey_looks_like"] = "near_PR00_list_EUR"
            elif err_street <= 0.15:
                row["survey_looks_like"] = "near_street_x1785"
            elif abs(isze / (pr * 1.40) - 1) <= 0.15:
                row["survey_looks_like"] = "near_trade_x140"
            else:
                row["survey_looks_like"] = "neither_list_nor_street"
        spot.append(row)

    ratios = []
    match_stats = {
        "n_with_both": 0,
        "near_list_15pct": 0,
        "near_street_15pct": 0,
        "near_trade_15pct": 0,
        "far": 0,
    }
    for pn, d in uk_unique.items():
        isze = d["isze_helper"]
        pr = pr00_map.get(pn) or newest_map.get(pn)
        if isze is None or not pr or pr <= 0:
            continue
        match_stats["n_with_both"] += 1
        ratios.append(isze / pr)
        if abs(isze / pr - 1) <= 0.15:
            match_stats["near_list_15pct"] += 1
        elif abs(isze / (pr * 1.785) - 1) <= 0.15:
            match_stats["near_street_15pct"] += 1
        elif abs(isze / (pr * 1.40) - 1) <= 0.15:
            match_stats["near_trade_15pct"] += 1
        else:
            match_stats["far"] += 1
    if ratios:
        match_stats["median_ratio"] = statistics.median(ratios)
        srt = sorted(ratios)
        match_stats["p10"] = srt[max(0, int(0.1 * len(srt)) - 1)]
        match_stats["p90"] = srt[int(0.9 * len(srt)) - 1]

    uk_set = set(d["isuzu_pn"] for d in uk)
    ituk_set = set(d["isuzu_pn"] for d in ituk)
    overlap = sorted(uk_set & ituk_set)
    uk_map = defaultdict(set)
    ituk_map = defaultdict(set)
    for d in uk:
        if d["comparability"] == "COMPARABLE" and d["peer_oe_pn"]:
            uk_map[(d["isuzu_pn"], d["peer_brand"])].add(d["peer_oe_pn"])
    for d in ituk:
        if d["comparability"] == "COMPARABLE" and d["peer_oe_pn"]:
            ituk_map[(d["isuzu_pn"], d["peer_brand"])].add(d["peer_oe_pn"])
    shared = set(uk_map) & set(ituk_map)
    conflicts = []
    agrees = 0
    for k in shared:
        if uk_map[k] == ituk_map[k]:
            agrees += 1
        else:
            conflicts.append(
                {"key": list(k), "uk": sorted(uk_map[k]), "ituk": sorted(ituk_map[k])}
            )

    survey_by_pn_brand = defaultdict(set)
    for d in uk + ituk:
        if d["comparability"] != "COMPARABLE":
            continue
        if not d["isuzu_pn"] or not d["peer_oe_pn"]:
            continue
        survey_by_pn_brand[(d["isuzu_pn"], d["peer_brand"])].add(d["peer_oe_pn"])

    mappable = []
    ambiguous = []
    no_hit = 0
    seen = set()
    for br in basket_rows:
        pn = pn_norm(br["isze_pn"])
        peer_oem = br["peer_oem"]
        peer_market = br["peer_market"]
        key = (pn, peer_oem, peer_market)
        if key in seen:
            continue
        seen.add(key)
        cands = set()
        for (spn, sbrand), peers in survey_by_pn_brand.items():
            if spn != pn:
                continue
            if brand_match(sbrand, peer_oem):
                cands |= peers
        if len(cands) == 1:
            mappable.append(
                {
                    "isze_pn": pn,
                    "peer_oem": peer_oem,
                    "peer_market": peer_market,
                    "peer_oe_pn": next(iter(cands)),
                    "status_now": br.get("status"),
                    "current_peer_oe": br.get("peer_oe_pn"),
                }
            )
        elif len(cands) > 1:
            ambiguous.append(
                {
                    "isze_pn": pn,
                    "peer_oem": peer_oem,
                    "peer_market": peer_market,
                    "candidates": sorted(cands),
                }
            )
        else:
            no_hit += 1

    clear_11 = []
    multi = []
    for (pn, brand), peers in survey_by_pn_brand.items():
        if brand not in ("Ford", "Toyota", "DAF"):
            continue
        if len(peers) == 1:
            clear_11.append([pn, brand, next(iter(peers))])
        else:
            multi.append([pn, brand, sorted(peers)])

    iam_ratios = [
        d["workshop_price"] / d["list_price"]
        for d in iam
        if d["list_price"] and d["workshop_price"] and d["list_price"] > 0
    ]

    wb = openpyxl.load_workbook(UK_OEM, data_only=False)
    ws = wb["Isuzu UK"]
    main_isze_types = Counter()
    sample_vals = []
    for i, row in enumerate(
        ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=16, max_col=16), start=2
    ):
        v = row[0].value
        if v is None:
            main_isze_types["None"] += 1
        elif isinstance(v, str) and v.startswith("="):
            main_isze_types["formula"] += 1
            if len(sample_vals) < 5:
                sample_vals.append([i, v])
        else:
            main_isze_types["value"] += 1
            if len(sample_vals) < 8:
                sample_vals.append([i, v])
    # Sheet1 formulas
    ws1 = wb["Sheet1"]
    e2 = ws1["E2"].value
    wb.close()

    implied = []
    for h in helper.values():
        if h["comp_eur_helper"] and h["comp_gbp_helper"] and h["comp_gbp_helper"] > 0:
            implied.append(h["comp_eur_helper"] / h["comp_gbp_helper"])

    iam_by_pn = defaultdict(list)
    for d in iam:
        if d["workshop_price"] is not None:
            iam_by_pn[d["isuzu_pn"]].append(d)

    iam_spot = []
    for pn in sorted(set(gap_map) & set(iam_by_pn))[:8]:
        workshops = [x["workshop_price"] for x in iam_by_pn[pn]]
        lists_ = [x["list_price"] for x in iam_by_pn[pn] if x["list_price"] is not None]
        pr = pr00_map.get(pn)
        iam_spot.append(
            {
                "pn": pn,
                "iam_workshop_avg": statistics.mean(workshops),
                "iam_list_avg": statistics.mean(lists_) if lists_ else None,
                "pr00": pr,
                "street_1785": pr * 1.785 if pr else None,
                "gap_avg_comp": float(gap.loc[gap["pn"] == pn, "Avg_Competitor_EUR"].iloc[0]),
            }
        )

    # Null rates
    integrity = {
        "uk_oem": {
            "rows": len(uk),
            "isuzu_pn_null": sum(1 for d in uk if not d["isuzu_pn"]),
            "peer_oe_null": sum(1 for d in uk if not d["peer_oe_pn"]),
            "isze_main_null": sum(1 for d in uk if d["isze_price_main"] is None),
            "comp_gbp_null": sum(1 for d in uk if d["comp_price_gbp"] is None),
            "isuzu_all_10digit": all(
                len(d["isuzu_pn"]) == 10 and d["isuzu_pn"].isdigit() for d in uk
            ),
            "currency": dict(Counter(d["currency"] for d in uk)),
            "comparability": dict(Counter(d["comparability"] for d in uk)),
            "brands": dict(Counter(d["peer_brand"] for d in uk)),
        },
        "ituk_oem": {
            "rows": len(ituk),
            "isuzu_pn_null": sum(1 for d in ituk if not d["isuzu_pn"]),
            "peer_oe_null": sum(1 for d in ituk if not d["peer_oe_pn"]),
            "isze_main_null": sum(1 for d in ituk if d["isze_price_main"] is None),
            "comp_gbp_null": sum(1 for d in ituk if d["comp_price_gbp"] is None),
            "isuzu_all_10digit": all(
                len(d["isuzu_pn"]) == 10 and d["isuzu_pn"].isdigit() for d in ituk
            ),
            "currency": dict(Counter(d["currency"] for d in ituk)),
            "comparability": dict(Counter(d["comparability"] for d in ituk)),
            "brands": dict(Counter(d["peer_brand"] for d in ituk)),
        },
        "iam": {
            "rows": len(iam),
            "isuzu_null": sum(1 for d in iam if not d["isuzu_pn"]),
            "iam_pn_null": sum(1 for d in iam if not d["iam_pn"]),
            "list_null": sum(1 for d in iam if d["list_price"] is None),
            "workshop_null": sum(1 for d in iam if d["workshop_price"] is None),
            "both_prices": sum(
                1
                for d in iam
                if d["list_price"] is not None and d["workshop_price"] is not None
            ),
            "not_available_comments": sum(
                1 for d in iam if "not available" in d["comment"].lower()
            ),
            "unique_isuzu": len(set(d["isuzu_pn"] for d in iam if d["isuzu_pn"])),
        },
    }

    qa = {
        "files": {
            "UK_OEM": str(UK_OEM),
            "UK_OEM_bytes": UK_OEM.stat().st_size,
            "ITUK_OEM": str(ITUK_OEM),
            "ITUK_OEM_bytes": ITUK_OEM.stat().st_size,
            "UK_IAM": str(UK_IAM),
            "UK_IAM_bytes": UK_IAM.stat().st_size,
        },
        "integrity": integrity,
        "main_isze_col_types": dict(main_isze_types),
        "main_isze_samples": sample_vals,
        "sheet1_E2_formula": e2,
        "helper_fx_median_eur_per_gbp": statistics.median(implied) if implied else None,
        "helper_formula": "Sheet1!E = D/C = Price_in_EUR / ISZE_price (competitiveness ratio). EXCHG rate header present but value cell empty; implied EUR/GBP from D/F.",
        "pricing_point": {
            "oem_CompetitorPrice": "GBP OEM peer list (Currency=GBP on every OEM row)",
            "oem_ISZE_price": "Helper Sheet1 col C filled for all unique PNs; main sheet mostly blank. Appears EUR list-side vs 2026 PR00 (see match_stats) — NOT street x1.785",
            "iam_Competitor_List_Price": "IAM list",
            "iam_Workshop_Purchase_Price": "IAM trade/workshop; median workshop/list=%.3f"
            % (statistics.median(iam_ratios) if iam_ratios else float("nan")),
            "no_1785_factor_in_sheets": True,
        },
        "match_stats_survey_isze_vs_pr00_2026": match_stats,
        "spot_checks": spot[:12],
        "uk_ituk_overlap_pn_count": len(overlap),
        "uk_ituk_overlap_pns": overlap,
        "uk_ituk_shared_brand_keys": len(shared),
        "uk_ituk_agree_peer": agrees,
        "uk_ituk_conflicts": conflicts[:15],
        "dup_uk": dup_analysis(uk),
        "dup_ituk": dup_analysis(ituk),
        "basket_mappable_1to1": len(mappable),
        "basket_ambiguous": len(ambiguous),
        "basket_no_hit": no_hit,
        "basket_mappable_sample": mappable,
        "basket_ambiguous_sample": ambiguous[:15],
        "survey_clear_11_isuzu_brand": len(clear_11),
        "survey_multi_peer_same_brand": len(multi),
        "multi_sample": multi[:10],
        "iam_spot_vs_gap": iam_spot,
    }
    OUT.write_text(json.dumps(fix(qa), indent=2, default=str), encoding="utf-8")
    print("WROTE", OUT)
    print("match_stats", match_stats)
    print(
        "basket 1:1",
        len(mappable),
        "amb",
        len(ambiguous),
        "nohit",
        no_hit,
        "clear11",
        len(clear_11),
        "multi",
        len(multi),
    )
    print("conflicts", len(conflicts), "agrees", agrees, "overlap_pn", len(overlap))


if __name__ == "__main__":
    main()
