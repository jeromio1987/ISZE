"""Ingest Jerome's SAP PR00 pull return into derived long CSV + close-out summary."""
from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
PRICING = Path(__file__).resolve().parent
RAW = PRICING / "_raw" / "SAP_Pull_PN_Validity_List_PR00_return_20260812.xlsx"
DER = ROOT / "01_data" / "00_derived"
OUT_LONG = DER / "pricing_pr00_gap_pull_return_20260812.csv"
OUT_META = DER / "pricing_pr00_gap_pull_return_20260812.meta.json"

WINDOWS = {
    "01/01/2023 - 31/12/2023": ("2023", "2023-01-01", "2023-12-31"),
    "01/01/2024 - 31/01/2025": ("2024", "2024-01-01", "2025-01-31"),
    "01/02/2025 - 17/03/2026": ("2025", "2025-02-01", "2026-03-17"),
}


def norm_pn(x) -> str | None:
    if x is None:
        return None
    s = str(x).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if s.isdigit():
        return s.zfill(10)
    return s


def parse_price(v):
    if v is None:
        return None, "blank"
    if isinstance(v, (int, float)):
        f = float(v)
        if f > 0:
            return f, "ok"
        if f == 0:
            return 0.0, "zero"
        return f, "ok"
    s = str(v).strip().replace(",", ".")
    if s.upper() in ("#N/A", "N/A", "NA", "-", ""):
        return None, "na"
    try:
        f = float(s)
        return (f, "ok") if f > 0 else (0.0, "zero")
    except ValueError:
        return None, f"bad:{s[:40]}"


def main() -> None:
    wb = load_workbook(RAW, data_only=True)
    ws = wb["PN_Validity"]
    rows_out = []
    by_win = defaultdict(Counter)
    pn_ok = set()
    pn_seen = set()

    for pn_raw, win, price in ws.iter_rows(min_row=2, values_only=True):
        pn = norm_pn(pn_raw)
        if not pn:
            continue
        pn_seen.add(pn)
        val, st = parse_price(price)
        label, vf, vt = WINDOWS.get(str(win).strip(), ("unknown", "", ""))
        by_win[str(win)][st] += 1
        if st == "ok":
            pn_ok.add(pn)
        rows_out.append(
            {
                "part_number": pn,
                "window_label": label,
                "valid_from": vf,
                "valid_to": vt,
                "window_text": win,
                "condition": "PR00",
                "price_eur": val if st == "ok" else "",
                "status": st,
                "source_file": RAW.name,
            }
        )

    DER.mkdir(parents=True, exist_ok=True)
    with OUT_LONG.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "part_number",
                "window_label",
                "valid_from",
                "valid_to",
                "window_text",
                "condition",
                "price_eur",
                "status",
                "source_file",
            ],
        )
        w.writeheader()
        w.writerows(rows_out)

    w2025 = "01/02/2025 - 17/03/2026"
    ok_2025 = sum(1 for r in rows_out if r["window_text"] == w2025 and r["status"] == "ok")
    na_2025 = sum(1 for r in rows_out if r["window_text"] == w2025 and r["status"] == "na")
    never = sorted(pn_seen - pn_ok)

    meta = {
        "source": str(RAW),
        "received": "2026-08-12",
        "condition": "PR00",
        "unique_pn": len(pn_seen),
        "pn_with_any_price": len(pn_ok),
        "pn_never_any_price": len(never),
        "pn_never_any_price_list": never,
        "rows": len(rows_out),
        "by_window": {k: dict(v) for k, v in by_win.items()},
        "window_2025_ok": ok_2025,
        "window_2025_na": na_2025,
        "note": (
            "#N/A = no PR00 in SAP for that PN/window — closed, do not re-request. "
            "Ingest usable prices; for remaining gaps use ZPR1 / NO_EVIDENCE hold."
        ),
        "derived_csv": str(OUT_LONG),
    }
    OUT_META.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print(json.dumps({k: meta[k] for k in meta if k != "pn_never_any_price_list"}, indent=2))
    print("never_any_price_sample", never[:15], "count", len(never))
    print("wrote", OUT_LONG)


if __name__ == "__main__":
    main()
