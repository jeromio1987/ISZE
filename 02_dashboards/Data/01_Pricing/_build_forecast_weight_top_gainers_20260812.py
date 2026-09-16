"""PROVISIONAL top share-gainers vs L12M from forecast_weight scaffold. Not an alloc."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
SCAFFOLD = ROOT / "01_data/00_derived/forecast_weight_scaffold_20260812.csv"
OUT_PN = ROOT / "01_data/00_derived/forecast_weight_top_gainers_vs_l12m_20260812.csv"
OUT_FAM = ROOT / "01_data/00_derived/forecast_weight_top_family_gainers_vs_l12m_20260812.csv"
OUT_XLSX = ROOT / "01_data/01_Pricing/forecast_weight_top_gainers_vs_l12m_20260812.xlsx"
META_OUT = ROOT / "01_data/00_derived/forecast_weight_top_gainers_vs_l12m_20260812.meta.json"

TOP_N = 35
TOP_F = 15


def sanitize_xlsx_for_excel_desktop(path: Path) -> None:
    """Patch openpyxl OOXML quirks that make Excel Desktop prompt Repair / hang."""
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    wb_xml = files["xl/workbook.xml"].decode("utf-8")
    wb_xml = re.sub(r"<workbookProtection\b[^/]*/>", "", wb_xml)
    wb_xml = re.sub(
        r"<workbookProtection\b[^>]*>.*?</workbookProtection>",
        "",
        wb_xml,
        flags=re.DOTALL,
    )
    files["xl/workbook.xml"] = wb_xml.encode("utf-8")
    styles = files["xl/styles.xml"].decode("utf-8")
    styles = styles.replace("<patternFill/>", '<patternFill patternType="none"/>')
    styles = styles.replace("<patternFill />", '<patternFill patternType="none"/>')
    styles = re.sub(r'rgb="00([0-9A-Fa-f]{6})"', r'rgb="FF\1"', styles)
    styles = re.sub(r'\s+pivotButton="0"', "", styles)
    styles = re.sub(r'\s+quotePrefix="0"', "", styles)
    files["xl/styles.xml"] = styles.encode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rels = rels.replace('Target="/xl/worksheets/', 'Target="worksheets/')
    rels = rels.replace('Target="/xl/theme/', 'Target="theme/')
    rels = rels.replace('Target="/xl/styles.xml"', 'Target="styles.xml"')
    rels = rels.replace('Target="/xl/sharedStrings.xml"', 'Target="sharedStrings.xml"')
    files["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    tmp = path.with_suffix(".sanitize.tmp.xlsx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


def flag_str(r: pd.Series) -> str:
    parts = []
    if r["flag_uti"]:
        parts.append("UTI_recommended")
    if r["flag_rg21"]:
        parts.append("RG21")
    if r["flag_p700"]:
        parts.append("P700")
    return "|".join(parts)


def write_sheet(wb: Workbook, name: str, frame: pd.DataFrame, widths: list[int] | None = None):
    ws = wb.create_sheet(name)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cols = list(frame.columns)
    for c, col in enumerate(cols, 1):
        cell = ws.cell(1, c, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    for r_i, row in enumerate(frame.itertuples(index=False), 2):
        for c_i, val in enumerate(row, 1):
            if isinstance(val, bool):
                val = int(val)
            if pd.isna(val):
                val = None
            cell = ws.cell(r_i, c_i, val)
            cell.border = thin
            if isinstance(val, float):
                colname = cols[c_i - 1]
                if "pct" in colname or "pp" in colname or "%" in colname:
                    cell.number_format = "0.000"
                elif "factor" in colname:
                    cell.number_format = "0.0000"
                elif "eur" in colname:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "0.0000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(frame) + 1}"
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    else:
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14
    return ws


def main() -> None:
    df = pd.read_csv(SCAFFOLD)
    df["pn"] = df["pn"].astype(str).str.replace(r"\.0$", "", regex=True)

    base = df[df["l12m_eur"].fillna(0) > 0].copy()
    sum_l12m = float(base["l12m_eur"].sum())
    sum_fw = float(base["forecast_weight_eur"].sum())
    base["share_l12m"] = base["l12m_eur"] / sum_l12m
    base["share_fw"] = base["forecast_weight_eur"] / sum_fw
    base["delta_share_pp"] = (base["share_fw"] - base["share_l12m"]) * 100.0

    flags = base["flags"].fillna("")
    base["flag_uti"] = (base["on_uti_p700_recommended"].fillna(0).astype(int) == 1) | flags.str.contains(
        "on_uti_p700_recommended=1", na=False
    )
    base["flag_rg21"] = flags.str.contains("in_rg21=1", na=False) | base["packs"].fillna("").str.contains(
        "RG21", case=False, na=False
    )
    base["flag_p700"] = flags.str.contains("in_p700=1", na=False) | base["packs"].fillna("").str.contains(
        "P700", case=False, na=False
    )
    base["flag_labels"] = base.apply(flag_str, axis=1)

    gainers = base[base["delta_share_pp"] > 0].sort_values("delta_share_pp", ascending=False)
    pn_top = gainers.head(TOP_N).copy()
    pn_top.insert(0, "rank", range(1, len(pn_top) + 1))

    pn_out = pn_top[
        [
            "rank",
            "pn",
            "description",
            "family_hint",
            "packs",
            "book_b_bucket",
            "growth_factor",
            "l12m_eur",
            "forecast_weight_eur",
            "share_l12m",
            "share_fw",
            "delta_share_pp",
            "flag_labels",
            "flag_uti",
            "flag_rg21",
            "flag_p700",
            "growth_rule_note",
        ]
    ].copy()
    pn_out["share_l12m_pct"] = pn_out["share_l12m"] * 100
    pn_out["share_fw_pct"] = pn_out["share_fw"] * 100
    pn_out = pn_out.drop(columns=["share_l12m", "share_fw"])
    pn_out = pn_out[
        [
            "rank",
            "pn",
            "description",
            "family_hint",
            "packs",
            "book_b_bucket",
            "growth_factor",
            "l12m_eur",
            "forecast_weight_eur",
            "share_l12m_pct",
            "share_fw_pct",
            "delta_share_pp",
            "flag_labels",
            "flag_uti",
            "flag_rg21",
            "flag_p700",
            "growth_rule_note",
        ]
    ]

    fam = base.copy()
    fam["family_key"] = fam["family_hint"].fillna("(unclassified)").astype(str).str.strip()
    fam["family_key"] = fam["family_key"].replace({"": "(unclassified)"})

    fam_agg = fam.groupby("family_key", as_index=False).agg(
        n_pn=("pn", "nunique"),
        l12m_eur=("l12m_eur", "sum"),
        forecast_weight_eur=("forecast_weight_eur", "sum"),
        mean_growth_factor_w=(
            "growth_factor",
            lambda s: (s * fam.loc[s.index, "l12m_eur"]).sum() / fam.loc[s.index, "l12m_eur"].sum(),
        ),
        n_uti=("flag_uti", "sum"),
        n_rg21=("flag_rg21", "sum"),
        n_p700=("flag_p700", "sum"),
    )
    fam_agg["share_l12m"] = fam_agg["l12m_eur"] / sum_l12m
    fam_agg["share_fw"] = fam_agg["forecast_weight_eur"] / sum_fw
    fam_agg["delta_share_pp"] = (fam_agg["share_fw"] - fam_agg["share_l12m"]) * 100.0
    fam_gain = fam_agg[fam_agg["delta_share_pp"] > 0].sort_values("delta_share_pp", ascending=False)
    fam_top = fam_gain.head(TOP_F).copy()
    fam_top.insert(0, "rank", range(1, len(fam_top) + 1))
    fam_out = fam_top.copy()
    fam_out["share_l12m_pct"] = fam_out["share_l12m"] * 100
    fam_out["share_fw_pct"] = fam_out["share_fw"] * 100
    fam_out = fam_out[
        [
            "rank",
            "family_key",
            "n_pn",
            "mean_growth_factor_w",
            "l12m_eur",
            "forecast_weight_eur",
            "share_l12m_pct",
            "share_fw_pct",
            "delta_share_pp",
            "n_uti",
            "n_rg21",
            "n_p700",
        ]
    ].rename(columns={"family_key": "family_hint"})
    fam_out.insert(1, "group_type", "family_hint")

    packs_rows = []
    for tag, mask in [
        ("pack:P700", base["flag_p700"]),
        ("pack:RG21", base["flag_rg21"]),
        ("pack:UTI_recommended", base["flag_uti"]),
        ("pack:ITUK", base["packs"].fillna("").str.contains("ITUK", case=False, na=False)),
    ]:
        sub = base[mask]
        if len(sub) == 0:
            continue
        l = float(sub["l12m_eur"].sum())
        f = float(sub["forecast_weight_eur"].sum())
        packs_rows.append(
            {
                "family_hint": tag,
                "n_pn": int(sub["pn"].nunique()),
                "mean_growth_factor_w": float((sub["growth_factor"] * sub["l12m_eur"]).sum() / l) if l else None,
                "l12m_eur": l,
                "forecast_weight_eur": f,
                "share_l12m_pct": 100 * l / sum_l12m,
                "share_fw_pct": 100 * f / sum_fw,
                "delta_share_pp": 100 * (f / sum_fw - l / sum_l12m),
                "n_uti": int(sub["flag_uti"].sum()),
                "n_rg21": int(sub["flag_rg21"].sum()),
                "n_p700": int(sub["flag_p700"].sum()),
            }
        )
    pack_df = pd.DataFrame(packs_rows).sort_values("delta_share_pp", ascending=False)
    pack_df.insert(0, "rank", range(1, len(pack_df) + 1))

    need_seed = df[df["flags"].fillna("").str.contains("NEED_SEED", na=False)]
    need_seed_n = int(len(need_seed))
    need_seed_uti = int(need_seed["on_uti_p700_recommended"].fillna(0).astype(int).sum())

    pn_out.to_csv(OUT_PN, index=False, float_format="%.6f")
    fam_out.to_csv(OUT_FAM, index=False, float_format="%.6f")

    wb = Workbook()
    ws_note = wb.active
    ws_note.title = "Note"
    ws_note["A1"] = (
        "PROVISIONAL — top share-gainers vs L12M from forecast_weight_scaffold_20260812 "
        "(growth-factor diagnostics only; NOT a locked +4/+5 alloc)."
    )
    ws_note["A1"].font = Font(bold=True, color="9C0006")
    ws_note["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_note.merge_cells("A1:G3")
    ws_note["A5"] = "Universe"
    ws_note["B5"] = f"{len(base):,} PNs with L12M EUR > 0"
    ws_note["A6"] = "sum L12M EUR"
    ws_note["B6"] = round(sum_l12m, 2)
    ws_note["A7"] = "sum forecast_weight EUR"
    ws_note["B7"] = round(sum_fw, 2)
    ws_note["A8"] = "NEED_SEED PNs (no L12M)"
    ws_note["B8"] = need_seed_n
    ws_note["C8"] = (
        f"of which UTI recommended: {need_seed_uti} — excluded from share-gain ranking (no L12M base)"
    )
    ws_note["A10"] = "Share is zero-sum: gainers rise only because others fall (many at x1.0)."
    ws_note.column_dimensions["A"].width = 32
    ws_note.column_dimensions["B"].width = 22
    ws_note.column_dimensions["C"].width = 70
    ws_note.row_dimensions[1].height = 48

    pn_xlsx = pn_out.rename(
        columns={
            "flag_labels": "flags",
            "share_l12m_pct": "share_l12m_%",
            "share_fw_pct": "share_fw_%",
            "delta_share_pp": "delta_pp",
        }
    ).drop(columns=["flag_uti", "flag_rg21", "flag_p700"])
    fam_xlsx = fam_out.drop(columns=["group_type"]).rename(
        columns={
            "share_l12m_pct": "share_l12m_%",
            "share_fw_pct": "share_fw_%",
            "delta_share_pp": "delta_pp",
        }
    )
    pack_xlsx = pack_df.rename(
        columns={
            "share_l12m_pct": "share_l12m_%",
            "share_fw_pct": "share_fw_%",
            "delta_share_pp": "delta_pp",
        }
    )

    write_sheet(
        wb,
        "PN_gainers",
        pn_xlsx,
        widths=[6, 12, 28, 14, 16, 22, 10, 12, 14, 12, 12, 10, 22, 18],
    )
    write_sheet(
        wb,
        "Family_gainers",
        fam_xlsx,
        widths=[6, 16, 8, 12, 12, 14, 12, 12, 10, 8, 8, 8],
    )
    write_sheet(
        wb,
        "PackTag_share",
        pack_xlsx,
        widths=[6, 22, 8, 12, 12, 14, 12, 12, 10, 8, 8, 8],
    )

    wb.save(OUT_XLSX)
    sanitize_xlsx_for_excel_desktop(OUT_XLSX)

    meta = {
        "built_utc": pd.Timestamp.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "status": "PROVISIONAL_share_gainers_vs_l12m",
        "note": "Not a locked +4/+5 alloc; diagnostic growth-factor weight shares only.",
        "inputs": {"scaffold": str(SCAFFOLD)},
        "outputs": {"pn_csv": str(OUT_PN), "family_csv": str(OUT_FAM), "xlsx": str(OUT_XLSX)},
        "universe": {
            "n_pn_l12m_gt0": int(len(base)),
            "sum_l12m_eur": sum_l12m,
            "sum_forecast_weight_eur": sum_fw,
            "n_pn_gainers": int((base["delta_share_pp"] > 0).sum()),
            "n_pn_losers": int((base["delta_share_pp"] < 0).sum()),
            "n_need_seed_excluded": need_seed_n,
            "n_need_seed_uti_rec": need_seed_uti,
            "sum_delta_pp_all_pn": float(base["delta_share_pp"].sum()),
        },
        "top_n_pn": TOP_N,
        "top_n_family": TOP_F,
        "pack_tag_delta_pp": {r["family_hint"]: round(r["delta_share_pp"], 4) for r in packs_rows},
    }
    META_OUT.write_text(json.dumps(meta, indent=2), encoding="utf-8")

    print("=== FILES ===")
    print(OUT_PN)
    print(OUT_FAM)
    print(OUT_XLSX)
    print(META_OUT)
    print("\n=== TOP 10 PN ===")
    for r in pn_out.head(10).itertuples(index=False):
        print(
            f"{r.rank:2d}. {r.pn} | {r.description} | fam={r.family_hint} | g={r.growth_factor} | "
            f"L12M={r.l12m_eur:,.0f} FW={r.forecast_weight_eur:,.0f} | "
            f"share {r.share_l12m_pct:.3f}%->{r.share_fw_pct:.3f}% | "
            f"d={r.delta_share_pp:+.3f}pp | {r.flag_labels} | {r.book_b_bucket}"
        )
    print("\n=== TOP 5 FAMILIES ===")
    for r in fam_out.head(5).itertuples(index=False):
        print(
            f"{r.rank:2d}. {r.family_hint} | n={r.n_pn} | g~{r.mean_growth_factor_w:.3f} | "
            f"L12M={r.l12m_eur:,.0f} FW={r.forecast_weight_eur:,.0f} | "
            f"{r.share_l12m_pct:.2f}%->{r.share_fw_pct:.2f}% | d={r.delta_share_pp:+.3f}pp | "
            f"UTI={r.n_uti} RG21={r.n_rg21} P700={r.n_p700}"
        )
    print("\n=== PACK TAGS ===")
    print(pack_df.to_string(index=False))
    print("\n=== NEED_SEED ===")
    print(
        f"NEED_SEED={need_seed_n} (UTI rec {need_seed_uti}) — no L12M so absent from share-gain ranking"
    )
    print(
        f"gainers={meta['universe']['n_pn_gainers']} losers={meta['universe']['n_pn_losers']} "
        f"sum_delta_pp={meta['universe']['sum_delta_pp_all_pn']:.6e}"
    )


if __name__ == "__main__":
    main()
