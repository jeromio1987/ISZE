"""Compare Duncan 18 Aug PMS pack vs ISZE Model B / sales_fact / PN audit."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
PACK = ROOT / r"01_data\90_INBOX\IML_ITUK_Meeting_20260813"
OUT = PACK / "_pms_anomaly_check.json"

DUNCAN_XLSX = PACK / "N75_N35_Part_Numbers_-_PMS_Project.xlsx"
AUDIT = ROOT / r"01_data\00_derived\ituk_pms_pn_audit.csv"
SALES = ROOT / r"01_data\00_derived\sales_fact.csv"
PEEK = ROOT / r"09_meetings\Preparation Meetings\IML_ITUK_PMS\peek_ISZE___ITUK__PN_Level_.json"

MODEL_SHEETS = [
    "NLR87   N35.125   RZ4E",
    "NLR88  N35.150  4JZ1",
    "NPR75  N75  4HK1",
    "NPR88  N75  4JZ1",
]


def load_duncan() -> pd.DataFrame:
    wb = openpyxl.load_workbook(DUNCAN_XLSX, read_only=True, data_only=True)
    rows = []
    for sheet in MODEL_SHEETS:
        ws = wb[sheet]
        header_row = None
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [v for v in row]
            joined = " ".join(str(x) for x in vals if x)
            if "PMS category" in joined and "Part" in joined:
                header_row = i
                break
        ws2 = wb[sheet]
        started = False
        for i, row in enumerate(ws2.iter_rows(values_only=True), 1):
            vals = list(row)
            if i < (header_row or 2):
                continue
            if i == header_row:
                started = True
                continue
            cat = vals[1]
            pn = vals[2]
            desc = vals[3]
            prio = vals[4]
            monthly = vals[5]
            ytd = vals[6]
            comments = vals[7] if len(vals) > 7 else None
            if pn is None:
                continue
            pn_s = str(pn).strip()
            if not pn_s.isdigit() and not pn_s.replace(".", "").isdigit():
                continue
            rows.append(
                {
                    "model_sheet": sheet.strip(),
                    "category": None if cat is None else str(cat).strip(),
                    "pn": pn_s.split(".")[0],
                    "desc": None if desc is None else str(desc).strip(),
                    "priority": None if prio is None else str(prio).strip(),
                    "monthly": float(monthly) if monthly not in (None, "") else None,
                    "ytd": float(ytd) if ytd not in (None, "") else None,
                    "comments": None if comments is None else str(comments).strip(),
                }
            )
    wb.close()
    return pd.DataFrame(rows)


def peek_ws2025() -> dict[str, dict]:
    data = json.loads(PEEK.read_text(encoding="utf-8"))
    out = {}
    last_cat = None
    for _, row in data["rows"]:
        if not row or row[0] in (None, "Category"):
            if row and row[0] in (
                "OIL FILTER",
                "FUEL FILTER",
                "AIR FILTER",
                "CLUTCH",
                "BRAKE",
            ):
                last_cat = row[0]
                pn = str(row[1]).strip() if row[1] else None
            else:
                continue
        else:
            pn = str(row[1]).strip() if row[1] else None
            if row[0] in ("OIL FILTER", "FUEL FILTER", "AIR FILTER", "CLUTCH", "BRAKE"):
                last_cat = row[0]
        if not pn or not pn[0].isdigit():
            continue
        out[pn] = {
            "category": last_cat,
            "desc": row[2],
            "model_ituk": row[3],
            "model_iml": row[4],
            "dvla_uio": row[5],
            "freq": row[6],
            "expected": row[7],
            "ws2025": row[8],
            "cover": row[9],
            "pr00": row[10],
        }
    return out


def main() -> None:
    d = load_duncan()
    d["pn"] = d["pn"].astype(str)
    print(f"duncan rows={len(d)} unique_pn={d['pn'].nunique()}", flush=True)

    # double count
    counts = d.groupby("pn").size().reset_index(name="n_models")
    multi = counts[counts["n_models"] > 1].merge(d, on="pn")
    dup_sum_if_added = (
        d.groupby("pn")
        .agg(ytd_sum=("ytd", "sum"), ytd_first=("ytd", "first"), n=("pn", "size"))
        .reset_index()
    )
    dup_sum_if_added = dup_sum_if_added[dup_sum_if_added["n"] > 1]
    inflated = float((dup_sum_if_added["ytd_sum"] - dup_sum_if_added["ytd_first"]).sum())

    peek = peek_ws2025()
    audit = pd.read_csv(AUDIT, dtype={"part_number": str})
    audit["part_number"] = audit["part_number"].astype(str).str.replace(r"\.0$", "", regex=True)
    a25 = audit[audit["year"] == 2025].copy()
    a25_map = {
        r.part_number: r
        for r in a25.itertuples(index=False)
    }

    pns = sorted(d["pn"].unique())
    print("loading sales_fact ITUK 2575…", flush=True)
    usecols = ["billing_date", "sold_to", "part_number", "billed_qty", "net_sales_value", "year"]
    sf = pd.read_csv(SALES, usecols=usecols, dtype={"sold_to": str, "part_number": str})
    sf["part_number"] = sf["part_number"].astype(str).str.replace(r"\.0$", "", regex=True)
    ituk = sf[sf["sold_to"].astype(str).str.contains("2575")].copy()
    print(f"ITUK rows={len(ituk)} sold_to sample={ituk['sold_to'].unique()[:8].tolist()}", flush=True)

    def agg_year(frame, year):
        sub = frame[frame["year"] == year]
        g = sub.groupby("part_number", as_index=False).agg(qty=("billed_qty", "sum"), eur=("net_sales_value", "sum"))
        return {r.part_number: {"qty": float(r.qty), "eur": float(r.eur)} for r in g.itertuples(index=False)}

    s25 = agg_year(ituk, 2025)
    s26 = agg_year(ituk, 2026)
    s24 = agg_year(ituk, 2024)

    # also all sold_to in case 2575 encoding differs
    if not s25:
        print("NO 2575 MATCH — unique sold_to", ituk["sold_to"].unique()[:20].tolist() if len(ituk) else sf["sold_to"].drop_duplicates().head(20).tolist(), flush=True)
        # try contains ITUK name? sold_to is code
        codes = sf["sold_to"].dropna().astype(str).value_counts().head(30)
        print(codes.to_string(), flush=True)

    rows = []
    for pn in pns:
        sub = d[d["pn"] == pn]
        ytd = float(sub["ytd"].iloc[0] or 0)
        monthly = float(sub["monthly"].iloc[0] or 0)
        models = sub["model_sheet"].tolist()
        cat = sub["category"].iloc[0]
        prio = sub["priority"].iloc[0]
        comments = "; ".join(x for x in sub["comments"].dropna().unique())
        pk = peek.get(pn, {})
        ws = pk.get("ws2025")
        try:
            ws_f = float(str(ws).replace(",", "")) if ws not in (None, "") else None
        except Exception:
            ws_f = None
        isze25 = s25.get(pn, {}).get("qty")
        isze26 = s26.get(pn, {}).get("qty")
        isze24 = s24.get(pn, {}).get("qty")
        au = a25[a25["part_number"] == pn]
        audit_qty = float(au["actual_qty"].sum()) if len(au) else None
        ann_from_monthly = monthly * 12 if monthly else None
        ytd_vs_ws = (ytd / ws_f) if ws_f else None
        monthly_ann_vs_ws = (ann_from_monthly / ws_f) if ws_f and ann_from_monthly else None
        ytd_vs_isze25 = (ytd / isze25) if isze25 else None
        rows.append(
            {
                "pn": pn,
                "category": cat,
                "priority": prio,
                "n_model_slides": int(len(models)),
                "models": models,
                "duncan_monthly": monthly,
                "duncan_ytd": ytd,
                "duncan_annualised_monthly": ann_from_monthly,
                "comments": comments,
                "our_ws2025": ws_f,
                "our_model_iml": pk.get("model_iml"),
                "our_cover_pn": pk.get("cover"),
                "isze_ituk_2024_qty": isze24,
                "isze_ituk_2025_qty": isze25,
                "isze_ituk_2026_qty": isze26,
                "audit_2025_qty": audit_qty,
                "ytd_vs_ws2025": ytd_vs_ws,
                "ann_monthly_vs_ws2025": monthly_ann_vs_ws,
                "ytd_vs_isze2025": ytd_vs_isze25,
                "zero_ytd": ytd == 0,
                "in_our_pn_sheet": pn in peek,
            }
        )

    # missing from Duncan that are big in our oil/fuel/air/brake
    duncan_set = set(pns)
    missing_big = []
    for pn, info in peek.items():
        try:
            ws_f = float(str(info.get("ws2025") or 0).replace(",", ""))
        except Exception:
            ws_f = 0
        if pn not in duncan_set and ws_f >= 200:
            missing_big.append({"pn": pn, "cat": info.get("category"), "ws2025": ws_f, "model": info.get("model_iml"), "desc": info.get("desc")})
    missing_big.sort(key=lambda x: -x["ws2025"])

    flagged_bvp = ["1876102860", "5876102500", "8983104491", "8983752250", "8982447940"]
    bvp_check = [r for r in rows if r["pn"] in flagged_bvp]

    summary = {
        "duncan_rows": int(len(d)),
        "unique_pns": int(d["pn"].nunique()),
        "duplicated_pns": int((counts["n_models"] > 1).sum()),
        "ytd_unique_sum": float(d.drop_duplicates("pn")["ytd"].sum()),
        "ytd_naive_sum_all_slides": float(d["ytd"].sum()),
        "ytd_inflation_if_summed": inflated,
        "ytd_over_ws_median": None,
        "ann_monthly_over_ws_median": None,
        "ituk_sales_matched_2025": int(sum(1 for r in rows if r["isze_ituk_2025_qty"])),
        "not_in_our_pn_sheet": [r["pn"] for r in rows if not r["in_our_pn_sheet"]],
        "zero_ytd": [r["pn"] for r in rows if r["zero_ytd"]],
        "missing_from_duncan_ws2025_ge_200": missing_big,
        "bvp_slide20": bvp_check,
        "pn_rows": rows,
        "duplicates": multi[["pn", "n_models", "model_sheet", "ytd", "category"]].to_dict("records"),
        "ituk_sold_to_unique": ituk["sold_to"].dropna().astype(str).unique().tolist()[:20] if len(ituk) else [],
        "sales_sold_to_top": sf["sold_to"].astype(str).value_counts().head(15).to_dict(),
    }
    ratios = [r["ytd_vs_ws2025"] for r in rows if r["ytd_vs_ws2025"] and r["duncan_ytd"]]
    ratios2 = [r["ann_monthly_vs_ws2025"] for r in rows if r["ann_monthly_vs_ws2025"] and r["duncan_ytd"]]
    if ratios:
        summary["ytd_over_ws_median"] = float(pd.Series(ratios).median())
        summary["ytd_over_ws_mean"] = float(pd.Series(ratios).mean())
    if ratios2:
        summary["ann_monthly_over_ws_median"] = float(pd.Series(ratios2).median())

    OUT.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print(f"wrote {OUT}", flush=True)
    print("unique", summary["unique_pns"], "dups", summary["duplicated_pns"], flush=True)
    print("ytd unique", round(summary["ytd_unique_sum"]), "naive", round(summary["ytd_naive_sum_all_slides"]), "inflate", round(summary["ytd_inflation_if_summed"]), flush=True)
    print("median ytd/ws", summary["ytd_over_ws_median"], "median ann/ws", summary["ann_monthly_over_ws_median"], flush=True)
    print("not in our sheet", summary["not_in_our_pn_sheet"], flush=True)
    print("missing big", missing_big[:12], flush=True)


if __name__ == "__main__":
    main()
