"""Extract PMS pptx + xlsx structure for anomaly check vs ISZE Model B."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\IML_ITUK_Meeting_20260813")
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def pptx_slides(path: Path) -> list[dict]:
    slides = []
    with zipfile.ZipFile(path) as z:
        names = sorted(
            n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")
        )
        for name in names:
            xml = z.read(name)
            root = ET.fromstring(xml)
            texts = [t.text.strip() for t in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t") if t.text and t.text.strip()]
            slides.append({"file": name, "texts": texts, "joined": " | ".join(texts)})
    return slides


def dump_pptx(name: str) -> None:
    path = ROOT / name
    slides = pptx_slides(path)
    out = ROOT / (path.stem + "_slides.json")
    out.write_text(json.dumps(slides, indent=2, ensure_ascii=False), encoding="utf-8")
    md = ROOT / (path.stem + "_slides.md")
    lines = [f"# {name}", ""]
    for i, s in enumerate(slides, 1):
        lines.append(f"## Slide {i}")
        lines.append("")
        for t in s["texts"]:
            lines.append(f"- {t}")
        lines.append("")
    md.write_text("\n".join(lines), encoding="utf-8")
    print(f"{name}: {len(slides)} slides -> {md.name}", flush=True)


def dump_xlsx(name: str) -> None:
    path = ROOT / name
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summary = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > 8:
                break
            rows.append([None if v is None else str(v)[:80] for v in row[:20]])
        # count
        n = 0
        headers = None
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i == 1:
                headers = [None if v is None else str(v) for v in row]
            n += 1
            if n >= 200000:
                break
        summary.append(
            {
                "sheet": ws.title,
                "max_row_sampled": n,
                "ncols": len(headers or []),
                "headers": headers,
                "first_rows": rows,
            }
        )
        print(f"  sheet={ws.title!r} rows~{n} cols={len(headers or [])}", flush=True)
    wb.close()
    out = ROOT / (path.stem + "_sheets.json")
    out.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"{name} -> {out.name}", flush=True)


def main() -> None:
    dump_pptx("PMS_Project_18_08_2026.pptx")
    dump_pptx("Lubrication_Project.pptx")
    dump_xlsx("N75_N35_Part_Numbers_-_PMS_Project.xlsx")
    dump_xlsx("Questionnaire_for_Lubs_centralize_purchasing.xlsx")


if __name__ == "__main__":
    main()
