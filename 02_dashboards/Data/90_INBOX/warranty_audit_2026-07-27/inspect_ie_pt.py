"""Deep dive Sheet4 structure + IE/PT price coverage from detail data."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

CV = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\06_Warranty\Warranty_CV_2025.xlsx")
OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\warranty_audit_2026-07-27")


def parse_price(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "#N/A", "#NA", "N/A"):
        return None
    s = s.replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def is_country_or_bucket(lbl) -> bool:
    if lbl is None:
        return False
    s = str(lbl).strip()
    if s.lower().startswith("grand"):
        return True
    if s.lower() in ("row labels",):
        return True
    # part numbers are long digits; "0" and country names are not
    if re.fullmatch(r"\d{8,}", s):
        return False
    if re.fullmatch(r"0+", s) or s == "0":
        return True  # ambiguous — treat as bucket for now
    # alphanumeric country codes / names
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9 _\-]{0,30}", s):
        return True
    return False


def sheet4_sections():
    wb = openpyxl.load_workbook(CV, data_only=True)
    ws = wb["Sheet4"]
    sections = []
    cur = None
    for r in range(1, ws.max_row + 1):
        a, b, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a is None:
            continue
        if is_country_or_bucket(a) and not re.fullmatch(r"\d{8,}", str(a).strip()):
            # also treat short numeric 0 as bucket
            s = str(a).strip()
            if re.fullmatch(r"\d{8,}", s):
                pass
            else:
                cur = {
                    "row": r,
                    "label": s,
                    "count": b,
                    "cell_c": c,
                    "parts": [],
                    "sum_qty_x_price": 0.0,
                    "priced_parts": 0,
                    "unpriced_parts": 0,
                    "na_price": 0,
                }
                sections.append(cur)
                continue
        if cur is None:
            continue
        price = parse_price(c)
        try:
            qty = float(b) if b is not None else 0.0
        except (TypeError, ValueError):
            qty = 0.0
        part = {"row": r, "part": a, "qty": qty, "price_raw": c, "price": price}
        cur["parts"].append(part)
        if price is None:
            cur["unpriced_parts"] += 1
            if c is not None and "#N/A" in str(c).upper():
                cur["na_price"] += 1
        else:
            cur["priced_parts"] += 1
            cur["sum_qty_x_price"] += qty * price

    # Also Sheet3
    ws3 = wb["Sheet3"]
    sheet3 = []
    for r in range(1, ws3.max_row + 1):
        vals = [ws3.cell(r, c).value for c in range(1, 5)]
        if any(v is not None for v in vals):
            sheet3.append({"row": r, "vals": vals})

    summary = []
    for s in sections:
        summary.append(
            {
                "row": s["row"],
                "label": s["label"],
                "count": s["count"],
                "cell_c": s["cell_c"],
                "n_parts": len(s["parts"]),
                "priced": s["priced_parts"],
                "unpriced": s["unpriced_parts"],
                "na": s["na_price"],
                "sum_eur": round(s["sum_qty_x_price"], 2),
                "sample_parts": s["parts"][:5],
                "sample_unpriced": [p for p in s["parts"] if p["price"] is None][:8],
            }
        )

    print("=== Sheet4 sections ===")
    for s in summary:
        print(
            f"r{s['row']:4d} [{s['label']:12s}] count={s['count']} parts={s['n_parts']} "
            f"priced={s['priced']} unpriced={s['unpriced']} sum€={s['sum_eur']:,.2f} cell_c={s['cell_c']}"
        )

    print("\n=== Sheet3 (all non-empty) ===")
    for row in sheet3[:80]:
        print(row)

    # Zero buckets detail
    zeros = [s for s in sections if str(s["label"]).strip() in ("0", "0.0")]
    print(f"\n=== '0' buckets: {len(zeros)} ===")
    for z in zeros:
        print(f"  row {z['row']}: count={z['count']} parts={len(z['parts'])} sum={z['sum_qty_x_price']:.2f}")
        print(f"    first 8 parts: {z['parts'][:8]}")
        print(f"    last 3 parts: {z['parts'][-3:]}")

    # Ireland / Portugal detail
    for name in ("IRELAND", "PORTUGAL"):
        sec = next((s for s in sections if s["label"].upper() == name), None)
        print(f"\n=== {name} ===")
        if not sec:
            print("  NOT FOUND as section")
            continue
        print(f"  header count={sec['count']} cell_c={sec['cell_c']} direct_parts={len(sec['parts'])} sum={sec['sum_qty_x_price']:.2f}")
        for p in sec["parts"][:20]:
            print(f"    {p}")

    out = {
        "sections": summary,
        "zero_total_eur": round(sum(s["sum_qty_x_price"] for s in zeros), 2),
        "sheet3": sheet3[:100],
    }
    (OUT / "_sheet4_sections.json").write_text(json.dumps(out, indent=2, default=str), encoding="utf-8")
    wb.close()
    return sections


def ie_pt_from_detail():
    """Find Ireland/Portugal in Sheet1 / DATA WITHOUT MIDI and price coverage vs ITUK/UTI maps."""
    wb = openpyxl.load_workbook(CV, data_only=True)
    # Build price maps from ITUK + UTI tabs
    price_map = {}
    for tab, part_col, price_col in (("ITUK", 1, 4), ("UTI", 1, 4)):
        ws = wb[tab]
        for r in range(2, ws.max_row + 1):
            part = ws.cell(r, part_col).value
            price = ws.cell(r, price_col).value
            if part is None:
                continue
            key = str(part).strip().lstrip("0") or "0"
            if isinstance(price, (int, float)) and price > 0:
                price_map[key] = float(price)
            # also keep with leading zeros
            price_map[str(part).strip()] = float(price) if isinstance(price, (int, float)) else None

    # Also store clean map
    clean_map = {}
    for tab in ("ITUK", "UTI"):
        ws = wb[tab]
        for r in range(2, ws.max_row + 1):
            part = ws.cell(r, 1).value
            price = ws.cell(r, 4).value
            if part is None:
                continue
            if isinstance(price, (int, float)) and price > 0:
                clean_map[str(int(part)) if isinstance(part, float) else str(part).strip()] = float(price)
                # normalized without leading zeros
                clean_map[str(part).strip().lstrip("0") or "0"] = float(price)

    print(f"\nPrice map size (ITUK+UTI): {len(clean_map)}")

    results = {}
    for sheet_name in ("Sheet1", "DATA WITHOUT MIDI"):
        ws = wb[sheet_name]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        # find cols
        def col(name):
            for i, h in enumerate(headers, 1):
                if h and str(h).strip().upper() == name.upper():
                    return i
            return None

        cty_i = col("CTY_Name") or col("CTY")
        part_i = col("PARTNBR")
        qty_i = col("QTY1")
        billy_i = col("BILLYMD")
        dst_i = col("DST_Name")
        print(f"\n{sheet_name}: CTY={cty_i} PART={part_i} QTY={qty_i} BILLY={billy_i}")

        for country in ("IRELAND", "PORTUGAL", "UK", "ITALY", "ISRAEL"):
            cells_total = 0
            cells_priced = 0
            qty_total = 0.0
            qty_priced = 0.0
            value_priced = 0.0
            samples_missing = []
            samples_priced = []
            cy2025 = 0
            for r in range(2, ws.max_row + 1):
                cty = ws.cell(r, cty_i).value if cty_i else None
                if cty is None or str(cty).strip().upper() != country:
                    continue
                billy = ws.cell(r, billy_i).value if billy_i else None
                if billy is not None:
                    bs = str(int(billy)) if isinstance(billy, float) else str(billy)
                    if not bs.startswith("2025"):
                        continue
                cy2025 += 1
                part = ws.cell(r, part_i).value if part_i else None
                qty = ws.cell(r, qty_i).value if qty_i else None
                try:
                    q = float(qty) if qty is not None else 0.0
                except (TypeError, ValueError):
                    q = 0.0
                if part is None or str(part).strip() in ("", "0000000000", "0"):
                    continue
                cells_total += 1
                qty_total += q
                pk = str(part).strip()
                pk2 = pk.lstrip("0") or "0"
                # try int form
                try:
                    pk3 = str(int(pk))
                except ValueError:
                    pk3 = pk2
                price = clean_map.get(pk) or clean_map.get(pk2) or clean_map.get(pk3)
                if price:
                    cells_priced += 1
                    qty_priced += q
                    value_priced += q * price
                    if len(samples_priced) < 5:
                        samples_priced.append(
                            {
                                "row": r,
                                "part": pk,
                                "qty": q,
                                "price": price,
                                "dst": ws.cell(r, dst_i).value if dst_i else None,
                            }
                        )
                else:
                    if len(samples_missing) < 12:
                        samples_missing.append(
                            {
                                "row": r,
                                "part": pk,
                                "qty": q,
                                "dst": ws.cell(r, dst_i).value if dst_i else None,
                                "billy": billy,
                                "partamt": ws.cell(r, col("PARTAMT1")).value if col("PARTAMT1") else None,
                            }
                        )

            if cy2025 or cells_total:
                cov = round(100 * cells_priced / cells_total, 1) if cells_total else None
                results[f"{sheet_name}:{country}"] = {
                    "cy2025_rows": cy2025,
                    "part_cells": cells_total,
                    "priced_cells": cells_priced,
                    "coverage_pct": cov,
                    "qty_total": qty_total,
                    "qty_priced": qty_priced,
                    "value_priced_eur": round(value_priced, 2),
                    "samples_missing": samples_missing,
                    "samples_priced": samples_priced,
                }
                print(
                    f"  {country}: rows={cy2025} part_cells={cells_total} priced={cells_priced} "
                    f"cov={cov}% value€={value_priced:,.2f}"
                )

    (OUT / "_ie_pt_pricing.json").write_text(json.dumps(results, indent=2, default=str), encoding="utf-8")
    wb.close()
    return results


if __name__ == "__main__":
    sheet4_sections()
    ie_pt_from_detail()
