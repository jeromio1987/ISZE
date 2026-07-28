"""One-off warranty audit for Jerome — UMI CY2025 + CV Sheet4 blank bucket."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

DATA = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data")
WARR = DATA / "06_Warranty"
OUT = DATA / "90_INBOX" / "warranty_audit_2026-07-27"
OUT.mkdir(parents=True, exist_ok=True)

results: dict = {}


def audit_umi() -> dict:
    umi_paths = [
        WARR / "UMI warranty data 202301 - 202605.xlsx",
        DATA / "UMI warranty data 202301 - 202605 (1).xlsx",
    ]
    umi_path = next((p for p in umi_paths if p.exists()), None)
    umi: dict = {"file": str(umi_path) if umi_path else None}
    if not umi_path:
        return umi

    wb = openpyxl.load_workbook(umi_path, read_only=True, data_only=True)
    umi["sheets"] = wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    umi["headers"] = hdr[:25]

    date_col = next(
        (i for i, h in enumerate(hdr) if "Written Date" in h or h.upper() == "WRITTEN DATE_YMD"),
        None,
    )
    cost_col = next(
        (i for i, h in enumerate(hdr) if "Approved part cost" in h or "Approved Part Cost" in h),
        None,
    )
    claim_col = next((i for i, h in enumerate(hdr) if "Claim" in h and "No" in h), None)

    cy_total = 0.0
    cy_lines = 0
    claims: set[str] = set()
    sample: list[dict] = []
    for r in rows[1:]:
        if not r or all(x is None for x in r):
            continue
        dt = r[date_col] if date_col is not None else None
        if dt is None:
            continue
        if isinstance(dt, (int, float)) and not isinstance(dt, bool):
            ds = str(int(dt))
        else:
            ds = str(dt).replace("-", "")[:8]
        if len(ds) >= 4 and ds[:4] == "2025":
            try:
                c = float(r[cost_col] or 0) if cost_col is not None else 0.0
            except (TypeError, ValueError):
                c = 0.0
            cy_total += c
            cy_lines += 1
            if claim_col is not None and r[claim_col] is not None:
                claims.add(str(r[claim_col]))
            if len(sample) < 5:
                sample.append({hdr[i]: r[i] if i < len(r) else None for i in range(min(10, len(hdr)))})

    umi.update(
        {
            "cy2025_eur": round(cy_total, 2),
            "cy2025_lines": cy_lines,
            "cy2025_claims": len(claims),
            "sample_rows": sample,
        }
    )
    wb.close()
    return umi


def audit_lcv() -> dict:
    lcv_path = WARR / "Warranty_LCV_2025.xlsx"
    lcv: dict = {"file": str(lcv_path)}
    if not lcv_path.exists():
        return lcv

    wb = openpyxl.load_workbook(lcv_path, read_only=True, data_only=True)
    lcv["sheets"] = wb.sheetnames
    sheet = next((s for s in wb.sheetnames if "Export" in s or "Claim" in s), wb.sheetnames[0])
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    dist_i = next((i for i, h in enumerate(hdr) if "DistributorName" in h or h == "Distributor Name"), None)
    cost_i = next((i for i, h in enumerate(hdr) if "Approved" in h and "cost" in h.lower()), None)

    dist_counts: dict[str, int] = {}
    israel_rows = umi_rows = 0
    israel_eur = 0.0
    israel_names: set[str] = set()
    for r in rows[1:]:
        if not r:
            continue
        d = str(r[dist_i] or "").strip() if dist_i is not None else ""
        if not d:
            continue
        dist_counts[d] = dist_counts.get(d, 0) + 1
        dl = d.lower()
        if "israel" in dl or "umi" in dl or "universal motor" in dl:
            israel_rows += 1
            israel_names.add(d)
            if "umi" in dl or "universal motor" in dl:
                umi_rows += 1
            if cost_i is not None:
                try:
                    israel_eur += float(r[cost_i] or 0)
                except (TypeError, ValueError):
                    pass

    lcv.update(
        {
            "export_sheet": sheet,
            "israel_rows": israel_rows,
            "umi_rows": umi_rows,
            "israel_eur_approx": round(israel_eur, 2),
            "israel_distributor_names": sorted(israel_names),
            "distributors_top": sorted(dist_counts.items(), key=lambda x: -x[1])[:20],
        }
    )
    wb.close()
    return lcv


def audit_cv_sheet4() -> dict:
    cv_path = WARR / "Warranty_CV_2025.xlsx"
    cv: dict = {"file": str(cv_path)}
    if not cv_path.exists():
        return cv

    wb = openpyxl.load_workbook(cv_path, read_only=True, data_only=True)
    cv["sheets"] = wb.sheetnames
    ws = wb["Sheet4"]
    cur = None
    sums: dict[str, float] = {}
    preview_rows: list[dict] = []
    for r in ws.iter_rows(min_row=1, max_col=4, values_only=True):
        lbl, cnt, val = (list(r) + [None, None, None])[:3]
        if lbl is None:
            continue
        s = str(lbl).strip()
        if not re.match(r"^\d{8,}", s) and not s.lower().startswith("grand"):
            cur = s
            sums.setdefault(cur, 0.0)
            if len(preview_rows) < 30:
                preview_rows.append({"type": "country_header", "label": s, "count": cnt, "cell_c": val})
            continue
        if cur and cnt is not None:
            v = val
            if isinstance(v, str):
                v = v.replace("\xa0", "").replace(" ", "").replace(",", ".")
                try:
                    v = float(v)
                except ValueError:
                    v = None
            if isinstance(v, (int, float)):
                try:
                    c = float(cnt)
                except (TypeError, ValueError):
                    c = 1.0
                sums[cur] += c * v

    cv["sheet4_buckets"] = {k: round(v, 2) for k, v in sorted(sums.items(), key=lambda x: -x[1])}
    blank = None
    for k, v in sums.items():
        if str(k).strip() in ("0", "0.0", "", "blank") or str(k).strip().lower() == "0":
            blank = round(v, 2)
    cv["blank_bucket_eur"] = blank
    cv["ireland_sheet4_eur"] = round(sums.get("IRELAND", sums.get("Ireland", 0)), 2)
    cv["portugal_sheet4_eur"] = round(sums.get("PORTUGAL", sums.get("Portugal", 0)), 2)
    cv["preview_rows"] = preview_rows
    wb.close()
    return cv


def audit_ie_pt_cv_pricing() -> dict:
    """Rough price-map coverage on CV truck sheets if Ireland/Portugal tabs exist."""
    cv_path = WARR / "Warranty_CV_2025.xlsx"
    out: dict = {"file": str(cv_path), "tabs": {}}
    if not cv_path.exists():
        return out

    wb = openpyxl.load_workbook(cv_path, read_only=True, data_only=True)
    for tab in wb.sheetnames:
        tl = tab.lower()
        if "ireland" not in tl and "portugal" not in tl and tab not in ("Ireland", "Portugal", "Imotors"):
            continue
        ws = wb[tab]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
        qty_i = next((i for i, h in enumerate(hdr) if h.upper() in ("QTY", "QUANTITY")), None)
        price_i = next((i for i, h in enumerate(hdr) if "PRICE" in h.upper() or "EUR" in h.upper()), None)
        if qty_i is None:
            continue
        total = priced = 0
        for r in rows[1:]:
            if not r:
                continue
            try:
                q = float(r[qty_i] or 0)
            except (TypeError, ValueError):
                continue
            if q == 0:
                continue
            total += 1
            p = r[price_i] if price_i is not None else None
            if p not in (None, "", 0, 0.0):
                try:
                    if float(p) > 0:
                        priced += 1
                except (TypeError, ValueError):
                    pass
        if total:
            out["tabs"][tab] = {
                "rows_with_qty": total,
                "rows_with_price": priced,
                "coverage_pct": round(100 * priced / total, 1),
            }
    wb.close()
    return out


results["umi_dedicated"] = audit_umi()
results["lcv_network"] = audit_lcv()
results["cv_sheet4"] = audit_cv_sheet4()
results["cv_ie_pt_pricing"] = audit_ie_pt_cv_pricing()

gps = pd.read_csv(DATA / "00_derived" / "gps_spu_by_distributor.csv")
umi25 = gps[(gps["year"] == 2025) & (gps["distributor_code"] == "UMI")].iloc[0]
results["pipeline_umi_2025"] = {
    "warranty_parts_eur": float(umi25["warranty_parts_eur"]),
    "warranty_data_gap": bool(umi25["warranty_data_gap"]),
    "sales_ex_warranty_eur": float(umi25["sales_ex_warranty_eur"]),
    "total_sales_eur": float(umi25["total_sales_eur"]),
}

json_path = DATA / "03_analytics" / "Isuzu" / "prism" / "data" / "gps_warranty_cy2025.json"
if not json_path.exists():
    json_path = Path(r"C:\Users\jerom\Desktop\claude\ISZE\03_analytics\Isuzu\prism\data\gps_warranty_cy2025.json")
with json_path.open(encoding="utf-8") as f:
    j = json.load(f)
results["json_warranty_umi"] = j["warranty_eur"].get("UMI")
results["json_data_gaps"] = j.get("data_gaps")

(OUT / "audit_results.json").write_text(json.dumps(results, indent=2, default=str), encoding="utf-8")
print(json.dumps(results, indent=2, default=str))
