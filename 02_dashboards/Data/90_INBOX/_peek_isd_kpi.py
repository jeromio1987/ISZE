"""Save the 3 Aug ISD KPI workbook via sender Restrict."""
from pathlib import Path

import openpyxl
import win32com.client

OUT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data\90_INBOX\_isd_std_mail_20260818\attachments")
dest = OUT / "2026-08-03_Tomoe_KPI_SPI_Check_Sheet_Monthly_2026_ISD.xlsx"

ol = win32com.client.Dispatch("Outlook.Application")
ns = ol.GetNamespace("MAPI")
items = ns.GetDefaultFolder(6).Items
items.Sort("[ReceivedTime]", True)
sub = items.Restrict("[SenderName] = 'Tomoe.Baumann@isuzu.de'")
c = 0
saved = False
for mail in sub:
    c += 1
    if c > 80:
        break
    if mail.Class != 43:
        continue
    for i in range(1, mail.Attachments.Count + 1):
        att = mail.Attachments.Item(i)
        name = str(att.FileName or "")
        if "ISD" in name.upper() and name.lower().endswith((".xlsx", ".xls")):
            att.SaveAsFile(str(dest))
            print("SAVED", dest)
            print("from", mail.ReceivedTime, mail.Subject, name)
            saved = True
            break
    if saved:
        break
print("scanned", c, "saved", saved)

if dest.exists():
    wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
    print("sheets", wb.sheetnames)
    print("size", dest.stat().st_size)
    for sn in wb.sheetnames[:8]:
        ws = wb[sn]
        print("---", sn, "---")
        for i, row in enumerate(ws.iter_rows(max_row=10, max_col=10, values_only=True)):
            print(" | ".join("" if x is None else str(x)[:36] for x in row)[:220])
    wb.close()
