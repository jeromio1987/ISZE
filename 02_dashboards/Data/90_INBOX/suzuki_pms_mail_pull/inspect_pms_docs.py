"""Inspect PMS SPA / study workbooks and compare ITUK PNs + fig keys."""
from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
SPA = ROOT / r"01_data\03_PMS\PMS SPA study_ISZE (1).xlsx"
STUDY = ROOT / r"09_meetings\Preparation Meetings\IML_ITUK_PMS\ISZE_ITUK_UTI_PMS_Study_CY2025_for_IML.xlsx"
ITUK_AUDIT = ROOT / r"01_data\00_derived\ituk_pms_pn_audit.csv"
AIM = ROOT / r"01_data\00_derived\aim_pn_classification.csv"
FIG_PEEK = ROOT / r"09_meetings\Preparation Meetings\IML_ITUK_PMS\peek_FIG-KEY_Parts.json"
SALES_FIG = ROOT / r"01_data\00_derived\sales_by_db_fig_key_year.csv"
CAPTURE = ROOT / r"01_data\00_derived\pms_capture_by_db_year.csv"


def sheet_preview(path: Path, max_rows: int = 6) -> None:
    print(f"\n#### {path.name}")
    if not path.exists():
        print("MISSING", path)
        return
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print("sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n-- {name} --")
        for i, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True), 1):
            vals = [v for v in row if v is not None]
            if vals:
                print(i, vals[:12])
    wb.close()


def extract_pns_from_workbook(path: Path) -> set[str]:
    pns: set[str] = set()
    if not path.exists():
        return pns
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pn_re = re.compile(r"^\d{7,12}$")
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is None:
                    continue
                s = str(v).strip()
                if pn_re.match(s):
                    pns.add(s)
    wb.close()
    return pns


def extract_fig_keys_from_workbook(path: Path) -> set[str]:
    figs: set[str] = set()
    if not path.exists():
        return figs
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    fig_re = re.compile(r"^\d{3}-\d{3}$")
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is None:
                    continue
                s = str(v).strip()
                if fig_re.match(s):
                    figs.add(s)
                # also "FIG-KEY 050-002" style
                m = re.search(r"(\d{3}-\d{3})", s)
                if m and ("FIG" in s.upper() or fig_re.match(s)):
                    figs.add(m.group(1))
    wb.close()
    return figs


def main() -> None:
    sheet_preview(SPA, max_rows=10)
    sheet_preview(STUDY, max_rows=5)

    spa_pns = extract_pns_from_workbook(SPA)
    study_pns = extract_pns_from_workbook(STUDY)
    spa_figs = extract_fig_keys_from_workbook(SPA)
    study_figs = extract_fig_keys_from_workbook(STUDY)

    print("\n#### PN counts")
    print("SPA PNs:", len(spa_pns))
    print("STUDY PNs:", len(study_pns))
    print("SPA fig keys:", sorted(spa_figs))
    print("STUDY fig keys:", sorted(study_figs))

    # ITUK audit PNs
    ituk_pns: set[str] = set()
    ituk_by_cat: dict[str, set[str]] = defaultdict(set)
    if ITUK_AUDIT.exists():
        with ITUK_AUDIT.open(encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                pn = (row.get("part_number") or row.get("pn") or "").strip()
                cat = (row.get("pms_category") or row.get("category") or "").strip()
                if pn:
                    ituk_pns.add(pn)
                    ituk_by_cat[cat].add(pn)
    print("ITUK audit PNs:", len(ituk_pns), "cats:", {k: len(v) for k, v in sorted(ituk_by_cat.items())})

    # Check SPA PNs sold/known for ITUK via sales fig key file + aim
    aim_fig: dict[str, str] = {}
    if AIM.exists():
        with AIM.open(encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                pn = (row.get("part_number") or "").strip()
                fk = (row.get("fig_key") or "").strip()
                if pn:
                    aim_fig[pn] = fk

    spa_in_aim = [pn for pn in sorted(spa_pns) if pn in aim_fig]
    spa_missing_aim = [pn for pn in sorted(spa_pns) if pn not in aim_fig]
    spa_with_fig = [pn for pn in spa_in_aim if aim_fig.get(pn)]
    spa_no_fig = [pn for pn in spa_in_aim if not aim_fig.get(pn)]
    print("\n#### SPA vs catalog")
    print("SPA in AIM/catalog:", len(spa_in_aim), "/", len(spa_pns))
    print("SPA missing from AIM:", len(spa_missing_aim), spa_missing_aim[:20])
    print("SPA with fig_key:", len(spa_with_fig))
    print("SPA blank fig_key:", len(spa_no_fig), spa_no_fig[:20])

    spa_figs_from_aim = sorted({aim_fig[pn] for pn in spa_with_fig if aim_fig[pn]})
    print("SPA fig keys via AIM:", spa_figs_from_aim)

    # ITUK sales by fig key (2025)
    ituk_figs_2025: set[str] = set()
    if SALES_FIG.exists():
        with SALES_FIG.open(encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                if row.get("distributor_code") == "ITUK" and row.get("year") == "2025":
                    fk = (row.get("fig_key") or "").strip()
                    if fk:
                        ituk_figs_2025.add(fk)
    print("\n#### ITUK 2025 sold fig keys:", len(ituk_figs_2025))
    missing_figs = [fk for fk in spa_figs_from_aim if fk not in ituk_figs_2025]
    print("SPA fig keys with NO ITUK 2025 sales:", missing_figs)

    # Check target PN 5876102500
    target = "5876102500"
    print("\n#### Target PN", target)
    print("in SPA:", target in spa_pns)
    print("in STUDY:", target in study_pns)
    print("in ITUK audit:", target in ituk_pns)
    print("AIM fig:", aim_fig.get(target))

    # Capture by DB summary
    if CAPTURE.exists():
        print("\n#### pms_capture_by_db_year (header + ITUK/UTI/IUK sample)")
        with CAPTURE.open(encoding="utf-8", newline="") as f:
            r = csv.DictReader(f)
            print("cols:", r.fieldnames)
            n = 0
            for row in r:
                if row.get("distributor_code") in {"ITUK", "UTI", "IUK", "IRELAND", "Imotors"} or row.get("db") in {
                    "ITUK",
                    "UTI",
                    "IUK",
                }:
                    print(row)
                    n += 1
                    if n > 30:
                        break


if __name__ == "__main__":
    main()
