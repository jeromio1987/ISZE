"""Build FY27 GPS stretch proposal — kickback distributors (+1.5pp flat; ITUK excluded)."""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from load_bp_fy27 import load_bp_fy27_reviewed, _load_incentive_sheet
from gps_v9_loaders import RATE_FMD, RATE_N35, RATE_N75, build_gps_uio_package

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
BUILD = ROOT / "01_data" / "91_build"
DATA = ROOT / "01_data"
OUT_XLSX = BUILD / "ISZE_GPS_Kickback_Stretch_Proposal_FY27.xlsx"
OUT_HTML = ROOT / "02_dashboards" / "gps" / "ISZE_GPS_Kickback_Stretch_Proposal_FY27.html"

ITUK_STUDY_PATHS = [
    DATA / "ITUK Parts Sales Target Study.xlsx",
    ROOT.parent / "ITUK Parts Sales Target Study.xlsx",
    DATA / "90_INBOX" / "ITUK Parts Sales Target Study.xlsx",
]

# Study sheet retail £/vehicle (rows 14–16 col C)
ITUK_RETAIL_GBP = {"N35": 400.0, "N75_LD": 1000.0, "F_MD": 1200.0}
ITUK_RATES_EUR = {"N35": RATE_N35, "N75_LD": RATE_N75, "F_MD": RATE_FMD}  # 296 / 739 / 887
ITUK_STUDY_COL_FY27 = "J"  # calendar/FY27 stretch column in study

RED = "CC0E1E"
NAVY = "1B2A4A"
TERRA = "C0552D"
HDR_FILL = PatternFill("solid", start_color=RED)
ALT_FILL = PatternFill("solid", start_color="F5F5F5")
DGR = "808080"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Target Study — dealer sell-out stretch % of retail potential by year column
ITUK_DEALER_STRETCH_PCT = {2025: 0.03, 2026: 0.07, 2027: 0.09}
ITUK_BRIDGE = 0.66 * 1.12  # GP haircut × GBP→EUR

# Legacy Kosha top-down proxy — audit diff only (build_gps_v6 BP vs Kosha tab)
KOSHA_PROXY_FY27 = {
    "ITUK": 7_425_511,
    "IUK": 7_004_984,
    "UTI": 3_830_292,
    "UMI": 3_815_301,
    "GERMANY": 2_902_192,
    "BELGIUM": 2_413_780,
    "PORTUGAL": 1_928_355,
    "SWEDEN": 1_451_769,
    "NORWAY": 1_619_512,
    "IRELAND": 1_366_406,
}

KICKBACK = [
    ("IUK", "IUK", "United Kingdom (LCV)"),
    ("UTI", "UTI", "Israel (trucks)"),
    ("UMI", "UMI", "Israel (LCV)"),
    ("GERMANY", "ISD", "Germany"),
    ("BELGIUM", "IBX", "Benelux"),
    ("PORTUGAL", "Imotors", "Portugal"),
    ("SWEDEN", "Sverige", "Sweden"),
    ("NORWAY", "Norway", "Norway"),
    ("IRELAND", "Ireland", "Ireland"),
]
KICKBACK_CODES = {code for code, _, _ in KICKBACK}

# Reviewed BP customers outside GPS_DBS / UIO v19 — max spend left open on sheet 2
NO_UIO_EXTRAS: list[tuple[str, str, str]] = [
    ("KAZAKHSTAN", "KTK", "Kazakhstan"),
    ("IMG", "IMG (IC)", "IMG (IC)"),
    ("MIDI", "MIDI", "MIDI Europe"),
    ("GRAND_AVTO", "Grand Avto", "Georgia / CIS"),
    ("CBC_TRANS", "CBC-Trans", "CIS"),
    ("KIMA", "KIMA", "CIS"),
]

OTHER_DB_LABELS: dict[str, str] = {
    "ITUK": "United Kingdom (CV) — Target Study",
    "CYPRUS": "Cyprus",
    "DENMARK": "Denmark",
    "FINLAND": "Finland",
    "GREECE": "Greece",
    "HUNGARY": "Hungary",
    "ICELAND": "Iceland",
    "MALTA": "Malta",
    "ROMANIA": "Romania",
}

STRETCH_PP = 1.5  # flat +pp on plan GPS for all kickback DBs (FY27)

# Last complete fiscal year — baseline for gap sizing vs FY27 plan
ACTUAL_FY = 2026
ACTUAL_FY_LABEL = "FY26"


def _fiscal_gps_now(d: dict) -> tuple[float | None, dict]:
    """GPS ex-warranty for gap sizing; prefer last complete FY, then fallbacks."""
    fiscal = d["actuals"]["fiscal"]
    fy_row = fiscal.get(str(ACTUAL_FY)) or fiscal.get(ACTUAL_FY) or {}
    gps = fy_row.get("gps_ex_warranty")
    if gps is not None:
        return gps, fy_row
    fy25 = fiscal.get("2025") or {}
    gps = fy25.get("gps_ex_warranty")
    if gps is not None:
        return gps, fy25
    cy = d["actuals"]["calendar"].get("2025") or {}
    return cy.get("gps_ex_warranty"), cy


def _round_pp(pp: float) -> float:
    """Legacy helper — stretch is flat STRETCH_PP; kept for ITUK study helpers."""
    return STRETCH_PP


def _fmt_pp(pp: float) -> str:
    return f"+{int(pp)}" if pp == int(pp) else f"+{pp:.1f}"


def _stretch_pp_and_rationale(gps_code: str, gps_now: float, gps_plan: float) -> tuple[float, str]:
    pp = STRETCH_PP
    fy = gps_now * 100
    plan = gps_plan * 100
    return pp, (
        f"Flat network stretch {_fmt_pp(pp)} on FY27 plan GPS "
        f"({ACTUAL_FY_LABEL} actual {fy:.0f}% vs plan {plan:.0f}%). "
        f"ITUK excluded — see reviewed BP only."
    )


def _resolve_ituk_study_path() -> Path | None:
    for p in ITUK_STUDY_PATHS:
        if p.is_file():
            return p
    return None


def _ituk_study_fy27_detail() -> dict:
    """Rebuild study row maths for FY27 column J (dealer +9% stretch)."""
    pkg = build_gps_uio_package()
    buckets = pkg["distributors"]["ITUK"]["uio"][2027]
    uio = {k: float(buckets.get(k, 0)) for k in ITUK_RETAIL_GBP}
    dealer_pot = {k: uio[k] * ITUK_RETAIL_GBP[k] for k in ITUK_RETAIL_GBP}
    j22 = sum(dealer_pot.values())
    # Row 33–36: UIO × canonical €/UIO rates (not £ retail × bridge)
    isze_wallet = {k: uio[k] * ITUK_RATES_EUR[k] for k in ITUK_RETAIL_GBP}
    j36 = sum(isze_wallet.values())
    pct = ITUK_DEALER_STRETCH_PCT[2027]
    j28 = j22 * pct
    j39 = j28 * ITUK_BRIDGE  # stretch increment via 0.7392 bridge
    return {
        "uio": uio,
        "rates_eur": dict(ITUK_RATES_EUR),
        "dealer_pot_gbp": dealer_pot,
        "j22_gbp": j22,
        "isze_wallet_eur": isze_wallet,
        "j36_eur": j36,
        "stretch_pct": pct,
        "j28_gbp": j28,
        "j39_eur": j39,
        "uio_source": pkg.get("source_uio", ""),
    }


def _copy_sheet_values(src_ws, wb: openpyxl.Workbook, title: str, max_row: int = 48, max_col: int = 13) -> None:
    """Copy study sheet values into proposal workbook (audit reference)."""
    safe = title[:31]
    if safe in wb.sheetnames:
        del wb[safe]
    ws = wb.create_sheet(safe)
    ws.sheet_view.showGridLines = False
    for r in range(1, min(max_row, src_ws.max_row or max_row) + 1):
        for c in range(1, min(max_col, src_ws.max_column or max_col) + 1):
            ws.cell(r, c, src_ws.cell(r, c).value)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    for col in range(3, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12


def _embed_ituk_study_sheets(wb: openpyxl.Workbook) -> str:
    path = _resolve_ituk_study_path()
    if path:
        src = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sn in ("Sales Share and Stretch", "ITUK UIO Projection"):
            if sn in src.sheetnames:
                _copy_sheet_values(src[sn], wb, f"Study · {sn[:18]}")
        src.close()
        return str(path)
    return ""


def _write_ituk_study_excerpt(ws, detail: dict, study_path: str) -> None:
    """Key study rows (FY27 col J) when source xlsx is not on disk."""
    ws.sheet_view.showGridLines = False
    ws["A1"] = "ITUK Parts Sales Target Study — excerpt (FY27 col J)"
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color=RED)
    note = (
        f"Embedded from: {study_path}"
        if study_path
        else "Source file not found on disk — rows below rebuilt from study formulas + UIO v19 (same maths as col J)."
    )
    ws["A2"] = note
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws.merge_cells("A2:F2")

    hdr = ["Row", "Study label", "Formula (col J)", "Value"]
    hr = 4
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.border = BORDER

    col = ITUK_STUDY_COL_FY27
    lines = [
        ("9", "UIO units — N35", f"{col}8", detail["uio"]["N35"], "#,##0"),
        ("9", "UIO units — N75/LD", f"{col}9", detail["uio"]["N75_LD"], "#,##0"),
        ("9", "UIO units — F/MD", f"{col}10", detail["uio"]["F_MD"], "#,##0"),
        ("14", "Retail £/UIO — N35", "C14", ITUK_RETAIL_GBP["N35"], "£#,##0"),
        ("15", "Retail £/UIO — N75", "C15", ITUK_RETAIL_GBP["N75_LD"], "£#,##0"),
        ("16", "Retail £/UIO — F/MD", "C16", ITUK_RETAIL_GBP["F_MD"], "£#,##0"),
        ("—", "ISZE rate €/UIO — N35", "gps_v9_loaders", ITUK_RATES_EUR["N35"], "€#,##0"),
        ("—", "ISZE rate €/UIO — N75/LD", "gps_v9_loaders", ITUK_RATES_EUR["N75_LD"], "€#,##0"),
        ("—", "ISZE rate €/UIO — F/MD", "gps_v9_loaders", ITUK_RATES_EUR["F_MD"], "€#,##0"),
        ("19", "Dealer potential — N35", f"{col}19", detail["dealer_pot_gbp"]["N35"], "£#,##0"),
        ("20", "Dealer potential — N75/LD", f"{col}20", detail["dealer_pot_gbp"]["N75_LD"], "£#,##0"),
        ("21", "Dealer potential — F/MD", f"{col}21", detail["dealer_pot_gbp"]["F_MD"], "£#,##0"),
        ("22", "Dealer retail potential (TTL)", f"{col}22 = SUM(19:21)", detail["j22_gbp"], "£#,##0"),
        ("28", f"Stretch uplift ({detail['stretch_pct']:.0%} of potential)", f"{col}28 = {col}22×{detail['stretch_pct']:.0%}", detail["j28_gbp"], "£#,##0"),
        ("33", "ISZE wallet — N35", f"{col}33 = {col}8×€296", detail["isze_wallet_eur"]["N35"], "€#,##0"),
        ("34", "ISZE wallet — N75/LD", f"{col}34 = {col}9×€739", detail["isze_wallet_eur"]["N75_LD"], "€#,##0"),
        ("35", "ISZE wallet — F/MD", f"{col}35 = {col}10×€887", detail["isze_wallet_eur"]["F_MD"], "€#,##0"),
        ("36", "ISZE max spend / wallet (TTL EUR)", f"{col}36 = SUM(33:35)", detail["j36_eur"], "€#,##0"),
        ("39", "ISZE stretch € increment", f"{col}39 = {col}28×0.66×1.12", detail["j39_eur"], "€#,##0"),
    ]
    r = hr + 1
    for row_no, label, formula, val, fmt in lines:
        ws.cell(r, 1, row_no).border = BORDER
        ws.cell(r, 2, label).border = BORDER
        ws.cell(r, 3, formula).border = BORDER
        c = ws.cell(r, 4, val)
        c.number_format = fmt
        c.font = Font(bold=True, color=TERRA)
        c.border = BORDER
        r += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 16


def _load_ituk_row(dual_by_code: dict, bp_map: dict[str, float]) -> dict:
    """ITUK stretch from Parts Sales Target Study — dealer sell-out, not +pp kickback rule."""
    d = dual_by_code["ITUK"]
    study = _ituk_study_fy27_detail()
    max27 = study["j36_eur"]  # row 36 EUR wallet (UIO v19 × €296/739/887)
    gps_now, fy_actual = _fiscal_gps_now(d)
    base = bp_map["ITUK"]
    gps_plan = base / max27
    dealer_stretch = ITUK_DEALER_STRETCH_PCT[2027]
    incr_eur = study["j39_eur"]
    pp_up = dealer_stretch * 100
    gps_stretch = gps_plan + dealer_stretch
    stretch_bp_eur = base + incr_eur
    return {
        "gps_code": "ITUK",
        "kosha_code": "ITUK",
        "country": "United Kingdom (CV) — Target Study",
        "gps_now": gps_now,
        "gps_plan": gps_plan,
        "gps_stretch": gps_stretch,
        "base_bp_eur": base,
        "gap_pp": (gps_plan - gps_now) * 100,
        "stretch_pp": pp_up,
        "stretch_pp_label": _fmt_pp(pp_up),
        "rationale": (
            f"Target Study FY27: dealer sell-out +{dealer_stretch*100:.0f}% of UK retail parts potential "
            f"(row 28 col J). ISZE increment = potential × {dealer_stretch*100:.0f}% × 0.66 × 1.12 "
            f"= +{_fmt_pp(pp_up)} on plan GPS wallet. Only DB on dealer-data stretch."
        ),
        "method": "target_study",
        "dealer_stretch_pct": dealer_stretch,
        "eur_per_pp": max27 / 100,
        "increment_eur": incr_eur,
        "stretch_bp_eur": stretch_bp_eur,
        "max27": max27,
        "fy_actual_sales_ex": fy_actual.get("sales_ex_warranty_eur"),
        "dealer_potential_gbp": study["j22_gbp"],
        "dealer_stretch_gbp": study["j28_gbp"],
        "study": study,
        "study_path": _resolve_ituk_study_path(),
    }


def _gps_row(
    *,
    gps_code: str,
    kosha_code: str,
    country: str,
    dual: dict,
    bp_map: dict[str, float],
    method: str,
    rationale: str,
    has_uio: bool = True,
    stretch_pp: float | None = None,
    increment_eur: float | None = None,
    stretch_bp_eur: float | None = None,
    gps_stretch: float | None = None,
) -> dict:
    if has_uio:
        max27 = dual["max_spend"]["fy2027"]
        gps_now, fy_actual = _fiscal_gps_now(dual)
        base = bp_map[gps_code]
        gps_plan = base / max27
        pp_up = stretch_pp if stretch_pp is not None else STRETCH_PP
        if increment_eur is None:
            increment_eur = pp_up / 100 * max27
        if gps_stretch is None:
            gps_stretch = gps_plan + pp_up / 100
        if stretch_bp_eur is None:
            stretch_bp_eur = base + increment_eur
        return {
            "gps_code": gps_code,
            "kosha_code": kosha_code,
            "country": country,
            "has_uio": True,
            "gps_now": gps_now,
            "gps_plan": gps_plan,
            "gps_stretch": gps_stretch,
            "base_bp_eur": base,
            "gap_pp": (gps_plan - gps_now) * 100 if gps_now is not None else None,
            "stretch_pp": pp_up,
            "stretch_pp_label": _fmt_pp(pp_up),
            "rationale": rationale,
            "method": method,
            "eur_per_pp": max27 / 100,
            "increment_eur": increment_eur,
            "stretch_bp_eur": stretch_bp_eur,
            "max27": max27,
            "fy_actual_sales_ex": fy_actual.get("sales_ex_warranty_eur"),
        }
    base = bp_map[gps_code]
    return {
        "gps_code": gps_code,
        "kosha_code": kosha_code,
        "country": country,
        "has_uio": False,
        "gps_now": None,
        "gps_plan": None,
        "gps_stretch": None,
        "base_bp_eur": base,
        "gap_pp": None,
        "stretch_pp": stretch_pp if stretch_pp is not None else STRETCH_PP,
        "stretch_pp_label": _fmt_pp(stretch_pp if stretch_pp is not None else STRETCH_PP),
        "rationale": rationale,
        "method": method,
        "eur_per_pp": None,
        "increment_eur": None,
        "stretch_bp_eur": None,
        "max27": None,
        "fy_actual_sales_ex": None,
    }


def _load_kickback_rows(dual_by_code: dict, bp_map: dict[str, float]) -> list[dict]:
    rows = []
    for gps_code, kosha_code, country in KICKBACK:
        d = dual_by_code[gps_code]
        gps_now, _ = _fiscal_gps_now(d)
        gps_plan = bp_map[gps_code] / d["max_spend"]["fy2027"]
        pp_up, rationale = _stretch_pp_and_rationale(gps_code, gps_now, gps_plan)
        rows.append(
            _gps_row(
                gps_code=gps_code,
                kosha_code=kosha_code,
                country=country,
                dual=d,
                bp_map=bp_map,
                method="kickback_pp",
                rationale=rationale,
                stretch_pp=pp_up,
            )
        )
    return rows


def _load_no_uio_rows() -> list[dict]:
    incentive = _load_incentive_sheet()
    reviewed = json.loads((BUILD / "bp_fy27_reviewed.json").read_text(encoding="utf-8"))
    non_gps = reviewed.get("non_gps_customers_eur", {})
    rows: list[dict] = []
    for gps_code, kosha_code, country in NO_UIO_EXTRAS:
        if gps_code == "KAZAKHSTAN":
            base = incentive.get("KAZAKHSTAN")
        else:
            base = non_gps.get(kosha_code)
        if base is None:
            continue
        rows.append(
            _gps_row(
                gps_code=gps_code,
                kosha_code=kosha_code,
                country=country,
                dual={},
                bp_map={gps_code: base},
                method="no_uio",
                rationale="No UIO v19 parc — enter max spend (col K) when available; GPS cols stay blank until K is filled.",
                has_uio=False,
            )
        )
    return rows


def _load_other_rows(dual_by_code: dict, bp_map: dict[str, float]) -> list[dict]:
    rows: list[dict] = []
    for gps_code in sorted(dual_by_code):
        if gps_code in KICKBACK_CODES or gps_code == "ITUK":
            continue
        d = dual_by_code[gps_code]
        country = OTHER_DB_LABELS.get(gps_code, d.get("scope", gps_code))
        gps_now, _ = _fiscal_gps_now(d)
        gps_plan = bp_map[gps_code] / d["max_spend"]["fy2027"]
        rationale = (
            f"Reviewed BP FY27 — not on kickback list. "
            f"{ACTUAL_FY_LABEL} GPS {gps_now*100:.0f}% vs plan {gps_plan*100:.0f}%. "
            f"Default +pp from B3 (override in col E)."
        )
        rows.append(
            _gps_row(
                gps_code=gps_code,
                kosha_code=gps_code,
                country=country,
                dual=d,
                bp_map=bp_map,
                method="other_gps",
                rationale=rationale,
            )
        )
    rows.extend(_load_no_uio_rows())
    return rows


def _load_ituk_only(dual_by_code: dict, bp_map: dict[str, float]) -> dict | None:
    """ITUK Target Study — never mixed into kickback or Other DBs tabs."""
    if "ITUK" not in dual_by_code:
        return None
    return _load_ituk_row(dual_by_code, bp_map)


def _load_rows() -> tuple[list[dict], list[dict], dict | None]:
    dual = json.loads((BUILD / "gps_dual_basis.json").read_text(encoding="utf-8"))
    bp_map = load_bp_fy27_reviewed()
    by_code = {d["code"]: d for d in dual["distributors"]}
    return (
        _load_kickback_rows(by_code, bp_map),
        _load_other_rows(by_code, bp_map),
        _load_ituk_only(by_code, bp_map),
    )


def _style_hdr(ws, hr: int, hdr: list[str]) -> None:
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[hr].height = 28


INPUT_FILL = PatternFill("solid", start_color="FFFDE7")  # light yellow — editable inputs


def _write_proposal_sheet(
    ws,
    rows: list[dict],
    wb: openpyxl.Workbook,
    *,
    title: str,
    subtitle: str,
    total_label: str,
    register_default_pp: bool = False,
) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=RED)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws.merge_cells("A2:K2")

    # ── Parameters (row 3) ──
    ws["A3"] = "Default +pp (all DBs)"
    ws["A3"].font = Font(name="Arial", size=10, bold=True, color=NAVY)
    c_pp = ws["B3"]
    c_pp.value = STRETCH_PP
    c_pp.number_format = "0.0"
    c_pp.font = Font(name="Arial", size=11, bold=True, color=TERRA)
    c_pp.fill = INPUT_FILL
    c_pp.border = BORDER
    ws["C3"] = "→ col E per DB (=B3 unless you type an override)"
    ws["C3"].font = Font(name="Arial", size=9, italic=True, color=DGR)
    ws.merge_cells("C3:K3")

    hdr = [
        "DB", f"GPS {ACTUAL_FY_LABEL} ex", "Plan GPS FY27", "Stretch GPS FY27", "+pp on plan",
        "BP FY27 base", "+€ GPS stretch", "Stretch BP FY27", "Δ vs base", "Notes",
        "Max spend FY27",
    ]
    hr = 5
    _style_hdr(ws, hr, hdr)

    data_start = hr + 1
    r = data_start
    for row in rows:
        ws.cell(r, 1, row["gps_code"]).font = Font(name="Arial", size=10, bold=True)
        ws.cell(r, 1).border = BORDER
        if not row.get("has_uio", True):
            ws.cell(r, 1).font = Font(name="Arial", size=10, bold=True, color=DGR)

        # B — FY26 GPS ex (input; blank when no UIO)
        c_b = ws.cell(r, 2, row["gps_now"] if row.get("has_uio", True) else None)
        c_b.number_format = "0.0%"
        c_b.fill = INPUT_FILL
        c_b.border = BORDER

        # F — BP base (input)
        c_f = ws.cell(r, 6, row["base_bp_eur"])
        c_f.number_format = "€#,##0"
        c_f.fill = INPUT_FILL
        c_f.border = BORDER

        # K — max spend (input; open when no UIO)
        c_k = ws.cell(r, 11, row["max27"] if row.get("has_uio", True) else None)
        c_k.number_format = "€#,##0"
        c_k.fill = INPUT_FILL
        c_k.border = BORDER

        # E — +pp (default B3; ITUK Target Study uses fixed +9pp)
        if row.get("method") == "target_study":
            c_e = ws.cell(r, 5, row["stretch_pp"])
        else:
            c_e = ws.cell(r, 5, f"=$B$3")
        c_e.number_format = "0.0"
        c_e.fill = INPUT_FILL
        c_e.border = BORDER

        # C — Plan GPS = BP ÷ max spend
        c_c = ws.cell(r, 3, f"=IF(K{r}>0,F{r}/K{r},0)")
        c_c.number_format = "0.0%"
        c_c.border = BORDER

        # D — Stretch GPS = plan + pp
        c_d = ws.cell(r, 4, f"=C{r}+E{r}/100")
        c_d.number_format = "0.0%"
        c_d.border = BORDER

        # G — € increment = pp% × max spend
        c_g = ws.cell(r, 7, f"=E{r}/100*K{r}")
        c_g.number_format = "€#,##0"
        c_g.border = BORDER

        # H — Stretch BP
        c_h = ws.cell(r, 8, f"=F{r}+G{r}")
        c_h.number_format = "€#,##0"
        c_h.border = BORDER

        # I — delta %
        c_i = ws.cell(r, 9, f"=IF(F{r}>0,G{r}/F{r},0)")
        c_i.number_format = "0.0%"
        c_i.border = BORDER

        # J — notes
        if row.get("method") == "target_study":
            note_txt = row.get("rationale", "Target Study — +9% dealer sell-out (not kickback +1.5pp)")
        elif row.get("has_uio", True):
            note_txt = (
                f'=TEXT(B{r},"0.0%")&" actual · plan "&TEXT(C{r},"0.0%")&" · stretch +pp in E (default B3)"'
            )
        else:
            note_txt = row.get("rationale", "No UIO — fill max spend (K) when available")
        c_j = ws.cell(r, 10, note_txt)
        c_j.font = Font(name="Arial", size=9, italic=True)
        c_j.alignment = Alignment(wrap_text=True, vertical="top")
        c_j.border = BORDER

        ws.row_dimensions[r].height = 28
        r += 1

    tr = r
    ws.cell(tr, 1, total_label).font = Font(bold=True)
    for col in (6, 7, 8):
        letter = get_column_letter(col)
        ws.cell(tr, col, f"=SUM({letter}{data_start}:{letter}{tr-1})").font = Font(bold=True)
        ws.cell(tr, col).number_format = "€#,##0"
        ws.cell(tr, col).border = BORDER
    ws.cell(tr, 9, f"=IF(F{tr}>0,G{tr}/F{tr},0)").font = Font(bold=True)
    ws.cell(tr, 9).number_format = "0.0%"
    ws.cell(tr, 9).border = BORDER

    note = tr + 2
    notes = [
        "Yellow cells = inputs you can change. B3 sets +pp for all DBs (col E); type in E to override one DB.",
        "C = F÷K (plan GPS).  D = C+E/100.  G = E/100×K.  H = F+G.  I = G÷F.",
        "BP base (F) from budget_per_DB_reviewed_24042026.xlsx · Max spend (K) from UIO v19 FY27 wallet.",
        f"GPS {ACTUAL_FY_LABEL} ex (B) = fiscal Apr 2025–Mar 2026, ex-warranty.",
        "Grey DB codes = no UIO v19 — leave K open until parc/wallet is known.",
    ]
    for i, txt in enumerate(notes):
        ws.cell(note + i, 1, f"• {txt}").font = Font(name="Arial", size=9, italic=True)
        ws.merge_cells(start_row=note + i, start_column=1, end_row=note + i, end_column=11)

    widths = [10, 10, 10, 11, 9, 12, 12, 12, 9, 36, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    if register_default_pp:
        wb.defined_names.add(
            DefinedName("STRETCH_PP_DEFAULT", attr_text=f"'{ws.title}'!$B$3")
        )


def _write_ituk_sheet(ws, ituk: dict) -> None:
    ws.sheet_view.showGridLines = False
    study = ituk["study"]
    study_path = ituk.get("study_path")
    path_txt = str(study_path) if study_path else "(not on disk — see tab 'Study · excerpt')"
    ws["A1"] = "ITUK stretch FY27 — step-by-step from Target Study"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=RED)
    ws["A2"] = (
        f"Source: ITUK Parts Sales Target Study.xlsx · sheet 'Sales Share and Stretch' · FY27 column {ITUK_STUDY_COL_FY27}. "
        f"File: {path_txt}"
    )
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws.merge_cells("A2:H2")

    def title(r: int, txt: str) -> None:
        ws.cell(r, 1, txt).font = Font(name="Arial", size=11, bold=True, color=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)

    def step_row(r: int, n: int, what: str, study_ref: str, calc: str, val, fmt: str | None = None) -> None:
        ws.cell(r, 1, n).font = Font(bold=True, color=TERRA)
        ws.cell(r, 2, what).font = Font(size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 4, study_ref).font = Font(size=9, italic=True)
        ws.cell(r, 5, calc).font = Font(size=9)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        if val is not None:
            c = ws.cell(r, 7, val)
            c.font = Font(bold=True, color=NAVY)
            if fmt:
                c.number_format = fmt
        for c in range(1, 8):
            ws.cell(r, c).border = BORDER

    r = 4
    title(r, "Walkthrough — dealer layer → ISZE layer → BP stretch")
    r += 1
    hdr = ["#", "Step", "", "Study ref", "Calculation", "", "Result"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.border = BORDER
    r += 1

    col = ITUK_STUDY_COL_FY27
    u = study["uio"]
    rt = study.get("rates_eur", ITUK_RATES_EUR)
    steps = [
        (
            1,
            "UK dealer parc (UIO by model, FY27)",
            f"Row 8–10 col {col}",
            f"DVLA + UIO v19 · N35 {u['N35']:,.0f} · N75/LD {u['N75_LD']:,.0f} · F/MD {u['F_MD']:,.0f}",
            None,
            None,
        ),
        (
            2,
            "Retail parts potential £/vehicle (dealer layer)",
            "Rows 14–16 col C",
            "N35 £400 · N75 £1,000 · F/MD £1,200 retail GP",
            None,
            None,
        ),
        (
            3,
            "ISZE max-spend rates €/UIO (gps_v9_loaders)",
            "UIO v19",
            f"N35 €{rt['N35']:.0f} · N75/LD €{rt['N75_LD']:.0f} · F/MD €{rt['F_MD']:.0f}",
            None,
            None,
        ),
        (
            4,
            "Dealer retail parts potential (TTL £) — row 22",
            f"Row 22 col {col}",
            f"{col}22 = Σ(UIO × £/veh) — ~£{study['j22_gbp']/1e6:.1f}M not ISZE wallet",
            study["j22_gbp"],
            "£#,##0",
        ),
        (
            5,
            "FY27 stretch rule on dealer sell-out",
            f"Row 28 col {col}",
            f"{col}28 = {col}22 × {study['stretch_pct']:.0%}  (+9% of retail potential)",
            study["j28_gbp"],
            "£#,##0",
        ),
        (
            6,
            "ISZE max spend / wallet FY27 (EUR) — row 36",
            f"Row 36 col {col}",
            f"{col}36 = Σ(UIO × €296/739/887) — ~€{study['j36_eur']/1e6:.1f}M not £17M",
            study["j36_eur"],
            "€#,##0",
        ),
        (
            7,
            "Bridge for stretch increment only",
            "Row 39",
            f"£ stretch × 0.7392  (0.66 GP × 1.12 FX)",
            ITUK_BRIDGE,
            "0.0000",
        ),
        (
            8,
            "ISZE stretch € increment",
            f"Row 39 col {col}",
            f"{col}39 = {col}28 × 0.7392  (= 9% × {col}36)",
            study["j39_eur"],
            "€#,##0",
        ),
        (
            9,
            "Reviewed BP FY27 (base)",
            "budget_per_DB_reviewed_24042026",
            "Incentive FY27 · VP turnover 24 Apr 2026",
            ituk["base_bp_eur"],
            "€#,##0",
        ),
        (
            10,
            "Plan GPS @ BP",
            "ITUK FY27 tab",
            "Plan GPS = BP ÷ wallet (row 36)",
            ituk["gps_plan"],
            "0.0%",
        ),
        (
            11,
            "Stretch BP FY27",
            "ITUK FY27 tab",
            "Stretch BP = base BP + study increment (row 39)",
            ituk["stretch_bp_eur"],
            "€#,##0",
        ),
        (
            12,
            "Stretch GPS FY27",
            "ITUK FY27 tab",
            f"Stretch GPS = stretch BP ÷ {col}36  (= plan + 9pp on wallet)",
            ituk["gps_stretch"],
            "0.0%",
        ),
        (
            13,
            "Check vs kickback DBs",
            "Tab 1",
            f"Only ITUK uses Target Study; kickback DBs use flat +{STRETCH_PP:g}pp",
            ituk["increment_eur"],
            "€#,##0",
        ),
    ]
    for n, what, ref, calc, val, fmt in steps:
        step_row(r, n, what, ref, calc, val, fmt)
        ws.row_dimensions[r].height = 28 if n in (1, 2) else 22
        r += 1

    r += 1
    title(r, "Result summary")
    r += 1
    ws.cell(
        r,
        1,
        f"  Stretch BP FY27 = €{ituk['stretch_bp_eur']/1e6:.2f}M "
        f"(base €{ituk['base_bp_eur']/1e6:.2f}M + study +€{ituk['increment_eur']/1e3:.0f}K). "
        f"GPS {ituk['gps_plan']*100:.1f}% → {ituk['gps_stretch']*100:.1f}% (+{ituk['stretch_pp']:.0f}pp). "
        f"Dealer stretch: +£{study['j28_gbp']/1e6:.2f}M on £{study['j22_gbp']/1e6:.1f}M potential.",
    ).font = Font(name="Arial", size=10)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)

    r += 2
    title(r, "Study rows not used for ISZE BP (dealer-only context)")
    r += 1
    for txt in [
        f"Row 25–26: dealer actual sales £ and dealer GPS % ({col}25/{col}22) — sell-out baseline in study, not ISZE billing.",
        f"Row 29–30: stretched dealer target {col}29 = {col}25+{col}28 and dealer GPS {col}30 = {col}29/{col}22.",
        f"Row 37–38: ISZE actual parts and GPS % — compare to {ACTUAL_FY_LABEL} column in Proposal (ex-warranty).",
        "Full study grid: tab 'Study · Sales Share…' (embedded file) or 'Study · excerpt' (rebuilt rows).",
    ]:
        ws.cell(r, 1, f"  • {txt}").font = Font(name="Arial", size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 36
    ws.column_dimensions["F"].width = 4
    ws.column_dimensions["G"].width = 14


def build_xlsx(kickback_rows: list[dict], other_rows: list[dict], ituk: dict | None) -> None:
    wb = openpyxl.Workbook()
    ws_prop = wb.active
    ws_prop.title = "Proposal FY27"
    _write_proposal_sheet(
        ws_prop,
        kickback_rows,
        wb,
        title="ISZE GPS Stretch Proposal — Kickback distributors FY2027",
        subtitle=(
            f"Editable: default +pp (B3), per-DB +pp (col E), BP base (F), max spend (K). "
            f"Formulas in C, D, G, H, I. ITUK excluded — see tab 'ITUK FY27'. Built {date.today().isoformat()}."
        ),
        total_label=f"TOTAL ({len(kickback_rows)} kickback DBs)",
        register_default_pp=True,
    )
    ws_other = wb.create_sheet("Other DBs FY27")
    _write_proposal_sheet(
        ws_other,
        other_rows,
        wb,
        title="ISZE GPS — Other distributors FY2027",
        subtitle=(
            f"Non-kickback GPS DBs + no-UIO lines. Default +pp from B3. "
            f"ITUK is NOT on this tab — Target Study on 'ITUK FY27'. Built {date.today().isoformat()}."
        ),
        total_label=f"TOTAL ({len(other_rows)} other DBs)",
    )
    if ituk and ituk.get("study"):
        ws_ituk = wb.create_sheet("ITUK FY27")
        _write_ituk_sheet(ws_ituk, ituk)
        study_path = _embed_ituk_study_sheets(wb)
        ws_ex = wb.create_sheet("Study · excerpt")
        _write_ituk_study_excerpt(ws_ex, ituk["study"], study_path)
    try:
        wb.save(OUT_XLSX)
        saved = OUT_XLSX
    except PermissionError:
        alt = OUT_XLSX.with_name(OUT_XLSX.stem + "_build.xlsx")
        wb.save(alt)
        saved = alt
        print("WARN: could not overwrite (file open?) — wrote", alt)
    print("written", saved)


def _fmt_eur(x: float) -> str:
    if abs(x) >= 1e6:
        return f"€{x/1e6:.2f}M"
    return f"€{x/1e3:.0f}K"


def _gps_comparison_svg(rows: list[dict]) -> str:
    row_h = 28
    y0 = 36
    bar_h = 7
    max_x = 420
    parts = [
        f'<text x="0" y="16" class="tit">GPS % — {ACTUAL_FY_LABEL} actual · FY27 plan · stretch</text>',
        f'<rect x="72" y="22" width="12" height="8" fill="{NAVY}"/><text x="88" y="30" class="leg">{ACTUAL_FY_LABEL}</text>',
        f'<rect x="130" y="22" width="12" height="8" fill="#999"/><text x="146" y="30" class="leg">Plan GPS</text>',
        f'<rect x="220" y="22" width="12" height="8" fill="{TERRA}"/><text x="236" y="30" class="leg">Stretch</text>',
    ]
    for i, row in enumerate(rows):
        y = y0 + i * row_h
        parts.append(f'<text x="0" y="{y+14}" class="lbl">{row["gps_code"]}</text>')
        for j, (val, col) in enumerate(
            [(row["gps_now"], NAVY), (row["gps_plan"], "#999"), (row["gps_stretch"], TERRA)]
        ):
            w = int(max_x * min(val, 1.0))
            parts.append(
                f'<rect x="72" y="{y + j*bar_h}" width="{w}" height="{bar_h-1}" fill="{col}" rx="1"/>'
            )
        parts.append(
            f'<text x="{max_x+80}" y="{y+14}" class="val">'
            f'{row["gps_now"]*100:.0f}% → {row["gps_plan"]*100:.0f}% → {row["gps_stretch"]*100:.0f}%'
            f' ({row["stretch_pp_label"]})</text>'
        )
    h = y0 + len(rows) * row_h + 10
    return f'<svg width="700" height="{h}" viewBox="0 0 700 {h}" xmlns="http://www.w3.org/2000/svg">\n' + "\n".join(parts) + "\n</svg>"


def build_html(rows: list[dict]) -> None:
    tot_base = sum(r["base_bp_eur"] for r in rows)
    tot_incr = sum(r["increment_eur"] for r in rows)
    tot_stretch = sum(r["stretch_bp_eur"] for r in rows)
    svg_gps = _gps_comparison_svg(rows)

    table_rows = ""
    for row in rows:
        d_pct = 100 * row["increment_eur"] / row["base_bp_eur"]
        hi = ""
        table_rows += f"""<tr{hi}>
<td><b>{row['gps_code']}</b></td>
<td>{row['gps_now']*100:.1f}%</td>
<td>{row['gps_plan']*100:.1f}%</td>
<td class="stretch"><b>{row['gps_stretch']*100:.1f}%</b></td>
<td class="pp">{row['stretch_pp_label']}</td>
<td>{_fmt_eur(row['base_bp_eur'])}</td>
<td class="inc">+{_fmt_eur(row['increment_eur'])}</td>
<td><b>{_fmt_eur(row['stretch_bp_eur'])}</b></td>
<td>+{d_pct:.1f}%</td>
<td class="why">{row['rationale']}</td></tr>"""

    html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>ISZE GPS Stretch Proposal FY27</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:'Segoe UI',Arial,sans-serif;background:#f4f4f5;color:#1a1a1a;line-height:1.5}}
.wrap{{max-width:1000px;margin:0 auto;padding:28px 24px 48px}}
header{{border-left:6px solid {RED};padding-left:16px;margin-bottom:24px}}
header h1{{font-size:22px;color:{NAVY}}}
header p{{font-size:13px;color:#555;margin-top:6px}}
.kpis{{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin-bottom:24px}}
.kpi{{background:#fff;border:1px solid #e0e0e0;border-top:4px solid {TERRA};border-radius:8px;padding:14px 16px}}
.kpi .v{{font-size:26px;font-weight:700;color:{NAVY}}}
.kpi .l{{font-size:12px;color:#444;margin-top:4px}}
.card{{background:#fff;border:1px solid #e0e0e0;border-radius:8px;padding:18px 20px;margin-bottom:20px}}
.card h2{{font-size:15px;color:{NAVY};margin-bottom:12px}}
.tit{{font-size:12px;font-weight:700;fill:{NAVY}}}
.lbl{{font-size:11px;font-weight:600;fill:{NAVY}}}
.leg{{font-size:10px;fill:#555}}
.val{{font-size:10px;fill:#444}}
table{{width:100%;border-collapse:collapse;font-size:12px}}
th{{background:{RED};color:#fff;padding:8px 10px;text-align:right}}
th:first-child,td:first-child{{text-align:left}}
td{{padding:7px 10px;border-bottom:1px solid #eee;text-align:right}}
tr:nth-child(even) td{{background:#fafafa}}
.inc{{color:{TERRA};font-weight:700}}
.pp,.stretch{{color:{NAVY};font-weight:600}}
.why{{text-align:left;font-size:11px;color:#444}}
.callout{{background:#fff8f0;border:1px solid #e8c4a8;border-radius:8px;padding:14px 16px;font-size:13px;margin-bottom:20px}}
.footer{{font-size:10px;color:#888;margin-top:20px}}
</style></head><body>
<div class="wrap">
<header>
<h1>GPS Stretch Proposal — Kickback distributors</h1>
<p>FY2027 · Reviewed BP 24 Apr 2026 · {len(rows)} distributors · flat +{STRETCH_PP:g}pp</p>
</header>

<div class="callout">
<strong>Kickback DBs</strong>: flat +{STRETCH_PP:g}pp on plan GPS. Stretch BP = base + (pp × max spend).
<strong>ITUK</strong> is excluded from this table — Target Study (+9% dealer sell-out) on workbook tab <em>ITUK FY27</em> only.
</div>

<div class="kpis">
<div class="kpi"><div class="v">{_fmt_eur(tot_base)}</div><div class="l">BP FY27 base</div></div>
<div class="kpi"><div class="v">+{_fmt_eur(tot_incr)}</div><div class="l">GPS stretch €</div></div>
<div class="kpi"><div class="v">{_fmt_eur(tot_stretch)}</div><div class="l">Stretch BP FY27</div></div>
</div>

<div class="card">
<h2>GPS comparison</h2>
{svg_gps}
</div>

<div class="card">
<h2>Proposal — {len(rows)} kickback distributors</h2>
<table>
<thead><tr>
<th>DB</th><th>GPS {ACTUAL_FY_LABEL} ex</th><th>Plan GPS FY27</th><th>Stretch GPS FY27</th><th>+pp</th>
<th>BP FY27</th><th>+€ stretch</th><th>Stretch BP</th><th>Δ</th><th>Rationale</th>
</tr></thead>
<tbody>
{table_rows}
<tr style="font-weight:700;background:#f0f0f0">
<td>TOTAL</td><td colspan="4"></td>
<td>{_fmt_eur(tot_base)}</td>
<td class="inc">+{_fmt_eur(tot_incr)}</td><td>{_fmt_eur(tot_stretch)}</td>
<td>+{100*tot_incr/tot_base:.1f}%</td><td></td></tr>
</tbody></table>
</div>

<p class="footer">ISZE internal · 9 kickback DBs · ITUK separate (Target Study tab) · flat +{STRETCH_PP:g}pp</p>
</div></body></html>"""

    OUT_HTML.write_text(html, encoding="utf-8")
    print("written", OUT_HTML)


def main() -> None:
    kickback_rows, other_rows, ituk = _load_rows()
    build_xlsx(kickback_rows, other_rows, ituk)
    build_html(kickback_rows)
    tot_base = sum(r["base_bp_eur"] for r in kickback_rows)
    tot_incr = sum(r["increment_eur"] for r in kickback_rows)
    print(
        f"{len(kickback_rows)} kickback DBs: base {_fmt_eur(tot_base)} + GPS increment {_fmt_eur(tot_incr)} "
        f"= stretch {_fmt_eur(tot_base + tot_incr)}"
    )
    print(f"{len(other_rows)} other DBs on tab 2 (ITUK excluded — own tab)")
    if ituk:
        print(
            f"ITUK Target Study: wallet {_fmt_eur(ituk['max27'])} · stretch BP {_fmt_eur(ituk['stretch_bp_eur'])}"
        )
    print(f"{sum(1 for r in other_rows if not r.get('has_uio', True))} no-UIO rows on tab 2")
    print("GPS stretch:", ", ".join(f"{r['gps_code']} {r['stretch_pp_label']}" for r in kickback_rows))


if __name__ == "__main__":
    main()
