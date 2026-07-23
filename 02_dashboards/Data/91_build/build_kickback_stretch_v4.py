# -*- coding: utf-8 -*-
"""Promote Kickback Stretch workbook V3 → V4 (UIO v19 UTI wallet + Changelog tab)."""
from __future__ import annotations

import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "09_meetings" / "Preparation Meetings" / "UTI" / "_build"))
from uti_gps_figures import STRETCH_PP, _max_spend_fy27, load_evolution_rows  # noqa: E402

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE\01_data")
V3 = ROOT / "ISZE_GPS_Kickback_Stretch_Proposal_FY27 _V3.xlsx"
V4 = ROOT / "ISZE_GPS_Kickback_Stretch_Proposal_FY27 _V4.xlsx"
V4_DL = Path(r"C:\Users\jerom\Downloads\ISZE_GPS_Kickback_Stretch_Proposal_FY27 _V4.xlsx")

HDR_FILL = PatternFill("solid", start_color="CC0E1E")
WRAP = Alignment(wrap_text=True, vertical="top")


def _uti_v19_row() -> dict:
    rows = load_evolution_rows()
    r26, r27 = rows[2], rows[3]
    max27 = int(round(_max_spend_fy27()))
    bp = float(r27["billing_eur"])
    plan_gps = bp / max27
    stretch_gps = plan_gps + STRETCH_PP / 100
    stretch_eur = round(STRETCH_PP / 100 * max27)
    return {
        "gps_fy26": round(r26["gps_net_pct"] / 100, 4),
        "max_spend": max27,
        "bp_base": bp,
        "plan_gps": plan_gps,
        "stretch_gps": stretch_gps,
        "stretch_eur": stretch_eur,
        "stretch_bp": bp + stretch_eur,
        "uio_fy27": r27["uio"],
    }


def _write_changelog(wb: openpyxl.Workbook) -> None:
    if "Changelog" in wb.sheetnames:
        del wb["Changelog"]
    ws = wb.create_sheet("Changelog", 0)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 88

    uti = _uti_v19_row()
    lines = [
        ("ISZE GPS Kickback Stretch Proposal — Changelog", ""),
        ("", ""),
        ("Version", "V4 — 2026-07-06"),
        ("Prior version", "V3 — 2026-07-03 (Jerome manual edit: default +pp raised to 1.8)"),
        ("", ""),
        ("What went wrong in V3 / UTI meeting pack", ""),
        (
            "1 · Wrong UIO vintage on UTI evolution slide",
            "Subtitle said UIO v19 but car-park numbers came from uio_2026.json built off the "
            "old v17 workbook (~10,668–11,253 units). v19 Israel UTI is 10,049 → 10,972.",
        ),
        (
            "2 · UTI max-spend wallet not v19",
            f"V3 tab 'Proposal FY27' UTI max spend was €8,485,129 — calibrated so "
            f"€3.10M BP ÷ wallet ≈ 36.5% plan GPS. True v19 wallet (LD+MD × €739/€887) "
            f"is €{uti['max_spend']:,} → plan GPS {uti['plan_gps']*100:.1f}% on same BP.",
        ),
        (
            "3 · Stretch € understated",
            "GPS stretch € used the old wallet: +1.8pp × €8.49M ≈ €153K. "
            f"V4: +1.8pp × €{uti['max_spend']:,} = €{uti['stretch_eur']:,}.",
        ),
        (
            "4 · Billing label confusion",
            "Speaker notes said 'gross IE40 billing'. Bars already use gps_billing_clean.json "
            "(parts only — transport, rebates & credit notes out; RUSH/VOR surcharges normalised). "
            "IE40 is internal SAP sales-org code, not a separate customer-facing metric.",
        ),
        (
            "5 · Chart overlap (evolution slide)",
            "FY27 plan column had €3.10M and 10,972 unit labels drawn on top of each other — fixed in slide rebuild.",
        ),
        (
            "6 · Two different 'plan GPS' concepts",
            "Kosha BP ÷ v19 wallet ≈ 35.1% (formula col C). Commercial negotiated target in meeting "
            "deck can still be quoted as 36.5% → 38.3% (+1.8pp) — this workbook uses formula plan GPS.",
        ),
        ("", ""),
        ("V4 changes (Proposal FY27 · UTI row)", ""),
        ("+pp on plan (col E)", f"{STRETCH_PP} pp — unchanged (default B3)"),
        ("GPS FY26 ex (col B)", f"{uti['gps_fy26']*100:.1f}% (was 32.0% hardcoded / 31.99% in V3)"),
        ("Max spend FY27 (col L)", f"€{uti['max_spend']:,} (UIO v19 FY27: {uti['uio_fy27']:,} units)"),
        ("Plan GPS FY27 (col C)", f"= BP ÷ wallet → {uti['plan_gps']*100:.1f}%"),
        ("Stretch GPS FY27 (col D)", f"= plan + {STRETCH_PP}pp → {uti['stretch_gps']*100:.1f}%"),
        ("+€ GPS stretch (col G)", f"€{uti['stretch_eur']:,}"),
        ("Stretch BP FY27 (col I)", f"€{uti['stretch_bp']:,.0f}"),
        ("", ""),
        ("Sources", ""),
        ("UIO v19", "06_uio/_CURRENT/ISUZU_UIO_Overview_2026.xlsx → 01_data/91_build/uio_2026.json"),
        ("Billing", "01_data/00_derived/gps_billing_clean.json"),
        ("BP base", "budget_per_DB_reviewed_24042026.xlsx (unchanged)"),
        ("Rebuild scripts", "refresh_uti_kosha_wallets.py · uti_gps_figures.py · this script"),
    ]

    for i, (a, b) in enumerate(lines, start=1):
        ca, cb = ws.cell(i, 1, a), ws.cell(i, 2, b)
        ca.alignment = WRAP
        cb.alignment = WRAP
        if i == 1:
            ca.font = Font(bold=True, size=14, color="CC0E1E")
        elif a and not b:
            ca.font = Font(bold=True, size=11, color="1B2A4A")
        elif a in ("Version", "Prior version"):
            ca.font = Font(bold=True)

    ws.freeze_panes = "A2"


def _patch_proposal_tab(wb: openpyxl.Workbook, uti: dict) -> None:
    ws = wb["Proposal FY27"]
    ws["A2"].value = (
        f"Editable: default +pp (B3), per-DB +pp (col E), BP base (F), max spend (L). "
        f"V4 — UIO v19 UTI wallet · built {date.today().isoformat()}."
    )
    # UTI row = 8
    ws["B8"].value = uti["gps_fy26"]
    ws["L8"].value = uti["max_spend"]
    # Formulas C,D,G,I,J,K unchanged — Excel recalculates on open


def main() -> None:
    if not V3.is_file():
        raise FileNotFoundError(V3)
    shutil.copy2(V3, V4)
    wb = openpyxl.load_workbook(V4)
    uti = _uti_v19_row()
    _write_changelog(wb)
    _patch_proposal_tab(wb, uti)
    wb.save(V4)
    try:
        shutil.copy2(V4, V4_DL)
        print("copied to", V4_DL)
    except OSError as e:
        print("Downloads copy skipped:", e)
    print("written", V4)
    print(
        f"UTI: FY26 ex {uti['gps_fy26']*100:.1f}% · max €{uti['max_spend']:,} · "
        f"plan {uti['plan_gps']*100:.1f}% · stretch {uti['stretch_gps']*100:.1f}% "
        f"(+{STRETCH_PP}pp)"
    )


if __name__ == "__main__":
    main()
