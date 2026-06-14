"""Build ISUZU_UIO_Dashboard_2026_v18.html from ISUZU_UIO_Overview_2026_v18.xlsx."""
from __future__ import annotations

import datetime
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Users\jerom\Desktop\claude\ISZE")
UIO_V18 = ROOT / "Data" / "ISUZU_UIO_Overview_2026_v18.xlsx"
OUT = ROOT / "Dashboards" / "ISUZU_UIO_Dashboard_2026_v18.html"
STAMP = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

V13_DIST = {
    "UTI": "UTI", "UMI": "UMI", "BL": "ICELAND", "RSA": "NORWAY",
    "Isuzu Sverige": "SWEDEN", "Isuzu Denmark": "DENMARK", "ITUK": "ITUK",
    "IUK": "IUK", "Isuzu Ireland": "IRELAND", "IBX": "BELGIUM",
    "ISD-Lang": "GERMANY", "Imotors": "PORTUGAL", "GasanZammit": "MALTA",
    "Nordic Auto.": "FINLAND", "AutoWallis": "HUNGARY",
    "Petros Petropoulos": "GREECE", "Isuzu Rom Auto": "ROMANIA",
    "Loutsios & Sons": "CYPRUS",
}
SEGMENT_ONLY = {"ITUK": "CV", "IUK": "LCV", "UTI": "CV", "UMI": "LCV"}
CTRY_LABEL = {
    "U KING": "UK", "PORTUGL": "Portugal", "NETHLDS": "Netherlands",
    "ISRAEL": "Israel", "GERMANY": "Germany", "BELGIUM": "Belgium",
}


def _extract() -> dict:
    wb = openpyxl.load_workbook(UIO_V18, read_only=True, data_only=True)
    ws = wb["1.UIO Overview"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r[0] is not None and "Ctry" in str(r[0]))
    hdr = rows[hdr_i]
    ycols = {
        int(m.group(1)): j
        for j, h in enumerate(hdr)
        if h and (m := re.match(r"^\s*(\d{4})[EF]?\s*$", str(h)))
    }
    gt = rows[1]
    hist_years = [y for y in sorted(ycols) if y <= 2024]
    fore_years = [y for y in sorted(ycols) if y >= 2025]

    by_ctry: dict[str, dict[int, float]] = defaultdict(lambda: defaultdict(float))
    cv24 = cv25 = lcv24 = lcv25 = 0.0

    for r in rows[hdr_i + 1 :]:
        if r[1] is None or r[5] is None:
            continue
        dist_lbl = str(r[5]).strip()
        if "MIDI" in dist_lbl.upper():
            continue
        code = V13_DIST.get(dist_lbl)
        if not code:
            continue
        cv = str(r[2] or "").strip()
        if code in SEGMENT_ONLY and cv != SEGMENT_ONLY[code]:
            continue
        cty = str(r[1]).strip()
        for y in (2024, 2025):
            v = r[ycols[y]]
            if isinstance(v, (int, float)):
                by_ctry[cty][y] += float(v)
                if y == 2024:
                    if cv == "CV":
                        cv24 += v
                    elif cv == "LCV":
                        lcv24 += v
                if y == 2025:
                    if cv == "CV":
                        cv25 += v
                    elif cv == "LCV":
                        lcv25 += v

    top = sorted(by_ctry.items(), key=lambda x: -x[1][2024])[:20]
    countries = [CTRY_LABEL.get(c, c.title()) for c, _ in top]
    uio2024 = [int(v[2024]) for _, v in top]
    uio2025e = [int(v[2025]) for _, v in top]

    ws4 = wb["4. ITUK UIO Projection"]
    ituk_proj: list[int] = []
    ituk_models: dict[str, list[int]] = {}
    for row in ws4.iter_rows(min_row=20, max_row=28, values_only=True):
        name = str(row[1] or "").strip()
        if not name:
            continue
        if name == "Total ITUK UIO":
            ituk_proj = [int(round(float(row[i] or 0))) for i in range(2, 8)]
            continue
        if not isinstance(row[2], (int, float)):
            continue
        ituk_models[name] = [int(round(float(row[i] or 0))) for i in range(2, 8)]

    wb.close()

    u24 = int(gt[ycols[2024]] or 0)
    u25 = int(round(float(gt[ycols[2025]] or 0)))
    u26 = int(round(float(gt[ycols[2026]] or 0)))
    u30 = int(round(float(gt[ycols[2030]] or 0)))
    u23 = int(gt[ycols[2023]] or 0)
    pct = round((u24 / u23 - 1) * 100, 1) if u23 else 0

    return {
        "historical_years": hist_years,
        "historical_values": [int(gt[ycols[y]] or 0) for y in hist_years],
        "forecast_years": fore_years,
        "forecast_values": [int(round(float(gt[ycols[y]] or 0))) for y in fore_years],
        "countries": countries,
        "uio2024": uio2024,
        "uio2025e": uio2025e,
        "cv_lcv": [int(cv24), int(cv25), int(lcv24), int(lcv25)],
        "ituk_proj": ituk_proj,
        "ituk_models": ituk_models,
        "kpi": {"u24": u24, "u25": u25, "u26": u26, "u30": u30, "pct": pct},
    }


def main() -> None:
    d = _extract()
    k = d["kpi"]
    payload = {
        "historical_years": d["historical_years"],
        "historical_values": d["historical_values"],
        "forecast_years": d["forecast_years"],
        "forecast_values": d["forecast_values"],
        "countries": d["countries"],
        "uio2024": d["uio2024"],
        "uio2025e": d["uio2025e"],
        "source": f"ISUZU_UIO_Overview_2026_v18.xlsx | {STAMP}",
    }
    data_js = json.dumps(payload, indent=2)

    html = (ROOT / "Dashboards" / "ISUZU_UIO_Dashboard_2026_v17.html").read_text(encoding="utf-8")
    html = re.sub(
        r"const UIO_DATA = \{[\s\S]*?\};\s*\n\ninitDashboard",
        f"const UIO_DATA = {data_js};\n\ninitDashboard",
        html,
        count=1,
    )
    html = html.replace("<title>ISUZU UIO Dashboard 2026 v17</title>", "<title>ISUZU UIO Dashboard 2026 v18</title>")
    html = html.replace("ISUZU_UIO_Overview_2026_v17.xlsx", "ISUZU_UIO_Overview_2026_v18.xlsx")
    html = html.replace('<span class="badge">v17 — May 2026</span>', '<span class="badge">v18 — Jun 2026</span>')
    html = html.replace("What's new in v17 (May 2026)", "What's new in v18 (Jun 2026)")
    html = re.sub(r"<div class=\"kpi-value\">254,117</div>", f'<div class="kpi-value">{k["u24"]:,}</div>', html, count=1)
    html = re.sub(r"<div class=\"kpi-value\">260,449</div>", f'<div class="kpi-value">{k["u25"]:,}</div>', html, count=1)
    html = re.sub(r"<div class=\"kpi-value\">262,559</div>", f'<div class="kpi-value">{k["u30"]:,}</div>', html, count=1)
    html = html.replace("▲ +3.9% vs 2023", f"▲ {k['pct']:+.1f}% vs 2023")
    html = html.replace("Last updated: May 2026", f"Last updated: Jun 2026 ({STAMP})")
    html = html.replace("ISUZU UIO Overview 2026 v17", "ISUZU UIO Overview 2026 v18")

    OUT.write_text(html, encoding="utf-8")
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
