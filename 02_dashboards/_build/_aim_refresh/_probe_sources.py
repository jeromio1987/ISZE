import csv, os, sys, json
sys.stdout.reconfigure(encoding="utf-8")
os.chdir(r"C:\Users\jerom\Desktop\claude\ISZE")

D = "01_data/00_derived/"

def load_pns(path, pncol):
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            pn = (r.get(pncol) or "").strip()
            if pn:
                out[pn] = r
    return out

# 1. the newly classified PNs, from the working folder answer
newpns = set()
cls = "01_data/91_build/export_aim_eq_20260924/classified.csv"
with open(cls, encoding="utf-8-sig", newline="") as f:
    rd = csv.DictReader(f)
    print("classified cols:", rd.fieldnames)
    for r in rd:
        pn = (r.get("part_number") or r.get("pn") or "").strip()
        if pn:
            newpns.add(pn.zfill(10))
print("classified rows PNs:", len(newpns))

# 2. price-list PNs from EXPORT.XLSX
pl = set()
try:
    import openpyxl
    wb = openpyxl.load_workbook("01_data/EXPORT.XLSX", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    print("EXPORT headers:", hdr[:16])
    idx = None
    for i, h in enumerate(hdr):
        if h and str(h).strip().lower() in ("part number", "part_number", "partno", "material"):
            idx = i
    print("pn col idx", idx)
    if idx is not None:
        for r in rows:
            v = r[idx]
            if v:
                pl.add(str(v).strip().zfill(10))
except Exception as e:
    print("EXPORT read failed:", e)
print("price-list PNs:", len(pl))


def probe(name, path, cand_cols):
    if not os.path.exists(path):
        print(f"  {name}: MISSING"); return
    have = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        rd = csv.DictReader(f)
        cols = rd.fieldnames or []
        col = next((c for c in cand_cols if c in cols), None)
        if col is None:
            print(f"  {name}: no aim col of {cand_cols} in {cols[:14]}"); return
        for r in rd:
            pn = (r.get("part_number") or "").strip()
            if not pn: continue
            have[pn] = (r.get(col) or "").strip()
    def real(v): return bool(v) and v.lower() not in ("unknown", "nan", "none")
    inter_pl = [p for p in pl if p in have]
    inter_new = [p for p in newpns if p in have]
    print(f"  {name}: [{col}] rows={len(have):,}")
    print(f"      price-list PNs present {len(inter_pl):,}/{len(pl):,}  -> real family {sum(1 for p in inter_pl if real(have[p])):,}")
    print(f"      newly-classified PNs present {len(inter_new):,}/{len(newpns):,} -> real family {sum(1 for p in inter_new if real(have[p])):,}")


print("\n--- source files, looking for today's labels")
probe("aim_master_v2.csv", D + "aim_master_v2.csv", ["aim_l1", "aim_l1_fixed"])
probe("aim_pn_classification_20260920.csv", D + "aim_pn_classification_20260920.csv", ["aim_l1_fixed", "aim_l1"])
probe("aim_pn_classification.csv", D + "aim_pn_classification.csv", ["aim_l1"])
probe("AIM_Master_PN_Categories_Figkey_20260924.csv", D + "AIM_Master_PN_Categories_Figkey_20260924.csv", ["aim_level_1"])
probe("AIM_Master_PN_Categories_Figkey_20260921.csv", D + "AIM_Master_PN_Categories_Figkey_20260921.csv", ["aim_level_1"])
